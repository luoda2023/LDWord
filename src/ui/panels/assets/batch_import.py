"""Batch profile import helpers for the assets panel."""

from __future__ import annotations

import copy
import csv
from datetime import date, datetime
import json
from pathlib import Path
import re
from typing import Any, Mapping, Sequence

from src.config.entity import EntityProfile
from src.config.materials import ASSET_ROLE_ALIASES, resolve_asset_role_alias
from src.ui.panels.assets.fields import (
    _field_sources_from_imported_keys,
    _normalized_field_aliases,
    _normalized_field_sources,
    _normalized_import_key,
)
from src.ui.panels.assets.items import (
    _normalized_asset_item_payloads,
    _normalized_asset_metadata,
)

def _profile_label(profile: EntityProfile) -> str:
    return profile.profile_name or profile.fields.get("company_name") or "未命名"


def _profile_item_label(index: int, profile: EntityProfile, checked: bool, current_index: int) -> str:
    pieces = [
        _profile_label(profile),
        _profile_role_label(index, profile),
    ]
    if index == current_index:
        pieces.append("当前编辑")
    pieces.append("本次生成" if checked else "未选择")
    return " · ".join(piece for piece in pieces if piece)


def _profile_role_label(index: int, profile: EntityProfile) -> str:
    text = f"{profile.profile_name} {profile.fields.get('company_name', '')}".lower()
    if "联合" in text or "成员" in text:
        return "联合体成员"
    if "分公司" in text or "branch" in text:
        return "分公司"
    if "客户" in text or "client" in text:
        return "客户"
    if any(marker in text for marker in ("版本", "甲方", "乙方", "归档", "version")):
        return "项目版本"
    if index == 0:
        return "主体公司"
    return "这一份"


def _load_batch_profiles_from_path(path: str | Path) -> list[EntityProfile]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".json":
        return _load_batch_profiles_from_json(source)
    if suffix == ".csv":
        return _profiles_from_table_rows(_load_batch_csv_rows(source))
    if suffix in {".xlsx", ".xlsm"}:
        return _profiles_from_table_rows(_load_batch_excel_rows(source))
    raise ValueError(f"Unsupported batch material file: {source}")


def _load_batch_profiles_from_json(path: Path) -> list[EntityProfile]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, Mapping) and isinstance(data.get("profiles"), list):
        return [
            _profile_from_mapping(raw_profile, index=index)
            for index, raw_profile in enumerate(data["profiles"], start=1)
            if isinstance(raw_profile, Mapping)
        ]
    if isinstance(data, Mapping) and isinstance(data.get("rows"), list):
        return _profiles_from_table_rows(
            [row for row in data["rows"] if isinstance(row, Mapping)]
        )
    if isinstance(data, list):
        return _profiles_from_table_rows(
            [row for row in data if isinstance(row, Mapping)]
        )
    return []


def _load_batch_csv_rows(path: Path) -> list[Mapping[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        _validate_unique_table_headers(reader.fieldnames or [])
        return list(reader)


def _load_batch_excel_rows(path: Path) -> list[Mapping[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel batch material imports require openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        primary_sheet = workbook.active
        rows = _excel_table_rows(primary_sheet)
        shared_values: dict[str, Any] = {}
        for sheet in workbook.worksheets:
            if sheet is primary_sheet:
                continue
            auxiliary_rows = _excel_table_rows(sheet)
            if len(auxiliary_rows) != 1:
                continue
            for key, value in auxiliary_rows[0].items():
                if _cell_text(value):
                    shared_values[key] = value
    finally:
        workbook.close()
    if not shared_values:
        return rows
    return [
        {
            **row,
            **{
                key: value
                for key, value in shared_values.items()
                if not _cell_text(row.get(key))
            },
        }
        for row in rows
    ]


def _excel_table_rows(sheet) -> list[dict[str, Any]]:
    """Read one header table without assigning cross-sheet business meaning.

    A secondary sheet is considered shared input only by the caller when this
    helper returns exactly one data row.  Multi-row sheets are intentionally
    left independent instead of creating an implicit Cartesian product.
    """

    raw_rows = list(sheet.iter_rows(values_only=True))
    if not raw_rows:
        return []
    headers = [str(value or "").strip() for value in raw_rows[0]]
    _validate_unique_table_headers(headers)
    rows: list[dict[str, Any]] = []
    for raw_row in raw_rows[1:]:
        row = {
            header: raw_row[index] if index < len(raw_row) else None
            for index, header in enumerate(headers)
            if header
        }
        if _row_has_values(row):
            rows.append(row)
    return rows


def _profiles_from_table_rows(rows: list[Mapping[str, Any]]) -> list[EntityProfile]:
    profiles: list[EntityProfile] = []
    for index, row in enumerate(rows, start=1):
        if not _row_has_values(row):
            continue
        profiles.append(_profile_from_mapping(row, index=index))
    return profiles


def _profile_from_mapping(raw_profile: Mapping[str, Any], *, index: int) -> EntityProfile:
    profile_id = _row_value(raw_profile, _PROFILE_ID_ALIASES)
    profile_name = _row_value(raw_profile, _PROFILE_NAME_ALIASES)
    assets_dir = _row_value(raw_profile, _ASSETS_DIR_ALIASES)
    fields = _mapping_value(raw_profile, "fields")
    normalized_fields = _normalized_profile_fields(fields if isinstance(fields, Mapping) else {})
    normalized_fields.update(_row_fields(raw_profile))
    declared_field_keys = [
        field_key
        for key in (fields.keys() if isinstance(fields, Mapping) else ())
        if (field_key := _direct_material_field_key(key))
    ]
    declared_field_keys = list(
        dict.fromkeys((*declared_field_keys, *_row_field_keys(raw_profile)))
    )

    field_sources = _mapping_value(raw_profile, "field_sources")
    normalized_field_sources = _normalized_field_sources(field_sources if isinstance(field_sources, Mapping) else {})
    if not normalized_field_sources:
        normalized_field_sources = _field_sources_from_imported_keys(
            set(declared_field_keys) or set(normalized_fields),
        )
    field_aliases = _mapping_value(raw_profile, "field_aliases")
    normalized_field_aliases = _normalized_field_aliases(field_aliases if isinstance(field_aliases, Mapping) else {})
    timeline_plans = _mapping_value(raw_profile, "timeline_plans")
    normalized_timeline_plans = (
        copy.deepcopy(dict(timeline_plans))
        if isinstance(timeline_plans, Mapping)
        else {}
    )

    asset_paths = _mapping_value(raw_profile, "asset_paths")
    normalized_asset_paths = _normalized_asset_paths(asset_paths if isinstance(asset_paths, Mapping) else {})
    normalized_asset_paths.update(_row_asset_paths(raw_profile))
    asset_bindings = _mapping_value(raw_profile, "asset_bindings")
    normalized_asset_bindings = (
        copy.deepcopy(dict(asset_bindings))
        if isinstance(asset_bindings, Mapping)
        else {}
    )
    normalized_asset_bindings.update(_row_asset_bindings(raw_profile))
    asset_metadata = _mapping_value(raw_profile, "asset_metadata")
    normalized_asset_metadata = _normalized_asset_metadata(
        asset_metadata if isinstance(asset_metadata, Mapping) else {}
    )
    for role, metadata in _row_asset_metadata(raw_profile).items():
        normalized_asset_metadata[role] = {
            **normalized_asset_metadata.get(role, {}),
            **metadata,
        }
    asset_items = _mapping_value(raw_profile, "asset_items")
    normalized_asset_items = [
        *_normalized_asset_item_payloads(
            asset_items
            if isinstance(asset_items, Sequence) and not isinstance(asset_items, (str, bytes))
            else []
        ),
        *_row_asset_item_payloads(raw_profile),
    ]
    asset_token_specs = _mapping_value(raw_profile, "asset_token_specs")
    normalized_asset_token_specs = (
        list(asset_token_specs)
        if isinstance(asset_token_specs, Sequence)
        and not isinstance(asset_token_specs, (str, bytes))
        else []
    )
    image_material_rules = _mapping_value(raw_profile, "image_material_rules")
    normalized_image_material_rules = (
        copy.deepcopy(dict(image_material_rules))
        if isinstance(image_material_rules, Mapping)
        else {}
    )
    content_bindings = _mapping_value(raw_profile, "content_bindings")
    normalized_content_bindings = (
        copy.deepcopy(dict(content_bindings))
        if isinstance(content_bindings, Mapping)
        else {}
    )
    content_rules = _mapping_value(raw_profile, "content_rules")
    normalized_content_rules = (
        copy.deepcopy(list(content_rules))
        if isinstance(content_rules, Sequence)
        and not isinstance(content_rules, (str, bytes))
        else []
    )
    attachment_bindings = _mapping_value(raw_profile, "attachment_bindings")
    normalized_attachment_bindings = (
        copy.deepcopy(dict(attachment_bindings))
        if isinstance(attachment_bindings, Mapping)
        else {}
    )
    attachment_role_specs = _mapping_value(raw_profile, "attachment_role_specs")
    normalized_attachment_role_specs = (
        copy.deepcopy(list(attachment_role_specs))
        if isinstance(attachment_role_specs, Sequence)
        and not isinstance(attachment_role_specs, (str, bytes))
        else []
    )

    if not profile_name:
        profile_name = normalized_fields.get("company_name") or f"第 {index} 份"

    return EntityProfile(
        profile_id=profile_id,
        profile_name=profile_name,
        fields=normalized_fields,
        assets_dir=assets_dir,
        timeline_plans=normalized_timeline_plans,
        declared_field_keys=declared_field_keys,
        field_sources=normalized_field_sources,
        field_aliases=normalized_field_aliases,
        asset_paths=normalized_asset_paths,
        asset_bindings=normalized_asset_bindings,
        asset_metadata=normalized_asset_metadata,
        asset_token_specs=normalized_asset_token_specs,
        asset_items=normalized_asset_items,
        asset_item_history=[],
        image_material_rules=normalized_image_material_rules,
        content_bindings=normalized_content_bindings,
        content_rules=normalized_content_rules,
        attachment_role_specs=normalized_attachment_role_specs,
        attachment_bindings=normalized_attachment_bindings,
    )


def _row_fields(row: Mapping[str, Any]) -> dict[str, str]:
    fields: dict[str, str] = {}
    for raw_key, raw_value in row.items():
        key = _row_field_key(raw_key)
        value = _cell_text(raw_value)
        if not key or not value:
            continue
        fields[key] = value
    return fields


def _row_field_keys(row: Mapping[str, Any]) -> tuple[str, ...]:
    """Return imported material columns even when a row has no value."""

    return tuple(
        dict.fromkeys(
            key
        for raw_key in row
        if (key := _row_field_key(raw_key))
        )
    )


def _row_field_key(raw_key: object) -> str:
    key = str(raw_key or "").strip()
    if not key:
        return ""
    if (
        _is_profile_meta_key(key)
        or _asset_role_for_import_key(key)
        or _asset_metadata_role_for_import_key(key)[0]
        or _repeated_asset_item_key(key)
        or _asset_binding_role_for_import_key(key)
    ):
        return ""
    if _normalized_import_key(key) in {
        "fields",
        "requiredfields",
        "fieldsources",
        "fieldaliases",
        "assetpaths",
        "assetmetadata",
        "assetitems",
        "assetbindings",
        "assettokenspecs",
        "attachmentrolespecs",
        "imagematerialrules",
        "contentbindings",
        "contentrules",
        "attachmentbindings",
    }:
        return ""
    return _direct_material_field_key(key)


def _row_asset_paths(row: Mapping[str, Any]) -> dict[str, str]:
    asset_paths: dict[str, str] = {}
    for raw_key, raw_value in row.items():
        role = _asset_role_for_import_key(str(raw_key or ""))
        value = _cell_text(raw_value)
        if role and value:
            asset_paths[role] = value
    return asset_paths


def _row_asset_bindings(row: Mapping[str, Any]) -> dict[str, dict[str, object]]:
    bindings: dict[str, dict[str, object]] = {}
    for raw_key, raw_value in row.items():
        role = _asset_binding_role_for_import_key(str(raw_key or ""))
        value = _cell_text(raw_value)
        if role and value:
            bindings[role] = {
                "role": role,
                "cardinality": "multiple",
                "source_kind": "directory",
                "source_path": value,
                "recursive": True,
                "order_policy": "natural_path",
                "naming_template": "{role}_{sequence:03d}",
                "min_items": 0,
                "max_items": None,
            }
    return bindings


def _row_asset_metadata(row: Mapping[str, Any]) -> dict[str, dict[str, str]]:
    metadata: dict[str, dict[str, str]] = {}
    for raw_key, raw_value in row.items():
        role, meta_key = _asset_metadata_role_for_import_key(str(raw_key or ""))
        value = _cell_text(raw_value)
        if role and meta_key and value:
            metadata.setdefault(role, {})[meta_key] = value
    return metadata


def _row_asset_item_payloads(row: Mapping[str, Any]) -> list[dict[str, object]]:
    by_index: dict[tuple[int, int | None], dict[str, object]] = {}
    for raw_key, raw_value in row.items():
        key = str(raw_key or "")
        value = _cell_text(raw_value)
        if not value:
            continue
        path_key = _repeated_question_figure_path_key(key)
        if path_key is not None:
            item = by_index.setdefault(path_key, {"metadata": {}})
            item["path"] = value
            continue
        figure_key, meta_key = _repeated_question_figure_metadata_key(key)
        if figure_key is not None and meta_key:
            item = by_index.setdefault(figure_key, {"metadata": {}})
            metadata = item.setdefault("metadata", {})
            if isinstance(metadata, dict):
                metadata[meta_key] = value

    payloads: list[dict[str, object]] = []
    for index, figure_order in sorted(by_index, key=_question_figure_import_sort_key):
        item = by_index[(index, figure_order)]
        path = str(item.get("path") or "").strip()
        metadata = dict(item.get("metadata", {}) or {})
        asset_id = str(metadata.get("asset_id") or metadata.get("assetId") or "").strip()
        if not path and not asset_id:
            continue
        metadata.setdefault("question_index", str(index))
        if figure_order is not None:
            metadata.setdefault("figure_order", str(figure_order))
        item_id = (
            f"question_figure_{index}_{figure_order}"
            if figure_order is not None
            else f"question_figure_{index}"
        )
        label = asset_id or (
            f"Question figure {index}-{figure_order}"
            if figure_order is not None
            else f"Question figure {index}"
        )
        payloads.append(
            {
                "item_id": item_id,
                "label": label,
                "role": "question_figure",
                "path": path,
                "metadata": metadata,
            }
        )
    return payloads


def _repeated_asset_item_key(key: str) -> bool:
    if _repeated_question_figure_path_key(key) is not None:
        return True
    figure_key, _meta_key = _repeated_question_figure_metadata_key(key)
    return figure_key is not None


def _repeated_question_figure_path_index(key: str) -> int | None:
    figure_key = _repeated_question_figure_path_key(key)
    return figure_key[0] if figure_key is not None else None


def _repeated_question_figure_path_key(key: str) -> tuple[int, int | None] | None:
    normalized = _normalized_import_key(key)
    if normalized.endswith("path"):
        normalized = normalized[:-4]
    if normalized.endswith("文件"):
        normalized = normalized[:-2]
    return _question_figure_key_from_base(normalized)


def _repeated_question_figure_metadata_key(key: str) -> tuple[tuple[int, int | None] | None, str]:
    normalized = _normalized_import_key(key)
    for suffix, meta_key in _ASSET_METADATA_IMPORT_SUFFIXES:
        if not normalized.endswith(suffix):
            continue
        base = normalized[: -len(suffix)]
        figure_key = _question_figure_key_from_base(base)
        if figure_key is not None:
            return figure_key, meta_key
    return None, ""


def _question_figure_index_from_base(normalized: str) -> int | None:
    figure_key = _question_figure_key_from_base(normalized)
    return figure_key[0] if figure_key is not None else None


def _question_figure_key_from_base(normalized: str) -> tuple[int, int | None] | None:
    for alias in _QUESTION_FIGURE_IMPORT_BASES:
        if not normalized.startswith(alias):
            continue
        suffix = normalized[len(alias):]
        key = _parse_question_figure_suffix(suffix)
        if key is not None:
            return key
    return None


def _parse_question_figure_suffix(suffix: str) -> tuple[int, int | None] | None:
    text = str(suffix or "").strip()
    if text.isdigit():
        index = int(text)
        return (index, None) if index > 0 else None
    match = re.fullmatch(
        r"(\d+)(?:图|figure|fig|image|img|pic|picture|asset)(\d+)",
        text,
    )
    if not match:
        return None
    question_index = int(match.group(1))
    figure_order = int(match.group(2))
    if question_index <= 0 or figure_order <= 0:
        return None
    return question_index, figure_order


def _question_figure_import_sort_key(key: tuple[int, int | None]) -> tuple[int, int, int]:
    question_index, figure_order = key
    return (
        question_index,
        figure_order if figure_order is not None else 0,
        0 if figure_order is None else 1,
    )


def _normalized_profile_fields(fields: Mapping[str, Any]) -> dict[str, str]:
    normalized: dict[str, str] = {}
    for raw_key, raw_value in fields.items():
        key = str(raw_key or "").strip()
        value = _cell_text(raw_value)
        if not key or not value:
            continue
        normalized[_direct_material_field_key(key)] = value
    return normalized


def _direct_material_field_key(value: object) -> str:
    key = str(value or "").strip()
    if key.startswith("{{") and key.endswith("}}"):
        key = key[2:-2].strip()
    return key


def _validate_unique_table_headers(headers: Sequence[object]) -> None:
    seen: set[str] = set()
    duplicates: list[str] = []
    for raw_header in headers:
        key = _direct_material_field_key(raw_header)
        if not key:
            continue
        if key in seen and key not in duplicates:
            duplicates.append(key)
        seen.add(key)
    if duplicates:
        raise ValueError(
            "资料表存在重复字段，无法导入：" + "、".join(duplicates)
        )


def _normalized_asset_paths(asset_paths: Mapping[str, Any]) -> dict[str, str]:
    normalized: dict[str, str] = {}
    for raw_key, raw_value in asset_paths.items():
        role = _asset_role_for_import_key(str(raw_key or "")) or str(raw_key or "").strip()
        value = _cell_text(raw_value)
        if role and value:
            normalized[role] = value
    return normalized


def _row_value(row: Mapping[str, Any], aliases: set[str]) -> str:
    for raw_key, raw_value in row.items():
        if _normalized_import_key(str(raw_key or "")) in aliases:
            return _cell_text(raw_value)
    return ""


def _mapping_value(row: Mapping[str, Any], target_key: str) -> Any:
    for raw_key, raw_value in row.items():
        if _normalized_import_key(str(raw_key or "")) == _normalized_import_key(target_key):
            return raw_value
    return None


def _row_has_values(row: Mapping[str, Any]) -> bool:
    return any(_cell_text(value) for value in row.values())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time().isoformat() == "00:00:00":
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_profile_meta_key(key: str) -> bool:
    normalized = _normalized_import_key(key)
    return normalized in (
        _PROFILE_ID_ALIASES
        | _PROFILE_NAME_ALIASES
        | _ASSETS_DIR_ALIASES
        | {"timelineplans", "timelineplan"}
    )


def _asset_role_for_import_key(key: str) -> str:
    normalized = _normalized_import_key(key)
    if normalized.endswith("path"):
        normalized = normalized[:-4]
    if normalized.endswith("文件"):
        normalized = normalized[:-2]
    return resolve_asset_role_alias(normalized)


def _asset_binding_role_for_import_key(key: str) -> str:
    normalized = _normalized_import_key(key)
    for suffix in ("folder", "directory", "dir", "文件夹", "目录"):
        if not normalized.endswith(suffix):
            continue
        base = normalized[: -len(suffix)]
        return resolve_asset_role_alias(base)
    return ""


def _asset_metadata_role_for_import_key(key: str) -> tuple[str, str]:
    normalized = _normalized_import_key(key)
    for suffix, meta_key in _ASSET_METADATA_IMPORT_SUFFIXES:
        if not normalized.endswith(suffix):
            continue
        base = normalized[: -len(suffix)]
        role = _asset_role_for_import_key(base) or resolve_asset_role_alias(base)
        if role:
            return role, meta_key
    return "", ""


_PROFILE_ID_ALIASES = {
    "profileid",
    "资料编号",
    "这一份编号",
}


_PROFILE_NAME_ALIASES = {
    "profilename",
    "这一份名称",
    "生成对象",
    "生成对象名称",
}


_ASSETS_DIR_ALIASES = {
    "assetsdir",
    "assetdir",
    "materialdir",
    "图片目录",
    "素材目录",
    "资料目录",
}


_ASSET_IMPORT_ALIASES = ASSET_ROLE_ALIASES


_QUESTION_FIGURE_IMPORT_BASES = tuple(
    alias
    for alias, role in ASSET_ROLE_ALIASES.items()
    if role == "question_figure"
)


_ASSET_METADATA_IMPORT_SUFFIXES = (
    ("assetid", "asset_id"),
    ("source", "source"),
    ("librarypackageid", "library_package_id"),
    ("packageid", "package_id"),
    ("materialpackageid", "material_package_id"),
    ("assetpackageid", "asset_package_id"),
    ("sourcepackageid", "source_package_id"),
    ("librarypackagename", "library_package_name"),
    ("packagename", "package_name"),
    ("materialpackagename", "material_package_name"),
    ("assetpackagename", "asset_package_name"),
    ("assetversion", "asset_version"),
    ("libraryversion", "library_version"),
    ("remoteversion", "remote_version"),
    ("masterversion", "master_version"),
    ("version", "asset_version"),
    ("assetetag", "asset_etag"),
    ("remoteetag", "remote_etag"),
    ("libraryetag", "library_etag"),
    ("etag", "asset_etag"),
    ("contenthash", "content_hash"),
    ("assetupdatedat", "asset_updated_at"),
    ("remoteupdatedat", "remote_updated_at"),
    ("libraryupdatedat", "library_updated_at"),
    ("updatedat", "asset_updated_at"),
    ("modifiedat", "modified_at"),
    ("lastmodified", "last_modified"),
    ("organizationid", "organization_id"),
    ("orgid", "org_id"),
    ("role", "role"),
    ("approvedby", "approved_by"),
    ("approver", "approver"),
    ("localcachepath", "cache_path"),
    ("cachepath", "cache_path"),
    ("cachedpath", "cache_path"),
    ("localpath", "cache_path"),
    ("resolvedpath", "cache_path"),
    ("assetpath", "cache_path"),
    ("thumbnailpath", "thumbnail_path"),
    ("thumbnailcachepath", "thumbnail_path"),
    ("thumbpath", "thumbnail_path"),
    ("previewpath", "thumbnail_path"),
    ("previewcachepath", "thumbnail_path"),
    ("figureorder", "figure_order"),
    ("figureindex", "figure_order"),
    ("imageorder", "figure_order"),
    ("assetorder", "figure_order"),
    ("sortorder", "figure_order"),
    ("alttext", "alt_text"),
    ("alt", "alt_text"),
    ("description", "alt_text"),
    ("caption", "caption"),
    ("图片说明", "alt_text"),
    ("说明", "alt_text"),
    ("描述", "alt_text"),
    ("题注", "caption"),
)


__all__ = [
    '_profile_label',
    '_profile_item_label',
    '_profile_role_label',
    '_load_batch_profiles_from_path',
    '_load_batch_profiles_from_json',
    '_load_batch_csv_rows',
    '_load_batch_excel_rows',
    '_profiles_from_table_rows',
    '_profile_from_mapping',
    '_row_fields',
    '_row_field_keys',
    '_row_asset_paths',
    '_row_asset_bindings',
    '_row_asset_metadata',
    '_row_asset_item_payloads',
    '_repeated_asset_item_key',
    '_repeated_question_figure_path_index',
    '_repeated_question_figure_path_key',
    '_repeated_question_figure_metadata_key',
    '_question_figure_index_from_base',
    '_question_figure_key_from_base',
    '_parse_question_figure_suffix',
    '_question_figure_import_sort_key',
    '_normalized_profile_fields',
    '_normalized_asset_paths',
    '_asset_binding_role_for_import_key',
    '_direct_material_field_key',
    '_validate_unique_table_headers',
    '_row_value',
    '_mapping_value',
    '_row_has_values',
    '_cell_text',
    '_is_profile_meta_key',
    '_asset_role_for_import_key',
    '_asset_metadata_role_for_import_key',
    '_PROFILE_ID_ALIASES',
    '_PROFILE_NAME_ALIASES',
    '_ASSETS_DIR_ALIASES',
    '_ASSET_IMPORT_ALIASES',
    '_QUESTION_FIGURE_IMPORT_BASES',
    '_ASSET_METADATA_IMPORT_SUFFIXES',
]
