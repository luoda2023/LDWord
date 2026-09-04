"""Explicit workbook-mapping workflow before canonical package materialization."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from src.domain.materials import (
    MaterialGroup,
    MaterialPackage,
    MaterialRecord,
    MaterialScope,
    generate_group_id,
    generate_package_id,
    generate_record_id,
)


@dataclass(frozen=True, slots=True)
class ImportMappingProfile:
    """User-confirmed workbook-to-contract mapping."""

    primary_sheet: str = ""
    field_key_map: Mapping[str, str] = field(default_factory=dict)
    group_source_fields: tuple[str, ...] = ()
    group_target_field: str = ""
    record_name_source_fields: tuple[str, ...] = ()
    activation_required_fields: tuple[str, ...] = ()
    auxiliary_single_row_as_shared: bool = True

    def canonical_key(self, source_key: str) -> str:
        return str(self.field_key_map.get(source_key, source_key))


@dataclass(frozen=True, slots=True)
class ImportCandidate:
    record_id: str
    display_name: str
    group_id: str
    lifecycle: str
    scope: MaterialScope
    origin: Mapping[str, object]
    missing_activation_fields: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class MaterialImportDraft:
    package_id: str
    display_name: str
    work_mode_id: str
    material_contract_id: str
    shared_scope: MaterialScope
    groups: tuple[MaterialGroup, ...]
    candidates: tuple[ImportCandidate, ...]
    source_path: str
    mapping: ImportMappingProfile
    issues: tuple[str, ...] = ()

    @property
    def candidate_count(self) -> int:
        return len(self.candidates)

    @property
    def active_count(self) -> int:
        return sum(item.lifecycle == "active" for item in self.candidates)

    @property
    def draft_count(self) -> int:
        return sum(item.lifecycle == "draft" for item in self.candidates)

    def materialize(self) -> MaterialPackage:
        if self.issues:
            raise ValueError(
                "material_import_draft_has_issues:" + "|".join(self.issues)
            )
        return MaterialPackage(
            package_id=self.package_id,
            display_name=self.display_name,
            work_mode_id=self.work_mode_id,
            material_contract_id=self.material_contract_id,
            shared_scope=self.shared_scope,
            groups=self.groups,
            records=tuple(
                MaterialRecord(
                    record_id=item.record_id,
                    display_name=item.display_name,
                    group_id=item.group_id,
                    lifecycle=item.lifecycle,
                    scope=item.scope,
                    origin=item.origin,
                )
                for item in self.candidates
            ),
            metadata={
                "import_source_kind": "xlsx",
            },
        )


def inspect_material_workbook_headers(
    path: str | Path,
) -> Mapping[str, tuple[str, ...]]:
    """Return exact source headers for an explicit user mapping step."""

    from openpyxl import load_workbook

    source = Path(path).resolve()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        result: dict[str, tuple[str, ...]] = {}
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            try:
                row = next(iterator)
            except StopIteration:
                result[sheet.title] = ()
                continue
            headers = tuple(str(value or "").strip() for value in row)
            nonempty = tuple(item for item in headers if item)
            if len(set(nonempty)) != len(nonempty):
                raise ValueError(
                    f"material_import_headers_duplicate:{sheet.title}"
                )
            result[sheet.title] = nonempty
        return result
    finally:
        workbook.close()


def inspect_material_workbook(
    path: str | Path,
    *,
    work_mode_id: str,
    material_contract_id: str,
    mapping: ImportMappingProfile,
) -> MaterialImportDraft:
    """Inspect workbook rows without loading or repairing a package."""

    from openpyxl import load_workbook

    source = Path(path).resolve()
    workbook = load_workbook(source, read_only=True, data_only=True)
    issues: list[str] = []
    try:
        primary = (
            workbook[mapping.primary_sheet]
            if mapping.primary_sheet
            else workbook.active
        )
        _headers, rows = _worksheet_rows(primary)
        shared_fields: dict[str, str] = {}
        shared_provenance: dict[str, str] = {}
        if mapping.auxiliary_single_row_as_shared:
            for sheet in workbook.worksheets:
                if sheet is primary:
                    continue
                _, auxiliary_rows = _worksheet_rows(sheet)
                if not auxiliary_rows:
                    continue
                if len(auxiliary_rows) != 1:
                    issues.append(
                        f"auxiliary_sheet_requires_single_row:{sheet.title}"
                    )
                    continue
                row_number, values = auxiliary_rows[0]
                for source_key, value in values.items():
                    if value == "":
                        continue
                    key = mapping.canonical_key(source_key)
                    previous = shared_fields.get(key)
                    if previous is not None and previous != value:
                        issues.append(
                            f"shared_field_conflict:{key}:{sheet.title}"
                        )
                        continue
                    shared_fields[key] = value
                    shared_provenance[key] = (
                        f"xlsx:{sheet.title}!{_column_for_header(sheet, source_key)}"
                        f"{row_number}"
                    )
    finally:
        workbook.close()

    groups: list[MaterialGroup] = []
    group_ids_by_value: dict[str, str] = {}
    candidates: list[ImportCandidate] = []
    for row_number, row in rows:
        canonical_row = {
            mapping.canonical_key(key): value
            for key, value in row.items()
        }
        group_value = _first_value(row, mapping.group_source_fields)
        group_id = ""
        if group_value:
            normalized_group = group_value.casefold()
            group_id = group_ids_by_value.get(normalized_group, "")
            if not group_id:
                group_id = generate_group_id()
                group_ids_by_value[normalized_group] = group_id
                group_fields = (
                    {mapping.group_target_field: group_value}
                    if mapping.group_target_field
                    else {}
                )
                groups.append(
                    MaterialGroup(
                        group_id=group_id,
                        display_name=group_value,
                        scope=MaterialScope(fields=group_fields),
                        metadata={
                            "import_sheet": primary.title,
                            "first_row": row_number,
                        },
                    )
                )
        missing = tuple(
            key
            for key in mapping.activation_required_fields
            if canonical_row.get(key, "") == ""
        )
        display_name = (
            _first_value(row, mapping.record_name_source_fields)
            or f"第 {len(candidates) + 1} 条"
        )
        record_fields = {
            key: value
            for key, value in canonical_row.items()
            if value != ""
            and (
                not mapping.group_target_field
                or key != mapping.group_target_field
            )
        }
        provenance = {
            key: f"xlsx:{primary.title}!row:{row_number}"
            for key in record_fields
        }
        candidates.append(
            ImportCandidate(
                record_id=generate_record_id(),
                display_name=display_name,
                group_id=group_id,
                lifecycle="active" if not missing else "draft",
                scope=MaterialScope(
                    fields=record_fields,
                    provenance=provenance,
                ),
                origin={
                    "source_kind": "xlsx",
                    "source_name": source.name,
                    "sheet": primary.title,
                    "row": row_number,
                },
                missing_activation_fields=missing,
            )
        )
    if not candidates:
        issues.append("primary_sheet_has_no_candidate_rows")
    return MaterialImportDraft(
        package_id=generate_package_id(),
        display_name=source.stem,
        work_mode_id=work_mode_id,
        material_contract_id=material_contract_id,
        shared_scope=MaterialScope(
            fields=shared_fields,
            provenance=shared_provenance,
        ),
        groups=tuple(groups),
        candidates=tuple(candidates),
        source_path=str(source),
        mapping=mapping,
        issues=tuple(issues),
    )


def _worksheet_rows(sheet) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    iterator = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        return [], []
    headers = [str(value or "").strip() for value in raw_headers]
    nonempty_headers = [item for item in headers if item]
    duplicates = sorted(
        {item for item in nonempty_headers if nonempty_headers.count(item) > 1}
    )
    if duplicates:
        raise ValueError(
            "material_import_headers_duplicate:" + ",".join(duplicates)
        )
    rows: list[tuple[int, dict[str, str]]] = []
    for row_number, raw_row in enumerate(iterator, start=2):
        values = {
            header: _cell_text(
                raw_row[index] if index < len(raw_row) else None
            )
            for index, header in enumerate(headers)
            if header
        }
        if any(value != "" for value in values.values()):
            rows.append((row_number, values))
    return headers, rows


def _column_for_header(sheet, header: str) -> str:
    for cell in sheet[1]:
        if str(cell.value or "").strip() == header:
            return cell.column_letter
    return ""


def _first_value(
    values: Mapping[str, str],
    keys: tuple[str, ...],
) -> str:
    for key in keys:
        value = values.get(key, "")
        if value:
            return value
    return ""


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


__all__ = [
    "ImportCandidate",
    "ImportMappingProfile",
    "MaterialImportDraft",
    "inspect_material_workbook",
    "inspect_material_workbook_headers",
]
