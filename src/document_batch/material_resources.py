"""Materialize V1 file, image, and attachment roles into one DOCX copy."""

from __future__ import annotations

import re
from collections.abc import Mapping
from copy import deepcopy
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from types import MappingProxyType

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image

from src.application.materials import ExecutionMaterialRecord, ExecutionResource
from src.config.attachment_materials import (
    AttachmentBinding,
    AttachmentCardinality,
    AttachmentItem,
    AttachmentSourceKind,
)
from src.config.content_artifacts import ContentMaterialBinding
from src.config.content_materials import (
    ContentInsertionRule,
    FileAssetRef,
    content_anchor_token,
)
from src.config.image_materials import (
    ImageWatermarkPolicy,
    ImageWatermarkTextSource,
    ResolvedImageWatermark,
)
from src.services.material_assets.image_transformer import (
    prepare_material_image,
    resolve_image_watermark,
)
from src.services.material_content.artifact_repository import (
    ContentArtifactRepository,
)
from src.services.material_content.compiler import compile_content_material
from src.services.material_content.composer import (
    ContentComposeInputs,
    ContentComposeRequest,
    ContentMaterialComposer,
)
from src.shared.engine.docx_material_tokens import (
    extract_docx_material_token_blocks,
)
from src.shared.engine.material_token_contract import (
    MATERIAL_TOKEN_PATTERN,
    MaterialTokenKind,
    MaterialTokenNamespace,
    material_token,
    parse_material_token,
)

_DOMAIN_BY_KIND = {
    MaterialTokenKind.CONTENT: "content",
    MaterialTokenKind.IMAGE: "image",
    MaterialTokenKind.ATTACHMENT: "attachment",
}

_EMU_PER_INCH = 914400
_IMAGE_BLOCK_VERTICAL_RESERVE_EMU = _EMU_PER_INCH


@dataclass(frozen=True, slots=True)
class AttachmentDeliveryItem:
    source_path: str
    role: str
    relative_path: str


def inspect_template_resource_roles(
    source_paths: tuple[Path, ...],
) -> Mapping[str, str]:
    """Return role -> resource domain for canonical tokens in templates."""

    roles: dict[str, str] = {}
    for source in source_paths:
        document = Document(source)
        for block in extract_docx_material_token_blocks(document).blocks:
            for match in MATERIAL_TOKEN_PATTERN.finditer(block.combined_text):
                try:
                    ref = parse_material_token(match.group(0))
                except (TypeError, ValueError):
                    continue
                domain = _DOMAIN_BY_KIND.get(ref.kind)
                if domain is None:
                    continue
                prior = roles.setdefault(ref.identifier, domain)
                if prior != domain:
                    raise ValueError(
                        f"document_batch_resource_token_conflict:{ref.identifier}"
                    )
    return MappingProxyType(roles)


def materialize_record_resources(
    source_docx: Path,
    output_docx: Path,
    *,
    record: ExecutionMaterialRecord,
    resource_domains: Mapping[str, str],
    image_policy: Mapping[str, object],
    work_dir: Path,
) -> tuple[AttachmentDeliveryItem, ...]:
    """Consume all resource tokens on an already field-filled staging DOCX."""

    template_roles = inspect_template_resource_roles((source_docx,))
    content_roles = tuple(
        role
        for role in _roles_for_domain(record, resource_domains, "content")
        if template_roles.get(role) == "content"
    )
    attachment_roles = _roles_for_domain(record, resource_domains, "attachment")
    image_roles = tuple(
        role
        for role in _roles_for_domain(record, resource_domains, "image")
        if template_roles.get(role) == "image"
    )

    current = source_docx
    if content_roles or attachment_roles:
        composed = work_dir / "composed.docx"
        repository = ContentArtifactRepository(work_dir / "content-artifacts")
        content_bindings: list[ContentMaterialBinding] = []
        content_rules: list[ContentInsertionRule] = []
        for role_index, role in enumerate(content_roles):
            resources = record.resources.get(role, ())
            if len(resources) > 1:
                expanded = work_dir / f"content-anchors-{role_index:04d}.docx"
                _expand_content_anchor(
                    current,
                    expanded,
                    role=role,
                    count=len(resources),
                )
                current = expanded
            for item_index, resource in enumerate(resources, start=1):
                content_id = (
                    role
                    if len(resources) == 1
                    else f"{role}__{item_index:04d}"
                )
                result = compile_content_material(resource.source_path, repository)
                if result.blocked or result.artifact_ref is None:
                    codes = ",".join(item.code for item in result.findings)
                    raise ValueError(
                        f"document_batch_content_compile_failed:{role}:{codes}"
                    )
                content_bindings.append(
                    ContentMaterialBinding(
                        content_id=content_id,
                        label=resource.object_ref.original_name or role,
                        artifact_ref=result.artifact_ref,
                    )
                )
                content_rules.append(
                    ContentInsertionRule(
                        rule_id=f"content:{content_id}",
                        content_id=content_id,
                        anchor_token=content_anchor_token(content_id),
                    )
                )
        attachment_bindings = tuple(
            _attachment_binding(role, record.resources.get(role, ()))
            for role in attachment_roles
        )
        ContentMaterialComposer(repository=repository).compose(
            ContentComposeRequest(
                source_docx_path=str(current),
                output_docx_path=str(composed),
                inputs=ContentComposeInputs(
                    content_bindings=tuple(content_bindings),
                    content_rules=tuple(content_rules),
                    attachment_bindings=attachment_bindings,
                ),
            )
        )
        current = composed

    if image_roles:
        image_output = work_dir / "images.docx"
        _insert_images(
            current,
            image_output,
            record=record,
            roles=image_roles,
            image_policy=image_policy,
            cache_dir=work_dir / "image-cache",
        )
        current = image_output

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    if current != output_docx:
        output_docx.write_bytes(current.read_bytes())
    Document(output_docx)
    return plan_attachment_delivery_items(
        record,
        resource_domains=resource_domains,
    )


def plan_attachment_delivery_items(
    record: ExecutionMaterialRecord,
    *,
    resource_domains: Mapping[str, str],
) -> tuple[AttachmentDeliveryItem, ...]:
    items: list[AttachmentDeliveryItem] = []
    for role in _roles_for_domain(record, resource_domains, "attachment"):
        resources = record.resources.get(role, ())
        for resource, relative_path in zip(
            resources,
            _attachment_relative_paths(resources),
        ):
            items.append(
                AttachmentDeliveryItem(
                    source_path=resource.source_path,
                    role=role,
                    relative_path=relative_path,
                )
            )
    return tuple(items)


def _expand_content_anchor(
    source_docx: Path,
    output_docx: Path,
    *,
    role: str,
    count: int,
) -> None:
    document = Document(source_docx)
    blocks = extract_docx_material_token_blocks(document)
    source_token = content_anchor_token(role)
    matches = [
        block
        for block in blocks.blocks
        if source_token in block.combined_text
    ]
    if len(matches) != 1:
        raise ValueError(
            f"document_batch_content_token_occurrence_invalid:{role}:{len(matches)}"
        )
    block = matches[0]
    direct = blocks.direct_body_blocks.get(block.block_id)
    if direct is None or block.combined_text.strip() != source_token:
        raise ValueError(
            f"document_batch_content_token_not_isolated_body:{role}"
        )
    original, _body_index = direct
    cursor = original
    for index in range(1, count + 1):
        paragraph = original if index == 1 else deepcopy(original)
        text_nodes = tuple(paragraph.iter(qn("w:t")))
        if not text_nodes:
            raise ValueError(f"document_batch_content_token_text_missing:{role}")
        text_nodes[0].text = content_anchor_token(f"{role}__{index:04d}")
        for text_node in text_nodes[1:]:
            text_node.text = ""
        if index > 1:
            cursor.addnext(paragraph)
            cursor = paragraph
    document.save(output_docx)
    Document(output_docx)


def _roles_for_domain(
    record: ExecutionMaterialRecord,
    resource_domains: Mapping[str, str],
    domain: str,
) -> tuple[str, ...]:
    return tuple(
        role
        for role in sorted(record.resources)
        if record.resources[role] and resource_domains.get(role) == domain
    )


def _file_ref(resource: ExecutionResource) -> FileAssetRef:
    path = Path(resource.source_path)
    payload = path.read_bytes()
    return FileAssetRef(
        source_path=str(path),
        original_name=resource.object_ref.original_name or path.name,
        media_type=resource.object_ref.media_type,
        content_sha256=sha256(payload).hexdigest(),
        byte_size=len(payload),
    )


def _attachment_binding(
    role: str,
    resources: tuple[ExecutionResource, ...],
) -> AttachmentBinding:
    multiple = len(resources) != 1
    return AttachmentBinding(
        role=role,
        label=role,
        source_kind=(
            AttachmentSourceKind.FILE_SET
            if multiple
            else AttachmentSourceKind.SINGLE_FILE
        ),
        cardinality=(
            AttachmentCardinality.MULTIPLE
            if multiple
            else AttachmentCardinality.SINGLE
        ),
        max_items=None if multiple else 1,
        items=tuple(
            AttachmentItem(
                item_id=f"{role}-{index:04d}",
                label=resource.object_ref.original_name,
                file_ref=_file_ref(resource),
                relative_path=relative_path,
            )
            for index, (resource, relative_path) in enumerate(
                zip(resources, _attachment_relative_paths(resources)),
                start=1,
            )
        ),
    )


def _attachment_relative_paths(
    resources: tuple[ExecutionResource, ...],
) -> tuple[str, ...]:
    names = tuple(
        Path(resource.object_ref.original_name or resource.source_path).name
        or f"attachment-{index:04d}"
        for index, resource in enumerate(resources, start=1)
    )
    counts: dict[str, int] = {}
    for name in names:
        counts[name.casefold()] = counts.get(name.casefold(), 0) + 1
    output: list[str] = []
    for index, name in enumerate(names, start=1):
        if counts[name.casefold()] == 1:
            output.append(name)
            continue
        path = Path(name)
        output.append(f"{path.stem}-{index:04d}{path.suffix}")
    return tuple(output)


def _insert_images(
    source_docx: Path,
    output_docx: Path,
    *,
    record: ExecutionMaterialRecord,
    roles: tuple[str, ...],
    image_policy: Mapping[str, object],
    cache_dir: Path,
) -> None:
    document = Document(source_docx)
    blocks = extract_docx_material_token_blocks(document)
    adaptive = bool(image_policy.get("adaptive", True))
    legacy_show_image_name = bool(image_policy.get("show_image_name", False))
    show_single_image_name = bool(
        image_policy.get("show_single_image_name", legacy_show_image_name)
    )
    show_multi_image_name = bool(
        image_policy.get("show_multi_image_name", legacy_show_image_name)
    )
    watermark = _resolved_watermark(image_policy, record.field_values)
    for role in roles:
        token = material_token(MaterialTokenNamespace.IMAGE, role)
        matches = [
            block
            for block in blocks.blocks
            if token in block.combined_text
        ]
        if len(matches) != 1:
            raise ValueError(
                f"document_batch_image_token_occurrence_invalid:{role}:{len(matches)}"
            )
        block = matches[0]
        direct = blocks.direct_body_blocks.get(block.block_id)
        if direct is None or block.combined_text.strip() != token:
            raise ValueError(
                f"document_batch_image_token_not_isolated_body:{role}"
            )
        paragraph_element, _body_index = direct
        paragraph = Paragraph(paragraph_element, document._body)
        for child in tuple(paragraph_element):
            if child.tag != qn("w:pPr"):
                paragraph_element.remove(child)
        available_width, available_height = _available_image_area(document)
        resources = record.resources.get(role, ())
        if len(resources) == 1:
            show_image_name = show_single_image_name
        elif len(resources) > 1:
            show_image_name = show_multi_image_name
        else:
            show_image_name = False
        # Preserve pagination semantics from the user's template.  Inline
        # pictures already participate in normal text flow; inferring a
        # heading from the previous paragraph can instead attach an unrelated
        # (often empty) paragraph and expose Word/WPS pagination markers.
        for index, resource in enumerate(resources):
            if show_image_name:
                image_name = Path(
                    resource.object_ref.original_name or resource.source_path
                ).name
                paragraph.add_run(image_name or f"{role}-{index + 1:04d}").add_break()
            prepared = prepare_material_image(
                resource.source_path,
                _file_ref(resource),
                watermark,
                cache_dir=cache_dir,
            )
            run = paragraph.add_run()
            width = (
                _adaptive_width(
                    Path(prepared.output_path),
                    available_width,
                    available_height,
                )
                if adaptive
                else None
            )
            run.add_picture(prepared.output_path, width=width)
            if index + 1 < len(resources):
                run.add_break()
    document.save(output_docx)
    Document(output_docx)


def _resolved_watermark(
    image_policy: Mapping[str, object],
    field_values: Mapping[str, str],
) -> ResolvedImageWatermark:
    if not bool(image_policy.get("watermark_enabled", False)):
        return ResolvedImageWatermark.disabled()
    source = str(image_policy.get("watermark_source", "fixed"))
    fixed_text = str(image_policy.get("watermark_text", ""))
    fixed_text = re.sub(
        r"\{\{([^{}:@\s]+)\}\}",
        lambda match: material_token(
            MaterialTokenNamespace.TEXT,
            match.group(1),
        ),
        fixed_text,
    )
    policy = ImageWatermarkPolicy(
        enabled=True,
        text_source=(
            ImageWatermarkTextSource.WORKBENCH_FREE_FIELD
            if source == "free"
            else ImageWatermarkTextSource.FIXED_FIELD
        ),
        text_template=(
            "" if source == "free" else fixed_text
        ),
    )
    return resolve_image_watermark(
        policy,
        field_values,
        runtime_text=str(image_policy.get("runtime_watermark_text", "")),
    )


def _available_image_area(document) -> tuple[int, int]:
    section = document.sections[0]
    available_width = int(
        section.page_width - section.left_margin - section.right_margin
    )
    available_height = int(
        section.page_height
        - section.top_margin
        - section.bottom_margin
        - _IMAGE_BLOCK_VERTICAL_RESERVE_EMU
    )
    return available_width, max(1, available_height)


def _adaptive_width(
    path: Path,
    available_width: int,
    available_height: int,
):
    with Image.open(path) as image:
        dpi = image.info.get("dpi", (96, 96))
        horizontal_dpi = float(dpi[0] or 96)
        vertical_dpi = float(dpi[1] or horizontal_dpi or 96)
        natural_width = int(image.width / horizontal_dpi * _EMU_PER_INCH)
        natural_height = int(image.height / vertical_dpi * _EMU_PER_INCH)
    width_for_height = int(
        natural_width * available_height / max(1, natural_height)
    )
    return max(1, min(natural_width, available_width, width_for_height))


__all__ = [
    "AttachmentDeliveryItem",
    "inspect_template_resource_roles",
    "materialize_record_resources",
    "plan_attachment_delivery_items",
]
