"""Atomic document-batch generation from one execution material snapshot."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document

from src.application.materials import (
    ExecutionMaterialFinalizeRequest,
    ExecutionMaterialSnapshot,
    finalize_execution_material_snapshot,
)
from src.document_batch.material_resources import (
    inspect_template_resource_roles,
    materialize_record_resources,
    plan_attachment_delivery_items,
)
from src.shared.engine.exact_material_placeholders import (
    replace_document_exact_placeholders,
    replace_document_literal_placeholders,
    scan_document_exact_placeholders,
)
from src.shared.engine.material_token_contract import MaterialTokenNamespace
from src.shared.files.content_hash import file_content_revision

_UNSAFE_NAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


@dataclass(frozen=True, slots=True)
class DocumentBatchRecipe:
    recipe_id: str = "document_batch"
    version: int = 1
    conflict_policy: str = "fail"

    def __post_init__(self) -> None:
        if self.recipe_id != "document_batch" or self.version != 1:
            raise ValueError("document_batch_recipe_unsupported")
        if self.conflict_policy != "fail":
            raise ValueError("document_batch_conflict_policy_unsupported")


@dataclass(frozen=True, slots=True)
class DocumentBatchRequest:
    source_paths: tuple[str, ...]
    output_root: str

    def __post_init__(self) -> None:
        if type(self.source_paths) is not tuple or not self.source_paths:
            raise ValueError("document_batch_source_paths_required")
        if type(self.output_root) is not str or not self.output_root:
            raise ValueError("document_batch_output_root_required")


@dataclass(frozen=True, slots=True)
class DocumentBatchArtifact:
    record_id: str
    source_path: str
    final_path: str
    field_values: tuple[tuple[str, str], ...]
    timeline_field_keys: tuple[str, ...] = ()
    attachment_outputs: tuple[tuple[str, str, str, str], ...] = ()


@dataclass(frozen=True, slots=True)
class DocumentBatchPlan:
    plan_id: str
    execution_snapshot_id: str
    execution_snapshot: ExecutionMaterialSnapshot
    output_root: str
    artifacts: tuple[DocumentBatchArtifact, ...]
    issues: tuple[str, ...] = ()

    @property
    def ok(self) -> bool:
        return bool(self.artifacts) and not self.issues

    def run(self) -> dict[str, object]:
        if not self.ok or not self.execution_snapshot.execution_ready:
            return {
                "status": "failed",
                "error_text": "；".join(self.issues) or "document_batch_plan_empty",
                "items": [],
            }
        output_root = Path(self.output_root)
        output_root.mkdir(parents=True, exist_ok=True)
        staging = Path(
            tempfile.mkdtemp(prefix=".alavette-document-batch-", dir=output_root)
        )
        staged: list[
            tuple[
                DocumentBatchArtifact,
                Path,
                tuple[tuple[Path, Path], ...],
            ]
        ] = []
        committed: list[Path] = []
        created_dirs: list[Path] = []

        def ensure_directory(directory: Path) -> None:
            missing: list[Path] = []
            cursor = directory
            while not cursor.exists():
                missing.append(cursor)
                cursor = cursor.parent
            directory.mkdir(parents=True, exist_ok=True)
            created_dirs.extend(reversed(missing))

        try:
            for index, artifact in enumerate(self.artifacts):
                source = Path(artifact.source_path)
                work_dir = staging / f"{index:04d}.work"
                work_dir.mkdir(parents=True)
                fields_stage = work_dir / "fields.docx"
                stage = staging / f"{index:04d}.docx"
                document = Document(source)
                replace_document_exact_placeholders(
                    document,
                    dict(artifact.field_values),
                    timeline_field_keys=artifact.timeline_field_keys,
                )
                replacements = {
                    "{{" + key + "}}": value
                    for key, value in artifact.field_values
                }
                replace_document_literal_placeholders(document, replacements)
                document.save(fields_stage)
                record = next(
                    item
                    for item in self.execution_snapshot.records
                    if item.record_id == artifact.record_id
                )
                deliveries = materialize_record_resources(
                    fields_stage,
                    stage,
                    record=record,
                    resource_domains=self.execution_snapshot.resource_domains,
                    image_policy=self.execution_snapshot.image_policy,
                    work_dir=work_dir,
                )
                # Reopen before any final path becomes visible.
                Document(stage)
                planned = {
                    (source_path, role, relative_path): Path(final_path)
                    for source_path, role, relative_path, final_path
                    in artifact.attachment_outputs
                }
                attachment_stages: list[tuple[Path, Path]] = []
                for delivery_index, delivery in enumerate(deliveries):
                    key = (
                        delivery.source_path,
                        delivery.role,
                        delivery.relative_path,
                    )
                    final_attachment = planned.get(key)
                    if final_attachment is None:
                        raise ValueError(
                            "document_batch_attachment_plan_mismatch:"
                            f"{delivery.role}:{delivery.relative_path}"
                        )
                    attachment_stage = (
                        staging
                        / f"{index:04d}.attachments"
                        / delivery.role
                        / delivery.relative_path
                    )
                    attachment_stage.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(delivery.source_path, attachment_stage)
                    attachment_stages.append(
                        (attachment_stage, final_attachment)
                    )
                staged.append((artifact, stage, tuple(attachment_stages)))
            for artifact, stage, attachment_stages in staged:
                final = Path(artifact.final_path)
                if not final.parent.exists():
                    ensure_directory(final.parent)
                if final.exists():
                    raise FileExistsError(f"document_batch_output_exists:{final}")
                os.replace(stage, final)
                committed.append(final)
                for attachment_stage, final_attachment in attachment_stages:
                    if not final_attachment.parent.exists():
                        ensure_directory(final_attachment.parent)
                    if final_attachment.exists():
                        raise FileExistsError(
                            f"document_batch_output_exists:{final_attachment}"
                        )
                    os.replace(attachment_stage, final_attachment)
                    committed.append(final_attachment)
        except Exception as exc:  # noqa: BLE001 - transaction must roll back any stage failure
            for final in reversed(committed):
                try:
                    final.unlink()
                except OSError:
                    pass
            for directory in reversed(created_dirs):
                try:
                    directory.rmdir()
                except OSError:
                    pass
            return {
                "status": "failed",
                "error_text": f"{type(exc).__name__}: {exc}",
                "items": [],
                "rolled_back_paths": [str(item) for item in committed],
            }
        finally:
            shutil.rmtree(staging, ignore_errors=True)
        items = [
            {
                "record_id": artifact.record_id,
                "source_path": artifact.source_path,
                "output_path": artifact.final_path,
                "attachment_paths": [
                    final_path
                    for _source, _role, _relative, final_path
                    in artifact.attachment_outputs
                ],
                "status": "success",
            }
            for artifact in self.artifacts
        ]
        return {
            "status": "success",
            "summary": f"已原子发布 {len(items)} 份文档",
            "output_path": self.output_root,
            "output_paths": {
                f"{index:04d}": item["output_path"]
                for index, item in enumerate(items)
            },
            "items": items,
            "document_batch_receipt": {
                "plan_id": self.plan_id,
                "execution_snapshot_id": self.execution_snapshot_id,
                "artifact_count": len(items),
            },
        }


def compile_document_batch_plan(
    snapshot: ExecutionMaterialSnapshot,
    request: DocumentBatchRequest,
    *,
    recipe: DocumentBatchRecipe | None = None,
) -> DocumentBatchPlan:
    if not isinstance(snapshot, ExecutionMaterialSnapshot):
        raise TypeError("document_batch_execution_snapshot_required")
    active_recipe = recipe or DocumentBatchRecipe()
    if snapshot.recipe_id != active_recipe.recipe_id:
        raise ValueError("document_batch_snapshot_recipe_mismatch")
    output_root = Path(request.output_root).resolve()
    issues: list[str] = []
    sources: list[Path] = []
    for raw in request.source_paths:
        source = Path(raw).resolve()
        if not source.is_file():
            issues.append(f"document_batch_source_missing:{source}")
        elif source.suffix.casefold() != ".docx":
            issues.append(f"document_batch_source_unsupported:{source}")
        sources.append(source)

    source_field_tokens: dict[Path, tuple[tuple[str, str], ...]] = {}
    for source in sources:
        if not source.is_file() or source.suffix.casefold() != ".docx":
            continue
        try:
            source_field_tokens[source] = tuple(
                (item.key, item.namespace.value)
                for item in scan_document_exact_placeholders(Document(source))
            )
        except Exception as exc:  # noqa: BLE001 - corrupt templates become preflight issues
            issues.append(
                f"document_batch_template_scan_failed:{source}:{type(exc).__name__}"
            )

    try:
        template_resource_roles = dict(
            inspect_template_resource_roles(tuple(sources))
        )
    except ValueError as exc:
        template_resource_roles = {}
        issues.append(str(exc))
    for role, domain in template_resource_roles.items():
        declared_domain = snapshot.resource_domains.get(role)
        if declared_domain is None:
            issues.append(f"document_batch_resource_role_unknown:{role}")
        elif declared_domain != domain:
            issues.append(
                "document_batch_resource_domain_mismatch:"
                f"{role}:{declared_domain}:{domain}"
            )

    artifacts: list[DocumentBatchArtifact] = []
    target_keys: set[str] = set()
    group_outputs_by_record = len(snapshot.records) > 1
    for record in snapshot.records:
        for role in template_resource_roles:
            if not record.resources.get(role, ()):
                issues.append(
                    f"document_batch_resource_required_missing:{record.record_id}:{role}"
                )
        record_output_root = output_root
        if group_outputs_by_record:
            record_output_root = output_root / _safe_name(
                record.display_name,
                fallback=record.record_id,
            )
        for source in sources:
            timeline_keys = set(record.timeline_field_keys)
            for key, namespace in source_field_tokens.get(source, ()):
                if key not in record.field_values:
                    issues.append(
                        f"document_batch_field_required_missing:{record.record_id}:{key}"
                    )
                    continue
                expected_namespace = (
                    MaterialTokenNamespace.TIME.value
                    if key in timeline_keys
                    else MaterialTokenNamespace.TEXT.value
                )
                if namespace != expected_namespace:
                    issues.append(
                        "document_batch_field_namespace_mismatch:"
                        f"{record.record_id}:{key}:{namespace}:{expected_namespace}"
                    )
            final = (
                record_output_root
                / f"{_safe_name(source.stem, fallback='document')}.docx"
            ).resolve()
            try:
                final.relative_to(output_root)
            except ValueError:
                issues.append(f"document_batch_output_outside_root:{final}")
                continue
            key = str(final).casefold()
            if key in target_keys:
                issues.append(f"document_batch_output_collision:{final}")
                continue
            target_keys.add(key)
            if final.exists():
                issues.append(f"document_batch_output_exists:{final}")
            attachment_outputs: list[tuple[str, str, str, str]] = []
            for item in plan_attachment_delivery_items(
                record,
                resource_domains=snapshot.resource_domains,
            ):
                attachment_final = (
                    final.parent
                    / f"{final.stem}_attachments"
                    / item.role
                    / item.relative_path
                ).resolve()
                try:
                    attachment_final.relative_to(output_root)
                except ValueError:
                    issues.append(
                        f"document_batch_output_outside_root:{attachment_final}"
                    )
                    continue
                attachment_key = str(attachment_final).casefold()
                if attachment_key in target_keys:
                    issues.append(
                        f"document_batch_output_collision:{attachment_final}"
                    )
                    continue
                target_keys.add(attachment_key)
                if attachment_final.exists():
                    issues.append(
                        f"document_batch_output_exists:{attachment_final}"
                    )
                attachment_outputs.append(
                    (
                        item.source_path,
                        item.role,
                        item.relative_path,
                        str(attachment_final),
                    )
                )
            artifacts.append(
                DocumentBatchArtifact(
                    record_id=record.record_id,
                    source_path=str(source),
                    final_path=str(final),
                    field_values=tuple(sorted(record.field_values.items())),
                    timeline_field_keys=record.timeline_field_keys,
                    attachment_outputs=tuple(attachment_outputs),
                )
            )
    finalized_snapshot = snapshot
    if not issues:
        template_revision = "sha256:" + hashlib.sha256(
            json.dumps(
                [file_content_revision(item) for item in sources],
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        master_revision = "sha256:" + hashlib.sha256(
            b"alavette.document_batch.master.v1"
        ).hexdigest()
        finalized = finalize_execution_material_snapshot(
            ExecutionMaterialFinalizeRequest(
                snapshot=snapshot,
                template_id="document_batch_source_set",
                template_revision=template_revision,
                master_id="document_batch_master",
                master_revision=master_revision,
                recipe_version=active_recipe.version,
                output_root=str(output_root),
                output_paths=tuple(
                    path
                    for item in artifacts
                    for path in (
                        item.final_path,
                        *(
                            output
                            for _source, _role, _relative, output
                            in item.attachment_outputs
                        ),
                    )
                ),
                supported_field_keys=tuple(
                    sorted(
                        {
                            key
                            for record in snapshot.records
                            for key in record.field_values
                        }
                    )
                ),
                supported_resource_roles=tuple(
                    sorted(snapshot.resource_domains)
                ),
            )
        )
        if not finalized.ok or finalized.snapshot is None:
            issues.extend(item.code for item in finalized.issues)
        else:
            finalized_snapshot = finalized.snapshot
    payload = {
        "snapshot_id": finalized_snapshot.snapshot_id,
        "recipe": [active_recipe.recipe_id, active_recipe.version],
        "artifacts": [
            [
                item.record_id,
                item.source_path,
                item.final_path,
                list(item.timeline_field_keys),
                [list(value) for value in item.attachment_outputs],
            ]
            for item in artifacts
        ],
    }
    plan_id = "sha256:" + hashlib.sha256(
        json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    return DocumentBatchPlan(
        plan_id=plan_id,
        execution_snapshot_id=finalized_snapshot.snapshot_id,
        execution_snapshot=finalized_snapshot,
        output_root=str(output_root),
        artifacts=tuple(artifacts),
        issues=tuple(issues),
    )


def _safe_name(value: str, *, fallback: str) -> str:
    normalized = _UNSAFE_NAME.sub("_", str(value or "")).strip(" .")
    return normalized[:120] or fallback
