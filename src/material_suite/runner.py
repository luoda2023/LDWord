"""Atomic multi-artifact generation for one frozen material-suite plan."""

from __future__ import annotations

import json
import os
import shutil
import tempfile
import uuid
from collections.abc import Callable, Mapping
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

from src.config.asset_resolution import file_content_revision
from src.material_suite.plan import (
    MaterialSuiteRunPlan,
    SuiteArtifactPlan,
    SuiteRecordPlan,
    scan_suite_template_placeholders,
)
from src.shared.engine.exact_material_placeholders import (
    replace_document_literal_placeholders,
)

ProgressCallback = Callable[[int, int, str], None]
CancelCheck = Callable[[], bool]


class MaterialSuiteGenerationRunner:
    """Generate each record in staging and publish it only as a complete set."""

    def __init__(self, plan: MaterialSuiteRunPlan):
        self.plan = plan

    def run(
        self,
        progress_cb: ProgressCallback,
        cancel_check: CancelCheck,
    ) -> dict[str, object]:
        if not self.plan.ok:
            return _preflight_failure(self.plan)

        run_id = f"suite-{uuid.uuid4().hex[:12]}"
        total_steps = max(1, self.plan.artifact_count)
        completed_steps = 0
        item_results: list[dict[str, object]] = []
        cancelled = False
        output_root = Path(self.plan.output_root)
        output_root.mkdir(parents=True, exist_ok=True)

        source_issue = _revision_issue(self.plan)
        if source_issue:
            return _failed_payload(
                source_issue,
                run_id=run_id,
                output_root=output_root,
                records=self.plan.execution_units,
            )

        execution_units = self.plan.execution_units
        for record in execution_units:
            if cancel_check():
                cancelled = True
                break
            progress_cb(
                completed_steps,
                total_steps,
                f"准备 {record.profile_name}",
            )
            result, generated_steps = self._generate_record(
                record,
                progress_cb=progress_cb,
                cancel_check=cancel_check,
                completed_steps=completed_steps,
                total_steps=total_steps,
                run_id=run_id,
            )
            completed_steps += generated_steps
            item_results.append(result)
            if result.get("status") == "cancelled":
                cancelled = True
                break

        pending_ids = [
            record.profile_id for record in execution_units[len(item_results) :]
        ]
        failed_ids = [
            str(item.get("profile_id") or "")
            for item in item_results
            if item.get("status") == "failed"
        ]
        interrupted_ids = [
            str(item.get("profile_id") or "")
            for item in item_results
            if item.get("status") == "cancelled"
        ]
        retry_ids = list(dict.fromkeys([*failed_ids, *interrupted_ids, *pending_ids]))
        succeeded = [item for item in item_results if item.get("status") == "success"]
        reports = _write_run_reports(
            output_root,
            run_id=run_id,
            plan=self.plan,
            items=item_results,
            pending_profile_ids=pending_ids,
            cancelled=cancelled,
        )
        if cancelled:
            status = "cancelled"
        elif failed_ids and succeeded:
            status = "partial_success"
        elif failed_ids:
            status = "failed"
        else:
            status = "success"
        error_text = ""
        if failed_ids:
            error_text = f"{len(failed_ids)} 条资料生成失败，可按失败记录重试"
        elif cancelled:
            error_text = "已取消；已完整发布的记录保留，当前半成品已清理"
        summary = (
            f"成套生成完成：成功 {len(succeeded)} 条，"
            f"失败 {len(failed_ids)} 条，待执行 {len(pending_ids)} 条"
        )
        return {
            "status": status,
            "summary": summary,
            "output_path": str(output_root),
            "output_paths": {"suite_root": str(output_root)},
            "report_paths": reports,
            "failed_count": len(retry_ids),
            "error_text": error_text,
            "items": item_results,
            "pending_profile_ids": pending_ids,
            "failed_items": [
                item for item in item_results if item.get("status") == "failed"
            ],
            "batch_isolation": {
                "history_run_id": run_id,
                "attempt_number": 1,
                "profile_count": len(self.plan.records),
                "execution_unit_count": len(execution_units),
                "succeeded_count": len(succeeded),
                "failed_count": len(failed_ids),
                "interrupted_count": len(interrupted_ids),
                "pending_count": len(pending_ids),
                "retry_eligible_profile_ids": retry_ids,
                "details": [
                    (
                        f"{item.get('profile_name')}: "
                        f"{item.get('status')} "
                        f"{item.get('error_text') or ''}"
                    ).strip()
                    for item in item_results
                ],
            },
            "material_suite_receipt": {
                "package_id": self.plan.package_id,
                "package_revision": self.plan.package_revision,
                "template_bundle_revision": self.plan.template_bundle.revision,
                "record_count": len(self.plan.records),
                "shared_unit_count": len(self.plan.shared_units),
                "artifact_count": self.plan.artifact_count,
                "recipe_id": self.plan.recipe.recipe_id,
                "selected_record_ids": list(
                    self.plan.request.selected_record_ids
                ),
            },
        }

    def _generate_record(
        self,
        record: SuiteRecordPlan,
        *,
        progress_cb: ProgressCallback,
        cancel_check: CancelCheck,
        completed_steps: int,
        total_steps: int,
        run_id: str,
    ) -> tuple[dict[str, object], int]:
        target = Path(record.output_dir)
        target.parent.mkdir(parents=True, exist_ok=True)
        staging = Path(
            tempfile.mkdtemp(
                prefix=f".alavette-{run_id}-",
                dir=str(target.parent),
            )
        )
        generated: list[dict[str, str]] = []
        steps = 0
        try:
            if target.exists():
                raise FileExistsError("目标记录目录已存在，未覆盖")
            for artifact in record.artifacts:
                if cancel_check():
                    return (
                        _record_result(
                            record,
                            status="cancelled",
                            error_text="生成期间取消，暂存结果已清理",
                        ),
                        steps,
                    )
                current = completed_steps + steps
                progress_cb(
                    current,
                    total_steps,
                    f"{record.profile_name} · {artifact.label}",
                )
                source = Path(artifact.source_path)
                if file_content_revision(source) != artifact.source_revision:
                    raise RuntimeError(f"模板已变化：{artifact.label}")
                staged_target = staging / artifact.target_relative_path
                staged_target.parent.mkdir(parents=True, exist_ok=True)
                _generate_artifact(artifact, staged_target)
                unresolved = scan_suite_template_placeholders(
                    staged_target,
                    artifact.kind,
                )
                if unresolved:
                    raise RuntimeError(
                        f"{artifact.label} 仍有未替换字段：" + "、".join(unresolved)
                    )
                generated.append(
                    {
                        "artifact_id": artifact.artifact_id,
                        "label": artifact.label,
                        "kind": artifact.kind,
                        "path": str(target / artifact.target_relative_path),
                        "source_revision": artifact.source_revision,
                        "output_revision": file_content_revision(staged_target),
                    }
                )
                steps += 1
                progress_cb(
                    completed_steps + steps,
                    total_steps,
                    f"{record.profile_name} · 已完成 {artifact.label}",
                )

            manifest = {
                "kind": "alavette.material_suite_record_receipt",
                "version": 1,
                "run_id": run_id,
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "package_id": self.plan.package_id,
                "package_name": self.plan.package_name,
                "package_revision": self.plan.package_revision,
                "template_bundle_revision": self.plan.template_bundle.revision,
                "profile_id": record.profile_id,
                "profile_name": record.profile_name,
                "group_id": record.group_id,
                "emit_scope": record.emit_scope,
                "route_key": record.route_key,
                "timeline_field_keys": list(record.timeline_field_keys),
                "frozen_values": dict(record.frozen_values),
                "value_provenance": [
                    {
                        "field": key,
                        "scope": scope,
                        "owner_id": owner_id,
                        "source": source,
                    }
                    for key, scope, owner_id, source in record.provenance
                ],
                "source_locator": dict(record.source_locator),
                "artifacts": generated,
            }
            (staging / "生成清单.json").write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            os.replace(staging, target)
            return (
                _record_result(
                    record,
                    status="success",
                    output_path=str(target),
                    artifacts=generated,
                ),
                steps,
            )
        except Exception as exc:  # noqa: BLE001 - isolate one record failure
            return (
                _record_result(
                    record,
                    status="failed",
                    error_text=f"{type(exc).__name__}: {exc}",
                ),
                steps,
            )
        finally:
            if staging.exists():
                shutil.rmtree(staging)


def _generate_artifact(
    artifact: SuiteArtifactPlan,
    target: Path,
) -> None:
    replacements = artifact.replacement_map()
    if artifact.missing_placeholders:
        raise ValueError("缺少字段：" + "、".join(artifact.missing_placeholders))
    if artifact.kind == "docx":
        document = Document(artifact.source_path)
        _replace_word_tokens(document, replacements)
        document.save(str(target))
        return
    if artifact.kind == "xlsx":
        source = Path(artifact.source_path)
        workbook = load_workbook(
            source,
            data_only=False,
            keep_vba=source.suffix.lower() == ".xlsm",
        )
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if not isinstance(cell.value, str):
                            continue
                        value = cell.value
                        for token, replacement in replacements.items():
                            value = value.replace(token, replacement)
                        cell.value = value
            workbook.save(target)
        finally:
            workbook.close()
        return
    raise ValueError(f"不支持的成套模板类型：{artifact.kind}")


def _replace_word_tokens(document, replacements: Mapping[str, str]) -> None:
    """Replace across Word runs and emit VBA-compatible manual line breaks."""

    normalized = {
        str(token): _normalize_line_breaks(value)
        for token, value in replacements.items()
    }
    replace_document_literal_placeholders(document, normalized)
    _materialize_manual_breaks(document)


def _normalize_line_breaks(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n")


def _materialize_manual_breaks(document) -> None:
    """Translate embedded newlines into Word ``w:br`` (Shift+Enter)."""

    roots = [document.element]
    roots.extend(
        root
        for part in document.part.package.parts
        if (root := getattr(part, "_element", None)) is not None
    )
    seen: set[int] = set()
    for root in roots:
        if id(root) in seen:
            continue
        seen.add(id(root))
        for text_node in tuple(root.iter(qn("w:t"))):
            raw = str(text_node.text or "")
            if "\n" not in raw and "\v" not in raw:
                continue
            parts = raw.replace("\v", "\n").split("\n")
            parent = text_node.getparent()
            if parent is None or parent.tag != qn("w:r"):
                text_node.text = " ".join(parts)
                continue
            index = parent.index(text_node)
            parent.remove(text_node)
            for part_index, part in enumerate(parts):
                replacement_text = OxmlElement("w:t")
                if part.startswith(" ") or part.endswith(" "):
                    replacement_text.set(
                        "{http://www.w3.org/XML/1998/namespace}space",
                        "preserve",
                    )
                replacement_text.text = part
                parent.insert(index, replacement_text)
                index += 1
                if part_index < len(parts) - 1:
                    parent.insert(index, OxmlElement("w:br"))
                    index += 1


def _revision_issue(plan: MaterialSuiteRunPlan) -> str:
    source = Path(plan.package_source_path)
    if (
        str(source)
        and source.is_file()
        and file_content_revision(source) != plan.package_revision
    ):
        # The package may be JSON, CSV, or XLSX; all are file-revision based.
        return "资料包在预检后已发生变化，请重新预检"
    for artifact in plan.template_bundle.artifacts:
        source = Path(artifact.source_path)
        if not source.is_file():
            return f"模板已不存在：{artifact.label}"
        if file_content_revision(source) != artifact.source_revision:
            return f"模板在预检后已发生变化：{artifact.label}"
    return ""


def _record_result(
    record: SuiteRecordPlan,
    *,
    status: str,
    output_path: str = "",
    error_text: str = "",
    artifacts: list[dict[str, str]] | None = None,
) -> dict[str, object]:
    return {
        "profile_id": record.profile_id,
        "profile_name": record.profile_name,
        "route_key": record.route_key,
        "group_id": record.group_id,
        "emit_scope": record.emit_scope,
        "status": status,
        "output_path": output_path,
        "error_text": error_text,
        "artifacts": list(artifacts or []),
    }


def _preflight_failure(plan: MaterialSuiteRunPlan) -> dict[str, object]:
    record_issues = [
        f"{record.profile_name}：{'；'.join(record.issues)}"
        for record in plan.execution_units
        if record.issues
    ]
    issues = [*plan.issues, *record_issues]
    return {
        "status": "failed",
        "summary": "成套生成预检未通过",
        "output_path": "",
        "report_paths": [],
        "failed_count": len(plan.execution_units),
        "error_text": "；".join(issues),
        "diagnostics_count": len(issues),
        "diagnostics_summary": f"成套生成预检发现 {len(issues)} 项阻断",
        "failed_items": [
            {
                "profile_id": record.profile_id,
                "profile_name": record.profile_name,
                "status": "failed",
                "error_text": "；".join(record.issues),
            }
            for record in plan.execution_units
            if record.issues
        ],
        "batch_isolation": {
            "profile_count": len(plan.records),
            "execution_unit_count": len(plan.execution_units),
            "failed_count": len(plan.execution_units),
            "retry_eligible_profile_ids": [
                record.profile_id for record in plan.execution_units
            ],
        },
    }


def _failed_payload(
    error_text: str,
    *,
    run_id: str,
    output_root: Path,
    records: tuple[SuiteRecordPlan, ...],
) -> dict[str, object]:
    return {
        "status": "failed",
        "summary": "成套生成源文件校验失败",
        "output_path": str(output_root),
        "report_paths": [],
        "failed_count": len(records),
        "error_text": error_text,
        "batch_isolation": {
            "history_run_id": run_id,
            "profile_count": len(records),
            "failed_count": len(records),
            "retry_eligible_profile_ids": [item.profile_id for item in records],
        },
    }


def _write_run_reports(
    output_root: Path,
    *,
    run_id: str,
    plan: MaterialSuiteRunPlan,
    items: list[dict[str, object]],
    pending_profile_ids: list[str],
    cancelled: bool,
) -> list[str]:
    payload = {
        "kind": "alavette.material_suite_run_report",
        "version": 1,
        "run_id": run_id,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "cancelled": cancelled,
        "package_id": plan.package_id,
        "package_name": plan.package_name,
        "package_revision": plan.package_revision,
        "template_bundle": asdict(plan.template_bundle),
        "generation_recipe": asdict(plan.recipe),
        "run_request": asdict(plan.request),
        "package_inspection": asdict(plan.inspection),
        "output_root": str(output_root),
        "items": items,
        "pending_profile_ids": pending_profile_ids,
    }
    report_dir = output_root / "_生成报告"
    report_dir.mkdir(parents=True, exist_ok=True)
    json_path = report_dir / f"{run_id}.json"
    md_path = report_dir / f"{run_id}.md"
    json_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    lines = [
        "# 成套生成报告",
        "",
        f"- 运行编号：{run_id}",
        f"- 资料包：{plan.package_name}",
        f"- 记录数：{len(plan.records)}",
        f"- 模板产物数：{plan.artifact_count}",
        "",
        "## 记录结果",
        "",
    ]
    for item in items:
        lines.append(
            f"- {item.get('profile_name')}：{item.get('status')}"
            + (f"（{item.get('error_text')}）" if item.get("error_text") else "")
        )
    if pending_profile_ids:
        lines.extend(["", f"待执行：{', '.join(pending_profile_ids)}"])
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return [str(json_path), str(md_path)]


__all__ = ["MaterialSuiteGenerationRunner"]
