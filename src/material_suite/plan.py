"""Compile layered packages into frozen, auditable suite run plans."""

from __future__ import annotations

import hashlib
import json
import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field, replace
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from src.application.materials import (
    ExecutionMaterialFinalizeRequest,
    ExecutionMaterialSnapshot,
    finalize_execution_material_snapshot,
)
from src.shared.files.content_hash import file_content_revision
from src.shared.engine.docx_material_tokens import (
    extract_docx_material_token_blocks,
)
from src.shared.engine.material_token_contract import parse_material_token

SUITE_TOKEN_PATTERN = re.compile(r"\{\{([^{}\r\n]+)\}\}")
SUPPORTED_WORD_SUFFIXES = frozenset({".docx"})
SUPPORTED_EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm"})
TEMPLATE_SUITE_KIND = "alavette.template_suite"
TEMPLATE_SUITE_VERSION = 1
EMIT_SCOPES = frozenset({"package_once", "group_once", "record_once"})

_ROUTE_FIELD_KEYS = (
    "产品名称",
    "product_name",
    "product",
    "产品类型",
    "product_type",
    "路线",
    "route",
)
_ROUTE_IDENTITY_FIELD_KEYS = (
    "产品编码",
    "产品编号",
    "product_code",
    "product_id",
    "route_id",
)
_PROJECT_FIELD_KEYS = (
    "项目名称",
    "project_name",
    "项目编号",
    "project_code",
    "entity_name",
)
_COMMON_FIELD_ALIASES = {
    "项目名称": ("project_name", "entity_name"),
    "项目编号": ("project_code",),
    "项目来源": ("project_source",),
    "项目开始日期": ("project_start_date", "start_date"),
    "项目结束日期": ("project_end_date", "end_date"),
    "目标成本": ("target_cost",),
    "产品名称": ("product_name", "product"),
    "产品编码": ("产品编号", "product_code", "product_id", "route_id"),
}


@dataclass(frozen=True, slots=True)
class GenerationRecipe:
    """Business definition of one package-driven generation run."""

    recipe_id: str = "material-suite.default"
    recipe_name: str = "默认成套生成"
    active_states: tuple[str, ...] = ("active",)
    required_record_fields: tuple[str, ...] = ()
    route_field_keys: tuple[str, ...] = _ROUTE_FIELD_KEYS
    record_name_field_keys: tuple[str, ...] = _PROJECT_FIELD_KEYS
    output_path_template: str = "{route}/{record_name}"
    failure_policy: str = "isolate_record"


@dataclass(frozen=True, slots=True)
class MaterialSuiteRunRequest:
    output_root: str = ""
    requested_by: str = "workbench"


@dataclass(frozen=True, slots=True)
class PackageRecordInspection:
    record_id: str
    record_name: str
    group_id: str
    lifecycle_state: str
    selected: bool
    readiness: str
    issues: tuple[str, ...] = ()
    source_locator: tuple[tuple[str, str], ...] = ()
    route_key: str = ""
    output_dir: str = ""
    artifact_count: int = 0
    skipped_artifact_count: int = 0
    timeline_count: int = 0


@dataclass(frozen=True, slots=True)
class PackageInspection:
    candidate_count: int = 0
    active_count: int = 0
    draft_count: int = 0
    disabled_count: int = 0
    archived_count: int = 0
    selected_count: int = 0
    executable_count: int = 0
    blocked_count: int = 0
    records: tuple[PackageRecordInspection, ...] = ()


@dataclass(frozen=True, slots=True)
class SuiteArtifactSpec:
    artifact_id: str
    label: str
    source_path: str
    kind: str
    group_name: str
    route_key: str = ""
    route_aliases: tuple[str, ...] = ()
    required: bool = True
    source_revision: str = ""
    emit_scope: str = "record_once"
    target_relative_path: str = ""


@dataclass(frozen=True, slots=True)
class SuiteTemplateBundle:
    root_path: str = ""
    bundle_id: str = ""
    bundle_name: str = ""
    revision: str = ""
    artifacts: tuple[SuiteArtifactSpec, ...] = ()
    issues: tuple[str, ...] = ()
    manifest_path: str = ""
    schema_version: int = TEMPLATE_SUITE_VERSION
    discovery_mode: str = "legacy_adapter"

    @property
    def route_keys(self) -> tuple[str, ...]:
        return tuple(
            dict.fromkeys(item.route_key for item in self.artifacts if item.route_key)
        )


@dataclass(frozen=True, slots=True)
class SuiteArtifactPlan:
    artifact_id: str
    label: str
    source_path: str
    source_revision: str
    target_relative_path: str
    kind: str
    group_name: str
    required: bool
    emit_scope: str = "record_once"
    placeholders: tuple[str, ...] = ()
    replacements: tuple[tuple[str, str], ...] = ()
    missing_placeholders: tuple[str, ...] = ()
    status: str = "ready"
    issues: tuple[str, ...] = ()

    def replacement_map(self) -> dict[str, str]:
        return dict(self.replacements)

    @property
    def executable(self) -> bool:
        return self.status == "ready"


@dataclass(frozen=True, slots=True)
class SuiteRecordPlan:
    record_id: str
    record_name: str
    route_key: str
    output_dir: str
    frozen_values: tuple[tuple[str, str], ...]
    timeline_field_keys: tuple[str, ...]
    artifacts: tuple[SuiteArtifactPlan, ...]
    issues: tuple[str, ...] = ()
    group_id: str = ""
    emit_scope: str = "record_once"
    provenance: tuple[tuple[str, str, str, str], ...] = ()
    source_locator: tuple[tuple[str, str], ...] = ()

    @property
    def unit_id(self) -> str:
        return self.record_id

    @property
    def unit_name(self) -> str:
        return self.record_name

    def value_map(self) -> dict[str, str]:
        return dict(self.frozen_values)

    @property
    def ready_artifacts(self) -> tuple[SuiteArtifactPlan, ...]:
        return tuple(item for item in self.artifacts if item.executable)

    @property
    def artifact_count(self) -> int:
        return len(self.ready_artifacts)

    @property
    def skipped_artifact_count(self) -> int:
        return sum(item.status == "skipped" for item in self.artifacts)


@dataclass(frozen=True, slots=True)
class MaterialSuiteRunPlan:
    package_id: str
    package_name: str
    package_revision: str
    execution_snapshot_id: str
    execution_snapshot: ExecutionMaterialSnapshot
    template_bundle: SuiteTemplateBundle
    output_root: str
    records: tuple[SuiteRecordPlan, ...] = ()
    issues: tuple[str, ...] = ()
    shared_units: tuple[SuiteRecordPlan, ...] = ()
    inspection: PackageInspection = field(default_factory=PackageInspection)
    recipe: GenerationRecipe = field(default_factory=GenerationRecipe)
    request: MaterialSuiteRunRequest = field(default_factory=MaterialSuiteRunRequest)

    @property
    def ok(self) -> bool:
        return (
            not self.issues
            and bool(self.execution_units)
            and not self.blocked_units
            and self.execution_snapshot.execution_ready
        )

    @property
    def planned_units(self) -> tuple[SuiteRecordPlan, ...]:
        return (*self.shared_units, *self.records)

    @property
    def execution_units(self) -> tuple[SuiteRecordPlan, ...]:
        return tuple(
            item
            for item in self.planned_units
            if not item.issues and item.ready_artifacts
        )

    @property
    def blocked_units(self) -> tuple[SuiteRecordPlan, ...]:
        return tuple(
            item
            for item in self.planned_units
            if item.issues or (item.artifacts and not item.ready_artifacts)
        )

    @property
    def artifact_count(self) -> int:
        return sum(item.artifact_count for item in self.execution_units)

    @property
    def output_directory_count(self) -> int:
        return len(self.execution_units)


def discover_material_suite_bundle(root_path: str | Path) -> SuiteTemplateBundle:
    """Load a versioned manifest or adapt the legacy directory convention."""

    root = Path(root_path).expanduser()
    if not root.is_dir():
        return SuiteTemplateBundle(
            root_path=str(root),
            issues=("模板套件目录不存在",),
        )
    template_root = root / "模板文件" if (root / "模板文件").is_dir() else root
    manifest = template_root / "template-suite.json"
    if manifest.is_file():
        return _load_manifest_bundle(template_root, manifest)
    return _discover_legacy_bundle(template_root)


def compile_material_suite_plan(
    snapshot: ExecutionMaterialSnapshot,
    template_bundle: SuiteTemplateBundle,
    *,
    output_root: str | Path,
    recipe: GenerationRecipe | None = None,
    request: MaterialSuiteRunRequest | None = None,
) -> MaterialSuiteRunPlan:
    """Compile one already-bound execution snapshot into a suite plan."""

    generation_recipe = recipe or GenerationRecipe()
    output = Path(output_root).expanduser()
    run_request = request or MaterialSuiteRunRequest(
        output_root=str(output),
    )
    global_issues = list(template_bundle.issues)
    if not isinstance(snapshot, ExecutionMaterialSnapshot):
        raise TypeError("material_suite_execution_snapshot_required")
    if snapshot.recipe_id != "material_suite":
        global_issues.append(
            f"执行快照配方不匹配：{snapshot.recipe_id}"
        )
    if not snapshot.records:
        global_issues.append("执行快照中没有资料记录")
    if not str(output_root or "").strip():
        global_issues.append("尚未选择输出目录")
    elif output.exists() and not output.is_dir():
        global_issues.append("输出位置不是目录")

    scans, scan_issues = _scan_template_bundle_once(template_bundle)
    global_issues.extend(scan_issues)
    has_record_artifacts = any(
        item.emit_scope == "record_once" for item in template_bundle.artifacts
    )
    inspection_rows: list[PackageRecordInspection] = []
    records: list[SuiteRecordPlan] = []
    seen_outputs: set[str] = set()
    groups_by_id = {item.group_id: item for item in snapshot.groups}
    for record in snapshot.records:
        readiness_issues: list[str] = []
        aliases = _build_alias_values(record.field_values, {})
        readiness_issues.extend(
            f"缺少必需字段：{key}"
            for key in generation_recipe.required_record_fields
            if not str(_lookup_value(key, aliases) or "").strip()
        )
        values = dict(record.field_values)
        values.setdefault("record_id", record.record_id)
        values.setdefault("record_name", record.display_name)
        values.setdefault("package_id", snapshot.package_ref.package_id)
        values.setdefault("package_name", snapshot.package_display_name)
        aliases = _build_alias_values(values, {})
        group = groups_by_id.get(record.group_id)
        route_key = (
            _first_value(group.field_values, generation_recipe.route_field_keys)
            if group is not None
            else _first_value(aliases, generation_recipe.route_field_keys)
        )
        route_candidates = _route_candidates(
            route_key,
            group_name=(group.display_name if group is not None else ""),
            aliases=aliases,
            recipe=generation_recipe,
        )
        project_name = (
            _first_value(aliases, generation_recipe.record_name_field_keys)
            or record.display_name
            or record.record_id
        )
        target_dir = _record_output_dir(
            output,
            recipe=generation_recipe,
            route_key=route_key,
            record_name=project_name,
            record_id=record.record_id,
            group_id=(group.group_id if group is not None else ""),
        )
        record_issues = list(readiness_issues)
        normalized_output = str(target_dir.resolve(strict=False)).casefold()
        if normalized_output in seen_outputs:
            record_issues.append("输出目录与另一条资料重复")
        seen_outputs.add(normalized_output)
        if target_dir.exists():
            record_issues.append("目标记录目录已存在；为防止覆盖，请更换输出目录")
        matched = _matched_artifacts(
            template_bundle.artifacts,
            emit_scope="record_once",
            route_candidates=route_candidates,
        )
        if has_record_artifacts:
            _append_route_issues(
                record_issues,
                template_bundle,
                route_key=route_key,
                route_candidates=route_candidates,
                matched=matched,
                emit_scope="record_once",
            )
        artifact_plans = _compile_artifact_plans(
            matched,
            scans=scans,
            aliases=aliases,
            issues=record_issues,
        )
        if artifact_plans and not any(item.executable for item in artifact_plans):
            record_issues.append("该资料没有可执行的成套产物")
        records.append(
            SuiteRecordPlan(
                record_id=record.record_id,
                record_name=record.display_name or project_name,
                route_key=route_key,
                output_dir=str(target_dir),
                frozen_values=tuple(sorted(values.items())),
                timeline_field_keys=record.timeline_field_keys,
                artifacts=artifact_plans,
                issues=tuple(dict.fromkeys(record_issues)),
                group_id=(group.group_id if group is not None else ""),
                provenance=tuple(
                    (key, owner, "", "")
                    for key, owner in sorted(record.field_owners.items())
                ),
                source_locator=(),
            )
        )
        inspection_rows.append(
            PackageRecordInspection(
                record_id=record.record_id,
                record_name=record.display_name,
                group_id=(group.group_id if group is not None else ""),
                lifecycle_state="active",
                selected=True,
                readiness=("blocked" if record_issues else "ready"),
                issues=tuple(dict.fromkeys(record_issues)),
                route_key=route_key,
                output_dir=str(target_dir),
                artifact_count=sum(
                    item.executable for item in artifact_plans
                ),
                skipped_artifact_count=sum(
                    item.status == "skipped" for item in artifact_plans
                ),
            )
        )

    shared_units = _compile_shared_units(
        snapshot,
        template_bundle=template_bundle,
        scans=scans,
        output=output,
    )
    all_units = (*shared_units, *records)
    output_owners: dict[str, list[str]] = {}
    for unit in all_units:
        output_key = str(
            Path(unit.output_dir).resolve(strict=False)
        ).casefold()
        output_owners.setdefault(output_key, []).append(unit.unit_id)
    duplicate_outputs = {
        output_key
        for output_key, owner_ids in output_owners.items()
        if len(owner_ids) > 1
    }

    def mark_output_collision(unit: SuiteRecordPlan) -> SuiteRecordPlan:
        output_key = str(
            Path(unit.output_dir).resolve(strict=False)
        ).casefold()
        if output_key not in duplicate_outputs:
            return unit
        return replace(
            unit,
            issues=tuple(
                dict.fromkeys(
                    (*unit.issues, "输出目录与另一交付单元重复")
                )
            ),
        )

    records = [mark_output_collision(item) for item in records]
    shared_units = tuple(
        mark_output_collision(item) for item in shared_units
    )
    record_plans_by_id = {item.record_id: item for item in records}
    inspection_rows = [
        replace(
            item,
            readiness="blocked",
            issues=record_plans_by_id[item.record_id].issues,
            route_key=record_plans_by_id[item.record_id].route_key,
            output_dir=record_plans_by_id[item.record_id].output_dir,
            artifact_count=record_plans_by_id[item.record_id].artifact_count,
            skipped_artifact_count=(
                record_plans_by_id[item.record_id].skipped_artifact_count
            ),
            timeline_count=len(
                record_plans_by_id[item.record_id].timeline_field_keys
            ),
        )
        if item.record_id in record_plans_by_id
        and record_plans_by_id[item.record_id].issues
        else replace(
            item,
            route_key=record_plans_by_id[item.record_id].route_key,
            output_dir=record_plans_by_id[item.record_id].output_dir,
            artifact_count=record_plans_by_id[item.record_id].artifact_count,
            skipped_artifact_count=(
                record_plans_by_id[item.record_id].skipped_artifact_count
            ),
            timeline_count=len(
                record_plans_by_id[item.record_id].timeline_field_keys
            ),
        )
        if item.record_id in record_plans_by_id
        else item
        for item in inspection_rows
    ]
    selected_count = len(inspection_rows)
    blocked_count = sum(
        item.selected and item.readiness != "ready" for item in inspection_rows
    )
    executable_count = sum(
        not item.issues and bool(item.ready_artifacts) for item in records
    )
    executable_shared_count = sum(
        not item.issues and bool(item.ready_artifacts) for item in shared_units
    )
    if selected_count == 0:
        global_issues.append("本次运行没有选中的活动记录")
    elif executable_count + executable_shared_count == 0:
        global_issues.append("本次选择中没有通过预检的活动记录")
    inspection = PackageInspection(
        candidate_count=len(snapshot.records),
        active_count=len(snapshot.records),
        selected_count=selected_count,
        executable_count=executable_count,
        blocked_count=blocked_count,
        records=tuple(inspection_rows),
    )
    finalized_snapshot = snapshot
    if not global_issues and not any(item.issues for item in all_units):
        output_paths = tuple(
            str((Path(unit.output_dir) / artifact.target_relative_path).resolve())
            for unit in all_units
            for artifact in unit.ready_artifacts
        )
        bundle_revision = template_bundle.revision
        if not re.fullmatch(r"sha256:[0-9a-f]{64}", bundle_revision):
            bundle_revision = "sha256:" + hashlib.sha256(
                json.dumps(
                    [
                        item.source_revision
                        for item in template_bundle.artifacts
                    ],
                    sort_keys=True,
                    separators=(",", ":"),
                ).encode("utf-8")
            ).hexdigest()
        finalized = finalize_execution_material_snapshot(
            ExecutionMaterialFinalizeRequest(
                snapshot=snapshot,
                template_id=template_bundle.bundle_id or "material_suite_bundle",
                template_revision=bundle_revision,
                master_id="material_suite_master",
                master_revision="sha256:"
                + hashlib.sha256(
                    b"alavette.material_suite.master.v1"
                ).hexdigest(),
                recipe_version=1,
                output_root=str(output.resolve()),
                output_paths=output_paths,
                supported_field_keys=tuple(
                    sorted(
                        {
                            key
                            for record in snapshot.records
                            for key in record.field_values
                        }
                    )
                ),
            )
        )
        if finalized.ok and finalized.snapshot is not None:
            finalized_snapshot = finalized.snapshot
        else:
            global_issues.extend(item.code for item in finalized.issues)
    return MaterialSuiteRunPlan(
        package_id=snapshot.package_ref.package_id,
        package_name=snapshot.package_display_name,
        package_revision=snapshot.package_ref.revision,
        execution_snapshot_id=finalized_snapshot.snapshot_id,
        execution_snapshot=finalized_snapshot,
        template_bundle=template_bundle,
        output_root=str(output),
        records=tuple(records),
        shared_units=shared_units,
        issues=tuple(dict.fromkeys(global_issues)),
        inspection=inspection,
        recipe=generation_recipe,
        request=MaterialSuiteRunRequest(
            output_root=str(output),
            requested_by=run_request.requested_by,
        ),
    )


def scan_suite_template_placeholders(
    source_path: str | Path,
    kind: str,
) -> tuple[str, ...]:
    """Return all exact and legacy placeholders in first-seen order."""

    source = Path(source_path)
    texts: Iterable[str]
    if kind == "docx":
        document = Document(str(source))
        texts = (
            block.combined_text
            for block in extract_docx_material_token_blocks(document).blocks
        )
    elif kind == "xlsx":
        workbook = load_workbook(
            source,
            read_only=True,
            data_only=False,
            keep_vba=source.suffix.lower() == ".xlsm",
        )
        try:
            texts = tuple(
                str(cell.value)
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str)
            )
        finally:
            workbook.close()
    else:
        raise ValueError(f"unsupported suite artifact kind: {kind}")
    found: list[str] = []
    for text in texts:
        for match in SUITE_TOKEN_PATTERN.finditer(str(text or "")):
            token = match.group(0)
            if token not in found:
                found.append(token)
    return tuple(found)


def _scan_template_bundle_once(
    bundle: SuiteTemplateBundle,
) -> tuple[dict[str, tuple[str, ...]], list[str]]:
    scans: dict[str, tuple[str, ...]] = {}
    issues: list[str] = []
    for spec in bundle.artifacts:
        try:
            scans[spec.artifact_id] = scan_suite_template_placeholders(
                spec.source_path,
                spec.kind,
            )
        except Exception as exc:  # noqa: BLE001 - template boundary
            issues.append(f"模板“{spec.label}”无法读取：{type(exc).__name__}: {exc}")
    return scans, issues


def _compile_shared_units(
    snapshot: ExecutionMaterialSnapshot,
    *,
    template_bundle: SuiteTemplateBundle,
    scans: Mapping[str, tuple[str, ...]],
    output: Path,
) -> tuple[SuiteRecordPlan, ...]:
    units: list[SuiteRecordPlan] = []
    package_specs = [
        item for item in template_bundle.artifacts if item.emit_scope == "package_once"
    ]
    if package_specs and snapshot.records:
        values = {
            **dict(snapshot.package_field_values),
            "package_id": snapshot.package_ref.package_id,
            "package_name": snapshot.package_display_name,
        }
        issues: list[str] = []
        aliases = _build_alias_values(values, {})
        artifacts = _compile_artifact_plans(
            package_specs,
            scans=scans,
            aliases=aliases,
            issues=issues,
        )
        if artifacts and not any(item.executable for item in artifacts):
            issues.append("资料包公共产物均不满足生成条件")
        target = output / "_资料包公共"
        if target.exists():
            issues.append("资料包公共产物目录已存在；为防止覆盖，请更换输出目录")
        units.append(
            SuiteRecordPlan(
                record_id=f"package:{snapshot.package_ref.package_id}",
                record_name=snapshot.package_display_name,
                route_key="",
                output_dir=str(target),
                frozen_values=tuple(sorted(values.items())),
                timeline_field_keys=(),
                artifacts=artifacts,
                issues=tuple(dict.fromkeys(issues)),
                emit_scope="package_once",
            )
        )

    group_specs = [
        item for item in template_bundle.artifacts if item.emit_scope == "group_once"
    ]
    if not group_specs:
        return tuple(units)
    for group in snapshot.groups:
        values = {
            **dict(group.field_values),
            "package_id": snapshot.package_ref.package_id,
            "package_name": snapshot.package_display_name,
            "group_id": group.group_id,
            "group_name": group.display_name,
        }
        aliases = _build_alias_values(values, {})
        issues: list[str] = []
        route_key = _first_value(aliases, _ROUTE_FIELD_KEYS)
        route_candidates = _route_candidates(
            route_key,
            group_name=group.display_name,
            aliases=aliases,
            recipe=GenerationRecipe(),
        )
        matched = _matched_artifacts(
            group_specs,
            emit_scope="group_once",
            route_candidates=route_candidates,
        )
        _append_route_issues(
            issues,
            template_bundle,
            route_key=route_key,
            route_candidates=route_candidates,
            matched=matched,
            emit_scope="group_once",
        )
        artifacts = _compile_artifact_plans(
            matched,
            scans=scans,
            aliases=aliases,
            issues=issues,
        )
        if artifacts and not any(item.executable for item in artifacts):
            issues.append("该分组没有可执行的公共产物")
        target = (
            output / _safe_component(route_key or group.display_name) / "_分组公共"
        )
        if target.exists():
            issues.append("分组公共产物目录已存在；为防止覆盖，请更换输出目录")
        units.append(
            SuiteRecordPlan(
                record_id=f"group:{group.group_id}",
                record_name=group.display_name,
                route_key=route_key,
                output_dir=str(target),
                frozen_values=tuple(sorted(values.items())),
                timeline_field_keys=(),
                artifacts=artifacts,
                issues=tuple(dict.fromkeys(issues)),
                group_id=group.group_id,
                emit_scope="group_once",
            )
        )
    return tuple(units)


def _compile_artifact_plans(
    specs: Sequence[SuiteArtifactSpec],
    *,
    scans: Mapping[str, tuple[str, ...]],
    aliases: Mapping[str, str],
    issues: list[str],
) -> tuple[SuiteArtifactPlan, ...]:
    plans: list[SuiteArtifactPlan] = []
    seen_targets: set[str] = set()
    for spec in specs:
        if spec.artifact_id not in scans:
            continue
        placeholders = scans[spec.artifact_id]
        replacements: list[tuple[str, str]] = []
        missing: list[str] = []
        for token in placeholders:
            key = _token_key(token)
            value = _lookup_value(key, aliases)
            if value is None:
                missing.append(token)
            else:
                replacements.append((token, value))
        artifact_issues: list[str] = []
        if missing and spec.required:
            missing_message = (
                f"模板“{spec.label}”缺少字段："
                + "、".join(_token_key(token) for token in missing)
            )
            artifact_issues.append(missing_message)
            issues.append(
                missing_message
            )
        target_relative_path = _artifact_target_path(spec)
        target_key = target_relative_path.casefold()
        if target_key in seen_targets:
            issues.append(f"多个模板会写入同一目标文件：{target_relative_path}")
        seen_targets.add(target_key)
        plans.append(
            SuiteArtifactPlan(
                artifact_id=spec.artifact_id,
                label=spec.label,
                source_path=spec.source_path,
                source_revision=spec.source_revision,
                target_relative_path=target_relative_path,
                kind=spec.kind,
                group_name=spec.group_name,
                required=spec.required,
                emit_scope=spec.emit_scope,
                placeholders=placeholders,
                replacements=tuple(replacements),
                missing_placeholders=tuple(missing),
                status=(
                    "blocked"
                    if missing and spec.required
                    else "skipped"
                    if missing
                    else "ready"
                ),
                issues=tuple(artifact_issues),
            )
        )
    return tuple(plans)


def _matched_artifacts(
    artifacts: Sequence[SuiteArtifactSpec],
    *,
    emit_scope: str,
    route_candidates: Sequence[str],
) -> list[SuiteArtifactSpec]:
    candidate_keys = {
        _normalized_route_key(value)
        for value in route_candidates
        if _normalized_route_key(value)
    }
    return [
        spec
        for spec in artifacts
        if spec.emit_scope == emit_scope
        and (
            not spec.route_key
            or bool(
                candidate_keys
                & {
                    _normalized_route_key(value)
                    for value in (spec.route_key, *spec.route_aliases)
                    if _normalized_route_key(value)
                }
            )
        )
    ]


def _append_route_issues(
    issues: list[str],
    bundle: SuiteTemplateBundle,
    *,
    route_key: str,
    route_candidates: Sequence[str],
    matched: Sequence[SuiteArtifactSpec],
    emit_scope: str,
) -> None:
    scoped_route_keys = tuple(
        dict.fromkeys(
            route_value
            for item in bundle.artifacts
            if item.emit_scope == emit_scope and item.route_key and item.required
            for route_value in (item.route_key, *item.route_aliases)
        )
    )
    candidate_keys = {
        _normalized_route_key(value)
        for value in route_candidates
        if _normalized_route_key(value)
    }
    scoped_keys = {
        _normalized_route_key(value)
        for value in scoped_route_keys
        if _normalized_route_key(value)
    }
    if scoped_route_keys and not candidate_keys:
        issues.append("资料缺少产品名称/路线字段，无法映射对应的测试表")
    elif (
        candidate_keys
        and scoped_route_keys
        and not candidate_keys & scoped_keys
    ):
        issues.append(
            f"产品路线“{route_key or route_candidates[0]}”没有对应的成套模板"
        )
    if not matched:
        issues.append("该资料没有可映射的成套模板")


def _route_candidates(
    route_key: str,
    *,
    group_name: str,
    aliases: Mapping[str, str],
    recipe: GenerationRecipe,
) -> tuple[str, ...]:
    values = [
        _first_value(aliases, _ROUTE_IDENTITY_FIELD_KEYS),
        route_key,
        group_name,
        _first_value(aliases, recipe.route_field_keys),
    ]
    return tuple(
        dict.fromkeys(
            value
            for item in values
            if (value := str(item or "").strip())
        )
    )


def _normalized_route_key(value: object) -> str:
    return re.sub(r"[\s_-]+", "", str(value or "")).casefold()


def _record_output_dir(
    output: Path,
    *,
    recipe: GenerationRecipe,
    route_key: str,
    record_name: str,
    record_id: str,
    group_id: str,
) -> Path:
    replacements = {
        "route": _safe_component(route_key) if route_key else "",
        "record_name": _safe_component(record_name),
        "record_id": _safe_component(record_id),
        "group_id": _safe_component(group_id) if group_id else "",
    }
    try:
        rendered = recipe.output_path_template.format_map(replacements)
    except (KeyError, ValueError, IndexError) as exc:
        raise ValueError(f"生成配方输出目录模板无效：{exc}") from exc
    parts = [
        _safe_component(part)
        for part in re.split(r"[/\\]+", rendered)
        if str(part).strip()
    ]
    if not parts:
        parts = [replacements["record_name"]]
    target = output.joinpath(*parts)
    root_resolved = output.resolve(strict=False)
    try:
        target.resolve(strict=False).relative_to(root_resolved)
    except ValueError as exc:
        raise ValueError("生成配方输出目录越界") from exc
    return target


def _artifact_target_path(spec: SuiteArtifactSpec) -> str:
    if spec.target_relative_path:
        rendered = spec.target_relative_path.replace(
            "{source_name}",
            Path(spec.source_path).name,
        )
        parts = [
            _safe_component(part)
            for part in re.split(r"[/\\]+", rendered)
            if str(part).strip()
        ]
        if not parts:
            raise ValueError(f"模板“{spec.label}”的目标路径为空")
        return str(Path(*parts))
    return str(Path(spec.group_name) / Path(spec.source_path).name)


def _load_manifest_bundle(
    template_root: Path,
    manifest: Path,
) -> SuiteTemplateBundle:
    issues: list[str] = []
    try:
        payload = json.loads(manifest.read_text(encoding="utf-8"))
    except Exception as exc:  # noqa: BLE001 - selected manifest boundary
        return SuiteTemplateBundle(
            root_path=str(template_root),
            manifest_path=str(manifest),
            discovery_mode="manifest",
            issues=(f"模板套件清单无法读取：{type(exc).__name__}: {exc}",),
        )
    if not isinstance(payload, Mapping):
        issues.append("模板套件清单根节点必须是对象")
        payload = {}
    if payload.get("kind") != TEMPLATE_SUITE_KIND:
        issues.append("模板套件清单 kind 无效")
    if payload.get("version") != TEMPLATE_SUITE_VERSION:
        issues.append("模板套件清单版本不受支持")
    raw_artifacts = payload.get("artifacts", [])
    if not isinstance(raw_artifacts, list):
        issues.append("模板套件清单 artifacts 必须是列表")
        raw_artifacts = []
    artifacts: list[SuiteArtifactSpec] = []
    seen_ids: set[str] = set()
    root_resolved = template_root.resolve()
    for index, raw in enumerate(raw_artifacts, start=1):
        if not isinstance(raw, Mapping):
            issues.append(f"模板套件第 {index} 个产物定义不是对象")
            continue
        source_text = str(raw.get("source", "") or "").strip()
        source = (template_root / source_text).resolve()
        try:
            source.relative_to(root_resolved)
        except ValueError:
            issues.append(f"模板套件第 {index} 个产物路径越界")
            continue
        if not source.is_file():
            issues.append(f"模板不存在：{source_text}")
            continue
        artifact_id = str(raw.get("artifact_id", "") or "").strip()
        if not artifact_id:
            artifact_id = hashlib.sha1(source_text.encode("utf-8")).hexdigest()[:16]
        if artifact_id in seen_ids:
            issues.append(f"模板产物编号重复：{artifact_id}")
            continue
        seen_ids.add(artifact_id)
        kind = str(raw.get("kind", "") or "").strip().lower()
        if not kind:
            kind = "docx" if source.suffix.lower() == ".docx" else "xlsx"
        if kind not in {"docx", "xlsx"}:
            issues.append(f"模板“{source_text}”类型不支持：{kind}")
            continue
        emit_scope = str(raw.get("emit_scope", "record_once") or "record_once").strip()
        if emit_scope not in EMIT_SCOPES:
            issues.append(f"模板“{source_text}”生成作用域无效：{emit_scope}")
            continue
        raw_route_aliases = raw.get("route_aliases", [])
        if not isinstance(raw_route_aliases, list):
            issues.append(f"模板“{source_text}”的 route_aliases 必须是列表")
            continue
        route_aliases = tuple(
            dict.fromkeys(
                alias
                for item in raw_route_aliases
                if (alias := str(item or "").strip())
            )
        )
        artifacts.append(
            SuiteArtifactSpec(
                artifact_id=artifact_id,
                label=str(raw.get("label", "") or source.stem),
                source_path=str(source),
                kind=kind,
                group_name=str(
                    raw.get("group_name", "")
                    or ("Word文档" if kind == "docx" else "Excel测试表")
                ),
                route_key=str(raw.get("route_key", "") or "").strip(),
                route_aliases=route_aliases,
                required=bool(raw.get("required", True)),
                source_revision=file_content_revision(source),
                emit_scope=emit_scope,
                target_relative_path=str(raw.get("target", "") or "").strip(),
            )
        )
    if not artifacts:
        issues.append("模板套件中没有可生成的 DOCX/XLSX 模板")
    digest = hashlib.sha256(manifest.read_bytes())
    for item in artifacts:
        digest.update(item.artifact_id.encode("utf-8"))
        digest.update(item.source_revision.encode("ascii", errors="ignore"))
    return SuiteTemplateBundle(
        root_path=str(template_root),
        bundle_id=str(
            payload.get("bundle_id", "") or _safe_component(template_root.name)
        ),
        bundle_name=str(payload.get("bundle_name", "") or template_root.name),
        revision=digest.hexdigest(),
        artifacts=tuple(artifacts),
        issues=tuple(dict.fromkeys(issues)),
        manifest_path=str(manifest),
        discovery_mode="manifest",
    )


def _discover_legacy_bundle(template_root: Path) -> SuiteTemplateBundle:
    issues: list[str] = []
    artifacts: list[SuiteArtifactSpec] = []
    word_roots = _existing_dirs(
        template_root / "Word文档",
        template_root / "Word",
        template_root / "文档模板",
    )
    for word_root in word_roots:
        for source in _iter_template_files(word_root, SUPPORTED_WORD_SUFFIXES):
            artifacts.append(
                _artifact_spec(
                    source,
                    template_root=template_root,
                    kind="docx",
                    group_name="Word文档",
                )
            )
    excel_roots = _existing_dirs(
        template_root / "Excel测试表",
        template_root / "Excel",
        template_root / "表格模板",
    )
    for excel_root in excel_roots:
        for source in _iter_template_files(excel_root, SUPPORTED_EXCEL_SUFFIXES):
            artifacts.append(
                _artifact_spec(
                    source,
                    template_root=template_root,
                    kind="xlsx",
                    group_name="Excel测试表",
                )
            )
    for child in sorted(template_root.iterdir(), key=lambda item: item.name.casefold()):
        if not child.is_dir() or not child.name.endswith("_测试表"):
            continue
        route_key = child.name[: -len("_测试表")].strip()
        for source in _iter_template_files(child, SUPPORTED_EXCEL_SUFFIXES):
            artifacts.append(
                _artifact_spec(
                    source,
                    template_root=template_root,
                    kind="xlsx",
                    group_name="Excel测试表",
                    route_key=route_key,
                )
            )
    if not artifacts:
        for source in sorted(
            template_root.iterdir(),
            key=lambda item: item.name.casefold(),
        ):
            suffix = source.suffix.lower()
            if not source.is_file():
                continue
            if suffix in SUPPORTED_WORD_SUFFIXES:
                artifacts.append(
                    _artifact_spec(
                        source,
                        template_root=template_root,
                        kind="docx",
                        group_name="Word文档",
                    )
                )
            elif suffix in SUPPORTED_EXCEL_SUFFIXES:
                artifacts.append(
                    _artifact_spec(
                        source,
                        template_root=template_root,
                        kind="xlsx",
                        group_name="Excel测试表",
                    )
                )
    if not artifacts:
        issues.append("模板套件中没有可生成的 DOCX/XLSX 模板")
    if not any(item.kind == "docx" for item in artifacts):
        issues.append("模板套件缺少 Word 文档模板")
    artifacts = _deduplicate_artifacts(artifacts)
    digest = hashlib.sha256()
    for item in artifacts:
        digest.update(item.artifact_id.encode("utf-8"))
        digest.update(item.source_revision.encode("ascii", errors="ignore"))
    bundle_name = template_root.name
    return SuiteTemplateBundle(
        root_path=str(template_root),
        bundle_id=_safe_component(bundle_name),
        bundle_name=bundle_name,
        revision=digest.hexdigest(),
        artifacts=tuple(artifacts),
        issues=tuple(issues),
        discovery_mode="legacy_adapter",
    )


def _artifact_spec(
    source: Path,
    *,
    template_root: Path,
    kind: str,
    group_name: str,
    route_key: str = "",
) -> SuiteArtifactSpec:
    relative = source.relative_to(template_root).as_posix()
    artifact_id = hashlib.sha1(relative.encode("utf-8")).hexdigest()[:16]
    return SuiteArtifactSpec(
        artifact_id=artifact_id,
        label=source.stem,
        source_path=str(source),
        kind=kind,
        group_name=group_name,
        route_key=route_key,
        source_revision=file_content_revision(source),
    )


def _iter_template_files(root: Path, suffixes: frozenset[str]) -> list[Path]:
    return sorted(
        (
            item
            for item in root.rglob("*")
            if item.is_file()
            and item.suffix.lower() in suffixes
            and not item.name.startswith("~$")
        ),
        key=lambda item: item.as_posix().casefold(),
    )


def _existing_dirs(*candidates: Path) -> tuple[Path, ...]:
    return tuple(item for item in candidates if item.is_dir())


def _deduplicate_artifacts(
    artifacts: Sequence[SuiteArtifactSpec],
) -> list[SuiteArtifactSpec]:
    result: list[SuiteArtifactSpec] = []
    seen: set[str] = set()
    for item in artifacts:
        key = str(Path(item.source_path).resolve(strict=False)).casefold()
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def _token_key(token: str) -> str:
    try:
        return parse_material_token(token).identifier
    except (TypeError, ValueError):
        return str(token)[2:-2].strip()


def _build_alias_values(
    values: Mapping[str, str],
    field_aliases: Mapping[str, str],
) -> dict[str, str]:
    result = {str(key): str(value) for key, value in values.items()}
    for alias, canonical in field_aliases.items():
        if canonical in result:
            result.setdefault(str(alias), result[canonical])
    for alias, candidates in _COMMON_FIELD_ALIASES.items():
        value = _first_value(result, candidates)
        if value:
            result.setdefault(alias, value)
    return result


def _lookup_value(key: str, values: Mapping[str, str]) -> str | None:
    if key in values:
        return str(values[key])
    folded = key.casefold()
    matches = [
        str(value) for name, value in values.items() if name.casefold() == folded
    ]
    return matches[0] if len(matches) == 1 else None


def _first_value(values: Mapping[str, str], keys: Sequence[str]) -> str:
    for key in keys:
        value = _lookup_value(key, values)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _safe_component(value: object) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", str(value or "").strip())
    text = text.rstrip(". ")
    return text[:120] or "未命名"


__all__ = [
    "EMIT_SCOPES",
    "GenerationRecipe",
    "MaterialSuiteRunPlan",
    "MaterialSuiteRunRequest",
    "PackageInspection",
    "PackageRecordInspection",
    "SuiteArtifactPlan",
    "SuiteArtifactSpec",
    "SuiteRecordPlan",
    "SuiteTemplateBundle",
    "compile_material_suite_plan",
    "discover_material_suite_bundle",
    "scan_suite_template_placeholders",
]
