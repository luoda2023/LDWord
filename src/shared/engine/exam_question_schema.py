"""Runtime validation for structured exam question sources."""

from __future__ import annotations

import ast
import copy
import json
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field, replace
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4

from docx import Document
from docx.shared import Inches

from src.config.fixed_layout_contract import FixedLayoutRowHeightPolicy
from src.config.master_preflight import check_master_preflight
from src.config.material_schema_registry import get_material_schema, resolve_material_schema_ids
from src.shared.io.artifact_publication import StagedArtifact, publish_staged_artifacts
from src.shared.engine.fixed_layout_tables import (
    apply_fixed_layout_row_height_policy,
    row_height_state,
)
from src.shared.engine.exam_paper_style import write_exam_paper_docx_files
from src.shared.engine.ooxml_ops import set_inline_shape_alt_text
from src.shared.engine.content_visibility import (
    executable_content_visibility_rule_selector,
)


EXAM_SCHEMA_ID = "exam_items_v1"
EXAM_FAMILY_ID = "exam_teaching"

_SOURCE_KEYS: tuple[str, ...] = (
    "exam_items",
    "exam_paper",
    "paper",
    "paper_json",
    "question_schema",
    "questions_json",
    "questions",
    "sections",
)
_TITLE_KEYS: tuple[str, ...] = ("paper_title", "title", "name")
_SECTION_KEYS: tuple[str, ...] = ("sections", "parts")
_QUESTION_KEYS: tuple[str, ...] = ("questions", "items")
_ANSWER_KEYS: tuple[str, ...] = ("answer", "answers", "correct_answer", "correctAnswers")
_ANALYSIS_KEYS: tuple[str, ...] = ("analysis", "explanation", "solution")
_KNOWLEDGE_KEYS: tuple[str, ...] = ("knowledge_points", "knowledgePoints", "tags")

_MARKDOWN_TITLE_RE = re.compile(r"^\s*#\s+(.+?)\s*$")
_MARKDOWN_SECTION_RE = re.compile(r"^\s*##+\s+(.+?)\s*$")
_MARKDOWN_QUESTION_RE = re.compile(r"^\s*(\d+)[\.\u3001]\s*(.+?)\s*$")
_MARKDOWN_OPTION_RE = re.compile(r"^\s*([A-Ha-h])[\.\uFF0E\u3001]\s*(.+?)\s*$")
_MARKDOWN_FIELD_RE = re.compile(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s*[:：]\s*(.*?)\s*$")
_MARKDOWN_SCORE_RE = re.compile(r"[\(（]\s*(\d+(?:\.\d+)?)\s*分\s*[\)）]")
_MARKDOWN_METADATA_PAIR_RE = re.compile(
    r"(科目|学科|年级|考试时间|时间|满分|总分)\s*[:：]\s*([^　]+?)(?=(?:　+| {2,}|\s*(?:科目|学科|年级|考试时间|时间|满分|总分)\s*[:：])|$)"
)
_MARKDOWN_ANSWER_SPLIT_RE = re.compile(r"^(.*?)(?:[。；;]?\s*解析\s*[:：]\s*(.+))$")
_MARKDOWN_ANSWER_GROUP_RE = re.compile(r"^\s*[一二三四五六七八九十]+[\u3001\.\uFF0E]\s*.+")
_MARKDOWN_ALLOWED_ANSWER_AREA_KINDS = {"lines", "free", "blank"}


@dataclass(frozen=True, slots=True)
class ExamQuestionSchemaIssue:
    """One validation issue in a structured exam source."""

    path: str
    kind: str
    message: str
    severity: str = "warning"
    expected: str = ""
    observed: str = ""


@dataclass(frozen=True, slots=True)
class ExamQuestionSchemaSummary:
    """Compact counts used by reports and Workbench summaries."""

    section_count: int = 0
    question_count: int = 0
    answered_question_count: int = 0
    analysis_count: int = 0
    knowledge_point_count: int = 0
    figure_count: int = 0
    scored_question_count: int = 0
    declared_total_score: float | None = None
    computed_total_score: float | None = None
    question_types: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class ExamQuestionSchemaValidationResult:
    """Structured runtime evidence for exam source validation."""

    schema_id: str = ""
    family_id: str = ""
    status: str = "not_applicable"
    source_key: str = ""
    summary: ExamQuestionSchemaSummary = field(default_factory=ExamQuestionSchemaSummary)
    issues: tuple[ExamQuestionSchemaIssue, ...] = ()

    @property
    def has_issues(self) -> bool:
        return bool(self.issues)

    @property
    def error_count(self) -> int:
        return sum(1 for issue in self.issues if issue.severity == "error")

    @property
    def warning_count(self) -> int:
        return sum(1 for issue in self.issues if issue.severity != "error")

    @property
    def manual_confirmation_required(self) -> bool:
        return bool(self.issues)

    def to_dict(self) -> dict[str, object]:
        return {
            "schema_id": self.schema_id,
            "family_id": self.family_id,
            "status": self.status,
            "source_key": self.source_key,
            "manual_confirmation_required": self.manual_confirmation_required,
            "summary": {
                "section_count": self.summary.section_count,
                "question_count": self.summary.question_count,
                "answered_question_count": self.summary.answered_question_count,
                "analysis_count": self.summary.analysis_count,
                "knowledge_point_count": self.summary.knowledge_point_count,
                "figure_count": self.summary.figure_count,
                "scored_question_count": self.summary.scored_question_count,
                "declared_total_score": self.summary.declared_total_score,
                "computed_total_score": self.summary.computed_total_score,
                "question_types": list(self.summary.question_types),
            },
            "issue_count": len(self.issues),
            "error_count": self.error_count,
            "warning_count": self.warning_count,
            "issues": [
                {
                    "path": issue.path,
                    "kind": issue.kind,
                    "severity": issue.severity,
                    "expected": issue.expected,
                    "observed": issue.observed,
                    "message": issue.message,
                }
                for issue in self.issues
            ],
        }


@dataclass(frozen=True, slots=True)
class ExamMarkdownImportIssue:
    """One issue found while importing a Markdown exam source."""

    path: str
    kind: str
    message: str
    severity: str = "warning"
    observed: str = ""

    def to_dict(self) -> dict[str, object]:
        return {
            "path": self.path,
            "kind": self.kind,
            "severity": self.severity,
            "observed": self.observed,
            "message": self.message,
        }


@dataclass(frozen=True, slots=True)
class ExamMarkdownImportResult:
    """Normalized result for one Markdown exam source import."""

    payload: dict[str, object] = field(default_factory=dict)
    summary: ExamQuestionSchemaSummary = field(default_factory=ExamQuestionSchemaSummary)
    issues: tuple[ExamMarkdownImportIssue, ...] = ()
    source_path: str = ""
    source_excerpt: str = ""
    status: str = "ok"

    @property
    def has_issues(self) -> bool:
        return bool(self.issues)

    @property
    def error_count(self) -> int:
        return sum(1 for issue in self.issues if issue.severity == "error")

    @property
    def warning_count(self) -> int:
        return sum(1 for issue in self.issues if issue.severity != "error")

    def to_dict(self) -> dict[str, object]:
        return {
            "status": self.status,
            "source_path": self.source_path,
            "source_excerpt": self.source_excerpt,
            "summary": {
                "section_count": self.summary.section_count,
                "question_count": self.summary.question_count,
                "answered_question_count": self.summary.answered_question_count,
                "analysis_count": self.summary.analysis_count,
                "knowledge_point_count": self.summary.knowledge_point_count,
                "figure_count": self.summary.figure_count,
                "scored_question_count": self.summary.scored_question_count,
                "declared_total_score": self.summary.declared_total_score,
                "computed_total_score": self.summary.computed_total_score,
                "question_types": list(self.summary.question_types),
            },
            "issue_count": len(self.issues),
            "error_count": self.error_count,
            "warning_count": self.warning_count,
            "issues": [issue.to_dict() for issue in self.issues],
            "payload": copy.deepcopy(self.payload),
        }


@dataclass(frozen=True, slots=True)
class ExamRenderedVersion:
    """One Word version rendered from the structured exam source."""

    preset_id: str
    label: str
    docx_path: str
    hidden_selectors: tuple[str, ...] = ()
    visible_question_count: int = 0
    visible_answer_count: int = 0
    visible_analysis_count: int = 0
    fixed_layout_kind: str = ""
    fixed_layout_row_count: int = 0
    fixed_layout_column_count: int = 0
    fixed_layout_row_height_twips: int | None = None
    fixed_layout_row_height_rule: str = ""
    question_asset_count: int = 0
    rendered_question_asset_count: int = 0
    missing_question_asset_count: int = 0
    question_asset_alt_text_count: int = 0
    rendered_question_asset_alt_text_count: int = 0
    missing_question_asset_alt_text_count: int = 0

    def to_dict(self) -> dict[str, object]:
        return {
            "preset_id": self.preset_id,
            "label": self.label,
            "docx_path": self.docx_path,
            "hidden_selectors": list(self.hidden_selectors),
            "visible_question_count": self.visible_question_count,
            "visible_answer_count": self.visible_answer_count,
            "visible_analysis_count": self.visible_analysis_count,
            "fixed_layout_kind": self.fixed_layout_kind,
            "fixed_layout_row_count": self.fixed_layout_row_count,
            "fixed_layout_column_count": self.fixed_layout_column_count,
            "fixed_layout_row_height_twips": self.fixed_layout_row_height_twips,
            "fixed_layout_row_height_rule": self.fixed_layout_row_height_rule,
            "question_asset_count": self.question_asset_count,
            "rendered_question_asset_count": self.rendered_question_asset_count,
            "missing_question_asset_count": self.missing_question_asset_count,
            "question_asset_alt_text_count": self.question_asset_alt_text_count,
            "rendered_question_asset_alt_text_count": (
                self.rendered_question_asset_alt_text_count
            ),
            "missing_question_asset_alt_text_count": (
                self.missing_question_asset_alt_text_count
            ),
        }


@dataclass(frozen=True, slots=True)
class ExamMasterEvidence:
    """Traceability for the master used by exam-paper assembly."""

    mode_id: str = "exam"
    master_id: str = ""
    master_label: str = ""
    master_source_type: str = ""
    master_docx_path: str = ""
    master_version: str = ""
    manifest_path: str = ""
    placeholder_contract_status: str = "not_checked"
    required_placeholders: tuple[str, ...] = ()
    optional_placeholders: tuple[str, ...] = ()
    runtime_fields: tuple[str, ...] = ()
    generated_fields: tuple[str, ...] = ()
    missing_required_placeholders: tuple[str, ...] = ()

    def to_dict(self) -> dict[str, object]:
        return {
            "mode_id": self.mode_id,
            "master_id": self.master_id,
            "master_label": self.master_label,
            "master_source_type": self.master_source_type,
            "master_docx_path": self.master_docx_path,
            "master_version": self.master_version,
            "manifest_path": self.manifest_path,
            "placeholder_contract_status": self.placeholder_contract_status,
            "required_placeholders": list(self.required_placeholders),
            "optional_placeholders": list(self.optional_placeholders),
            "runtime_fields": list(self.runtime_fields),
            "generated_fields": list(self.generated_fields),
            "missing_required_placeholders": list(self.missing_required_placeholders),
        }


@dataclass(frozen=True, slots=True)
class ExamDeliveryRuntimeResult:
    """Runtime evidence for JSON -> Markdown preview -> Word versions."""

    schema_id: str = ""
    family_id: str = ""
    status: str = "not_applicable"
    source_key: str = ""
    markdown_preview_path: str = ""
    markdown_preview_excerpt: str = ""
    rendered_versions: tuple[ExamRenderedVersion, ...] = ()
    master_evidence: ExamMasterEvidence | None = None
    skipped_reason: str = ""

    @property
    def version_count(self) -> int:
        return len(self.rendered_versions)

    def to_dict(self) -> dict[str, object]:
        return {
            "schema_id": self.schema_id,
            "family_id": self.family_id,
            "status": self.status,
            "source_key": self.source_key,
            "markdown_preview_path": self.markdown_preview_path,
            "markdown_preview_excerpt": self.markdown_preview_excerpt,
            "version_count": self.version_count,
            "skipped_reason": self.skipped_reason,
            "master_evidence": (
                self.master_evidence.to_dict()
                if self.master_evidence is not None
                else {}
            ),
            "rendered_versions": [
                version.to_dict() for version in self.rendered_versions
            ],
        }


def inspect_exam_question_schema(config) -> ExamQuestionSchemaValidationResult:
    """Validate the exam question payload carried by a resolved config.

    The first runtime slice is intentionally schema-driven and report-only. It
    runs for ``exam_items_v1`` / ``exam_teaching`` configs, validates that one
    structured source can support multiple output versions, and leaves AI
    content quality or complex diagrams to plugin/manual-gate boundaries.
    """

    schema_ids = _config_material_schema_ids(config)
    family_id = _exam_family_for_schema_ids(schema_ids)
    if EXAM_SCHEMA_ID not in schema_ids and family_id != EXAM_FAMILY_ID:
        return ExamQuestionSchemaValidationResult()

    payload, source_key = _extract_exam_payload(getattr(config, "entity_data", {}) or {})
    if payload is None:
        return ExamQuestionSchemaValidationResult(
            schema_id=EXAM_SCHEMA_ID,
            family_id=family_id or EXAM_FAMILY_ID,
            status="error",
            issues=(
                ExamQuestionSchemaIssue(
                    path="entity_data",
                    kind="missing_structured_source",
                    severity="error",
                    expected="exam_items/paper_json/questions/sections",
                    message=(
                        "Exam schema requires one structured question source; "
                        "answers must remain in that source for delivery presets."
                    ),
                ),
            ),
        )

    normalized_payload = _normalize_exam_payload(payload)
    return _validate_exam_payload(
        normalized_payload,
        source_key=source_key,
        family_id=family_id or EXAM_FAMILY_ID,
        metadata=getattr(config, "entity_data", {}) or {},
    )


def build_exam_delivery_runtime(
    config,
    *,
    output_dir: Path | str,
    source_stem: str = "exam",
    validation: ExamQuestionSchemaValidationResult | None = None,
    master=None,
) -> ExamDeliveryRuntimeResult:
    """Render first-slice exam artifacts from one structured question source.

    This intentionally stays conservative: it only renders when the structured
    source has no validation errors.  AI content quality, complex diagrams, and
    subject-specific asset replacement remains plugin/manual or later work.
    """

    validation = validation or inspect_exam_question_schema(config)
    if validation.status == "not_applicable":
        return ExamDeliveryRuntimeResult()

    payload, source_key = _exam_payload_from_config(config)
    if payload is None:
        return ExamDeliveryRuntimeResult(
            schema_id=EXAM_SCHEMA_ID,
            family_id=validation.family_id or EXAM_FAMILY_ID,
            status="blocked",
            source_key=validation.source_key or source_key,
            skipped_reason="missing_structured_source",
        )

    safe_stem = _safe_artifact_stem(source_stem)
    markdown = render_exam_markdown_preview(payload)
    exam_paper_config = getattr(config, "exam_paper", None)
    master_evidence = _exam_master_evidence(master)

    if validation.error_count:
        return ExamDeliveryRuntimeResult(
            schema_id=EXAM_SCHEMA_ID,
            family_id=validation.family_id or EXAM_FAMILY_ID,
            status="blocked",
            source_key=validation.source_key or source_key,
            markdown_preview_excerpt=_preview_excerpt(markdown),
            master_evidence=master_evidence,
            skipped_reason="schema_validation_errors",
        )

    if exam_paper_config is not None and master is None:
        return ExamDeliveryRuntimeResult(
            schema_id=EXAM_SCHEMA_ID,
            family_id=validation.family_id or EXAM_FAMILY_ID,
            status="blocked",
            source_key=validation.source_key or source_key,
            markdown_preview_excerpt=_preview_excerpt(markdown),
            skipped_reason="execution_master_required",
        )

    output_root = Path(output_dir).expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    artifact_dir = output_root / "exam_runtime"
    with TemporaryDirectory(prefix=".exam-runtime-", dir=str(output_root)) as temp_dir:
        staging_dir = Path(temp_dir).resolve()
        staged_markdown_path = staging_dir / f"{safe_stem}_preview.md"
        staged_markdown_path.write_text(markdown, encoding="utf-8")

        versions: list[ExamRenderedVersion] = []
        if exam_paper_config is not None:
            versions.extend(
                _render_exam_master_docx(
                    payload,
                    exam_paper_config,
                    staging_dir,
                    source_stem=safe_stem,
                    summary=getattr(validation, "summary", None),
                    master=master,
                )
            )
        else:
            for preset in _delivery_presets_for_runtime(config):
                if not _preset_final_docx_enabled(preset):
                    continue
                preset_id = _preset_value(preset, "preset_id", "final").strip() or "final"
                label = _preset_value(preset, "label", preset_id).strip() or preset_id
                docx_path = (
                    staging_dir
                    / f"{safe_stem}_{_safe_artifact_stem(preset_id)}.docx"
                )
                versions.append(
                    _render_exam_version_docx(payload, preset, docx_path, label=label)
                )

        if not versions:
            return ExamDeliveryRuntimeResult(
                schema_id=EXAM_SCHEMA_ID,
                family_id=validation.family_id or EXAM_FAMILY_ID,
                status="blocked",
                source_key=validation.source_key or source_key,
                markdown_preview_excerpt=_preview_excerpt(markdown),
                master_evidence=master_evidence,
                skipped_reason="no_final_docx_delivery_preset",
            )

        staged_artifacts = [
            StagedArtifact(
                artifact_id="preview",
                stage_path=staged_markdown_path,
                final_path=artifact_dir / staged_markdown_path.name,
            )
        ]
        version_artifact_ids: list[str] = []
        for index, version in enumerate(versions):
            staged_docx = Path(version.docx_path).expanduser().resolve()
            try:
                relative_docx = staged_docx.relative_to(staging_dir)
            except ValueError as exc:
                raise RuntimeError(
                    f"exam renderer escaped its staging directory: {staged_docx}"
                ) from exc
            artifact_id = f"version-{index}-{_safe_artifact_stem(version.preset_id)}"
            version_artifact_ids.append(artifact_id)
            staged_artifacts.append(
                StagedArtifact(
                    artifact_id=artifact_id,
                    stage_path=staged_docx,
                    final_path=artifact_dir / relative_docx,
                )
            )

        published_paths = publish_staged_artifacts(
            staged_artifacts,
            execution_id=f"exam-{uuid4().hex}",
            work_root=output_root,
        )
        published_versions = tuple(
            replace(version, docx_path=published_paths[artifact_id])
            for version, artifact_id in zip(versions, version_artifact_ids)
        )

    return ExamDeliveryRuntimeResult(
        schema_id=EXAM_SCHEMA_ID,
        family_id=validation.family_id or EXAM_FAMILY_ID,
        status="warning" if validation.warning_count else "ok",
        source_key=validation.source_key or source_key,
        markdown_preview_path=published_paths["preview"],
        markdown_preview_excerpt=_preview_excerpt(markdown),
        rendered_versions=published_versions,
        master_evidence=master_evidence,
    )


def render_exam_markdown_preview(payload: object) -> str:
    """Build a human-reviewable Markdown preview from structured exam JSON."""

    normalized = _normalize_exam_payload(payload)
    lines: list[str] = [f"# {_exam_title(normalized)}", ""]
    for key, label in (
        ("subject", "Subject"),
        ("grade", "Grade"),
        ("duration", "Duration"),
        ("total_score", "Total score"),
    ):
        value = _text(normalized.get(key))
        if value:
            lines.append(f"- {label}: {value}")
    if len(lines) > 2:
        lines.append("")

    for section_index, section in enumerate(_exam_sections(normalized), start=1):
        title = _text(section.get("title")) or f"Section {section_index}"
        lines.extend((f"## {title}", ""))
        for question_number, question in enumerate(_section_questions(section), start=1):
            lines.append(f"{question_number}. {_question_stem(question)}")
            for option in _question_options(question):
                lines.append(f"   - {option}")
            answer = _question_answer(question)
            if answer:
                lines.append(f"   - Answer: {answer}")
            analysis = _question_analysis(question)
            if analysis:
                lines.append(f"   - Analysis: {analysis}")
            knowledge = _question_knowledge(question)
            if knowledge:
                lines.append(f"   - Knowledge: {knowledge}")
            score = _text(question.get("score"))
            if score:
                lines.append(f"   - Score: {score}")
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def parse_exam_markdown_source(
    markdown: str,
    *,
    source_path: str | Path = "",
) -> ExamMarkdownImportResult:
    """Parse the supported exam Markdown draft format into ``exam_items_v1``."""

    text = str(markdown or "")
    issues: list[ExamMarkdownImportIssue] = []
    payload: dict[str, object] = {}
    sections: list[dict[str, object]] = []
    current_section: dict[str, object] | None = None
    current_question: dict[str, object] | None = None
    in_answer_block = False
    answer_groups: list[dict[str, object]] = []
    current_answer_group: dict[str, object] | None = None
    current_answer_item: dict[str, object] | None = None

    def ensure_section(title: str = "") -> dict[str, object]:
        nonlocal current_section
        if current_section is None:
            current_section = {
                "title": title or "一、题目区",
                "questions": [],
            }
            sections.append(current_section)
        return current_section

    def ensure_answer_group(title: str = "") -> dict[str, object]:
        nonlocal current_answer_group
        if current_answer_group is None or title:
            current_answer_group = {
                "title": title,
                "answers": {},
            }
            answer_groups.append(current_answer_group)
        return current_answer_group

    def append_stem_line(line: str) -> None:
        if current_question is None:
            return
        existing = str(current_question.get("stem", "") or "").rstrip()
        current_question["stem"] = f"{existing}\n{line.strip()}".strip()

    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        stripped = raw_line.strip()
        if not stripped:
            continue

        section_match = _MARKDOWN_SECTION_RE.match(stripped)
        if section_match:
            section_title = section_match.group(1).strip()
            if _markdown_is_answer_heading(section_title):
                in_answer_block = True
                current_question = None
                current_answer_group = None
                current_answer_item = None
                continue
            in_answer_block = False
            current_section = {
                "title": section_title,
                "questions": [],
            }
            sections.append(current_section)
            current_question = None
            continue

        if in_answer_block:
            answer_match = _MARKDOWN_QUESTION_RE.match(stripped)
            if answer_match:
                group = ensure_answer_group()
                answer_number = int(answer_match.group(1))
                answer_payload = _parse_markdown_answer_text(answer_match.group(2))
                group_answers = group.setdefault("answers", {})
                if isinstance(group_answers, dict):
                    if answer_number in group_answers:
                        issues.append(
                            _markdown_issue(
                                f"line.{line_number}",
                                "duplicate_answer_number",
                                "同一答案分组内的答案题号不能重复。",
                                severity="error",
                                observed=str(answer_number),
                            )
                        )
                    group_answers[answer_number] = answer_payload
                current_answer_item = answer_payload
                continue

            if _markdown_looks_like_answer_group(stripped):
                current_answer_group = ensure_answer_group(stripped)
                current_answer_item = None
                continue

            if current_answer_item is not None:
                extra = stripped
                if "analysis" in current_answer_item:
                    current_answer_item["analysis"] = (
                        f"{current_answer_item['analysis']}\n{extra}".strip()
                    )
                else:
                    current_answer_item["answer"] = (
                        f"{current_answer_item.get('answer', '')}\n{extra}".strip()
                    )
                continue

            issues.append(
                _markdown_issue(
                    f"line.{line_number}",
                    "unrecognized_answer_line",
                    "答案速查区域存在未识别内容。",
                    observed=stripped,
                )
            )
            continue

        title_match = _MARKDOWN_TITLE_RE.match(stripped)
        if title_match and "paper_title" not in payload:
            payload["paper_title"] = title_match.group(1).strip()
            continue

        if stripped.startswith(">"):
            _merge_markdown_metadata(payload, stripped.lstrip("> "))
            continue

        question_match = _MARKDOWN_QUESTION_RE.match(stripped)
        if question_match:
            section = ensure_section()
            question_number = int(question_match.group(1))
            existing_questions = section.get("questions")
            expected_number = (
                len(existing_questions) + 1
                if isinstance(existing_questions, list)
                else 1
            )
            if question_number != expected_number:
                issues.append(
                    _markdown_issue(
                        f"line.{line_number}",
                        "question_number_sequence_invalid",
                        "每个大题内的小题必须从 1 开始连续编号。",
                        severity="error",
                        observed=str(question_number),
                    )
                )
            question: dict[str, object] = {
                "number": question_number,
                "stem": question_match.group(2).strip(),
            }
            questions = section.setdefault("questions", [])
            if isinstance(questions, list):
                questions.append(question)
            current_question = question
            continue

        option_match = _MARKDOWN_OPTION_RE.match(stripped)
        if option_match and current_question is not None:
            options = current_question.setdefault("options", [])
            if isinstance(options, list):
                option_label = option_match.group(1).upper()
                option_text = option_match.group(2).strip()
                options.append(f"{option_label}. {option_text}")
            continue

        field_match = _MARKDOWN_FIELD_RE.match(stripped)
        if field_match and current_question is not None:
            field_key = field_match.group(1).strip()
            field_value = field_match.group(2).strip()
            if field_key == "answer_area_kind":
                normalized_kind = field_value.lower()
                if normalized_kind in _MARKDOWN_ALLOWED_ANSWER_AREA_KINDS:
                    current_question["answer_area_kind"] = normalized_kind
                else:
                    issues.append(
                        _markdown_issue(
                            f"line.{line_number}",
                            "unknown_answer_area_kind",
                            "答题区类型未识别。",
                            observed=field_value,
                        )
                    )
                continue
            if field_key == "answer_lines":
                try:
                    current_question["answer_lines"] = max(0, int(float(field_value)))
                except ValueError:
                    issues.append(
                        _markdown_issue(
                            f"line.{line_number}",
                            "invalid_answer_lines",
                            "答题区行数必须是数字。",
                            observed=field_value,
                        )
                    )
                continue

        if current_question is not None:
            append_stem_line(stripped)
        else:
            issues.append(
                _markdown_issue(
                    f"line.{line_number}",
                    "unrecognized_line",
                    "题稿中存在未挂接到试卷结构的内容。",
                    observed=stripped,
                )
            )

    payload["sections"] = sections
    _finalize_markdown_questions(payload)
    _apply_markdown_answers(payload, answer_groups, issues)

    validation = _validate_exam_payload(
        _normalize_exam_payload(payload),
        source_key="markdown",
        family_id=EXAM_FAMILY_ID,
        metadata=payload,
    )
    issues.extend(_markdown_issues_from_validation(validation))

    status = "ok"
    if any(issue.severity == "error" for issue in issues):
        status = "error"
    elif issues:
        status = "warning"

    return ExamMarkdownImportResult(
        payload=payload,
        summary=validation.summary,
        issues=tuple(issues),
        source_path=str(source_path or ""),
        source_excerpt=_preview_excerpt(text),
        status=status,
    )


def parse_exam_markdown_file(path: str | Path) -> ExamMarkdownImportResult:
    """Read and parse one Markdown exam source file."""

    source = Path(path)
    try:
        markdown = source.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        markdown = source.read_text(encoding="utf-8-sig")
    return parse_exam_markdown_source(markdown, source_path=str(source))


def exam_markdown_import_entity_data(
    result: ExamMarkdownImportResult,
) -> dict[str, str]:
    """Return ``entity_data`` values produced by a Markdown import."""

    payload = dict(result.payload or {})
    data = {
        "exam_items": json.dumps(payload, ensure_ascii=False),
    }
    for source_key, target_key in (
        ("paper_title", "title"),
        ("subject", "subject"),
        ("grade", "grade"),
        ("duration", "duration"),
        ("total_score", "total_score"),
    ):
        value = str(payload.get(source_key, "") or "").strip()
        if value:
            data[target_key] = value
    return data


def _merge_markdown_metadata(payload: dict[str, object], text: str) -> None:
    for match in _MARKDOWN_METADATA_PAIR_RE.finditer(str(text or "")):
        key = match.group(1).strip()
        value = match.group(2).strip()
        if not value:
            continue
        target = {
            "科目": "subject",
            "学科": "subject",
            "年级": "grade",
            "考试时间": "duration",
            "时间": "duration",
            "满分": "total_score",
            "总分": "total_score",
        }.get(key)
        if target:
            payload[target] = value


def _markdown_is_answer_heading(title: str) -> bool:
    normalized = str(title or "").strip()
    return bool(normalized) and any(
        token in normalized
        for token in ("答案速查", "参考答案", "答案与解析", "答案")
    )


def _markdown_looks_like_answer_group(line: str) -> bool:
    text = str(line or "").strip()
    if not text:
        return False
    if _MARKDOWN_QUESTION_RE.match(text):
        return False
    if _MARKDOWN_OPTION_RE.match(text):
        return False
    return bool(_MARKDOWN_ANSWER_GROUP_RE.match(text)) or text.endswith("题")


def _parse_markdown_answer_text(text: str) -> dict[str, object]:
    raw = str(text or "").strip()
    match = _MARKDOWN_ANSWER_SPLIT_RE.match(raw)
    if match and match.group(2):
        answer = match.group(1).strip(" 。；;")
        analysis = match.group(2).strip()
        return {
            "answer": answer,
            "analysis": analysis,
        }
    return {"answer": raw}


def _finalize_markdown_questions(payload: dict[str, object]) -> None:
    for section in list(payload.get("sections") or []):
        if not isinstance(section, Mapping):
            continue
        for question in list(section.get("questions") or []):
            if not isinstance(question, dict):
                continue
            stem = str(question.get("stem", "") or "").strip()
            matches = list(_MARKDOWN_SCORE_RE.finditer(stem))
            if matches and "score" not in question:
                question["score"] = matches[-1].group(1)
            if matches:
                stem = _MARKDOWN_SCORE_RE.sub("", stem).strip()
            question["stem"] = _normalize_markdown_inline_text(stem)
            if "options" in question and isinstance(question["options"], list):
                question["options"] = [
                    _normalize_markdown_inline_text(str(option or ""))
                    for option in question["options"]
                    if str(option or "").strip()
                ]


def _normalize_markdown_inline_text(text: str) -> str:
    normalized = str(text or "").strip()
    normalized = re.sub(r"`([^`]+)`", r"\1", normalized)
    normalized = re.sub(r"[ \t]+\n", "\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _apply_markdown_answers(
    payload: dict[str, object],
    answer_groups: list[dict[str, object]],
    issues: list[ExamMarkdownImportIssue],
) -> None:
    sections = [
        section
        for section in list(payload.get("sections") or [])
        if isinstance(section, dict)
    ]
    if not answer_groups:
        issues.append(
            _markdown_issue(
                "answers",
                "missing_answer_block",
                "Markdown 题稿缺少“答案速查”区域。",
                severity="error",
            )
        )
        return

    grouped = [
        group
        for group in answer_groups
        if str(group.get("title", "") or "").strip()
    ]
    if grouped:
        for group_index, group in enumerate(grouped):
            if group_index >= len(sections):
                issues.append(
                    _markdown_issue(
                        f"answers.{group_index}",
                        "unmatched_answer_group",
                        "答案分组多于题目大题，未能全部回填。",
                        severity="error",
                        observed=str(group.get("title", "") or ""),
                    )
                )
                continue
            group_family = _markdown_question_family(
                str(group.get("title", "") or "")
            )
            section_family = _markdown_question_family(
                str(sections[group_index].get("title", "") or "")
            )
            if (
                group_family
                and section_family
                and group_family != section_family
            ):
                issues.append(
                    _markdown_issue(
                        f"answers.{group_index}",
                        "answer_group_section_mismatch",
                        "答案分组顺序或题型与对应大题不一致。",
                        severity="error",
                        observed=str(group.get("title", "") or ""),
                    )
                )
            _apply_answer_group_to_section(
                group,
                sections[group_index],
                issues,
                group_path=f"answers.{group_index}",
            )
        return

    flat_questions: list[dict[str, object]] = []
    for section in sections:
        flat_questions.extend(
            question
            for question in list(section.get("questions") or [])
            if isinstance(question, dict)
        )
    answers: dict[int, object] = {}
    for group in answer_groups:
        group_answers = group.get("answers", {})
        if isinstance(group_answers, Mapping):
            answers.update(group_answers)
    for answer_number, answer_payload in answers.items():
        target_index = int(answer_number) - 1
        if target_index < 0 or target_index >= len(flat_questions):
            issues.append(
                _markdown_issue(
                    f"answers.{answer_number}",
                    "unmatched_answer",
                    "答案题号无法匹配到题目。",
                    severity="error",
                    observed=str(answer_number),
                )
            )
            continue
        _merge_answer_payload(flat_questions[target_index], answer_payload)


def _apply_answer_group_to_section(
    group: Mapping[str, object],
    section: dict[str, object],
    issues: list[ExamMarkdownImportIssue],
    *,
    group_path: str,
) -> None:
    questions = [
        question
        for question in list(section.get("questions") or [])
        if isinstance(question, dict)
    ]
    answers = group.get("answers", {})
    if not isinstance(answers, Mapping):
        return
    for answer_number, answer_payload in answers.items():
        target_index = int(answer_number) - 1
        if target_index < 0 or target_index >= len(questions):
            issues.append(
                _markdown_issue(
                    f"{group_path}.{answer_number}",
                    "unmatched_answer",
                    "答案题号无法匹配到当前大题的小题。",
                    severity="error",
                    observed=str(answer_number),
                )
            )
            continue
        _merge_answer_payload(questions[target_index], answer_payload)


def _merge_answer_payload(
    question: dict[str, object],
    answer_payload: object,
) -> None:
    if not isinstance(answer_payload, Mapping):
        question["answer"] = str(answer_payload or "").strip()
        return
    answer = str(answer_payload.get("answer", "") or "").strip()
    analysis = str(answer_payload.get("analysis", "") or "").strip()
    if answer:
        question["answer"] = _normalize_markdown_inline_text(answer)
    if analysis:
        question["analysis"] = _normalize_markdown_inline_text(analysis)


def _markdown_question_family(value: str) -> str:
    normalized = str(value or "").casefold()
    groups = (
        ("choice", ("选择题", "单选题", "多选题", "选择")),
        ("judgment", ("判断题", "判断")),
        ("fill", ("填空题", "填空")),
        (
            "calculation",
            ("计算题", "解答题", "应用题", "证明题", "计算", "解答", "证明"),
        ),
        ("reading", ("阅读题", "阅读理解", "阅读")),
        ("writing", ("作文题", "写作题", "作文", "写作")),
        ("short_answer", ("简答题", "问答题", "简答", "问答")),
    )
    for family, tokens in groups:
        if any(token in normalized for token in tokens):
            return family
    return ""


def _markdown_issues_from_validation(
    validation: ExamQuestionSchemaValidationResult,
) -> list[ExamMarkdownImportIssue]:
    issues: list[ExamMarkdownImportIssue] = []
    for item in validation.issues:
        issues.append(
            _markdown_issue(
                item.path,
                item.kind,
                item.message,
                severity=item.severity,
                observed=item.observed,
            )
        )
    return issues


def _markdown_issue(
    path: str,
    kind: str,
    message: str,
    *,
    severity: str = "warning",
    observed: str = "",
) -> ExamMarkdownImportIssue:
    return ExamMarkdownImportIssue(
        path=str(path or ""),
        kind=str(kind or ""),
        message=str(message or ""),
        severity=str(severity or "warning"),
        observed=str(observed or ""),
    )


def _config_material_schema_ids(config) -> tuple[str, ...]:
    profile = getattr(config, "input_source_profile", None)
    return resolve_material_schema_ids(
        str(getattr(profile, "material_schema_id", "") or "").strip(),
        list(getattr(profile, "material_schema_ids", []) or []),
    )


def _exam_family_for_schema_ids(schema_ids: tuple[str, ...]) -> str:
    for schema_id in schema_ids:
        try:
            family = get_material_schema(schema_id).family
        except KeyError:
            continue
        if family == EXAM_FAMILY_ID:
            return family
    return ""


def _extract_exam_payload(entity_data: object) -> tuple[object | None, str]:
    if not isinstance(entity_data, Mapping):
        return None, ""

    for key in _SOURCE_KEYS:
        if key not in entity_data:
            continue
        parsed = _parse_structured_value(entity_data.get(key))
        if parsed is None:
            continue
        if key in {"questions", "questions_json"} and _is_sequence(parsed):
            return {"sections": [{"title": "Questions", "type": "mixed", "questions": parsed}]}, f"entity_data.{key}"
        if key == "sections" and _is_sequence(parsed):
            payload = dict(entity_data)
            payload["sections"] = parsed
            return payload, f"entity_data.{key}"
        return parsed, f"entity_data.{key}"

    if any(str(key or "").strip() in {*_TITLE_KEYS, "subject", "grade", "duration", "total_score"} for key in entity_data):
        return dict(entity_data), "entity_data"
    return None, ""


def _parse_structured_value(value: object) -> object | None:
    if isinstance(value, Mapping):
        return dict(value)
    if _is_sequence(value):
        return list(value)
    text = str(value or "").strip()
    if not text:
        return None
    if text[:1] not in {"{", "["}:
        return None
    for parser in (json.loads, ast.literal_eval):
        try:
            parsed = parser(text)
        except (ValueError, SyntaxError, TypeError, json.JSONDecodeError):
            continue
        if isinstance(parsed, Mapping):
            return dict(parsed)
        if _is_sequence(parsed):
            return list(parsed)
    return None


def _exam_payload_from_config(config) -> tuple[dict[str, object] | None, str]:
    payload, source_key = _extract_exam_payload(getattr(config, "entity_data", {}) or {})
    if payload is None:
        return None, source_key
    normalized = _normalize_exam_payload(payload)
    return _merge_runtime_fields_into_payload(normalized, config), source_key


def _merge_runtime_fields_into_payload(
    payload: dict[str, object],
    config,
) -> dict[str, object]:
    entity_data = getattr(config, "entity_data", {}) or {}
    if not isinstance(entity_data, Mapping):
        return payload

    exam_paper = getattr(config, "exam_paper", None)
    runtime_fields = list(getattr(exam_paper, "runtime_fields", []) or [])
    if not runtime_fields:
        runtime_fields = ["title", "subject", "grade", "duration", "total_score"]

    merged = dict(payload)
    for field_id in runtime_fields:
        key = str(field_id or "").strip()
        if not key:
            continue
        value = _runtime_field_value(entity_data, key)
        if value is None:
            continue
        if key == "title":
            merged["paper_title"] = value
        else:
            merged[key] = value
    return merged


def _runtime_field_value(entity_data: Mapping[str, object], field_id: str) -> object | None:
    aliases = {
        "title": ("paper_title", "title", "exam_title", "试卷标题"),
        "subject": ("subject", "科目"),
        "grade": ("grade", "年级"),
        "duration": ("duration", "time", "exam_duration", "考试时间"),
        "total_score": ("total_score", "score", "full_score", "满分"),
        "class_name": ("class_name", "class", "班级"),
        "teacher": ("teacher", "命题人", "任课教师"),
        "exam_date": ("exam_date", "date", "日期"),
    }.get(field_id, (field_id,))
    for key in aliases:
        if key not in entity_data:
            continue
        value = entity_data.get(key)
        if value is None:
            continue
        text = str(value).strip()
        if text:
            return value
    return None


def _normalize_exam_payload(payload: object) -> dict[str, object]:
    if _is_sequence(payload):
        return {"sections": [{"title": "Questions", "type": "mixed", "questions": list(payload)}]}
    if not isinstance(payload, Mapping):
        return {}
    normalized = dict(payload)
    paper = normalized.get("paper")
    if isinstance(paper, Mapping):
        merged = dict(paper)
        for key, value in normalized.items():
            if key not in merged:
                merged[key] = value
        normalized = merged

    sections = _first_present(normalized, _SECTION_KEYS)
    parsed_sections = _parse_structured_value(sections)
    if _is_sequence(parsed_sections):
        normalized["sections"] = list(parsed_sections)
    elif sections is not None:
        normalized["sections"] = sections

    questions = _first_present(normalized, _QUESTION_KEYS)
    parsed_questions = _parse_structured_value(questions)
    if "sections" not in normalized and _is_sequence(parsed_questions):
        normalized["sections"] = [
            {"title": "Questions", "type": "mixed", "questions": list(parsed_questions)}
        ]
    return normalized


def _validate_exam_payload(
    payload: dict[str, object],
    *,
    source_key: str,
    family_id: str,
    metadata: object,
) -> ExamQuestionSchemaValidationResult:
    issues: list[ExamQuestionSchemaIssue] = []
    question_types: list[str] = []
    score_values: list[float] = []
    section_count = 0
    question_count = 0
    answered_count = 0
    analysis_count = 0
    knowledge_count = 0
    figure_count = 0
    scored_count = 0

    if not _first_text(payload, _TITLE_KEYS) and not _metadata_text(metadata, _TITLE_KEYS):
        issues.append(
            _issue(
                "paper_title",
                "missing_paper_title",
                "Paper title is required for an exam paper source.",
                severity="warning",
            )
        )

    declared_total = _number(_first_present(payload, ("total_score", "totalScore")))
    if declared_total is None:
        declared_total = _number(_metadata_value(metadata, "total_score"))

    sections = payload.get("sections")
    if not _is_sequence(sections):
        issues.append(
            _issue(
                "sections",
                "missing_sections",
                "Exam source must include a non-empty sections list.",
                severity="error",
                expected="list[section]",
                observed=type(sections).__name__,
            )
        )
        sections = []

    if _is_sequence(sections) and not sections:
        issues.append(
            _issue(
                "sections",
                "empty_sections",
                "Exam source must include at least one section.",
                severity="error",
            )
        )

    for section_index, section in enumerate(list(sections or [])):
        section_path = f"sections.{section_index}"
        if not isinstance(section, Mapping):
            issues.append(
                _issue(
                    section_path,
                    "invalid_section",
                    "Each section must be an object with questions.",
                    severity="error",
                    expected="object",
                    observed=type(section).__name__,
                )
            )
            continue
        section_count += 1
        section_type = _text(section.get("type") or section.get("question_type"))
        if section_type:
            question_types.append(section_type)
        if not _text(section.get("title")):
            issues.append(
                _issue(
                    f"{section_path}.title",
                    "missing_section_title",
                    "Section title is recommended for stable Word rendering.",
                )
            )

        questions = _first_present(section, _QUESTION_KEYS)
        parsed_questions = _parse_structured_value(questions)
        if _is_sequence(parsed_questions):
            questions = parsed_questions
        if not _is_sequence(questions):
            issues.append(
                _issue(
                    f"{section_path}.questions",
                    "missing_questions",
                    "Each section must include a questions list.",
                    severity="error",
                    expected="list[question]",
                    observed=type(questions).__name__,
                )
            )
            continue
        if not questions:
            issues.append(
                _issue(
                    f"{section_path}.questions",
                    "empty_questions",
                    "Section has no questions.",
                    severity="warning",
                )
            )
        for question_index, question in enumerate(list(questions or [])):
            question_path = f"{section_path}.questions.{question_index}"
            result = _validate_question(
                question,
                path=question_path,
                section_type=section_type,
            )
            issues.extend(result["issues"])
            question_count += int(result["question_count"])
            answered_count += int(result["answered_count"])
            analysis_count += int(result["analysis_count"])
            knowledge_count += int(result["knowledge_count"])
            figure_count += int(result["figure_count"])
            scored_count += int(result["scored_count"])
            score = result["score"]
            if isinstance(score, float):
                score_values.append(score)
            q_type = str(result["question_type"] or "")
            if q_type:
                question_types.append(q_type)

    computed_total = sum(score_values) if score_values else None
    if declared_total is not None and computed_total is not None and scored_count == question_count:
        if abs(declared_total - computed_total) > 0.01:
            issues.append(
                _issue(
                    "total_score",
                    "total_score_mismatch",
                    "Declared total score does not equal the sum of question scores.",
                    severity="error",
                    expected=_format_number(declared_total),
                    observed=_format_number(computed_total),
                )
            )

    status = "ok"
    if any(issue.severity == "error" for issue in issues):
        status = "error"
    elif issues:
        status = "warning"

    return ExamQuestionSchemaValidationResult(
        schema_id=EXAM_SCHEMA_ID,
        family_id=family_id,
        status=status,
        source_key=source_key,
        summary=ExamQuestionSchemaSummary(
            section_count=section_count,
            question_count=question_count,
            answered_question_count=answered_count,
            analysis_count=analysis_count,
            knowledge_point_count=knowledge_count,
            figure_count=figure_count,
            scored_question_count=scored_count,
            declared_total_score=declared_total,
            computed_total_score=computed_total,
            question_types=tuple(_unique_texts(question_types)),
        ),
        issues=tuple(issues),
    )


def _validate_question(
    question: object,
    *,
    path: str,
    section_type: str,
) -> dict[str, object]:
    issues: list[ExamQuestionSchemaIssue] = []
    if not isinstance(question, Mapping):
        return {
            "issues": [
                _issue(
                    path,
                    "invalid_question",
                    "Each question must be an object.",
                    severity="error",
                    expected="object",
                    observed=type(question).__name__,
                )
            ],
            "question_count": 0,
            "answered_count": 0,
            "analysis_count": 0,
            "knowledge_count": 0,
            "figure_count": 0,
            "scored_count": 0,
            "score": None,
            "question_type": "",
        }

    q_type = _text(question.get("type") or question.get("question_type") or section_type)
    if not _text(question.get("stem") or question.get("question") or question.get("prompt")):
        issues.append(
            _issue(
                f"{path}.stem",
                "missing_stem",
                "Question stem is required.",
                severity="error",
            )
        )

    answer = _first_present(question, _ANSWER_KEYS)
    answered = _has_value(answer)
    if not answered:
        issues.append(
            _issue(
                f"{path}.answer",
                "missing_answer",
                "Answer must remain in the structured source; delivery presets decide visibility.",
                severity="error",
            )
        )

    score = _number(question.get("score"))
    scored = score is not None
    if not scored:
        issues.append(
            _issue(
                f"{path}.score",
                "missing_or_invalid_score",
                "Question score is recommended for total-score checks and answer sheets.",
                expected="number",
                observed=_text(question.get("score")),
            )
        )

    analysis = _first_present(question, _ANALYSIS_KEYS)
    knowledge = _first_present(question, _KNOWLEDGE_KEYS)
    figure = question.get("figure") or question.get("image")
    figure_count, figure_issues = _inspect_question_figure(figure, path=path)
    issues.extend(figure_issues)

    options = question.get("options")
    if _looks_like_choice(q_type) and not _has_value(options):
        issues.append(
            _issue(
                f"{path}.options",
                "missing_choice_options",
                "Choice questions should include options in the structured source.",
            )
        )
    if options is not None and not (isinstance(options, Mapping) or _is_sequence(options) or isinstance(options, str)):
        issues.append(
            _issue(
                f"{path}.options",
                "invalid_options",
                "Question options should be a list, mapping, or text block.",
                expected="list|object|string",
                observed=type(options).__name__,
            )
        )

    return {
        "issues": issues,
        "question_count": 1,
        "answered_count": 1 if answered else 0,
        "analysis_count": 1 if _has_value(analysis) else 0,
        "knowledge_count": _knowledge_count(knowledge),
        "figure_count": figure_count,
        "scored_count": 1 if scored else 0,
        "score": score,
        "question_type": q_type,
    }


def _inspect_question_figure(
    figure: object,
    *,
    path: str,
) -> tuple[int, list[ExamQuestionSchemaIssue]]:
    if figure is None or figure == "":
        return 0, []
    figures = list(figure) if _is_sequence(figure) else [figure]
    issues: list[ExamQuestionSchemaIssue] = []
    for index, item in enumerate(figures):
        item_path = f"{path}.figure" if len(figures) == 1 else f"{path}.figure.{index}"
        if not isinstance(item, Mapping):
            issues.append(
                _issue(
                    item_path,
                    "invalid_figure",
                    "Question figure should be an object with source/path metadata.",
                    expected="object",
                    observed=type(item).__name__,
                )
            )
            continue
        source = _text(item.get("source"))
        path_value = _text(item.get("path") or item.get("asset_id"))
        if source in {"upload", "asset", "file"} and not path_value:
            issues.append(
                _issue(
                    f"{item_path}.path",
                    "missing_figure_path",
                    "Uploaded question figures need a path or asset id.",
                )
            )
    return len(figures), issues


def _delivery_presets_for_runtime(config) -> tuple[object, ...]:
    presets = tuple(getattr(config, "delivery_presets", []) or ())
    if presets:
        return presets
    return (
        {
            "preset_id": "final",
            "label": "Final DOCX",
            "content_visibility_rules": (),
            "artifacts": {"final_docx": True},
        },
    )


def _exam_master_evidence(master) -> ExamMasterEvidence | None:
    if master is None:
        return None
    preflight = check_master_preflight(master)
    return ExamMasterEvidence(
        mode_id=master.mode_id,
        master_id=master.master_id,
        master_label=master.label,
        master_source_type=master.source_type,
        master_docx_path=str(master.docx_path),
        master_version=master.master_version,
        manifest_path=str(master.manifest_path or ""),
        placeholder_contract_status=preflight.status,
        required_placeholders=master.placeholder_contract.required,
        optional_placeholders=master.placeholder_contract.optional,
        runtime_fields=master.placeholder_contract.runtime_fields,
        generated_fields=master.placeholder_contract.generated_fields,
        missing_required_placeholders=preflight.missing_required_placeholders,
    )


def _render_exam_master_docx(
    payload: dict[str, object],
    exam_paper_config,
    artifact_dir: Path,
    *,
    source_stem: str,
    summary: ExamQuestionSchemaSummary | None = None,
    master,
) -> tuple[ExamRenderedVersion, ...]:
    answer_policy = str(getattr(exam_paper_config, "answer_policy", "") or "").strip()
    include_student_version = answer_policy != "answer_only"
    include_answer_version = answer_policy != "student_only"
    outputs = write_exam_paper_docx_files(
        str(getattr(master, "master_id", "") or ""),
        artifact_dir,
        payload=payload,
        config=exam_paper_config,
        include_student=include_student_version,
        include_answer_key=include_answer_version,
        filename_stem=source_stem,
        master_docx_path=getattr(master, "docx_path", ""),
    )

    question_count = int(getattr(summary, "question_count", 0) or 0)
    answer_count = int(getattr(summary, "answered_question_count", 0) or 0)
    figure_count, figure_alt_text_count = _exam_payload_question_asset_counts(payload)
    if not figure_count:
        figure_count = int(getattr(summary, "figure_count", 0) or 0)
    versions = []
    if include_student_version and outputs.student_docx is not None:
        rendered_count, rendered_alt_text_count = _docx_inline_shape_asset_counts(
            outputs.student_docx
        )
        versions.append(
            ExamRenderedVersion(
                preset_id="student",
                label="学生卷",
                docx_path=str(outputs.student_docx),
                hidden_selectors=("answer", "analysis", "solution", "knowledge_points"),
                visible_question_count=question_count,
                visible_answer_count=0,
                visible_analysis_count=0,
                question_asset_count=figure_count,
                rendered_question_asset_count=rendered_count,
                missing_question_asset_count=max(0, figure_count - rendered_count),
                question_asset_alt_text_count=figure_alt_text_count,
                rendered_question_asset_alt_text_count=min(
                    rendered_alt_text_count,
                    figure_alt_text_count,
                ),
                missing_question_asset_alt_text_count=max(
                    0,
                    figure_alt_text_count - rendered_alt_text_count,
                ),
            )
        )
    if include_answer_version and outputs.answer_key_docx is not None:
        analysis_count = int(getattr(summary, "analysis_count", 0) or 0)
        versions.append(
            ExamRenderedVersion(
                preset_id="answer_key",
                label="答案速查",
                docx_path=str(outputs.answer_key_docx),
                hidden_selectors=("question_body", "knowledge_points"),
                visible_question_count=0,
                visible_answer_count=answer_count,
                visible_analysis_count=analysis_count,
                question_asset_count=0,
                rendered_question_asset_count=0,
                missing_question_asset_count=0,
            )
        )
    return tuple(versions)


def _exam_payload_question_asset_counts(payload: Mapping[str, object]) -> tuple[int, int]:
    asset_count = 0
    alt_text_count = 0
    for _section_index, _section_title, question in _flatten_exam_questions(payload):
        figures = _question_figures(question)
        asset_count += len(figures)
        alt_text_count += sum(1 for figure in figures if _question_figure_alt_text(figure))
    return asset_count, alt_text_count


def _docx_inline_shape_asset_counts(docx_path: Path | str) -> tuple[int, int]:
    try:
        document = Document(str(docx_path))
    except Exception:
        return 0, 0
    inline_shapes = list(getattr(document, "inline_shapes", []) or [])
    alt_text_count = 0
    for inline_shape in inline_shapes:
        doc_pr = getattr(getattr(inline_shape, "_inline", None), "docPr", None)
        if doc_pr is not None and _text(doc_pr.get("descr")):
            alt_text_count += 1
    return len(inline_shapes), alt_text_count


def _render_exam_version_docx(
    payload: dict[str, object],
    preset: object,
    docx_path: Path,
    *,
    label: str,
) -> ExamRenderedVersion:
    hidden_selectors = _preset_hidden_selectors(preset)
    preset_id = _preset_value(preset, "preset_id", "final").strip() or "final"
    show_question = not {"question_body", "question_only"} & set(hidden_selectors)
    show_answer = "answer" not in hidden_selectors
    show_analysis = not {"analysis", "solution"} & set(hidden_selectors)
    show_knowledge = "knowledge_points" not in hidden_selectors
    answer_sheet = _is_answer_sheet_preset(preset, hidden_selectors)

    document = Document()
    document.add_heading(_exam_title(payload), level=0)
    if label:
        document.add_paragraph(f"Version: {label}")
    if answer_sheet:
        return _render_exam_answer_sheet_docx(
            document,
            payload,
            preset_id=preset_id,
            label=label,
            docx_path=docx_path,
            hidden_selectors=hidden_selectors,
        )

    visible_questions = 0
    visible_answers = 0
    visible_analysis = 0
    question_assets = 0
    rendered_question_assets = 0
    missing_question_assets = 0
    question_asset_alt_texts = 0
    rendered_question_asset_alt_texts = 0
    missing_question_asset_alt_texts = 0
    global_question_number = 0

    for section_index, section in enumerate(_exam_sections(payload), start=1):
        section_title = _text(section.get("title")) or f"Section {section_index}"
        document.add_heading(section_title, level=1)
        for question in _section_questions(section):
            global_question_number += 1
            if show_question:
                visible_questions += 1
                document.add_paragraph(
                    f"{global_question_number}. {_question_stem(question)}"
                )
                for option in _question_options(question):
                    document.add_paragraph(option, style=None)
                asset_state = _render_question_figures(document, question)
                question_assets += asset_state[0]
                rendered_question_assets += asset_state[1]
                missing_question_assets += asset_state[2]
                question_asset_alt_texts += asset_state[3]
                rendered_question_asset_alt_texts += asset_state[4]
                missing_question_asset_alt_texts += asset_state[5]
            answer = _question_answer(question)
            if show_answer and answer:
                visible_answers += 1
                document.add_paragraph(f"Answer: {answer}")

            analysis = _question_analysis(question)
            if show_analysis and analysis:
                visible_analysis += 1
                document.add_paragraph(f"Analysis: {analysis}")

            knowledge = _question_knowledge(question)
            if show_knowledge and knowledge:
                document.add_paragraph(f"Knowledge: {knowledge}")

    document.save(str(docx_path))
    return ExamRenderedVersion(
        preset_id=preset_id,
        label=label,
        docx_path=str(docx_path),
        hidden_selectors=hidden_selectors,
        visible_question_count=visible_questions,
        visible_answer_count=visible_answers,
        visible_analysis_count=visible_analysis,
        question_asset_count=question_assets,
        rendered_question_asset_count=rendered_question_assets,
        missing_question_asset_count=missing_question_assets,
        question_asset_alt_text_count=question_asset_alt_texts,
        rendered_question_asset_alt_text_count=rendered_question_asset_alt_texts,
        missing_question_asset_alt_text_count=missing_question_asset_alt_texts,
    )


def _render_exam_answer_sheet_docx(
    document: Document,
    payload: dict[str, object],
    *,
    preset_id: str,
    label: str,
    docx_path: Path,
    hidden_selectors: tuple[str, ...],
) -> ExamRenderedVersion:
    questions = _flatten_exam_questions(payload)
    document.add_paragraph("Answer sheet")
    table = document.add_table(rows=max(1, len(questions) + 1), cols=4)
    table.style = "Table Grid"
    for column_index, header in enumerate(("No.", "Question type", "Answer area", "Score")):
        table.cell(0, column_index).text = header

    for row_index, (question_number, section_title, question) in enumerate(
        questions,
        start=1,
    ):
        table.cell(row_index, 0).text = str(question_number)
        table.cell(row_index, 1).text = _question_type_label(question, section_title)
        table.cell(row_index, 2).text = "____________________________"
        table.cell(row_index, 3).text = _text(question.get("score"))

    policy = FixedLayoutRowHeightPolicy(
        policy_id="exam_teaching.answer_sheet_row_height",
        family_id=EXAM_FAMILY_ID,
        label="答题卡固定行高",
        mode="enforce_exact",
        row_height_pt=22.0,
        parameter_path="exam_teaching.answer_sheet.row_height_pt",
        applies_to=("answer_sheet_table",),
        rationale=(
            "Answer-sheet tables are fixed-layout delivery artifacts and write "
            "Word w:trHeight outside generic TableConfig."
        ),
    )
    apply_fixed_layout_row_height_policy(table, policy)
    first_data_row = table.rows[1] if len(table.rows) > 1 else table.rows[0]
    height_state = row_height_state(first_data_row)
    document.save(str(docx_path))
    return ExamRenderedVersion(
        preset_id=preset_id,
        label=label,
        docx_path=str(docx_path),
        hidden_selectors=hidden_selectors,
        visible_question_count=0,
        visible_answer_count=0,
        visible_analysis_count=0,
        fixed_layout_kind="answer_sheet",
        fixed_layout_row_count=len(table.rows),
        fixed_layout_column_count=4,
        fixed_layout_row_height_twips=height_state.height_twips,
        fixed_layout_row_height_rule=height_state.rule,
    )


def _is_answer_sheet_preset(
    preset: object,
    hidden_selectors: tuple[str, ...],
) -> bool:
    preset_id = _preset_value(preset, "preset_id", "").strip().lower()
    if preset_id == "answer_sheet":
        return True
    selectors = set(hidden_selectors)
    return "question_body" in selectors and "answer" in selectors


def _preset_hidden_selectors(preset: object) -> tuple[str, ...]:
    selectors: list[str] = []
    for rule in list(_preset_value_raw(preset, "content_visibility_rules", ()) or ()):
        selector = executable_content_visibility_rule_selector(rule)
        if selector:
            selectors.append(selector)
    return tuple(dict.fromkeys(selectors))


def _preset_final_docx_enabled(preset: object) -> bool:
    artifacts = _preset_value_raw(preset, "artifacts", None)
    if isinstance(artifacts, Mapping):
        return bool(artifacts.get("final_docx", True))
    if artifacts is None:
        return True
    return bool(getattr(artifacts, "final_docx", True))


def _preset_value(preset: object, key: str, default: str = "") -> str:
    return _text(_preset_value_raw(preset, key, default))


def _preset_value_raw(preset: object, key: str, default: object = "") -> object:
    if isinstance(preset, Mapping):
        return preset.get(key, default)
    return getattr(preset, key, default)


def _exam_title(payload: Mapping[str, object]) -> str:
    return _first_text(payload, _TITLE_KEYS) or "Exam Paper"


def _exam_sections(payload: Mapping[str, object]) -> list[Mapping[str, object]]:
    sections = payload.get("sections")
    if not _is_sequence(sections):
        return []
    return [section for section in list(sections) if isinstance(section, Mapping)]


def _flatten_exam_questions(
    payload: Mapping[str, object],
) -> list[tuple[int, str, Mapping[str, object]]]:
    rows: list[tuple[int, str, Mapping[str, object]]] = []
    question_number = 0
    for section_index, section in enumerate(_exam_sections(payload), start=1):
        section_title = _text(section.get("title")) or f"Section {section_index}"
        for question in _section_questions(section):
            question_number += 1
            rows.append((question_number, section_title, question))
    return rows


def _section_questions(section: Mapping[str, object]) -> list[Mapping[str, object]]:
    questions = _first_present(section, _QUESTION_KEYS)
    parsed = _parse_structured_value(questions)
    if _is_sequence(parsed):
        questions = parsed
    if not _is_sequence(questions):
        return []
    return [question for question in list(questions) if isinstance(question, Mapping)]


def _question_type_label(
    question: Mapping[str, object],
    section_title: str,
) -> str:
    return (
        _text(question.get("type"))
        or _text(question.get("question_type"))
        or section_title
    )


def _render_question_figures(
    document: Document,
    question: Mapping[str, object],
) -> tuple[int, int, int, int, int, int]:
    figures = _question_figures(question)
    rendered = 0
    missing = 0
    alt_text_count = 0
    rendered_alt_text_count = 0
    missing_alt_text_count = 0
    for figure in figures:
        alt_text = _question_figure_alt_text(figure)
        if alt_text:
            alt_text_count += 1
        figure_path = _resolve_question_figure_path(figure)
        if not figure_path:
            missing += 1
            continue
        try:
            inline_shape = document.add_picture(str(figure_path), width=Inches(2.0))
            if alt_text:
                set_inline_shape_alt_text(
                    inline_shape,
                    alt_text=alt_text,
                    title=_question_figure_title(figure),
                )
                rendered_alt_text_count += 1
            else:
                missing_alt_text_count += 1
            rendered += 1
        except Exception:
            missing += 1
    return (
        len(figures),
        rendered,
        missing,
        alt_text_count,
        rendered_alt_text_count,
        missing_alt_text_count,
    )


def _question_figures(question: Mapping[str, object]) -> list[Mapping[str, object]]:
    figure = question.get("figure") or question.get("image")
    if figure is None or figure == "":
        return []
    figures = list(figure) if _is_sequence(figure) else [figure]
    return [item for item in figures if isinstance(item, Mapping)]


def _resolve_question_figure_path(figure: Mapping[str, object]) -> Path | None:
    raw_path = _text(figure.get("path") or figure.get("asset_path") or figure.get("asset_id"))
    if not raw_path:
        return None
    path = Path(raw_path).expanduser()
    candidates = [path] if path.is_absolute() else [path, Path.cwd() / path]
    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def _question_figure_alt_text(figure: Mapping[str, object]) -> str:
    return _first_text(
        figure,
        ("alt", "alt_text", "altText", "description", "caption"),
    )


def _question_figure_title(figure: Mapping[str, object]) -> str:
    return _first_text(
        figure,
        ("title", "caption", "asset_id", "source"),
    )


def _question_stem(question: Mapping[str, object]) -> str:
    return _text(question.get("stem") or question.get("question") or question.get("prompt")) or "(missing stem)"


def _question_options(question: Mapping[str, object]) -> list[str]:
    options = question.get("options")
    if isinstance(options, Mapping):
        return [
            f"{_text(key)}. {_text(value)}".strip()
            for key, value in options.items()
            if _text(key) or _text(value)
        ]
    if _is_sequence(options):
        return [_text(option) for option in list(options) if _text(option)]
    if isinstance(options, str):
        return [line.strip() for line in options.splitlines() if line.strip()]
    return []


def _question_answer(question: Mapping[str, object]) -> str:
    answer = _first_present(question, _ANSWER_KEYS)
    return _format_value(answer)


def _question_analysis(question: Mapping[str, object]) -> str:
    analysis = _first_present(question, _ANALYSIS_KEYS)
    return _format_value(analysis)


def _question_knowledge(question: Mapping[str, object]) -> str:
    knowledge = _first_present(question, _KNOWLEDGE_KEYS)
    return _format_value(knowledge)


def _format_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, Mapping):
        parts = [
            f"{_text(key)}={_text(item)}"
            for key, item in value.items()
            if _text(key) or _text(item)
        ]
        return "; ".join(parts)
    if _is_sequence(value):
        return ", ".join(_text(item) for item in list(value) if _text(item))
    return _text(value)


def _safe_artifact_stem(value: object) -> str:
    stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", _text(value)).strip("._")
    return stem or "exam"


def _preview_excerpt(markdown: str, *, max_length: int = 500) -> str:
    text = re.sub(r"\s+", " ", markdown).strip()
    if len(text) <= max_length:
        return text
    return text[: max_length - 3].rstrip() + "..."


def _first_present(mapping: Mapping[str, object], keys: tuple[str, ...]) -> object:
    for key in keys:
        if key in mapping:
            return mapping.get(key)
    return None


def _first_text(mapping: Mapping[str, object], keys: tuple[str, ...]) -> str:
    for key in keys:
        text = _text(mapping.get(key))
        if text:
            return text
    return ""


def _metadata_value(metadata: object, key: str) -> object:
    if not isinstance(metadata, Mapping):
        return None
    return metadata.get(key)


def _metadata_text(metadata: object, keys: tuple[str, ...]) -> str:
    if not isinstance(metadata, Mapping):
        return ""
    return _first_text(metadata, keys)


def _has_value(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, Mapping):
        return bool(value)
    if _is_sequence(value):
        return bool(list(value))
    return True


def _is_sequence(value: object) -> bool:
    return isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray))


def _text(value: object) -> str:
    return str(value or "").strip()


def _number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value)
    if not text:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def _knowledge_count(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, str):
        return 1 if value.strip() else 0
    if isinstance(value, Mapping):
        return len(value)
    if _is_sequence(value):
        return len(list(value))
    return 1


def _looks_like_choice(question_type: str) -> bool:
    normalized = question_type.casefold()
    return any(token in normalized for token in ("choice", "select", "选择", "单选", "多选"))


def _issue(
    path: str,
    kind: str,
    message: str,
    *,
    severity: str = "warning",
    expected: str = "",
    observed: str = "",
) -> ExamQuestionSchemaIssue:
    return ExamQuestionSchemaIssue(
        path=path,
        kind=kind,
        message=message,
        severity=severity,
        expected=expected,
        observed=observed,
    )


def _unique_texts(values: Sequence[str]) -> list[str]:
    result: list[str] = []
    for value in values:
        normalized = str(value or "").strip()
        if normalized and normalized not in result:
            result.append(normalized)
    return result


def _format_number(value: float) -> str:
    if value.is_integer():
        return str(int(value))
    return f"{value:.2f}".rstrip("0").rstrip(".")


__all__ = [
    "EXAM_FAMILY_ID",
    "EXAM_SCHEMA_ID",
    "ExamQuestionSchemaIssue",
    "ExamQuestionSchemaSummary",
    "ExamQuestionSchemaValidationResult",
    "ExamMarkdownImportIssue",
    "ExamMarkdownImportResult",
    "ExamRenderedVersion",
    "ExamMasterEvidence",
    "ExamDeliveryRuntimeResult",
    "build_exam_delivery_runtime",
    "exam_markdown_import_entity_data",
    "inspect_exam_question_schema",
    "parse_exam_markdown_file",
    "parse_exam_markdown_source",
    "render_exam_markdown_preview",
]
