"""Deterministically distill an existing exam DOCX into a reusable user master."""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from src.shared.engine.exam_paper_style import (
    MASTER_METADATA_PLACEHOLDER,
    MASTER_QUESTION_INSERT_MARKER,
    MASTER_TITLE_PLACEHOLDER,
)

_SECTION_HEADING_RE = re.compile(
    r"^\s*[一二三四五六七八九十百]+\s*[、.．]\s*\S+"
)
_QUESTION_RE = re.compile(r"^\s*\d+\s*[.．、]\s*\S+")
_OPTION_RE = re.compile(r"^\s*[A-HＡ-Ｈ]\s*[.．、]\s*\S+", re.IGNORECASE)
_ANSWER_SECTION_RE = re.compile(r"答案\s*(?:与|及)?\s*(?:解析|解答)|参考答案")
_GRAPHIC_MARKERS = ("<w:drawing", "<w:pict", "<w:object", "<v:textbox")


class ExamMasterConversionError(ValueError):
    """The source cannot be safely converted without guessing its structure."""


@dataclass(frozen=True, slots=True)
class ExamMasterConversionResult:
    source_path: Path
    output_path: Path
    source_title: str
    first_section_title: str
    answer_section_detected: bool
    removed_body_block_count: int
    preserved_graphic_run_count: int
    installed_style_names: tuple[str, ...]


def convert_exam_docx_to_user_master(
    source_docx_path: Path | str,
    output_docx_path: Path | str,
) -> ExamMasterConversionResult:
    """Convert one complete exam into a placeholder-based master.

    The source is never modified.  The conversion keeps all page furniture
    before the first real section (including VML text boxes anchored to the
    title), replaces variable metadata, and removes every body block from the
    first section onward.
    """

    source = Path(source_docx_path).expanduser().resolve()
    output = Path(output_docx_path).expanduser().resolve()
    if source.suffix.casefold() != ".docx" or not source.is_file():
        raise ExamMasterConversionError("source_exam_docx_required")
    if output.suffix.casefold() != ".docx":
        raise ExamMasterConversionError("output_master_must_be_docx")
    if source == output:
        raise ExamMasterConversionError("source_exam_must_remain_unchanged")

    try:
        document = Document(str(source))
    except Exception as exc:
        raise ExamMasterConversionError("source_exam_docx_unreadable") from exc

    paragraphs = list(document.paragraphs)
    if not paragraphs:
        raise ExamMasterConversionError("source_exam_has_no_body_paragraphs")
    first_section = _find_first_section_paragraph(paragraphs)
    if first_section is None:
        raise ExamMasterConversionError("first_exam_section_not_detected")
    title = _find_title_paragraph(paragraphs, first_section)
    if title is None:
        raise ExamMasterConversionError("exam_title_not_detected")
    metadata = _find_metadata_paragraph(paragraphs, title, first_section)

    source_title = title.text.strip()
    first_section_title = first_section.text.strip()
    answer_detected = any(
        _ANSWER_SECTION_RE.search(paragraph.text or "")
        for paragraph in paragraphs
    )
    preserved_graphic_runs = sum(
        1 for paragraph in paragraphs for run in paragraph.runs if _run_has_graphics(run)
    )

    representatives = _representative_paragraphs(paragraphs, first_section)
    installed_styles = _install_exam_styles(document, representatives)

    _replace_visible_paragraph_text(title, MASTER_TITLE_PLACEHOLDER)
    if metadata is None:
        metadata = _insert_paragraph_after(title, MASTER_METADATA_PLACEHOLDER)
    else:
        _replace_visible_paragraph_text(metadata, MASTER_METADATA_PLACEHOLDER)
    _replace_visible_paragraph_text(first_section, MASTER_QUESTION_INSERT_MARKER)
    removed_count = _remove_body_blocks_after(first_section)
    _enable_field_updates(document)
    document.core_properties.comments = "distilled-user-exam-master-v1"

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    _validate_converted_master(output)
    return ExamMasterConversionResult(
        source_path=source,
        output_path=output,
        source_title=source_title,
        first_section_title=first_section_title,
        answer_section_detected=answer_detected,
        removed_body_block_count=removed_count,
        preserved_graphic_run_count=preserved_graphic_runs,
        installed_style_names=installed_styles,
    )


def _find_first_section_paragraph(paragraphs: list[Paragraph]) -> Paragraph | None:
    for paragraph in paragraphs:
        if _SECTION_HEADING_RE.match(paragraph.text or ""):
            return paragraph
    return None


def _find_title_paragraph(
    paragraphs: list[Paragraph],
    first_section: Paragraph,
) -> Paragraph | None:
    candidates = paragraphs[: paragraphs.index(first_section)]
    nonempty = [paragraph for paragraph in candidates if paragraph.text.strip()]
    for paragraph in nonempty:
        style_name = str(getattr(paragraph.style, "name", "") or "").casefold()
        text = paragraph.text.strip()
        if style_name.startswith("heading") or "试卷" in text or "测试卷" in text:
            return paragraph
    return nonempty[0] if nonempty else None


def _find_metadata_paragraph(
    paragraphs: list[Paragraph],
    title: Paragraph,
    first_section: Paragraph,
) -> Paragraph | None:
    start = paragraphs.index(title) + 1
    stop = paragraphs.index(first_section)
    for paragraph in paragraphs[start:stop]:
        text = paragraph.text.strip()
        if text and ("考试时间" in text or "满分" in text):
            return paragraph
    return None


def _representative_paragraphs(
    paragraphs: list[Paragraph],
    first_section: Paragraph,
) -> dict[str, Paragraph]:
    start = paragraphs.index(first_section)
    after = paragraphs[start + 1 :]
    question = next(
        (paragraph for paragraph in after if _QUESTION_RE.match(paragraph.text or "")),
        first_section,
    )
    option = next(
        (paragraph for paragraph in after if _OPTION_RE.match(paragraph.text or "")),
        question,
    )
    return {
        "Exam Section Heading": first_section,
        "Exam Question": question,
        "Exam Question Stem": question,
        "Exam Option": option,
        "Exam Answer Space": question,
    }


def _install_exam_styles(
    document: Document,
    representatives: dict[str, Paragraph],
) -> tuple[str, ...]:
    installed: list[str] = []
    for style_name, representative in representatives.items():
        try:
            style = document.styles[style_name]
        except KeyError:
            style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        try:
            style.base_style = document.styles["Normal"]
        except KeyError:
            pass
        _replace_style_formatting(style.element, representative)
        installed.append(style_name)
    return tuple(installed)


def _replace_style_formatting(style_element, representative: Paragraph) -> None:
    for tag in (qn("w:pPr"), qn("w:rPr")):
        existing = style_element.find(tag)
        if existing is not None:
            style_element.remove(existing)

    source_ppr = representative._p.pPr
    if source_ppr is not None:
        copied_ppr = deepcopy(source_ppr)
        paragraph_style = copied_ppr.find(qn("w:pStyle"))
        if paragraph_style is not None:
            copied_ppr.remove(paragraph_style)
        style_element.append(copied_ppr)

    source_rpr = next(
        (
            run._r.rPr
            for run in representative.runs
            if run.text.strip() and run._r.rPr is not None
        ),
        None,
    )
    if source_rpr is not None:
        style_element.append(deepcopy(source_rpr))


def _replace_visible_paragraph_text(paragraph: Paragraph, text: str) -> None:
    editable_runs = [run for run in paragraph.runs if not _run_has_graphics(run)]
    target = next((run for run in editable_runs if run.text), None)
    if target is None:
        target = paragraph.add_run()
        editable_runs.append(target)
    target.text = text
    for run in editable_runs:
        if run is target:
            continue
        run._r.getparent().remove(run._r)


def _run_has_graphics(run) -> bool:
    xml = run._r.xml
    return any(marker in xml for marker in _GRAPHIC_MARKERS)


def _insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.add_run(text)
    return inserted


def _remove_body_blocks_after(paragraph: Paragraph) -> int:
    body = paragraph._p.getparent()
    removed = 0
    sibling = paragraph._p.getnext()
    while sibling is not None:
        next_sibling = sibling.getnext()
        if sibling.tag != qn("w:sectPr"):
            body.remove(sibling)
            removed += 1
        sibling = next_sibling
    return removed


def _enable_field_updates(document: Document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _validate_converted_master(path: Path) -> None:
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ExamMasterConversionError("converted_master_unreadable") from exc
    texts = [paragraph.text.strip() for paragraph in document.paragraphs]
    if sum(MASTER_TITLE_PLACEHOLDER in text for text in texts) != 1:
        raise ExamMasterConversionError("converted_master_title_placeholder_invalid")
    if texts.count(MASTER_QUESTION_INSERT_MARKER) != 1:
        raise ExamMasterConversionError("converted_master_question_placeholder_invalid")
    if any(_ANSWER_SECTION_RE.search(text) for text in texts):
        raise ExamMasterConversionError("converted_master_contains_answer_section")


__all__ = [
    "ExamMasterConversionError",
    "ExamMasterConversionResult",
    "convert_exam_docx_to_user_master",
]
