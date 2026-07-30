"""Exact namespaced material-field analysis and replacement.

The public syntax is ``{{@text:id}}`` or ``{{@time:id}}``.  Identifiers remain
user-defined and match frozen field IDs exactly.  Preview and runtime share
this scanner so neither side can silently reintroduce legacy bare tokens.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Mapping, Sequence

from docx.oxml.ns import qn

from src.shared.engine.docx_material_tokens import extract_docx_material_token_blocks
from src.shared.engine.material_token_contract import (
    MATERIAL_TOKEN_PATTERN,
    MaterialTokenKind,
    MaterialTokenNamespace,
    material_token,
    parse_material_token,
)


EXACT_PLACEHOLDER_PATTERN = MATERIAL_TOKEN_PATTERN


@dataclass(frozen=True, slots=True)
class ExactMaterialPlaceholder:
    """One exact placeholder discovered on replaceable Word surfaces."""

    key: str
    placeholder: str
    occurrence_count: int = 1
    namespace: MaterialTokenNamespace = MaterialTokenNamespace.TEXT


@dataclass(frozen=True, slots=True)
class ExactMaterialReplacementResult:
    """Actual replacements made by the exact material runtime."""

    total_replacements: int = 0
    replaced_keys: tuple[str, ...] = ()


def scan_document_exact_placeholders(document) -> list[ExactMaterialPlaceholder]:
    """Return exact placeholders in first-seen order for one loaded DOCX."""

    token_keys: list[str] = []
    raw_placeholders: dict[str, str] = {}
    counts: dict[str, int] = {}
    for text in _iter_replaceable_text(document):
        for match in EXACT_PLACEHOLDER_PATTERN.finditer(text or ""):
            try:
                ref = parse_material_token(match.group(0))
            except (TypeError, ValueError):
                continue
            if ref.kind is not MaterialTokenKind.FIELD:
                continue
            if ref.key not in counts:
                token_keys.append(ref.key)
                raw_placeholders[ref.key] = ref.token
                counts[ref.key] = 0
            counts[ref.key] += 1
    return [
        ExactMaterialPlaceholder(
            key=parse_material_token(token_key).identifier,
            placeholder=raw_placeholders[token_key],
            occurrence_count=counts[token_key],
            namespace=parse_material_token(token_key).namespace,
        )
        for token_key in token_keys
    ]


def replace_document_exact_placeholders(
    document,
    values: Mapping[str, object],
    *,
    timeline_field_keys: Sequence[str] = (),
) -> ExactMaterialReplacementResult:
    """Replace strict field tokens while keeping business value keys short.

    Mapping keys may be canonical ``@text:/@time:`` tokens (dependency-index
    execution) or plain logical field ids (normal field execution).  Plain ids
    use ``@time:`` only when explicitly owned by ``timeline_field_keys``.
    """

    timeline_ids = {
        str(item or "").strip() for item in timeline_field_keys if str(item or "").strip()
    }
    replacements: dict[str, str] = {}
    token_identifiers: dict[str, str] = {}
    for raw_key, value in dict(values or {}).items():
        key = str(raw_key or "").strip()
        if not key or key != str(raw_key or ""):
            continue
        try:
            ref = parse_material_token(key)
        except (TypeError, ValueError):
            namespace = (
                MaterialTokenNamespace.TIME
                if key in timeline_ids
                else MaterialTokenNamespace.TEXT
            )
            token = material_token(namespace, key)
            identifier = key
        else:
            if ref.kind is not MaterialTokenKind.FIELD:
                continue
            token = ref.token
            identifier = ref.identifier
        replacements[token] = str(value or "")
        token_identifiers[token] = identifier
    if not replacements:
        return ExactMaterialReplacementResult()

    counts: dict[str, int] = {token: 0 for token in replacements}
    for text_items in _iter_ooxml_paragraph_text_items(document):
        paragraph_counts = _replace_tokens_across_text_items(text_items, replacements)
        for token, count in paragraph_counts.items():
            counts[token] += count
    replaced_keys = [
        token_identifiers[token]
        for token, count in counts.items()
        if count
    ]
    return ExactMaterialReplacementResult(
        total_replacements=sum(counts.values()),
        replaced_keys=tuple(replaced_keys),
    )


def _iter_replaceable_text(document):
    for block in extract_docx_material_token_blocks(document).blocks:
        yield block.combined_text


def _iter_ooxml_paragraph_text_items(document):
    """Yield each Word paragraph's text nodes on the scanner's full surface."""

    roots = [document.element]
    roots.extend(
        root
        for part in document.part.package.parts
        if (root := getattr(part, "_element", None)) is not None
    )
    seen_roots: set[int] = set()
    for root in roots:
        identity = id(root)
        if identity in seen_roots:
            continue
        seen_roots.add(identity)
        for paragraph in root.iter(qn("w:p")):
            text_items = tuple(paragraph.iter(qn("w:t")))
            if text_items:
                yield text_items


def _replace_tokens_across_text_items(
    items: Sequence,
    replacements: Mapping[str, str],
) -> dict[str, int]:
    """Replace token spans while retaining surrounding run formatting."""

    texts = [str(getattr(item, "text", "") or "") for item in items]
    full_text = "".join(texts)
    if not full_text:
        return {}

    matches: list[tuple[int, int, str, str]] = []
    occupied: list[tuple[int, int]] = []
    counts: dict[str, int] = {}
    for token, replacement in replacements.items():
        start = 0
        while True:
            index = full_text.find(token, start)
            if index < 0:
                break
            end = index + len(token)
            if not any(index < used_end and end > used_start for used_start, used_end in occupied):
                matches.append((index, end, token, str(replacement or "")))
                occupied.append((index, end))
                counts[token] = counts.get(token, 0) + 1
            start = end

    if not matches:
        return {}

    starts: list[int] = []
    offset = 0
    for text in texts:
        starts.append(offset)
        offset += len(text)

    for start, end, _token, replacement in sorted(matches, reverse=True):
        start_index, start_offset = _item_position(starts, texts, start)
        end_index, end_offset = _item_position(starts, texts, end - 1)
        end_offset += 1
        if start_index == end_index:
            current = str(getattr(items[start_index], "text", "") or "")
            items[start_index].text = (
                current[:start_offset] + replacement + current[end_offset:]
            )
            continue

        first_text = str(getattr(items[start_index], "text", "") or "")
        last_text = str(getattr(items[end_index], "text", "") or "")
        items[start_index].text = first_text[:start_offset] + replacement
        for item_index in range(start_index + 1, end_index):
            items[item_index].text = ""
        items[end_index].text = last_text[end_offset:]
    return counts


def _item_position(
    starts: Sequence[int],
    texts: Sequence[str],
    position: int,
) -> tuple[int, int]:
    for index in range(len(texts) - 1, -1, -1):
        if position >= starts[index] and texts[index]:
            return index, position - starts[index]
    return 0, max(0, position)


__all__ = [
    "EXACT_PLACEHOLDER_PATTERN",
    "ExactMaterialPlaceholder",
    "ExactMaterialReplacementResult",
    "replace_document_exact_placeholders",
    "scan_document_exact_placeholders",
]
