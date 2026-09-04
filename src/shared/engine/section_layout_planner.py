"""Auditable planning and execution for Word section topology.

Section properties are page-layout ownership boundaries, not disposable blank
paragraphs.  This module therefore separates read-only inventory, planning,
mutation, and post-mutation verification.  Destructive cleanup is deliberately
conservative and only runs under an explicit policy.
"""

from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from hashlib import sha256
from typing import TYPE_CHECKING, Any

from lxml import etree

from src.shared.engine.ooxml_ops import (
    clone_element,
    find_or_create_before,
    get_or_add_paragraph_properties_first,
    qn,
)
from src.shared.engine.page_number_planner import build_page_number_execution_plan

if TYPE_CHECKING:
    from docx import Document

    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext


_BREAK_TYPES = frozenset({"nextPage", "continuous", "evenPage", "oddPage"})
_BOUNDARY_MODES = frozenset({"preserve_source", "semantic_rebuild", "normalize_all"})
_CLEANUP_POLICIES = frozenset({"preserve", "remove_proven_redundant"})
_CAPTION_PATTERN = re.compile(r"^(?:图|表|figure|table)\s*[-—－:：]?[\d一二三四五六七八九十]+", re.IGNORECASE)
_PROTECTED_TAGS = frozenset(
    {
        qn("w:drawing"),
        qn("w:object"),
        qn("w:pict"),
        qn("w:fldChar"),
        qn("w:instrText"),
        qn("w:sdt"),
        qn("w:bookmarkStart"),
        qn("w:bookmarkEnd"),
        qn("w:commentRangeStart"),
        qn("w:commentRangeEnd"),
        qn("w:commentReference"),
        qn("w:ins"),
        qn("w:del"),
        qn("w:moveFrom"),
        qn("w:moveTo"),
        qn("w:altChunk"),
    }
)


@dataclass(frozen=True, slots=True)
class SectionBoundarySnapshot:
    ordinal: int
    anchor_kind: str
    body_index: int
    paragraph_index: int | None
    element_path: tuple[int, ...]
    break_type: str
    paragraph_text: str
    layout_signature: str
    full_signature: str
    page_size: tuple[tuple[str, str], ...] = ()
    page_margins: tuple[tuple[str, str], ...] = ()
    header_footer_refs: tuple[tuple[str, str, str], ...] = ()
    title_page: bool = False
    page_numbering: tuple[tuple[str, str], ...] = ()
    columns: tuple[tuple[str, str], ...] = ()
    protected_markup: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class SectionInventory:
    boundaries: tuple[SectionBoundarySnapshot, ...]
    paragraph_count: int
    table_count: int
    body_child_count: int
    digest: str

    @property
    def section_count(self) -> int:
        return len(self.boundaries)


@dataclass(frozen=True, slots=True)
class SectionOperation:
    action: str
    reason: str
    body_index: int = -1
    paragraph_index: int = -1
    element_path: tuple[int, ...] = ()
    target_break_type: str = ""
    priority: int = 0
    blocked_reason: str = ""

    @property
    def blocked(self) -> bool:
        return bool(self.blocked_reason)


@dataclass(frozen=True, slots=True)
class SectionExecutionPlan:
    boundary_mode: str
    source_digest: str
    source_section_count: int
    operations: tuple[SectionOperation, ...] = ()
    required_semantic_starts: tuple[int, ...] = ()
    semantic_break_type: str = "nextPage"
    notes: tuple[str, ...] = ()

    @property
    def blocked_operations(self) -> tuple[SectionOperation, ...]:
        return tuple(operation for operation in self.operations if operation.blocked)


@dataclass(frozen=True, slots=True)
class SectionExecutionReceipt:
    boundary_mode: str
    source_digest: str
    source_section_count: int
    final_digest: str
    final_section_count: int
    applied_operations: tuple[SectionOperation, ...] = ()
    blocked_operations: tuple[SectionOperation, ...] = ()
    validation_errors: tuple[str, ...] = ()

    @property
    def inserted_count(self) -> int:
        return sum(
            operation.action in {"insert_break", "insert_carrier_break"}
            for operation in self.applied_operations
        )

    @property
    def removed_count(self) -> int:
        return sum(operation.action == "remove_break" for operation in self.applied_operations)


def collect_section_inventory(doc: Document) -> SectionInventory:
    """Capture active body section boundaries without mutating the document.

    ``python-docx`` only exposes direct-body section properties.  Word files can
    also contain a paragraph section boundary inside a block content control or
    ``customXml`` container.  Those boundaries must be inventoried even when a
    later mutation is conservatively blocked.
    """

    body = doc.element.body
    children = list(body)
    paragraph_elements = [paragraph._element for paragraph in doc.paragraphs]
    paragraph_indexes = {element: index for index, element in enumerate(paragraph_elements)}
    boundaries: list[SectionBoundarySnapshot] = []

    for sect_pr in _active_section_property_elements(doc):
        element_path = _element_path_from_body(body, sect_pr)
        body_index = element_path[0] if element_path else -1
        anchor_kind = "body"
        paragraph_index: int | None = None
        paragraph_text = ""
        protected_markup: tuple[str, ...] = ()
        parent = sect_pr.getparent()
        paragraph_element = parent.getparent() if parent is not None and parent.tag == qn("w:pPr") else None
        if paragraph_element is not None and paragraph_element.tag == qn("w:p"):
            is_direct = paragraph_element.getparent() is body
            anchor_kind = "paragraph" if is_direct else "nested_paragraph"
            paragraph_index = paragraph_indexes.get(paragraph_element)
            paragraph_text = _paragraph_text(paragraph_element)
            protected_markup = tuple(
                sorted(
                    set(_protected_markup_names(paragraph_element, ignore_section_properties=True))
                    | set(_protected_ancestor_names(paragraph_element, body))
                )
            )

        boundaries.append(
            SectionBoundarySnapshot(
                ordinal=len(boundaries) + 1,
                anchor_kind=anchor_kind,
                body_index=body_index,
                paragraph_index=paragraph_index,
                element_path=element_path,
                break_type=_break_type(sect_pr),
                paragraph_text=paragraph_text,
                layout_signature=_section_signature(sect_pr, ignore_break_type=True),
                full_signature=_section_signature(sect_pr, ignore_break_type=False),
                page_size=_element_attributes(sect_pr.find(qn("w:pgSz"))),
                page_margins=_element_attributes(sect_pr.find(qn("w:pgMar"))),
                header_footer_refs=_header_footer_refs(sect_pr),
                title_page=sect_pr.find(qn("w:titlePg")) is not None,
                page_numbering=_element_attributes(sect_pr.find(qn("w:pgNumType"))),
                columns=_element_attributes(sect_pr.find(qn("w:cols"))),
                protected_markup=tuple(
                    sorted(
                        set(protected_markup)
                        | set(_protected_markup_names(sect_pr, ignore_section_properties=False))
                    )
                ),
            )
        )

    digest_payload = {
        "boundaries": [asdict(boundary) for boundary in boundaries],
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "body_child_count": len(children),
    }
    digest = sha256(
        json.dumps(digest_payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()
    return SectionInventory(
        boundaries=tuple(boundaries),
        paragraph_count=len(doc.paragraphs),
        table_count=len(doc.tables),
        body_child_count=len(children),
        digest=digest,
    )


def iter_active_sections(doc: Document) -> tuple[Any, ...]:
    """Return ``Section`` wrappers for every active section boundary.

    This includes boundaries nested in block-level containers, which are not
    returned by ``doc.sections``.
    """

    from docx.section import Section

    return tuple(
        Section(sect_pr, doc.part)
        for sect_pr in _active_section_property_elements(doc)
    )


def build_section_execution_plan(
    doc: Document,
    config: ResolvedConfig,
    context: PipelineContext,
    *,
    inventory: SectionInventory | None = None,
) -> SectionExecutionPlan:
    """Plan section changes from explicit policy and semantic requirements."""

    source = inventory or collect_section_inventory(doc)
    section_config = config.section
    boundary_mode = effective_boundary_mode(section_config)
    operations: list[SectionOperation] = []
    notes: list[str] = []

    page_plan = build_page_number_execution_plan(doc, context, config.header_footer)
    required_starts = tuple(
        sorted(
            {
                boundary.start_index
                for boundary in page_plan.boundaries
                if boundary.requires_section_break
            }
        )
    )
    protected_body_indexes = {
        _body_index_for_paragraph(doc, start_index - 1)
        for start_index in required_starts
        if start_index > 0
    }

    cleanup_policy = str(getattr(section_config, "empty_break_policy", "preserve") or "preserve")
    caption_policy = str(
        getattr(config.caption, "table_break_policy", "preserve") or "preserve"
    )
    cleanup_candidates: set[int] = set()
    if cleanup_policy == "remove_proven_redundant":
        cleanup_candidates.update(
            _proven_empty_break_body_indexes(doc, source, protected_body_indexes)
        )
    if caption_policy == "remove_proven_redundant":
        cleanup_candidates.update(
            _proven_caption_table_break_body_indexes(doc, source, protected_body_indexes)
        )
    for body_index in sorted(cleanup_candidates, reverse=True):
        cleanup_boundary = _boundary_for_body_index(source, body_index)
        operations.append(
            SectionOperation(
                action="remove_break",
                reason="explicit cleanup policy proved the boundary redundant",
                body_index=body_index,
                element_path=(
                    cleanup_boundary.element_path if cleanup_boundary is not None else ()
                ),
                priority=10,
            )
        )

    target_type = str(getattr(section_config, "section_break_type", "") or "")
    semantic_break_type = (
        target_type
        if boundary_mode == "normalize_all"
        and target_type in {"nextPage", "oddPage", "evenPage"}
        else "nextPage"
    )
    if boundary_mode == "normalize_all" and target_type:
        for boundary in source.boundaries:
            if boundary.body_index in cleanup_candidates or boundary.break_type == target_type:
                continue
            blocked_reason = _boundary_mutation_block_reason(boundary)
            operations.append(
                SectionOperation(
                    action="set_break_type",
                    reason="normalize_all policy",
                    body_index=boundary.body_index,
                    element_path=boundary.element_path,
                    target_break_type=target_type,
                    priority=20,
                    blocked_reason=blocked_reason,
                )
            )

    if boundary_mode in {"semantic_rebuild", "normalize_all"}:
        for start_index in required_starts:
            previous_index = start_index - 1
            body_index = _body_index_for_paragraph(doc, previous_index)
            existing = _boundary_for_body_index(source, body_index)
            if existing is not None and existing.body_index not in cleanup_candidates:
                if boundary_mode == "normalize_all":
                    requires_type_change = existing.break_type != semantic_break_type
                else:
                    requires_type_change = existing.break_type not in {
                        "nextPage",
                        "oddPage",
                        "evenPage",
                    }
                if requires_type_change:
                    operations.append(
                        SectionOperation(
                            action="set_break_type",
                            reason="semantic/page-number boundary must start on a new page",
                            body_index=body_index,
                            paragraph_index=previous_index,
                            element_path=existing.element_path,
                            target_break_type=semantic_break_type,
                            priority=40,
                            blocked_reason=_boundary_mutation_block_reason(existing),
                        )
                    )
                continue
            blocking = _blocking_body_content_between(doc, previous_index, start_index)
            can_use_carrier = bool(blocking) and {
                item.strip() for item in blocking.split(",") if item.strip()
            } == {"tbl"}
            operations.append(
                SectionOperation(
                    action=(
                        "insert_carrier_break" if can_use_carrier else "insert_break"
                    ),
                    reason=(
                        "semantic/page-number boundary after a direct-body table"
                        if can_use_carrier
                        else "semantic/page-number boundary"
                    ),
                    body_index=(
                        _body_index_for_paragraph(doc, start_index)
                        if can_use_carrier
                        else body_index
                    ),
                    paragraph_index=(start_index if can_use_carrier else previous_index),
                    target_break_type=semantic_break_type,
                    priority=30,
                    blocked_reason=(
                        f"cannot insert a section break across body content: {blocking}"
                        if blocking and not can_use_carrier
                        else ""
                    ),
                )
            )
    elif required_starts:
        notes.append(
            "semantic boundaries were inventoried but preserve_source forbids topology insertion"
        )

    operations = _deduplicate_operations(operations)
    return SectionExecutionPlan(
        boundary_mode=boundary_mode,
        source_digest=source.digest,
        source_section_count=source.section_count,
        operations=tuple(sorted(operations, key=lambda item: (item.priority, -item.body_index))),
        required_semantic_starts=required_starts,
        semantic_break_type=semantic_break_type,
        notes=tuple(notes),
    )


def execute_section_execution_plan(
    doc: Document,
    plan: SectionExecutionPlan,
) -> SectionExecutionReceipt:
    """Execute a non-stale plan and verify its declared topology effects."""

    current = collect_section_inventory(doc)
    if current.digest != plan.source_digest:
        raise RuntimeError("section execution plan is stale; document topology changed after planning")
    if plan.blocked_operations:
        messages = tuple(operation.blocked_reason for operation in plan.blocked_operations)
        return SectionExecutionReceipt(
            boundary_mode=plan.boundary_mode,
            source_digest=plan.source_digest,
            source_section_count=plan.source_section_count,
            final_digest=current.digest,
            final_section_count=current.section_count,
            blocked_operations=plan.blocked_operations,
            validation_errors=messages,
        )

    applied: list[SectionOperation] = []
    carrier_boundaries: dict[int, Any] = {}
    for operation in plan.operations:
        changed = False
        if operation.action == "remove_break":
            changed = _remove_boundary_at_path(doc, operation.element_path, operation.body_index)
        elif operation.action == "set_break_type":
            changed = _set_break_type_at_path(
                doc,
                operation.element_path,
                operation.body_index,
                operation.target_break_type,
            )
        elif operation.action == "insert_break":
            changed = _insert_break_after_paragraph(
                doc,
                operation.paragraph_index,
                operation.target_break_type,
            )
        elif operation.action == "insert_carrier_break":
            carrier_boundary = _insert_break_carrier_before_paragraph(
                doc,
                operation.paragraph_index,
                operation.target_break_type,
            )
            changed = carrier_boundary is not None
            if carrier_boundary is not None:
                carrier_boundaries[operation.paragraph_index] = carrier_boundary
        if changed:
            applied.append(operation)

    final = collect_section_inventory(doc)
    expected_count = (
        plan.source_section_count
        + sum(
            operation.action in {"insert_break", "insert_carrier_break"}
            for operation in applied
        )
        - sum(operation.action == "remove_break" for operation in applied)
    )
    errors: list[str] = []
    if final.section_count != expected_count:
        errors.append(
            f"section count mismatch: expected {expected_count}, actual {final.section_count}"
        )
    for start_index in plan.required_semantic_starts:
        if plan.boundary_mode not in {"semantic_rebuild", "normalize_all"}:
            continue
        carrier_boundary = carrier_boundaries.get(start_index)
        if carrier_boundary is not None:
            if carrier_boundary.getparent() is None:
                errors.append(
                    f"missing carrier semantic boundary before paragraph {start_index}"
                )
            elif _break_type(carrier_boundary) != plan.semantic_break_type:
                errors.append(
                    "invalid carrier semantic boundary before paragraph "
                    f"{start_index}: {_break_type(carrier_boundary)}"
                )
            continue
        previous_body_index = _body_index_for_paragraph(doc, start_index - 1)
        boundary = _boundary_for_body_index(final, previous_body_index)
        if boundary is None:
            errors.append(f"missing semantic boundary before paragraph {start_index}")
            continue
        if plan.boundary_mode == "normalize_all":
            valid_type = boundary.break_type == plan.semantic_break_type
        else:
            valid_type = boundary.break_type in {"nextPage", "oddPage", "evenPage"}
        if not valid_type:
            errors.append(
                f"invalid semantic boundary before paragraph {start_index}: {boundary.break_type}"
            )

    return SectionExecutionReceipt(
        boundary_mode=plan.boundary_mode,
        source_digest=plan.source_digest,
        source_section_count=plan.source_section_count,
        final_digest=final.digest,
        final_section_count=final.section_count,
        applied_operations=tuple(applied),
        blocked_operations=plan.blocked_operations,
        validation_errors=tuple(errors),
    )


def effective_boundary_mode(section_config: Any) -> str:
    mode = str(getattr(section_config, "boundary_mode", "semantic_rebuild") or "semantic_rebuild")
    # Compatibility: the old UI exposed only section_break_type.  A concrete
    # value remains an explicit request to normalize all existing boundaries.
    if mode == "preserve_source" and getattr(section_config, "section_break_type", None):
        return "normalize_all"
    return mode


def validate_section_policy(
    section_config: Any,
    caption_config: Any | None = None,
) -> list[tuple[str, str]]:
    issues: list[tuple[str, str]] = []
    mode = str(
        getattr(section_config, "boundary_mode", "semantic_rebuild")
        or "semantic_rebuild"
    )
    if mode not in _BOUNDARY_MODES:
        issues.append(("section.boundary_mode", f"unsupported boundary_mode: {mode}"))
    break_type = getattr(section_config, "section_break_type", None)
    if break_type not in (None, "") and str(break_type) not in _BREAK_TYPES:
        issues.append(
            ("section.section_break_type", f"unsupported section_break_type: {break_type}")
        )
    value = str(getattr(section_config, "empty_break_policy", "preserve") or "preserve")
    if value not in _CLEANUP_POLICIES:
        issues.append(
            (
                "section.empty_break_policy",
                f"unsupported empty_break_policy: {value}",
            )
        )
    if caption_config is not None:
        caption_value = str(
            getattr(caption_config, "table_break_policy", "preserve") or "preserve"
        )
        if caption_value not in _CLEANUP_POLICIES:
            issues.append(
                (
                    "caption.table_break_policy",
                    f"unsupported table_break_policy: {caption_value}",
                )
            )
    return issues


def ensure_section_break_before_paragraph(
    doc: Document,
    paragraph_index: int,
    break_type: str = "nextPage",
) -> bool:
    """Compatibility-safe imperative helper backed by the local-section clone."""

    if paragraph_index <= 0 or paragraph_index >= len(doc.paragraphs):
        return False
    blocking = _blocking_body_content_between(doc, paragraph_index - 1, paragraph_index)
    if blocking:
        if {
            item.strip() for item in blocking.split(",") if item.strip()
        } == {"tbl"}:
            return (
                _insert_break_carrier_before_paragraph(
                    doc,
                    paragraph_index,
                    break_type,
                )
                is not None
            )
        return False
    return _insert_break_after_paragraph(
        doc,
        paragraph_index - 1,
        break_type,
    )


def _proven_empty_break_body_indexes(
    doc: Document,
    inventory: SectionInventory,
    protected_body_indexes: set[int],
) -> set[int]:
    candidates: set[int] = set()
    boundaries = inventory.boundaries
    previous_body_index = -1
    for index, boundary in enumerate(boundaries[:-1]):
        next_boundary = boundaries[index + 1]
        if (
            boundary.anchor_kind != "paragraph"
            or boundary.body_index in protected_body_indexes
            or boundary.paragraph_text.strip()
            or boundary.protected_markup
            or boundary.layout_signature != next_boundary.layout_signature
        ):
            previous_body_index = boundary.body_index
            continue
        if _body_region_is_proven_empty(doc, previous_body_index + 1, boundary.body_index):
            candidates.add(boundary.body_index)
        previous_body_index = boundary.body_index
    return candidates


def _proven_caption_table_break_body_indexes(
    doc: Document,
    inventory: SectionInventory,
    protected_body_indexes: set[int],
) -> set[int]:
    candidates: set[int] = set()
    children = list(doc.element.body)
    paragraph_by_element = {paragraph._element: paragraph for paragraph in doc.paragraphs}
    for index, boundary in enumerate(inventory.boundaries[:-1]):
        if (
            boundary.anchor_kind != "paragraph"
            or boundary.body_index in protected_body_indexes
            or boundary.paragraph_text.strip()
            or boundary.protected_markup
            or boundary.layout_signature != inventory.boundaries[index + 1].layout_signature
        ):
            continue
        if boundary.body_index <= 0 or boundary.body_index + 1 >= len(children):
            continue
        previous = children[boundary.body_index - 1]
        following = children[boundary.body_index + 1]
        if previous.tag != qn("w:p") or following.tag != qn("w:tbl"):
            continue
        paragraph = paragraph_by_element.get(previous)
        style_name = paragraph.style.name if paragraph is not None and paragraph.style else ""
        text = _paragraph_text(previous).strip()
        if "caption" in style_name.casefold() or "题注" in style_name or _CAPTION_PATTERN.match(text):
            candidates.add(boundary.body_index)
    return candidates


def _body_region_is_proven_empty(doc: Document, start: int, end: int) -> bool:
    children = list(doc.element.body)
    for child in children[max(0, start) : end + 1]:
        if child.tag != qn("w:p"):
            return False
        if not _paragraph_is_structurally_blank(child):
            return False
    return True


def _paragraph_is_structurally_blank(paragraph_element) -> bool:
    if _paragraph_text(paragraph_element).strip():
        return False
    if _protected_markup_names(paragraph_element, ignore_section_properties=True):
        return False
    ppr = paragraph_element.find(qn("w:pPr"))
    if ppr is not None:
        for child in ppr:
            if child.tag != qn("w:sectPr"):
                return False
    return all(child.tag == qn("w:pPr") for child in paragraph_element)


def _blocking_body_content_between(
    doc: Document,
    left_paragraph_index: int,
    right_paragraph_index: int,
) -> str:
    if left_paragraph_index < 0 or right_paragraph_index >= len(doc.paragraphs):
        return "paragraph index outside the document body"
    children = list(doc.element.body)
    left = doc.paragraphs[left_paragraph_index]._element
    right = doc.paragraphs[right_paragraph_index]._element
    try:
        left_index = children.index(left)
        right_index = children.index(right)
    except ValueError:
        return "paragraph is not a direct body child"
    blocking = [
        etree.QName(child).localname
        for child in children[left_index + 1 : right_index]
        if child.tag != qn("w:p")
    ]
    return ", ".join(blocking)


def _insert_break_after_paragraph(
    doc: Document,
    paragraph_index: int,
    break_type: str,
) -> bool:
    if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
        return False
    paragraph = doc.paragraphs[paragraph_index]
    ppr = get_or_add_paragraph_properties_first(paragraph._element)
    sect_pr = ppr.find(qn("w:sectPr"))
    created = sect_pr is None
    if sect_pr is None:
        sect_pr = _clone_containing_section_properties(doc, paragraph_index + 1)
        # A copied relationship id would make the new and existing sections
        # explicitly point at the same physical header/footer part.  That is
        # not Word's "Link to Previous" semantics and allows a later write to
        # one section to corrupt the other.  Preserve existing source refs on
        # existing boundaries; a new boundary begins by inheritance and the
        # header/footer planner may then create independent parts as required.
        _strip_explicit_header_footer_references(sect_pr)
        ppr.append(sect_pr)
    old_type = _break_type(sect_pr)
    if old_type == break_type:
        return created
    _set_break_type(sect_pr, break_type)
    return True


def _insert_break_carrier_before_paragraph(
    doc: Document,
    target_paragraph_index: int,
    break_type: str,
):
    """Insert a minimal section-bearing paragraph after intervening tables.

    Word stores a section boundary on a paragraph mark.  When a semantic start
    immediately follows a table there is no paragraph after the table that can
    end the previous section.  A hidden one-twentieth-point carrier provides
    that legal anchor without moving the table into the new section.
    """

    if target_paragraph_index <= 0 or target_paragraph_index >= len(doc.paragraphs):
        return None
    target = doc.paragraphs[target_paragraph_index]._element
    body = doc.element.body
    if target.getparent() is not body:
        return None

    carrier = etree.Element(qn("w:p"))
    ppr = etree.SubElement(carrier, qn("w:pPr"))
    spacing = etree.SubElement(ppr, qn("w:spacing"))
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "1")
    spacing.set(qn("w:lineRule"), "exact")
    mark_properties = etree.SubElement(ppr, qn("w:rPr"))
    etree.SubElement(mark_properties, qn("w:vanish"))
    size = etree.SubElement(mark_properties, qn("w:sz"))
    size.set(qn("w:val"), "2")
    size_cs = etree.SubElement(mark_properties, qn("w:szCs"))
    size_cs.set(qn("w:val"), "2")

    sect_pr = _clone_containing_section_properties(doc, target_paragraph_index)
    _strip_explicit_header_footer_references(sect_pr)
    _set_break_type(sect_pr, break_type)
    ppr.append(sect_pr)
    target.addprevious(carrier)
    return sect_pr


def _strip_explicit_header_footer_references(sect_pr) -> None:
    for ref_tag in ("w:headerReference", "w:footerReference"):
        for reference in list(sect_pr.findall(qn(ref_tag))):
            sect_pr.remove(reference)


def _clone_containing_section_properties(doc: Document, target_paragraph_index: int):
    target_body_index = _body_index_for_paragraph(doc, target_paragraph_index)
    for body_index, child in enumerate(doc.element.body):
        if body_index < target_body_index:
            continue
        if child.tag == qn("w:p"):
            ppr = child.find(qn("w:pPr"))
            sect_pr = ppr.find(qn("w:sectPr")) if ppr is not None else None
            if sect_pr is not None:
                return clone_element(sect_pr)
        elif child.tag == qn("w:sectPr"):
            return clone_element(child)
    return etree.Element(qn("w:sectPr"))


def _remove_boundary_at_path(
    doc: Document,
    element_path: tuple[int, ...],
    body_index: int,
) -> bool:
    sect_pr = _resolve_section_property(doc, element_path, body_index)
    if sect_pr is None or sect_pr.getparent() is None or sect_pr.getparent().tag != qn("w:pPr"):
        return False
    sect_pr.getparent().remove(sect_pr)
    return True


def _set_break_type_at_path(
    doc: Document,
    element_path: tuple[int, ...],
    body_index: int,
    break_type: str,
) -> bool:
    sect_pr = _resolve_section_property(doc, element_path, body_index)
    if sect_pr is None or _break_type(sect_pr) == break_type:
        return False
    _set_break_type(sect_pr, break_type)
    return True


def _set_break_type(sect_pr, break_type: str) -> None:
    type_element = find_or_create_before(
        sect_pr,
        "w:type",
        (
            "w:pgSz",
            "w:pgMar",
            "w:paperSrc",
            "w:pgBorders",
            "w:lnNumType",
            "w:pgNumType",
            "w:cols",
            "w:formProt",
            "w:vAlign",
            "w:noEndnote",
            "w:titlePg",
            "w:textDirection",
            "w:bidi",
            "w:rtlGutter",
            "w:docGrid",
            "w:printerSettings",
            "w:sectPrChange",
        ),
    )
    type_element.set(qn("w:val"), break_type)


def _deduplicate_operations(operations: list[SectionOperation]) -> list[SectionOperation]:
    # Later/higher-priority semantic operations own the final type for a boundary.
    by_key: dict[tuple[Any, ...], SectionOperation] = {}
    for operation in operations:
        if operation.action in {"remove_break", "set_break_type"}:
            key = (operation.action, operation.element_path or (operation.body_index,))
        else:
            key = (operation.action, operation.body_index, operation.paragraph_index)
        existing = by_key.get(key)
        if existing is None or operation.priority >= existing.priority:
            by_key[key] = operation
    return list(by_key.values())


def _boundary_for_body_index(
    inventory: SectionInventory,
    body_index: int,
) -> SectionBoundarySnapshot | None:
    matches = [
        boundary for boundary in inventory.boundaries if boundary.body_index == body_index
    ]
    return next(
        (boundary for boundary in matches if boundary.anchor_kind == "paragraph"),
        matches[0] if matches else None,
    )


def _active_section_property_elements(doc: Document) -> tuple[Any, ...]:
    body = doc.element.body
    active: list[Any] = []
    for element in body.iter(qn("w:sectPr")):
        parent = element.getparent()
        if parent is body:
            active.append(element)
            continue
        if parent is None or parent.tag != qn("w:pPr"):
            continue
        paragraph = parent.getparent()
        if paragraph is not None and paragraph.tag == qn("w:p"):
            active.append(element)
    return tuple(active)


def _element_path_from_body(body, element) -> tuple[int, ...]:
    path: list[int] = []
    current = element
    while current is not body:
        parent = current.getparent()
        if parent is None:
            return ()
        path.append(parent.index(current))
        current = parent
    return tuple(reversed(path))


def _resolve_element_path(body, path: tuple[int, ...]):
    current = body
    for index in path:
        if index < 0 or index >= len(current):
            return None
        current = current[index]
    return current


def _resolve_section_property(
    doc: Document,
    element_path: tuple[int, ...],
    body_index: int,
):
    body = doc.element.body
    if element_path:
        element = _resolve_element_path(body, element_path)
        if element is not None and element.tag == qn("w:sectPr"):
            return element
    children = list(body)
    if body_index < 0 or body_index >= len(children):
        return None
    child = children[body_index]
    if child.tag == qn("w:p"):
        ppr = child.find(qn("w:pPr"))
        return ppr.find(qn("w:sectPr")) if ppr is not None else None
    return child if child.tag == qn("w:sectPr") else None


def _protected_ancestor_names(paragraph_element, body) -> tuple[str, ...]:
    names: set[str] = set()
    current = paragraph_element.getparent()
    protected_ancestors = {
        qn("w:sdt"),
        qn("w:customXml"),
        qn("w:ins"),
        qn("w:del"),
        qn("w:moveFrom"),
        qn("w:moveTo"),
    }
    while current is not None and current is not body:
        if current.tag in protected_ancestors:
            names.add(etree.QName(current).localname)
        current = current.getparent()
    return tuple(sorted(names))


def _boundary_mutation_block_reason(boundary: SectionBoundarySnapshot) -> str:
    if "sectPrChange" in boundary.protected_markup:
        return "section properties contain tracked sectPrChange markup"
    if boundary.anchor_kind == "nested_paragraph":
        containers = ", ".join(boundary.protected_markup) or "nested container"
        return f"section boundary is nested inside protected content: {containers}"
    return ""


def _body_index_for_paragraph(doc: Document, paragraph_index: int) -> int:
    if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
        return -1
    element = doc.paragraphs[paragraph_index]._element
    for body_index, child in enumerate(doc.element.body):
        if child is element:
            return body_index
    return -1


def _break_type(sect_pr) -> str:
    type_element = sect_pr.find(qn("w:type"))
    if type_element is None:
        return "nextPage"
    return str(type_element.get(qn("w:val"), "nextPage") or "nextPage")


def _section_signature(sect_pr, *, ignore_break_type: bool) -> str:
    cloned = clone_element(sect_pr)
    if ignore_break_type:
        type_element = cloned.find(qn("w:type"))
        if type_element is not None:
            cloned.remove(type_element)
    for element in cloned.iter():
        for attribute in list(element.attrib):
            if etree.QName(attribute).localname.startswith("rsid"):
                del element.attrib[attribute]
    raw = etree.tostring(cloned, method="c14n", with_comments=False)
    return sha256(raw).hexdigest()


def _element_attributes(element) -> tuple[tuple[str, str], ...]:
    if element is None:
        return ()
    return tuple(
        sorted((etree.QName(name).localname, str(value)) for name, value in element.attrib.items())
    )


def _header_footer_refs(sect_pr) -> tuple[tuple[str, str, str], ...]:
    refs: list[tuple[str, str, str]] = []
    for tag, kind in (("w:headerReference", "header"), ("w:footerReference", "footer")):
        for element in sect_pr.findall(qn(tag)):
            refs.append(
                (
                    kind,
                    str(element.get(qn("w:type"), "default") or "default"),
                    str(element.get(qn("r:id"), "") or ""),
                )
            )
    return tuple(sorted(refs))


def _protected_markup_names(element, *, ignore_section_properties: bool) -> tuple[str, ...]:
    names: set[str] = set()
    for descendant in element.iter():
        if ignore_section_properties and descendant.tag == qn("w:sectPr"):
            continue
        if descendant.tag in _PROTECTED_TAGS or descendant.tag == qn("w:sectPrChange"):
            names.add(etree.QName(descendant).localname)
    return tuple(sorted(names))


def _paragraph_text(paragraph_element) -> str:
    return "".join(node.text or "" for node in paragraph_element.iter(qn("w:t")))


__all__ = [
    "SectionBoundarySnapshot",
    "SectionExecutionPlan",
    "SectionExecutionReceipt",
    "SectionInventory",
    "SectionOperation",
    "build_section_execution_plan",
    "collect_section_inventory",
    "effective_boundary_mode",
    "ensure_section_break_before_paragraph",
    "execute_section_execution_plan",
    "iter_active_sections",
    "validate_section_policy",
]
