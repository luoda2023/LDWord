"""Official-document assembly from master contracts."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Callable, Mapping
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.config.master_library import (
    MasterSpec,
    resolve_official_master_for_contract,
)
from src.config.master_preflight import check_master_preflight
from src.config.official_document_profiles import (
    OfficialDocumentAssemblyContract,
    get_official_document_assembly_contract,
    get_official_document_profile,
)
from src.config.official_material_form import official_material_field_label
from src.shared.io.artifact_publication import (
    StagedArtifact,
    publish_staged_artifacts,
)
from src.shared.engine.fixed_layout_text import replace_fixed_layout_placeholders
from src.shared.engine.run_ops import (
    get_full_text,
    replace_run_text,
    set_run_fonts,
)
from src.shared.engine.material_token_contract import (
    MATERIAL_TOKEN_PATTERN,
    MaterialTokenNamespace,
    material_token,
    parse_material_token,
)


_REFERENCE_PAGE_HEADINGS = ("版式占位符说明", "母版占位符说明")
_OFFICIAL_BODY_TOKEN = material_token(
    MaterialTokenNamespace.TEXT,
    "official_body",
)
_OFFICIAL_DISPLAY_TOKENS = (
    material_token(MaterialTokenNamespace.TEXT, "official_organization"),
    material_token(MaterialTokenNamespace.TEXT, "official_title"),
)
_BODY_LEVEL_PATTERNS = (
    (1, re.compile(r"^[一二三四五六七八九十百]+、")),
    (2, re.compile(r"^（[一二三四五六七八九十百]+）")),
    (3, re.compile(r"^\d+[.．]")),
    (4, re.compile(r"^（\d+）")),
)

DocxPdfRenderer = Callable[[Path, Path], str | None]
DocxPdfRendererEntry = tuple[str, DocxPdfRenderer]
DocxPdfRendererResolver = Callable[[], DocxPdfRendererEntry | None]


@dataclass(frozen=True, slots=True)
class _ReviewPdfOutcome:
    status: str
    path: Path | None = None
    renderer: str = ""
    issue: str = ""


@dataclass(frozen=True, slots=True)
class OfficialDocumentAssemblyResult:
    status: str
    profile_id: str
    master_id: str = ""
    material_schema_ids: tuple[str, ...] = ()
    docx_path: Path | None = None
    internal_review_docx_path: Path | None = None
    archive_manifest_path: Path | None = None
    archive_manifest_markdown_path: Path | None = None
    review_pdf_path: Path | None = None
    review_pdf_status: str = "not_requested"
    review_pdf_renderer: str = ""
    review_pdf_issue: str = ""
    replaced_placeholders: tuple[str, ...] = ()
    missing_required_fields: tuple[str, ...] = ()
    unresolved_placeholders: tuple[str, ...] = ()

    @property
    def ok(self) -> bool:
        return self.status == "ok"

    @property
    def output_paths(self) -> dict[str, str]:
        paths: dict[str, str] = {}
        if self.docx_path is not None:
            paths["official_docx"] = str(self.docx_path)
        if self.internal_review_docx_path is not None:
            paths["internal_review_docx"] = str(self.internal_review_docx_path)
        if self.archive_manifest_path is not None:
            paths["archive_manifest"] = str(self.archive_manifest_path)
        if self.archive_manifest_markdown_path is not None:
            paths["archive_manifest_md"] = str(self.archive_manifest_markdown_path)
        if (
            self.review_pdf_status == "generated"
            and self.review_pdf_path is not None
            and self.review_pdf_path.is_file()
        ):
            paths["review_pdf"] = str(self.review_pdf_path)
        return paths

    def to_dict(self) -> dict[str, object]:
        return {
            "status": self.status,
            "profile_id": self.profile_id,
            "master_id": self.master_id,
            "material_schema_ids": list(self.material_schema_ids),
            "docx_path": str(self.docx_path) if self.docx_path is not None else "",
            "internal_review_docx_path": (
                str(self.internal_review_docx_path)
                if self.internal_review_docx_path is not None
                else ""
            ),
            "archive_manifest_path": (
                str(self.archive_manifest_path)
                if self.archive_manifest_path is not None
                else ""
            ),
            "archive_manifest_markdown_path": (
                str(self.archive_manifest_markdown_path)
                if self.archive_manifest_markdown_path is not None
                else ""
            ),
            "review_pdf_path": str(self.review_pdf_path or ""),
            "review_pdf_status": self.review_pdf_status,
            "review_pdf_renderer": self.review_pdf_renderer,
            "review_pdf_issue": self.review_pdf_issue,
            "output_paths": self.output_paths,
            "replaced_placeholders": list(self.replaced_placeholders),
            "missing_required_fields": list(self.missing_required_fields),
            "unresolved_placeholders": list(self.unresolved_placeholders),
        }


def assemble_official_document_docx(
    profile_id: str,
    entity_data: Mapping[str, object],
    output_dir: Path | str,
    *,
    field_aliases: Mapping[str, str] | None = None,
    filename: str | None = None,
    master: MasterSpec | None = None,
    removed_field_keys: set[str] | tuple[str, ...] | list[str] = (),
    remove_reference_page: bool = True,
    generate_review_pdf: bool = False,
    review_pdf_renderer_resolver: DocxPdfRendererResolver | None = None,
    delivery_versions: tuple[str, ...] | None = None,
) -> OfficialDocumentAssemblyResult:
    """Assemble one official-document DOCX from the registered master contract."""

    contract = get_official_document_assembly_contract(profile_id)
    normalized_profile_id = str(profile_id or "").strip()
    if contract is None:
        return OfficialDocumentAssemblyResult(
            status="unknown_profile",
            profile_id=normalized_profile_id,
            review_pdf_status=_blocked_review_pdf_status(generate_review_pdf),
        )

    requested_versions = tuple(
        dict.fromkeys(
            str(item or "").strip()
            for item in (
                contract.delivery_versions
                if delivery_versions is None
                else delivery_versions
            )
            if str(item or "").strip()
        )
    )
    supported_versions = {
        "official_docx",
        "internal_review_docx",
        "archive_manifest",
        "review_pdf",
    }
    if not requested_versions or any(
        item not in supported_versions for item in requested_versions
    ):
        return OfficialDocumentAssemblyResult(
            status="delivery_versions_invalid",
            profile_id=contract.profile_id,
            material_schema_ids=contract.material_schema_ids,
            review_pdf_status=_blocked_review_pdf_status(generate_review_pdf),
        )

    master_spec = resolve_official_master_for_contract(
        contract,
        requested=master,
    )
    if master_spec is None:
        return OfficialDocumentAssemblyResult(
            status="master_not_found",
            profile_id=contract.profile_id,
            material_schema_ids=contract.material_schema_ids,
            review_pdf_status=_blocked_review_pdf_status(generate_review_pdf),
        )
    if (
        master_spec.supported_assembly_types
        and contract.profile_id not in master_spec.supported_assembly_types
    ):
        return OfficialDocumentAssemblyResult(
            status="master_incompatible",
            profile_id=contract.profile_id,
            master_id=master_spec.master_id,
            material_schema_ids=contract.material_schema_ids,
            review_pdf_status=_blocked_review_pdf_status(generate_review_pdf),
        )

    preflight = check_master_preflight(master_spec)
    if preflight.status != "ok":
        return OfficialDocumentAssemblyResult(
            status=preflight.status,
            profile_id=contract.profile_id,
            master_id=master_spec.master_id,
            material_schema_ids=contract.material_schema_ids,
            missing_required_fields=preflight.missing_required_placeholders,
            review_pdf_status=_blocked_review_pdf_status(generate_review_pdf),
        )

    missing_required = _missing_required_fields(
        contract,
        entity_data,
        removed_field_keys=removed_field_keys,
    )
    if missing_required:
        return OfficialDocumentAssemblyResult(
            status="missing_required_fields",
            profile_id=contract.profile_id,
            master_id=master_spec.master_id,
            material_schema_ids=contract.material_schema_ids,
            missing_required_fields=missing_required,
            review_pdf_status=_blocked_review_pdf_status(generate_review_pdf),
        )

    target = _target_docx_path(output_dir, filename, entity_data, contract.profile_id)
    internal_review_target = _internal_review_docx_path(target)
    archive_manifest_target = _archive_manifest_path(target)
    archive_markdown_target = _archive_manifest_markdown_path(target)
    review_pdf_target = _review_pdf_path(target)

    with TemporaryDirectory(prefix="lark-official-assembly-") as temporary:
        stage_root = Path(temporary)
        stage_target = stage_root / target.name
        stage_internal_review = stage_root / internal_review_target.name
        stage_archive_manifest = stage_root / archive_manifest_target.name
        stage_archive_markdown = stage_root / archive_markdown_target.name
        stage_review_pdf = stage_root / review_pdf_target.name

        assembly_values = _effective_official_entity_data(contract, entity_data)
        if master_spec.source_type != "builtin":
            assembly_values["title"] = str(entity_data.get("title", "") or "")
        document = Document(str(master_spec.docx_path))
        _repair_official_display_line_heights(document)
        if remove_reference_page:
            _remove_placeholder_reference_page(document)
        _remove_empty_optional_paragraphs(document, contract, assembly_values)
        _remove_empty_optional_table_rows(document, contract, assembly_values)
        _insert_profile_specific_paragraphs(document, contract, assembly_values)
        _compact_profile_specific_layout(document, contract)
        replacements = _placeholder_replacements(
            contract,
            assembly_values,
            field_aliases=field_aliases,
        )
        replaced = _replace_placeholders(document, replacements)
        _normalize_terminal_table_separators(document)
        if master_spec.source_type == "builtin":
            _fit_official_redhead_marks(document)
        _stabilize_heading_blocks(document)
        _stabilize_trailing_closing_block(document)
        unresolved = _unresolved_official_placeholders(document)
        if unresolved:
            return OfficialDocumentAssemblyResult(
                status="unresolved_placeholders",
                profile_id=contract.profile_id,
                master_id=master_spec.master_id,
                material_schema_ids=contract.material_schema_ids,
                review_pdf_status=_blocked_review_pdf_status(generate_review_pdf),
                replaced_placeholders=replaced,
                unresolved_placeholders=unresolved,
            )

        document.save(str(stage_target))
        needs_internal_stage = bool(
            "internal_review_docx" in requested_versions
            or "archive_manifest" in requested_versions
            or generate_review_pdf
        )
        if needs_internal_stage:
            _insert_internal_review_notice(document)
            document.save(str(stage_internal_review))
        review_pdf_outcome = _ReviewPdfOutcome(status="not_requested")
        if generate_review_pdf:
            review_pdf_outcome = _render_optional_review_pdf(
                stage_internal_review,
                stage_review_pdf,
                renderer_resolver=review_pdf_renderer_resolver,
            )
        published_review_pdf = (
            review_pdf_target
            if review_pdf_outcome.status == "generated"
            and review_pdf_outcome.path is not None
            else None
        )
        if "archive_manifest" in requested_versions:
            _write_archive_manifests(
                stage_archive_manifest,
                stage_archive_markdown,
                contract,
                master_spec,
                entity_data,
                formal_docx_path=target,
                internal_review_docx_path=internal_review_target,
                published_archive_manifest_path=archive_manifest_target,
                published_archive_manifest_markdown_path=archive_markdown_target,
                review_pdf_path=published_review_pdf,
                review_pdf_status=review_pdf_outcome.status,
                review_pdf_renderer=review_pdf_outcome.renderer,
                review_pdf_issue=review_pdf_outcome.issue,
                replaced_placeholders=replaced,
                delivery_versions=requested_versions,
            )
        staged_artifacts: list[StagedArtifact] = []
        if "official_docx" in requested_versions:
            staged_artifacts.append(
                StagedArtifact("official_docx", stage_target, target)
            )
        if "internal_review_docx" in requested_versions:
            staged_artifacts.append(
                StagedArtifact(
                    "internal_review_docx",
                    stage_internal_review,
                    internal_review_target,
                )
            )
        if "archive_manifest" in requested_versions:
            staged_artifacts.extend(
                (
                    StagedArtifact(
                        "archive_manifest",
                        stage_archive_manifest,
                        archive_manifest_target,
                    ),
                    StagedArtifact(
                        "archive_manifest_md",
                        stage_archive_markdown,
                        archive_markdown_target,
                    ),
                )
            )
        if published_review_pdf is not None:
            staged_artifacts.append(
                StagedArtifact("review_pdf", stage_review_pdf, review_pdf_target)
            )
        publish_staged_artifacts(
            staged_artifacts,
            execution_id=f"official-assembly-{uuid4().hex}",
        )

    return OfficialDocumentAssemblyResult(
        status="ok",
        profile_id=contract.profile_id,
        master_id=master_spec.master_id,
        material_schema_ids=contract.material_schema_ids,
        docx_path=(target if "official_docx" in requested_versions else None),
        internal_review_docx_path=(
            internal_review_target
            if "internal_review_docx" in requested_versions
            else None
        ),
        archive_manifest_path=(
            archive_manifest_target
            if "archive_manifest" in requested_versions
            else None
        ),
        archive_manifest_markdown_path=(
            archive_markdown_target
            if "archive_manifest" in requested_versions
            else None
        ),
        review_pdf_path=published_review_pdf,
        review_pdf_status=review_pdf_outcome.status,
        review_pdf_renderer=review_pdf_outcome.renderer,
        review_pdf_issue=review_pdf_outcome.issue,
        replaced_placeholders=replaced,
        unresolved_placeholders=(),
    )


def _missing_required_fields(
    contract: OfficialDocumentAssemblyContract,
    entity_data: Mapping[str, object],
    *,
    removed_field_keys: set[str] | tuple[str, ...] | list[str] = (),
) -> tuple[str, ...]:
    values = entity_data if isinstance(entity_data, Mapping) else {}
    removed = {str(key) for key in removed_field_keys}
    missing: list[str] = []
    for binding in contract.field_bindings:
        if not binding.required or binding.field_key in removed:
            continue
        value = values.get(binding.field_key, "")
        if str(value or "").strip():
            continue
        missing.append(binding.field_key)
    return tuple(dict.fromkeys(missing))


def _placeholder_replacements(
    contract: OfficialDocumentAssemblyContract,
    entity_data: Mapping[str, object],
    *,
    field_aliases: Mapping[str, str] | None = None,
) -> dict[str, str]:
    values = entity_data if isinstance(entity_data, Mapping) else {}
    replacements: dict[str, str] = {}
    for binding in contract.field_bindings:
        token = material_token(MaterialTokenNamespace.TEXT, binding.placeholder_id)
        value = (
            str(values.get(binding.field_key, "") or "")
            if binding.applicable
            else ""
        )
        if binding.field_key == "body":
            value = _strip_duplicate_recipient_line(
                value,
                str(values.get("recipient", "") or ""),
            )
        elif binding.field_key == "recipient":
            # The built-in master owns the full-width colon. Accept pasted or
            # model-produced recipient values with punctuation without
            # producing a visible ``：：`` suffix in the formal document.
            value = re.sub(r"[:：]+\s*$", "", value).strip()
        replacements[token] = value
        label = official_material_field_label(binding.field_key)
        if label:
            replacements[material_token(MaterialTokenNamespace.TEXT, label)] = value
            replacements[material_token(MaterialTokenNamespace.TEXT, label + "1")] = value
    for raw_key, raw_value in values.items():
        key = str(raw_key or "").strip()
        if key:
            replacements.setdefault(
                material_token(MaterialTokenNamespace.TEXT, key),
                str(raw_value or ""),
            )
    for raw_alias, raw_field_key in dict(field_aliases or {}).items():
        alias = str(raw_alias or "").strip()
        field_key = str(raw_field_key or "").strip()
        if alias and field_key in values:
            replacements[
                material_token(MaterialTokenNamespace.TEXT, alias)
            ] = str(values.get(field_key, "") or "")
    return replacements


def _effective_official_entity_data(
    contract: OfficialDocumentAssemblyContract,
    entity_data: Mapping[str, object],
) -> dict[str, object]:
    """Derive safe display values without inventing administrative facts.

    The form should ask once for facts that cannot be inferred.  Presentation
    details owned by a master (文种后缀、标点、印发尾注) are normalized here,
    and an omitted落款机关 safely reuses the already supplied发文机关 instead
    of producing a date-only closing block.
    """

    values = dict(entity_data or {})
    profile_id = contract.profile_id
    organization = str(values.get("organization", "") or "").strip()
    if profile_id == "order":
        organization = re.sub(r"(?:命令（令）|命令|令)\s*$", "", organization).strip()
    elif profile_id == "minutes":
        organization = re.sub(r"纪要\s*$", "", organization).strip()
    values["organization"] = organization

    issuer = str(values.get("issuer", "") or "").strip()
    if not issuer and organization and profile_id != "order":
        values["issuer"] = organization

    values["title"] = _balance_official_title(
        str(values.get("title", "") or "")
    )
    values["copy_scope"] = re.sub(
        r"[。；;：:、，,]+\s*$",
        "",
        str(values.get("copy_scope", "") or "").strip(),
    )
    printing_date = re.sub(
        r"\s*印发\s*$",
        "",
        str(values.get("printing_date", "") or "").strip(),
    )
    values["printing_date"] = f"{printing_date}印发" if printing_date else ""
    return values


def _balance_official_title(value: str) -> str:
    """Insert stable, balanced breaks only for titles that must wrap."""

    title = re.sub(r"[\t\u3000 ]+", "", str(value or "").strip())
    if not title or "\n" in title or len(title) <= 22:
        return title
    line_count = 2 if len(title) <= 38 else 3
    remaining = title
    lines: list[str] = []
    forbidden_endings = frozenset("和与及、的并暨")
    for lines_left in range(line_count, 1, -1):
        target = round(len(remaining) / lines_left)
        lower = max(2, target - 3)
        upper = min(len(remaining) - (lines_left - 1) * 2, target + 3)
        candidates = range(lower, upper + 1)
        split_at = min(
            candidates,
            key=lambda index: (
                remaining[index - 1] in forbidden_endings,
                abs(index - target),
            ),
        )
        lines.append(remaining[:split_at])
        remaining = remaining[split_at:]
    lines.append(remaining)
    return "\n".join(lines)


def _strip_duplicate_recipient_line(body: str, recipient: str) -> str:
    """Remove a model-repeated recipient already owned by the master field."""

    normalized_recipient = _recipient_comparison_key(recipient)
    if not normalized_recipient:
        return str(body or "")
    lines = re.split(r"\r\n?|\n", str(body or ""))
    first_index = next(
        (index for index, line in enumerate(lines) if line.strip()),
        None,
    )
    if first_index is None:
        return str(body or "")
    first = lines[first_index].strip()
    if not first.endswith(("：", ":")):
        return str(body or "")
    if _recipient_comparison_key(first) != normalized_recipient:
        return str(body or "")
    del lines[first_index]
    return "\n".join(lines).strip()


def _recipient_comparison_key(value: str) -> str:
    return re.sub(r"[\s\u3000，,、；;：:]+", "", str(value or ""))


def _remove_empty_optional_paragraphs(
    document,
    contract: OfficialDocumentAssemblyContract,
    entity_data: Mapping[str, object],
) -> None:
    values = entity_data if isinstance(entity_data, Mapping) else {}
    preserve_blank_layout_tokens = {
        "{{@text:official_security_level}}",
        "{{@text:official_urgency}}",
    }
    if contract.profile_id in {"bulletin", "announcement", "notice_public"}:
        # These public-facing types may legitimately omit a document number,
        # but its grid line still reserves the standard mark-to-rule geometry.
        preserve_blank_layout_tokens.add("{{@text:official_document_no}}")
    empty_tokens = {
        material_token(MaterialTokenNamespace.TEXT, binding.placeholder_id)
        for binding in contract.field_bindings
        if not binding.required
        and (
            not binding.applicable
            or not str(values.get(binding.field_key, "") or "").strip()
        )
        and material_token(
            MaterialTokenNamespace.TEXT,
            binding.placeholder_id,
        ) not in preserve_blank_layout_tokens
    }
    if not empty_tokens:
        return
    for paragraph in list(document.paragraphs):
        text = get_full_text(paragraph)
        if not any(token in text for token in empty_tokens):
            continue
        parent = paragraph._p.getparent()
        if parent is not None:
            parent.remove(paragraph._p)


def _remove_empty_optional_table_rows(
    document,
    contract: OfficialDocumentAssemblyContract,
    entity_data: Mapping[str, object],
) -> None:
    values = entity_data if isinstance(entity_data, Mapping) else {}
    empty_tokens = {
        material_token(MaterialTokenNamespace.TEXT, binding.placeholder_id)
        for binding in contract.field_bindings
        if not binding.required
        and (
            not binding.applicable
            or not str(values.get(binding.field_key, "") or "").strip()
        )
    }
    if not empty_tokens:
        return
    for table in list(document.tables):
        for row in list(table.rows):
            text = "\n".join(cell.text for cell in row.cells)
            if not any(token in text for token in empty_tokens):
                continue
            # Metadata rows containing another non-empty/applicable placeholder
            # must survive. Imprint rows contain exactly one value placeholder.
            row_tokens = set(_official_placeholder_ids(text))
            if any(
                material_token(MaterialTokenNamespace.TEXT, token) not in empty_tokens
                for token in row_tokens
            ):
                continue
            parent = row._tr.getparent()
            if parent is not None:
                parent.remove(row._tr)
        if table.rows:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in list(cell.paragraphs):
                        paragraph_tokens = set(
                            _official_placeholder_ids(get_full_text(paragraph))
                        )
                        if not paragraph_tokens or len(cell.paragraphs) <= 1:
                            continue
                        if all(
                            material_token(MaterialTokenNamespace.TEXT, token)
                            in empty_tokens
                            for token in paragraph_tokens
                        ):
                            paragraph._p.getparent().remove(paragraph._p)
            continue
        parent = table._tbl.getparent()
        if parent is not None:
            parent.remove(table._tbl)


def _insert_profile_specific_paragraphs(
    document,
    contract: OfficialDocumentAssemblyContract,
    entity_data: Mapping[str, object],
) -> None:
    if contract.profile_id != "minutes":
        return
    existing_text = "\n".join(
        get_full_text(paragraph) for paragraph in _iter_all_paragraphs(document)
    )
    if (
        "{{@text:official_meeting_time}}" in existing_text
        or "{{@text:official_meeting_attendees}}" in existing_text
    ):
        return
    body_paragraph = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if "{{@text:official_body}}" in get_full_text(paragraph)
        ),
        None,
    )
    if body_paragraph is None:
        return

    values = entity_data if isinstance(entity_data, Mapping) else {}
    rows = (
        ("meeting_date", "会议时间", "official_meeting_time"),
        ("participants", "参会人员", "official_meeting_attendees"),
    )
    for field_key, label, placeholder_id in rows:
        if not str(values.get(field_key, "") or "").strip():
            continue
        paragraph = body_paragraph.insert_paragraph_before(
            f"{label}："
            + material_token(MaterialTokenNamespace.TEXT, placeholder_id),
            style=body_paragraph.style,
        )
        paragraph.paragraph_format.first_line_indent = (
            body_paragraph.paragraph_format.first_line_indent
        )


def _compact_profile_specific_layout(
    document,
    contract: OfficialDocumentAssemblyContract,
) -> None:
    """Keep compact official documents from creating a trailing page.

    The master intentionally leaves a spacer immediately before the printing
    table.  Word can use that spacer to push the last table row (or only the
    end-of-document marker) onto a second page for request and approval
    documents.  Removing the empty spacer preserves the visible hierarchy and
    keeps these short, single-page documents stable across Word renderers.
    """

    if contract.profile_id not in {"minutes", "request", "approval"} or not document.tables:
        return
    printing_table = next(
        (
            table._tbl
            for table in document.tables
            if "official_printing_" in "\n".join(
                cell.text for row in table.rows for cell in row.cells
            )
        ),
        None,
    )
    if printing_table is None:
        return
    preceding = printing_table.getprevious()
    if (
        preceding is not None
        and preceding.tag.endswith("p")
        and not "".join(preceding.itertext()).strip()
        and not preceding.xpath(".//w:sectPr")
    ):
        preceding.getparent().remove(preceding)

    # Request documents commonly include an attachment line.  The master also
    # leaves a blank paragraph between that line and the issuer block; together
    # they can leave too little room for Word's end-of-document marker even
    # when the whole printing table is visibly on page one.
    if contract.profile_id != "request":
        return
    issue_date = printing_table.getprevious()
    issuer = issue_date.getprevious() if issue_date is not None else None
    request_spacer = issuer.getprevious() if issuer is not None else None
    if (
        request_spacer is not None
        and request_spacer.tag.endswith("p")
        and not "".join(request_spacer.itertext()).strip()
        and not request_spacer.xpath(".//w:sectPr")
    ):
        request_spacer.getparent().remove(request_spacer)


def _replace_placeholders(document, replacements: Mapping[str, str]) -> tuple[str, ...]:
    replaced: set[str] = set()
    for token, value in replacements.items():
        count = 0
        if token == _OFFICIAL_BODY_TOKEN:
            count += _replace_official_body_paragraphs(document, token, value)
        for paragraph in _iter_all_paragraphs(document):
            if token not in get_full_text(paragraph):
                continue
            count += replace_run_text(paragraph, token, value)
        fixed_layout_result = replace_fixed_layout_placeholders(document, {token: value})
        count += fixed_layout_result.total_replacements
        if count:
            replaced.add(_placeholder_id(token))
    return tuple(sorted(replaced))


def _replace_official_body_paragraphs(document, token: str, value: str) -> int:
    """Expand AI body lines into real semantic Word paragraphs.

    A newline stored inside one run becomes ``w:br`` and makes every line
    inherit the ordinary-body font.  Splitting at assembly time preserves
    GB/T heading roles: level 1 Heiti, level 2 Kaiti, levels 3/4 Fangsong.
    """

    candidates = [
        paragraph
        for paragraph in list(document.paragraphs)
        if get_full_text(paragraph).strip() == token
    ]
    if not candidates:
        return 0

    raw_lines = re.split(r"\r\n?|\n", str(value or ""))
    lines = [line.strip() for line in raw_lines if line.strip()]
    if not lines:
        lines = [""]

    for placeholder_paragraph in candidates:
        for line in lines:
            level = _official_body_level(line)
            style_name = "Normal" if level == 0 else f"Heading {level}"
            paragraph = placeholder_paragraph.insert_paragraph_before(
                line,
                style=style_name,
            )
            _format_official_body_paragraph(paragraph, level=level)
        placeholder_paragraph._element.getparent().remove(
            placeholder_paragraph._element
        )
    return len(candidates)


def _stabilize_trailing_closing_block(document) -> None:
    """Keep an issuing block from becoming a page by itself.

    Word paginates the issuer and issue date independently from the final body
    paragraph by default.  On a nearly full page that leaves an otherwise
    empty trailing page containing only the two closing lines.  Bind the final
    semantic body unit, the optional spacer, issuer, and date as one short
    pagination group.  Unlike the old Heading-style default, this does not
    chain every numbered heading together or move the whole body block.
    """

    paragraphs = list(document.paragraphs)
    closing_pair: tuple[int, int] | None = None
    for index in range(len(paragraphs) - 1):
        issuer = paragraphs[index]
        issue_date = paragraphs[index + 1]
        if (
            issuer.alignment == WD_ALIGN_PARAGRAPH.RIGHT
            and issue_date.alignment == WD_ALIGN_PARAGRAPH.RIGHT
            and get_full_text(issuer).strip()
            and get_full_text(issue_date).strip()
        ):
            closing_pair = (index, index + 1)
    if closing_pair is None:
        return

    issuer_index, date_index = closing_pair
    chain_indices = [issuer_index]
    body_index = issuer_index - 1
    if body_index >= 0 and not get_full_text(paragraphs[body_index]).strip():
        chain_indices.insert(0, body_index)
        body_index -= 1
    if body_index >= 0 and get_full_text(paragraphs[body_index]).strip():
        chain_indices.insert(0, body_index)
        previous_index = body_index - 1
        if previous_index >= 0:
            previous_style = str(
                getattr(paragraphs[previous_index].style, "name", "") or ""
            )
            if previous_style.startswith("Heading "):
                chain_indices.insert(0, previous_index)

    for index in chain_indices:
        paragraphs[index].paragraph_format.keep_with_next = True
    paragraphs[issuer_index].paragraph_format.keep_together = True
    paragraphs[date_index].paragraph_format.keep_together = True


def _stabilize_heading_blocks(document) -> None:
    """Keep every numbered heading with at least its following paragraph."""

    for paragraph in document.paragraphs:
        style_name = str(getattr(paragraph.style, "name", "") or "")
        if not style_name.startswith("Heading "):
            continue
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.page_break_before = False


def _fit_official_redhead_marks(document) -> None:
    """Fit long red issuing-agency marks to at most two balanced lines."""

    for paragraph in document.paragraphs:
        colored_runs = [
            run
            for run in paragraph.runs
            if run.font.color.rgb is not None
            and str(run.font.color.rgb).upper() == "C00000"
            and (run.font.size is None or run.font.size.pt >= 30)
        ]
        if not colored_runs:
            continue
        text = "".join(run.text for run in colored_runs).strip()
        if not text:
            continue
        weighted_length = sum(
            0.55 if ord(character) < 128 else 1.0
            for character in text
            if not character.isspace()
        )
        if weighted_length <= 0:
            continue
        base_size = max(
            (float(run.font.size.pt) for run in colored_runs if run.font.size),
            default=50.0,
        )
        available_width_pt = 482.0 if base_size <= 40 else 442.0
        minimum_size = 28.0 if base_size <= 40 else 32.0
        target_lines = 1.0 if weighted_length <= 12.0 else 2.0
        fitted_size = max(
            minimum_size,
            min(
                base_size,
                available_width_pt * target_lines / weighted_length,
            ),
        )
        if fitted_size >= base_size - 0.1:
            continue
        for run in colored_runs:
            run.font.size = Pt(fitted_size)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        paragraph.paragraph_format.line_spacing = Pt(fitted_size * 1.2)


def _normalize_terminal_table_separators(document) -> None:
    """Remove an internal imprint rule when only one imprint line remains."""

    for table in document.tables:
        if len(table.rows) != 1 or len(table.columns) != 1:
            continue
        if table._tbl.tblPr.find(qn("w:tblpPr")) is None:
            continue
        cell = table.rows[0].cells[0]
        content_paragraphs = [
            paragraph
            for paragraph in cell.paragraphs
            if get_full_text(paragraph).strip()
        ]
        if len(content_paragraphs) > 1:
            continue
        for paragraph in list(cell.paragraphs):
            if get_full_text(paragraph).strip():
                continue
            p_pr = paragraph._p.get_or_add_pPr()
            borders = p_pr.find(qn("w:pBdr"))
            if borders is None or len(cell.paragraphs) <= 1:
                continue
            paragraph._p.getparent().remove(paragraph._p)


def _repair_official_display_line_heights(document) -> None:
    """Repair clipping without replacing a user template's selected font.

    User masters remain the typography authority. The only automatic change
    here is replacing an undersized exact line box on the organization mark or
    document title with an ``AT_LEAST`` line box. This preserves an explicit
    小标宋 family while allowing Word/WPS to honor that family's real metrics.
    """

    for paragraph in _iter_all_paragraphs(document):
        text = get_full_text(paragraph)
        if not any(token in text for token in _OFFICIAL_DISPLAY_TOKENS):
            continue
        font_size_pt = _effective_paragraph_font_size_pt(paragraph)
        if font_size_pt < 20:
            continue
        line_rule, line_height_pt = _effective_paragraph_line_height(paragraph)
        minimum_pt = font_size_pt * 1.2
        if (
            line_rule != WD_LINE_SPACING.EXACTLY
            or line_height_pt is None
            or line_height_pt >= minimum_pt
        ):
            continue
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        paragraph.paragraph_format.line_spacing = Pt(minimum_pt)


def _effective_paragraph_font_size_pt(paragraph) -> float:
    sizes = [
        float(run.font.size.pt)
        for run in paragraph.runs
        if run.font.size is not None
    ]
    style = getattr(paragraph, "style", None)
    while style is not None:
        if getattr(style.font, "size", None) is not None:
            sizes.append(float(style.font.size.pt))
            break
        style = getattr(style, "base_style", None)
    return max(sizes, default=0.0)


def _effective_paragraph_line_height(paragraph):
    containers = [paragraph]
    style = getattr(paragraph, "style", None)
    while style is not None:
        containers.append(style)
        style = getattr(style, "base_style", None)
    for container in containers:
        paragraph_format = container.paragraph_format
        rule = paragraph_format.line_spacing_rule
        spacing = paragraph_format.line_spacing
        if rule is None and spacing is None:
            continue
        height = getattr(spacing, "pt", None)
        return rule, float(height) if height is not None else None
    return None, None


def _official_body_level(text: str) -> int:
    normalized = str(text or "").strip()
    for level, pattern in _BODY_LEVEL_PATTERNS:
        if pattern.match(normalized):
            return level
    return 0


def _format_official_body_paragraph(paragraph, *, level: int) -> None:
    role = {
        1: "黑体",
        2: "楷体",
        3: "仿宋",
        4: "仿宋",
    }.get(level, "仿宋")
    paragraph.paragraph_format.first_line_indent = Cm(1.1)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(28)
    # Heading 1/2 are semantic font roles here, not navigation headings.
    # Explicitly cancel Word template pagination controls so a run of numbered
    # body levels can split naturally across pages.
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.page_break_before = False
    for run in paragraph.runs:
        set_run_fonts(
            run,
            font_cn=role,
            font_en=role,
            size_pt=16,
            bold=False,
        )
        run.font.color.rgb = RGBColor(0, 0, 0)


def _placeholder_id(token: str) -> str:
    return parse_material_token(str(token or "")).identifier


def _unresolved_official_placeholders(document) -> tuple[str, ...]:
    found: set[str] = set()
    for paragraph in _iter_all_paragraphs(document):
        found.update(_official_placeholder_ids(get_full_text(paragraph)))
    return tuple(sorted(found))


def _official_placeholder_ids(text: object) -> tuple[str, ...]:
    """Project official ``@text`` surfaces through the central token parser."""

    found: list[str] = []
    for match in MATERIAL_TOKEN_PATTERN.finditer(str(text or "")):
        try:
            ref = parse_material_token(match.group(0))
        except (TypeError, ValueError):
            continue
        if (
            ref.namespace is MaterialTokenNamespace.TEXT
            and ref.identifier.startswith("official_")
        ):
            found.append(ref.identifier)
    return tuple(found)


def _iter_all_paragraphs(container):
    for paragraph in list(getattr(container, "paragraphs", []) or []):
        yield paragraph
    for table in list(getattr(container, "tables", []) or []):
        for row in list(getattr(table, "rows", []) or []):
            for cell in list(getattr(row, "cells", []) or []):
                yield from _iter_all_paragraphs(cell)
    if hasattr(container, "sections"):
        for section in list(getattr(container, "sections", []) or []):
            for part in (
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            ):
                yield from _iter_all_paragraphs(part)


def _target_docx_path(
    output_dir: Path | str,
    filename: str | None,
    entity_data: Mapping[str, object],
    profile_id: str,
) -> Path:
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    target_name = str(filename or "").strip()
    if not target_name:
        title = str(entity_data.get("title", "") if isinstance(entity_data, Mapping) else "").strip()
        target_name = f"{_safe_stem(title or profile_id)}.docx"
    if not target_name.lower().endswith(".docx"):
        target_name = f"{target_name}.docx"
    return _available_output_path(target_dir / Path(target_name).name)


def _available_output_path(path: Path) -> Path:
    """Return a non-overwriting sibling path for one assembly artifact."""

    if not path.exists():
        return path
    for index in range(2, 1000):
        candidate = path.with_name(f"{path.stem}_{index}{path.suffix}")
        if not candidate.exists():
            return candidate
    return path.with_name(f"{path.stem}_1000{path.suffix}")


def _internal_review_docx_path(formal_docx_path: Path) -> Path:
    return formal_docx_path.with_name(f"{formal_docx_path.stem}_internal_review.docx")


def _review_pdf_path(formal_docx_path: Path) -> Path:
    return formal_docx_path.with_name(f"{formal_docx_path.stem}_review.pdf")


def _blocked_review_pdf_status(generate_review_pdf: bool) -> str:
    return "assembly_blocked" if generate_review_pdf else "not_requested"


def _render_optional_review_pdf(
    source_docx_path: Path,
    target_pdf_path: Path,
    *,
    renderer_resolver: DocxPdfRendererResolver | None = None,
) -> _ReviewPdfOutcome:
    resolver = renderer_resolver or _default_review_pdf_renderer_resolver
    try:
        renderer_entry = resolver()
    except Exception as exc:
        return _ReviewPdfOutcome(
            status="renderer_unavailable",
            issue=f"renderer_resolution_failed: {exc}",
        )
    if renderer_entry is None:
        return _ReviewPdfOutcome(
            status="renderer_unavailable",
            issue="docx_to_pdf_renderer_unavailable",
        )

    renderer_name, renderer = renderer_entry
    try:
        renderer(source_docx_path, target_pdf_path)
    except Exception as exc:
        _remove_partial_review_pdf(target_pdf_path)
        return _ReviewPdfOutcome(
            status="render_failed",
            renderer=str(renderer_name or ""),
            issue=f"review_pdf_render_failed: {exc}",
        )
    if not target_pdf_path.is_file() or target_pdf_path.stat().st_size <= 0:
        _remove_partial_review_pdf(target_pdf_path)
        return _ReviewPdfOutcome(
            status="output_missing",
            renderer=str(renderer_name or ""),
            issue="review_pdf_missing_after_render",
        )
    return _ReviewPdfOutcome(
        status="generated",
        path=target_pdf_path,
        renderer=str(renderer_name or ""),
    )


def _default_review_pdf_renderer_resolver() -> DocxPdfRendererEntry | None:
    from src.shared.engine.docx_page_renderer import (
        available_docx_pdf_renderer,
    )

    return available_docx_pdf_renderer()


def _remove_partial_review_pdf(path: Path) -> None:
    try:
        if path.exists():
            path.unlink()
    except OSError:
        return


def _archive_manifest_path(formal_docx_path: Path) -> Path:
    return formal_docx_path.with_name(f"{formal_docx_path.stem}_archive_manifest.json")


def _archive_manifest_markdown_path(formal_docx_path: Path) -> Path:
    return formal_docx_path.with_name(f"{formal_docx_path.stem}_archive_manifest.md")


def _safe_stem(value: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|\r\n\t]+", "_", str(value or "").strip())
    cleaned = re.sub(r"\s+", "_", cleaned).strip("._ ")
    return cleaned or "official_document"


def _insert_internal_review_notice(document) -> None:
    paragraph = document.add_paragraph(
        "内部审阅稿：仅供内部流转和格式核验，不代表正式签发、盖章或归档审批有效。"
    )
    body = document._body._element
    body.remove(paragraph._p)
    body.insert(0, paragraph._p)


def _write_archive_manifests(
    json_path: Path,
    markdown_path: Path,
    contract: OfficialDocumentAssemblyContract,
    master: MasterSpec,
    entity_data: Mapping[str, object],
    *,
    formal_docx_path: Path,
    internal_review_docx_path: Path,
    published_archive_manifest_path: Path,
    published_archive_manifest_markdown_path: Path,
    review_pdf_path: Path | None,
    review_pdf_status: str,
    review_pdf_renderer: str,
    review_pdf_issue: str,
    replaced_placeholders: tuple[str, ...],
    delivery_versions: tuple[str, ...],
) -> None:
    profile = get_official_document_profile(contract.profile_id)
    values = entity_data if isinstance(entity_data, Mapping) else {}
    output_paths: dict[str, str] = {}
    if "official_docx" in delivery_versions:
        output_paths["official_docx"] = str(formal_docx_path)
    if "internal_review_docx" in delivery_versions:
        output_paths["internal_review_docx"] = str(internal_review_docx_path)
    if "archive_manifest" in delivery_versions:
        output_paths["archive_manifest"] = str(
            published_archive_manifest_path or json_path
        )
        output_paths["archive_manifest_md"] = str(
            published_archive_manifest_markdown_path or markdown_path
        )
    if review_pdf_path is not None and review_pdf_status == "generated":
        output_paths["review_pdf"] = str(review_pdf_path)
    archive_fields = {
        key: str(values.get(key, "") or "")
        for key in (
            "title",
            "document_no",
            "organization",
            "issue_date",
            "issuer",
            "recipient",
            "security_level",
            "urgency",
            "copy_scope",
            "printing_org",
            "printing_date",
            "archive_status",
            "archive_no",
            "retention_period",
        )
    }
    payload = {
        "status": "ok",
        "profile_id": contract.profile_id,
        "profile_label": getattr(profile, "label", "") if profile is not None else "",
        "master_id": master.master_id,
        "material_schema_ids": list(contract.material_schema_ids),
        "delivery_versions": list(delivery_versions),
        "output_paths": output_paths,
        "review_pdf": {
            "status": review_pdf_status,
            "path": str(review_pdf_path or ""),
            "renderer": review_pdf_renderer,
            "issue": review_pdf_issue,
        },
        "archive_fields": archive_fields,
        "replaced_placeholders": list(replaced_placeholders),
        "boundaries": [
            *list(contract.boundaries),
            "system-generated archive manifest is evidence only",
            "does not certify official release validity",
            "does not verify seal legality",
            "does not replace archive-office approval",
        ],
    }
    json_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    markdown_path.write_text(_archive_manifest_markdown(payload), encoding="utf-8")


def _archive_manifest_markdown(payload: Mapping[str, object]) -> str:
    archive_fields = payload.get("archive_fields", {})
    if not isinstance(archive_fields, Mapping):
        archive_fields = {}
    output_paths = payload.get("output_paths", {})
    if not isinstance(output_paths, Mapping):
        output_paths = {}
    boundaries = [
        str(item or "").strip()
        for item in list(payload.get("boundaries", []) or [])
        if str(item or "").strip()
    ]
    lines = [
        "# 公文归档清单",
        "",
        f"- 文种: {payload.get('profile_label') or payload.get('profile_id') or '-'}",
        f"- 标题: {archive_fields.get('title') or '-'}",
        f"- 文号: {archive_fields.get('document_no') or '-'}",
        f"- 发文机关: {archive_fields.get('organization') or '-'}",
        f"- 成文日期: {archive_fields.get('issue_date') or '-'}",
        f"- 归档号: {archive_fields.get('archive_no') or '-'}",
        f"- 保管期限: {archive_fields.get('retention_period') or '-'}",
        "",
        "## 交付产物",
        "",
    ]
    for key, value in output_paths.items():
        lines.append(f"- {key}: `{value}`")
    if boundaries:
        lines.extend(["", "## 边界说明", ""])
        lines.extend(f"- {item}" for item in boundaries[:8])
    lines.append("")
    return "\n".join(lines)


def _remove_placeholder_reference_page(document) -> None:
    body = document._body._element
    heading_element = None
    for paragraph in document.paragraphs:
        if str(paragraph.text or "").strip() in _REFERENCE_PAGE_HEADINGS:
            heading_element = paragraph._p
            break
    if heading_element is None:
        return

    children = list(body)
    try:
        start_index = children.index(heading_element)
    except ValueError:
        return

    # ``Document.add_section(WD_SECTION.NEW_PAGE)`` stores the section break on
    # the empty paragraph immediately before the reference-page heading.  If
    # that paragraph survives, Word renders an empty trailing page even after
    # the reference-page content has been removed.
    if start_index > 0:
        preceding = children[start_index - 1]
        if (
            preceding.tag.endswith("p")
            and preceding.xpath(".//w:sectPr")
            and not "".join(preceding.itertext()).strip()
        ):
            body.remove(preceding)
            start_index -= 1

    for child in list(body)[start_index:]:
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


__all__ = [
    "OfficialDocumentAssemblyResult",
    "assemble_official_document_docx",
]
