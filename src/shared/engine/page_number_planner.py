"""Shared page-number planning for section boundaries and header/footer output."""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Iterable

from src.config.feature_configs import default_continuous_page_number_phases
from src.config.special_title_rules import (
    parse_special_title_selector,
    special_title_selector_label,
)
from src.shared.engine.section_semantics import canonicalize_section_type

if TYPE_CHECKING:
    from docx import Document
    from src.pipeline.context import PipelineContext


COVER_SECTION_TYPES = frozenset({"cover"})
PRE_NUMBERING_SECTION_TYPES = COVER_SECTION_TYPES
FRONT_MATTER_SECTION_TYPES = frozenset({"abstract_cn", "abstract_en", "toc"})
BACK_MATTER_SECTION_TYPES = frozenset(
    {"references", "errata", "appendix", "acknowledgment", "resume"}
)
NON_NUMBERED_HEADING_SECTION_TYPES = BACK_MATTER_SECTION_TYPES
ALL_NUMBERED_SECTION_TYPES = FRONT_MATTER_SECTION_TYPES | {"body"} | BACK_MATTER_SECTION_TYPES
BOUNDARY_SECTION_TYPES = PRE_NUMBERING_SECTION_TYPES | FRONT_MATTER_SECTION_TYPES | {"body"} | BACK_MATTER_SECTION_TYPES

_SELECTOR_SECTION_TYPES: dict[str, frozenset[str]] = {
    "cover": COVER_SECTION_TYPES,
    "pre_numbering": PRE_NUMBERING_SECTION_TYPES,
    "front_matter": FRONT_MATTER_SECTION_TYPES,
    "abstracts": frozenset({"abstract_cn", "abstract_en"}),
    "toc": frozenset({"toc"}),
    "body": frozenset({"body"}),
    "back_matter": BACK_MATTER_SECTION_TYPES,
    "references": frozenset({"references"}),
    "errata": frozenset({"errata"}),
    "appendix": frozenset({"appendix"}),
    "acknowledgment": frozenset({"acknowledgment"}),
    "resume": frozenset({"resume"}),
    "all_numbered_content": ALL_NUMBERED_SECTION_TYPES,
}


@dataclass(slots=True)
class PageNumberPhaseRule:
    phase_id: str
    selectors: tuple[str, ...]
    matched_section_types: frozenset[str]
    visible: bool
    number_format: str
    start_mode: str
    start_value: int


@dataclass(slots=True)
class LogicalSectionBoundary:
    start_index: int
    section_type: str
    phase_id: str | None
    reasons: tuple[str, ...] = ()

    @property
    def requires_section_break(self) -> bool:
        return self.start_index > 0


@dataclass(slots=True)
class PageNumberSectionPlan:
    section_index: int
    start_index: int
    section_type: str
    phase_id: str | None
    page_number_visible: bool
    number_format: str
    start_value: int | None
    header_visible: bool
    footer_text_visible: bool
    special_title_selector: str | None = None

    @property
    def hide_header_footer(self) -> bool:
        """Compatibility projection for callers not yet migrated."""

        return not self.header_visible and not self.footer_text_visible


@dataclass(slots=True)
class PageNumberExecutionPlan:
    boundaries: list[LogicalSectionBoundary] = field(default_factory=list)
    sections: list[PageNumberSectionPlan] = field(default_factory=list)
    diagnostics: list["PageNumberDiagnostic"] = field(default_factory=list)
    missing_doc_tree: bool = False


@dataclass(slots=True)
class PageNumberDiagnostic:
    level: str
    message: str
    suggestion: str = ""
    location: str = ""


_SECTION_TYPE_LABELS: dict[str, str] = {
    "front_matter": "前置部分",
    "body": "正文部分",
    "back_matter": "后置部分",
    "cover": "封面",
    "pre_numbering": "封面",
    "toc": "目录",
    "references": "参考文献",
    "appendix": "附录",
    "acknowledgment": "致谢",
    "resume": "简历",
    "errata": "勘误",
    "abstract_cn": "中文摘要",
    "abstract_en": "英文摘要",
}


def build_page_number_execution_plan(
    doc: Document,
    context: PipelineContext,
    header_footer,
) -> PageNumberExecutionPlan:
    rules, diagnostics = _resolve_phase_rules(header_footer)
    hidden_header_types = resolve_hidden_header_section_types(header_footer)
    hidden_footer_types = resolve_hidden_footer_section_types(header_footer)
    missing_doc_tree = getattr(context, "doc_tree", None) is None
    boundaries = _collect_logical_boundaries(doc, context, rules)
    section_starts = _section_start_indices(doc)

    sections: list[PageNumberSectionPlan] = []
    previous_phase_id: str | None = None
    for section_index, start_index in enumerate(section_starts):
        section_type = _resolve_section_type(context, start_index)
        scope_keys = _resolve_scope_keys(context, start_index)
        phase = _match_phase_rule(rules, scope_keys)
        phase_id = phase.phase_id if phase is not None else None
        header_visible = bool(getattr(header_footer, "header_enabled", True)) and (
            not _scope_is_hidden(scope_keys, hidden_header_types)
        )
        footer_text_visible = bool(getattr(header_footer, "footer_enabled", True)) and (
            not _scope_is_hidden(scope_keys, hidden_footer_types)
        )
        special_title_selector = next(
            (
                key
                for key in scope_keys
                if parse_special_title_selector(key) is not None
            ),
            None,
        )
        page_number_visible = bool(header_footer.page_number_enabled)
        number_format = "decimal"
        start_value: int | None = None

        if phase is not None:
            page_number_visible = page_number_visible and bool(phase.visible)
            number_format = phase.number_format
            if phase.phase_id != previous_phase_id and phase.start_mode == "restart":
                start_value = phase.start_value
        else:
            page_number_visible = False

        sections.append(
            PageNumberSectionPlan(
                section_index=section_index,
                start_index=start_index,
                section_type=section_type,
                phase_id=phase_id,
                page_number_visible=page_number_visible,
                number_format=number_format,
                start_value=start_value,
                header_visible=header_visible,
                footer_text_visible=footer_text_visible,
                special_title_selector=special_title_selector,
            )
        )
        previous_phase_id = phase_id

    return PageNumberExecutionPlan(
        boundaries=boundaries,
        sections=sections,
        diagnostics=diagnostics + _collect_missing_phase_diagnostics(context, rules),
        missing_doc_tree=missing_doc_tree,
    )


def _resolve_phase_rules(
    header_footer,
) -> tuple[list[PageNumberPhaseRule], list[PageNumberDiagnostic]]:
    phases = list(getattr(getattr(header_footer, "page_number_plan", None), "phases", []) or [])
    if not phases:
        phases = default_continuous_page_number_phases()

    diagnostics: list[PageNumberDiagnostic] = []
    rules: list[PageNumberPhaseRule] = []
    claimed_by_section_type: dict[str, list[str]] = {}
    claimed_phase_ids: dict[str, int] = {}

    for index, phase in enumerate(phases):
        selectors = tuple(
            str(selector or "").strip()
            for selector in (getattr(phase, "selectors", None) or [])
            if str(selector or "").strip()
        )
        phase_id = str(getattr(phase, "phase_id", "") or "").strip() or f"phase_{index + 1}"
        claimed_phase_ids[phase_id] = claimed_phase_ids.get(phase_id, 0) + 1
        matched = expand_page_number_selectors(selectors)
        rule = PageNumberPhaseRule(
            phase_id=phase_id,
            selectors=selectors,
            matched_section_types=matched,
            visible=bool(getattr(phase, "visible", True)),
            number_format=_normalize_number_format(getattr(phase, "number_format", "decimal")),
            start_mode=_normalize_start_mode(getattr(phase, "start_mode", "continue")),
            start_value=max(1, int(getattr(phase, "start_value", 1) or 1)),
        )
        rules.append(rule)
        for section_type in matched:
            claimed_by_section_type.setdefault(section_type, []).append(phase_id)

    for section_type, phase_ids in claimed_by_section_type.items():
        if len(phase_ids) > 1:
            diagnostics.append(
                _build_phase_overlap_diagnostic(
                    section_type=section_type,
                    phase_ids=phase_ids,
                    level="error",
                    location="header_footer.page_number_plan",
                )
            )

    for phase_id, count in claimed_phase_ids.items():
        if count <= 1:
            continue
        diagnostics.append(
            PageNumberDiagnostic(
                level="error",
                message=f"编号分组名称“{phase_id}”重复。",
                suggestion="为每个编号分组使用不同名称。",
                location="header_footer.page_number_plan",
            )
        )

    return rules, diagnostics


def expand_page_number_selectors(selectors: Iterable[str]) -> frozenset[str]:
    matched: set[str] = set()
    for selector in selectors:
        canonical = canonicalize_section_type(selector)
        if canonical in _SELECTOR_SECTION_TYPES:
            matched.update(_SELECTOR_SECTION_TYPES[canonical])
        elif selector in _SELECTOR_SECTION_TYPES:
            matched.update(_SELECTOR_SECTION_TYPES[selector])
        elif canonical:
            matched.add(canonical)
    return frozenset(matched)


def _resolve_hidden_channel_section_types(
    channel_config,
) -> frozenset[str]:
    selectors_raw = getattr(channel_config, "hidden_selectors", None)
    selectors = tuple(
        str(selector or "").strip()
        for selector in (selectors_raw or ())
        if str(selector or "").strip()
    )
    return expand_page_number_selectors(selectors)


def resolve_hidden_header_section_types(header_footer) -> frozenset[str]:
    return _resolve_hidden_channel_section_types(
        getattr(header_footer, "header", None),
    )


def resolve_hidden_footer_section_types(header_footer) -> frozenset[str]:
    return _resolve_hidden_channel_section_types(
        getattr(header_footer, "footer", None),
    )


def resolve_suppressed_header_footer_section_types(header_footer) -> frozenset[str]:
    """Legacy common-scope projection; new code should use channel resolvers."""

    return (
        resolve_hidden_header_section_types(header_footer)
        & resolve_hidden_footer_section_types(header_footer)
    )


def collect_page_number_diagnostics(
    doc: Document,
    context: PipelineContext,
    header_footer,
) -> list[PageNumberDiagnostic]:
    page_number_enabled = bool(getattr(header_footer, "page_number_enabled", True))
    has_scoped_output = bool(
        resolve_hidden_header_section_types(header_footer)
        or resolve_hidden_footer_section_types(header_footer)
    )
    if not page_number_enabled and not has_scoped_output:
        return []

    plan = build_page_number_execution_plan(doc, context, header_footer)
    plan_cfg = getattr(header_footer, "page_number_plan", None)
    validation_mode = str(getattr(plan_cfg, "validation_mode", "strict") or "strict")
    on_missing_doc_tree = str(
        getattr(plan_cfg, "on_missing_doc_tree", "warn_and_fallback") or "warn_and_fallback"
    )

    diagnostics: list[PageNumberDiagnostic] = []
    if plan.missing_doc_tree and on_missing_doc_tree == "warn_and_fallback":
        message = "未识别到文档结构，页眉、页脚与页码范围将按默认正文结构处理。"
        suggestion = "先运行标题/结构识别，可获得更准确的页码范围。"
        if not page_number_enabled and has_scoped_output:
            message = "未识别到文档结构，页眉和页脚文字范围将按默认正文结构处理。"
            suggestion = "先运行标题/结构识别，或清空对应通道的隐藏范围。"
        diagnostics.append(
            PageNumberDiagnostic(
                level="warning",
                message=message,
                suggestion=suggestion,
                location="context.doc_tree",
            )
        )

    if not page_number_enabled:
        return diagnostics

    diagnostic_level = "error" if validation_mode == "strict" else "warning"
    for item in plan.diagnostics:
        diagnostics.append(
            PageNumberDiagnostic(
                level="error" if item.level == "error" else diagnostic_level,
                message=item.message,
                suggestion=item.suggestion,
                location=item.location or "header_footer.page_number_plan",
            )
        )

    return diagnostics


def collect_static_page_number_diagnostics(header_footer) -> list[PageNumberDiagnostic]:
    if not bool(getattr(header_footer, "page_number_enabled", True)):
        return []

    diagnostics: list[PageNumberDiagnostic] = []
    validation_mode = str(
        getattr(getattr(header_footer, "page_number_plan", None), "validation_mode", "strict") or "strict"
    )
    level = "error" if validation_mode == "strict" else "warning"
    location = "header_footer.page_number_plan"
    phases = list(getattr(getattr(header_footer, "page_number_plan", None), "phases", []) or [])
    if not phases:
        phases = default_continuous_page_number_phases()

    for index, phase in enumerate(phases, start=1):
        phase_id = str(getattr(phase, "phase_id", "") or "").strip() or f"phase_{index}"
        selectors = [
            str(selector or "").strip()
            for selector in (getattr(phase, "selectors", None) or [])
            if str(selector or "").strip()
        ]
        if not selectors:
            diagnostics.append(
                PageNumberDiagnostic(
                    level=level,
                    message=f"第 {index} 个编号分组尚未选择范围。",
                    suggestion="选择包含部分，或删除这个空分组。",
                    location=location,
                )
            )
            continue
        if not expand_page_number_selectors(selectors):
            diagnostics.append(
                PageNumberDiagnostic(
                    level=level,
                    message=f"编号分组“{phase_id}”没有匹配到已知分区。",
                    suggestion="检查范围名称，或确认结构识别会产出对应分区。",
                    location=location,
                )
            )

    _rules, overlap_messages = _resolve_phase_rules(header_footer)
    diagnostics.extend(
        PageNumberDiagnostic(
            level="error" if item.level == "error" else level,
            message=item.message,
            suggestion=item.suggestion,
            location=item.location or location,
        )
        for item in overlap_messages
    )
    return diagnostics


def _collect_logical_boundaries(
    doc: Document,
    context: PipelineContext,
    rules: list[PageNumberPhaseRule],
) -> list[LogicalSectionBoundary]:
    doc_tree = getattr(context, "doc_tree", None)
    if doc_tree is None:
        return []

    section_boundary_starts: set[int] = set()
    for section in getattr(doc_tree, "sections", []) or []:
        section_type = canonicalize_section_type(getattr(section, "section_type", ""))
        start_index = int(getattr(section, "start_index", -1))
        if section_type in BOUNDARY_SECTION_TYPES and start_index > 0:
            section_boundary_starts.add(start_index)

    special_range_starts: set[int] = set()
    special_range_ends: set[int] = set()
    total = len(doc.paragraphs)
    for section in getattr(doc_tree, "special_title_ranges", []) or []:
        start_index = int(getattr(section, "start_index", -1))
        end_index = int(getattr(section, "end_index", -1))
        if 0 < start_index < total:
            special_range_starts.add(start_index)
        if 0 < end_index < total:
            special_range_ends.add(end_index)

    start_points = sorted(
        section_boundary_starts | special_range_starts | special_range_ends
    )
    boundaries: list[LogicalSectionBoundary] = []
    previous_phase_id: str | None = _match_phase_id(
        rules,
        _resolve_scope_keys(context, 0),
    )

    for start_index in start_points:
        section_type = _resolve_section_type(context, start_index)
        phase_id = _match_phase_id(
            rules,
            _resolve_scope_keys(context, start_index),
        )
        reasons: list[str] = []
        if start_index in section_boundary_starts:
            reasons.append("section_start")
        if start_index in special_range_starts:
            reasons.append("special_title_start")
        if start_index in special_range_ends:
            reasons.append("special_title_end")
        if phase_id != previous_phase_id:
            reasons.append("phase_change")
        boundaries.append(
            LogicalSectionBoundary(
                start_index=start_index,
                section_type=section_type,
                phase_id=phase_id,
                reasons=tuple(reasons),
            )
        )
        previous_phase_id = phase_id

    return boundaries


def _collect_missing_phase_diagnostics(
    context: PipelineContext,
    rules: list[PageNumberPhaseRule],
) -> list[PageNumberDiagnostic]:
    doc_tree = getattr(context, "doc_tree", None)
    if doc_tree is None:
        return []

    diagnostics: list[PageNumberDiagnostic] = []
    seen_section_types: set[str] = set()
    for section in getattr(doc_tree, "sections", []) or []:
        section_type = canonicalize_section_type(getattr(section, "section_type", ""))
        if section_type in seen_section_types or section_type not in BOUNDARY_SECTION_TYPES:
            continue
        seen_section_types.add(section_type)
        if _match_phase_rule(rules, section_type) is None:
            diagnostics.append(
                _build_uncovered_section_diagnostic(
                    section_type=section_type,
                    level="warning",
                    location="header_footer.page_number_plan",
                )
            )

    return diagnostics


def format_page_number_diagnostic_text(diagnostic: PageNumberDiagnostic) -> str:
    message = str(diagnostic.message or "").strip()
    suggestion = str(diagnostic.suggestion or "").strip()
    if not suggestion:
        return message
    return f"{message}\n处理：{suggestion}"


def _section_type_label(section_type: str) -> str:
    return _SECTION_TYPE_LABELS.get(
        section_type,
        special_title_selector_label(section_type) or "未知分区",
    )


def _build_phase_overlap_diagnostic(
    *,
    section_type: str,
    phase_ids: list[str],
    level: str,
    location: str,
) -> PageNumberDiagnostic:
    section_label = _section_type_label(section_type)
    phase_text = "、".join(str(phase_id or "").strip() for phase_id in phase_ids if str(phase_id or "").strip())
    return PageNumberDiagnostic(
        level=level,
        message=f"“{section_label}”同时属于多个编号分组：{phase_text}。",
        suggestion="让同一部分只属于一个编号分组；如需续号，请保留一组并把起号方式改为“延续前段”。",
        location=location,
    )


def _build_uncovered_section_diagnostic(
    *,
    section_type: str,
    level: str,
    location: str,
) -> PageNumberDiagnostic:
    section_label = _section_type_label(section_type)
    return PageNumberDiagnostic(
        level=level,
        message=f"“{section_label}”没有对应编号分组。",
        suggestion="新增编号分组覆盖这一部分，或把它并入现有分组。",
        location=location,
    )


def _resolve_section_type(context: PipelineContext, para_index: int) -> str:
    doc_tree = getattr(context, "doc_tree", None)
    if doc_tree is None:
        return "body"

    getter = getattr(doc_tree, "get_section_for_paragraph", None)
    try:
        section_type = getter(para_index) if callable(getter) else "body"
    except Exception:
        section_type = "body"
    return canonicalize_section_type(section_type)


def _resolve_scope_keys(
    context: PipelineContext,
    para_index: int,
) -> frozenset[str]:
    keys = {_resolve_section_type(context, para_index)}
    doc_tree = getattr(context, "doc_tree", None)
    getter = getattr(
        doc_tree,
        "get_special_title_selectors_for_paragraph",
        None,
    )
    if callable(getter):
        try:
            keys.update(
                str(selector or "").strip()
                for selector in getter(para_index)
                if str(selector or "").strip()
            )
        except Exception:
            pass
    return frozenset(keys)


def _section_start_indices(doc: Document) -> list[int]:
    from src.shared.engine.ooxml_ops import qn

    starts = [0]
    total = len(doc.paragraphs)
    for index, para in enumerate(doc.paragraphs):
        if index + 1 >= total:
            continue
        p_pr = para._element.find(qn("w:pPr"))
        if p_pr is None or p_pr.find(qn("w:sectPr")) is None:
            continue
        starts.append(index + 1)

    deduped: list[int] = []
    for start in starts:
        if start not in deduped:
            deduped.append(start)
    return deduped


def _match_phase_rule(
    rules: list[PageNumberPhaseRule],
    scope_keys: str | Iterable[str],
) -> PageNumberPhaseRule | None:
    keys = (
        {scope_keys}
        if isinstance(scope_keys, str)
        else {
            str(scope_key or "").strip()
            for scope_key in scope_keys
            if str(scope_key or "").strip()
        }
    )
    dynamic_keys = {
        key
        for key in keys
        if parse_special_title_selector(key) is not None
    }
    if dynamic_keys:
        for rule in rules:
            if dynamic_keys.intersection(rule.matched_section_types):
                return rule
    for rule in rules:
        if keys.intersection(rule.matched_section_types):
            return rule
    return None


def _scope_is_hidden(
    scope_keys: Iterable[str],
    hidden_scope_keys: Iterable[str],
) -> bool:
    """Give a matched user rule precedence over its underlying base range."""

    keys = {
        str(scope_key or "").strip()
        for scope_key in scope_keys
        if str(scope_key or "").strip()
    }
    hidden = {
        str(scope_key or "").strip()
        for scope_key in hidden_scope_keys
        if str(scope_key or "").strip()
    }
    dynamic_keys = {
        key
        for key in keys
        if parse_special_title_selector(key) is not None
    }
    if dynamic_keys:
        return bool(dynamic_keys.intersection(hidden))
    return bool(keys.intersection(hidden))


def _match_phase_id(
    rules: list[PageNumberPhaseRule],
    scope_keys: str | Iterable[str],
) -> str | None:
    rule = _match_phase_rule(rules, scope_keys)
    return rule.phase_id if rule is not None else None


def _normalize_number_format(value) -> str:
    raw = str(value or "decimal")
    if raw in {"decimal", "upperRoman", "lowerRoman"}:
        return raw
    return "decimal"


def _normalize_start_mode(value) -> str:
    raw = str(value or "continue")
    if raw in {"continue", "restart"}:
        return raw
    return "continue"
