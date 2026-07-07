"""Material field consistency inspection for scene-specific source facts."""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Mapping

from docx.document import Document as DocxDocument

from src.config.material_schema_registry import get_material_schema
from src.shared.engine.fixed_layout_text import iter_fixed_layout_text_blocks


CONTRACT_FIELD_KEYS: tuple[str, ...] = (
    "party_a",
    "party_b",
    "contract_amount",
    "signing_date",
    "contract_no",
)

FIELD_LABELS: dict[str, tuple[str, ...]] = {
    "party_a": ("Party A", "甲方"),
    "party_b": ("Party B", "乙方"),
    "contract_amount": ("Contract amount", "合同金额", "金额", "价款"),
    "signing_date": ("Signing date", "签署日期", "签约日期", "日期"),
    "contract_no": ("Contract number", "合同编号", "编号"),
}


@dataclass(slots=True, frozen=True)
class MaterialFieldConsistencyIssue:
    field_key: str
    kind: str
    message: str
    expected: str = ""
    observed: str = ""
    location: str = ""
    severity: str = "warning"


@dataclass(slots=True, frozen=True)
class MaterialFieldConsistencyItem:
    field_key: str
    expected: str
    occurrence_count: int = 0
    placeholders_remaining: tuple[str, ...] = ()
    issues: tuple[MaterialFieldConsistencyIssue, ...] = ()

    @property
    def status(self) -> str:
        return "warning" if self.issues else "ok"


@dataclass(slots=True, frozen=True)
class MaterialFieldConsistencyResult:
    schema_id: str = ""
    family_id: str = ""
    status: str = "not_applicable"
    items: tuple[MaterialFieldConsistencyItem, ...] = ()

    @property
    def issues(self) -> tuple[MaterialFieldConsistencyIssue, ...]:
        return tuple(issue for item in self.items for issue in item.issues)

    @property
    def has_issues(self) -> bool:
        return bool(self.issues)


def inspect_material_field_consistency(
    doc: DocxDocument,
    *,
    schema_id: str,
    entity_data: Mapping[str, str] | None,
) -> MaterialFieldConsistencyResult:
    """Inspect whether material field values are consistently present in DOCX text.

    The first closed slice is intentionally narrow: ``contract_parties_v1``.
    Other schemas remain not-applicable until their own field semantics are defined.
    """

    normalized_schema_id = str(schema_id or "").strip()
    if normalized_schema_id != "contract_parties_v1":
        return MaterialFieldConsistencyResult(schema_id=normalized_schema_id)

    try:
        schema = get_material_schema(normalized_schema_id)
    except KeyError:
        return MaterialFieldConsistencyResult(schema_id=normalized_schema_id)

    fields = {
        str(key): str(value).strip()
        for key, value in dict(entity_data or {}).items()
        if str(key or "").strip() and str(value or "").strip()
    }
    target_keys = [
        field.key
        for field in schema.fields
        if field.key in CONTRACT_FIELD_KEYS and field.key in fields
    ]
    if not target_keys:
        return MaterialFieldConsistencyResult(
            schema_id=schema.schema_id,
            family_id=schema.family,
        )

    blocks = _document_text_blocks(doc)
    full_text = "\n".join(blocks)
    items = tuple(
        _inspect_field(
            field_key=field_key,
            expected=fields[field_key],
            blocks=blocks,
            full_text=full_text,
        )
        for field_key in target_keys
    )
    return MaterialFieldConsistencyResult(
        schema_id=schema.schema_id,
        family_id=schema.family,
        status="warning" if any(item.issues for item in items) else "ok",
        items=items,
    )


def _inspect_field(
    *,
    field_key: str,
    expected: str,
    blocks: list[str],
    full_text: str,
) -> MaterialFieldConsistencyItem:
    placeholders = tuple(
        placeholder
        for placeholder in (f"{{{{{field_key}}}}}", "${" + field_key + "}")
        if placeholder in full_text
    )
    occurrence_count = _count_compact_occurrences(full_text, expected)
    issues: list[MaterialFieldConsistencyIssue] = []

    if placeholders:
        issues.append(
            MaterialFieldConsistencyIssue(
                field_key=field_key,
                kind="unresolved_placeholder",
                expected=expected,
                observed=", ".join(placeholders),
                message=f"Unresolved placeholder remains for {field_key}.",
            )
        )

    if occurrence_count == 0:
        issues.append(
            MaterialFieldConsistencyIssue(
                field_key=field_key,
                kind="expected_value_missing",
                expected=expected,
                message=f"Expected value for {field_key} was not found in the document.",
            )
        )

    conflict = _first_label_conflict(field_key, expected, blocks)
    if conflict:
        location, observed = conflict
        issues.append(
            MaterialFieldConsistencyIssue(
                field_key=field_key,
                kind="label_value_conflict",
                expected=expected,
                observed=observed,
                location=location,
                message=f"{field_key} label appears with a value different from material data.",
            )
        )

    return MaterialFieldConsistencyItem(
        field_key=field_key,
        expected=expected,
        occurrence_count=occurrence_count,
        placeholders_remaining=placeholders,
        issues=tuple(issues),
    )


def _document_text_blocks(doc: DocxDocument) -> list[str]:
    blocks: list[str] = []
    for paragraph in _iter_paragraphs(doc):
        text = str(getattr(paragraph, "text", "") or "").strip()
        if text:
            blocks.append(text)
    for block in iter_fixed_layout_text_blocks(doc):
        if block.text:
            blocks.append(block.text)
    return blocks


def _iter_paragraphs(container):
    for paragraph in list(getattr(container, "paragraphs", []) or []):
        yield paragraph
    for table in list(getattr(container, "tables", []) or []):
        for row in list(getattr(table, "rows", []) or []):
            for cell in list(getattr(row, "cells", []) or []):
                yield from _iter_paragraphs(cell)


def _count_compact_occurrences(text: str, value: str) -> int:
    compact_text = _compact_for_compare(text)
    compact_value = _compact_for_compare(value)
    if not compact_text or not compact_value:
        return 0
    return compact_text.count(compact_value)


def _first_label_conflict(
    field_key: str,
    expected: str,
    blocks: list[str],
) -> tuple[str, str] | None:
    labels = FIELD_LABELS.get(field_key, ())
    if not labels:
        return None
    compact_expected = _compact_for_compare(expected)
    expected_digits = _digits_only(expected)

    for index, block in enumerate(blocks):
        if not any(label.casefold() in block.casefold() for label in labels):
            continue
        if compact_expected and compact_expected in _compact_for_compare(block):
            continue
        if expected_digits and expected_digits in _digits_only(block):
            continue
        if any(placeholder in block for placeholder in (f"{{{{{field_key}}}}}", "${" + field_key + "}")):
            continue
        observed = _value_after_separator(block)
        if observed:
            return f"paragraph:{index + 1}", observed
    return None


def _value_after_separator(text: str) -> str:
    parts = re.split(r"[:：]", text, maxsplit=1)
    if len(parts) != 2:
        return ""
    return parts[1].strip()


def _compact_for_compare(value: str) -> str:
    return re.sub(r"[\s,，。.:：;；/\\\-年月日￥¥$()（）]+", "", str(value or "").casefold())


def _digits_only(value: str) -> str:
    return re.sub(r"\D+", "", str(value or ""))


__all__ = [
    "MaterialFieldConsistencyIssue",
    "MaterialFieldConsistencyItem",
    "MaterialFieldConsistencyResult",
    "inspect_material_field_consistency",
]
