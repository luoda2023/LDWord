"""Sample verification artifacts for official-document masters."""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Mapping

from docx import Document
from PIL import Image, ImageChops

from src.config.master_library import MasterSpec
from src.config.official_document_profiles import (
    get_official_document_assembly_contract,
)
from src.shared.engine.docx_page_renderer import render_docx_pages
from src.shared.engine.official_document_assembly import (
    OfficialDocumentAssemblyResult,
    assemble_official_document_docx,
)
from src.shared.engine.official_document_sample_data import (
    official_sample_entity_data,
)


_REFERENCE_PAGE_HEADINGS = ("版式占位符说明", "母版占位符说明")


@dataclass(frozen=True, slots=True)
class OfficialDocumentSampleVerificationResult:
    """Generated official sample and optional PDF/PNG rendering evidence."""

    profile_id: str
    sample_docx_path: Path | None
    manifest_path: Path
    status: str
    assembly_status: str = ""
    master_id: str = ""
    material_schema_ids: tuple[str, ...] = ()
    replaced_placeholders: tuple[str, ...] = ()
    missing_required_fields: tuple[str, ...] = ()
    unresolved_placeholders: tuple[str, ...] = ()
    pdf_path: Path | None = None
    png_paths: tuple[Path, ...] = ()
    renderer: str = ""
    output_paths: Mapping[str, str] = field(default_factory=dict)
    baseline_png_path: Path | None = None
    baseline_status: str = ""
    baseline_difference_bbox: tuple[int, int, int, int] | tuple[()] = ()
    issues: tuple[str, ...] = ()

    @property
    def ok(self) -> bool:
        return self.status in {
            "docx_verified",
            "renderer_unavailable",
            "visual_png_ok",
            "visual_baseline_ok",
        }

    def to_dict(self) -> dict[str, object]:
        return {
            "profile_id": self.profile_id,
            "sample_docx_path": str(self.sample_docx_path or ""),
            "manifest_path": str(self.manifest_path),
            "status": self.status,
            "assembly_status": self.assembly_status,
            "master_id": self.master_id,
            "material_schema_ids": list(self.material_schema_ids),
            "replaced_placeholders": list(self.replaced_placeholders),
            "missing_required_fields": list(self.missing_required_fields),
            "unresolved_placeholders": list(self.unresolved_placeholders),
            "pdf_path": str(self.pdf_path or ""),
            "png_paths": [str(path) for path in self.png_paths],
            "renderer": self.renderer,
            "output_paths": dict(self.output_paths),
            "baseline_png_path": str(self.baseline_png_path or ""),
            "baseline_status": self.baseline_status,
            "baseline_difference_bbox": list(self.baseline_difference_bbox),
            "issues": list(self.issues),
        }


def verify_official_document_sample_visual(
    profile_id: str,
    output_dir: Path | str,
    *,
    entity_data: Mapping[str, object] | None = None,
    filename: str | None = None,
    master: MasterSpec | None = None,
    attempt_render: bool = True,
    baseline_dir: Path | str | None = None,
) -> OfficialDocumentSampleVerificationResult:
    """Generate and verify one official-document sample DOCX."""

    normalized_profile_id = str(profile_id or "notice").strip() or "notice"
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = target_dir / f"{_safe_id(normalized_profile_id)}_visual_manifest.json"
    sample_data = official_sample_entity_data(normalized_profile_id)
    if entity_data is not None:
        sample_data.update(dict(entity_data))
    sample_data.setdefault("document_type", normalized_profile_id)

    assembly = assemble_official_document_docx(
        normalized_profile_id,
        sample_data,
        target_dir,
        filename=filename or f"{normalized_profile_id}_official_sample.docx",
        master=master,
    )
    if not assembly.ok or assembly.docx_path is None:
        result = _result_from_assembly(
            normalized_profile_id,
            manifest_path,
            assembly,
            status=assembly.status,
            issues=(f"assembly_{assembly.status}",),
        )
        _write_manifest(result)
        return result

    structure_issues = _docx_structure_issues(
        assembly.docx_path,
        expected_values=_expected_visible_sample_values(
            normalized_profile_id,
            sample_data,
        ),
        unresolved=assembly.unresolved_placeholders,
    )
    if structure_issues:
        result = _result_from_assembly(
            normalized_profile_id,
            manifest_path,
            assembly,
            status="docx_verification_failed",
            issues=structure_issues,
        )
        _write_manifest(result)
        return result

    rendered = render_docx_pages(
        assembly.docx_path,
        target_dir,
        attempt_render=attempt_render,
    )
    if rendered.status == "docx_only":
        result = _result_from_assembly(
            normalized_profile_id,
            manifest_path,
            assembly,
            status="docx_verified",
            renderer="skipped",
            baseline_status="not_evaluated" if baseline_dir is not None else "",
            baseline_png_path=(
                _baseline_png_path(normalized_profile_id, Path(baseline_dir))
                if baseline_dir is not None
                else None
            ),
        )
        _write_manifest(result)
        return result

    if not rendered.ready and rendered.status != "png_blank":
        issues = list(rendered.issues)
        if baseline_dir is not None:
            issues.append("visual_baseline_not_evaluated")
        result = _result_from_assembly(
            normalized_profile_id,
            manifest_path,
            assembly,
            status=rendered.status,
            pdf_path=rendered.pdf_path,
            renderer=rendered.renderer,
            issues=tuple(issues),
            baseline_status="not_evaluated" if baseline_dir is not None else "",
            baseline_png_path=(
                _baseline_png_path(normalized_profile_id, Path(baseline_dir))
                if baseline_dir is not None
                else None
            ),
        )
        _write_manifest(result)
        return result

    png_paths = rendered.page_paths
    issues = list(rendered.issues)
    status = "png_blank" if rendered.status == "png_blank" else "visual_png_ok"

    baseline_png_path = None
    baseline_status = ""
    baseline_difference_bbox: tuple[int, int, int, int] | tuple[()] = ()
    if baseline_dir is not None and png_paths:
        comparison = _compare_visual_baseline(
            normalized_profile_id,
            png_paths[0],
            Path(baseline_dir),
        )
        baseline_png_path = comparison.baseline_png_path
        baseline_status = comparison.status
        baseline_difference_bbox = comparison.difference_bbox
        issues.extend(comparison.issues)
        if status == "visual_png_ok":
            if comparison.status == "baseline_ok":
                status = "visual_baseline_ok"
            elif comparison.status == "baseline_missing":
                status = "visual_baseline_missing"
            elif comparison.status == "baseline_mismatch":
                status = "visual_baseline_mismatch"

    result = _result_from_assembly(
        normalized_profile_id,
        manifest_path,
        assembly,
        status=status,
        pdf_path=rendered.pdf_path,
        png_paths=png_paths,
        renderer=rendered.renderer,
        baseline_png_path=baseline_png_path,
        baseline_status=baseline_status,
        baseline_difference_bbox=baseline_difference_bbox,
        issues=tuple(issues),
    )
    _write_manifest(result)
    return result


def _result_from_assembly(
    profile_id: str,
    manifest_path: Path,
    assembly: OfficialDocumentAssemblyResult,
    *,
    status: str,
    pdf_path: Path | None = None,
    png_paths: tuple[Path, ...] = (),
    renderer: str = "",
    baseline_png_path: Path | None = None,
    baseline_status: str = "",
    baseline_difference_bbox: tuple[int, int, int, int] | tuple[()] = (),
    issues: tuple[str, ...] = (),
) -> OfficialDocumentSampleVerificationResult:
    return OfficialDocumentSampleVerificationResult(
        profile_id=profile_id,
        sample_docx_path=assembly.docx_path,
        manifest_path=manifest_path,
        status=status,
        assembly_status=assembly.status,
        master_id=assembly.master_id,
        material_schema_ids=assembly.material_schema_ids,
        replaced_placeholders=assembly.replaced_placeholders,
        missing_required_fields=assembly.missing_required_fields,
        unresolved_placeholders=assembly.unresolved_placeholders,
        pdf_path=pdf_path,
        png_paths=png_paths,
        renderer=renderer,
        output_paths=assembly.output_paths,
        baseline_png_path=baseline_png_path,
        baseline_status=baseline_status,
        baseline_difference_bbox=baseline_difference_bbox,
        issues=issues,
    )


@dataclass(frozen=True, slots=True)
class _VisualBaselineComparison:
    status: str
    baseline_png_path: Path
    difference_bbox: tuple[int, int, int, int] | tuple[()] = ()
    issues: tuple[str, ...] = ()


def _compare_visual_baseline(
    profile_id: str,
    actual_png_path: Path,
    baseline_dir: Path,
) -> _VisualBaselineComparison:
    baseline_png_path = _baseline_png_path(profile_id, baseline_dir)
    if not baseline_png_path.is_file():
        return _VisualBaselineComparison(
            "baseline_missing",
            baseline_png_path,
            issues=(f"visual_baseline_missing: {baseline_png_path}",),
        )

    with Image.open(actual_png_path) as actual_image:
        actual = actual_image.convert("RGB")
    with Image.open(baseline_png_path) as baseline_image:
        baseline = baseline_image.convert("RGB")

    if actual.size != baseline.size:
        return _VisualBaselineComparison(
            "baseline_mismatch",
            baseline_png_path,
            issues=(
                "visual_baseline_size_mismatch: "
                f"actual={actual.size[0]}x{actual.size[1]}; "
                f"baseline={baseline.size[0]}x{baseline.size[1]}",
            ),
        )

    bbox = ImageChops.difference(actual, baseline).getbbox()
    if bbox is None:
        return _VisualBaselineComparison("baseline_ok", baseline_png_path)
    return _VisualBaselineComparison(
        "baseline_mismatch",
        baseline_png_path,
        difference_bbox=tuple(int(value) for value in bbox),
        issues=(f"visual_baseline_mismatch: {baseline_png_path}",),
    )


def _baseline_png_path(profile_id: str, baseline_dir: Path) -> Path:
    return baseline_dir / f"{_safe_id(profile_id)}_official_sample_page-1.png"


def _docx_structure_issues(
    path: Path,
    *,
    expected_values: tuple[str, ...],
    unresolved: tuple[str, ...],
) -> tuple[str, ...]:
    text = _all_docx_text(path)
    issues: list[str] = []
    if unresolved:
        issues.append("unresolved_official_placeholders")
    if "{{official_" in text:
        issues.append("official_placeholder_text_residue")
    if any(heading in text for heading in _REFERENCE_PAGE_HEADINGS):
        issues.append("reference_page_not_removed")
    for value in expected_values:
        if value and value not in text:
            issues.append(f"missing_sample_text: {value}")
    return tuple(dict.fromkeys(issues))


def _expected_visible_sample_values(
    profile_id: str,
    sample_data: Mapping[str, object],
) -> tuple[str, ...]:
    keys = [
        "organization",
        "document_no",
        "title",
        "recipient",
        "body",
        "attachment_note",
        "issuer",
        "issue_date",
        "copy_scope",
        "printing_org",
        "printing_date",
        "security_level",
        "urgency",
        "signer",
    ]
    if str(profile_id or "").strip() == "minutes":
        keys.extend(("meeting_date", "participants"))
    contract = get_official_document_assembly_contract(profile_id)
    if contract is not None:
        applicable_fields = set(contract.applicable_material_field_keys)
        keys = [key for key in keys if key in applicable_fields]
    return tuple(
        str(sample_data.get(key, "") or "").strip()
        for key in keys
        if str(sample_data.get(key, "") or "").strip()
    )


def _all_docx_text(path: Path) -> str:
    document = Document(str(path))
    parts: list[str] = []

    def collect(container) -> None:
        parts.extend(paragraph.text for paragraph in getattr(container, "paragraphs", []) or [])
        for table in getattr(container, "tables", []) or []:
            for row in table.rows:
                for cell in row.cells:
                    collect(cell)
        if hasattr(container, "sections"):
            for section in container.sections:
                for part in (
                    section.header,
                    section.first_page_header,
                    section.even_page_header,
                    section.footer,
                    section.first_page_footer,
                    section.even_page_footer,
                ):
                    collect(part)

    collect(document)
    return "\n".join(parts)


def _write_manifest(result: OfficialDocumentSampleVerificationResult) -> None:
    result.manifest_path.write_text(
        json.dumps(result.to_dict(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _safe_id(value: str) -> str:
    normalized = str(value or "").strip() or "notice"
    return "".join(ch if ch.isalnum() or ch in {"_", "-"} else "_" for ch in normalized)


__all__ = [
    "OfficialDocumentSampleVerificationResult",
    "official_sample_entity_data",
    "verify_official_document_sample_visual",
]
