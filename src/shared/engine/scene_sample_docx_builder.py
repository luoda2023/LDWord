"""Build openable DOCX samples for scene sample fixture specs."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from src.config.scene_sample_fixture_registry import (
    SceneSampleFixtureSpec,
    list_scene_sample_fixtures,
)
from src.config.scene_request_cell_fixture_registry import (
    SceneRequestCellFixtureSpec,
    build_scene_request_cell_fixture_summary,
    list_scene_request_cell_fixtures,
)
from src.shared.engine.fixed_layout_tables import set_fixed_layout_row_height


DEFAULT_SCENE_SAMPLE_FIXTURE_DIR = Path("artifacts") / "scene_sample_fixtures"
DEFAULT_SCENE_SAMPLE_MANIFEST_NAME = "manifest.json"
DEFAULT_SCENE_REQUEST_CELL_REPORT_NAME = "request_cell_report.md"


@dataclass(frozen=True, slots=True)
class SceneSampleDocxArtifact:
    fixture_id: str
    pack_id: str
    path: str
    docx_surfaces: tuple[str, ...]
    expected_preflight_findings: tuple[str, ...]
    expected_behaviors: tuple[str, ...]
    report_expectations: tuple[str, ...]
    manual_gate_id: str = ""
    boundary_notes: tuple[str, ...] = ()
    request_sample_ids: tuple[str, ...] = ()

    def to_payload(self) -> dict[str, object]:
        return {
            "fixture_id": self.fixture_id,
            "pack_id": self.pack_id,
            "path": self.path,
            "docx_surfaces": list(self.docx_surfaces),
            "expected_preflight_findings": list(self.expected_preflight_findings),
            "expected_behaviors": list(self.expected_behaviors),
            "report_expectations": list(self.report_expectations),
            "manual_gate_id": self.manual_gate_id,
            "boundary_notes": list(self.boundary_notes),
            "request_sample_ids": list(self.request_sample_ids),
        }


@dataclass(frozen=True, slots=True)
class SceneSampleDocxLibrary:
    output_dir: str
    manifest_path: str
    request_cell_report_path: str
    artifacts: tuple[SceneSampleDocxArtifact, ...]
    request_cells: tuple[SceneRequestCellFixtureSpec, ...]
    request_cell_summary: dict[str, object]

    def to_payload(self) -> dict[str, object]:
        return {
            "output_dir": self.output_dir,
            "manifest_path": self.manifest_path,
            "request_cell_report_path": self.request_cell_report_path,
            "artifact_count": len(self.artifacts),
            "artifacts": [artifact.to_payload() for artifact in self.artifacts],
            "request_cell_count": len(self.request_cells),
            "request_cells": [
                _request_cell_payload(cell, self.request_cell_report_path)
                for cell in self.request_cells
            ],
            "request_cell_summary": self.request_cell_summary,
        }


def build_scene_sample_docx(
    output_dir: str | Path,
    fixture: SceneSampleFixtureSpec,
) -> Path:
    """Create one real DOCX package for a scene sample fixture."""

    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    source = target_dir / f"{fixture.fixture_id}.docx"
    doc = Document()
    doc.add_paragraph(fixture.label)
    if "table_grid" in fixture.docx_surfaces or "fixed_row_height" in fixture.docx_surfaces:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Name"
        table.cell(1, 1).text = "Example"
        if "fixed_row_height" in fixture.docx_surfaces:
            set_fixed_layout_row_height(table.rows[0], 20.0, rule="atLeast")
    doc.save(source)
    parts = _sample_parts_for_surfaces(fixture.docx_surfaces)
    if parts:
        with ZipFile(source, "a") as package:
            for name, data in parts.items():
                package.writestr(name, data)
    return source


def build_scene_sample_docx_library(
    output_dir: str | Path,
    *,
    fixtures: tuple[SceneSampleFixtureSpec, ...] | None = None,
    manifest_name: str = DEFAULT_SCENE_SAMPLE_MANIFEST_NAME,
) -> SceneSampleDocxLibrary:
    """Create all scene sample DOCX files and a JSON-compatible manifest."""

    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    selected = fixtures if fixtures is not None else list_scene_sample_fixtures()
    request_cells = list_scene_request_cell_fixtures()
    artifacts: list[SceneSampleDocxArtifact] = []
    for fixture in selected:
        path = build_scene_sample_docx(target_dir, fixture)
        artifacts.append(
            SceneSampleDocxArtifact(
                fixture_id=fixture.fixture_id,
                pack_id=fixture.pack_id,
                path=str(path),
                docx_surfaces=fixture.docx_surfaces,
                expected_preflight_findings=fixture.expected_preflight_findings,
                expected_behaviors=fixture.expected_behaviors,
                report_expectations=fixture.report_expectations,
                manual_gate_id=fixture.manual_gate_id,
                boundary_notes=fixture.boundary_notes,
                request_sample_ids=fixture.request_sample_ids,
            )
        )
    manifest_path = target_dir / manifest_name
    request_cell_report_path = target_dir / DEFAULT_SCENE_REQUEST_CELL_REPORT_NAME
    library = SceneSampleDocxLibrary(
        output_dir=str(target_dir),
        manifest_path=str(manifest_path),
        request_cell_report_path=str(request_cell_report_path),
        artifacts=tuple(artifacts),
        request_cells=request_cells,
        request_cell_summary=build_scene_request_cell_fixture_summary().to_payload(),
    )
    _write_request_cell_report(request_cell_report_path, request_cells)
    _write_json_payload(manifest_path, library.to_payload())
    return library


def scene_sample_fixture_manifest_paths(
    output_dir: str | Path = DEFAULT_SCENE_SAMPLE_FIXTURE_DIR,
) -> dict[str, str]:
    manifest_path = Path(output_dir) / DEFAULT_SCENE_SAMPLE_MANIFEST_NAME
    if not manifest_path.exists() or not manifest_path.is_file():
        return {}
    return {"fixture_manifest": str(manifest_path)}


def _sample_parts_for_surfaces(surfaces: tuple[str, ...]) -> dict[str, bytes]:
    parts: dict[str, bytes] = {}
    if "fields" in surfaces:
        parts["word/sample-fields.xml"] = _xml_body(
            b'<w:p><w:fldSimple w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>'
        )
    if "comments" in surfaces:
        parts["word/comments.xml"] = (
            b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            b'<w:comment w:id="0"><w:p><w:r><w:t>Review</w:t></w:r></w:p></w:comment>'
            b"</w:comments>"
        )
    if "tracked_changes" in surfaces:
        parts["word/sample-tracked-changes.xml"] = _xml_body(
            b'<w:p><w:ins><w:r><w:t>Inserted</w:t></w:r></w:ins></w:p>'
            b'<w:p><w:del><w:r><w:delText>Deleted</w:delText></w:r></w:del></w:p>'
        )
    if "hidden_text" in surfaces:
        parts["word/sample-hidden-text.xml"] = _xml_body(
            b'<w:p><w:r><w:rPr><w:vanish /></w:rPr><w:t>Hidden</w:t></w:r></w:p>'
        )
    if "textboxes" in surfaces:
        parts["word/sample-textbox.xml"] = _xml_body(
            b'<w:p><w:r><w:txbxContent><w:p><w:r><w:t>Box</w:t></w:r></w:p></w:txbxContent></w:r></w:p>'
        )
    if "content_controls" in surfaces:
        parts["word/sample-content-control.xml"] = _xml_body(
            b'<w:sdt><w:sdtContent><w:p><w:r><w:t>Name</w:t></w:r></w:p></w:sdtContent></w:sdt>'
        )
    if "ole_objects" in surfaces:
        parts["word/embeddings/oleObject1.bin"] = b"ole"
        parts["word/_rels/sample-ole.rels"] = _rels(
            b'<Relationship Id="rIdOle" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="../embeddings/oleObject1.bin"/>'
        )
    if "embedded_workbooks" in surfaces:
        parts["word/embeddings/workbook.xlsx"] = b"not-a-real-xlsx"
    if "visio_drawings" in surfaces:
        parts["word/embeddings/diagram.vsdx"] = b"visio"
    if "embedded_packages" in surfaces:
        parts["word/embeddings/package.bin"] = b"package"
        parts["word/_rels/sample-package.rels"] = _rels(
            b'<Relationship Id="rIdPkg" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" Target="../embeddings/package.bin"/>'
        )
    if "drawing_media_rels" in surfaces:
        parts["word/media/seal-placeholder.png"] = (
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
            b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02"
            b"\x00\x00\x00\x90wS\xde\x00\x00\x00\x00IEND\xaeB`\x82"
        )
        parts["word/_rels/sample-drawing.rels"] = _rels(
            b'<Relationship Id="rIdSeal" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/seal-placeholder.png"/>'
        )
    if "headers_footers" in surfaces:
        parts["word/header1.xml"] = _xml_body(
            b"<w:p><w:r><w:t>{{seal}}</w:t></w:r></w:p>"
        )
        parts["word/_rels/sample-header.rels"] = _rels(
            b'<Relationship Id="rIdHeader" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="../header1.xml"/>'
        )
    if "package_relationships" in surfaces:
        parts["word/_rels/sample-package-root.rels"] = _rels(
            b'<Relationship Id="rIdArchive" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="../docProps/core.xml"/>'
        )
    if "macros" in surfaces:
        parts["word/vbaProject.bin"] = b"macro"
        parts["word/_rels/sample-vba.rels"] = _rels(
            b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
        )
    return parts


def _xml_body(inner: bytes) -> bytes:
    return (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b"<w:body>"
        + inner
        + b"</w:body></w:document>"
    )


def _rels(inner: bytes) -> bytes:
    return (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + inner
        + b"</Relationships>"
    )


def _write_json_payload(path: Path, payload: dict[str, object]) -> None:
    import json

    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _request_cell_payload(
    cell: SceneRequestCellFixtureSpec,
    report_path: str,
) -> dict[str, object]:
    payload = cell.to_payload()
    payload["report_path"] = report_path
    payload["report_anchor"] = _request_cell_anchor(cell.sample_id)
    payload["report_section_title"] = cell.sample_id
    return payload


def _write_request_cell_report(
    path: Path,
    request_cells: tuple[SceneRequestCellFixtureSpec, ...],
) -> None:
    lines = [
        "# Scene Request-Cell Report",
        "",
        f"- Request cells: {len(request_cells)}",
        "",
    ]
    for cell in request_cells:
        lines.extend(_request_cell_report_section(cell))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _request_cell_report_section(
    cell: SceneRequestCellFixtureSpec,
) -> list[str]:
    lines = [
        f'<a id="{_request_cell_anchor(cell.sample_id)}"></a>',
        f"## {cell.sample_id}",
        "",
        f"- Request: {cell.request_text}",
        f"- Expected status: {cell.expected_status}",
        f"- Coverage: {cell.coverage_level}",
        f"- Packs: {_join_report_values(cell.expected_pack_ids)}",
        f"- Families: {_join_report_values(cell.expected_family_ids)}",
        f"- Fixtures: {_join_report_values(cell.fixture_ids)}",
        f"- Manual gates: {_join_report_values(cell.manual_gate_ids)}",
        f"- Boundary notes: {_join_report_values(cell.boundary_notes)}",
        "- Disambiguation required: "
        + ("yes" if cell.disambiguation_required else "no"),
        "",
    ]
    return lines


def _request_cell_anchor(sample_id: str) -> str:
    normalized = "".join(
        char.lower() if char.isalnum() else "-"
        for char in str(sample_id or "").strip()
    ).strip("-")
    while "--" in normalized:
        normalized = normalized.replace("--", "-")
    return f"request-cell-{normalized or 'unknown'}"


def _join_report_values(values: tuple[str, ...]) -> str:
    return ", ".join(str(value) for value in values if str(value).strip()) or "-"


__all__ = [
    "DEFAULT_SCENE_SAMPLE_FIXTURE_DIR",
    "DEFAULT_SCENE_SAMPLE_MANIFEST_NAME",
    "DEFAULT_SCENE_REQUEST_CELL_REPORT_NAME",
    "SceneSampleDocxArtifact",
    "SceneSampleDocxLibrary",
    "build_scene_sample_docx",
    "build_scene_sample_docx_library",
    "scene_sample_fixture_manifest_paths",
]
