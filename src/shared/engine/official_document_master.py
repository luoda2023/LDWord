"""Generate the built-in GB/T 9704 official-document master families."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.shared.engine.font_resolver import resolve_font


OFFICIAL_MASTER_VERSION = "official-master-family-gbt-2026-08-02-r5"

OFFICIAL_MASTER_FAMILIES: dict[str, tuple[str, ...]] = {
    "common": (
        "resolution",
        "decision",
        "bulletin",
        "announcement",
        "notice_public",
        "opinion",
        "notice",
        "circular",
        "approval",
    ),
    "upward": ("report", "request", "proposal"),
    # A notice or approval may deliberately use letter form. Automatic routing
    # only selects this family for ``letter``; the extra ids allow explicit use.
    "letter": ("letter", "notice", "approval"),
    "minutes": ("minutes",),
    "order": ("order",),
}


def write_official_gbt_master_docx(target: Path | str) -> Path:
    """Backward-compatible writer for the ordinary redhead master."""

    return write_official_master_family_docx(target, family_id="common")


def write_official_master_family_docx(
    target: Path | str,
    *,
    family_id: str,
) -> Path:
    """Write one clean runtime master for a distinct GB/T layout family."""

    family = str(family_id or "").strip().lower()
    if family not in OFFICIAL_MASTER_FAMILIES:
        raise ValueError(f"unknown official master family: {family_id}")

    path = Path(target)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.core_properties.title = f"GB/T 9704 official master: {family}"
    document.core_properties.subject = "LDWord official-document master"
    document.core_properties.keywords = f"alavette:official-layout={family}"
    document.core_properties.comments = OFFICIAL_MASTER_VERSION

    _configure_page(document, hide_first_page_number=family == "letter")
    _configure_styles(document)
    {
        "common": _add_common_layout,
        "upward": _add_upward_layout,
        "letter": _add_letter_layout,
        "minutes": _add_minutes_layout,
        "order": _add_order_layout,
    }[family](document)

    document.save(str(path))
    return path


def _configure_page(document: Document, *, hide_first_page_number: bool) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.5)
    # GB/T 9704 places the page-number baseline 7 mm below the lower edge of
    # the type area.  With the 35 mm bottom margin and the 14 pt number this
    # footer origin renders at that position in Word/WPS.
    section.footer_distance = Cm(2.35)
    section.different_first_page_header_footer = bool(hide_first_page_number)
    document.settings.odd_and_even_pages_header_footer = True
    _configure_page_number_footer(section.footer, odd=True)
    _configure_page_number_footer(section.even_page_footer, odd=False)
    if hide_first_page_number:
        first = section.first_page_footer.paragraphs[0]
        first.text = ""


def _configure_styles(document: Document) -> None:
    body_font = _official_font("仿宋")
    normal = document.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(16)
    _set_style_font(normal, body_font)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(28)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    heading_roles = (
        ("Heading 1", "黑体"),
        ("Heading 2", "楷体"),
        ("Heading 3", "仿宋"),
        ("Heading 4", "仿宋"),
    )
    for name, role in heading_roles:
        font_name = _official_font(role)
        style = document.styles[name]
        style.font.name = font_name
        style.font.size = Pt(16)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        _set_style_font(style, font_name)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(28)
        # These built-in styles provide the GB/T font roles only.  Word's
        # default Heading styles keep consecutive headings together, which can
        # move an otherwise fitting body block to the next page and leave half
        # of the first page blank.
        style.paragraph_format.keep_with_next = False
        style.paragraph_format.keep_together = False
        style.paragraph_format.page_break_before = False


def _add_common_layout(document: Document) -> None:
    _add_classification_reserve(document)
    _add_redhead_mark(document, "{{@text:official_organization}}")
    _add_document_number(document, "{{@text:official_document_no}}")
    _add_separator(document)
    _add_title_and_body(document, recipient=True, attachment=True)
    _add_closing(document)
    _add_imprint(document)


def _add_upward_layout(document: Document) -> None:
    _add_classification_reserve(document)
    _add_redhead_mark(document, "{{@text:official_organization}}")

    row = document.add_table(rows=1, cols=2)
    row.alignment = WD_TABLE_ALIGNMENT.CENTER
    row.autofit = False
    _set_table_grid(row, (Cm(7.8), Cm(7.8)))
    _set_table_borders(row, edges=(), color="FFFFFF")
    left, right = row.rows[0].cells
    left.text = "{{@text:official_document_no}}"
    right.text = ""
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
        cell.paragraphs[0].paragraph_format.line_spacing = Pt(28)
    signer_label = right.paragraphs[0].add_run("签发人：")
    _set_run_font(
        signer_label,
        east_asia=_official_font("仿宋"),
        ascii_font=_official_font("仿宋"),
        size=16,
    )
    signer_value = right.paragraphs[0].add_run("{{@text:official_signer}}")
    _set_run_font(
        signer_value,
        east_asia=_official_font("楷体"),
        ascii_font=_official_font("楷体"),
        size=16,
    )

    _add_separator(document)
    _add_title_and_body(document, recipient=True, attachment=True)
    _add_closing(document)
    _add_imprint(document)


def _add_letter_layout(document: Document) -> None:
    # Letter form is measured from the physical page top rather than from the
    # ordinary 37 mm text-area top.  The 23.5 mm paragraph origin compensates
    # for the display font's ascent so the rendered mark starts at about 30 mm.
    section = document.sections[0]
    section.top_margin = Cm(2.35)
    section.footer_distance = Cm(2.0)

    mark = document.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark.paragraph_format.left_indent = Cm(-0.7)
    mark.paragraph_format.right_indent = Cm(-0.7)
    mark.paragraph_format.space_after = Pt(0)
    _set_large_mark_line_height(mark, font_size_pt=36)
    _add_formatted_run(
        mark,
        "{{@text:official_organization}}",
        east_asia=_official_font("小标宋"),
        ascii_font=_official_font("小标宋"),
        size=36,
        color="C00000",
        bold=False,
    )
    _set_double_bottom_border(
        mark,
        color="C00000",
        space="0",
        style="thickThinSmallGap",
    )

    metadata = document.add_table(rows=1, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.autofit = False
    _set_table_grid(metadata, (Cm(7.8), Cm(7.8)))
    _set_table_borders(metadata, edges=(), color="FFFFFF")
    left, right = metadata.rows[0].cells
    left.text = "{{@text:official_security_level}}"
    urgency = left.add_paragraph("{{@text:official_urgency}}")
    right.text = "{{@text:official_document_no}}"
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
        for index, paragraph in enumerate(cell.paragraphs):
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(4 if index == 0 else 0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph.paragraph_format.line_spacing = Pt(28)
    for paragraph in left.paragraphs:
        _set_paragraph_font(
            paragraph,
            east_asia=_official_font("黑体"),
            size=16,
        )

    _add_title_and_body(document, recipient=True, attachment=True, title_before=28)
    _add_closing(document)
    _add_letter_copy_scope(document)
    footers = (
        section.first_page_footer,
        section.footer,
        section.even_page_footer,
    )
    for footer in footers:
        paragraph = footer.paragraphs[0]
        paragraph.paragraph_format.left_indent = Cm(-0.7)
        paragraph.paragraph_format.right_indent = Cm(-0.7)
        _set_double_bottom_border(
            paragraph,
            color="C00000",
            space="2",
            style="thinThickSmallGap",
        )


def _add_minutes_layout(document: Document) -> None:
    # A minutes mark follows the ordinary 35 mm distance from the top of the
    # text area, unlike the 20 mm command/order layout.
    _add_classification_reserve(document)
    mark = document.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mark.paragraph_format.space_after = Pt(2)
    _set_large_mark_line_height(mark, font_size_pt=45)
    _add_formatted_run(
        mark,
        "{{@text:official_organization}}纪要",
        east_asia=_official_font("小标宋"),
        ascii_font=_official_font("小标宋"),
        size=45,
        color="C00000",
        bold=False,
    )

    issue = document.add_paragraph("{{@text:official_document_no}}")
    issue.alignment = WD_ALIGN_PARAGRAPH.CENTER
    issue.paragraph_format.space_after = Pt(6)

    meta = document.add_table(rows=1, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    _set_table_grid(meta, (Cm(9.5), Cm(6.1)))
    _set_table_borders(meta, edges=(), color="FFFFFF")
    meta.rows[0].cells[0].text = "{{@text:official_issuer}}"
    meta.rows[0].cells[1].text = "{{@text:official_issue_date}}"
    meta.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for cell in meta.rows[0].cells:
        _set_cell_margins(cell, top=0, bottom=0, start=0, end=0)

    _add_separator(document)
    _add_title(document, before=Pt(56), after=Pt(18))

    meeting_time = document.add_paragraph("会议时间：{{@text:official_meeting_time}}")
    meeting_time.paragraph_format.first_line_indent = Cm(0)
    attendees = document.add_paragraph("出席：{{@text:official_meeting_attendees}}")
    attendees.paragraph_format.first_line_indent = Cm(0)
    body = document.add_paragraph("{{@text:official_body}}")
    body.paragraph_format.first_line_indent = Cm(1.1)
    _add_imprint(document)


def _add_order_layout(document: Document) -> None:
    mark = document.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Command/order form starts 20 mm below the ordinary text-area top.  A
    # 31 pt paragraph offset plus the selected display font's ascent renders
    # the glyph top at that position without retaining empty optional fields.
    mark.paragraph_format.space_before = Pt(31)
    mark.paragraph_format.space_after = Pt(4)
    _set_large_mark_line_height(mark, font_size_pt=50)
    _add_formatted_run(
        mark,
        "{{@text:official_organization}}令",
        east_asia=_official_font("小标宋"),
        ascii_font=_official_font("小标宋"),
        size=50,
        color="C00000",
        bold=False,
    )
    doc_no = document.add_paragraph("{{@text:official_document_no}}")
    doc_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # The title begins two 28 pt body lines below the order number.
    doc_no.paragraph_format.space_after = Pt(52)
    _add_title(document, before=Pt(0), after=Pt(20))
    body = document.add_paragraph("{{@text:official_body}}")
    body.paragraph_format.first_line_indent = Cm(1.1)
    signer = document.add_paragraph()
    signer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    signer.paragraph_format.right_indent = Cm(1.1)
    signer.paragraph_format.space_before = Pt(20)
    _add_formatted_run(
        signer,
        "{{@text:official_signer}}",
        east_asia=_official_font("楷体"),
        ascii_font=_official_font("楷体"),
        size=16,
    )
    date = document.add_paragraph("{{@text:official_issue_date}}")
    date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date.paragraph_format.right_indent = Cm(1.1)


def _add_classification_reserve(document: Document, *, compact: bool = False) -> None:
    if not compact:
        spacer = document.add_paragraph("")
        spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        spacer.paragraph_format.line_spacing = Pt(18)
    security = document.add_paragraph("{{@text:official_security_level}}")
    urgency = document.add_paragraph("{{@text:official_urgency}}")
    for paragraph in (security, urgency):
        paragraph.paragraph_format.first_line_indent = Cm(0)
        _set_paragraph_font(
            paragraph,
            east_asia=_official_font("黑体"),
            size=16,
        )


def _add_redhead_mark(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    _set_large_mark_line_height(paragraph, font_size_pt=50)
    _add_formatted_run(
        paragraph,
        text,
        east_asia=_official_font("小标宋"),
        ascii_font=_official_font("小标宋"),
        size=50,
        color="C00000",
        bold=False,
    )


def _add_document_number(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(2)


def _add_separator(document: Document) -> None:
    separator = document.add_paragraph()
    separator.paragraph_format.space_before = Pt(3)
    separator.paragraph_format.space_after = Pt(0)
    separator.paragraph_format.line_spacing = Pt(4)
    _set_bottom_border(separator, color="C00000", size="18")


def _add_title_and_body(
    document: Document,
    *,
    recipient: bool,
    attachment: bool,
    title_before: int = 56,
) -> None:
    _add_title(document, before=Pt(title_before), after=Pt(18))
    if recipient:
        target = document.add_paragraph("{{@text:official_recipient}}：")
        target.paragraph_format.first_line_indent = Cm(0)
    body = document.add_paragraph("{{@text:official_body}}")
    body.paragraph_format.first_line_indent = Cm(1.1)
    if attachment:
        attachment_line = document.add_paragraph(
            "附件：{{@text:official_attachment_note}}"
        )
        attachment_line.paragraph_format.first_line_indent = Cm(1.1)


def _add_title(document: Document, *, before: Pt, after: Pt) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = before
    paragraph.paragraph_format.space_after = after
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(32)
    _add_formatted_run(
        paragraph,
        "{{@text:official_title}}",
        east_asia=_official_font("小标宋"),
        ascii_font=_official_font("小标宋"),
        size=22,
        bold=False,
    )


def _add_closing(document: Document) -> None:
    spacer = document.add_paragraph("")
    spacer.paragraph_format.line_spacing = Pt(14)
    issuer = document.add_paragraph("{{@text:official_issuer}}")
    issuer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    issuer.paragraph_format.right_indent = Cm(1.1)
    issuer.paragraph_format.space_after = Pt(0)
    issue_date = document.add_paragraph("{{@text:official_issue_date}}")
    issue_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    issue_date.paragraph_format.right_indent = Cm(1.1)


def _add_imprint(document: Document) -> None:
    # A single non-splitting row owns the full imprint.  The copy line and the
    # printing line are two paragraphs in one cell, so Word cannot strand only
    # the second line on a new page.
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_grid(table, (Cm(15.6),))
    row = table.rows[0]
    _prevent_row_split(row)
    copy_cell = row.cells[0]
    copy_cell.text = "抄送：{{@text:official_copy_scope}}。"
    separator = copy_cell.add_paragraph("")
    printing = copy_cell.add_paragraph(
        "{{@text:official_printing_org}}\t{{@text:official_printing_date}}"
    )
    _set_right_tab_stop(printing, position_twips=8200)
    copy_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(copy_cell, top=0, bottom=0, start=0, end=0)
    for paragraph in (copy_cell.paragraphs[0], printing):
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0.55)
        paragraph.paragraph_format.right_indent = Cm(0.55)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(28)
        _set_paragraph_font(
            paragraph,
            east_asia=_official_font("仿宋"),
            size=14,
        )
    separator.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    separator.paragraph_format.line_spacing = Pt(2)
    separator.paragraph_format.space_before = Pt(0)
    separator.paragraph_format.space_after = Pt(0)
    _set_bottom_border(separator, color="000000", size="6")
    _set_table_borders(
        table,
        edges=("top", "bottom"),
        color="000000",
        sizes={"top": "8", "bottom": "8"},
    )
    _set_table_floating_bottom(table)


def _add_letter_copy_scope(document: Document) -> None:
    """Place the letter-form copy scope on the final page, not page one."""

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_grid(table, (Cm(17.0),))
    _set_table_borders(table, edges=(), color="FFFFFF")
    row = table.rows[0]
    _prevent_row_split(row)
    cell = row.cells[0]
    cell.text = "抄送：{{@text:official_copy_scope}}"
    _set_cell_margins(cell, top=0, bottom=0, start=320, end=320)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(28)
    _set_paragraph_font(
        paragraph,
        east_asia=_official_font("仿宋"),
        size=14,
    )
    _set_table_floating_bottom(table)


def _configure_page_number_footer(footer, *, odd: bool) -> None:
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if odd else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.right_indent = Cm(0.55) if odd else Cm(0)
    paragraph.paragraph_format.left_indent = Cm(0.55) if not odd else Cm(0)
    run = paragraph.add_run("— ")
    page_font = _official_font("宋体")
    _set_run_font(run, east_asia=page_font, ascii_font=page_font, size=14)
    _append_page_field(paragraph)
    run = paragraph.add_run(" —")
    _set_run_font(run, east_asia=page_font, ascii_font=page_font, size=14)


def _append_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    page_font = _official_font("宋体")
    _set_run_font(run, east_asia=page_font, ascii_font=page_font, size=14)


def _add_formatted_run(
    paragraph,
    text: str,
    *,
    east_asia: str,
    ascii_font: str,
    size: float,
    color: str = "000000",
    bold: bool = False,
) -> None:
    run = paragraph.add_run(text)
    _set_run_font(
        run,
        east_asia=east_asia,
        ascii_font=ascii_font,
        size=size,
        color=color,
        bold=bold,
    )


def _set_paragraph_font(
    paragraph,
    *,
    east_asia: str,
    size: float,
    bold: bool = False,
) -> None:
    for run in paragraph.runs:
        _set_run_font(
            run,
            east_asia=east_asia,
            ascii_font=east_asia,
            size=size,
            bold=bold,
        )


def _set_run_font(
    run,
    *,
    east_asia: str,
    ascii_font: str,
    size: float,
    color: str = "000000",
    bold: bool = False,
) -> None:
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _set_style_font(style, font_name: str) -> None:
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def _official_font(role: str) -> str:
    """Resolve one GB/T semantic font role to an installed physical family."""

    return resolve_font(role, lang="cn")


def _set_large_mark_line_height(paragraph, *, font_size_pt: float) -> None:
    """Prevent large redhead glyphs from inheriting Normal's 28 pt exact grid.

    CJK display fonts commonly have ascent/descent metrics taller than their
    nominal point size. ``AT_LEAST`` keeps the GB/T body grid independent while
    allowing Word/WPS to expand the line for an explicitly selected template
    font instead of clipping its top or bottom.
    """

    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    paragraph.paragraph_format.line_spacing = Pt(float(font_size_pt) * 1.2)


def _set_bottom_border(paragraph, *, color: str, size: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def _set_double_bottom_border(
    paragraph,
    *,
    color: str,
    space: str = "2",
    style: str = "double",
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), style)
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _set_top_border(
    paragraph,
    *,
    color: str,
    size: str,
    style: str = "single",
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), style)
    top.set(qn("w:sz"), size)
    top.set(qn("w:space"), "2")
    top.set(qn("w:color"), color)
    p_bdr.append(top)


def _set_right_tab_stop(paragraph, *, position_twips: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        p_pr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(position_twips)))
    tabs.append(tab)


def _set_cell_margins(cell, *, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (
        ("top", top),
        ("bottom", bottom),
        ("start", start),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_grid(table, widths: tuple[Cm, ...]) -> None:
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    total_twips = 0
    for index, width in enumerate(widths):
        twips = int(width.twips)
        total_twips += twips
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(twips))
        grid.append(col)
        for row in table.rows:
            row.cells[index].width = width
            tc_w = row.cells[index]._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(twips))
            tc_w.set(qn("w:type"), "dxa")
    tbl_w = table._tbl.tblPr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table._tbl.tblPr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    layout = table._tbl.tblPr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table._tbl.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_table_borders(
    table,
    *,
    edges: tuple[str, ...],
    color: str,
    sizes: dict[str, str] | None = None,
) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    enabled = set(edges)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single" if edge in enabled else "nil")
        node.set(qn("w:sz"), str((sizes or {}).get(edge, "4")))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def _set_table_floating_bottom(table) -> None:
    """Anchor a terminal record table to the bottom of the text area.

    The table remains part of the body story, so Word places it on the actual
    final page.  Positioning relative to the page margins keeps the record at
    the lower edge of the type area without inserting manual page breaks.
    """

    tbl_pr = table._tbl.tblPr
    position = tbl_pr.first_child_found_in("w:tblpPr")
    if position is None:
        position = OxmlElement("w:tblpPr")
        tbl_pr.append(position)
    for key, value in (
        ("horzAnchor", "margin"),
        ("tblpXSpec", "center"),
        ("vertAnchor", "margin"),
        ("tblpYSpec", "bottom"),
        ("leftFromText", "0"),
        ("rightFromText", "0"),
        ("topFromText", "0"),
        ("bottomFromText", "0"),
    ):
        position.set(qn(f"w:{key}"), value)
    overlap = tbl_pr.first_child_found_in("w:tblOverlap")
    if overlap is None:
        overlap = OxmlElement("w:tblOverlap")
        tbl_pr.append(overlap)
    overlap.set(qn("w:val"), "never")


__all__ = [
    "OFFICIAL_MASTER_FAMILIES",
    "OFFICIAL_MASTER_VERSION",
    "write_official_gbt_master_docx",
    "write_official_master_family_docx",
]
