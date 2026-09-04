"""Isolated capability probe for exact Word/WPS document geometry.

This module is deliberately diagnostic infrastructure, not the production layout
adapter.  A provider is considered usable only when every capability required by
``OfficeProviderReceipt.exact_layout_ready`` is proven in a fresh Office process.

The live checks always run in a child Python process and use ``DispatchEx``.  They
never attach to an Office instance owned by the user and never open a user file:
the parent creates a synthetic DOCX and the child edits only a temporary shadow
copy.  The parent owns the timeout and cleans up the exact PID reported by the
child if normal COM shutdown does not finish.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import datetime, timezone
from enum import Enum
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
from src.shared.win_process import popen_hidden, run_hidden
import sys
import tempfile
import time
from typing import Any, Iterable, Mapping, Sequence

from src.shared.engine.office_broker_command import (
    OFFICE_LAYOUT_PROBE_CHILD_FLAG,
    build_office_broker_child_command,
)


PROBE_SCHEMA_VERSION = 1


class OfficeProvider(str, Enum):
    WORD = "word"
    WPS = "wps"


class OfficeCapability(str, Enum):
    REGISTRATION = "registration"
    APPLICATION_CREATED = "application_created"
    DEDICATED_PROCESS = "dedicated_process"
    SHADOW_DOCUMENT = "shadow_document"
    PRINT_LAYOUT = "print_layout"
    FIELDS_UPDATE = "fields_update"
    TOC_UPDATE = "toc_update"
    REPAGINATE = "repaginate"
    PAGE_NUMBER = "page_number"
    VERTICAL_POSITION = "vertical_position"
    CURRENT_COLUMN = "current_column"
    CONTAINER_WIDTH = "container_width"
    SEQUENTIAL_GEOMETRY = "sequential_geometry"
    CLEAN_SHUTDOWN = "clean_shutdown"


class CapabilityOutcome(str, Enum):
    PASSED = "passed"
    FAILED = "failed"
    UNAVAILABLE = "unavailable"
    NOT_RUN = "not_run"


class ProviderOutcome(str, Enum):
    QUALIFIED = "qualified"
    DISQUALIFIED = "disqualified"
    UNAVAILABLE = "unavailable"
    DISCOVERED = "discovered"
    TIMED_OUT = "timed_out"
    CHILD_FAILED = "child_failed"


EvidenceValue = str | int | float | bool | None


@dataclass(frozen=True, slots=True)
class OfficeProviderSpec:
    provider: OfficeProvider
    prog_id: str
    process_names: tuple[str, ...]


DEFAULT_PROVIDER_SPECS: tuple[OfficeProviderSpec, ...] = (
    OfficeProviderSpec(
        provider=OfficeProvider.WORD,
        prog_id="Word.Application",
        process_names=("WINWORD",),
    ),
    OfficeProviderSpec(
        provider=OfficeProvider.WPS,
        prog_id="KWPS.Application",
        process_names=("wps",),
    ),
)


_LIVE_CAPABILITIES: tuple[OfficeCapability, ...] = (
    OfficeCapability.APPLICATION_CREATED,
    OfficeCapability.DEDICATED_PROCESS,
    OfficeCapability.SHADOW_DOCUMENT,
    OfficeCapability.PRINT_LAYOUT,
    OfficeCapability.FIELDS_UPDATE,
    OfficeCapability.TOC_UPDATE,
    OfficeCapability.REPAGINATE,
    OfficeCapability.PAGE_NUMBER,
    OfficeCapability.VERTICAL_POSITION,
    OfficeCapability.CURRENT_COLUMN,
    OfficeCapability.CONTAINER_WIDTH,
    OfficeCapability.SEQUENTIAL_GEOMETRY,
    OfficeCapability.CLEAN_SHUTDOWN,
)

_EXACT_LAYOUT_REQUIRED: frozenset[OfficeCapability] = frozenset(
    (OfficeCapability.REGISTRATION, *_LIVE_CAPABILITIES)
)


@dataclass(frozen=True, slots=True)
class CapabilityReceipt:
    capability: OfficeCapability
    outcome: CapabilityOutcome
    elapsed_ms: float = 0.0
    detail: str = ""
    evidence: tuple[tuple[str, EvidenceValue], ...] = ()

    @property
    def passed(self) -> bool:
        return self.outcome is CapabilityOutcome.PASSED

    def to_dict(self) -> dict[str, Any]:
        return {
            "capability": self.capability.value,
            "outcome": self.outcome.value,
            "elapsed_ms": round(float(self.elapsed_ms), 3),
            "detail": self.detail,
            "evidence": {key: value for key, value in self.evidence},
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> CapabilityReceipt:
        raw_evidence = data.get("evidence") or {}
        evidence = tuple(
            (str(key), _coerce_evidence_value(value))
            for key, value in sorted(dict(raw_evidence).items())
        )
        return cls(
            capability=OfficeCapability(str(data["capability"])),
            outcome=CapabilityOutcome(str(data["outcome"])),
            elapsed_ms=float(data.get("elapsed_ms", 0.0) or 0.0),
            detail=str(data.get("detail", "") or ""),
            evidence=evidence,
        )


@dataclass(frozen=True, slots=True)
class OfficeProviderReceipt:
    provider: OfficeProvider
    prog_id: str
    registered: bool
    outcome: ProviderOutcome
    elapsed_ms: float
    application_pid: int | None = None
    capabilities: tuple[CapabilityReceipt, ...] = ()
    issues: tuple[str, ...] = ()

    def capability(self, name: OfficeCapability) -> CapabilityReceipt | None:
        return next(
            (item for item in self.capabilities if item.capability is name),
            None,
        )

    @property
    def failed_capabilities(self) -> tuple[OfficeCapability, ...]:
        return tuple(
            item.capability
            for item in self.capabilities
            if item.outcome is CapabilityOutcome.FAILED
        )

    @property
    def exact_layout_ready(self) -> bool:
        receipts = {item.capability: item for item in self.capabilities}
        return self.outcome is ProviderOutcome.QUALIFIED and all(
            receipts.get(name) is not None and receipts[name].passed
            for name in _EXACT_LAYOUT_REQUIRED
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "provider": self.provider.value,
            "prog_id": self.prog_id,
            "registered": self.registered,
            "outcome": self.outcome.value,
            "exact_layout_ready": self.exact_layout_ready,
            "elapsed_ms": round(float(self.elapsed_ms), 3),
            "application_pid": self.application_pid,
            "capabilities": [item.to_dict() for item in self.capabilities],
            "failed_capabilities": [item.value for item in self.failed_capabilities],
            "issues": list(self.issues),
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> OfficeProviderReceipt:
        raw_pid = data.get("application_pid")
        return cls(
            provider=OfficeProvider(str(data["provider"])),
            prog_id=str(data["prog_id"]),
            registered=bool(data.get("registered", False)),
            outcome=ProviderOutcome(str(data["outcome"])),
            elapsed_ms=float(data.get("elapsed_ms", 0.0) or 0.0),
            application_pid=int(raw_pid) if raw_pid is not None else None,
            capabilities=tuple(
                CapabilityReceipt.from_dict(item)
                for item in data.get("capabilities", ())
            ),
            issues=tuple(str(item) for item in data.get("issues", ())),
        )


@dataclass(frozen=True, slots=True)
class OfficeLayoutProbeReceipt:
    schema_version: int
    started_at_utc: str
    platform: str
    python_version: str
    timeout_seconds: float
    elapsed_ms: float
    providers: tuple[OfficeProviderReceipt, ...]

    @property
    def exact_layout_ready(self) -> bool:
        return any(provider.exact_layout_ready for provider in self.providers)

    @property
    def qualified_providers(self) -> tuple[OfficeProvider, ...]:
        return tuple(
            provider.provider
            for provider in self.providers
            if provider.exact_layout_ready
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "schema_version": self.schema_version,
            "started_at_utc": self.started_at_utc,
            "platform": self.platform,
            "python_version": self.python_version,
            "timeout_seconds": self.timeout_seconds,
            "elapsed_ms": round(float(self.elapsed_ms), 3),
            "exact_layout_ready": self.exact_layout_ready,
            "qualified_providers": [item.value for item in self.qualified_providers],
            "providers": [item.to_dict() for item in self.providers],
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> OfficeLayoutProbeReceipt:
        return cls(
            schema_version=int(data.get("schema_version", PROBE_SCHEMA_VERSION)),
            started_at_utc=str(data.get("started_at_utc", "") or ""),
            platform=str(data.get("platform", "") or ""),
            python_version=str(data.get("python_version", "") or ""),
            timeout_seconds=float(data.get("timeout_seconds", 0.0) or 0.0),
            elapsed_ms=float(data.get("elapsed_ms", 0.0) or 0.0),
            providers=tuple(
                OfficeProviderReceipt.from_dict(item)
                for item in data.get("providers", ())
            ),
        )


def probe_office_layout(
    *,
    timeout_seconds: float = 45.0,
    provider_specs: Sequence[OfficeProviderSpec] = DEFAULT_PROVIDER_SPECS,
    execute_live: bool = True,
) -> OfficeLayoutProbeReceipt:
    """Probe registered Office providers without touching a user document.

    ``timeout_seconds`` is a hard wall-clock limit for each provider child.  Set
    ``execute_live=False`` for registration-only discovery (useful in diagnostics
    and tests); discovery never starts an Office process.
    """

    if timeout_seconds <= 0:
        raise ValueError("timeout_seconds must be positive")

    started_wall = datetime.now(timezone.utc).isoformat()
    started = time.perf_counter()
    receipts: list[OfficeProviderReceipt] = []
    for spec in provider_specs:
        registered, registration_detail = _registration_status(spec.prog_id)
        registration = CapabilityReceipt(
            capability=OfficeCapability.REGISTRATION,
            outcome=(
                CapabilityOutcome.PASSED
                if registered
                else CapabilityOutcome.UNAVAILABLE
            ),
            detail=registration_detail,
            evidence=(("prog_id", spec.prog_id),),
        )
        if not registered:
            receipts.append(
                OfficeProviderReceipt(
                    provider=spec.provider,
                    prog_id=spec.prog_id,
                    registered=False,
                    outcome=ProviderOutcome.UNAVAILABLE,
                    elapsed_ms=0.0,
                    capabilities=(
                        registration,
                        *_not_run_capabilities("provider is not registered"),
                    ),
                    issues=(registration_detail,),
                )
            )
            continue

        if not execute_live:
            receipts.append(
                OfficeProviderReceipt(
                    provider=spec.provider,
                    prog_id=spec.prog_id,
                    registered=True,
                    outcome=ProviderOutcome.DISCOVERED,
                    elapsed_ms=0.0,
                    capabilities=(
                        registration,
                        *_not_run_capabilities("live checks disabled"),
                    ),
                )
            )
            continue

        live_receipt = _run_provider_child(spec, timeout_seconds)
        receipts.append(
            replace(
                live_receipt,
                capabilities=(registration, *live_receipt.capabilities),
            )
        )

    return OfficeLayoutProbeReceipt(
        schema_version=PROBE_SCHEMA_VERSION,
        started_at_utc=started_wall,
        platform=sys.platform,
        python_version=sys.version.split()[0],
        timeout_seconds=float(timeout_seconds),
        elapsed_ms=(time.perf_counter() - started) * 1000.0,
        providers=tuple(receipts),
    )


def _not_run_capabilities(reason: str) -> tuple[CapabilityReceipt, ...]:
    return tuple(
        CapabilityReceipt(
            capability=name,
            outcome=CapabilityOutcome.NOT_RUN,
            detail=reason,
        )
        for name in _LIVE_CAPABILITIES
    )


def _registration_status(prog_id: str) -> tuple[bool, str]:
    if os.name != "nt":
        return False, "Office COM registration requires Windows"
    try:
        import winreg
    except ImportError:
        return False, "winreg is unavailable"

    access_modes = (0, winreg.KEY_WOW64_64KEY, winreg.KEY_WOW64_32KEY)
    errors: list[str] = []
    for access in access_modes:
        try:
            with winreg.OpenKey(
                winreg.HKEY_CLASSES_ROOT,
                rf"{prog_id}\CLSID",
                0,
                winreg.KEY_READ | access,
            ) as key:
                clsid, _kind = winreg.QueryValueEx(key, None)
            if str(clsid or "").strip():
                return True, f"registered ({clsid})"
        except OSError as exc:
            errors.append(str(exc))
    detail = errors[-1] if errors else "CLSID key not found"
    return False, f"{prog_id} is not registered: {detail}"


def _run_provider_child(
    spec: OfficeProviderSpec,
    timeout_seconds: float,
) -> OfficeProviderReceipt:
    started = time.perf_counter()
    with tempfile.TemporaryDirectory(prefix="ldword-office-layout-probe-") as raw_dir:
        work_dir = Path(raw_dir)
        source_path = work_dir / "synthetic-source.docx"
        shadow_path = work_dir / "probe-shadow.docx"
        image_path = work_dir / "probe-image.png"
        result_path = work_dir / "child-result.json"
        pid_path = work_dir / "office-owner-pid.json"
        try:
            _build_probe_fixture(source_path, image_path)
            shutil.copy2(source_path, shadow_path)
        except Exception as exc:
            return _child_failed_receipt(
                spec,
                f"failed to prepare synthetic shadow DOCX: {exc}",
                started,
            )

        source_digest = _sha256(source_path)
        before_pids = _process_ids_for_names(spec.process_names)
        request = {
            "provider": spec.provider.value,
            "prog_id": spec.prog_id,
            "source_path": str(source_path.resolve()),
            "shadow_path": str(shadow_path.resolve()),
            "image_path": str(image_path.resolve()),
            "result_path": str(result_path.resolve()),
            "pid_path": str(pid_path.resolve()),
            "before_pids": sorted(before_pids),
        }
        request_path = work_dir / "request.json"
        request_path.write_text(
            json.dumps(request, ensure_ascii=False),
            encoding="utf-8",
        )

        command = build_office_broker_child_command(
            module_name="src.shared.engine.office_layout_probe",
            frozen_flag=OFFICE_LAYOUT_PROBE_CHILD_FLAG,
            request_path=request_path,
        )
        creationflags = 0
        if os.name == "nt":
            creationflags = getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
        try:
            process = popen_hidden(
                command,
                cwd=str(Path(__file__).resolve().parents[3]),
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                creationflags=creationflags,
            )
        except Exception as exc:
            return _child_failed_receipt(
                spec,
                f"failed to start provider child: {exc}",
                started,
            )

        timed_out = False
        stdout = ""
        stderr = ""
        try:
            stdout, stderr = process.communicate(timeout=timeout_seconds)
        except subprocess.TimeoutExpired:
            timed_out = True
            _terminate_child_process(process)
            try:
                stdout, stderr = process.communicate(timeout=3)
            except Exception:
                pass

        owned_pid = _read_owned_pid(pid_path)
        cleanup_detail = _ensure_owned_process_stopped(owned_pid)

        if timed_out:
            elapsed_ms = (time.perf_counter() - started) * 1000.0
            issue = f"provider probe timed out after {timeout_seconds:g}s"
            if cleanup_detail:
                issue = f"{issue}; {cleanup_detail}"
            return OfficeProviderReceipt(
                provider=spec.provider,
                prog_id=spec.prog_id,
                registered=True,
                outcome=ProviderOutcome.TIMED_OUT,
                elapsed_ms=elapsed_ms,
                application_pid=owned_pid,
                capabilities=(
                    *_not_run_capabilities(issue)[:-1],
                    _cleanup_receipt(owned_pid, cleanup_detail),
                ),
                issues=(issue,),
            )

        if process.returncode != 0 or not result_path.is_file():
            detail = (stderr or stdout or "").strip()
            issue = detail or f"provider child exited with code {process.returncode}"
            receipt = _child_failed_receipt(spec, issue, started, owned_pid)
            return _with_cleanup(receipt, cleanup_detail)

        try:
            child_data = json.loads(result_path.read_text(encoding="utf-8"))
            receipt = OfficeProviderReceipt.from_dict(child_data)
        except Exception as exc:
            receipt = _child_failed_receipt(
                spec,
                f"invalid provider child receipt: {exc}",
                started,
                owned_pid,
            )
            return _with_cleanup(receipt, cleanup_detail)

        source_unchanged = _sha256(source_path) == source_digest
        shadow_result = CapabilityReceipt(
            capability=OfficeCapability.SHADOW_DOCUMENT,
            outcome=(
                CapabilityOutcome.PASSED
                if source_unchanged and shadow_path.is_file()
                else CapabilityOutcome.FAILED
            ),
            detail=(
                "only the temporary shadow was opened"
                if source_unchanged and shadow_path.is_file()
                else "synthetic source changed or shadow disappeared"
            ),
            evidence=(
                ("source_sha256_unchanged", source_unchanged),
                ("shadow_exists", shadow_path.is_file()),
            ),
        )
        receipt = _replace_capability(receipt, shadow_result)
        receipt = _with_cleanup(receipt, cleanup_detail)
        outcome = (
            ProviderOutcome.QUALIFIED
            if _all_required_live_capabilities_pass(receipt.capabilities)
            else ProviderOutcome.DISQUALIFIED
        )
        return replace(
            receipt,
            outcome=outcome,
            elapsed_ms=(time.perf_counter() - started) * 1000.0,
        )


def _child_failed_receipt(
    spec: OfficeProviderSpec,
    issue: str,
    started: float,
    application_pid: int | None = None,
) -> OfficeProviderReceipt:
    return OfficeProviderReceipt(
        provider=spec.provider,
        prog_id=spec.prog_id,
        registered=True,
        outcome=ProviderOutcome.CHILD_FAILED,
        elapsed_ms=(time.perf_counter() - started) * 1000.0,
        application_pid=application_pid,
        capabilities=_not_run_capabilities(issue),
        issues=(issue,),
    )


def _replace_capability(
    receipt: OfficeProviderReceipt,
    capability: CapabilityReceipt,
) -> OfficeProviderReceipt:
    items = [
        item
        for item in receipt.capabilities
        if item.capability is not capability.capability
    ]
    items.append(capability)
    order = {name: index for index, name in enumerate(_LIVE_CAPABILITIES)}
    items.sort(key=lambda item: order.get(item.capability, len(order)))
    return replace(receipt, capabilities=tuple(items))


def _with_cleanup(
    receipt: OfficeProviderReceipt,
    cleanup_detail: str,
) -> OfficeProviderReceipt:
    return _replace_capability(
        receipt,
        _cleanup_receipt(receipt.application_pid, cleanup_detail),
    )


def _cleanup_receipt(
    owned_pid: int | None,
    cleanup_detail: str,
) -> CapabilityReceipt:
    stopped = owned_pid is not None and not _pid_exists(owned_pid)
    return CapabilityReceipt(
        capability=OfficeCapability.CLEAN_SHUTDOWN,
        outcome=(
            CapabilityOutcome.PASSED if stopped else CapabilityOutcome.FAILED
        ),
        detail=cleanup_detail or "Office process exited after Quit",
        evidence=(("owned_pid", owned_pid), ("process_stopped", stopped)),
    )


def _all_required_live_capabilities_pass(
    capabilities: Iterable[CapabilityReceipt],
) -> bool:
    receipts = {item.capability: item for item in capabilities}
    return all(
        receipts.get(name) is not None and receipts[name].passed
        for name in _LIVE_CAPABILITIES
    )


def _terminate_child_process(process: subprocess.Popen[str]) -> None:
    if process.poll() is not None:
        return
    if os.name == "nt":
        try:
            run_hidden(
                ["taskkill", "/PID", str(process.pid), "/T", "/F"],
                capture_output=True,
                text=True,
                timeout=5,
                check=False,
            )
            return
        except Exception:
            pass
    try:
        process.kill()
    except Exception:
        pass


def _read_owned_pid(path: Path) -> int | None:
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
        pid = int(raw.get("pid", 0))
    except Exception:
        return None
    return pid if pid > 0 else None


def _ensure_owned_process_stopped(pid: int | None) -> str:
    if pid is None:
        return "Office PID was not reported; no unrelated process was targeted"
    deadline = time.monotonic() + 2.0
    while time.monotonic() < deadline:
        if not _pid_exists(pid):
            return "Office process exited after Quit"
        time.sleep(0.05)
    if os.name == "nt":
        try:
            completed = run_hidden(
                ["taskkill", "/PID", str(pid), "/T", "/F"],
                capture_output=True,
                text=True,
                timeout=5,
                check=False,
            )
            if completed.returncode == 0:
                return "owned Office PID required forced cleanup"
            if not _pid_exists(pid):
                return "Office process exited while cleanup was being checked"
            return "owned Office PID cleanup failed"
        except Exception as exc:
            return f"owned Office PID cleanup failed: {exc}"
    return "owned Office PID remained alive"


def _pid_exists(pid: int) -> bool:
    if pid <= 0:
        return False
    if os.name == "nt":
        try:
            import ctypes

            process = ctypes.windll.kernel32.OpenProcess(0x1000, False, pid)
            if not process:
                return False
            try:
                exit_code = ctypes.c_ulong()
                if not ctypes.windll.kernel32.GetExitCodeProcess(
                    process,
                    ctypes.byref(exit_code),
                ):
                    return False
                return int(exit_code.value) == 259  # STILL_ACTIVE
            finally:
                ctypes.windll.kernel32.CloseHandle(process)
        except Exception:
            return False
    try:
        os.kill(pid, 0)
    except OSError:
        return False
    return True


def _process_ids_for_names(names: Sequence[str]) -> set[int]:
    if os.name != "nt" or not names:
        return set()
    quoted_names = ",".join(
        "'" + str(name).replace("'", "''") + "'" for name in names
    )
    script = (
        f"Get-Process -Name {quoted_names} -ErrorAction SilentlyContinue | "
        "ForEach-Object { $_.Id }"
    )
    try:
        completed = run_hidden(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            capture_output=True,
            text=True,
            timeout=5,
            check=False,
        )
    except Exception:
        return set()
    result: set[int] = set()
    for line in (completed.stdout or "").splitlines():
        try:
            pid = int(line.strip())
        except ValueError:
            continue
        if pid > 0:
            result.add(pid)
    return result


def _build_probe_fixture(docx_path: Path, image_path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (800, 480), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((4, 4, 795, 475), outline="black", width=4)
    draw.text((40, 40), "LDWord layout probe", fill="black")
    image.save(image_path, format="PNG")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    first_section = doc.sections[0]
    first_section.page_width = Mm(210)
    first_section.page_height = Mm(297)
    first_section.top_margin = Mm(18)
    first_section.bottom_margin = Mm(18)
    first_section.left_margin = Mm(18)
    first_section.right_margin = Mm(18)

    doc.add_heading("Office Layout Capability Probe", level=1)
    doc.add_paragraph("COLUMN_ANCHOR_1")
    for index in range(3):
        doc.add_paragraph(f"First column filler {index}")
    break_para = doc.add_paragraph()
    break_para.add_run().add_break(WD_BREAK.COLUMN)
    doc.add_paragraph("COLUMN_ANCHOR_2")
    toc_para = doc.add_paragraph()
    _append_complex_field(
        toc_para,
        ' TOC \\o "1-2" \\h \\z \\u ',
        "Table of contents pending update",
        OxmlElement,
        qn,
    )

    second_section = doc.add_section(WD_SECTION.NEW_PAGE)
    second_section.page_width = Mm(210)
    second_section.page_height = Mm(297)
    second_section.top_margin = Mm(18)
    second_section.bottom_margin = Mm(18)
    second_section.left_margin = Mm(18)
    second_section.right_margin = Mm(18)

    # ``add_section`` moves/copies section properties.  Reacquire the wrappers;
    # mutating a wrapper retained from before the split can address the new final
    # section rather than the paragraph-level properties for the first section.
    first_section = doc.sections[0]
    second_section = doc.sections[1]
    cols = first_section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        first_section._sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "720")
    cols.set(qn("w:equalWidth"), "1")

    second_cols = second_section._sectPr.find(qn("w:cols"))
    if second_cols is None:
        second_cols = OxmlElement("w:cols")
        second_section._sectPr.append(second_cols)
    second_cols.set(qn("w:num"), "1")
    second_cols.set(qn("w:equalWidth"), "1")

    doc.add_heading("Sequential geometry", level=1)
    doc.add_paragraph("SEQUENTIAL_IMAGE_SLOT")
    for index in range(4):
        doc.add_paragraph(f"Short lead paragraph {index}")
    doc.add_paragraph("SEQUENTIAL_LATER_ANCHOR")
    doc.add_paragraph("PAGE_GEOMETRY_ANCHOR")
    numpages_para = doc.add_paragraph("Page count: ")
    _append_complex_field(
        numpages_para,
        " NUMPAGES ",
        "1",
        OxmlElement,
        qn,
    )
    for index in range(130):
        if index in {30, 70, 110}:
            doc.add_heading(f"Probe heading {index}", level=2)
        doc.add_paragraph(f"Pagination filler line {index:03d}")
    doc.save(docx_path)


def _append_complex_field(
    paragraph: Any,
    instruction: str,
    placeholder: str,
    element_factory: Any,
    qn: Any,
) -> None:
    run = paragraph.add_run()
    begin = element_factory("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = element_factory("w:instrText")
    instruction_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instruction_node.text = instruction
    separate = element_factory("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = element_factory("w:t")
    text.text = placeholder
    end = element_factory("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction_node, separate, text, end):
        run._r.append(node)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _coerce_evidence_value(value: Any) -> EvidenceValue:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


# The following functions execute only in the isolated child process.


def _safe_application_value(application: Any, name: str) -> str:
    try:
        return str(getattr(application, name, "") or "")
    except Exception:
        return ""


def _child_probe(request: Mapping[str, Any]) -> OfficeProviderReceipt:
    provider = OfficeProvider(str(request["provider"]))
    prog_id = str(request["prog_id"])
    shadow_path = Path(str(request["shadow_path"]))
    image_path = Path(str(request["image_path"]))
    pid_path = Path(str(request["pid_path"]))
    before_pids = {int(item) for item in request.get("before_pids", ())}
    started = time.perf_counter()
    capabilities: list[CapabilityReceipt] = []
    issues: list[str] = []
    application_pid: int | None = None
    application: Any = None
    document: Any = None

    try:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        create_started = time.perf_counter()
        try:
            # DispatchEx is non-negotiable: Dispatch may bind to a user's instance.
            application = win32com.client.DispatchEx(prog_id)
            capabilities.append(
                _passed(
                    OfficeCapability.APPLICATION_CREATED,
                    create_started,
                    "created via DispatchEx",
                    dispatch="DispatchEx",
                    office_name=_safe_application_value(application, "Name"),
                    office_version=_safe_application_value(application, "Version"),
                    office_build=_safe_application_value(application, "Build"),
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.APPLICATION_CREATED, create_started, exc)
            )
            raise

        try:
            application.Visible = False
        except Exception:
            pass
        try:
            application.DisplayAlerts = 0
        except Exception:
            pass
        try:
            application.AutomationSecurity = 3
        except Exception:
            pass

        open_started = time.perf_counter()
        try:
            document = application.Documents.Open(
                str(shadow_path.resolve()),
                ReadOnly=False,
                AddToRecentFiles=False,
                Visible=False,
            )
            capabilities.append(
                _passed(
                    OfficeCapability.SHADOW_DOCUMENT,
                    open_started,
                    "temporary shadow opened read/write",
                    shadow_name=shadow_path.name,
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.SHADOW_DOCUMENT, open_started, exc)
            )
            raise

        # Word does not necessarily expose a usable window handle before the
        # first document window exists.  Resolve the PID immediately after the
        # shadow opens, before any layout operation.
        pid_started = time.perf_counter()
        application_pid = _application_process_id(application, document)
        if application_pid is None:
            capabilities.append(
                _failed(
                    OfficeCapability.DEDICATED_PROCESS,
                    pid_started,
                    RuntimeError("application HWND/PID is unavailable"),
                )
            )
        else:
            _atomic_json_write(pid_path, {"pid": application_pid})
            is_new = application_pid not in before_pids
            capabilities.append(
                CapabilityReceipt(
                    capability=OfficeCapability.DEDICATED_PROCESS,
                    outcome=(
                        CapabilityOutcome.PASSED
                        if is_new
                        else CapabilityOutcome.FAILED
                    ),
                    elapsed_ms=(time.perf_counter() - pid_started) * 1000.0,
                    detail=(
                        "DispatchEx returned a new process PID"
                        if is_new
                        else "Office PID existed before the probe"
                    ),
                    evidence=(
                        ("application_pid", application_pid),
                        ("pid_was_present_before_probe", not is_new),
                    ),
                )
            )

        view_started = time.perf_counter()
        try:
            document.ActiveWindow.View.Type = 3  # wdPrintView
            view_type = int(document.ActiveWindow.View.Type)
            if view_type != 3:
                raise RuntimeError(f"unexpected view type {view_type}")
            capabilities.append(
                _passed(
                    OfficeCapability.PRINT_LAYOUT,
                    view_started,
                    "ActiveWindow is in Print Layout",
                    view_type=view_type,
                    application_visible=bool(application.Visible),
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.PRINT_LAYOUT, view_started, exc)
            )

        fields_started = time.perf_counter()
        try:
            field_count = int(document.Fields.Count)
            if field_count < 1:
                raise RuntimeError("fixture field was not detected")
            document.Fields.Update()
            capabilities.append(
                _passed(
                    OfficeCapability.FIELDS_UPDATE,
                    fields_started,
                    "document Fields.Update completed",
                    field_count=field_count,
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.FIELDS_UPDATE, fields_started, exc)
            )

        toc_started = time.perf_counter()
        try:
            toc_count = int(document.TablesOfContents.Count)
            if toc_count < 1:
                raise RuntimeError("fixture TOC was not detected")
            for index in range(1, toc_count + 1):
                toc = document.TablesOfContents.Item(index)
                toc.Update()
                toc.UpdatePageNumbers()
            capabilities.append(
                _passed(
                    OfficeCapability.TOC_UPDATE,
                    toc_started,
                    "TOC update and page-number update completed",
                    toc_count=toc_count,
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.TOC_UPDATE, toc_started, exc)
            )

        repaginate_started = time.perf_counter()
        try:
            document.Repaginate()
            page_count = int(document.ComputeStatistics(2))  # wdStatisticPages
            if page_count < 2:
                raise RuntimeError(f"unexpected page count {page_count}")
            capabilities.append(
                _passed(
                    OfficeCapability.REPAGINATE,
                    repaginate_started,
                    "Repaginate completed",
                    page_count=page_count,
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.REPAGINATE, repaginate_started, exc)
            )

        geometry_range = _find_marker(document, "PAGE_GEOMETRY_ANCHOR")
        _scroll_into_view(document, geometry_range)
        page_started = time.perf_counter()
        try:
            page_number = _range_information_int(geometry_range, 3)
            if page_number < 1:
                raise RuntimeError(f"invalid page number {page_number}")
            capabilities.append(
                _passed(
                    OfficeCapability.PAGE_NUMBER,
                    page_started,
                    "Range.Information returned a page number",
                    page_number=page_number,
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.PAGE_NUMBER, page_started, exc)
            )

        vertical_started = time.perf_counter()
        try:
            y = _range_information_float(geometry_range, 6)
            if not math.isfinite(y) or y < 0:
                raise RuntimeError(f"invalid vertical position {y}")
            capabilities.append(
                _passed(
                    OfficeCapability.VERTICAL_POSITION,
                    vertical_started,
                    "Range.Information returned page-relative Y",
                    vertical_position_pt=round(y, 3),
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.VERTICAL_POSITION, vertical_started, exc)
            )

        column_started = time.perf_counter()
        try:
            first = _column_geometry(
                document,
                _find_marker(document, "COLUMN_ANCHOR_1"),
            )
            second = _column_geometry(
                document,
                _find_marker(document, "COLUMN_ANCHOR_2"),
            )
            if first["column_count"] != 2 or second["column_count"] != 2:
                raise RuntimeError("two-column fixture was not preserved")
            if first["column_index"] != 1 or second["column_index"] != 2:
                raise RuntimeError(
                    "provider could not distinguish the first and second columns"
                )
            capabilities.append(
                _passed(
                    OfficeCapability.CURRENT_COLUMN,
                    column_started,
                    "current page column was inferred from live X geometry",
                    first_column_index=first["column_index"],
                    second_column_index=second["column_index"],
                    column_count=first["column_count"],
                )
            )
            widths = (first["container_width_pt"], second["container_width_pt"])
            if any(not math.isfinite(value) or value <= 0 for value in widths):
                raise RuntimeError(f"invalid column widths {widths}")
            capabilities.append(
                _passed(
                    OfficeCapability.CONTAINER_WIDTH,
                    column_started,
                    "live section column widths are positive",
                    first_width_pt=round(widths[0], 3),
                    second_width_pt=round(widths[1], 3),
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.CURRENT_COLUMN, column_started, exc)
            )
            capabilities.append(
                _failed(OfficeCapability.CONTAINER_WIDTH, column_started, exc)
            )

        sequence_started = time.perf_counter()
        try:
            later_range = _find_marker(document, "SEQUENTIAL_LATER_ANCHOR")
            _scroll_into_view(document, later_range)
            before_page = _range_information_int(later_range, 3)
            before_y = _range_information_float(later_range, 6)
            if before_page < 1 or before_y < 0:
                raise RuntimeError(
                    f"invalid pre-insert geometry page={before_page}, y={before_y}"
                )

            slot_range = _find_marker(document, "SEQUENTIAL_IMAGE_SLOT")
            slot_range.Text = ""
            slot_range.Collapse(1)  # wdCollapseStart
            shape = document.InlineShapes.AddPicture(
                str(image_path.resolve()),
                False,
                True,
                slot_range,
            )
            shape.LockAspectRatio = True
            shape.Width = 250.0

            # Deliberately no Repaginate here.  The later range must expose the
            # geometry caused by the preceding insertion when queried immediately.
            _scroll_into_view(document, later_range)
            after_page = _range_information_int(later_range, 3)
            after_y = _range_information_float(later_range, 6)
            if after_page < 1 or after_y < 0:
                raise RuntimeError(
                    f"invalid post-insert geometry page={after_page}, y={after_y}"
                )
            page_height = float(later_range.Sections.Item(1).PageSetup.PageHeight)
            before_absolute = (before_page - 1) * page_height + before_y
            after_absolute = (after_page - 1) * page_height + after_y
            delta = after_absolute - before_absolute
            if delta <= 20.0:
                raise RuntimeError(
                    f"later anchor did not reflect insertion immediately (delta={delta})"
                )
            capabilities.append(
                _passed(
                    OfficeCapability.SEQUENTIAL_GEOMETRY,
                    sequence_started,
                    "later anchor changed without an explicit full Repaginate",
                    before_page=before_page,
                    before_y_pt=round(before_y, 3),
                    after_page=after_page,
                    after_y_pt=round(after_y, 3),
                    absolute_delta_pt=round(delta, 3),
                    explicit_repaginate_between_queries=False,
                )
            )
        except Exception as exc:
            capabilities.append(
                _failed(OfficeCapability.SEQUENTIAL_GEOMETRY, sequence_started, exc)
            )

        try:
            document.Close(False)
        finally:
            document = None
        try:
            application.Quit()
        finally:
            application = None
        capabilities.append(
            CapabilityReceipt(
                capability=OfficeCapability.CLEAN_SHUTDOWN,
                outcome=CapabilityOutcome.PASSED,
                detail="Quit returned; parent will verify process exit",
                evidence=(("application_pid", application_pid),),
            )
        )
    except Exception as exc:
        issues.append(f"provider probe aborted: {type(exc).__name__}: {exc}")
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if application is not None:
            try:
                application.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()  # type: ignore[name-defined]
        except Exception:
            pass

    present = {item.capability for item in capabilities}
    for name in _LIVE_CAPABILITIES:
        if name not in present:
            capabilities.append(
                CapabilityReceipt(
                    capability=name,
                    outcome=CapabilityOutcome.NOT_RUN,
                    detail="provider probe aborted before this check",
                )
            )
    order = {name: index for index, name in enumerate(_LIVE_CAPABILITIES)}
    capabilities.sort(key=lambda item: order[item.capability])
    passed = _all_required_live_capabilities_pass(capabilities)
    return OfficeProviderReceipt(
        provider=provider,
        prog_id=prog_id,
        registered=True,
        outcome=(
            ProviderOutcome.QUALIFIED if passed else ProviderOutcome.DISQUALIFIED
        ),
        elapsed_ms=(time.perf_counter() - started) * 1000.0,
        application_pid=application_pid,
        capabilities=tuple(capabilities),
        issues=tuple(issues),
    )


def _passed(
    capability: OfficeCapability,
    started: float,
    detail: str,
    **evidence: EvidenceValue,
) -> CapabilityReceipt:
    return CapabilityReceipt(
        capability=capability,
        outcome=CapabilityOutcome.PASSED,
        elapsed_ms=(time.perf_counter() - started) * 1000.0,
        detail=detail,
        evidence=tuple(sorted(evidence.items())),
    )


def _failed(
    capability: OfficeCapability,
    started: float,
    error: object,
) -> CapabilityReceipt:
    return CapabilityReceipt(
        capability=capability,
        outcome=CapabilityOutcome.FAILED,
        elapsed_ms=(time.perf_counter() - started) * 1000.0,
        detail=f"{type(error).__name__}: {error}",
    )


def _application_process_id(
    application: Any,
    document: Any | None = None,
) -> int | None:
    if os.name != "nt":
        return None
    hwnd: int | None = None
    owners = [application]
    for owner in (document,):
        if owner is None:
            continue
        try:
            owners.append(owner.ActiveWindow)
        except Exception:
            pass
    try:
        owners.append(application.ActiveWindow)
    except Exception:
        pass
    for owner in owners:
        for name in ("Hwnd", "HWND"):
            try:
                candidate = int(getattr(owner, name))
            except Exception:
                continue
            if candidate > 0:
                hwnd = candidate
                break
        if hwnd is not None:
            break
    if hwnd is None:
        return None
    try:
        import ctypes
        from ctypes import wintypes

        pid = wintypes.DWORD()
        ctypes.windll.user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
        return int(pid.value) if pid.value > 0 else None
    except Exception:
        return None


def _find_marker(document: Any, marker: str) -> Any:
    target = document.Content.Duplicate
    finder = target.Find
    finder.ClearFormatting()
    found = bool(finder.Execute(FindText=marker, Forward=True, Wrap=0))
    if not found:
        raise RuntimeError(f"marker not found: {marker}")
    return target


def _scroll_into_view(document: Any, target: Any) -> None:
    try:
        document.ActiveWindow.ScrollIntoView(target, True)
    except Exception:
        pass


def _range_information_int(target: Any, constant: int) -> int:
    return int(target.Information(constant))


def _range_information_float(target: Any, constant: int) -> float:
    return float(target.Information(constant))


def _column_geometry(document: Any, target: Any) -> dict[str, float | int]:
    _scroll_into_view(document, target)
    x = _range_information_float(target, 5)  # relative to page
    if not math.isfinite(x) or x < 0:
        raise RuntimeError(f"invalid horizontal position {x}")
    section = target.Sections.Item(1)
    setup = section.PageSetup
    columns = setup.TextColumns
    count = int(columns.Count)
    if count < 1:
        raise RuntimeError(f"invalid TextColumns.Count {count}")
    left = float(setup.LeftMargin)
    starts: list[float] = []
    widths: list[float] = []
    cursor = left
    for index in range(1, count + 1):
        column = columns.Item(index)
        width = float(column.Width)
        if not math.isfinite(width) or width <= 0:
            raise RuntimeError(f"invalid column width {width}")
        starts.append(cursor)
        widths.append(width)
        try:
            space_after = float(column.SpaceAfter)
        except Exception:
            space_after = 0.0
        cursor += width + max(0.0, space_after)

    best_index = min(
        range(count),
        key=lambda index: abs(x - starts[index]),
    )
    for index, (start, width) in enumerate(zip(starts, widths, strict=True)):
        if start - 1.0 <= x <= start + width + 1.0:
            best_index = index
            break
    return {
        "column_index": best_index + 1,
        "column_count": count,
        "container_width_pt": widths[best_index],
        "horizontal_position_pt": x,
    }


def _atomic_json_write(path: Path, data: Mapping[str, Any]) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    os.replace(temporary, path)


def _child_main(request_path: Path) -> int:
    request = json.loads(request_path.read_text(encoding="utf-8"))
    result_path = Path(str(request["result_path"]))
    try:
        receipt = _child_probe(request)
        _atomic_json_write(result_path, receipt.to_dict())
        return 0
    except Exception as exc:
        provider = OfficeProvider(str(request["provider"]))
        failed = OfficeProviderReceipt(
            provider=provider,
            prog_id=str(request["prog_id"]),
            registered=True,
            outcome=ProviderOutcome.CHILD_FAILED,
            elapsed_ms=0.0,
            capabilities=_not_run_capabilities(str(exc)),
            issues=(f"unhandled child error: {type(exc).__name__}: {exc}",),
        )
        _atomic_json_write(result_path, failed.to_dict())
        return 1


def run_office_layout_probe_child(request_path: str | Path) -> int:
    """Run the internal broker child before any GUI initialization."""

    return _child_main(Path(request_path))


def _module_main(argv: Sequence[str]) -> int:
    if len(argv) == 2 and argv[0] == "--child-request":
        return _child_main(Path(argv[1]))
    raise SystemExit("office_layout_probe is an internal module; use the probe script")


if __name__ == "__main__":
    raise SystemExit(_module_main(sys.argv[1:]))


__all__ = [
    "CapabilityOutcome",
    "CapabilityReceipt",
    "DEFAULT_PROVIDER_SPECS",
    "OfficeCapability",
    "OfficeLayoutProbeReceipt",
    "OfficeProvider",
    "OfficeProviderReceipt",
    "OfficeProviderSpec",
    "PROBE_SCHEMA_VERSION",
    "ProviderOutcome",
    "probe_office_layout",
    "run_office_layout_probe_child",
]
