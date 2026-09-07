# -*- coding: utf-8 -*-
"""Turn an attached engineering sample document into per-chapter knowledge
samples used by the authoring pipeline.

The AI assistant writes engineering documents chapter by chapter.  When a
real-world sample (docx/doc/wps) is attached, the pipeline already learns its
table of contents (chapter titles).  This module additionally slices the
sample's *body text* chapter by chapter so the per-chapter prompt can carry a
budget-capped excerpt of the real prose (reference sample), letting the model
absorb the trade's wording and structure instead of only the headings.
"""
from __future__ import annotations

from pathlib import Path

_MAX_SAMPLE_CHARS = 6000


def extract_knowledge_samples(
    source_path: str | Path,
    *,
    max_total_chars: int = _MAX_SAMPLE_CHARS,
) -> tuple[dict[str, object], ...]:
    """Slice one document into (source, chapter_title, content) samples.

    If the document has detectable top-level chapters, returns one sample per
    chapter (body prose only, headings excluded).  Otherwise returns a single
    whole-document sample.  Content is capped so the sum stays below
    ``max_total_chars``; a text-only fallback is used for unreadable files.
    """
    from src.assistant.application.chapter_document_editor import (
        ChapterEditError,
        chapter_body_text,
        detect_chapter_outline,
    )

    path = Path(str(source_path or "")).expanduser()
    if not path.is_file():
        return ()
    try:
        chapters = detect_chapter_outline(path)
    except (ChapterEditError, OSError, ValueError):
        chapters = ()

    samples: list[dict[str, object]] = []
    used = 0
    budget = int(max_total_chars)
    source = path.name
    if chapters:
        for chapter in chapters:
            if used >= budget:
                break
            try:
                body = chapter_body_text(path, chapter).strip()
            except (ChapterEditError, OSError, ValueError):
                body = ""
            if not body:
                continue
            remaining = budget - used
            piece = body[:remaining]
            samples.append(
                {
                    "source": source,
                    "chapter_title": chapter.title,
                    "content": piece,
                }
            )
            used += len(piece)
        if samples:
            return tuple(samples)
    # Whole-document fallback.
    try:
        text = _read_whole_document(path).strip()
    except (OSError, ValueError):
        return ()
    if not text:
        return ()
    return (
        {
            "source": source,
            "chapter_title": "",
            "content": text[:budget],
        },
    )


def _read_whole_document(path: Path) -> str:
    suffix = path.suffix.casefold()
    if suffix in {".md", ".markdown"}:
        return path.read_text(encoding="utf-8-sig")
    from docx import Document

    if suffix in {".doc", ".wps"}:
        from src.services.legacy_word_import import ensure_editable_docx

        path = ensure_editable_docx(path).docx_path
    document = Document(str(path))
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            chunks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                chunks.append("\t".join(cells))
    return "\n".join(chunks)


__all__ = ["extract_knowledge_samples"]
