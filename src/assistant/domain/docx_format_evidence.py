"""Deterministic DOCX format evidence for reference-sample analysis.

The model may summarize this evidence, but it must never infer Word geometry
from extracted body text.  This module deliberately returns plain, bounded
data without local paths so it is safe to place behind the existing disclosure
grant.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Iterable, Mapping
import hashlib
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

from src.assistant.contracts.task_plan import (
    SOURCE_ROLE_PRODUCTION_INPUT,
    SOURCE_ROLE_REFERENCE_MATERIAL,
    SOURCE_ROLE_STANDARD_FORMAT_REFERENCE,
    SOURCE_ROLE_STRUCTURED_SOURCE,
)

DOCX_FORMAT_EVIDENCE_SCHEMA_VERSION = "docx-format-evidence-v1"
STANDARD_FORMAT_REFERENCE_ROLE = SOURCE_ROLE_STANDARD_FORMAT_REFERENCE
FORMAT_EVIDENCE_DISCLOSURE_FIELD = "document_format_evidence"

_FORMAT_REQUIREMENT_TERMS = (
    "格式要求",
    "格式规则",
    "格式规范",
    "样式要求",
    "样式规则",
    "样式规范",
    "版式要求",
    "版式规则",
    "版式规范",
    "排版要求",
    "排版规则",
    "排版规范",
)
_FORMAT_ANALYSIS_TERMS = (
    "确定格式",
    "确认格式",
    "分析格式",
    "识别格式",
    "提取格式",
    "看看格式",
    "检查格式",
)
_REFERENCE_TERMS = (
    "标准",
    "样稿",
    "范例",
    "参考",
    "模板",
    "这份",
    "这个",
    "文件",
    "文档",
)
_MUTATION_TERMS = (
    "修复",
    "修改",
    "调整",
    "改成",
    "统一格式",
    "套用",
    "应用",
    "按此格式",
    "按照这个格式",
    "套用到",
    "应用到",
    "处理另一",
    "修改另一",
    "格式化另一",
    "统一另一",
)
_TARGET_APPLICATION_TERMS = (
    "另一份",
    "另一个",
    "其他文档",
    "其它文档",
    "目标文档",
    "目标文件",
    "套用到",
    "应用到",
    "处理另一",
    "修改另一",
    "格式化另一",
    "统一另一",
)
_MAX_STYLE_ROWS = 24
_MAX_SECTION_ROWS = 12
_MAX_TABLE_STYLE_ROWS = 12
_MAX_SAMPLE_TEXT = 120


def is_format_requirements_request(query: str) -> bool:
    """Return whether the user asks to inspect a reference file's format rules."""

    normalized = _normalize_query(query)
    if not normalized:
        return False
    if any(term in normalized for term in _MUTATION_TERMS):
        return is_format_reference_application_request(normalized)
    has_requirement = any(term in normalized for term in _FORMAT_REQUIREMENT_TERMS)
    has_analysis = any(term in normalized for term in _FORMAT_ANALYSIS_TERMS)
    has_reference = any(term in normalized for term in _REFERENCE_TERMS)
    return bool(has_analysis or (has_requirement and has_reference))


def is_format_reference_application_request(query: str) -> bool:
    """Identify a sample-analysis request whose production target is absent."""

    normalized = _normalize_query(query)
    if not normalized:
        return False
    has_mutation = any(term in normalized for term in _MUTATION_TERMS)
    has_target = any(
        term in normalized for term in _TARGET_APPLICATION_TERMS
    )
    has_requirement = any(
        term in normalized for term in _FORMAT_REQUIREMENT_TERMS
    )
    has_analysis = any(
        term in normalized for term in _FORMAT_ANALYSIS_TERMS
    )
    has_reference = any(term in normalized for term in _REFERENCE_TERMS)
    return bool(
        has_mutation
        and has_target
        and has_reference
        and (has_requirement or has_analysis)
    )


def bind_attachment_semantic_roles(
    refs: Iterable[Mapping[str, object]],
    query: str,
) -> tuple[dict[str, object], ...]:
    """Bind per-turn semantic roles without leaking intent into later turns."""

    reference_request = is_format_requirements_request(query)
    rows = [dict(value) for value in refs]
    has_explicit_reference = any(
        str(row.get("semantic_role") or "") == STANDARD_FORMAT_REFERENCE_ROLE
        and row.get("semantic_role_source") != "assistant_intent"
        for row in rows
    )
    infer_single_reference = (
        reference_request
        and len(rows) == 1
        and not has_explicit_reference
    )
    for row in rows:
        if infer_single_reference:
            row["semantic_role"] = STANDARD_FORMAT_REFERENCE_ROLE
            row["semantic_role_source"] = "assistant_intent"
            if is_format_reference_application_request(query):
                row["target_attachment_required"] = True
        elif (
            not reference_request or len(rows) != 1
        ) and row.get("semantic_role_source") == "assistant_intent":
            row.pop("semantic_role", None)
            row.pop("semantic_role_source", None)
            row.pop("target_attachment_required", None)
    return tuple(rows)


def attachment_disclosure_fields(
    refs: Iterable[Mapping[str, object]],
) -> tuple[str, ...]:
    fields: list[str] = []
    for item in refs:
        role = str(item.get("semantic_role") or "")
        if role == STANDARD_FORMAT_REFERENCE_ROLE:
            fields.append(FORMAT_EVIDENCE_DISCLOSURE_FIELD)
        elif role in {
            SOURCE_ROLE_PRODUCTION_INPUT,
            SOURCE_ROLE_STRUCTURED_SOURCE,
        }:
            continue
        elif role in {"", SOURCE_ROLE_REFERENCE_MATERIAL}:
            fields.append("document_text")
    return tuple(dict.fromkeys(fields))


def extract_docx_format_evidence(path: str | Path) -> dict[str, object]:
    """Extract bounded, auditable formatting evidence from one DOCX."""

    source = Path(path).expanduser()
    if not source.is_file():
        raise FileNotFoundError("Format reference does not exist")
    if source.suffix.casefold() != ".docx":
        raise ValueError("Only DOCX format references are supported")

    document = Document(str(source))
    paragraphs = tuple(_iter_all_paragraphs(document))
    style_counts = Counter(
        str(getattr(paragraph.style, "style_id", "") or "")
        for paragraph in paragraphs
        if getattr(paragraph, "style", None) is not None
    )
    style_rows = _style_evidence_rows(document, style_counts)
    section_rows = tuple(
        _section_evidence(index, section)
        for index, section in enumerate(tuple(document.sections)[:_MAX_SECTION_ROWS])
    )
    table_style_counts = Counter(
        str(getattr(table.style, "name", "") or "无显式表格样式")
        for table in document.tables
    )
    direct_paragraph_count = sum(
        1 for paragraph in paragraphs if _has_direct_paragraph_format(paragraph)
    )
    direct_run_count = sum(
        1
        for paragraph in paragraphs
        for run in paragraph.runs
        if _has_direct_run_format(run)
    )
    numbered_paragraph_count = sum(
        1 for paragraph in paragraphs if _effective_paragraph_numbering(paragraph)
    )
    direct_by_style = Counter(
        str(getattr(paragraph.style, "name", "") or "无样式")
        for paragraph in paragraphs
        if _has_direct_paragraph_format(paragraph)
    )
    settings = document.settings.element
    document_element = document.element
    return {
        "schema_version": DOCX_FORMAT_EVIDENCE_SCHEMA_VERSION,
        "source": {
            "name": source.name,
            "sha256": _sha256(source),
            "media_type": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            "semantic_role": STANDARD_FORMAT_REFERENCE_ROLE,
        },
        "inventory": {
            "paragraph_count": len(paragraphs),
            "body_paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
            "section_count": len(document.sections),
            "used_paragraph_style_count": len(style_counts),
            "numbered_paragraph_count": numbered_paragraph_count,
            "direct_paragraph_format_count": direct_paragraph_count,
            "direct_run_format_count": direct_run_count,
            "inline_shape_count": len(document.inline_shapes),
            "drawing_count": sum(1 for _ in document_element.iter(qn("w:drawing"))),
            "formula_count": sum(1 for _ in document_element.iter(qn("m:oMath"))),
            "field_instruction_count": sum(
                1 for _ in document_element.iter(qn("w:instrText"))
            ),
            "tracked_insertion_count": sum(
                1 for _ in document_element.iter(qn("w:ins"))
            ),
            "tracked_deletion_count": sum(
                1 for _ in document_element.iter(qn("w:del"))
            ),
        },
        "document_settings": {
            "even_and_odd_headers": bool(settings.find(qn("w:evenAndOddHeaders")) is not None),
            "track_revisions": bool(settings.find(qn("w:trackRevisions")) is not None),
            "default_tab_stop_twips": _xml_int(
                settings.find(qn("w:defaultTabStop")),
                "w:val",
            ),
        },
        "sections": list(section_rows),
        "paragraph_styles": list(style_rows),
        "table_styles": [
            {"name": name, "usage_count": count}
            for name, count in table_style_counts.most_common(_MAX_TABLE_STYLE_ROWS)
        ],
        "table_format_profiles": list(_table_format_profiles(document)),
        "numbering_profiles": list(_numbering_profiles(document, paragraphs)),
        "direct_formatting": {
            "paragraphs_by_style": [
                {"style": name, "count": count}
                for name, count in direct_by_style.most_common(_MAX_STYLE_ROWS)
            ],
            "interpretation": (
                "这些计数表示样式之外还存在直接段落或字符格式；"
                "在提升为标准规则前需要检查其是否为稳定重复模式。"
            ),
        },
        "evidence_policy": {
            "explicit_values_are_authoritative": True,
            "null_means_inherited_or_not_declared": True,
            "body_text_is_not_format_evidence": True,
            "limitations": [
                "主题字体、域代码、文本框和绘图对象可能需要 Word 渲染复核。",
                "直接格式覆盖与样式基线并存时，应先确认哪一项属于标准规则。",
                "只观察到一次的局部样式不能自动提升为全局格式要求。",
            ],
        },
    }


def format_requirements_system_instruction() -> str:
    return (
        "\n\n当用户把 DOCX 声明为标准文件、参考样稿并要求确定格式要求时，"
        "你正在执行只读的“格式规范分析”，不是格式化或生成文档。"
        "必须以用户材料消息中 <document_format_evidence> 的结构化证据为准；"
        "正文文字本身不能证明字体、字号、行距、页边距或编号规则。"
        "请按“已确认的格式要求 / 局部例外或冲突 / 仍需用户确认 / 可执行下一步”"
        "组织回答。不得补猜证据中为 null 的属性，也不得声称已经保存模板或修改文件。"
        "如果同一请求还要求把格式套用到另一份文档，而本轮只提供了一份标准样稿，"
        "只完成样稿分析并明确要求用户另行提供目标文档；不得把标准样稿当作生产输入。"
        "如果没有 document_format_evidence，明确要求用户添加 DOCX 标准样稿。"
    )


def _iter_all_paragraphs(document) -> Iterable[Any]:
    seen: set[int] = set()

    def unique(values: Iterable[Any]) -> Iterable[Any]:
        for paragraph in values:
            identity = id(paragraph._p)
            if identity in seen:
                continue
            seen.add(identity)
            yield paragraph

    yield from unique(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from unique(cell.paragraphs)
    for section in document.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            yield from unique(container.paragraphs)


def _style_evidence_rows(document, counts: Counter[str]) -> tuple[dict[str, object], ...]:
    style_map = {
        str(getattr(style, "style_id", "") or ""): style
        for style in document.styles
        if style.type == WD_STYLE_TYPE.PARAGRAPH
    }
    rows: list[dict[str, object]] = []
    for style_id, usage_count in counts.most_common(_MAX_STYLE_ROWS):
        style = style_map.get(style_id)
        if style is None:
            continue
        chain = tuple(_style_chain(style))
        font_values, font_sources = _effective_font(chain)
        paragraph_values, paragraph_sources = _effective_paragraph(chain)
        rows.append(
            {
                "style_id": style_id,
                "name": str(getattr(style, "name", "") or style_id),
                "usage_count": usage_count,
                "based_on": str(
                    getattr(getattr(style, "base_style", None), "name", "") or ""
                ),
                "semantic_level": _semantic_style_level(style),
                "font": font_values,
                "paragraph": paragraph_values,
                "numbering": _style_numbering(style),
                "value_sources": {
                    "font": font_sources,
                    "paragraph": paragraph_sources,
                },
            }
        )
    return tuple(rows)


def _style_chain(style) -> Iterable[Any]:
    seen: set[str] = set()
    current = style
    while current is not None:
        identity = str(getattr(current, "style_id", "") or id(current))
        if identity in seen:
            return
        seen.add(identity)
        yield current
        current = getattr(current, "base_style", None)


def _effective_font(chain: tuple[Any, ...]) -> tuple[dict[str, object], dict[str, str]]:
    readers = {
        "ascii_font": lambda style: _style_font_attr(style, "w:ascii"),
        "east_asia_font": lambda style: _style_font_attr(style, "w:eastAsia"),
        "ascii_theme": lambda style: _style_font_attr(style, "w:asciiTheme"),
        "east_asia_theme": lambda style: _style_font_attr(style, "w:eastAsiaTheme"),
        "size_pt": lambda style: _length_pt(getattr(style.font, "size", None)),
        "bold": lambda style: getattr(style.font, "bold", None),
        "italic": lambda style: getattr(style.font, "italic", None),
        "underline": lambda style: _enum_text(getattr(style.font, "underline", None)),
        "color": lambda style: _font_color(style),
    }
    return _resolve_chain_values(chain, readers)


def _effective_paragraph(
    chain: tuple[Any, ...],
) -> tuple[dict[str, object], dict[str, str]]:
    readers = {
        "alignment": lambda style: _enum_text(style.paragraph_format.alignment),
        "line_spacing": lambda style: _line_spacing(style.paragraph_format),
        "space_before_pt": lambda style: _length_pt(style.paragraph_format.space_before),
        "space_after_pt": lambda style: _length_pt(style.paragraph_format.space_after),
        "left_indent_cm": lambda style: _length_cm(style.paragraph_format.left_indent),
        "right_indent_cm": lambda style: _length_cm(style.paragraph_format.right_indent),
        "first_line_indent_cm": lambda style: _length_cm(
            style.paragraph_format.first_line_indent
        ),
        "keep_with_next": lambda style: style.paragraph_format.keep_with_next,
        "keep_together": lambda style: style.paragraph_format.keep_together,
        "page_break_before": lambda style: style.paragraph_format.page_break_before,
        "widow_control": lambda style: style.paragraph_format.widow_control,
    }
    return _resolve_chain_values(chain, readers)


def _resolve_chain_values(
    chain: tuple[Any, ...],
    readers: Mapping[str, Any],
) -> tuple[dict[str, object], dict[str, str]]:
    values: dict[str, object] = {}
    sources: dict[str, str] = {}
    for key, reader in readers.items():
        value = None
        source = ""
        for style in chain:
            candidate = reader(style)
            if candidate is not None and candidate != "":
                value = candidate
                source = str(getattr(style, "name", "") or getattr(style, "style_id", ""))
                break
        values[key] = value
        if source:
            sources[key] = source
    return values, sources


def _section_evidence(index: int, section) -> dict[str, object]:
    orientation = "landscape" if section.orientation == WD_ORIENT.LANDSCAPE else "portrait"
    sect_pr = section._sectPr
    columns = sect_pr.find(qn("w:cols"))
    column_count = _xml_int(columns, "w:num") or 1
    return {
        "index": index + 1,
        "orientation": orientation,
        "page_width_cm": _length_cm(section.page_width),
        "page_height_cm": _length_cm(section.page_height),
        "margins_cm": {
            "top": _length_cm(section.top_margin),
            "right": _length_cm(section.right_margin),
            "bottom": _length_cm(section.bottom_margin),
            "left": _length_cm(section.left_margin),
            "gutter": _length_cm(section.gutter),
        },
        "header_distance_cm": _length_cm(section.header_distance),
        "footer_distance_cm": _length_cm(section.footer_distance),
        "different_first_page": bool(section.different_first_page_header_footer),
        "column_count": column_count,
        "start_type": _enum_text(section.start_type),
        "headers": {
            "default": _container_evidence(section.header),
            "first": _container_evidence(section.first_page_header),
            "even": _container_evidence(section.even_page_header),
        },
        "footers": {
            "default": _container_evidence(section.footer),
            "first": _container_evidence(section.first_page_footer),
            "even": _container_evidence(section.even_page_footer),
        },
    }


def _container_evidence(container) -> dict[str, object]:
    text = " ".join(
        paragraph.text.strip()
        for paragraph in container.paragraphs
        if paragraph.text.strip()
    )
    return {
        "linked_to_previous": bool(container.is_linked_to_previous),
        "has_text": bool(text),
        "text_sample": text[:_MAX_SAMPLE_TEXT],
        "paragraph_count": len(container.paragraphs),
        "field_instructions": list(
            dict.fromkeys(
                str(node.text or "").strip()
                for node in container._element.iter(qn("w:instrText"))
                if str(node.text or "").strip()
            )
        )[:8],
    }


def _semantic_style_level(style) -> int | None:
    name = str(getattr(style, "name", "") or "")
    match = re.search(r"(?:heading|标题)\s*([1-9])", name, flags=re.IGNORECASE)
    if match:
        return int(match.group(1))
    p_pr = getattr(getattr(style, "element", None), "pPr", None)
    outline = p_pr.find(qn("w:outlineLvl")) if p_pr is not None else None
    value = _xml_int(outline, "w:val")
    return value + 1 if value is not None and 0 <= value <= 8 else None


def _style_numbering(style) -> dict[str, int] | None:
    p_pr = getattr(getattr(style, "element", None), "pPr", None)
    num_pr = p_pr.find(qn("w:numPr")) if p_pr is not None else None
    if num_pr is None:
        return None
    num_id = _xml_int(num_pr.find(qn("w:numId")), "w:val")
    level = _xml_int(num_pr.find(qn("w:ilvl")), "w:val")
    if num_id is None and level is None:
        return None
    return {"num_id": num_id if num_id is not None else -1, "level": level or 0}


def _paragraph_numbering(paragraph) -> dict[str, int] | None:
    p_pr = paragraph._p.pPr
    num_pr = p_pr.find(qn("w:numPr")) if p_pr is not None else None
    if num_pr is None:
        return None
    num_id = _xml_int(num_pr.find(qn("w:numId")), "w:val")
    level = _xml_int(num_pr.find(qn("w:ilvl")), "w:val")
    return {"num_id": num_id if num_id is not None else -1, "level": level or 0}


def _effective_paragraph_numbering(paragraph) -> dict[str, int] | None:
    direct = _paragraph_numbering(paragraph)
    if direct is not None:
        return direct
    style = getattr(paragraph, "style", None)
    if style is None:
        return None
    for inherited in _style_chain(style):
        numbering = _style_numbering(inherited)
        if numbering is not None:
            return numbering
    return None


def _numbering_profiles(document, paragraphs: tuple[Any, ...]) -> tuple[dict[str, object], ...]:
    usage = Counter(
        (
            int(numbering["num_id"]),
            int(numbering["level"]),
        )
        for paragraph in paragraphs
        if (numbering := _effective_paragraph_numbering(paragraph)) is not None
        and int(numbering["num_id"]) >= 0
    )
    if not usage:
        return ()
    try:
        numbering_root = document.part.numbering_part.element
    except (AttributeError, KeyError):
        return tuple(
            {
                "num_id": num_id,
                "level": level,
                "usage_count": count,
            }
            for (num_id, level), count in usage.most_common(24)
        )
    num_to_abstract: dict[int, int] = {}
    for node in numbering_root.findall(qn("w:num")):
        num_id = _xml_int(node, "w:numId")
        abstract = _xml_int(node.find(qn("w:abstractNumId")), "w:val")
        if num_id is not None and abstract is not None:
            num_to_abstract[num_id] = abstract
    abstract_map: dict[int, Any] = {}
    for node in numbering_root.findall(qn("w:abstractNum")):
        abstract_id = _xml_int(node, "w:abstractNumId")
        if abstract_id is not None:
            abstract_map[abstract_id] = node
    rows: list[dict[str, object]] = []
    for (num_id, level), count in usage.most_common(24):
        abstract_id = num_to_abstract.get(num_id)
        abstract = abstract_map.get(abstract_id)
        level_node = None
        if abstract is not None:
            level_node = next(
                (
                    node
                    for node in abstract.findall(qn("w:lvl"))
                    if _xml_int(node, "w:ilvl") == level
                ),
                None,
            )
        rows.append(
            {
                "num_id": num_id,
                "abstract_num_id": abstract_id,
                "level": level,
                "usage_count": count,
                "format": _xml_text(level_node, "w:numFmt", "w:val"),
                "level_text": _xml_text(level_node, "w:lvlText", "w:val"),
                "start": _xml_int(
                    level_node.find(qn("w:start")) if level_node is not None else None,
                    "w:val",
                ),
                "suffix": _xml_text(level_node, "w:suff", "w:val"),
            }
        )
    return tuple(rows)


def _table_format_profiles(document) -> tuple[dict[str, object], ...]:
    profiles: Counter[str] = Counter()
    payloads: dict[str, dict[str, object]] = {}
    for table in document.tables:
        payload = _table_format_payload(table)
        identity = repr(payload)
        profiles[identity] += 1
        payloads[identity] = payload
    rows: list[dict[str, object]] = []
    for identity, count in profiles.most_common(_MAX_TABLE_STYLE_ROWS):
        rows.append({**payloads[identity], "usage_count": count})
    return tuple(rows)


def _table_format_payload(table) -> dict[str, object]:
    table_pr = table._tbl.tblPr
    table_width = table_pr.find(qn("w:tblW"))
    width_type = (
        str(table_width.get(qn("w:type")) or "")
        if table_width is not None
        else ""
    )
    width_value = _xml_int(table_width, "w:w")
    layout = table_pr.find(qn("w:tblLayout"))
    look = table_pr.find(qn("w:tblLook"))
    borders_node = table_pr.find(qn("w:tblBorders"))
    border_payload: dict[str, dict[str, object]] = {}
    if borders_node is not None:
        for edge in ("top", "right", "bottom", "left", "insideH", "insideV"):
            node = borders_node.find(qn(f"w:{edge}"))
            if node is None:
                continue
            border_payload[edge] = {
                "style": str(node.get(qn("w:val")) or ""),
                "size_eighth_pt": _xml_int(node, "w:sz"),
                "color": str(node.get(qn("w:color")) or ""),
            }
    repeat_header_rows = 0
    cant_split_rows = 0
    for row in table.rows:
        tr_pr = row._tr.trPr
        if tr_pr is None:
            continue
        if tr_pr.find(qn("w:tblHeader")) is not None:
            repeat_header_rows += 1
        if tr_pr.find(qn("w:cantSplit")) is not None:
            cant_split_rows += 1
    return {
        "style": str(getattr(table.style, "name", "") or "无显式表格样式"),
        "row_count": len(table.rows),
        "column_count": len(table.columns),
        "alignment": _enum_text(table.alignment),
        "autofit": bool(table.autofit),
        "layout": (
            str(layout.get(qn("w:type")) or "")
            if layout is not None
            else ""
        ),
        "width": {
            "type": width_type,
            "value": width_value,
            "cm": (
                round(float(width_value) * 2.54 / 1440, 3)
                if width_value is not None and width_type == "dxa"
                else None
            ),
        },
        "look": (
            str(look.get(qn("w:val")) or "")
            if look is not None
            else ""
        ),
        "repeat_header_row_count": repeat_header_rows,
        "cant_split_row_count": cant_split_rows,
        "borders": border_payload,
    }


def _has_direct_paragraph_format(paragraph) -> bool:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return False
    ignored = {qn("w:pStyle"), qn("w:sectPr")}
    return any(child.tag not in ignored for child in p_pr)


def _has_direct_run_format(run) -> bool:
    r_pr = run._r.rPr
    return bool(r_pr is not None and len(r_pr))


def _style_font_attr(style, attribute: str) -> str | None:
    r_pr = getattr(getattr(style, "element", None), "rPr", None)
    fonts = r_pr.find(qn("w:rFonts")) if r_pr is not None else None
    return str(fonts.get(qn(attribute)) or "") or None if fonts is not None else None


def _font_color(style) -> str | None:
    color = getattr(style.font, "color", None)
    rgb = getattr(color, "rgb", None)
    if rgb is not None:
        return str(rgb)
    theme = getattr(color, "theme_color", None)
    return _enum_text(theme)


def _line_spacing(paragraph_format) -> dict[str, object] | None:
    value = paragraph_format.line_spacing
    rule = _enum_text(paragraph_format.line_spacing_rule)
    if value is None and not rule:
        return None
    if hasattr(value, "pt"):
        rendered: object = {"kind": "points", "value": round(float(value.pt), 3)}
    elif value is None:
        rendered = None
    else:
        rendered = {"kind": "multiple", "value": round(float(value), 3)}
    return {"rule": rule, "value": rendered}


def _length_pt(value) -> float | None:
    return round(float(value.pt), 3) if value is not None else None


def _length_cm(value) -> float | None:
    return round(float(value.cm), 3) if value is not None else None


def _enum_text(value) -> str | None:
    if value is None:
        return None
    name = getattr(value, "name", None)
    return str(name).casefold() if name else str(value)


def _xml_int(element, attribute: str) -> int | None:
    if element is None:
        return None
    raw = element.get(qn(attribute))
    try:
        return int(raw) if raw is not None else None
    except (TypeError, ValueError):
        return None


def _xml_text(parent, child_name: str, attribute: str) -> str | None:
    if parent is None:
        return None
    child = parent.find(qn(child_name))
    if child is None:
        return None
    value = str(child.get(qn(attribute)) or "").strip()
    return value or None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalize_query(value: str) -> str:
    return re.sub(
        r"[\s\-_/\\:：,，.。;；、!！?？()（）\[\]【】{}<>《》]+",
        "",
        str(value or "").casefold(),
    )


__all__ = [
    "DOCX_FORMAT_EVIDENCE_SCHEMA_VERSION",
    "FORMAT_EVIDENCE_DISCLOSURE_FIELD",
    "STANDARD_FORMAT_REFERENCE_ROLE",
    "attachment_disclosure_fields",
    "bind_attachment_semantic_roles",
    "extract_docx_format_evidence",
    "format_requirements_system_instruction",
    "is_format_reference_application_request",
    "is_format_requirements_request",
]
