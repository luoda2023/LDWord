"""Compile AI Markdown into DocumentFragment and compose a local draft DOCX."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass, replace
from hashlib import sha256
import os
from pathlib import Path
import re
from typing import TypeAlias
from uuid import uuid4

from docx import Document

from src.config.content_artifacts import ContentMaterialBinding
from src.config.content_materials import ContentInsertionRule, content_anchor_token
from src.services.material_content.artifact_repository import ContentArtifactRepository
from src.services.material_content.compiler import compile_content_material
from src.services.material_content.composer import (
    ContentComposeInputs,
    ContentComposeRequest,
    ContentMaterialComposer,
)
from src.assistant.domain.exam_authoring_contract import (
    generated_exam_blockers,
    generated_exam_warnings,
    resolve_exam_blueprint,
)
from src.assistant.contracts.task_plan import (
    ARTIFACT_KIND_EXAM,
    ARTIFACT_KIND_NARRATIVE,
    ARTIFACT_KIND_OFFICIAL,
    SOURCE_ROLE_PRODUCTION_INPUT,
    SOURCE_ROLE_STRUCTURED_SOURCE,
    SourceArtifactRef,
)
from src.shared.engine.exam_question_schema import parse_exam_markdown_source
from src.shared.engine.material_token_contract import (
    MaterialTokenNamespace,
    material_token,
)
from src.services.official_draft_source import (
    OFFICIAL_DRAFT_SCHEMA_ID,
    OfficialDraftSource,
    compile_official_draft_text,
    complete_official_draft_field,
    load_official_draft_source,
    write_official_draft_source,
)


BIDDING_DOCUMENT_PROFILE_ID = "assistant.bidding-local-assembly.v1"
_BIDDING_PROMPT_PROFILE_ID = "assistant.bidding-markdown.v1"


@dataclass(frozen=True, slots=True)
class GeneratedContentDraft:
    draft_id: str
    session_id: str
    markdown_path: str
    document_path: str
    fragment: dict[str, object]
    fragment_digest: str
    artifact_id: str
    manifest_sha256: str
    compose_receipt_id: str
    document_profile_id: str = ""

    @property
    def artifact_kind(self) -> str:
        return ARTIFACT_KIND_NARRATIVE

    @property
    def production_input_path(self) -> str:
        return self.document_path

    @property
    def preview_path(self) -> str:
        return self.document_path

    def to_dict(self) -> dict[str, object]:
        return {
            "contract_kind": "assistant_generated_content_draft",
            "artifact_kind": self.artifact_kind,
            "draft_id": self.draft_id,
            "session_id": self.session_id,
            "markdown_path": self.markdown_path,
            "document_path": self.document_path,
            "fragment": dict(self.fragment),
            "fragment_digest": self.fragment_digest,
            "artifact_id": self.artifact_id,
            "manifest_sha256": self.manifest_sha256,
            "compose_receipt_id": self.compose_receipt_id,
            "document_profile_id": self.document_profile_id,
        }


@dataclass(frozen=True, slots=True)
class GeneratedExamDraft:
    """Validated exam Markdown retained as the authoritative production input."""

    draft_id: str
    session_id: str
    markdown_path: str
    source_digest: str
    artifact_id: str
    schema_id: str
    validation_summary: dict[str, object]
    issues: tuple[dict[str, object], ...] = ()

    @property
    def artifact_kind(self) -> str:
        return ARTIFACT_KIND_EXAM

    @property
    def production_input_path(self) -> str:
        return self.markdown_path

    @property
    def preview_path(self) -> str:
        return self.markdown_path

    def to_dict(self) -> dict[str, object]:
        return {
            "contract_kind": "assistant_generated_exam_draft",
            "artifact_kind": self.artifact_kind,
            "draft_id": self.draft_id,
            "session_id": self.session_id,
            "markdown_path": self.markdown_path,
            "source_digest": self.source_digest,
            "artifact_id": self.artifact_id,
            "schema_id": self.schema_id,
            "validation_summary": dict(self.validation_summary),
            "issues": [dict(item) for item in self.issues],
        }


@dataclass(frozen=True, slots=True)
class GeneratedOfficialDraft:
    """Validated public-document fields retained as the production source."""

    draft_id: str
    session_id: str
    source_path: str
    preview_docx_path: str
    source_digest: str
    artifact_id: str
    schema_id: str
    document_type_id: str
    field_values: Mapping[str, str]
    field_provenance: Mapping[str, str]
    missing_user_fields: tuple[str, ...] = ()
    provisional_fields: tuple[str, ...] = ()
    warnings: tuple[str, ...] = ()

    @property
    def artifact_kind(self) -> str:
        return ARTIFACT_KIND_OFFICIAL

    @property
    def production_input_path(self) -> str:
        return self.source_path

    @property
    def preview_path(self) -> str:
        return self.preview_docx_path

    def to_dict(self) -> dict[str, object]:
        return {
            "contract_kind": "assistant_generated_official_draft",
            "artifact_kind": self.artifact_kind,
            "draft_id": self.draft_id,
            "session_id": self.session_id,
            "source_path": self.source_path,
            "preview_docx_path": self.preview_docx_path,
            "source_digest": self.source_digest,
            "artifact_id": self.artifact_id,
            "schema_id": self.schema_id,
            "document_type_id": self.document_type_id,
            "field_values": dict(self.field_values),
            "field_provenance": dict(self.field_provenance),
            "missing_user_fields": list(self.missing_user_fields),
            "provisional_fields": list(self.provisional_fields),
            "warnings": list(self.warnings),
        }


GeneratedDraft: TypeAlias = (
    GeneratedContentDraft | GeneratedExamDraft | GeneratedOfficialDraft
)


class AssistantContentGenerationAdapter:
    def __init__(self, root: Path) -> None:
        self.root = Path(root)
        self.drafts_root = self.root / "drafts"
        self.repository = ContentArtifactRepository(self.root / "content_artifacts")
        self.composer = ContentMaterialComposer(repository=self.repository)

    def compile_and_compose(
        self,
        *,
        session_id: str,
        markdown: str,
        prompt_profile_id: str = "",
        cancelled=None,
    ) -> GeneratedContentDraft:
        normalized = _validate_generated_markdown(markdown)
        if str(prompt_profile_id or "").strip() == _BIDDING_PROMPT_PROFILE_ID:
            _validate_bidding_markdown(normalized)
        draft_id = uuid4().hex
        session_root = self.drafts_root / _safe_id(session_id)
        session_root.mkdir(parents=True, exist_ok=True)
        markdown_path = session_root / f"{draft_id}.md"
        _atomic_write_text(markdown_path, normalized)
        compiled = compile_content_material(
            markdown_path,
            self.repository,
            cancelled=cancelled,
        )
        if compiled.blocked or compiled.artifact_ref is None or compiled.fragment is None:
            codes = [str(item.code) for item in compiled.findings]
            raise ValueError("content_compile_blocked:" + ",".join(codes))
        content_id = "assistant_body"
        binding = ContentMaterialBinding(
            content_id=content_id,
            label="AI 生成内容",
            artifact_ref=compiled.artifact_ref,
        )
        rule = ContentInsertionRule(
            rule_id="assistant-content-body",
            content_id=content_id,
            anchor_token=content_anchor_token(content_id),
        )
        compose_inputs = ContentComposeInputs(
            content_bindings=(binding,),
            content_rules=(rule,),
        )
        anchor_path = session_root / f"{draft_id}.anchor.docx"
        composed_path = session_root / f"{draft_id}.composed.docx"
        _write_anchor_docx(anchor_path, rule.anchor_token)
        receipt = self.composer.compose(
            ContentComposeRequest(
                source_docx_path=str(anchor_path),
                output_docx_path=str(composed_path),
                inputs=compose_inputs,
            ),
            cancel_check=cancelled,
        )
        document_path = composed_path
        document_profile_id = ""
        if str(prompt_profile_id or "").strip() == _BIDDING_PROMPT_PROFILE_ID:
            document_path = session_root / f"{draft_id}.bidding.docx"
            _write_bidding_document_profile(
                source_path=composed_path,
                output_path=document_path,
            )
            document_profile_id = BIDDING_DOCUMENT_PROFILE_ID
        return GeneratedContentDraft(
            draft_id=draft_id,
            session_id=session_id,
            markdown_path=str(markdown_path),
            document_path=str(document_path),
            fragment=compiled.fragment.to_dict(),
            fragment_digest=compiled.fragment.digest,
            artifact_id=compiled.artifact_ref.artifact_id,
            manifest_sha256=compiled.artifact_ref.manifest_sha256,
            compose_receipt_id=receipt.receipt_id,
            document_profile_id=document_profile_id,
        )

    def compile_generated(
        self,
        *,
        session_id: str,
        markdown: str,
        artifact_kind: str,
        intent: str = "",
        scene_id: str = "",
        scale_profile_id: str = "",
        prompt_profile_id: str = "",
        document_type_id: str = "",
        authoritative_fields: Mapping[str, object] | None = None,
        cancelled=None,
    ) -> GeneratedDraft:
        """Dispatch through an explicit artifact compiler, never by work-mode guess."""

        compilers = {
            ARTIFACT_KIND_NARRATIVE: self.compile_and_compose,
            ARTIFACT_KIND_EXAM: self.compile_exam_markdown,
            ARTIFACT_KIND_OFFICIAL: self.compile_official_draft,
        }
        compiler = compilers.get(str(artifact_kind or "").strip())
        if compiler is None:
            raise ValueError(f"generated_artifact_kind_unsupported:{artifact_kind}")
        arguments = {
            "session_id": session_id,
            "markdown": markdown,
            "cancelled": cancelled,
        }
        if artifact_kind == ARTIFACT_KIND_EXAM:
            arguments["intent"] = intent
            arguments["scene_id"] = scene_id
            arguments["scale_profile_id"] = scale_profile_id
        elif artifact_kind == ARTIFACT_KIND_OFFICIAL:
            arguments["intent"] = intent
            arguments["document_type_id"] = document_type_id
            arguments["authoritative_fields"] = authoritative_fields
        else:
            arguments["prompt_profile_id"] = prompt_profile_id
        return compiler(**arguments)

    def compile_official_draft(
        self,
        *,
        session_id: str,
        markdown: str,
        intent: str = "",
        document_type_id: str = "",
        authoritative_fields: Mapping[str, object] | None = None,
        cancelled=None,
    ) -> GeneratedOfficialDraft:
        if cancelled is not None and cancelled():
            raise RuntimeError("content_generation_cancelled")
        source = compile_official_draft_text(
            markdown,
            intent=intent,
            document_type_id=document_type_id,
            authoritative_fields=authoritative_fields,
        )
        draft_id = uuid4().hex
        session_root = self.drafts_root / _safe_id(session_id)
        session_root.mkdir(parents=True, exist_ok=True)
        source_path = session_root / f"{draft_id}.official.json"
        preview_path = session_root / f"{draft_id}.official.preview.docx"
        source = replace(source, preview_docx_path=str(preview_path))
        _write_official_preview_docx(preview_path, source)
        write_official_draft_source(source_path, source)
        return _generated_official_draft(
            source_path,
            source,
            draft_id=draft_id,
            session_id=session_id,
        )

    def compile_exam_markdown(
        self,
        *,
        session_id: str,
        markdown: str,
        intent: str = "",
        scene_id: str = "",
        scale_profile_id: str = "",
        cancelled=None,
    ) -> GeneratedExamDraft:
        normalized = _validate_generated_markdown(markdown)
        if cancelled is not None and cancelled():
            raise RuntimeError("content_generation_cancelled")
        draft_id = uuid4().hex
        session_root = self.drafts_root / _safe_id(session_id)
        session_root.mkdir(parents=True, exist_ok=True)
        markdown_path = session_root / f"{draft_id}.exam.md"
        result = parse_exam_markdown_source(
            normalized,
            source_path=str(markdown_path),
        )
        blockers = list(
            generated_exam_blockers(
                normalized,
                result,
                intent=intent,
                scene_id=scene_id,
                scale_profile_id=scale_profile_id,
            )
        )
        if result.error_count or blockers:
            codes = [
                str(issue.kind or "unknown")
                for issue in result.issues
                if str(issue.severity or "").casefold() == "error"
            ]
            codes.extend(blockers)
            diagnostic_path = self.persist_rejected_exam(
                session_id=session_id,
                markdown=normalized,
                phase="compile",
            )
            raise ValueError(
                "exam_content_compile_blocked:"
                + ",".join(dict.fromkeys(codes))
                + f"|diagnostic={diagnostic_path}"
            )
        _atomic_write_text(markdown_path, normalized)
        digest = sha256(normalized.encode("utf-8")).hexdigest()
        payload = result.to_dict()
        validation_summary = dict(payload.get("summary") or {})
        validation_summary["blueprint"] = resolve_exam_blueprint(
            intent,
            scene_id=scene_id,
            scale_profile_id=scale_profile_id,
        ).to_dict()
        contract_warnings = generated_exam_warnings(
            normalized,
            result,
            intent=intent,
            scene_id=scene_id,
            scale_profile_id=scale_profile_id,
        )
        return GeneratedExamDraft(
            draft_id=draft_id,
            session_id=session_id,
            markdown_path=str(markdown_path),
            source_digest=digest,
            artifact_id=f"exam-{draft_id}",
            schema_id="exam_items_v1",
            validation_summary=validation_summary,
            issues=tuple(
                dict(item)
                for item in payload.get("issues", ())
                if isinstance(item, dict)
            )
            + tuple(
                {
                    "path": "quality",
                    "kind": code,
                    "message": "题稿可继续生产，但建议在交付前复核该质量项。",
                    "severity": "warning",
                }
                for code in contract_warnings
            ),
        )

    def persist_rejected_exam(
        self,
        *,
        session_id: str,
        markdown: str,
        phase: str,
    ) -> str:
        """Retain rejected provider text locally for repair and diagnostics."""

        session_root = self.drafts_root / _safe_id(session_id)
        session_root.mkdir(parents=True, exist_ok=True)
        safe_phase = re.sub(r"[^A-Za-z0-9_-]+", "_", str(phase or "validation"))
        path = session_root / f"{uuid4().hex}.{safe_phase}.rejected.exam.md"
        _atomic_write_text(path, str(markdown or ""))
        return str(path)


def is_generated_draft(value: object) -> bool:
    return isinstance(
        value,
        (GeneratedContentDraft, GeneratedExamDraft, GeneratedOfficialDraft),
    )


def generated_draft_source_ref(draft: GeneratedDraft) -> SourceArtifactRef:
    if isinstance(draft, GeneratedOfficialDraft):
        return SourceArtifactRef(
            artifact_id=draft.artifact_id,
            role=SOURCE_ROLE_STRUCTURED_SOURCE,
            media_type="application/json",
            path=draft.source_path,
            name=Path(draft.source_path).name,
            schema_id=draft.schema_id,
            digest=f"sha256:{draft.source_digest}",
            source_kind="assistant_generated",
        )
    if isinstance(draft, GeneratedExamDraft):
        return SourceArtifactRef(
            artifact_id=draft.artifact_id,
            role=SOURCE_ROLE_STRUCTURED_SOURCE,
            media_type="text/markdown",
            path=draft.markdown_path,
            name=Path(draft.markdown_path).name,
            schema_id=draft.schema_id,
            digest=f"sha256:{draft.source_digest}",
            source_kind="assistant_generated",
        )
    return SourceArtifactRef(
        artifact_id=draft.artifact_id,
        role=SOURCE_ROLE_PRODUCTION_INPUT,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        path=draft.document_path,
        name=Path(draft.document_path).name,
        schema_id="assistant-generated-content-v1",
        digest=f"sha256:{_file_sha256(Path(draft.document_path))}",
        source_kind="assistant_generated",
    )


def generated_draft_trace_refs(
    draft: GeneratedDraft,
) -> tuple[dict[str, object], ...]:
    if isinstance(draft, GeneratedOfficialDraft):
        return (
            {
                "draft_id": draft.draft_id,
                "artifact_kind": draft.artifact_kind,
                "artifact_id": draft.artifact_id,
                "source_digest": draft.source_digest,
                "schema_id": draft.schema_id,
                "document_type_id": draft.document_type_id,
                "field_provenance": dict(draft.field_provenance),
                "provisional_fields": list(draft.provisional_fields),
                "warnings": list(draft.warnings),
            },
        )
    if isinstance(draft, GeneratedExamDraft):
        return (
            {
                "draft_id": draft.draft_id,
                "artifact_kind": draft.artifact_kind,
                "artifact_id": draft.artifact_id,
                "source_digest": draft.source_digest,
                "schema_id": draft.schema_id,
                "validation_summary": dict(draft.validation_summary),
                "issues": [dict(item) for item in draft.issues],
            },
        )
    return (
        {
            "draft_id": draft.draft_id,
            "artifact_kind": draft.artifact_kind,
            "fragment_digest": draft.fragment_digest,
            "artifact_id": draft.artifact_id,
            "manifest_sha256": draft.manifest_sha256,
            "compose_receipt_id": draft.compose_receipt_id,
            "document_profile_id": draft.document_profile_id,
            "production_input_digest": _file_sha256(Path(draft.document_path)),
        },
    )


def complete_generated_official_draft_field(
    source_path: Path | str,
    *,
    field_key: str,
    value: str,
) -> GeneratedOfficialDraft:
    path = Path(source_path)
    source = complete_official_draft_field(
        load_official_draft_source(path),
        field_key=field_key,
        value=value,
    )
    preview_path = Path(source.preview_docx_path)
    _write_official_preview_docx(preview_path, source)
    write_official_draft_source(path, source)
    return _generated_official_draft(path, source)


def load_generated_official_draft(
    source_path: Path | str,
) -> GeneratedOfficialDraft:
    path = Path(source_path)
    return _generated_official_draft(
        path,
        load_official_draft_source(path),
    )


def _generated_official_draft(
    source_path: Path,
    source: OfficialDraftSource,
    *,
    draft_id: str = "",
    session_id: str = "",
) -> GeneratedOfficialDraft:
    normalized_draft_id = (
        str(draft_id or "").strip()
        or source_path.name.removesuffix(".official.json")
    )
    return GeneratedOfficialDraft(
        draft_id=normalized_draft_id,
        session_id=str(session_id or source_path.parent.name),
        source_path=str(source_path),
        preview_docx_path=source.preview_docx_path,
        source_digest=_file_sha256(source_path),
        artifact_id=f"official-{normalized_draft_id}",
        schema_id=OFFICIAL_DRAFT_SCHEMA_ID,
        document_type_id=source.document_type_id,
        field_values=dict(source.fields),
        field_provenance=dict(source.field_provenance),
        missing_user_fields=source.missing_user_fields,
        provisional_fields=source.provisional_fields,
        warnings=source.warnings,
    )


def _write_official_preview_docx(
    path: Path,
    source: OfficialDraftSource,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    title = str(source.fields.get("title") or "公文内容草稿（标题待补充）")
    document.add_heading(title, level=1)
    recipient = str(source.fields.get("recipient") or "").strip()
    if recipient:
        document.add_paragraph(f"{recipient}：")
    body = str(source.fields.get("body") or "正文待补充。")
    for paragraph in re.split(r"\n+", body):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    document.add_paragraph(
        "发文机关："
        + str(source.fields.get("organization") or "待确认")
    )
    document.add_paragraph(
        "发文字号："
        + str(source.fields.get("document_no") or "待编")
    )
    document.add_paragraph(
        "成文日期："
        + str(source.fields.get("issue_date") or "待确认")
    )
    temporary = path.with_name(f".{path.name}.tmp")
    document.save(temporary)
    os.replace(temporary, path)


def _validate_generated_markdown(markdown: str) -> str:
    text = str(markdown or "").strip()
    if not text:
        raise ValueError("generated_content_empty")
    if "\x00" in text:
        raise ValueError("generated_content_contains_nul")
    lowered = text.casefold()
    if "<w:document" in lowered or "word/document.xml" in lowered:
        raise ValueError("direct_ooxml_is_forbidden")
    if text.startswith("```") and text.endswith("```"):
        lines = text.splitlines()
        if len(lines) >= 3:
            text = "\n".join(lines[1:-1]).strip()
    return text + "\n"


def _validate_bidding_markdown(markdown: str) -> None:
    """Reject structurally incomplete or material-ownership-breaking bid drafts."""

    text = str(markdown or "")
    codes: list[str] = []
    if not re.search(r"(?m)^#\s+\S", text):
        codes.append("level_1_title_missing")
    required_tokens = {
        "company_name": "{{@text:company_name}}",
        "project_name": "{{@text:project_name}}",
        "legal_person": "{{@text:legal_person}}",
    }
    for key, token in required_tokens.items():
        if token not in text:
            codes.append(f"material_token_missing:{key}")
    required_sections = {
        "project_understanding": ("项目理解",),
        "response_content": ("响应内容", "项目响应", "需求响应"),
        "implementation_plan": ("实施方案", "实施计划"),
        "commitments": ("承诺事项", "服务承诺", "投标承诺"),
    }
    headings = tuple(
        match.group(1).strip()
        for match in re.finditer(r"(?m)^#{2,3}\s+(.+?)\s*$", text)
    )
    for key, aliases in required_sections.items():
        if not any(
            alias in heading
            for heading in headings
            for alias in aliases
        ):
            codes.append(f"required_section_missing:{key}")
    if re.search(r"\{\{@img:[^}]+\}\}", text, flags=re.IGNORECASE):
        codes.append("provider_owned_image_token_forbidden")
    if codes:
        raise ValueError(
            "bidding_content_compile_blocked:"
            + ",".join(dict.fromkeys(codes))
        )


def _safe_id(value: str) -> str:
    normalized = str(value or "").strip()
    if not normalized or not all(char.isalnum() or char in "_-" for char in normalized):
        raise ValueError("unsafe assistant content session id")
    return normalized


def _atomic_write_text(path: Path, text: str) -> None:
    temporary = path.with_suffix(f"{path.suffix}.{os.getpid()}.tmp")
    try:
        with temporary.open("w", encoding="utf-8", newline="\n") as handle:
            handle.write(text)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def _file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_anchor_docx(path: Path, token: str) -> None:
    temporary = path.with_suffix(f".docx.{os.getpid()}.tmp")
    try:
        document = Document()
        document.add_paragraph(token)
        document.save(temporary)
        os.replace(temporary, path)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def _write_bidding_document_profile(
    *,
    source_path: Path,
    output_path: Path,
) -> None:
    """Add local-only image slots after content compilation.

    Generated Markdown intentionally cannot carry image material tokens.  The
    bidding document profile adds the two package-owned slots to the DOCX
    boundary, so the content compiler stays non-recursive while production can
    still bind the current bidding package deterministically.
    """

    document = Document(source_path)
    anchor = next(
        (paragraph for paragraph in document.paragraphs if paragraph.text.strip()),
        None,
    )
    if anchor is None:
        raise ValueError("bidding_document_profile_requires_content")

    logo_label = _insert_paragraph_after(anchor, "企业标志")
    logo_label.paragraph_format.keep_with_next = True
    logo_token = _insert_paragraph_after(
        logo_label,
        material_token(MaterialTokenNamespace.IMAGE, "LOGO1"),
    )
    logo_token.alignment = 1

    seal_label = document.add_paragraph("企业公章")
    seal_label.paragraph_format.keep_with_next = True
    seal_token = document.add_paragraph(
        material_token(MaterialTokenNamespace.IMAGE, "公章1")
    )
    seal_token.alignment = 1

    temporary = output_path.with_suffix(f".docx.{os.getpid()}.tmp")
    try:
        document.save(temporary)
        os.replace(temporary, output_path)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def _insert_paragraph_after(paragraph, text: str):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.add_run(text)
    return inserted


__all__ = [
    "BIDDING_DOCUMENT_PROFILE_ID",
    "AssistantContentGenerationAdapter",
    "GeneratedContentDraft",
    "GeneratedDraft",
    "GeneratedExamDraft",
    "GeneratedOfficialDraft",
    "complete_generated_official_draft_field",
    "generated_draft_source_ref",
    "generated_draft_trace_refs",
    "is_generated_draft",
    "load_generated_official_draft",
]
