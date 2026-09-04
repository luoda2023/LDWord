"""Synchronous production-execution facade for GUI, CLI, and assistant."""

from __future__ import annotations

import copy
import hashlib
import json
from collections.abc import Callable, Mapping
from dataclasses import asdict, dataclass, field, is_dataclass, replace
from pathlib import Path

from docx import Document

from src.application.materials import (
    ExecutionMaterialFinalizeRequest,
    ExecutionMaterialSnapshot,
    finalize_execution_material_snapshot,
)
from src.config.document_structure_contract import (
    DocumentStructureEvidence,
    RegionDecision,
)
from src.config.execution_feature_state import execution_plan_is_enabled
from src.config.execution_target import resolve_execution_target
from src.config.resolver import resolve_config
from src.config.scene import SceneWorkspace
from src.config.scene_surface_registry import (
    scene_uses_exam_paper_surface,
    scene_uses_official_document_surface,
)
from src.config.template import TemplateConfig
from src.config.work_mode import execution_work_mode_issue, resolve_work_mode_id
from src.document_batch.material_resources import inspect_template_resource_roles
from src.pipeline.runner import plan_pipeline_output_paths
from src.services.document_structure_evidence import (
    build_document_structure_evidence,
    document_structure_evidence_is_current,
)
from src.services.execution_session.support import file_sha256
from src.services.production_runtime.execution_runtime import (
    WorkbenchProductionRunner,
)
from src.services.production_runtime.material_artifacts import (
    plan_material_artifact_paths,
)
from src.shared.engine.exact_material_placeholders import (
    scan_document_exact_placeholders,
)

ProgressCallback = Callable[[int, int, str], None]
CancelCheck = Callable[[], bool]


@dataclass(frozen=True, slots=True)
class ProductionExecutionRequest:
    input_path: Path
    output_root: Path
    mode_id: str
    scene: SceneWorkspace
    template: TemplateConfig
    plan_id: str
    plan_path: str = ""
    plan_source_type: str = ""
    template_id: str = ""
    template_path: str = ""
    template_source_type: str = ""
    document_type_id: str = ""
    output_suffix: str = "_formatted"
    exam_scale_profile_id: str = ""
    exam_visual_quality_required: bool = False
    require_format_change: bool = False
    session_overrides: Mapping[str, object] = field(default_factory=dict)
    material_snapshot: ExecutionMaterialSnapshot | None = None
    document_structure_evidence: DocumentStructureEvidence | None = None
    document_scope_decisions: tuple[RegionDecision, ...] = ()


def execute_production_request(
    request: ProductionExecutionRequest,
    *,
    progress_callback: ProgressCallback | None = None,
    cancel_check: CancelCheck | None = None,
) -> dict[str, object]:
    input_path = Path(request.input_path).expanduser().resolve()
    output_root = Path(request.output_root).expanduser().resolve()
    if not input_path.is_file():
        return _terminal("failed", f"input_document_missing:{input_path}")
    suffix = input_path.suffix.casefold()
    exam_markdown = (
        suffix in {".md", ".markdown"}
        and scene_uses_exam_paper_surface(
            request.scene,
            mode_id=request.mode_id,
        )
    )
    official_draft = (
        suffix == ".json"
        and scene_uses_official_document_surface(
            request.scene,
            mode_id=request.mode_id,
        )
    )
    if suffix != ".docx" and not exam_markdown and not official_draft:
        return _terminal(
            "failed",
            f"input_document_unsupported:{input_path.suffix}",
        )
    if output_root.exists() and not output_root.is_dir():
        return _terminal(
            "failed",
            f"output_root_is_not_directory:{output_root}",
        )
    scene = copy.deepcopy(request.scene)
    mode_issue = execution_work_mode_issue(
        scene,
        requested_mode_id=request.mode_id,
    )
    if mode_issue:
        return _terminal("failed", mode_issue)
    mode_id = resolve_work_mode_id(
        scene,
        requested_mode_id=request.mode_id,
    )
    plan_enabled = execution_plan_is_enabled(scene)
    structure_evidence = (
        request.document_structure_evidence if plan_enabled else None
    )
    scope_decisions = (
        tuple(request.document_scope_decisions or ()) if plan_enabled else ()
    )
    if (
        plan_enabled
        and suffix == ".docx"
        and mode_id not in {"official", "exam"}
    ):
        if structure_evidence is None:
            structure_evidence = build_document_structure_evidence(input_path)
        elif not document_structure_evidence_is_current(structure_evidence, input_path):
            return _terminal("failed", "document_structure_evidence_stale")
        if not structure_evidence.ready:
            return _terminal(
                "failed",
                ";".join(structure_evidence.issues)
                or "document_structure_evidence_not_ready",
            )
    snapshot = request.material_snapshot
    execution_request = request
    if snapshot is not None:
        if not snapshot.execution_ready:
            execution_request = _request_with_available_material_outputs(
                request,
                snapshot,
                input_path=input_path,
                output_root=output_root,
                scene=scene,
            )
            snapshot, issues = finalize_production_material_snapshot(
                execution_request,
                snapshot,
                input_path=input_path,
                output_root=output_root,
                scene=scene,
                mode_id=mode_id,
            )
            if snapshot is None:
                return _terminal(
                    "failed",
                    ",".join(issues)
                    or "execution_material_snapshot_not_finalized",
                )
        if snapshot.work_mode_id != mode_id:
            return _terminal("failed", "execution_material_mode_mismatch")
        if (
            request.document_type_id
            and snapshot.document_type
            and request.document_type_id != snapshot.document_type
        ):
            return _terminal(
                "failed",
                "execution_material_document_type_mismatch",
            )
    output_root.mkdir(parents=True, exist_ok=True)
    runner = WorkbenchProductionRunner(
        doc_path=str(input_path),
        template=copy.deepcopy(execution_request.template),
        scene=scene,
        session_overrides=dict(execution_request.session_overrides),
        material_snapshot=snapshot,
        output_dir=output_root,
        output_suffix=execution_request.output_suffix,
        document_type_id=execution_request.document_type_id,
        exam_scale_profile_id=execution_request.exam_scale_profile_id,
        exam_visual_quality_required=(
            execution_request.exam_visual_quality_required
        ),
        require_format_change=execution_request.require_format_change,
        document_scope_context=(structure_evidence, scope_decisions),
    )
    try:
        return runner.run(
            progress_callback or (lambda _current, _total, _message: None),
            cancel_check or (lambda: False),
        )
    except KeyboardInterrupt:
        return _terminal("cancelled", "execution_cancelled_by_keyboard_interrupt")
    except Exception as exc:
        return _terminal("failed", f"{type(exc).__name__}: {exc}")


def finalize_production_material_snapshot(
    request: ProductionExecutionRequest,
    snapshot: ExecutionMaterialSnapshot,
    *,
    input_path: Path | None = None,
    output_root: Path | None = None,
    scene: SceneWorkspace | None = None,
    mode_id: str = "",
) -> tuple[ExecutionMaterialSnapshot | None, tuple[str, ...]]:
    """Finalize one bound material record against the real production plan."""

    if not isinstance(snapshot, ExecutionMaterialSnapshot):
        return None, ("execution_material_snapshot_type_invalid",)
    if len(snapshot.records) != 1:
        return None, ("single_document_execution_requires_one_material_record",)
    source = Path(input_path or request.input_path).expanduser().resolve()
    destination = Path(output_root or request.output_root).expanduser().resolve()
    active_scene = copy.deepcopy(scene or request.scene)
    active_mode = str(mode_id or request.mode_id or "").strip()
    record = snapshot.records[0]
    values = dict(record.field_values)
    owners = dict(record.field_owners)
    try:
        config = resolve_config(
            copy.deepcopy(request.template),
            active_scene,
            dict(request.session_overrides),
            entity_data=values,
            field_scopes=owners,
            exact_material_placeholders=True,
        )
        planned = plan_pipeline_output_paths(
            source,
            config,
            output_dir=destination,
            output_suffix=request.output_suffix,
        )
        planned.update(
            plan_material_artifact_paths(
                input_path=source,
                output_dir=destination,
                config=config,
            )
        )
    except (OSError, RuntimeError, TypeError, ValueError) as exc:
        return None, (f"execution_output_plan_failed:{type(exc).__name__}",)
    if not planned:
        return None, ("execution_output_plan_empty",)

    supported_fields, supported_roles, material_issues = (
        production_material_template_contract(source, snapshot)
    )
    if material_issues:
        return None, material_issues

    target = resolve_execution_target(
        mode_id=active_mode,
        scene=active_scene,
        document_path=str(source),
        official_document_type_id=request.document_type_id,
    )
    if target.issues:
        return None, tuple(target.issues)
    template_revision = _production_template_revision(request)
    if target.master_path:
        master_path = Path(target.master_path).expanduser()
        if not master_path.is_file():
            return None, ("approved_master_missing",)
        master_id = target.master_id
        master_revision = "sha256:" + file_sha256(master_path)
    else:
        master_id = "pipeline.generic.v1"
        master_revision = "sha256:" + hashlib.sha256(
            b"ldword.pipeline.generic.master.v1"
        ).hexdigest()
    finalized = finalize_execution_material_snapshot(
        ExecutionMaterialFinalizeRequest(
            snapshot=snapshot,
            template_id=(
                str(request.template_id or "").strip()
                or str(getattr(request.template, "name", "") or "").strip()
                or "workbench.inline.template"
            ),
            template_revision=template_revision,
            master_id=master_id,
            master_revision=master_revision,
            recipe_version=1,
            output_root=str(destination),
            output_paths=tuple(str(path) for path in planned.values()),
            supported_field_keys=tuple(sorted(supported_fields)),
            supported_resource_roles=tuple(sorted(supported_roles)),
        )
    )
    if not finalized.ok or finalized.snapshot is None:
        return None, tuple(item.code for item in finalized.issues)
    return finalized.snapshot, ()


def _request_with_available_material_outputs(
    request: ProductionExecutionRequest,
    snapshot: ExecutionMaterialSnapshot,
    *,
    input_path: Path,
    output_root: Path,
    scene: SceneWorkspace,
) -> ProductionExecutionRequest:
    """Version material outputs instead of making a visible retry fail."""

    base_suffix = str(request.output_suffix or "_formatted")
    previous_paths: tuple[str, ...] | None = None
    for version in range(1, 10_001):
        output_suffix = (
            base_suffix if version == 1 else f"{base_suffix} ({version})"
        )
        candidate = replace(request, output_suffix=output_suffix)
        try:
            paths = tuple(
                sorted(
                    _plan_material_output_paths(
                        candidate,
                        snapshot,
                        input_path=input_path,
                        output_root=output_root,
                        scene=scene,
                    ).values()
                )
            )
        except (OSError, RuntimeError, TypeError, ValueError):
            return request
        if paths and not any(Path(path).exists() for path in paths):
            return candidate
        if previous_paths is not None and paths == previous_paths:
            # Delivery presets can own filenames independently of output_suffix.
            return request
        previous_paths = paths
    return request


def _plan_material_output_paths(
    request: ProductionExecutionRequest,
    snapshot: ExecutionMaterialSnapshot,
    *,
    input_path: Path,
    output_root: Path,
    scene: SceneWorkspace,
) -> dict[str, str]:
    record = snapshot.records[0]
    config = resolve_config(
        copy.deepcopy(request.template),
        scene,
        dict(request.session_overrides),
        entity_data=dict(record.field_values),
        field_scopes=dict(record.field_owners),
        exact_material_placeholders=True,
    )
    planned = plan_pipeline_output_paths(
        input_path,
        config,
        output_dir=output_root,
        output_suffix=request.output_suffix,
    )
    planned.update(
        plan_material_artifact_paths(
            input_path=input_path,
            output_dir=output_root,
            config=config,
        )
    )
    return planned


def production_material_template_contract(
    source: Path,
    snapshot: ExecutionMaterialSnapshot,
) -> tuple[set[str], set[str], tuple[str, ...]]:
    """Validate bound values/resources against tokens in the real DOCX."""

    record = snapshot.records[0]
    if source.suffix.casefold() != ".docx":
        supported_roles = {
            role
            for role, domain in snapshot.resource_domains.items()
            if domain == "attachment"
        }
        present_roles = {
            role for role, resources in record.resources.items() if resources
        }
        unsupported = sorted(present_roles - supported_roles)
        issues = (
            ("material.bind.template_resources_unsupported:" + ",".join(unsupported),)
            if unsupported
            else ()
        )
        return set(), supported_roles, issues

    try:
        document = Document(source)
        template_fields = {
            item.key for item in scan_document_exact_placeholders(document)
        }
        template_roles = dict(inspect_template_resource_roles((source,)))
    except Exception as exc:  # noqa: BLE001 - corrupt input is a preflight failure
        return set(), set(), (
            f"material.bind.template_scan_failed:{type(exc).__name__}",
        )

    issues: list[str] = []
    present_fields = set(record.field_values)
    missing_fields = sorted(template_fields - present_fields)
    if missing_fields:
        issues.append(
            "material.bind.template_fields_required_missing:"
            + ",".join(missing_fields)
        )

    for role, domain in sorted(template_roles.items()):
        declared_domain = snapshot.resource_domains.get(role)
        if declared_domain is None:
            issues.append(f"material.bind.template_resource_role_unknown:{role}")
        elif declared_domain != domain:
            issues.append(
                "material.bind.template_resource_domain_mismatch:"
                f"{role}:{declared_domain}:{domain}"
            )
        elif not record.resources.get(role, ()):
            issues.append(
                f"material.bind.template_resource_required_missing:{role}"
            )

    present_roles = {
        role for role, resources in record.resources.items() if resources
    }
    attachment_roles = {
        role
        for role, domain in snapshot.resource_domains.items()
        if domain == "attachment"
    }
    supported_roles = set(template_roles) | attachment_roles
    unsupported_roles = sorted(present_roles - supported_roles)
    if unsupported_roles:
        issues.append(
            "material.bind.template_resources_unsupported:"
            + ",".join(unsupported_roles)
        )
    # A package can intentionally carry reusable fields that one document does
    # not consume.  Missing document tokens are unsafe; surplus package fields
    # are not.
    supported_fields = template_fields | present_fields
    return supported_fields, supported_roles, tuple(issues)


def _production_template_revision(
    request: ProductionExecutionRequest,
) -> str:
    configured = str(request.template_path or "").strip()
    if configured:
        path = Path(configured).expanduser()
        if path.is_file():
            return "sha256:" + file_sha256(path)
    template = request.template
    payload = asdict(template) if is_dataclass(template) else repr(template)
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        allow_nan=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return "sha256:" + hashlib.sha256(encoded).hexdigest()


def _terminal(status: str, error_text: str) -> dict[str, object]:
    return {
        "status": status,
        "output_path": "",
        "output_paths": {},
        "report_paths": [],
        "failed_count": 0 if status == "cancelled" else 1,
        "artifact_failure_count": 0,
        "error_text": error_text,
    }


__all__ = [
    "ProductionExecutionRequest",
    "execute_production_request",
    "production_material_template_contract",
]
