"""Atomic end-to-end composition of frozen content materials into DOCX.

The composer owns the file transaction.  Importers remain source-specific and
the renderer remains source-neutral; this service joins them only after every
source and every original-target anchor has passed preflight.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, replace
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document

from src.config.attachment_materials import AttachmentBinding
from src.config.content_artifacts import ContentMaterialBinding
from src.config.content_materials import ContentInsertionRule, DocumentFragment
from src.config.library import CONFIG_LIBRARY_ROOT
from src.modules.structure.heading_recognition import rebuild_document_index
from src.pipeline.context import PipelineContext
from src.services.material_content.docx_renderer import (
    ContentDocxRenderer,
    ContentImageJobDraft,
    ContentRenderBlockedError,
    ContentRenderReceipt,
)
from src.services.material_content.artifact_repository import (
    ContentArtifactRepository,
    ContentArtifactRepositoryError,
)
from src.shared.io.safe_docx_package import (
    DocxPackageError,
    SafeDocxPackage,
    capture_bounded_file,
)
from src.services.material_content.resource_materializer import (
    ContentResourceMaterializationError,
    ContentResourceMaterializationReceipt,
    materialize_content_fragment_resources,
)
from src.services.material_attachments.docx_renderer import (
    AttachmentReferenceDocxRenderer,
    AttachmentRenderBlockedError,
    AttachmentRenderReceipt,
)
from src.shared.engine.material_token_contract import (
    MATERIAL_TOKEN_PATTERN,
    MaterialTokenKind,
    parse_material_token,
)


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_W_P = f"{{{_W_NS}}}p"
_W_T = f"{{{_W_NS}}}t"
_W_DRAWING = f"{{{_W_NS}}}drawing"
_W_BOOKMARK_START = f"{{{_W_NS}}}bookmarkStart"
_WP_DOCPR = f"{{{_WP_NS}}}docPr"
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)


CancelCheck = Callable[[], bool]


@dataclass(frozen=True, slots=True)
class ContentComposeInputs:
    """Recipe-local inputs for one atomic composition operation."""

    input_id: str = ""
    content_bindings: tuple[ContentMaterialBinding, ...] = ()
    content_rules: tuple[ContentInsertionRule, ...] = ()
    attachment_bindings: tuple[AttachmentBinding, ...] = ()

    def __post_init__(self) -> None:
        for name, values, expected_type in (
            ("content_bindings", self.content_bindings, ContentMaterialBinding),
            ("content_rules", self.content_rules, ContentInsertionRule),
            ("attachment_bindings", self.attachment_bindings, AttachmentBinding),
        ):
            if type(values) is not tuple or any(
                not isinstance(item, expected_type) for item in values
            ):
                raise TypeError(f"{name}_invalid")
        binding_ids = tuple(item.content_id for item in self.content_bindings)
        rule_ids = tuple(item.rule_id for item in self.content_rules)
        roles = tuple(item.role for item in self.attachment_bindings)
        for name, values in (
            ("content_binding", binding_ids),
            ("content_rule", rule_ids),
            ("attachment_role", roles),
        ):
            if len(values) != len(set(values)):
                raise ValueError(f"{name}_duplicate")
        known_bindings = set(binding_ids)
        if any(item.content_id not in known_bindings for item in self.content_rules):
            raise ValueError("content_rule_binding_missing")
        payload = {
            "content_bindings": [
                item.to_dict() for item in self.content_bindings
            ],
            "content_rules": [item.to_dict() for item in self.content_rules],
            "attachment_bindings": [
                item.to_dict() for item in self.attachment_bindings
            ],
        }
        computed = "sha256:" + sha256(
            json.dumps(
                payload,
                ensure_ascii=False,
                allow_nan=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        if self.input_id and self.input_id != computed:
            raise ValueError("content_compose_input_id_mismatch")
        object.__setattr__(self, "input_id", computed)


@dataclass(frozen=True, slots=True)
class ContentComposeRequest:
    source_docx_path: str
    output_docx_path: str
    inputs: ContentComposeInputs

    def __post_init__(self) -> None:
        if not isinstance(self.source_docx_path, str) or not self.source_docx_path.strip():
            raise ValueError("source_docx_path must not be empty")
        if not isinstance(self.output_docx_path, str) or not self.output_docx_path.strip():
            raise ValueError("output_docx_path must not be empty")
        if not isinstance(self.inputs, ContentComposeInputs):
            raise TypeError("inputs must be ContentComposeInputs")


@dataclass(frozen=True, slots=True)
class ContentComposeDiagnostic:
    code: str
    message: str
    stage: str
    content_id: str = ""
    rule_id: str = ""
    path: str = ""


class ContentComposeError(RuntimeError):
    def __init__(self, diagnostics):
        self.diagnostics = tuple(diagnostics)
        detail = "\n".join(
            f"- [{item.stage}:{item.code}] {item.message}"
            for item in self.diagnostics
        )
        super().__init__(f"Content composition failed:\n{detail}")


class ContentComposeCancelledError(ContentComposeError):
    @property
    def stage(self) -> str:
        return self.diagnostics[0].stage


@dataclass(frozen=True, slots=True)
class ContentArtifactEvidence:
    content_id: str
    artifact_id: str
    manifest_sha256: str
    fragment_sha256: str
    compile_receipt_id: str

    def to_dict(self) -> dict[str, object]:
        return {
            "content_id": self.content_id,
            "artifact_id": self.artifact_id,
            "manifest_sha256": self.manifest_sha256,
            "fragment_sha256": self.fragment_sha256,
            "compile_receipt_id": self.compile_receipt_id,
        }


@dataclass(frozen=True, slots=True)
class ContentRenderSummary:
    rule_id: str
    content_id: str
    occurrence_count: int
    rendered_block_count: int
    inserted_body_element_count: int
    dropped_page_break_count: int
    occurrence_ids: tuple[str, ...]
    image_job_ids: tuple[str, ...]
    diagnostic_codes: tuple[str, ...]

    def to_dict(self) -> dict[str, object]:
        return {
            "rule_id": self.rule_id,
            "content_id": self.content_id,
            "occurrence_count": self.occurrence_count,
            "rendered_block_count": self.rendered_block_count,
            "inserted_body_element_count": self.inserted_body_element_count,
            "dropped_page_break_count": self.dropped_page_break_count,
            "occurrence_ids": list(self.occurrence_ids),
            "image_job_ids": list(self.image_job_ids),
            "diagnostic_codes": list(self.diagnostic_codes),
        }


@dataclass(frozen=True, slots=True)
class DocumentIndexEvidence:
    paragraph_count: int
    table_count: int
    heading_count: int
    heading_map: tuple[tuple[int, int], ...]
    section_ranges: tuple[tuple[str, int, int], ...]
    paragraph_text_sha256: str
    index_sha256: str

    def to_dict(self) -> dict[str, object]:
        return {
            "paragraph_count": self.paragraph_count,
            "table_count": self.table_count,
            "heading_count": self.heading_count,
            "heading_map": [list(item) for item in self.heading_map],
            "section_ranges": [list(item) for item in self.section_ranges],
            "paragraph_text_sha256": self.paragraph_text_sha256,
            "index_sha256": self.index_sha256,
        }


@dataclass(frozen=True, slots=True)
class ComposeReceipt:
    receipt_id: str
    compose_input_id: str
    source_docx_path: str
    output_docx_path: str
    source_sha256: str
    output_sha256: str
    content_artifacts: tuple[ContentArtifactEvidence, ...]
    render_receipts: tuple[ContentRenderSummary, ...]
    attachment_render_receipts: tuple[AttachmentRenderReceipt, ...]
    deferred_image_drafts: tuple[ContentImageJobDraft, ...]
    resource_materialization: ContentResourceMaterializationReceipt
    index_evidence: DocumentIndexEvidence

    def __post_init__(self) -> None:
        if not isinstance(
            self.resource_materialization,
            ContentResourceMaterializationReceipt,
        ):
            raise TypeError(
                "resource_materialization must be a "
                "ContentResourceMaterializationReceipt"
            )
        supplied_id = str(self.receipt_id or "").strip()
        computed_id = sha256(self.canonical_json().encode("utf-8")).hexdigest()
        if supplied_id and supplied_id != computed_id:
            raise ValueError("receipt_id does not match canonical receipt payload")
        object.__setattr__(self, "receipt_id", computed_id)

    def canonical_payload(self) -> dict[str, object]:
        return {
            "compose_input_id": self.compose_input_id,
            "source_docx_path": self.source_docx_path,
            "output_docx_path": self.output_docx_path,
            "source_sha256": self.source_sha256,
            "output_sha256": self.output_sha256,
            "content_artifacts": [item.to_dict() for item in self.content_artifacts],
            "render_receipts": [item.to_dict() for item in self.render_receipts],
            "attachment_render_receipts": [
                item.to_dict() for item in self.attachment_render_receipts
            ],
            "deferred_image_drafts": [
                _image_draft_to_dict(item) for item in self.deferred_image_drafts
            ],
            "resource_materialization": self.resource_materialization.to_dict(),
            "index_evidence": self.index_evidence.to_dict(),
        }

    def canonical_json(self) -> str:
        return json.dumps(
            self.canonical_payload(),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )

    def to_dict(self) -> dict[str, object]:
        return {"receipt_id": self.receipt_id, **self.canonical_payload()}


@dataclass(frozen=True, slots=True)
class _PackageInventory:
    content_token_locations: tuple[str, ...]
    bookmark_names: tuple[str, ...]
    drawing_count: int
    max_docpr_id: int


class ContentMaterialComposer:
    """Artifact resolver plus one atomic target-DOCX composition transaction."""

    def __init__(
        self,
        *,
        repository: ContentArtifactRepository | None = None,
        renderer: ContentDocxRenderer | None = None,
        attachment_renderer: AttachmentReferenceDocxRenderer | None = None,
    ) -> None:
        self._repository = repository or ContentArtifactRepository(
            CONFIG_LIBRARY_ROOT / "content_artifacts"
        )
        self._renderer = renderer or ContentDocxRenderer()
        self._attachment_renderer = (
            attachment_renderer or AttachmentReferenceDocxRenderer()
        )

    def compose(
        self,
        request: ContentComposeRequest,
        *,
        cancel_check: CancelCheck | None = None,
    ) -> ComposeReceipt:
        if not isinstance(request, ContentComposeRequest):
            raise TypeError("request must be a ContentComposeRequest")
        _check_cancel(cancel_check, "start")
        source, output = _validate_paths(request)
        inputs = request.inputs
        input_diagnostics = _validate_inputs(inputs, self._repository)
        if input_diagnostics:
            raise ContentComposeError(input_diagnostics)

        try:
            source_payload = capture_bounded_file(source)
            source_package = SafeDocxPackage.open(source_payload)
            source_sha256 = source_package.source_sha256
            source_inventory = _inspect_docx_package(source_package)
            document = Document(BytesIO(source_payload))
            del source_package, source_payload
        except DocxPackageError as exc:
            raise ContentComposeError(
                (
                    _diagnostic(
                        f"target_{exc.code}",
                        str(exc),
                        "target",
                        path=str(source),
                    ),
                )
            ) from exc
        except Exception as exc:
            raise ContentComposeError(
                (_diagnostic("target_open_failed", str(exc), "target", path=str(source)),)
            ) from exc

        binding_by_id = {item.content_id: item for item in inputs.content_bindings}
        fragments: dict[str, DocumentFragment] = {}
        import_diagnostics: list[ContentComposeDiagnostic] = []
        for binding in inputs.content_bindings:
            try:
                fragment = self._repository.load_fragment(binding.artifact_ref)
                fragments[binding.content_id] = fragment
            except ContentArtifactRepositoryError as exc:
                import_diagnostics.append(
                    _diagnostic(
                        exc.code,
                        "compiled content artifact is missing or invalid",
                        "artifact_resolve",
                        content_id=binding.content_id,
                    )
                )
        if import_diagnostics:
            raise ContentComposeError(import_diagnostics)
        _check_cancel(cancel_check, "imported")

        preflight_diagnostics: list[ContentComposeDiagnostic] = []
        for rule in inputs.content_rules:
            preflight = self._renderer.preflight(
                document, fragments[rule.content_id], rule
            )
            if preflight.ready:
                continue
            for finding in preflight.diagnostics:
                if str(getattr(finding.severity, "value", finding.severity)) != "error":
                    continue
                preflight_diagnostics.append(
                    _diagnostic(
                        f"renderer_{finding.code}",
                        finding.message,
                        "preflight",
                        content_id=rule.content_id,
                        rule_id=rule.rule_id,
                        path=str(source),
                    )
                )
        for binding in inputs.attachment_bindings:
            preflight = self._attachment_renderer.preflight(document, binding)
            for finding in preflight.diagnostics:
                preflight_diagnostics.append(
                    _diagnostic(
                        f"attachment_renderer_{finding.code}",
                        finding.message,
                        "preflight",
                        path=str(source),
                    )
                )
        if preflight_diagnostics:
            raise ContentComposeError(preflight_diagnostics)
        _check_cancel(cancel_check, "preflighted")

        try:
            resource_materialization = materialize_content_fragment_resources(
                inputs.content_bindings,
                fragments,
                self._repository,
            )
        except ContentResourceMaterializationError as exc:
            raise ContentComposeError(
                tuple(
                    _diagnostic(
                        f"resource_{item.code}",
                        item.message,
                        "materialize",
                        content_id=item.content_id,
                        path=item.path,
                    )
                    for item in exc.diagnostics
                )
            ) from exc
        render_receipts: list[ContentRenderReceipt] = []
        attachment_render_receipts: list[AttachmentRenderReceipt] = []
        try:
            for rule in inputs.content_rules:
                render_receipts.append(
                    self._renderer.render(document, fragments[rule.content_id], rule)
                )
            for binding in inputs.attachment_bindings:
                attachment_render_receipts.append(
                    self._attachment_renderer.render(document, binding)
                )
        except ContentRenderBlockedError as exc:
            raise ContentComposeError(
                tuple(
                    _diagnostic(
                        f"renderer_{item.code}",
                        item.message,
                        "render",
                        rule_id=item.rule_id,
                        path=str(source),
                    )
                    for item in exc.diagnostics
                )
            ) from exc
        except AttachmentRenderBlockedError as exc:
            raise ContentComposeError(
                tuple(
                    _diagnostic(
                        f"attachment_renderer_{item.code}",
                        item.message,
                        "render",
                        path=str(source),
                    )
                    for item in exc.diagnostics
                )
            ) from exc
        _check_cancel(cancel_check, "rendered")

        drafts = _normalize_deferred_drafts(
            render_receipts, source_inventory.max_docpr_id
        )
        staging = _new_staging_path(output)
        published = False
        try:
            document.save(str(staging))
            _canonicalize_docx_package(staging)
            _fsync_file(staging)
            _check_cancel(cancel_check, "staged")

            first_verified = _open_and_validate_staging(
                staging,
                drafts,
                expected_drawing_count=source_inventory.drawing_count,
            )
            first_context = PipelineContext(source_doc_path=str(source))
            rebuild_document_index(first_verified, first_context)
            first_verified.save(str(staging))
            _canonicalize_docx_package(staging)
            _fsync_file(staging)

            final_verified = _open_and_validate_staging(
                staging,
                drafts,
                expected_drawing_count=source_inventory.drawing_count,
            )
            before_reindex_xml = final_verified.element.xml
            final_context = PipelineContext(source_doc_path=str(source))
            doc_tree = rebuild_document_index(final_verified, final_context)
            if final_verified.element.xml != before_reindex_xml:
                raise ContentComposeError(
                    (
                        _diagnostic(
                            "document_index_not_stable",
                            "a second index rebuild still changed document XML",
                            "verify",
                            path=str(staging),
                        ),
                    )
                )
            index_evidence = _build_index_evidence(final_verified, doc_tree)

            if _file_sha256(source) != source_sha256:
                raise ContentComposeError(
                    (
                        _diagnostic(
                            "target_changed_during_compose",
                            "source target DOCX changed during composition",
                            "verify",
                            path=str(source),
                        ),
                    )
                )
            output_sha256 = _file_sha256(staging)
            receipt = ComposeReceipt(
                receipt_id="",
                compose_input_id=inputs.input_id,
                source_docx_path=str(source),
                output_docx_path=str(output),
                source_sha256=source_sha256,
                output_sha256=output_sha256,
                content_artifacts=tuple(
                    _content_artifact_evidence(
                        binding_by_id[rule.content_id],
                        fragments[rule.content_id],
                        self._repository,
                    )
                    for rule in inputs.content_rules
                ),
                render_receipts=tuple(
                    _render_summary(item) for item in render_receipts
                ),
                attachment_render_receipts=tuple(attachment_render_receipts),
                deferred_image_drafts=drafts,
                resource_materialization=resource_materialization,
                index_evidence=index_evidence,
            )
            _check_cancel(cancel_check, "before_publish")
            try:
                os.replace(staging, output)
            except OSError as exc:
                raise ContentComposeError(
                    (_diagnostic("atomic_publish_failed", str(exc), "publish", path=str(output)),)
                ) from exc
            published = True
            return receipt
        finally:
            if not published:
                staging.unlink(missing_ok=True)


def compose_content_materials(
    source_docx_path: str | Path,
    output_docx_path: str | Path,
    inputs: ContentComposeInputs,
    *,
    repository: ContentArtifactRepository | None = None,
    cancel_check: CancelCheck | None = None,
) -> ComposeReceipt:
    request = ContentComposeRequest(
        source_docx_path=str(source_docx_path),
        output_docx_path=str(output_docx_path),
        inputs=inputs,
    )
    return ContentMaterialComposer(repository=repository).compose(
        request, cancel_check=cancel_check
    )


def _validate_paths(request):
    source = Path(request.source_docx_path).expanduser().resolve()
    output = Path(request.output_docx_path).expanduser().resolve()
    diagnostics: list[ContentComposeDiagnostic] = []
    if source.suffix.casefold() != ".docx":
        diagnostics.append(_diagnostic("source_not_docx", "source must be .docx", "request", path=str(source)))
    if output.suffix.casefold() != ".docx":
        diagnostics.append(_diagnostic("output_not_docx", "output must be .docx", "request", path=str(output)))
    if not source.is_file():
        diagnostics.append(_diagnostic("source_missing", "source target DOCX does not exist", "request", path=str(source)))
    if not output.parent.is_dir():
        diagnostics.append(_diagnostic("output_parent_missing", "output parent directory does not exist", "request", path=str(output.parent)))
    same_path = os.path.normcase(str(source)) == os.path.normcase(str(output))
    if not same_path and output.exists():
        try:
            same_path = os.path.samefile(source, output)
        except OSError:
            pass
    if same_path:
        diagnostics.append(_diagnostic("source_output_same", "output path must not equal or alias source path", "request", path=str(output)))
    if diagnostics:
        raise ContentComposeError(diagnostics)
    return source, output


def _validate_inputs(inputs, repository):
    diagnostics: list[ContentComposeDiagnostic] = []
    bindings = {item.content_id: item for item in inputs.content_bindings}
    rule_counts: dict[str, int] = {}
    for rule in inputs.content_rules:
        rule_counts[rule.content_id] = rule_counts.get(rule.content_id, 0) + 1
    for binding in inputs.content_bindings:
        count = rule_counts.get(binding.content_id, 0)
        if count != 1:
            diagnostics.append(
                _diagnostic(
                    "binding_rule_cardinality_invalid",
                    f"content binding requires exactly one rule, found {count}",
                    "inputs",
                    content_id=binding.content_id,
                )
            )
        try:
            repository.validate(binding.artifact_ref)
        except ContentArtifactRepositoryError as exc:
            diagnostics.append(
                _diagnostic(
                    exc.code,
                    "compiled content artifact is missing or invalid",
                    "inputs",
                    content_id=binding.content_id,
                )
            )
    for content_id, count in rule_counts.items():
        if content_id not in bindings:
            diagnostics.append(
                _diagnostic("rule_binding_missing", "content rule has no binding", "inputs", content_id=content_id)
            )
        elif count > 1:
            diagnostics.append(
                _diagnostic(
                    "duplicate_content_rules",
                    f"content id has {count} rules with the same canonical anchor",
                    "inputs",
                    content_id=content_id,
                )
            )
    return diagnostics


def _new_staging_path(output):
    descriptor, raw_path = tempfile.mkstemp(
        prefix=f".{output.name}.compose-",
        suffix=".docx",
        dir=str(output.parent),
    )
    os.close(descriptor)
    return Path(raw_path)


def _canonicalize_docx_package(path):
    canonical_path = path.with_name(path.name + ".canonical")
    try:
        source_package = SafeDocxPackage.open_path(path)
        with ZipFile(
            canonical_path,
            "w",
            compression=ZIP_DEFLATED,
            compresslevel=9,
        ) as target_zip:
            for name in sorted(source_package.part_names):
                info = ZipInfo(name, date_time=_FIXED_ZIP_TIME)
                info.compress_type = ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = 0o100644 << 16
                target_zip.writestr(info, source_package.read_part(name))
        os.replace(canonical_path, path)
    finally:
        canonical_path.unlink(missing_ok=True)


def _open_and_validate_staging(path, drafts, *, expected_drawing_count):
    try:
        payload = capture_bounded_file(path)
        package = SafeDocxPackage.open(payload)
        inventory = _inspect_docx_package(package)
        document = Document(BytesIO(payload))
        del package, payload
    except Exception as exc:
        raise ContentComposeError(
            (_diagnostic("staging_reopen_failed", str(exc), "verify", path=str(path)),)
        ) from exc
    diagnostics: list[ContentComposeDiagnostic] = []
    if inventory.content_token_locations:
        diagnostics.append(
            _diagnostic(
                "structural_material_token_residual",
                "file/attachment tokens remain in: "
                + ", ".join(inventory.content_token_locations),
                "verify",
                path=str(path),
            )
        )
    if inventory.drawing_count != expected_drawing_count:
        diagnostics.append(
            _diagnostic(
                "content_image_was_inlined",
                "content composition changed drawing count instead of deferring images",
                "verify",
                path=str(path),
            )
        )
    marker_counts = {
        marker: inventory.bookmark_names.count(marker)
        for marker in (item.stable_marker_id for item in drafts)
    }
    invalid_markers = sorted(marker for marker, count in marker_counts.items() if count != 1)
    if invalid_markers:
        diagnostics.append(
            _diagnostic(
                "image_sentinel_missing_or_duplicate",
                "invalid image sentinels: " + ", ".join(invalid_markers),
                "verify",
                path=str(path),
            )
        )
    reserved_docpr_ids = [item.reserved_docpr_id for item in drafts]
    if len(reserved_docpr_ids) != len(set(reserved_docpr_ids)) or any(
        item <= inventory.max_docpr_id for item in reserved_docpr_ids
    ):
        diagnostics.append(
            _diagnostic(
                "reserved_docpr_id_collision",
                "deferred image docPr ids collide with the staged package",
                "verify",
                path=str(path),
            )
        )
    if diagnostics:
        raise ContentComposeError(diagnostics)
    return document


def _inspect_docx_package(path):
    token_locations: list[str] = []
    bookmark_names: list[str] = []
    drawing_count = 0
    max_docpr_id = 0
    package = (
        path
        if isinstance(path, SafeDocxPackage)
        else SafeDocxPackage.open_path(path)
    )
    for name in sorted(package.part_names):
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        root = package.parse_xml(name)
        for paragraph in root.iter(_W_P):
            text = "".join(node.text or "" for node in paragraph.iter(_W_T))
            if any(
                _is_structural_token(match.group(0))
                for match in MATERIAL_TOKEN_PATTERN.finditer(text)
            ):
                token_locations.append(name)
                break
        drawing_count += sum(1 for _ in root.iter(_W_DRAWING))
        for bookmark in root.iter(_W_BOOKMARK_START):
            value = bookmark.get(f"{{{_W_NS}}}name")
            if value:
                bookmark_names.append(value)
        for docpr in root.iter(_WP_DOCPR):
            try:
                max_docpr_id = max(max_docpr_id, int(docpr.get("id", "0")))
            except ValueError:
                continue
    return _PackageInventory(
        tuple(token_locations), tuple(bookmark_names), drawing_count, max_docpr_id
    )


def _is_structural_token(value: str) -> bool:
    try:
        return parse_material_token(value).kind in {
            MaterialTokenKind.CONTENT,
            MaterialTokenKind.ATTACHMENT,
        }
    except (TypeError, ValueError):
        return False


def _normalize_deferred_drafts(render_receipts, existing_max_docpr_id):
    next_docpr_id = existing_max_docpr_id + 1
    normalized: list[ContentImageJobDraft] = []
    seen_jobs: set[str] = set()
    seen_markers: set[str] = set()
    for receipt in render_receipts:
        for draft in receipt.image_job_drafts:
            if draft.job_id in seen_jobs or draft.stable_marker_id in seen_markers:
                raise ContentComposeError(
                    (_diagnostic("deferred_image_identity_collision", "deferred image identities collide", "render"),)
                )
            seen_jobs.add(draft.job_id)
            seen_markers.add(draft.stable_marker_id)
            normalized.append(
                replace(
                    draft,
                    sequence=len(normalized),
                    reserved_docpr_id=next_docpr_id,
                )
            )
            next_docpr_id += 1
    return tuple(normalized)


def _build_index_evidence(document, doc_tree):
    heading_map = tuple(sorted((int(key), int(value)) for key, value in doc_tree.heading_map.items()))
    section_ranges = tuple(
        sorted(
            (str(name), int(bounds[0]), int(bounds[1]))
            for name, bounds in doc_tree.section_ranges.items()
        )
    )
    paragraph_texts = [paragraph.text for paragraph in document.paragraphs]
    paragraph_text_sha256 = sha256(
        json.dumps(paragraph_texts, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    payload = {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "heading_map": heading_map,
        "section_ranges": section_ranges,
        "paragraph_text_sha256": paragraph_text_sha256,
    }
    index_sha256 = sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return DocumentIndexEvidence(
        paragraph_count=len(document.paragraphs),
        table_count=len(document.tables),
        heading_count=len(heading_map),
        heading_map=heading_map,
        section_ranges=section_ranges,
        paragraph_text_sha256=paragraph_text_sha256,
        index_sha256=index_sha256,
    )


def _content_artifact_evidence(binding, fragment, repository):
    resolved = repository.validate(binding.artifact_ref)
    receipt = repository.load_receipt(binding.artifact_ref)
    fragment_sha256 = sha256(
        json.dumps(
            fragment.to_dict(),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    if fragment_sha256 != resolved.manifest.fragment.sha256:
        raise ContentComposeError(
            (
                _diagnostic(
                    "artifact_fragment_digest_mismatch",
                    "artifact fragment digest changed after resolution",
                    "verify",
                    content_id=binding.content_id,
                ),
            )
        )
    return ContentArtifactEvidence(
        content_id=binding.content_id,
        artifact_id=binding.artifact_ref.artifact_id,
        manifest_sha256=binding.artifact_ref.manifest_sha256,
        fragment_sha256=fragment_sha256,
        compile_receipt_id=receipt.receipt_id,
    )


def _render_summary(receipt):
    return ContentRenderSummary(
        rule_id=receipt.rule_id,
        content_id=receipt.content_id,
        occurrence_count=receipt.occurrence_count,
        rendered_block_count=receipt.rendered_block_count,
        inserted_body_element_count=receipt.inserted_body_element_count,
        dropped_page_break_count=receipt.dropped_page_break_count,
        occurrence_ids=tuple(item.occurrence_id for item in receipt.occurrences),
        image_job_ids=tuple(item.job_id for item in receipt.image_job_drafts),
        diagnostic_codes=tuple(
            f"{getattr(item.severity, 'value', item.severity)}:{item.code}"
            for item in receipt.diagnostics
        ),
    )


def _image_draft_to_dict(item):
    return {
        "job_id": item.job_id,
        "content_id": item.content_id,
        "resource_id": item.resource_id,
        "stable_marker_id": item.stable_marker_id,
        "occurrence_id": item.occurrence_id,
        "sequence": item.sequence,
        "reserved_docpr_id": item.reserved_docpr_id,
        "alt_text": item.alt_text,
        "width_px": item.width_px,
        "height_px": item.height_px,
    }


def _file_sha256(path):
    digest = sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _fsync_file(path):
    # Windows requires a writable descriptor for FlushFileBuffers, which is
    # what Python's os.fsync delegates to there.
    with Path(path).open("rb+") as stream:
        os.fsync(stream.fileno())


def _check_cancel(check, stage):
    if check is None:
        return
    try:
        cancelled = bool(check())
    except Exception as exc:
        raise ContentComposeError(
            (_diagnostic("cancel_check_failed", str(exc), stage),)
        ) from exc
    if cancelled:
        raise ContentComposeCancelledError(
            (_diagnostic("compose_cancelled", "composition was cancelled", stage),)
        )


def _diagnostic(code, message, stage, *, content_id="", rule_id="", path=""):
    return ContentComposeDiagnostic(
        code=code,
        message=str(message),
        stage=stage,
        content_id=content_id,
        rule_id=rule_id,
        path=path,
    )


__all__ = [
    "CancelCheck",
    "ComposeReceipt",
    "ContentComposeCancelledError",
    "ContentComposeDiagnostic",
    "ContentComposeError",
    "ContentComposeInputs",
    "ContentComposeRequest",
    "ContentArtifactEvidence",
    "ContentMaterialComposer",
    "ContentRenderSummary",
    "DocumentIndexEvidence",
    "compose_content_materials",
]
