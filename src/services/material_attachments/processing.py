"""Transaction-safe processing of complete attachment bindings."""

from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field, replace
from hashlib import sha256
import json
from pathlib import Path
import shutil
from types import MappingProxyType
from uuid import uuid4

from docx import Document

from src.config.atomic_io import atomic_write_text
from src.config.attachment_materials import (
    AttachmentBinding,
    AttachmentBundleReceipt,
    AttachmentDocumentPlan,
    AttachmentFileReceipt,
    AttachmentFileStatus,
    AttachmentProcessingMode,
)
from src.config.material_snapshot import MaterialSnapshot
from src.config.image_materials import (
    FrozenImageMaterialRule,
    ImageSourceBinding,
)
from src.services.material_assets.image_document_executor import (
    ImageDocumentExecutionError,
    ImageDocumentExecutionRequest,
    ImageDocumentExecutor,
)
from src.services.material_assets.image_transform_batch import (
    MAX_IMAGE_TRANSFORM_WORKERS,
)
from src.services.material_delivery.transaction import (
    DirectoryPublishError,
    publish_staged_directory,
    remove_owned_path,
)
from src.shared.engine.docx_material_tokens import extract_docx_material_token_blocks
from src.shared.engine.exact_material_placeholders import (
    replace_document_exact_placeholders,
)
from src.shared.engine.material_dependency_index import (
    DependencyDiagnosticSeverity,
    MaterialConsumerKind,
    MaterialConsumerRef,
    MaterialDependencyIndex,
    MaterialDependencyOccurrence,
)
from src.shared.engine.material_token_router import (
    MaterialTokenKind,
    route_material_tokens,
)
from src.shared.engine.office_image_layout_contracts import OfficeImageProvider


_RECEIPT_FILENAME = "attachment_bundle_receipt.json"


@dataclass(frozen=True, slots=True)
class AttachmentProcessingDiagnostic:
    code: str
    message: str
    consumer_id: str = ""
    item_id: str = ""
    relative_path: str = ""

    def __post_init__(self) -> None:
        if not isinstance(self.code, str) or not self.code.strip():
            raise ValueError("diagnostic code must not be empty")
        if not isinstance(self.message, str) or not self.message.strip():
            raise ValueError("diagnostic message must not be empty")

    def to_dict(self) -> dict[str, str]:
        return {
            "code": self.code,
            "message": self.message,
            "consumer_id": self.consumer_id,
            "item_id": self.item_id,
            "relative_path": self.relative_path,
        }


class AttachmentBundleProcessingError(RuntimeError):
    """A binding failed and its staged directory was not published."""

    def __init__(self, diagnostics: Sequence[AttachmentProcessingDiagnostic]) -> None:
        self.diagnostics = tuple(diagnostics)
        if not self.diagnostics:
            raise ValueError("processing error requires diagnostics")
        detail = "; ".join(
            f"{item.code}:{item.consumer_id or item.relative_path}:{item.message}"
            for item in self.diagnostics
        )
        super().__init__(f"attachment bundle processing failed: {detail}")

    def to_dict(self) -> dict[str, object]:
        return {
            "error": "attachment_bundle_processing_failed",
            "diagnostics": [item.to_dict() for item in self.diagnostics],
        }


@dataclass(frozen=True, slots=True)
class AttachmentBundleBuildRequest:
    snapshot: MaterialSnapshot
    dependency_index: MaterialDependencyIndex
    binding_role: str
    final_directory: str
    image_cache_dir: str = ""
    watermark_font_path: str = ""
    max_image_transform_workers: int = 1
    office_provider: OfficeImageProvider = OfficeImageProvider.WORD
    layout_options: Mapping[str, object] = field(default_factory=dict)

    def __post_init__(self) -> None:
        if not isinstance(self.snapshot, MaterialSnapshot):
            raise TypeError("snapshot must be a MaterialSnapshot")
        if not isinstance(self.dependency_index, MaterialDependencyIndex):
            raise TypeError("dependency_index must be a MaterialDependencyIndex")
        if not isinstance(self.binding_role, str) or not self.binding_role.strip():
            raise ValueError("binding_role must not be empty")
        if not isinstance(self.final_directory, str) or not self.final_directory.strip():
            raise ValueError("final_directory must not be empty")
        for name in ("image_cache_dir", "watermark_font_path"):
            if not isinstance(getattr(self, name), str):
                raise TypeError(f"{name} must be a string")
        if (
            isinstance(self.max_image_transform_workers, bool)
            or not isinstance(self.max_image_transform_workers, int)
            or self.max_image_transform_workers < 1
            or self.max_image_transform_workers > MAX_IMAGE_TRANSFORM_WORKERS
        ):
            raise ValueError(
                "max_image_transform_workers must be between 1 and "
                f"{MAX_IMAGE_TRANSFORM_WORKERS}"
            )
        try:
            provider = (
                self.office_provider
                if isinstance(self.office_provider, OfficeImageProvider)
                else OfficeImageProvider(str(self.office_provider))
            )
        except ValueError as exc:
            raise ValueError("unsupported Office image provider") from exc
        object.__setattr__(self, "office_provider", provider)
        if not isinstance(self.layout_options, Mapping):
            raise TypeError("layout_options must be a mapping")
        object.__setattr__(
            self,
            "layout_options",
            MappingProxyType(dict(self.layout_options)),
        )


class AttachmentBundleService:
    """Process every declared item in staging, then publish one binding atomically."""

    def __init__(
        self,
        document_loader: Callable[[str], object] = Document,
        image_executor: ImageDocumentExecutor | None = None,
    ) -> None:
        if not callable(document_loader):
            raise TypeError("document_loader must be callable")
        self._document_loader = document_loader
        if image_executor is not None and not isinstance(
            image_executor,
            ImageDocumentExecutor,
        ):
            raise TypeError("image_executor must be an ImageDocumentExecutor")
        self._image_executor = image_executor or ImageDocumentExecutor()

    def process(
        self,
        request: AttachmentBundleBuildRequest,
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> AttachmentBundleReceipt:
        if not isinstance(request, AttachmentBundleBuildRequest):
            raise TypeError("request must be an AttachmentBundleBuildRequest")
        binding = _binding_for_role(request.snapshot, request.binding_role)
        final = Path(request.final_directory).expanduser().resolve(strict=False)
        if final.is_symlink():
            raise _error("attachment_final_symlink", "final directory cannot be a symlink")
        final.parent.mkdir(parents=True, exist_ok=True)
        _validate_final_does_not_own_sources(final, binding)
        plans = _build_plans(request, binding)
        for plan in plans:
            _verify_input_ref(plan)
        transaction_id = uuid4().hex
        staging = final.parent / f".{final.name}.{transaction_id}.staging"
        image_cache = (
            Path(request.image_cache_dir).expanduser().resolve()
            if request.image_cache_dir.strip()
            else final.parent / ".lark-material-image-cache"
        )
        receipts: list[AttachmentFileReceipt] = []
        try:
            staging.mkdir(parents=False, exist_ok=False)
            for plan in plans:
                if cancel_check is not None and cancel_check():
                    raise _error(
                        "attachment_processing_cancelled",
                        "attachment processing was cancelled",
                        plan=plan,
                    )
                receipts.append(
                    self._process_file(
                        plan,
                        binding=binding,
                        request=request,
                        staging=staging,
                        image_cache=image_cache,
                        transaction_id=transaction_id,
                        cancel_check=cancel_check,
                    )
                )
            for plan in plans:
                _verify_input_ref(plan)
            receipt = AttachmentBundleReceipt(
                snapshot_id=request.snapshot.snapshot_id,
                binding_role=binding.role,
                binding_revision=binding.binding_revision,
                dependency_index_id=request.dependency_index.index_id,
                output_directory=str(final),
                files=tuple(receipts),
            )
            atomic_write_text(
                staging / _RECEIPT_FILENAME,
                json.dumps(receipt.to_dict(), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            publish_staged_directory(
                staging,
                final,
                transaction_id=transaction_id,
            )
            return receipt
        except AttachmentBundleProcessingError:
            raise
        except DirectoryPublishError as exc:
            raise _error(
                "attachment_bundle_publish_failed",
                str(exc),
            ) from exc
        except Exception as exc:
            raise _error(
                "attachment_bundle_internal_error",
                f"{type(exc).__name__}: {exc}",
            ) from exc
        finally:
            remove_owned_path(staging, final.parent)

    def _process_file(
        self,
        plan: AttachmentDocumentPlan,
        *,
        binding: AttachmentBinding,
        request: AttachmentBundleBuildRequest,
        staging: Path,
        image_cache: Path,
        transaction_id: str,
        cancel_check: Callable[[], bool] | None,
    ) -> AttachmentFileReceipt:
        snapshot = request.snapshot
        dependency_index = request.dependency_index
        source = Path(plan.input_ref.source_path)
        target = staging.joinpath(*plan.output_relative_path.split("/"))
        _assert_within(target, staging)
        target.parent.mkdir(parents=True, exist_ok=True)
        dependencies = dependency_index.occurrences_for_consumer(plan.consumer_id)
        diagnostics = dependency_index.diagnostics_for_consumer(plan.consumer_id)
        unresolved = tuple(
            dict.fromkeys(
                [item.token for item in dependencies]
                + [item.token for item in diagnostics if item.token]
            )
        )

        if binding.processing_mode is AttachmentProcessingMode.PASSTHROUGH:
            shutil.copy2(source, target)
            return _file_receipt(
                plan,
                target,
                status=AttachmentFileStatus.PASSTHROUGH,
                unresolved_tokens=unresolved,
            )
        if Path(plan.output_relative_path).suffix.casefold() != ".docx":
            shutil.copy2(source, target)
            return _file_receipt(
                plan,
                target,
                status=AttachmentFileStatus.NOT_APPLICABLE,
            )

        content_dependencies = tuple(
            item for item in dependencies if item.kind is MaterialTokenKind.CONTENT
        )
        if content_dependencies:
            raise _error(
                "attachment_content_token_unsupported",
                "attachment DOCX cannot consume nested content tokens",
                plan=plan,
            )
        attachment_dependencies = tuple(
            item
            for item in dependencies
            if item.kind is MaterialTokenKind.ATTACHMENT
        )
        if attachment_dependencies:
            raise _error(
                "attachment_nested_bundle_unsupported",
                "attachment DOCX cannot consume nested attachment tokens",
                plan=plan,
            )
        image_dependencies = tuple(
            item
            for item in dependencies
            if item.kind is MaterialTokenKind.IMAGE and item.replaceable
        )
        field_dependencies = tuple(
            item
            for item in dependencies
            if item.kind is MaterialTokenKind.FIELD and item.replaceable
        )
        if not field_dependencies and not image_dependencies:
            shutil.copy2(source, target)
            return _file_receipt(
                plan,
                target,
                status=AttachmentFileStatus.NOT_APPLICABLE,
            )

        replacement_count = 0
        replaced_field_keys: tuple[str, ...] = ()
        if field_dependencies:
            try:
                document = self._document_loader(str(source))
            except Exception as exc:
                raise _error(
                    "attachment_docx_open_failed",
                    str(exc) or type(exc).__name__,
                    plan=plan,
                ) from exc
            replacements = {
                item.token: snapshot.field_values[item.resource_key]
                for item in field_dependencies
            }
            field_result = replace_document_exact_placeholders(
                document,
                replacements,
            )
            if field_result.total_replacements != len(field_dependencies):
                raise _error(
                    "attachment_field_replacement_count_mismatch",
                    f"expected {len(field_dependencies)}, "
                    f"replaced {field_result.total_replacements}",
                    plan=plan,
                )
            try:
                document.save(str(target))
            except Exception as exc:
                raise _error(
                    "attachment_docx_save_failed",
                    str(exc) or type(exc).__name__,
                    plan=plan,
                ) from exc
            replacement_count = field_result.total_replacements
            replaced_field_keys = tuple(field_result.replaced_keys)
        else:
            try:
                shutil.copy2(source, target)
            except OSError as exc:
                raise _error(
                    "attachment_copy_failed",
                    str(exc),
                    plan=plan,
                ) from exc

        if image_dependencies:
            frozen_rules, source_bindings = _select_image_inputs(
                snapshot,
                image_dependencies,
                plan=plan,
            )
            image_transaction_id = (
                f"attachment-{transaction_id}-"
                + sha256(plan.consumer_id.encode("utf-8")).hexdigest()[:16]
            )
            try:
                image_result = self._image_executor.execute(
                    ImageDocumentExecutionRequest(
                        owned_docx_path=str(target),
                        snapshot_id=snapshot.snapshot_id,
                        consumer_id=plan.consumer_id,
                        frozen_rules=frozen_rules,
                        source_bindings=source_bindings,
                        cache_dir=str(image_cache),
                        transaction_id=image_transaction_id,
                        office_provider=request.office_provider,
                        watermark_font_path=request.watermark_font_path,
                        max_transform_workers=(
                            request.max_image_transform_workers
                        ),
                        layout_options=request.layout_options,
                    ),
                    cancel_check=cancel_check,
                )
            except ImageDocumentExecutionError as exc:
                raise _error(
                    f"attachment_{exc.code}",
                    str(exc),
                    plan=plan,
                ) from exc
            plan = replace(
                plan,
                plan_id="",
                resolved_image_plans=image_result.plans,
            )

        try:
            verified_document = self._document_loader(str(target))
        except Exception as exc:
            raise _error(
                "attachment_docx_verify_open_failed",
                str(exc) or type(exc).__name__,
                plan=plan,
            ) from exc
        remaining = route_material_tokens(
            extract_docx_material_token_blocks(verified_document).blocks,
            image_tokens=(rule.anchor_token for rule in snapshot.frozen_image_rules),
            content_rules=snapshot.content_rules,
            consumer_id=plan.consumer_id,
        )
        remaining_tokens = tuple(
            dict.fromkeys(item.token for item in remaining.occurrences)
        )
        if remaining_tokens:
            raise _error(
                "attachment_unresolved_tokens",
                "unresolved tokens remain after substitution: "
                + ", ".join(remaining_tokens),
                plan=plan,
            )
        return _file_receipt(
            plan,
            target,
            status=AttachmentFileStatus.SUBSTITUTED,
            field_replacement_count=replacement_count,
            replaced_field_keys=replaced_field_keys,
        )


def _select_image_inputs(
    snapshot: MaterialSnapshot,
    dependencies: Sequence[MaterialDependencyOccurrence],
    *,
    plan: AttachmentDocumentPlan,
) -> tuple[
    tuple[FrozenImageMaterialRule, ...],
    tuple[ImageSourceBinding, ...],
]:
    tokens = frozenset(item.token for item in dependencies)
    rules = tuple(
        rule
        for rule in snapshot.frozen_image_rules
        if rule.anchor_token in tokens
    )
    selected_tokens = {rule.anchor_token for rule in rules}
    missing_tokens = sorted(tokens - selected_tokens, key=str.casefold)
    if missing_tokens:
        raise _error(
            "attachment_image_rule_selection_failed",
            "dependency index references missing image rules: "
            + ", ".join(missing_tokens),
            plan=plan,
        )
    selected_roles = {rule.source_role for rule in rules}
    sources = tuple(
        source
        for source in snapshot.image_source_bindings
        if source.role in selected_roles
    )
    roles_with_sources = {source.role for source in sources}
    missing_roles = sorted(selected_roles - roles_with_sources, key=str.casefold)
    if missing_roles:
        raise _error(
            "attachment_image_source_selection_failed",
            "selected image rules have no frozen source: "
            + ", ".join(missing_roles),
            plan=plan,
        )
    return rules, sources


def process_attachment_bundle(
    request: AttachmentBundleBuildRequest,
    *,
    cancel_check: Callable[[], bool] | None = None,
) -> AttachmentBundleReceipt:
    return AttachmentBundleService().process(request, cancel_check=cancel_check)


def _binding_for_role(snapshot: MaterialSnapshot, role: str) -> AttachmentBinding:
    normalized = str(role or "").strip()
    binding = next(
        (item for item in snapshot.attachment_bindings if item.role == normalized),
        None,
    )
    if binding is None:
        raise _error(
            "attachment_binding_missing",
            f"snapshot has no attachment binding {normalized!r}",
        )
    return binding


def _build_plans(
    request: AttachmentBundleBuildRequest,
    binding: AttachmentBinding,
) -> tuple[AttachmentDocumentPlan, ...]:
    plans: list[AttachmentDocumentPlan] = []
    for item in binding.items:
        if item.relative_path.casefold() == _RECEIPT_FILENAME.casefold():
            raise _error(
                "attachment_reserved_path_collision",
                f"{item.relative_path!r} is reserved for the bundle receipt",
            )
        consumer = MaterialConsumerRef(
            MaterialConsumerKind.ATTACHMENT_ITEM,
            binding.role,
            item.item_id,
        )
        indexed_revision = request.dependency_index.source_revisions.get(
            consumer.consumer_id,
            "",
        )
        if indexed_revision != item.file_ref.content_sha256:
            raise _error(
                "attachment_dependency_revision_mismatch",
                "dependency index does not match the frozen attachment source",
                consumer_id=consumer.consumer_id,
                item_id=item.item_id,
                relative_path=item.relative_path,
            )
        blocking = tuple(
            diagnostic
            for diagnostic in request.dependency_index.diagnostics_for_consumer(
                consumer.consumer_id
            )
            if diagnostic.severity is DependencyDiagnosticSeverity.ERROR
        )
        if binding.processing_mode is AttachmentProcessingMode.SUBSTITUTE_COPY and blocking:
            raise _error(
                "attachment_dependency_invalid",
                ", ".join(item.code for item in blocking),
                consumer_id=consumer.consumer_id,
                item_id=item.item_id,
                relative_path=item.relative_path,
            )
        plans.append(
            AttachmentDocumentPlan(
                plan_id="",
                snapshot_id=request.snapshot.snapshot_id,
                consumer_id=consumer.consumer_id,
                binding_role=binding.role,
                item_id=item.item_id,
                input_ref=item.file_ref,
                output_relative_path=item.relative_path,
                dependency_index_id=request.dependency_index.index_id,
            )
        )
    return tuple(plans)


def _verify_input_ref(plan: AttachmentDocumentPlan) -> None:
    path = Path(plan.input_ref.source_path)
    if not path.is_file():
        raise _error(
            "attachment_source_missing",
            f"source file is missing: {path}",
            plan=plan,
        )
    try:
        byte_size = path.stat().st_size
        digest = _sha256_file(path)
    except OSError as exc:
        raise _error(
            "attachment_source_unreadable",
            str(exc),
            plan=plan,
        ) from exc
    if byte_size != plan.input_ref.byte_size:
        raise _error(
            "attachment_source_size_drift",
            "source size changed after Freeze",
            plan=plan,
        )
    if digest != plan.input_ref.content_sha256:
        raise _error(
            "attachment_source_hash_drift",
            "source hash changed after Freeze",
            plan=plan,
        )


def _file_receipt(
    plan: AttachmentDocumentPlan,
    target: Path,
    *,
    status: AttachmentFileStatus,
    field_replacement_count: int = 0,
    replaced_field_keys: tuple[str, ...] = (),
    unresolved_tokens: tuple[str, ...] = (),
) -> AttachmentFileReceipt:
    return AttachmentFileReceipt(
        consumer_id=plan.consumer_id,
        item_id=plan.item_id,
        relative_path=plan.output_relative_path,
        input_sha256=plan.input_ref.content_sha256,
        output_sha256=_sha256_file(target),
        output_byte_size=target.stat().st_size,
        status=status,
        field_replacement_count=field_replacement_count,
        replaced_field_keys=replaced_field_keys,
        image_job_count=len(plan.resolved_image_plans),
        unresolved_tokens=unresolved_tokens,
    )


def _validate_final_does_not_own_sources(
    final: Path,
    binding: AttachmentBinding,
) -> None:
    for item in binding.items:
        source = Path(item.file_ref.source_path).resolve(strict=False)
        if source == final or final in source.parents:
            raise _error(
                "attachment_final_owns_source",
                "final directory cannot contain an immutable source attachment",
                item_id=item.item_id,
                relative_path=item.relative_path,
            )


def _assert_within(path: Path, root: Path) -> None:
    try:
        path.resolve(strict=False).relative_to(root.resolve(strict=False))
    except ValueError as exc:
        raise _error(
            "attachment_output_path_escape",
            f"output path escapes staging: {path}",
        ) from exc


def _sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _error(
    code: str,
    message: str,
    *,
    plan: AttachmentDocumentPlan | None = None,
    consumer_id: str = "",
    item_id: str = "",
    relative_path: str = "",
) -> AttachmentBundleProcessingError:
    return AttachmentBundleProcessingError(
        (
            AttachmentProcessingDiagnostic(
                code=code,
                message=message,
                consumer_id=(plan.consumer_id if plan is not None else consumer_id),
                item_id=(plan.item_id if plan is not None else item_id),
                relative_path=(
                    plan.output_relative_path if plan is not None else relative_path
                ),
            ),
        )
    )


__all__ = [
    "AttachmentBundleBuildRequest",
    "AttachmentBundleProcessingError",
    "AttachmentBundleService",
    "AttachmentProcessingDiagnostic",
    "process_attachment_bundle",
]
