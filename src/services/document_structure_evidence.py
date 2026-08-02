"""Immutable evidence and current-document decisions for logical regions."""

from __future__ import annotations

from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import BadZipFile, LargeZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError

from src.shared.files.content_hash import file_content_revision
from src.config.document_structure_contract import (
    DetectedRegion,
    DocumentStructureEvidence,
    ParagraphAnchor,
    RegionDecision,
    StructureReviewItem,
    pending_document_structure_review_roles,
    validate_region_decisions,
)
from src.config.section_semantics import canonicalize_section_type
from src.modules.structure.heading_recognition import analyze_document_tree


DETECTOR_REVISION = "logical-regions:v2"
_DIGEST_RE = re.compile(r"^sha256:[0-9a-f]{64}$")
_DOCUMENT_ANALYSIS_ERRORS = (
    OSError,
    ValueError,
    KeyError,
    BadZipFile,
    LargeZipFile,
    PackageNotFoundError,
    XMLSyntaxError,
)


def build_document_structure_evidence(
    source_path: str | Path,
) -> DocumentStructureEvidence:
    path = _canonical_path(source_path)
    if not path or Path(path).suffix.lower() != ".docx":
        return _failed_evidence(path, "document_structure_source_invalid")

    source = Path(path)
    revision_before = _read_revision(source)
    if not revision_before:
        return _failed_evidence(path, "document_structure_source_unreadable")

    try:
        doc = Document(str(source))
    except _DOCUMENT_ANALYSIS_ERRORS:
        return _failed_evidence(
            path,
            "document_structure_analysis_failed",
            source_revision=revision_before,
        )

    revision_after = _read_revision(source)
    if not revision_after:
        return _failed_evidence(path, "document_structure_source_unreadable")
    if revision_after != revision_before:
        return _failed_evidence(
            path,
            "document_structure_source_changed",
            source_revision=revision_after,
        )

    try:
        return _build_evidence_from_document(path, revision_after, doc)
    except _DOCUMENT_ANALYSIS_ERRORS:
        return _failed_evidence(
            path,
            "document_structure_analysis_failed",
            source_revision=revision_after,
        )


def build_document_structure_evidence_from_bytes(
    source_path: str | Path,
    source_bytes: bytes | None,
    *,
    source_revision: str = "",
) -> DocumentStructureEvidence:
    """Analyse the exact input bytes already frozen for an execution session."""

    path = _canonical_path(source_path)
    payload = bytes(source_bytes or b"")
    revision = str(source_revision or "").strip()
    if not path or Path(path).suffix.lower() != ".docx":
        return _failed_evidence(path, "document_structure_source_invalid")
    if not payload:
        return _failed_evidence(path, "document_structure_source_unreadable")
    if not revision:
        revision = "sha256:" + sha256(payload).hexdigest()
    try:
        doc = Document(BytesIO(payload))
        return _build_evidence_from_document(path, revision, doc)
    except _DOCUMENT_ANALYSIS_ERRORS:
        return _failed_evidence(
            path,
            "document_structure_analysis_failed",
            source_revision=revision,
        )


def _build_evidence_from_document(
    path: str,
    source_revision: str,
    doc,
) -> DocumentStructureEvidence:
    tree = analyze_document_tree(doc)
    paragraphs = tuple(str(paragraph.text or "").strip() for paragraph in doc.paragraphs)
    anchors = _paragraph_anchors(paragraphs)
    regions: list[DetectedRegion] = []
    review_items: list[StructureReviewItem] = []
    for section in tree.sections:
        start_index = int(section.start_index)
        end_index = int(section.end_index)
        if start_index < 0 or start_index >= len(anchors) or end_index <= start_index:
            continue
        role_id = canonicalize_section_type(section.section_type)
        accepted = bool(section.title_confident) and float(section.confidence) >= 8.0
        status = "accepted" if accepted else "review"
        codes = (
            ("title_evidence",)
            if accepted
            else ("content_or_position_fallback",)
        )
        region = DetectedRegion(
            role_id=role_id,
            start_anchor=anchors[start_index],
            end_anchor=(anchors[end_index] if end_index < len(anchors) else None),
            detection_status=status,
            evidence_codes=codes,
        )
        regions.append(region)
        if status == "review":
            review_items.append(
                StructureReviewItem(
                    role_id=role_id,
                    code="region_requires_confirmation",
                )
            )

    payload = {
        "source_path": path,
        "source_revision": source_revision,
        "detector_revision": DETECTOR_REVISION,
        "regions": [asdict(region) for region in regions],
        "review_items": [asdict(item) for item in review_items],
        "issues": [],
    }
    return DocumentStructureEvidence(
        source_path=path,
        source_revision=source_revision,
        detector_revision=DETECTOR_REVISION,
        evidence_digest=_payload_digest(payload),
        regions=tuple(regions),
        review_items=tuple(review_items),
    )


def document_structure_evidence_is_current(
    evidence: DocumentStructureEvidence,
    source_path: str | Path,
) -> bool:
    if not isinstance(evidence, DocumentStructureEvidence) or not evidence.ready:
        return False
    path = _canonical_path(source_path)
    if path != evidence.source_path:
        return False
    if evidence.detector_revision != DETECTOR_REVISION:
        return False
    if not _DIGEST_RE.fullmatch(evidence.evidence_digest):
        return False
    return _read_revision(Path(path)) == evidence.source_revision


def read_document_paragraph_anchors(
    source_path: str | Path,
) -> tuple[ParagraphAnchor, ...]:
    """Return current non-empty paragraph choices for the review dialog."""

    path = _canonical_path(source_path)
    if not path:
        return ()
    try:
        doc = Document(path)
    except _DOCUMENT_ANALYSIS_ERRORS:
        return ()
    texts = tuple(str(paragraph.text or "").strip() for paragraph in doc.paragraphs)
    return tuple(
        anchor
        for anchor in _paragraph_anchors(texts)
        if anchor.preview_text
    )


def _paragraph_anchors(paragraphs: tuple[str, ...]) -> tuple[ParagraphAnchor, ...]:
    digests = tuple(_text_digest(text) for text in paragraphs)
    occurrences: dict[str, int] = {}
    anchors: list[ParagraphAnchor] = []
    for index, text in enumerate(paragraphs):
        digest = digests[index]
        occurrence = occurrences.get(digest, 0)
        occurrences[digest] = occurrence + 1
        anchors.append(
            ParagraphAnchor(
                source_index=index,
                text_digest=digest,
                previous_text_digest=digests[index - 1] if index > 0 else "",
                next_text_digest=digests[index + 1] if index + 1 < len(digests) else "",
                occurrence=occurrence,
                preview_text=text[:48],
            )
        )
    return tuple(anchors)


def _failed_evidence(
    source_path: str,
    issue: str,
    *,
    source_revision: str = "",
) -> DocumentStructureEvidence:
    payload = {
        "source_path": source_path,
        "source_revision": source_revision,
        "detector_revision": DETECTOR_REVISION,
        "regions": [],
        "review_items": [],
        "issues": [issue],
    }
    return DocumentStructureEvidence(
        source_path=source_path,
        source_revision=source_revision,
        detector_revision=DETECTOR_REVISION,
        evidence_digest=_payload_digest(payload),
        regions=(),
        review_items=(),
        issues=(issue,),
    )


def _payload_digest(payload: dict[str, object]) -> str:
    canonical = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return "sha256:" + sha256(canonical).hexdigest()


def _text_digest(text: str) -> str:
    normalized = " ".join(str(text or "").split()).casefold().encode("utf-8")
    return sha256(normalized).hexdigest()


def _canonical_path(source_path: str | Path) -> str:
    cleaned = str(source_path or "").strip()
    if not cleaned:
        return ""
    try:
        return str(Path(cleaned).resolve())
    except (OSError, RuntimeError):
        return str(Path(cleaned))


def _read_revision(path: Path) -> str:
    try:
        return file_content_revision(path) if path.is_file() else ""
    except OSError:
        return ""


__all__ = [
    "DETECTOR_REVISION",
    "DetectedRegion",
    "DocumentStructureEvidence",
    "ParagraphAnchor",
    "RegionDecision",
    "StructureReviewItem",
    "build_document_structure_evidence",
    "build_document_structure_evidence_from_bytes",
    "document_structure_evidence_is_current",
    "pending_document_structure_review_roles",
    "read_document_paragraph_anchors",
    "validate_region_decisions",
]
