"""Pipeline execution entry point."""

from __future__ import annotations

import copy
import os
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Callable, Mapping
from uuid import uuid4

from docx import Document

from src.config.delivery_preset_display import delivery_preset_display_name
from src.config.execution_config_integrity import (
    delivery_preset_identity_issue,
    execution_config_integrity_issue,
)
from src.config.master_library import MasterSpec
from src.config.material_schema_registry import (
    get_material_schema,
    resolve_material_schema_ids,
)
from src.config.official_document_profiles import get_official_document_profile
from src.config.plugin_manual_gate import plugin_manual_gate_payload
from src.config.resolved import ResolvedConfig
from src.config.document_scope import (
    DocumentScopePolicy,
    document_scope_policy_issue,
)
from src.config.document_structure_contract import (
    DocumentStructureEvidence,
    RegionDecision,
)
from src.config.scene_product_coverage_manifest import (
    SceneCoveragePack,
    coverage_packs_for_config,
)
from src.config.scene_family_registry import get_planned_scene_family
from src.modules.base import BaseModule
from src.modules.structure.heading_recognition import rebuild_document_index
from src.pipeline.context import PipelineContext
from src.pipeline.journal_governance_mixin import PipelineJournalGovernanceMixin
from src.pipeline.module_phases import build_module_phase_graph, order_modules_by_phase
from src.pipeline.result import PipelineResult
from src.shared.engine.field_refresh import (
    document_has_toc,
    refresh_doc_fields_with_word,
)
from src.shared.engine.document_scope_runtime import bind_document_scope
from src.shared.engine.content_visibility import (
    extract_content_visibility_rule_selectors,
    preview_content_visibility_effects,
    scan_content_visibility_markers,
)
from src.shared.engine.block_visibility import apply_block_visibility
from src.shared.engine.material_field_consistency import (
    inspect_material_field_consistency,
)
from src.shared.engine.exam_question_schema import inspect_exam_question_schema
from src.shared.engine.exam_question_schema import build_exam_delivery_runtime
from src.shared.engine.journal_citation_schema import inspect_journal_citations
from src.shared.engine.official_numbering_preservation import (
    inspect_official_numbering_preservation,
)
from src.shared.engine.official_document_assembly import (
    assemble_official_document_docx,
)
from src.shared.engine.technical_chapter_inventory import (
    inspect_technical_chapter_inventory,
)
from src.shared.engine.application_section_word_limits import (
    inspect_application_section_word_limits,
)
from src.pipeline.scheduler import (
    validate_data_flow,
    validate_schema_contract,
)
from src.pipeline.tracker import ChangeTracker
from src.shared.io.artifact_transaction import (
    OwnedAssemblyTransaction,
    stable_file_evidence,
)
from src.shared.io.safe_docx_package import (
    DocxPackageError,
    SafeDocxPackage,
    capture_bounded_file,
)
from src.shared.engine.object_preflight import (
    inspect_docx_package,
    object_preflight_module_skips,
    object_preflight_targets_for_touchpoints,
)
from src.shared.engine.scene_journey_runtime import (
    build_scene_journey_runtime_evidence,
)


class _PipelineCancellationRequested(RuntimeError):
    """Abort an uncommitted output transaction as a cancellation, not a failure."""


@dataclass(slots=True)
class OutputTargetIssue:
    kind: str = ""
    severity: str = "warning"
    message: str = ""


@dataclass(slots=True)
class OutputTargetPreflightItem:
    preset_id: str = ""
    label: str = ""
    path: str = ""
    artifact: str = "final_docx"
    exists: bool = False
    parent_path: str = ""
    parent_exists: bool = False
    parent_writable: bool = True
    path_length: int = 0
    issues: list[OutputTargetIssue] = field(default_factory=list)

    @property
    def has_issues(self) -> bool:
        return bool(self.issues)


@dataclass(slots=True)
class OutputTargetPreflightResult:
    items: list[OutputTargetPreflightItem] = field(default_factory=list)

    @property
    def has_issues(self) -> bool:
        return any(item.has_issues for item in self.items)

    @property
    def issue_count(self) -> int:
        return sum(len(item.issues) for item in self.items)

    @property
    def has_errors(self) -> bool:
        return any(
            _output_target_issue_is_blocking(issue)
            for item in self.items
            for issue in item.issues
        )

    @property
    def error_count(self) -> int:
        return sum(
            1
            for item in self.items
            for issue in item.issues
            if _output_target_issue_is_blocking(issue)
        )


def build_coverage_boundary_report_items(config) -> list[dict[str, object]]:
    """Build report-ready coverage/plugin boundary evidence for a config."""

    return [
        _coverage_boundary_payload(pack)
        for pack in coverage_packs_for_config(config)
        if _coverage_boundary_should_report(pack)
    ]


def _coverage_boundary_should_report(pack: SceneCoveragePack) -> bool:
    if pack.plugin_boundary:
        return True
    return pack.pack_id == "contract_delivery"


def _coverage_boundary_payload(pack: SceneCoveragePack) -> dict[str, object]:
    return {
        "pack_id": pack.pack_id,
        "label": pack.label,
        "boundary": pack.boundary,
        "primary_landings": list(pack.primary_landings),
        "secondary_landings": list(pack.secondary_landings),
        "capability_axis_ids": list(pack.capability_axis_ids),
        "implemented_closures": list(pack.implemented_closures),
        "missing_closures": list(pack.missing_closures),
        "closure_tasks": [
            {
                "summary": task.summary,
                "priority": task.priority,
                "owner": task.owner,
                "target_phase": task.target_phase,
                "validation_commands": list(task.validation_commands),
            }
            for task in pack.closure_tasks
        ],
        "plugin_manual_gate": plugin_manual_gate_payload(pack.pack_id),
    }


class Pipeline(PipelineJournalGovernanceMixin):
    """Execute registered modules against a document."""

    def __init__(
        self,
        modules: list[BaseModule],
        config: ResolvedConfig,
        *,
        output_dir: str | None = None,
        output_suffix: str = "_new",
        progress_callback: Callable[[int, int, str], None] | None = None,
        cancel_check: Callable[[], bool] | None = None,
        force_delivery_presets: bool = False,
        official_master: MasterSpec | None = None,
        exam_master: MasterSpec | None = None,
        official_document_type_id: str = "",
        logical_source_path: str | Path | None = None,
        output_path_overrides: Mapping[str, str | Path] | None = None,
        output_paths_are_owned_stages: bool = False,
        defer_field_refresh: bool = False,
        document_structure_evidence: DocumentStructureEvidence | None = None,
        document_scope_decisions: tuple[RegionDecision, ...] = (),
    ) -> None:
        self._config = config
        self._output_dir = Path(output_dir) if output_dir else None
        self._output_suffix = output_suffix
        self._progress_callback = progress_callback
        self._cancel_check = cancel_check
        self._force_delivery_presets = bool(force_delivery_presets)
        self._official_master = official_master
        self._exam_master = exam_master
        self._official_document_type_id = str(
            official_document_type_id or ""
        ).strip()
        self._logical_source_path = (
            Path(logical_source_path) if logical_source_path is not None else None
        )
        self._output_path_overrides = {
            str(key): Path(value)
            for key, value in dict(output_path_overrides or {}).items()
            if str(key or "").strip()
        }
        self._output_paths_are_owned_stages = bool(output_paths_are_owned_stages)
        self._defer_field_refresh = bool(defer_field_refresh)
        self._document_structure_evidence = document_structure_evidence
        self._document_scope_decisions = tuple(document_scope_decisions or ())
        self._tracker = ChangeTracker()

        registered_modules = tuple(modules)
        self._phase_graph = build_module_phase_graph(registered_modules)
        self._phase_steps_by_name = {
            step.name: step for step in self._phase_graph.steps
        }
        self._modules = list(order_modules_by_phase(registered_modules))

        schema_errors = validate_schema_contract(self._modules)
        if schema_errors:
            raise ValueError(
                "Pipeline schema contract validation failed:\n"
                + "\n".join(f"  - {error}" for error in schema_errors)
            )

        data_flow_errors = validate_data_flow(self._modules)
        if data_flow_errors:
            raise ValueError(
                "Pipeline data-flow validation failed:\n"
                + "\n".join(f"  - {error}" for error in data_flow_errors)
            )

    def execute(self, doc_path: str) -> PipelineResult:
        """Run the full pipeline for one .docx file."""
        if self._is_cancelled():
            return self._cancelled_result()

        scope_issue = document_scope_policy_issue(
            getattr(self._config, "document_scope", None),
            mode_id=str(getattr(self._config, "mode_id", "") or "custom"),
        )
        if scope_issue:
            return PipelineResult(
                success=False,
                status="failed",
                error=scope_issue,
                config=self._config,
            )

        configuration_issue = pipeline_configuration_integrity_issue(self._config)
        if configuration_issue:
            return PipelineResult(
                success=False,
                status="failed",
                error=configuration_issue,
                config=self._config,
            )

        path = Path(doc_path)
        logical_path = self._logical_source_path or path
        terminal_owner = pipeline_terminal_assembly_owner(self._config)
        if not path.exists():
            return PipelineResult(
                success=False,
                error=f"File not found: {path}",
                config=self._config,
            )

        total_steps = (0 if terminal_owner else len(self._modules)) + 3
        step = 0

        self._emit(step, total_steps, "Loading document")
        if self._is_cancelled():
            return self._cancelled_result()

        try:
            source_payload = capture_bounded_file(path)
            safe_package = SafeDocxPackage.open(source_payload)
            safe_package.validate_xml_parts()
            doc = Document(BytesIO(source_payload))
        except DocxPackageError as exc:
            return PipelineResult(
                success=False,
                status="failed",
                error=f"unsafe_docx_package:{exc.code}",
                config=self._config,
            )
        except Exception as exc:
            return PipelineResult(
                success=False,
                error=f"Failed to open document: {exc}",
                config=self._config,
            )

        original_doc = copy.deepcopy(doc)
        step += 1

        ctx = PipelineContext(
            source_doc_path=str(logical_path),
            source_doc_dir=str(logical_path.parent),
            working_doc_path=str(path),
            working_doc_dir=str(path.parent),
            mode_id=str(getattr(self._config, "mode_id", "") or "custom"),
            document_scope=copy.deepcopy(
                getattr(self._config, "document_scope", None)
                or DocumentScopePolicy()
            ),
            document_scope_gate_active=True,
        )
        ctx.terminal_assembly_owner = terminal_owner
        ctx.coverage_boundaries = build_coverage_boundary_report_items(self._config)
        self._run_scene_journey_runtime(ctx)
        if terminal_owner in {"exam", "conflict"}:
            self._run_exam_question_schema_validation(ctx)

        if terminal_owner:
            return self._execute_terminal_assembly(
                terminal_owner=terminal_owner,
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )

        return self._execute_generic_pipeline(
            doc=doc,
            original_doc=original_doc,
            ctx=ctx,
            logical_path=logical_path,
            safe_package=safe_package,
            step=step,
            total_steps=total_steps,
        )

    def _execute_generic_pipeline(
        self,
        *,
        doc: Document,
        original_doc: Document,
        ctx: PipelineContext,
        logical_path: Path,
        safe_package: SafeDocxPackage,
        step: int,
        total_steps: int,
    ) -> PipelineResult:
        try:
            ctx.document_structure_evidence = self._document_structure_evidence
            ctx.document_scope_decisions = self._document_scope_decisions
            if self._document_structure_evidence is not None:
                ctx.document_scope_binding = bind_document_scope(
                    doc,
                    self._document_structure_evidence,
                    self._document_scope_decisions,
                    ctx.document_scope,
                    mode_id=ctx.mode_id,
                )
            rebuild_document_index(doc, ctx)
            self._capture_document_scope_receipt(ctx)
        except Exception as exc:
            return self._failed_result(
                error=f"document_scope_binding_failed:{exc}",
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )

        preflight_skips, preflight_result = self._run_generic_preflights(
            doc=doc,
            original_doc=original_doc,
            ctx=ctx,
            logical_path=logical_path,
            safe_package=safe_package,
            step=step,
            total_steps=total_steps,
        )
        if preflight_result is not None:
            return preflight_result
        step, module_result = self._execute_registered_modules(
            doc=doc,
            original_doc=original_doc,
            ctx=ctx,
            preflight_skips=preflight_skips,
            step=step,
            total_steps=total_steps,
        )
        if module_result is not None:
            return module_result
        return self._complete_generic_pipeline(
            doc=doc,
            original_doc=original_doc,
            ctx=ctx,
            logical_path=logical_path,
            step=step,
            total_steps=total_steps,
            )

    def _capture_document_scope_receipt(self, ctx: PipelineContext) -> None:
        evidence = ctx.document_structure_evidence
        tree = ctx.doc_tree
        if not isinstance(evidence, DocumentStructureEvidence) or tree is None:
            return
        decisions = tuple(ctx.document_scope_decisions or ())
        included_roles = set(ctx.document_scope.included_roles(ctx.mode_id))
        effective_roles = {
            section.section_type
            for section in tuple(getattr(tree, "sections", ()) or ())
            if not bool(getattr(section, "excluded", False))
            and section.section_type in included_roles
        }
        corrected_count = sum(
            decision.action in {"exclude", "set_start"}
            for decision in decisions
        )
        excluded_roles = {
            decision.role_id
            for decision in decisions
            if decision.action == "exclude"
        }
        review_roles = {
            item.role_id
            for item in evidence.review_items
            if item.role_id in included_roles
        }
        skipped_uncertain_count = len(review_roles & excluded_roles)
        ctx.document_scope_receipt = {
            "plan_mode": ctx.document_scope.mode,
            "detected_region_count": len(evidence.regions),
            "effective_region_count": len(effective_roles),
            "confirmation_status": (
                "user_confirmed" if decisions else "auto_accepted"
            ),
            "corrected_region_count": corrected_count,
            "skipped_uncertain_count": skipped_uncertain_count,
        }
        summaries: list[str] = []
        if corrected_count:
            summaries.append(f"{corrected_count} 项已校正")
        if skipped_uncertain_count:
            summaries.append(f"{skipped_uncertain_count} 项已跳过")
        if summaries:
            self._tracker.record(
                rule_name="document_scope",
                target="文档范围",
                section="global",
                change_type="scope",
                before="",
                after="，".join(summaries),
            )

    def _run_generic_preflights(
        self,
        *,
        doc: Document,
        original_doc: Document,
        ctx: PipelineContext,
        logical_path: Path,
        safe_package: SafeDocxPackage,
        step: int,
        total_steps: int,
    ) -> tuple[dict[str, dict[str, object]], PipelineResult | None]:
        self._run_journal_rule_source_governance(ctx)
        self._run_journal_citation_validation(doc, ctx)
        self._emit(step, total_steps, "Object preflight")
        preflight_blocked, preflight_skips = self._run_object_preflight(
            Path(ctx.working_doc_path),
            ctx,
            safe_package=safe_package,
        )
        if preflight_blocked:
            return preflight_skips, self._failed_result(
                error="Object preflight blocked execution due to fragile embedded content.",
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )
        visibility_preflight_error = self._run_content_visibility_preflight(doc, ctx)
        if visibility_preflight_error:
            return preflight_skips, self._failed_result(
                error=visibility_preflight_error,
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )
        output_preflight_error = self._run_output_target_preflight(logical_path, ctx)
        if output_preflight_error:
            return preflight_skips, self._failed_result(
                error=output_preflight_error,
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )
        return preflight_skips, None

    def _execute_registered_modules(
        self,
        *,
        doc: Document,
        original_doc: Document,
        ctx: PipelineContext,
        preflight_skips: Mapping[str, Mapping[str, object]],
        step: int,
        total_steps: int,
    ) -> tuple[int, PipelineResult | None]:
        for mod in self._modules:
            if self._is_cancelled():
                return step, self._cancelled_result()

            skip_entry = preflight_skips.get(mod.meta.name)
            if skip_entry:
                skip_reason = str(skip_entry.get("reason", "") or "")
                self._record_parameter_runtime_consumption(
                    ctx,
                    mod,
                    status="skipped",
                    reason=skip_reason,
                )
                self._tracker.record(
                    rule_name=mod.meta.name,
                    target="object_preflight",
                    section="global",
                    change_type="skip",
                    before="",
                    after=skip_reason,
                    paragraph_index=-1,
                    success=True,
                )
                step += 1
                continue

            self._emit(step, total_steps, f"Running module: {mod.meta.description}")

            try:
                issues = mod.validate(doc, self._config, ctx) or []
                fatal_issues = self._record_validation_issues(mod, issues)
                if fatal_issues:
                    error = (
                        f"Module '{mod.meta.name}' validation failed: "
                        f"{fatal_issues[0].message}"
                    )
                    self._record_parameter_runtime_consumption(
                        ctx,
                        mod,
                        status="validation_failed",
                        reason=fatal_issues[0].message,
                    )
                    if self._is_strict_mode():
                        return step, self._failed_result(
                            error=error,
                            doc=doc,
                            original_doc=original_doc,
                            ctx=ctx,
                        )
                    step += 1
                    continue

                mod.apply(doc, self._config, self._tracker, ctx)
                phase_step = self._phase_steps_by_name[mod.meta.name]
                if phase_step.invalidates_document_index:
                    self._rebuild_document_index_after_module(doc, ctx, mod)
                self._record_parameter_runtime_consumption(
                    ctx,
                    mod,
                    status="executed",
                )
            except Exception as exc:
                self._record_parameter_runtime_consumption(
                    ctx,
                    mod,
                    status="failed",
                    reason=str(exc),
                )
                self._tracker.record(
                    rule_name=mod.meta.name,
                    target="pipeline",
                    section="global",
                    change_type="error",
                    paragraph_index=-1,
                    success=False,
                    failure_reason=str(exc),
                )
                if self._is_strict_mode():
                    return step, self._failed_result(
                        error=f"Module '{mod.meta.name}' execution failed: {exc}",
                        doc=doc,
                        original_doc=original_doc,
                        ctx=ctx,
                    )

            step += 1
        return step, None

    def _complete_generic_pipeline(
        self,
        *,
        doc: Document,
        original_doc: Document,
        ctx: PipelineContext,
        logical_path: Path,
        step: int,
        total_steps: int,
    ) -> PipelineResult:
        if self._is_cancelled():
            return self._cancelled_result()

        self._run_official_numbering_preservation_validation(doc, original_doc, ctx)
        self._run_technical_chapter_inventory_validation(doc, ctx)
        self._run_application_section_word_limits_validation(doc, ctx)
        self._run_material_field_consistency(doc, ctx)
        if self._is_cancelled():
            return self._cancelled_result(
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )

        self._emit(step, total_steps, "Writing report")
        self._emit(step, total_steps, "Saving output")
        try:
            output_paths = self._save_outputs(doc, logical_path, ctx=ctx)
        except _PipelineCancellationRequested:
            return self._cancelled_result(
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )
        except Exception as exc:
            return self._failed_result(
                error=f"Output staging failed: {exc}",
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )
        output_paths = (
            _exam_delivery_runtime_output_paths(ctx)
            or _official_document_assembly_output_paths(ctx)
            or output_paths
        )
        if self._is_cancelled():
            return self._cancelled_result(
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
                output_paths=output_paths,
            )
        self._run_journal_submission_package_validation(output_paths, ctx)
        step += 1

        self._emit(step, total_steps, "Completed")
        return self._build_result(
            doc=doc,
            original_doc=original_doc,
            output_paths=output_paths,
            ctx=ctx,
        )

    def _execute_terminal_assembly(
        self,
        *,
        terminal_owner: str,
        doc: Document,
        original_doc: Document,
        ctx: PipelineContext,
    ) -> PipelineResult:
        """Let one specialized assembler own mutation and publication end to end."""

        if terminal_owner == "conflict":
            error = (
                "Terminal assembly ownership is ambiguous: exam and official "
                "schemas cannot execute in the same pipeline."
            )
            self._record_terminal_assembly_owner(
                terminal_owner,
                success=False,
                error=error,
            )
            return self._failed_result(
                error=error,
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )

        if self._output_paths_are_owned_stages:
            error = (
                f"Terminal assembler '{terminal_owner}' cannot run inside the "
                "generic caller-owned stage transaction."
            )
            self._record_terminal_assembly_owner(
                terminal_owner,
                success=False,
                error=error,
            )
            return self._failed_result(
                error=error,
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )

        if self._is_cancelled():
            return self._cancelled_result()

        try:
            if terminal_owner == "exam":
                self._emit(1, 3, "Building exam delivery artifacts")
                self._run_exam_delivery_runtime(ctx)
                runtime = ctx.exam_delivery_runtime
                output_paths = _exam_delivery_runtime_output_paths(ctx)
                status = str(getattr(runtime, "status", "") or "")
                accepted = status in {"ok", "warning"} and bool(output_paths)
                reason = str(getattr(runtime, "skipped_reason", "") or status)
            elif terminal_owner == "official":
                self._emit(1, 3, "Building official document artifacts")
                self._run_official_document_assembly_runtime(ctx)
                runtime = ctx.official_document_assembly
                output_paths = _official_document_assembly_output_paths(ctx)
                status = str(getattr(runtime, "status", "") or "")
                accepted = status == "ok" and bool(output_paths)
                reason = status
            else:
                raise RuntimeError(f"unsupported terminal assembly owner: {terminal_owner}")
        except Exception as exc:
            error = f"Terminal assembler '{terminal_owner}' failed: {exc}"
            self._record_terminal_assembly_owner(
                terminal_owner,
                success=False,
                error=error,
            )
            return self._failed_result(
                error=error,
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )

        if self._is_cancelled():
            return self._cancelled_result(
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
                output_paths=output_paths,
            )

        if not accepted:
            error = (
                f"Terminal assembler '{terminal_owner}' blocked delivery: "
                f"{reason or 'no publishable artifacts'}"
            )
            self._record_terminal_assembly_owner(
                terminal_owner,
                success=False,
                error=error,
            )
            return self._failed_result(
                error=error,
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
            )

        self._record_terminal_assembly_owner(terminal_owner, success=True)
        self._emit(2, 3, "Completed")
        return self._build_result(
            doc=doc,
            original_doc=original_doc,
            output_paths=output_paths,
            ctx=ctx,
        )

    def _record_terminal_assembly_owner(
        self,
        owner: str,
        *,
        success: bool,
        error: str = "",
    ) -> None:
        self._tracker.record(
            rule_name="terminal_assembly_owner",
            target=owner,
            section="pipeline",
            change_type=(
                "terminal_assembly_owned"
                if success
                else "terminal_assembly_blocked"
            ),
            before=f"generic_modules={len(self._modules)}",
            after=(
                f"owner={owner}; generic modules and ordinary output publication skipped"
                if success
                else error
            ),
            paragraph_index=-1,
            success=success,
            failure_reason=error or None,
        )

    def _rebuild_document_index_after_module(
        self,
        doc: Document,
        ctx: PipelineContext,
        mod: BaseModule,
    ) -> None:
        """Consume the graph's invalidation contract immediately on success."""

        previous_tree = ctx.doc_tree
        previous_heading_map = ctx.heading_map
        previous_count = len(getattr(previous_tree, "headings", ()) or ())
        # Invalidate before rebuilding so a failed rebuild can never leave a
        # stale index available to later modules in non-strict mode.
        ctx.doc_tree = None
        ctx.heading_map = None
        rebuilt = rebuild_document_index(doc, ctx)
        self._tracker.record(
            rule_name=mod.meta.name,
            target="document_index",
            section="pipeline",
            change_type="document_index_rebuild",
            before=(
                f"headings={previous_count}; "
                f"heading_map={'present' if previous_heading_map is not None else 'absent'}"
            ),
            after=(
                f"headings={len(getattr(rebuilt, 'headings', ()) or ())}; "
                f"paragraphs={len(doc.paragraphs)}"
            ),
            paragraph_index=-1,
            success=True,
        )

    def _emit(self, current: int, total: int, message: str) -> None:
        if self._progress_callback:
            self._progress_callback(current, total, message)

    def _is_cancelled(self) -> bool:
        if self._cancel_check is None:
            return False
        try:
            return bool(self._cancel_check())
        except Exception:
            return False

    def _is_strict_mode(self) -> bool:
        return bool(getattr(self._config, "strict_mode", False))

    def _cancelled_result(
        self,
        *,
        doc: Document | None = None,
        original_doc: Document | None = None,
        ctx: PipelineContext | None = None,
        output_paths: Mapping[str, str | Path] | None = None,
    ) -> PipelineResult:
        """Return cancellation without discarding already-published evidence."""

        return PipelineResult(
            success=False,
            doc=doc,
            original_doc=original_doc,
            tracker=self._tracker,
            context=ctx,
            output_paths={
                str(key): str(value)
                for key, value in dict(output_paths or {}).items()
            },
            failed_items=self._collect_failures(),
            error="Cancelled by user.",
            cancelled=True,
            config=self._config,
        )

    def _failed_result(
        self,
        *,
        error: str,
        doc: Document | None = None,
        original_doc: Document | None = None,
        ctx: PipelineContext | None = None,
    ) -> PipelineResult:
        return PipelineResult(
            success=False,
            status="failed",
            doc=doc,
            original_doc=original_doc,
            tracker=self._tracker,
            context=ctx,
            failed_items=self._collect_failures(),
            error=error,
            config=self._config,
        )

    def _collect_failures(self) -> list[dict]:
        return [
            {
                "rule_name": record.rule_name,
                "target": record.target,
                "section": record.section,
                "change_type": record.change_type,
                "paragraph_index": record.paragraph_index,
                "reason": record.failure_reason or "",
            }
            for record in self._tracker.get_failures()
        ]

    def _record_validation_issues(self, mod: BaseModule, issues: list) -> list:
        fatal_issues = []

        for issue in issues:
            is_error = getattr(issue, "level", "") == "error"
            if is_error:
                fatal_issues.append(issue)

            self._tracker.record(
                rule_name=mod.meta.name,
                target=getattr(issue, "location", "") or "validate",
                section="global",
                change_type=f"validate_{getattr(issue, 'level', 'info')}",
                after=getattr(issue, "message", ""),
                paragraph_index=-1,
                success=not is_error,
                failure_reason=getattr(issue, "message", "") if is_error else None,
            )

        return fatal_issues

    def _record_parameter_runtime_consumption(
        self,
        ctx: PipelineContext | None,
        mod: BaseModule,
        *,
        status: str,
        reason: str = "",
    ) -> None:
        if ctx is None:
            return
        if ctx.parameter_runtime_consumption is None:
            ctx.parameter_runtime_consumption = []
        ctx.parameter_runtime_consumption.append(
            {
                "module_name": mod.meta.name,
                "description": mod.meta.description,
                "category": mod.meta.category,
                "status": status,
                "reason": str(reason or ""),
                "config_paths": [
                    str(path or "").strip()
                    for path in getattr(mod.meta, "requires_config", ()) or ()
                    if str(path or "").strip()
                ],
            }
        )

    def _run_scene_journey_runtime(self, ctx: PipelineContext) -> None:
        result = build_scene_journey_runtime_evidence(self._config)
        if not result.is_applicable:
            return
        ctx.scene_journey_runtime = result
        self._tracker.record(
            rule_name="scene_journey_runtime",
            target=";".join(result.pack_ids) or "scene_matrix",
            section="scene",
            change_type="scene_journey_runtime",
            before="",
            after=(
                f"status={result.status}; scope={result.evidence_scope}; "
                f"contracts={result.contract_count}; "
                f"static_audit={result.static_audit_status}; "
                f"reports={len(result.report_expectations)}; "
                f"artifacts={len(result.artifact_channel_ids)}; "
                f"repair={len(result.repair_target_types)}"
            ),
            paragraph_index=-1,
            success=True,
        )

    def _run_object_preflight(
        self,
        path: Path,
        ctx: PipelineContext,
        *,
        safe_package: SafeDocxPackage,
    ) -> tuple[bool, dict[str, dict[str, object]]]:
        compliance = getattr(self._config, "compliance_profile", None)
        policy = getattr(compliance, "object_preflight", None)
        if policy is None or not bool(getattr(policy, "enabled", True)):
            return False, {}

        ctx.object_preflight_policy = self._build_object_preflight_policy_snapshot(policy)
        result = inspect_docx_package(path, policy, safe_package=safe_package)
        ctx.object_preflight = result
        for finding in result.findings:
            is_error = finding.severity == "error"
            self._tracker.record(
                rule_name="object_preflight",
                target=finding.kind,
                section=finding.location,
                change_type=f"preflight_{finding.severity}",
                before="",
                after=finding.message,
                paragraph_index=-1,
                success=not is_error,
                failure_reason=finding.message if is_error else None,
            )

        preflight_skips = object_preflight_module_skips(result.findings, policy)
        ctx.object_preflight_module_skips = list(preflight_skips.values())
        failure_policy = str(getattr(compliance, "failure_policy", "") or "").strip()
        blocked = not result.inspection_succeeded or (
            bool(result.blocking_findings)
            and (self._is_strict_mode() or failure_policy == "block")
        )
        return blocked, preflight_skips

    def _build_object_preflight_policy_snapshot(self, policy) -> dict[str, object]:
        input_profile = getattr(self._config, "input_source_profile", None)
        schema_ids = resolve_material_schema_ids(
            str(getattr(input_profile, "material_schema_id", "") or "").strip(),
            list(getattr(input_profile, "material_schema_ids", []) or []),
        )
        schema_id = schema_ids[0] if schema_ids else ""
        family_id = ""
        ooxml_touchpoints: tuple[str, ...] = ()
        recommended_targets: tuple[str, ...] = ()
        for candidate_schema_id in schema_ids:
            try:
                family_id = get_material_schema(candidate_schema_id).family
                family = get_planned_scene_family(family_id)
                ooxml_touchpoints = family.ooxml_touchpoints
                recommended_targets = object_preflight_targets_for_touchpoints(ooxml_touchpoints)
                break
            except KeyError:
                family_id = ""
                ooxml_touchpoints = ()
                recommended_targets = ()
        return {
            "enabled": bool(getattr(policy, "enabled", True)),
            "preservation_mode": str(getattr(policy, "preservation_mode", "") or ""),
            "scan_targets": _string_list(getattr(policy, "scan_targets", []) or []),
            "block_on": _string_list(getattr(policy, "block_on", []) or []),
            "skip_high_risk_modules": bool(
                getattr(policy, "skip_high_risk_modules", True)
            ),
            "skip_modules_by_finding": {
                str(kind): _string_list(modules)
                for kind, modules in dict(
                    getattr(policy, "skip_modules_by_finding", {}) or {}
                ).items()
            },
            "material_schema_id": schema_id,
            "material_schema_ids": list(schema_ids),
            "planning_family_id": family_id,
            "planning_ooxml_touchpoints": list(ooxml_touchpoints),
            "recommended_scan_targets": list(recommended_targets),
        }

    def _run_content_visibility_preflight(
        self,
        doc: Document,
        ctx: PipelineContext,
    ) -> str | None:
        scan = scan_content_visibility_markers(
            doc,
            getattr(self._config, "delivery_presets", []) or [],
        )
        ctx.content_visibility_scan = scan
        if scan.has_blocking_issues:
            ctx.content_visibility_preview = []
            diagnostics = list(scan.blocking_diagnostics or [])
            diagnostic_summary = "; ".join(
                f"[{item.get('code', 'invalid')}] {item.get('message', '')}"
                for item in diagnostics
            )
            error = (
                "Content visibility preflight blocked execution before module "
                f"mutation: {diagnostic_summary}"
            )
            self._tracker.record(
                rule_name="delivery_visibility_preflight",
                target="delivery_presets",
                section="delivery",
                change_type="visibility_preflight_blocked",
                before="document body unchanged",
                after=diagnostic_summary,
                paragraph_index=-1,
                success=False,
                failure_reason=error,
            )
            return error

        ctx.content_visibility_preview = preview_content_visibility_effects(
            doc,
            getattr(self._config, "delivery_presets", []) or [],
        )
        if not scan.selectors and not scan.used_rule_selectors:
            return None
        before = (
            "document selectors: "
            + (", ".join(scan.selectors) if scan.selectors else "-")
        )
        after = (
            "preset selectors: "
            + (", ".join(scan.used_rule_selectors) if scan.used_rule_selectors else "-")
        )
        issue_messages = scan.issue_messages()
        if issue_messages:
            after += "; " + "; ".join(issue_messages)
        else:
            after += "; selectors matched"
        preview_text = _content_visibility_preview_summary(
            list(ctx.content_visibility_preview or [])
        )
        if preview_text:
            after += "; preview: " + preview_text
        self._tracker.record(
            rule_name="delivery_visibility_preflight",
            target="delivery_presets",
            section="delivery",
            change_type=(
                "visibility_preflight_warning"
                if issue_messages
                else "visibility_preflight"
            ),
            before=before,
            after=after,
            paragraph_index=-1,
            success=True,
        )
        return None

    def _run_exam_question_schema_validation(self, ctx: PipelineContext) -> None:
        result = inspect_exam_question_schema(self._config)
        if str(getattr(result, "status", "") or "") == "not_applicable":
            return
        ctx.exam_question_schema = result
        summary = getattr(result, "summary", None)
        after = (
            f"status={getattr(result, 'status', '') or '-'}; "
            f"sections={int(getattr(summary, 'section_count', 0) or 0)}; "
            f"questions={int(getattr(summary, 'question_count', 0) or 0)}; "
            f"errors={int(getattr(result, 'error_count', 0) or 0)}; "
            f"warnings={int(getattr(result, 'warning_count', 0) or 0)}"
        )
        self._tracker.record(
            rule_name="exam_question_schema",
            target=getattr(result, "schema_id", "") or "exam_items_v1",
            section="input",
            change_type=(
                "schema_validation_warning"
                if getattr(result, "has_issues", False)
                else "schema_validation"
            ),
            before=getattr(result, "source_key", "") or "structured source",
            after=after,
            paragraph_index=-1,
            success=True,
        )

    def _run_exam_delivery_runtime(self, ctx: PipelineContext) -> None:
        result = ctx.exam_question_schema
        if result is None:
            result = inspect_exam_question_schema(self._config)
            if str(getattr(result, "status", "") or "") == "not_applicable":
                return
            ctx.exam_question_schema = result
        runtime_output_dir = self._output_dir or Path(ctx.source_doc_dir or ".")
        runtime = build_exam_delivery_runtime(
            self._config,
            output_dir=runtime_output_dir,
            source_stem=Path(ctx.source_doc_path or "exam").stem,
            validation=result,
            master=self._exam_master,
        )
        if str(getattr(runtime, "status", "") or "") == "not_applicable":
            return
        ctx.exam_delivery_runtime = runtime
        self._tracker.record(
            rule_name="exam_delivery_runtime",
            target=getattr(runtime, "schema_id", "") or "exam_items_v1",
            section="input_to_delivery",
            change_type=(
                "exam_delivery_runtime_blocked"
                if getattr(runtime, "status", "") == "blocked"
                else "exam_delivery_runtime"
            ),
            before=getattr(runtime, "source_key", "") or "structured source",
            after=(
                f"status={getattr(runtime, 'status', '') or '-'}; "
                f"preview={getattr(runtime, 'markdown_preview_path', '') or '-'}; "
                f"versions={int(getattr(runtime, 'version_count', 0) or 0)}"
            ),
            paragraph_index=-1,
            success=str(getattr(runtime, "status", "") or "") in {"ok", "warning"},
            failure_reason=(
                None
                if str(getattr(runtime, "status", "") or "") in {"ok", "warning"}
                else str(getattr(runtime, "skipped_reason", "") or "runtime_blocked")
            ),
        )

    def _run_official_document_assembly_runtime(self, ctx: PipelineContext) -> None:
        profile_id = _official_document_assembly_profile_id(
            self._config,
            self._official_document_type_id,
        )
        if not profile_id:
            return

        runtime_output_dir = self._output_dir or Path(ctx.source_doc_dir or ".")
        output_config = getattr(self._config, "output", None)
        result = assemble_official_document_docx(
            profile_id,
            getattr(self._config, "entity_data", {}) or {},
            output_dir=runtime_output_dir,
            field_aliases=getattr(self._config, "field_aliases", {}) or {},
            filename=f"{Path(ctx.source_doc_path or 'official_document').stem}_official.docx",
            master=self._official_master,
            removed_field_keys={
                str(key)
                for key, scope in dict(
                    getattr(self._config, "field_scopes", {}) or {}
                ).items()
                if str(scope) == "removed"
            },
            generate_review_pdf=bool(
                getattr(output_config, "review_pdf", False)
            ),
        )
        ctx.official_document_assembly = result
        missing_required = ", ".join(
            str(value) for value in getattr(result, "missing_required_fields", ()) or ()
        )
        unresolved = ", ".join(
            str(value) for value in getattr(result, "unresolved_placeholders", ()) or ()
        )
        after = (
            f"status={getattr(result, 'status', '') or '-'}; "
            f"master={getattr(result, 'master_id', '') or '-'}; "
            f"outputs={len(getattr(result, 'output_paths', {}) or {})}; "
            f"docx={getattr(result, 'docx_path', '') or '-'}; "
            f"review_pdf={getattr(result, 'review_pdf_status', '') or '-'}"
        )
        if missing_required:
            after += f"; missing={missing_required}"
        if unresolved:
            after += f"; unresolved={unresolved}"
        self._tracker.record(
            rule_name="official_document_assembly",
            target=profile_id,
            section="input_to_delivery",
            change_type=(
                "official_document_assembly"
                if getattr(result, "status", "") == "ok"
                else "official_document_assembly_blocked"
            ),
            before="entity_data official material fields",
            after=after,
            paragraph_index=-1,
            success=getattr(result, "status", "") == "ok",
        )

    def _run_journal_citation_validation(self, doc: Document, ctx: PipelineContext) -> None:
        result = inspect_journal_citations(self._config, doc)
        if str(getattr(result, "status", "") or "") == "not_applicable":
            return
        ctx.journal_citations = result
        summary = getattr(result, "summary", None)
        after = (
            f"status={getattr(result, 'status', '') or '-'}; "
            f"references={int(getattr(summary, 'reference_count', 0) or 0)}; "
            f"citations={int(getattr(summary, 'citation_count', 0) or 0)}; "
            f"matched={int(getattr(summary, 'matched_citation_count', 0) or 0)}; "
            f"missing={int(getattr(summary, 'missing_reference_count', 0) or 0)}; "
            f"errors={int(getattr(result, 'error_count', 0) or 0)}; "
            f"warnings={int(getattr(result, 'warning_count', 0) or 0)}"
        )
        self._tracker.record(
            rule_name="journal_citations",
            target=getattr(result, "schema_id", "") or "journal_submission_materials_v1",
            section="input",
            change_type=(
                "citation_source_warning"
                if getattr(result, "has_issues", False)
                else "citation_source_validation"
            ),
            before=getattr(result, "source_key", "") or "BibTeX/CSL source",
            after=after,
            paragraph_index=-1,
            success=True,
        )

    def _run_official_numbering_preservation_validation(
        self,
        doc: Document,
        original_doc: Document,
        ctx: PipelineContext,
    ) -> None:
        result = inspect_official_numbering_preservation(
            self._config,
            original_doc=original_doc,
            current_doc=doc,
            heading_map=ctx.heading_map,
            tracker_records=self._tracker.get_all(),
        )
        if str(getattr(result, "status", "") or "") == "not_applicable":
            return
        ctx.official_numbering_preservation = result
        summary = getattr(result, "summary", None)
        after = (
            f"status={getattr(result, 'status', '') or '-'}; "
            f"strategy={getattr(result, 'strategy', '') or '-'}; "
            f"headings={int(getattr(summary, 'heading_count', 0) or 0)}; "
            f"changed_headings={int(getattr(summary, 'changed_heading_count', 0) or 0)}; "
            f"numbering_records={int(getattr(summary, 'heading_numbering_record_count', 0) or 0)}"
        )
        if not getattr(result, "has_issues", False):
            return
        self._tracker.record(
            rule_name="official_numbering_preservation",
            target="numbering.xml",
            section="official_policy",
            change_type="preflight_warning",
            before="official numbering preserve evidence",
            after=after,
            paragraph_index=-1,
            success=True,
        )

    def _run_technical_chapter_inventory_validation(
        self,
        doc: Document,
        ctx: PipelineContext,
    ) -> None:
        result = inspect_technical_chapter_inventory(
            self._config,
            doc,
            doc_tree=ctx.doc_tree,
            heading_map=ctx.heading_map,
        )
        if str(getattr(result, "status", "") or "") == "not_applicable":
            return
        ctx.technical_chapter_inventory = result
        summary = getattr(result, "summary", None)
        after = (
            f"status={getattr(result, 'status', '') or '-'}; "
            f"chapters={int(getattr(summary, 'chapter_count', 0) or 0)}; "
            f"headings={int(getattr(summary, 'heading_count', 0) or 0)}; "
            f"appendix={int(getattr(summary, 'appendix_count', 0) or 0)}; "
            f"tables={int(getattr(summary, 'table_count', 0) or 0)}; "
            f"figures={int(getattr(summary, 'figure_count', 0) or 0)}"
        )
        if not getattr(result, "has_issues", False):
            return
        self._tracker.record(
            rule_name="technical_chapter_inventory",
            target="chapter_inventory",
            section="technical_long_docs",
            change_type="preflight_warning",
            before="technical chapter inventory",
            after=after,
            paragraph_index=-1,
            success=True,
        )

    def _run_application_section_word_limits_validation(
        self,
        doc: Document,
        ctx: PipelineContext,
    ) -> None:
        result = inspect_application_section_word_limits(
            self._config,
            doc,
            doc_tree=ctx.doc_tree,
            heading_map=ctx.heading_map,
        )
        if str(getattr(result, "status", "") or "") == "not_applicable":
            return
        ctx.application_section_word_limits = result
        summary = getattr(result, "summary", None)
        after = (
            f"status={getattr(result, 'status', '') or '-'}; "
            f"matched_sections={int(getattr(summary, 'matched_section_count', 0) or 0)}; "
            f"exceeded={int(getattr(summary, 'exceeded_section_count', 0) or 0)}; "
            f"required_missing={int(getattr(summary, 'required_missing_count', 0) or 0)}"
        )
        if not getattr(result, "has_issues", False):
            return
        self._tracker.record(
            rule_name="application_section_word_limits",
            target="section_limits",
            section="application_reports",
            change_type="preflight_warning",
            before="application section word limit evidence",
            after=after,
            paragraph_index=-1,
            success=True,
        )

    def _run_output_target_preflight(
        self,
        src_path: Path,
        ctx: PipelineContext,
    ) -> str | None:
        result = _build_output_target_preflight(
            src_path=src_path,
            config=self._config,
            output_dir=self._output_dir,
            output_suffix=self._output_suffix,
            force_delivery_presets=self._force_delivery_presets,
            output_path_overrides=self._output_path_overrides,
        )
        ctx.output_target_preflight = result
        if not result.has_issues:
            return None
        before = "; ".join(
            f"{_output_target_display_label(item)}: {item.path}"
            for item in result.items
        )
        blocking_issues = [
            issue
            for item in result.items
            for issue in item.issues
            if _output_target_issue_is_blocking(issue)
        ]
        advisory_issues = [
            issue
            for item in result.items
            for issue in item.issues
            if not _output_target_issue_is_blocking(issue)
        ]
        if advisory_issues:
            self._tracker.record(
                rule_name="output_target_preflight",
                target="output",
                section="delivery",
                change_type="output_target_warning",
                before=before,
                after="; ".join(
                    issue.message for issue in advisory_issues if issue.message
                ),
                paragraph_index=-1,
                success=True,
            )
        if not blocking_issues:
            return None

        issue_text = "; ".join(
            issue.message for issue in blocking_issues if issue.message
        )
        error = f"Output target preflight blocked execution: {issue_text}"
        self._tracker.record(
            rule_name="output_target_preflight",
            target="output",
            section="delivery",
            change_type="output_target_blocked",
            before=before,
            after=issue_text,
            paragraph_index=-1,
            success=False,
            failure_reason=error,
        )
        return error

    def _run_material_field_consistency(self, doc: Document, ctx: PipelineContext) -> None:
        input_profile = getattr(self._config, "input_source_profile", None)
        schema_id = str(getattr(input_profile, "material_schema_id", "") or "").strip()
        entity_data = getattr(self._config, "entity_data", {}) or {}
        result = inspect_material_field_consistency(
            doc,
            schema_id=schema_id,
            entity_data=entity_data,
        )
        if str(getattr(result, "status", "") or "") == "not_applicable":
            return
        ctx.material_field_consistency = result

    def _build_result(
        self,
        *,
        doc: Document,
        original_doc: Document,
        output_paths: dict[str, str],
        ctx: PipelineContext | None = None,
    ) -> PipelineResult:
        if self._is_cancelled():
            return self._cancelled_result(
                doc=doc,
                original_doc=original_doc,
                ctx=ctx,
                output_paths=output_paths,
            )
        failed_items = self._collect_failures()

        if not failed_items:
            return PipelineResult(
                success=True,
                status="success",
                doc=doc,
                original_doc=original_doc,
                tracker=self._tracker,
                context=ctx,
                output_paths=output_paths,
                config=self._config,
            )

        return PipelineResult(
            success=True,
            status="partial_success",
            doc=doc,
            original_doc=original_doc,
            tracker=self._tracker,
            context=ctx,
            output_paths=output_paths,
            failed_items=failed_items,
            error=f"{len(failed_items)} module operation(s) failed.",
            config=self._config,
        )

    def _save_outputs(
        self,
        doc: Document,
        src_path: Path,
        *,
        ctx: PipelineContext | None = None,
    ) -> dict[str, str]:
        self._raise_if_cancelled()
        stem = src_path.stem
        suffix = src_path.suffix or ".docx"
        out_dir = self._output_dir or src_path.parent
        out_dir.mkdir(parents=True, exist_ok=True)

        presets = list(getattr(self._config, "delivery_presets", []) or [])
        if len(presets) > 1 or (self._force_delivery_presets and presets):
            return self._save_delivery_outputs(
                doc,
                src_path,
                presets,
                suffix,
                ctx=ctx,
            )

        output_cfg = getattr(self._config, "output", None)
        final_enabled = bool(getattr(output_cfg, "final_docx", True))
        if not final_enabled:
            return {}

        final_path = self._output_path_overrides.get(
            "final",
            out_dir / f"{stem}{self._output_suffix}{suffix}",
        )
        return self._publish_output_documents(
            [("final", final_path, doc)],
        )

    def _save_delivery_outputs(
        self,
        doc: Document,
        src_path: Path,
        presets: list,
        suffix: str,
        *,
        ctx: PipelineContext | None = None,
    ) -> dict[str, str]:
        delivery_issue = delivery_preset_identity_issue(self._config)
        if delivery_issue:
            raise ValueError(f"pipeline_configuration_invalid:{delivery_issue}")

        outputs: list[tuple[str, Path, Document]] = []
        for preset in presets:
            self._raise_if_cancelled()
            preset_id = str(getattr(preset, "preset_id", "") or "").strip()
            if not preset_id:
                continue
            artifacts = getattr(preset, "artifacts", None)
            if not bool(getattr(artifacts, "final_docx", False)):
                continue

            override = self._output_path_overrides.get(preset_id)
            if override is None:
                out_dir = self._delivery_output_dir(preset, src_path)
                filename = self._delivery_filename(preset, src_path)
                final_path = out_dir / _ensure_suffix(filename, suffix)
            else:
                final_path = override
            preset_doc = _clone_document(doc)
            visibility_receipt = _apply_content_visibility_rules(preset_doc, preset)
            if ctx is not None:
                if ctx.content_visibility_receipts is None:
                    ctx.content_visibility_receipts = {}
                ctx.content_visibility_receipts[preset_id] = copy.deepcopy(
                    visibility_receipt
                )
            self._record_content_visibility_summary(preset, visibility_receipt)
            outputs.append((preset_id, final_path, preset_doc))
        if not outputs:
            return {}
        return self._publish_output_documents(outputs)

    def _publish_output_documents(
        self,
        outputs: list[tuple[str, Path, Document]],
    ) -> dict[str, str]:
        """Write caller-owned stages or atomically publish a complete final set."""

        _validate_output_document_identities(outputs)

        if self._output_paths_are_owned_stages:
            output_paths: dict[str, str] = {}
            for output_id, owned_stage_path, output_doc in outputs:
                self._raise_if_cancelled()
                owned_stage_path.parent.mkdir(parents=True, exist_ok=True)
                self._save_staged_document(output_doc, owned_stage_path)
                self._raise_if_cancelled()
                output_paths[output_id] = str(owned_stage_path)
            return output_paths

        self._raise_if_cancelled()
        transaction = OwnedAssemblyTransaction(
            execution_id=f"pipeline-{uuid4().hex}",
            final_paths=tuple(final_path for _, final_path, _ in outputs),
            work_root=None,
        )
        try:
            candidates = {}
            for output_id, final_path, output_doc in outputs:
                self._raise_if_cancelled()
                stage_path = transaction.allocate_stage(output_id, final_path)
                self._save_staged_document(output_doc, stage_path)
                self._raise_if_cancelled()
                candidates[final_path] = stable_file_evidence(stage_path)
            self._raise_if_cancelled()
            transaction.publish(candidates)
            # Backups still exist here, so cancellation observed during the
            # atomic-set publish can restore every previous final.
            self._raise_if_cancelled()
            transaction.release_backups()
        except BaseException:
            try:
                transaction.restore_finals()
            finally:
                transaction.cleanup_owned()
            raise
        else:
            transaction.cleanup_owned()

        return {
            output_id: str(final_path)
            for output_id, final_path, _ in outputs
        }

    def _raise_if_cancelled(self) -> None:
        if self._is_cancelled():
            raise _PipelineCancellationRequested("Cancelled by user.")

    def _save_staged_document(self, doc: Document, stage_path: Path) -> None:
        doc.save(str(stage_path))
        self._best_effort_refresh_fields(doc, stage_path)

    def _record_content_visibility_summary(self, preset, receipt: dict[str, object]) -> None:
        removed_count = int(receipt.get("removed_paragraph_count") or 0)
        removed_body_count = int(receipt.get("removed_body_element_count") or 0)
        removed_table_count = int(receipt.get("removed_table_count") or 0)
        stripped_count = int(receipt.get("stripped_marker_paragraph_count") or 0)
        if not removed_body_count and not removed_count and not stripped_count:
            return
        preset_id = str(getattr(preset, "preset_id", "") or "final").strip() or "final"
        removed_keys = ", ".join(
            str(value) for value in receipt.get("remove_selectors", []) or []
        )
        after = (
            f"Removed {removed_body_count} body element(s) "
            f"({removed_count} paragraph(s), {removed_table_count} table(s))"
        )
        if removed_keys:
            after += f" for {removed_keys}"
        if stripped_count:
            after += f"; stripped {stripped_count} marker paragraph(s)"
        self._tracker.record(
            rule_name="delivery_visibility",
            target=f"preset:{preset_id}",
            section="delivery",
            change_type="content_visibility",
            before="",
            after=after,
            paragraph_index=-1,
            success=True,
        )

    def _delivery_output_dir(self, preset, src_path: Path) -> Path:
        return _delivery_output_dir_for_preset(
            preset,
            src_path=src_path,
            base_output_dir=self._output_dir or src_path.parent,
        )

    def _delivery_filename(self, preset, src_path: Path) -> str:
        return _delivery_filename_for_preset(
            preset,
            src_path=src_path,
            base_output_dir=self._output_dir or src_path.parent,
        )

    def _best_effort_refresh_fields(self, doc: Document, final_path: Path) -> None:
        if self._defer_field_refresh:
            return
        if not document_has_toc(doc):
            return
        try:
            refresh_doc_fields_with_word(str(final_path), timeout_sec=30)
        except Exception:
            return


class _SafeFormatDict(dict):
    def __missing__(self, key: str) -> str:
        return "{" + key + "}"


def _render_delivery_template(
    template: str,
    *,
    src_path: Path,
    preset,
    base_output_dir: Path,
) -> str:
    values = _SafeFormatDict(
        {
            "document_dir": str(src_path.parent),
            "output_dir": str(base_output_dir),
            "stem": src_path.stem,
            "suffix": src_path.suffix or ".docx",
            "preset_id": str(getattr(preset, "preset_id", "") or ""),
            "preset_label": delivery_preset_display_name(preset),
        }
    )
    return str(template or "").format_map(values)


def _content_visibility_preview_summary(previews: list[object]) -> str:
    parts: list[str] = []
    for preview in previews:
        removed_paragraphs = int(
            getattr(preview, "removed_paragraph_count", 0) or 0
        )
        removed_body_elements = int(
            getattr(
                preview,
                "removed_body_element_count",
                removed_paragraphs,
            )
            or 0
        )
        removed_tables = int(getattr(preview, "removed_table_count", 0) or 0)
        stripped = int(getattr(preview, "stripped_marker_paragraph_count", 0) or 0)
        if not removed_body_elements and not removed_paragraphs and not stripped:
            continue
        preset_id = str(getattr(preview, "preset_id", "") or "preset").strip()
        text = (
            f"{preset_id}: remove body={removed_body_elements}, "
            f"paragraphs={removed_paragraphs}, tables={removed_tables}"
        )
        blocks = int(getattr(preview, "removed_block_count", 0) or 0)
        if removed_body_elements and blocks:
            text += f" in {blocks} block(s)"
        sample = _content_visibility_preview_sample(preview)
        if removed_body_elements and sample:
            text += f" sample={sample}"
        if stripped:
            text += f", strip {stripped}"
        parts.append(text)
    return "; ".join(parts)


def _content_visibility_preview_sample(preview: object, *, max_length: int = 40) -> str:
    for block in list(getattr(preview, "removed_blocks", []) or []):
        for sample in list(getattr(block, "text_samples", []) or []):
            cleaned = str(sample or "").strip()
            if cleaned:
                return _trim_content_visibility_preview(cleaned, max_length=max_length)
    for sample in list(getattr(preview, "removed_text_samples", []) or []):
        cleaned = str(sample or "").strip()
        if cleaned:
            return _trim_content_visibility_preview(cleaned, max_length=max_length)
    return ""


def _trim_content_visibility_preview(text: str, *, max_length: int) -> str:
    if len(text) <= max_length:
        return text
    return text[:max_length].rstrip() + "..."


def _build_output_target_preflight(
    *,
    src_path: Path,
    config,
    output_dir: Path | None,
    output_suffix: str,
    force_delivery_presets: bool,
    output_path_overrides: Mapping[str, str | Path] | None = None,
) -> OutputTargetPreflightResult:
    base_output_dir = output_dir or src_path.parent
    targets = _planned_output_targets(
        src_path=src_path,
        config=config,
        base_output_dir=base_output_dir,
        output_suffix=output_suffix,
        force_delivery_presets=force_delivery_presets,
        output_path_overrides=output_path_overrides,
    )
    path_counts: dict[str, int] = {}
    for item in targets:
        key = _output_target_key(item.path)
        identity_count = sum(
            1
            for candidate in targets
            if _paths_refer_to_same_file(item.path, candidate.path)
        )
        path_counts[key] = max(path_counts.get(key, 0), identity_count)
    for item in targets:
        _append_output_target_file_issues(
            item,
            path_counts,
            source_path=src_path,
        )
    return OutputTargetPreflightResult(items=targets)


def _planned_output_targets(
    *,
    src_path: Path,
    config,
    base_output_dir: Path,
    output_suffix: str,
    force_delivery_presets: bool,
    output_path_overrides: Mapping[str, str | Path] | None = None,
) -> list[OutputTargetPreflightItem]:
    delivery_issue = delivery_preset_identity_issue(config)
    if delivery_issue:
        raise ValueError(f"pipeline_configuration_invalid:{delivery_issue}")

    suffix = src_path.suffix or ".docx"
    overrides = {
        str(key): Path(value)
        for key, value in dict(output_path_overrides or {}).items()
        if str(key or "").strip()
    }
    presets = list(getattr(config, "delivery_presets", []) or [])
    if len(presets) > 1 or (force_delivery_presets and presets):
        items: list[OutputTargetPreflightItem] = []
        for preset in presets:
            preset_id = str(getattr(preset, "preset_id", "") or "").strip()
            if not preset_id:
                continue
            artifacts = getattr(preset, "artifacts", None)
            if not bool(getattr(artifacts, "final_docx", False)):
                continue
            override = overrides.get(preset_id)
            if override is None:
                out_dir = _delivery_output_dir_for_preset(
                    preset,
                    src_path=src_path,
                    base_output_dir=base_output_dir,
                )
                filename = _delivery_filename_for_preset(
                    preset,
                    src_path=src_path,
                    base_output_dir=base_output_dir,
                )
                path = out_dir / _ensure_suffix(filename, suffix)
            else:
                path = override
            items.append(
                _output_target_item(
                    preset_id=preset_id,
                    label=delivery_preset_display_name(preset),
                    path=path,
                )
            )
        return items

    output_cfg = getattr(config, "output", None)
    if not bool(getattr(output_cfg, "final_docx", True)):
        return []
    return [
        _output_target_item(
            preset_id="final",
            label="Final DOCX",
            path=overrides.get(
                "final",
                base_output_dir / f"{src_path.stem}{output_suffix}{suffix}",
            ),
        )
    ]


def plan_pipeline_output_paths(
    input_path: str | Path,
    config,
    *,
    output_dir: str | Path | None = None,
    output_suffix: str = "_new",
    force_delivery_presets: bool = False,
) -> dict[str, str]:
    """Return the logical final DOCX targets without creating any artifact."""

    source = Path(input_path)
    base = Path(output_dir) if output_dir is not None else source.parent
    targets = _planned_output_targets(
        src_path=source,
        config=config,
        base_output_dir=base,
        output_suffix=output_suffix,
        force_delivery_presets=force_delivery_presets,
    )
    planned: dict[str, str] = {}
    for item in targets:
        preset_id = str(item.preset_id or "").strip()
        if not preset_id:
            raise ValueError("planned output target is missing preset_id")
        if preset_id in planned:
            raise ValueError(f"duplicate planned output preset_id: {preset_id}")
        planned[preset_id] = str(item.path)
    return planned


def pipeline_output_target_matches_source(
    source_path: str | Path,
    output_path: str | Path,
) -> bool:
    """Return whether an output resolves to the source filesystem identity."""

    return pipeline_output_paths_share_identity(source_path, output_path)


def pipeline_output_paths_share_identity(
    left: str | Path,
    right: str | Path,
) -> bool:
    """Return whether two logical output paths resolve to one filesystem file."""

    return _paths_refer_to_same_file(left, right)


def _delivery_output_dir_for_preset(
    preset,
    *,
    src_path: Path,
    base_output_dir: Path,
) -> Path:
    template = str(getattr(preset, "output_dir_template", "") or "").strip()
    if not template or template == "{document_dir}/output":
        return base_output_dir

    rendered = _render_delivery_template(
        template,
        src_path=src_path,
        preset=preset,
        base_output_dir=base_output_dir,
    )
    path = Path(rendered)
    return path if path.is_absolute() else base_output_dir / path


def _delivery_filename_for_preset(
    preset,
    *,
    src_path: Path,
    base_output_dir: Path,
) -> str:
    template = str(getattr(preset, "filename_template", "") or "").strip()
    if not template:
        template = "{stem}_{preset_id}"
    rendered = _render_delivery_template(
        template,
        src_path=src_path,
        preset=preset,
        base_output_dir=base_output_dir,
    )
    return rendered or f"{src_path.stem}_{getattr(preset, 'preset_id', 'final')}"


def _output_target_item(
    *,
    preset_id: str,
    label: str,
    path: Path,
) -> OutputTargetPreflightItem:
    parent = path.parent
    nearest_parent = _nearest_existing_parent(parent)
    parent_writable = True
    if nearest_parent is not None:
        parent_writable = os.access(str(nearest_parent), os.W_OK)
    return OutputTargetPreflightItem(
        preset_id=preset_id,
        label=label,
        path=str(path),
        exists=path.exists(),
        parent_path=str(parent),
        parent_exists=parent.exists(),
        parent_writable=parent_writable,
        path_length=len(str(path)),
    )


def _append_output_target_file_issues(
    item: OutputTargetPreflightItem,
    path_counts: dict[str, int],
    *,
    source_path: Path,
) -> None:
    path = Path(item.path)
    display_label = _output_target_display_label(item)
    if path_counts.get(_output_target_key(item.path), 0) > 1:
        item.issues.append(
            OutputTargetIssue(
                kind="duplicate_target",
                severity="error",
                message=f"{display_label} 与其它交付指向同一输出文件: {item.path}",
            )
        )
    if pipeline_output_target_matches_source(source_path, path):
        item.issues.append(
            OutputTargetIssue(
                kind="source_overwrite",
                severity="error",
                message=(
                    f"{display_label} output resolves to the input document and "
                    f"would overwrite the source: {item.path}"
                ),
            )
        )
    if item.exists:
        kind = "target_is_directory" if path.is_dir() else "target_exists"
        severity = "error" if path.is_dir() else "warning"
        message = (
            f"{display_label} 输出目标是目录: {item.path}"
            if path.is_dir()
            else f"{display_label} 输出文件已存在，将被覆盖: {item.path}"
        )
        item.issues.append(OutputTargetIssue(kind=kind, severity=severity, message=message))
    if not item.parent_writable:
        item.issues.append(
            OutputTargetIssue(
                kind="parent_not_writable",
                severity="error",
                message=f"{display_label} 输出目录不可写: {item.parent_path}",
            )
        )
    if item.path_length >= 260:
        item.issues.append(
            OutputTargetIssue(
                kind="path_too_long",
                severity="warning",
                message=f"{display_label} 输出路径长度 {item.path_length}，可能超过 Windows 限制",
            )
        )
    elif item.path_length >= 240:
        item.issues.append(
            OutputTargetIssue(
                kind="path_near_limit",
                severity="warning",
                message=f"{display_label} 输出路径长度 {item.path_length}，接近 Windows 限制",
            )
        )


def _output_target_display_label(item: OutputTargetPreflightItem) -> str:
    raw_label = str(item.label or "").strip()
    if raw_label:
        return raw_label
    return delivery_preset_display_name(
        str(item.preset_id or item.artifact or "output").strip()
    )


def _nearest_existing_parent(path: Path) -> Path | None:
    current = path
    while True:
        if current.exists():
            return current
        if current.parent == current:
            return None
        current = current.parent


_OUTPUT_TARGET_ADVISORY_SEVERITIES = frozenset({"warning", "info"})


def _output_target_issue_is_blocking(issue: OutputTargetIssue) -> bool:
    """Only explicitly advisory severities may pass preflight."""

    return str(issue.severity or "").strip().casefold() not in (
        _OUTPUT_TARGET_ADVISORY_SEVERITIES
    )


def _paths_refer_to_same_file(left: str | Path, right: str | Path) -> bool:
    left_path = Path(left).expanduser()
    right_path = Path(right).expanduser()
    try:
        if left_path.exists() and right_path.exists():
            return os.path.samefile(left_path, right_path)
    except OSError:
        pass
    return _output_target_key(left_path) == _output_target_key(right_path)


def _output_target_key(path: str | Path) -> str:
    candidate = Path(path).expanduser()
    try:
        resolved = candidate.resolve(strict=False)
    except (OSError, RuntimeError):
        resolved = Path(os.path.abspath(str(candidate)))
    return os.path.normcase(str(resolved)).casefold()


def _string_list(values) -> list[str]:
    return [
        str(value).strip()
        for value in list(values or [])
        if str(value or "").strip()
    ]


def _ensure_suffix(filename: str, suffix: str) -> Path:
    path = Path(str(filename or "").strip() or "output")
    expected = str(suffix or ".docx")
    if not expected.startswith("."):
        expected = "." + expected
    if path.suffix.casefold() == expected.casefold():
        return path
    # A dot inside a logical stem (for example ``draft.bidding_original``)
    # is not a valid replacement for the required DOCX extension.
    return Path(str(path) + expected)


def _validate_output_document_identities(
    outputs: list[tuple[str, Path, Document]],
) -> None:
    """Reject ambiguous publication keys before staging any document."""

    seen: set[str] = set()
    for index, (output_id, _final_path, _document) in enumerate(outputs):
        normalized = str(output_id or "").strip()
        if not normalized:
            raise ValueError(f"output_id_empty:{index}")
        if normalized in seen:
            raise ValueError(f"duplicate_output_id:{normalized}")
        seen.add(normalized)


_OFFICIAL_ASSEMBLY_SCHEMA_IDS = {
    "official_document_v1",
    "administrative_meeting_fields_v1",
}

_EXAM_ASSEMBLY_SCHEMA_ID = "exam_items_v1"
_EXAM_ASSEMBLY_FAMILY_ID = "exam_teaching"


def pipeline_configuration_integrity_issue(config) -> str:
    """Project the shared config integrity gate into Pipeline error wording."""

    issue = execution_config_integrity_issue(config)
    return f"pipeline_configuration_invalid:{issue}" if issue else ""


def pipeline_terminal_assembly_owner(config) -> str:
    """Return the unique specialized publication owner for an execution config.

    ``conflict`` is deliberately explicit and fail-closed.  Specialized
    assemblers consume structured material and publish their own artifact set;
    they must never be followed by the generic module/output pipeline.
    """

    if pipeline_configuration_integrity_issue(config):
        return "invalid"

    input_profile = getattr(config, "input_source_profile", None)
    schema_ids = resolve_material_schema_ids(
        str(getattr(input_profile, "material_schema_id", "") or ""),
        list(getattr(input_profile, "material_schema_ids", []) or []),
    )
    exam_owner = _EXAM_ASSEMBLY_SCHEMA_ID in schema_ids
    if not exam_owner:
        for schema_id in schema_ids:
            try:
                family = str(getattr(get_material_schema(schema_id), "family", "") or "")
            except (KeyError, TypeError, ValueError):
                return "invalid"
            if family == _EXAM_ASSEMBLY_FAMILY_ID:
                exam_owner = True
                break

    official_owner = bool(_OFFICIAL_ASSEMBLY_SCHEMA_IDS.intersection(schema_ids))
    if exam_owner and official_owner:
        return "conflict"
    if exam_owner:
        return "exam"
    if official_owner:
        return "official"
    return ""


def _official_document_assembly_profile_id(
    config,
    official_document_type_id: str,
) -> str:
    input_profile = getattr(config, "input_source_profile", None)
    schema_ids = resolve_material_schema_ids(
        str(getattr(input_profile, "material_schema_id", "") or ""),
        list(getattr(input_profile, "material_schema_ids", []) or []),
    )
    if not _OFFICIAL_ASSEMBLY_SCHEMA_IDS.intersection(schema_ids):
        return ""

    profile_id = str(official_document_type_id or "").strip()
    if not profile_id:
        raise RuntimeError(
            "execution_resource_unproven:official_document_type:required"
        )
    if profile_id == "per_item":
        raise RuntimeError(
            "execution_resource_unproven:official_document_type:per_item"
        )
    if get_official_document_profile(profile_id) is None:
        raise RuntimeError(
            f"official_document_type_unknown:{profile_id}"
        )
    return profile_id
def _official_document_assembly_output_paths(ctx: PipelineContext | None) -> dict[str, str]:
    if ctx is None:
        return {}
    result = getattr(ctx, "official_document_assembly", None)
    if result is None or str(getattr(result, "status", "") or "") != "ok":
        return {}
    output_paths = getattr(result, "output_paths", None)
    if isinstance(output_paths, Mapping):
        return {
            str(key): str(value)
            for key, value in output_paths.items()
            if str(key).strip() and str(value).strip()
        }
    docx_path = str(getattr(result, "docx_path", "") or "").strip()
    return {"official_docx": docx_path} if docx_path else {}


def _exam_delivery_runtime_output_paths(ctx: PipelineContext | None) -> dict[str, str]:
    if ctx is None:
        return {}
    runtime = getattr(ctx, "exam_delivery_runtime", None)
    if runtime is None:
        return {}
    output_paths: dict[str, str] = {}
    for version in list(getattr(runtime, "rendered_versions", ()) or ()):
        preset_id = str(getattr(version, "preset_id", "") or "").strip()
        docx_path = str(getattr(version, "docx_path", "") or "").strip()
        if preset_id and docx_path:
            output_paths[preset_id] = docx_path
    return output_paths


def _clone_document(doc: Document) -> Document:
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return Document(buffer)


def _apply_content_visibility_rules(doc: Document, preset) -> dict[str, object]:
    scan = scan_content_visibility_markers(doc, [preset])
    if scan.has_blocking_issues:
        diagnostic_summary = "; ".join(
            f"[{item.get('code', 'invalid')}] {item.get('message', '')}"
            for item in scan.blocking_diagnostics
        )
        raise ValueError(
            "content_visibility_preflight_required_before_mutation:"
            f"{diagnostic_summary}"
        )
    remove_selectors = extract_content_visibility_rule_selectors([preset])
    receipt = apply_block_visibility(
        doc,
        remove_selectors=remove_selectors,
    )
    return receipt.to_dict()
