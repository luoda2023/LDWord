"""Load material-token mappings from simple JSON, CSV, and Excel files."""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping

@dataclass(slots=True)
class MaterialMappingPayload:
    entity_data: dict[str, str] = field(default_factory=dict)
    records: list[dict[str, str]] = field(default_factory=list)
    source_format: str = ""
    table_shape: str = ""
    record_count: int = 0

    def is_empty(self) -> bool:
        return not self.entity_data


def load_material_mapping(path: str | Path) -> MaterialMappingPayload:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".json":
        return _load_json_mapping(source)
    if suffix == ".csv":
        return _load_csv_mapping(source)
    if suffix in {".xlsx", ".xlsm"}:
        return _load_excel_mapping(source)
    raise ValueError(f"Unsupported material mapping file: {source}")


def _load_json_mapping(path: Path) -> MaterialMappingPayload:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, Mapping):
        return MaterialMappingPayload(source_format="json")

    entity_data = data.get("entity_data")
    if "replacements" in data:
        raise ValueError("资料映射不再支持替换规则；请导入 Token 字段和值")
    if entity_data is None:
        entity_data = data
    return MaterialMappingPayload(
        entity_data=_normalize_entity_data(entity_data),
        source_format="json",
        record_count=1,
    )


def _load_csv_mapping(path: Path) -> MaterialMappingPayload:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        _ensure_unique_mapping_headers(reader.fieldnames or [])
        rows = list(reader)

    return _payload_from_table_rows(rows, source_format="csv")


def _load_excel_mapping(path: Path) -> MaterialMappingPayload:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel material mappings require openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        raw_rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not raw_rows:
        return MaterialMappingPayload(source_format=path.suffix.lower().lstrip("."))

    headers = [str(value or "").strip() for value in raw_rows[0]]
    _ensure_unique_mapping_headers(headers)
    rows: list[dict[str, Any]] = []
    for raw_row in raw_rows[1:]:
        if not any(value is not None and str(value).strip() for value in raw_row):
            continue
        row: dict[str, Any] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            row[header] = raw_row[index] if index < len(raw_row) else None
        if row:
            rows.append(row)

    return _payload_from_table_rows(
        rows,
        source_format=path.suffix.lower().lstrip("."),
    )


def _payload_from_table_rows(
    rows: list[Mapping[str, Any]],
    *,
    source_format: str = "",
) -> MaterialMappingPayload:

    if not rows:
        return MaterialMappingPayload(source_format=source_format)

    field_names = {str(name or "").strip().lower() for name in (rows[0].keys() or [])}
    if {"old", "new"}.issubset(field_names):
        raise ValueError("资料表不再支持 old/new 替换规则；请使用 key/value Token 表")
    if {"key", "value"}.issubset(field_names):
        return MaterialMappingPayload(
            entity_data=_csv_entity_data(rows),
            source_format=source_format,
            table_shape="key_value",
            record_count=1,
        )

    return MaterialMappingPayload(
        entity_data={
            str(key): str(value)
            for key, value in rows[0].items()
            if str(key or "").strip() and value is not None
        },
        records=[
            {
                str(key): str(value)
                for key, value in row.items()
                if str(key or "").strip() and value is not None
            }
            for row in rows
        ],
        source_format=source_format,
        table_shape="records",
        record_count=len(rows),
    )


def _normalize_entity_data(value: Any) -> dict[str, str]:
    if not isinstance(value, Mapping):
        return {}
    return {
        str(key): str(raw_value)
        for key, raw_value in value.items()
        if str(key or "").strip()
    }


def _csv_entity_data(rows: list[Mapping[str, Any]]) -> dict[str, str]:
    entity_data: dict[str, str] = {}
    for row in rows:
        key = str(_get_case_insensitive(row, "key") or "").strip()
        if not key:
            continue
        entity_data[key] = str(_get_case_insensitive(row, "value") or "")
    return entity_data


def _get_case_insensitive(row: Mapping[str, Any], key: str) -> Any:
    for raw_key, value in row.items():
        if str(raw_key or "").strip().lower() == key:
            return value
    return None


def _ensure_unique_mapping_headers(headers: list[object] | tuple[object, ...]) -> None:
    seen: set[str] = set()
    duplicates: list[str] = []
    for raw_header in headers:
        header = str(raw_header or "").strip()
        if header.startswith("{{") and header.endswith("}}"):
            header = header[2:-2].strip()
        if not header:
            continue
        if header in seen and header not in duplicates:
            duplicates.append(header)
        seen.add(header)
    if duplicates:
        raise ValueError(
            "资料表存在重复字段，无法导入：" + "、".join(duplicates)
        )


__all__ = [
    "MaterialMappingPayload",
    "_ensure_unique_mapping_headers",
    "load_material_mapping",
]
