"""Mode-scoped master library wrappers backed by ``config_library``."""

from __future__ import annotations

import json
import hashlib
import re
from shutil import copy2
from dataclasses import dataclass
from collections.abc import Mapping
from pathlib import Path

from docx import Document

from src.app_paths import master_library_data_root
from src.config.scene import ExamPaperConfig, coerce_exam_paper_config
from src.shared.engine.exam_paper_style import (
    BUILTIN_EXAM_BLANK_STYLE_IDS,
    BUILTIN_EXAM_MASTER_DIR,
    USER_EXAM_MASTER_DIR,
    ExamBlankStyleSpec,
    resolve_exam_blank_style,
    sync_user_exam_blank_master_files,
)
from src.config.master_placeholder_index import (
    MASTER_PLACEHOLDER_SCANNER_VERSION,
    placeholder_identifier,
    scan_master_placeholder_index,
)


_PROJECT_ROOT = Path(__file__).resolve().parents[2]
MASTER_MANIFEST_ROOT = _PROJECT_ROOT / "config_library" / "masters"
OFFICIAL_USER_MASTER_DIR = master_library_data_root() / "official" / "user"
_OFFICIAL_BUILTIN_MASTER_IDS = (
    "official_gbt_standard",
    "official_gbt_upward",
    "official_gbt_letter",
    "official_gbt_minutes",
    "official_gbt_order",
)
_OFFICIAL_MASTER_ID_BY_LAYOUT = {
    "common": "official_gbt_standard",
    "upward": "official_gbt_upward",
    "letter": "official_gbt_letter",
    "minutes": "official_gbt_minutes",
    "order": "official_gbt_order",
}


@dataclass(frozen=True, slots=True)
class PlaceholderContract:
    """Placeholder fields a master must preserve for assembly."""

    required: tuple[str, ...] = ()
    optional: tuple[str, ...] = ()
    runtime_fields: tuple[str, ...] = ()
    generated_fields: tuple[str, ...] = ()
    runtime_inserted_placeholders: tuple[str, ...] = ()
    legacy_aliases: tuple[str, ...] = ()

    @property
    def placeholders(self) -> tuple[str, ...]:
        return (*self.required, *self.optional)


@dataclass(frozen=True, slots=True)
class MasterSpec:
    """One selectable master shell scoped to a work mode."""

    master_id: str
    mode_id: str
    label: str
    family: str
    source_type: str
    docx_path: Path
    summary: str = ""
    status: str = ""
    readonly: bool = True
    base_master_id: str = ""
    template_config_id: str = ""
    compatible_template_config_ids: tuple[str, ...] = ()
    master_version: str = ""
    placeholder_contract: PlaceholderContract = PlaceholderContract()
    manifest_path: Path | None = None
    preview_sample_path: Path | None = None
    supported_assembly_types: tuple[str, ...] = ()
    boundaries: tuple[str, ...] = ()
    execution_frozen: bool = False

    @property
    def qualified_id(self) -> str:
        return f"{self.mode_id}/{self.master_id}"

    @property
    def placeholders(self) -> tuple[str, ...]:
        return self.placeholder_contract.placeholders


class MasterManifestError(ValueError):
    """A canonical built-in master manifest is missing or invalid."""


EXAM_PLACEHOLDER_CONTRACT = PlaceholderContract(
    required=("af_title", "af_questions"),
    optional=(
        "af_version",
        "af_subject",
        "af_grade",
        "af_duration",
        "af_total_score",
        "af_answer_area",
    ),
    runtime_fields=("title", "subject", "grade", "duration", "total_score"),
    generated_fields=("version",),
    legacy_aliases=(
        "legacy_exam_title",
        "legacy_exam_version",
        "legacy_exam_metadata",
        "legacy_exam_questions_marker",
        "legacy_exam_answer_area_marker",
    ),
)


def list_masters(
    mode_id: str,
    *,
    exam_config: ExamPaperConfig | dict | None = None,
    user_master_dir: Path | str | None = None,
) -> tuple[MasterSpec, ...]:
    """Return masters visible for one work mode.

    Exam and official modes are backed today. Other modes intentionally return
    an empty tuple until their mother/master assets have real contracts.
    """

    normalized_mode = str(mode_id or "").strip()
    if normalized_mode == "exam":
        return list_exam_masters(
            exam_config=exam_config, user_master_dir=user_master_dir
        )
    if normalized_mode == "official":
        return list_official_masters(user_master_dir=user_master_dir)
    return ()


def get_master(
    master_id: str,
    mode_id: str | None = None,
    *,
    exam_config: ExamPaperConfig | dict | None = None,
    user_master_dir: Path | str | None = None,
) -> MasterSpec | None:
    """Return a master by local id or by ``mode/master`` qualified id."""

    requested_mode = str(mode_id or "").strip()
    requested_id = str(master_id or "").strip()
    if "/" in requested_id:
        requested_mode, requested_id = requested_id.split("/", 1)
    if not requested_mode:
        return None
    for spec in list_masters(
        requested_mode,
        exam_config=exam_config,
        user_master_dir=user_master_dir,
    ):
        if spec.master_id == requested_id or spec.qualified_id == requested_id:
            return spec
    return None


def default_master(
    mode_id: str,
    *,
    exam_config: ExamPaperConfig | dict | None = None,
    user_master_dir: Path | str | None = None,
) -> MasterSpec | None:
    """Return the default master for a mode."""

    normalized_mode = str(mode_id or "").strip()
    if normalized_mode == "exam":
        return get_master(
            "default_exam",
            "exam",
            exam_config=exam_config,
            user_master_dir=user_master_dir,
        )
    if normalized_mode == "official":
        return get_master("official_gbt_standard", "official")
    return None


def resolve_official_master_for_contract(
    contract,
    *,
    requested: MasterSpec | None,
) -> MasterSpec | None:
    """Resolve one official contract without discarding explicit user choices."""

    if requested is not None:
        if requested.execution_frozen:
            return requested
        supported = tuple(requested.supported_assembly_types or ())
        if not supported or str(contract.profile_id or "").strip() in supported:
            return requested
        if requested.source_type != "builtin":
            return requested
    return (
        get_master(str(contract.master_id or "").strip(), "official")
        or requested
        or default_master("official")
    )


def list_exam_masters(
    *,
    exam_config: ExamPaperConfig | dict | None = None,
    user_master_dir: Path | str | None = None,
) -> tuple[MasterSpec, ...]:
    """Wrap existing exam blank-style masters as mode-scoped master specs."""

    config = coerce_exam_paper_config(exam_config)
    sync_user_exam_blank_master_files(
        config,
        Path(user_master_dir) if user_master_dir is not None else USER_EXAM_MASTER_DIR,
    )

    specs = [
        _exam_spec_from_blank_style(resolve_exam_blank_style(config, style_id))
        for style_id in BUILTIN_EXAM_BLANK_STYLE_IDS
    ]
    for style in config.custom_blank_styles:
        if not str(style.style_id or "").strip():
            continue
        specs.append(
            _exam_spec_from_blank_style(
                resolve_exam_blank_style(config, style.style_id),
                master_docx_path=style.master_docx_path,
                source_type=_exam_user_source_type(style.style_id),
            )
        )
    return tuple(specs)


def list_official_masters(
    *,
    user_master_dir: Path | str | None = None,
) -> tuple[MasterSpec, ...]:
    """Return built-in layout families plus valid user-dropped DOCX files."""

    builtins = _official_builtin_masters()
    if not builtins:
        return ()
    folder = (
        Path(user_master_dir)
        if user_master_dir is not None
        else OFFICIAL_USER_MASTER_DIR
    )
    users = _discover_user_official_masters(folder, fallbacks=builtins)
    return (*builtins, *users)


def create_official_master_copy(
    source: MasterSpec,
    *,
    output_dir: Path | str | None = None,
) -> MasterSpec:
    """Copy one official DOCX into the user library and return its discovered spec."""

    if not source.docx_path.is_file():
        raise FileNotFoundError(source.docx_path)
    folder = Path(output_dir) if output_dir is not None else OFFICIAL_USER_MASTER_DIR
    folder.mkdir(parents=True, exist_ok=True)
    target = _unique_official_docx_path(folder / f"{source.docx_path.stem}_副本.docx")
    copy2(source.docx_path, target)
    spec = _official_user_master_spec(target, fallback=source)
    if spec is None:
        target.unlink(missing_ok=True)
        raise ValueError("公文版式副本缺少必要占位符")
    return spec


def discover_user_official_masters(
    directory: Path | str | None = None,
) -> tuple[MasterSpec, ...]:
    """Discover valid user DOCX files without exposing or requiring JSON metadata."""

    builtins = _official_builtin_masters()
    if not builtins:
        return ()
    folder = Path(directory) if directory is not None else OFFICIAL_USER_MASTER_DIR
    return _discover_user_official_masters(folder, fallbacks=builtins)


def _official_builtin_masters() -> tuple[MasterSpec, ...]:
    return tuple(
        master
        for master_id in _OFFICIAL_BUILTIN_MASTER_IDS
        if (master := _official_builtin_master(master_id)) is not None
    )


def _official_builtin_master(
    master_id: str = "official_gbt_standard",
) -> MasterSpec | None:
    """Load one built-in GB/T 9704 layout-family manifest."""

    manifest_path, manifest = _master_manifest("official", master_id)
    contract = _placeholder_contract_from_manifest(
        manifest.get("placeholder_contract"),
        manifest_path=manifest_path,
    )
    resolved_master_id = _manifest_text(manifest, "master_id", manifest_path)
    return MasterSpec(
        master_id=resolved_master_id,
        mode_id="official",
        label=_manifest_text(manifest, "label", manifest_path),
        family=_manifest_text(manifest, "family", manifest_path),
        source_type=_manifest_text(manifest, "source_type", manifest_path),
        docx_path=_resolve_stored_path(
            _manifest_text(manifest, "docx_path", manifest_path)
        ),
        summary=_manifest_text(manifest, "summary", manifest_path),
        status=_manifest_text(manifest, "status", manifest_path),
        readonly=True,
        base_master_id="",
        template_config_id=_manifest_text(
            manifest, "template_config_id", manifest_path
        ),
        compatible_template_config_ids=_manifest_sequence(
            manifest,
            "compatible_template_config_ids",
            manifest_path,
        ),
        master_version=_manifest_text(manifest, "master_version", manifest_path),
        placeholder_contract=contract,
        manifest_path=manifest_path,
        supported_assembly_types=_manifest_sequence(
            manifest,
            "supported_assembly_types",
            manifest_path,
        ),
        boundaries=_manifest_sequence(manifest, "boundaries", manifest_path),
    )


def _discover_user_official_masters(
    folder: Path,
    *,
    fallbacks: tuple[MasterSpec, ...],
) -> tuple[MasterSpec, ...]:
    if not folder.is_dir():
        return ()
    specs: list[MasterSpec] = []
    for path in sorted(folder.glob("*.docx"), key=lambda item: item.name.casefold()):
        if path.name.startswith("~$"):
            continue
        fallback = _official_user_master_fallback(path, fallbacks=fallbacks)
        spec = _official_user_master_spec(path, fallback=fallback)
        if spec is not None:
            specs.append(spec)
    return tuple(specs)


def _official_user_master_fallback(
    path: Path,
    *,
    fallbacks: tuple[MasterSpec, ...],
) -> MasterSpec:
    by_id = {master.master_id: master for master in fallbacks}
    try:
        keywords = str(Document(str(path)).core_properties.keywords or "")
    except (OSError, ValueError):
        keywords = ""
    marker = "ldword:official-layout="
    layout_id = ""
    for token in keywords.split(","):
        token = token.strip()
        if token.startswith(marker):
            layout_id = token[len(marker) :].strip().lower()
            break
    target_id = _OFFICIAL_MASTER_ID_BY_LAYOUT.get(layout_id, "official_gbt_standard")
    return by_id.get(target_id) or fallbacks[0]


def _official_user_master_spec(
    path: Path,
    *,
    fallback: MasterSpec,
) -> MasterSpec | None:
    try:
        index = scan_master_placeholder_index(path)
    except (OSError, ValueError):
        return None
    available = {
        placeholder_identifier(key) for key in index.replaceable_placeholder_ids
    }
    if any(
        placeholder not in available
        for placeholder in fallback.placeholder_contract.required
    ):
        return None
    return MasterSpec(
        master_id=_official_user_master_id(path.stem),
        mode_id="official",
        label=path.stem,
        family="official",
        source_type="user",
        docx_path=path,
        summary="用户 DOCX 版式；视觉样式由 Word 文件本身决定。",
        status="用户版式，可用 Word 直接修改",
        readonly=False,
        base_master_id=fallback.master_id,
        template_config_id=fallback.template_config_id,
        compatible_template_config_ids=fallback.compatible_template_config_ids,
        master_version=f"sha256:{index.sha256}",
        placeholder_contract=fallback.placeholder_contract,
        supported_assembly_types=fallback.supported_assembly_types,
        boundaries=fallback.boundaries,
    )


def _official_user_master_id(stem: str) -> str:
    normalized = str(stem or "").strip().casefold()
    readable = re.sub(r"[^a-z0-9]+", "_", normalized).strip("_")[:32]
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:10]
    return f"user_official_{readable or 'master'}_{digest}"


def _unique_official_docx_path(path: Path) -> Path:
    if not path.exists():
        return path
    index = 2
    while True:
        candidate = path.with_name(f"{path.stem}_{index}{path.suffix}")
        if not candidate.exists():
            return candidate
        index += 1


def _exam_spec_from_blank_style(
    spec: ExamBlankStyleSpec,
    *,
    master_docx_path: str = "",
    source_type: str | None = None,
) -> MasterSpec:
    source = source_type or ("builtin" if spec.readonly else "user")
    manifest_path: Path | None = None
    manifest: dict[str, object] | None = None
    if source == "builtin":
        manifest_path, manifest = _master_manifest("exam", spec.style_id)
        contract = _placeholder_contract_from_manifest(
            manifest.get("placeholder_contract"),
            manifest_path=manifest_path,
        )
        label = _manifest_text(manifest, "label", manifest_path)
        family = _manifest_text(manifest, "family", manifest_path)
        docx_path = _resolve_stored_path(
            _manifest_text(manifest, "docx_path", manifest_path)
        )
        summary = _manifest_text(manifest, "summary", manifest_path)
        status = _manifest_text(manifest, "status", manifest_path)
        template_config_id = _manifest_text(
            manifest,
            "template_config_id",
            manifest_path,
        )
        compatible_template_config_ids = _manifest_sequence(
            manifest,
            "compatible_template_config_ids",
            manifest_path,
        )
        master_version = _manifest_text(manifest, "master_version", manifest_path)
        boundaries = _manifest_sequence(manifest, "boundaries", manifest_path)
    else:
        contract = EXAM_PLACEHOLDER_CONTRACT
        label = spec.label
        family = "exam"
        docx_path = _exam_master_docx_path(spec, master_docx_path)
        summary = spec.summary
        status = spec.status
        template_config_id = "default"
        compatible_template_config_ids = ("default",)
        master_version = ""
        boundaries = (
            "does not judge exam content quality",
            "does not replace manual review of scoring rules",
        )
    return MasterSpec(
        master_id=spec.style_id,
        mode_id="exam",
        label=label,
        family=family,
        source_type=source,
        docx_path=docx_path,
        summary=summary,
        status=status,
        readonly=spec.readonly,
        base_master_id=spec.base_style_id,
        template_config_id=template_config_id,
        compatible_template_config_ids=compatible_template_config_ids,
        master_version=master_version,
        placeholder_contract=contract,
        manifest_path=manifest_path,
        boundaries=boundaries,
    )


def _exam_master_docx_path(spec: ExamBlankStyleSpec, master_docx_path: str) -> Path:
    if spec.readonly:
        return _exam_builtin_docx_path(spec.master_filename or "default_exam_v20.docx")
    if master_docx_path:
        return _resolve_stored_path(master_docx_path)
    if spec.master_filename:
        return _resolve_stored_path(spec.master_filename)
    return _exam_builtin_docx_path("default_exam_v20.docx")


def _exam_builtin_docx_path(filename: str) -> Path:
    return BUILTIN_EXAM_MASTER_DIR / filename


def _resolve_stored_path(raw_path: str) -> Path:
    path = Path(str(raw_path or "").strip())
    if path.is_absolute():
        return path
    return _PROJECT_ROOT / path


def _master_manifest(
    mode_id: str,
    master_id: str,
) -> tuple[Path, dict[str, object]]:
    path = MASTER_MANIFEST_ROOT / mode_id / "builtin" / f"{master_id}.master.json"
    if not path.is_file():
        raise MasterManifestError(f"master_manifest_missing:{path}")
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise MasterManifestError(f"master_manifest_unreadable:{path}: {exc}") from exc
    if not isinstance(raw, dict):
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{path}: root must be an object"
        )
    manifest_mode = _manifest_text(raw, "mode_id", path)
    manifest_id = _manifest_text(raw, "master_id", path)
    if manifest_mode != mode_id:
        raise MasterManifestError(
            f"master_manifest_identity_mismatch:{path}: "
            f"mode_id={manifest_mode!r}, expected {mode_id!r}"
        )
    if manifest_id != master_id:
        raise MasterManifestError(
            f"master_manifest_identity_mismatch:{path}: "
            f"master_id={manifest_id!r}, expected {master_id!r}"
        )
    if _manifest_text(raw, "source_type", path) != "builtin":
        raise MasterManifestError(
            f"master_manifest_identity_mismatch:{path}: source_type must be 'builtin'"
        )
    if _manifest_text(raw, "family", path) != mode_id:
        raise MasterManifestError(
            f"master_manifest_identity_mismatch:{path}: family must equal mode_id"
        )
    for key in (
        "label",
        "docx_path",
        "summary",
        "status",
        "template_config_id",
        "master_version",
    ):
        _manifest_text(raw, key, path)
    template_config_id = _manifest_text(raw, "template_config_id", path)
    compatible_template_config_ids = _manifest_sequence(
        raw, "compatible_template_config_ids", path
    )
    if template_config_id not in compatible_template_config_ids:
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{path}: "
            "compatible_template_config_ids must include template_config_id"
        )
    _manifest_sequence(raw, "boundaries", path)
    contract = _placeholder_contract_from_manifest(
        raw.get("placeholder_contract"),
        manifest_path=path,
    )
    placeholder_index = _validate_placeholder_index(
        raw.get("placeholder_index"), manifest_path=path
    )
    indexed_identifiers = {
        placeholder_identifier(item)
        for item in _manifest_sequence(placeholder_index, "placeholder_ids", path)
    }
    missing_contract_placeholders = tuple(
        placeholder
        for placeholder in contract.required
        if placeholder_identifier(placeholder) not in indexed_identifiers
    )
    if missing_contract_placeholders:
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{path}: placeholder_index omits required "
            f"contract ids {missing_contract_placeholders!r}"
        )
    if mode_id == "official":
        _manifest_sequence(raw, "supported_assembly_types", path)
    docx_path = _validate_manifest_docx_path(raw, manifest_path=path, mode_id=mode_id)
    try:
        actual_docx_sha256 = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    except OSError as exc:
        raise MasterManifestError(
            f"master_manifest_asset_invalid:{path}: master DOCX is unreadable: {exc}"
        ) from exc
    if placeholder_index["sha256"] != actual_docx_sha256:
        raise MasterManifestError(
            f"master_manifest_index_stale:{path}: placeholder_index does not match DOCX"
        )
    return path, dict(raw)


def _placeholder_contract_from_manifest(
    value: object,
    *,
    manifest_path: Path,
) -> PlaceholderContract:
    if not isinstance(value, Mapping):
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: "
            "placeholder_contract must be an object"
        )
    expected_keys = {
        "required",
        "optional",
        "runtime_fields",
        "generated_fields",
        "runtime_inserted_placeholders",
        "legacy_aliases",
    }
    if set(value) != expected_keys:
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: "
            "placeholder_contract must contain exactly the canonical fields"
        )
    return PlaceholderContract(
        required=_manifest_sequence(value, "required", manifest_path),
        optional=_manifest_sequence(value, "optional", manifest_path, allow_empty=True),
        runtime_fields=_manifest_sequence(
            value, "runtime_fields", manifest_path, allow_empty=True
        ),
        generated_fields=_manifest_sequence(
            value, "generated_fields", manifest_path, allow_empty=True
        ),
        runtime_inserted_placeholders=_manifest_sequence(
            value,
            "runtime_inserted_placeholders",
            manifest_path,
            allow_empty=True,
        ),
        legacy_aliases=_manifest_sequence(
            value, "legacy_aliases", manifest_path, allow_empty=True
        ),
    )


def _manifest_text(
    manifest: Mapping[str, object],
    key: str,
    manifest_path: Path,
) -> str:
    value = manifest.get(key)
    if not isinstance(value, str) or not value.strip():
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: {key} must be non-empty text"
        )
    return value.strip()


def _manifest_sequence(
    manifest: Mapping[str, object],
    key: str,
    manifest_path: Path,
    *,
    allow_empty: bool = False,
) -> tuple[str, ...]:
    value = manifest.get(key)
    if not isinstance(value, list):
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: {key} must be an array"
        )
    if any(not isinstance(item, str) or not item.strip() for item in value):
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: "
            f"{key} must contain only non-empty text"
        )
    normalized = tuple(item.strip() for item in value)
    if len(set(normalized)) != len(normalized):
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: {key} contains duplicates"
        )
    if not normalized and not allow_empty:
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: {key} must not be empty"
        )
    return normalized


def _validate_placeholder_index(
    value: object,
    *,
    manifest_path: Path,
) -> Mapping[str, object]:
    if not isinstance(value, Mapping):
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: "
            "placeholder_index must be an object"
        )
    sha256 = value.get("sha256")
    if not isinstance(sha256, str) or re.fullmatch(r"[0-9a-f]{64}", sha256) is None:
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: "
            "placeholder_index.sha256 must be a lowercase SHA-256 digest"
        )
    scanner_version = value.get("scanner_version")
    if not isinstance(scanner_version, str) or not scanner_version.strip():
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: "
            "placeholder_index.scanner_version is required"
        )
    if scanner_version != MASTER_PLACEHOLDER_SCANNER_VERSION:
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: unsupported "
            f"placeholder_index.scanner_version {scanner_version!r}"
        )
    placeholder_ids = set(_manifest_sequence(value, "placeholder_ids", manifest_path))
    count_keys: dict[str, set[str]] = {}
    for key in ("raw_occurrence_counts", "replaceable_occurrence_counts"):
        counts = value.get(key)
        if not isinstance(counts, Mapping) or any(
            not isinstance(item_key, str)
            or not item_key.strip()
            or not isinstance(count, int)
            or isinstance(count, bool)
            or count <= 0
            for item_key, count in counts.items()
        ):
            raise MasterManifestError(
                f"master_manifest_schema_invalid:{manifest_path}: "
                f"placeholder_index.{key} must map ids to positive integers"
            )
        count_keys[key] = set(counts)
        if not count_keys[key] <= placeholder_ids:
            raise MasterManifestError(
                f"master_manifest_schema_invalid:{manifest_path}: "
                f"placeholder_index.{key} contains undeclared ids"
            )
    if (
        not count_keys["replaceable_occurrence_counts"]
        <= count_keys["raw_occurrence_counts"]
    ):
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: replaceable placeholder "
            "ids must also exist in raw_occurrence_counts"
        )
    if count_keys["raw_occurrence_counts"] != placeholder_ids:
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: placeholder_ids must "
            "match raw_occurrence_counts"
        )
    if not isinstance(value.get("reference_page_excluded"), bool):
        raise MasterManifestError(
            f"master_manifest_schema_invalid:{manifest_path}: "
            "placeholder_index.reference_page_excluded must be boolean"
        )
    return value


def _validate_manifest_docx_path(
    manifest: Mapping[str, object],
    *,
    manifest_path: Path,
    mode_id: str,
) -> Path:
    raw_path = _manifest_text(manifest, "docx_path", manifest_path)
    stored_path = Path(raw_path)
    if stored_path.is_absolute() or stored_path.suffix.casefold() != ".docx":
        raise MasterManifestError(
            f"master_manifest_asset_invalid:{manifest_path}: "
            "docx_path must be a project-relative .docx path"
        )
    resolved = (_PROJECT_ROOT / stored_path).resolve()
    allowed_root = (
        _PROJECT_ROOT / "config_library" / "masters" / mode_id / "builtin"
    ).resolve()
    try:
        resolved.relative_to(allowed_root)
    except ValueError as exc:
        raise MasterManifestError(
            f"master_manifest_asset_invalid:{manifest_path}: "
            "docx_path escapes the canonical built-in master directory"
        ) from exc
    if not resolved.is_file():
        raise MasterManifestError(
            f"master_manifest_asset_invalid:{manifest_path}: "
            f"master DOCX is missing: {resolved}"
        )
    return resolved


def _exam_user_source_type(style_id: str) -> str:
    normalized = str(style_id or "").strip()
    if normalized.startswith("user_file_"):
        return "discovered"
    if normalized.startswith("user_imported_exam"):
        return "imported"
    return "user"


__all__ = [
    "EXAM_PLACEHOLDER_CONTRACT",
    "MasterManifestError",
    "OFFICIAL_USER_MASTER_DIR",
    "MasterSpec",
    "MASTER_MANIFEST_ROOT",
    "PlaceholderContract",
    "create_official_master_copy",
    "default_master",
    "discover_user_official_masters",
    "get_master",
    "list_exam_masters",
    "list_masters",
    "list_official_masters",
    "resolve_official_master_for_contract",
]
