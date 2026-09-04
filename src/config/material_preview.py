"""Preview helpers for material placeholder matching."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from src.shared.engine.exact_material_placeholders import (
    EXACT_PLACEHOLDER_PATTERN,
    ExactMaterialPlaceholder,
    scan_document_exact_placeholders,
)


PLACEHOLDER_PATTERN = EXACT_PLACEHOLDER_PATTERN
MaterialPlaceholder = ExactMaterialPlaceholder


def scan_docx_placeholders(path: str | Path) -> list[str]:
    return [item.key for item in scan_docx_placeholder_inventory(path)]


def scan_docx_placeholder_inventory(path: str | Path) -> list[MaterialPlaceholder]:
    target = Path(str(path or ""))
    if not target.exists() or target.suffix.lower() != ".docx":
        return []

    return scan_document_exact_placeholders(Document(str(target)))


__all__ = [
    "MaterialPlaceholder",
    "PLACEHOLDER_PATTERN",
    "scan_docx_placeholder_inventory",
    "scan_docx_placeholders",
]
