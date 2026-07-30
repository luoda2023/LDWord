"""Spreadsheet inspection before material-package records are activated."""

from __future__ import annotations

import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook

from src.config.entity import EntityProfile
from src.config.material_package_v6 import (
    MaterialGroup,
    MaterialPackageV6,
    MaterialRecord,
    MaterialValueScope,
)
from src.shared.engine.material_timeline import (
    default_timeline_plan,
    evenly_distributed_nodes,
    normalize_timeline_plans,
)


@dataclass(frozen=True, slots=True)
class ImportMappingProfile:
    """Explicit meaning assigned to a workbook table."""

    primary_sheet: str = ""
    group_fields: tuple[str, ...] = (
        "产品名称",
        "product_name",
        "product",
        "产品类型",
        "product_type",
        "路线",
        "route",
    )
    record_name_fields: tuple[str, ...] = (
        "项目名称",
        "project_name",
        "entity_name",
    )
    record_key_fields: tuple[str, ...] = (
        "项目编号",
        "project_code",
        "profile_id",
        "record_id",
    )
    activation_required_fields: tuple[str, ...] = ("项目名称", "项目编号")
    auxiliary_single_row_as_shared: bool = True


@dataclass(slots=True)
class ImportCandidate:
    record_id: str
    record_name: str
    group_id: str
    lifecycle_state: str
    values: MaterialValueScope
    source_locator: dict[str, Any]
    missing_activation_fields: tuple[str, ...] = ()


@dataclass(slots=True)
class PackageImportDraft:
    package_id: str
    package_name: str
    shared_scope: MaterialValueScope
    groups: list[MaterialGroup]
    candidates: list[ImportCandidate]
    source_path: str
    mapping: ImportMappingProfile
    issues: tuple[str, ...] = ()

    @property
    def candidate_count(self) -> int:
        return len(self.candidates)

    @property
    def active_count(self) -> int:
        return sum(
            item.lifecycle_state == "active" for item in self.candidates
        )

    @property
    def draft_count(self) -> int:
        return sum(item.lifecycle_state == "draft" for item in self.candidates)

    @property
    def disabled_count(self) -> int:
        return sum(
            item.lifecycle_state == "disabled" for item in self.candidates
        )

    def materialize(self) -> MaterialPackageV6:
        records: list[MaterialRecord] = []
        for candidate in self.candidates:
            timeline_plans = _legacy_timeline_from_fields(candidate.values.fields)
            values = copy.deepcopy(candidate.values)
            if timeline_plans and not values.timeline_plans:
                values.timeline_plans = timeline_plans
            profile = EntityProfile(
                profile_id=candidate.record_id,
                profile_name=candidate.record_name,
                fields=dict(values.fields),
                field_aliases=dict(values.field_aliases),
                field_functions=copy.deepcopy(values.field_functions),
                timeline_plans=copy.deepcopy(values.timeline_plans),
                declared_field_keys=list(values.fields),
                field_sources=dict(values.field_sources),
            )
            records.append(
                MaterialRecord(
                    record_id=candidate.record_id,
                    record_name=candidate.record_name,
                    group_id=candidate.group_id,
                    lifecycle_state=candidate.lifecycle_state,
                    values=values,
                    profile=profile,
                    source_locator=copy.deepcopy(candidate.source_locator),
                )
            )
        return MaterialPackageV6(
            package_id=self.package_id,
            package_name=self.package_name,
            shared_scope=copy.deepcopy(self.shared_scope),
            groups=copy.deepcopy(self.groups),
            records=records,
            source_path=self.source_path,
        )


def inspect_material_workbook(
    path: str | Path,
    *,
    mapping: ImportMappingProfile | None = None,
) -> PackageImportDraft:
    """Read every nonempty primary row as a candidate, never as an implicit run."""

    source = Path(path)
    profile = mapping or ImportMappingProfile()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        primary = (
            workbook[profile.primary_sheet]
            if profile.primary_sheet
            else workbook.active
        )
        headers, rows = _worksheet_rows(primary)
        shared_fields: dict[str, str] = {}
        shared_sources: dict[str, str] = {}
        issues: list[str] = []
        if profile.auxiliary_single_row_as_shared:
            for sheet in workbook.worksheets:
                if sheet is primary:
                    continue
                _, auxiliary_rows = _worksheet_rows(sheet)
                if not auxiliary_rows:
                    continue
                if len(auxiliary_rows) != 1:
                    issues.append(
                        f"辅助工作表“{sheet.title}”有 {len(auxiliary_rows)} 行，"
                        "未自动解释为共享层"
                    )
                    continue
                row_number, values = auxiliary_rows[0]
                for key, value in values.items():
                    if not value:
                        continue
                    if key in shared_fields and shared_fields[key] != value:
                        issues.append(
                            f"共享字段“{key}”在多个工作表中值不一致"
                        )
                        continue
                    shared_fields[key] = value
                    shared_sources[key] = (
                        f"excel:{sheet.title}!{_column_for_header(sheet, key)}"
                        f"{row_number}"
                    )
    finally:
        workbook.close()

    groups: list[MaterialGroup] = []
    group_ids_by_value: dict[str, str] = {}
    candidates: list[ImportCandidate] = []
    used_group_ids: set[str] = set()
    for candidate_index, (row_number, row) in enumerate(rows, start=1):
        group_value = _first_value(row, profile.group_fields)
        group_id = ""
        if group_value:
            normalized_group = group_value.casefold()
            group_id = group_ids_by_value.get(normalized_group, "")
            if not group_id:
                group_id = _unique_indexed_id(
                    "group",
                    len(groups) + 1,
                    used_group_ids,
                )
                used_group_ids.add(group_id)
                group_ids_by_value[normalized_group] = group_id
                groups.append(
                    MaterialGroup(
                        group_id=group_id,
                        group_name=group_value,
                        route_id=group_value,
                        values=MaterialValueScope(
                            fields={"产品名称": group_value},
                            field_sources={
                                "产品名称": f"excel:{primary.title}!row:{row_number}"
                            },
                        ),
                        source_locator={
                            "sheet": primary.title,
                            "first_row": row_number,
                        },
                    )
                )
        missing = tuple(
            key
            for key in profile.activation_required_fields
            if not _mapped_required_value(row, key)
        )
        lifecycle = "active" if not missing else "draft"
        record_id = f"record-{candidate_index:04d}"
        record_name = (
            _first_value(row, profile.record_name_fields)
            or _first_value(row, profile.record_key_fields)
            or f"第 {candidate_index} 份"
        )
        record_fields = {
            key: value
            for key, value in row.items()
            if value
            and key not in profile.group_fields
        }
        field_sources = {
            key: f"excel:{primary.title}!row:{row_number}"
            for key in record_fields
        }
        candidates.append(
            ImportCandidate(
                record_id=record_id,
                record_name=record_name,
                group_id=group_id,
                lifecycle_state=lifecycle,
                values=MaterialValueScope(
                    fields=record_fields,
                    field_sources=field_sources,
                ),
                source_locator={
                    "kind": "excel",
                    "path": str(source.resolve()),
                    "sheet": primary.title,
                    "row": row_number,
                },
                missing_activation_fields=missing,
            )
        )
    if not candidates:
        issues.append("主工作表中没有资料候选行")
    return PackageImportDraft(
        package_id=source.stem,
        package_name=source.stem,
        shared_scope=MaterialValueScope(
            fields=shared_fields,
            field_sources=shared_sources,
        ),
        groups=groups,
        candidates=candidates,
        source_path=str(source.resolve()),
        mapping=profile,
        issues=tuple(issues),
    )


def _worksheet_rows(sheet) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    iterator = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        return [], []
    headers = [str(value or "").strip() for value in raw_headers]
    nonempty_headers = [item for item in headers if item]
    duplicates = sorted(
        {
            item
            for item in nonempty_headers
            if nonempty_headers.count(item) > 1
        }
    )
    if duplicates:
        raise ValueError(
            "material_import_headers_duplicate:" + ",".join(duplicates)
        )
    rows: list[tuple[int, dict[str, str]]] = []
    for row_number, raw_row in enumerate(iterator, start=2):
        values = {
            header: _cell_text(
                raw_row[index] if index < len(raw_row) else None
            )
            for index, header in enumerate(headers)
            if header
        }
        if any(values.values()):
            rows.append((row_number, values))
    return headers, rows


def _column_for_header(sheet, header: str) -> str:
    for cell in sheet[1]:
        if str(cell.value or "").strip() == header:
            return cell.column_letter
    return ""


def _mapped_required_value(row: Mapping[str, str], key: str) -> str:
    aliases = {
        "项目名称": ("项目名称", "project_name", "entity_name"),
        "项目编号": ("项目编号", "project_code", "record_id", "profile_id"),
        "产品名称": ("产品名称", "product_name", "product"),
    }
    return _first_value(row, aliases.get(key, (key,)))


def _first_value(
    values: Mapping[str, str],
    keys: tuple[str, ...],
) -> str:
    for key in keys:
        value = str(values.get(key, "") or "").strip()
        if value:
            return value
    return ""


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


_LEGACY_TIMELINE_OUTPUTS = (
    "节点_设计策划",
    "节点_编制计划书",
    "节点_设计输入",
    "节点_性能评审",
    "节点_设计输出",
    "节点_系统评审",
    "节点_设计验证",
    "节点_试产可行性",
    "节点_试产总结",
    "节点_设计确认",
    "节点_设计正稿",
)


def _legacy_timeline_from_fields(
    fields: Mapping[str, str],
) -> dict[str, dict[str, object]]:
    start_key = _first_value_key(
        fields,
        ("项目开始日期", "project_start_date", "start_date"),
    )
    end_key = _first_value_key(
        fields,
        ("项目结束日期", "project_end_date", "end_date"),
    )
    if not start_key or not end_key:
        return {}
    plan = default_timeline_plan()
    plan["start_field"] = start_key
    plan["end_field"] = end_key
    plan["calendar"] = {
        **dict(plan.get("calendar", {})),
        "basis": "calendar_day",
        "weekend_adjust": "forward",
    }
    nodes = evenly_distributed_nodes(len(_LEGACY_TIMELINE_OUTPUTS))
    for node, output_key in zip(
        nodes,
        _LEGACY_TIMELINE_OUTPUTS,
        strict=True,
    ):
        node["outputs"] = [{"field": output_key, "format": "yyyy-MM-dd"}]
    plan["nodes"] = nodes
    return normalize_timeline_plans({"legacy_iso_timeline": plan})


def _first_value_key(
    fields: Mapping[str, str],
    keys: tuple[str, ...],
) -> str:
    for key in keys:
        if str(fields.get(key, "") or "").strip():
            return key
    return ""


def _unique_indexed_id(prefix: str, index: int, used: set[str]) -> str:
    candidate = f"{prefix}-{index:04d}"
    while candidate in used:
        index += 1
        candidate = f"{prefix}-{index:04d}"
    return candidate


__all__ = [
    "ImportCandidate",
    "ImportMappingProfile",
    "PackageImportDraft",
    "inspect_material_workbook",
]
