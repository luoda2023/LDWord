"""Header/footer formatting with section-aware page-number strategy."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING

from docx.oxml import OxmlElement

from src.modules.base import BaseModule, Issue, ModuleMeta
from src.shared.engine.field_builder import build_complex_field, iter_field_instructions
from src.shared.engine.font_resolver import resolve_font
from src.shared.engine.ooxml_ops import find_or_create, find_or_create_before, qn
from src.shared.engine.page_number_planner import (
    NON_NUMBERED_HEADING_SECTION_TYPES,
    build_page_number_execution_plan,
    collect_page_number_diagnostics,
    format_page_number_diagnostic_text,
)
from src.shared.engine.run_ops import set_run_east_asian_font
from src.shared.engine.section_layout_planner import collect_section_inventory

if TYPE_CHECKING:
    from docx import Document

    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


_PAGE_NUMBER_FIELD_INSTRUCTIONS = {
    "decimal": " PAGE ",
    "upperRoman": " PAGE \\* ROMAN ",
    "lowerRoman": " PAGE \\* roman ",
    "upperLetter": " PAGE \\* ALPHABETIC ",
    "lowerLetter": " PAGE \\* alphabetic ",
    "ordinal": " PAGE \\* Ordinal ",
    "cardinalText": " PAGE \\* CardText ",
    "ordinalText": " PAGE \\* OrdText ",
}

_PAGE_NUMBER_W_FORMATS = {
    "decimal": "decimal",
    "upperRoman": "upperRoman",
    "lowerRoman": "lowerRoman",
    "upperLetter": "upperLetter",
    "lowerLetter": "lowerLetter",
    "ordinal": "ordinal",
    "cardinalText": "cardinalText",
    "ordinalText": "ordinalText",
    "decimalZero": "decimalZero",
    "decimalFullWidth": "decimalFullWidth",
}

_FOOTER_ALIGNMENT_VALUES = {
    "left": "left",
    "center": "center",
    "right": "right",
}

_CONTENT_TEMPLATE_TOKEN_RE = re.compile(r"\{(page|pages|section_pages|text|styleref|styleref_number)\}")


class HeaderFooterModule(BaseModule):
    meta = ModuleMeta(
        name="header_footer",
        description="页眉页脚",
        category="basic",
        requires_config=("header_footer",),
        soft_after=("heading_recognition", "section_format"),
        soft_consumes=("doc_tree",),
    )

    def validate(
        self,
        doc: Document,
        config: ResolvedConfig,
        context: PipelineContext,
    ) -> list[Issue]:
        issues: list[Issue] = []
        protected_boundaries = _protected_section_boundaries(doc)
        if protected_boundaries:
            issues.append(
                Issue(
                    level="error",
                    module_name=self.meta.name,
                    message=(
                        "检测到内容控件/customXml 内的嵌套分节或带修订的分节属性；"
                        "当前页眉页脚映射不能安全改写这些分节"
                    ),
                    location="document.sections.protected",
                )
            )
        issues.extend(
            [
            Issue(
                level=item.level,
                module_name=self.meta.name,
                message=format_page_number_diagnostic_text(item),
                location=item.location,
            )
            for item in collect_page_number_diagnostics(doc, context, config.header_footer)
            ]
        )
        return issues

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        hf_cfg = config.header_footer
        count = 0
        protected_boundaries = _protected_section_boundaries(doc)
        if protected_boundaries:
            raise RuntimeError(
                "protected section properties cannot be safely rewritten by header_footer"
            )

        header_enabled = _header_enabled(hf_cfg)
        footer_enabled = _footer_enabled(hf_cfg)
        header_mode = str(hf_cfg.header_mode or "styleref")
        if not header_enabled:
            header_mode = "none"
        header_border = bool(hf_cfg.header_border) and header_enabled
        page_number_enabled = bool(hf_cfg.page_number_enabled)
        section_plan = build_page_number_execution_plan(doc, context, hf_cfg)
        _apply_document_header_footer_behavior(doc, hf_cfg)

        for section, plan in zip(doc.sections, section_plan.sections):
            section_type = plan.section_type
            _apply_section_header_footer_behavior(section, hf_cfg)
            if page_number_enabled and plan.phase_id is not None:
                _set_page_number_format(
                    section,
                    plan.number_format,
                    start=plan.start_value,
                )

            if _uses_advanced_default_variant(hf_cfg):
                count += _render_header_footer_variant(
                    section,
                    hf_cfg,
                    "default",
                    plan,
                    config,
                )
            else:
                if not plan.header_visible or header_mode == "none":
                    _clear_header(section)
                    count += 1
                elif header_mode == "fixed":
                    _set_fixed_header(section, hf_cfg)
                    count += 1
                else:
                    style_ref = None
                    include_number = True
                    if (
                        section_type in NON_NUMBERED_HEADING_SECTION_TYPES
                        or plan.special_title_selector is not None
                    ):
                        style_ref = getattr(config.heading_model, "non_numbered_heading_style_name", "") or None
                        include_number = False
                    _set_styleref_header(section, hf_cfg, style_ref=style_ref, include_number=include_number)
                    count += 1

                _set_header_border(
                    section,
                    header_border and plan.header_visible and header_mode != "none",
                    hf_cfg,
                )

                footer_content = _resolve_variant_content(hf_cfg, "default", "footer")
                if _render_footer_content(
                    section,
                    section.footer,
                    footer_content,
                    hf_cfg,
                    plan,
                    variant_name="default",
                    footer_text_visible=footer_enabled and plan.footer_text_visible,
                ):
                    count += 1

            if _different_first_page_enabled(hf_cfg):
                count += _render_header_footer_variant(
                    section,
                    hf_cfg,
                    "first",
                    plan,
                    config,
                )
            if _different_odd_even_pages_enabled(hf_cfg):
                count += _render_header_footer_variant(
                    section,
                    hf_cfg,
                    "even",
                    plan,
                    config,
                )

            # Font unification must run after header/footer content is rebuilt.
            _format_header_footer_font(section, hf_cfg)

        if count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{len(list(doc.sections))} 个节",
                section="global",
                change_type="format",
                before="(mixed)",
                after=(
                    f"header={header_mode}, "
                    f"border={header_border}, "
                    f"page_num={page_number_enabled}"
                ),
            )
        context.final_section_inventory = collect_section_inventory(doc)


def _protected_section_boundaries(doc):
    return tuple(
        boundary
        for boundary in collect_section_inventory(doc).boundaries
        if boundary.anchor_kind == "nested_paragraph"
        or "sectPrChange" in boundary.protected_markup
    )

def _apply_document_header_footer_behavior(doc, hf_cfg) -> None:
    settings = getattr(doc, "settings", None)
    if settings is not None and hasattr(settings, "odd_and_even_pages_header_footer"):
        settings.odd_and_even_pages_header_footer = _different_odd_even_pages_enabled(hf_cfg)


def _apply_section_header_footer_behavior(section, hf_cfg) -> None:
    if hasattr(section, "different_first_page_header_footer"):
        section.different_first_page_header_footer = _different_first_page_enabled(hf_cfg)


def _different_first_page_enabled(hf_cfg) -> bool:
    behavior = getattr(hf_cfg, "behavior", None)
    return bool(getattr(behavior, "different_first_page", False))


def _different_odd_even_pages_enabled(hf_cfg) -> bool:
    behavior = getattr(hf_cfg, "behavior", None)
    return bool(getattr(behavior, "different_odd_even_pages", False))


def _header_enabled(hf_cfg) -> bool:
    return bool(getattr(getattr(hf_cfg, "header", None), "enabled", True))


def _footer_enabled(hf_cfg) -> bool:
    return bool(getattr(getattr(hf_cfg, "footer", None), "enabled", True))


def _uses_advanced_default_variant(hf_cfg) -> bool:
    variant = _get_variant_config(hf_cfg, "default")
    return _content_mode(getattr(variant, "header", None)) != "inherit" or _content_mode(
        getattr(variant, "footer", None)
    ) != "inherit"


def _get_variant_config(hf_cfg, variant_name: str):
    variants = getattr(hf_cfg, "variants", None)
    return getattr(variants, variant_name, None) if variants is not None else None


def _content_mode(content) -> str:
    return str(getattr(content, "mode", "inherit") or "inherit").strip().lower()


def _variant_parts(section, variant_name: str):
    if variant_name == "first":
        return section.first_page_header, section.first_page_footer
    if variant_name == "even":
        return section.even_page_header, section.even_page_footer
    return section.header, section.footer


def _render_header_footer_variant(section, hf_cfg, variant_name: str, plan, config) -> int:
    header_part, footer_part = _variant_parts(section, variant_name)
    header_enabled = _header_enabled(hf_cfg)
    footer_enabled = _footer_enabled(hf_cfg)
    count = 0
    header_content = _resolve_variant_content(hf_cfg, variant_name, "header")
    footer_content = _resolve_variant_content(hf_cfg, variant_name, "footer")
    if not header_enabled or not plan.header_visible:
        if _clear_part(header_part):
            count += 1
        _set_part_header_border(header_part, False, hf_cfg)
    elif _render_header_content(header_part, header_content, hf_cfg, plan, config):
        _set_part_header_border(header_part, _content_mode(header_content) != "none", hf_cfg)
        count += 1
    if _render_footer_content(
        section,
        footer_part,
        footer_content,
        hf_cfg,
        plan,
        variant_name=variant_name,
        footer_text_visible=footer_enabled and plan.footer_text_visible,
    ):
        count += 1
    return count


def _resolve_variant_content(hf_cfg, variant_name: str, part_name: str):
    variant = _get_variant_config(hf_cfg, variant_name)
    content = getattr(variant, part_name, None) if variant is not None else None
    if _content_mode(content) != "inherit":
        return content
    if variant_name != "default":
        default_variant = _get_variant_config(hf_cfg, "default")
        default_content = getattr(default_variant, part_name, None) if default_variant is not None else None
        if _content_mode(default_content) != "inherit":
            return default_content
    return _legacy_content_config(hf_cfg, part_name)


def _legacy_content_config(hf_cfg, part_name: str):
    from src.config.feature_configs import HeaderFooterContentConfig

    if part_name == "header":
        return HeaderFooterContentConfig(
            mode=str(hf_cfg.header_mode or "styleref"),
            fixed_text=str(hf_cfg.header_text or ""),
            alignment=str(getattr(hf_cfg, "header_alignment", "center") or "center"),
            styleref_level=int(hf_cfg.styleref_level or 1),
        )
    footer_text = str(getattr(hf_cfg, "footer_text", "") or "")
    mode = str(
        getattr(getattr(hf_cfg, "footer", None), "content_mode", "none") or "none"
    ).strip().lower()
    if mode == "page_number_with_text":
        mode = "fixed"
    elif mode == "page_number":
        mode = "none"
    if not _footer_enabled(hf_cfg):
        mode = "none"
    return HeaderFooterContentConfig(
        mode=mode,
        fixed_text=footer_text,
        alignment=str(getattr(hf_cfg, "footer_alignment", "center") or "center"),
    )


def _render_header_content(part, content, hf_cfg, plan, config) -> bool:
    return _render_content_part(part, content, hf_cfg, plan, config, part_name="header")


def _render_footer_content(
    section,
    part,
    content,
    hf_cfg,
    plan,
    *,
    variant_name: str,
    footer_text_visible: bool,
) -> bool:
    if not _prepare_part_for_write(part, hf_cfg):
        return False
    if _global_preserve_existing_content(hf_cfg) and _part_has_content(part):
        return False

    mode = _content_mode(content)
    if mode == "preserve":
        return False
    text_visible = footer_text_visible and _footer_content_has_output(content)
    page_number_visible = _page_number_visible_for_variant(
        hf_cfg,
        variant_name,
        plan,
    )
    if not text_visible and not page_number_visible:
        _clear_part(part)
        return True

    _clear_part(part)
    para = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
    text_alignment = _normalize_footer_alignment(getattr(content, "alignment", "center"))
    page_alignment = _page_number_alignment_for_variant(hf_cfg, variant_name)
    page_template = _page_number_template_for_variant(hf_cfg, variant_name)

    if text_visible and page_number_visible:
        _write_combined_footer_outputs(
            section,
            para,
            content,
            text_alignment=text_alignment,
            page_alignment=page_alignment,
            page_template=page_template,
            plan=plan,
        )
    elif page_number_visible:
        _set_paragraph_alignment(para, page_alignment)
        _add_template_to_paragraph(
            para,
            page_template,
            text="",
            num_format=plan.number_format,
            styleref_level=1,
            styleref_style="",
            styleref_include_number=True,
        )
    else:
        _set_paragraph_alignment(para, text_alignment)
        _add_footer_text_content(para, content, plan)
    return True


def _page_number_variant_config(hf_cfg, variant_name: str):
    if variant_name not in {"first", "even"}:
        return None
    page_number_plan = getattr(hf_cfg, "page_number_plan", None)
    return getattr(page_number_plan, variant_name, None)


def _page_number_visible_for_variant(hf_cfg, variant_name: str, plan) -> bool:
    if not bool(getattr(hf_cfg, "page_number_enabled", True)):
        return False
    variant = _page_number_variant_config(hf_cfg, variant_name)
    visibility = str(getattr(variant, "visibility", "inherit") or "inherit").strip().lower()
    if visibility == "hide":
        return False
    if visibility == "show":
        return True
    return bool(plan.page_number_visible)


def _page_number_template_for_variant(hf_cfg, variant_name: str) -> str:
    variant = _page_number_variant_config(hf_cfg, variant_name)
    template = str(getattr(variant, "template", "") or "").strip()
    if not template:
        template = str(getattr(hf_cfg, "page_number_template", "{page}") or "{page}").strip()
    if "{page}" not in template:
        template = f"{template} {{page}}".strip()
    return template or "{page}"


def _page_number_alignment_for_variant(hf_cfg, variant_name: str) -> str:
    variant = _page_number_variant_config(hf_cfg, variant_name)
    alignment = str(getattr(variant, "alignment", "inherit") or "inherit").strip().lower()
    if alignment not in _FOOTER_ALIGNMENT_VALUES:
        alignment = str(getattr(hf_cfg, "page_number_alignment", "center") or "center")
    return _normalize_footer_alignment(alignment)


def _normalize_footer_alignment(alignment: str | None) -> str:
    raw = str(alignment or "center").strip().lower()
    return _FOOTER_ALIGNMENT_VALUES.get(raw, "center")


def _footer_content_has_output(content) -> bool:
    mode = _content_mode(content)
    if mode in {"fixed", "page_number_with_text"}:
        return bool(str(getattr(content, "fixed_text", "") or ""))
    if mode == "template":
        return bool(str(getattr(content, "template", "") or ""))
    return False


def _write_combined_footer_outputs(
    section,
    para,
    content,
    *,
    text_alignment: str,
    page_alignment: str,
    page_template: str,
    plan,
) -> None:
    if text_alignment == page_alignment:
        _set_paragraph_alignment(para, page_alignment)
        _add_template_to_paragraph(
            para,
            page_template,
            text="",
            num_format=plan.number_format,
            styleref_level=1,
            styleref_style="",
            styleref_include_number=True,
        )
        para.add_run(" ")
        _add_footer_text_content(para, content, plan)
        return

    _set_paragraph_alignment(para, "left")
    segments = sorted(
        ((text_alignment, "text"), (page_alignment, "page")),
        key=lambda item: {"left": 0, "center": 1, "right": 2}[item[0]],
    )
    _set_footer_tab_stops(para, section, {alignment for alignment, _kind in segments})
    for index, (alignment, kind) in enumerate(segments):
        if index > 0 or alignment != "left":
            para.add_run("\t")
        if kind == "page":
            _add_template_to_paragraph(
                para,
                page_template,
                text="",
                num_format=plan.number_format,
                styleref_level=1,
                styleref_style="",
                styleref_include_number=True,
            )
        else:
            _add_footer_text_content(para, content, plan)


def _set_footer_tab_stops(para, section, alignments: set[str]) -> None:
    p_pr = find_or_create(para._element, "w:pPr")
    tabs = find_or_create_before(
        p_pr,
        "w:tabs",
        (
            "w:suppressAutoHyphens",
            "w:kinsoku",
            "w:wordWrap",
            "w:overflowPunct",
            "w:topLinePunct",
            "w:autoSpaceDE",
            "w:autoSpaceDN",
            "w:bidi",
            "w:adjustRightInd",
            "w:snapToGrid",
            "w:spacing",
            "w:ind",
            "w:contextualSpacing",
            "w:mirrorIndents",
            "w:suppressOverlap",
            "w:jc",
            "w:textDirection",
            "w:textAlignment",
            "w:textboxTightWrap",
            "w:outlineLvl",
            "w:divId",
            "w:cnfStyle",
            "w:rPr",
            "w:sectPr",
            "w:pPrChange",
        ),
    )
    for child in list(tabs):
        tabs.remove(child)
    page_width = int(getattr(section, "page_width", 0) or 0)
    left_margin = int(getattr(section, "left_margin", 0) or 0)
    right_margin = int(getattr(section, "right_margin", 0) or 0)
    usable_twips = max(1440, int((page_width - left_margin - right_margin) / 635))
    for alignment, position in (
        ("center", usable_twips // 2),
        ("right", usable_twips),
    ):
        if alignment not in alignments:
            continue
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), alignment)
        tab.set(qn("w:pos"), str(position))
        tabs.append(tab)


def _add_footer_text_content(para, content, plan) -> None:
    mode = _content_mode(content)
    if mode == "template":
        _add_template_to_paragraph(
            para,
            str(getattr(content, "template", "") or ""),
            text=str(getattr(content, "fixed_text", "") or ""),
            num_format=plan.number_format,
            styleref_level=int(getattr(content, "styleref_level", 1) or 1),
            styleref_style=str(getattr(content, "styleref_style", "") or ""),
            styleref_include_number=bool(getattr(content, "styleref_include_number", True)),
        )
        return
    text = str(getattr(content, "fixed_text", "") or "")
    if text:
        para.add_run(text)


def _render_content_part(
    part,
    content,
    hf_cfg,
    plan,
    config=None,
    *,
    part_name: str,
    footer_text_visible: bool = True,
    page_number_visible: bool = True,
) -> bool:
    if not _prepare_part_for_write(part, hf_cfg):
        return False

    mode = _content_mode(content)
    if mode == "preserve":
        return False
    if _global_preserve_existing_content(hf_cfg) and _part_has_content(part):
        return False
    if mode == "none":
        _clear_part(part)
        return True
    if part_name == "footer" and not page_number_visible:
        if footer_text_visible and str(getattr(content, "fixed_text", "") or ""):
            return _write_text_part(
                part,
                str(getattr(content, "fixed_text", "") or ""),
                content,
            )
        _clear_part(part)
        return True
    if mode == "fixed":
        text = str(getattr(content, "fixed_text", "") or "") if footer_text_visible else ""
        return _write_text_part(part, text, content)
    if mode == "styleref":
        return _write_styleref_part(part, content, plan, config)
    if mode in {"page_number", "page_number_with_text"}:
        return _write_template_part(
            part,
            _page_number_template_for_content(content, mode),
            content,
            plan,
            include_text=footer_text_visible,
        )
    if mode == "template":
        template = str(getattr(content, "template", "") or "")
        return _write_template_part(
            part,
            template,
            content,
            plan,
            include_text=footer_text_visible,
        )

    _clear_part(part)
    return True


def _prepare_part_for_write(part, hf_cfg) -> bool:
    behavior = getattr(hf_cfg, "behavior", None)
    link_mode = str(getattr(behavior, "link_to_previous", "never") or "never").strip().lower()
    if link_mode == "always":
        part.is_linked_to_previous = True
        return False
    if link_mode == "never":
        part.is_linked_to_previous = False
    return True


def _global_preserve_existing_content(hf_cfg) -> bool:
    behavior = getattr(hf_cfg, "behavior", None)
    return bool(getattr(behavior, "preserve_existing_content", False))


def _part_has_content(part) -> bool:
    for para in part.paragraphs:
        if para.text:
            return True
        if any(True for _kind, _elem, _instr in iter_field_instructions(para._element)):
            return True
    return False


def _write_text_part(part, text: str, content) -> bool:
    _clear_part(part)
    para = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
    _set_paragraph_alignment(para, getattr(content, "alignment", "center"))
    if text:
        para.add_run(text)
    return True


def _write_styleref_part(part, content, plan, config) -> bool:
    _clear_part(part)
    para = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
    _set_paragraph_alignment(para, getattr(content, "alignment", "center"))
    style_ref = str(getattr(content, "styleref_style", "") or "").strip()
    include_number = bool(getattr(content, "styleref_include_number", True))
    if (
        plan.section_type in NON_NUMBERED_HEADING_SECTION_TYPES
        or plan.special_title_selector is not None
    ) and config is not None:
        style_ref = getattr(config.heading_model, "non_numbered_heading_style_name", "") or style_ref
        include_number = False
    level = int(getattr(content, "styleref_level", 1) or 1)
    _add_styleref_to_paragraph(
        para,
        style_name=style_ref,
        level=level,
        include_number=include_number,
    )
    return True


def _page_number_template_for_content(content, mode: str) -> str:
    template = str(getattr(content, "template", "") or "{page}")
    text = str(getattr(content, "fixed_text", "") or "")
    if mode == "page_number_with_text" and text and "{text}" not in template:
        return f"{template} {{text}}"
    return template


def _write_template_part(
    part,
    template: str,
    content,
    plan,
    *,
    include_text: bool = True,
) -> bool:
    _clear_part(part)
    para = part.paragraphs[0] if part.paragraphs else part.add_paragraph()
    _set_paragraph_alignment(para, getattr(content, "alignment", "center"))
    _add_template_to_paragraph(
        para,
        template,
        text=(
            str(getattr(content, "fixed_text", "") or "")
            if include_text
            else ""
        ),
        num_format=plan.number_format,
        styleref_level=int(getattr(content, "styleref_level", 1) or 1),
        styleref_style=str(getattr(content, "styleref_style", "") or ""),
        styleref_include_number=bool(getattr(content, "styleref_include_number", True)),
    )
    return True


def _add_template_to_paragraph(
    para,
    template: str,
    *,
    text: str,
    num_format: str,
    styleref_level: int,
    styleref_style: str,
    styleref_include_number: bool,
) -> None:
    source = str(template or "{page}")
    cursor = 0
    for match in _CONTENT_TEMPLATE_TOKEN_RE.finditer(source):
        if match.start() > cursor:
            para.add_run(source[cursor:match.start()])
        token = match.group(1)
        if token == "page":
            _add_field_to_paragraph(para, _page_field_instruction(num_format))
        elif token == "pages":
            _add_field_to_paragraph(para, " NUMPAGES ")
        elif token == "section_pages":
            _add_field_to_paragraph(para, " SECTIONPAGES ")
        elif token == "text" and text:
            para.add_run(text)
        elif token in {"styleref", "styleref_number"}:
            style_name = styleref_style.strip()
            level = max(1, int(styleref_level or 1))
            if token == "styleref":
                _add_styleref_to_paragraph(
                    para,
                    style_name=style_name,
                    level=level,
                    include_number=styleref_include_number,
                )
            elif style_name:
                _add_field_to_paragraph(
                    para,
                    f' STYLEREF "{style_name}" \\n ',
                )
            else:
                _add_field_to_paragraph(para, f" STYLEREF {level} \\n ")
        cursor = match.end()
    if cursor < len(source):
        para.add_run(source[cursor:])


def _page_field_instruction(num_format: str) -> str:
    return _PAGE_NUMBER_FIELD_INSTRUCTIONS.get(
        str(num_format or "decimal"),
        _PAGE_NUMBER_FIELD_INSTRUCTIONS["decimal"],
    )


def _set_styleref_header(section, hf_cfg, *, style_ref: str | None = None, include_number: bool = True) -> None:
    header = section.header
    header.is_linked_to_previous = False

    for para in header.paragraphs:
        for run in list(para.runs):
            run._element.getparent().remove(run._element)

    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    _set_paragraph_alignment(para, getattr(hf_cfg, "header_alignment", "center"))
    level = int(hf_cfg.styleref_level or 1)
    style_name = str(style_ref or "").strip()

    _add_styleref_to_paragraph(
        para,
        style_name=style_name,
        level=level,
        include_number=include_number,
    )


def _add_styleref_to_paragraph(
    para,
    *,
    style_name: str,
    level: int,
    include_number: bool,
) -> None:
    normalized_style_name = str(style_name or "").strip()
    if not normalized_style_name:
        # Numeric built-in style identifiers are locale-independent. The
        # heading-numbering module writes the visible number into the heading
        # text, so a single STYLEREF returns both number and title.
        _add_field_to_paragraph(para, f" STYLEREF {max(1, int(level or 1))} ")
        return

    if include_number:
        _add_field_to_paragraph(
            para,
            f' STYLEREF "{normalized_style_name}" \\n ',
        )
    _add_field_to_paragraph(para, f' STYLEREF "{normalized_style_name}" ')


def _set_fixed_header(section, hf_cfg) -> None:
    header = section.header
    header.is_linked_to_previous = False

    for para in header.paragraphs:
        for run in list(para.runs):
            run._element.getparent().remove(run._element)

    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    _set_paragraph_alignment(para, getattr(hf_cfg, "header_alignment", "center"))
    text = str(hf_cfg.header_text or "")
    if text:
        para.add_run(text)


def _clear_header(section) -> None:
    header = section.header
    header.is_linked_to_previous = False

    if not header.paragraphs:
        header.add_paragraph()
    for para in header.paragraphs:
        _clear_paragraph_runs(para)


def _clear_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False

    if not footer.paragraphs:
        footer.add_paragraph()
    for para in footer.paragraphs:
        _clear_paragraph_runs(para)


def _set_header_border(section, enable: bool, hf_cfg=None) -> None:
    header = section.header
    _set_part_header_border(header, enable, hf_cfg)


def _set_part_header_border(header, enable: bool, hf_cfg=None) -> None:
    if not header.paragraphs:
        return

    para = header.paragraphs[0]
    p_pr = find_or_create(para._element, "w:pPr")
    p_bdr = find_or_create(p_pr, "w:pBdr")
    bottom = find_or_create(p_bdr, "w:bottom")
    border_cfg = getattr(getattr(hf_cfg, "header", None), "border_style", None) if hf_cfg is not None else None

    if enable and (border_cfg is None or bool(getattr(border_cfg, "enabled", True))):
        width_pt = max(0.25, float(getattr(border_cfg, "width_pt", 0.5) or 0.5))
        spacing_pt = max(0.0, float(getattr(border_cfg, "spacing_pt", 1.0) or 0.0))
        bottom.set(qn("w:val"), str(getattr(border_cfg, "line_style", "single") or "single"))
        bottom.set(qn("w:sz"), str(max(2, int(round(width_pt * 8)))))
        bottom.set(qn("w:space"), str(int(round(spacing_pt))))
        bottom.set(qn("w:color"), str(getattr(border_cfg, "color", "auto") or "auto"))
    else:
        bottom.set(qn("w:val"), "none")
        bottom.set(qn("w:sz"), "0")


def _set_page_number(
    section,
    hf_cfg=None,
    *,
    num_format: str = "decimal",
    include_footer_text: bool = True,
) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False

    for para in footer.paragraphs:
        _clear_paragraph_runs(para)

    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p_pr = find_or_create(para._element, "w:pPr")
    jc = find_or_create(p_pr, "w:jc")
    jc.set(qn("w:val"), _page_number_alignment(hf_cfg))

    footer_text = (
        str(getattr(hf_cfg, "footer_text", "") or "")
        if hf_cfg is not None and include_footer_text
        else ""
    )
    template = str(getattr(hf_cfg, "page_number_template", "{page}") or "{page}") if hf_cfg is not None else "{page}"
    if footer_text and "{text}" not in template:
        template = f"{template} {{text}}"
    _add_template_to_paragraph(
        para,
        template,
        text=footer_text,
        num_format=num_format,
        styleref_level=1,
        styleref_style="",
        styleref_include_number=True,
    )


def _set_fixed_footer(section, hf_cfg) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False

    for para in footer.paragraphs:
        _clear_paragraph_runs(para)

    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p_pr = find_or_create(para._element, "w:pPr")
    jc = find_or_create(p_pr, "w:jc")
    jc.set(qn("w:val"), _footer_alignment(hf_cfg))
    text = str(getattr(hf_cfg, "footer_text", "") or "")
    if text:
        para.add_run(text)


def _footer_alignment(hf_cfg) -> str:
    raw = str(getattr(hf_cfg, "footer_alignment", "center") or "center").strip().lower()
    return _FOOTER_ALIGNMENT_VALUES.get(raw, "center")


def _page_number_alignment(hf_cfg) -> str:
    raw = str(
        getattr(hf_cfg, "page_number_alignment", "center") or "center"
    ).strip().lower()
    return _FOOTER_ALIGNMENT_VALUES.get(raw, "center")


def _set_page_number_format(section, fmt: str = "decimal", *, start: int | None = None) -> None:
    sect_pr = section._sectPr
    pg_num_type = find_or_create_before(
        sect_pr,
        "w:pgNumType",
        (
            "w:cols",
            "w:formProt",
            "w:vAlign",
            "w:noEndnote",
            "w:titlePg",
            "w:textDirection",
            "w:bidi",
            "w:rtlGutter",
            "w:docGrid",
            "w:printerSettings",
            "w:sectPrChange",
        ),
    )

    normalized_fmt = _PAGE_NUMBER_W_FORMATS.get(str(fmt or "decimal"), "decimal")
    pg_num_type.set(qn("w:fmt"), normalized_fmt)
    if start is not None:
        pg_num_type.set(qn("w:start"), str(int(start)))
    else:
        pg_num_type.attrib.pop(qn("w:start"), None)


def _clear_page_number_fields(section) -> bool:
    footer = section.footer
    footer.is_linked_to_previous = False

    changed = False
    for para in footer.paragraphs:
        if not _paragraph_has_field(para, "PAGE"):
            continue
        _clear_paragraph_runs(para)
        changed = True
    return changed


def _format_header_footer_font(section, hf_cfg) -> None:
    header_parts = [section.header]
    footer_parts = [section.footer]
    if _different_first_page_enabled(hf_cfg):
        header_parts.append(section.first_page_header)
        footer_parts.append(section.first_page_footer)
    if _different_odd_even_pages_enabled(hf_cfg):
        header_parts.append(section.even_page_header)
        footer_parts.append(section.even_page_footer)

    _format_parts_font(header_parts, getattr(getattr(hf_cfg, "header", None), "typography", None))
    _format_parts_font(footer_parts, getattr(getattr(hf_cfg, "footer", None), "typography", None))


def _format_parts_font(parts, typography) -> None:
    font_cn = getattr(typography, "font_cn", None)
    font_en = getattr(typography, "font_en", None)
    size_pt = getattr(typography, "size_pt", None)
    bold = bool(getattr(typography, "bold", False))
    italic = bool(getattr(typography, "italic", False))

    if not any((font_cn, font_en, size_pt, bold, italic)):
        return

    from docx.shared import Pt

    for part in parts:
        for para in part.paragraphs:
            for run in para.runs:
                if font_en:
                    run.font.name = resolve_font(font_en, lang="en")
                if font_cn:
                    set_run_east_asian_font(run, font_cn)
                if size_pt:
                    run.font.size = Pt(size_pt)
                run.font.bold = bold
                run.font.italic = italic


def _add_field_to_paragraph(para, instr: str) -> None:
    for elem in build_complex_field(instr, result_text=" "):
        para._element.append(elem)


def _clear_paragraph_runs(para) -> None:
    for child in list(para._element):
        if child.tag != qn("w:pPr"):
            para._element.remove(child)


def _clear_part(part) -> bool:
    if not _prepare_part_for_clear(part):
        return False
    if not part.paragraphs:
        part.add_paragraph()
    for para in part.paragraphs:
        _clear_paragraph_runs(para)
    return True


def _prepare_part_for_clear(part) -> bool:
    part.is_linked_to_previous = False
    return True


def _set_paragraph_alignment(para, alignment: str | None) -> None:
    raw = str(alignment or "center").strip().lower()
    value = _FOOTER_ALIGNMENT_VALUES.get(raw, "center")
    p_pr = find_or_create(para._element, "w:pPr")
    jc = find_or_create(p_pr, "w:jc")
    jc.set(qn("w:val"), value)


def _paragraph_has_field(para, field_keyword: str) -> bool:
    keyword = field_keyword.strip().upper()
    for _kind, _elem, instr in iter_field_instructions(para._element):
        normalized = " ".join((instr or "").upper().split())
        if normalized.startswith(keyword):
            return True
    return False
