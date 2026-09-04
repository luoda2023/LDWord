"""Standalone formula typography independent from equation-table structure."""

from __future__ import annotations

import re
from collections.abc import Iterable
from copy import deepcopy
from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from lxml import etree

from src.formula_core import parse_document_formulas
from src.formula_core.semantics import (
    BOLD_ITALIC_MARKER_COMMANDS,
    BOLD_ROMAN_MARKER_COMMANDS,
    FUNCTION_NAMES,
    GREEK_COMMAND_NAMES,
    ITALIC_MARKER_COMMANDS,
    ROMAN_MARKER_COMMANDS,
    UPRIGHT_FAMILY_MARKER_COMMANDS,
)
from src.modules.base import BaseModule, ModuleMeta
from src.modules.special.formula_typography_ops import (
    apply_math_run_typography,
    apply_paragraph_alignment,
    apply_paragraph_spacing,
    apply_word_run_typography,
)
from src.modules.table.table_format import top_level_table_anchor_positions
from src.shared.engine.document_scope_runtime import document_scope_allows_paragraph
from src.shared.engine.ooxml_ops import qn

if TYPE_CHECKING:
    from docx import Document
    from docx.text.paragraph import Paragraph
    from lxml.etree import _Element

    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


_FUNCTIONS = frozenset(str(name).lower() for name in FUNCTION_NAMES)
_MATH_TOKEN_RE = re.compile(
    r"\s+|\\[A-Za-z]+|[0-9]+(?:\.[0-9]+)?|[A-Za-z]+|"
    r"[\u0391-\u03ff]+|[\u1f00-\u1fff]+|."
)
_NUMBER_RE = re.compile(r"^[0-9]+(?:\.[0-9]+)?$")
_IDENTIFIER_RE = re.compile(r"^[A-Za-z\u0391-\u03ff\u1f00-\u1fff]+$")
_GREEK_COMMANDS = frozenset(GREEK_COMMAND_NAMES)
_COMMAND_RE = re.compile(r"^\\([A-Za-z]+)$")


class FormulaStyleModule(BaseModule):
    """Apply formula font, size, spacing and semantic math-run style."""

    meta = ModuleMeta(
        name="formula_style",
        description="统一独立公式与行内公式样式",
        category="special",
        requires_config=("formula_table", "formula_style"),
        soft_after=("formula_to_table", "equation_table_format"),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
        *,
        body_scope_positions: dict[object, int] | None = None,
        table_scope_positions: dict[object, int] | None = None,
    ) -> None:
        visual = config.formula_table
        policy = config.formula_style
        font_name = (
            str(visual.formula_font_name or "").strip()
            if bool(policy.unify_font)
            else None
        )
        size_pt = (
            float(visual.formula_font_size_pt)
            if bool(policy.unify_size) and visual.formula_font_size_pt
            else None
        )
        parse_result = parse_document_formulas(doc)
        block_paragraph_elements = {
            occurrence.paragraph._p
            for occurrence in parse_result.occurrences
            if _occurrence_is_block_formula(occurrence)
        }

        formula_paragraphs = 0
        semantic_runs = 0
        eligible_paragraphs = list(
            _iter_style_paragraphs(
                doc,
                context,
                body_scope_positions=body_scope_positions,
                table_scope_positions=table_scope_positions,
            )
        )
        eligible_paragraph_elements = {
            paragraph._p for paragraph in eligible_paragraphs
        }
        for paragraph in eligible_paragraphs:
            math_nodes = paragraph._p.findall(f".//{qn('m:oMath')}")
            if not math_nodes:
                continue
            formula_paragraphs += 1
            if paragraph._p in block_paragraph_elements:
                apply_paragraph_alignment(
                    paragraph,
                    str(visual.block_alignment or "center"),
                )
                if bool(policy.unify_spacing):
                    apply_paragraph_spacing(
                        paragraph,
                        line_spacing=float(visual.formula_line_spacing or 1.0),
                        space_before_pt=float(visual.formula_space_before_pt or 0.0),
                        space_before_unit=str(visual.formula_space_before_unit or "pt"),
                        space_after_pt=float(visual.formula_space_after_pt or 0.0),
                        space_after_unit=str(visual.formula_space_after_unit or "pt"),
                    )

            apply_math_run_typography(
                paragraph._p,
                font_name=font_name,
                size_pt=size_pt,
            )
            for math_run in list(paragraph._p.findall(f".//{qn('m:r')}")):
                if _apply_semantic_math_style(math_run):
                    semantic_runs += 1

        # ``keep_source`` is a supported conversion policy.  Formula-only
        # source paragraphs still receive visual typography and layout, while
        # ordinary data-table cells stay outside the formula rule boundary.
        styled_source_elements: set[object] = set()
        for occurrence in parse_result.occurrences:
            paragraph = occurrence.paragraph
            if (
                occurrence.source_type == "word_native"
                or not occurrence.is_formula_only
                or paragraph._p in styled_source_elements
                or paragraph._p not in eligible_paragraph_elements
            ):
                continue
            styled_source_elements.add(paragraph._p)
            formula_paragraphs += 1
            if _occurrence_is_block_formula(occurrence):
                apply_paragraph_alignment(
                    paragraph,
                    str(visual.block_alignment or "center"),
                )
                if bool(policy.unify_spacing):
                    apply_paragraph_spacing(
                        paragraph,
                        line_spacing=float(visual.formula_line_spacing or 1.0),
                        space_before_pt=float(visual.formula_space_before_pt or 0.0),
                        space_before_unit=str(visual.formula_space_before_unit or "pt"),
                        space_after_pt=float(visual.formula_space_after_pt or 0.0),
                        space_after_unit=str(visual.formula_space_after_unit or "pt"),
                    )
            apply_word_run_typography(
                paragraph.runs,
                font_name=font_name,
                size_pt=size_pt,
            )

        runtime = dict(context.formula_runtime or {})
        runtime.update(
            {
                "style_enabled": True,
                "styled_formula_paragraphs": formula_paragraphs,
                "semantic_runs": semantic_runs,
            }
        )
        context.formula_runtime = runtime
        if formula_paragraphs:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{formula_paragraphs} 个公式段落",
                section="formula",
                change_type="format",
                before="混合公式样式",
                after=(
                    f"font={font_name or 'preserve'}, size={size_pt or 'preserve'}, "
                    f"semantic_runs={semantic_runs}"
                ),
            )


def _iter_style_paragraphs(
    doc: Document,
    context: PipelineContext,
    *,
    body_scope_positions: dict[object, int] | None = None,
    table_scope_positions: dict[object, int] | None = None,
) -> Iterable[Paragraph]:
    # Body formulae are in scope.  Table formulae are in scope only when the
    # table has the structural shape / marker of an equation table.
    resolved_body_positions = body_scope_positions or {
        paragraph._p: index for index, paragraph in enumerate(doc.paragraphs)
    }
    for paragraph in doc.paragraphs:
        position = resolved_body_positions.get(paragraph._p, -1)
        if document_scope_allows_paragraph(context, position):
            yield paragraph

    if table_scope_positions is None:
        anchors = top_level_table_anchor_positions(doc)
        resolved_table_positions = {
            table._element: anchors[index] if index < len(anchors) else -1
            for index, table in enumerate(doc.tables)
        }
    else:
        resolved_table_positions = table_scope_positions

    for table in doc.tables:
        position = resolved_table_positions.get(table._element, -1)
        if not document_scope_allows_paragraph(context, position):
            continue
        yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table) -> Iterable[Paragraph]:
    from src.modules.table.table_format import _is_equation_table

    allowed = _is_equation_table(table)
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            identity = id(cell._tc)
            if identity in seen:
                continue
            seen.add(identity)
            if allowed:
                yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _is_formula_only_paragraph(paragraph: Paragraph) -> bool:
    outside = paragraph._p.xpath(
        ".//w:t[not(ancestor::m:oMath) and not(ancestor::m:oMathPara)]/text()"
    )
    return not "".join(str(item) for item in outside).strip()


def _occurrence_is_block_formula(occurrence) -> bool:
    if bool(getattr(occurrence, "is_block", False)):
        return True
    return (
        occurrence.paragraph.paragraph_format.alignment
        == WD_ALIGN_PARAGRAPH.CENTER
    )


def _paragraph_is_style_allowed(paragraph: Paragraph) -> bool:
    node = paragraph._p
    while node is not None:
        parent = node.getparent()
        if parent is None:
            return False
        local = _local_name(parent)
        if local == "body":
            return True
        if local == "tbl":
            from src.modules.table.table_format import _is_equation_table

            return _is_equation_table(parent)
        node = parent
    return False


def _apply_semantic_math_style(math_run: _Element) -> bool:
    text_nodes = [child for child in list(math_run) if child.tag == qn("m:t")]
    text = "".join(str(node.text or "") for node in text_nodes)
    if not text:
        return False
    # Explicit m:sty is author intent; inferred semantics must not overwrite it.
    if _explicit_math_style(math_run):
        return _remove_conflicting_word_toggles(math_run)

    segments = _semantic_segments(
        text,
        in_script=_run_in_script(math_run),
        superscript_base=_run_is_superscript_base(math_run),
    )
    if not segments:
        return False
    if len(segments) == 1 and segments[0][0] == text and len(text_nodes) == 1:
        return _set_inferred_math_style(math_run, segments[0][1])
    return _split_math_run(math_run, segments)


def _local_name(node) -> str:
    tag = str(getattr(node, "tag", "") or "")
    return tag.rsplit("}", 1)[-1]


def _run_has_math_ancestor(math_run, slot: str, parents: set[str]) -> bool:
    node = math_run
    while node is not None:
        if _local_name(node) == slot and _local_name(node.getparent()) in parents:
            return True
        node = node.getparent()
    return False


def _run_in_script(math_run) -> bool:
    return _run_has_math_ancestor(
        math_run, "sub", {"sSub", "sSubSup", "sPre", "nary"}
    ) or _run_has_math_ancestor(
        math_run, "sup", {"sSup", "sSubSup", "sPre", "nary"}
    )


def _run_is_superscript_base(math_run) -> bool:
    return _run_has_math_ancestor(math_run, "e", {"sSup", "sSubSup"})


def _explicit_math_style(math_run) -> str:
    run_pr = math_run.find(qn("m:rPr"))
    style = run_pr.find(qn("m:sty")) if run_pr is not None else None
    if style is None:
        return ""
    return str(style.get(qn("m:val")) or style.get("val") or "").strip().lower()


def _remove_conflicting_word_toggles(math_run) -> bool:
    run_pr = math_run.find(qn("m:rPr"))
    word_properties = [math_run.find(qn("w:rPr"))]
    if run_pr is not None:
        word_properties.append(run_pr.find(qn("w:rPr")))
    changed = False
    for word_pr in word_properties:
        if word_pr is None:
            continue
        for tag in ("w:i", "w:iCs", "w:b", "w:bCs"):
            child = word_pr.find(qn(tag))
            if child is not None:
                word_pr.remove(child)
                changed = True
    return changed


def _significant(tokens: list[str], index: int, direction: int) -> str:
    cursor = index + direction
    while 0 <= cursor < len(tokens):
        if not tokens[cursor].isspace():
            return tokens[cursor]
        cursor += direction
    return ""


def _command_name(token: str) -> str | None:
    match = _COMMAND_RE.match(str(token or ""))
    return match.group(1).strip().lower() if match else None


def _marker_style_for_command(command: str | None) -> str | None:
    if command in BOLD_ITALIC_MARKER_COMMANDS:
        return "bi"
    if command in BOLD_ROMAN_MARKER_COMMANDS:
        return "b"
    if command in ROMAN_MARKER_COMMANDS or command in UPRIGHT_FAMILY_MARKER_COMMANDS:
        return "p"
    if command in ITALIC_MARKER_COMMANDS:
        return "i"
    return None


def _explicit_marker_style(tokens: list[str], index: int) -> str | None:
    token = tokens[index]
    if _command_name(token) is None and not _IDENTIFIER_RE.fullmatch(token):
        return None
    previous_index = index - 1
    while previous_index >= 0 and tokens[previous_index].isspace():
        previous_index -= 1
    if previous_index >= 0:
        direct = _marker_style_for_command(_command_name(tokens[previous_index]))
        if direct is not None:
            return direct

    depth = 0
    open_index: int | None = None
    cursor = index - 1
    while cursor >= 0:
        if tokens[cursor] == "}":
            depth += 1
        elif tokens[cursor] == "{":
            if depth == 0:
                open_index = cursor
                break
            depth -= 1
        cursor -= 1
    if open_index is None:
        return None
    marker_index = open_index - 1
    while marker_index >= 0 and tokens[marker_index].isspace():
        marker_index -= 1
    return (
        _marker_style_for_command(_command_name(tokens[marker_index]))
        if marker_index >= 0
        else None
    )


def _is_differential(tokens: list[str], index: int) -> bool:
    if str(tokens[index] or "").lower() != "d":
        return False
    following = _significant(tokens, index, 1)
    if not _IDENTIFIER_RE.fullmatch(following):
        return False
    previous = _significant(tokens, index, -1)
    if previous in {"/", "∂", ")", "]"}:
        return True
    if _IDENTIFIER_RE.fullmatch(previous) or _NUMBER_RE.fullmatch(previous):
        return True
    # ``dx/du`` at the beginning of a flattened run.
    next_index = index + 1
    while next_index < len(tokens) and tokens[next_index].isspace():
        next_index += 1
    next_index += 1
    while next_index < len(tokens) and tokens[next_index].isspace():
        next_index += 1
    return not previous and next_index < len(tokens) and tokens[next_index] == "/"


def _is_imaginary_unit(tokens: list[str], index: int) -> bool:
    if str(tokens[index] or "").lower() not in {"i", "j"}:
        return False
    previous = _significant(tokens, index, -1)
    following = _significant(tokens, index, 1)
    if following == "^":
        return True
    return previous in {"=", "+", "-", "(", ","} and (
        not following or not _IDENTIFIER_RE.fullmatch(following)
    )


def _is_natural_constant(tokens: list[str], index: int) -> bool:
    if str(tokens[index] or "").lower() != "e":
        return False
    if _significant(tokens, index, 1) != "^":
        return False
    previous = _significant(tokens, index, -1)
    return not (
        _NUMBER_RE.fullmatch(previous)
        or _IDENTIFIER_RE.fullmatch(previous)
        or previous in {")", "]", "}"}
    )


def _semantic_segments(
    text: str,
    *,
    in_script: bool,
    superscript_base: bool,
) -> list[tuple[str, str]]:
    raw = [match.group(0) for match in _MATH_TOKEN_RE.finditer(text)]
    tokens: list[str] = []
    for token in raw:
        if (
            not in_script
            and _IDENTIFIER_RE.fullmatch(token)
            and token.lower() not in _FUNCTIONS
            and len(token) > 1
        ):
            tokens.extend(token)
        else:
            tokens.append(token)

    styled: list[tuple[str, str]] = []
    for index, token in enumerate(tokens):
        low = token.lower()
        previous = _significant(tokens, index, -1)
        following = _significant(tokens, index, 1)
        marker_style = _explicit_marker_style(tokens, index)
        if token.isspace() or _NUMBER_RE.fullmatch(token):
            style = "p"
        elif token.startswith("\\"):
            style = "i" if low[1:] in _GREEK_COMMANDS else "p"
        elif not _IDENTIFIER_RE.fullmatch(token):
            style = "p"
        elif marker_style is not None:
            style = marker_style
        elif low in _FUNCTIONS and (
            following in {"(", "[", "{"}
            or bool(_IDENTIFIER_RE.fullmatch(following))
        ):
            style = "p"
        elif in_script:
            style = "i" if len(token) == 1 else "p"
        elif _is_differential(tokens, index):
            style = "p"
        elif _is_natural_constant(tokens, index) or (
            low == "e" and superscript_base and len(token) == 1
        ):
            style = "p"
        elif _is_imaginary_unit(tokens, index):
            style = "p"
        else:
            style = "i" if len(token) == 1 else "p"
        if styled and styled[-1][1] == style:
            styled[-1] = (styled[-1][0] + token, style)
        else:
            styled.append((token, style))
    return styled


def _set_inferred_math_style(math_run, style: str) -> bool:
    run_pr = math_run.find(qn("m:rPr"))
    if run_pr is None:
        run_pr = OxmlElement("m:rPr")
        math_run.insert(0, run_pr)
    style_el = run_pr.find(qn("m:sty"))
    if style_el is None:
        style_el = OxmlElement("m:sty")
        run_pr.append(style_el)
    if style_el.get(qn("m:val")) == style:
        return False
    style_el.set(qn("m:val"), style)
    return True


def _split_math_run(math_run, segments: list[tuple[str, str]]) -> bool:
    parent = math_run.getparent()
    if parent is None:
        return False
    insert_at = parent.index(math_run)
    template_children = [
        deepcopy(child) for child in list(math_run) if child.tag != qn("m:t")
    ]
    for text, style in segments:
        new_run = etree.Element(math_run.tag)
        for child in template_children:
            new_run.append(deepcopy(child))
        text_el = etree.SubElement(new_run, qn("m:t"))
        text_el.text = text
        _set_inferred_math_style(new_run, style)
        parent.insert(insert_at, new_run)
        insert_at += 1
    parent.remove(math_run)
    return True


__all__ = ["FormulaStyleModule"]
