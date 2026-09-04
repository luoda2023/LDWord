"""equation_table_format - formula-table formatting module."""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

from src.config.style_semantics import normalize_spacing_unit, spacing_value_to_pt
from src.modules.base import BaseModule, ModuleMeta
from src.modules.special.formula_typography_ops import (
    apply_math_run_typography,
    apply_paragraph_alignment,
    apply_paragraph_spacing,
    apply_word_run_typography,
)
from src.modules.table.table_format import (
    _is_equation_table,
    _looks_like_formula_cell_text,
    top_level_table_anchor_positions,
)
from src.shared.engine.document_scope_runtime import document_scope_allows_paragraph
from src.shared.engine.font_resolver import resolve_font
from src.shared.engine.line_spacing_ops import apply_line_spacing, sync_spacing_ooxml
from src.shared.engine.ooxml_ops import find_or_create, qn
from src.shared.engine.sequence_numbering import (
    build_heading_chapter_ranges,
    parse_chapter_numbering_format,
    resolve_chapter_number,
)
from src.shared.engine.table_builder import clear_cell_border_overrides

if TYPE_CHECKING:
    from docx import Document
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    from lxml.etree import _Element

    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


_PARAGRAPH_ALIGNMENTS: dict[str, WD_ALIGN_PARAGRAPH] = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}

_TABLE_ALIGNMENTS = frozenset({"left", "center", "right"})
_WRAPPER_PAIRS = {
    "(": ")",
    "（": "）",
    "[": "]",
    "【": "】",
}
_EQUATION_NUMBER_RE = re.compile(
    r"^\s*(?P<open>[\(\[（【])?\s*(?P<head>\d+)"
    r"(?:\s*(?P<sep>[.\-:—–_/])\s*(?P<tail>\d+))?\s*(?P<close>[\)\]）】])?\s*$"
)
_EQUATION_NUMBER_TAIL_RE = re.compile(
    r"(?P<open>[\(\[（【])\s*\d+"
    r"(?:\s*[.\-:—–/]\s*\d+)?\s*(?P<close>[\)\]）】])\s*$"
)
_OFFICE_NS = "urn:schemas-microsoft-com:office:office"
_EQUATION_OLE_HINT_RE = re.compile(r"(?:equation|mathtype|eqn)", re.IGNORECASE)


@dataclass(frozen=True)
class EquationTablePlan:
    table_alignment: str
    formula_alignment: str
    number_alignment: str
    formula_font_name: str | None
    formula_font_size_pt: float | None
    formula_line_spacing: float
    formula_space_before_pt: float
    formula_space_before_unit: str
    formula_space_after_pt: float
    formula_space_after_unit: str
    number_font_name: str | None
    number_font_size_pt: float | None
    auto_shrink_number_column: bool
    unify_spacing: bool
    include_chapter_number: bool
    numbering_separator: str
    numbering_format: str


@dataclass
class EquationNumberingState:
    global_sequence: int = 0
    chapter_sequences: dict[int, int] = field(default_factory=dict)


@dataclass(frozen=True)
class NumberCellFormatResult:
    normalized: bool = False
    skipped: bool = False


@dataclass
class EquationTableStats:
    normalized_number_count: int = 0
    skipped_number_count: int = 0


@dataclass(frozen=True)
class NumberParagraphMatch:
    paragraph_index: int
    open_wrapper: str
    close_wrapper: str


class EquationTableFormatModule(BaseModule):
    """Execute the thesis-owned formula-table workflow."""

    meta = ModuleMeta(
        name="equation_table_format",
        description="公式表格格式",
        category="special",
        requires_config=("formula_table", "formula_style", "equation_numbering"),
        soft_after=("table_format",),
        soft_consumes=("heading_map", "doc_tree"),
        modifies_structure=True,
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        heading_map = getattr(context, "heading_map", None) or {}
        doc_tree = getattr(context, "doc_tree", None)
        chapter_ranges = build_heading_chapter_ranges(
            heading_map,
            len(doc.paragraphs),
            doc_tree=doc_tree,
        )
        body_scope_positions = {
            paragraph._p: index for index, paragraph in enumerate(doc.paragraphs)
        }
        original_table_anchors = top_level_table_anchor_positions(doc)
        table_scope_positions = {
            table._element: (
                original_table_anchors[index]
                if index < len(original_table_anchors)
                else -1
            )
            for index, table in enumerate(doc.tables)
        }

        # In the thesis plan, standalone native equations become two-column
        # equation tables before numbering/layout is applied. Other plans keep
        # their existing structure while still receiving formula typography.
        if (
            str(getattr(config, "mode_id", "") or "").strip() == "thesis"
            and bool(config.formula_to_table.enabled)
        ):
            from src.modules.special.formula_to_table import FormulaToTableModule

            table_scope_positions.update(
                FormulaToTableModule().apply(doc, config, tracker, context)
            )

        plan = _resolve_equation_table_plan(config)
        numbering_state = EquationNumberingState()
        inferred_document_chapter = _dominant_existing_equation_chapter(doc.tables)

        count = 0
        normalized_number_count = 0
        skipped_number_count = 0

        if bool(config.equation_numbering.enabled):
            for table_index, table in enumerate(doc.tables):
                if not _is_equation_table(table):
                    continue

                table_para_index = table_scope_positions.get(table._element, -1)
                if not document_scope_allows_paragraph(context, table_para_index):
                    continue
                stats = _format_equation_table(
                    table,
                    plan,
                    table_para_index=table_para_index,
                    chapter_ranges=chapter_ranges,
                    numbering_state=numbering_state,
                    fallback_chapter=inferred_document_chapter,
                )
                normalized_number_count += stats.normalized_number_count
                skipped_number_count += stats.skipped_number_count
                count += 1

        if count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{count} 个公式表格",
                section="global",
                change_type="format",
                before="(mixed)",
                after=(
                    f"table={plan.table_alignment}, "
                    f"formula={plan.formula_alignment}, "
                    f"number={plan.number_alignment}, "
                    f"numbering={plan.numbering_format}"
                ),
            )
        if normalized_number_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{normalized_number_count} 个公式编号",
                section="global",
                change_type="format",
                before="existing numbering text",
                after=f"normalized to {plan.numbering_format}",
            )
        if skipped_number_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{skipped_number_count} 个公式编号",
                section="global",
                change_type="skip",
                before="chapter-aware numbering normalization",
                after="kept for review because numbering could not be normalized safely",
            )

        runtime = dict(context.formula_runtime or {})
        runtime.update(
            {
                "formatted_equation_tables": count,
                "normalized_numbers": normalized_number_count,
                "skipped_numbers": skipped_number_count,
            }
        )
        context.formula_runtime = runtime

        # Style is deliberately independent from table recognition: this pass
        # also covers native inline and standalone formulas not in a table.
        from src.modules.special.formula_style import FormulaStyleModule

        if bool(config.formula_style.enabled):
            FormulaStyleModule().apply(
                doc,
                config,
                tracker,
                context,
                body_scope_positions=body_scope_positions,
                table_scope_positions=table_scope_positions,
            )


def _resolve_equation_table_plan(config: ResolvedConfig) -> EquationTablePlan:
    formula_cfg = config.formula_table
    style_cfg = config.formula_style
    numbering_cfg = config.equation_numbering

    block_alignment = _normalize_alignment_keyword(formula_cfg.block_alignment, default="center")
    table_alignment = _resolve_block_alignment(
        formula_cfg.table_alignment,
        block_alignment,
        default="center",
    )
    formula_alignment = _resolve_block_alignment(
        formula_cfg.formula_cell_alignment,
        block_alignment,
        default="center",
    )
    number_alignment = _normalize_alignment_keyword(formula_cfg.number_alignment, default="right")

    formula_font_name = (
        _clean_font_name(formula_cfg.formula_font_name)
        if bool(style_cfg.unify_font)
        else None
    )
    formula_font_size_pt = (
        _clean_optional_positive_float(formula_cfg.formula_font_size_pt)
        if bool(style_cfg.unify_size)
        else None
    )
    formula_line_spacing = _clean_positive_float(formula_cfg.formula_line_spacing, default=1.0)
    formula_space_before_pt = _clean_nonnegative_float(
        formula_cfg.formula_space_before_pt,
        default=0.0,
    )
    formula_space_before_unit = normalize_spacing_unit(
        getattr(formula_cfg, "formula_space_before_unit", "pt")
    )
    formula_space_after_pt = _clean_nonnegative_float(
        formula_cfg.formula_space_after_pt,
        default=0.0,
    )
    formula_space_after_unit = normalize_spacing_unit(
        getattr(formula_cfg, "formula_space_after_unit", "pt")
    )

    number_font_name = _clean_font_name(formula_cfg.number_font_name)
    # A number-specific size is an explicit thesis-plan requirement (for example
    # 10.5 pt equation numbers beside 12 pt formulae).  ``unify_size`` is only
    # a fallback when that value is absent; it must not overwrite it.
    number_font_size_pt = _clean_optional_positive_float(
        formula_cfg.number_font_size_pt
    )
    if number_font_size_pt is None and bool(style_cfg.unify_size):
        number_font_size_pt = formula_font_size_pt

    numbering_format = str(getattr(numbering_cfg, "numbering_format", "") or "chapter.seq")
    include_chapter_number, numbering_separator = parse_chapter_numbering_format(numbering_format)

    return EquationTablePlan(
        table_alignment=table_alignment,
        formula_alignment=formula_alignment,
        number_alignment=number_alignment,
        formula_font_name=formula_font_name,
        formula_font_size_pt=formula_font_size_pt,
        formula_line_spacing=formula_line_spacing,
        formula_space_before_pt=formula_space_before_pt,
        formula_space_before_unit=formula_space_before_unit,
        formula_space_after_pt=formula_space_after_pt,
        formula_space_after_unit=formula_space_after_unit,
        number_font_name=number_font_name,
        number_font_size_pt=number_font_size_pt,
        auto_shrink_number_column=bool(formula_cfg.auto_shrink_number_column),
        unify_spacing=bool(style_cfg.unify_spacing),
        include_chapter_number=include_chapter_number,
        numbering_separator=numbering_separator,
        numbering_format=numbering_format,
    )


def _format_equation_table(
    table: Table,
    plan: EquationTablePlan,
    *,
    table_para_index: int,
    chapter_ranges: list[tuple[int, int, int]],
    numbering_state: EquationNumberingState,
    fallback_chapter: int = 0,
) -> EquationTableStats:
    _clear_table_borders(table)
    _set_table_alignment(table._element, plan.table_alignment)

    chapter_num = resolve_chapter_number(table_para_index, chapter_ranges)
    if chapter_num <= 0:
        chapter_num = _existing_equation_table_chapter(table) or fallback_chapter
    if plan.include_chapter_number and chapter_num <= 0:
        # Preserve the 0.2 fallback for documents whose heading map is absent
        # or incomplete: a valid equation row can still start safely at 1.1.
        chapter_num = 1
    stats = EquationTableStats()
    allow_plain_formula_rows = _has_explicit_equation_table_marker(table)

    for row in table.rows:
        cells = row.cells
        if len(cells) >= 2:
            number_cell_is_empty = _cell_is_effectively_empty(cells[-1])
            row_has_formula = _row_has_formula_content(
                cells[:-1],
                allow_plain_text=allow_plain_formula_rows,
            )
            # A table is classified once, but formatting/numbering is a row-level
            # decision.  Do not let one genuine equation row pull unrelated
            # notes, image rows, or metadata rows into the formula rule.
            if not row_has_formula:
                continue
            moved_wrapper: tuple[str, str] | None = None
            blocked_number_move = False
            if number_cell_is_empty and row_has_formula:
                moved_wrapper, blocked_number_move = (
                    _extract_trailing_number_wrapper(cells[:-1])
                )
            for cell in cells[:-1]:
                _format_formula_cell(
                    cell,
                    plan,
                    allow_plain_text=allow_plain_formula_rows,
                )
            number_result = _format_number_cell(
                cells[-1],
                plan,
                chapter_num=chapter_num,
                numbering_state=numbering_state,
                create_if_empty=(
                    number_cell_is_empty
                    and not blocked_number_move
                    and row_has_formula
                ),
                preferred_wrapper=moved_wrapper,
                number_required=row_has_formula,
            )
            stats.normalized_number_count += int(number_result.normalized)
            stats.skipped_number_count += int(number_result.skipped)
            continue

        if not _row_has_formula_content(
            cells,
            allow_plain_text=allow_plain_formula_rows,
        ):
            continue
        for cell in cells:
            _format_formula_cell(
                cell,
                plan,
                allow_plain_text=allow_plain_formula_rows,
            )

    if plan.auto_shrink_number_column:
        _apply_equation_column_geometry(table)
    return stats


def _existing_equation_table_chapter(table: Table) -> int:
    for row in table.rows:
        if len(row.cells) < 2:
            continue
        for paragraph in row.cells[-1].paragraphs:
            match = _EQUATION_NUMBER_RE.match(str(paragraph.text or "").strip())
            if match is None or match.group("tail") is None:
                continue
            try:
                return max(0, int(match.group("head")))
            except (TypeError, ValueError):
                continue
    return 0


def _dominant_existing_equation_chapter(tables) -> int:
    chapters = [
        chapter
        for table in tables
        if _is_equation_table(table)
        for chapter in [_existing_equation_table_chapter(table)]
        if chapter > 0
    ]
    if not chapters:
        return 0
    return int(Counter(chapters).most_common(1)[0][0])


def _clear_table_borders(table: Table) -> None:
    tbl_el = table._element
    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    for tag in ("w:tblStyle", "w:tblLook"):
        inherited_style = tbl_pr.find(qn(tag))
        if inherited_style is not None:
            tbl_pr.remove(inherited_style)
    tbl_borders = find_or_create(tbl_pr, "w:tblBorders")

    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = find_or_create(tbl_borders, f"w:{side}")
        border.set(qn("w:val"), "none")
        border.set(qn("w:sz"), "0")
        border.set(qn("w:space"), "0")

    clear_cell_border_overrides(tbl_el)
    for cell in tbl_el.iter(qn("w:tc")):
        cell_properties = cell.find(qn("w:tcPr"))
        if cell_properties is None:
            cell_properties = OxmlElement("w:tcPr")
            cell.insert(0, cell_properties)
        vertical_alignment = find_or_create(cell_properties, "w:vAlign")
        vertical_alignment.set(qn("w:val"), "center")


def _set_table_alignment(tbl_el: _Element, alignment: str) -> None:
    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    jc = find_or_create(tbl_pr, "w:jc")
    jc.set(qn("w:val"), alignment)


def _apply_equation_column_geometry(table: Table) -> None:
    """Make the final number column compact while preserving total width."""

    tbl_el = table._element
    grid = tbl_el.find(qn("w:tblGrid"))
    grid_cols = list(grid.findall(qn("w:gridCol"))) if grid is not None else []
    column_count = len(grid_cols)
    if column_count < 2:
        column_count = max((len(row.cells) for row in table.rows), default=0)
    if column_count < 2:
        return

    current_widths: list[int] = []
    for col in grid_cols:
        try:
            current_widths.append(int(col.get(qn("w:w")) or 0))
        except (TypeError, ValueError):
            current_widths.append(0)
    total_width = sum(width for width in current_widths if width > 0)
    if total_width <= 0:
        total_width = 8640

    number_length = max(
        (
            len((row.cells[-1].text or "").strip())
            for row in table.rows
            if len(row.cells) >= 2
        ),
        default=5,
    )
    number_width = max(900, min(1800, 720 + number_length * 110))
    number_width = min(number_width, max(900, total_width // 4))
    if column_count == 2:
        widths = [max(900, total_width - number_width), number_width]
    else:
        left_width = number_width
        middle_count = column_count - 2
        middle_total = max(900 * middle_count, total_width - left_width - number_width)
        middle_width = middle_total // middle_count
        widths = [left_width] + [middle_width] * middle_count + [number_width]
        widths[-2] += total_width - sum(widths)

    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl_pr = tbl_el.find(qn("w:tblPr"))
        insert_at = tbl_el.index(tbl_pr) + 1 if tbl_pr is not None else 0
        tbl_el.insert(insert_at, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    tbl_width = find_or_create(tbl_pr, "w:tblW")
    tbl_width.set(qn("w:type"), "dxa")
    tbl_width.set(qn("w:w"), str(sum(widths)))
    layout = find_or_create(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    for row in table.rows:
        for index, cell in enumerate(row.cells[:column_count]):
            tc_pr = find_or_create(cell._tc, "w:tcPr")
            tc_width = find_or_create(tc_pr, "w:tcW")
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(widths[index]))


def _format_formula_cell(
    cell: _Cell,
    plan: EquationTablePlan,
    *,
    allow_plain_text: bool = False,
) -> None:
    nonempty_paragraphs = [
        para for para in cell.paragraphs if (para.text or "").strip()
    ]
    allow_single_plain_paragraph = (
        allow_plain_text and len(nonempty_paragraphs) == 1
    )

    for para in cell.paragraphs:
        # Equation-table geometry owns cell alignment, matching the legacy
        # visual rule. Typography and spacing, however, belong only to an
        # evidenced formula paragraph; otherwise an explanatory paragraph in
        # the same cell is silently converted to formula styling.
        apply_paragraph_alignment(para, plan.formula_alignment)
        formula_evidence, formula_only, has_omml = (
            _paragraph_formula_style_evidence(
                para,
                allow_plain_text=(
                    allow_single_plain_paragraph
                    and para is nonempty_paragraphs[0]
                ),
            )
        )
        if not formula_evidence:
            continue
        if formula_only and plan.unify_spacing:
            apply_paragraph_spacing(
                para,
                line_spacing=plan.formula_line_spacing,
                space_before_pt=plan.formula_space_before_pt,
                space_before_unit=plan.formula_space_before_unit,
                space_after_pt=plan.formula_space_after_pt,
                space_after_unit=plan.formula_space_after_unit,
            )
        if formula_only:
            apply_word_run_typography(
                para.runs,
                font_name=plan.formula_font_name,
                size_pt=plan.formula_font_size_pt,
            )
        if has_omml:
            apply_math_run_typography(
                para._element,
                font_name=plan.formula_font_name,
                size_pt=plan.formula_font_size_pt,
            )


def _paragraph_formula_style_evidence(
    paragraph: Paragraph,
    *,
    allow_plain_text: bool = False,
) -> tuple[bool, bool, bool]:
    """Return (has_formula, formula_only, has_omml) for one paragraph."""

    element = paragraph._p
    has_omml = any(
        element.find(f".//{qn(tag)}") is not None
        for tag in ("m:oMath", "m:oMathPara")
    )
    has_equation_ole = _element_has_equation_ole(element)
    text = str(paragraph.text or "").strip()
    has_formula_text = bool(text and _looks_like_formula_cell_text(text))
    has_plain_formula_text = bool(text and allow_plain_text)
    has_formula = (
        has_omml
        or has_equation_ole
        or has_formula_text
        or has_plain_formula_text
    )
    formula_only = (
        has_formula_text
        or has_plain_formula_text
        or ((has_omml or has_equation_ole) and not text)
    )
    return has_formula, formula_only, has_omml


def _format_number_cell(
    cell: _Cell,
    plan: EquationTablePlan,
    *,
    chapter_num: int,
    numbering_state: EquationNumberingState,
    create_if_empty: bool = False,
    preferred_wrapper: tuple[str, str] | None = None,
    number_required: bool = False,
) -> NumberCellFormatResult:
    match = _find_number_paragraph(cell)
    if match is None and create_if_empty and cell.paragraphs:
        match = NumberParagraphMatch(
            paragraph_index=0,
            open_wrapper=(preferred_wrapper or ("(", ")"))[0],
            close_wrapper=(preferred_wrapper or ("(", ")"))[1],
        )
    target_paragraph = (
        cell.paragraphs[match.paragraph_index]
        if match is not None
        else None
    )
    if target_paragraph is not None:
        if _paragraph_has_numbering_fields_or_bookmarks(target_paragraph):
            result = NumberCellFormatResult(skipped=number_required)
        elif plan.include_chapter_number and chapter_num <= 0:
            result = NumberCellFormatResult(skipped=True)
        else:
            normalized_text = _next_equation_number_text(
                plan,
                chapter_num=chapter_num,
                numbering_state=numbering_state,
                match=match,
            )
            _replace_paragraph_text(target_paragraph, normalized_text)
            result = NumberCellFormatResult(normalized=True)
    else:
        result = NumberCellFormatResult(skipped=number_required)

    if target_paragraph is None:
        return result
    apply_paragraph_alignment(target_paragraph, plan.number_alignment)
    if plan.unify_spacing:
        apply_paragraph_spacing(
            target_paragraph,
            line_spacing=plan.formula_line_spacing,
            space_before_pt=plan.formula_space_before_pt,
            space_before_unit=plan.formula_space_before_unit,
            space_after_pt=plan.formula_space_after_pt,
            space_after_unit=plan.formula_space_after_unit,
        )
    apply_word_run_typography(
        target_paragraph.runs,
        font_name=plan.number_font_name,
        size_pt=plan.number_font_size_pt,
    )
    return result


def _row_has_formula_content(
    cells: list[_Cell],
    *,
    allow_plain_text: bool = False,
) -> bool:
    for cell in cells:
        element = cell._tc
        for tag in ("m:oMath", "m:oMathPara"):
            if element.find(f".//{qn(tag)}") is not None:
                return True
        if _cell_has_equation_ole(cell):
            return True
        text = " ".join(
            (paragraph.text or "").strip()
            for paragraph in cell.paragraphs
            if (paragraph.text or "").strip()
        )
        if text and (
            allow_plain_text
            or _looks_like_formula_cell_text(text)
        ):
            return True
    return False


def _cell_has_equation_ole(cell: _Cell) -> bool:
    """Return True only for Equation/MathType OLE, not arbitrary objects."""

    return _element_has_equation_ole(cell._tc)


def _element_has_equation_ole(element: _Element) -> bool:
    for ole in element.findall(f".//{{{_OFFICE_NS}}}OLEObject"):
        prog_id = str(
            ole.get("ProgID")
            or ole.get(f"{{{_OFFICE_NS}}}ProgID")
            or ""
        )
        if _EQUATION_OLE_HINT_RE.search(prog_id):
            return True
    return False


def _has_explicit_equation_table_marker(table: Table) -> bool:
    table_properties = table._element.find(qn("w:tblPr"))
    if table_properties is None:
        return False
    for tag in ("w:tblCaption", "w:tblDescription"):
        marker = table_properties.find(qn(tag))
        value = (marker.get(qn("w:val")) or "") if marker is not None else ""
        if value.strip().casefold() == "ldword-equation-table":
            return True
    return False


def _cell_is_effectively_empty(cell: _Cell) -> bool:
    if any((paragraph.text or "").strip() for paragraph in cell.paragraphs):
        return False
    return not any(
        cell._tc.find(f".//{qn(tag)}") is not None
        for tag in (
            "m:oMath",
            "m:oMathPara",
            "w:object",
            "w:drawing",
            "w:pict",
            "w:fldChar",
            "w:instrText",
            "w:bookmarkStart",
            "w:bookmarkEnd",
        )
    )


def _extract_trailing_number_wrapper(
    formula_cells: list[_Cell],
) -> tuple[tuple[str, str] | None, bool]:
    """Move a same-cell trailing equation number without flattening OMML.

    The boolean result marks a detected tail that could not be moved safely;
    callers then avoid creating a duplicate number in the empty final cell.
    """

    for cell in reversed(formula_cells):
        for paragraph in reversed(cell.paragraphs):
            text_nodes = list(paragraph._p.findall(f".//{qn('w:t')}"))
            text = "".join(str(node.text or "") for node in text_nodes)
            match = _EQUATION_NUMBER_TAIL_RE.search(text)
            if match is None:
                continue
            if _paragraph_has_numbering_fields_or_bookmarks(paragraph):
                return None, True

            open_wrapper = match.group("open") or "("
            close_wrapper = match.group("close") or ")"
            if _WRAPPER_PAIRS.get(open_wrapper) != close_wrapper:
                open_wrapper, close_wrapper = "(", ")"
            retained_text = text[: match.start()].rstrip()
            remaining = retained_text
            for node in text_nodes:
                original_length = len(str(node.text or ""))
                node.text = remaining[:original_length]
                remaining = remaining[original_length:]
            return (open_wrapper, close_wrapper), False
    return None, False


def _paragraph_has_numbering_fields_or_bookmarks(para: Paragraph) -> bool:
    element = para._element
    return any(
        element.find(f".//{qn(tag)}") is not None
        for tag in ("w:fldChar", "w:instrText", "w:bookmarkStart", "w:bookmarkEnd")
    )


def _find_number_paragraph(cell: _Cell) -> NumberParagraphMatch | None:
    for paragraph_index, para in enumerate(cell.paragraphs):
        match = _EQUATION_NUMBER_RE.match((para.text or "").strip())
        if match is None:
            continue
        open_wrapper = match.group("open") or ""
        close_wrapper = match.group("close") or ""
        if open_wrapper and _WRAPPER_PAIRS.get(open_wrapper) != close_wrapper:
            open_wrapper = ""
            close_wrapper = ""
        return NumberParagraphMatch(
            paragraph_index=paragraph_index,
            open_wrapper=open_wrapper,
            close_wrapper=close_wrapper,
        )
    return None


def _next_equation_number_text(
    plan: EquationTablePlan,
    *,
    chapter_num: int,
    numbering_state: EquationNumberingState,
    match: NumberParagraphMatch,
) -> str:
    if plan.include_chapter_number and chapter_num > 0:
        next_seq = numbering_state.chapter_sequences.get(chapter_num, 0) + 1
        numbering_state.chapter_sequences[chapter_num] = next_seq
        core = f"{chapter_num}{plan.numbering_separator}{next_seq}" if plan.numbering_separator else f"{chapter_num}{next_seq}"
    else:
        numbering_state.global_sequence += 1
        core = str(numbering_state.global_sequence)

    return f"{match.open_wrapper}{core}{match.close_wrapper}"


def _replace_paragraph_text(para: Paragraph, new_text: str) -> None:
    for child in list(para._element):
        if child.tag != qn("w:pPr"):
            para._element.remove(child)
    para.add_run(new_text)


def _apply_paragraph_alignment(para: Paragraph, alignment: str) -> None:
    para.paragraph_format.alignment = _PARAGRAPH_ALIGNMENTS.get(
        alignment,
        WD_ALIGN_PARAGRAPH.CENTER,
    )


def _apply_paragraph_spacing(
    para: Paragraph,
    *,
    line_spacing: float,
    space_before_pt: float,
    space_before_unit: str = "pt",
    space_after_pt: float,
    space_after_unit: str = "pt",
) -> None:
    pf = para.paragraph_format
    apply_line_spacing(pf, "multiple", line_spacing)
    before_abs_pt = spacing_value_to_pt(space_before_pt, space_before_unit)
    after_abs_pt = spacing_value_to_pt(space_after_pt, space_after_unit)
    if before_abs_pt is not None:
        from docx.shared import Pt
        pf.space_before = Pt(before_abs_pt)
    if after_abs_pt is not None:
        from docx.shared import Pt
        pf.space_after = Pt(after_abs_pt)
    sync_spacing_ooxml(
        para._element,
        space_before_pt=space_before_pt,
        space_before_unit=space_before_unit,
        space_after_pt=space_after_pt,
        space_after_unit=space_after_unit,
        line_spacing_type="multiple",
        line_spacing_value=line_spacing,
    )


def _apply_word_run_typography(
    runs: list[Run],
    *,
    font_name: str | None,
    size_pt: float | None,
) -> None:
    if not runs or (font_name is None and size_pt is None):
        return

    resolved_font = resolve_font(font_name, lang="en") if font_name else None
    half_points = _size_to_half_points(size_pt)

    for run in runs:
        r_pr = _find_or_insert_word_run_pr(run._element)
        if resolved_font:
            r_fonts = find_or_create(r_pr, "w:rFonts")
            r_fonts.set(qn("w:ascii"), resolved_font)
            r_fonts.set(qn("w:hAnsi"), resolved_font)
            r_fonts.set(qn("w:cs"), resolved_font)
        if half_points is not None:
            find_or_create(r_pr, "w:sz").set(qn("w:val"), half_points)
            find_or_create(r_pr, "w:szCs").set(qn("w:val"), half_points)


def _apply_math_run_typography(
    container: _Element,
    *,
    font_name: str | None,
    size_pt: float | None,
) -> None:
    if font_name is None and size_pt is None:
        return

    resolved_font = resolve_font(font_name, lang="en") if font_name else None
    half_points = _size_to_half_points(size_pt)

    for math_run in container.findall(f".//{qn('m:r')}"):
        r_pr = _find_or_insert_math_word_run_pr(math_run)
        if resolved_font:
            r_fonts = find_or_create(r_pr, "w:rFonts")
            r_fonts.set(qn("w:ascii"), resolved_font)
            r_fonts.set(qn("w:hAnsi"), resolved_font)
            r_fonts.set(qn("w:cs"), resolved_font)
        if half_points is not None:
            find_or_create(r_pr, "w:sz").set(qn("w:val"), half_points)
            find_or_create(r_pr, "w:szCs").set(qn("w:val"), half_points)


def _find_or_insert_word_run_pr(run_el: _Element) -> _Element:
    r_pr = run_el.find(qn("w:rPr"))
    if r_pr is not None:
        return r_pr
    r_pr = OxmlElement("w:rPr")
    run_el.insert(0, r_pr)
    return r_pr


def _find_or_insert_math_word_run_pr(math_run_el: _Element) -> _Element:
    r_pr = math_run_el.find(qn("w:rPr"))
    if r_pr is not None:
        return r_pr

    r_pr = OxmlElement("w:rPr")
    children = list(math_run_el)
    insert_at = 0
    for index, child in enumerate(children):
        if child.tag == qn("m:rPr"):
            insert_at = index + 1
            break
    math_run_el.insert(insert_at, r_pr)
    return r_pr


def _resolve_block_alignment(primary: str | None, block_alignment: str, *, default: str) -> str:
    primary_alignment = _normalize_alignment_keyword(primary, default=default)
    if primary_alignment == default and block_alignment != default:
        return block_alignment
    return primary_alignment


def _normalize_alignment_keyword(value: str | None, *, default: str) -> str:
    normalized = str(value or "").strip().lower()
    if normalized in _TABLE_ALIGNMENTS:
        return normalized
    return default


def _clean_font_name(value: str | None) -> str | None:
    normalized = str(value or "").strip()
    return normalized or None


def _clean_optional_positive_float(value) -> float | None:
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    if parsed <= 0:
        return None
    return parsed


def _clean_positive_float(value, *, default: float) -> float:
    parsed = _clean_optional_positive_float(value)
    return parsed if parsed is not None else default


def _clean_nonnegative_float(value, *, default: float) -> float:
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return default
    return parsed if parsed >= 0 else default


def _size_to_half_points(size_pt: float | None) -> str | None:
    if size_pt is None:
        return None
    return str(int(round(size_pt * 2)))
