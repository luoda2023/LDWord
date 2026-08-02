"""Scope-aware chemical typography and sub/superscript recovery."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING

from src.modules.base import BaseModule, ModuleMeta
from src.modules.table.table_format import (
    _is_equation_table,
    top_level_table_anchor_positions,
)
from src.shared.engine.chem_marks import (
    apply_chem_typography_to_paragraph,
    normalize_explicit_run_font_hints_in_paragraph,
    paragraph_has_chem_typography_evidence,
)
from src.shared.engine.document_scope_runtime import document_scope_allows_paragraph
from src.shared.engine.paragraph_iter import iter_table_cells

if TYPE_CHECKING:
    from docx import Document
    from docx.text.paragraph import Paragraph
    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


_SCOPE_KEYS = (
    "references", "body", "headings", "abstract_cn", "abstract_en", "captions", "tables",
)
_HEADING_STYLE_RE = re.compile(r"^(heading\s*\d+|heading|标题)", re.IGNORECASE)
_CAPTION_RE = re.compile(r"^(?:图|表|Figure|Table|Fig\.?|Tab\.?)\s*", re.IGNORECASE)
_BODY_LIKE_SECTIONS = {"body", "errata", "acknowledgment", "appendix", "resume"}


class ChemTypographyModule(BaseModule):
    meta = ModuleMeta(
        name="chem_typography",
        description="化学式排版",
        category="special",
        requires_config=("chem_typography",),
        soft_after=("paragraph_style", "heading_recognition", "caption"),
        soft_consumes=("doc_tree", "heading_map"),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        chem_config = config.chem_typography
        if not bool(chem_config.enabled):
            return
        active_scopes = _resolve_active_scopes(chem_config)
        paragraph_count = 0
        font_chars = 0
        mark_chars = 0
        hint_count = 0

        for paragraph_index, paragraph in enumerate(doc.paragraphs):
            if not str(paragraph.text or "").strip():
                continue
            if not _should_process_paragraph(
                paragraph_index,
                paragraph,
                context=context,
                active_scopes=active_scopes,
            ):
                continue
            if not paragraph_has_chem_typography_evidence(
                paragraph,
                chem_cfg=chem_config,
            ):
                continue
            hint_count += normalize_explicit_run_font_hints_in_paragraph(
                paragraph
            )
            changed_fonts, changed_marks = apply_chem_typography_to_paragraph(
                paragraph, chem_cfg=chem_config
            )
            if changed_fonts or changed_marks:
                paragraph_count += 1
                font_chars += changed_fonts
                mark_chars += changed_marks

        if "tables" in active_scopes:
            anchors = top_level_table_anchor_positions(doc)
            for table_index, table in enumerate(doc.tables):
                # Formula typography owns equation tables. Running chemical
                # script recovery there can reinterpret equation indices or
                # charges a second time and was explicitly excluded in 0.2.
                if _is_equation_table(table):
                    continue
                anchor = anchors[table_index] if table_index < len(anchors) else -1
                if not document_scope_allows_paragraph(context, anchor):
                    continue
                for _, _, cell in iter_table_cells(table):
                    for paragraph in cell.paragraphs:
                        if not str(paragraph.text or "").strip():
                            continue
                        if not paragraph_has_chem_typography_evidence(
                            paragraph,
                            chem_cfg=chem_config,
                        ):
                            continue
                        hint_count += normalize_explicit_run_font_hints_in_paragraph(
                            paragraph
                        )
                        changed_fonts, changed_marks = apply_chem_typography_to_paragraph(
                            paragraph, chem_cfg=chem_config
                        )
                        if changed_fonts or changed_marks:
                            paragraph_count += 1
                            font_chars += changed_fonts
                            mark_chars += changed_marks

        if paragraph_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{paragraph_count} 个段落",
                section="global",
                change_type="format",
                before="化学式格式异常",
                after=f"恢复上下角标 {mark_chars} 个字符，西文字体 {font_chars} 个字符",
            )
        if hint_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{hint_count} 个字体提示",
                section="global",
                change_type="format",
                before="显式西文字体仍带 EastAsia 主题提示",
                after="仅在选中范围内改为显式字体优先",
            )


def _resolve_active_scopes(chem_config) -> set[str]:
    raw = getattr(chem_config, "scopes", {}) or {}
    active = {key for key in _SCOPE_KEYS if bool(raw.get(key, False))}
    if bool(raw.get("abstract", False)):
        active.update({"abstract_cn", "abstract_en"})
    return active


def _apply_chem_typography(
    paragraph: Paragraph,
    chem_config,
) -> tuple[int, int]:
    """Compatibility entry for the migrated paragraph-level engine."""

    return apply_chem_typography_to_paragraph(
        paragraph,
        chem_cfg=chem_config,
    )


def _get_section_type(paragraph_index: int, context: PipelineContext) -> str:
    tree = getattr(context, "doc_tree", None)
    getter = getattr(tree, "get_section_for_paragraph", None)
    if callable(getter):
        try:
            section = getter(paragraph_index)
            if section:
                return str(section)
        except Exception:
            pass
    return "body"


def _is_heading_paragraph(
    paragraph_index: int,
    paragraph: Paragraph,
    context: PipelineContext,
) -> bool:
    if paragraph_index in (getattr(context, "heading_map", None) or {}):
        return True
    tree = getattr(context, "doc_tree", None)
    getter = getattr(tree, "get_heading_level", None)
    if callable(getter):
        try:
            if getter(paragraph_index) is not None:
                return True
        except Exception:
            pass
    style = paragraph.style
    return bool(_HEADING_STYLE_RE.match((style.name if style else "") or ""))


def _should_process_paragraph(
    paragraph_index: int,
    paragraph: Paragraph,
    *,
    context: PipelineContext,
    active_scopes: set[str],
) -> bool:
    if not document_scope_allows_paragraph(context, paragraph_index):
        return False
    if _is_heading_paragraph(paragraph_index, paragraph, context):
        return "headings" in active_scopes
    if _CAPTION_RE.match(str(paragraph.text or "").strip()):
        return "captions" in active_scopes
    section = _get_section_type(paragraph_index, context)
    if section == "references":
        return "references" in active_scopes
    if section in {"abstract_cn", "abstract_en"}:
        return section in active_scopes
    if section in _BODY_LIKE_SECTIONS:
        return "body" in active_scopes
    return False


__all__ = ["ChemTypographyModule", "_apply_chem_typography"]
