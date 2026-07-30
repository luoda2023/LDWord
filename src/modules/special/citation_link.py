"""Link numeric body citations to numbered reference entries."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING

from docx.oxml import OxmlElement

from src.modules.base import BaseModule, ModuleMeta
from src.shared.engine.document_scope_runtime import document_scope_allows_role
from src.shared.engine.ooxml_ops import qn

if TYPE_CHECKING:
    from docx import Document
    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


REFERENCE_ENTRY_RE = re.compile(r"^\s*[\[\uff3b(（]?\s*(?P<num>\d{1,4})\s*[\]\uff3d)）\.]?\s*")
CITATION_TOKEN_RE = re.compile(r"(?P<open>\[|\uff3b)(?P<inner>[^\[\]\uff3b\uff3d]+)(?P<close>\]|\uff3d)")
NUM_TOKEN_RE = re.compile(r"\d{1,4}")
REF_ENTRY_BOOKMARK_PREFIX = "_RefEntry_"
REF_NUM_BOOKMARK_PREFIX = "_RefNum_"


class CitationLinkModule(BaseModule):
    meta = ModuleMeta(
        name="citation_link",
        description="\u5f15\u6587\u4e0e\u53c2\u8003\u6587\u732e\u5173\u8054",
        category="special",
        consumes=("doc_tree",),
        soft_after=("heading_recognition", "reference_format"),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        if not (
            document_scope_allows_role(context, "body")
            and document_scope_allows_role(context, "references")
        ):
            return
        citation_cfg = config.citation_link
        doc_tree = getattr(context, "doc_tree", None)

        ref_range = _resolve_reference_range(doc, doc_tree)
        if ref_range is None:
            tracker.record(
                rule_name=self.meta.name,
                target="references",
                section="global",
                change_type="skip",
                before="reference section detection",
                after="skipped because reference section was not found",
            )
            return

        ref_start, ref_end = ref_range
        body_start, body_end = _resolve_body_range(doc, doc_tree, ref_start)
        next_bookmark_id = _max_bookmark_id(doc) + 1

        (
            num_to_target,
            next_bookmark_id,
            entry_bookmarks_added,
            number_bookmarks_added,
            duplicated_numbers,
        ) = _build_reference_targets(
            doc,
            ref_start=ref_start,
            ref_end=ref_end,
            next_bookmark_id=next_bookmark_id,
            auto_number_reference_entries=bool(citation_cfg.auto_number_reference_entries),
        )

        linked_total = 0
        unresolved_numbers: set[int] = set()
        for para_index in range(body_start, body_end + 1):
            if para_index < 0 or para_index >= len(doc.paragraphs):
                continue
            para = doc.paragraphs[para_index]
            if _paragraph_has_field(para):
                continue
            linked, unresolved = _replace_citation_tokens_in_paragraph(para, num_to_target)
            if linked > 0:
                linked_total += linked
                tracker.record(
                    rule_name=self.meta.name,
                    target=f"paragraph #{para_index}",
                    section="body",
                    change_type="field",
                    before="plain citation",
                    after=f"linked citation fields ({linked})",
                    paragraph_index=para_index,
                )
            unresolved_numbers.update(unresolved)

        if entry_bookmarks_added or number_bookmarks_added or linked_total:
            tracker.record(
                rule_name=self.meta.name,
                target="summary",
                section="global",
                change_type="field",
                before=(
                    f"entry_bookmarks={entry_bookmarks_added}, "
                    f"number_bookmarks={number_bookmarks_added}, "
                    f"references={len(num_to_target)}"
                ),
                after=(
                    f"linked={linked_total}, "
                    f"unresolved={len(unresolved_numbers)}, "
                    f"duplicates={len(duplicated_numbers)}"
                ),
                paragraph_index=-1,
            )


def _resolve_reference_range(doc: Document, doc_tree) -> tuple[int, int] | None:
    if doc_tree is not None:
        section = getattr(doc_tree, "get_section", lambda *_: None)("references")
        if section is not None:
            start = int(getattr(section, "start_index", -1))
            end = int(getattr(section, "end_index", -1))
            if start >= 0 and end > start:
                return (start, end - 1)

    return None


def _resolve_body_range(doc: Document, doc_tree, reference_start: int | None) -> tuple[int, int]:
    total = len(doc.paragraphs)
    start = 0
    end = total - 1

    if doc_tree is not None:
        section = getattr(doc_tree, "get_section", lambda *_: None)("body")
        if section is not None:
            start = max(0, int(getattr(section, "start_index", 0)))
            end = min(total - 1, int(getattr(section, "end_index", total)) - 1)

    if isinstance(reference_start, int) and reference_start > 0:
        end = min(end, reference_start - 1)
    return (start, end)


def _paragraph_has_field(para) -> bool:
    return (
        para._element.find(f".//{qn('w:instrText')}") is not None
        or para._element.find(f".//{qn('w:fldChar')}") is not None
    )


def _max_bookmark_id(doc: Document) -> int:
    max_id = 0
    for bookmark in doc.element.body.iter(qn("w:bookmarkStart")):
        try:
            max_id = max(max_id, int(bookmark.get(qn("w:id"), "0")))
        except Exception:
            continue
    return max_id


def _ensure_paragraph_bookmark(para, bookmark_name: str, bookmark_id: int) -> bool:
    for bookmark in para._element.iter(qn("w:bookmarkStart")):
        if bookmark.get(qn("w:name")) == bookmark_name:
            return False

    ppr = para._element.find(qn("w:pPr"))
    insert_idx = list(para._element).index(ppr) + 1 if ppr is not None else 0

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))

    para._element.insert(insert_idx, start)
    para._element.insert(insert_idx + 1, end)
    return True


def _build_reference_targets(
    doc: Document,
    *,
    ref_start: int,
    ref_end: int,
    next_bookmark_id: int,
    auto_number_reference_entries: bool,
) -> tuple[dict[int, str], int, int, int, set[int]]:
    num_to_target: dict[int, str] = {}
    entry_bookmarks_added = 0
    number_bookmarks_added = 0
    duplicated_numbers: set[int] = set()

    for para_index in range(max(0, ref_start), min(len(doc.paragraphs) - 1, ref_end) + 1):
        para = doc.paragraphs[para_index]
        text = (para.text or "").strip()
        match = REFERENCE_ENTRY_RE.match(text)
        if match is None:
            continue
        number = int(match.group("num"))

        entry_bookmark = f"{REF_ENTRY_BOOKMARK_PREFIX}{number}"
        if _ensure_paragraph_bookmark(para, entry_bookmark, next_bookmark_id):
            entry_bookmarks_added += 1
            next_bookmark_id += 1

        target_bookmark = entry_bookmark
        if auto_number_reference_entries:
            number_bookmark = f"{REF_NUM_BOOKMARK_PREFIX}{number}"
            if _ensure_paragraph_bookmark(para, number_bookmark, next_bookmark_id):
                number_bookmarks_added += 1
                next_bookmark_id += 1
            target_bookmark = number_bookmark

        if number in duplicated_numbers:
            continue
        if number in num_to_target:
            duplicated_numbers.add(number)
            num_to_target.pop(number, None)
            continue
        num_to_target[number] = target_bookmark

    return num_to_target, next_bookmark_id, entry_bookmarks_added, number_bookmarks_added, duplicated_numbers


def _make_text_run(text: str, *, superscript: bool = False):
    run = OxmlElement("w:r")
    if superscript:
        rpr = OxmlElement("w:rPr")
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        rpr.append(vert)
        run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(text_el)
    return run


def _make_field_run(field_type: str, *, superscript: bool = False, dirty: bool = False):
    run = OxmlElement("w:r")
    if superscript:
        rpr = OxmlElement("w:rPr")
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        rpr.append(vert)
        run.append(rpr)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), field_type)
    if dirty:
        fld.set(qn("w:dirty"), "true")
    run.append(fld)
    return run


def _make_instr_run(instr_text: str, *, superscript: bool = False):
    run = OxmlElement("w:r")
    if superscript:
        rpr = OxmlElement("w:rPr")
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        rpr.append(vert)
        run.append(rpr)
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = instr_text
    run.append(instr)
    return run


def _build_ref_field_runs(bookmark_name: str, display_text: str) -> list:
    return [
        _make_field_run("begin", superscript=True, dirty=True),
        _make_instr_run(f" REF {bookmark_name} \\h \\* CHARFORMAT \\* MERGEFORMAT ", superscript=True),
        _make_field_run("separate", superscript=True),
        _make_text_run(display_text, superscript=True),
        _make_field_run("end", superscript=True),
    ]


def _replace_citation_tokens_in_paragraph(para, num_to_target: dict[int, str]) -> tuple[int, set[int]]:
    text = para.text or ""
    if not text:
        return 0, set()

    matches = list(CITATION_TOKEN_RE.finditer(text))
    if not matches:
        return 0, set()

    new_nodes = []
    cursor = 0
    linked = 0
    unresolved: set[int] = set()

    for match in matches:
        start, end = match.span()
        if start > cursor:
            new_nodes.append(_make_text_run(text[cursor:start]))

        open_bracket = match.group("open")
        inner = match.group("inner")
        close_bracket = match.group("close")
        parts = list(re.finditer(r"\d{1,4}|[^\d]+", inner))
        numbers = [int(part.group(0)) for part in parts if part.group(0).isdigit()]
        if numbers and all(number in num_to_target for number in numbers):
            new_nodes.append(_make_text_run(open_bracket, superscript=True))
            for part in parts:
                token = part.group(0)
                if token.isdigit():
                    linked += 1
                    new_nodes.extend(_build_ref_field_runs(num_to_target[int(token)], token))
                else:
                    new_nodes.append(_make_text_run(token, superscript=True))
            new_nodes.append(_make_text_run(close_bracket, superscript=True))
        else:
            unresolved.update(number for number in numbers if number not in num_to_target)
            new_nodes.append(_make_text_run(text[start:end]))
        cursor = end

    if cursor < len(text):
        new_nodes.append(_make_text_run(text[cursor:]))

    if linked <= 0:
        return 0, unresolved

    for child in list(para._element):
        if child.tag != qn("w:pPr"):
            para._element.remove(child)
    for node in new_nodes:
        para._element.append(node)
    return linked, unresolved
