"""caption — 图表题注编号、格式与保守自动插入."""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass, field
from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

from src.config.style_semantics import apply_style_special_indent
from src.config.template import StyleConfig
from src.modules.base import BaseModule, ModuleMeta
from src.modules.table.table_format import _is_equation_table
from src.shared.engine.field_builder import build_complex_field
from src.shared.engine.font_resolver import resolve_font
from src.shared.engine.indent_ops import apply_style_config_indents
from src.shared.engine.line_spacing_ops import apply_line_spacing, apply_paragraph_spacing, sync_spacing_ooxml
from src.shared.engine.media_ops import paragraph_has_image
from src.shared.engine.ooxml_ops import find_or_create, qn
from src.shared.engine.run_ops import set_run_east_asian_font
from src.shared.engine.sequence_numbering import (
    build_heading_chapter_ranges,
    parse_chapter_numbering_format,
    resolve_chapter_number,
)
from src.shared.engine.document_scope_runtime import (
    document_scope_allows_paragraph,
)

if TYPE_CHECKING:
    from docx import Document
    from docx.text.paragraph import Paragraph
    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


_RE_FIG_CAPTION = re.compile(
    r"^(?P<prefix>图|Figure|Fig\.?)\s*"
    r"(?P<num>[一二三四五六七八九十百零\d]+(?:[.\-]\d+)*)"
    r"(?P<sep>\s+|[\s\u3000])"
    r"(?P<title>.*)",
    re.IGNORECASE,
)

_RE_TBL_CAPTION = re.compile(
    r"^(?P<prefix>表|Table|Tab\.?)\s*"
    r"(?P<num>[一二三四五六七八九十百零\d]+(?:[.\-]\d+)*)"
    r"(?P<sep>\s+|[\s\u3000])"
    r"(?P<title>.*)",
    re.IGNORECASE,
)

_RE_FIG_CAPTION_CONT_LINE = re.compile(r"^\s*[（(]\s*[A-Za-z]\s*[）)]")

ALIGNMENT_MAP: dict[str, int] = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "distribute": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
}


@dataclass
class CaptionInfo:
    """Detected caption metadata."""

    para_index: int
    kind: str
    prefix: str
    old_number: str
    title: str
    new_number: str = ""


@dataclass
class _FigureAnchor:
    para_index: int
    element: object
    caption_para_index: int | None = None
    caption_para_indices: list[int] = field(default_factory=list)


@dataclass
class _TableAnchor:
    table_index: int
    para_index: int
    element: object
    caption_para_index: int | None = None


class CaptionModule(BaseModule):
    """图表题注模块。

    职责：
    - 扫描已存在的图/表题注
    - 对图片后、表格前的缺失题注做保守自动插入
    - 按章节或全局编号统一题注文本
    - 应用题注段落格式
    """

    meta = ModuleMeta(
        name="caption",
        description="题注编号",
        category="table",
        requires_config=("caption",),
        soft_after=("heading_recognition",),
        soft_consumes=("doc_tree", "heading_map"),
        provides=("caption_counters",),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        caption_cfg = config.caption
        heading_map = context.heading_map or {}
        chapter_ranges = _build_chapter_ranges(
            heading_map,
            len(doc.paragraphs),
            getattr(context, "doc_tree", None),
        )

        existing_captions = _scan_captions(doc)
        body_items = _build_body_items(doc, context)
        figures = _detect_figures(doc, body_items)
        tables = _detect_tables(doc, body_items)

        fig_counter = 0
        tbl_counter = 0
        chapter_fig_counters: dict[int, int] = {}
        chapter_tbl_counters: dict[int, int] = {}
        handled_caption_indices: set[int] = set()
        insert_operations: list[dict] = []
        count = 0
        skipped_count = 0

        for fig in figures:
            chap = _get_chapter_num(fig.para_index, chapter_ranges)
            if _should_skip_caption_chapter_numbering(
                numbering_mode=caption_cfg.numbering_mode,
                numbering_format=caption_cfg.numbering_format,
                chapter_num=chap,
                has_chapter_context=bool(chapter_ranges),
            ):
                if fig.caption_para_index is not None:
                    cap_block_indices = fig.caption_para_indices or [fig.caption_para_index]
                    handled_caption_indices.update(cap_block_indices)
                    fig_style = _resolve_caption_style(config, "figure")
                    for cap_i in cap_block_indices:
                        if 0 <= cap_i < len(doc.paragraphs):
                            _apply_caption_style(doc.paragraphs[cap_i], fig_style)
                    count += len(cap_block_indices)
                    skipped_count += 1
                elif getattr(caption_cfg, "auto_insert", True):
                    skipped_count += 1
                continue
            seq_text = _next_caption_sequence_text(
                kind="figure",
                chapter_num=chap,
                numbering_mode=caption_cfg.numbering_mode,
                numbering_format=caption_cfg.numbering_format,
                fig_counter_ref=chapter_fig_counters,
                tbl_counter_ref=chapter_tbl_counters,
                global_fig_counter=fig_counter,
                global_tbl_counter=tbl_counter,
            )
            if caption_cfg.numbering_mode == "chapter" and chap > 0 and _parse_caption_numbering_format(caption_cfg.numbering_format)[0]:
                chapter_fig_counters[chap] = chapter_fig_counters.get(chap, 0) + 1
            else:
                fig_counter += 1

            if fig.caption_para_index is not None:
                cap_block_indices = fig.caption_para_indices or [fig.caption_para_index]
                handled_caption_indices.update(cap_block_indices)
                para = doc.paragraphs[fig.caption_para_index]
                title = _extract_caption_title(para.text or "", _RE_FIG_CAPTION)
                if getattr(caption_cfg, "format_inserted", False):
                    _rewrite_caption_paragraph_as_field(
                        para,
                        kind="figure",
                        caption_cfg=caption_cfg,
                        sequence_text=seq_text,
                        title=title,
                        config=config,
                    )
                else:
                    _rewrite_caption_paragraph(para, "figure", caption_cfg, seq_text, title, config)
                fig_style = _resolve_caption_style(config, "figure")
                for cap_i in cap_block_indices[1:]:
                    if 0 <= cap_i < len(doc.paragraphs):
                        _apply_caption_style(doc.paragraphs[cap_i], fig_style)
                count += len(cap_block_indices)
            elif getattr(caption_cfg, "auto_insert", True):
                insert_operations.append(
                    {
                        "sort_key": fig.para_index,
                        "anchor": fig.element,
                        "position": "after",
                        "kind": "figure",
                        "number_text": seq_text,
                        "title": str(getattr(caption_cfg, "placeholder", "[待补充]") or "[待补充]").strip(),
                    }
                )
                count += 1

        for tbl in tables:
            chap = _get_chapter_num(tbl.para_index, chapter_ranges)
            if _should_skip_caption_chapter_numbering(
                numbering_mode=caption_cfg.numbering_mode,
                numbering_format=caption_cfg.numbering_format,
                chapter_num=chap,
                has_chapter_context=bool(chapter_ranges),
            ):
                if tbl.caption_para_index is not None:
                    handled_caption_indices.add(tbl.caption_para_index)
                    _apply_caption_style(
                        doc.paragraphs[tbl.caption_para_index],
                        _resolve_caption_style(config, "table"),
                    )
                    count += 1
                    skipped_count += 1
                elif getattr(caption_cfg, "auto_insert", True):
                    skipped_count += 1
                continue
            seq_text = _next_caption_sequence_text(
                kind="table",
                chapter_num=chap,
                numbering_mode=caption_cfg.numbering_mode,
                numbering_format=caption_cfg.numbering_format,
                fig_counter_ref=chapter_fig_counters,
                tbl_counter_ref=chapter_tbl_counters,
                global_fig_counter=fig_counter,
                global_tbl_counter=tbl_counter,
            )
            if caption_cfg.numbering_mode == "chapter" and chap > 0 and _parse_caption_numbering_format(caption_cfg.numbering_format)[0]:
                chapter_tbl_counters[chap] = chapter_tbl_counters.get(chap, 0) + 1
            else:
                tbl_counter += 1

            if tbl.caption_para_index is not None:
                handled_caption_indices.add(tbl.caption_para_index)
                para = doc.paragraphs[tbl.caption_para_index]
                title = _extract_caption_title(para.text or "", _RE_TBL_CAPTION)
                if getattr(caption_cfg, "format_inserted", False):
                    _rewrite_caption_paragraph_as_field(
                        para,
                        kind="table",
                        caption_cfg=caption_cfg,
                        sequence_text=seq_text,
                        title=title,
                        config=config,
                    )
                else:
                    _rewrite_caption_paragraph(para, "table", caption_cfg, seq_text, title, config)
                count += 1
            elif getattr(caption_cfg, "auto_insert", True):
                insert_operations.append(
                    {
                        "sort_key": tbl.para_index,
                        "anchor": tbl.element,
                        "position": "before",
                        "kind": "table",
                        "number_text": seq_text,
                        "title": str(getattr(caption_cfg, "placeholder", "[待补充]") or "[待补充]").strip(),
                    }
                )
                count += 1

        # Orphan captions: keep legacy behavior and renumber after anchor-backed captions.
        for cap in existing_captions:
            if cap.para_index in handled_caption_indices:
                continue
            chap = _get_chapter_num(cap.para_index, chapter_ranges)
            if _should_skip_caption_chapter_numbering(
                numbering_mode=caption_cfg.numbering_mode,
                numbering_format=caption_cfg.numbering_format,
                chapter_num=chap,
                has_chapter_context=bool(chapter_ranges),
            ):
                _apply_caption_style(
                    doc.paragraphs[cap.para_index],
                    _resolve_caption_style(config, cap.kind),
                )
                skipped_count += 1
                count += 1
                continue
            if cap.kind == "figure":
                if caption_cfg.numbering_mode == "chapter" and chap > 0 and _parse_caption_numbering_format(caption_cfg.numbering_format)[0]:
                    chapter_fig_counters[chap] = chapter_fig_counters.get(chap, 0) + 1
                    seq_text = f"{chap}.{chapter_fig_counters[chap]}"
                else:
                    fig_counter += 1
                    seq_text = str(fig_counter)
            else:
                if caption_cfg.numbering_mode == "chapter" and chap > 0 and _parse_caption_numbering_format(caption_cfg.numbering_format)[0]:
                    chapter_tbl_counters[chap] = chapter_tbl_counters.get(chap, 0) + 1
                    seq_text = f"{chap}.{chapter_tbl_counters[chap]}"
                else:
                    tbl_counter += 1
                    seq_text = str(tbl_counter)
            para = doc.paragraphs[cap.para_index]
            if getattr(caption_cfg, "format_inserted", False):
                _rewrite_caption_paragraph_as_field(
                    para,
                    kind=cap.kind,
                    caption_cfg=caption_cfg,
                    sequence_text=seq_text,
                    title=cap.title,
                    config=config,
                )
            else:
                _rewrite_caption_paragraph(para, cap.kind, caption_cfg, seq_text, cap.title, config)
            count += 1

        _apply_insert_operations(doc, config, caption_cfg, insert_operations)

        context.caption_counters = {
            "figure": max(fig_counter, sum(chapter_fig_counters.values())),
            "table": max(tbl_counter, sum(chapter_tbl_counters.values())),
        }

        if count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{count} 个题注",
                section="global",
                change_type="format",
                before="(mixed)",
                after=f"mode={caption_cfg.numbering_mode}",
            )
        if skipped_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{skipped_count} captions",
                section="global",
                change_type="skip",
                before="chapter-aware caption numbering",
                after="skipped due to missing chapter context",
            )


def _resolve_caption_style(config: ResolvedConfig, kind: str) -> StyleConfig:
    specific_key = "figure_caption" if kind == "figure" else "table_caption"
    specific = config.styles.get(specific_key)
    shared = config.styles.get("caption")
    source = specific or shared or config.styles.get("body") or config.styles.get("normal") or StyleConfig()

    style = deepcopy(source)
    if specific is None and shared is None:
        style.alignment = "center"
        style.left_indent_chars = 0.0
        style.left_indent_unit = "chars"
        style.right_indent_chars = 0.0
        style.right_indent_unit = "chars"
        apply_style_special_indent(style, "none", 0, "chars")
    return style


def _caption_prefix(caption_cfg, kind_or_cap) -> str:
    kind = kind_or_cap.kind if isinstance(kind_or_cap, CaptionInfo) else str(kind_or_cap)
    configured = getattr(caption_cfg, "figure_prefix", None) if kind == "figure" else getattr(caption_cfg, "table_prefix", None)
    text = str(configured or "").strip()
    if text:
        return text
    return "图" if kind == "figure" else "表"


def _parse_caption_numbering_format(numbering_format: str | None) -> tuple[bool, str]:
    return parse_chapter_numbering_format(numbering_format)


def _compose_caption_number_text(*, chapter_num: int, sequence_text: str, numbering_mode: str, numbering_format: str | None) -> str:
    include_chapter, separator = _parse_caption_numbering_format(numbering_format)
    if str(numbering_mode or "").strip().lower() != "chapter" or chapter_num <= 0 or not include_chapter:
        return str(sequence_text)
    sequence = str(sequence_text or "").strip()
    if separator:
        suffix = sequence.split(separator, 1)[-1] if separator in sequence else sequence.rsplit(".", 1)[-1]
    else:
        suffix = sequence
    return f"{chapter_num}{separator}{suffix}"


def _build_run_element(text: str):
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    run.append(text_el)
    text_el.text = text
    text_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return run


def _clear_paragraph_content(para: Paragraph) -> None:
    for child in list(para._element):
        if child.tag != qn("w:pPr"):
            para._element.remove(child)


def _append_caption_number_field_runs(
    para: Paragraph,
    *,
    prefix: str,
    chapter_num: int,
    sequence_text: str,
    numbering_mode: str,
    numbering_format: str | None,
    separator: str,
) -> None:
    include_chapter, chapter_sep = _parse_caption_numbering_format(numbering_format)
    para._element.append(_build_run_element(prefix))

    if str(numbering_mode or "").strip().lower() == "chapter" and chapter_num > 0 and include_chapter:
        for elem in build_complex_field(" STYLEREF 1 \\s ", result_text=str(chapter_num)):
            para._element.append(elem)
        if chapter_sep:
            para._element.append(_build_run_element(chapter_sep))
        seq_suffix = str(sequence_text).split(".", 1)[-1] if "." in str(sequence_text) else str(sequence_text)
        for elem in build_complex_field(" SEQ Figure \\* ARABIC \\s 1 ", result_text=seq_suffix):
            para._element.append(elem)
    else:
        seq_result = str(sequence_text)
        for elem in build_complex_field(" SEQ Figure \\* ARABIC ", result_text=seq_result):
            para._element.append(elem)

    para._element.append(_build_run_element(separator))


def _rewrite_caption_paragraph_as_field(
    para: Paragraph,
    *,
    kind: str,
    caption_cfg,
    sequence_text: str,
    title: str,
    config: ResolvedConfig,
) -> None:
    prefix = _caption_prefix(caption_cfg, kind)
    separator = str(getattr(caption_cfg, "separator", "\u3000") or "\u3000")
    chapter_num = int(sequence_text.split(".", 1)[0]) if "." in sequence_text else 0
    seq_name = prefix or ("Figure" if kind == "figure" else "Table")

    _clear_paragraph_content(para)
    para._element.append(_build_run_element(prefix))
    include_chapter, chapter_sep = _parse_caption_numbering_format(caption_cfg.numbering_format)
    if str(caption_cfg.numbering_mode or "").strip().lower() == "chapter" and chapter_num > 0 and include_chapter:
        for elem in build_complex_field(" STYLEREF 1 \\s ", result_text=str(chapter_num)):
            para._element.append(elem)
        if chapter_sep:
            para._element.append(_build_run_element(chapter_sep))
        seq_suffix = str(sequence_text).split(".", 1)[-1] if "." in str(sequence_text) else str(sequence_text)
        for elem in build_complex_field(f" SEQ {seq_name} \\* ARABIC \\s 1 ", result_text=seq_suffix):
            para._element.append(elem)
    else:
        for elem in build_complex_field(f" SEQ {seq_name} \\* ARABIC ", result_text=str(sequence_text)):
            para._element.append(elem)
    para._element.append(_build_run_element(separator))
    para._element.append(_build_run_element(str(title or "").strip()))
    _apply_caption_style(para, _resolve_caption_style(config, kind))


def _apply_caption_style(para: Paragraph, style_config: StyleConfig) -> None:
    pf = para.paragraph_format
    align_key = str(getattr(style_config, "alignment", "") or "").strip().lower()
    pf.alignment = ALIGNMENT_MAP.get(align_key, WD_ALIGN_PARAGRAPH.CENTER)

    apply_paragraph_spacing(pf, style_config, para._element)
    apply_line_spacing(pf, style_config.line_spacing_type, style_config.line_spacing_pt)
    sync_spacing_ooxml(
        para._element,
        space_before_pt=style_config.space_before_pt,
        space_before_unit=getattr(style_config, "space_before_unit", "pt"),
        space_after_pt=style_config.space_after_pt,
        space_after_unit=getattr(style_config, "space_after_unit", "pt"),
        line_spacing_type=style_config.line_spacing_type,
        line_spacing_value=style_config.line_spacing_pt,
    )
    apply_style_config_indents(pf, para._element, style_config, size_pt=style_config.size_pt)

    for run in para.runs:
        _apply_run_style(run, style_config)

    p_pr = para._element.find(qn("w:pPr"))
    if p_pr is not None:
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            p_pr.remove(num_pr)


def _apply_run_style(run, style_config: StyleConfig) -> None:
    r_pr = find_or_create(run._element, "w:rPr")
    if style_config.font_en:
        resolved_en = resolve_font(style_config.font_en, lang="en")
        r_fonts = find_or_create(r_pr, "w:rFonts")
        r_fonts.set(qn("w:ascii"), resolved_en)
        r_fonts.set(qn("w:hAnsi"), resolved_en)
    if style_config.font_cn:
        set_run_east_asian_font(run, style_config.font_cn)
    if style_config.size_pt:
        half_points = str(int(round(float(style_config.size_pt) * 2)))
        find_or_create(r_pr, "w:sz").set(qn("w:val"), half_points)
        find_or_create(r_pr, "w:szCs").set(qn("w:val"), half_points)
    _set_toggle(r_pr, "w:b", bool(style_config.bold))
    _set_toggle(r_pr, "w:bCs", bool(style_config.bold))
    _set_toggle(r_pr, "w:i", bool(style_config.italic))
    _set_toggle(r_pr, "w:iCs", bool(style_config.italic))


def _set_toggle(parent, tag: str, enabled: bool) -> None:
    element = find_or_create(parent, tag)
    if enabled:
        element.attrib.pop(qn("w:val"), None)
    else:
        element.set(qn("w:val"), "0")


def _scan_captions(doc: Document) -> list[CaptionInfo]:
    captions: list[CaptionInfo] = []
    for i, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if not text:
            continue

        fig_match = _RE_FIG_CAPTION.match(text)
        if fig_match:
            captions.append(
                CaptionInfo(
                    para_index=i,
                    kind="figure",
                    prefix=fig_match.group("prefix"),
                    old_number=fig_match.group("num"),
                    title=fig_match.group("title").strip(),
                )
            )
            continue

        tbl_match = _RE_TBL_CAPTION.match(text)
        if tbl_match:
            captions.append(
                CaptionInfo(
                    para_index=i,
                    kind="table",
                    prefix=tbl_match.group("prefix"),
                    old_number=tbl_match.group("num"),
                    title=tbl_match.group("title").strip(),
                )
            )

    return captions


def _build_body_items(doc: Document, context: PipelineContext) -> list[tuple[str, int, object, int]]:
    body = doc.element.body
    body_items: list[tuple[str, int, object, int]] = []
    para_idx = 0
    table_idx = 0
    last_para_idx = -1
    doc_tree = getattr(context, "doc_tree", None)

    def section_ok(index: int) -> bool:
        section = (
            doc_tree.get_section_for_paragraph(index)
            if doc_tree is not None
            else "body"
        )
        return (
            document_scope_allows_paragraph(context, index)
            and section
            not in {"cover", "toc", "references", "abstract_cn", "abstract_en"}
        )

    for child in body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            last_para_idx = para_idx
            if section_ok(para_idx):
                body_items.append(("para", para_idx, child, para_idx))
            para_idx += 1
        elif tag == "tbl":
            if last_para_idx >= 0 and section_ok(last_para_idx):
                body_items.append(("table", table_idx, child, last_para_idx))
            table_idx += 1
    return body_items


def _detect_figures(doc: Document, body_items: list[tuple[str, int, object, int]]) -> list[_FigureAnchor]:
    anchors: list[_FigureAnchor] = []
    for idx, (item_type, para_index, element, _virtual_para) in enumerate(body_items):
        if item_type != "para":
            continue
        para = doc.paragraphs[para_index]
        if not paragraph_has_image(para):
            continue

        caption_para_index = None
        caption_para_indices: list[int] = []
        for probe in range(idx + 1, len(body_items)):
            next_type, next_para_index, _next_element, _ = body_items[probe]
            if next_type != "para":
                break
            next_text = (doc.paragraphs[next_para_index].text or "").strip()
            if not next_text:
                continue
            if _RE_FIG_CAPTION.match(next_text):
                caption_para_index = next_para_index
                caption_para_indices = [next_para_index]
                caption_para_indices.extend(
                    _collect_figure_caption_continuations(doc, body_items, probe)
                )
            break

        anchors.append(
            _FigureAnchor(
                para_index=para_index,
                element=element,
                caption_para_index=caption_para_index,
                caption_para_indices=caption_para_indices,
            )
        )
    return anchors


def _is_figure_caption_continuation(text: str) -> bool:
    raw = (text or "").strip()
    if not raw or len(raw) > 180:
        return False
    return bool(_RE_FIG_CAPTION_CONT_LINE.match(raw))


def _collect_figure_caption_continuations(
    doc: Document,
    body_items: list[tuple[str, int, object, int]],
    caption_pos: int,
) -> list[int]:
    indices: list[int] = []
    for probe in range(caption_pos + 1, len(body_items)):
        item_type, para_index, _element, _virtual_para = body_items[probe]
        if item_type != "para":
            break
        text = (doc.paragraphs[para_index].text or "").strip()
        if not text:
            break
        if not _is_figure_caption_continuation(text):
            break
        indices.append(para_index)
    return indices


def _detect_tables(doc: Document, body_items: list[tuple[str, int, object, int]]) -> list[_TableAnchor]:
    anchors: list[_TableAnchor] = []
    for idx, (item_type, table_index, element, virtual_para_index) in enumerate(body_items):
        if item_type != "table":
            continue
        table = doc.tables[table_index]
        if _is_equation_table(table):
            continue

        caption_para_index = None
        for probe in range(idx - 1, -1, -1):
            prev_type, prev_index, _prev_element, _ = body_items[probe]
            if prev_type == "table":
                break
            prev_text = (doc.paragraphs[prev_index].text or "").strip()
            if not prev_text:
                continue
            if _RE_TBL_CAPTION.match(prev_text):
                caption_para_index = prev_index
            break

        anchors.append(
            _TableAnchor(
                table_index=table_index,
                para_index=virtual_para_index,
                element=element,
                caption_para_index=caption_para_index,
            )
        )
    return anchors


def _next_caption_sequence_text(
    *,
    kind: str,
    chapter_num: int,
    numbering_mode: str,
    numbering_format: str,
    fig_counter_ref: dict[int, int],
    tbl_counter_ref: dict[int, int],
    global_fig_counter: int,
    global_tbl_counter: int,
) -> str:
    include_chapter, separator = _parse_caption_numbering_format(numbering_format)
    if str(numbering_mode or "").strip().lower() == "chapter" and chapter_num > 0 and include_chapter:
        if kind == "figure":
            seq = fig_counter_ref.get(chapter_num, 0) + 1
        else:
            seq = tbl_counter_ref.get(chapter_num, 0) + 1
        return f"{chapter_num}.{seq}" if separator != "" else f"{chapter_num}{seq}"
    if kind == "figure":
        return str(global_fig_counter + 1)
    return str(global_tbl_counter + 1)


def _should_skip_caption_chapter_numbering(
    *,
    numbering_mode: str,
    numbering_format: str | None,
    chapter_num: int,
    has_chapter_context: bool,
) -> bool:
    include_chapter, _separator = _parse_caption_numbering_format(numbering_format)
    return (
        str(numbering_mode or "").strip().lower() == "chapter"
        and include_chapter
        and has_chapter_context
        and chapter_num <= 0
    )


def _extract_caption_title(text: str, pattern: re.Pattern[str]) -> str:
    match = pattern.match((text or "").strip())
    if match:
        return match.group("title").strip()
    return (text or "").strip()


def _rewrite_caption_paragraph(
    para: Paragraph,
    kind: str,
    caption_cfg,
    sequence_text: str,
    title: str,
    config: ResolvedConfig,
) -> None:
    prefix = _caption_prefix(caption_cfg, kind)
    separator = str(getattr(caption_cfg, "separator", "\u3000") or "\u3000")
    number_text = _compose_caption_number_text(
        chapter_num=int(sequence_text.split(".", 1)[0]) if "." in sequence_text else 0,
        sequence_text=sequence_text,
        numbering_mode=caption_cfg.numbering_mode,
        numbering_format=caption_cfg.numbering_format,
    )
    _replace_caption_text(para, f"{prefix}{number_text}{separator}{title}".strip())
    _apply_caption_style(para, _resolve_caption_style(config, kind))


def _insert_caption_element(anchor, position: str):
    p = OxmlElement("w:p")
    if position == "after":
        anchor.addnext(p)
    else:
        anchor.addprevious(p)
    return p


def _apply_insert_operations(doc: Document, config: ResolvedConfig, caption_cfg, operations: list[dict]) -> None:
    from docx.text.paragraph import Paragraph

    operations.sort(key=lambda item: item["sort_key"], reverse=True)
    for operation in operations:
        inserted = _insert_caption_element(operation["anchor"], operation["position"])
        inserted_para = Paragraph(inserted, doc._body)
        if getattr(caption_cfg, "format_inserted", False):
            _rewrite_caption_paragraph_as_field(
                inserted_para,
                kind=operation["kind"],
                caption_cfg=caption_cfg,
                sequence_text=operation["number_text"],
                title=operation["title"],
                config=config,
            )
        else:
            _rewrite_caption_paragraph(
                inserted_para,
                operation["kind"],
                caption_cfg,
                operation["number_text"],
                operation["title"],
                config,
            )


def _build_chapter_ranges(
    heading_map: dict[int, int],
    total_paragraphs: int,
    doc_tree=None,
) -> list[tuple[int, int, int]]:
    return build_heading_chapter_ranges(
        heading_map,
        total_paragraphs,
        doc_tree=doc_tree,
    )


def _get_chapter_num(para_index: int, chapter_ranges) -> int:
    return resolve_chapter_number(para_index, chapter_ranges)


def _replace_caption_text(para: Paragraph, new_text: str) -> None:
    if not para.runs:
        para.add_run(new_text)
        return
    para.runs[0].text = new_text
    for run in para.runs[1:]:
        run._element.getparent().remove(run._element)
