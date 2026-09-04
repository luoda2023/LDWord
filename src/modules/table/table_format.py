"""Table formatting module."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from lxml import etree

from src.formula_core.normalize import (
    looks_like_bibliographic_reference_text,
    looks_like_caption_text,
    looks_like_formula_text,
)
from src.modules.base import BaseModule, ModuleMeta
from src.config.feature_configs import normalize_table_smart_levels
from src.config.table_style_presets import color_palette, color_variant
from src.shared.engine.font_resolver import resolve_font
from src.shared.engine.ooxml_ops import find_or_create, qn
from src.shared.engine.paragraph_iter import iter_table_cells
from src.shared.engine.run_ops import set_run_east_asian_font
from src.shared.engine.document_scope_runtime import (
    document_scope_allows_paragraph,
)
from src.shared.engine.table_builder import (
    clear_cell_border_overrides,
    clear_repeat_header_row,
    set_repeat_header_row,
    set_table_borders,
)

if TYPE_CHECKING:
    from docx import Document
    from docx.table import Table
    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


_DEFAULT_PAGE_WIDTH_TWIPS = 11906
_MIN_COL_WIDTH = 850
_CELL_PADDING = 216
_CJK_CHAR_W = 210
_ASCII_CHAR_W = 115
_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_O_NS = "urn:schemas-microsoft-com:office:office"

_RE_EQUATION_NUMBER = re.compile(
    r"^\s*[\(\[（【]?\s*\d+(?:\s*[.\-:—–/]\s*\d+)?\s*[\)\]）】]?\s*$"
)
_RE_EQUATION_OLE = re.compile(r"(?:equation|mathtype|eqn)", re.IGNORECASE)
_RE_EQUATION_SOURCE = re.compile(
    r"(?:\$\$.+?\$\$|(?<!\$)\$[^$\r\n]+?\$(?!\$)|\\\(.+?\\\)|\\\[.+?\\\])"
)
_RE_EQUATION_NUMBER_TAIL = re.compile(
    r"[\(\[（【]\s*\d+(?:\s*[.\-:—–/]\s*\d+)?\s*[\)\]）】]\s*$"
)

_RE_EQUATION_ARROW = re.compile(r"(?:→|←|↔|⇌|->|<-|<->)")
_RE_CJK_TEXT = re.compile(r"[\u4e00-\u9fff]")
_RE_PROSE_PUNCT = re.compile(r"[，。；：！？、]")
_RE_OMML_FORMULA_ANCHOR = re.compile(
    r"(=|≈|≠|≤|≥|∑|∫|∏|√|→|←|↔|⇌|±|∂)"
)
_RE_COMPACT_FORMULA_TEXT = re.compile(
    r"^[A-Za-z\u0391-\u03a9\u03b1-\u03c90-9\s()+\-*/^_\\{}\[\].,]+$"
)
_RE_COMPACT_PLUS_TIMES_EXPR = re.compile(
    r"[A-Za-z\u0391-\u03a9\u03b1-\u03c90-9\)\]}]\s*[+*]\s*"
    r"[A-Za-z\u0391-\u03a9\u03b1-\u03c90-9\(\[{\\]"
)
_RE_COMPACT_SLASH_EXPR = re.compile(
    r"[A-Za-z\u0391-\u03a9\u03b1-\u03c90-9\)\]}]\s*/\s*"
    r"[A-Za-z\u0391-\u03a9\u03b1-\u03c90-9\(\[{\\]"
)
_MAX_UNMARKED_EQUATION_TABLE_ROWS = 20
_RE_HEADER_TRAILING_UNIT = re.compile(r"^(?P<head>.+?)\s*\((?P<unit>[^()]{1,80})\)\s*$")
_RE_TRAILING_PAREN_CHUNK = re.compile(r"^(?P<head>.+?)(?P<paren>[（(][^()（）]{1,80}[）)])\s*$")
_TABLE_ALIGNMENTS = frozenset({"left", "center", "right"})


class TableFormatModule(BaseModule):
    meta = ModuleMeta(
        name="table_format",
        description="Table formatting",
        category="table",
        requires_config=("table",),
        soft_consumes=("doc_tree",),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        tbl_cfg = config.table
        avail_w = _available_width(config)

        layout_mode = str(tbl_cfg.layout_mode or "smart").strip().lower()
        if layout_mode not in {"smart", "compact", "full", "keep"}:
            layout_mode = "smart"
        applies_layout = layout_mode != "keep"
        table_alignment = _normalize_table_alignment(getattr(tbl_cfg, "table_alignment", "center"))

        smart_levels = normalize_table_smart_levels(tbl_cfg.smart_levels)

        def _is_table_in_scope(pos: int) -> bool:
            return document_scope_allows_paragraph(context, pos)

        tbl_para_pos = top_level_table_anchor_positions(doc)

        smart_width_plan: dict[int, int] = {}
        if layout_mode == "smart":
            raw_targets: list[tuple[int, int]] = []
            for index, table in enumerate(doc.tables):
                pos = tbl_para_pos[index] if index < len(tbl_para_pos) else -1
                if not _is_table_in_scope(pos):
                    continue
                tbl_el = table._element
                if _table_has_drawing_or_object(tbl_el):
                    continue
                if _is_equation_table(table):
                    continue

                grid_widths = _get_grid_cols(tbl_el)
                num_cols = len(grid_widths) or _count_cols(tbl_el)
                if num_cols <= 0:
                    continue

                first_min_ws = _first_row_min_widths(tbl_el, num_cols)
                body_min_ws = _non_first_row_min_widths(tbl_el, num_cols)
                raw_targets.append(
                    (
                        index,
                        _target_normal_table_width(
                            avail_w,
                            num_cols,
                            first_min_ws,
                            body_min_ws,
                        ),
                    )
                )

            smart_width_plan = _build_smart_width_plan(
                raw_targets,
                avail_w,
                max_levels=smart_levels,
            )

        count = 0
        header_unit_break_count = 0
        body_paren_break_count = 0
        rich_content_table_skip_count = 0

        for index, table in enumerate(doc.tables):
            pos = tbl_para_pos[index] if index < len(tbl_para_pos) else -1
            if not _is_table_in_scope(pos):
                continue

            tbl_el = table._element
            if _table_has_drawing_or_object(tbl_el):
                rich_content_table_skip_count += 1
                continue
            if _is_equation_table(table):
                continue

            if applies_layout:
                _normalize_first_row_span_pattern(tbl_el)

                grid_widths = _get_grid_cols(tbl_el)
                num_cols = len(grid_widths) or _count_cols(tbl_el)
                if num_cols <= 0:
                    continue

                weights = _analyze_col_weights(tbl_el, num_cols)
                first_min_ws = _first_row_min_widths(tbl_el, num_cols)
                body_min_ws = _non_first_row_min_widths(tbl_el, num_cols)

                table_w = avail_w
                if layout_mode == "compact":
                    table_w = _target_normal_table_width(
                        avail_w,
                        num_cols,
                        first_min_ws,
                        body_min_ws,
                    )
                elif layout_mode == "smart":
                    table_w = smart_width_plan.get(
                        index,
                        _target_normal_table_width(
                            avail_w,
                            num_cols,
                            first_min_ws,
                            body_min_ws,
                        ),
                    )

                widths = _distribute_widths(weights, table_w, first_min_ws)
                widths = _second_pass_rebalance_first_row(
                    tbl_el,
                    widths,
                    weights,
                    body_min_widths=body_min_ws,
                )

                unit_breaks = _normalize_first_row_unit_breaks(tbl_el, widths)
                body_breaks = _normalize_body_paren_breaks(tbl_el, widths)
                header_unit_break_count += unit_breaks
                body_paren_break_count += body_breaks

                if unit_breaks or body_breaks:
                    weights = _analyze_col_weights(tbl_el, num_cols)
                    first_min_ws = _first_row_min_widths(tbl_el, num_cols)
                    body_min_ws = _non_first_row_min_widths(tbl_el, num_cols)
                    widths = _distribute_widths(weights, table_w, first_min_ws)
                    widths = _second_pass_rebalance_first_row(
                        tbl_el,
                        widths,
                        weights,
                        body_min_widths=body_min_ws,
                    )

                _set_table_width(tbl_el, table_w)
                _set_grid_cols(tbl_el, widths)
                _update_cell_widths(tbl_el, widths)
                _clear_row_width_exceptions(tbl_el)
                self._set_fixed_layout(tbl_el)

            if table_alignment is not None:
                _set_table_alignment(tbl_el, table_alignment)

            border_mode = str(tbl_cfg.border_mode or "three_line").strip().lower()
            if border_mode == "three_line":
                _apply_three_line_borders(tbl_el, tbl_cfg)
            elif border_mode == "full_grid":
                set_table_borders(
                    table,
                    mode="full_grid",
                    width_pt=float(tbl_cfg.border_width_pt or 0.5),
                    header_width_pt=float(tbl_cfg.three_line_header_width_pt or 1.0),
                    bottom_width_pt=float(tbl_cfg.three_line_bottom_width_pt or 0.5),
                )
            elif border_mode == "color_table":
                _apply_color_table_style(table, tbl_cfg)
            elif border_mode == "none":
                _apply_no_table_borders(tbl_el)
            elif border_mode not in {"keep", ""}:
                set_table_borders(
                    table,
                    mode=border_mode,
                    width_pt=float(tbl_cfg.border_width_pt or 0.5),
                    header_width_pt=float(tbl_cfg.three_line_header_width_pt or 1.0),
                    bottom_width_pt=float(tbl_cfg.three_line_bottom_width_pt or 0.5),
                )

            if tbl_cfg.repeat_header:
                set_repeat_header_row(table, row=0)
            else:
                clear_repeat_header_row(table, row=0)

            _format_table_cells(table, tbl_cfg)

            count += 1

        if count:
            layout_desc = (
                f"{layout_mode}(levels={smart_levels})" if layout_mode == "smart" else layout_mode
            )
            tracker.record(
                rule_name=self.meta.name,
                target=f"{count} 个表格",
                section="global",
                change_type="format",
                before="(mixed)",
                after=f"border={tbl_cfg.border_mode}, layout={layout_desc}",
            )
        if header_unit_break_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{header_unit_break_count} 个表头单元格",
                section="body",
                change_type="format",
                before="unit kept inline",
                after="normalized to title and unit line break",
            )
        if body_paren_break_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{body_paren_break_count} 个正文单元格",
                section="body",
                change_type="format",
                before="trailing parenthesis kept inline",
                after="normalized to body line break",
            )
        if rich_content_table_skip_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{rich_content_table_skip_count} 个表格",
                section="body",
                change_type="skip",
                before="regular table formatting",
                after="skipped rich-content tables",
            )

    @staticmethod
    def _set_fixed_layout(tbl_el) -> None:
        tbl_pr = find_or_create(tbl_el, "w:tblPr")
        layout = find_or_create(tbl_pr, "w:tblLayout")
        layout.set(qn("w:type"), "fixed")


def _available_width(config) -> int:
    """Return available table width in twips."""
    ps = config.page_setup
    margin = ps.margin
    left_cm = margin.left_cm
    right_cm = margin.right_cm
    gutter_cm = ps.gutter_cm

    try:
        from src.modules.basic.page_setup import PAPER_SIZES

        paper_key = (ps.paper_size or "A4").strip().upper()
        paper_w, paper_h = PAPER_SIZES.get(paper_key, (21.0, 29.7))
        orientation = getattr(ps, "orientation", "portrait")
        page_w_cm = paper_h if orientation == "landscape" else paper_w
        page_w_twips = _cm_to_twips(page_w_cm)
    except Exception:
        page_w_twips = _DEFAULT_PAGE_WIDTH_TWIPS

    return page_w_twips - _cm_to_twips(left_cm) - _cm_to_twips(right_cm) - _cm_to_twips(gutter_cm)


def _cm_to_twips(cm: float) -> int:
    return int(cm * 567)


def top_level_table_anchor_positions(doc: Document) -> list[int]:
    """Map each top-level table to the nearest paragraph index."""
    positions: list[int] = []
    pending_leading_tables = 0
    para_index = -1

    for child in doc.element.body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            para_index += 1
            if pending_leading_tables:
                positions.extend([para_index] * pending_leading_tables)
                pending_leading_tables = 0
        elif tag == "tbl":
            if para_index >= 0:
                positions.append(para_index)
            else:
                pending_leading_tables += 1

    if pending_leading_tables:
        positions.extend([-1] * pending_leading_tables)

    return positions


def _count_cols(tbl_el) -> int:
    for tr in tbl_el.findall(qn("w:tr")):
        count = 0
        for tc in tr.findall(qn("w:tc")):
            span = 1
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is not None:
                grid_span = tc_pr.find(qn("w:gridSpan"))
                if grid_span is not None:
                    span = int(grid_span.get(qn("w:val"), "1"))
            count += span
        if count > 0:
            return count
    return 0


def _get_grid_cols(tbl_el) -> list[int]:
    grid = tbl_el.find(qn("w:tblGrid"))
    if grid is None:
        return []
    return [int(col.get(qn("w:w"), "0")) for col in grid.findall(qn("w:gridCol"))]


def _set_grid_cols(tbl_el, widths: list[int]) -> None:
    grid = tbl_el.find(qn("w:tblGrid"))
    if grid is None:
        tbl_pr = tbl_el.find(qn("w:tblPr"))
        index = list(tbl_el).index(tbl_pr) + 1 if tbl_pr is not None else 0
        grid = etree.Element(qn("w:tblGrid"))
        tbl_el.insert(index, grid)
    else:
        for old in list(grid.findall(qn("w:gridCol"))):
            grid.remove(old)

    for width in widths:
        col = etree.SubElement(grid, qn("w:gridCol"))
        col.set(qn("w:w"), str(width))


def _set_table_width(tbl_el, total_w: int) -> None:
    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    tbl_w = find_or_create(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(total_w))
    tbl_w.set(qn("w:type"), "dxa")


def _set_table_alignment(tbl_el, align: str) -> None:
    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    jc = find_or_create(tbl_pr, "w:jc")
    jc.set(qn("w:val"), align)


def _normalize_table_alignment(value: str | None) -> str | None:
    normalized = str(value or "").strip().lower()
    if normalized in _TABLE_ALIGNMENTS:
        return normalized
    return None


def _table_has_drawing_or_object(tbl_el) -> bool:
    for tc in tbl_el.iter(qn("w:tc")):
        if tc.findall(f".//{qn('w:object')}"):
            return True
        if tc.findall(f".//{qn('w:drawing')}"):
            return True
        if tc.findall(f".//{qn('w:pict')}"):
            return True
    return False


def _analyze_col_weights(tbl_el, num_cols: int) -> list[float]:
    weights = [0.0] * num_cols

    for tr in tbl_el.findall(qn("w:tr")):
        grid_index = 0
        tr_pr = tr.find(qn("w:trPr"))
        if tr_pr is not None:
            grid_before = tr_pr.find(qn("w:gridBefore"))
            if grid_before is not None:
                grid_index = int(grid_before.get(qn("w:val"), "0"))

        for tc in tr.findall(qn("w:tc")):
            span = 1
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is not None:
                grid_span = tc_pr.find(qn("w:gridSpan"))
                if grid_span is not None:
                    span = int(grid_span.get(qn("w:val"), "1"))

            weight = sum(2.0 if ord(ch) > 0x7F else 1.0 for ch in _get_cell_text(tc))
            per_col = weight / span if span > 0 else 0.0
            for offset in range(span):
                col_index = grid_index + offset
                if col_index < num_cols:
                    weights[col_index] = max(weights[col_index], per_col)

            grid_index += span

    return weights


def _text_width_twips(text: str) -> int:
    width = 0
    for ch in text:
        width += _CJK_CHAR_W if ord(ch) > 0x7F else _ASCII_CHAR_W
    return width + _CELL_PADDING


def _iter_cell_paragraphs(tc):
    return tc.iter(qn("w:p"))


def _extract_cell_lines(tc) -> list[str]:
    lines: list[str] = []
    for para in _iter_cell_paragraphs(tc):
        current: list[str] = []
        for node in para.iter():
            tag = node.tag.split("}")[-1] if "}" in node.tag else node.tag
            if tag == "t" and node.text:
                parts = node.text.replace("\r", "\n").split("\n")
                for index, part in enumerate(parts):
                    if index:
                        lines.append("".join(current))
                        current = []
                    current.append(part)
            elif tag == "br":
                lines.append("".join(current))
                current = []
        lines.append("".join(current))

    if lines:
        return lines
    return [(_get_cell_text(tc) or "")]


def _first_row_min_widths(tbl_el, num_cols: int) -> list[int]:
    mins = [_MIN_COL_WIDTH] * num_cols
    rows = tbl_el.findall(qn("w:tr"))
    if not rows:
        return mins

    tr = rows[0]
    grid_index = 0
    tr_pr = tr.find(qn("w:trPr"))
    if tr_pr is not None:
        grid_before = tr_pr.find(qn("w:gridBefore"))
        if grid_before is not None:
            grid_index = int(grid_before.get(qn("w:val"), "0"))

    for tc in tr.findall(qn("w:tc")):
        span = 1
        tc_pr = tc.find(qn("w:tcPr"))
        if tc_pr is not None:
            grid_span = tc_pr.find(qn("w:gridSpan"))
            if grid_span is not None:
                span = int(grid_span.get(qn("w:val"), "1"))

        lines = _extract_cell_lines(tc)
        non_empty = [line for line in lines if line.strip()]
        if not non_empty:
            need = 0
        elif len(non_empty) >= 2:
            need = max(_text_width_twips(line) for line in non_empty)
        else:
            need = _text_width_twips(non_empty[0])

        per_col = need // span if span > 0 else need
        for offset in range(span):
            col_index = grid_index + offset
            if col_index < num_cols:
                mins[col_index] = max(mins[col_index], per_col)

        grid_index += span

    return mins


def _non_first_row_min_widths(tbl_el, num_cols: int) -> list[int]:
    mins = [0] * num_cols
    rows = tbl_el.findall(qn("w:tr"))
    if len(rows) <= 1:
        return mins

    for tr in rows[1:]:
        grid_index = 0
        tr_pr = tr.find(qn("w:trPr"))
        if tr_pr is not None:
            grid_before = tr_pr.find(qn("w:gridBefore"))
            if grid_before is not None:
                grid_index = int(grid_before.get(qn("w:val"), "0"))

        for tc in tr.findall(qn("w:tc")):
            span = 1
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is not None:
                grid_span = tc_pr.find(qn("w:gridSpan"))
                if grid_span is not None:
                    span = int(grid_span.get(qn("w:val"), "1"))

            lines = [line for line in _extract_cell_lines(tc) if line.strip()]
            need = max((_text_width_twips(line) for line in lines), default=0)
            per_col = need // span if span > 0 else need
            for offset in range(span):
                col_index = grid_index + offset
                if col_index < num_cols:
                    mins[col_index] = max(mins[col_index], per_col)

            grid_index += span

    return mins


def _distribute_widths(
    weights: list[float],
    total: int,
    min_widths: list[int] | None = None,
) -> list[int]:
    if not weights:
        return []

    normalized = [max(1.0, weight) for weight in weights]
    floors = [_MIN_COL_WIDTH] * len(normalized)
    if min_widths:
        for index, value in enumerate(min_widths[: len(floors)]):
            floors[index] = max(floors[index], value)

    floor_sum = sum(floors)
    if floor_sum >= total:
        scale = total / floor_sum if floor_sum else 1.0
        result = [max(1, int(width * scale)) for width in floors]
        diff = total - sum(result)
        for index in range(abs(diff)):
            result[index % len(result)] += 1 if diff > 0 else -1
        return result

    remaining = total - floor_sum
    total_weight = sum(normalized) or 1.0
    result = [
        int(floors[index] + remaining * normalized[index] / total_weight)
        for index in range(len(normalized))
    ]

    diff = total - sum(result)
    for index in range(abs(diff)):
        result[index % len(result)] += 1 if diff > 0 else -1

    return result


def _row_spans(tr) -> list[int]:
    spans: list[int] = []
    for tc in tr.findall(qn("w:tc")):
        span = 1
        tc_pr = tc.find(qn("w:tcPr"))
        if tc_pr is not None:
            grid_span = tc_pr.find(qn("w:gridSpan"))
            if grid_span is not None:
                span = int(grid_span.get(qn("w:val"), "1"))
        spans.append(span)
    return spans


def _set_tc_grid_span(tc, span: int) -> None:
    tc_pr = tc.find(qn("w:tcPr"))
    if tc_pr is None:
        tc_pr = etree.SubElement(tc, qn("w:tcPr"))
        tc.insert(0, tc_pr)
    grid_span = tc_pr.find(qn("w:gridSpan"))
    if span <= 1:
        if grid_span is not None:
            tc_pr.remove(grid_span)
        return
    if grid_span is None:
        grid_span = etree.SubElement(tc_pr, qn("w:gridSpan"))
    grid_span.set(qn("w:val"), str(span))


def _normalize_first_row_span_pattern(tbl_el) -> bool:
    rows = tbl_el.findall(qn("w:tr"))
    if len(rows) < 2:
        return False

    first = rows[0]
    first_spans = _row_spans(first)
    if not first_spans:
        return False

    first_total = sum(first_spans)
    first_len = len(first_spans)
    patterns: dict[tuple[int, ...], int] = {}

    for tr in rows[1:]:
        spans = _row_spans(tr)
        if len(spans) != first_len or sum(spans) != first_total:
            continue
        key = tuple(spans)
        patterns[key] = patterns.get(key, 0) + 1

    if not patterns:
        return False

    reference = list(max(patterns.items(), key=lambda item: item[1])[0])
    if reference == first_spans:
        return False

    diff = [index for index, pair in enumerate(zip(first_spans, reference)) if pair[0] != pair[1]]
    if len(diff) != 2 or diff[1] != diff[0] + 1:
        return False

    left, right = diff
    if first_spans[left] != reference[right] or first_spans[right] != reference[left]:
        return False
    if 1 not in (first_spans[left], first_spans[right]):
        return False

    cells = first.findall(qn("w:tc"))
    _set_tc_grid_span(cells[left], reference[left])
    _set_tc_grid_span(cells[right], reference[right])
    return True


def _first_row_cell_specs(tbl_el) -> list[dict[str, int | bool]]:
    rows = tbl_el.findall(qn("w:tr"))
    if not rows:
        return []

    tr = rows[0]
    specs: list[dict[str, int | bool]] = []
    grid_index = 0
    tr_pr = tr.find(qn("w:trPr"))
    if tr_pr is not None:
        grid_before = tr_pr.find(qn("w:gridBefore"))
        if grid_before is not None:
            grid_index = int(grid_before.get(qn("w:val"), "0"))

    for tc in tr.findall(qn("w:tc")):
        span = 1
        tc_pr = tc.find(qn("w:tcPr"))
        if tc_pr is not None:
            grid_span = tc_pr.find(qn("w:gridSpan"))
            if grid_span is not None:
                span = int(grid_span.get(qn("w:val"), "1"))

        lines = _extract_cell_lines(tc)
        non_empty = [line for line in lines if line.strip()]
        multiline = len(non_empty) >= 2
        if not non_empty:
            need = 0
        elif multiline:
            need = max(_text_width_twips(line) for line in non_empty)
        else:
            need = _text_width_twips(non_empty[0])

        specs.append(
            {
                "start": grid_index,
                "span": span,
                "need": need,
                "multiline": multiline,
            }
        )
        grid_index += span

    return specs


def _second_pass_rebalance_first_row(
    tbl_el,
    widths: list[int],
    weights: list[float],
    body_min_widths: list[int] | None = None,
) -> list[int]:
    if not widths:
        return widths

    specs = _first_row_cell_specs(tbl_el)
    if not specs:
        return widths

    donor_cap = [0] * len(widths)
    recv_need = [0] * len(widths)
    multiline_mask = [False] * len(widths)

    if body_min_widths:
        for index, need in enumerate(body_min_widths[: len(widths)]):
            deficit = need - widths[index]
            if deficit > 40:
                recv_need[index] = deficit

    for spec in specs:
        start = int(spec["start"])
        span = int(spec["span"])
        end = min(len(widths), start + span)
        if start >= end:
            continue

        actual = sum(widths[start:end])
        need = int(spec["need"])

        if spec["multiline"]:
            for col_index in range(start, end):
                multiline_mask[col_index] = True
            keep = max(need, _MIN_COL_WIDTH * span)
            surplus = actual - keep
        elif sum(recv_need) > 0:
            keep_2line = _CELL_PADDING + int(max(0, need - _CELL_PADDING) / 2) + 60
            keep = max(_MIN_COL_WIDTH * span, keep_2line)
            surplus = actual - keep
        else:
            deficit = need - actual
            if deficit > 40:
                per = deficit // span
                rem = deficit - per * span
                for offset, col_index in enumerate(range(start, end)):
                    recv_need[col_index] += per + (1 if offset < rem else 0)
            continue

        if surplus <= 40:
            continue

        block_sum = sum(widths[start:end]) or 1
        allocated = 0
        for col_index in range(start, end):
            cap = int(surplus * widths[col_index] / block_sum)
            room = max(0, widths[col_index] - _MIN_COL_WIDTH)
            if recv_need[col_index] > 0:
                room = 0
            cap = min(cap, room)
            donor_cap[col_index] += cap
            allocated += cap

        rem = surplus - allocated
        for col_index in range(start, end):
            if rem <= 0:
                break
            room = max(0, widths[col_index] - _MIN_COL_WIDTH - donor_cap[col_index])
            if recv_need[col_index] > 0 or room <= 0:
                continue
            delta = min(room, rem)
            donor_cap[col_index] += delta
            rem -= delta

    total_donor = sum(donor_cap)
    if total_donor <= 0:
        return widths

    total_recv = sum(recv_need)
    if total_recv <= 0:
        recv_weights = [0.0 if multiline_mask[index] else max(1.0, weights[index]) for index in range(len(widths))]
        weight_sum = sum(recv_weights)
        if weight_sum <= 0:
            return widths
        recv_need = [int(total_donor * weight / weight_sum) for weight in recv_weights]
        rem = total_donor - sum(recv_need)
        order = sorted(range(len(widths)), key=lambda index: recv_weights[index], reverse=True)
        for col_index in order:
            if rem <= 0 or recv_weights[col_index] <= 0:
                break
            recv_need[col_index] += 1
            rem -= 1
        total_recv = sum(recv_need)

    move = min(total_donor, total_recv)
    donor_rem = donor_cap[:]
    recv_rem = recv_need[:]
    while move > 0:
        recv_candidates = [index for index, need in enumerate(recv_rem) if need > 0]
        donor_candidates = [index for index, cap in enumerate(donor_rem) if cap > 0 and recv_rem[index] <= 0]
        if not recv_candidates or not donor_candidates:
            break

        recv_index = max(recv_candidates, key=lambda index: recv_rem[index])
        donor_candidates = [index for index in donor_candidates if index != recv_index]
        if not donor_candidates:
            break

        donor_index = max(donor_candidates, key=lambda index: donor_rem[index])
        step = min(move, donor_rem[donor_index], recv_rem[recv_index], 20)
        if step <= 0:
            break
        if widths[donor_index] - step < _MIN_COL_WIDTH:
            donor_rem[donor_index] = max(0, widths[donor_index] - _MIN_COL_WIDTH)
            continue

        widths[donor_index] -= step
        widths[recv_index] += step
        donor_rem[donor_index] -= step
        recv_rem[recv_index] -= step
        move -= step

    return widths


def _target_normal_table_width(
    avail_w: int,
    num_cols: int,
    first_min_ws: list[int],
    body_min_ws: list[int],
) -> int:
    if avail_w <= 0 or num_cols <= 0:
        return avail_w

    mins: list[int] = []
    for index in range(num_cols):
        need = _MIN_COL_WIDTH
        if index < len(first_min_ws):
            need = max(need, first_min_ws[index])
        if index < len(body_min_ws):
            need = max(need, body_min_ws[index])
        mins.append(need)

    content_total = sum(mins)
    if content_total >= int(avail_w * 0.9):
        return avail_w

    buffer_w = min(120 * num_cols, 600)
    target = content_total + buffer_w
    target = max(target, int(avail_w * 0.48))
    target = min(target, int(avail_w * 0.92))
    target = max(target, content_total)
    return min(target, avail_w)


def _build_smart_width_palette(raw_widths: list[int], avail_w: int, max_levels: int = 4) -> list[int]:
    if not raw_widths or avail_w <= 0:
        return []

    values = [max(_MIN_COL_WIDTH, min(avail_w, int(width))) for width in raw_widths if int(width) > 0]
    if not values:
        return []
    values.sort()

    levels = max(1, min(max_levels, len(values)))
    if levels == 1:
        return [values[0]]

    anchors: list[int] = []
    for index in range(levels):
        pos = int(round(index * (len(values) - 1) / (levels - 1)))
        anchors.append(values[pos])

    tol = max(120, int(avail_w * 0.03))
    palette = [anchors[0]]
    for width in anchors[1:-1]:
        if abs(width - palette[-1]) > tol:
            palette.append(width)
    if abs(anchors[-1] - palette[-1]) > tol:
        palette.append(anchors[-1])
    else:
        palette[-1] = anchors[-1]
    return palette


def _nearest_palette_width(raw_w: int, palette: list[int]) -> int:
    if not palette:
        return raw_w

    values = sorted(int(width) for width in palette if int(width) > 0)
    if not values:
        return raw_w
    if raw_w <= values[0]:
        return values[0]

    shrink_penalty = 1.25
    for index in range(1, len(values)):
        hi = values[index]
        if raw_w > hi:
            continue
        lo = values[index - 1]
        if hi <= lo:
            return hi
        up_gap = hi - raw_w
        down_gap = raw_w - lo
        return hi if up_gap <= down_gap * shrink_penalty else lo

    return values[-1]


def _build_smart_width_plan(
    raw_targets: list[tuple[int, int]],
    avail_w: int,
    max_levels: int = 4,
) -> dict[int, int]:
    if not raw_targets:
        return {}
    palette = _build_smart_width_palette(
        [width for _, width in raw_targets],
        avail_w,
        max_levels=max_levels,
    )
    return {index: _nearest_palette_width(width, palette) for index, width in raw_targets}


def _header_text_with_unit_break(text: str) -> str | None:
    value = (text or "").strip()
    if not value:
        return None
    match = _RE_HEADER_TRAILING_UNIT.match(value)
    if not match:
        return None
    head = match.group("head").strip()
    unit = match.group("unit").strip()
    if not head or not unit:
        return None
    if not re.search(r"[A-Za-z]", head):
        return None
    if len(head) < 8:
        return None
    return f"{head}\n({unit})"


def _body_text_with_paren_break(text: str) -> str | None:
    value = (text or "").strip()
    if not value:
        return None
    match = _RE_TRAILING_PAREN_CHUNK.match(value)
    if not match:
        return None
    head = match.group("head").strip()
    paren = match.group("paren").strip()
    if not head or not paren or len(head) < 6:
        return None
    return f"{head}\n{paren}"


def _normalize_first_row_unit_breaks(tbl_el, widths: list[int]) -> int:
    if not widths:
        return 0

    rows = tbl_el.findall(qn("w:tr"))
    if not rows:
        return 0

    changed = 0
    tr = rows[0]
    grid_index = 0
    tr_pr = tr.find(qn("w:trPr"))
    if tr_pr is not None:
        grid_before = tr_pr.find(qn("w:gridBefore"))
        if grid_before is not None:
            grid_index = int(grid_before.get(qn("w:val"), "0"))

    for tc in tr.findall(qn("w:tc")):
        span = 1
        tc_pr = tc.find(qn("w:tcPr"))
        if tc_pr is not None:
            grid_span = tc_pr.find(qn("w:gridSpan"))
            if grid_span is not None:
                span = int(grid_span.get(qn("w:val"), "1"))

        start = grid_index
        end = min(len(widths), start + span)
        cell_w = sum(widths[start:end]) if start < end else 0

        lines = [line for line in _extract_cell_lines(tc) if line.strip()]
        if len(lines) >= 2:
            grid_index += span
            continue

        text = (lines[0] if lines else (_get_cell_text(tc) or "")).strip()
        rewritten = _header_text_with_unit_break(text)
        if rewritten is None or _cell_has_rich_run_typography(tc):
            grid_index += span
            continue
        if cell_w > 0 and _text_width_twips(text) <= int(cell_w * 1.03):
            grid_index += span
            continue

        _set_cell_plain_text(tc, rewritten)
        changed += 1
        grid_index += span

    return changed


def _normalize_body_paren_breaks(tbl_el, widths: list[int]) -> int:
    if not widths:
        return 0

    rows = tbl_el.findall(qn("w:tr"))
    if len(rows) <= 1:
        return 0

    changed = 0
    for tr in rows[1:]:
        grid_index = 0
        tr_pr = tr.find(qn("w:trPr"))
        if tr_pr is not None:
            grid_before = tr_pr.find(qn("w:gridBefore"))
            if grid_before is not None:
                grid_index = int(grid_before.get(qn("w:val"), "0"))

        for tc in tr.findall(qn("w:tc")):
            span = 1
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is not None:
                grid_span = tc_pr.find(qn("w:gridSpan"))
                if grid_span is not None:
                    span = int(grid_span.get(qn("w:val"), "1"))

            start = grid_index
            end = min(len(widths), start + span)
            cell_w = sum(widths[start:end]) if start < end else 0

            lines = [line for line in _extract_cell_lines(tc) if line.strip()]
            if len(lines) >= 2:
                grid_index += span
                continue

            text = (lines[0] if lines else (_get_cell_text(tc) or "")).strip()
            rewritten = _body_text_with_paren_break(text)
            if rewritten is None or _cell_has_rich_run_typography(tc):
                grid_index += span
                continue
            if cell_w > 0 and _text_width_twips(text) < int(cell_w * 0.90):
                grid_index += span
                continue

            _set_cell_plain_text(tc, rewritten)
            changed += 1
            grid_index += span

    return changed


def _update_cell_widths(tbl_el, widths: list[int]) -> None:
    for tr in tbl_el.findall(qn("w:tr")):
        grid_index = 0
        tr_pr = tr.find(qn("w:trPr"))
        if tr_pr is not None:
            grid_before = tr_pr.find(qn("w:gridBefore"))
            if grid_before is not None:
                grid_index = int(grid_before.get(qn("w:val"), "0"))

        for tc in tr.findall(qn("w:tc")):
            span = 1
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is not None:
                grid_span = tc_pr.find(qn("w:gridSpan"))
                if grid_span is not None:
                    span = int(grid_span.get(qn("w:val"), "1"))

            cell_w = 0
            for offset in range(span):
                col_index = grid_index + offset
                if col_index < len(widths):
                    cell_w += widths[col_index]

            if tc_pr is None:
                tc_pr = etree.SubElement(tc, qn("w:tcPr"))
                tc.insert(0, tc_pr)
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = etree.SubElement(tc_pr, qn("w:tcW"))
            tc_w.set(qn("w:w"), str(cell_w))
            tc_w.set(qn("w:type"), "dxa")

            grid_index += span


def _clear_row_width_exceptions(tbl_el) -> None:
    for tr in tbl_el.findall(qn("w:tr")):
        tr_pr = tr.find(qn("w:trPr"))
        if tr_pr is None:
            continue
        for tbl_pr_ex in tr_pr.findall(qn("w:tblPrEx")):
            for tag in ("w:tblW", "w:tblInd", "w:tblCellSpacing"):
                extra = tbl_pr_ex.find(qn(tag))
                if extra is not None:
                    tbl_pr_ex.remove(extra)


def _apply_three_line_borders(tbl_el, tbl_cfg) -> None:
    outer_w = max(1, int(float(tbl_cfg.three_line_header_width_pt or 1.0) * 8))
    inner_w = max(1, int(float(tbl_cfg.three_line_bottom_width_pt or 0.5) * 8))

    _clear_table_style_inheritance(tbl_el)
    clear_cell_border_overrides(tbl_el)

    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    tbl_borders = find_or_create(tbl_pr, "w:tblBorders")

    for tag in ("w:top", "w:bottom"):
        border = find_or_create(tbl_borders, tag)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(outer_w))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")

    for tag in ("w:left", "w:right", "w:insideH", "w:insideV"):
        border = find_or_create(tbl_borders, tag)
        border.set(qn("w:val"), "none")
        border.set(qn("w:sz"), "0")

    rows = tbl_el.findall(qn("w:tr"))
    if not rows:
        return

    for tc in rows[0].findall(qn("w:tc")):
        tc_pr = find_or_create(tc, "w:tcPr")
        tc_borders = find_or_create(tc_pr, "w:tcBorders")
        bottom = find_or_create(tc_borders, "w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(inner_w))
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")


def _apply_color_table_style(table: Table, tbl_cfg) -> None:
    palette = color_palette(getattr(tbl_cfg, "color_table_accent", "blue"))
    variant = color_variant(getattr(tbl_cfg, "color_table_variant", "header_grid"))
    tbl_el = table._element
    line_w = max(1, int(float(getattr(tbl_cfg, "border_width_pt", 0.5) or 0.5) * 8))
    outer_w = max(1, int(float(getattr(tbl_cfg, "three_line_header_width_pt", 1.0) or 1.0) * 8))
    header_rule_w = max(1, int(float(getattr(tbl_cfg, "three_line_bottom_width_pt", 0.5) or 0.5) * 8))
    border_w = outer_w if variant.header_rule_only else line_w

    _clear_table_style_inheritance(tbl_el)
    clear_cell_border_overrides(tbl_el)
    _apply_color_table_borders(tbl_el, variant, palette.accent, palette.accent_light, border_w)

    rows = list(table.rows)
    if not rows:
        return

    header_fill = palette.accent if variant.header_fill else ""
    for row_index, row in enumerate(rows):
        for cell in row.cells:
            if row_index == 0 and header_fill:
                _set_cell_shading(cell._element, header_fill)
                _set_cell_text_color(cell._element, palette.header_text)
                _set_cell_bold(cell._element, True)
                if variant.header_rule_only:
                    _set_cell_border(cell._element, "bottom", "single", header_rule_w, palette.accent)
            elif variant.zebra and row_index % 2 == 1:
                _set_cell_shading(cell._element, palette.accent_soft)
            else:
                _remove_cell_shading(cell._element)


def _apply_color_table_borders(tbl_el, variant, accent: str, light: str, width_eighth: int) -> None:
    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    tbl_borders = find_or_create(tbl_pr, "w:tblBorders")

    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _set_table_border(tbl_borders, side, "none", 0, light)

    if variant.header_rule_only:
        _set_table_border(tbl_borders, "top", "single", width_eighth, accent)
        _set_table_border(tbl_borders, "bottom", "single", width_eighth, accent)
        return

    if variant.show_horizontal:
        _set_table_border(tbl_borders, "top", "single", width_eighth, accent)
        _set_table_border(tbl_borders, "bottom", "single", width_eighth, accent)
        _set_table_border(tbl_borders, "insideH", "single", width_eighth, light)

    if variant.show_vertical:
        _set_table_border(tbl_borders, "left", "single", width_eighth, accent)
        _set_table_border(tbl_borders, "right", "single", width_eighth, accent)
        _set_table_border(tbl_borders, "insideV", "single", width_eighth, light)


def _apply_no_table_borders(tbl_el) -> None:
    _clear_table_style_inheritance(tbl_el)
    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    tbl_borders = find_or_create(tbl_pr, "w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _set_table_border(tbl_borders, side, "none", 0)

    for tc in tbl_el.iter(qn("w:tc")):
        tc_pr = find_or_create(tc, "w:tcPr")
        tc_borders = find_or_create(tc_pr, "w:tcBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = find_or_create(tc_borders, f"w:{side}")
            border.set(qn("w:val"), "none")
            border.set(qn("w:sz"), "0")


def _clear_table_style_inheritance(tbl_el) -> None:
    tbl_pr = find_or_create(tbl_el, "w:tblPr")
    for tag in ("w:tblStyle", "w:tblLook"):
        child = tbl_pr.find(qn(tag))
        if child is not None:
            tbl_pr.remove(child)


def _set_table_border(tbl_borders, side: str, value: str, size: int, color: str = "000000") -> None:
    border = find_or_create(tbl_borders, f"w:{side}")
    border.set(qn("w:val"), value)
    border.set(qn("w:sz"), str(max(0, int(size))))
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), color)


def _set_cell_border(tc, side: str, value: str, size: int, color: str = "000000") -> None:
    tc_pr = find_or_create(tc, "w:tcPr")
    tc_borders = find_or_create(tc_pr, "w:tcBorders")
    border = find_or_create(tc_borders, f"w:{side}")
    border.set(qn("w:val"), value)
    border.set(qn("w:sz"), str(max(0, int(size))))
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), color)


def _set_cell_shading(tc, fill: str) -> None:
    tc_pr = find_or_create(tc, "w:tcPr")
    shd = find_or_create(tc_pr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _remove_cell_shading(tc) -> None:
    tc_pr = tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return
    shd = tc_pr.find(qn("w:shd"))
    if shd is not None:
        tc_pr.remove(shd)


def _set_cell_text_color(tc, color: str) -> None:
    for run in tc.iter(qn("w:r")):
        run_pr = find_or_create(run, "w:rPr")
        color_node = find_or_create(run_pr, "w:color")
        color_node.set(qn("w:val"), color)


def _set_cell_bold(tc, enabled: bool) -> None:
    for run in tc.iter(qn("w:r")):
        run_pr = find_or_create(run, "w:rPr")
        for tag in ("w:b", "w:bCs"):
            node = find_or_create(run_pr, tag)
            node.set(qn("w:val"), "1" if enabled else "0")


def _set_run_toggle(run, tags: tuple[str, str], enabled: bool) -> None:
    run_pr = find_or_create(run._element, "w:rPr")
    for tag in tags:
        node = find_or_create(run_pr, tag)
        node.set(qn("w:val"), "1" if enabled else "0")


def _format_table_cells(table: Table, tbl_cfg) -> None:
    font_cn = tbl_cfg.font_cn
    font_en = tbl_cfg.font_en
    size_pt = tbl_cfg.size_pt
    alignment = tbl_cfg.cell_alignment
    line_spacing = tbl_cfg.line_spacing_mode
    bold = bool(getattr(tbl_cfg, "bold", False))
    italic = bool(getattr(tbl_cfg, "italic", False))
    first_row_bold = bool(getattr(tbl_cfg, "first_row_bold", False))
    border_mode = str(getattr(tbl_cfg, "border_mode", "") or "").strip().lower()
    variant = color_variant(getattr(tbl_cfg, "color_table_variant", "header_grid"))
    color_header_bold = border_mode == "color_table" and bool(variant.header_fill)

    for row_index, _col_index, cell in iter_table_cells(table):
        for para in cell.paragraphs:
            pf = para.paragraph_format

            if alignment == "center":
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif alignment == "left":
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif alignment == "right":
                pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            if line_spacing == "single":
                pf.line_spacing = 1.0
            elif line_spacing == "one_half":
                pf.line_spacing = 1.5
            elif line_spacing == "double":
                pf.line_spacing = 2.0

            pf.space_before = Pt(0)
            pf.space_after = Pt(0)

            for run in para.runs:
                if font_en:
                    run.font.name = resolve_font(font_en, lang="en")
                if font_cn:
                    set_run_east_asian_font(run, font_cn)
                if size_pt:
                    run.font.size = Pt(size_pt)
                run_bold = bold or (row_index == 0 and (first_row_bold or color_header_bold))
                run.font.bold = run_bold
                run.font.italic = italic
                _set_run_toggle(run, ("w:b", "w:bCs"), run_bold)
                _set_run_toggle(run, ("w:i", "w:iCs"), italic)


def _get_cell_text(tc) -> str:
    parts: list[str] = []
    for para in tc.findall(qn("w:p")):
        for run in para.findall(qn("w:r")):
            for child in run:
                local = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                if local == "t" and child.text:
                    parts.append(child.text)
                elif local in {"br", "cr"}:
                    parts.append("\n")
                elif local == "tab":
                    parts.append("\t")
    return "".join(parts)


def _cell_has_math_content(tc) -> bool:
    return bool(tc.findall(f".//{{{_M_NS}}}oMath") or tc.findall(f".//{{{_M_NS}}}oMathPara"))


def _cell_has_non_text_content(tc) -> bool:
    if _cell_has_math_content(tc):
        return True
    if tc.findall(f".//{qn('w:object')}"):
        return True
    if tc.findall(f".//{qn('w:drawing')}"):
        return True
    if tc.findall(f".//{qn('w:pict')}"):
        return True
    return False


def _cell_has_rich_run_typography(tc) -> bool:
    runs = list(tc.iter(qn("w:r")))
    signatures: list[tuple[bool, bool, str, bool, bool]] = []
    for run in runs:
        text = "".join((node.text or "") for node in run.findall(qn("w:t")))
        if not text:
            continue
        run_pr = run.find(qn("w:rPr"))
        if run_pr is None:
            signatures.append((False, False, "", False, False))
            continue

        def _onoff(tag: str) -> bool:
            node = run_pr.find(qn(f"w:{tag}"))
            if node is None:
                return False
            value = (node.get(qn("w:val")) or "").strip().lower()
            return value not in {"0", "false", "off", "none"}

        vert_align = run_pr.find(qn("w:vertAlign"))
        if vert_align is not None:
            value = (vert_align.get(qn("w:val")) or "").strip()
            if value in {"subscript", "superscript"}:
                return True

        position = run_pr.find(qn("w:position"))
        position_value = (position.get(qn("w:val")) or "").strip() if position is not None else ""
        if position_value not in {"", "0"}:
            return True

        underline = run_pr.find(qn("w:u"))
        underline_value = ""
        if underline is not None:
            underline_value = (underline.get(qn("w:val")) or "").strip().lower() or "single"
            if underline_value == "none":
                underline_value = ""

        signatures.append(
            (
                _onoff("b") or _onoff("bCs"),
                _onoff("i") or _onoff("iCs"),
                underline_value,
                _onoff("strike") or _onoff("dstrike"),
                _onoff("caps") or _onoff("smallCaps"),
            )
        )

    return len(set(signatures)) > 1


def _set_cell_plain_text(tc, text: str) -> None:
    text_nodes = tc.findall(f".//{qn('w:t')}")
    if text_nodes:
        text_nodes[0].text = text
        for node in text_nodes[1:]:
            node.text = ""
        for run in tc.iter(qn("w:r")):
            for child in list(run):
                local = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                if local in {"tab", "br", "cr"}:
                    run.remove(child)
        return

    paras = tc.findall(qn("w:p"))
    para = paras[0] if paras else etree.SubElement(tc, qn("w:p"))
    run = etree.SubElement(para, qn("w:r"))
    text_node = etree.SubElement(run, qn("w:t"))
    text_node.text = text


def _is_equation_table(table: Table | object) -> bool:
    """Return whether *table* has the structural shape of an equation table.

    Native math by itself is deliberately not sufficient: ordinary data
    tables may legitimately contain equations.  An equation table must either
    carry our explicit marker, or place formula content before a final empty /
    equation-number cell.
    """
    tbl_el = getattr(table, "_element", table)
    rows = tbl_el.findall(qn("w:tr"))

    tbl_pr = tbl_el.find(qn("w:tblPr"))
    has_explicit_marker = False
    if tbl_pr is not None:
        for tag in ("w:tblCaption", "w:tblDescription"):
            hint = tbl_pr.find(qn(tag))
            value = (hint.get(qn("w:val")) or "") if hint is not None else ""
            if value.strip().casefold() == "alavette-equation-table":
                has_explicit_marker = True
                break
    if has_explicit_marker:
        return True
    if not rows or len(rows) > _MAX_UNMARKED_EQUATION_TABLE_ROWS:
        return False

    for tr in rows:
        cells = tr.findall(qn("w:tc"))
        if len(cells) < 2:
            continue
        number_text = _get_cell_text(cells[-1]).strip()
        has_number_slot = not number_text or bool(_RE_EQUATION_NUMBER.match(number_text))
        if not has_number_slot:
            continue
        for formula_cell in cells[:-1]:
            has_native_math = _cell_has_formula_math_anchor(formula_cell)
            has_equation_ole = any(
                _RE_EQUATION_OLE.search(
                    str(ole.get("ProgID") or ole.get(f"{{{_O_NS}}}ProgID") or "")
                )
                for ole in formula_cell.findall(f".//{{{_O_NS}}}OLEObject")
            )
            if has_native_math or has_equation_ole:
                return True
            if (
                number_text
                and _RE_EQUATION_NUMBER.match(number_text)
                and _looks_like_formula_cell_text(_get_cell_text(formula_cell))
            ):
                return True

        if not number_text:
            formula_text = " ".join(
                _get_cell_text(tc).strip() for tc in cells[:-1]
            )
            tail_match = _RE_EQUATION_NUMBER_TAIL.search(formula_text)
            if tail_match is not None:
                prefix = formula_text[: tail_match.start()].strip()
                if _looks_like_formula_cell_text(prefix):
                    return True
            # A bare inline wrapper such as ``$x$`` is valid formula content
            # but not sufficient structural evidence for an equation table.
            # Unwrapped high-confidence formula text retains the 0.2 fallback.
            if (
                formula_text
                and not _RE_EQUATION_SOURCE.search(formula_text)
                and _looks_like_formula_cell_text(formula_text)
            ):
                return True

    # A text-only fallback requires both a mathematical operator in the
    # formula region and an equation-number-looking final cell.  Ordinary data
    # tables that merely contain "=" must not be routed to this module.
    for tr in rows:
        cells = tr.findall(qn("w:tc"))
        if len(cells) < 2:
            continue
        formula_text = " ".join(_get_cell_text(tc).strip() for tc in cells[:-1])
        number_text = _get_cell_text(cells[-1]).strip()
        if (
            formula_text
            and _looks_like_formula_cell_text(formula_text)
            and _RE_EQUATION_NUMBER.match(number_text)
        ):
            return True

    return False


def _looks_like_equation_text(text: str) -> bool:
    value = str(text or "").strip()
    if not value:
        return False
    if _RE_EQUATION_ARROW.search(value):
        return True
    if "=" in value and not _RE_CJK_TEXT.search(value):
        left, right = (part.strip() for part in value.split("=", 1))
        compact = bool(left and right and not re.search(r"\s", left + right))
        math_signal = bool(
            re.search(r"[0-9\^_{}\(\)\[\]+\-*/Α-Ͽ]", left + right)
            or len(left) <= 2
        )
        if compact and math_signal:
            return True
    return bool(
        re.search(r"\s\+\s", value)
        and re.search(r"[A-Za-z\u0391-\u03a9\u03b1-\u03c9]", value)
    )


def _looks_like_formula_cell_text(text: str) -> bool:
    """Conservative text-only formula evidence used by table and row gates."""

    value = str(text or "").strip()
    if not value:
        return False
    if looks_like_caption_text(value) or looks_like_bibliographic_reference_text(value):
        return False
    if _RE_EQUATION_SOURCE.search(value) or _looks_like_equation_text(value):
        return True
    if _RE_COMPACT_FORMULA_TEXT.fullmatch(value):
        if _RE_COMPACT_PLUS_TIMES_EXPR.search(value):
            return True
        if (
            "/" in value
            and re.search(r"[(){}\[\]]", value)
            and _RE_COMPACT_SLASH_EXPR.search(value)
        ):
            return True

    matched, confidence, _source_type = looks_like_formula_text(value)
    if not matched:
        return False
    if value.startswith("\\"):
        return float(confidence) >= 0.60
    return float(confidence) >= 0.72


def _extract_cell_omml_linear_text(tc) -> str:
    return "".join(
        str(value)
        for value in tc.xpath(
            ".//*[namespace-uri()='%s' and local-name()='t']/text()" % _M_NS
        )
    ).strip()


def _cell_has_formula_math_anchor(tc) -> bool:
    has_math = bool(
        tc.findall(f".//{{{_M_NS}}}oMath")
        or tc.findall(f".//{{{_M_NS}}}oMathPara")
    )
    if not has_math:
        return False

    linear = _extract_cell_omml_linear_text(tc)
    if not linear:
        return True
    if looks_like_caption_text(linear) or looks_like_bibliographic_reference_text(linear):
        return False
    if len(linear) > 240:
        return False
    if _RE_OMML_FORMULA_ANCHOR.search(linear):
        return True

    chinese_count = len(_RE_CJK_TEXT.findall(linear))
    prose_punct_count = len(_RE_PROSE_PUNCT.findall(linear))
    matched, confidence, _source_type = looks_like_formula_text(linear)
    if (
        matched
        and float(confidence) >= 0.78
        and chinese_count <= max(2, int(len(linear) * 0.12))
    ):
        return True
    if chinese_count >= max(6, int(len(linear) * 0.16)) and prose_punct_count:
        return False
    if chinese_count >= max(12, int(len(linear) * 0.28)):
        return False
    if any(
        marker in linear
        for marker in (
            "体积比为",
            "混合液中",
            "搅拌",
            "透析",
            "分别得到",
            "分别表示",
            "其中，",
        )
    ):
        return False
    return True
