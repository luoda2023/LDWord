"""
entity_fill — 实体信息填充模块

从 EntityArchive 中读取实体数据（作者、学校、日期等），
填充到文档中的占位符或指定位置。
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from src.modules.base import BaseModule, ModuleMeta
from src.config.material_schema_registry import resolve_material_schema_ids
from src.shared.engine.run_ops import replace_run_text, get_full_text, iter_story_paragraphs
from src.shared.engine.fixed_layout_text import (
    replace_fixed_layout_mapped_fields,
    replace_fixed_layout_placeholders,
)
from src.shared.engine.exact_material_placeholders import (
    replace_document_exact_placeholders,
)
from src.shared.engine.material_token_contract import (
    MaterialTokenNamespace,
    material_token,
)

if TYPE_CHECKING:
    from docx import Document
    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


def _entity_replacement_items(
    entity_data: dict[str, str],
    field_aliases: dict[str, str],
    timeline_fields: set[str],
) -> list[tuple[str, str, str]]:
    items: list[tuple[str, str, str]] = []
    for key, value in dict(entity_data or {}).items():
        field_key = str(key or "").strip()
        if not field_key:
            continue
        for token in _field_tokens(field_key, timeline_fields):
            items.append((token, str(value or ""), field_key))
    items.extend(
        _alias_replacement_items(entity_data, field_aliases, timeline_fields)
    )
    return items


def _alias_replacement_items(
    entity_data: dict[str, str],
    field_aliases: dict[str, str],
    timeline_fields: set[str],
) -> list[tuple[str, str, str]]:
    items: list[tuple[str, str, str]] = []
    values = dict(entity_data or {})
    for alias, field_key in dict(field_aliases or {}).items():
        alias_key = str(alias or "").strip()
        target_key = str(field_key or "").strip()
        if not alias_key or not target_key or target_key not in values:
            continue
        for token in _field_tokens(
            alias_key,
            timeline_fields,
            owner_key=target_key,
        ):
            items.append((token, str(values.get(target_key, "") or ""), target_key))
    return items


def _field_tokens(
    key: str,
    timeline_fields: set[str],
    *,
    owner_key: str | None = None,
) -> tuple[str]:
    namespace = (
        MaterialTokenNamespace.TIME
        if str(owner_key or key) in timeline_fields
        else MaterialTokenNamespace.TEXT
    )
    return (material_token(namespace, key),)


class EntityFillModule(BaseModule):
    """实体信息填充模块。

    职责：
    - 读取 config.entity_data 中的键值对
    - 替换文档中的 {{@text:key}} / {{@time:key}} 占位符
    - 输出已填充的键值对到 context.entity_values
    """

    meta = ModuleMeta(
        name="entity_fill",
        description="实体信息填充",
        category="fill",
        requires_config=("entity_data",),
        provides=("entity_values",),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        entity_data = config.entity_data
        if not entity_data:
            return
        timeline_fields = set(getattr(config, "timeline_field_keys", ()) or ())

        if bool(getattr(config, "exact_material_placeholders", False)):
            exact_values = dict(entity_data)
            source_keys = {str(key): str(key) for key in entity_data}
            for alias, field_key in dict(
                getattr(config, "field_aliases", {}) or {}
            ).items():
                alias_key = str(alias or "").strip()
                source_key = str(field_key or "").strip()
                if not alias_key or source_key not in entity_data:
                    continue
                exact_values[alias_key] = entity_data[source_key]
                source_keys[alias_key] = source_key
            result = replace_document_exact_placeholders(
                doc,
                exact_values,
                timeline_field_keys=timeline_fields,
            )
            replaced_source_keys = {
                source_keys.get(key, key) for key in result.replaced_keys
            }
            context.entity_values = {
                key: str(entity_data.get(key, "") or "")
                for key in replaced_source_keys
                if key in entity_data
            }
            if result.total_replacements:
                tracker.record(
                    rule_name=self.meta.name,
                    target=f"{result.total_replacements} 处精确资料占位符",
                    section="global",
                    change_type="replace",
                    before="{{@text:用户字段}}",
                    after=f"已填入 {len(result.replaced_keys)} 个字段",
                )
            return

        total_replaced = 0
        filled_values: dict[str, str] = {}

        for para in doc.paragraphs:
            text = get_full_text(para)
            if not text:
                continue

            for key, value in entity_data.items():
                # Strict public text namespace; no legacy bare/dollar aliases.
                for placeholder in _field_tokens(key, timeline_fields):
                    if placeholder in text:
                        count = replace_run_text(para, placeholder, value)
                        total_replaced += count
                        filled_values[key] = value
                        text = get_full_text(para)

        # 同样处理页眉页脚
        for section in doc.sections:
            for part in (section.header, section.footer):
                for para in part.paragraphs:
                    text = get_full_text(para)
                    if not text:
                        continue
                    for key, value in entity_data.items():
                        for fmt in _field_tokens(key, timeline_fields):
                            if fmt in text:
                                replace_run_text(para, fmt, value)
                                total_replaced += 1
                                filled_values[key] = value
                                text = get_full_text(para)

        for para in iter_story_paragraphs(doc):
            text = get_full_text(para)
            if not text:
                continue
            for token, value, key in _entity_replacement_items(
                entity_data,
                getattr(config, "field_aliases", {}) or {},
                timeline_fields,
            ):
                if token in text:
                    count = replace_run_text(para, token, value)
                    total_replaced += count
                    filled_values[key] = value
                    text = get_full_text(para)

        fixed_layout_replacements: dict[str, str] = {}
        token_to_key: dict[str, str] = {}
        token_to_value: dict[str, str] = {}
        for key, value in entity_data.items():
            for token in _field_tokens(key, timeline_fields):
                fixed_layout_replacements[token] = value
                token_to_key[token] = key
                token_to_value[token] = value
        for token, value, key in _alias_replacement_items(
            entity_data,
            getattr(config, "field_aliases", {}) or {},
            timeline_fields,
        ):
            fixed_layout_replacements[token] = value
            token_to_key[token] = key
            token_to_value[token] = value
        fixed_layout_result = replace_fixed_layout_placeholders(
            doc,
            fixed_layout_replacements,
        )
        total_replaced += fixed_layout_result.total_replacements
        for token in fixed_layout_result.replaced_tokens:
            key = token_to_key.get(token)
            if key:
                filled_values[key] = token_to_value.get(token, entity_data.get(key, ""))

        input_profile = getattr(config, "input_source_profile", None)
        schema_ids = resolve_material_schema_ids(
            str(getattr(input_profile, "material_schema_id", "") or "").strip(),
            list(getattr(input_profile, "material_schema_ids", []) or []),
        )
        mapped_result = replace_fixed_layout_mapped_fields(
            doc,
            entity_data,
            field_aliases=getattr(config, "field_aliases", {}) or {},
            schema_ids=schema_ids,
        )
        total_replaced += mapped_result.total_replacements
        for key in mapped_result.replaced_fields:
            filled_values[key] = entity_data[key]
        if mapped_result.has_replacements:
            tracker.record(
                rule_name=self.meta.name,
                target=(
                    f"{mapped_result.total_replacements} 处固定版位字段映射"
                ),
                section="fixed_layout",
                change_type="fixed_layout_field_mapping",
                before=", ".join(mapped_result.matched_identifiers),
                after=", ".join(mapped_result.replaced_fields),
            )

        context.entity_values = filled_values

        if total_replaced:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{total_replaced} 处占位符/字段",
                section="global",
                change_type="replace",
                before="{{@text:...}} / {{@time:...}}",
                after=f"已填充 {len(filled_values)} 个字段",
            )
