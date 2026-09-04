"""TOC generation and formatting module."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING

from docx.oxml import OxmlElement

from src.modules.base import BaseModule, ModuleMeta
from src.shared.engine.document_scope_runtime import document_scope_includes_role
from src.shared.engine.document_text_heuristics import (
    looks_like_numbered_toc_entry_with_page_suffix,
    looks_like_toc_entry_line,
)
from src.shared.engine.field_builder import (
    build_complex_field,
    build_toc_instruction,
    iter_field_instructions,
)
from src.shared.engine.field_refresh import (
    document_has_toc,
    ensure_update_fields_on_open,
)
from src.shared.engine.ooxml_ops import clone_element, find_or_create, qn
from src.shared.engine.toc_style_ops import (
    TOC_HEADING_STYLE_CANDIDATES,
    apply_toc_paragraph_style,
    resolve_toc_entry_style_config,
    resolve_toc_title_style_config,
    sync_toc_styles,
    toc_heading_key,
    toc_heading_level,
    toc_word_style_for_heading_level,
)

if TYPE_CHECKING:
    from docx import Document
    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


TOC_TITLE_TEXT = "\u76ee\u5f55"
TOC_FIELD_PLACEHOLDER = "\u8bf7\u66f4\u65b0\u57df\u4ee5\u663e\u793a\u76ee\u5f55"
TOC_TITLE_RE = re.compile(r"^(\u76ee\u5f55|\u76ee\s*\u5f55|contents|tableofcontents)$", re.IGNORECASE)
TOC_LEVEL_STYLE_RE = re.compile(r"^(toc|\u76ee\u5f55)\s*(\d+)$", re.IGNORECASE)
RE_TOC_LEVEL1_CN = re.compile(r"^\u7b2c[\u4e00-\u9fff\d]+(?:\u7ae0|\u7bc7)")
RE_TOC_LEVEL2_CN = re.compile(r"^\u7b2c[\u4e00-\u9fff\d]+\u8282")
RE_LEVEL4 = re.compile(r"^\d+\.\d+\.\d+\.\d+(?:[\.、．])?\s*\S")
RE_LEVEL3 = re.compile(r"^\d+\.\d+\.\d+(?:[\.、．])?\s*\S")
RE_LEVEL2 = re.compile(r"^\d+\.\d+(?:[\.、．])?\s*\S")
BACK_MATTER_TITLES = {
    "references": "\u53c2\u8003\u6587\u732e",
    "errata": "\u52d8\u8bef",
    "appendix": "\u9644\u5f55",
    "acknowledgment": "\u81f4\u8c22",
    "resume": "\u4e2a\u4eba\u7b80\u5386",
}


class TocModule(BaseModule):
    meta = ModuleMeta(
        name="toc",
        description="\u76ee\u5f55\u751f\u6210",
        category="structure",
        requires_config=("toc",),
        depends_on=("heading_recognition",),
        soft_after=("heading_numbering",),
        consumes=("heading_map", "doc_tree"),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        if not document_scope_includes_role(context, "toc"):
            return
        toc_cfg = config.toc
        max_level = _clamp_toc_level(toc_cfg.max_level)
        insert_position = toc_cfg.insert_position
        mode = str(getattr(toc_cfg, "mode", "word_native") or "word_native").strip().lower()
        style_sync_count = sync_toc_styles(doc, config, max_level=max_level)
        existing_plain_range = _resolve_existing_toc_range(doc, context)

        if mode == "plain":
            entries = _collect_plain_toc_entries(doc, config, context, max_level=max_level)
            if existing_plain_range is not None:
                start, end = existing_plain_range
                inserted = _replace_paragraph_toc_range_with_plain(doc, start, end, entries)
                _format_inserted_plain_toc(doc, config, inserted, entries)
                tracker.record(
                    rule_name=self.meta.name,
                    target=f"plain toc {len(entries)} entries",
                    section="global",
                    change_type="rebuild",
                    before=f"range=({start},{end})",
                    after="plain toc rebuilt",
                )
            else:
                insert_idx = _find_insert_position(doc, insert_position, context)
                if insert_idx is not None:
                    inserted = _insert_plain_toc(doc, insert_idx, entries)
                    _format_inserted_plain_toc(doc, config, inserted, entries)
                    tracker.record(
                        rule_name=self.meta.name,
                        target=f"paragraph {insert_idx}",
                        section="global",
                        change_type="insert",
                        before="no toc",
                        after=f"plain toc ({len(entries)} entries)",
                    )
        else:
            has_existing_toc = _has_existing_toc(doc)
            if has_existing_toc:
                updated_fields = _sync_toc_field_instructions(doc, max_level)
                formatted_count = _format_existing_toc_paragraphs(doc, config, context)
                ensure_update_fields_on_open(doc)
                tracker.record(
                    rule_name=self.meta.name,
                    target="existing toc",
                    section="global",
                    change_type="format",
                    before="toc stale",
                    after=f"depth=1-{max_level}, updated {updated_fields} fields and formatted {formatted_count} paragraphs",
                )
            else:
                insert_idx = _find_insert_position(doc, insert_position, context)
                if insert_idx is not None:
                    _insert_toc(doc, insert_idx, max_level)
                    _format_inserted_toc_title(doc, config, insert_idx)
                    ensure_update_fields_on_open(doc)
                    tracker.record(
                        rule_name=self.meta.name,
                        target=f"paragraph {insert_idx}",
                        section="global",
                        change_type="insert",
                        before="no toc",
                        after=f"native toc (1-{max_level})",
                    )

        if style_sync_count:
            tracker.record(
                rule_name=self.meta.name,
                target=f"{style_sync_count} toc style defs",
                section="global",
                change_type="style",
                before="toc styles unsynced",
                after="toc title/entry styles synced",
            )


def _has_existing_toc(doc: Document) -> bool:
    return document_has_toc(doc)


def _clamp_toc_level(value) -> int:
    try:
        level = int(value or 3)
    except (TypeError, ValueError):
        level = 3
    return max(1, min(level, 6))


def _mark_toc_for_update(doc: Document) -> None:
    body = doc.element.body
    for _kind, elem, instr in iter_field_instructions(body):
        if _is_toc_instruction(instr):
            elem.set(qn("w:dirty"), "true")


def _sync_toc_field_instructions(doc: Document, max_level: int) -> int:
    body = doc.element.body
    instruction = build_toc_instruction(max_level=max(1, int(max_level or 1)))
    changed = 0
    for kind, elem, instr in iter_field_instructions(body):
        if not _is_toc_instruction(instr):
            continue
        elem.set(qn("w:dirty"), "true")
        if kind == "simple":
            if elem.get(qn("w:instr"), "") != instruction:
                elem.set(qn("w:instr"), instruction)
                changed += 1
            continue
        if _replace_complex_field_instruction(elem, instruction):
            changed += 1
    return changed


def _replace_complex_field_instruction(begin_elem, instruction: str) -> bool:
    begin_run = begin_elem.getparent()
    if begin_run is None:
        return False
    container = begin_run.getparent()
    if container is None:
        return False

    children = list(container)
    try:
        start_index = children.index(begin_run)
    except ValueError:
        return False

    instr_nodes = []
    nested_depth = 0
    for child in children[start_index + 1:]:
        fld_chars = list(child.iter(qn("w:fldChar")))
        if any(
            fld.get(qn("w:fldCharType")) in {"separate", "end"}
            for fld in fld_chars
            if nested_depth == 0
        ):
            break
        instr_nodes.extend(child.iter(qn("w:instrText")))
        for fld in fld_chars:
            fld_type = fld.get(qn("w:fldCharType"))
            if fld_type == "begin":
                nested_depth += 1
            elif fld_type == "end" and nested_depth > 0:
                nested_depth -= 1

    if not instr_nodes:
        return False

    current = "".join(node.text or "" for node in instr_nodes)
    instr_nodes[0].text = instruction
    for node in instr_nodes[1:]:
        node.text = ""
    return current != instruction


def _norm_no_space(text: str) -> str:
    return re.sub(r"\s+", "", text or "").strip()


def _is_toc_title_paragraph(para) -> bool:
    raw = _norm_no_space(para.text)
    if not raw:
        return False
    if TOC_TITLE_RE.match(raw):
        return True
    style_name = str(para.style.name if para.style else "").strip()
    return any(style_name == candidate for candidate in TOC_HEADING_STYLE_CANDIDATES)


def _is_toc_level_style_para(para) -> bool:
    style_name = str(para.style.name if para.style else "").strip()
    style_id = str(getattr(para.style, "style_id", "") if para.style else "").strip()
    return bool(TOC_LEVEL_STYLE_RE.match(style_name) or TOC_LEVEL_STYLE_RE.match(style_id))


def _infer_toc_level_from_paragraph(para) -> str:
    style_name = str(para.style.name if para.style else "").strip()
    style_id = str(getattr(para.style, "style_id", "") if para.style else "").strip()
    for value in (style_name, style_id):
        match = TOC_LEVEL_STYLE_RE.match(value)
        if match:
            return toc_heading_key(int(match.group(2)))

    raw = (para.text or "").strip()
    if RE_TOC_LEVEL2_CN.match(raw):
        return "heading2"
    if RE_TOC_LEVEL1_CN.match(raw):
        return "heading1"
    if RE_LEVEL4.match(raw):
        return "heading4"
    if RE_LEVEL3.match(raw):
        return "heading3"
    if RE_LEVEL2.match(raw):
        return "heading2"
    return "heading1"


def _toc_section_is_suspicious(doc: Document, start: int, end: int) -> bool:
    total = len(doc.paragraphs)
    if start < 0 or end <= start or start >= total:
        return True

    entry_like = 0
    heading_like = 0
    nonempty = 0
    for index in range(max(0, start), min(total, end)):
        para = doc.paragraphs[index]
        raw = (para.text or "").strip()
        if not raw:
            continue
        nonempty += 1
        if _is_toc_title_paragraph(para):
            continue
        if _is_toc_level_style_para(para) or looks_like_toc_entry_line(raw) or looks_like_numbered_toc_entry_with_page_suffix(raw):
            entry_like += 1
            continue
        style_name = str(para.style.name if para.style else "").strip().lower()
        if style_name.startswith("heading") or "\u6807\u9898" in style_name or raw.endswith(("\u3002", "\uff1b", "!", "\uff1f", ".", ";", ":")):
            heading_like += 1

    span = max(0, end - start)
    span_ratio = span / max(total, 1)
    if entry_like == 0:
        return True
    if heading_like >= 3 and entry_like == 0:
        return True
    if span_ratio >= 0.80 and (nonempty - entry_like) >= 2:
        return True
    return False


def _resolve_existing_toc_range(doc: Document, context: PipelineContext) -> tuple[int, int] | None:
    doc_tree = getattr(context, "doc_tree", None)
    if doc_tree is not None:
        toc_section = getattr(doc_tree, "get_section", lambda *_: None)("toc")
        if toc_section is not None:
            start = max(0, int(toc_section.start_index))
            end = min(len(doc.paragraphs), int(toc_section.end_index))
            if not _toc_section_is_suspicious(doc, start, end):
                return (start, end)
    return None


def _format_existing_toc_paragraphs(doc: Document, config: ResolvedConfig, context: PipelineContext) -> int:
    resolved_range = _resolve_existing_toc_range(doc, context)
    if resolved_range is None:
        return 0
    start, end = resolved_range

    formatted = 0
    for index in range(start, end):
        para = doc.paragraphs[index]
        raw = (para.text or "").strip()
        if not raw:
            continue
        if _is_toc_title_paragraph(para):
            style_config = resolve_toc_title_style_config(config)
            if style_config is not None:
                apply_toc_paragraph_style(para, style_config, is_title=True)
                try:
                    para.style = doc.styles[TOC_HEADING_STYLE_CANDIDATES[0]]
                except Exception:
                    pass
                formatted += 1
            continue

        if _is_toc_level_style_para(para) or looks_like_toc_entry_line(raw) or looks_like_numbered_toc_entry_with_page_suffix(raw):
            level = _infer_toc_level_from_paragraph(para)
            style_config = resolve_toc_entry_style_config(config, level)
            if style_config is None:
                continue
            apply_toc_paragraph_style(para, style_config, is_title=False)
            try:
                para.style = doc.styles[toc_word_style_for_heading_level(level)]
            except Exception:
                pass
            formatted += 1
    return formatted


def _format_inserted_toc_title(doc: Document, config: ResolvedConfig, insert_idx: int) -> None:
    if insert_idx < 0 or insert_idx >= len(doc.paragraphs):
        return
    para = doc.paragraphs[insert_idx]
    if not _is_toc_title_paragraph(para):
        return
    style_config = resolve_toc_title_style_config(config)
    if style_config is None:
        return
    apply_toc_paragraph_style(para, style_config, is_title=True)
    try:
        para.style = doc.styles[TOC_HEADING_STYLE_CANDIDATES[0]]
    except Exception:
        pass


def _find_insert_position(doc: Document, mode: str, context: PipelineContext) -> int | None:
    del mode
    doc_tree = getattr(context, "doc_tree", None)
    if doc_tree is None:
        return None
    resolver = getattr(doc_tree, "insertion_index_for_role", None)
    if not callable(resolver):
        return None
    position = resolver("toc")
    if position is None:
        return None
    return max(0, min(int(position), len(doc.paragraphs)))


def _build_text_paragraph_element(text: str):
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    paragraph.append(run)
    run.append(text_el)
    text_el.text = text
    text_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return paragraph


def _insert_paragraph_elements(doc: Document, position: int, elements: list) -> list:
    body = doc.element.body
    paragraphs = body.findall(qn("w:p"))
    if position < len(paragraphs):
        insert_at = list(body).index(paragraphs[position])
        for offset, elem in enumerate(elements):
            body.insert(insert_at + offset, elem)
    else:
        for elem in elements:
            body.append(elem)
    return elements


def _extract_trailing_section_break(paragraph_elements: list) -> object | None:
    for para in reversed(paragraph_elements):
        ppr = para.find(qn("w:pPr"))
        if ppr is None:
            continue
        sect_pr = ppr.find(qn("w:sectPr"))
        if sect_pr is not None:
            return clone_element(sect_pr)
    return None


def _apply_section_break_to_last_paragraph(elements: list, sect_pr) -> None:
    if not elements or sect_pr is None:
        return
    ppr = find_or_create(elements[-1], "w:pPr")
    existing = ppr.find(qn("w:sectPr"))
    if existing is not None:
        ppr.remove(existing)
    ppr.append(sect_pr)


def _insert_toc(doc: Document, position: int, max_level: int) -> None:
    title_para = _build_text_paragraph_element(TOC_TITLE_TEXT)
    field_para = OxmlElement("w:p")
    instr = build_toc_instruction(max_level=max_level)
    for elem in build_complex_field(instr, result_text=TOC_FIELD_PLACEHOLDER, mark_dirty=True):
        field_para.append(elem)
    _insert_paragraph_elements(doc, position, [title_para, field_para])


def _heading_included_in_toc(config: ResolvedConfig, level_num: int) -> bool:
    heading_numbering = getattr(config, "heading_numbering", None)
    level_bindings = getattr(heading_numbering, "level_bindings", None) or {}
    binding = level_bindings.get(f"heading{level_num}")
    if binding is None:
        return True
    return bool(getattr(binding, "include_in_toc", True))


def _collect_plain_toc_entries(
    doc: Document,
    config: ResolvedConfig,
    context: PipelineContext,
    *,
    max_level: int,
) -> list[dict]:
    doc_tree = getattr(context, "doc_tree", None)
    heading_map = getattr(context, "heading_map", None) or {}
    entries: list[dict] = []
    seen_para_indices: set[int] = set()

    def push(level: str, title: str, para_index: int) -> None:
        text = str(title or "").strip()
        if not text or para_index < 0 or para_index >= len(doc.paragraphs):
            return
        if para_index in seen_para_indices:
            return
        seen_para_indices.add(para_index)
        entries.append({"level": level, "title": text, "para_index": para_index})

    if doc_tree is not None:
        for sec_type, fallback_title in (("abstract_cn", "\u6458\u8981"), ("abstract_en", "Abstract")):
            section = getattr(doc_tree, "get_section", lambda *_: None)(sec_type)
            if section is None:
                continue
            para_index = int(getattr(section, "start_index", -1))
            title = (doc.paragraphs[para_index].text or "").strip() if 0 <= para_index < len(doc.paragraphs) else fallback_title
            push("heading1", title or fallback_title, para_index)

    for para_index in sorted(heading_map):
        level_num = int(heading_map[para_index])
        if level_num < 1 or level_num > max_level:
            continue
        if not _heading_included_in_toc(config, level_num):
            continue
        if doc_tree is not None:
            sec_type = getattr(doc_tree, "get_section_for_paragraph", lambda *_: "body")(para_index)
            if sec_type != "body":
                continue
        title = (doc.paragraphs[para_index].text or "").strip()
        if title:
            push(toc_heading_key(level_num), title, para_index)

    if doc_tree is not None:
        for sec_type, fallback_title in BACK_MATTER_TITLES.items():
            section = getattr(doc_tree, "get_section", lambda *_: None)(sec_type)
            if section is None:
                continue
            para_index = int(getattr(section, "start_index", -1))
            title = (doc.paragraphs[para_index].text or "").strip() if 0 <= para_index < len(doc.paragraphs) else fallback_title
            push("heading1", title or fallback_title, para_index)

    return entries


def _insert_plain_toc(doc: Document, position: int, entries: list[dict]) -> list:
    elements = [_build_text_paragraph_element(TOC_TITLE_TEXT)]
    for entry in entries:
        elements.append(_build_text_paragraph_element(str(entry.get("title", "") or "").strip()))
    return _insert_paragraph_elements(doc, position, elements)


def _replace_paragraph_toc_range_with_plain(doc: Document, start: int, end: int, entries: list[dict]) -> list:
    body = doc.element.body
    paragraphs = body.findall(qn("w:p"))
    if start < 0 or start >= len(paragraphs):
        return _insert_plain_toc(doc, len(doc.paragraphs), entries)

    existing_range = paragraphs[start:min(end, len(paragraphs))]
    trailing_sect_pr = _extract_trailing_section_break(existing_range)
    insert_at = list(body).index(paragraphs[start])
    for para in existing_range:
        body.remove(para)

    elements = [_build_text_paragraph_element(TOC_TITLE_TEXT)]
    for entry in entries:
        elements.append(_build_text_paragraph_element(str(entry.get("title", "") or "").strip()))
    _apply_section_break_to_last_paragraph(elements, trailing_sect_pr)
    for offset, elem in enumerate(elements):
        body.insert(insert_at + offset, elem)
    return elements


def _format_inserted_plain_toc(doc: Document, config: ResolvedConfig, inserted_elements: list, entries: list[dict]) -> None:
    el_to_para = {para._element: para for para in doc.paragraphs}

    title_para = el_to_para.get(inserted_elements[0]) if inserted_elements else None
    if title_para is not None:
        title_style = resolve_toc_title_style_config(config)
        if title_style is not None:
            apply_toc_paragraph_style(title_para, title_style, is_title=True)
            try:
                title_para.style = doc.styles[TOC_HEADING_STYLE_CANDIDATES[0]]
            except Exception:
                pass

    for elem, entry in zip(inserted_elements[1:], entries):
        para = el_to_para.get(elem)
        if para is None:
            continue
        level = toc_heading_level(entry.get("level", "heading6"))
        style_config = resolve_toc_entry_style_config(config, level)
        if style_config is None:
            continue
        apply_toc_paragraph_style(para, style_config, is_title=False)
        try:
            para.style = doc.styles[toc_word_style_for_heading_level(level)]
        except Exception:
            pass


def _is_toc_instruction(instr: str) -> bool:
    normalized = " ".join((instr or "").upper().split())
    return normalized.startswith("TOC ")
