"""Detect document headings and logical section ranges."""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.config.special_title_rules import match_special_title_model
from src.config.section_semantics import canonicalize_section_type
from src.modules.base import BaseModule, ModuleMeta
from src.shared.engine.docx_heading_semantics import get_paragraph_outline_level
from src.shared.engine.document_text_heuristics import (
    looks_like_date_placeholder_line,
    looks_like_numbered_toc_entry_with_page_suffix,
    looks_like_reference_entry_line,
    looks_like_toc_entry_line,
)
from src.shared.engine.document_structure_model import DocSection, DocTree, HeadingInfo
from src.shared.engine.style_resolver import get_heading_level

if TYPE_CHECKING:
    from docx import Document
    from docx.text.paragraph import Paragraph
    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


LEVEL_PATTERNS: list[tuple[int, re.Pattern[str]]] = [
    (8, re.compile(r"^\d+(?:\.\d+){7}\.?\s+")),
    (7, re.compile(r"^\d+(?:\.\d+){6}\.?\s+")),
    (6, re.compile(r"^\d+(?:\.\d+){5}\.?\s+")),
    (5, re.compile(r"^\d+(?:\.\d+){4}\.?\s+")),
    (4, re.compile(r"^\d+\.\d+\.\d+\.\d+[\s\u3000\t]")),
    (3, re.compile(r"^\d+\.\d+\.\d+[\s\u3000\t]")),
    (2, re.compile(r"^\d+\.\d+[\s\u3000\t]")),
    (1, re.compile(r"^\d+[\s\u3000\t]")),
    # 统一中文章节约定：第X(单元|部分|章|篇|部|卷|分)=顶级(H1)，第X节=章内小节(H2)
    (
        1,
        re.compile(
            r"^\u7b2c\s*[\u4e00-\u9fff\d]+\s*(?:\u5355\u5143|\u90e8\u5206|[\u7ae0\u7bc7\u90e8\u5377\u5206])[\s\u3000]"
        ),
    ),
    (
        2,
        re.compile(r"^\u7b2c\s*[\u4e00-\u9fff\d]+\s*\u8282[\s\u3000]"),
    ),
    (2, re.compile(r"^[\u4e00-\u9fff]{1,4}[、.)）][\s\u3000]?")),
    (3, re.compile(r"^[（(][\u4e00-\u9fff]{1,4}[)）]")),
    (4, re.compile(r"^\d+[)）][\s\u3000\t]")),
    (1, re.compile(r"^[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ][、.\s\u3000]")),
    (3, re.compile(r"^[A-Z][、.)）][\s\u3000]?")),
    (4, re.compile(r"^[a-z][、.)）][\s\u3000]?")),
]

BODY_START_PATTERNS: list[re.Pattern[str]] = [
    re.compile(r"^#{1,6}\s+"),
    re.compile(
        r"^\u7b2c\s*[\u4e00-\u9fff\d]+\s*(?:\u5355\u5143|\u90e8\u5206|[\u7ae0\u7bc7\u90e8\u5377\u5206\u8282])"
    ),
    re.compile(r"^\d+[\.、．]\S"),
    re.compile(r"^\d+\s+\S"),
]

NARRATIVE_PUNCT = re.compile(r"[。；;！？!?]")
SENTENCE_END_PUNCT = {"。", ".", "！", "!", "？", "?", "；", ";", "：", ":"}
NUMERIC_CHAIN_HEADING_RE = re.compile(r"^(\d+(?:\.\d+){1,7})\.?[\s\u3000\t]+\S")
TOC_LEVEL_STYLE_RE = re.compile(r"^(toc|\u76ee\u5f55)\s*\d+$", re.IGNORECASE)
TOC_TITLE_RE = re.compile(r"^(\u76ee\u5f55|\u76ee\s*\u5f55|contents|tableofcontents)$", re.IGNORECASE)
BROKEN_TOC_BOOKMARK_HINTS = (
    "error! bookmark not defined",
    "bookmark not defined",
    "\u672a\u5b9a\u4e49\u4e66\u7b7e",
)
CHAPTER_REF_RE = re.compile(r"\u7b2c[\u4e00-\u9fff\d]+\u7ae0")
SECTION_REF_RE = re.compile(r"\u7b2c[\u4e00-\u9fff\d]+\u8282")
TABULAR_NUMERIC_TOKEN_RE = re.compile(r"^[+\-]?\d+(?:\.\d+)?(?:[eE][+\-]?\d+)?%?$")
TABULAR_CITATION_TOKEN_RE = re.compile(r"^\[\d{1,4}\]$")
INLINE_ABSTRACT_CN_RE = re.compile(r"^\s*\u6458\u8981\s*[\uff1a:]\s*\S")
INLINE_ABSTRACT_EN_RE = re.compile(r"^\s*abstract\s*[\uff1a:]\s*\S", re.IGNORECASE)
SHORT_TABULAR_LITERAL_TOKENS = {
    "none",
    "o2",
    "n2",
    "air",
    "this work",
    "ref",
    "ref.",
    "xe lamp",
    "simulated sunlight",
    "seawater",
    "ethanol",
}

SPECIAL_SECTION_TITLES: dict[str, str] = {
    "\u5b66\u4f4d\u8bba\u6587": "cover",
    "\u535a\u58eb\u5b66\u4f4d\u8bba\u6587": "cover",
    "\u7855\u58eb\u5b66\u4f4d\u8bba\u6587": "cover",
    "\u672c\u79d1\u6bd5\u4e1a\u8bba\u6587": "cover",
    "\u6bd5\u4e1a\u8bba\u6587": "cover",
    "\u672c\u79d1\u6bd5\u4e1a\u8bbe\u8ba1": "cover",
    "\u6bd5\u4e1a\u8bbe\u8ba1": "cover",
    "\u53c2\u8003\u6587\u732e": "references",
    "references": "references",
    "bibliography": "references",
    "\u6458\u8981": "abstract_cn",
    "abstract": "abstract_en",
    "\u81f4\u8c22": "acknowledgment",
    "acknowledgement": "acknowledgment",
    "acknowledgments": "acknowledgment",
    "acknowledgements": "acknowledgment",
    "\u9644\u5f55": "appendix",
    "appendix": "appendix",
    "\u76ee\u5f55": "toc",
    "\u7eea\u8bba": "body",
    "\u5f15\u8a00": "body",
    "\u52d8\u8bef": "errata",
    "\u52d8\u8bef\u9875": "errata",
    "\u4e2a\u4eba\u7b80\u5386": "resume",
    "\u5728\u5b66\u671f\u95f4\u53d1\u8868\u7684\u5b66\u672f\u8bba\u6587\u4e0e\u7814\u7a76\u6210\u679c": "resume",
}

SECTION_ANCHORS: dict[str, tuple[str, ...]] = {
    "cover": (
        "\u5b66\u4f4d\u8bba\u6587",
        "\u535a\u58eb\u5b66\u4f4d",
        "\u7855\u58eb\u5b66\u4f4d",
        "\u672c\u79d1\u6bd5\u4e1a\u8bba\u6587",
        "\u6bd5\u4e1a\u8bba\u6587",
        "\u672c\u79d1\u6bd5\u4e1a\u8bbe\u8ba1",
        "\u6bd5\u4e1a\u8bbe\u8ba1",
    ),
    "abstract_cn": ("\u6458\u8981", "\u6458 \u8981"),
    "abstract_en": ("abstract",),
    "toc": ("\u76ee\u5f55", "\u76ee \u5f55", "contents", "table of contents"),
    "references": ("\u53c2\u8003\u6587\u732e", "references", "bibliography"),
    "errata": ("\u52d8\u8bef\u9875", "\u52d8\u8bef"),
    "appendix": ("\u9644\u5f55", "appendix"),
    "acknowledgment": ("\u81f4\u8c22", "acknowledgement", "acknowledgments", "acknowledgements"),
    "resume": (
        "\u4e2a\u4eba\u7b80\u5386",
        "\u7b80\u5386",
        "\u5728\u5b66\u671f\u95f4\u53d1\u8868\u7684\u5b66\u672f\u8bba\u6587\u4e0e\u7814\u7a76\u6210\u679c",
    ),
}

SECTION_ORDER = [
    "cover",
    "abstract_cn",
    "abstract_en",
    "toc",
    "body",
    "references",
    "errata",
    "appendix",
    "acknowledgment",
    "resume",
]

SECTION_MIN_ACCEPT_SCORE = {
    "cover": 8.0,
    "abstract_cn": 8.0,
    "abstract_en": 8.0,
    "toc": 8.0,
    "references": 8.0,
    "errata": 8.0,
    "appendix": 8.0,
    "acknowledgment": 8.0,
    "resume": 8.0,
}
SECTION_TITLE_CONFIDENCE = {key: 10.0 for key in SECTION_MIN_ACCEPT_SCORE}
PRE_BODY_TYPES = {"cover", "abstract_cn", "abstract_en", "toc"}
POST_BODY_TYPES = {"references", "errata", "appendix", "acknowledgment", "resume"}
POST_BODY_MIN_RATIO = 0.35
PRE_BODY_MAX_RATIO = 0.70
HIGH_CONF_OVERRIDE = 12.0
MIN_SECTION_SCORE = 4.0
GENERIC_COVER_DATE_RE = re.compile(
    r"(?:[二〇零○一一二三四五六七八九十\d]{4}\s*年(?:\s*[一二三四五六七八九十\d]{1,3}\s*月)?|20\d{2}\s*年)"
)


@dataclass
class _SectionCandidate:
    para_index: int
    section_type: str
    score: float


class HeadingRecognitionModule(BaseModule):
    meta = ModuleMeta(
        name="heading_recognition",
        description="\u6807\u9898\u8bc6\u522b",
        category="structure",
        requires_config=(),
        provides=("doc_tree", "heading_map"),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        doc_tree = rebuild_document_index(doc, context, config)
        headings = doc_tree.headings

        if headings:
            counts: dict[int, int] = {}
            for heading in headings:
                counts[heading.level] = counts.get(heading.level, 0) + 1
            tracker.record(
                rule_name=self.meta.name,
                target=f"{len(headings)} headings",
                section="global",
                change_type="detect",
                before="(none)",
                after=", ".join(f"H{level}={count}" for level, count in sorted(counts.items())),
            )


def rebuild_document_index(
    doc: Document,
    context: PipelineContext,
    config: ResolvedConfig | None = None,
) -> DocTree:
    """Rebuild the shared heading/section index after a structural mutation.

    The pipeline may call this between modules without replaying the heading
    recognition tracker event.  Keeping the index construction here prevents
    later consumers from observing stale paragraph indices after content,
    caption, TOC, section, or image mutations.
    """

    if config is not None:
        context.heading_model_config = getattr(config, "heading_model", None)
    doc_tree = analyze_document_tree(
        doc,
        heading_model=getattr(context, "heading_model_config", None),
    )
    from src.shared.engine.document_scope_runtime import project_document_scope_tree

    doc_tree = project_document_scope_tree(
        doc,
        doc_tree,
        getattr(context, "document_scope_binding", None),
        getattr(context, "document_scope", None),
        mode_id=str(getattr(context, "mode_id", "") or "custom"),
    )
    context.doc_tree = doc_tree
    context.heading_map = dict(doc_tree.heading_map)
    return doc_tree


def analyze_document_tree(
    doc: Document,
    heading_model=None,
) -> DocTree:
    """Return the canonical read-only heading and logical-region analysis."""

    candidate_headings, special_title_matches = _scan_heading_infos(
        doc,
        heading_model,
    )
    sections, detection_log = _build_sections(doc, candidate_headings)
    headings = _filter_body_headings(candidate_headings, sections)
    heading_map = {heading.para_index: heading.level for heading in headings}
    section_ranges = {
        section.section_type: (section.start_index, section.end_index)
        for section in sections
    }
    special_title_ranges = _build_special_title_ranges(
        len(doc.paragraphs),
        candidate_headings,
        special_title_matches,
    )
    return DocTree(
        headings=headings,
        heading_map=heading_map,
        section_ranges=section_ranges,
        sections=sections,
        special_title_matches=special_title_matches,
        special_title_ranges=special_title_ranges,
        detection_log=detection_log,
    )


def _scan_heading_infos(
    doc: Document,
    heading_model=None,
) -> tuple[list[HeadingInfo], dict[int, str]]:
    headings: list[HeadingInfo] = []
    special_title_matches: dict[int, str] = {}
    previous_level: int | None = None
    for index, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if not text:
            continue
        special_match = match_special_title_model(text, heading_model)
        if special_match is not None:
            special_title_matches[index] = special_match.selector
        level = 1 if special_match is not None else _detect_heading(para, text)
        if level is None:
            level = _detect_heading_from_context(para, text, previous_level)
        if level is not None:
            headings.append(HeadingInfo(index, level, text[:80]))
            previous_level = level
    return headings, special_title_matches


def _build_special_title_ranges(
    total_paragraphs: int,
    headings: list[HeadingInfo],
    matches: dict[int, str],
) -> list[DocSection]:
    """End each matched range at the next special title or ordinary H1."""

    level_one_indices = sorted(
        heading.para_index
        for heading in headings
        if heading.level == 1
    )
    ranges: list[DocSection] = []
    for start_index, selector in sorted(matches.items()):
        end_index = next(
            (
                heading_index
                for heading_index in level_one_indices
                if heading_index > start_index
            ),
            total_paragraphs,
        )
        if end_index <= start_index:
            continue
        ranges.append(
            DocSection(
                section_type=selector,
                start_index=start_index,
                end_index=end_index,
                confidence=10.0,
                title_confident=True,
            )
        )
    return ranges


def _section_type_for_paragraph(sections: list[DocSection], para_index: int) -> str:
    for section in sections:
        if section.start_index <= para_index < section.end_index:
            return canonicalize_section_type(section.section_type)
    return "body"


def _filter_body_headings(headings: list[HeadingInfo], sections: list[DocSection]) -> list[HeadingInfo]:
    if not sections:
        return list(headings)
    return [
        heading
        for heading in headings
        if _section_type_for_paragraph(sections, heading.para_index) == "body"
    ]


def _norm_no_space(text: str) -> str:
    return re.sub(r"\s+", "", text or "").strip()


def _is_toc_style_para(para: Paragraph) -> bool:
    style = para.style
    style_name = ((style.name if style else "") or "").strip()
    style_id = ((getattr(style, "style_id", "") if style else "") or "").strip()
    return bool(TOC_LEVEL_STYLE_RE.match(style_name) or TOC_LEVEL_STYLE_RE.match(style_id))


def _para_has_pageref_field(para: Paragraph) -> bool:
    from src.shared.engine.field_builder import iter_field_instructions

    for _kind, _elem, instr in iter_field_instructions(para._element):
        normalized = " ".join((instr or "").upper().split())
        if normalized.startswith("PAGEREF"):
            return True
    return False


def _looks_like_tabular_numeric_line(text: str) -> bool:
    raw = (text or "").strip()
    if "\t" not in raw:
        return False
    tokens = [token.strip() for token in raw.split("\t") if token.strip()]
    if len(tokens) < 2:
        return False
    numeric_like = 0
    for token in tokens:
        if re.fullmatch(r"[+\-]?\d+(?:\.\d+)?(?:[eE][+\-]?\d+)?%?", token) or re.fullmatch(r"\d+:\d+", token):
            numeric_like += 1
    return numeric_like >= 2 and not re.search(r"[A-Za-z\u4e00-\u9fff]", "".join(tokens))


def _looks_like_tabular_structured_line(text: str) -> bool:
    raw = (text or "").strip()
    if "\t" not in raw:
        return False
    tokens = [token.strip() for token in raw.split("\t") if token.strip()]
    if len(tokens) < 2:
        return False
    if any(len(token) > 80 for token in tokens):
        return False
    if any(NARRATIVE_PUNCT.search(token) for token in tokens):
        return False

    structured = 0
    numericish = 0
    shortish = 0
    for token in tokens:
        if len(token) <= 24:
            shortish += 1
        lowered = token.lower()
        if TABULAR_NUMERIC_TOKEN_RE.fullmatch(token) or re.fullmatch(r"\d+:\d+", token):
            structured += 1
            numericish += 1
            continue
        if TABULAR_CITATION_TOKEN_RE.fullmatch(token):
            structured += 1
            continue
        if lowered in SHORT_TABULAR_LITERAL_TOKENS:
            structured += 1
            continue
        has_digit = bool(re.search(r"\d", token))
        has_unitish = bool(re.search(r"[A-Za-z\u00b5\u03bc\u03a9\u03c9\u03bb\u039b\u00b0\u2103/%<>=\-−]", token))
        if has_digit and has_unitish and len(token) <= 32:
            structured += 1
            numericish += 1

    if numericish <= 0:
        return False
    if structured >= max(2, len(tokens) - 1):
        return True
    if len(tokens) >= 3 and shortish >= len(tokens) - 1 and structured >= 2:
        return True
    return False


def _looks_like_broken_toc_entry_line(text: str) -> bool:
    raw = (text or "").strip()
    if not raw or "\t" not in raw:
        return False
    lower = raw.lower()
    return any(hint in lower for hint in BROKEN_TOC_BOOKMARK_HINTS) or (
        "\u4e66\u7b7e" in lower and "\u672a\u5b9a\u4e49" in lower
    )


def _looks_like_chapter_outline_sentence(text: str) -> bool:
    raw = (text or "").strip()
    if not raw:
        return False
    total_hits = len(CHAPTER_REF_RE.findall(raw)) + len(SECTION_REF_RE.findall(raw))
    if total_hits >= 2:
        return True
    if total_hits == 0 or len(raw) < 24 or not NARRATIVE_PUNCT.search(raw):
        return False
    sep_count = raw.count("，") + raw.count(",") + raw.count("；") + raw.count(";") + raw.count("。") + raw.count("、")
    return sep_count >= 2


def _has_heading_visual_traits(para: Paragraph) -> bool:
    from docx.shared import Pt

    if get_paragraph_outline_level(para) is not None:
        return True
    if para.runs:
        first_run = para.runs[0]
        if first_run.bold:
            return True
        if first_run.font.size and first_run.font.size >= Pt(14):
            return True
    return False


def _detect_by_pattern(text: str) -> int | None:
    for level, pattern in LEVEL_PATTERNS:
        if pattern.match(text):
            return level
    return None


def _detect_numeric_chain_level(text: str) -> int | None:
    match = NUMERIC_CHAIN_HEADING_RE.match((text or "").strip())
    if not match:
        return None
    level = match.group(1).count(".") + 1
    return level if 2 <= level <= 8 else None


def _is_disqualified_heading_candidate(para: Paragraph, raw: str) -> bool:
    if not raw:
        return True
    if len(raw) > 80 and NARRATIVE_PUNCT.search(raw):
        return True
    if _is_toc_style_para(para):
        return True
    if looks_like_toc_entry_line(raw) or looks_like_numbered_toc_entry_with_page_suffix(raw):
        return True
    if _para_has_pageref_field(para):
        return True
    if looks_like_reference_entry_line(raw) or looks_like_date_placeholder_line(raw):
        return True
    if _looks_like_broken_toc_entry_line(raw):
        return True
    if _looks_like_tabular_numeric_line(raw) or _looks_like_tabular_structured_line(raw):
        return True
    if _looks_like_chapter_outline_sentence(raw):
        return True
    return False


def _detect_heading_from_context(para: Paragraph, text: str, previous_level: int | None) -> int | None:
    raw = (text or "").strip()
    if previous_level is None or _is_disqualified_heading_candidate(para, raw):
        return None
    if raw[-1:] in SENTENCE_END_PUNCT:
        return None
    level = _detect_numeric_chain_level(raw)
    if level is None:
        return None
    if level > previous_level + 1:
        return None
    return level


def _detect_special_section(text: str) -> str | None:
    normalized = _norm_no_space(text).lower()
    if not normalized:
        return None
    for title, section in SPECIAL_SECTION_TITLES.items():
        if normalized == _norm_no_space(title).lower():
            return canonicalize_section_type(section)
    if normalized.startswith("\u9644\u5f55") or normalized.startswith("appendix"):
        return "appendix"
    return None


def _detect_heading(para: Paragraph, text: str) -> int | None:
    raw = (text or "").strip()
    if not raw:
        return None
    if len(raw) > 80 and NARRATIVE_PUNCT.search(raw):
        return None
    if _is_toc_style_para(para):
        return None
    if looks_like_toc_entry_line(raw) or looks_like_numbered_toc_entry_with_page_suffix(raw):
        return None
    if _para_has_pageref_field(para):
        return None
    if looks_like_reference_entry_line(raw) or looks_like_date_placeholder_line(raw):
        return None
    if _looks_like_broken_toc_entry_line(raw):
        return None
    if _looks_like_tabular_numeric_line(raw) or _looks_like_tabular_structured_line(raw):
        return None
    if _looks_like_chapter_outline_sentence(raw):
        return None

    style_level = get_heading_level(para)
    if style_level is not None:
        return style_level

    outline_level = get_paragraph_outline_level(para)
    if outline_level is not None and 0 <= outline_level <= 8 and 4 <= len(raw) <= 80:
        return outline_level + 1

    special_section = _detect_special_section(raw)
    pattern_level = _detect_by_pattern(raw)
    if special_section is not None and len(raw) <= 30:
        return 1
    if pattern_level is not None and len(raw) <= 60 and _has_heading_visual_traits(para):
        return pattern_level
    return None


def _score_section_anchor(para: Paragraph, section_type: str, anchors: tuple[str, ...], total: int, para_index: int) -> float:
    text = (para.text or "").strip()
    if not text:
        return 0.0
    ratio = para_index / max(total, 1)
    if section_type == "abstract_cn" and ratio < 0.35 and len(text) > 20 and INLINE_ABSTRACT_CN_RE.match(text):
        return 9.0
    if section_type == "abstract_en" and ratio < 0.45 and len(text) > 20 and INLINE_ABSTRACT_EN_RE.match(text):
        return 9.0
    if len(text) > 80:
        return 0.0
    if section_type != "toc" and (looks_like_toc_entry_line(text) or looks_like_numbered_toc_entry_with_page_suffix(text)):
        return 0.0
    if section_type != "toc" and _para_has_pageref_field(para):
        return 0.0
    if section_type != "references" and looks_like_reference_entry_line(text):
        return 0.0

    text_no_space = _norm_no_space(text).lower()
    best = 0.0
    matched = ""
    for anchor in anchors:
        anchor_no_space = _norm_no_space(anchor).lower()
        local = 0.0
        if text_no_space == anchor_no_space:
            local = 10.0
        elif text_no_space.startswith(anchor_no_space):
            local = 6.0
            if section_type == "appendix" and re.match(
                r"^(?:\u9644\u5f55[A-Za-z\uff21-\uff3a\uff41-\uff5a0-9\u4e00\u4e8c\u4e09\u56db\u4e94\u516d\u4e03\u516b\u4e5d\u5341]|appendix[A-Za-z0-9])",
                text_no_space,
                re.IGNORECASE,
            ):
                local = 10.0
        elif anchor_no_space and anchor_no_space in text_no_space:
            local = 2.0
        if local > best:
            best = local
            matched = anchor
    if not matched:
        return 0.0

    score = best
    if get_heading_level(para) is not None or get_paragraph_outline_level(para) is not None:
        score += 5.0
    elif _has_heading_visual_traits(para):
        score += 3.0
    if len(text) <= len(matched) + 4 or (section_type == "appendix" and best >= 10.0 and len(text) <= 24):
        score += 3.0
    if score >= 6.0:
        if section_type == "cover" and ratio < 0.15:
            score += 2.0
        elif section_type in {"abstract_cn", "abstract_en"} and ratio < 0.35:
            score += 2.0
        elif section_type == "toc" and ratio < 0.60:
            score += 2.0
        elif section_type in (POST_BODY_TYPES - {"appendix"}) and ratio > 0.50:
            score += 2.0
    return score


def _is_candidate_credible(candidate: _SectionCandidate, total: int) -> bool:
    if candidate.score < SECTION_MIN_ACCEPT_SCORE.get(candidate.section_type, MIN_SECTION_SCORE):
        return False
    ratio = candidate.para_index / max(total, 1)
    if candidate.section_type in POST_BODY_TYPES and ratio < POST_BODY_MIN_RATIO and candidate.score < HIGH_CONF_OVERRIDE:
        return False
    if candidate.section_type in PRE_BODY_TYPES and ratio > PRE_BODY_MAX_RATIO and candidate.score < HIGH_CONF_OVERRIDE:
        return False
    return True


def _scan_section_candidates(
    doc: Document,
    total: int,
) -> list[_SectionCandidate]:
    candidates: list[_SectionCandidate] = []
    for index, para in enumerate(doc.paragraphs):
        if not (para.text or "").strip():
            continue
        for section_type, anchors in SECTION_ANCHORS.items():
            score = _score_section_anchor(para, section_type, anchors, total, index)
            if score >= MIN_SECTION_SCORE:
                candidates.append(_SectionCandidate(index, section_type, score))
    return candidates


def _pick_best_section_candidates(candidates: list[_SectionCandidate], total: int, detection_log: list[str]) -> dict[str, DocSection]:
    best: dict[str, _SectionCandidate] = {}
    for candidate in candidates:
        if not _is_candidate_credible(candidate, total):
            detection_log.append(f"drop weak section candidate: {candidate.section_type}@{candidate.para_index}")
            continue
        current = best.get(candidate.section_type)
        if current is None or candidate.score > current.score:
            best[candidate.section_type] = candidate
            continue
        if candidate.score == current.score:
            if candidate.section_type in PRE_BODY_TYPES and candidate.para_index < current.para_index:
                best[candidate.section_type] = candidate
            elif candidate.section_type == "appendix" and candidate.para_index < current.para_index:
                best[candidate.section_type] = candidate
            elif (
                candidate.section_type in POST_BODY_TYPES
                and candidate.section_type != "appendix"
                and candidate.para_index > current.para_index
            ):
                best[candidate.section_type] = candidate
    return {
        section_type: DocSection(
            section_type=section_type,
            start_index=candidate.para_index,
            confidence=candidate.score,
            title_confident=candidate.score >= SECTION_TITLE_CONFIDENCE.get(section_type, 10.0),
        )
        for section_type, candidate in best.items()
    }


def _order_compatible(prev_type: str, cur_type: str, order_rank: dict[str, int]) -> bool:
    if prev_type == cur_type:
        return False
    prev_is_pre = prev_type in PRE_BODY_TYPES
    cur_is_pre = cur_type in PRE_BODY_TYPES
    if prev_is_pre and cur_is_pre:
        return True
    if prev_is_pre and not cur_is_pre:
        return True
    if not prev_is_pre and cur_is_pre:
        return False
    prev_is_post = prev_type in POST_BODY_TYPES
    cur_is_post = cur_type in POST_BODY_TYPES
    if prev_is_post and cur_is_post:
        return True
    if prev_is_post and not cur_is_post:
        return False
    if not prev_is_post and cur_is_post:
        return True
    return order_rank.get(prev_type, 10_000) < order_rank.get(cur_type, 10_000)


def _order_validate_sections(
    sections_by_type: dict[str, DocSection],
    detection_log: list[str],
) -> list[DocSection]:
    if "cover" in sections_by_type:
        sections_by_type["cover"].start_index = 0
    sorted_sections = sorted(sections_by_type.values(), key=lambda section: section.start_index)
    if not sorted_sections:
        return []
    order_rank = {name: index for index, name in enumerate(SECTION_ORDER)}
    size = len(sorted_sections)
    scores = [sorted_sections[i].confidence for i in range(size)]
    previous = [-1] * size
    for right in range(size):
        for left in range(right):
            if not _order_compatible(sorted_sections[left].section_type, sorted_sections[right].section_type, order_rank):
                continue
            candidate_score = scores[left] + sorted_sections[right].confidence
            if candidate_score > scores[right]:
                scores[right] = candidate_score
                previous[right] = left
    best_index = max(range(size), key=lambda index: scores[index])
    picked: list[int] = []
    cursor = best_index
    while cursor != -1:
        picked.append(cursor)
        cursor = previous[cursor]
    picked.reverse()
    validated = [sorted_sections[index] for index in picked]
    keep = {(section.section_type, section.start_index) for section in validated}
    for section in sorted_sections:
        if (section.section_type, section.start_index) not in keep:
            detection_log.append(f"drop order-conflicting section: {section.section_type}@{section.start_index}")
    return validated


def _detect_toc_by_style(doc: Document, total: int) -> DocSection | None:
    start = None
    end = None
    for index, para in enumerate(doc.paragraphs):
        raw = (para.text or "").strip()
        if _is_toc_style_para(para):
            if start is None:
                start = index
            end = index + 1
            continue
        if start is None:
            continue
        if not raw:
            end = index + 1
            continue
        if looks_like_toc_entry_line(raw) or looks_like_numbered_toc_entry_with_page_suffix(raw):
            end = index + 1
            continue
        break
    if start is not None and end is not None and end > start:
        if start > 0 and TOC_TITLE_RE.match(_norm_no_space(doc.paragraphs[start - 1].text or "")):
            start -= 1
        return DocSection("toc", start, min(end, total), confidence=8.0, title_confident=True)
    return None


def _detect_toc_by_title_and_entries(doc: Document, total: int) -> DocSection | None:
    for index, para in enumerate(doc.paragraphs):
        if not TOC_TITLE_RE.match(_norm_no_space(para.text)):
            continue
        end = index + 1
        entry_count = 0
        for probe in range(index + 1, total):
            raw = (doc.paragraphs[probe].text or "").strip()
            if not raw:
                end = probe + 1
                continue
            if _is_toc_style_para(doc.paragraphs[probe]) or looks_like_toc_entry_line(raw) or looks_like_numbered_toc_entry_with_page_suffix(raw):
                end = probe + 1
                entry_count += 1
                continue
            break
        if entry_count >= 2:
            return DocSection("toc", index, min(end, total), confidence=7.5, title_confident=True)
    return None


def _detect_references_by_content(doc: Document, total: int) -> DocSection | None:
    if total <= 0:
        return None
    start_search = max(0, int(total * 0.35))
    cluster_start = None
    cluster_size = 0
    best_start = None
    best_size = 0
    for index in range(start_search, total):
        raw = (doc.paragraphs[index].text or "").strip()
        if looks_like_reference_entry_line(raw):
            if cluster_start is None:
                cluster_start = index
            cluster_size += 1
            if cluster_size > best_size:
                best_size = cluster_size
                best_start = cluster_start
            continue
        if raw:
            cluster_start = None
            cluster_size = 0
    if best_start is not None and best_size >= 2:
        return DocSection("references", best_start, total, confidence=7.5, title_confident=False)
    return None


def _detect_cover_by_content(doc: Document, total: int) -> DocSection | None:
    for index in range(min(total, 12)):
        raw = (doc.paragraphs[index].text or "").strip()
        if not raw or len(raw) > 40 or NARRATIVE_PUNCT.search(raw):
            continue
        if INLINE_ABSTRACT_CN_RE.match(raw) or INLINE_ABSTRACT_EN_RE.match(raw):
            continue
        if any(keyword in raw for keyword in SECTION_ANCHORS["cover"]):
            return DocSection("cover", 0, 1, confidence=8.0, title_confident=False)
    return None


def _detect_cover_by_first_section(doc: Document, total: int) -> DocSection | None:
    """Recognize a sparse, centered first page terminated by a page section break."""

    from src.shared.engine.ooxml_ops import qn

    if total < 4:
        return None
    boundary = None
    for index, paragraph in enumerate(doc.paragraphs[:-1]):
        paragraph_properties = paragraph._p.pPr
        section_properties = (
            paragraph_properties.find(qn("w:sectPr"))
            if paragraph_properties is not None
            else None
        )
        if section_properties is None:
            continue
        section_type = section_properties.find(qn("w:type"))
        type_value = (
            str(section_type.get(qn("w:val")) or "").strip()
            if section_type is not None
            else "nextPage"
        )
        if type_value in {"nextPage", "oddPage", "evenPage"}:
            boundary = index + 1
        break
    if boundary is None:
        return None

    early_limit = min(80, max(30, int(total * 0.35)))
    if boundary < 3 or boundary > early_limit:
        return None
    nonempty = [
        (index, paragraph, str(paragraph.text or "").strip())
        for index, paragraph in enumerate(doc.paragraphs[:boundary])
        if str(paragraph.text or "").strip()
    ]
    if not 1 <= len(nonempty) <= 12 or len(nonempty) / boundary > 0.55:
        return None

    centered = sum(
        paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
        for _index, paragraph, _text in nonempty
    )
    has_title = any(
        4 <= len(text) <= 80
        and not NARRATIVE_PUNCT.search(text)
        and not looks_like_date_placeholder_line(text)
        for _index, _paragraph, text in nonempty
    )
    has_date = any(GENERIC_COVER_DATE_RE.search(text) for _index, _paragraph, text in nonempty)
    if centered < 2 or not has_title or not has_date:
        return None
    return DocSection(
        "cover",
        0,
        boundary,
        confidence=9.0,
        title_confident=True,
        boundary_confident=True,
    )


def _scan_body_start_range(
    doc: Document,
    headings: list[HeadingInfo],
    start: int,
    end: int,
) -> int | None:
    for heading in headings:
        if not (start <= heading.para_index < end):
            continue
        section_type = _detect_special_section(heading.text)
        if section_type in {None, "body"}:
            return heading.para_index
    for index in range(start, min(end, len(doc.paragraphs))):
        para = doc.paragraphs[index]
        raw = (para.text or "").strip()
        if not raw:
            continue
        if _is_toc_style_para(para):
            continue
        if looks_like_toc_entry_line(raw) or looks_like_numbered_toc_entry_with_page_suffix(raw):
            continue
        if looks_like_reference_entry_line(raw) or looks_like_date_placeholder_line(raw):
            continue
        if TOC_TITLE_RE.match(_norm_no_space(raw)):
            continue
        style_level = get_heading_level(para)
        if style_level is not None and style_level <= 2:
            return index
        outline_level = get_paragraph_outline_level(para)
        if outline_level in (0, 1) and raw[-1:] not in {"：", ":", "；", ";", "。", "!", "！", "?", "？"}:
            return index
        if any(pattern.match(raw) for pattern in BODY_START_PATTERNS):
            return index
    return None


def _insert_body_section(
    doc: Document,
    headings: list[HeadingInfo],
    special_sections: list[DocSection],
    total: int,
    detection_log: list[str],
) -> list[DocSection]:
    if not special_sections:
        return [DocSection("body", 0, total, confidence=10.0)] if total > 0 else []

    sections = sorted(special_sections, key=lambda section: section.start_index)
    pre_sections = [section for section in sections if section.section_type in PRE_BODY_TYPES]
    post_sections = [section for section in sections if section.section_type in POST_BODY_TYPES]

    body_scan_start = max(
        (
            section.end_index
            if section.boundary_confident and section.end_index > section.start_index
            else section.start_index + 1
            for section in pre_sections
        ),
        default=0,
    )
    first_post = next((section for section in post_sections if section.start_index >= body_scan_start), None)
    body_scan_end = first_post.start_index if first_post is not None else total

    body_start = _scan_body_start_range(doc, headings, body_scan_start, body_scan_end)
    if body_start is None:
        body_start = body_scan_start

    if pre_sections:
        last_pre = max(pre_sections, key=lambda section: section.start_index)
        if last_pre.start_index < body_start:
            last_pre.end_index = min(last_pre.end_index, body_start)

    body_end = first_post.start_index if first_post is not None else total
    if body_start < body_end:
        sections.append(DocSection("body", body_start, body_end, confidence=10.0))
        detection_log.append(f"insert body section: [{body_start}, {body_end})")
    return sorted([section for section in sections if section.end_index > section.start_index], key=lambda section: section.start_index)


def _build_sections(
    doc: Document,
    headings: list[HeadingInfo],
) -> tuple[list[DocSection], list[str]]:
    total = len(doc.paragraphs)
    detection_log: list[str] = []

    candidates = _scan_section_candidates(doc, total)
    anchors = _pick_best_section_candidates(candidates, total, detection_log)

    cover_section = _detect_cover_by_first_section(doc, total) or _detect_cover_by_content(doc, total)
    if "cover" not in anchors and cover_section is not None:
        anchors["cover"] = cover_section
        detection_log.append(f"fallback cover by content: @{cover_section.start_index}")

    toc_section = _detect_toc_by_style(doc, total) or _detect_toc_by_title_and_entries(doc, total)
    if "toc" not in anchors and toc_section is not None:
        anchors["toc"] = toc_section
        detection_log.append(f"fallback toc: [{toc_section.start_index}, {toc_section.end_index})")

    references_section = _detect_references_by_content(doc, total)
    current_references = anchors.get("references")
    if references_section is not None and current_references is None:
        anchors["references"] = references_section
        detection_log.append(f"fallback references by cluster: @{references_section.start_index}")
    elif (
        references_section is not None
        and current_references is not None
        and (
            not current_references.title_confident
            or (
                current_references.start_index / max(total, 1) < POST_BODY_MIN_RATIO
                and references_section.start_index > current_references.start_index
            )
        )
    ):
        anchors["references"] = references_section
        detection_log.append(
            f"replace weak references anchor: {current_references.start_index} -> {references_section.start_index}"
        )

    ordered = _order_validate_sections(anchors, detection_log)
    for index, section in enumerate(ordered):
        next_start = ordered[index + 1].start_index if index + 1 < len(ordered) else total
        if section.boundary_confident and section.end_index > section.start_index:
            section.end_index = min(section.end_index, next_start)
        else:
            section.end_index = next_start
    finalized = [section for section in ordered if section.end_index > section.start_index]

    sections = _insert_body_section(doc, headings, finalized, total, detection_log)
    if not sections and total > 0:
        sections = [DocSection("body", 0, total, confidence=10.0)]
    return sections, detection_log


def _has_list_numpr(para: Paragraph) -> bool:
    """Check if paragraph has w:numPr (Word list numbering)."""
    from src.shared.engine.ooxml_ops import qn

    p_pr = para._element.find(qn("w:pPr"))
    if p_pr is None:
        return False
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        return False
    num_id = num_pr.find(qn("w:numId"))
    return num_id is not None and (num_id.get(qn("w:val")) or "0") != "0"
