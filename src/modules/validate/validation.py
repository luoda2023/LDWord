"""validation — lightweight document validation module."""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from src.config.special_title_rules import match_special_title
from src.config.style_semantics import resolve_style_size_pt
from src.modules.base import BaseModule, Issue, ModuleMeta
from src.modules.basic.page_setup import (
    PAPER_SIZES,
    _desired_landscape,
    _desired_margin_config,
    _desired_paper_key,
    _section_is_landscape,
    _section_orientation_contradiction,
    _section_start_indices,
)
from src.modules.special.equation_table_format import (
    _find_number_paragraph,
    _has_explicit_equation_table_marker,
    _row_has_formula_content,
)
from src.modules.special.formula_convert import paragraph_has_formula_source
from src.modules.structure.heading_numbering import has_literal_heading_number_prefix
from src.modules.table.table_format import (
    _is_equation_table,
    top_level_table_anchor_positions,
)
from src.shared.engine.count_engine import count_document
from src.shared.engine.document_scope_runtime import document_scope_allows_paragraph
from src.shared.engine.heading_numbering_ooxml import effective_numbering
from src.shared.engine.ooxml_ops import qn
from src.shared.engine.section_layout_planner import (
    collect_section_inventory,
    iter_active_sections,
)

if TYPE_CHECKING:
    from docx import Document

    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


_BROKEN_TOC_BOOKMARK_HINTS = (
    "error! bookmark not defined",
    "bookmark not defined",
    "未定义书签",
)
_FORMULA_OLE_PROGID_HINTS = ("equation", "mathtype", "eqn")
_OFFICE_NS = "urn:schemas-microsoft-com:office:office"
_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


class ValidationModule(BaseModule):
    """Validate final document state against core formatting expectations."""

    meta = ModuleMeta(
        name="validation",
        description="结果校验",
        category="validate",
        depends_on=("heading_recognition",),
        consumes=("doc_tree", "heading_map"),
        provides=("validation_issues", "count_result"),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        context.validation_issues = self.validate(doc, config, context)
        count_profile_id = str(
            getattr(getattr(config, "compliance_profile", None), "count_profile_id", "")
            or ""
        ).strip()
        if count_profile_id:
            count_result = count_document(doc, profile_id=count_profile_id)
            context.count_result = count_result
            tracker.record(
                rule_name="count_engine",
                target=count_profile_id,
                section="global",
                change_type="count_summary",
                after=count_result.summary(),
                paragraph_index=-1,
                success=True,
            )

    def validate(
        self,
        doc: Document,
        config: ResolvedConfig,
        context: PipelineContext,
    ) -> list[Issue]:
        issues: list[Issue] = []
        if _module_policy_enabled(config, "page_setup"):
            issues.extend(self._check_page_setup(doc, config, context))
        issues.extend(self._check_section_layout(doc, config, context))
        issues.extend(self._check_heading_styles(doc, config, context))
        issues.extend(self._check_heading_numbering_ownership(doc, config, context))
        issues.extend(self._check_body_typography(doc, config, context))
        issues.extend(self._check_broken_toc_placeholders(doc))
        issues.extend(self._check_formula_chain(doc, config, context))
        return issues

    def _check_page_setup(self, doc, config, context) -> list[Issue]:
        """Verify every page-setup ownership policy against the final XML."""

        issues: list[Issue] = []
        ps = config.page_setup
        issue_level = "error" if bool(getattr(config, "strict_mode", True)) else "warning"
        sections = iter_active_sections(doc)
        final_inventory = collect_section_inventory(doc)
        source_inventory = getattr(context, "page_setup_source_inventory", None)
        paper_mode = str(
            getattr(ps, "paper_size_mode", "force_template") or "force_template"
        )
        orientation_mode = str(
            getattr(ps, "orientation_mode", "preserve_source") or "preserve_source"
        )
        margin_mode = str(getattr(ps, "margin_mode", "force_template") or "force_template")
        template_landscape = str(getattr(ps, "orientation", "portrait")) == "landscape"
        section_starts = _section_start_indices(doc, len(sections))

        for index, section in enumerate(sections):
            source_boundary = _inventory_boundary(source_inventory, index)
            final_boundary = _inventory_boundary(final_inventory, index)
            source_landscape = (
                _inventory_boundary_is_landscape(source_boundary)
                if source_boundary is not None
                else _section_is_landscape(section, template_landscape)
            )
            if source_landscape is None:
                source_landscape = _section_is_landscape(section, template_landscape)
            desired_landscape, orientation_owned = _desired_landscape(
                ps,
                context,
                section_index=index,
                start_index=section_starts[index],
                source_landscape=source_landscape,
                template_landscape=template_landscape,
                orientation_mode=orientation_mode,
            )
            paper_key, paper_owned = _desired_paper_key(
                ps,
                context,
                section_index=index,
                start_index=section_starts[index],
                paper_mode=paper_mode,
            )
            margin_config, margin_owned = _desired_margin_config(
                ps,
                context,
                section_index=index,
                start_index=section_starts[index],
                margin_mode=margin_mode,
            )
            actual_landscape = _section_is_landscape(section, template_landscape)

            if _section_orientation_contradiction(section):
                issues.append(
                    Issue(
                        level=issue_level,
                        module_name=self.meta.name,
                        message="页面宽高与显式 w:orient 相互矛盾",
                        location=f"Section {index + 1}",
                    )
                )

            if (orientation_owned or orientation_mode == "force_template") and (
                actual_landscape != desired_landscape
            ):
                expected = "landscape" if desired_landscape else "portrait"
                issues.append(
                    Issue(
                        level=issue_level,
                        module_name=self.meta.name,
                        message=f"页面方向不符：期望 {expected}",
                        location=f"Section {index + 1}",
                    )
                )
            elif (
                orientation_mode == "preserve_source"
                and source_boundary is not None
                and actual_landscape != source_landscape
            ):
                issues.append(
                    Issue(
                        level=issue_level,
                        module_name=self.meta.name,
                        message="preserve_source 策略下页面方向改变",
                        location=f"Section {index + 1}",
                    )
                )

            if paper_owned:
                paper_width, paper_height = PAPER_SIZES[paper_key]
                expected_width_cm, expected_height_cm = (
                    (paper_height, paper_width)
                    if desired_landscape
                    else (paper_width, paper_height)
                )
                if not _length_matches(section.page_width, expected_width_cm) or not _length_matches(
                    section.page_height,
                    expected_height_cm,
                ):
                    issues.append(
                        Issue(
                            level=issue_level,
                            module_name=self.meta.name,
                            message=f"纸张尺寸不符：期望 {paper_key}",
                            location=f"Section {index + 1}",
                        )
                    )
            elif source_boundary is not None and final_boundary is not None:
                if not _same_preserved_paper(
                    source_boundary.page_size,
                    final_boundary.page_size,
                    allow_rotation=orientation_owned,
                ):
                    issues.append(
                        Issue(
                            level=issue_level,
                            module_name=self.meta.name,
                            message="preserve_source 策略下纸张规格改变",
                            location=f"Section {index + 1}",
                        )
                    )

            if margin_owned and margin_config is not None:
                expected_margins = {
                    "top_margin": (margin_config.top_cm, "上边距"),
                    "bottom_margin": (margin_config.bottom_cm, "下边距"),
                    "left_margin": (margin_config.left_cm, "左边距"),
                    "right_margin": (margin_config.right_cm, "右边距"),
                    "gutter": (margin_config.gutter_cm, "装订线"),
                    "header_distance": (
                        margin_config.header_distance_cm,
                        "页眉距边界",
                    ),
                    "footer_distance": (
                        margin_config.footer_distance_cm,
                        "页脚距边界",
                    ),
                }
                for attribute, (expected_cm, label) in expected_margins.items():
                    if _length_matches(getattr(section, attribute), expected_cm):
                        continue
                    issues.append(
                        Issue(
                            level=issue_level,
                            module_name=self.meta.name,
                            message=f"{label}不符：期望 {expected_cm:g}cm",
                            location=f"Section {index + 1}",
                        )
                    )
            elif (
                source_boundary is not None
                and final_boundary is not None
                and source_boundary.page_margins != final_boundary.page_margins
            ):
                issues.append(
                    Issue(
                        level=issue_level,
                        module_name=self.meta.name,
                        message="preserve_source 策略下页边距或页眉页脚距离改变",
                        location=f"Section {index + 1}",
                    )
                )
        return issues

    def _check_section_layout(self, doc, config, context) -> list[Issue]:
        """Verify topology receipts and explicit page-layout ownership policies."""

        issues: list[Issue] = []
        issue_level = "error" if bool(getattr(config, "strict_mode", True)) else "warning"
        inventory = collect_section_inventory(doc)
        context.final_section_inventory = inventory
        if not _module_policy_enabled(config, "section_format"):
            return issues
        source = getattr(context, "section_inventory", None)
        receipt = getattr(context, "section_execution_receipt", None)
        if receipt is not None:
            for message in receipt.validation_errors:
                issues.append(
                    Issue(
                        level=issue_level,
                        module_name=self.meta.name,
                        message=f"分节执行校验失败：{message}",
                        location="document.sections",
                    )
                )
            if inventory.section_count != receipt.final_section_count:
                issues.append(
                    Issue(
                        level=issue_level,
                        module_name=self.meta.name,
                        message=(
                            "分节数量在执行后发生未声明变化："
                            f"计划结果 {receipt.final_section_count}，最终 {inventory.section_count}"
                        ),
                        location="document.sections",
                    )
                )

        section_config = config.section
        boundary_mode = str(getattr(section_config, "boundary_mode", "") or "")
        if boundary_mode == "preserve_source" and source is not None:
            if inventory.section_count != source.section_count:
                issues.append(
                    Issue(
                        level=issue_level,
                        module_name=self.meta.name,
                        message=(
                            "preserve_source 策略下分节数量改变："
                            f"原 {source.section_count}，现 {inventory.section_count}"
                        ),
                        location="document.sections",
                    )
                )
            header_footer_enabled = _module_policy_enabled(config, "header_footer")
            behavior = getattr(config.header_footer, "behavior", None)
            link_mode = str(
                getattr(behavior, "link_to_previous", "preserve") or "preserve"
            )
            if (
                not header_footer_enabled
                and link_mode == "preserve"
                and inventory.section_count == source.section_count
            ):
                before_refs = tuple(item.header_footer_refs for item in source.boundaries)
                after_refs = tuple(item.header_footer_refs for item in inventory.boundaries)
                if before_refs != after_refs:
                    issues.append(
                        Issue(
                            level=issue_level,
                            module_name=self.meta.name,
                            message="preserve_source 策略下页眉页脚引用关系被改写",
                            location="document.sections.header_footer_refs",
                        )
                    )

        return issues

    def _check_heading_numbering_ownership(self, doc, config, context) -> list[Issue]:
        """Reject literal/native double numbering and missing native numbering."""

        issues: list[Issue] = []
        heading_map = getattr(context, "heading_map", None) or {}
        level_bindings = getattr(config.heading_numbering, "level_bindings", {}) or {}
        max_levels = int(getattr(config.heading_model, "max_heading_levels", 0) or 0)
        exact_values = set(config.heading_model.non_numbered_title_texts or [])
        prefix_values = list(config.heading_model.non_numbered_prefixes or [])
        for para_index, level in heading_map.items():
            if (
                not isinstance(para_index, int)
                or para_index < 0
                or para_index >= len(doc.paragraphs)
                or not isinstance(level, int)
                or level < 1
                or level > max_levels
            ):
                continue
            paragraph = doc.paragraphs[para_index]
            if not document_scope_allows_paragraph(context, para_index):
                continue
            if _is_non_numbered_title(paragraph, exact_values, prefix_values):
                continue
            binding = level_bindings.get(f"heading{level}")
            if binding is None or not bool(getattr(binding, "enabled", False)):
                continue
            native_numbering = effective_numbering(paragraph)
            literal_numbering = has_literal_heading_number_prefix(paragraph.text)
            if native_numbering is not None and literal_numbering:
                issues.append(
                    Issue(
                        level="error",
                        module_name=self.meta.name,
                        message="标题同时存在可见文字编号与 Word 原生编号，可能产生重复编号",
                        location=f"段落 #{para_index}",
                    )
                )
            elif native_numbering is None:
                issues.append(
                    Issue(
                        level="error",
                        module_name=self.meta.name,
                        message="标题缺少 Word 原生多级编号",
                        location=f"段落 #{para_index}",
                    )
                )
        return issues

    def _check_body_typography(self, doc, config, context) -> list[Issue]:
        """Verify that ordinary body paragraphs resolve to the configured font size."""

        style_config = config.styles.get("body") or config.styles.get("normal")
        expected_size = resolve_style_size_pt(style_config) if style_config else None
        if expected_size is None:
            return []

        heading_map = getattr(context, "heading_map", None) or {}
        doc_tree = getattr(context, "doc_tree", None)
        issues: list[Issue] = []
        for para_index, paragraph in enumerate(doc.paragraphs):
            if para_index in heading_map or not (paragraph.text or "").strip():
                continue
            if not document_scope_allows_paragraph(context, para_index):
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            style_lower = style_name.lower()
            if (
                style_lower.startswith("toc")
                or "caption" in style_lower
                or "header" in style_lower
                or "footer" in style_lower
                or "目录" in style_name
                or "题注" in style_name
                or "页眉" in style_name
                or "页脚" in style_name
            ):
                continue
            if doc_tree is not None:
                section_getter = getattr(doc_tree, "get_section_for_paragraph", None)
                if callable(section_getter) and section_getter(para_index) != "body":
                    continue
            actual_size = _first_visible_run_size_pt(paragraph)
            if actual_size is None or abs(actual_size - expected_size) <= 0.1:
                continue
            issues.append(
                Issue(
                    level="error",
                    module_name=self.meta.name,
                    message=(
                        "正文字号不符: "
                        f"期望 {expected_size:g} 磅，实际 {actual_size:g} 磅"
                    ),
                    location=f"段落 #{para_index}",
                )
            )
            if len(issues) >= 20:
                break
        return issues

    def _check_heading_styles(self, doc, config, context) -> list[Issue]:
        issues: list[Issue] = []
        heading_map = getattr(context, "heading_map", None) or {}
        style_map = getattr(config.heading_model, "level_to_word_style", {}) or {}
        for para_index, level in heading_map.items():
            if not isinstance(para_index, int) or para_index < 0 or para_index >= len(doc.paragraphs):
                continue
            para = doc.paragraphs[para_index]
            if not document_scope_allows_paragraph(context, para_index):
                continue
            if _is_non_numbered_title(
                para,
                set(config.heading_model.non_numbered_title_texts or []),
                list(config.heading_model.non_numbered_prefixes or []),
            ):
                continue
            expected = style_map.get(f"heading{level}", "")
            actual = para.style.name if para.style else ""
            if expected and actual != expected:
                issues.append(
                    Issue(
                        level="error",
                        module_name=self.meta.name,
                        message=f"标题样式不符: 期望 '{expected}', 实际 '{actual}'",
                        location=f"段落 #{para_index}",
                    )
                )
        return issues

    def _check_broken_toc_placeholders(self, doc) -> list[Issue]:
        issues: list[Issue] = []
        for index, para in enumerate(doc.paragraphs):
            raw = (para.text or "").strip()
            if not raw or "\t" not in raw:
                continue
            low = raw.lower()
            if any(hint in low for hint in _BROKEN_TOC_BOOKMARK_HINTS):
                issues.append(
                    Issue(
                        level="error",
                        module_name=self.meta.name,
                        message="目录条目存在损坏的书签占位文本",
                        location=f"段落 #{index}",
                    )
                )
        return issues

    def _check_formula_chain(self, doc, config, context) -> list[Issue]:
        issues: list[Issue] = []
        formula_convert_enabled = config.is_module_enabled("formula_convert")
        equation_format_enabled = config.is_module_enabled("equation_table_format")
        if not formula_convert_enabled and not equation_format_enabled:
            return issues

        if formula_convert_enabled:
            scoped_paragraphs = list(
                _iter_scoped_document_paragraphs(doc, context)
            )
            for index, paragraph in enumerate(scoped_paragraphs):
                text = paragraph.text or ""
                if paragraph_has_formula_source(text):
                    issues.append(
                        Issue(
                            level="warning",
                            module_name=self.meta.name,
                            message="公式源码未安全转换，已保留原文并需要复核",
                            location=f"公式段落 #{index}",
                        )
                    )
            formula_ole_count = sum(
                1
                for paragraph in scoped_paragraphs
                for element in paragraph._p.iter(
                    f"{{{_OFFICE_NS}}}OLEObject"
                )
                if _is_formula_ole_element(element)
            )
            if formula_ole_count:
                issues.append(
                    Issue(
                        level="warning",
                        module_name=self.meta.name,
                        message=(
                            f"检测到 {formula_ole_count} 个旧式 OLE 公式对象；"
                            "已保留显示，但不能宣称已转为 Word 原生 OMML"
                        ),
                        location="document.xml",
                    )
                )

        if not equation_format_enabled:
            return issues

        expected_alignment = _ALIGNMENTS.get(
            str(config.formula_table.number_alignment or "right").lower(),
            WD_ALIGN_PARAGRAPH.RIGHT,
        )
        expected_size = config.formula_table.number_font_size_pt
        expected_half_points = (
            str(int(round(float(expected_size) * 2)))
            if expected_size not in (None, "")
            else ""
        )
        table_anchors = top_level_table_anchor_positions(doc)
        for table_index, table in enumerate(doc.tables):
            if not _is_equation_table(table):
                continue
            table_anchor = (
                table_anchors[table_index]
                if table_index < len(table_anchors)
                else -1
            )
            if not document_scope_allows_paragraph(context, table_anchor):
                continue
            allow_plain_formula_rows = _has_explicit_equation_table_marker(table)
            for row_index, row in enumerate(table.rows):
                if len(row.cells) < 2:
                    continue
                number_cell = row.cells[-1]
                number_text = (number_cell.text or "").strip()
                row_has_formula = _row_has_formula_content(
                    row.cells[:-1],
                    allow_plain_text=allow_plain_formula_rows,
                )
                if row_has_formula and not number_text:
                    issues.append(
                        Issue(
                            level="warning",
                            module_name=self.meta.name,
                            message="公式行缺少编号",
                            location=f"表格 #{table_index + 1} 行 #{row_index + 1}",
                        )
                    )
                    continue
                if row_has_formula and _find_number_paragraph(number_cell) is None:
                    issues.append(
                        Issue(
                            level="warning",
                            module_name=self.meta.name,
                            message="公式编号不是可识别的序号格式，已保留原文并需要复核",
                            location=f"表格 #{table_index + 1} 行 #{row_index + 1}",
                        )
                    )
                    continue
                for paragraph in number_cell.paragraphs:
                    if not (paragraph.text or "").strip():
                        continue
                    if paragraph.paragraph_format.alignment != expected_alignment:
                        issues.append(
                            Issue(
                                level="warning",
                                module_name=self.meta.name,
                                message="公式编号对齐方式与模板不一致",
                                location=f"表格 #{table_index + 1} 行 #{row_index + 1}",
                            )
                        )
                    if expected_half_points:
                        size_node = paragraph._element.find(
                            f".//{qn('w:rPr')}/{qn('w:sz')}"
                        )
                        actual_size = (
                            size_node.get(qn("w:val")) if size_node is not None else ""
                        )
                        if actual_size != expected_half_points:
                            issues.append(
                                Issue(
                                    level="warning",
                                    module_name=self.meta.name,
                                    message=(
                                        "公式编号字号与模板不一致："
                                        f"期望 {float(expected_size):g} 磅"
                                    ),
                                    location=f"表格 #{table_index + 1} 行 #{row_index + 1}",
                                )
                            )
                    break
        return issues


def _iter_scoped_document_paragraphs(doc, context):
    seen: set[int] = set()
    for index, paragraph in enumerate(doc.paragraphs):
        if not document_scope_allows_paragraph(context, index):
            continue
        identity = id(paragraph._p)
        if identity in seen:
            continue
        seen.add(identity)
        yield paragraph

    table_anchors = top_level_table_anchor_positions(doc)
    for table_index, table in enumerate(doc.tables):
        anchor = (
            table_anchors[table_index]
            if table_index < len(table_anchors)
            else -1
        )
        if not document_scope_allows_paragraph(context, anchor):
            continue
        yield from _iter_unique_table_paragraphs(table, seen)


def _iter_unique_table_paragraphs(table, seen: set[int]):
    seen_cells: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            cell_identity = id(cell._tc)
            if cell_identity in seen_cells:
                continue
            seen_cells.add(cell_identity)
            for paragraph in cell.paragraphs:
                identity = id(paragraph._p)
                if identity in seen:
                    continue
                seen.add(identity)
                yield paragraph
            for nested_table in cell.tables:
                yield from _iter_unique_table_paragraphs(nested_table, seen)


def _is_formula_ole_element(element) -> bool:
    prog_id = str(
        element.get("ProgID")
        or element.get(f"{{{_OFFICE_NS}}}ProgID")
        or ""
    ).lower()
    return any(hint in prog_id for hint in _FORMULA_OLE_PROGID_HINTS)


def _first_visible_run_size_pt(paragraph) -> float | None:
    for run in paragraph.runs:
        if not (run.text or "").strip():
            continue
        if run.font.size is not None:
            return float(run.font.size.pt)

    style = paragraph.style
    seen: set[int] = set()
    while style is not None and id(style) not in seen:
        seen.add(id(style))
        if style.font.size is not None:
            return float(style.font.size.pt)
        style = style.base_style
    return None


def _is_non_numbered_title(
    paragraph,
    exact_values: set[str],
    prefix_values: list[str],
) -> bool:
    style_name = paragraph.style.name if paragraph.style else ""
    if "unnumbered" in style_name.lower():
        return True
    return (
        match_special_title(
            paragraph.text,
            exact_values=exact_values,
            prefix_values=prefix_values,
        )
        is not None
    )


def _module_policy_enabled(config, module_name: str) -> bool:
    """Honor explicit switches without breaking direct unit configurations."""

    switches = dict(getattr(config, "module_switches", {}) or {})
    if module_name in switches:
        return bool(switches[module_name])
    return True


def _inventory_boundary(inventory, index: int):
    boundaries = tuple(getattr(inventory, "boundaries", ()) or ())
    if 0 <= index < len(boundaries):
        return boundaries[index]
    return None


def _inventory_boundary_is_landscape(boundary) -> bool | None:
    if boundary is None:
        return None
    attributes = dict(getattr(boundary, "page_size", ()) or ())
    orientation = str(attributes.get("orient", "") or "")
    if orientation == "landscape":
        return True
    if orientation == "portrait":
        return False
    try:
        return int(attributes["w"]) > int(attributes["h"])
    except (KeyError, TypeError, ValueError):
        return None


def _same_preserved_paper(
    before_attributes,
    after_attributes,
    *,
    allow_rotation: bool,
) -> bool:
    before = dict(before_attributes or ())
    after = dict(after_attributes or ())
    if before == after:
        return True
    if not allow_rotation:
        return False
    before_other = {
        key: value for key, value in before.items() if key not in {"w", "h", "orient"}
    }
    after_other = {
        key: value for key, value in after.items() if key not in {"w", "h", "orient"}
    }
    if before_other != after_other:
        return False
    try:
        before_stock = sorted((int(before["w"]), int(before["h"])))
        after_stock = sorted((int(after["w"]), int(after["h"])))
    except (KeyError, TypeError, ValueError):
        return False
    return before_stock == after_stock


def _length_matches(actual, expected_cm: float, *, tolerance: int = 5000) -> bool:
    if actual is None:
        return False
    return abs(int(actual) - int(Cm(float(expected_cm)))) <= tolerance
