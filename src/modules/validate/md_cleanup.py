"""Clean Markdown/LaTeX/HTML residue without flattening formula payloads."""

from __future__ import annotations

import re
import unicodedata
from collections import Counter
from functools import cache
from pathlib import Path
from typing import TYPE_CHECKING

from src.formula_core.repair import repair_formula_text
from src.modules.base import BaseModule, ModuleMeta
from src.modules.special.formula_convert import (
    paragraph_has_formula_source,
    paragraph_has_protected_formula_content,
)
from src.shared.engine.run_ops import get_full_text

if TYPE_CHECKING:
    from docx import Document

    from src.config.resolved import ResolvedConfig
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker


_MD_PATTERNS: list[tuple[re.Pattern, str]] = [
    (re.compile(r"\*\*(.+?)\*\*"), r"\1"),
    (re.compile(r"\*(.+?)\*"), r"\1"),
    (re.compile(r"(?<!\w)__(.+?)__(?!\w)"), r"\1"),
    (re.compile(r"(?<!\w)_(?!_)(.+?)(?<!_)_(?!\w)"), r"\1"),
    (re.compile(r"~~(.+?)~~"), r"\1"),
    (re.compile(r"`(.+?)`"), r"\1"),
    (re.compile(r"^\s*#{1,6}\s+"), ""),
    (re.compile(r"^\s*[-*+]\s+"), ""),
    (re.compile(r"^\s*\d+\.\s+"), ""),
    (re.compile(r"\[([^\]]+)\]\([^)]+\)"), r"\1"),
    (re.compile(r"!\[[^\]]*\]\([^)]+\)"), ""),
    (re.compile(r"^>{1,3}\s?"), ""),
    (re.compile(r"^-{3,}$"), ""),
    (re.compile(r"^\*{3,}$"), ""),
]
_LATEX_PATTERNS: list[tuple[re.Pattern, str]] = [
    (re.compile(r"\\textbf\{([^}]+)\}"), r"\1"),
    (re.compile(r"\\textit\{([^}]+)\}"), r"\1"),
    (re.compile(r"\\emph\{([^}]+)\}"), r"\1"),
    (re.compile(r"\\underline\{([^}]+)\}"), r"\1"),
    (re.compile(r"\\text\{([^}]+)\}"), r"\1"),
    (re.compile(r"\\cite\{([^}]+)\}"), r"[\1]"),
    (re.compile(r"\\ref\{([^}]+)\}"), r"\1"),
    (re.compile(r"\\label\{[^}]+\}"), ""),
    (re.compile(r"\$([^$]+)\$"), r"\1"),
    (re.compile(r"\\(?:section|subsection|subsubsection)\*?\{([^}]+)\}"), r"\1"),
    (re.compile(r"\\(?:begin|end)\{[^}]+\}"), ""),
    (re.compile(r"\\\\"), ""),
    (re.compile(r"\\[a-zA-Z]+(?:\[[^\]]*\])?\s*"), ""),
]
_HTML_PATTERNS: list[tuple[re.Pattern, str]] = [
    (re.compile(r"<br\s*/?>"), ""),
    (re.compile(r"</?(?:p|div|span|b|i|em|strong)[^>]*>"), ""),
    (re.compile(r"&nbsp;"), " "),
    (re.compile(r"&amp;"), "&"),
    (re.compile(r"&lt;"), "<"),
    (re.compile(r"&gt;"), ">"),
]
ALL_PATTERNS = _MD_PATTERNS + _LATEX_PATTERNS + _HTML_PATTERNS


class MdCleanupModule(BaseModule):
    meta = ModuleMeta(
        name="md_cleanup",
        description="Markdown/LaTeX 清理",
        category="validate",
        requires_config=(),
        soft_after=("formula_convert",),
    )

    def apply(
        self,
        doc: Document,
        config: ResolvedConfig,
        tracker: ChangeTracker,
        context: PipelineContext,
    ) -> None:
        options = getattr(config, "md_cleanup", None)
        formula_repairs, fake_markers_removed = _cleanup_formula_copy_artifacts(
            doc,
            tracker,
            normalize_noise=bool(
                getattr(options, "formula_copy_noise_cleanup", True)
            ),
            suppress_fake_lists=bool(
                getattr(options, "suppress_formula_fake_lists", True)
            ),
        )

        count = 0
        if _markup_cleanup_applies(config, context):
            count = sum(
                1 for paragraph in doc.paragraphs if _clean_paragraph(paragraph)
            )
            for table in doc.tables:
                seen: set[int] = set()
                for row in table.rows:
                    for cell in row.cells:
                        if id(cell._tc) in seen:
                            continue
                        seen.add(id(cell._tc))
                        count += sum(
                            1
                            for paragraph in cell.paragraphs
                            if _clean_paragraph(paragraph)
                        )

        if count or formula_repairs or fake_markers_removed:
            tracker.record(
                rule_name=self.meta.name,
                target=(
                    f"通用清理 {count} 段；公式复制噪声 {formula_repairs} 段；"
                    f"伪列表标记 {fake_markers_removed} 段"
                ),
                section="global",
                change_type="cleanup",
                before="含 Markdown/LaTeX/HTML 或公式复制噪声",
                after="已按开关安全清理",
            )


def _markup_cleanup_applies(config, context) -> bool:
    """Return whether generic markup stripping applies to this actual source.

    Formula-copy repairs remain separately controlled by ``MdCleanupOptions``.
    Generic Markdown/LaTeX/HTML regex cleanup is intentionally limited to a
    Markdown source so an ordinary Word document can safely contain literal
    markup, code samples, or LaTeX examples.
    """

    profile = getattr(config, "input_source_profile", None)
    policy = str(getattr(profile, "markdown_policy", "disabled") or "disabled")
    if policy == "disabled":
        return False

    source_path = str(getattr(context, "source_doc_path", "") or "").strip()
    if not source_path:
        # Direct module callers do not always carry execution provenance.  Keep
        # their explicit cleanup request working while production executions
        # always provide a concrete source path.
        return True
    return Path(source_path).suffix.casefold() in {".md", ".markdown"}


def _clean_paragraph(paragraph) -> bool:
    text = get_full_text(paragraph)
    if not text:
        return False
    if paragraph_has_protected_formula_content(paragraph) or paragraph_has_formula_source(text):
        return False
    cleaned = text
    for pattern, replacement in ALL_PATTERNS:
        cleaned = pattern.sub(replacement, cleaned)
    if cleaned == text:
        return False
    _replace_paragraph_text(paragraph, cleaned)
    return True


def _cleanup_formula_copy_artifacts(
    doc,
    tracker,
    *,
    normalize_noise: bool,
    suppress_fake_lists: bool,
) -> tuple[int, int]:
    normalized = 0
    removed = 0
    if normalize_noise:
        for index, paragraph in enumerate(list(doc.paragraphs)):
            if paragraph_has_protected_formula_content(paragraph):
                continue
            raw = str(paragraph.text or "").strip()
            cleaned = _normalize_formula_copy_noise_text(raw)
            if not cleaned or cleaned == raw:
                continue
            _replace_paragraph_text(paragraph, cleaned)
            normalized += 1
            tracker.record(
                rule_name="md_cleanup",
                target=f"段落 #{index}",
                section="formula",
                change_type="cleanup",
                before=raw[:120],
                after=cleaned[:120],
                paragraph_index=index,
            )

    if suppress_fake_lists:
        for index in range(len(doc.paragraphs) - 2, 0, -1):
            paragraph = doc.paragraphs[index]
            if paragraph_has_protected_formula_content(paragraph):
                continue
            marker = str(paragraph.text or "").strip()
            if not _is_fake_formula_list_marker_line(marker):
                continue
            if not (
                _looks_like_formula_paragraph(doc.paragraphs[index - 1])
                and _looks_like_formula_paragraph(doc.paragraphs[index + 1])
            ):
                continue
            parent = paragraph._p.getparent()
            if parent is None:
                continue
            parent.remove(paragraph._p)
            removed += 1
            tracker.record(
                rule_name="md_cleanup",
                target=f"段落 #{index}",
                section="formula",
                change_type="cleanup",
                before=marker,
                after="移除夹在公式行之间的伪列表标记",
                paragraph_index=index,
            )
    return normalized, removed


def _replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            parent = run._element.getparent()
            if parent is not None:
                parent.remove(run._element)
    else:
        paragraph.add_run(text)


def _normalize_formula_copy_noise_text(text: str) -> str | None:
    raw = str(text or "").strip()
    if not raw or (not paragraph_has_formula_source(raw) and "=" not in raw):
        return None
    repaired = repair_formula_text(raw)
    repaired_text = str(repaired.text or "").strip()
    if repaired_text and repaired_text != raw:
        unresolved = repaired_text.count("ESC_") + repaired_text.count("__ESC")
        if unresolved <= 2 and (
            "\\" in repaired_text
            or len(repaired_text) <= len(raw) * 0.75
            or float(repaired.confidence or 0.0) >= 0.9
        ):
            return repaired_text

    compact = _compact_formula_text(raw)
    segments = _split_formula_segments(compact)
    if len(segments) < 2:
        return None
    keys = [_formula_skeleton_key(segment) for segment in segments]
    keys = [key for key in keys if key]
    if not keys:
        return None
    dominant, count = Counter(keys).most_common(1)[0]
    if count < 2:
        return None
    candidates = [
        segment for segment in segments if _formula_skeleton_key(segment) == dominant
    ]
    return max(candidates, key=_formula_detail_score) if candidates else None


def _split_formula_segments(compact: str) -> list[str]:
    if not compact or len(compact) > 240 or compact.count("=") < 2:
        return []

    @cache
    def search(start: int) -> tuple[str, ...] | None:
        if start == len(compact):
            return ()
        if start >= len(compact) or not (compact[start].isalpha() or compact[start] == "("):
            return None
        best: tuple[str, ...] | None = None
        for end in range(start + 5, min(len(compact), start + 120) + 1):
            segment = compact[start:end]
            if segment.count("=") > 1:
                break
            if not _is_formula_segment_candidate(segment):
                continue
            rest = search(end)
            if rest is not None and (best is None or len(rest) + 1 > len(best)):
                best = (segment,) + rest
        return best

    result = search(0)
    return list(result) if result and len(result) >= 2 else []


def _is_formula_segment_candidate(text: str) -> bool:
    if not (5 <= len(text) <= 120) or text.count("=") != 1:
        return False
    if re.search(r"[\u4e00-\u9fff]", text):
        return False
    left, right = text.split("=", 1)
    if not left or not right or not any(char.isalpha() for char in text):
        return False
    if not re.fullmatch(r"[A-Za-z0-9+\-*/^_=()\[\]{},.\\<>!~]+", text):
        return False
    return bool(re.search(r"[+\-*/^_\d\[\]]|[A-Za-z]\(", text))


def _compact_formula_text(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(text or ""))
    translations = {
        "−": "-", "×": "*", "·": "*", "→": "->", "←": "<-",
        "∞": "inf", "∑": "sum", "∫": "int", "∏": "prod", "√": "sqrt",
        "≈": "~=", "≠": "!=", "≤": "<=", "≥": ">=",
    }
    for source, target in translations.items():
        normalized = normalized.replace(source, target)
    return re.sub(r"\s+", "", normalized)


def _formula_skeleton_key(text: str) -> str:
    return re.sub(r"[\^_{}()\s]", "", str(text or "").lower())


def _formula_detail_score(text: str) -> int:
    return (
        text.count("^") * 4
        + text.count("_") * 4
        + text.count("{")
        + text.count("}")
        + text.count("\\")
        + len(text) // 20
    )


def _is_fake_formula_list_marker_line(text: str) -> bool:
    return bool(re.fullmatch(r"(?:\d+[.)、]|[（(]\d+[)）])", str(text or "").strip()))


def _looks_like_formula_paragraph(paragraph) -> bool:
    if paragraph_has_protected_formula_content(paragraph):
        return True
    text = str(paragraph.text or "").strip()
    return paragraph_has_formula_source(text) or _is_formula_segment_candidate(
        _compact_formula_text(text)
    )


__all__ = ["MdCleanupModule"]
