"""Deterministic, private-data-free ISO attachment integration fixtures."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import Workbook


TEXT_IDENTIFIERS = tuple(f"field_{index:02d}" for index in range(1, 31))
TIMELINE_IDENTIFIERS = tuple(f"节点_{index:02d}" for index in range(1, 12))


@dataclass(frozen=True, slots=True)
class IsoAttachmentFixture:
    docx_directory: Path
    workbook_path: Path


def build_iso_attachment_fixture(root: Path) -> IsoAttachmentFixture:
    """Build six templates and eighteen profiles with the former fixture depth.

    The generated package deliberately preserves the integration dimensions of
    the original developer-local sample: 41 unique strict tokens, 114 total
    occurrences, 30 direct fields, 11 timeline candidates, and 18 profiles.
    """

    fixture_root = root / "iso_attachment_fixture"
    docx_directory = fixture_root / "templates"
    docx_directory.mkdir(parents=True, exist_ok=True)

    tokens = (
        *(f"{{{{@text:{identifier}}}}}" for identifier in TEXT_IDENTIFIERS),
        *(f"{{{{@time:{identifier}}}}}" for identifier in TIMELINE_IDENTIFIERS),
    )
    occurrences = (*tokens, *(tokens[index % len(tokens)] for index in range(73)))
    assert len(tokens) == 41
    assert len(occurrences) == 114

    for document_index in range(6):
        start = document_index * 19
        document = Document()
        document.add_heading(f"Anonymous ISO template {document_index + 1}", level=1)
        document.add_paragraph(" ".join(occurrences[start : start + 19]))
        document.save(docx_directory / f"template_{document_index + 1:02d}.docx")

    workbook_path = fixture_root / "profiles.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "profiles"
    sheet.append(("profile_id", "profile_name", *TEXT_IDENTIFIERS))
    for profile_index in range(1, 19):
        sheet.append(
            (
                f"profile_{profile_index:02d}",
                f"Anonymous profile {profile_index:02d}",
                *(
                    f"fixture-{profile_index:02d}-{identifier}"
                    for identifier in TEXT_IDENTIFIERS
                ),
            )
        )
    workbook.save(workbook_path)
    workbook.close()

    return IsoAttachmentFixture(
        docx_directory=docx_directory,
        workbook_path=workbook_path,
    )


__all__ = [
    "IsoAttachmentFixture",
    "TEXT_IDENTIFIERS",
    "TIMELINE_IDENTIFIERS",
    "build_iso_attachment_fixture",
]
