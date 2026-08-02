from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

from src.config.migration import (
    normalize_module_switches,
    normalize_scene_payload,
    normalize_template_payload,
)
from src.config.dataclass_utils import dict_to_dataclass
from src.config.resolved import ResolvedConfig
from src.config.resolver import resolve_config
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.modules.default_switches import default_module_switches
from src.modules.special.chem_typography import ChemTypographyModule
from src.modules.special.formula_convert import FormulaConvertModule
from src.modules.special.formula_style import FormulaStyleModule
from src.modules.table.table_format import _is_equation_table
from src.modules.validate.md_cleanup import MdCleanupModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.document_structure_model import DocSection, DocTree
from src.shared.engine.font_resolver import resolve_font
from src.shared.engine.ooxml_ops import qn


def _mark_equation_table(table) -> None:
    description = OxmlElement("w:tblDescription")
    description.set(qn("w:val"), "alavette-equation-table")
    table._tbl.tblPr.append(description)


def _append_math_run(paragraph, text: str, style: str | None = None):
    math = paragraph._p.find(qn("m:oMath"))
    if math is None:
        math = OxmlElement("m:oMath")
        paragraph._p.append(math)
    run = OxmlElement("m:r")
    if style:
        run_pr = OxmlElement("m:rPr")
        style_node = OxmlElement("m:sty")
        style_node.set(qn("m:val"), style)
        run_pr.append(style_node)
        run.append(run_pr)
    text_node = OxmlElement("m:t")
    text_node.text = text
    run.append(text_node)
    math.append(run)
    return run


def test_formula_conversion_does_not_promote_an_ordinary_data_table():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "变量"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "$x_i$"
    table.cell(1, 1).text = "样本 A"

    FormulaConvertModule().apply(
        document, ResolvedConfig(), ChangeTracker(), PipelineContext()
    )

    assert table.cell(1, 0).text == "$x_i$"
    assert table._tbl.find(f".//{qn('m:oMath')}") is None
    assert _is_equation_table(table) is False


def test_formula_conversion_is_allowed_inside_a_marked_equation_table():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    _mark_equation_table(table)
    table.cell(0, 0).text = "$$x_i^2$$"
    table.cell(0, 1).text = ""

    FormulaConvertModule().apply(
        document, ResolvedConfig(), ChangeTracker(), PipelineContext()
    )

    assert table.cell(0, 0)._tc.find(f".//{qn('m:oMath')}") is not None
    assert _is_equation_table(table) is True


def test_plain_formula_source_with_an_empty_data_cell_is_not_an_equation_table():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "$x_i$"
    table.cell(0, 1).text = ""

    assert _is_equation_table(table) is False


def test_semantic_formula_style_preserves_explicit_msty_and_splits_functions():
    document = Document()
    explicit_paragraph = document.add_paragraph()
    explicit_run = _append_math_run(explicit_paragraph, "x", style="b")
    function_paragraph = document.add_paragraph()
    _append_math_run(function_paragraph, "sin(x)")
    config = ResolvedConfig()
    config.formula_style.unify_font = False
    config.formula_style.unify_size = False
    config.formula_style.unify_spacing = False

    FormulaStyleModule().apply(
        document, config, ChangeTracker(), PipelineContext()
    )

    explicit_style = explicit_run.find(f"{qn('m:rPr')}/{qn('m:sty')}")
    assert explicit_style is not None
    assert explicit_style.get(qn("m:val")) == "b"
    styled = [
        (
            str(run.find(qn("m:t")).text or ""),
            run.find(f"{qn('m:rPr')}/{qn('m:sty')}").get(qn("m:val")),
        )
        for run in function_paragraph._p.findall(f".//{qn('m:r')}")
    ]
    assert styled == [("sin(", "p"), ("x", "i"), (")", "p")]


def test_semantic_formula_style_restores_explicit_marker_families():
    document = Document()
    paragraph = document.add_paragraph()
    _append_math_run(paragraph, r"\mathbf{A}+\vec{v}")
    config = ResolvedConfig()
    config.formula_style.unify_font = False
    config.formula_style.unify_size = False

    FormulaStyleModule().apply(
        document, config, ChangeTracker(), PipelineContext()
    )

    styles = {
        str(run.find(qn("m:t")).text or ""): run.find(
            f"{qn('m:rPr')}/{qn('m:sty')}"
        ).get(qn("m:val"))
        for run in paragraph._p.findall(f".//{qn('m:r')}")
    }
    assert styles["A"] == "b"
    assert styles["v"] == "bi"


def test_formula_style_applies_native_math_typography_only_inside_reviewed_scope():
    document = Document()
    cover_paragraph = document.add_paragraph()
    cover_run = _append_math_run(cover_paragraph, "x")
    body_paragraph = document.add_paragraph()
    body_run = _append_math_run(body_paragraph, "y")
    config = ResolvedConfig()
    config.formula_table.formula_font_name = "Courier New"
    config.formula_table.formula_font_size_pt = 15.0
    context = PipelineContext(
        document_scope_gate_active=True,
        doc_tree=DocTree(
            sections=[
                DocSection("cover", 0, 1),
                DocSection("body", 1, 2),
            ],
            writable_roles=frozenset({"body"}),
        ),
    )

    FormulaStyleModule().apply(document, config, ChangeTracker(), context)

    assert cover_run.find(qn("w:rPr")) is None
    body_r_pr = body_run.find(qn("w:rPr"))
    assert body_r_pr is not None
    assert body_r_pr.find(qn("w:rFonts")).get(qn("w:ascii")) == resolve_font(
        "Courier New",
        lang="en",
    )
    assert body_r_pr.find(qn("w:sz")).get(qn("w:val")) == "30"


def test_formula_convert_only_rewrites_formula_sources_inside_reviewed_scope():
    document = Document()
    cover_paragraph = document.add_paragraph(r"cover $x_1$ text")
    body_paragraph = document.add_paragraph(r"body $y_2$ text")
    context = PipelineContext(
        document_scope_gate_active=True,
        doc_tree=DocTree(
            sections=[
                DocSection("cover", 0, 1),
                DocSection("body", 1, 2),
            ],
            writable_roles=frozenset({"body"}),
        ),
    )

    FormulaConvertModule().apply(
        document,
        ResolvedConfig(),
        ChangeTracker(),
        context,
    )

    assert cover_paragraph.text == r"cover $x_1$ text"
    assert cover_paragraph._p.find(f".//{qn('m:oMath')}") is None
    assert body_paragraph.text == "body  text"
    assert body_paragraph._p.find(f".//{qn('m:oMath')}") is not None
    assert context.formula_runtime["matched"] == 1
    assert context.formula_runtime["converted"] == 1


def test_formula_convert_keeps_out_of_scope_table_formula_and_latex_wrapper():
    document = Document()
    cover_wrapper = document.add_paragraph(r"\documentclass{article}")
    cover_table = document.add_table(rows=1, cols=2)
    _mark_equation_table(cover_table)
    cover_table.cell(0, 0).text = r"$$x_1$$"
    body_paragraph = document.add_paragraph(r"$$y_2$$")
    body_wrapper = document.add_paragraph(r"\end{document}")
    context = PipelineContext(
        document_scope_gate_active=True,
        doc_tree=DocTree(
            sections=[
                DocSection("cover", 0, 1),
                DocSection("body", 1, 3),
            ],
            writable_roles=frozenset({"body"}),
        ),
    )
    config = ResolvedConfig()
    config.formula_convert.output_mode = "latex"

    FormulaConvertModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    assert cover_wrapper._p.getparent() is not None
    assert cover_wrapper.text == r"\documentclass{article}"
    assert cover_table.cell(0, 0).text == r"$$x_1$$"
    assert cover_table.cell(0, 0)._tc.find(f".//{qn('m:oMath')}") is None
    assert body_paragraph._p.find(f".//{qn('m:oMath')}") is not None
    assert body_wrapper._p.getparent() is None
    assert context.formula_runtime["matched"] == 1
    assert context.formula_runtime["latex_wrappers_removed"] == 1


def test_thesis_formula_containerization_keeps_original_reviewed_scope_positions():
    document = Document()
    cover_paragraph = document.add_paragraph()
    cover_run = _append_math_run(cover_paragraph, "cover_x")
    body_paragraph = document.add_paragraph()
    _append_math_run(body_paragraph, "body_y")
    body_paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    config = ResolvedConfig(mode_id="thesis")
    config.equation_numbering.numbering_format = "global"
    context = PipelineContext(
        mode_id="thesis",
        document_scope_gate_active=True,
        doc_tree=DocTree(
            sections=[
                DocSection("cover", 0, 1),
                DocSection("body", 1, 2),
            ],
            writable_roles=frozenset({"body"}),
        ),
    )

    from src.modules.special.equation_table_format import EquationTableFormatModule

    EquationTableFormatModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    assert len(document.paragraphs) == 1
    assert len(document.tables) == 1
    assert cover_run.find(qn("w:rPr")) is None
    assert "cover_x" in "".join(document.paragraphs[0]._p.itertext())
    assert "body_y" in "".join(document.tables[0]._element.itertext())
    assert document.tables[0].cell(0, 1).text == "(1)"
    assert context.formula_runtime["tables_created"] == 1


def test_formula_table_conversion_keeps_formula_only_inline_fragments_inline():
    document = Document()
    document.add_paragraph(r"$x_i$")
    document.add_paragraph(r"$$y_i$$")
    config = ResolvedConfig(mode_id="thesis")
    config.equation_numbering.numbering_format = "global"
    context = PipelineContext(mode_id="thesis")

    FormulaConvertModule().apply(document, config, ChangeTracker(), context)

    from src.modules.special.equation_table_format import EquationTableFormatModule

    EquationTableFormatModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    assert len(document.paragraphs) == 1
    assert "x" in "".join(document.paragraphs[0]._p.itertext())
    assert len(document.tables) == 1
    assert "y" in "".join(document.tables[0]._element.itertext())

    EquationTableFormatModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    assert len(document.paragraphs) == 1
    assert document.paragraphs[0].paragraph_format.alignment is None
    assert len(document.tables) == 1


def test_formula_table_preserves_and_wraps_block_source_when_conversion_is_off():
    document = Document()
    document.add_paragraph(r"$$x_i^2$$")
    config = ResolvedConfig(mode_id="thesis")
    config.module_switches["formula_convert"] = False
    config.module_switches["equation_table_format"] = True
    config.equation_numbering.numbering_format = "global"
    context = PipelineContext(mode_id="thesis")

    from src.modules.special.equation_table_format import EquationTableFormatModule

    EquationTableFormatModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    assert document.paragraphs == []
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == r"$$x_i^2$$"
    assert document.tables[0].cell(0, 1).text == "(1)"
    assert context.formula_runtime["tables_created"] == 1


def test_formula_table_wraps_block_source_under_explicit_keep_source_policy():
    document = Document()
    document.add_paragraph(r"$$x_i^2$$")
    config = ResolvedConfig(mode_id="thesis")
    config.formula_convert.output_mode = "keep_source"
    config.equation_numbering.numbering_format = "global"
    context = PipelineContext(mode_id="thesis")

    FormulaConvertModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )
    from src.modules.special.equation_table_format import EquationTableFormatModule

    EquationTableFormatModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == r"$$x_i^2$$"
    assert document.tables[0].cell(0, 1).text == "(1)"
    assert context.formula_runtime["tables_created"] == 1


def test_latex_logic_roundtrips_to_editable_omml_and_removes_project_wrappers():
    document = Document()
    document.add_paragraph(r"\documentclass{article}")
    document.add_paragraph(r"\begin{document}")
    document.add_paragraph(r"$$x_i^2$$")
    document.add_paragraph(r"\end{document}")
    config = ResolvedConfig()
    config.formula_convert.output_mode = "latex"
    context = PipelineContext()

    FormulaConvertModule().apply(document, config, ChangeTracker(), context)

    assert document.element.body.find(f".//{qn('m:oMath')}") is not None
    assert all("document" not in paragraph.text for paragraph in document.paragraphs)
    assert context.formula_runtime["latex_logic_editable_output"] is True
    assert context.formula_runtime["latex_exchange"]
    assert context.formula_runtime["latex_wrappers_removed"] == 3


def test_keep_source_policy_still_reports_formula_detection_evidence():
    document = Document()
    document.add_paragraph(r"$$x_i^2$$")
    config = ResolvedConfig()
    config.formula_convert.output_mode = "keep_source"
    context = PipelineContext()

    FormulaConvertModule().apply(document, config, ChangeTracker(), context)

    assert context.formula_runtime["matched"] == 1
    assert context.formula_runtime["source_counts"]["latex"] == 1
    assert context.formula_runtime["preserved"] == 1
    assert context.formula_runtime["converted"] == 0
    assert context.formula_runtime["manual_review_required"] is False


def test_retired_image_fallback_reports_preserved_formulas_for_review():
    document = Document()
    document.add_paragraph(r"$$x_i^2$$")
    config = ResolvedConfig()
    config.formula_convert.output_mode = "image_fallback"
    context = PipelineContext()
    tracker = ChangeTracker()

    FormulaConvertModule().apply(document, config, tracker, context)

    assert context.formula_runtime["matched"] == 1
    assert context.formula_runtime["skipped"] == 1
    assert context.formula_runtime["manual_review_required"] is True
    assert context.formula_runtime["diagnostics"][0]["reason"] == (
        "image_fallback_unavailable"
    )
    assert tracker.get_by_module("formula_convert")[-1].change_type == "review"


def test_formula_fake_list_cleanup_obeys_both_strategy_switches():
    document = Document()
    document.add_paragraph("x_1=2")
    document.add_paragraph("1.")
    document.add_paragraph("y_1=3")
    config = ResolvedConfig()
    config.md_cleanup.suppress_formula_fake_lists = True

    MdCleanupModule().apply(document, config, ChangeTracker(), PipelineContext())
    assert [paragraph.text for paragraph in document.paragraphs] == ["x_1=2", "y_1=3"]

    preserved = Document()
    preserved.add_paragraph("x_1=2")
    preserved.add_paragraph("1.")
    preserved.add_paragraph("y_1=3")
    config.md_cleanup.formula_copy_noise_cleanup = False
    config.md_cleanup.suppress_formula_fake_lists = False
    MdCleanupModule().apply(preserved, config, ChangeTracker(), PipelineContext())
    assert [paragraph.text for paragraph in preserved.paragraphs] == [
        "x_1=2", "1.", "y_1=3"
    ]


def test_empty_chemistry_scope_means_no_processing():
    document = Document()
    paragraph = document.add_paragraph("H2O")
    config = ResolvedConfig()

    ChemTypographyModule().apply(
        document, config, ChangeTracker(), PipelineContext()
    )
    assert not any(run.font.subscript for run in paragraph.runs)

    config.chem_typography.enabled = True
    config.chem_typography.scopes["body"] = True
    ChemTypographyModule().apply(
        document, config, ChangeTracker(), PipelineContext()
    )
    assert [run.text for run in paragraph.runs if run.font.subscript] == ["2"]


def test_legacy_formula_blocks_migrate_to_the_nested_thesis_rule():
    normalized = normalize_scene_payload(
        {
            "category": "thesis",
            "formula_to_table": {"enabled": True, "block_only": True},
            "formula_style": {"enabled": True, "unify_font": False},
        }
    )
    assert "equation_table_format" not in normalized["module_switches"]
    assert normalized["thesis_formula_rules"]["formula_to_table"]["enabled"] is True
    assert normalized["thesis_formula_rules"]["formula_to_table"]["block_only"] is True
    assert normalized["thesis_formula_rules"]["formula_style"]["unify_font"] is False

    template = normalize_template_payload(
        {
            "formula_table": {
                "formula_space_before_value": 0.5,
                "formula_space_before_unit": "cm",
            }
        }
    )
    assert "formula_table" not in template
    switches = normalize_module_switches(
        {"formula_to_table": True, "formula_style": False}
    )
    assert switches["equation_table_format"] is True
    canonical_override = normalize_module_switches(
        {"formula_to_table": True, "equation_table_format": False}
    )
    assert canonical_override["equation_table_format"] is False


def test_legacy_formula_policy_and_visuals_keep_their_new_owners():
    normalized = normalize_scene_payload(
        {
            "category": "thesis",
            "formula_convert": {
                "enabled": True,
                "output_mode": "latex",
            },
            "formula_to_table": {
                "enabled": True,
                "block_only": True,
            },
            "formula_style": {
                "enabled": False,
                "unify_font": False,
            },
            "equation_table_format": {
                "enabled": True,
                "numbering_format": "chapter-seq",
            },
            "formula_table": {
                "formula_font_name": "Cambria Math",
                "number_font_size_pt": 10.5,
            },
        }
    )

    assert "formula_convert" not in normalized["module_switches"]
    assert normalized["mode_id"] == "thesis"
    assert "equation_table_format" not in normalized["module_switches"]
    rules = normalized["thesis_formula_rules"]
    assert rules["formula_convert"]["enabled"] is True
    assert rules["formula_convert"]["output_mode"] == "latex"
    assert rules["formula_to_table"] == {
        "enabled": True,
        "block_only": True,
    }
    assert rules["equation_numbering"]["enabled"] is True
    assert rules["formula_style"]["unify_font"] is False
    assert rules["formula_style"]["enabled"] is False
    assert rules["equation_numbering"]["numbering_format"] == "chapter-seq"
    assert rules["formula_table"]["formula_font_name"] == "Cambria Math"
    assert rules["formula_table"]["number_font_size_pt"] == 10.5
    assert not any(
        key.startswith(
            (
                "formula_convert.",
                "formula_table.",
                "formula_style.",
                "equation_numbering.",
                "chem_typography.",
            )
        )
        for key in normalized.get("template_overrides", {})
    )

    scene = dict_to_dataclass(SceneWorkspace, normalized)
    resolved = resolve_config(TemplateConfig(), scene)
    assert resolved.mode_id == "thesis"
    assert resolved.formula_table.formula_font_name == "Cambria Math"
    assert resolved.equation_numbering.numbering_format == "chapter-seq"


def test_legacy_formula_subswitches_keep_independent_nested_state():
    normalized = normalize_scene_payload(
        {
            "category": "thesis",
            "formula_to_table": {"enabled": True},
            "formula_style": {"enabled": False},
        }
    )

    assert normalized["thesis_formula_rules"]["formula_to_table"]["enabled"] is True
    assert normalized["thesis_formula_rules"]["formula_style"]["enabled"] is False


def test_canonical_formula_rules_win_over_stale_legacy_roots():
    normalized = normalize_scene_payload(
        {
            "mode_id": "thesis",
            "thesis_formula_rules": {
                "formula_enabled": True,
                "formula_convert": {
                    "enabled": False,
                    "output_mode": "keep_source",
                },
            },
            "formula_convert": {
                "enabled": True,
                "output_mode": "word_native",
            },
        }
    )

    conversion = normalized["thesis_formula_rules"]["formula_convert"]
    assert conversion["enabled"] is False
    assert conversion["output_mode"] == "keep_source"


def test_legacy_formula_dotted_overrides_move_wholly_into_the_thesis_rule():
    normalized = normalize_scene_payload(
        {
            "category": "thesis",
            "template_overrides": {
                "formula_to_table.block_only": False,
                "equation_table_format.enabled": False,
                "equation_table_format.numbering_format": "global",
                "table.layout_mode": "full",
            },
        }
    )

    rules = normalized["thesis_formula_rules"]
    assert rules["formula_to_table"]["block_only"] is False
    assert rules["equation_numbering"]["enabled"] is False
    assert rules["equation_numbering"]["numbering_format"] == "global"
    assert normalized["template_overrides"] == {"table.layout_mode": "full"}


def test_legacy_formula_aliases_survive_real_scene_switch_normalization():
    for payload in (
        {
            "category": "thesis",
            "module_switches": {
                "formula_to_table": True,
                "formula_style": False,
            }
        },
        {
            "category": "thesis",
            "capabilities": {
                "formula_to_table": True,
                "formula_style": False,
            }
        },
        {
            "category": "thesis",
            "pipeline": ["formula_to_table"],
        },
    ):
        normalized = normalize_scene_payload(payload)
        assert "equation_table_format" not in normalized["module_switches"]
        rules = normalized["thesis_formula_rules"]
        assert rules["formula_to_table"]["enabled"] is True
        assert rules["formula_style"]["enabled"] is False


def test_formula_table_format_is_not_enabled_by_missing_generic_switch():
    assert default_module_switches()["equation_table_format"] is False
