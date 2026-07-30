from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from src.config.entity import EntityArchive, EntityProfile
from src.material_suite.plan import (
    compile_material_suite_plan,
    discover_material_suite_bundle,
)
from src.material_suite.runner import MaterialSuiteGenerationRunner
from src.shared.engine.material_timeline import default_timeline_segment


def _write_docx(path: Path, *texts: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    document.save(path)


def _write_xlsx(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active["A1"] = value
    workbook.save(path)
    workbook.close()


def _archive() -> EntityArchive:
    first = default_timeline_segment(
        1,
        node_count=2,
        start_value="2026-01-01",
        end_value="2026-01-05",
    )
    second = default_timeline_segment(
        2,
        node_count=2,
        start_value="2026-02-01",
        end_value="2026-02-03",
    )
    return EntityArchive(
        archive_id="iso-projects",
        archive_name="ISO 项目资料包",
        profiles=[
            EntityProfile(
                profile_id="record-1",
                profile_name="项目甲",
                fields={
                    "产品名称": "产品A",
                    "项目名称": "项目甲",
                    "项目编号": "ISO-001",
                    "自定义文本1": "第一行\n第二行",
                },
                timeline_plans={"segment_1": first, "segment_2": second},
            )
        ],
    )


def test_suite_compiler_maps_common_word_and_route_excel_and_freezes_timelines(
    tmp_path,
):
    template_root = tmp_path / "模板文件"
    _write_docx(
        template_root / "Word文档" / "01 计划书.docx",
        "{{项目名称}}",
        "{{@time:时间节点1-2}} / {{时间节点2-2}}",
    )
    _write_xlsx(
        template_root / "产品A_测试表" / "测试记录表1.xlsx",
        "{{项目编号}} - {{项目名称}}",
    )

    bundle = discover_material_suite_bundle(template_root)
    plan = compile_material_suite_plan(
        _archive(),
        bundle,
        output_root=tmp_path / "输出",
    )

    assert plan.ok
    assert len(plan.records) == 1
    record = plan.records[0]
    assert record.route_key == "产品A"
    assert Path(record.output_dir) == tmp_path / "输出" / "产品A" / "项目甲"
    assert len(record.artifacts) == 2
    assert dict(record.frozen_values)["时间节点1-2"] == "2026-01-05"
    assert dict(record.frozen_values)["时间节点2-2"] == "2026-02-03"


def test_suite_runner_publishes_complete_record_and_preserves_manual_line_breaks(
    tmp_path,
):
    template_root = tmp_path / "模板文件"
    _write_docx(
        template_root / "Word文档" / "01 计划书.docx",
        "项目：{{项目名称}}",
        "说明：{{自定义文本1}}",
        "二段结束：{{@time:时间节点2-2}}",
    )
    _write_xlsx(
        template_root / "产品A_测试表" / "测试记录表1.xlsx",
        "{{项目编号}} - {{项目名称}}",
    )
    plan = compile_material_suite_plan(
        _archive(),
        discover_material_suite_bundle(template_root),
        output_root=tmp_path / "输出",
    )
    progress = []

    result = MaterialSuiteGenerationRunner(plan).run(
        lambda current, total, stage: progress.append((current, total, stage)),
        lambda: False,
    )

    assert result["status"] == "success"
    record_dir = tmp_path / "输出" / "产品A" / "项目甲"
    word_path = record_dir / "Word文档" / "01 计划书.docx"
    excel_path = record_dir / "Excel测试表" / "测试记录表1.xlsx"
    assert word_path.is_file()
    assert excel_path.is_file()
    assert (record_dir / "生成清单.json").is_file()
    document = Document(word_path)
    assert "项目：项目甲" in "\n".join(item.text for item in document.paragraphs)
    assert "第一行\n第二行" in "\n".join(item.text for item in document.paragraphs)
    workbook = load_workbook(excel_path, data_only=False)
    try:
        assert workbook.active["A1"].value == "ISO-001 - 项目甲"
    finally:
        workbook.close()
    assert progress[-1][0] == progress[-1][1] == 2


def test_suite_preflight_blocks_missing_fields_and_leaves_no_partial_directory(
    tmp_path,
):
    template_root = tmp_path / "模板文件"
    _write_docx(
        template_root / "Word文档" / "必需文档.docx",
        "{{不存在的字段}}",
    )
    plan = compile_material_suite_plan(
        _archive(),
        discover_material_suite_bundle(template_root),
        output_root=tmp_path / "输出",
    )

    assert not plan.ok
    assert "不存在的字段" in "；".join(plan.records[0].issues)
    result = MaterialSuiteGenerationRunner(plan).run(
        lambda *_args: None,
        lambda: False,
    )
    assert result["status"] == "failed"
    assert not (tmp_path / "输出" / "产品A" / "项目甲").exists()


def test_suite_does_not_map_another_products_excel_templates(tmp_path):
    template_root = tmp_path / "模板文件"
    _write_docx(template_root / "Word文档" / "通用.docx", "{{项目名称}}")
    _write_xlsx(
        template_root / "产品A_测试表" / "A.xlsx",
        "{{项目编号}}",
    )
    _write_xlsx(
        template_root / "产品B_测试表" / "B.xlsx",
        "{{项目编号}}",
    )

    plan = compile_material_suite_plan(
        _archive(),
        discover_material_suite_bundle(template_root),
        output_root=tmp_path / "输出",
    )

    assert [item.label for item in plan.records[0].artifacts] == ["通用", "A"]
