from __future__ import annotations

from hashlib import sha256

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.config.attachment_materials import (
    AttachmentBinding,
    AttachmentItem,
    AttachmentSourceKind,
)
from src.config.content_materials import FileAssetRef
from src.services.material_attachments.docx_renderer import (
    AttachmentRenderBlockedError,
    render_attachment_references,
)


def _binding(
    tmp_path,
    *,
    with_item=True,
    role="qualification_evidence",
):
    source_name = (
        "evidence.pdf"
        if role == "qualification_evidence"
        else f"{role}.pdf"
    )
    source = tmp_path / source_name
    source.write_bytes(b"pdf")
    item = AttachmentItem(
        item_id=f"{role}-1",
        label="ISO9001质量管理体系认证证书",
        file_ref=FileAssetRef(
            original_name=source.name,
            media_type="application/pdf",
            byte_size=source.stat().st_size,
            content_sha256=sha256(source.read_bytes()).hexdigest(),
            source_path=str(source),
        ),
    )
    return AttachmentBinding(
        role=role,
        label="资质附件",
        source_kind=AttachmentSourceKind.SINGLE_FILE,
        items=(item,) if with_item else (),
    )


def test_attachment_token_is_derived_and_inventory_renders_at_isolated_anchor(tmp_path):
    binding = _binding(tmp_path)
    document = Document()
    document.add_paragraph("附件清单")
    document.add_paragraph(binding.anchor_token)

    receipts = render_attachment_references(document, (binding,))

    assert binding.anchor_token == "{{@attach:qualification_evidence}}"
    assert receipts[0].rendered_item_count == 1
    assert [item.text for item in document.paragraphs] == [
        "附件清单",
        "附件：ISO9001质量管理体系认证证书（evidence.pdf）",
    ]


def test_absent_optional_attachment_anchor_is_a_noop(tmp_path):
    document = Document()
    document.add_paragraph("正文")

    receipt = render_attachment_references(document, (_binding(tmp_path),))[0]

    assert receipt.occurrence_count == 0
    assert document.paragraphs[0].text == "正文"


def test_present_anchor_with_empty_binding_fails_closed(tmp_path):
    binding = _binding(tmp_path, with_item=False)
    document = Document()
    document.add_paragraph(binding.anchor_token)

    with pytest.raises(AttachmentRenderBlockedError) as raised:
        render_attachment_references(document, (binding,))

    assert raised.value.diagnostics[0].code == "attachment_inventory_empty"


def test_attachment_anchor_in_table_is_rejected(tmp_path):
    binding = _binding(tmp_path)
    document = Document()
    document.add_table(1, 1).cell(0, 0).text = binding.anchor_token

    with pytest.raises(AttachmentRenderBlockedError) as raised:
        render_attachment_references(document, (binding,))

    assert raised.value.diagnostics[0].code == "attachment_anchor_not_isolated_body"


def test_existing_header_anchor_is_rejected_without_preflight_mutation(tmp_path):
    binding = _binding(tmp_path)
    document = Document()
    header = document.sections[0].header
    header.add_paragraph(binding.anchor_token)
    document_xml_before = document.element.xml
    header_xml_before = header._element.xml

    with pytest.raises(AttachmentRenderBlockedError) as raised:
        render_attachment_references(document, (binding,))

    assert "attachment_anchor_forbidden_surface" in {
        item.code for item in raised.value.diagnostics
    }
    assert document.element.xml == document_xml_before
    assert header._element.xml == header_xml_before


@pytest.mark.parametrize(
    ("corruption", "expected_code"),
    (
        ("bookmark", "attachment_anchor_not_isolated_body"),
        ("field", "attachment_anchor_not_isolated_body"),
        ("section", "attachment_anchor_section_boundary"),
    ),
)
def test_non_plain_attachment_anchor_is_blocked_atomically(
    tmp_path,
    corruption,
    expected_code,
):
    binding = _binding(tmp_path)
    document = Document()
    paragraph = document.add_paragraph(binding.anchor_token)
    if corruption == "bookmark":
        bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), "7")
        bookmark.set(qn("w:name"), "must-not-be-deleted")
        paragraph._p.append(bookmark)
    elif corruption == "field":
        run = OxmlElement("w:r")
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "begin")
        run.append(field)
        paragraph._p.append(run)
    else:
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    before = document.element.xml

    with pytest.raises(AttachmentRenderBlockedError) as raised:
        render_attachment_references(document, (binding,))

    assert expected_code in {
        item.code for item in raised.value.diagnostics
    }
    assert document.element.xml == before


def test_all_attachment_bindings_preflight_before_any_render(tmp_path):
    valid = _binding(tmp_path, role="a_valid")
    invalid = _binding(tmp_path, role="z_invalid")
    document = Document()
    document.add_paragraph(valid.anchor_token)
    paragraph = document.add_paragraph(invalid.anchor_token)
    paragraph._p.append(OxmlElement("w:bookmarkStart"))
    before = document.element.xml

    with pytest.raises(AttachmentRenderBlockedError):
        render_attachment_references(document, (valid, invalid))

    assert document.element.xml == before
    assert valid.anchor_token in document.paragraphs[0].text
