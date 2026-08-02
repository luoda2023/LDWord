from __future__ import annotations

from pathlib import Path

from docx import Document

from src.application.materials import MaterialRunBindRequest, bind_material_run
from src.domain.materials import (
    MaterialContract,
    MaterialFieldContract,
    MaterialPackage,
    MaterialPackageRef,
    MaterialRecord,
    MaterialRunSelection,
    MaterialScope,
    generate_package_id,
    generate_record_id,
)
from src.infrastructure.materials.codec import material_package_revision
from src.material_suite.plan import (
    compile_material_suite_plan,
    discover_material_suite_bundle,
)
from src.material_suite.runner import MaterialSuiteGenerationRunner


def _write_docx(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def _snapshot():
    package = MaterialPackage(
        package_id=generate_package_id(),
        display_name="成套资料",
        work_mode_id="custom",
        material_contract_id="suite_fields_v1",
        records=(
            MaterialRecord(
                record_id=generate_record_id(),
                display_name="项目甲",
                lifecycle="active",
                scope=MaterialScope(
                    fields={
                        "产品名称": "产品A",
                        "项目名称": "项目甲",
                        "项目编号": "ISO-001",
                    }
                ),
            ),
        ),
    )
    contract = MaterialContract(
        contract_id="suite_fields_v1",
        work_mode_id="custom",
        label="成套资料",
        fields=tuple(
            MaterialFieldContract(
                key=key,
                label=key,
                required=True,
                allowed_scopes=("record",),
            )
            for key in ("产品名称", "项目名称", "项目编号")
        ),
    )
    revision = material_package_revision(package)
    selection = MaterialRunSelection(
        package_ref=MaterialPackageRef(
            package_id=package.package_id,
            revision=revision,
        ),
        selected_record_ids=(package.records[0].record_id,),
    )
    result = bind_material_run(
        MaterialRunBindRequest(
            package=package,
            package_revision=revision,
            source_type="user",
            selection=selection,
            contract=contract,
            recipe_id="material_suite",
            scene_id="custom",
            work_mode_id="custom",
        ),
        object_path=lambda _item: "",
    )
    assert result.ok
    assert result.snapshot is not None
    return result.snapshot


def test_suite_consumes_only_frozen_execution_snapshot(tmp_path) -> None:
    template_root = tmp_path / "模板文件"
    _write_docx(
        template_root / "Word文档" / "计划书.docx",
        "{{项目编号}} / {{项目名称}}",
    )
    snapshot = _snapshot()

    plan = compile_material_suite_plan(
        snapshot,
        discover_material_suite_bundle(template_root),
        output_root=tmp_path / "输出",
    )
    result = MaterialSuiteGenerationRunner(plan).run(
        lambda *_args: None,
        lambda: False,
    )

    assert plan.ok
    assert plan.execution_snapshot.execution_ready
    assert (
        plan.execution_snapshot.preflight_receipt["material_snapshot_id"]
        == snapshot.snapshot_id
    )
    assert plan.execution_snapshot_id == plan.execution_snapshot.snapshot_id
    assert plan.records[0].record_id == snapshot.records[0].record_id
    assert result["status"] == "success"
    generated = (
        tmp_path
        / "输出"
        / "产品A"
        / "项目甲"
        / "Word文档"
        / "计划书.docx"
    )
    assert generated.is_file()
    text = "\n".join(item.text for item in Document(generated).paragraphs)
    assert "ISO-001 / 项目甲" in text
    assert result["material_suite_receipt"]["execution_snapshot_id"] == (
        plan.execution_snapshot.snapshot_id
    )
