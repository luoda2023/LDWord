import ast
import json
import sys
from pathlib import Path

from docx import Document
from PySide6.QtTest import QTest


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.qt_api import QApplication, QPoint, Qt, QWidget
from src.shared.ui.icons.catalog import get_icon_names
from src.shared.ui.form_row import FormRow
from src.shared.ui.styled_combo_box import SOURCE_BADGE_TEXT_ROLE
from src.shared.ui.theme import get_theme
from src.config import library as config_library
from src.config import material_package_library
from src.config.material_context import MaterialExecutionContext
from src.config.scene import SceneWorkspace
from src.config.builtin_templates import create_builtin_template
from src.config.scene_family_application import apply_planned_scene_family_defaults
from src.ui.bridge import PanelBridge
import src.ui.panels.scene_panel as scene_panel_module
import src.shared.engine.official_document_material_package as material_package_module
from src.ui.panels.scene_panel import ScenePanel, _set_combo_by_data
from src.ui.panels.scene_overview_projection import FIRST_SCREEN_BANNED_TERMS
from src.ui.panels.scene_scope_sections import DocumentScopeSection
from src.ui.panels.template_navigation_context import (
    build_template_navigation_context,
)
from src.ui.panels.scene_summary_projection import (
    build_academic_rule_source_evidence_summary_items,
    build_application_report_evidence_summary_items,
    build_compliance_summary_items,
    build_contract_field_evidence_summary_items,
    build_control_contract_summary_items,
    build_coverage_summary_items,
    build_delivery_summary_items,
    build_fixed_layout_batch_evidence_summary_items,
    build_input_profile_summary_items,
    build_journal_submission_evidence_summary_items,
    build_parameter_ownership_summary_items,
    build_product_readiness_summary_items,
    build_scene_request_cell_summary_items,
    build_scene_overview_summary_items,
    build_scene_scope_summary_items,
    build_scene_sample_fixture_detail_text,
    build_scene_sample_fixture_summary_items,
    recommended_object_preflight_targets_for_scene,
    scene_request_cell_filter_options,
    scene_request_cell_fixture_specs_for_scene,
    scene_request_cell_list_item_projection,
    scene_request_cell_matches_filter,
)
from src.config.scene_presets import (
    build_scene_profile_summary,
    create_bidding_scene,
    create_exam_scene,
    create_official_scene,
    create_technical_scene,
    create_thesis_scene,
)


def test_scene_surfaces_do_not_hide_invalid_default_delivery_with_first_preset():
    scene = create_official_scene()
    first_id = scene.delivery_presets[0].preset_id
    scene.default_delivery_preset_id = "missing_delivery"

    items = {item.key: item for item in build_delivery_summary_items(scene)}

    assert first_id != "missing_delivery"
    assert items["default_delivery"].value == "无效引用：missing_delivery"
    assert items["default_delivery"].variant == "warning"
    assert items["default_artifacts"].value == "未设置"
    assert items["default_artifacts"].variant == "warning"
    assert "交付 missing_delivery" in build_scene_profile_summary(scene)


def _app():
    return QApplication.instance() or QApplication([])


def test_scene_detail_implementations_keep_one_way_module_ownership():
    panel_path = ROOT / "src/ui/panels/scene_panel.py"
    panel_source = panel_path.read_text(encoding="utf-8")
    panel_tree = ast.parse(panel_source)
    panel_classes = {
        node.name for node in panel_tree.body if isinstance(node, ast.ClassDef)
    }

    owned_details = {
        "scene_detail_base.py": {"_SimpleFormDetail"},
        "scene_content_detail.py": {"_ContentDetail"},
        "scene_output_detail.py": {"_OutputDetail"},
        "scene_exam_detail.py": {"ExamPaperDetail"},
    }
    assert panel_classes.isdisjoint(
        class_name
        for class_names in owned_details.values()
        for class_name in class_names
    )

    for module_name, class_names in owned_details.items():
        module_source = (ROOT / "src/ui/panels" / module_name).read_text(
            encoding="utf-8"
        )
        module_tree = ast.parse(module_source)
        assert class_names.issubset({
            node.name for node in module_tree.body if isinstance(node, ast.ClassDef)
        })
        assert all(
            not (
                isinstance(node, ast.ImportFrom)
                and node.module == "src.ui.panels.scene_panel"
            )
            for node in ast.walk(module_tree)
        )

    assert "from src.ui.panels.scene_content_detail import _ContentDetail" in panel_source
    assert "from src.ui.panels.scene_output_detail import _OutputDetail" in panel_source
    assert "from src.ui.panels.scene_exam_detail import (" in panel_source
    assert "self._exam_paper = ExamPaperDetail(self.bridge)" in panel_source
    exam_detail_source = (ROOT / "src/ui/panels/scene_exam_detail.py").read_text(
        encoding="utf-8"
    )
    assert "ExamPaperPreviewDialog" not in exam_detail_source


def test_scene_activation_and_edit_transactions_have_one_real_owner():
    panel_source = (ROOT / "src/ui/panels/scene_panel.py").read_text(
        encoding="utf-8"
    )
    owner_source = (
        ROOT / "src/ui/panels/scene_session_coordinator.py"
    ).read_text(encoding="utf-8")
    owner_tree = ast.parse(owner_source)

    assert "SceneSessionCoordinator" in {
        node.name for node in owner_tree.body if isinstance(node, ast.ClassDef)
    }
    assert "class PreparedSceneChanges" in owner_source
    assert "class SceneActivationSnapshot" in owner_source
    assert "capture_state_snapshot()" in owner_source
    assert "restore_state_snapshot(" in owner_source
    assert "def resolve_template_binding(" in owner_source
    assert "def activate(" in owner_source
    assert "def commit_prepared_changes(" in owner_source
    assert "def rollback_prepared_changes(" in owner_source
    assert "scene_user_target_path(" in owner_source
    assert "def _assert_prepared_revision(" in owner_source
    assert "def _load_canonical_committed_scene(" in owner_source
    assert "target = prepared.committed_path or prepared.target_path" in owner_source
    assert "from src.ui.panels.scene_panel" not in owner_source

    assert "self._scene_session.activate(" in panel_source
    assert "self._scene_session.prepare_pending_changes(" in panel_source
    assert "self._scene_session.commit_prepared_changes()" in panel_source
    assert "def _capture_scene_activation_snapshot(" not in panel_source
    assert "def _restore_scene_activation_snapshot(" not in panel_source
    assert "def _adopt_committed_scene(" not in panel_source
    assert "self._prepared_scene_changes" not in panel_source


def test_exam_detail_separates_surface_builders_and_official_projection():
    detail_source = (ROOT / "src/ui/panels/scene_exam_detail.py").read_text(
        encoding="utf-8"
    )
    detail_tree = ast.parse(detail_source)
    detail_class = next(
        node
        for node in detail_tree.body
        if isinstance(node, ast.ClassDef) and node.name == "ExamPaperDetail"
    )
    methods = {
        node.name: node
        for node in detail_class.body
        if isinstance(node, ast.FunctionDef)
    }

    assert methods["__init__"].end_lineno - methods["__init__"].lineno + 1 <= 25
    assert (
        methods["_refresh_official_preview"].end_lineno
        - methods["_refresh_official_preview"].lineno
        + 1
        <= 25
    )
    assert "def _build_exam_master_controls(" in detail_source
    assert "def _build_official_preview_controls(" in detail_source
    assert "build_official_plan_preview_projection(" in detail_source
    assert "check_master_preflight(" not in detail_source


def test_scene_plan_preview_uses_authoritative_official_mode_for_custom_scene():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = SceneWorkspace(
        scene_id="custom_plan",
        category="custom",
        template_id="official_gbt",
    )
    bridge.set_current_scene(scene, config_id="custom_plan", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._exam_paper._active_preview_provider_id == "official"
        assert panel._exam_paper._plan_preview_card.isHidden() is False
        assert panel._exam_paper._preview_mode_control.isHidden() is True
        assert panel._nav_cards["scn_exam_paper"].isHidden() is False
    finally:
        panel.close()
        app.processEvents()


def test_scene_ui_hot_imports_do_not_depend_on_release_or_static_dashboards():
    ui_paths = (
        ROOT / "src/ui/panels/scene_panel.py",
        ROOT / "src/ui/panels/scene_summary_projection.py",
    )

    for path in ui_paths:
        source = path.read_text(encoding="utf-8")
        tree = ast.parse(source)
        top_level_modules = {
            node.module
            for node in ast.walk(tree)
            if isinstance(node, ast.ImportFrom) and node.module
        }
        top_level_modules.update(
            alias.name
            for node in ast.walk(tree)
            if isinstance(node, ast.Import)
            for alias in node.names
        )
        forbidden = {
            module
            for module in top_level_modules
            if module.startswith("src.config.scene_matrix_dashboard")
            or module == "src.config.scene_matrix_drilldown"
            or (
                module.startswith("src.config.scene_")
                and "release" in module
            )
        }
        assert forbidden == set(), f"{path.name}: {sorted(forbidden)}"

    summary_source = ui_paths[1].read_text(encoding="utf-8")
    assert "build_scene_matrix_dashboard_summary_items" not in summary_source
    assert "_scene_matrix_dashboard_tooltip" not in summary_source


def test_scene_panel_removes_page_elements_detail_from_scene_surface():
    source = (ROOT / "src/ui/panels/scene_panel.py").read_text(encoding="utf-8")

    assert "self._page_number = ToggleSwitch" not in source
    assert "hf.page_number_enabled =" not in source
    assert "_PageElementsDetail" not in source
    assert "_page_elem" not in source
    assert "scn_page_elem" not in source


def test_document_scope_section_contains_only_plan_intent_controls():
    _app()
    section = DocumentScopeSection()
    try:
        section.set_mode_id("thesis")
        assert section.mode_control.count() == 3
        assert [
            section.mode_control.segment_text(index)
            for index in range(section.mode_control.count())
        ] == ["全部内容", "仅正文", "指定区域"]
        assert set(section._role_checks) == {
            "abstract_cn",
            "abstract_en",
            "toc",
            "body",
            "references",
            "appendix",
            "acknowledgment",
            "resume",
        }

        section.set_scope("selected", ("references", "appendix"))
        assert section.mode() == "selected"
        assert section.selected_roles() == ["references", "appendix"]
        assert section._roles_widget.isHidden() is False

        section.set_scope("all", ())
        assert section._roles_widget.isHidden() is True
        assert section.findChild(QWidget, "document_path") is None
    finally:
        section.close()


def test_scene_panel_projects_scene_profiles_with_shared_summary_grid():
    source = (ROOT / "src/ui/panels/scene_panel.py").read_text(encoding="utf-8")
    surface_source = "\n".join(
        (
            source,
            (ROOT / "src/ui/panels/scene_content_detail.py").read_text(
                encoding="utf-8"
            ),
            (ROOT / "src/ui/panels/scene_output_detail.py").read_text(
                encoding="utf-8"
            ),
        )
    )
    product_projection_source = (
        ROOT / "src/ui/panels/scene_product_summary_projection.py"
    ).read_text(encoding="utf-8")
    engineering_projection_source = (
        ROOT / "src/ui/panels/scene_summary_projection.py"
    ).read_text(encoding="utf-8")

    assert "scene_product_summary_projection" in surface_source
    assert "SummaryGrid" in surface_source
    assert "build_scene_overview_spec" in surface_source
    assert "DocumentScopeSection" in surface_source
    assert "build_input_profile_summary_items" in surface_source
    assert "build_compliance_summary_items" in surface_source
    assert (
        "build_parameter_ownership_summary_items"
        in engineering_projection_source
    )
    assert "build_coverage_summary_items" in engineering_projection_source
    assert "build_scene_scope_summary_items" in product_projection_source
    assert "build_delivery_summary_items" in surface_source
    assert "input_source_profile" in surface_source
    assert "compliance_profile" in surface_source
    assert "default_delivery_preset_id" in surface_source


def test_scene_navigation_cards_use_readable_status_subtitles():
    _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        category="contract_delivery",
        template_id="default",
    )
    scene.input_source_profile.accepted_formats = ["docx", "xlsx"]
    scene.input_source_profile.required_material_fields = [
        "company_name",
        "contract_no",
    ]
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.compliance_profile.rule_family = "contract_delivery"
    scene.compliance_profile.count_profile_id = "contract_fields"
    panel = ScenePanel(bridge)
    try:
        snapshots = panel._navigation_card_snapshots(scene)
        content_subtitle = snapshots["scn_content"]["subtitle"]
        rules_subtitle = snapshots["scn_rules"]["subtitle"]

        assert "scn_features" not in snapshots
        assert "scn_table_chart" not in snapshots
        assert "scn_formula" not in snapshots
        assert "scn_citation" not in snapshots
        assert "scn_scope" not in snapshots
        assert "scn_style_rules" not in snapshots
        assert "scn_reference" not in snapshots
        assert "scn_cleanup" not in snapshots
        assert "scn_output" not in snapshots
        assert "Word 文档、Excel 表格" in content_subtitle
        assert "资料字段 2 个" in content_subtitle
        assert "资料规则 1 项" in content_subtitle
        assert "schema" not in content_subtitle.lower()
        assert "/" not in content_subtitle
        assert "_" not in content_subtitle
        assert "/" not in rules_subtitle
        assert "全部内容" in rules_subtitle
        assert "产物" in rules_subtitle
        assert snapshots["scn_rules"]["badge_text"] == "全部内容"
        assert snapshots["scn_content"]["badge_text"] == "已配置"
    finally:
        panel.close()


def test_scene_panel_does_not_mount_template_page_number_summary_in_scene():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="thesis",
        template_id="thesis_gbt",
        compatible_template_ids=["thesis_gbt"],
    )
    template = create_builtin_template("thesis_gbt")
    bridge.set_current_scene(scene, config_id="thesis", emit_signal=False)
    bridge.set_current_template(template, config_id="thesis_gbt", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert hasattr(panel, "_page_elem") is False
        assert "scn_page_elem" not in panel._detail_map
        assert "scn_page_elem" not in panel._nav_cards
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_template_change_skips_removed_page_number_detail():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="thesis",
        template_id="thesis_gbt",
        compatible_template_ids=["thesis_gbt"],
    )
    template = create_builtin_template("thesis_gbt")
    bridge.set_current_scene(scene, config_id="thesis", emit_signal=False)
    bridge.set_current_template(template, config_id="thesis_gbt", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        updated_template = create_builtin_template("thesis_gbt")
        updated_template.header_footer.page_number_enabled = False
        bridge.set_current_template(updated_template, config_id="thesis_gbt")
        app.processEvents()

        assert hasattr(panel, "_page_elem") is False
        assert "scn_page_elem" not in panel._detail_map
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_profile_summaries_reflect_current_scene():
    app = _app()
    bridge = PanelBridge()
    scene = create_bidding_scene()
    template = create_builtin_template("bid_engineering")
    bridge.set_current_scene(scene, config_id="bidding", emit_signal=False)
    bridge.set_current_template(
        template, config_id="bid_engineering", emit_signal=False
    )

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert not hasattr(panel._overview, "_profile_summary")
        assert not hasattr(panel._overview, "_advanced_evidence_card")
        assert panel._content._input_summary.value_for("materials") == "必需"
        material_detail = panel._content._input_summary.detail_for("materials")
        assert "标书资料" in material_detail
        assert "bid_materials_v1" not in material_detail
        input_projection = {
            item.key: item for item in build_input_profile_summary_items(scene)
        }
        assert "bid_materials_v1" in input_projection["materials"].tooltip
        assert panel._content._input_summary.value_for("material_fields") == "3 个字段"
        assert "公司名称" in panel._content._input_summary.detail_for("material_fields")
        assert panel._content._input_summary.value_for("image_roles") == "2 个角色"
        assert "印章" in panel._content._input_summary.detail_for("image_roles")
        assert panel._cleanup._compliance_summary.value_for("object_policy") == "启用"
        assert "严格" in panel._cleanup._compliance_summary.detail_for("object_policy")
        assert (
            panel._cleanup._compliance_summary.value_for("scan_targets") == "11 个目标"
        )
        assert "content_controls" in panel._cleanup._compliance_summary.detail_for(
            "scan_targets"
        )
        assert "hidden_text" in panel._cleanup._compliance_summary.detail_for(
            "scan_targets"
        )
        assert (
            panel._output._delivery_summary.value_for("preset_count") == "3 个输出版本"
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_overview_summary_projects_planning_family_governance():
    scene = SceneWorkspace(scene_id="contract_delivery")
    scene.input_source_profile.material_schema_id = "contract_parties_v1"

    items = {item.key: item for item in build_scene_overview_summary_items(scene)}

    assert items["planning_family"].value == "合同交付"
    assert "P1" in items["planning_family"].detail
    assert "核对字段一致性" in items["planning_workflows"].detail
    assert "修订保护" in items["planning_first_slice"].detail
    assert "批注" in items["planning_ooxml"].detail
    assert "修订" in items["planning_ooxml"].detail
    assert "revision" in items["planning_ooxml"].tooltip


def test_scene_overview_summary_projects_parameter_ownership_registry():
    scene = SceneWorkspace(scene_id="contract_delivery")
    items = {item.key: item for item in build_scene_overview_summary_items(scene)}
    ownership_items = {
        item.key: item for item in build_parameter_ownership_summary_items(scene)
    }

    assert ownership_items["parameter_ownership"].value == "归属清楚"
    assert items["parameter_ownership"].value == "归属清楚"
    assert "模板管样式" in items["parameter_ownership"].detail
    assert "方案管流程" in items["parameter_ownership"].detail
    assert "资料管输入" in items["parameter_ownership"].detail
    assert "输出管版本" in items["parameter_ownership"].detail
    assert "13" not in items["parameter_ownership"].detail
    assert "37" not in items["parameter_ownership"].detail
    assert "归属计数" in items["parameter_ownership"].tooltip
    assert "SceneWorkspace" in items["parameter_ownership"].tooltip
    assert items["parameter_boundary"].value == "按字段应用"
    assert "固定行高" in items["parameter_boundary"].detail
    assert "固定版式" in items["parameter_boundary"].detail
    assert "待裁决" not in items["parameter_boundary"].detail


def test_scene_overview_summary_projects_control_contract_registry():
    scene = SceneWorkspace(scene_id="contract_delivery")
    items = {item.key: item for item in build_scene_overview_summary_items(scene)}
    control_items = {item.key: item for item in build_control_contract_summary_items()}

    assert control_items["control_contract"].value == "已统一"
    assert items["control_contract"].value == "已统一"
    assert items["control_contract"].detail == (
        "缩进、段落间距、固定版式行高、公式、水印等已归一"
    )
    assert "左缩进" in items["control_contract"].tooltip
    assert "右缩进" in items["control_contract"].tooltip
    assert "特殊缩进" in items["control_contract"].tooltip
    assert "段前" in items["control_contract"].tooltip
    assert "段后" in items["control_contract"].tooltip
    assert "固定版位行高" in items["control_contract"].tooltip
    assert "公式策略" in items["control_contract"].tooltip
    assert "水印状态" in items["control_contract"].tooltip
    assert "资料规则" in items["control_contract"].tooltip
    assert "输出版本" in items["control_contract"].tooltip
    assert "内容显隐" in items["control_contract"].tooltip
    assert "插件人工确认" in items["control_contract"].tooltip
    assert items["control_contract_scope"].value == "16 类控件"
    assert items["control_contract_scope"].detail == "5 类控件归属已分开"
    assert "模板管样式 9" in items["control_contract_scope"].tooltip
    assert "方案管流程 3" in items["control_contract_scope"].tooltip
    assert "资料管输入 1" in items["control_contract_scope"].tooltip
    assert "输出管版本 2" in items["control_contract_scope"].tooltip
    assert "插件管人工确认 1" in items["control_contract_scope"].tooltip
    assert "同名格式参数" in items["control_contract_scope"].tooltip


def test_scene_panel_content_watermark_controls_write_scene_contracts():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(scene_id="technical")
    bridge.set_current_scene(scene, config_id="technical", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert not hasattr(panel, "_formula")

        panel._content._watermark_enabled.setChecked(True)
        panel._content._watermark_text.setText("征求意见")
        panel._content._on_edited()

        assert scene.watermark.enabled is True
        assert scene.watermark.text == "征求意见"
        assert panel._content._watermark_text.isEnabled() is True

        panel._content._watermark_enabled.setChecked(False)
        panel._content._on_edited()

        assert scene.watermark.enabled is False
        assert scene.watermark.text == "征求意见"
        assert panel._content._watermark_text.isEnabled() is False
    finally:
        panel.close()
        app.processEvents()


def test_scene_overview_summary_omits_completed_coverage_closure_tasks():
    scene = SceneWorkspace(scene_id="contract_delivery")
    scene.input_source_profile.material_schema_id = "contract_parties_v1"

    items = {item.key: item for item in build_scene_overview_summary_items(scene)}
    coverage_items = {item.key: item for item in build_coverage_summary_items(scene)}

    assert coverage_items["coverage_pack"].value == "合同交付"
    assert "不提供法律意见" in coverage_items["coverage_pack"].detail
    assert items["coverage_pack"].value == "合同交付"
    assert "coverage_next_closure" not in items


def test_scene_overview_summary_projects_sample_fixture_coverage():
    scene = SceneWorkspace(scene_id="contract_delivery", category="contract_delivery")
    sample_items = {
        item.key: item for item in build_scene_sample_fixture_summary_items(scene)
    }
    request_items = {
        item.key: item for item in build_scene_request_cell_summary_items(scene)
    }
    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}
    detail_text = build_scene_sample_fixture_detail_text(scene)
    request_cell_specs = scene_request_cell_fixture_specs_for_scene(scene)

    assert sample_items["sample_fixture_coverage"].value == "已有样本文档"
    assert "2 个样本" in sample_items["sample_fixture_coverage"].detail
    assert "修订" in sample_items["sample_fixture_coverage"].tooltip
    assert (
        "scene_sample_fixture_registry"
        not in sample_items["sample_fixture_coverage"].tooltip
    )
    assert "覆盖 pack" not in sample_items["sample_fixture_coverage"].tooltip
    assert overview["sample_fixture_coverage"].value == "已有样本文档"
    assert overview["sample_fixture_boundary"].value == "有边界说明"
    assert "法律审查" in overview["sample_fixture_boundary"].detail
    assert request_items["request_cell_coverage"].value == "常见说法已覆盖"
    assert "4 个请求" in request_items["request_cell_coverage"].detail
    assert "0 个借用样本" in request_items["request_cell_coverage"].detail
    assert (
        "常见说法登记：高频用户说法" in request_items["request_cell_coverage"].tooltip
    )
    assert (
        "scene_request_cell_fixture_registry"
        not in request_items["request_cell_coverage"].tooltip
    )
    assert "fixture-backed" not in request_items["request_cell_coverage"].tooltip
    assert overview["request_cell_coverage"].value == "常见说法已覆盖"
    assert "4 个请求" in overview["request_cell_coverage"].detail
    assert "资料包：合同交付" in detail_text
    assert "合同交付样本 1" in detail_text
    assert "常见说法：" in detail_text
    assert "样本：contract_delivery_revisions" not in detail_text
    assert "样本：contract_signing_consistency" not in detail_text
    assert "contract_delivery_revisions" not in detail_text
    assert "contract_signing_consistency" not in detail_text
    assert "用户说法：合同签署包字段一致性" in detail_text
    assert "覆盖方式：直接证据" in detail_text
    assert "证据样本：1 个" in detail_text
    assert "request-cells:" not in detail_text
    assert "fixture=" not in detail_text
    assert "cell=" not in detail_text
    assert "level=" not in detail_text
    assert {cell.sample_id for cell in request_cell_specs} == {
        "contract_signing_consistency",
        "contract_review_revisions",
        "contract_signature_package_fields",
        "ambiguous_contract_legal_review",
    }
    assert {cell.coverage_level for cell in request_cell_specs} == {
        "direct_family_fixture",
        "ambiguous_fixture_set",
    }
    assert ("Word 对象：正文片段 / 域 / 批注 / 修订 / 隐藏文字 / 嵌入附件") in (
        detail_text
    )
    assert "边界：合同样本只验证格式和字段，不代表法律审查" in detail_text


def test_scene_overview_summary_projects_contract_green_with_field_evidence():
    scene = SceneWorkspace(scene_id="contract_delivery", category="contract_delivery")

    readiness_items = {
        item.key: item for item in build_product_readiness_summary_items(scene)
    }
    field_items = {
        item.key: item for item in build_contract_field_evidence_summary_items(scene)
    }
    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}

    assert readiness_items["product_readiness"].value == "可直接使用"
    assert readiness_items["product_readiness_family"].value == "可直接使用"
    assert "product_readiness_gaps" not in readiness_items
    assert field_items["contract_field_evidence"].value == "已打通"
    assert "合同方字段资料" in field_items["contract_field_evidence"].detail
    assert "字段一致性报告" in field_items["contract_field_evidence"].detail
    assert "contract_parties_v1" in field_items["contract_field_evidence"].tooltip
    assert "contract_field_consistency_report" in (
        field_items["contract_field_evidence"].tooltip
    )
    assert "contract_delivery_signature_fields_degraded" in (
        field_items["contract_field_evidence"].tooltip
    )
    assert overview["product_readiness"].value == "可直接使用"
    assert "product_readiness_gaps" not in overview
    assert overview["contract_field_evidence"].value == "已打通"


def test_scene_overview_summary_projects_plugin_boundary_for_exam_family():
    scene = SceneWorkspace(scene_id="exam_teaching")
    scene.category = "exam_teaching"

    items = {item.key: item for item in build_scene_overview_summary_items(scene)}

    assert items["coverage_pack"].value == "试卷/教学资料"
    assert "coverage_next_closure" not in items
    assert items["coverage_plugin_boundary"].value == "1 项需确认"
    assert "AI 内容质量" in items["coverage_plugin_boundary"].detail


def test_scene_family_application_uses_shared_content_visibility_rule_labels():
    scene = SceneWorkspace(scene_id="exam_teaching", category="exam_teaching")

    result = apply_planned_scene_family_defaults(scene)

    assert result.applied is True
    student_version = next(
        preset
        for preset in scene.delivery_presets
        if preset.preset_id == "student_version"
    )
    assert student_version.label == "学生版"
    assert [
        (rule.selector, rule.label, rule.action)
        for rule in student_version.content_visibility_rules
    ] == [
        ("answer", "答案", "remove"),
        ("analysis", "解析", "remove"),
        ("solution", "解题过程", "remove"),
        ("teacher_note", "教师备注", "remove"),
        ("knowledge_points", "知识点", "remove"),
    ]

    answer_sheet = next(
        preset
        for preset in scene.delivery_presets
        if preset.preset_id == "answer_sheet"
    )
    assert answer_sheet.label == "答题卡"
    assert [rule.label for rule in answer_sheet.content_visibility_rules] == [
        "答案",
        "解析",
        "解题过程",
        "教师备注",
        "知识点",
        "题干正文",
    ]
    assert all(":" not in rule.label for rule in answer_sheet.content_visibility_rules)


def test_scene_panel_family_delivery_combo_uses_shared_display_labels():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="exam_teaching",
        name="试卷/教学资料",
        category="exam_teaching",
        template_id="default",
        compatible_template_ids=["default"],
    )
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="exam_teaching", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        family_delivery_tooltip = panel._output._apply_family_delivery_btn.toolTip()
        assert "将新增：学生版、教师版、答案速查等 5 个" in family_delivery_tooltip
        assert "Student version" not in family_delivery_tooltip

        panel._output._apply_family_delivery_btn.click()
        app.processEvents()

        student_index = panel._output._default_delivery.findData("student_version")
        answer_index = panel._output._default_delivery.findData("answer_key")
        assert student_index >= 0
        assert answer_index >= 0
        assert panel._output._default_delivery.itemText(student_index) == "学生版"
        assert panel._output._default_delivery.itemText(answer_index) == "答案速查"
        assert panel._output._delivery_label.text() == "学生版"
        assert "Student version" not in panel._output._delivery_label.text()
        student_preset = next(
            preset
            for preset in scene.delivery_presets
            if preset.preset_id == "student_version"
        )
        assert student_preset.label == "学生版"
    finally:
        panel.close()
        app.processEvents()


def test_scene_overview_summary_projects_official_archive_profile_defaults():
    scene = create_official_scene()
    apply_planned_scene_family_defaults(scene)

    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}
    compliance = {item.key: item for item in build_compliance_summary_items(scene)}
    delivery = {item.key: item for item in build_delivery_summary_items(scene)}

    assert overview["planning_family"].value == "公文/会议材料"
    assert "正式归档交付" in overview["planning_workflows"].detail
    assert "归档信息" in overview["planning_first_slice"].detail
    assert "公文/会议纪要" in overview["coverage_pack"].detail
    assert "coverage_next_closure" not in overview
    assert overview["official_policy_evidence"].label == "公文元数据证据"
    assert overview["official_policy_evidence"].value == "已打通"
    assert "公文字段资料" in overview["official_policy_evidence"].detail
    assert "报告 4 项" in overview["official_policy_evidence"].detail
    assert "official_metadata_report" in overview["official_policy_evidence"].tooltip
    assert "formal_internal_archive_manifest" in (
        overview["official_policy_evidence"].tooltip
    )
    assert "official_policy_metadata_archive_report" in (
        overview["official_policy_evidence"].tooltip
    )
    assert (
        "Official metadata evidence" not in overview["official_policy_evidence"].label
    )
    assert "ready" not in overview["official_policy_evidence"].value
    assert overview["product_readiness"].value == "可直接使用"
    assert "product_readiness_gaps" not in overview
    assert delivery["default_delivery"].value == "正式纪要"
    assert delivery["default_delivery"].detail == "执行时默认生成"
    assert "formal_minutes" in delivery["default_delivery"].tooltip
    assert "资料清单" in delivery["default_artifacts"].value
    assert "政策资料归档" in delivery["preset_count"].detail
    assert "归档清单" in delivery["preset_count"].detail
    assert "policy_collection" in delivery["preset_count"].tooltip
    assert "archive_manifest" in delivery["preset_count"].tooltip
    assert compliance["count_profile"].value == "administrative_sections"
    assert "Administrative section" in compliance["count_profile"].detail
    assert compliance["planning_scan_targets"].value in {"已应用", "已覆盖"}


def test_scene_panel_official_plan_preview_uses_official_master_contract(monkeypatch):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()
        assert panel._nav_cards["scn_exam_paper"].isHidden() is False
        assert panel._nav_cards["scn_exam_paper"]._title.text() == "方案概览"

        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()

        detail = panel._exam_paper
        assert detail._plan_preview_card.isHidden() is False
        assert detail._master_card.isHidden() is True
        assert detail._preview_mode_control.isHidden() is True
        assert detail._prompt_card.isHidden() is True
        assert not hasattr(detail, "_official_card")
        assert not hasattr(detail, "_preview_card")
        assert not hasattr(detail, "_official_word_preview_controller")
        assert not hasattr(detail, "_exam_word_preview_controller")
        assert detail._active_preview_provider_id == "official"
        assert detail._official_profile.findData("notice") >= 0
        assert detail._official_profile.findData("minutes") >= 0
        assert (
            detail._official_material_sample.findData("builtin/notice_archive_check")
            >= 0
        )
        assert (
            detail._official_material_sample.findData("builtin/letter_material_request")
            < 0
        )
        assert (
            detail._official_material_sample.findData("builtin/minutes_coordination")
            < 0
        )
        assert (
            detail._official_material_sample.currentData()
            == "builtin/notice_archive_check"
        )
        assert detail._official_master.findData("official_gbt_standard") >= 0
        assert "official_gbt" in detail._official_template.text()
        assert detail._official_profile.isHidden()
        assert detail._official_material_sample.isHidden()
        assert detail._official_master.isHidden()
        assert detail._official_action_row.isHidden()
        assert not hasattr(detail, "_official_navigation_row")
        assert not hasattr(detail, "_official_material_nav_btn")
        assert not hasattr(detail, "_official_master_nav_btn")
        assert "资料状态：尚未填写" in detail._official_material_status.text()
        assert not hasattr(detail, "_official_preview_page")
        assert not hasattr(detail, "_official_preview_mode_control")
        assert (
            detail._document_word_preview.toolbar_widget().parent()
            is detail._plan_preview_card._header_widget
        )
        assert "official_document_v1" in detail._official_contract.text()
        assert "official_title" in detail._official_contract.text()
        assert "装配检查" in detail._official_contract.text()
        assert "母版" not in detail._official_contract.text()

        _set_combo_by_data(detail._official_profile, "minutes")
        app.processEvents()
        assert detail._official_profile.currentData() == "minutes"
        assert (
            detail._official_material_sample.findData("builtin/minutes_coordination")
            >= 0
        )
        assert (
            detail._official_material_sample.findData("builtin/notice_archive_check")
            < 0
        )
        assert (
            detail._official_material_sample.currentData()
            == "builtin/minutes_coordination"
        )
        assert detail._official_master.currentData() == "official_gbt_minutes"
        assert "administrative_meeting_fields_v1" in detail._official_contract.text()

        snapshot = panel._navigation_card_snapshots(scene)["scn_exam_paper"]
        assert snapshot["badge_text"] == "公文版式"
        assert "GB/T 9704 通用红头公文版式" in snapshot["subtitle"]
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_navigation_uses_bridge_mode_when_scene_identity_is_stale():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = ScenePanel(bridge)
    try:
        stale_scene = SceneWorkspace(
            scene_id="custom",
            category="custom",
            mode_id="custom",
            template_id="default",
        )
        panel._current_scene = stale_scene
        panel._update_dynamic_card_visibility()
        app.processEvents()

        assert bridge.current_work_mode_id() == "official"
        assert panel._nav_cards["scn_exam_paper"].isHidden() is False
        assert panel._nav_cards["scn_overview"].isHidden() is True
    finally:
        panel.close()
        app.processEvents()


def test_official_open_scene_folder_button_keeps_label_but_targets_user_docx_pool(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        scene_panel_module,
        "OFFICIAL_USER_MASTER_DIR",
        tmp_path / "config_library" / "masters" / "official" / "user",
    )
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._overview._open_scene_folder_btn.text() == "打开方案文件夹"
        assert panel._current_scene_folder() == (
            tmp_path / "config_library" / "masters" / "official" / "user"
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_plan_minutes_type_flows_to_preview_and_material_fill(
    monkeypatch,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = config_library.load_scene_from_library("official", mode_id="official")
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_official_document_type_id("minutes", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()
        preview_detail = panel._exam_paper
        assert preview_detail._official_profile.currentData() == "minutes"
        assert (
            preview_detail._official_material_sample.currentData()
            == "builtin/minutes_coordination"
        )
        assert (
            "administrative_meeting_fields_v1"
            in preview_detail._official_contract.text()
        )

        panel._nav_rail.select_card("scn_content")
        app.processEvents()
        content_detail = panel._content
        assert content_detail._official_material_profile.currentData() == "minutes"
        assert (
            content_detail._official_material_rows["meeting_date"].isHidden() is False
        )
        assert (
            content_detail._official_material_rows["participants"].isHidden() is False
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_plan_letter_type_uses_builtin_material_sample(
    monkeypatch,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = config_library.load_scene_from_library("official", mode_id="official")
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_official_document_type_id("letter", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()
        detail = panel._exam_paper

        assert detail._official_profile.currentData() == "letter"
        assert (
            detail._official_material_sample.findData("builtin/letter_material_request")
            >= 0
        )
        assert (
            detail._official_material_sample.findData("builtin/notice_archive_check")
            < 0
        )
        assert (
            detail._official_material_sample.findData("builtin/minutes_coordination")
            < 0
        )
        assert (
            detail._official_material_sample.currentData()
            == "builtin/letter_material_request"
        )
        assert "函件资料包样例" in detail._official_material_sample.currentText()
        assert detail._official_apply_sample_btn.isEnabled() is True
        assert "official_document_v1" in detail._official_contract.text()
        assert (
            "administrative_meeting_fields_v1" not in detail._official_contract.text()
        )

        detail._apply_official_sample_material()
        app.processEvents()
        context = bridge.current_material_context()
        assert context.profile_id == "official:letter"
        assert context.profile_name == "函"
        assert context.entity_data["document_type"] == "letter"
        assert context.entity_data["title"] == "关于商请协助提供归档材料的函"
        assert "必填字段已齐" in detail._official_material_status.text()

        _set_combo_by_data(detail._official_profile, "notice")
        app.processEvents()
        assert detail._official_profile.currentData() == "notice"
        assert (
            detail._official_material_sample.findData("builtin/notice_archive_check")
            >= 0
        )
        assert (
            detail._official_material_sample.findData("builtin/letter_material_request")
            < 0
        )
        assert (
            detail._official_material_sample.currentData()
            == "builtin/notice_archive_check"
        )
        assert detail._official_apply_sample_btn.isEnabled() is True

        _set_combo_by_data(detail._official_profile, "letter")
        app.processEvents()
        assert detail._official_profile.currentData() == "letter"
        assert (
            detail._official_material_sample.currentData()
            == "builtin/letter_material_request"
        )
        assert "函件资料包样例" in detail._official_material_sample.currentText()
        assert detail._official_apply_sample_btn.isEnabled() is True
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_plan_preview_applies_sample_and_generates_docx(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr("src.ui.panels.scene_exam_detail.PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.Toast.show_success",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.Toast.show_warning",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.Toast.show_error",
        lambda *args, **kwargs: None,
    )
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    opened_paths: list[Path] = []
    try:
        app.processEvents()
        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()

        detail = panel._exam_paper
        detail._open_sample_path_handler = lambda path: (
            opened_paths.append(Path(path)) or True
        )
        _set_combo_by_data(detail._official_profile, "minutes")
        app.processEvents()
        assert (
            detail._official_material_sample.currentData()
            == "builtin/minutes_coordination"
        )

        detail._apply_official_sample_material()
        app.processEvents()

        context = bridge.current_material_context()
        assert context.profile_id == "official:minutes"
        assert context.profile_name == "纪要"
        assert context.entity_data["document_type"] == "minutes"
        assert context.entity_data["title"] == "专题协调会议纪要"
        assert context.entity_data["meeting_date"]
        assert context.entity_data["participants"]
        assert "必填字段已齐" in detail._official_material_status.text()
        assert "administrative_meeting_fields_v1" in detail._official_contract.text()
        assert bridge.current_scene() is scene
        assert bridge.current_scene_id() == "official"
        assert scene.default_material_profile_id == "official:notice"

        detail._generate_official_sample()
        app.processEvents()

        output_dir = tmp_path / "output" / "official_master_samples"
        assert (output_dir / "minutes_official_sample.docx").is_file()
        assert opened_paths == [output_dir]
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_plan_preview_rejects_ambiguous_material_sample(
    monkeypatch,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.Toast.show_warning",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.load_official_document_material_package_sample",
        lambda _sample_id: (
            material_package_module.OfficialDocumentMaterialPackageResult(
                status="sample_ambiguous",
                profile_id="notice_archive_check",
                material_schema_ids=(),
                context=MaterialExecutionContext(),
                issues=("ambiguous",),
            )
        ),
    )
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    bridge.set_current_material_context(
        MaterialExecutionContext(
            profile_id="official:notice",
            entity_data={"document_type": "notice", "title": "Keep current task"},
        ),
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()
        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()

        panel._exam_paper._apply_official_sample_material()
        app.processEvents()

        context = bridge.current_material_context()
        assert context.profile_id == "official:notice"
        assert context.entity_data["title"] == "Keep current task"
        assert bridge.current_scene() is scene
        assert scene.default_material_profile_id == "official:notice"
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_plan_preview_imports_and_exports_material_package(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_success", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_warning", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_error", lambda *args, **kwargs: None
    )
    import_path = tmp_path / "letter_material.json"
    export_path = tmp_path / "exported_letter_material.json"
    material_package_module.export_official_document_material_package(
        MaterialExecutionContext(
            profile_id="official:letter",
            entity_data={
                "title": "关于材料补充的函",
                "body": "请补充相关材料。",
                "organization": "示例办公室",
                "document_no": "示函〔2026〕1号",
                "issue_date": "2026年7月10日",
                "recipient": "项目单位",
            },
        ),
        import_path,
        profile_id="letter",
    )

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()
        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()

        detail = panel._exam_paper
        assert not hasattr(detail, "_official_import_material_btn")
        assert not hasattr(detail, "_official_export_material_btn")

        detail._load_official_material_package(import_path)
        app.processEvents()

        context = bridge.current_material_context()
        assert detail._official_profile.currentData() == "letter"
        assert context.profile_id == "official:letter"
        assert context.entity_data["document_type"] == "letter"
        assert context.entity_data["title"] == "关于材料补充的函"
        assert "必填字段已齐" in detail._official_material_status.text()

        detail._export_official_material_package_to_path(export_path)
        app.processEvents()

        payload = json.loads(export_path.read_text(encoding="utf-8"))
        assert payload["kind"] == "alavette.material_package"
        assert payload["version"] == 5
        assert payload["profiles"][0]["profile_id"] == "official:letter"
        assert payload["profiles"][0]["fields"]["title"] == "关于材料补充的函"
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_material_fill_imports_and_exports_material_package(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_success", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_warning", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_error", lambda *args, **kwargs: None
    )
    import_path = tmp_path / "letter_material.json"
    export_path = tmp_path / "exported_letter_material.json"
    material_package_module.export_official_document_material_package(
        MaterialExecutionContext(
            profile_id="official:letter",
            entity_data={
                "title": "关于资料补交的函",
                "body": "请于本周内补交资料。",
                "organization": "示例办公室",
                "document_no": "示函〔2026〕2号",
                "issue_date": "2026年7月10日",
                "recipient": "项目单位",
            },
        ),
        import_path,
        profile_id="letter",
    )

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    template = create_builtin_template("official_gbt")
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(template, config_id="official_gbt", emit_signal=False)
    panel = ScenePanel(bridge)
    scene_signals = []
    template_signals = []
    bridge.scene_changed.connect(lambda value: scene_signals.append(value))
    bridge.template_changed.connect(lambda value: template_signals.append(value))
    try:
        app.processEvents()
        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()
        preview_detail = panel._exam_paper
        assert preview_detail._official_profile.currentData() == "notice"
        assert (
            preview_detail._official_material_sample.currentData()
            == "builtin/notice_archive_check"
        )

        panel._nav_rail.select_card("scn_content")
        app.processEvents()

        detail = panel._content
        assert not hasattr(detail, "_official_import_material_btn")
        assert not hasattr(detail, "_official_export_material_btn")
        assert detail._official_material_profile.currentData() == "notice"

        detail._load_official_material_package(import_path)
        app.processEvents()

        context = bridge.current_material_context()
        assert detail._official_material_profile.currentData() == "letter"
        assert detail._official_material_fields["title"].text() == "关于资料补交的函"
        assert detail._official_material_fields["recipient"].text() == "项目单位"
        assert (
            detail._official_material_fields["body"].get_text()
            == "请于本周内补交资料。"
        )
        assert context.profile_id == "official:letter"
        assert context.entity_data["document_type"] == "letter"
        assert context.entity_data["title"] == "关于资料补交的函"
        assert preview_detail._official_profile.currentData() == "letter"
        assert (
            preview_detail._official_material_sample.currentData()
            == "builtin/letter_material_request"
        )
        assert bridge.current_scene() is scene
        assert bridge.current_template() == template
        assert scene.default_material_profile_id == "official:notice"
        assert scene_signals == []
        assert template_signals == []

        detail._official_material_fields["body"].set_text("请于三个工作日内补交资料。")
        app.processEvents()
        detail._export_official_material_package_to_path(export_path)
        app.processEvents()

        payload = json.loads(export_path.read_text(encoding="utf-8"))
        assert payload["kind"] == "alavette.material_package"
        assert payload["version"] == 5
        profile = payload["profiles"][0]
        assert profile["profile_id"] == "official:letter"
        assert profile["fields"]["title"] == "关于资料补交的函"
        assert profile["fields"]["body"] == "请于三个工作日内补交资料。"
        assert bridge.current_scene() is scene
        assert bridge.current_template() == template
        assert scene.default_material_profile_id == "official:notice"
        assert scene_signals == []
        assert template_signals == []
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_material_fill_imports_csv_and_syncs_task_profile(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_success", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_warning", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_error", lambda *args, **kwargs: None
    )
    import_path = tmp_path / "minutes_material.csv"
    import_path.write_text(
        "文种,标题,正文,发文机关,发文字号,成文日期,会议时间,参会人员\n"
        "minutes,档案工作协调会纪要,会议研究了归档安排。,综合办公室,办纪〔2026〕2号,"
        "2026-07-10,2026-07-10 09:00,张三、李四\n",
        encoding="utf-8-sig",
    )
    multiple_path = tmp_path / "multiple_materials.csv"
    multiple_path.write_text(
        "document_type,title,body,organization,document_no,issue_date\n"
        "notice,Notice A,Body A,Office,A-1,2026-07-10\n"
        "notice,Notice B,Body B,Office,B-1,2026-07-11\n",
        encoding="utf-8",
    )

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    template = create_builtin_template("official_gbt")
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(template, config_id="official_gbt", emit_signal=False)
    panel = ScenePanel(bridge)
    scene_signals = []
    template_signals = []
    bridge.scene_changed.connect(scene_signals.append)
    bridge.template_changed.connect(template_signals.append)
    try:
        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()
        preview_detail = panel._exam_paper

        panel._nav_rail.select_card("scn_content")
        app.processEvents()
        detail = panel._content
        assert not hasattr(detail, "_official_import_material_btn")
        bridge.set_current_official_document_type_id("minutes")
        app.processEvents()
        detail._load_official_material_package(import_path)
        app.processEvents()

        context = bridge.current_material_context()
        assert detail._official_material_profile.currentData() == "minutes"
        assert detail._official_material_fields["title"].text() == "档案工作协调会纪要"
        assert (
            detail._official_material_fields["meeting_date"].text()
            == "2026-07-10 09:00"
        )
        assert detail._official_material_fields["participants"].text() == "张三、李四"
        assert detail._official_material_rows["meeting_date"].isHidden() is False
        assert context.profile_id == "official:minutes"
        assert context.entity_data["document_type"] == "minutes"
        assert preview_detail._official_profile.currentData() == "minutes"
        assert (
            preview_detail._official_material_sample.currentData()
            == "builtin/minutes_coordination"
        )
        assert bridge.current_scene() is scene
        assert bridge.current_template() == template
        assert scene.default_material_profile_id == "official:notice"
        assert scene_signals == []
        assert template_signals == []

        before = bridge.current_material_context()
        rejected = detail._load_official_material_package(multiple_path)
        after = bridge.current_material_context()
        assert rejected.status == "multiple_records"
        assert after.profile_id == before.profile_id
        assert after.entity_data == before.entity_data
        assert detail._official_material_profile.currentData() == "minutes"
        assert preview_detail._official_profile.currentData() == "minutes"
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_material_batch_import_is_transient_and_does_not_replace_single_task(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_success", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_warning", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_error", lambda *args, **kwargs: None
    )
    batch_path = tmp_path / "official_batch.csv"
    batch_path.write_text(
        "task_id,document_type,title,body,organization,document_no,issue_date\n"
        "notice_1,notice,Notice A,Body A,Office,A-1,2026-07-10\n"
        "letter_1,letter,Letter B,Body B,Office,,2026-07-11\n",
        encoding="utf-8",
    )

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    template = create_builtin_template("official_gbt")
    single_context = MaterialExecutionContext(
        profile_id="official:notice",
        profile_name="Current notice",
        entity_data={
            "document_type": "notice",
            "title": "Current single task",
        },
    )
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(template, config_id="official_gbt", emit_signal=False)
    bridge.set_current_material_context(single_context, emit_signal=False)
    panel = ScenePanel(bridge)
    scene_signals = []
    template_signals = []
    batch_signals = []
    bridge.scene_changed.connect(scene_signals.append)
    bridge.template_changed.connect(template_signals.append)
    bridge.material_batch_selection_changed.connect(batch_signals.append)
    try:
        panel._nav_rail.select_card("scn_content")
        app.processEvents()
        detail = panel._content

        assert not hasattr(detail, "_official_import_batch_btn")
        loaded = detail._load_official_material_batch(batch_path)
        app.processEvents()

        assert loaded.status == "ready_with_issues"
        assert loaded.invalid_count == 1
        assert len(batch_signals) == 1
        selection = bridge.current_material_batch_selection()
        assert selection.source_kind == "official_document_table"
        assert selection.source_path == str(batch_path)
        assert selection.profile_ids == ["notice_1", "letter_1"]
        assert selection.archive.profiles[0].fields["document_type"] == "notice"
        assert selection.archive.profiles[1].fields["document_type"] == "letter"
        assert "已载入 2 份公文资料" in detail._official_material_batch_status.text()
        assert detail._official_material_batch_status.isHidden() is False
        assert detail._official_batch_profile_row.isHidden() is False
        assert detail._official_batch_profile.currentData() == "notice_1"

        letter_index = detail._official_batch_profile.findData("letter_1")
        assert letter_index >= 0
        detail._official_batch_profile.setCurrentIndex(letter_index)
        app.processEvents()
        assert detail._official_material_profile.currentData() == "letter"
        assert detail._official_material_fields["title"].text() == "Letter B"
        assert detail._official_material_fields["document_no"].text() == ""

        detail._official_material_fields["document_no"].setText("B-2")
        QTest.qWait(320)
        app.processEvents()
        edited_selection = bridge.current_material_batch_selection()
        edited_letter = edited_selection.archive.get_profile("letter_1")
        assert edited_letter.fields["document_no"] == "B-2"
        assert not hasattr(edited_letter, "required_fields")
        assert edited_selection.item_metadata["letter_1"]["import_status"] == "ok"
        assert edited_selection.item_metadata["letter_1"]["edited_in_batch"] is True
        assert len(batch_signals) > 1

        after = bridge.current_material_context()
        assert after.profile_id == single_context.profile_id
        assert after.profile_name == single_context.profile_name
        assert after.entity_data == single_context.entity_data
        assert bridge.current_scene() is scene
        assert bridge.current_template() == template
        assert scene.default_material_profile_id == "official:notice"
        assert scene_signals == []
        assert template_signals == []
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_material_rejects_legacy_without_context_mutation(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_success", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_warning", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_error", lambda *args, **kwargs: None
    )
    source_path = tmp_path / "legacy_notice.json"
    source_payload = {
        "profile_id": "official:notice",
        "entity_data": {
            "title": "Legacy title",
            "body": "Legacy body",
            "organization": "Archive Office",
            "document_no": "N-1",
            "issue_date": "2026-07-10",
        },
    }
    original_text = json.dumps(source_payload, ensure_ascii=False, indent=2)
    source_path.write_text(original_text, encoding="utf-8")

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    template = create_builtin_template("official_gbt")
    context = MaterialExecutionContext(
        profile_id="official:letter",
        profile_name="Current letter",
        entity_data={"document_type": "letter", "title": "Current task"},
    )
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(template, config_id="official_gbt", emit_signal=False)
    bridge.set_current_material_context(context, emit_signal=False)
    panel = ScenePanel(bridge)
    scene_signals = []
    template_signals = []
    bridge.scene_changed.connect(scene_signals.append)
    bridge.template_changed.connect(template_signals.append)
    try:
        panel._nav_rail.select_card("scn_content")
        app.processEvents()
        detail = panel._content

        assert not hasattr(detail, "_official_convert_material_btn")
        assert not hasattr(detail, "_convert_official_material_package_to_path")
        result = detail._load_official_material_package(source_path)
        app.processEvents()

        assert result.status == "invalid_package"
        assert result.context.is_empty()
        assert source_path.read_text(encoding="utf-8") == original_text

        after = bridge.current_material_context()
        assert after.profile_id == context.profile_id
        assert after.profile_name == context.profile_name
        assert after.entity_data == context.entity_data
        assert bridge.current_scene() is scene
        assert bridge.current_template() == template
        assert scene_signals == []
        assert template_signals == []
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_material_fill_saves_user_library_material_package(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_success", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_warning", lambda *args, **kwargs: None
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_panel.Toast.show_error", lambda *args, **kwargs: None
    )
    library_root = tmp_path / "config_library" / "material_packages"
    monkeypatch.setattr(
        material_package_library,
        "MATERIAL_PACKAGE_LIBRARY_DIR",
        library_root,
    )
    user_dir = material_package_library.material_package_user_dir("official")

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    template = create_builtin_template("official_gbt")
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(template, config_id="official_gbt", emit_signal=False)
    panel = ScenePanel(bridge)
    scene_signals = []
    template_signals = []
    bridge.scene_changed.connect(lambda value: scene_signals.append(value))
    bridge.template_changed.connect(lambda value: template_signals.append(value))
    try:
        app.processEvents()
        panel._nav_rail.select_card("scn_content")
        app.processEvents()

        detail = panel._content
        assert not hasattr(detail, "_official_save_material_btn")
        detail._official_material_fields["title"].setText("Notice title")
        detail._official_material_fields["body"].set_text("Notice body")
        detail._official_material_fields["organization"].setText("Archive Office")
        detail._official_material_fields["document_no"].setText("A-2026-1")
        detail._official_material_fields["issue_date"].setText("2026-07-10")
        app.processEvents()

        detail._save_official_material_package_to_library_with_label(
            "saved_notice_package"
        )
        app.processEvents()

        package_path = user_dir / "saved_notice_package" / "package.json"
        payload = json.loads(package_path.read_text(encoding="utf-8"))
        assert payload["kind"] == "alavette.material_package"
        assert payload["version"] == 5
        assert payload["package_id"] == "saved_notice_package"
        assert payload["archive_name"] == "saved_notice_package"
        profile = payload["profiles"][0]
        assert profile["profile_id"] == "official:notice"
        assert profile["fields"]["title"] == "Notice title"
        assert profile["fields"]["body"] == "Notice body"
        assert bridge.current_scene() is scene
        assert bridge.current_template() == template
        assert scene.default_material_profile_id == "official:notice"
        assert scene_signals == []
        assert template_signals == []
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_official_material_fill_writes_material_context(monkeypatch):
    monkeypatch.setattr(
        config_library, "SCENE_LIBRARY_DIR", ROOT / "config_library" / "plans"
    )
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()
        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()
        preview_detail = panel._exam_paper
        assert "资料状态：尚未填写" in preview_detail._official_material_status.text()

        panel._nav_rail.select_card("scn_content")
        app.processEvents()

        detail = panel._content
        assert detail._official_material_section.isHidden() is False
        assert detail._official_material_profile.currentData() == "notice"
        assert detail._official_material_rows["meeting_date"].isHidden() is True
        assert detail._official_material_rows["recipient"].isHidden() is True
        assert detail._official_material_rows["copy_scope"].isHidden() is True
        assert detail._official_material_rows["title"].label_text == "标题 *"
        assert detail._official_material_contract_row.isHidden() is True

        detail._official_optional_fields_toggle.click()
        app.processEvents()
        assert detail._official_material_rows["recipient"].isHidden() is False
        assert detail._official_material_rows["attachment_note"].isHidden() is False
        assert detail._official_material_rows["copy_scope"].isHidden() is True

        detail._official_advanced_fields_toggle.click()
        app.processEvents()
        assert detail._official_material_rows["copy_scope"].isHidden() is False
        assert detail._official_material_contract_row.isHidden() is False

        material_emissions = []
        bridge.material_context_changed.connect(material_emissions.append)
        detail._official_material_fields["title"].setText("关于开展资料归档检查的通知")
        detail._official_material_fields["organization"].setText("示例市档案局")
        detail._official_material_fields["document_no"].setText("示档发〔2026〕1号")
        detail._official_material_fields["issue_date"].setText("2026年7月10日")
        detail._official_material_fields["recipient"].setText("各部门")
        detail._official_material_fields["body"].set_text("请各部门完成归档自查。")
        QTest.qWait(320)
        app.processEvents()

        context = bridge.current_material_context()
        assert len(material_emissions) == 1
        assert context.profile_id == "official:notice"
        assert context.profile_name == "公文资料"
        assert context.entity_data["document_type"] == "notice"
        assert context.entity_data["title"] == "关于开展资料归档检查的通知"
        assert context.entity_data["recipient"] == "各部门"
        assert context.entity_data["body"] == "请各部门完成归档自查。"
        assert "必填字段已齐" in detail._official_material_status.text()
        assert "必填字段已齐" in preview_detail._official_material_status.text()
        assert "official_document_v1" in detail._official_material_contract.text()
        assert "标题 -> official_title" in detail._official_material_contract.text()

        _set_combo_by_data(detail._official_material_profile, "minutes")
        app.processEvents()
        assert detail._official_material_rows["recipient"].isHidden() is True
        assert detail._official_material_rows["attachment_note"].isHidden() is True
        assert detail._official_material_rows["meeting_date"].isHidden() is False
        assert detail._official_material_rows["participants"].isHidden() is False
        assert detail._official_material_rows["meeting_date"].label_text == "会议时间"

        detail._official_material_fields["meeting_date"].setText("2026年7月10日上午")
        detail._official_material_fields["participants"].setText("张三、李四")
        QTest.qWait(320)
        app.processEvents()

        context = bridge.current_material_context()
        assert context.profile_id == "official:minutes"
        assert context.entity_data["document_type"] == "minutes"
        assert context.entity_data["meeting_date"] == "2026年7月10日上午"
        assert context.entity_data["participants"] == "张三、李四"
        assert context.entity_data["recipient"] == "各部门"
        assert (
            "administrative_meeting_fields_v1"
            in detail._official_material_contract.text()
        )
        assert preview_detail._official_profile.currentData() == "minutes"
        assert (
            preview_detail._official_material_sample.currentData()
            == "builtin/minutes_coordination"
        )

        bridge.set_current_material_context(MaterialExecutionContext())
        app.processEvents()
        assert bridge.current_official_document_type_id() == "minutes"
        assert detail._official_material_profile.currentData() == "minutes"
        assert preview_detail._official_profile.currentData() == "minutes"
        assert (
            preview_detail._official_material_sample.currentData()
            == "builtin/minutes_coordination"
        )
        assert bridge.current_scene() is scene
        assert scene.default_material_profile_id == "official:notice"
    finally:
        panel.close()
        app.processEvents()


def test_scene_overview_summary_projects_academic_rule_source_evidence():
    scene = create_thesis_scene()
    apply_planned_scene_family_defaults(scene)

    evidence = {
        item.key: item
        for item in build_academic_rule_source_evidence_summary_items(scene)
    }
    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}

    assert overview["planning_family"].value == "中文论文/课程论文"
    assert "中文论文/课程论文" in overview["coverage_pack"].detail
    assert evidence["academic_rule_source_evidence"].label == "学术规则证据"
    assert evidence["academic_rule_source_evidence"].value == "已打通"
    assert "学校论文规则资料" in evidence["academic_rule_source_evidence"].detail
    assert (
        "thesis_school_rule_context_v1"
        in evidence["academic_rule_source_evidence"].tooltip
    )
    assert "school_rule_source_selection" in (
        evidence["academic_rule_source_evidence"].tooltip
    )
    assert "section_classifier_confirmation" in (
        evidence["academic_rule_source_evidence"].tooltip
    )
    assert "chinese_academic_school_rule_section_confirmation" in (
        evidence["academic_rule_source_evidence"].tooltip
    )
    assert overview["academic_rule_source_evidence"].value == "已打通"
    assert (
        "Academic rule evidence" not in evidence["academic_rule_source_evidence"].label
    )
    assert overview["product_readiness"].value == "可直接使用"
    assert "product_readiness_gaps" not in overview


def test_scene_overview_summary_projects_journal_submission_evidence():
    scene = SceneWorkspace(scene_id="journal_en", category="journal_en")
    apply_planned_scene_family_defaults(scene)

    evidence = {
        item.key: item
        for item in build_journal_submission_evidence_summary_items(scene)
    }
    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}
    input_items = {item.key: item for item in build_input_profile_summary_items(scene)}

    assert overview["planning_family"].value == "英文期刊投稿"
    assert overview["coverage_pack"].value == "英文期刊投稿"
    assert evidence["journal_submission_evidence"].label == "期刊投稿证据"
    assert evidence["journal_submission_evidence"].value == "已打通"
    assert "期刊投稿资料" in evidence["journal_submission_evidence"].detail
    assert "期刊资料" in evidence["journal_submission_evidence"].detail
    assert (
        "journal_submission_materials_v1"
        in evidence["journal_submission_evidence"].tooltip
    )
    assert "journal_materials_v1" in evidence["journal_submission_evidence"].tooltip
    assert "journal_submission_reviewed_generic_rules" in (
        evidence["journal_submission_evidence"].tooltip
    )
    assert "reviewed_journal_profile_update" in (
        evidence["journal_submission_evidence"].tooltip
    )
    assert "submission_artifact_manifest" in (
        evidence["journal_submission_evidence"].tooltip
    )
    assert "english_journal_bibtex_csl_degraded" in (
        evidence["journal_submission_evidence"].tooltip
    )
    assert overview["journal_submission_evidence"].value == "已打通"
    assert (
        "Journal submission evidence"
        not in evidence["journal_submission_evidence"].label
    )
    assert overview["product_readiness"].value == "可直接使用"
    assert "product_readiness_gaps" not in overview
    assert input_items["materials"].detail == "期刊投稿资料"
    assert "journal_submission_materials_v1" in input_items["materials"].tooltip


def test_scene_overview_summary_projects_technical_long_doc_evidence():
    scene = create_technical_scene()
    apply_planned_scene_family_defaults(scene)

    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}

    assert overview["planning_family"].value == "技术长文档"
    assert "技术长文档" in overview["coverage_pack"].detail
    assert overview["technical_long_doc_evidence"].label == "技术长文档证据"
    assert overview["technical_long_doc_evidence"].value == "已打通"
    assert "技术文档资料" in overview["technical_long_doc_evidence"].detail
    assert "长文档元数据" in overview["technical_long_doc_evidence"].detail
    assert "technical_document_v1" in overview["technical_long_doc_evidence"].tooltip
    assert (
        "long_document_metadata_v1" in overview["technical_long_doc_evidence"].tooltip
    )
    assert "index_appendix_inventory" in overview["technical_long_doc_evidence"].tooltip
    assert "multi_file_merge_boundary_report" in (
        overview["technical_long_doc_evidence"].tooltip
    )
    assert "technical_long_docs_index_appendix_merge_boundary" in (
        overview["technical_long_doc_evidence"].tooltip
    )
    assert (
        "Technical long-doc evidence"
        not in overview["technical_long_doc_evidence"].label
    )
    assert overview["product_readiness"].value == "可直接使用"
    assert "product_readiness_gaps" not in overview


def test_scene_overview_summary_projects_fixed_layout_batch_evidence():
    scene = SceneWorkspace(
        scene_id="form_batch_documents",
        category="form_batch_documents",
    )
    apply_planned_scene_family_defaults(scene)

    evidence = {
        item.key: item
        for item in build_fixed_layout_batch_evidence_summary_items(scene)
    }
    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}
    input_items = {item.key: item for item in build_input_profile_summary_items(scene)}

    assert overview["planning_family"].value == "批量表单/套打"
    assert overview["coverage_pack"].value == "批量表单/套打"
    assert evidence["fixed_layout_batch_evidence"].label == "固定版式批量证据"
    assert evidence["fixed_layout_batch_evidence"].value == "已打通"
    assert "表单字段资料" in evidence["fixed_layout_batch_evidence"].detail
    assert "人员记录资料" in evidence["fixed_layout_batch_evidence"].detail
    assert "form_batch_fields_v1" in evidence["fixed_layout_batch_evidence"].tooltip
    assert "personnel_records_v1" in evidence["fixed_layout_batch_evidence"].tooltip
    assert "fixed_layout_profile_browser" in (
        evidence["fixed_layout_batch_evidence"].tooltip
    )
    assert "answer_sheet_reuse_path" in evidence["fixed_layout_batch_evidence"].tooltip
    assert "batch_forms_fixed_layout" in evidence["fixed_layout_batch_evidence"].tooltip
    assert "通道：12/12" in evidence["fixed_layout_batch_evidence"].detail
    assert "channels=" not in evidence["fixed_layout_batch_evidence"].detail
    assert overview["fixed_layout_batch_evidence"].value == "已打通"
    assert (
        "Fixed-layout batch evidence"
        not in evidence["fixed_layout_batch_evidence"].label
    )
    assert overview["product_readiness"].value == "可直接使用"
    assert "product_readiness_gaps" not in overview
    assert input_items["materials"].detail == "表单字段资料"
    assert "form_batch_fields_v1" in input_items["materials"].tooltip


def test_scene_overview_summary_projects_product_sales_package_defaults():
    scene = SceneWorkspace(
        scene_id="product_sales_documents",
        category="product_sales_documents",
    )
    apply_planned_scene_family_defaults(scene)

    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}
    input_items = {item.key: item for item in build_input_profile_summary_items(scene)}
    compliance = {item.key: item for item in build_compliance_summary_items(scene)}
    delivery = {item.key: item for item in build_delivery_summary_items(scene)}

    assert overview["planning_family"].value == "产品/售前材料"
    assert "客户/内部版本交付" in overview["planning_workflows"].detail
    assert overview["coverage_pack"].value == "项目申报/产品材料"
    assert "coverage_next_closure" not in overview
    assert input_items["materials"].value == "必需"
    assert "产品资料" in input_items["materials"].detail
    assert "案例资料" in input_items["materials"].detail
    assert "product_assets_v1" in input_items["materials"].tooltip
    assert "case_study_assets_v1" in input_items["materials"].tooltip
    assert "产品图片" in input_items["image_roles"].detail
    assert compliance["count_profile"].value == "product_asset_inventory"
    assert "Product and pre-sales asset" in compliance["count_profile"].detail
    assert compliance["planning_scan_targets"].value in {"已应用", "已覆盖"}
    assert delivery["default_delivery"].value == "客户版"
    assert "customer_copy" in delivery["default_delivery"].tooltip
    assert "售前资料包" in delivery["preset_count"].detail
    assert "pre_sales_package" in delivery["preset_count"].tooltip
    assert "资料清单" in delivery["default_artifacts"].value


def test_scene_overview_summary_projects_application_report_evidence():
    scene = SceneWorkspace(
        scene_id="product_sales_documents",
        category="product_sales_documents",
    )
    apply_planned_scene_family_defaults(scene)

    evidence = {
        item.key: item
        for item in build_application_report_evidence_summary_items(scene)
    }
    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}
    input_items = {item.key: item for item in build_input_profile_summary_items(scene)}

    assert overview["coverage_pack"].value == "项目申报/产品材料"
    assert evidence["application_report_evidence"].label == "申报/材料证据"
    assert evidence["application_report_evidence"].value == "已打通"
    assert "项目申报资料" in evidence["application_report_evidence"].detail
    assert "产品资料" in evidence["application_report_evidence"].detail
    assert "案例资料" in evidence["application_report_evidence"].detail
    assert "project_application_materials_v1" in (
        evidence["application_report_evidence"].tooltip
    )
    assert "product_assets_v1" in evidence["application_report_evidence"].tooltip
    assert "case_study_assets_v1" in evidence["application_report_evidence"].tooltip
    assert "project_application_rule_defaults" in (
        evidence["application_report_evidence"].tooltip
    )
    assert "submission_system_boundary_report" in (
        evidence["application_report_evidence"].tooltip
    )
    assert "asset_consistency_report" in (
        evidence["application_report_evidence"].tooltip
    )
    assert "quote_body_disambiguation" in (
        evidence["application_report_evidence"].tooltip
    )
    assert "application_reports_product_sales_assets" in (
        evidence["application_report_evidence"].tooltip
    )
    assert overview["application_report_evidence"].value == "已打通"
    assert (
        "Application/report evidence"
        not in evidence["application_report_evidence"].label
    )
    assert overview["product_readiness"].value == "可直接使用"
    assert "product_readiness_gaps" not in overview
    assert "产品资料" in input_items["materials"].detail
    assert "product_assets_v1" in input_items["materials"].tooltip


def test_scene_overview_summary_projects_regulated_disclosure_archive_defaults():
    scene = SceneWorkspace(
        scene_id="regulated_disclosure_documents",
        category="regulated_disclosure_documents",
    )
    apply_planned_scene_family_defaults(scene)

    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}
    input_items = {item.key: item for item in build_input_profile_summary_items(scene)}
    compliance = {item.key: item for item in build_compliance_summary_items(scene)}
    delivery = {item.key: item for item in build_delivery_summary_items(scene)}

    assert overview["planning_family"].value == "披露/审阅材料"
    assert "归档打包" in overview["planning_workflows"].detail
    assert overview["coverage_pack"].value == "专业披露/审阅"
    assert "coverage_next_closure" not in overview
    assert input_items["materials"].value == "必需"
    assert input_items["materials"].detail == "披露材料资料"
    assert "regulated_disclosure_materials_v1" in input_items["materials"].tooltip
    assert compliance["count_profile"].value == "disclosure_section_inventory"
    assert "Regulated disclosure section" in compliance["count_profile"].detail
    assert compliance["planning_scan_targets"].value in {"已应用", "已覆盖"}
    assert "hidden_text" in compliance["planning_scan_targets"].detail
    assert delivery["default_delivery"].value == "董事会审阅稿"
    assert "board_review_copy" in delivery["default_delivery"].tooltip
    assert "披露归档包" in delivery["preset_count"].detail
    assert "归档清单" in delivery["preset_count"].detail
    assert "disclosure_archive_package" in delivery["preset_count"].tooltip
    assert "archive_manifest" in delivery["preset_count"].tooltip
    assert "资料清单" in delivery["default_artifacts"].value


def test_scene_overview_summary_projects_n2128_family_application_parity_defaults():
    cases = (
        (
            "hr_batch_documents",
            "batch_forms",
            "personnel_records_v1",
            "batch_item_inventory",
            "per_person_docx",
            "failed_items_report",
            "content_controls",
        ),
        (
            "form_batch_documents",
            "batch_forms",
            "form_batch_fields_v1",
            "batch_item_inventory",
            "per_record_docx",
            "residue_check_report",
            "textboxes",
        ),
        (
            "finance_quote_documents",
            "professional_disclosure",
            "finance_quote_fields_v1",
            "finance_attachment_inventory",
            "customer_quote",
            "attachment_report",
            "embedded_workbooks",
        ),
        (
            "bilingual_translation_documents",
            "professional_disclosure",
            "bilingual_terms_v1",
            "bilingual_parallel_text",
            "bilingual_review_copy",
            "term_consistency_report",
            "tracked_changes",
        ),
    )

    for (
        family_id,
        coverage_pack,
        schema_id,
        count_profile_id,
        default_preset_id,
        extra_preset_id,
        scan_target,
    ) in cases:
        scene = SceneWorkspace(scene_id=family_id, category=family_id)
        apply_planned_scene_family_defaults(scene)

        overview = {
            item.key: item for item in build_scene_overview_summary_items(scene)
        }
        input_items = {
            item.key: item for item in build_input_profile_summary_items(scene)
        }
        compliance = {item.key: item for item in build_compliance_summary_items(scene)}
        delivery = {item.key: item for item in build_delivery_summary_items(scene)}

        family_labels = {
            "hr_batch_documents": "人事批量文档",
            "form_batch_documents": "批量表单/套打",
            "finance_quote_documents": "报价/财务材料",
            "bilingual_translation_documents": "双语审阅材料",
        }
        coverage_labels = {
            "batch_forms": "批量表单/套打",
            "professional_disclosure": "专业披露/审阅",
        }
        schema_labels = {
            "personnel_records_v1": "人员记录资料",
            "form_batch_fields_v1": "表单字段资料",
            "finance_quote_fields_v1": "报价字段资料",
            "bilingual_terms_v1": "双语术语资料",
        }
        delivery_labels = {
            "per_person_docx": "按人员生成",
            "per_record_docx": "按记录生成",
            "customer_quote": "客户报价稿",
            "bilingual_review_copy": "双语审阅稿",
            "failed_items_report": "失败项报告",
            "residue_check_report": "残留检查报告",
            "attachment_report": "附件报告",
            "term_consistency_report": "术语一致性报告",
        }
        assert overview["planning_family"].value == family_labels[family_id]
        assert overview["coverage_pack"].value == coverage_labels[coverage_pack]
        assert "coverage_next_closure" not in overview
        assert input_items["materials"].value == "必需"
        assert schema_labels[schema_id] in input_items["materials"].detail
        assert schema_id in input_items["materials"].tooltip
        assert compliance["count_profile"].value == count_profile_id
        assert compliance["planning_scan_targets"].value in {"已应用", "已覆盖"}
        assert scan_target in compliance["planning_scan_targets"].detail
        assert delivery["default_delivery"].value == delivery_labels[default_preset_id]
        assert default_preset_id in delivery["default_delivery"].tooltip
        assert delivery_labels[extra_preset_id] in delivery["preset_count"].detail
        assert extra_preset_id in delivery["preset_count"].tooltip


def test_scene_overview_summary_keeps_ip_patent_as_plugin_manual_boundary():
    scene = SceneWorkspace(
        scene_id="ip_patent_documents",
        category="ip_patent_documents",
    )
    result = apply_planned_scene_family_defaults(scene)

    overview = {item.key: item for item in build_scene_overview_summary_items(scene)}

    assert result.applied is False
    assert overview["planning_family"].value == "知识产权/专利材料"
    assert overview["coverage_pack"].value == "专业披露/审阅"
    assert overview["coverage_plugin_boundary"].value == "1 项需确认"
    assert "审计、法律、专利" in overview["coverage_plugin_boundary"].detail
    assert overview["product_readiness"].value == "需人工/插件把关"
    assert overview["product_readiness_family"].value == "需人工/插件把关"
    assert "知识产权/专利插件" in overview["product_readiness_gaps"].detail


def test_scene_compliance_summary_recommends_planning_family_scan_targets():
    scene = SceneWorkspace(scene_id="contract_delivery")
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.compliance_profile.object_preflight.scan_targets = ["fields"]

    recommended = recommended_object_preflight_targets_for_scene(scene)
    items = {item.key: item for item in build_compliance_summary_items(scene)}

    assert "tracked_changes" in recommended
    assert "ole_objects" in recommended
    assert items["planning_scan_targets"].value == "需同步"
    assert "tracked_changes" in items["planning_scan_targets"].detail

    scene.compliance_profile.object_preflight.scan_targets = list(recommended)
    items = {item.key: item for item in build_compliance_summary_items(scene)}

    assert items["planning_scan_targets"].value == "已应用"


def test_scene_panel_overview_shows_planning_family_governance(tmp_path):
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        description="字段一致性、审阅稿和签署稿规划族。",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        overview_layout = panel._overview.layout()
        assert overview_layout.indexOf(
            panel._overview._scene_card
        ) < overview_layout.indexOf(panel._overview._settings_card)
        assert overview_layout.indexOf(
            panel._overview._settings_card
        ) < overview_layout.indexOf(panel._overview._run_preview_card)
        assert panel._overview._combo.property("fullWidthMode") is True
        assert panel._overview._new_scene_btn.text() == "新建方案"
        assert panel._overview._duplicate_scene_btn.text() == "创建副本"
        assert panel._overview._rename_scene_btn.text() == "重命名方案"
        assert panel._overview._delete_scene_btn.text() == "删除方案"
        assert panel._overview._rename_scene_btn.isEnabled() is False
        assert panel._overview._delete_scene_btn.isEnabled() is False
        assert not hasattr(panel._overview, "_save_scene_btn")
        assert "save" not in panel._overview._scene_action_row._buttons
        assert "保存为我的方案" not in {
            button.text()
            for button in panel._overview._scene_action_row._buttons.values()
        }
        assert not hasattr(panel._overview, "_fork_scene_btn")
        assert panel._overview._open_scene_folder_btn.text() == "打开方案文件夹"
        assert panel._overview._execute_btn.parent() is not panel._overview._scene_card
        assert panel._overview._task_title.isHidden()
        assert panel._overview._task_meta.isHidden()
        assert panel._overview._task_title.text() == "合同交付"
        assert "模板：默认格式" in panel._overview._task_meta.text()
        assert "读取文件" in panel._overview._run_steps_label.text()
        assert "检查风险" in panel._overview._run_steps_label.text()
        assert "统一主要样式" in panel._overview._run_steps_label.text()
        assert "统一页面、正文、标题、表格和页眉页脚" not in (
            panel._overview._run_steps_label.text()
        )
        assert panel._overview._summary.isHidden()
        assert "全部内容" in panel._overview._summary.text()
        assert "/" not in panel._overview._summary.text()
        assert panel._overview._run_steps_label.isHidden()
        run_step_texts = [
            row.step_text()
            for row in panel._overview._run_step_rows
            if not row.isHidden()
        ]
        assert len(run_step_texts) == 5
        assert any("读取文件" in text for text in run_step_texts)
        assert any("检查风险" in text for text in run_step_texts)
        assert any("统一主要样式" in text for text in run_step_texts)
        assert all("\n" not in text for text in run_step_texts)
        assert all(
            not row._icon.pixmap().isNull()
            for row in panel._overview._run_step_rows
            if not row.isHidden()
        )
        assert "materials" in panel._overview._setting_rows
        assert all(
            not panel._overview._setting_rows[key]._icon.pixmap().isNull()
            for key in ("materials", "scope")
        )
        assert "已开启" in panel._overview._setting_rows["materials"].summary_text()
        assert panel._overview._setting_rows["scope"].summary_text() == "全部内容"
        assert "style_source" not in panel._overview._setting_rows
        assert "risk_confirmation" not in panel._overview._setting_rows
        assert "manual_confirmation" not in panel._overview._setting_rows
        first_screen_text = "\n".join(
            [
                panel._overview._task_title.text(),
                panel._overview._task_meta.text(),
                panel._overview._task_suitable.text(),
                panel._overview._task_boundary.text(),
                panel._overview._run_steps_label.text(),
                *(
                    row.summary_text()
                    for row in panel._overview._setting_rows.values()
                    if not row.isHidden()
                ),
            ]
        ).lower()
        assert not [
            term
            for term in FIRST_SCREEN_BANNED_TERMS
            if term.lower() in first_screen_text
        ]
        assert "scn_evidence" not in panel._nav_cards
        assert "scn_evidence" not in panel._navigation_card_snapshots(scene)
        assert not hasattr(panel._overview, "_advanced_evidence_card")
        assert not hasattr(panel._overview, "_profile_summary")
        assert not hasattr(panel._overview, "_sample_fixture_list")
        assert not hasattr(panel._overview, "_request_cell_list")
        assert "scn_overview" in panel._nav_cards
        assert panel._nav_cards["scn_overview"].isHidden() is False
        assert panel._details.current_detail is panel._overview
        engineering_overview = {
            item.key: item for item in build_scene_overview_summary_items(scene)
        }
        assert engineering_overview["planning_family"].value == "合同交付"
        assert engineering_overview["sample_fixture_coverage"].value == "已有样本文档"
        assert "Word 对象" in engineering_overview["sample_fixture_coverage"].detail
        assert build_scene_sample_fixture_detail_text(scene)
    finally:
        panel.close()


def test_scene_panel_overview_actions_emit_contextual_navigation_intents():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(scene_id="custom", template_id="default")
    bridge.set_current_scene(scene, config_id="custom", emit_signal=False)
    panel = ScenePanel(bridge)
    intents: list[object] = []
    bridge.navigate_to_intent.connect(intents.append)
    try:
        app.processEvents()

        panel._overview._execute_btn.click()

        assert intents[-1]["panel_id"] == "workbench"
        assert intents[-1]["card_id"] == "quick_execute"
        assert intents[-1]["return_panel_id"] == "scene"

        panel._show_detail_from_overview("tpl_overview")

        assert intents[-1]["panel_id"] == "template"
        assert intents[-1]["card_id"] == "tpl_overview"
        assert intents[-1]["return_panel_id"] == "scene"
        assert (
            intents[-1]["payload"]["entry_context_title"] == "来自方案：核对模板与样式"
        )
        assert "模板：" in intents[-1]["payload"]["entry_context_detail"]
        assert intents[-1]["payload"]["entry_context_action"] == (
            "核对页面、正文、标题、表格、页眉页脚、目录和题注"
        )
        assert intents[-1]["payload"]["template_preview_coverage"] == (
            "页面",
            "正文",
            "标题",
            "表格",
            "页眉页脚",
            "目录",
            "题注",
        )
        assert intents[-1]["payload"]["template_preview_groups"] == (
            "tpl_page",
            "tpl_style",
            "tpl_heading",
            "tpl_table",
            "tpl_header_footer",
            "tpl_toc",
            "tpl_caption",
        )
    finally:
        panel.close()


def test_scene_overview_duplicate_creates_mode_scoped_user_plan_and_selects_it(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "plans")
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam", emit_signal=False)
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        original_scene_id = bridge.current_scene_id()
        panel._overview._duplicate_scene_btn.click()
        app.processEvents()

        duplicated_scene_id = bridge.current_scene_id()
        duplicated_path = (
            config_library.scene_user_dir("exam") / f"{duplicated_scene_id}.json"
        )
        assert duplicated_scene_id
        assert duplicated_scene_id != original_scene_id
        assert duplicated_path.is_file()
        assert panel._overview._combo.currentData() == duplicated_scene_id
        assert bridge.current_scene_source_type() == "user"
        assert panel._overview._rename_scene_btn.isEnabled() is True
        assert panel._overview._delete_scene_btn.isEnabled() is True

        monkeypatch.setattr(
            scene_panel_module,
            "input_text",
            lambda *args, **kwargs: "校级期中方案",
        )
        panel._overview._rename_scene_btn.click()
        app.processEvents()

        assert bridge.current_scene().name == "校级期中方案"
        assert duplicated_path.is_file()

        monkeypatch.setattr(
            scene_panel_module,
            "confirm",
            lambda *args, **kwargs: True,
        )
        panel._overview._delete_scene_btn.click()
        app.processEvents()

        assert duplicated_path.exists() is False
        assert bridge.current_scene_id() != duplicated_scene_id
        assert panel._overview._rename_scene_btn.isEnabled() is False
        assert panel._overview._delete_scene_btn.isEnabled() is False
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_exposes_complete_plan_management_handlers():
    app = _app()
    bridge = PanelBridge()
    panel = ScenePanel(bridge)
    try:
        app.processEvents()
        assert hasattr(panel, "_create_scene_copy")
        assert hasattr(panel, "_on_new_scene_requested")
        assert hasattr(panel, "_on_duplicate_scene_requested")
        assert hasattr(panel, "_on_rename_scene_requested")
        assert hasattr(panel, "_on_delete_scene_requested")
    finally:
        panel.close()
        app.processEvents()


def test_scene_overview_strategy_edit_stays_local(tmp_path, monkeypatch):
    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(scene_id="custom", template_id="default")
    bridge.set_current_scene(
        scene, config_id="custom", source="library", emit_signal=False
    )
    bridge.set_current_template(
        create_builtin_template("default"), config_id="default", emit_signal=False
    )
    panel = ScenePanel(bridge)
    scene_changed: list[SceneWorkspace] = []
    apply_calls: list[SceneWorkspace] = []
    bridge.scene_changed.connect(scene_changed.append)
    try:
        app.processEvents()
        monkeypatch.setattr(
            panel,
            "_apply_scene",
            lambda updated_scene: apply_calls.append(updated_scene),
        )

        panel._overview._on_strategy_changed(True)
        app.processEvents()

        assert bridge.is_scene_dirty() is True
        assert scene_changed == []
        assert apply_calls == []
    finally:
        panel.close()
        app.processEvents()


def test_scene_overview_shows_exam_assembly_strategy_with_canonical_schema(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam", emit_signal=False)
    scene = create_exam_scene()
    bridge.set_current_scene(scene, config_id="exam", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("default"), config_id="default", emit_signal=False
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._overview._setting_rows["exam_blank_style"].isHidden() is False
        assert panel._overview._setting_rows["exam_runtime_fields"].isHidden() is False
        assert "exam_question_structure" not in panel._overview._setting_rows
        assert "exam_answer_handling" not in panel._overview._setting_rows
        assert panel._overview._setting_rows["materials"].isHidden() is True
        assert panel._overview._setting_rows["scope"].isHidden() is True
        assert "style_source" not in panel._overview._setting_rows
        assert scene.input_source_profile.material_schema_id == "exam_items_v1"
        assert panel._nav_cards["scn_content"].isHidden() is False
        assert "input_material" not in panel._nav_section_headers
        rules_snapshot = panel._navigation_card_snapshots(scene)["scn_rules"]
        assert "仅正文" not in rules_snapshot["subtitle"]
        assert rules_snapshot["badge_text"] == "生成结果"

        panel._nav_rail.select_card("scn_rules")
        app.processEvents()
        assert panel._rules._scope.isHidden() is True
        assert panel._rules._output_rules.isHidden() is False

        assert (
            panel._overview._setting_rows["exam_blank_style"]._label.text()
            == "当前方案"
        )
        assert (
            panel._overview._setting_rows["exam_runtime_fields"]._label.text()
            == "本次信息"
        )
        assert panel._nav_cards["scn_exam_paper"].isHidden() is False
        assert panel._nav_cards["scn_exam_paper"]._title.text() == "方案概览"
        exam_snapshot = panel._navigation_card_snapshots(scene)["scn_exam_paper"]
        assert exam_snapshot["subtitle"].endswith("· 当前方案")
        assert (
            panel._overview._setting_rows["exam_runtime_fields"]._jump_btn.isHidden()
            is False
        )
        assert (
            "工作台"
            in panel._overview._setting_rows["exam_runtime_fields"].summary_text()
        )
        panel._overview._setting_rows["exam_blank_style"]._jump_btn.click()
        assert panel._nav_rail.selected_card_id() == "scn_exam_paper"

        run_step_texts = [
            row.step_text()
            for row in panel._overview._run_step_rows
            if not row.isHidden()
        ]
        assert any("导入试卷内容" in text for text in run_step_texts)
        assert any("填写考试信息" in text for text in run_step_texts)
        assert any("生成 Word 试卷" in text for text in run_step_texts)
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_hides_scope_for_official_master_assembly(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = create_official_scene()
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("official_gbt"),
        config_id="official_gbt",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._overview._setting_rows["scope"].isHidden() is True
        assert "style_source" not in panel._overview._setting_rows

        rules_snapshot = panel._navigation_card_snapshots(scene)["scn_rules"]
        assert rules_snapshot["badge_text"] == "生成结果"

        panel._nav_rail.select_card("scn_rules")
        app.processEvents()
        assert panel._rules._scope.isHidden() is True
        assert panel._rules._output_rules.isHidden() is False
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_keeps_material_card_for_exam_with_material_contract(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam", emit_signal=False)
    scene = create_exam_scene()
    scene.input_source_profile.required_material_fields = ["exam_title"]
    bridge.set_current_scene(scene, config_id="exam", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("default"), config_id="default", emit_signal=False
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._nav_cards["scn_content"].isHidden() is False
        assert "input_material" not in panel._nav_section_headers
        assert (
            panel._navigation_card_snapshots(scene)["scn_content"]["badge_text"]
            == "已配置"
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_exam_paper_detail_updates_config_and_answer_delivery(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    monkeypatch.setattr(
        "src.shared.engine.exam_paper_style.USER_EXAM_MASTER_DIR",
        tmp_path / "user_masters",
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.USER_EXAM_MASTER_DIR",
        tmp_path / "user_masters",
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.Toast.show_success",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.Toast.show_warning",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.Toast.show_error",
        lambda *args, **kwargs: None,
    )
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam", emit_signal=False)
    scene = create_exam_scene()
    bridge.set_current_scene(scene, config_id="exam", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("default"), config_id="default", emit_signal=False
    )
    panel = ScenePanel(bridge)
    try:
        panel.resize(1330, 1120)
        panel.show()
        app.processEvents()
        panel._nav_rail.select_card("scn_exam_paper")
        app.processEvents()
        app.processEvents()

        detail = panel._exam_paper
        expected_gap = get_theme().template_detail_section_gap
        assert detail.layout().spacing() == expected_gap
        assert (
            detail._plan_preview_card.y()
            - (detail._attached_plan_card.y() + detail._attached_plan_card.height())
            == expected_gap
        )
        assert (
            detail._prompt_card.y()
            - (detail._plan_preview_card.y() + detail._plan_preview_card.height())
            == expected_gap
        )

        assert detail._card.isHidden()
        assert detail._detail_summary.isHidden()
        assert detail._master_card.isHidden() is True
        assert detail._attached_plan_card is panel._overview._scene_card
        assert detail._attached_plan_card.isHidden() is False
        assert panel._overview._combo.currentData() == "exam"
        assert not hasattr(detail, "_master_intro")
        assert not hasattr(detail, "_style_status")
        assert detail._live_summary.isHidden()
        assert not hasattr(detail, "_question_structure")
        assert not hasattr(detail, "_answer_policy")
        assert detail._blank_style.findData("compact_exam") == -1
        assert detail._blank_style.findData("class_quiz_exam") == -1
        assert detail._blank_style.findData("term_exam") == -1
        assert (
            detail._blank_style.itemText(detail._blank_style.findData("default_exam"))
            == "A4 标准卷面"
        )
        assert detail._blank_style.property("fullWidthMode") is True
        assert detail._style_action_row.isHidden()
        assert not hasattr(detail, "_copy_style_btn")
        assert not hasattr(detail, "_rename_master_btn")
        assert not hasattr(detail, "_delete_master_btn")
        assert not hasattr(detail, "_open_master_btn")
        assert not hasattr(detail, "_sample_docx_btn")
        assert not hasattr(detail, "_import_style_btn")
        assert detail._plan_preview_card.isHidden() is False
        assert detail._plan_preview_card._title_label.text() == "方案概览"
        assert detail._preview_mode_control.current_data() == "student"
        assert detail._active_preview_provider_id == "exam"
        assert detail._document_word_preview.objectName() == (
            "scn_plan_document_preview"
        )
        assert (
            detail._document_word_preview.toolbar_widget().parent()
            is detail._plan_preview_card._header_widget
        )
        assert detail._document_word_preview.page_count() == 0
        assert "当前环境未启用真实 Word 渲染" in (
            detail._document_word_preview._status.text()
        )
        assert not hasattr(detail, "_preview_page")
        assert not hasattr(detail, "_preview_canvas")
        assert detail._prompt_card.isHidden() is False
        assert detail._prompt_card._title_label.text() == "AI 提示词"
        assert (
            detail._prompt_card._content_layout.indexOf(detail._prompt_mode_control)
            == -1
        )
        assert (
            detail._prompt_card._header_layout.indexOf(detail._prompt_mode_control) >= 0
        )
        assert detail._prompt_mode_control.current_data() == "markdown"
        assert "只输出 Markdown 正文" in detail._prompt_area.get_text()
        assert "不要编写页眉、页脚、页码、密封线" in detail._prompt_area.get_text()
        assert "{{af_questions}}" not in detail._prompt_area.get_text()

        detail._prompt_mode_control.set_current_index(1)
        app.processEvents()
        assert detail._prompt_mode_control.current_data() == "master"
        assert "{{af_title}}" in detail._prompt_area.get_text()
        assert "{{af_questions}}" in detail._prompt_area.get_text()
        assert "页眉页脚只作为卷面版式存在" in detail._prompt_area.get_text()
        detail._copy_prompt_btn.click()
        assert app.clipboard().text() == detail._prompt_area.get_text()

        detail._preview_mode_control.set_current_index(1)
        app.processEvents()
        assert detail._preview_mode_control.current_data() == "answer"
        assert bridge.current_scene().exam_paper.answer_policy == "student_plus_answer"

        current_scene = bridge.current_scene()
        assert current_scene.master_id == "default_exam"
        assert current_scene.exam_paper.question_structure_mode == "markdown_headings"
        assert current_scene.exam_paper.answer_policy == "student_plus_answer"
        assert detail._document_word_preview.page_count() == 0
        assert not hasattr(detail, "_master_intro")
        assert not hasattr(detail, "_style_status")
        preset_ids = {
            str(getattr(preset, "preset_id", "") or "")
            for preset in current_scene.delivery_presets
        }
        assert "student" in preset_ids

    finally:
        panel.close()
        app.processEvents()


def test_scene_exam_paper_detail_discovers_manual_user_master_docx(
    tmp_path, monkeypatch
):
    user_master_dir = tmp_path / "user_masters"
    user_master_dir.mkdir(parents=True)
    manual_master = user_master_dir / "AI生成母版.docx"
    document = Document()
    document.add_paragraph("{{af_title}}")
    document.add_paragraph("科目：{{af_subject}}")
    document.add_paragraph("{{af_questions}}")
    document.save(manual_master)

    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    monkeypatch.setattr(
        "src.shared.engine.exam_paper_style.USER_EXAM_MASTER_DIR",
        user_master_dir,
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.USER_EXAM_MASTER_DIR",
        user_master_dir,
    )

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam", emit_signal=False)
    scene = create_exam_scene()
    bridge.set_current_scene(scene, config_id="exam", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("default"), config_id="default", emit_signal=False
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()
        detail = panel._exam_paper

        assert detail._blank_style.findText("AI生成母版") >= 0
        assert [
            style.label
            for style in bridge.current_scene().exam_paper.custom_blank_styles
        ] == ["AI生成母版"]
        discovered = bridge.current_scene().exam_paper.custom_blank_styles[0]
        assert discovered.master_docx_path.endswith("AI生成母版.docx")
    finally:
        panel.close()
        app.processEvents()


def test_scene_exam_paper_open_folder_uses_mode_scoped_user_dir(
    tmp_path, monkeypatch
):
    active_user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"

    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    monkeypatch.setattr(
        "src.shared.engine.exam_paper_style.USER_EXAM_MASTER_DIR",
        active_user_dir,
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.USER_EXAM_MASTER_DIR",
        active_user_dir,
    )
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.Toast.show_success",
        lambda *args, **kwargs: None,
    )

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam", emit_signal=False)
    scene = create_exam_scene()
    bridge.set_current_scene(scene, config_id="exam", emit_signal=False)
    bridge.set_current_template(
        create_builtin_template("default"), config_id="default", emit_signal=False
    )
    panel = ScenePanel(bridge)
    try:
        app.processEvents()
        detail = panel._exam_paper
        opened_paths = []
        detail._open_sample_path_handler = lambda path: (
            opened_paths.append(path) or True
        )

        detail._open_master_folder_btn.click()

        assert opened_paths == [active_user_dir]
        assert active_user_dir.is_dir()
    finally:
        panel.close()
        app.processEvents()


def test_scene_overview_selector_places_source_badges_on_real_scene_options(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    config_library.ensure_config_library()
    config_library.save_scene_to_library(
        SceneWorkspace(
            scene_id="exam_default_copy",
            name="试卷-期中",
            template_id="default",
            compatible_template_ids=["default"],
        ),
        scene_id="exam_default_copy",
    )
    app = _app()
    bridge = PanelBridge()
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        combo = panel._overview._combo
        texts = [combo.itemText(index) for index in range(combo.count())]
        assert "我的方案" not in texts
        assert "内置方案" not in texts
        assert texts.index("试卷-期中方案") < texts.index("自定义方案")

        user_index = texts.index("试卷-期中方案")
        builtin_index = texts.index("自定义方案")
        assert combo.itemData(user_index, SOURCE_BADGE_TEXT_ROLE) == "自定"
        assert combo.itemData(builtin_index, SOURCE_BADGE_TEXT_ROLE) == "内置"
        assert combo.model().item(user_index).isEnabled() is True
        assert combo.model().item(builtin_index).isEnabled() is True
    finally:
        panel.close()
        app.processEvents()


def test_scene_overview_distinct_user_plan_gets_user_management_permissions(
    tmp_path,
    monkeypatch,
):
    monkeypatch.setattr(config_library, "SCENE_LIBRARY_DIR", tmp_path / "scenes")
    config_library.ensure_config_library()
    user_scene_id = "exam_school_midterm"
    user_scene = config_library.load_scene_from_library("exam", mode_id="exam")
    user_scene.scene_id = user_scene_id
    user_scene.name = "校级期中试卷"
    config_library.save_scene_to_library(
        user_scene,
        scene_id=user_scene_id,
        mode_id="exam",
    )

    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam")
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        descriptor = panel._descriptor_for_scene_id(user_scene_id)
        assert descriptor is not None
        assert descriptor.source_type == "user"
        assert descriptor.name == "校级期中试卷"
        assert panel._current_scene_id() == "exam"
        user_index = panel._overview._combo.findData(user_scene_id)
        assert user_index >= 0
        panel._overview._combo.setCurrentIndex(user_index)
        app.processEvents()

        assert panel._current_scene_id() == user_scene_id
        combo_texts = [
            panel._overview._combo.itemText(index)
            for index in range(panel._overview._combo.count())
        ]
        assert "校级期中试卷方案" in combo_texts
        assert "默认试卷方案" in combo_texts
        assert panel._current_scene_is_builtin() is False
        assert bridge.current_scene_source() == "library"
        assert bridge.current_scene_source_type() == "user"
        assert panel._overview._rename_scene_btn.isEnabled() is True
        assert panel._overview._delete_scene_btn.isEnabled() is True
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_scope_summary_renders_without_overlap_at_narrow_width():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)
    panel = ScenePanel(bridge)
    try:
        panel.resize(620, 900)
        panel.show()
        app.processEvents()
        panel._show_detail("scn_overview")
        app.processEvents()
        app.processEvents()

        row = panel._overview._setting_rows["scope"]
        summary = row._summary
        button = row._jump_btn
        summary_right = summary.mapTo(row, QPoint(summary.width(), 0)).x()
        button_left = button.mapTo(row, QPoint(0, 0)).x()
        button_right = button.mapTo(row, QPoint(button.width(), 0)).x()
        required_summary_height = summary.heightForWidth(summary.width())
        if required_summary_height < 0:
            required_summary_height = summary.sizeHint().height()

        assert row.summary_text() == "全部内容"
        assert summary.wordWrap() is True
        assert summary.width() > 0
        assert summary.height() >= required_summary_height
        assert summary_right <= button_left - 8
        assert button_right <= row.width()
        assert row.height() >= row.sizeHint().height()
        assert panel._detail_scroll.horizontalScrollBar().maximum() == 0

        image = row.grab().toImage()
        dark_samples = 0
        for y in range(0, image.height(), max(1, image.height() // 12)):
            for x in range(0, image.width(), max(1, image.width() // 24)):
                color = image.pixelColor(x, y)
                if color.alpha() and color.red() + color.green() + color.blue() < 600:
                    dark_samples += 1
        assert dark_samples >= 4
    finally:
        panel.close()


def test_scene_panel_output_advanced_naming_and_reports_are_collapsed_by_default():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(scene_id="contract_delivery", template_id="default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        output = panel._output
        assert output._advanced_output_title.text() == "高级命名、规则与报告"
        assert output._advanced_output_container.isHidden() is True
        assert output._advanced_output_toggle_btn.text() == "展开"
        assert output._final_docx.isHidden() is False
        assert output._review_pdf_row.isHidden() is True
        assert output._material_package.isHidden() is False

        output._advanced_output_toggle_btn.click()
        app.processEvents()

        assert output._advanced_output_container.isHidden() is False
        assert output._advanced_output_toggle_btn.text() == "收起"
    finally:
        panel.close()


def test_scene_panel_official_review_pdf_toggle_survives_delivery_switch():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    scene = config_library.load_scene_from_library("official", mode_id="official")
    bridge.set_current_scene(scene, config_id="official", emit_signal=False)
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        output = panel._output
        compact_review = panel._rules._output_rules._general_checks["review_pdf"]
        assert output._review_pdf_row.isHidden() is False
        assert output._review_pdf.isChecked() is True
        assert compact_review.isHidden() is False
        assert compact_review.isChecked() is True
        assert scene.default_delivery_preset().artifacts.review_pdf is True

        internal_review_index = output._default_delivery.findData("internal_review")
        assert internal_review_index >= 0
        output._default_delivery.setCurrentIndex(internal_review_index)
        app.processEvents()

        assert scene.default_delivery_preset_id == "internal_review"
        assert scene.default_delivery_preset().artifacts.review_pdf is True
        assert output._review_pdf.isChecked() is True

        compact_review.click()
        app.processEvents()

        assert scene.default_delivery_preset().artifacts.review_pdf is False
        assert output._review_pdf.isChecked() is False
    finally:
        panel.close()


def test_scene_panel_output_rules_card_edits_delivery_artifacts_directly():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(scene_id="contract_delivery", template_id="default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        output_rules = panel._rules._output_rules
        assert panel._output.isHidden() is True
        assert output_rules._general_checks["final_docx"].isHidden() is False
        assert output_rules._general_checks["review_pdf"].isHidden() is True
        assert output_rules._exam_student_check.isHidden() is True
        assert output_rules._general_checks["report"].isChecked() is True

        output_rules._general_checks["report"].click()
        app.processEvents()

        assert scene.default_delivery_preset().artifacts.report_json is False
        assert scene.default_delivery_preset().artifacts.report_markdown is False
        assert scene.delivery_presets[0].artifacts.report_json is False
        assert scene.delivery_presets[0].artifacts.report_markdown is False

        output_rules._general_checks["material_package"].click()
        app.processEvents()

        assert scene.default_delivery_preset().artifacts.material_package is True
        assert scene.delivery_presets[0].artifacts.material_package is True
    finally:
        panel.close()


def test_scene_panel_output_rules_card_controls_exam_answer_delivery():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam", emit_signal=False)
    scene = SceneWorkspace(
        scene_id="exam_default",
        mode_id="exam",
        category="exam",
        template_id="default",
    )
    bridge.set_current_scene(scene, config_id="exam_default", emit_signal=False)
    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        output_rules = panel._rules._output_rules
        assert output_rules._exam_student_check.isHidden() is False
        assert output_rules._general_checks["final_docx"].isHidden() is True
        assert output_rules._exam_student_check.isChecked() is True
        assert output_rules._exam_student_check.isEnabled() is True
        assert output_rules._exam_answer_check.isChecked() is True
        assert scene.exam_paper.answer_policy == "student_plus_answer"
        assert {preset.preset_id for preset in scene.delivery_presets} == {"final"}
        assert scene.default_delivery_preset_id == "final"

        output_rules._exam_answer_check.click()
        app.processEvents()

        assert scene.exam_paper.answer_policy == "student_only"
        assert {preset.preset_id for preset in scene.delivery_presets} == {"student"}

        output_rules._exam_answer_check.click()
        app.processEvents()
        output_rules._exam_student_check.click()
        app.processEvents()

        assert scene.exam_paper.answer_policy == "answer_only"
        assert {preset.preset_id for preset in scene.delivery_presets} == {"answer_key"}
        assert scene.default_delivery_preset_id == "answer_key"
    finally:
        panel.close()


def test_scene_panel_navigation_intent_shows_return_to_execution_action():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(scene_id="contract_delivery", template_id="default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    panel = ScenePanel(bridge)
    intents: list[object] = []
    bridge.navigate_to_intent.connect(intents.append)
    try:
        app.processEvents()

        panel.handle_navigation_intent(
            {
                "panel_id": "scene",
                "card_id": "scn_output",
                "return_panel_id": "workbench",
                "return_card_id": "quick_execute",
            }
        )

        assert panel._return_bar.isHidden() is False
        assert panel._nav_rail.selected_card_id() == "scn_rules"

        panel._return_btn.click()

        assert intents[-1] == {
            "panel_id": "workbench",
            "card_id": "quick_execute",
        }
        assert panel._return_bar.isHidden() is True

        panel.handle_navigation_intent(
            {
                "panel_id": "scene",
                "card_id": "scn_output",
                "field_id": "review",
                "active_issue_id": "output_target.review.warning",
                "return_panel_id": "workbench",
                "return_card_id": "quick_execute",
                "payload": {
                    "issue_title": "输出目标预检",
                    "issue_item_id": "output_target.review.warning",
                },
            }
        )

        assert panel._return_label.text() == "从执行问题进入：输出目标预检"
        assert panel._output.isHidden() is False

        panel._return_btn.click()

        assert intents[-1]["panel_id"] == "workbench"
        assert intents[-1]["card_id"] == "quick_execute"
        assert intents[-1]["active_issue_id"] == "output_target.review.warning"
        assert intents[-1]["payload"]["active_issue_id"] == (
            "output_target.review.warning"
        )

        panel.handle_navigation_intent(
            {
                "panel_id": "scene",
                "card_id": "scn_output",
                "field_id": "filename_template",
                "return_panel_id": "workbench",
                "return_card_id": "quick_execute",
            }
        )
        app.processEvents()

        assert panel._nav_rail.selected_card_id() == "scn_rules"
        assert panel._output._advanced_output_container.isHidden() is False
        assert (
            panel._output._delivery_filename.property("navigation_field_highlight")
            is True
        )
        assert "文件名规则" in panel._output._delivery_filename.toolTip()
        assert "filename_template" not in panel._output._delivery_filename.toolTip()
        assert panel._output._delivery_filename.property(
            "navigation_field_raw_label"
        ) == ("filename_template")
        assert panel._output._delivery_filename.objectName() == (
            "scn_output_filename_template"
        )

        panel._output._advanced_output_expanded = False
        panel._output._sync_advanced_output_visibility()
        assert panel._output._advanced_output_container.isHidden() is True

        panel.handle_navigation_intent(
            {
                "panel_id": "scene",
                "card_id": "scn_output",
                "field_id": "preset_id",
                "return_panel_id": "workbench",
                "return_card_id": "quick_execute",
            }
        )
        app.processEvents()

        assert panel._nav_rail.selected_card_id() == "scn_rules"
        assert panel._output._advanced_output_container.isHidden() is False
        assert (
            panel._output._delivery_preset_id.property("navigation_field_highlight")
            is True
        )
        assert "交付编号" in panel._output._delivery_preset_id.toolTip()
        assert panel._output._delivery_preset_id.property(
            "navigation_field_raw_label"
        ) == ("preset_id")
    finally:
        panel.close()
        app.processEvents()


def test_scene_request_cell_projection_filters_boundary_levels_without_product_ui():
    scene = SceneWorkspace(
        scene_id="professional_disclosure",
        name="专业披露",
        category="professional_disclosure",
        template_id="default",
        compatible_template_ids=["default"],
    )
    cells = scene_request_cell_fixture_specs_for_scene(scene)
    assert len(cells) == 9
    manual = tuple(
        cell
        for cell in cells
        if scene_request_cell_matches_filter(cell, "manual_boundary_fixture")
    )
    ambiguous = tuple(
        cell
        for cell in cells
        if scene_request_cell_matches_filter(cell, "ambiguous_fixture_set")
    )
    direct = tuple(
        cell
        for cell in cells
        if scene_request_cell_matches_filter(cell, "direct_family_fixture")
    )
    assert len(manual) == 6
    assert len(ambiguous) == 3
    assert not direct
    manual_projection = scene_request_cell_list_item_projection(manual[0])
    assert "人工确认" in manual_projection.text
    assert "professional_disclosure_review_gate" in manual_projection.tooltip
    ambiguous_projection = scene_request_cell_list_item_projection(ambiguous[0])
    assert "容易误解" in ambiguous_projection.text
    assert "需要先澄清" in ambiguous_projection.tooltip
    assert "候选资料包" in ambiguous_projection.tooltip
    assert "候选路由" not in ambiguous_projection.tooltip


def test_scene_panel_can_apply_planning_family_preflight_targets():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.compliance_profile.object_preflight.scan_targets = ["fields"]
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        recommended = recommended_object_preflight_targets_for_scene(scene)
        assert panel._cleanup._apply_family_preflight_btn.isEnabled() is True
        assert (
            panel._cleanup._compliance_summary.value_for("planning_scan_targets")
            == "需同步"
        )

        panel._cleanup._apply_family_preflight_btn.click()
        app.processEvents()

        assert scene.compliance_profile.object_preflight.scan_targets == list(
            recommended
        )
        assert (
            "tracked_changes" in scene.compliance_profile.object_preflight.scan_targets
        )
        assert panel._cleanup._scan_target_checks["tracked_changes"].isChecked() is True
        assert panel._cleanup._scan_target_checks["fields"].isChecked() is True
        assert (
            panel._cleanup._compliance_summary.value_for("planning_scan_targets")
            == "已应用"
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_scan_target_checkboxes_write_back_to_scene():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    scene.compliance_profile.object_preflight.scan_targets = ["comments", "fields"]
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._cleanup._scan_target_checks["fields"].isChecked() is True
        assert panel._cleanup._scan_target_checks["macros"].isChecked() is False
        assert (
            panel._cleanup._compliance_summary.value_for("scan_targets") == "2 个目标"
        )

        panel._cleanup._scan_target_checks["fields"].click()
        panel._cleanup._scan_target_checks["tracked_changes"].click()
        app.processEvents()

        assert scene.compliance_profile.object_preflight.scan_targets == [
            "tracked_changes",
            "comments",
        ]
        assert (
            panel._cleanup._compliance_summary.value_for("scan_targets") == "2 个目标"
        )
        target_detail = panel._cleanup._compliance_summary.detail_for("scan_targets")
        assert "tracked_changes" in target_detail
        assert "fields" not in target_detail
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_profile_controls_write_back_to_scene():
    app = _app()
    bridge = PanelBridge()
    scene = create_thesis_scene()
    template = create_builtin_template("thesis_gbt")
    bridge.set_current_scene(scene, config_id="thesis", emit_signal=False)
    bridge.set_current_template(template, config_id="thesis_gbt", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        disabled_index = panel._content._markdown_policy.findData("disabled")
        panel._content._markdown_policy.setCurrentIndex(disabled_index)
        app.processEvents()
        assert scene.input_source_profile.markdown_policy == "disabled"
        assert "MD 禁用" in panel._content._input_summary.value_for("input_policy")

        panel._cleanup._skip_high_risk.click()
        app.processEvents()
        assert scene.compliance_profile.object_preflight.skip_high_risk_modules is False
        assert panel._cleanup._compliance_summary.value_for("skip_policy") == "关闭"

        review_index = panel._output._default_delivery.findData("review")
        panel._output._default_delivery.setCurrentIndex(review_index)
        app.processEvents()
        assert scene.default_delivery_preset_id == "review"
        assert scene.default_delivery_preset().artifacts.compare_docx is True
        assert panel._output._delivery_summary.value_for("default_delivery") == "审阅稿"

        panel._output._delivery_label.setText("Review package")
        panel._output._delivery_output_dir.setText("review/{preset_id}")
        panel._output._delivery_filename.setText("{stem}_reviewed")
        app.processEvents()

        review_preset = next(
            preset for preset in scene.delivery_presets if preset.preset_id == "review"
        )
        assert review_preset.label == "Review package"
        assert review_preset.output_dir_template == "review/{preset_id}"
        assert review_preset.filename_template == "{stem}_reviewed"
        assert (
            panel._output._default_delivery.itemText(
                panel._output._default_delivery.currentIndex()
            )
            == "Review package"
        )
        assert (
            panel._output._delivery_summary.value_for("default_delivery")
            == "Review package"
        )
        assert (
            panel._output._delivery_summary.detail_for("default_delivery")
            == "执行时默认生成"
        )
        assert (
            panel._output._delivery_validation_summary.value_for(
                "delivery_template_status"
            )
            == "通过"
        )
        variable_detail = panel._output._delivery_validation_summary.detail_for(
            "delivery_variables"
        )
        assert "交付编号" in variable_detail
        assert "{preset_id}" not in variable_detail
        assert "{preset_id}" in panel._output._delivery_validation_summary.tooltip_for(
            "delivery_variables"
        )
        assert "review/review" in panel._output._delivery_validation_summary.detail_for(
            "delivery_template_status"
        )

        panel._output._material_manifest.click()
        panel._output._material_package.click()
        app.processEvents()

        assert scene.default_delivery_preset().artifacts.material_manifest is True
        assert scene.default_delivery_preset().artifacts.material_package is True
        assert review_preset.artifacts.material_manifest is True
        assert review_preset.artifacts.material_package is True
        assert "资料清单" in panel._output._delivery_summary.value_for(
            "default_artifacts"
        )
        assert "资料包" in panel._output._delivery_summary.value_for(
            "default_artifacts"
        )

        panel._output._visibility_rules.set_text("answer=remove\nanalysis=remove")
        app.processEvents()

        assert [rule.selector for rule in review_preset.content_visibility_rules] == [
            "answer",
            "analysis",
        ]
        assert [rule.label for rule in review_preset.content_visibility_rules] == [
            "答案",
            "解析",
        ]
        assert [rule.action for rule in review_preset.content_visibility_rules] == [
            "remove",
            "remove",
        ]
        rule_status_detail = panel._output._delivery_validation_summary.detail_for(
            "delivery_rule_status"
        )
        assert "答案：删除块" in rule_status_detail
        assert "解析：删除块" in rule_status_detail
        rule_status_tooltip = panel._output._delivery_validation_summary.tooltip_for(
            "delivery_rule_status"
        )
        assert "answer=remove" in rule_status_tooltip
        assert "analysis=remove" in rule_status_tooltip
        assert "显隐规则 2 条" in panel._output._delivery_summary.detail_for(
            "default_artifacts"
        )

        original_count = len(scene.delivery_presets)
        panel._output._copy_preset_btn.click()
        app.processEvents()

        copied_preset = scene.delivery_presets[-1]
        assert len(scene.delivery_presets) == original_count + 1
        assert copied_preset.preset_id == "review_copy"
        assert copied_preset.label == "Review package Copy"
        assert copied_preset.output_dir_template == "review/{preset_id}"
        assert copied_preset.filename_template == "{stem}_reviewed"
        assert copied_preset.artifacts.material_package is True
        assert [rule.selector for rule in copied_preset.content_visibility_rules] == [
            "answer",
            "analysis",
        ]
        assert [rule.label for rule in copied_preset.content_visibility_rules] == [
            "答案",
            "解析",
        ]
        assert scene.default_delivery_preset_id == "review_copy"
        assert (
            panel._output._delivery_summary.value_for("default_delivery")
            == "Review package Copy"
        )
        assert (
            panel._output._delivery_summary.detail_for("default_delivery")
            == "执行时默认生成"
        )

        panel._output._delivery_preset_id.setText("review public")
        panel._output._delivery_preset_id.editingFinished.emit()
        panel._output._delivery_target_template.setText("thesis_review_template")
        app.processEvents()

        assert copied_preset.preset_id == "review_public"
        assert copied_preset.target_template_id == "thesis_review_template"
        assert scene.default_delivery_preset_id == "review_public"
        assert panel._output._default_delivery.currentData() == "review_public"
        assert panel._output._delivery_preset_id.text() == "review_public"
        assert (
            panel._output._delivery_summary.value_for("default_delivery")
            == "Review package Copy"
        )
        assert (
            panel._output._delivery_summary.detail_for("default_delivery")
            == "执行时默认生成"
        )

        panel._output._move_up_preset_btn.click()
        app.processEvents()
        assert [preset.preset_id for preset in scene.delivery_presets][
            -2
        ] == "review_public"
        assert scene.default_delivery_preset_id == "review_public"
        assert panel._output._move_down_preset_btn.isEnabled()

        panel._output._move_down_preset_btn.click()
        app.processEvents()
        assert [preset.preset_id for preset in scene.delivery_presets][
            -1
        ] == "review_public"
        assert scene.default_delivery_preset_id == "review_public"

        panel._output._add_preset_btn.click()
        app.processEvents()

        added_preset = scene.delivery_presets[-1]
        assert added_preset.preset_id == "delivery"
        assert added_preset.output_dir_template == "{preset_id}"
        assert added_preset.filename_template == "{stem}_{preset_id}"
        assert scene.default_delivery_preset_id == "delivery"

        panel._output._remove_preset_btn.click()
        app.processEvents()

        assert all(preset.preset_id != "delivery" for preset in scene.delivery_presets)
        assert scene.default_delivery_preset_id == "review_public"
        assert (
            panel._output._delivery_summary.value_for("default_delivery")
            == "Review package Copy"
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_content_detail_edits_material_schema_contract():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        contract_index = panel._content._schema_registry_combo.findData(
            "contract_parties_v1"
        )
        assert contract_index >= 0
        panel._content._schema_registry_combo.setCurrentIndex(contract_index)
        panel._content._set_primary_schema_btn.click()
        app.processEvents()

        assert scene.input_source_profile.material_schema_id == "contract_parties_v1"
        assert scene.input_source_profile.material_schema_ids == []
        assert "合同方字段资料" in panel._content._input_summary.detail_for("materials")
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "通过"
        )
        assert "合同方字段资料" in panel._content._schema_validation_summary.detail_for(
            "schema_labels"
        )
        assert (
            "Contract party fields"
            in panel._content._schema_validation_summary.tooltip_for("schema_labels")
        )

        signature_index = panel._content._schema_registry_combo.findData(
            "signature_assets_v1"
        )
        assert signature_index >= 0
        panel._content._schema_registry_combo.setCurrentIndex(signature_index)
        panel._content._append_schema_btn.click()
        panel._content._required_material_fields.set_text("contract_no")
        panel._content._required_image_roles.set_text("legal_signature")
        app.processEvents()

        assert scene.input_source_profile.material_schema_ids == [
            "contract_parties_v1",
            "signature_assets_v1",
        ]
        assert scene.input_source_profile.required_material_fields == ["contract_no"]
        assert scene.input_source_profile.required_image_roles == ["legal_signature"]
        material_detail = panel._content._input_summary.detail_for("materials")
        assert "合同方字段资料" in material_detail
        assert "签章资料" in material_detail
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "通过"
        )
        assert (
            panel._content._schema_validation_summary.value_for("schema_rule_overview")
            == "已识别 2 个资料规则"
        )
        overview_detail = panel._content._schema_validation_summary.detail_for(
            "schema_rule_overview"
        )
        assert "字段：5 个必填项" in overview_detail
        assert "图片/签章：2 必填 / 3 全部" in overview_detail
        assert "附件：无" in overview_detail
        assert "contract_parties_v1" not in overview_detail
        assert "contract_no" not in overview_detail
        overview_tooltip = panel._content._schema_validation_summary.tooltip_for(
            "schema_rule_overview"
        )
        assert "contract_parties_v1" in overview_tooltip
        assert "contract_no" in overview_tooltip
        assert "签章资料" in panel._content._schema_validation_summary.detail_for(
            "schema_labels"
        )
        assert (
            "Signature and seal assets"
            in panel._content._schema_validation_summary.tooltip_for("schema_labels")
        )
        assert "合同交付" in panel._content._schema_validation_summary.detail_for(
            "schema_families"
        )
        assert (
            "contract_delivery"
            in panel._content._schema_validation_summary.tooltip_for("schema_families")
        )
        assert (
            panel._content._schema_validation_summary.value_for("schema_fields")
            == "5 个必填项"
        )
        field_detail = panel._content._schema_validation_summary.detail_for(
            "schema_fields"
        )
        assert "合同编号（必填）" in field_detail
        assert "contract_no" not in field_detail
        assert "profile override" not in field_detail
        assert "本方案补充要求" not in field_detail
        field_tooltip = panel._content._schema_validation_summary.tooltip_for(
            "schema_fields"
        )
        assert "contract_no · Contract number · 必填" in field_tooltip
        assert "本方案补充要求" in field_tooltip
        assert "profile override" not in field_tooltip
        assert (
            panel._content._schema_validation_summary.value_for("schema_image_roles")
            == "2 必填 / 3 全部"
        )
        image_detail = panel._content._schema_validation_summary.detail_for(
            "schema_image_roles"
        )
        assert "印章（必填）" in image_detail
        assert "法定代表签名（必填）" in image_detail
        image_tooltip = panel._content._schema_validation_summary.tooltip_for(
            "schema_image_roles"
        )
        assert "seal · Seal · 必填 · image" in image_tooltip
        assert (
            "legal_signature · Legal representative signature · 必填 · image"
            in image_tooltip
        )
        assert "本方案补充要求" in image_tooltip
        assert "profile override" not in image_tooltip
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_attachment_roles"
            )
            == "无"
        )
        assert panel._content._input_summary.value_for("material_fields") == "5 个字段"
        assert "合同编号" in panel._content._input_summary.detail_for("material_fields")
        assert panel._content._input_summary.value_for("image_roles") == "2 个角色"
        image_role_detail = panel._content._input_summary.detail_for("image_roles")
        assert "印章" in image_role_detail
        assert "法定代表签名" in image_role_detail

        panel._content._material_schema_ids.set_text(
            "contract_parties_v1\nmissing_schema_v1"
        )
        app.processEvents()

        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "未识别 1 个"
        )
        assert (
            panel._content._schema_validation_summary.value_for("schema_rule_overview")
            == "已识别 1 个，未识别 1 个"
        )
        assert panel._content._remove_unknown_schema_btn.isEnabled() is True
        assert (
            "missing_schema_v1"
            not in panel._content._schema_validation_summary.detail_for(
                "schema_rule_overview"
            )
        )
        assert (
            "missing_schema_v1"
            in panel._content._schema_validation_summary.tooltip_for(
                "schema_rule_overview"
            )
        )
        assert (
            "missing_schema_v1"
            not in panel._content._schema_validation_summary.detail_for(
                "schema_registry_status"
            )
        )
        assert (
            "missing_schema_v1"
            in panel._content._schema_validation_summary.tooltip_for(
                "schema_registry_status"
            )
        )

        panel._content._remove_unknown_schema_btn.click()
        app.processEvents()

        assert scene.input_source_profile.material_schema_id == "contract_parties_v1"
        assert scene.input_source_profile.material_schema_ids == ["contract_parties_v1"]
        assert "missing_schema_v1" not in panel._content._material_schema_ids.get_text()
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "通过"
        )
        assert panel._content._remove_unknown_schema_btn.isEnabled() is False
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_content_detail_focuses_material_schema_contract_controls():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="鍚堝悓浜や粯",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.input_source_profile.material_schema_ids = [
        "contract_parties_v1",
        "signature_assets_v1",
    ]
    scene.input_source_profile.required_material_fields = ["contract_no"]
    scene.input_source_profile.required_image_roles = ["legal_signature"]
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._content.focus_navigation_field(
            "input_source_profile.material_schema_id"
        )
        assert (
            panel._content._material_schema_id.property("navigation_field_highlight")
            is True
        )
        assert "主资料规则" in panel._content._material_schema_id.toolTip()
        assert (
            "input_source_profile.material_schema_id"
            not in panel._content._material_schema_id.toolTip()
        )
        assert (
            panel._content._material_schema_id.property("navigation_field_raw_label")
            == "input_source_profile.material_schema_id"
        )

        assert panel._content.focus_navigation_field("signature_assets_v1")
        assert (
            panel._content._material_schema_ids.property("navigation_field_highlight")
            is True
        )

        assert panel._content.focus_navigation_field("contract_no")
        assert (
            panel._content._required_material_fields.property(
                "navigation_field_highlight"
            )
            is True
        )

        assert panel._content.focus_navigation_field("legal_signature")
        assert (
            panel._content._required_image_roles.property("navigation_field_highlight")
            is True
        )

        assert (
            panel._content.focus_navigation_field("input_source_profile.unknown")
            is False
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_content_detail_repairs_unknown_primary_material_schema():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    scene.input_source_profile.material_schema_id = "missing_primary_schema_v1"
    scene.input_source_profile.material_schema_ids = [
        "missing_primary_schema_v1",
        "contract_parties_v1",
        "missing_extra_schema_v1",
    ]
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "未识别 2 个"
        )
        assert panel._content._remove_unknown_schema_btn.isEnabled() is True

        panel._content._remove_unknown_schema_btn.click()
        app.processEvents()

        assert scene.input_source_profile.material_schema_id == "contract_parties_v1"
        assert scene.input_source_profile.material_schema_ids == ["contract_parties_v1"]
        assert panel._content._material_schema_id.text() == "contract_parties_v1"
        assert panel._content._material_schema_ids.get_text() == "contract_parties_v1"
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "通过"
        )
        assert "合同方字段资料" in panel._content._schema_validation_summary.detail_for(
            "schema_labels"
        )
        assert (
            "Contract party fields"
            in panel._content._schema_validation_summary.tooltip_for("schema_labels")
        )
        assert panel._content._remove_unknown_schema_btn.isEnabled() is False
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_content_detail_replaces_unknown_material_schema_with_selected_registry_item():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    scene.input_source_profile.material_schema_id = "missing_primary_schema_v1"
    scene.input_source_profile.material_schema_ids = [
        "missing_primary_schema_v1",
        "contract_parties_v1",
        "missing_extra_schema_v1",
    ]
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "未识别 2 个"
        )
        assert panel._content._replace_unknown_schema_btn.isEnabled() is True
        assert "（推荐）" in panel._content._replace_unknown_schema_btn.toolTip()

        signature_index = panel._content._schema_registry_combo.findData(
            "signature_assets_v1"
        )
        assert signature_index >= 0
        panel._content._schema_registry_combo.setCurrentIndex(signature_index)
        app.processEvents()

        assert panel._content._replace_unknown_schema_btn.isEnabled() is True
        assert (
            "signature_assets_v1"
            in panel._content._replace_unknown_schema_btn.toolTip()
        )

        panel._content._replace_unknown_schema_btn.click()
        app.processEvents()

        assert scene.input_source_profile.material_schema_id == "signature_assets_v1"
        assert scene.input_source_profile.material_schema_ids == [
            "signature_assets_v1",
            "contract_parties_v1",
        ]
        assert panel._content._material_schema_id.text() == "signature_assets_v1"
        assert panel._content._material_schema_ids.get_text() == (
            "signature_assets_v1\ncontract_parties_v1"
        )
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "通过"
        )
        assert "签章资料" in panel._content._schema_validation_summary.detail_for(
            "schema_labels"
        )
        assert "合同方字段资料" in panel._content._schema_validation_summary.detail_for(
            "schema_labels"
        )
        labels_tooltip = panel._content._schema_validation_summary.tooltip_for(
            "schema_labels"
        )
        assert "Signature and seal assets" in labels_tooltip
        assert "Contract party fields" in labels_tooltip
        assert panel._content._replace_unknown_schema_btn.isEnabled() is False
        assert panel._content._remove_unknown_schema_btn.isEnabled() is False
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_content_detail_recommends_replacement_for_unknown_material_schema():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.input_source_profile.material_schema_ids = [
        "contract_parties_v1",
        "signature_assets_v2",
    ]
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._content._schema_registry_combo.currentData() == ""
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "未识别 1 个"
        )
        assert panel._content._replace_unknown_schema_btn.isEnabled() is True
        assert (
            "signature_assets_v1"
            in panel._content._replace_unknown_schema_btn.toolTip()
        )
        assert "（推荐）" in panel._content._replace_unknown_schema_btn.toolTip()

        panel._content._replace_unknown_schema_btn.click()
        app.processEvents()

        assert scene.input_source_profile.material_schema_id == "contract_parties_v1"
        assert scene.input_source_profile.material_schema_ids == [
            "contract_parties_v1",
            "signature_assets_v1",
        ]
        assert panel._content._material_schema_ids.get_text() == (
            "contract_parties_v1\nsignature_assets_v1"
        )
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "通过"
        )
        assert "签章资料" in panel._content._schema_validation_summary.detail_for(
            "schema_labels"
        )
        assert (
            "Signature and seal assets"
            in panel._content._schema_validation_summary.tooltip_for("schema_labels")
        )
        assert panel._content._replace_unknown_schema_btn.isEnabled() is False
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_content_detail_previews_attachment_schema_roles():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="qualification_archive_packages",
        name="资质附件包",
        category="qualification_archive_packages",
        template_id="default",
        compatible_template_ids=["default"],
    )
    template = create_builtin_template("default")
    bridge.set_current_scene(
        scene, config_id="qualification_archive_packages", emit_signal=False
    )
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        schema_index = panel._content._schema_registry_combo.findData(
            "qualification_archive_assets_v1"
        )
        assert schema_index >= 0
        panel._content._schema_registry_combo.setCurrentIndex(schema_index)
        panel._content._set_primary_schema_btn.click()
        app.processEvents()

        assert (
            scene.input_source_profile.material_schema_id
            == "qualification_archive_assets_v1"
        )
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_registry_status"
            )
            == "通过"
        )
        assert "标书/资质归档" in panel._content._schema_validation_summary.detail_for(
            "schema_families"
        )
        assert (
            "qualification_archive_packages"
            in panel._content._schema_validation_summary.tooltip_for("schema_families")
        )
        assert (
            panel._content._schema_validation_summary.value_for("schema_fields")
            == "2 必填 / 9 全部"
        )
        field_detail = panel._content._schema_validation_summary.detail_for(
            "schema_fields"
        )
        assert "机构名称（必填）" in field_detail
        assert "资料包名称（必填）" in field_detail
        assert "还有 4 项" in field_detail
        field_tooltip = panel._content._schema_validation_summary.tooltip_for(
            "schema_fields"
        )
        assert "organization · Organization · 必填" in field_tooltip
        assert "valid_until · Valid until · 可选" in field_tooltip
        assert "consortium_member_name · Consortium member name · 可选" in field_tooltip
        assert (
            panel._content._schema_validation_summary.value_for("schema_image_roles")
            == "无"
        )
        assert (
            panel._content._schema_validation_summary.value_for(
                "schema_attachment_roles"
            )
            == "2 必填 / 3 全部"
        )
        attachment_detail = panel._content._schema_validation_summary.detail_for(
            "schema_attachment_roles"
        )
        assert "资质证书（必填，图片 / PDF）" in attachment_detail
        assert "营业执照（必填，图片 / PDF）" in attachment_detail
        assert "附件（可选，图片 / PDF）" in attachment_detail
        attachment_tooltip = panel._content._schema_validation_summary.tooltip_for(
            "schema_attachment_roles"
        )
        assert "assets/01_certificates" in attachment_tooltip
        assert "assets/02_business_license" in attachment_tooltip
        assert (
            "certificate · Qualification certificate · 必填 · image/pdf"
            in attachment_tooltip
        )
        assert (
            "business_license · Business license · 必填 · image/pdf"
            in attachment_tooltip
        )
        assert (
            "attachment · Supporting attachment · 可选 · image/pdf"
            in attachment_tooltip
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_delivery_validation_summary_flags_template_and_rule_issues():
    app = _app()
    bridge = PanelBridge()
    scene = create_thesis_scene()
    template = create_builtin_template("thesis_gbt")
    bridge.set_current_scene(scene, config_id="thesis", emit_signal=False)
    bridge.set_current_template(template, config_id="thesis_gbt", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        review_index = panel._output._default_delivery.findData("review")
        panel._output._default_delivery.setCurrentIndex(review_index)
        app.processEvents()

        panel._output._delivery_filename.setText("{stem}:{unknown}")
        panel._output._visibility_rules.set_text("answer=drop\n=remove\nanswer=hide")
        app.processEvents()

        assert (
            panel._output._delivery_validation_summary.value_for(
                "delivery_template_status"
            )
            == "需处理"
        )
        template_detail = panel._output._delivery_validation_summary.detail_for(
            "delivery_template_status"
        )
        assert "未知变量: unknown" in template_detail
        assert "包含非法字符: :" in template_detail

        assert (
            panel._output._delivery_validation_summary.value_for("delivery_rule_status")
            == "需处理"
        )
        rule_detail = panel._output._delivery_validation_summary.detail_for(
            "delivery_rule_status"
        )
        assert "未知处理动作: drop" in rule_detail
        assert "未知处理动作: hide" in rule_detail
        assert "缺少内容块名" in rule_detail
        assert "重复内容块: 答案 (answer)" in rule_detail
        assert "selector" not in rule_detail

        panel._output._delivery_output_dir.setText("CON/{preset_id}.")
        panel._output._delivery_filename.setText("folder/{stem}")
        app.processEvents()

        path_detail = panel._output._delivery_validation_summary.detail_for(
            "delivery_template_status"
        )
        assert "包含 Windows 保留名: CON" in path_detail
        assert "路径段不能以空格或点结尾: preview." in path_detail
        assert "文件命名不应包含路径分隔符" in path_detail
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_delivery_business_preset_templates_create_common_versions():
    app = _app()
    bridge = PanelBridge()
    scene = create_thesis_scene()
    template = create_builtin_template("thesis_gbt")
    bridge.set_current_scene(scene, config_id="thesis", emit_signal=False)
    bridge.set_current_template(template, config_id="thesis_gbt", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        student_template_index = panel._output._delivery_preset_template_combo.findData(
            "student_version"
        )
        assert student_template_index >= 0
        panel._output._delivery_preset_template_combo.setCurrentIndex(
            student_template_index
        )
        panel._output._add_preset_template_btn.click()
        app.processEvents()

        student_preset = scene.delivery_presets[-1]
        assert student_preset.preset_id == "student_version"
        assert student_preset.label == "学生版"
        assert student_preset.output_dir_template == "{preset_id}"
        assert student_preset.filename_template == "{stem}_{preset_id}"
        assert scene.default_delivery_preset_id == "student_version"
        assert scene.default_delivery_preset().artifacts.compare_docx is False
        assert panel._output._visibility_rules.get_text() == (
            "answer=remove\nanalysis=remove\nsolution=remove"
        )
        assert [
            (rule.selector, rule.action)
            for rule in student_preset.content_visibility_rules
        ] == [
            ("answer", "remove"),
            ("analysis", "remove"),
            ("solution", "remove"),
        ]
        assert [rule.label for rule in student_preset.content_visibility_rules] == [
            "答案",
            "解析",
            "解题过程",
        ]
        assert "显隐规则 3 条" in panel._output._delivery_summary.detail_for(
            "default_artifacts"
        )

        panel._output._add_preset_template_btn.click()
        app.processEvents()

        duplicate_student = scene.delivery_presets[-1]
        assert duplicate_student.preset_id == "student_version_2"
        assert duplicate_student.label == "学生版 2"
        assert scene.default_delivery_preset_id == "student_version_2"

        archive_template_index = panel._output._delivery_preset_template_combo.findData(
            "material_archive"
        )
        assert archive_template_index >= 0
        archive_template_tooltip = (
            panel._output._delivery_preset_template_combo.itemData(
                archive_template_index,
                Qt.ToolTipRole,
            )
        )
        assert "新增交付版本：资料归档包" in archive_template_tooltip
        assert "资料清单" in archive_template_tooltip
        assert "资料包" in archive_template_tooltip
        assert "详细报告" in archive_template_tooltip
        assert "模板编号：material_archive" in archive_template_tooltip
        panel._output._delivery_preset_template_combo.setCurrentIndex(
            archive_template_index
        )
        panel._output._add_preset_template_btn.click()
        app.processEvents()

        archive_preset = scene.delivery_presets[-1]
        assert archive_preset.preset_id == "material_archive"
        assert archive_preset.artifacts.material_manifest is True
        assert archive_preset.artifacts.material_package is True
        assert archive_preset.report_level == "detailed"
        assert scene.default_delivery_preset().artifacts.material_manifest is True
        assert scene.default_delivery_preset().artifacts.material_package is True
        assert "资料清单" in panel._output._delivery_summary.value_for(
            "default_artifacts"
        )
        assert "资料包" in panel._output._delivery_summary.value_for(
            "default_artifacts"
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_applies_contract_family_recommended_delivery_defaults():
    app = _app()
    bridge = PanelBridge()
    scene = SceneWorkspace(
        scene_id="contract_delivery",
        name="合同交付",
        category="contract_delivery",
        template_id="default",
        compatible_template_ids=["default"],
    )
    template = create_builtin_template("default")
    bridge.set_current_scene(scene, config_id="contract_delivery", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        assert panel._output._apply_family_delivery_btn.isEnabled() is True
        family_delivery_tooltip = panel._output._apply_family_delivery_btn.toolTip()
        assert "方案类型：合同交付" in family_delivery_tooltip
        assert (
            "将新增：合同审阅稿、合同签署稿、字段一致性报告" in family_delivery_tooltip
        )
        assert "默认版本：合同审阅稿" in family_delivery_tooltip
        assert (
            "版本 ID：review_copy, signing_copy, field_consistency_report"
            in family_delivery_tooltip
        )

        panel._output._apply_family_delivery_btn.click()
        app.processEvents()

        preset_ids = [preset.preset_id for preset in scene.delivery_presets]
        assert "review_copy" in preset_ids
        assert "signing_copy" in preset_ids
        assert "field_consistency_report" in preset_ids
        assert scene.default_delivery_preset_id == "review_copy"
        assert scene.default_delivery_preset().artifacts.compare_docx is True

        assert scene.input_source_profile.material_schema_id == "contract_parties_v1"
        assert scene.input_source_profile.material_schema_ids == [
            "contract_parties_v1",
            "signature_assets_v1",
        ]
        assert "seal" in scene.input_source_profile.required_image_roles
        assert panel._content._input_summary.value_for("materials") == "必需"
        material_detail = panel._content._input_summary.detail_for("materials")
        assert "合同方字段资料" in material_detail
        assert "签章资料" in material_detail
        assert panel._content._input_summary.value_for("image_roles") == "1 个角色"
        assert "印章" in panel._content._input_summary.detail_for("image_roles")
        assert (
            panel._cleanup._compliance_summary.value_for("count_profile")
            == "contract_fields"
        )
        assert (
            panel._cleanup._compliance_summary.value_for("planning_scan_targets")
            == "已应用"
        )
        assert (
            panel._output._delivery_summary.value_for("default_delivery")
            == "合同审阅稿"
        )
        assert "对比稿" in panel._output._delivery_summary.value_for(
            "default_artifacts"
        )

        signing_index = panel._output._default_delivery.findData("signing_copy")
        assert signing_index >= 0
        panel._output._default_delivery.setCurrentIndex(signing_index)
        app.processEvents()

        assert scene.default_delivery_preset_id == "signing_copy"
        assert scene.default_delivery_preset().artifacts.material_manifest is True
        assert scene.default_delivery_preset().artifacts.material_package is True
        assert "资料清单" in panel._output._delivery_summary.value_for(
            "default_artifacts"
        )
        assert "资料包" in panel._output._delivery_summary.value_for(
            "default_artifacts"
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_delivery_variable_picker_inserts_into_templates():
    app = _app()
    bridge = PanelBridge()
    scene = create_thesis_scene()
    template = create_builtin_template("thesis_gbt")
    bridge.set_current_scene(scene, config_id="thesis", emit_signal=False)
    bridge.set_current_template(template, config_id="thesis_gbt", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        review_index = panel._output._default_delivery.findData("review")
        panel._output._default_delivery.setCurrentIndex(review_index)
        app.processEvents()

        review_preset = next(
            preset for preset in scene.delivery_presets if preset.preset_id == "review"
        )
        panel._output._delivery_output_dir.setText("review/")
        panel._output._delivery_output_dir.setCursorPosition(len("review/"))
        variable_index = panel._output._delivery_variable_combo.findData("preset_label")
        assert "preset_label" not in panel._output._delivery_variable_combo.itemText(
            variable_index
        )
        assert "{preset_label}" in panel._output._delivery_variable_combo.itemData(
            variable_index,
            Qt.ToolTipRole,
        )
        panel._output._delivery_variable_combo.setCurrentIndex(variable_index)
        panel._output._insert_output_variable_btn.click()
        app.processEvents()

        assert review_preset.output_dir_template == "review/{preset_label}"
        assert "review/审阅稿" in panel._output._delivery_validation_summary.detail_for(
            "delivery_template_status"
        )

        panel._output._delivery_filename.setText("{stem}_")
        panel._output._delivery_filename.setCursorPosition(len("{stem}_"))
        variable_index = panel._output._delivery_variable_combo.findData("preset_id")
        panel._output._delivery_variable_combo.setCurrentIndex(variable_index)
        panel._output._insert_filename_variable_btn.click()
        app.processEvents()

        assert review_preset.filename_template == "{stem}_{preset_id}"
        assert (
            "document_review"
            in panel._output._delivery_validation_summary.detail_for(
                "delivery_template_status"
            )
        )
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_delivery_visibility_rule_inserter_updates_current_preset():
    app = _app()
    bridge = PanelBridge()
    scene = create_thesis_scene()
    template = create_builtin_template("thesis_gbt")
    bridge.set_current_scene(scene, config_id="thesis", emit_signal=False)
    bridge.set_current_template(template, config_id="thesis_gbt", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        review_index = panel._output._default_delivery.findData("review")
        panel._output._default_delivery.setCurrentIndex(review_index)
        app.processEvents()

        review_preset = next(
            preset for preset in scene.delivery_presets if preset.preset_id == "review"
        )
        selector_index = panel._output._visibility_selector_combo.findData("answer")
        assert selector_index >= 0
        assert (
            panel._output._visibility_selector_combo.itemText(selector_index) == "答案"
        )
        assert "answer" in panel._output._visibility_selector_combo.itemData(
            selector_index,
            Qt.ToolTipRole,
        )
        panel._output._visibility_selector_combo.setCurrentIndex(selector_index)
        app.processEvents()
        assert panel._output._visibility_selector_input.text() == "answer"

        action_index = panel._output._visibility_action_combo.findData("remove")
        assert action_index >= 0
        panel._output._visibility_action_combo.setCurrentIndex(action_index)
        panel._output._insert_visibility_rule_btn.click()
        app.processEvents()

        assert panel._output._visibility_rules.get_text() == "answer=remove"
        assert panel._output._visibility_selector_combo.currentData() in ("", None)
        assert [
            (rule.selector, rule.action)
            for rule in review_preset.content_visibility_rules
        ] == [("answer", "remove")]
        assert [rule.label for rule in review_preset.content_visibility_rules] == [
            "答案"
        ]
        assert (
            panel._output._delivery_validation_summary.value_for("delivery_rule_status")
            == "1 条"
        )
        assert "显隐规则 1 条" in panel._output._delivery_summary.detail_for(
            "default_artifacts"
        )

        panel._output._visibility_selector_input.setText("{{#visibility:Answer}}")
        panel._output._insert_visibility_rule_btn.click()
        app.processEvents()

        assert panel._output._visibility_rules.get_text() == "answer=remove"
        assert len(review_preset.content_visibility_rules) == 1
    finally:
        panel.close()
        app.processEvents()


def test_scene_panel_delivery_target_template_combo_and_status():
    app = _app()
    bridge = PanelBridge()
    scene = create_thesis_scene()
    template = create_builtin_template("thesis_gbt")
    bridge.set_current_scene(scene, config_id="thesis", emit_signal=False)
    bridge.set_current_template(template, config_id="thesis_gbt", emit_signal=False)

    panel = ScenePanel(bridge)
    try:
        app.processEvents()

        review_index = panel._output._default_delivery.findData("review")
        panel._output._default_delivery.setCurrentIndex(review_index)
        app.processEvents()

        review_preset = next(
            preset for preset in scene.delivery_presets if preset.preset_id == "review"
        )
        thesis_template_index = panel._output._delivery_target_template_combo.findData(
            "thesis_gbt"
        )
        assert thesis_template_index >= 0
        assert (
            "thesis_gbt"
            not in panel._output._delivery_target_template_combo.itemText(
                thesis_template_index
            )
        )
        thesis_template_tooltip = (
            panel._output._delivery_target_template_combo.itemData(
                thesis_template_index,
                Qt.ToolTipRole,
            )
        )
        assert "thesis_gbt" in thesis_template_tooltip

        panel._output._delivery_target_template_combo.setCurrentIndex(
            thesis_template_index
        )
        app.processEvents()

        assert review_preset.target_template_id == "thesis_gbt"
        assert panel._output._delivery_target_template.text() == "thesis_gbt"
        assert (
            panel._output._delivery_validation_summary.value_for(
                "delivery_target_template"
            )
            == "兼容模板"
        )
        assert "thesis_gbt" in panel._output._delivery_validation_summary.detail_for(
            "delivery_target_template"
        )

        panel._output._delivery_target_template.setText("default")
        app.processEvents()

        assert (
            panel._output._delivery_validation_summary.value_for(
                "delivery_target_template"
            )
            == "需处理"
        )
        assert "default" in panel._output._delivery_validation_summary.detail_for(
            "delivery_target_template"
        )
        default_template_index = panel._output._delivery_target_template_combo.findData(
            "default"
        )
        assert default_template_index >= 0
        default_template_tooltip = (
            panel._output._delivery_target_template_combo.itemData(
                default_template_index,
                Qt.ToolTipRole,
            )
        )
        assert "default" in default_template_tooltip

        panel._output._delivery_target_template.setText("missing_template_id")
        app.processEvents()

        assert (
            panel._output._delivery_validation_summary.value_for(
                "delivery_target_template"
            )
            == "需处理"
        )
        assert (
            "模板不存在: missing_template_id"
            in panel._output._delivery_validation_summary.detail_for(
                "delivery_target_template"
            )
        )
    finally:
        panel.close()
        app.processEvents()
