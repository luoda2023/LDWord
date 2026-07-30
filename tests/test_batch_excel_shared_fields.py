from datetime import date
from pathlib import Path

from openpyxl import Workbook

from src.ui.panels.assets.batch_import import (
    _load_batch_excel_rows,
    _profiles_from_table_rows,
)


def test_excel_batch_import_merges_single_row_auxiliary_sheets(tmp_path: Path) -> None:
    workbook = Workbook()
    projects = workbook.active
    projects.title = "项目基础数据"
    projects.append(("项目名称", "项目结束日期"))
    projects.append(("项目甲", date(2026, 7, 26)))
    projects.append(("项目乙", None))

    people = workbook.create_sheet("人员配置")
    people.append(("姓名1", "姓名2"))
    people.append(("张三", "李四"))

    custom = workbook.create_sheet("自定义文本")
    custom.append(("自定义文本1",))
    custom.append(("统一说明",))

    ignored = workbook.create_sheet("多行明细")
    ignored.append(("不得合并",))
    ignored.append(("第一行",))
    ignored.append(("第二行",))

    path = tmp_path / "batch.xlsx"
    workbook.save(path)

    rows = _load_batch_excel_rows(path)
    profiles = _profiles_from_table_rows(rows)

    assert len(rows) == 2
    assert rows[0]["姓名1"] == "张三"
    assert rows[1]["自定义文本1"] == "统一说明"
    assert "不得合并" not in rows[0]
    assert profiles[0].fields["项目结束日期"] == "2026-07-26"
    assert profiles[1].fields["姓名2"] == "李四"
