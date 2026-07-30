from __future__ import annotations

from hashlib import sha256
import inspect
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import pytest

from src.config.content_artifacts import ArtifactFileRef, ContentArtifactResource
from src.config.content_materials import ContentResourceKey, FileAssetRef
from src.config.image_materials import (
    FrozenImageMaterialRule,
    ImageAnchorOrigin,
    ImageCardinality,
    ImageCoLocationGuard,
    ImageMaterialRule,
    ImageOccurrencePolicy,
    ImagePlacementMode,
    ImagePlacementPolicy,
    ImageSourceBinding,
    ImageSourceItem,
    ImageWatermarkPolicy,
    ImageWatermarkTextSource,
    ResolvedImageWatermark,
)
from src.services.material_assets.image_plan_builder import (
    ContentResourceImageSource,
    ImageInsertionPlanBuilder,
    ImagePlanBuildError,
    ImagePlanReceipt,
)
from src.services.material_assets.image_policy_freezer import (
    ImagePolicyFreezeError,
    freeze_image_policy,
)
from src.services.material_content.docx_renderer import ContentImageJobDraft


SNAPSHOT_ID = "a" * 64


def _image_ref(tmp_path: Path, name="asset.png", *, image_format="PNG", media_type="image/png"):
    path = tmp_path / name
    Image.new("RGB", (32, 24), (20, 90, 160)).save(path, format=image_format)
    payload = path.read_bytes()
    return FileAssetRef(
        source_path=str(path),
        original_name=name,
        media_type=media_type,
        content_sha256=sha256(payload).hexdigest(),
        byte_size=len(payload),
    )


def _disabled():
    return ResolvedImageWatermark.disabled()


def _fixed_rule(*, cardinality=ImageCardinality.SINGLE, occurrence=ImageOccurrencePolicy.EXACTLY_ONE):
    return FrozenImageMaterialRule(
        rule_id="rule-iso",
        source_role="iso",
        anchor_token="{{@img:ISO1}}",
        required=True,
        occurrence_policy=occurrence,
        cardinality=cardinality,
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIXED_BOX_FLOW,
            fixed_width_cm=15,
        ),
        watermark=_disabled(),
    )


def _strict_rule():
    return FrozenImageMaterialRule(
        rule_id="rule-iso",
        source_role="iso",
        anchor_token="{{@img:ISO1}}",
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
            max_width_cm=16,
            co_location_guard=ImageCoLocationGuard.PRECEDING_NONEMPTY_PARAGRAPH,
        ),
        watermark=_disabled(),
    )


def _source(ref, *, item_id="image-1", sequence=0):
    return ImageSourceBinding("iso", item_id, sequence, ref)


def _font_path():
    candidates = (
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibri.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    path = next((item for item in candidates if item.is_file()), None)
    if path is None:
        pytest.skip("no deterministic test font")
    return path


def _xml(document):
    return document.element.xml


def _codes(error):
    return {item.code for item in error.diagnostics}


def _bookmark_names(document):
    return [
        item.get(qn("w:name"))
        for item in document.element.iter(qn("w:bookmarkStart"))
        if item.get(qn("w:name"))
    ]


def _append_content_sentinel(document, marker, bookmark_id=40):
    paragraph = document.add_paragraph()._p
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), marker)
    paragraph.append(start)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:vanish"))
    rpr.append(OxmlElement("w:noProof"))
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = marker
    run.append(text)
    paragraph.append(run)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph.append(end)


def test_persisted_and_frozen_rule_contracts_are_distinct_and_strict():
    persisted = ImageMaterialRule(
        rule_id="rule-iso",
        source_role="iso",
        anchor_token="{{@img:ISO1}}",
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIXED_BOX_FLOW,
            fixed_width_cm=15,
        ),
    )
    assert ImageMaterialRule.from_dict(persisted.to_dict()) == persisted
    assert FrozenImageMaterialRule.from_dict(_fixed_rule().to_dict()) == _fixed_rule()
    with pytest.raises(ValueError, match="multiple images require a flow placement"):
        ImageMaterialRule(
            rule_id="bad",
            source_role="iso",
            anchor_token="{{@img:ISO1}}",
            cardinality=ImageCardinality.MULTIPLE,
            placement=ImagePlacementPolicy(
                mode=ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
                max_width_cm=16,
                co_location_guard=ImageCoLocationGuard.PRECEDING_NONEMPTY_PARAGRAPH,
            ),
        )


@pytest.mark.parametrize(
    "mode",
    (
        ImagePlacementMode.NATURAL_SIZE,
        ImagePlacementMode.FIT_CONTAINER_FLOW,
    ),
)
def test_global_flow_modes_support_multiple_images(mode: ImagePlacementMode) -> None:
    rule = ImageMaterialRule(
        rule_id=f"rule-{mode.value}",
        source_role="iso",
        anchor_token="{{@img:ISO1}}",
        cardinality=ImageCardinality.MULTIPLE,
        placement=ImagePlacementPolicy(mode=mode),
    )
    frozen = FrozenImageMaterialRule(
        rule_id=rule.rule_id,
        source_role=rule.source_role,
        anchor_token=rule.anchor_token,
        cardinality=rule.cardinality,
        placement=rule.placement,
        watermark=_disabled(),
    )

    assert ImageMaterialRule.from_dict(rule.to_dict()) == rule
    assert FrozenImageMaterialRule.from_dict(frozen.to_dict()) == frozen


def test_policy_freezer_resolves_field_and_font_once_and_planner_cannot_re_resolve(tmp_path):
    values = {"company": "Acme A"}
    persisted = ImageMaterialRule(
        rule_id="rule-iso",
        source_role="iso",
        anchor_token="{{@img:ISO1}}",
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIXED_BOX_FLOW,
            fixed_width_cm=15,
        ),
        watermark=ImageWatermarkPolicy(
            enabled=True,
            text_template="Company {{@text:company}}",
        ),
    )
    ref = _image_ref(tmp_path)
    frozen = freeze_image_policy(
        frozen_field_values=values,
        rules=(persisted,),
        source_items=(ImageSourceItem("iso", "image-1", 0, ref),),
        watermark_font_path=_font_path(),
        watermark_font_identity="test-font",
    )
    values["company"] = "Acme B"
    document = Document()
    document.add_paragraph("{{@img:ISO1}}")
    result = ImageInsertionPlanBuilder().build(
        document,
        snapshot_id=SNAPSHOT_ID,
        frozen_rules=frozen.rules,
        source_bindings=frozen.source_bindings,
    )
    watermark = result.plans[0].watermark
    assert watermark.resolved_text == "Company Acme A"
    assert watermark.resolved_font_identity == "test-font"
    parameters = inspect.signature(ImageInsertionPlanBuilder.build).parameters
    assert "field_values" not in parameters
    assert "watermark_font_path" not in parameters
    assert "watermark_font_identity" not in parameters


def test_policy_freezer_freezes_workbench_watermark_text_at_runtime_boundary(tmp_path):
    persisted = ImageMaterialRule(
        rule_id="rule-iso",
        source_role="iso",
        anchor_token="{{@img:ISO1}}",
        placement=ImagePlacementPolicy(mode=ImagePlacementMode.NATURAL_SIZE),
        watermark=ImageWatermarkPolicy(
            enabled=True,
            text_source=ImageWatermarkTextSource.WORKBENCH_FREE_FIELD,
        ),
    )
    ref = _image_ref(tmp_path)

    frozen = freeze_image_policy(
        frozen_field_values={},
        rules=(persisted,),
        source_items=(ImageSourceItem("iso", "image-1", 0, ref),),
        runtime_watermark_text="本次项目专用",
        watermark_font_path=_font_path(),
    )

    assert frozen.rules[0].watermark.resolved_text == "本次项目专用"
    with pytest.raises(ImagePolicyFreezeError) as caught:
        freeze_image_policy(
            frozen_field_values={},
            rules=(persisted,),
            source_items=(ImageSourceItem("iso", "image-1", 0, ref),),
            watermark_font_path=_font_path(),
        )
    assert getattr(caught.value, "code", "") == (
        "watermark_runtime_watermark_text_missing"
    )


def test_single_strict_plan_freezes_guard_without_layout_mutation(tmp_path):
    document = Document()
    heading = document.add_paragraph("2.1 资质证书")
    heading.paragraph_format.keep_with_next = True
    document.add_paragraph("{{@img:ISO1}}")
    before_ppr = heading._p.pPr.xml
    result = ImageInsertionPlanBuilder().build(
        document,
        snapshot_id=SNAPSHOT_ID,
        frozen_rules=(_strict_rule(),),
        source_bindings=(_source(_image_ref(tmp_path)),),
    )
    assert result.plans[0].anchor.guard_marker_id.startswith("LarkGuard_")
    assert heading._p.pPr.xml == before_ppr
    assert not list(document.element.iter(qn("w:drawing")))
    assert not list(document.element.iter(qn("w:br")))


def test_multiple_sources_are_ordered_and_get_distinct_sentinels(tmp_path):
    document = Document()
    document.add_paragraph("{{@img:ISO1}}")
    result = ImageInsertionPlanBuilder().build(
        document,
        snapshot_id=SNAPSHOT_ID,
        frozen_rules=(_fixed_rule(cardinality=ImageCardinality.MULTIPLE),),
        source_bindings=(
            _source(_image_ref(tmp_path, "second.png"), item_id="second", sequence=2),
            _source(_image_ref(tmp_path, "first.png"), item_id="first", sequence=1),
        ),
    )
    assert [item.sequence for item in result.plans] == [1, 2]
    assert len(_bookmark_names(document)) == 2


def test_content_resource_source_is_already_resolved_and_defaults_disabled(tmp_path):
    document = Document()
    marker = "LarkContentImage_123"
    _append_content_sentinel(document, marker)
    before = _xml(document)
    ref = _image_ref(tmp_path, "resource.png")
    source = ContentResourceImageSource(
        resource=ContentArtifactResource(
            resource_id="assets/resource.png",
            media_type="image/png",
            file=ArtifactFileRef(
                path="resources/resource.png",
                sha256=ref.content_sha256,
                byte_size=ref.byte_size,
            ),
        ),
        source_path=ref.source_path,
    )
    assert isinstance(source.watermark, ResolvedImageWatermark)
    assert not source.watermark.enabled
    result = ImageInsertionPlanBuilder().build(
        document,
        snapshot_id=SNAPSHOT_ID,
        frozen_rules=(),
        source_bindings=(),
        content_image_drafts=(
            ContentImageJobDraft(
                job_id="content-job",
                content_id="route",
                resource_id="assets/resource.png",
                stable_marker_id=marker,
                occurrence_id="content-occ",
                sequence=0,
                reserved_docpr_id=50,
            ),
        ),
        content_resource_sources={ContentResourceKey("route", "assets/resource.png"): source},
    )
    assert _xml(document) == before
    assert result.plans[0].anchor.origin is ImageAnchorOrigin.CONTENT_RESOURCE
    assert result.plans[0].watermark == source.watermark


def test_content_resource_mapping_requires_namespaced_key(tmp_path):
    document = Document()
    marker = "LarkContentImage_badkey"
    _append_content_sentinel(document, marker)
    ref = _image_ref(tmp_path)
    source = ContentResourceImageSource(
        ContentArtifactResource(
            "asset.png",
            "image/png",
            ArtifactFileRef("resources/asset.png", ref.content_sha256, ref.byte_size),
        ),
        ref.source_path,
    )
    with pytest.raises(TypeError, match="ContentResourceKey"):
        ImageInsertionPlanBuilder().build(
            document,
            snapshot_id=SNAPSHOT_ID,
            frozen_rules=(),
            source_bindings=(),
            content_resource_sources={"asset.png": source},
        )


@pytest.mark.parametrize(
    ("configure", "expected"),
    [
        (lambda doc: (doc.add_paragraph("{{@img:ISO1}}"), doc.add_paragraph("{{@img:ISO1}}")), "token_duplicate_occurrence"),
        (lambda doc: doc.add_paragraph("before {{@img:ISO1}} after"), "image_anchor_not_isolated"),
        (lambda doc: doc.add_table(1, 1).cell(0, 0).paragraphs[0].add_run("{{@img:ISO1}}"), "image_anchor_unsupported_surface"),
        (lambda doc: doc.sections[0].header.add_paragraph("{{@img:ISO1}}"), "image_anchor_unsupported_surface"),
    ],
)
def test_invalid_anchor_surfaces_are_atomic(tmp_path, configure, expected):
    document = Document()
    configure(document)
    before = _xml(document)
    with pytest.raises(ImagePlanBuildError) as raised:
        ImageInsertionPlanBuilder().build(
            document,
            snapshot_id=SNAPSHOT_ID,
            frozen_rules=(_fixed_rule(),),
            source_bindings=(_source(_image_ref(tmp_path)),),
        )
    assert expected in _codes(raised.value)
    assert _xml(document) == before


def test_image_anchor_with_non_plain_ooxml_is_blocked_atomically(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("{{@img:ISO1}}")
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "7")
    bookmark.set(qn("w:name"), "must-not-be-deleted")
    paragraph._p.append(bookmark)
    before = _xml(document)

    with pytest.raises(ImagePlanBuildError) as raised:
        ImageInsertionPlanBuilder().build(
            document,
            snapshot_id=SNAPSHOT_ID,
            frozen_rules=(_fixed_rule(),),
            source_bindings=(_source(_image_ref(tmp_path)),),
        )

    assert "image_anchor_not_isolated" in _codes(raised.value)
    assert _xml(document) == before


def test_image_anchor_section_boundary_keeps_specific_diagnostic(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("{{@img:ISO1}}")
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    before = _xml(document)

    with pytest.raises(ImagePlanBuildError) as raised:
        ImageInsertionPlanBuilder().build(
            document,
            snapshot_id=SNAPSHOT_ID,
            frozen_rules=(_fixed_rule(),),
            source_bindings=(_source(_image_ref(tmp_path)),),
        )

    assert "image_anchor_section_boundary" in _codes(raised.value)
    assert "image_anchor_not_isolated" not in _codes(raised.value)
    assert _xml(document) == before


def test_missing_anchor_and_missing_strict_guard_are_blocking(tmp_path):
    ref = _image_ref(tmp_path)
    for rule, text, expected in (
        (_fixed_rule(), "No token", "token_missing_occurrence"),
        (_strict_rule(), "{{@img:ISO1}}", "same_page_guard_missing"),
    ):
        document = Document()
        document.add_paragraph(text)
        before = _xml(document)
        with pytest.raises(ImagePlanBuildError) as raised:
            ImageInsertionPlanBuilder().build(
                document,
                snapshot_id=SNAPSHOT_ID,
                frozen_rules=(rule,),
                source_bindings=(_source(ref),),
            )
        assert expected in _codes(raised.value)
        assert _xml(document) == before


@pytest.mark.parametrize("failure", ["mime", "extension", "hash", "docx"])
def test_polluted_or_changed_image_files_are_blocked(tmp_path, failure):
    if failure == "docx":
        path = tmp_path / "pollution.docx"
        Document().save(path)
        payload = path.read_bytes()
        ref = FileAssetRef(
            str(path), path.name,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            sha256(payload).hexdigest(), len(payload),
        )
        # Strong ImageSourceBinding rejects this before planning.
        with pytest.raises(ValueError, match="image media type"):
            _source(ref)
        return
    if failure == "extension":
        ref = _image_ref(tmp_path, "wrong.jpg", image_format="PNG")
        expected = "image_extension_mismatch"
    else:
        ref = _image_ref(tmp_path, media_type="image/jpeg" if failure == "mime" else "image/png")
        expected = "image_mime_mismatch" if failure == "mime" else "image_source_hash_mismatch"
        if failure == "hash":
            payload = bytearray(Path(ref.source_path).read_bytes())
            payload[-1] ^= 1
            Path(ref.source_path).write_bytes(payload)
    document = Document()
    document.add_paragraph("{{@img:ISO1}}")
    with pytest.raises(ImagePlanBuildError) as raised:
        ImageInsertionPlanBuilder().build(
            document,
            snapshot_id=SNAPSHOT_ID,
            frozen_rules=(_fixed_rule(),),
            source_bindings=(_source(ref),),
        )
    assert expected in _codes(raised.value)


def test_plan_markers_and_receipt_are_deterministic(tmp_path):
    ref = _image_ref(tmp_path)
    def build_one():
        document = Document()
        document.add_paragraph("{{@img:ISO1}}")
        result = ImageInsertionPlanBuilder().build(
            document,
            snapshot_id=SNAPSHOT_ID,
            frozen_rules=(_fixed_rule(),),
            source_bindings=(_source(ref),),
        )
        return document, result
    first_doc, first = build_one()
    second_doc, second = build_one()
    assert first.plans == second.plans
    assert first.receipt == second.receipt
    assert _bookmark_names(first_doc) == _bookmark_names(second_doc)


def test_image_plan_receipt_cannot_be_mistaken_for_variant_identity():
    parameters = inspect.signature(ImagePlanReceipt).parameters
    assert "variant_id" not in parameters
    assert "variant_version" not in parameters
    assert "staged_docx_sha256" not in parameters
    assert "plan_id" not in parameters
    assert "Only ``DeliveryVariantPlan``" in (ImagePlanReceipt.__doc__ or "")
