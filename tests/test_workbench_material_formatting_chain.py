from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

from docx import Document
from PIL import Image

from src.application.materials import (
    ExecutionMaterialRecord,
    ExecutionMaterialSnapshot,
    project_execution_material_record_snapshot,
)
from src.config.builtin_templates import create_builtin_template
from src.config.execution_feature_state import project_execution_scene
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.domain.materials import MaterialPackageRef, MaterialRunSelection
from src.modules.fill.entity_fill import EntityFillModule
from src.services.document_structure_evidence import build_document_structure_evidence
from src.services.docx_format_change import compare_docx_formatting
from src.services.execution_run_log import append_execution_run
from src.services.production_execution import (
    ProductionExecutionRequest,
    execute_production_request,
)
from src.services.production_runtime.formatting_runtime import (
    apply_required_format_change_policy,
)
from src.ui.panels.workbench.document_input_manifest import (
    discover_document_inputs,
)
from src.ui.panels.workbench.execution_session_controller import (
    ExecutionBuildResult,
    WorkbenchExecutionSessionController,
    _batch_format_change_evidence,
    _FileProductionRunner,
    _ProductionRunner,
)


def _record(record_id: str, display_name: str) -> ExecutionMaterialRecord:
    return ExecutionMaterialRecord(
        record_id=record_id,
        display_name=display_name,
        group_id="",
        field_values={},
        field_owners={},
        resources={},
        resource_owners={},
    )


def _snapshot(*records: ExecutionMaterialRecord) -> ExecutionMaterialSnapshot:
    return ExecutionMaterialSnapshot(
        snapshot_id="",
        run_id="run-test",
        package_ref=MaterialPackageRef(
            package_id=f"pkg_{'a' * 32}",
            revision=f"sha256:{'b' * 64}",
        ),
        work_mode_id="custom",
        material_contract_id="generic_document_v1",
        recipe_id="document_batch",
        scene_id="custom",
        document_type="",
        package_display_name="Materials",
        package_field_values={},
        package_field_owners={},
        groups=(),
        records=records,
    )


def _selection() -> MaterialRunSelection:
    return MaterialRunSelection(
        package_ref=MaterialPackageRef(
            package_id=f"pkg_{'a' * 32}",
            revision=f"sha256:{'b' * 64}",
        ),
        selected_record_ids=(f"rec_{'c' * 32}",),
    )


def test_material_selection_builds_the_formatting_runner(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    snapshot = _snapshot(_record(f"rec_{'c' * 32}", "Record"))
    monkeypatch.setattr(
        "src.ui.panels.workbench.execution_session_controller."
        "bind_workbench_material",
        lambda *_args, **_kwargs: (snapshot, ()),
    )
    monkeypatch.setattr(
        "src.ui.panels.workbench.execution_session_controller."
        "build_threaded_runner",
        lambda runner, **_kwargs: ExecutionBuildResult(worker=runner),
    )
    controller = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    )
    mapped_root = tmp_path / "outputs" / "source"
    evidence = build_document_structure_evidence(source)

    build = controller.build_file_batch_worker(
        document_paths=(str(source),),
        template=TemplateConfig(name="Template"),
        scene=SceneWorkspace(mode_id="custom"),
        selection=_selection(),
        mode_id="custom",
        template_id="template",
        output_root=str(tmp_path / "outputs"),
        output_roots_by_path={str(source).casefold(): str(mapped_root)},
        document_structure_evidence=evidence,
    )

    assert isinstance(build.worker, _ProductionRunner)
    request = build.worker._request
    assert request.material_snapshot is not None
    assert request.require_format_change is True
    assert request.output_root == mapped_root
    assert request.output_suffix == "-formatted"
    assert request.document_structure_evidence is evidence


def test_plan_off_material_selection_does_not_require_format_change(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    snapshot = _snapshot(_record(f"rec_{'c' * 32}", "Record"))
    monkeypatch.setattr(
        "src.ui.panels.workbench.execution_session_controller."
        "bind_workbench_material",
        lambda *_args, **_kwargs: (snapshot, ()),
    )
    monkeypatch.setattr(
        "src.ui.panels.workbench.execution_session_controller."
        "build_threaded_runner",
        lambda runner, **_kwargs: ExecutionBuildResult(worker=runner),
    )
    scene = project_execution_scene(
        SceneWorkspace(mode_id="custom"),
        plan_enabled=False,
        template_enabled=True,
        material_enabled=True,
    )
    controller = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    )

    build = controller.build_file_batch_worker(
        document_paths=(str(source),),
        template=TemplateConfig(name="Template"),
        scene=scene,
        selection=_selection(),
        mode_id="custom",
        template_id="template",
        output_root=str(tmp_path / "outputs"),
    )

    assert isinstance(build.worker, _ProductionRunner)
    assert build.worker._request.require_format_change is False


def test_multiple_material_records_each_enter_production_formatting(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    first = _record(f"rec_{'c' * 32}", "First")
    second = _record(f"rec_{'d' * 32}", "Second")
    snapshot = _snapshot(first, second)
    monkeypatch.setattr(
        "src.ui.panels.workbench.execution_session_controller."
        "bind_workbench_material",
        lambda *_args, **_kwargs: (snapshot, ()),
    )
    monkeypatch.setattr(
        "src.ui.panels.workbench.execution_session_controller."
        "build_threaded_runner",
        lambda runner, **_kwargs: ExecutionBuildResult(worker=runner),
    )
    controller = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    )
    output_root = tmp_path / "outputs"

    build = controller.build_record_batch_worker(
        document_paths=(str(source),),
        template=TemplateConfig(name="Template"),
        scene=SceneWorkspace(mode_id="custom"),
        selection=_selection(),
        mode_id="custom",
        template_id="template",
        output_root=str(output_root),
    )

    assert isinstance(build.worker, _FileProductionRunner)
    assert [
        request.material_snapshot.records[0].record_id
        for request in build.worker._requests
    ] == [first.record_id, second.record_id]
    assert {
        request.output_root.name for request in build.worker._requests
    } == {"First", "Second"}


def test_material_fields_are_filled_inside_production_pipeline(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("{{@text:client}}")
    document.save(source)
    record = replace(
        _record(f"rec_{'c' * 32}", "Record"),
        field_values={"client": "ACME"},
        field_owners={"client": "record"},
    )
    monkeypatch.setattr(
        "src.services.production_runtime.execution_runtime.create_all_modules",
        lambda: [EntityFillModule()],
    )

    result = execute_production_request(
        ProductionExecutionRequest(
            input_path=source,
            output_root=tmp_path / "outputs",
            mode_id="custom",
            scene=SceneWorkspace(mode_id="custom"),
            template=create_builtin_template("default", mode_id="custom"),
            plan_id="custom",
            template_id="default",
            output_suffix="-formatted",
            material_snapshot=_snapshot(record),
        )
    )

    output = Path(str(result["output_path"]))
    assert result["status"] == "success"
    assert result["modules_enabled"] == 1
    assert Document(output).paragraphs[0].text == "ACME"


def test_record_projection_resets_stale_finalization():
    record = _record(f"rec_{'c' * 32}", "Record")
    snapshot = _snapshot(record)
    finalized_shape = replace(
        snapshot,
        template_id="old-template",
        template_revision=f"sha256:{'1' * 64}",
        master_id="old-master",
        master_revision=f"sha256:{'2' * 64}",
        output_root="C:\\old",
        output_paths=("C:\\old\\result.docx",),
        preflight_receipt={"status": "passed"},
    )

    projected = project_execution_material_record_snapshot(
        finalized_shape,
        record.record_id,
    )

    assert projected.records == (record,)
    assert not projected.execution_ready
    assert projected.template_id == ""
    assert projected.output_paths == ()


def test_output_root_cannot_contain_the_input_document(tmp_path, monkeypatch):
    output_root = tmp_path / "outputs"
    output_root.mkdir()
    source = output_root / "source.docx"
    Document().save(source)
    controller = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    )

    build = controller.build_file_batch_worker(
        document_paths=(str(source),),
        template=TemplateConfig(name="Template"),
        scene=SceneWorkspace(mode_id="custom"),
        selection=None,
        mode_id="custom",
        output_root=str(output_root),
    )

    assert build.worker is None
    assert build.error_text.startswith("workbench_output_root_contains_input:")


def test_output_planning_rejects_same_stem_collision(tmp_path):
    docx_source = tmp_path / "source.docx"
    markdown_source = tmp_path / "source.md"
    Document().save(docx_source)
    markdown_source.write_text("# source", encoding="utf-8")
    controller = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: docx_source,
    )

    build = controller.build_file_batch_worker(
        document_paths=(str(docx_source), str(markdown_source)),
        template=TemplateConfig(name="Template"),
        scene=SceneWorkspace(mode_id="custom"),
        selection=None,
        mode_id="custom",
        output_root=str(tmp_path / "outputs"),
    )

    assert build.worker is None
    assert build.error_text.startswith("workbench_output_path_collision:")


def test_manifest_does_not_reingest_alavette_output_directories(tmp_path):
    Document().save(tmp_path / "source.docx")
    generated = tmp_path / "Alavette-Output" / "通用文档"
    generated.mkdir(parents=True)
    Document().save(generated / "generated.docx")

    manifest = discover_document_inputs(
        (str(tmp_path),),
        accepted_suffixes=(".docx",),
    )

    assert [Path(item.source_path).name for item in manifest.documents] == [
        "source.docx"
    ]


def test_semantic_format_evidence_ignores_noop_package_rewrite(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Body")
    document.save(source)
    noop = tmp_path / "noop.docx"
    Document(source).save(noop)
    changed = tmp_path / "changed.docx"
    modified = Document(source)
    modified.paragraphs[0].runs[0].bold = True
    modified.save(changed)

    noop_evidence = compare_docx_formatting(source, noop)
    changed_evidence = compare_docx_formatting(source, changed)

    assert noop_evidence["format_changed"] is False
    assert changed_evidence["format_changed"] is True


def test_content_evidence_ignores_run_segmentation(tmp_path):
    source = tmp_path / "source.docx"
    original = Document()
    original.add_paragraph("Body")
    original.save(source)
    segmented = tmp_path / "segmented.docx"
    modified = Document(source)
    paragraph = modified.paragraphs[0]
    paragraph.clear()
    paragraph.add_run("Bo")
    paragraph.add_run("dy")
    modified.save(segmented)

    evidence = compare_docx_formatting(source, segmented)

    assert evidence["content_changed"] is False


def test_content_evidence_detects_inserted_images(tmp_path):
    source = tmp_path / "source.docx"
    original = Document()
    original.add_paragraph("{{@img:photo}}")
    original.save(source)
    image = tmp_path / "photo.png"
    Image.new("RGB", (64, 32), "navy").save(image)
    changed = tmp_path / "changed.docx"
    modified = Document(source)
    paragraph = modified.paragraphs[0]
    paragraph.clear()
    paragraph.add_run().add_picture(str(image))
    modified.save(changed)

    evidence = compare_docx_formatting(source, changed)

    assert evidence["content_changed"] is True


def test_required_format_change_downgrades_noop_success():
    payload = {"status": "success", "summary": "done"}
    evidence = {"status": "compared", "format_changed": False}

    apply_required_format_change_policy(payload, evidence, required=True)

    assert payload["status"] == "partial_success"
    assert payload["warnings"][0].startswith("formatting_noop:")


def test_batch_format_evidence_reports_any_unchanged_document():
    evidence = _batch_format_change_evidence(
        [
            {
                "format_change_evidence": {
                    "status": "compared",
                    "format_changed": True,
                }
            },
            {
                "format_change_evidence": {
                    "status": "compared",
                    "format_changed": False,
                }
            },
        ]
    )

    assert evidence["format_changed"] is False
    assert evidence["documents_unchanged"] == 1


def test_file_batch_aggregates_partial_artifact_failures_and_warnings(
    tmp_path,
    monkeypatch,
):
    requests = tuple(
        ProductionExecutionRequest(
            input_path=tmp_path / f"source-{index}.docx",
            output_root=tmp_path / "outputs",
            mode_id="custom",
            scene=SceneWorkspace(mode_id="custom"),
            template=TemplateConfig(name="Template"),
            plan_id="",
        )
        for index in (1, 2)
    )
    responses = iter(
        (
            {
                "status": "partial_success",
                "output_path": str(tmp_path / "outputs" / "one.docx"),
                "artifact_failure_count": 2,
                "artifact_failures": [
                    {"kind": "report", "error": "disk full"}
                ],
                "warnings": ["report_incomplete"],
                "error_text": "artifact publish incomplete",
            },
            {
                "status": "success",
                "output_path": str(tmp_path / "outputs" / "two.docx"),
                "artifact_failure_count": 0,
                "warnings": [],
                "error_text": "",
            },
        )
    )
    monkeypatch.setattr(
        "src.ui.panels.workbench.execution_session_controller."
        "execute_production_request",
        lambda *_args, **_kwargs: next(responses),
    )

    result = _FileProductionRunner(requests).run(
        lambda *_args: None,
        lambda: False,
    )

    assert result["status"] == "partial_success"
    assert result["summary"] == "完整完成 1/2 份文档；部分完成 1 份"
    assert result["failed_count"] == 0
    assert result["partial_success_count"] == 1
    assert result["artifact_failure_count"] == 2
    assert result["artifact_failures"][0]["source_path"].endswith(
        "source-1.docx"
    )
    assert result["warnings"] == ["source-1.docx: report_incomplete"]
    assert result["items"][0]["artifact_failure_count"] == 2


def test_file_batch_does_not_count_cancelled_document_as_complete(
    tmp_path,
    monkeypatch,
):
    requests = tuple(
        ProductionExecutionRequest(
            input_path=tmp_path / f"source-{index}.docx",
            output_root=tmp_path / "outputs",
            mode_id="custom",
            scene=SceneWorkspace(mode_id="custom"),
            template=TemplateConfig(name="Template"),
            plan_id="",
        )
        for index in (1, 2)
    )
    responses = iter(
        (
            {
                "status": "cancelled",
                "output_path": "",
                "error_text": "execution_cancelled",
            },
            {
                "status": "success",
                "output_path": str(tmp_path / "outputs" / "two.docx"),
                "error_text": "",
            },
        )
    )
    monkeypatch.setattr(
        "src.ui.panels.workbench.execution_session_controller."
        "execute_production_request",
        lambda *_args, **_kwargs: next(responses),
    )

    result = _FileProductionRunner(requests).run(
        lambda *_args: None,
        lambda: False,
    )

    assert result["status"] == "partial_success"
    assert result["summary"] == "完整完成 1/2 份文档；取消 1 份"
    assert result["cancelled_count"] == 1
    assert result["success_count"] == 1
    assert "1 份文档已取消" in result["error_text"]


def test_execution_receipt_persists_branch_and_format_evidence(tmp_path):
    path = append_execution_run(
        {
            "status": "partial_success",
            "execution_route": "production_with_material",
            "input_paths": ["C:/input.docx"],
            "output_path": "C:/output.docx",
            "format_change_evidence": {
                "schema_version": "docx-format-change-v1",
                "status": "compared",
                "format_changed": False,
            },
        },
        log_root=tmp_path,
    )

    payload = json.loads(path.read_text(encoding="utf-8"))
    assert payload["execution_route"] == "production_with_material"
    assert payload["format_change_evidence"]["format_changed"] is False
