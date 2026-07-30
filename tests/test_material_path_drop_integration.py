from __future__ import annotations

from docx import Document
from PIL import Image
from PySide6.QtCore import QMimeData, QPoint, QPointF, QUrl
from PySide6.QtGui import QDragEnterEvent, QDropEvent

from src.config.scene import SceneWorkspace
from src.qt_api import QApplication, Qt
from src.ui.bridge import PanelBridge
from src.ui.panels.assets import (
    asset_file_operations_presenter,
    asset_group_rows_presenter,
    content_materials_presenter,
)
from src.ui.panels.assets.specs import AttachmentRoleSpec
from src.ui.panels.assets_panel import AssetsPanel
from src.ui.panels.workbench import quick_execution_drop_area
from src.ui.panels.workbench.quick_execution_drop_area import QuickExecutionDropArea


def _drop_local_path(widget, path) -> None:
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(path))])
    drag = QDragEnterEvent(
        QPoint(2, 2),
        Qt.CopyAction,
        mime,
        Qt.LeftButton,
        Qt.NoModifier,
    )
    QApplication.sendEvent(widget, drag)
    assert drag.isAccepted()
    drop = QDropEvent(
        QPointF(2, 2),
        Qt.CopyAction,
        mime,
        Qt.LeftButton,
        Qt.NoModifier,
    )
    QApplication.sendEvent(widget, drop)
    assert drop.isAccepted()


def _bidding_panel() -> AssetsPanel:
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "bid_materials_v1"
    bridge = PanelBridge()
    bridge.set_current_scene(scene, config_id="bidding", emit_signal=False)
    return AssetsPanel(bridge)


def test_content_material_drop_runs_the_same_strict_binding_as_browse(
    tmp_path,
    qapp,
    monkeypatch,
):
    source = tmp_path / "content.docx"
    document = Document()
    document.add_heading("文件资料", level=1)
    document.add_paragraph("正文")
    document.save(source)
    browse_panel = AssetsPanel(PanelBridge())
    drop_panel = AssetsPanel(PanelBridge())
    try:
        browse_panel._request_add_content_material()
        drop_panel._request_add_content_material()
        browse_id = browse_panel._content_rules[-1].content_id
        drop_id = drop_panel._content_rules[-1].content_id
        monkeypatch.setattr(
            content_materials_presenter.QFileDialog,
            "getOpenFileName",
            staticmethod(lambda *_args: (str(source), "Word")),
        )

        browse_panel._select_content_material_file(browse_id)
        _drop_local_path(drop_panel._content_material_path_edits[drop_id], source)

        browse_binding = browse_panel._content_bindings[browse_id]
        drop_binding = drop_panel._content_bindings[drop_id]
        assert browse_binding == drop_binding
        artifact = drop_panel._content_repository().validate(drop_binding.artifact_ref)
        assert artifact.manifest.source.original_name == source.name
        assert drop_panel._content_material_path_edits[drop_id].text() == source.name
        assert drop_panel._content_material_thumbnail_labels[drop_id].text() == (
            "DOCX"
        )
    finally:
        browse_panel.close()
        drop_panel.close()


def test_invalid_docx_drop_keeps_content_binding_unchanged(tmp_path, qapp):
    source = tmp_path / "invalid.docx"
    source.write_bytes(b"not-an-openxml-package")
    panel = AssetsPanel(PanelBridge())
    try:
        panel._request_add_content_material()
        content_id = panel._content_rules[-1].content_id

        _drop_local_path(panel._content_material_path_edits[content_id], source)

        assert content_id not in panel._content_bindings
        assert panel._content_material_path_edits[content_id].text() == ""
        assert panel._content_material_thumbnail_labels[content_id].text() == "未选"
    finally:
        panel.close()


def test_single_image_and_multi_image_folder_use_shared_drop_acquisition(
    tmp_path,
    qapp,
    monkeypatch,
):
    image = tmp_path / "logo.png"
    Image.new("RGB", (80, 80), color="blue").save(image)
    folder = tmp_path / "qualification"
    folder.mkdir()
    Image.new("RGB", (80, 80), color="green").save(folder / "1.png")
    browse_panel = _bidding_panel()
    drop_panel = _bidding_panel()
    try:
        monkeypatch.setattr(
            asset_file_operations_presenter.QFileDialog,
            "getOpenFileName",
            staticmethod(lambda *_args: (str(image), "Image")),
        )
        monkeypatch.setattr(
            asset_group_rows_presenter.QFileDialog,
            "getExistingDirectory",
            staticmethod(lambda *_args: str(folder)),
        )

        browse_panel._select_asset_file("logo")
        browse_panel._select_asset_group_directory("qualification")
        _drop_local_path(drop_panel._asset_slot_path_edits["logo"], image)
        _drop_local_path(
            drop_panel._asset_group_path_labels["qualification"],
            folder,
        )

        assert browse_panel._asset_paths["logo"] == drop_panel._asset_paths["logo"]
        assert (
            browse_panel._asset_bindings["qualification"]
            == drop_panel._asset_bindings["qualification"]
        )
        assert drop_panel._asset_bindings["qualification"].source_path == str(
            folder.resolve()
        )
        assert len(drop_panel._asset_bindings["qualification"].items) == 1
    finally:
        browse_panel.close()
        drop_panel.close()


def test_attachment_browse_and_drop_run_the_same_intake(
    tmp_path,
    qapp,
    monkeypatch,
):
    pdf = tmp_path / "evidence.pdf"
    pdf.write_bytes(b"%PDF-1.4\nevidence")
    spec = AttachmentRoleSpec(
        role="evidence",
        label="证明材料",
        accepted_types=("pdf",),
    )
    browse_panel = AssetsPanel(PanelBridge())
    drop_panel = AssetsPanel(PanelBridge())
    try:
        for panel in (browse_panel, drop_panel):
            panel._attachment_role_specs = (spec,)
            panel._sync_attachment_role_rows()
        monkeypatch.setattr(
            asset_file_operations_presenter.QFileDialog,
            "getOpenFileName",
            staticmethod(lambda *_args: (str(pdf), "PDF")),
        )

        browse_panel._select_attachment_file("evidence")
        _drop_local_path(
            drop_panel._attachment_role_status_labels["evidence"],
            pdf,
        )

        binding = drop_panel._attachment_bindings["evidence"]
        assert browse_panel._attachment_bindings["evidence"] == binding
        assert binding.items[0].file_ref.source_path == str(pdf.resolve())
    finally:
        browse_panel.close()
        drop_panel.close()


def test_attachment_directory_browse_and_drop_preserve_the_same_package_tree(
    tmp_path,
    qapp,
    monkeypatch,
):
    folder = tmp_path / "qualification"
    (folder / "part-a").mkdir(parents=True)
    (folder / "part-b").mkdir()
    (folder / "part-a" / "proof.pdf").write_bytes(b"%PDF-1.4\na")
    (folder / "part-b" / "proof.pdf").write_bytes(b"%PDF-1.4\nb")
    spec = AttachmentRoleSpec(
        role="qualification_package",
        label="资质文件包",
        accepted_types=("pdf",),
        cardinality="multiple",
        source_kind="directory_package",
        recursive=True,
        max_items=None,
    )
    browse_panel = AssetsPanel(PanelBridge())
    drop_panel = AssetsPanel(PanelBridge())
    try:
        for panel in (browse_panel, drop_panel):
            panel._attachment_role_specs = (spec,)
            panel._sync_attachment_role_rows()
        monkeypatch.setattr(
            asset_file_operations_presenter.QFileDialog,
            "getExistingDirectory",
            staticmethod(lambda *_args: str(folder)),
        )

        browse_panel._select_attachment_file("qualification_package")
        _drop_local_path(
            drop_panel._attachment_role_status_labels["qualification_package"],
            folder,
        )

        binding = drop_panel._attachment_bindings["qualification_package"]
        assert browse_panel._attachment_bindings["qualification_package"] == binding
        assert [item.relative_path for item in binding.items] == [
            "part-a/proof.pdf",
            "part-b/proof.pdf",
        ]
        assert binding.source_path == str(folder.resolve())
    finally:
        browse_panel.close()
        drop_panel.close()


def test_workbench_browse_and_drop_use_the_same_selection(tmp_path, qapp, monkeypatch):
    source = tmp_path / "input.docx"
    source.write_bytes(b"docx")
    browse_area = QuickExecutionDropArea()
    drop_area = QuickExecutionDropArea()
    try:
        monkeypatch.setattr(
            quick_execution_drop_area.QFileDialog,
            "getOpenFileName",
            staticmethod(lambda *_args: (str(source), "Word")),
        )
        browse_area._pick_file()
        _drop_local_path(drop_area._hint_label, source)

        assert browse_area.file_path() == drop_area.file_path() == str(source.resolve())
        assert drop_area._selected_container.isHidden() is False
    finally:
        browse_area.close()
        drop_area.close()
