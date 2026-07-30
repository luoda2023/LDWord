from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from src.config.material_import_draft import (
    ImportMappingProfile,
    inspect_material_workbook,
)
from src.config.material_package_v6 import (
    MATERIAL_PACKAGE_VERSION,
    MaterialPackageV6,
    load_material_package_v6,
    migrate_entity_archive_to_v6,
    save_material_package_v6,
    synchronize_material_package_from_archive,
)
from src.config.material_scope import MaterialScopeResolver
from src.material_suite.plan import (
    GenerationRecipe,
    compile_material_suite_plan,
    discover_material_suite_bundle,
)
from src.material_suite.runner import MaterialSuiteGenerationRunner


def _workbook_with_six_products_and_three_slots(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "项目基础数据"
    sheet.append(
        [
            "产品名称",
            "项目序号",
            "项目名称",
            "项目编号",
            "项目来源",
            "项目开始日期",
            "项目结束日期",
        ]
    )
    products = [f"{index:02d} 产品{index}" for index in range(1, 7)]
    for product_index, product in enumerate(products, start=1):
        for sequence in range(1, 4):
            active = product_index == 1
            sheet.append(
                [
                    product,
                    sequence,
                    f"project-{sequence}" if active else "",
                    f"P-{sequence:03d}" if active else "",
                    f"source-{sequence}" if active else "",
                    f"2026-0{sequence}-01" if active else "",
                    f"2026-0{sequence}-10" if active else "",
                ]
            )
    shared = workbook.create_sheet("人员配置")
    shared.append(["负责人", "审核人"])
    shared.append(["张三", "李四"])
    workbook.save(path)
    workbook.close()


def _word_template(
    path: Path,
    text: str = "{{项目名称}} / {{项目编号}} / {{负责人}}",
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def test_workbook_inspection_distinguishes_slots_active_records_and_drafts(
    tmp_path: Path,
) -> None:
    source = tmp_path / "项目数据.xlsx"
    _workbook_with_six_products_and_three_slots(source)

    draft = inspect_material_workbook(
        source,
        mapping=ImportMappingProfile(
            activation_required_fields=("项目名称", "项目编号"),
        ),
    )

    assert draft.candidate_count == 18
    assert draft.active_count == 3
    assert draft.draft_count == 15
    assert len(draft.groups) == 6
    assert draft.shared_scope.fields == {"负责人": "张三", "审核人": "李四"}
    assert [item.source_locator["row"] for item in draft.candidates[:3]] == [2, 3, 4]
    assert all(item.lifecycle_state == "active" for item in draft.candidates[:3])
    assert all(item.lifecycle_state == "draft" for item in draft.candidates[3:])


def test_suite_compiles_only_active_records_and_drafts_do_not_block(
    tmp_path: Path,
) -> None:
    source = tmp_path / "项目数据.xlsx"
    _workbook_with_six_products_and_three_slots(source)
    package = inspect_material_workbook(source).materialize()
    template_root = tmp_path / "模板文件"
    _word_template(template_root / "Word文档" / "计划书.docx")

    plan = compile_material_suite_plan(
        package,
        discover_material_suite_bundle(template_root),
        output_root=tmp_path / "输出",
        recipe=GenerationRecipe(
            required_record_fields=("项目名称", "项目编号"),
        ),
    )

    assert plan.ok
    assert plan.inspection.candidate_count == 18
    assert plan.inspection.active_count == 3
    assert plan.inspection.draft_count == 15
    assert plan.inspection.selected_count == 3
    assert plan.inspection.executable_count == 3
    assert len(plan.records) == 3
    assert plan.artifact_count == 3
    assert all(record.group_id == package.groups[0].group_id for record in plan.records)
    assert all("负责人" in dict(record.frozen_values) for record in plan.records)


def test_scope_resolution_is_layered_and_keeps_value_provenance(tmp_path: Path) -> None:
    source = tmp_path / "项目数据.xlsx"
    _workbook_with_six_products_and_three_slots(source)
    package = inspect_material_workbook(source).materialize()
    record = package.records[0]

    resolved = MaterialScopeResolver(package).resolve_record(record.record_id)

    assert resolved.values["负责人"] == "张三"
    assert resolved.values["项目名称"] == "project-1"
    assert resolved.provenance["负责人"].scope == "package"
    assert resolved.provenance["项目名称"].scope == "record"


def test_v6_round_trip_and_v5_migration_are_explicit_and_non_destructive(
    tmp_path: Path,
) -> None:
    source = tmp_path / "项目数据.xlsx"
    _workbook_with_six_products_and_three_slots(source)
    package = inspect_material_workbook(source).materialize()
    target = tmp_path / "package-v6.json"

    save_material_package_v6(package, target)
    payload = json.loads(target.read_text(encoding="utf-8"))
    loaded = load_material_package_v6(target)

    assert payload["version"] == MATERIAL_PACKAGE_VERSION == 6
    assert isinstance(loaded, MaterialPackageV6)
    assert len(loaded.records) == 18
    assert loaded.shared_scope.fields["负责人"] == "张三"

    legacy = tmp_path / "legacy-v5.json"
    from src.config.entity import EntityArchive, EntityProfile, save_entity_archive

    save_entity_archive(
        EntityArchive(
            archive_id="legacy",
            archive_name="旧资料包",
            profiles=[
                EntityProfile(
                    profile_id="one",
                    profile_name="项目一",
                    fields={"产品名称": "产品A", "项目名称": "项目一"},
                )
            ],
        ),
        legacy,
    )
    migrated = tmp_path / "migrated-v6.json"
    report = migrate_entity_archive_to_v6(legacy, migrated)

    assert legacy.is_file()
    assert migrated.is_file()
    assert report.source_version == 5
    assert report.target_version == 6
    assert report.record_count == 1
    assert load_material_package_v6(migrated).records[0].record_id == "one"


def test_template_manifest_supports_package_group_and_record_emit_scopes(
    tmp_path: Path,
) -> None:
    source = tmp_path / "项目数据.xlsx"
    _workbook_with_six_products_and_three_slots(source)
    package = inspect_material_workbook(source).materialize()
    template_root = tmp_path / "模板文件"
    _word_template(template_root / "package.docx", "{{package_name}} / {{负责人}}")
    _word_template(template_root / "group.docx", "{{产品名称}} / {{负责人}}")
    _word_template(template_root / "record.docx")
    manifest = {
        "kind": "alavette.template_suite",
        "version": 1,
        "bundle_id": "scope-suite",
        "bundle_name": "作用域套件",
        "artifacts": [
            {
                "artifact_id": "package",
                "source": "package.docx",
                "kind": "docx",
                "emit_scope": "package_once",
                "target": "Word文档/{source_name}",
            },
            {
                "artifact_id": "group",
                "source": "group.docx",
                "kind": "docx",
                "emit_scope": "group_once",
                "target": "Word文档/{source_name}",
            },
            {
                "artifact_id": "record",
                "source": "record.docx",
                "kind": "docx",
                "emit_scope": "record_once",
                "target": "Word文档/{source_name}",
            },
        ],
    }
    (template_root / "template-suite.json").write_text(
        json.dumps(manifest, ensure_ascii=False),
        encoding="utf-8",
    )

    plan = compile_material_suite_plan(
        package,
        discover_material_suite_bundle(template_root),
        output_root=tmp_path / "输出",
    )

    assert plan.ok
    assert len(plan.records) == 3
    assert [item.emit_scope for item in plan.shared_units] == [
        "package_once",
        "group_once",
    ]
    assert plan.artifact_count == 5
    result = MaterialSuiteGenerationRunner(plan).run(
        lambda *_args: None,
        lambda: False,
    )
    assert result["status"] == "success"
    assert (tmp_path / "输出" / "_资料包公共" / "Word文档" / "package.docx").is_file()
    assert (
        tmp_path
        / "输出"
        / "01 产品1"
        / "_分组公共"
        / "Word文档"
        / "group.docx"
    ).is_file()


def test_template_files_are_scanned_once_per_compile_not_once_per_record(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "项目数据.xlsx"
    _workbook_with_six_products_and_three_slots(source)
    package = inspect_material_workbook(source).materialize()
    template_root = tmp_path / "模板文件"
    _word_template(template_root / "Word文档" / "计划书.docx")
    bundle = discover_material_suite_bundle(template_root)
    import src.material_suite.plan as plan_module

    original = plan_module.scan_suite_template_placeholders
    calls: list[str] = []

    def counted(source_path, kind):
        calls.append(str(source_path))
        return original(source_path, kind)

    monkeypatch.setattr(
        plan_module,
        "scan_suite_template_placeholders",
        counted,
    )
    plan = compile_material_suite_plan(
        package,
        bundle,
        output_root=tmp_path / "输出",
    )

    assert plan.ok
    assert len(plan.records) == 3
    assert calls == [bundle.artifacts[0].source_path]


def test_explicitly_selected_draft_is_visible_but_never_a_generation_unit(
    tmp_path: Path,
) -> None:
    source = tmp_path / "项目数据.xlsx"
    _workbook_with_six_products_and_three_slots(source)
    package = inspect_material_workbook(source).materialize()
    draft = package.records[3]
    template_root = tmp_path / "模板文件"
    _word_template(template_root / "Word文档" / "计划书.docx")

    plan = compile_material_suite_plan(
        package,
        discover_material_suite_bundle(template_root),
        output_root=tmp_path / "输出",
        profile_ids=[draft.record_id],
    )

    assert not plan.ok
    assert plan.records == ()
    assert plan.inspection.selected_count == 1
    assert plan.inspection.executable_count == 0
    inspected = next(
        item
        for item in plan.inspection.records
        if item.record_id == draft.record_id
    )
    assert inspected.readiness == "lifecycle_draft"


def test_legacy_editor_round_trip_preserves_hierarchy_and_lifecycle(
    tmp_path: Path,
) -> None:
    source = tmp_path / "项目数据.xlsx"
    _workbook_with_six_products_and_three_slots(source)
    package = inspect_material_workbook(source).materialize()
    archive = package.to_entity_archive()
    archive.profiles[3].fields["项目名称"] = "后来补全的项目"
    archive.profiles[3].fields["项目编号"] = "LATE-004"

    refreshed = synchronize_material_package_from_archive(package, archive)
    record = refreshed.records[3]

    assert record.lifecycle_state == "draft"
    assert record.values.fields["项目名称"] == "后来补全的项目"
    assert "负责人" not in record.values.fields
    assert refreshed.shared_scope.fields["负责人"] == "张三"
    assert refreshed.get_group(record.group_id).group_name == "02 产品2"
