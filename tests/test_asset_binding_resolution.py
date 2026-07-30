from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image

from src.config.asset_resolution import (
    asset_item_from_payload,
    asset_item_payload,
    directory_content_revision,
    natural_relative_path_key,
    normalize_asset_item_payloads,
    refresh_asset_binding,
    resolve_asset_binding,
    resolve_profile_assets,
)
from src.config import asset_resolution as asset_resolution_module
from src.config import material_batch as material_batch_module
from src.config.entity import (
    AssetBinding,
    EntityArchive,
    EntityProfile,
    load_entity_archive,
    save_entity_archive_bundle,
)
from src.config import entity_bundle as entity_bundle_module
from src.config.materials import (
    AssetInsertionRule,
    AssetItem,
    build_image_insertions,
    missing_required_asset_roles,
)
from src.config.material_context import MaterialExecutionContext
from src.config.material_batch import (
    build_material_batch_items,
    check_material_batch_preflight,
)
from src.config.resolved import ImageInsertionItem, ResolvedConfig
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.modules.insert.image_insertion import ImageInsertionModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.services.production_runtime.material_preflight import (
    image_anchor_diagnostics,
    material_requirement_diagnostics,
    undeclared_image_token_diagnostics,
)
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner
from src.ui.panels.assets.batch_import import _profile_from_mapping


def _image(path: Path, payload: bytes = b"image") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)
    return path


def test_asset_item_payload_normalization_preserves_snapshot_identity():
    raw = {
        "item_id": "qualification:007",
        "label": "Qualification 7",
        "role": "qualification",
        "path": "C:/materials/qualification/7.png",
        "mime_type": "image/png",
        "tags": ["primary", "signed"],
        "width_cm": 8.5,
        "metadata": {"asset_id": "remote-7", "source": "library"},
        "group_id": "qualification",
        "sequence": 7,
        "source_path": "C:/materials/qualification",
        "original_relative_path": "signed/7.png",
        "normalized_name": "qualification_007.png",
        "content_hash": "a" * 64,
    }

    normalized = normalize_asset_item_payloads([raw])

    assert normalized == [raw]
    assert asset_item_payload(asset_item_from_payload(normalized[0])) == raw


def test_asset_item_payload_does_not_read_a_legacy_sha256_alias():
    item = asset_item_from_payload(
        {
            "item_id": "logo:001",
            "role": "logo",
            "path": "logo.png",
            "sha256": "sha256:" + "a" * 64,
        }
    )

    assert item.content_hash == ""
    assert "sha256" not in asset_item_payload(item)


def test_directory_revision_preserves_file_boundaries(tmp_path):
    one_file = tmp_path / "one-file"
    two_files = tmp_path / "two-files"
    one_file.mkdir()
    two_files.mkdir()
    (one_file / "a").write_bytes(b"b\0")
    (two_files / "a").write_bytes(b"")
    (two_files / "b").write_bytes(b"")

    # The retired raw concatenation serialized both trees as b"a\0b\0".
    # The canonical revision hashes each file first, preserving boundaries.
    assert directory_content_revision(one_file) != directory_content_revision(
        two_files
    )


def test_bundle_producer_and_batch_validator_share_directory_revision_owner():
    assert (
        entity_bundle_module.directory_content_revision
        is asset_resolution_module.directory_content_revision
    )
    assert (
        material_batch_module.directory_content_revision
        is asset_resolution_module.directory_content_revision
    )
    assert not hasattr(entity_bundle_module, "_directory_content_revision")
    assert not hasattr(material_batch_module, "_directory_revision")


def test_directory_binding_recurses_naturally_and_assigns_canonical_identity(tmp_path):
    source = tmp_path / "qualification"
    ten = _image(source / "A2" / "10.jpg", b"ten")
    one = _image(source / "A2" / "1.jpg", b"one")
    two = _image(source / "A2" / "2.png", b"two")
    nested = _image(source / "A10" / "1.jpg", b"nested")
    (source / "A2" / "notes.txt").write_text("skip", encoding="utf-8")

    binding = AssetBinding(
        role="qualification",
        cardinality="multiple",
        source_kind="directory",
        source_path=str(source),
        recursive=True,
        min_items=1,
        max_items=None,
    )
    resolution = resolve_asset_binding(binding)

    assert [item.original_relative_path for item in resolution.items] == [
        "A2/1.jpg",
        "A2/2.png",
        "A2/10.jpg",
        "A10/1.jpg",
    ]
    assert [item.sequence for item in resolution.items] == [1, 2, 3, 4]
    assert [item.item_id for item in resolution.items] == [
        "qualification:001",
        "qualification:002",
        "qualification:003",
        "qualification:004",
    ]
    assert [item.normalized_name for item in resolution.items] == [
        "qualification_001.jpg",
        "qualification_002.png",
        "qualification_003.jpg",
        "qualification_004.jpg",
    ]
    assert all(item.group_id == "qualification" for item in resolution.items)
    assert any(item.code == "asset_file_type_unsupported" for item in resolution.diagnostics)
    assert all(path.is_file() for path in (one, two, ten, nested))
    assert one.name == nested.name == "1.jpg"


def test_directory_binding_can_be_non_recursive(tmp_path):
    source = tmp_path / "images"
    _image(source / "2.jpg", b"two")
    _image(source / "nested" / "1.jpg", b"one")

    resolution = resolve_asset_binding(
        AssetBinding(
            role="product_image",
            cardinality="multiple",
            source_kind="directory",
            source_path=str(source),
            recursive=False,
            max_items=None,
        )
    )

    assert [item.original_relative_path for item in resolution.items] == ["2.jpg"]


def test_directory_binding_reports_duplicate_content(tmp_path):
    source = tmp_path / "images"
    _image(source / "1.jpg", b"same")
    _image(source / "2.jpg", b"same")

    resolution = resolve_asset_binding(
        AssetBinding(
            role="case_image",
            cardinality="multiple",
            source_kind="directory",
            source_path=str(source),
            recursive=True,
            max_items=None,
        )
    )

    assert any(item.code == "asset_duplicate_content" for item in resolution.diagnostics)


def test_directory_binding_reports_missing_empty_and_unreadable_sources(
    tmp_path,
    monkeypatch,
):
    missing = resolve_asset_binding(
        AssetBinding(
            role="case_image",
            cardinality="multiple",
            source_kind="directory",
            source_path=str(tmp_path / "missing"),
            recursive=True,
            max_items=None,
        )
    )
    empty_dir = tmp_path / "empty"
    empty_dir.mkdir()
    empty = resolve_asset_binding(
        AssetBinding(
            role="case_image",
            cardinality="multiple",
            source_kind="directory",
            source_path=str(empty_dir),
            recursive=True,
            max_items=None,
        )
    )
    unreadable_path = _image(tmp_path / "unreadable" / "1.png", b"image")
    original_revision = asset_resolution_module.file_content_revision
    monkeypatch.setattr(
        asset_resolution_module,
        "file_content_revision",
        lambda path: "" if path == unreadable_path else original_revision(path),
    )
    unreadable = resolve_asset_binding(
        AssetBinding(
            role="case_image",
            cardinality="multiple",
            source_kind="directory",
            source_path=str(unreadable_path.parent),
            recursive=True,
            max_items=None,
        )
    )

    assert {item.code for item in missing.diagnostics} == {"asset_directory_missing"}
    assert {item.code for item in empty.diagnostics} == {"asset_directory_empty"}
    assert unreadable.items == []
    assert "asset_file_unreadable" in {item.code for item in unreadable.diagnostics}


def test_natural_relative_path_key_orders_numeric_segments():
    paths = ["A10/1.jpg", "A2/10.jpg", "A2/2.jpg", "A2/1.jpg"]
    assert sorted(paths, key=natural_relative_path_key) == [
        "A2/1.jpg",
        "A2/2.jpg",
        "A2/10.jpg",
        "A10/1.jpg",
    ]


def test_build_image_insertions_preserves_all_multiple_items_and_single_uses_first():
    items = [
        AssetItem(
            item_id="qualification:002",
            role="qualification",
            path="two.jpg",
            group_id="qualification",
            sequence=2,
            normalized_name="qualification_002.jpg",
        ),
        AssetItem(
            item_id="qualification:001",
            role="qualification",
            path="one.jpg",
            group_id="qualification",
            sequence=1,
            normalized_name="qualification_001.jpg",
        ),
    ]

    multiple = build_image_insertions(
        items,
        [
            AssetInsertionRule(
                asset_role="qualification",
                target="{{qualification}}",
                cardinality="multiple",
                max_items=None,
            )
        ],
    )
    single = build_image_insertions(
        items,
        [AssetInsertionRule(asset_role="qualification", target="{{qualification}}")],
    )

    assert [item.path for item in multiple] == ["one.jpg", "two.jpg"]
    assert [item.sequence for item in multiple] == [1, 2]
    assert [item.group_id for item in multiple] == ["qualification", "qualification"]
    assert [item.path for item in single] == ["one.jpg"]


def test_refresh_binding_persists_snapshot_order(tmp_path):
    source = tmp_path / "figures"
    _image(source / "10.png", b"ten")
    _image(source / "2.png", b"two")
    binding = AssetBinding(
        role="figure",
        cardinality="multiple",
        source_kind="directory",
        source_path=str(source),
        recursive=True,
        max_items=None,
    )

    refreshed, resolution = refresh_asset_binding(binding)
    snapshot_resolution = resolve_asset_binding(refreshed)

    assert [item.original_relative_path for item in resolution.items] == ["2.png", "10.png"]
    assert [item.sequence for item in snapshot_resolution.items] == [1, 2]
    assert refreshed.snapshot_revision.startswith("sha256:")
    assert len(refreshed.items) == 2


def test_binding_snapshot_reports_directory_changes_until_rescanned(tmp_path):
    source = tmp_path / "figures"
    _image(source / "1.png", b"one")
    refreshed, _resolution = refresh_asset_binding(
        AssetBinding(
            role="figure",
            cardinality="multiple",
            source_kind="directory",
            source_path=str(source),
            recursive=True,
            max_items=None,
        )
    )
    _image(source / "2.png", b"two")

    stale = resolve_asset_binding(refreshed)
    refreshed_again, current = refresh_asset_binding(refreshed)

    assert any(item.code == "asset_binding_snapshot_stale" for item in stale.diagnostics)
    assert len(stale.items) == 1
    assert len(current.items) == 2
    assert refreshed_again.snapshot_revision != refreshed.snapshot_revision


def test_binding_snapshot_does_not_emit_deleted_items(tmp_path):
    source = tmp_path / "figures"
    deleted = _image(source / "1.png", b"one")
    refreshed, _resolution = refresh_asset_binding(
        AssetBinding(
            role="figure",
            cardinality="multiple",
            source_kind="directory",
            source_path=str(source),
            recursive=True,
            max_items=None,
        )
    )
    deleted.unlink()

    resolution = resolve_asset_binding(refreshed)

    assert resolution.items == []
    assert any(
        item.code == "asset_file_missing" and item.path == str(deleted)
        for item in resolution.diagnostics
    )
    assert not resolution.ok


def test_image_binding_snapshot_ignores_non_image_directory_pollution(tmp_path):
    source = tmp_path / "figures"
    _image(source / "1.png", b"one")
    pollution = source / "notes.docx"
    pollution.write_bytes(b"first")
    refreshed, resolution = refresh_asset_binding(
        AssetBinding(
            role="figure",
            cardinality="multiple",
            source_kind="directory",
            source_path=str(source),
            recursive=True,
            max_items=None,
        )
    )
    revision = refreshed.snapshot_revision

    pollution.write_bytes(b"second")
    cached = resolve_asset_binding(refreshed)
    refreshed_again, rescanned = refresh_asset_binding(refreshed)

    assert [item.path for item in resolution.items] == [str(source / "1.png")]
    assert any(
        item.code == "asset_file_type_unsupported"
        and item.path == str(pollution)
        for item in resolution.diagnostics
    )
    assert not any(
        item.code == "asset_binding_snapshot_stale"
        for item in cached.diagnostics
    )
    assert refreshed_again.snapshot_revision == revision
    assert [item.path for item in rescanned.items] == [str(source / "1.png")]


def test_profile_resolution_isolates_non_image_legacy_role_but_keeps_attachment(
    tmp_path,
):
    wrong_logo = tmp_path / "logo.docx"
    attachment = tmp_path / "evidence.docx"
    wrong_logo.write_bytes(b"not-an-image")
    attachment.write_bytes(b"attachment")
    profile = EntityProfile(
        asset_paths={
            "logo": str(wrong_logo),
            "evidence": str(attachment),
        }
    )
    role_specs = (
        SimpleNamespace(
            role="logo",
            accepted_types=("image",),
            required=True,
            min_items=1,
            max_items=1,
            cardinality="single",
        ),
        SimpleNamespace(
            role="evidence",
            accepted_types=("docx", "pdf"),
            required=False,
            min_items=0,
            max_items=1,
            cardinality="single",
        ),
    )

    resolution = resolve_profile_assets(profile, role_specs=role_specs)

    assert resolution.items_for_role("logo") == []
    assert [item.path for item in resolution.items_for_role("evidence")] == [
        str(attachment)
    ]
    assert any(
        item.code == "asset_file_type_unsupported"
        and item.role == "logo"
        for item in resolution.diagnostics
    )


def test_image_planning_rejects_non_image_items_and_marks_required_role_missing():
    wrong = AssetItem(role="logo", path="logo.docx")
    valid = AssetItem(role="logo", path="logo.png")
    rule = AssetInsertionRule(
        asset_role="logo",
        target="{{logo}}",
        required=True,
    )

    assert build_image_insertions([wrong], [rule]) == []
    assert missing_required_asset_roles([wrong], [rule]) == ["logo"]
    assert [item.path for item in build_image_insertions([wrong, valid], [rule])] == [
        "logo.png"
    ]


def test_image_insertion_rejects_non_image_bytes_and_supports_webp(tmp_path):
    wrong = tmp_path / "wrong.png"
    wrong.write_bytes(b"not-an-image")
    webp = tmp_path / "valid.webp"
    Image.new("RGB", (24, 24), color="red").save(webp)
    module = ImageInsertionModule()
    wrong_config = ResolvedConfig(
        images=[ImageInsertionItem(path=str(wrong), position="end", width_cm=2.0)]
    )
    webp_config = ResolvedConfig(
        images=[ImageInsertionItem(path=str(webp), position="end", width_cm=2.0)]
    )

    wrong_document = Document()
    issues = module.validate(wrong_document, wrong_config, PipelineContext())
    module.apply(
        wrong_document,
        wrong_config,
        ChangeTracker(),
        PipelineContext(),
    )
    webp_document = Document()
    webp_context = PipelineContext()
    assert module.validate(webp_document, webp_config, webp_context) == []
    module.apply(webp_document, webp_config, ChangeTracker(), webp_context)

    assert len(issues) == 1
    assert issues[0].level == "error"
    assert "不是可读取的图片" in issues[0].message
    assert len(wrong_document.inline_shapes) == 0
    assert len(webp_document.inline_shapes) == 1
    assert webp_context.inserted_images[0]["path"] == str(webp)


def test_profile_resolution_rejects_multiple_source_owners(tmp_path):
    selected = _image(tmp_path / "selected.png", b"selected")
    legacy = _image(tmp_path / "legacy.png", b"legacy")
    profile = EntityProfile(
        asset_bindings={
            "logo": AssetBinding(role="logo", source_path=str(selected)),
        },
        asset_paths={"logo": str(legacy)},
        asset_items=[
            {"item_id": "explicit", "role": "logo", "path": str(legacy)}
        ],
    )

    resolution = resolve_profile_assets(profile)

    assert resolution.items_for_role("logo") == []
    assert "logo" not in resolution.owners
    conflict = next(
        item for item in resolution.diagnostics if item.code == "asset_source_conflict"
    )
    assert conflict.severity == "error"
    assert "asset_bindings, asset_items, asset_paths" in conflict.message
    assert not resolution.ok


def test_profile_resolution_records_one_owner_per_role(tmp_path):
    logo = _image(tmp_path / "logo.png", b"logo")
    seal = _image(tmp_path / "seal.png", b"seal")
    figure = _image(tmp_path / "figure.png", b"figure")
    collection = tmp_path / "collection"
    cover = _image(collection / "cover.png", b"cover")
    profile = EntityProfile(
        asset_bindings={
            "logo": AssetBinding(role="logo", source_path=str(logo)),
        },
        asset_paths={"seal": str(seal)},
        asset_items=[
            {"item_id": "figure", "role": "figure", "path": str(figure)}
        ],
        assets_dir=str(collection),
    )

    resolution = resolve_profile_assets(profile)

    assert resolution.ok
    assert resolution.owners == {
        "cover": "assets_dir",
        "figure": "asset_items",
        "logo": "asset_bindings",
        "seal": "asset_paths",
    }
    assert {item.path for item in resolution.items} == {
        str(cover),
        str(figure),
        str(logo),
        str(seal),
    }


@pytest.mark.parametrize("source_name", ["asset_items", "asset_paths"])
def test_profile_resolution_does_not_emit_missing_source_items(tmp_path, source_name):
    missing = tmp_path / "missing.png"
    kwargs = (
        {"asset_items": [{"role": "logo", "path": str(missing)}]}
        if source_name == "asset_items"
        else {"asset_paths": {"logo": str(missing)}}
    )

    resolution = resolve_profile_assets(EntityProfile(**kwargs))

    assert resolution.items_for_role("logo") == []
    assert resolution.owners == {"logo": source_name}
    assert any(
        item.code == "asset_file_missing" and item.role == "logo"
        for item in resolution.diagnostics
    )
    assert not resolution.ok


def test_profile_resolution_rejects_duplicate_aliases_in_one_mapping(tmp_path):
    first = _image(tmp_path / "first.png", b"first")
    second = _image(tmp_path / "second.png", b"second")

    resolution = resolve_profile_assets(
        EntityProfile(
            asset_paths={
                "logo": str(first),
                "徽标": str(second),
            }
        )
    )

    assert resolution.items_for_role("logo") == []
    assert "logo" not in resolution.owners
    assert any(
        item.code == "asset_source_role_duplicate" and item.role == "logo"
        for item in resolution.diagnostics
    )


def test_material_batch_rejects_profile_and_base_context_role_conflict(tmp_path):
    profile_logo = _image(tmp_path / "profile-logo.png", b"profile")
    shared_logo = _image(tmp_path / "shared-logo.png", b"shared")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                asset_paths={"logo": str(profile_logo)},
            )
        ],
    )
    base_context = MaterialExecutionContext(
        asset_items=[AssetItem(role="logo", path=str(shared_logo))]
    )

    item = build_material_batch_items(archive, base_context=base_context)[0]

    assert item.context.asset_items == []
    conflict = next(
        diagnostic
        for diagnostic in item.context.asset_diagnostics
        if diagnostic["code"] == "asset_source_conflict"
    )
    assert conflict["severity"] == "error"
    assert conflict["role"] == "logo"


def test_v1_package_is_rejected_without_legacy_migration(tmp_path):
    source = tmp_path / "legacy.json"
    source.write_text(
        json.dumps(
            {
                "kind": "alavette.material_package",
                "version": 1,
                "archive_id": "legacy",
                "profiles": [
                    {
                        "profile_id": "main",
                        "profile_name": "Main",
                        "fields": {},
                        "asset_paths": {"logo": "logo.png"},
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="material_package_version_unsupported:1"):
        load_entity_archive(source)


def test_v1_attachment_package_is_rejected_without_shape_guessing(tmp_path):
    attachment = tmp_path / "evidence.docx"
    attachment.write_bytes(b"attachment")
    source = tmp_path / "legacy_attachment.json"
    source.write_text(
        json.dumps(
            {
                "kind": "alavette.material_package",
                "version": 1,
                "archive_id": "legacy-attachment",
                "profiles": [
                    {
                        "profile_id": "main",
                        "profile_name": "Main",
                        "fields": {},
                        "asset_paths": {"evidence": str(attachment)},
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="material_package_version_unsupported:1"):
        load_entity_archive(source)


def test_bundle_copies_group_with_normalized_names_and_survives_move(tmp_path):
    source = tmp_path / "source"
    _image(source / "二级" / "10.jpg", b"ten")
    _image(source / "一级" / "2.png", b"two")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                asset_bindings={
                    "qualification": AssetBinding(
                        role="qualification",
                        cardinality="multiple",
                        source_kind="directory",
                        source_path=str(source),
                        recursive=True,
                        min_items=1,
                        max_items=None,
                    )
                },
            )
        ],
    )

    package_dir = tmp_path / "package"
    package_dir.mkdir()
    stale_marker = package_dir / "stale.txt"
    stale_marker.write_text("stale", encoding="utf-8")
    package_path = save_entity_archive_bundle(archive, package_dir)
    assert not stale_marker.exists()
    first_loaded = load_entity_archive(package_path)
    first_items = resolve_profile_assets(first_loaded.profiles[0]).items
    package_path = save_entity_archive_bundle(archive, package_dir)
    second_loaded = load_entity_archive(package_path)
    second_items = resolve_profile_assets(second_loaded.profiles[0]).items
    assert [
        (item.item_id, item.sequence, item.normalized_name, item.original_relative_path)
        for item in first_items
    ] == [
        (item.item_id, item.sequence, item.normalized_name, item.original_relative_path)
        for item in second_items
    ]
    moved_dir = tmp_path / "moved"
    package_dir.rename(moved_dir)
    loaded = load_entity_archive(moved_dir / package_path.name)
    binding = loaded.profiles[0].asset_bindings["qualification"]
    resolution = resolve_profile_assets(loaded.profiles[0])

    assert Path(binding.source_path).is_dir()
    assert sorted(path.name for path in Path(binding.source_path).iterdir()) == [
        "qualification_001.png",
        "qualification_002.jpg",
    ]
    assert [item.normalized_name for item in resolution.items] == [
        "qualification_001.png",
        "qualification_002.jpg",
    ]
    assert [item.original_relative_path for item in resolution.items] == [
        "一级/2.png",
        "二级/10.jpg",
    ]
    assert all(item.content_hash.startswith("sha256:") for item in resolution.items)
    assert all(Path(item.path).is_file() for item in resolution.items)
    assert (source / "一级" / "2.png").is_file()
    assert (source / "二级" / "10.jpg").is_file()


def test_bundle_collection_revision_uses_canonical_owner_and_blocks_tampering(
    tmp_path,
):
    collection = tmp_path / "collection"
    _image(collection / "nested" / "one.png", b"one")
    _image(collection / "two.png", b"two")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                assets_dir=str(collection),
            )
        ],
    )

    manifest = save_entity_archive_bundle(archive, tmp_path / "package")
    loaded = load_entity_archive(manifest)
    profile = loaded.profiles[0]
    expected = profile.asset_metadata["__collection__"]["sha256"]

    assert expected == directory_content_revision(Path(profile.assets_dir))
    assert check_material_batch_preflight(loaded).ok is True

    (Path(profile.assets_dir) / "nested" / "one.png").write_bytes(b"tampered")
    result = check_material_batch_preflight(loaded)

    assert result.ok is False
    assert "asset_hash_mismatch:main:__collection__" in result.issues


def test_bundle_asset_item_content_hash_blocks_tampering(tmp_path):
    source = _image(tmp_path / "source.png", b"original")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                asset_items=[
                    {
                        "item_id": "logo:001",
                        "role": "logo",
                        "path": str(source),
                    }
                ],
            )
        ],
    )

    manifest = save_entity_archive_bundle(archive, tmp_path / "package")
    loaded = load_entity_archive(manifest)
    item = loaded.profiles[0].asset_items[0]

    assert item["content_hash"].startswith("sha256:")
    assert "sha256" not in item
    assert check_material_batch_preflight(loaded).ok is True

    Path(item["path"]).write_bytes(b"tampered")
    result = check_material_batch_preflight(loaded)

    assert result.ok is False
    assert "asset_hash_mismatch:main:logo:001" in result.issues


def test_bundle_export_rolls_back_existing_directory_on_failure(tmp_path, monkeypatch):
    package_dir = tmp_path / "package"
    package_dir.mkdir()
    marker = package_dir / "existing.txt"
    marker.write_text("keep", encoding="utf-8")
    source = tmp_path / "logo.png"
    source.write_bytes(b"image")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                asset_bindings={
                    "logo": AssetBinding(role="logo", source_path=str(source))
                },
            )
        ],
    )

    def fail_bundle(*_args, **_kwargs):
        raise RuntimeError("copy failed")

    monkeypatch.setattr(
        entity_bundle_module,
        "_bundle_asset_binding",
        fail_bundle,
    )
    with pytest.raises(RuntimeError, match="copy failed"):
        save_entity_archive_bundle(archive, package_dir)

    assert marker.read_text(encoding="utf-8") == "keep"
    assert not list(tmp_path.glob(".package.staging-*"))
    assert not list(tmp_path.glob(".package.backup-*"))


def test_bundle_export_fails_closed_on_missing_asset_binding(tmp_path):
    package_dir = tmp_path / "package"
    package_dir.mkdir()
    marker = package_dir / "existing.txt"
    marker.write_text("keep", encoding="utf-8")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                asset_bindings={
                    "logo": AssetBinding(
                        role="logo",
                        source_path=str(tmp_path / "missing.png"),
                    )
                },
            )
        ],
    )

    with pytest.raises(
        ValueError,
        match="entity_bundle_asset_binding_invalid:logo:asset_file_missing",
    ):
        save_entity_archive_bundle(archive, package_dir)

    assert marker.read_text(encoding="utf-8") == "keep"
    assert not list(tmp_path.glob(".package.staging-*"))


def test_bundle_export_rejects_legacy_missing_asset_path(tmp_path):
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                asset_paths={"logo": str(tmp_path / "missing.png")},
            )
        ],
    )

    with pytest.raises(ValueError, match="entity_bundle_asset_file_missing"):
        save_entity_archive_bundle(archive, tmp_path / "package")


def test_bundle_export_rejects_symlinked_collection_file(tmp_path):
    collection = tmp_path / "collection"
    collection.mkdir()
    outside = tmp_path / "outside.png"
    outside.write_bytes(b"outside")
    linked = collection / "linked.png"
    try:
        linked.symlink_to(outside)
    except OSError:
        pytest.skip("symbolic links are unavailable in this Windows environment")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                assets_dir=str(collection),
            )
        ],
    )

    with pytest.raises(ValueError, match="entity_bundle_link_forbidden"):
        save_entity_archive_bundle(archive, tmp_path / "package")

    assert not (tmp_path / "package").exists()


def test_bundle_copy_detects_source_mutation_and_uses_streaming_comparison(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "logo.png"
    source.write_bytes(b"before")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                asset_paths={"logo": str(source)},
            )
        ],
    )
    original_copy = entity_bundle_module.copy2

    def mutate_after_copy(copy_source, target):
        result = original_copy(copy_source, target)
        Path(copy_source).write_bytes(b"after!")
        return result

    def fail_read_bytes(*_args):
        raise AssertionError("unbounded read_bytes")

    monkeypatch.setattr(entity_bundle_module.Path, "read_bytes", fail_read_bytes)
    monkeypatch.setattr(entity_bundle_module, "copy2", mutate_after_copy)

    with pytest.raises(ValueError, match="entity_bundle_source_changed"):
        save_entity_archive_bundle(archive, tmp_path / "package")

    assert not (tmp_path / "package").exists()


def test_bundle_export_restores_existing_directory_when_publish_rename_fails(
    tmp_path,
    monkeypatch,
):
    package_dir = tmp_path / "package"
    package_dir.mkdir()
    marker = package_dir / "existing.txt"
    marker.write_text("keep", encoding="utf-8")
    source = tmp_path / "logo.png"
    source.write_bytes(b"image")
    archive = EntityArchive(
        archive_id="assets",
        profiles=[
            EntityProfile(
                profile_id="main",
                asset_bindings={
                    "logo": AssetBinding(role="logo", source_path=str(source))
                },
            )
        ],
    )
    original_rename = entity_bundle_module.Path.rename

    def fail_staging_publish(path, target):
        if path.name.startswith(".package.staging-") and Path(target) == package_dir:
            raise OSError("injected bundle publication failure")
        return original_rename(path, target)

    monkeypatch.setattr(entity_bundle_module.Path, "rename", fail_staging_publish)

    with pytest.raises(OSError, match="bundle publication failure"):
        save_entity_archive_bundle(archive, package_dir)

    assert marker.read_text(encoding="utf-8") == "keep"
    assert not list(tmp_path.glob(".package.staging-*"))
    assert not list(tmp_path.glob(".package.backup-*"))


def test_image_insertion_keeps_group_together_at_one_anchor(tmp_path):
    first = tmp_path / "qualification_001.png"
    second = tmp_path / "qualification_002.png"
    Image.new("RGB", (40, 40), color="red").save(first)
    Image.new("RGB", (40, 40), color="blue").save(second)
    document = Document()
    document.add_paragraph("before")
    document.add_paragraph("{{qualification}}")
    document.add_paragraph("after")
    context = PipelineContext()
    config = ResolvedConfig(
        images=[
            ImageInsertionItem(
                path=str(first),
                position="{{qualification}}",
                width_cm=2.0,
                role="qualification",
                item_id="qualification:001",
                group_id="qualification",
                sequence=1,
                normalized_name=first.name,
            ),
            ImageInsertionItem(
                path=str(second),
                position="{{qualification}}",
                width_cm=2.0,
                role="qualification",
                item_id="qualification:002",
                group_id="qualification",
                sequence=2,
                normalized_name=second.name,
            ),
        ]
    )

    ImageInsertionModule().apply(document, config, ChangeTracker(), context)

    assert len(document.inline_shapes) == 2
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "before",
        "",
        "",
        "after",
    ]
    assert [item["sequence"] for item in context.inserted_images] == [1, 2]
    assert [item["normalized_name"] for item in context.inserted_images] == [
        first.name,
        second.name,
    ]


def test_material_preflight_reports_multi_image_contract_violations(tmp_path):
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "product_assets_v1"
    scene.input_source_profile.failure_policy = "warn"
    duplicate = tmp_path / "same.png"
    duplicate.write_bytes(b"image")
    context = MaterialExecutionContext(
        entity_data={"product_name": "Product", "product_version": "1.0"},
        asset_items=[
            AssetItem(
                item_id="product:001",
                role="product_image",
                path=str(duplicate),
                group_id="product_image",
                sequence=1,
                normalized_name="product_image_001.png",
            ),
            AssetItem(
                item_id="product:002",
                role="product_image",
                path=str(duplicate),
                group_id="product_image",
                sequence=1,
                normalized_name="product_image_001.png",
            ),
        ],
    )

    diagnostics = material_requirement_diagnostics(scene, context)
    change_types = {item["change_type"] for item in diagnostics}

    assert "preflight_asset_group_sequence_conflict" in change_types
    assert "preflight_asset_normalized_name_conflict" in change_types


def test_non_exam_multi_image_group_runs_through_word_and_manifest(tmp_path):
    first = tmp_path / "product_image_001.png"
    second = tmp_path / "product_image_002.png"
    Image.new("RGB", (40, 40), color="red").save(first)
    Image.new("RGB", (40, 40), color="blue").save(second)
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Product gallery")
    document.add_paragraph("{{product_image}}")
    document.add_paragraph("End")
    document.save(source)
    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["image_insertion"] = True
    scene.input_source_profile.material_schema_id = "product_assets_v1"
    scene.input_source_profile.failure_policy = "warn"
    scene.default_delivery_preset().artifacts.material_manifest = True
    items = [
        AssetItem(
            item_id="product_image:001",
            role="product_image",
            path=str(first),
            group_id="product_image",
            sequence=1,
            original_relative_path="gallery/1.png",
            normalized_name=first.name,
        ),
        AssetItem(
            item_id="product_image:002",
            role="product_image",
            path=str(second),
            group_id="product_image",
            sequence=2,
            original_relative_path="gallery/2.png",
            normalized_name=second.name,
        ),
    ]
    result = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            material_schema_ids=("product_assets_v1",),
            entity_data={"product_name": "Product"},
            asset_items=items,
            image_rules=[
                AssetInsertionRule(
                    rule_id="product_image",
                    asset_role="product_image",
                    target="{{product_image}}",
                    cardinality="multiple",
                    min_items=1,
                    max_items=None,
                )
            ],
        ),
    ).run(lambda *_args: None, lambda: False)

    output = Document(result["output_path"])
    manifest_path = Path(result["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    role = next(item for item in manifest["asset_roles"] if item["role"] == "product_image")

    assert result["status"] == "success"
    assert len(output.inline_shapes) == 2
    assert "{{product_image}}" not in "\n".join(paragraph.text for paragraph in output.paragraphs)
    assert role["cardinality"] == "multiple"
    assert role["item_count"] == 2
    assert [item["sequence"] for item in role["items"]] == [1, 2]
    assert [item["normalized_name"] for item in role["items"]] == [
        first.name,
        second.name,
    ]
    assert manifest["image_rules"][0]["cardinality"] == "multiple"


def test_image_anchor_preflight_reports_missing_and_ambiguous_targets(tmp_path):
    image = tmp_path / "product.png"
    Image.new("RGB", (20, 20), color="red").save(image)
    context = MaterialExecutionContext(
        asset_items=[AssetItem(role="product_image", path=str(image))],
        image_rules=[
            AssetInsertionRule(
                asset_role="product_image",
                target="{{product_image}}",
                cardinality="multiple",
                max_items=None,
            )
        ],
    )
    scene = SceneWorkspace()
    scene.module_switches["image_insertion"] = True
    scene.input_source_profile.failure_policy = "warn"
    missing_doc = tmp_path / "missing.docx"
    Document().save(missing_doc)
    ambiguous_doc = tmp_path / "ambiguous.docx"
    ambiguous = Document()
    ambiguous.add_paragraph("{{product_image}}")
    ambiguous.add_paragraph("{{product_image}}")
    ambiguous.save(ambiguous_doc)

    missing = image_anchor_diagnostics(missing_doc, scene, context)
    ambiguous_result = image_anchor_diagnostics(ambiguous_doc, scene, context)

    assert missing[0]["change_type"] == "preflight_asset_anchor_missing"
    assert missing[0]["repair_target_type"] == "image_anchor"
    assert ambiguous_result[0]["change_type"] == "preflight_asset_anchor_ambiguous"
    assert ambiguous_result[0]["occurrence_count"] == 2


def test_image_token_preflight_blocks_selected_assets_without_binding_rules(tmp_path):
    target = tmp_path / "unbound.docx"
    document = Document()
    document.add_paragraph("Company logo")
    document.add_paragraph("{{@img:LOGO1}}")
    document.save(target)
    scene = SceneWorkspace()
    scene.module_switches["image_insertion"] = True
    scene.input_source_profile.failure_policy = "block"
    context = MaterialExecutionContext(
        asset_items=[
            AssetItem(role="logo", path=str(tmp_path / "logo.png")),
        ],
    )

    diagnostics = undeclared_image_token_diagnostics(target, scene, context)

    assert len(diagnostics) == 1
    assert diagnostics[0]["change_type"] == "preflight_image_rule_missing"
    assert diagnostics[0]["target"] == "{{@img:LOGO1}}"
    assert diagnostics[0]["level"] == "error"


def test_batch_import_recognizes_multi_image_folder_columns(tmp_path):
    source = tmp_path / "qualifications"
    source.mkdir()
    (source / "1.jpg").write_bytes(b"image")

    profile = _profile_from_mapping(
        {
            "profile_id": "one",
            "qualificationFolder": str(source),
            "productImageDirectory": str(source),
        },
        index=1,
    )

    assert set(profile.asset_bindings) == {"qualification", "product_image"}
    assert profile.asset_bindings["qualification"].source_kind == "directory"
    assert profile.asset_bindings["qualification"].recursive is True
    assert "qualificationFolder" not in profile.fields
    assert "productImageDirectory" not in profile.fields


def test_directory_resolution_diagnostics_reach_execution_preflight(tmp_path):
    archive = EntityArchive(
        archive_id="products",
        profiles=[
            EntityProfile(
                profile_id="one",
                fields={"product_name": "Product"},
                asset_bindings={
                    "product_image": AssetBinding(
                        role="product_image",
                        cardinality="multiple",
                        source_kind="directory",
                        source_path=str(tmp_path / "missing"),
                        recursive=True,
                        min_items=1,
                        max_items=None,
                    )
                },
            )
        ],
    )
    base_context = MaterialExecutionContext(
        material_schema_ids=("product_assets_v1",),
        image_rules=[
            AssetInsertionRule(
                asset_role="product_image",
                target="{{product_image}}",
                required=True,
                cardinality="multiple",
                min_items=1,
                max_items=None,
            )
        ],
    )
    item = build_material_batch_items(archive, base_context=base_context)[0]
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "product_assets_v1"
    scene.input_source_profile.failure_policy = "warn"

    diagnostics = material_requirement_diagnostics(scene, item.context)

    assert any(
        diagnostic["change_type"] == "preflight_asset_directory_missing"
        for diagnostic in diagnostics
    )
    assert any(
        diagnostic["repair_target_type"] == "asset_binding"
        for diagnostic in diagnostics
    )
