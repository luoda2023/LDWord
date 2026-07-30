from __future__ import annotations

from pathlib import Path

from docx import Document
from PIL import Image

from src.config.content_artifacts import ContentMaterialBinding
from src.config.content_materials import ContentInsertionRule, content_anchor_token
from src.config.entity import EntityProfile
from src.config.material_snapshot import MaterialSnapshot
from src.services.material_execution.material_snapshot_builder import (
    MaterialSnapshotBuildRequest,
    MaterialSnapshotBuilder,
)
from src.services.material_content.artifact_repository import (
    ContentArtifactRepository,
)
from src.services.material_content.compiler import compile_content_material
from src.services.material_content.composer import (
    ContentComposeRequest,
    ContentMaterialComposer,
)


def test_composer_consumes_only_artifacts_after_external_sources_are_deleted(
    tmp_path: Path,
) -> None:
    material_dir = tmp_path / "material"
    material_dir.mkdir()
    image = material_dir / "diagram.png"
    Image.new("RGB", (20, 10), (30, 90, 150)).save(image, format="PNG")
    markdown = material_dir / "content.md"
    markdown.write_text(
        "# 技术路线\n\n正文内容\n\n![图](diagram.png)\n",
        encoding="utf-8",
    )
    repository = ContentArtifactRepository(tmp_path / "content-artifacts")
    compiled = compile_content_material(markdown, repository)
    assert compiled.blocked is False
    binding = ContentMaterialBinding("route", "技术路线", compiled.artifact_ref)
    rule = ContentInsertionRule(
        "content:route",
        "route",
        content_anchor_token("route"),
    )
    snapshot = MaterialSnapshot(
        content_bindings=(binding,),
        content_rules=(rule,),
        material_schema_id="test-schema",
        material_schema_version="1",
        rule_versions={"content": "content-ir-v2"},
    )
    markdown.unlink()
    image.unlink()
    target = tmp_path / "target.docx"
    document = Document()
    document.add_paragraph(rule.anchor_token)
    document.save(target)
    output = tmp_path / "output.docx"

    receipt = ContentMaterialComposer(repository=repository).compose(
        ContentComposeRequest(str(target), str(output), snapshot)
    )

    assert output.is_file()
    assert receipt.content_artifacts[0].artifact_id == binding.artifact_ref.artifact_id
    assert receipt.content_artifacts[0].manifest_sha256 == binding.artifact_ref.manifest_sha256
    assert len(receipt.resource_materialization.entries) == 1
    assert Path(
        receipt.resource_materialization.entries[0].source_path
    ).is_file()
    reopened = Document(output)
    assert [paragraph.text for paragraph in reopened.paragraphs[:2]] == [
        "技术路线",
        "正文内容",
    ]


def test_snapshot_builder_validates_artifact_without_freezing_source_paths(
    tmp_path: Path,
) -> None:
    markdown = tmp_path / "source.md"
    markdown.write_text("正文", encoding="utf-8")
    repository = ContentArtifactRepository(tmp_path / "content-artifacts")
    compiled = compile_content_material(markdown, repository)
    binding = ContentMaterialBinding("body", "正文", compiled.artifact_ref)
    rule = ContentInsertionRule(
        "content:body",
        "body",
        content_anchor_token("body"),
    )
    profile = EntityProfile(
        profile_id="profile",
        profile_name="测试",
        content_bindings={"body": binding},
        content_rules=[rule],
    )
    markdown.unlink()

    snapshot = MaterialSnapshotBuilder().build_or_raise(
        MaterialSnapshotBuildRequest(
            profile=profile,
            frozen_field_values={},
            material_schema_id="test-schema",
            material_schema_version="1",
            rule_versions={"content": "content-ir-v2"},
            content_artifact_root=repository.root,
        )
    )

    assert snapshot.content_bindings == (binding,)
    assert not any(key.startswith("content:") for key in snapshot.source_revisions)
    assert str(tmp_path) not in snapshot.canonical_json()
