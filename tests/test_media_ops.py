from docx import Document
from PIL import Image

from src.shared.engine.media_ops import paragraph_has_image


def test_paragraph_has_image_detects_inline_media_for_image_rules(tmp_path):
    image_path = tmp_path / "pixel.png"
    Image.new("RGB", (2, 2), "white").save(image_path)
    document = Document()
    plain = document.add_paragraph("plain")
    illustrated = document.add_paragraph()
    illustrated.add_run().add_picture(str(image_path))

    assert paragraph_has_image(plain) is False
    assert paragraph_has_image(illustrated) is True
