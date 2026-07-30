import json
from types import SimpleNamespace

from docx import Document
from docx.shared import Cm

from src.config.execution_config_integrity import execution_config_integrity_issues
from src.config.journal_rule_source_registry import (
    audit_journal_rule_source_registry,
    get_journal_rule_source,
    journal_rule_source_for_count_profile,
)
from src.config.resolved import ResolvedConfig
from src.config.scene import ComplianceProfile, DeliveryPreset, InputSourceProfile
from src.modules.basic.page_setup import PageSetupModule
from src.pipeline.context import PipelineContext
from src.pipeline.runner import Pipeline
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.journal_rule_source_governance import (
    JournalRuleSourceIssue,
    inspect_journal_rule_source_governance,
    journal_rule_source_governance_execution_issues,
)
import src.shared.engine.journal_rule_source_governance as journal_governance
from src.services.production_runtime.delivery_reporting import pipeline_context_payload


def _journal_config(
    entity_data: dict[str, object] | None = None,
    *,
    count_profile_id: str = "journal_words",
) -> ResolvedConfig:
    return ResolvedConfig(
        delivery_presets=[DeliveryPreset()],
        default_delivery_preset_id="final",
        input_source_profile=InputSourceProfile(
            accepted_formats=["docx", "bibtex", "csl_json"],
            structured_formats=["bibtex", "csl_json"],
            material_schema_id="journal_submission_materials_v1",
        ),
        compliance_profile=ComplianceProfile(
            profile_id="journal_submission",
            rule_family="journal_en",
            count_profile_id=count_profile_id,
        ),
        entity_data=dict(entity_data or {}),
    )


def test_journal_rule_source_registry_links_reviewed_count_profiles():
    assert audit_journal_rule_source_registry() == ()

    source = get_journal_rule_source("journal_en_default_rules")

    assert source.review_status == "reviewed_generic"
    assert source.is_reviewed is True
    assert source.source_reference == "count_profiles/builtin.json"
    assert {
        "journal_words",
        "journal_display_items",
    } <= set(source.count_profile_ids)
    assert journal_rule_source_for_count_profile("journal_words") == source
    assert journal_rule_source_for_count_profile("") is None


def test_journal_rule_source_governance_reports_reviewed_generic_default():
    result = inspect_journal_rule_source_governance(
        _journal_config({"journal_name": "Example Journal"})
    )

    assert result.status == "ok"
    assert result.family_id == "journal_en"
    assert result.rule_source_id == "journal_en_default_rules"
    assert result.count_profile_id == "journal_words"
    assert result.review_status == "reviewed_generic"
    assert result.summary.reviewed_source_count == 1
    assert result.summary.matched_count_profile_count == 1
    assert result.summary.target_journal_declared is True
    assert result.manual_confirmation_required is True
    assert any("publisher-specific" in note for note in result.boundary_notes)
    assert journal_rule_source_governance_execution_issues(
        _journal_config({"journal_name": "Example Journal"})
    ) == ()


def test_journal_rule_source_governance_rejects_unreviewed_count_profile():
    result = inspect_journal_rule_source_governance(
        _journal_config(count_profile_id="journal_unreviewed_runtime")
    )

    assert result.status == "error"
    assert result.error_count == 1
    assert result.issues[0].kind == "unregistered_count_profile"
    assert result.manual_confirmation_required is True


def test_journal_rule_source_governance_rejects_missing_count_profile_without_default():
    config = _journal_config(count_profile_id="")
    result = inspect_journal_rule_source_governance(config)
    expected = (
        "journal_rule_source_governance_error:count_profile_missing:-"
    )

    assert result.status == "error"
    assert result.count_profile_id == ""
    assert result.rule_source_id == ""
    assert result.error_count == 1
    assert result.issues[0].kind == "count_profile_missing"
    assert journal_rule_source_governance_execution_issues(config) == (expected,)
    assert execution_config_integrity_issues(config) == (expected,)


def test_journal_rule_source_warning_remains_non_blocking(monkeypatch):
    warning = JournalRuleSourceIssue(
        path="journal_rule_sources.generic.boundary",
        kind="manual_review_recommended",
        message="Publisher-specific review is recommended.",
        severity="warning",
        observed="reviewed_generic",
    )
    monkeypatch.setattr(
        journal_governance,
        "inspect_journal_rule_source_governance",
        lambda config, *, entity_data=None: SimpleNamespace(issues=(warning,)),
    )

    assert journal_governance.journal_rule_source_governance_execution_issues(
        object()
    ) == ()


def test_unknown_declared_journal_rule_source_is_one_blocking_config_error():
    config = _journal_config(
        {"journal_rule_source_id": "definitely_unknown"}
    )

    result = inspect_journal_rule_source_governance(config)
    expected = (
        "journal_rule_source_governance_error:"
        "unregistered_rule_source:definitely_unknown"
    )

    assert result.status == "error"
    assert result.error_count == 1
    assert result.issues[0].kind == "unregistered_rule_source"
    assert journal_rule_source_governance_execution_issues(config) == (expected,)
    assert execution_config_integrity_issues(config) == (expected,)


def test_pipeline_blocks_journal_rule_source_error_before_docx_mutation(tmp_path):
    source = tmp_path / "journal-invalid-source.docx"
    document = Document()
    document.sections[0].top_margin = Cm(1)
    document.add_paragraph("SOURCE_MARKER")
    document.save(source)
    before = source.read_bytes()
    output_dir = tmp_path / "outputs"

    result = Pipeline(
        [PageSetupModule()],
        _journal_config(
            {"journal_rule_source_id": "definitely_unknown"}
        ),
        output_dir=str(output_dir),
    ).execute(str(source))

    assert result.success is False
    assert result.status == "failed"
    assert result.error == (
        "pipeline_configuration_invalid:"
        "journal_rule_source_governance_error:"
        "unregistered_rule_source:definitely_unknown"
    )
    assert result.context is None
    assert result.tracker is None
    assert source.read_bytes() == before
    assert not output_dir.exists()


def test_pipeline_blocks_missing_journal_count_profile_before_docx_mutation(tmp_path):
    source = tmp_path / "journal-count-profile-missing.docx"
    document = Document()
    document.sections[0].top_margin = Cm(1)
    document.add_paragraph("SOURCE_MARKER")
    document.save(source)
    before = source.read_bytes()
    output_dir = tmp_path / "outputs"

    result = Pipeline(
        [PageSetupModule()],
        _journal_config(count_profile_id=""),
        output_dir=str(output_dir),
    ).execute(str(source))

    assert result.success is False
    assert result.status == "failed"
    assert result.error == (
        "pipeline_configuration_invalid:"
        "journal_rule_source_governance_error:count_profile_missing:-"
    )
    assert result.context is None
    assert result.tracker is None
    assert source.read_bytes() == before
    assert not output_dir.exists()


def test_journal_rule_source_tracker_marks_error_as_failure_when_called_directly():
    pipeline = Pipeline(
        [],
        _journal_config(
            {"journal_rule_source_id": "definitely_unknown"}
        ),
    )
    context = PipelineContext()

    pipeline._run_journal_rule_source_governance(context)

    records = pipeline._tracker.get_by_module(
        "journal_rule_source_governance"
    )
    assert len(records) == 1
    assert records[0].change_type == "journal_rule_source_error"
    assert records[0].success is False
    assert "not registered as reviewed" in str(records[0].failure_reason)


def test_pipeline_records_journal_rule_source_governance_and_reports(tmp_path):
    source = tmp_path / "journal.docx"
    doc = Document()
    doc.add_paragraph(r"This manuscript cites \cite{smith2020}.")
    doc.save(source)
    bibtex = """
@article{smith2020,
  title={A study},
  author={Smith, Jane},
  year={2020}
}
"""
    result = Pipeline(
        [],
        _journal_config(
            {
                "journal_name": "Example Journal",
                "bibtex": bibtex,
            }
        ),
        output_dir=str(tmp_path),
    ).execute(str(source))

    assert result.success is True
    assert result.context is not None
    governance = result.context.journal_rule_source_governance
    assert governance.status == "ok"
    assert governance.rule_source_id == "journal_en_default_rules"
    payload = pipeline_context_payload(result.context)
    assert payload["journal_rule_source_governance"]["rule_source_id"] == (
        "journal_en_default_rules"
    )
    records = result.tracker.get_by_module("journal_rule_source_governance")
    assert records
    assert records[0].change_type == "journal_rule_source_governance"
    assert "review=reviewed_generic" in records[0].after
    assert "count_profile=journal_words" in records[0].after

    report_json = tmp_path / "journal_rule_source_changes.json"
    report_md = tmp_path / "journal_rule_source_changes.md"
    write_json_report(
        result,
        input_path=source,
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )

    data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert data["journal_rule_source_governance"]["rule_source_id"] == (
        "journal_en_default_rules"
    )
    assert data["journal_rule_source_governance"]["review_status"] == (
        "reviewed_generic"
    )
    assert data["journal_rule_source_governance"][
        "manual_confirmation_required"
    ] is True
    assert "## 英文期刊规则源治理" in markdown
    assert "journal_en_default_rules" in markdown
    assert "reviewed_generic" in markdown
