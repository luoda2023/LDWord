import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.config.template import StyleConfig
from src.modules.basic.paragraph_style import ParagraphStyleModule
from src.modules.table.caption import (
    CaptionInfo,
    CaptionModule,
    _apply_caption_style,
    _caption_prefix,
    _compose_caption_number_text,
    _parse_caption_numbering_format,
    _resolve_caption_style,
)
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn


def test_caption_style_prefers_kind_specific_style():
    config = ResolvedConfig()
    config.styles["body"] = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=12,
        alignment="justify",
    )
    config.styles["caption"] = StyleConfig(
        font_cn="Heiti",
        font_en="Arial",
        size_pt=10.5,
        alignment="center",
        line_spacing_type="exact",
        line_spacing_pt=16,
    )
    config.styles["figure_caption"] = StyleConfig(
        font_cn="FangSong",
        font_en="Calibri",
        size_pt=9,
        alignment="right",
    )

    fig_style = _resolve_caption_style(config, "figure")
    tbl_style = _resolve_caption_style(config, "table")

    assert fig_style.font_cn == "FangSong"
    assert fig_style.font_en == "Calibri"
    assert fig_style.size_pt == 9
    assert fig_style.alignment == "right"

    assert tbl_style.font_cn == "Heiti"
    assert tbl_style.font_en == "Arial"
    assert tbl_style.size_pt == 10.5
    assert tbl_style.alignment == "center"


def test_caption_style_body_fallback_clears_body_indent():
    config = ResolvedConfig()
    config.styles["body"] = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=12,
        alignment="justify",
        left_indent_chars=2,
        left_indent_unit="chars",
        special_indent_mode="first_line",
        special_indent_value=2,
        special_indent_unit="chars",
        first_line_indent_chars=2,
        first_line_indent_unit="chars",
    )

    style = _resolve_caption_style(config, "figure")

    assert style.alignment == "center"
    assert style.left_indent_chars == 0
    assert style.right_indent_chars == 0
    assert style.special_indent_mode == "none"
    assert style.special_indent_value == 0


def test_caption_style_syncs_spacing_indent_and_font():
    document = Document()
    para = document.add_paragraph("Figure1 caption")
    style = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=10.5,
        bold=True,
        italic=False,
        alignment="center",
        line_spacing_type="exact",
        line_spacing_pt=16,
        space_before_pt=0,
        space_after_pt=6,
    )

    _apply_caption_style(para, style)

    spacing = para._element.find(qn("w:pPr")).find(qn("w:spacing"))
    ind = para._element.find(qn("w:pPr")).find(qn("w:ind"))

    assert spacing.get(qn("w:after")) == "120"
    assert spacing.get(qn("w:line")) == "320"
    assert spacing.get(qn("w:lineRule")) == "exact"
    assert ind.get(qn("w:left")) == "0"
    assert ind.get(qn("w:firstLine")) == "0"
    assert para.runs[0].font.size.pt == 10.5
    assert para.runs[0].font.bold is True


def test_paragraph_style_preserves_kind_specific_caption_spacing():
    document = Document()
    figure = document.add_paragraph("Figure 7 Overview", style="Caption")
    table = document.add_paragraph("Table 3 Results", style="Caption")
    config = ResolvedConfig()
    config.caption.numbering_mode = "global"
    config.caption.separator = " "
    config.styles["caption"] = StyleConfig(
        size_pt=10.5,
        alignment="center",
        space_before_pt=6,
        space_after_pt=6,
    )
    config.styles["figure_caption"] = StyleConfig(
        size_pt=10.5,
        alignment="center",
        space_before_pt=6,
        space_after_pt=12,
    )
    config.styles["table_caption"] = StyleConfig(
        size_pt=10.5,
        alignment="center",
        space_before_pt=6,
        space_after_pt=6,
    )

    CaptionModule().apply(
        document,
        config,
        ChangeTracker(),
        PipelineContext(),
    )
    ParagraphStyleModule().apply(
        document,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    assert figure.paragraph_format.space_after.pt == 12
    assert table.paragraph_format.space_after.pt == 6


def test_caption_numbering_format_supports_configured_chapter_separator():
    assert _parse_caption_numbering_format("chapter.seq") == (True, ".")
    assert _parse_caption_numbering_format("chapter-seq") == (True, "-")
    assert _parse_caption_numbering_format("seq") == (False, "")

    assert _compose_caption_number_text(
        chapter_num=3,
        sequence_text="3.2",
        numbering_mode="chapter",
        numbering_format="chapter-seq",
    ) == "3-2"
    assert _compose_caption_number_text(
        chapter_num=3,
        sequence_text="2",
        numbering_mode="global",
        numbering_format="chapter.seq",
    ) == "2"


def test_caption_prefix_prefers_configured_prefixes():
    config = ResolvedConfig()
    config.caption.figure_prefix = "Figure"
    config.caption.table_prefix = "Table"

    figure = CaptionInfo(0, "figure", "Figure", "1", "Title")
    table = CaptionInfo(1, "table", "Table", "1", "Title")

    assert _caption_prefix(config.caption, figure) == "Figure"
    assert _caption_prefix(config.caption, table) == "Table"


def test_caption_module_inserts_missing_figure_caption_after_image_anchor():
    document = Document()
    document.add_paragraph("Intro")
    image_para = document.add_paragraph()
    etree.SubElement(image_para.add_run()._element, qn("w:drawing"))
    document.add_paragraph("Tail")

    config = ResolvedConfig()
    config.caption.figure_prefix = "F"
    config.caption.separator = " "
    config.caption.placeholder = "TODO"
    config.caption.numbering_mode = "global"
    config.caption.auto_insert = True
    config.styles["caption"] = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=10.5,
        alignment="center",
    )

    CaptionModule().apply(document, config, ChangeTracker(), PipelineContext())

    assert [para.text for para in document.paragraphs] == ["Intro", "", "F1 TODO", "Tail"]


def test_caption_module_inserts_missing_table_caption_before_table_anchor():
    document = Document()
    document.add_paragraph("Intro")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "value"
    document.add_paragraph("Tail")

    config = ResolvedConfig()
    config.caption.table_prefix = "T"
    config.caption.separator = " "
    config.caption.placeholder = "TODO"
    config.caption.numbering_mode = "global"
    config.caption.auto_insert = True
    config.styles["caption"] = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=10.5,
        alignment="center",
    )

    CaptionModule().apply(document, config, ChangeTracker(), PipelineContext())

    assert [para.text for para in document.paragraphs] == ["Intro", "T1 TODO", "Tail"]


def test_caption_module_respects_auto_insert_false_for_missing_figure_caption():
    document = Document()
    document.add_paragraph("Intro")
    image_para = document.add_paragraph()
    etree.SubElement(image_para.add_run()._element, qn("w:drawing"))
    document.add_paragraph("Tail")

    config = ResolvedConfig()
    config.caption.figure_prefix = "F"
    config.caption.separator = " "
    config.caption.placeholder = "TODO"
    config.caption.numbering_mode = "global"
    config.caption.auto_insert = False

    CaptionModule().apply(document, config, ChangeTracker(), PipelineContext())

    assert [para.text for para in document.paragraphs] == ["Intro", "", "Tail"]


def test_caption_module_formats_figure_caption_continuation_lines():
    document = Document()
    document.add_paragraph("Intro")
    image_para = document.add_paragraph()
    etree.SubElement(image_para.add_run()._element, qn("w:drawing"))
    document.add_paragraph("Figure 7 Overview")
    document.add_paragraph("(a) Left panel detail")
    document.add_paragraph("Tail")

    config = ResolvedConfig()
    config.caption.figure_prefix = "Figure"
    config.caption.separator = " "
    config.caption.numbering_mode = "global"
    config.styles["caption"] = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=10.5,
        alignment="center",
        line_spacing_type="exact",
        line_spacing_pt=16,
    )

    CaptionModule().apply(document, config, ChangeTracker(), PipelineContext())

    assert document.paragraphs[2].text == "Figure1 Overview"
    assert document.paragraphs[3].text == "(a) Left panel detail"
    assert document.paragraphs[3].paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert document.paragraphs[3].runs[0].font.size.pt == 10.5


def test_caption_module_inserts_field_numbering_when_enabled():
    document = Document()
    document.add_paragraph("Intro")
    image_para = document.add_paragraph()
    etree.SubElement(image_para.add_run()._element, qn("w:drawing"))
    document.add_paragraph("Tail")

    config = ResolvedConfig()
    config.caption.figure_prefix = "Figure"
    config.caption.separator = " "
    config.caption.placeholder = "TODO"
    config.caption.numbering_mode = "global"
    config.caption.format_inserted = True
    config.styles["caption"] = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=10.5,
        alignment="center",
    )

    CaptionModule().apply(document, config, ChangeTracker(), PipelineContext())

    inserted = document.paragraphs[2]
    instr_texts = [elem.text or "" for elem in inserted._element.findall(f".//{qn('w:instrText')}")]

    assert inserted.text.endswith(" TODO")
    assert any("SEQ Figure" in text for text in instr_texts)


def test_caption_module_rewrites_existing_caption_with_field_numbering_when_enabled():
    document = Document()
    document.add_paragraph("Intro")
    image_para = document.add_paragraph()
    etree.SubElement(image_para.add_run()._element, qn("w:drawing"))
    caption_para = document.add_paragraph("Figure 7 Overview")
    document.add_paragraph("Tail")

    config = ResolvedConfig()
    config.caption.figure_prefix = "Figure"
    config.caption.separator = " "
    config.caption.numbering_mode = "global"
    config.caption.format_inserted = True
    config.styles["caption"] = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=10.5,
        alignment="center",
    )

    CaptionModule().apply(document, config, ChangeTracker(), PipelineContext())

    instr_texts = [elem.text or "" for elem in caption_para._element.findall(f".//{qn('w:instrText')}")]

    assert caption_para.text.endswith(" Overview")
    assert any("SEQ Figure" in text for text in instr_texts)


def test_caption_module_skips_existing_caption_chapter_numbering_without_heading_context():
    document = Document()
    document.add_paragraph("Intro")
    image_para = document.add_paragraph()
    etree.SubElement(image_para.add_run()._element, qn("w:drawing"))
    caption_para = document.add_paragraph("Figure 7 Overview")
    document.add_paragraph("第一章 绪论")

    config = ResolvedConfig()
    config.caption.figure_prefix = "Figure"
    config.caption.separator = " "
    config.caption.numbering_mode = "chapter"
    config.caption.numbering_format = "chapter.seq"
    config.caption.format_inserted = True
    tracker = ChangeTracker()
    context = PipelineContext(heading_map={3: 1})

    CaptionModule().apply(document, config, tracker, context)

    instr_texts = [elem.text or "" for elem in caption_para._element.findall(f".//{qn('w:instrText')}")]

    assert caption_para.text == "Figure 7 Overview"
    assert instr_texts == []
    assert any(record.change_type == "skip" for record in tracker.get_by_module("caption"))


def test_caption_module_skips_auto_insert_in_chapter_mode_without_heading_context():
    document = Document()
    document.add_paragraph("Intro")
    image_para = document.add_paragraph()
    etree.SubElement(image_para.add_run()._element, qn("w:drawing"))
    document.add_paragraph("第一章 绪论")
    document.add_paragraph("Tail")

    config = ResolvedConfig()
    config.caption.figure_prefix = "Figure"
    config.caption.separator = " "
    config.caption.placeholder = "TODO"
    config.caption.numbering_mode = "chapter"
    config.caption.numbering_format = "chapter.seq"
    config.caption.auto_insert = True
    tracker = ChangeTracker()
    context = PipelineContext(heading_map={2: 1})

    CaptionModule().apply(document, config, tracker, context)

    assert [para.text for para in document.paragraphs] == ["Intro", "", "第一章 绪论", "Tail"]
    assert any(record.change_type == "skip" for record in tracker.get_by_module("caption"))
