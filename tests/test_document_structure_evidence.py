from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.services.document_structure_evidence import (
    RegionDecision,
    build_document_structure_evidence,
    build_document_structure_evidence_from_bytes,
    document_structure_evidence_is_current,
    validate_region_decisions,
)


def _save_structured_doc(path):
    doc = Document()
    doc.add_paragraph("摘要")
    doc.add_paragraph("摘要内容。")
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容。")
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[1] Author. Title. 2024.")
    doc.save(path)


def test_sparse_centered_first_section_is_accepted_as_generic_cover(tmp_path):
    source = tmp_path / "generic-cover.docx"
    doc = Document()
    for _index in range(3):
        doc.add_paragraph("")
    title = doc.add_paragraph("\u56db\u5ddd\u7701\u591a\u5f0f\u8054\u8fd0\u57fa\u5730\u5efa\u8bbe\u89c4\u5212")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _index in range(14):
        doc.add_paragraph("")
    authority = doc.add_paragraph("\u4e50\u5c71\u5e02\u4eba\u6c11\u653f\u5e9c")
    authority.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date = doc.add_paragraph("\u4e8c\u25cb\u4e8c\u516d\u5e74\u4e8c\u6708")
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    section_boundary = len(doc.paragraphs) + 1
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("\u7b2c\u4e00\u7ae0 \u6b63\u6587", level=1)
    for index in range(60):
        doc.add_paragraph(f"\u6b63\u6587\u5185\u5bb9 {index}\u3002")
    doc.save(source)

    evidence = build_document_structure_evidence(source)
    cover = next(region for region in evidence.regions if region.role_id == "cover")
    body = next(region for region in evidence.regions if region.role_id == "body")

    assert cover.detection_status == "accepted"
    assert cover.start_anchor.source_index == 0
    assert cover.end_anchor is not None
    assert cover.end_anchor.source_index == section_boundary
    assert body.start_anchor.source_index >= section_boundary
    assert not evidence.requires_review


def test_document_structure_evidence_is_stable_and_bound_to_source_revision(tmp_path):
    source = tmp_path / "structured.docx"
    _save_structured_doc(source)

    first = build_document_structure_evidence(source)
    second = build_document_structure_evidence(source)

    assert first.ready
    assert first.evidence_digest == second.evidence_digest
    assert first.source_revision == second.source_revision
    assert document_structure_evidence_is_current(first, source)
    assert {region.role_id for region in first.regions} >= {
        "abstract_cn",
        "body",
        "references",
    }
    assert all(region.start_anchor.text_digest for region in first.regions)

    changed = Document(str(source))
    changed.add_paragraph("新增内容")
    changed.save(source)

    assert not document_structure_evidence_is_current(first, source)


def test_content_only_reference_detection_requires_review(tmp_path):
    source = tmp_path / "reference-fallback.docx"
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    for index in range(8):
        doc.add_paragraph(f"正文内容 {index}。")
    doc.add_paragraph("[1] Author. First title. 2023.")
    doc.add_paragraph("[2] Author. Second title. 2024.")
    doc.save(source)

    evidence = build_document_structure_evidence(source)
    reference = next(
        region for region in evidence.regions if region.role_id == "references"
    )

    assert reference.detection_status == "review"
    assert any(
        item.role_id == "references" for item in evidence.review_items
    )


def test_region_decisions_are_closed_to_supported_roles_and_actions(tmp_path):
    source = tmp_path / "structured.docx"
    _save_structured_doc(source)
    evidence = build_document_structure_evidence(source)
    body = next(region for region in evidence.regions if region.role_id == "body")

    assert (
        validate_region_decisions(
            evidence,
            (
                RegionDecision("body", "set_start", body.start_anchor),
                RegionDecision("references", "exclude"),
            ),
            allowed_roles=("body", "references"),
        )
        == ()
    )
    assert validate_region_decisions(
        evidence,
        (RegionDecision("resume", "invent"),),
        allowed_roles=("body", "references"),
    ) == (
        "document_structure_decision_role_unsupported:resume",
        "document_structure_decision_action_invalid:invent",
    )


def test_invalid_source_returns_failed_evidence(tmp_path):
    evidence = build_document_structure_evidence(tmp_path / "missing.docx")

    assert not evidence.ready
    assert evidence.issues == ("document_structure_source_unreadable",)


def test_malformed_docx_returns_failed_evidence_instead_of_raising(tmp_path):
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not-a-docx")

    from_path = build_document_structure_evidence(source)
    from_bytes = build_document_structure_evidence_from_bytes(
        source,
        source.read_bytes(),
    )

    assert from_path.issues == ("document_structure_analysis_failed",)
    assert from_bytes.issues == ("document_structure_analysis_failed",)
