from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.shared.engine.docx_material_tokens import (
    is_strict_material_token_paragraph,
)


TOKEN = "{{@file:route}}"


def test_strict_material_token_paragraph_accepts_plain_split_runs_and_properties():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    paragraph.add_run("  ")
    paragraph.add_run(TOKEN[:8])
    paragraph.add_run(TOKEN[8:])
    paragraph.add_run("   ")

    # The service owners still inspect sectPr themselves so they can emit their
    # specific diagnostics.
    assert is_strict_material_token_paragraph(paragraph._p, TOKEN) is True


def test_strict_material_token_paragraph_rejects_non_plain_ooxml():
    document = Document()
    paragraph = document.add_paragraph(TOKEN)
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "1")
    bookmark.set(qn("w:name"), "unsafe")
    paragraph._p.append(bookmark)

    assert is_strict_material_token_paragraph(paragraph._p, TOKEN) is False


def test_strict_material_token_paragraph_rejects_mixed_or_duplicate_text():
    document = Document()
    mixed = document.add_paragraph("before " + TOKEN)
    duplicate = document.add_paragraph(TOKEN + TOKEN)

    assert is_strict_material_token_paragraph(mixed._p, TOKEN) is False
    assert is_strict_material_token_paragraph(duplicate._p, TOKEN) is False
