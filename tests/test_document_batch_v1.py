from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from src.application.materials import (
    ExecutionMaterialRecord,
    ExecutionResource,
    MaterialRunBindRequest,
    bind_material_run,
)
from src.document_batch import (
    DocumentBatchRequest,
    compile_document_batch_plan,
)
from src.document_batch import recipe as batch_recipe
from src.document_batch.material_resources import _insert_images
from src.domain.materials import (
    MaterialContract,
    MaterialFieldContract,
    MaterialObjectRef,
    MaterialPackage,
    MaterialPackageRef,
    MaterialRecord,
    MaterialResourceBinding,
    MaterialResourceRoleContract,
    MaterialRunSelection,
    MaterialScope,
    MaterialTimelineSpec,
    clone_material_package,
    clone_material_record,
    generate_package_id,
    generate_record_id,
)
from src.infrastructure.materials.codec import material_package_revision
from tests.test_material_application_v1 import _contract, _package


def test_document_batch_uses_frozen_snapshot_and_publishes_as_one_set(tmp_path):
    package = _package()
    revision = material_package_revision(package)
    selection = MaterialRunSelection(
        package_ref=MaterialPackageRef(package.package_id, revision),
        selected_record_ids=(package.records[0].record_id,),
    )
    bound = bind_material_run(
        MaterialRunBindRequest(
            package=package,
            package_revision=revision,
            source_type="builtin",
            selection=selection,
            contract=_contract(),
            recipe_id="document_batch",
            scene_id="official",
            work_mode_id="official",
            document_type="notice",
        ),
        object_path=lambda _item: "",
    )
    assert bound.ok and bound.snapshot is not None
    assert bound.snapshot.image_policy["show_single_image_name"] is False
    assert bound.snapshot.image_policy["show_multi_image_name"] is False
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("{{title}}")
    document.save(source)
    output = tmp_path / "output"
    plan = compile_document_batch_plan(
        bound.snapshot,
        DocumentBatchRequest((str(source),), str(output)),
    )

    result = plan.run()

    assert result["status"] == "success"
    created = next(output.rglob("*.docx"))
    assert created.parent == output
    assert Document(created).paragraphs[0].text == "关于测试的通知"


def test_document_batch_only_groups_outputs_when_multiple_records_are_selected(
    tmp_path,
):
    original = _package()
    second = clone_material_record(
        original.records[0],
        record_id=generate_record_id(),
        display_name="复核通知",
    )
    package = clone_material_package(
        original,
        records=(*original.records, second),
    )
    revision = material_package_revision(package)
    bound = bind_material_run(
        MaterialRunBindRequest(
            package=package,
            package_revision=revision,
            source_type="builtin",
            selection=MaterialRunSelection(
                package_ref=MaterialPackageRef(package.package_id, revision),
                selected_record_ids=tuple(
                    record.record_id for record in package.records
                ),
            ),
            contract=_contract(),
            recipe_id="document_batch",
            scene_id="official",
            work_mode_id="official",
            document_type="notice",
        ),
        object_path=lambda _item: "",
    )
    assert bound.ok and bound.snapshot is not None
    source = tmp_path / "source.docx"
    Document().save(source)
    output = tmp_path / "output"

    plan = compile_document_batch_plan(
        bound.snapshot,
        DocumentBatchRequest((str(source),), str(output)),
    )

    assert plan.ok
    assert {
        Path(item.final_path).parent.name for item in plan.artifacts
    } == {"通知", "复核通知"}


def test_document_batch_rolls_back_the_atomic_set_on_publish_failure(
    tmp_path,
    monkeypatch,
):
    package = _package()
    revision = material_package_revision(package)
    bound = bind_material_run(
        MaterialRunBindRequest(
            package=package,
            package_revision=revision,
            source_type="builtin",
            selection=MaterialRunSelection(
                package_ref=MaterialPackageRef(package.package_id, revision),
                selected_record_ids=(package.records[0].record_id,),
            ),
            contract=_contract(),
            recipe_id="document_batch",
            scene_id="official",
            work_mode_id="official",
            document_type="notice",
        ),
        object_path=lambda _item: "",
    )
    assert bound.ok and bound.snapshot is not None
    sources = []
    for name in ("first.docx", "second.docx"):
        source = tmp_path / name
        Document().save(source)
        sources.append(str(source))
    output = tmp_path / "output"
    plan = compile_document_batch_plan(
        bound.snapshot,
        DocumentBatchRequest(tuple(sources), str(output)),
    )
    assert plan.ok
    real_replace = batch_recipe.os.replace
    replace_count = 0

    def _fail_second_replace(source, destination):
        nonlocal replace_count
        replace_count += 1
        if replace_count == 2:
            raise OSError("simulated second publish failure")
        real_replace(source, destination)

    monkeypatch.setattr(batch_recipe.os, "replace", _fail_second_replace)

    result = plan.run()

    assert result["status"] == "failed"
    assert len(result["rolled_back_paths"]) == 1
    assert list(output.rglob("*.docx")) == []


def test_image_insertion_does_not_add_pagination_flags_to_neighboring_paragraphs(
    tmp_path,
):
    image = tmp_path / "photo.png"
    Image.new("RGB", (640, 960), "navy").save(image)
    payload = image.read_bytes()
    object_ref = MaterialObjectRef(
        object_id="sha256:" + sha256(payload).hexdigest(),
        media_type="image/png",
        original_name=image.name,
        size=len(payload),
    )
    record = ExecutionMaterialRecord(
        record_id="record-1",
        display_name="Record 1",
        group_id="group-1",
        field_values={},
        field_owners={},
        resources={
            "photo": (ExecutionResource(object_ref, str(image.resolve())),),
        },
        resource_owners={"photo": ("record",)},
    )
    source = tmp_path / "source.docx"
    document = Document()
    heading = document.add_paragraph("Image section")
    heading.paragraph_format.keep_with_next = True
    document.add_paragraph()
    document.add_paragraph("{{@img:photo}}")
    document.save(source)

    output = tmp_path / "output.docx"
    _insert_images(
        source,
        output,
        record=record,
        roles=("photo",),
        image_policy={"adaptive": True},
        cache_dir=tmp_path / "image-cache",
    )

    rendered = Document(output)
    assert rendered.paragraphs[0].paragraph_format.keep_with_next is True
    assert rendered.paragraphs[1].text == ""
    assert rendered.paragraphs[1].paragraph_format.keep_with_next is None
    assert rendered.paragraphs[2].paragraph_format.keep_together is None
    assert len(rendered.inline_shapes) == 1


def test_image_names_use_role_label_once_and_page_break_follows_group(
    tmp_path,
):
    resources = []
    expected_colors = [(128, 0, 128), (0, 128, 128), (0, 0, 128)]
    for name, color in (
        ("photo1.png", "purple"),
        ("photo2.png", "teal"),
        ("photo10.png", "navy"),
    ):
        path = tmp_path / name
        Image.new("RGB", (32, 32), color).save(path)
        payload = path.read_bytes()
        resources.append(
            ExecutionResource(
                MaterialObjectRef(
                    object_id="sha256:" + sha256(payload).hexdigest(),
                    media_type="image/png",
                    original_name=name,
                    size=len(payload),
                ),
                str(path.resolve()),
            )
        )
    record = ExecutionMaterialRecord(
        record_id="record-gallery",
        display_name="Gallery",
        group_id="",
        field_values={},
        field_owners={},
        resources={"gallery": tuple(resources)},
        resource_owners={"gallery": ("record",)},
    )
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("{{@img:gallery}}")
    document.add_heading("Following title", level=1)
    document.save(source)

    output = tmp_path / "output.docx"
    _insert_images(
        source,
        output,
        record=record,
        roles=("gallery",),
        image_policy={
            "adaptive": True,
            "show_multi_image_name": True,
            "page_break_after_images": True,
            "role_label:gallery": "现场照片",
            "role_cardinality:gallery": "multiple",
        },
        cache_dir=tmp_path / "image-cache",
    )

    rendered = Document(output)
    image_paragraph = rendered.paragraphs[0]
    assert image_paragraph.text.count("现场照片") == 1
    assert all(name not in image_paragraph.text for name in (
        "photo1.png",
        "photo2.png",
        "photo10.png",
    ))
    page_breaks = [
        element
        for element in image_paragraph._p.iter(qn("w:br"))
        if element.get(qn("w:type")) == "page"
    ]
    assert len(page_breaks) == 1
    blips = tuple(image_paragraph._p.iter(qn("a:blip")))
    rendered_colors = [
        Image.open(
            BytesIO(
                rendered.part.related_parts[blip.get(qn("r:embed"))].blob
            )
        ).convert("RGB").getpixel((0, 0))
        for blip in blips
    ]
    assert rendered_colors == expected_colors


def test_document_batch_consumes_all_five_v1_material_token_domains(tmp_path):
    content = tmp_path / "body.md"
    content.write_text("# Inserted body\n\nStructured content.", encoding="utf-8")
    content_two = tmp_path / "body-two.md"
    content_two.write_text("Second folder item.", encoding="utf-8")
    image = tmp_path / "logo.png"
    Image.new("RGB", (320, 160), "navy").save(image)
    image_two = tmp_path / "cover.png"
    Image.new("RGB", (160, 320), "teal").save(image_two)
    single_image = tmp_path / "avatar.png"
    Image.new("RGB", (240, 2400), "purple").save(single_image)
    attachment = tmp_path / "appendix.pdf"
    attachment.write_bytes(b"attachment-payload")

    paths = (content, content_two, image, image_two, single_image, attachment)
    refs = {
        path.name: MaterialObjectRef(
            object_id="sha256:" + sha256(path.read_bytes()).hexdigest(),
            media_type={
                ".md": "text/markdown",
                ".png": "image/png",
                ".pdf": "application/pdf",
            }[path.suffix],
            original_name=path.name,
            size=path.stat().st_size,
        )
        for path in paths
    }
    record = MaterialRecord(
        record_id=generate_record_id(),
        display_name="Five domains",
        lifecycle="active",
        scope=MaterialScope(
            fields={"title": "Frozen title", "start": "2026-07-31"},
            resources={
                "body": MaterialResourceBinding(
                    "body",
                    (refs["body.md"], refs["body-two.md"]),
                ),
                "logo": MaterialResourceBinding(
                    "logo",
                    (refs["logo.png"], refs["cover.png"]),
                ),
                "portrait": MaterialResourceBinding(
                    "portrait",
                    (refs["avatar.png"],),
                ),
                "appendix": MaterialResourceBinding(
                    "appendix", (refs["appendix.pdf"],)
                ),
            },
            timelines={
                "deadline": MaterialTimelineSpec(
                    output_field="deadline",
                    preset_id="date_offset_days",
                    preset_version=1,
                    anchor_field="start",
                    parameters={"days": "6"},
                )
            },
        ),
    )
    contract = MaterialContract(
        contract_id="five_domains_v1",
        work_mode_id="custom",
        label="Five domains",
        fields=tuple(
            MaterialFieldContract(
                key=key,
                label=key,
                allowed_scopes=("record", "run"),
                allow_run_override=True,
            )
            for key in ("title", "start", "deadline")
        ),
        resource_roles=(
            MaterialResourceRoleContract(
                "body",
                "body",
                "content",
                max_items=None,
            ),
            MaterialResourceRoleContract(
                "logo",
                "logo",
                "image",
                max_items=None,
            ),
            MaterialResourceRoleContract("portrait", "portrait", "image"),
            MaterialResourceRoleContract("appendix", "appendix", "attachment"),
        ),
    )
    package = MaterialPackage(
        package_id=generate_package_id(),
        display_name="Five domains",
        work_mode_id="custom",
        material_contract_id=contract.contract_id,
        records=(record,),
        metadata={
            "material_contract_extensions": {
                "image_policy": {
                    "adaptive": True,
                    "watermark_enabled": True,
                    "watermark_source": "free",
                    "watermark_text": "",
                    "show_single_image_name": False,
                    "show_multi_image_name": True,
                }
            }
        },
    )
    revision = material_package_revision(package)
    by_id = {refs[path.name].object_id: path for path in paths}
    bound = bind_material_run(
        MaterialRunBindRequest(
            package=package,
            package_revision=revision,
            source_type="user",
            selection=MaterialRunSelection(
                package_ref=MaterialPackageRef(package.package_id, revision),
                selected_record_ids=(record.record_id,),
                runtime_image_watermark_text="INTERNAL",
            ),
            contract=contract,
            recipe_id="document_batch",
            scene_id="custom",
            work_mode_id="custom",
        ),
        object_path=lambda item: by_id[item.object_id],
    )
    assert bound.ok and bound.snapshot is not None
    assert bound.snapshot.records[0].timeline_field_keys == (
        "deadline",
        "start",
    )
    assert dict(bound.snapshot.resource_domains) == {
        "body": "content",
        "logo": "image",
        "portrait": "image",
        "appendix": "attachment",
    }

    source = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("{{@text:title}}")
    document.add_paragraph("{{@time:deadline}}")
    document.add_paragraph("{{@file:body}}")
    document.add_paragraph("Logo images")
    document.add_paragraph("{{@img:logo}}")
    document.add_paragraph("Portrait image")
    document.add_paragraph("{{@img:portrait}}")
    document.add_paragraph("{{@attach:appendix}}")
    document.save(source)
    output = tmp_path / "output"
    plan = compile_document_batch_plan(
        bound.snapshot,
        DocumentBatchRequest((str(source),), str(output)),
    )

    result = plan.run()

    assert result["status"] == "success", result
    created = next(output.rglob("template.docx"))
    rendered = Document(created)
    text = "\n".join(item.text for item in rendered.paragraphs)
    assert "Frozen title" in text
    assert "2026-08-06" in text
    assert "Inserted body" in text
    assert "Structured content." in text
    assert "Second folder item." in text
    assert "appendix.pdf" in text
    assert text.splitlines().count("logo") == 1
    assert "logo.png" not in text
    assert "cover.png" not in text
    assert "avatar.png" not in text
    assert "{{@" not in text
    assert len(tuple(rendered.element.body.iter(qn("w:drawing")))) == 3
    image_paragraph_index, image_paragraph = next(
        (index, paragraph)
        for index, paragraph in enumerate(rendered.paragraphs)
        if paragraph.text.strip() == "logo"
    )
    assert rendered.paragraphs[image_paragraph_index - 1].text == "Logo images"
    assert (
        rendered.paragraphs[image_paragraph_index - 1]
        .paragraph_format.keep_with_next
        is None
    )
    assert image_paragraph.paragraph_format.keep_together is None
    image_elements = tuple(image_paragraph._p.iter())
    drawing_positions = tuple(
        index
        for index, element in enumerate(image_elements)
        if element.tag == qn("w:drawing")
    )
    name_positions = [
        index
        for index, element in enumerate(image_elements)
        if element.tag == qn("w:t")
        and element.text == "logo"
    ]
    assert len(name_positions) == 1
    assert name_positions[0] < drawing_positions[0]
    assert drawing_positions[0] < drawing_positions[1]
    portrait_heading_index = next(
        index
        for index, paragraph in enumerate(rendered.paragraphs)
        if paragraph.text == "Portrait image"
    )
    portrait_paragraph = rendered.paragraphs[portrait_heading_index + 1]
    assert portrait_paragraph.paragraph_format.keep_together is None
    assert (
        rendered.paragraphs[portrait_heading_index]
        .paragraph_format.keep_with_next
        is None
    )
    available_height = int(
        rendered.sections[0].page_height
        - rendered.sections[0].top_margin
        - rendered.sections[0].bottom_margin
        - 914400
    )
    assert rendered.inline_shapes[2].height <= available_height
    copied = (
        output
        / "template_attachments"
        / "appendix"
        / "appendix.pdf"
    )
    assert copied.read_bytes() == attachment.read_bytes()

    single_name_package = clone_material_package(
        package,
        metadata={
            "material_contract_extensions": {
                "image_policy": {
                    "adaptive": True,
                    "watermark_enabled": False,
                    "watermark_source": "fixed",
                    "watermark_text": "",
                    "show_single_image_name": True,
                    "show_multi_image_name": False,
                }
            }
        },
    )
    single_name_revision = material_package_revision(single_name_package)
    single_name_bound = bind_material_run(
        MaterialRunBindRequest(
            package=single_name_package,
            package_revision=single_name_revision,
            source_type="user",
            selection=MaterialRunSelection(
                package_ref=MaterialPackageRef(
                    single_name_package.package_id,
                    single_name_revision,
                ),
                selected_record_ids=(record.record_id,),
            ),
            contract=contract,
            recipe_id="document_batch",
            scene_id="custom",
            work_mode_id="custom",
        ),
        object_path=lambda item: by_id[item.object_id],
    )
    assert single_name_bound.ok and single_name_bound.snapshot is not None
    single_name_output = tmp_path / "output-single-name"
    single_name_result = compile_document_batch_plan(
        single_name_bound.snapshot,
        DocumentBatchRequest((str(source),), str(single_name_output)),
    ).run()

    assert single_name_result["status"] == "success", single_name_result
    single_name_document = Document(
        next(single_name_output.rglob("template.docx"))
    )
    single_name_text = "\n".join(
        item.text for item in single_name_document.paragraphs
    )
    assert any(
        paragraph.text.strip() == "portrait"
        for paragraph in single_name_document.paragraphs
    )
    assert "avatar.png" not in single_name_text
    assert "logo.png" not in single_name_text
    assert "cover.png" not in single_name_text


def test_document_batch_fails_closed_for_an_unresolved_canonical_field(tmp_path):
    package = _package()
    revision = material_package_revision(package)
    bound = bind_material_run(
        MaterialRunBindRequest(
            package=package,
            package_revision=revision,
            source_type="builtin",
            selection=MaterialRunSelection(
                package_ref=MaterialPackageRef(package.package_id, revision),
                selected_record_ids=(package.records[0].record_id,),
            ),
            contract=_contract(),
            recipe_id="document_batch",
            scene_id="official",
            work_mode_id="official",
            document_type="notice",
        ),
        object_path=lambda _item: "",
    )
    assert bound.ok and bound.snapshot is not None
    source = tmp_path / "bad-template.docx"
    document = Document()
    document.add_paragraph("{{@text:unknown}}")
    document.save(source)

    plan = compile_document_batch_plan(
        bound.snapshot,
        DocumentBatchRequest((str(source),), str(tmp_path / "output")),
    )

    assert not plan.ok
    assert any(
        item.endswith(":unknown")
        for item in plan.issues
        if item.startswith("document_batch_field_required_missing:")
    )
    assert plan.run()["status"] == "failed"
