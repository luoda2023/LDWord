from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from src.application.materials import (
    ExecutionMaterialRecord,
    ExecutionMaterialSnapshot,
)
from src.config.execution_feature_state import project_execution_scene
from src.config.library import load_scene_from_library, load_template_from_library
from src.domain.materials import MaterialPackageRef
from src.services.production_runtime.execution_runtime import (
    WorkbenchProductionRunner,
)
from src.shared.engine.official_source_formatting import (
    detect_official_source_roles,
)
from src.ui.adapters.workbench_execution_gate import decide_execution_gate
from src.ui.panels.workbench.quick_execution_source_presenter import (
    build_official_document_readiness_projection,
    official_source_execution_gate,
    official_source_execution_issues,
)

NOTICE_PARAGRAPHS = (
    "某市人民政府文件",
    "某政发〔2026〕10号",
    "某市人民政府关于开展专项检查工作的通知",
    "各区人民政府，市政府各部门：",
    "为进一步规范管理，现将有关事项通知如下。",
    "一、工作目标",
    "坚持问题导向，全面开展专项检查。",
    "（一）全面排查",
    "各单位要认真组织实施。",
    "某市人民政府",
    "2026年8月13日",
)


def _misaligned_notice(path: Path) -> None:
    document = Document()
    for index, text in enumerate(NOTICE_PARAGRAPHS):
        style = "Heading 1" if index == 5 else "Heading 2" if index == 7 else "Normal"
        paragraph = document.add_paragraph(text, style=style)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.font.size = Pt(10)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "必须保留的表格"
    document.save(path)


def _official_material_snapshot() -> ExecutionMaterialSnapshot:
    fields = {
        "title": "关于开展专项检查工作的通知",
        "body": "一、工作任务\n请认真组织实施。",
        "organization": "某市人民政府",
        "document_no": "某政发〔2026〕10号",
        "issue_date": "2026年8月13日",
    }
    record = ExecutionMaterialRecord(
        record_id="record-1",
        display_name="通知",
        group_id="",
        field_values=fields,
        field_owners={key: "record" for key in fields},
        resources={},
        resource_owners={},
    )
    return ExecutionMaterialSnapshot(
        snapshot_id="",
        run_id="run-test",
        package_ref=MaterialPackageRef(
            package_id=f"pkg_{'a' * 32}",
            revision=f"sha256:{'b' * 64}",
        ),
        work_mode_id="official",
        material_contract_id="official_document_v1",
        recipe_id="document_batch",
        scene_id="official",
        document_type="notice",
        package_display_name="公文资料",
        package_field_values={},
        package_field_owners={},
        groups=(),
        records=(record,),
    )


def test_existing_official_docx_roles_are_mapped_without_material(tmp_path):
    source = tmp_path / "notice.docx"
    _misaligned_notice(source)

    document = Document(source)
    roles, warnings = detect_official_source_roles(
        document,
        document_type_id="notice",
    )

    assert warnings == ()
    assert {role.role: role.paragraph_index for role in roles if role.level == 0} == {
        "organization": 0,
        "document_no": 1,
        "title": 2,
        "recipient": 3,
        "body": 8,
        "issuer": 9,
        "issue_date": 10,
    }
    assert [(role.paragraph_index, role.level) for role in roles if role.level] == [
        (5, 1),
        (7, 2),
    ]


def test_official_workbench_auto_routes_plain_docx_to_source_formatting(tmp_path):
    source = tmp_path / "notice.docx"
    _misaligned_notice(source)
    scene = project_execution_scene(
        load_scene_from_library("official", mode_id="official"),
        plan_enabled=True,
        template_enabled=True,
        material_enabled=False,
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=load_template_from_library("official_gbt", mode_id="official"),
        scene=scene,
        output_dir=tmp_path / "outputs",
        document_type_id="notice",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert payload["error_text"] == ""
    assert payload["modules_enabled"] == 0
    evidence = payload["official_source_formatting"]
    assert evidence["status"] == "ok"
    assert evidence["strategy"] == "builtin_master_overlay"
    assert evidence["document_type_id"] == "notice"
    assert evidence["master_id"] == "official_gbt_standard"
    assert evidence["layout_family"] == "common"
    assert evidence["preserved_table_count"] == 1

    output = Document(Path(str(payload["output_path"])))
    assert tuple(paragraph.text for paragraph in output.paragraphs) == NOTICE_PARAGRAPHS
    assert len(output.tables) == 1
    assert output.tables[0].cell(0, 0).text == "必须保留的表格"
    assert output.paragraphs[2].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert output.paragraphs[2].runs[0].font.size.pt == pytest.approx(22)
    assert output.paragraphs[5].style.name == "Heading 1"
    assert output.paragraphs[5].runs[0].font.size.pt == pytest.approx(16)
    assert output.paragraphs[9].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert output.paragraphs[10].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert output.sections[0].top_margin.cm == pytest.approx(3.7, abs=0.01)
    assert output.sections[0].left_margin.cm == pytest.approx(2.8, abs=0.01)


def test_official_source_card_explains_source_formatting_without_material(tmp_path):
    source = tmp_path / "notice.docx"
    _misaligned_notice(source)

    summary, rows = build_official_document_readiness_projection(
        profile_id="notice",
        preview_snapshot=None,
        plan_label="公文基础方案",
        template_label="GB/T 9704 公文格式",
        document_path=str(source),
        material_enabled=False,
    )

    assert summary == "已有公文校版：已识别 11 个段落角色"
    assert rows["source"] == "保留原文校版：notice.docx"
    assert rows["assembly"] == "通知 · 套用GB/T 9704 通用红头公文版式"
    assert "文号、标题、主送机关、正文 5 段" in rows["fields"]
    assert "落款、日期" in rows["fields"]


def test_official_workbench_keeps_master_assembly_when_material_is_bound(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = project_execution_scene(
        load_scene_from_library("official", mode_id="official"),
        plan_enabled=True,
        template_enabled=True,
        material_enabled=True,
    )
    scene.default_delivery_preset().artifacts.review_pdf = False

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=load_template_from_library("official_gbt", mode_id="official"),
        scene=scene,
        output_dir=tmp_path / "outputs",
        document_type_id="notice",
        material_snapshot=_official_material_snapshot(),
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert payload.get("official_source_formatting") is None
    assert payload["official_document_assembly"]["status"] == "ok"
    assert payload["official_document_assembly"]["master_id"] == (
        "official_gbt_standard"
    )


def test_official_source_route_continues_when_title_cannot_be_mapped(tmp_path):
    source = tmp_path / "unstructured.docx"
    document = Document()
    document.add_paragraph("这是一份没有公文标题结构的普通材料。")
    document.add_paragraph("其中仅包含两段普通说明文字。")
    document.save(source)
    scene = project_execution_scene(
        load_scene_from_library("official", mode_id="official"),
        plan_enabled=True,
        template_enabled=True,
        material_enabled=False,
    )

    blockers, warnings = official_source_execution_issues(
        str(source),
        "notice",
    )
    gate = official_source_execution_gate(
        decide_execution_gate(),
        official_scene=True,
        material_enabled=False,
        document_path=str(source),
        profile_id="notice",
    )
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=load_template_from_library("official_gbt", mode_id="official"),
        scene=scene,
        output_dir=tmp_path / "outputs",
        document_type_id="notice",
    ).run(lambda *_args: None, lambda: False)

    assert blockers == ()
    assert "未识别到独立标题；其余内容仍按正文和落款规则套版" in warnings
    assert gate.can_run is True
    assert gate.warning_reasons == warnings
    assert payload["status"] == "success"
    assert payload["error_text"] == ""
    assert payload["official_source_formatting"]["status"] == "warning"
    output = Document(Path(str(payload["output_path"])))
    assert output.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert output.paragraphs[0].paragraph_format.first_line_indent.cm == pytest.approx(
        1.1,
        abs=0.01,
    )
    assert output.paragraphs[0].runs[0].font.size.pt == pytest.approx(16)


def test_existing_letter_uses_builtin_letter_master_layout(tmp_path):
    source = tmp_path / "letter.docx"
    document = Document()
    paragraphs = (
        "某市档案局",
        "某档函〔2026〕3号",
        "关于商请协助提供档案材料的函",
        "有关单位：",
        "为推进资料归集工作，现商请贵单位协助提供相关档案材料。",
        "某市档案局",
        "2026年8月14日",
        "抄送：办公室",
    )
    for text in paragraphs:
        paragraph = document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.runs[0].font.size = Pt(10)
    document.save(source)
    scene = project_execution_scene(
        load_scene_from_library("official", mode_id="official"),
        plan_enabled=True,
        template_enabled=True,
        material_enabled=False,
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=load_template_from_library("official_gbt", mode_id="official"),
        scene=scene,
        output_dir=tmp_path / "outputs",
        document_type_id="letter",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    evidence = payload["official_source_formatting"]
    assert evidence["master_id"] == "official_gbt_letter"
    assert evidence["layout_family"] == "letter"
    output = Document(Path(str(payload["output_path"])))
    section = output.sections[0]
    assert section.top_margin.cm == pytest.approx(2.35, abs=0.01)
    assert section.footer_distance.cm == pytest.approx(2.0, abs=0.01)
    assert section.different_first_page_header_footer is True
    mark = output.paragraphs[0]
    assert mark.runs[0].font.size.pt == pytest.approx(36)
    mark_border = mark._p.get_or_add_pPr().find(qn("w:pBdr"))
    assert mark_border.find(qn("w:bottom")).get(qn("w:val")) == (
        "thickThinSmallGap"
    )
    assert output.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert output.paragraphs[2].paragraph_format.space_before.pt == pytest.approx(28)
    assert output.paragraphs[4].paragraph_format.first_line_indent.cm == pytest.approx(
        1.1,
        abs=0.01,
    )
    assert output.paragraphs[5].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert output.paragraphs[6].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    footer_border = (
        section.first_page_footer.paragraphs[0]
        ._p.get_or_add_pPr()
        .find(qn("w:pBdr"))
    )
    assert footer_border.find(qn("w:bottom")).get(qn("w:val")) == (
        "thinThickSmallGap"
    )


def test_titleless_official_source_still_maps_recipient_body_and_closing(tmp_path):
    source = tmp_path / "titleless-letter.docx"
    document = Document()
    for text in (
        "某市档案局",
        "某档函〔2026〕8号",
        "有关单位：",
        "现就档案移交有关事项说明如下。",
        "某市档案局",
        "2026年8月14日",
    ):
        document.add_paragraph(text)
    document.save(source)

    roles, warnings = detect_official_source_roles(
        Document(source),
        document_type_id="letter",
    )

    assert "official_source_title_not_detected" in warnings
    role_by_index = {role.paragraph_index: role.role for role in roles}
    assert role_by_index == {
        0: "organization",
        1: "document_no",
        2: "recipient",
        3: "body",
        4: "issuer",
        5: "issue_date",
    }
