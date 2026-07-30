from __future__ import annotations

from hashlib import sha256
import json
from pathlib import Path

from docx import Document
from PIL import Image
import pytest

import src.services.material_delivery.transaction as directory_transaction_module
from src.config.attachment_materials import (
    AttachmentBinding,
    AttachmentCardinality,
    AttachmentFileStatus,
    AttachmentItem,
    AttachmentProcessingMode,
    AttachmentSourceKind,
)
from src.config.content_materials import FileAssetRef
from src.config.image_materials import (
    FrozenImageMaterialRule,
    ImageCardinality,
    ImageOccurrencePolicy,
    ImagePlacementMode,
    ImagePlacementPolicy,
    ImageSourceBinding,
    ResolvedImageWatermark,
)
from src.config.material_snapshot import MaterialSnapshot
from src.services.material_attachments import (
    AttachmentBundleBuildRequest,
    AttachmentBundleProcessingError,
    AttachmentBundleService,
    process_attachment_bundle,
)
from src.services.material_assets.image_document_executor import (
    ImageDocumentExecutionDependencies,
    ImageDocumentExecutor,
)
from src.services.material_execution.dependency_indexer import (
    MaterialDependencyIndexRequest,
    build_material_dependency_index_from_sources,
)
from src.shared.engine.exact_material_placeholders import (
    scan_document_exact_placeholders,
)
from src.shared.engine.office_image_layout import inventory_ooxml_images
from tests.material_image_layout_fake import FakeSuccessfulLayout


DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def _docx(path: Path, body: str, *, table: str = "", header: str = "") -> Path:
    document = Document()
    document.add_paragraph(body)
    if table:
        document.add_table(1, 1).cell(0, 0).text = table
    if header:
        document.sections[0].header.add_paragraph(header)
    document.save(path)
    return path


def _docx_paragraphs(path: Path, *paragraphs: str) -> Path:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return path


def _ref(path: Path) -> FileAssetRef:
    raw = path.read_bytes()
    return FileAssetRef(
        source_path=str(path.resolve()),
        original_name=path.name,
        media_type=DOCX_MEDIA_TYPE,
        content_sha256=sha256(raw).hexdigest(),
        byte_size=len(raw),
    )


def _snapshot(
    binding: AttachmentBinding,
    *,
    fields: dict[str, object] | None = None,
    token_bindings: dict[str, str] | None = None,
    image_rules: tuple[FrozenImageMaterialRule, ...] = (),
    image_sources: tuple[ImageSourceBinding, ...] = (),
) -> MaterialSnapshot:
    return MaterialSnapshot(
        field_values=fields or {},
        field_token_bindings=token_bindings or {},
        frozen_image_rules=image_rules,
        image_source_bindings=image_sources,
        attachment_bindings=(binding,),
        material_schema_id="attachments",
        material_schema_version="1",
        rule_versions={"attachments": "1"},
    )


def _image_material(
    tmp_path: Path,
    *,
    token: str,
    role: str,
    color: str,
) -> tuple[FrozenImageMaterialRule, ImageSourceBinding]:
    path = tmp_path / f"{role}.png"
    Image.new("RGB", (80, 48), color).save(path)
    payload = path.read_bytes()
    image_ref = FileAssetRef(
        source_path=str(path.resolve()),
        original_name=path.name,
        media_type="image/png",
        content_sha256=sha256(payload).hexdigest(),
        byte_size=len(payload),
    )
    rule = FrozenImageMaterialRule(
        rule_id=f"rule-{role}",
        source_role=role,
        anchor_token=token,
        required=True,
        occurrence_policy=ImageOccurrencePolicy.EXACTLY_ONE,
        cardinality=ImageCardinality.SINGLE,
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIXED_BOX_FLOW,
            fixed_width_cm=6.0,
        ),
        watermark=ResolvedImageWatermark.disabled(),
    )
    return rule, ImageSourceBinding(role, f"image-{role}", 0, image_ref)


def _index(snapshot: MaterialSnapshot, main: Path):
    return build_material_dependency_index_from_sources(
        MaterialDependencyIndexRequest(
            snapshot=snapshot,
            main_document_path=str(main),
        )
    )


def test_substitute_copy_replaces_fields_across_docx_surfaces_and_keeps_source(
    tmp_path: Path,
) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    source = _docx(
        tmp_path / "attachment.docx",
        "{{@text:company}}",
        table="{{@text:project}}",
        header="{{@text:company}}",
    )
    source_before = source.read_bytes()
    binding = AttachmentBinding(
        role="evidence",
        source_path=str(source),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        items=(AttachmentItem("evidence-1", _ref(source)),),
    )
    snapshot = _snapshot(
        binding,
        fields={"company_name": "Example Ltd", "project": "Alpha"},
        token_bindings={
            "@text:company_name": "company_name",
            "@text:company": "company_name",
            "@text:project": "project",
        },
    )
    index = _index(snapshot, main)
    final = tmp_path / "output" / "evidence"

    receipt = process_attachment_bundle(
        AttachmentBundleBuildRequest(
            snapshot=snapshot,
            dependency_index=index,
            binding_role="evidence",
            final_directory=str(final),
        )
    )

    assert source.read_bytes() == source_before
    output = final / "attachment.docx"
    assert output.is_file()
    assert scan_document_exact_placeholders(Document(output)) == []
    output_document = Document(output)
    assert output_document.paragraphs[0].text == "Example Ltd"
    assert output_document.tables[0].cell(0, 0).text == "Alpha"
    assert output_document.sections[0].header.paragraphs[-1].text == "Example Ltd"
    assert receipt.files[0].status is AttachmentFileStatus.SUBSTITUTED
    assert receipt.files[0].field_replacement_count == 3
    assert set(receipt.files[0].replaced_field_keys) == {"company", "project"}
    persisted = json.loads(
        (final / "attachment_bundle_receipt.json").read_text(encoding="utf-8")
    )
    assert persisted == receipt.to_dict()


def test_attachment_bundle_cancel_is_distinct_and_leaves_no_publication(
    tmp_path: Path,
) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    source = _docx(tmp_path / "attachment.docx", "evidence")
    binding = AttachmentBinding(
        role="evidence",
        source_path=str(source),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        items=(AttachmentItem("evidence-1", _ref(source)),),
    )
    snapshot = _snapshot(binding)
    final = tmp_path / "output" / "evidence"

    with pytest.raises(AttachmentBundleProcessingError) as captured:
        process_attachment_bundle(
            AttachmentBundleBuildRequest(
                snapshot=snapshot,
                dependency_index=_index(snapshot, main),
                binding_role="evidence",
                final_directory=str(final),
            ),
            cancel_check=lambda: True,
        )

    assert [item.code for item in captured.value.diagnostics] == [
        "attachment_processing_cancelled"
    ]
    assert not final.exists()
    assert not list(final.parent.glob("*.staging"))


def test_attachment_publish_failure_restores_previous_bundle(
    tmp_path: Path,
    monkeypatch,
) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    source = _docx(tmp_path / "attachment.docx", "evidence")
    binding = AttachmentBinding(
        role="evidence",
        source_path=str(source),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        items=(AttachmentItem("evidence-1", _ref(source)),),
    )
    snapshot = _snapshot(binding)
    final = tmp_path / "output" / "evidence"
    request = AttachmentBundleBuildRequest(
        snapshot=snapshot,
        dependency_index=_index(snapshot, main),
        binding_role="evidence",
        final_directory=str(final),
    )
    process_attachment_bundle(request)
    prior_tree = {
        path.relative_to(final).as_posix(): path.read_bytes()
        for path in final.rglob("*")
        if path.is_file()
    }
    original_replace = directory_transaction_module.os.replace

    def fail_staging_publication(source_path, destination_path):
        source = Path(source_path)
        destination = Path(destination_path)
        if source.name.endswith(".staging") and destination == final:
            raise OSError("injected attachment directory publication failure")
        return original_replace(source_path, destination_path)

    monkeypatch.setattr(
        directory_transaction_module.os,
        "replace",
        fail_staging_publication,
    )

    with pytest.raises(AttachmentBundleProcessingError) as captured:
        process_attachment_bundle(request)

    assert [item.code for item in captured.value.diagnostics] == [
        "attachment_bundle_publish_failed"
    ]
    assert {
        path.relative_to(final).as_posix(): path.read_bytes()
        for path in final.rglob("*")
        if path.is_file()
    } == prior_tree
    assert not list(final.parent.glob(".*.backup"))
    assert not list(final.parent.glob(".*.staging"))


def test_passthrough_copies_unknown_tokens_byte_for_byte_as_warning(tmp_path: Path) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    source = _docx(tmp_path / "passthrough.docx", "{{@text:unknown}}")
    binding = AttachmentBinding(
        role="originals",
        source_path=str(source),
        processing_mode=AttachmentProcessingMode.PASSTHROUGH,
        items=(AttachmentItem("original-1", _ref(source)),),
    )
    snapshot = _snapshot(binding)
    index = _index(snapshot, main)
    assert index.ok
    final = tmp_path / "output" / "originals"

    receipt = process_attachment_bundle(
        AttachmentBundleBuildRequest(
            snapshot=snapshot,
            dependency_index=index,
            binding_role="originals",
            final_directory=str(final),
        )
    )

    assert (final / source.name).read_bytes() == source.read_bytes()
    assert receipt.files[0].status is AttachmentFileStatus.PASSTHROUGH
    assert receipt.files[0].unresolved_tokens == ("{{@text:unknown}}",)
    assert receipt.files[0].field_replacement_count == 0


def test_unknown_token_blocks_substitute_copy_without_publishing(tmp_path: Path) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    source = _docx(tmp_path / "attachment.docx", "{{@text:typo_field}}")
    binding = AttachmentBinding(
        role="evidence",
        source_path=str(source),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        items=(AttachmentItem("evidence-1", _ref(source)),),
    )
    snapshot = _snapshot(binding)
    index = _index(snapshot, main)
    final = tmp_path / "output" / "evidence"

    with pytest.raises(
        AttachmentBundleProcessingError,
        match="attachment_dependency_invalid.*unknown_field_token",
    ):
        process_attachment_bundle(
            AttachmentBundleBuildRequest(
                snapshot=snapshot,
                dependency_index=index,
                binding_role="evidence",
                final_directory=str(final),
            )
        )

    assert not final.exists()
    assert not list(final.parent.glob("*.staging"))


@pytest.mark.parametrize(
    ("token", "reason"),
    (
        ("{{@file:nested_content}}", "content_token_unsupported"),
        ("{{@attach:evidence}}", "attachment_token_unsupported"),
    ),
)
def test_recursive_file_or_attachment_tokens_inside_attachment_are_blocked(
    tmp_path: Path,
    token: str,
    reason: str,
) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    source = _docx(tmp_path / "nested.docx", token)
    binding = AttachmentBinding(
        role="evidence",
        source_path=str(source),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        items=(AttachmentItem("evidence-1", _ref(source)),),
    )
    snapshot = _snapshot(binding)
    index = _index(snapshot, main)
    consumer_id = "attachment:evidence:evidence-1"

    assert reason in {
        item.code for item in index.diagnostics_for_consumer(consumer_id)
    }
    with pytest.raises(
        AttachmentBundleProcessingError,
        match=f"attachment_dependency_invalid.*{reason}",
    ):
        process_attachment_bundle(
            AttachmentBundleBuildRequest(
                snapshot=snapshot,
                dependency_index=index,
                binding_role="evidence",
                final_directory=str(tmp_path / "output" / "evidence"),
            )
        )


def test_source_drift_keeps_previous_published_bundle_untouched(tmp_path: Path) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    first = _docx(tmp_path / "first.docx", "{{@text:company}}")
    second = _docx(tmp_path / "second.docx", "{{@text:company}}")
    binding = AttachmentBinding(
        role="package",
        source_kind=AttachmentSourceKind.FILE_SET,
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        cardinality=AttachmentCardinality.MULTIPLE,
        items=(
            AttachmentItem("first", _ref(first)),
            AttachmentItem("second", _ref(second)),
        ),
        max_items=None,
    )
    snapshot = _snapshot(
        binding,
        fields={"company": "Example"},
        token_bindings={"@text:company": "company"},
    )
    index = _index(snapshot, main)
    final = tmp_path / "output" / "package"
    request = AttachmentBundleBuildRequest(
        snapshot=snapshot,
        dependency_index=index,
        binding_role="package",
        final_directory=str(final),
    )
    first_receipt = process_attachment_bundle(request)
    old_receipt_bytes = (final / "attachment_bundle_receipt.json").read_bytes()
    old_bundle_hash = first_receipt.bundle_sha256
    second.write_bytes(second.read_bytes() + b"drift")

    with pytest.raises(AttachmentBundleProcessingError, match="source_size_drift"):
        process_attachment_bundle(request)

    assert (final / "attachment_bundle_receipt.json").read_bytes() == old_receipt_bytes
    assert json.loads(old_receipt_bytes)["bundle_sha256"] == old_bundle_hash
    assert not list(final.parent.glob("*.staging"))
    assert not list(final.parent.glob("*.backup"))


def test_attachment_images_are_consumer_scoped_and_select_only_present_rules(
    tmp_path: Path,
) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    first = _docx_paragraphs(
        tmp_path / "first.docx",
        "{{@text:company}}",
        "{{@img:LOGO1}}",
    )
    second = _docx_paragraphs(tmp_path / "second.docx", "{{@img:LOGO1}}")
    first_before = first.read_bytes()
    second_before = second.read_bytes()
    logo_rule, logo_source = _image_material(
        tmp_path,
        token="{{@img:LOGO1}}",
        role="logo1",
        color="blue",
    )
    unused_rule, unused_source = _image_material(
        tmp_path,
        token="{{@img:LOGO2}}",
        role="logo2",
        color="red",
    )
    binding = AttachmentBinding(
        role="package",
        source_kind=AttachmentSourceKind.FILE_SET,
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        cardinality=AttachmentCardinality.MULTIPLE,
        items=(
            AttachmentItem("first", _ref(first)),
            AttachmentItem("second", _ref(second)),
        ),
        max_items=None,
    )
    snapshot = _snapshot(
        binding,
        fields={"company_name": "Example Ltd"},
        token_bindings={
            "@text:company": "company_name",
            "@text:company_name": "company_name",
        },
        image_rules=(logo_rule, unused_rule),
        image_sources=(logo_source, unused_source),
    )
    index = _index(snapshot, main)
    assert index.ok
    fake_layout = FakeSuccessfulLayout()
    executor = ImageDocumentExecutor(
        ImageDocumentExecutionDependencies(layout_runner=fake_layout)
    )
    service = AttachmentBundleService(image_executor=executor)
    final = tmp_path / "output" / "package"

    receipt = service.process(
        AttachmentBundleBuildRequest(
            snapshot=snapshot,
            dependency_index=index,
            binding_role="package",
            final_directory=str(final),
            image_cache_dir=str(tmp_path / "image-cache"),
        )
    )

    assert first.read_bytes() == first_before
    assert second.read_bytes() == second_before
    assert len(fake_layout.requests) == 2
    jobs = tuple(request.jobs[0].plan for request in fake_layout.requests)
    assert {job.source_role for job in jobs} == {"logo1"}
    assert len({job.job_id for job in jobs}) == 2
    assert len({job.occurrence_id for job in jobs}) == 2
    first_output = final / first.name
    second_output = final / second.name
    assert Document(first_output).paragraphs[0].text == "Example Ltd"
    assert scan_document_exact_placeholders(Document(first_output)) == []
    assert scan_document_exact_placeholders(Document(second_output)) == []
    assert len(inventory_ooxml_images(first_output)) == 1
    assert len(inventory_ooxml_images(second_output)) == 1
    assert [item.image_job_count for item in receipt.files] == [1, 1]
    assert [item.field_replacement_count for item in receipt.files] == [1, 0]
    assert all(
        item.status is AttachmentFileStatus.SUBSTITUTED
        for item in receipt.files
    )


def test_attachment_image_on_unsupported_surface_blocks_whole_bundle(
    tmp_path: Path,
) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    source = _docx(
        tmp_path / "table-logo.docx",
        "body",
        table="{{@img:LOGO1}}",
    )
    rule, image_source = _image_material(
        tmp_path,
        token="{{@img:LOGO1}}",
        role="logo1",
        color="green",
    )
    binding = AttachmentBinding(
        role="evidence",
        source_path=str(source),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        items=(AttachmentItem("evidence-1", _ref(source)),),
    )
    snapshot = _snapshot(
        binding,
        image_rules=(rule,),
        image_sources=(image_source,),
    )
    index = _index(snapshot, main)
    final = tmp_path / "output" / "evidence"

    with pytest.raises(
        AttachmentBundleProcessingError,
        match="attachment_dependency_invalid.*image_surface_unsupported",
    ):
        process_attachment_bundle(
            AttachmentBundleBuildRequest(
                snapshot=snapshot,
                dependency_index=index,
                binding_role="evidence",
                final_directory=str(final),
            )
        )

    assert not final.exists()
    assert not list(final.parent.glob("*.staging"))


def test_text_and_image_namespaces_with_same_identifier_remain_independent(
    tmp_path: Path,
) -> None:
    main = _docx(tmp_path / "main.docx", "main")
    source = _docx_paragraphs(
        tmp_path / "logo.docx",
        "{{@text:LOGO1}}",
        "{{@img:LOGO1}}",
    )
    source_before = source.read_bytes()
    rule, image_source = _image_material(
        tmp_path,
        token="{{@img:LOGO1}}",
        role="logo1",
        color="purple",
    )
    binding = AttachmentBinding(
        role="evidence",
        source_path=str(source),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        items=(AttachmentItem("evidence-1", _ref(source)),),
    )
    snapshot = _snapshot(
        binding,
        fields={"logo_text": "must-not-replace-image"},
        token_bindings={"@text:LOGO1": "logo_text"},
        image_rules=(rule,),
        image_sources=(image_source,),
    )
    index = _index(snapshot, main)
    final = tmp_path / "output" / "evidence"

    fake_layout = FakeSuccessfulLayout()
    receipt = AttachmentBundleService(
        image_executor=ImageDocumentExecutor(
            ImageDocumentExecutionDependencies(layout_runner=fake_layout)
        )
    ).process(
        AttachmentBundleBuildRequest(
            snapshot=snapshot,
            dependency_index=index,
            binding_role="evidence",
            final_directory=str(final),
            image_cache_dir=str(tmp_path / "image-cache"),
        )
    )

    assert source.read_bytes() == source_before
    output = final / source.name
    assert Document(output).paragraphs[0].text == "must-not-replace-image"
    assert len(inventory_ooxml_images(output)) == 1
    assert receipt.files[0].field_replacement_count == 1
    assert receipt.files[0].image_job_count == 1
