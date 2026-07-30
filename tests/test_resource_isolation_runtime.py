import json
from dataclasses import replace
from pathlib import Path
from shutil import copy2

import pytest
from docx import Document

from src.config.builtin_templates import create_builtin_template
from src.config.entity import (
    EntityArchive,
    EntityProfile,
    load_entity_archive,
    save_entity_archive_bundle,
)
from src.services.execution_session import resolution as execution_session_resolution
from src.services.execution_session import (
    build_execution_session_snapshot,
    execution_session_frozen_official_master,
)
from src.config.library import (
    load_scene_from_library,
    load_template_from_library,
    template_dependent_scene_descriptors,
)
from src.config.material_batch import (
    MaterialBatchItem,
    build_material_batch_items,
    check_material_batch_preflight,
)
from src.config.material_context import MaterialExecutionContext
from src.config import master_library
from src.config.master_library import get_master
from src.config.object_preflight_evidence import build_object_preflight_evidence
from src.config.official_document_profiles import (
    get_official_document_assembly_contract,
)
from src.shared.engine.official_document_assembly import assemble_official_document_docx
from src.services.production_runtime.execution_runtime import (
    WorkbenchBatchProductionRunner,
)
from src.services.production_runtime import execution_runtime
from src.ui.bridge import PanelBridge
from src.ui.panels.workbench.execution_worker import ExecutionWorker


def _object_preflight_receipt(scene, input_path) -> dict[str, str]:
    evidence = build_object_preflight_evidence(scene, input_path)
    return {
        "object_preflight_confirmation_revision": evidence.source_revision,
        "object_preflight_confirmation_digest": evidence.evidence_digest,
    }


def test_bridge_suspends_incompatible_material_on_work_mode_change():
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    bridge.set_current_scene(
        load_scene_from_library("official", mode_id="official"),
        config_id="official",
        emit_signal=False,
    )
    bridge.set_current_material_context(
        MaterialExecutionContext(
            mode_id="official",
            scene_id="official",
            package_id="notice_a",
            profile_id="official:notice",
            entity_data={"title": "旧公文"},
        ),
        emit_signal=False,
    )

    bridge.set_current_work_mode("exam", emit_signal=False)

    assert bridge.current_material_context().is_empty()
    suspended = bridge.suspended_material_states()
    assert len(suspended) == 1
    assert suspended[0]["reason"] == "work_mode_changed"
    assert suspended[0]["context"].package_id == "notice_a"


def test_bridge_keeps_official_document_type_independent_from_the_base_plan():
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    bridge.set_current_scene(
        load_scene_from_library("official", mode_id="official"),
        config_id="official",
        emit_signal=False,
    )
    bridge.set_current_material_context(
        MaterialExecutionContext(
            mode_id="official",
            scene_id="official",
            package_id="notice_a",
        ),
        emit_signal=False,
    )

    bridge.set_current_official_document_type_id(
        "letter",
        source="document_type_selector",
        emit_signal=False,
    )

    assert bridge.current_scene_id() == "official"
    assert bridge.current_official_document_type_id() == "letter"
    assert bridge.current_official_document_type_source() == "document_type_selector"
    assert bridge.current_material_context().package_id == "notice_a"
    assert bridge.suspended_material_states() == ()


def test_execution_snapshot_blocks_unresolved_requested_template(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = load_scene_from_library("official", mode_id="official")

    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=create_builtin_template("official_gbt"),
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=source,
        plan_id="official",
        template_id="missing_template",
        document_type_id="notice",
        **_object_preflight_receipt(scene, source),
    )

    assert snapshot.ready is False
    assert "template_ref_unresolved:missing_template" in snapshot.issues
    assert snapshot.template_ref.requested_id == "missing_template"


def test_execution_snapshot_keeps_runtime_template_ref_out_of_persisted_plan(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = load_scene_from_library("official", mode_id="official")
    persisted_identity = (
        scene.template_id,
        tuple(scene.compatible_template_ids),
        scene.master_id,
    )

    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=create_builtin_template("official_gbt"),
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=source,
        plan_id="official",
        template_id="official_gbt",
        document_type_id="notice",
        **_object_preflight_receipt(scene, source),
    )

    assert snapshot.ready is True
    assert snapshot.template_ref.requested_id == "official_gbt"
    assert snapshot.template_ref.effective_id == "official_gbt"
    assert (
        scene.template_id,
        tuple(scene.compatible_template_ids),
        scene.master_id,
    ) == persisted_identity
    assert not hasattr(scene, "template_ref")


def test_direct_batch_runner_rejects_custom_scene_in_official_mode_before_output(
    tmp_path,
):
    output_dir = tmp_path / "output"
    runner = WorkbenchBatchProductionRunner(
        doc_path=str(tmp_path / "missing.docx"),
        template=create_builtin_template("official_gbt"),
        scene=load_scene_from_library("custom", mode_id="custom"),
        archive=EntityArchive(
            profiles=[
                EntityProfile(
                    profile_id="row_1",
                    fields={"document_type": "notice"},
                )
            ]
        ),
        profile_ids=["row_1"],
        base_output_dir=output_dir,
        source_kind="legacy_generic_batch",
        execution_session=type(
            "Session",
            (),
            {"mode_id": "official"},
        )(),
    )

    payload = runner.run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["error_text"] == (
        "official_batch_source_not_supported:official_document_table_required"
    )
    assert output_dir.exists() is False


def test_direct_runner_blocks_official_terminal_assembly_on_custom_surface(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "output"
    scene = load_scene_from_library("official", mode_id="official")
    scene.mode_id = "custom"
    scene.master_id = "user_master_that_must_not_be_discarded"

    def must_not_run(*_args, **_kwargs):
        raise AssertionError("official assembly and live master fallback must not run")

    monkeypatch.setattr(execution_runtime.Pipeline, "execute", must_not_run)
    monkeypatch.setattr(execution_runtime, "get_master", must_not_run)

    payload = execution_runtime.WorkbenchProductionRunner(
        doc_path=str(source),
        template=create_builtin_template("official_gbt"),
        scene=scene,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["error_text"] == (
        "terminal_assembly_surface_mismatch:official:"
        "official_document_surface_required"
    )
    assert payload["output_paths"] == {}
    assert not list(output_dir.glob("*_official.docx"))


def test_generic_batch_cannot_bypass_official_terminal_surface_gate(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "output"
    scene = load_scene_from_library("official", mode_id="official")
    scene.mode_id = "custom"
    scene.master_id = "user_master_that_must_not_be_discarded"

    def must_not_run(*_args, **_kwargs):
        raise AssertionError("official assembly and live master fallback must not run")

    monkeypatch.setattr(execution_runtime.Pipeline, "execute", must_not_run)
    monkeypatch.setattr(execution_runtime, "get_master", must_not_run)

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=create_builtin_template("official_gbt"),
        scene=scene,
        archive=EntityArchive(
            profiles=[
                EntityProfile(
                    profile_id="row_1",
                    profile_name="row_1",
                    fields={"document_type": "notice"},
                )
            ]
        ),
        profile_ids=["row_1"],
        base_output_dir=output_dir,
        source_kind="legacy_generic_batch",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert (
        "terminal_assembly_surface_mismatch:official:"
        "official_document_surface_required"
    ) in payload["error_text"]
    assert payload["items"][0]["status"] == "failed"
    assert not list(output_dir.rglob("*_official.docx"))


def test_strict_template_load_does_not_substitute_mode_default():
    try:
        load_template_from_library(
            "missing_template",
            mode_id="official",
        )
    except FileNotFoundError as exc:
        assert "template_ref_unresolved" in str(exc)
    else:
        raise AssertionError("strict resource resolution must not fall back")


def test_official_assembly_never_overwrites_existing_output(tmp_path):
    values = {
        "title": "第一份",
        "body": "正文",
        "organization": "机关",
        "document_no": "1号",
        "issue_date": "2026年7月11日",
    }
    first = assemble_official_document_docx(
        "notice",
        values,
        tmp_path,
        filename="same.docx",
    )
    second = assemble_official_document_docx(
        "notice",
        {**values, "title": "第二份"},
        tmp_path,
        filename="same.docx",
    )

    assert first.docx_path == tmp_path / "same.docx"
    assert second.docx_path == tmp_path / "same_2.docx"
    assert first.docx_path.read_bytes() != second.docx_path.read_bytes()


def test_batch_preflight_blocks_duplicate_profile_and_output_identity(tmp_path):
    archive = EntityArchive(
        profiles=[
            EntityProfile(profile_id="same", profile_name="同名"),
            EntityProfile(profile_id="same", profile_name="同名"),
        ]
    )

    result = check_material_batch_preflight(
        archive,
        base_output_dir=tmp_path,
        output_dir_template="fixed",
    )

    assert result.ok is False
    assert "duplicate_profile_id:same" in result.issues
    assert any(issue.startswith("duplicate_output_path:") for issue in result.issues)


@pytest.mark.parametrize(
    "template",
    [
        "../escaped",
        r"..\escaped",
        r"safe\..\escaped",
        r"C:\Windows\Temp\escaped",
        r"\\server\share\escaped",
        "CON",
        "safe/PRN.txt",
        "{profile_name.__class__}",
        "{profile_name:>10}",
    ],
)
def test_batch_output_template_cannot_escape_its_namespace(
    tmp_path: Path,
    template: str,
) -> None:
    base = tmp_path / "session-output"
    archive = EntityArchive(
        archive_id="archive",
        profiles=[EntityProfile(profile_id="profile", profile_name="Profile")],
    )

    preflight = check_material_batch_preflight(
        archive,
        base_output_dir=base,
        output_dir_template=template,
    )

    assert preflight.ok is False
    assert any("output_dir_" in issue for issue in preflight.issues)
    with pytest.raises(ValueError, match="output_dir_"):
        build_material_batch_items(
            archive,
            base_output_dir=base,
            output_dir_template=template,
        )


def test_batch_output_placeholder_values_are_single_sanitized_segments(
    tmp_path: Path,
) -> None:
    base = tmp_path / "session-output"
    archive = EntityArchive(
        archive_id=r"..\archive",
        archive_name=r"C:\outside\archive",
        profiles=[
            EntityProfile(
                profile_id=r"..\profile",
                profile_name=r"..\escaped",
                fields={"entity_name": r"\\server\share\entity"},
            )
        ],
    )
    template = (
        "{archive_id}/{archive_name}/{profile_id}/{profile_name}/{entity_name}"
    )

    preflight = check_material_batch_preflight(
        archive,
        base_output_dir=base,
        output_dir_template=template,
    )
    [item] = build_material_batch_items(
        archive,
        base_output_dir=base,
        output_dir_template=template,
    )

    assert preflight.ok is True
    assert Path(item.output_dir).resolve().is_relative_to(base.resolve())
    relative_parts = Path(item.output_dir).resolve().relative_to(base.resolve()).parts
    assert len(relative_parts) == 5
    assert all(part not in {".", ".."} for part in relative_parts)


def test_batch_output_rejects_existing_symlink_escape(tmp_path: Path) -> None:
    base = tmp_path / "session-output"
    outside = tmp_path / "outside"
    base.mkdir()
    outside.mkdir()
    link = base / "linked"
    try:
        link.symlink_to(outside, target_is_directory=True)
    except OSError as exc:
        pytest.skip(f"directory symlinks are unavailable: {exc}")
    archive = EntityArchive(
        profiles=[EntityProfile(profile_id="profile", profile_name="Profile")]
    )

    preflight = check_material_batch_preflight(
        archive,
        base_output_dir=base,
        output_dir_template="linked/result",
    )

    assert preflight.ok is False
    assert any("output_dir_outside_base" in issue for issue in preflight.issues)


def test_direct_batch_runner_rechecks_output_namespace_before_building(
    tmp_path: Path,
) -> None:
    base = tmp_path / "session-output"
    escaped = tmp_path / "escaped"
    archive = EntityArchive(
        profiles=[EntityProfile(profile_id="profile", profile_name="Profile")]
    )
    runner = WorkbenchBatchProductionRunner(
        doc_path=str(tmp_path / "source.docx"),
        template=None,
        scene=None,
        archive=archive,
        base_output_dir=base,
        output_dir_template="../escaped",
    )

    payload = runner.run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert "output_dir_" in payload["error_text"]
    assert not escaped.exists()


def test_batch_runner_revalidates_built_item_before_any_writer(
    tmp_path: Path,
    monkeypatch,
) -> None:
    base = tmp_path / "session-output"
    outside = tmp_path / "outside"
    archive = EntityArchive(
        profiles=[EntityProfile(profile_id="profile", profile_name="Profile")]
    )
    monkeypatch.setattr(
        execution_runtime,
        "build_material_batch_items",
        lambda *_args, **_kwargs: [
            MaterialBatchItem(
                profile_id="profile",
                profile_name="Profile",
                output_dir=str(outside),
                context=MaterialExecutionContext(),
            )
        ],
    )
    runner = WorkbenchBatchProductionRunner(
        doc_path=str(tmp_path / "source.docx"),
        template=None,
        scene=None,
        archive=archive,
        base_output_dir=base,
        output_dir_template="{profile_name}",
    )

    payload = runner.run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["error_text"] == "output_dir_outside_base"
    assert not outside.exists()


def test_unversioned_entity_archive_is_rejected_non_destructively(
    tmp_path,
):
    assets = tmp_path / "assets"
    assets.mkdir()
    unversioned = tmp_path / "unversioned.json"
    unversioned.write_text(
        '{"archive_id":"a","profiles":[{"profile_id":"p","assets_dir":"assets"}]}',
        encoding="utf-8",
    )
    original = unversioned.read_text(encoding="utf-8")

    with pytest.raises(ValueError, match="material_package_version_unsupported:0"):
        load_entity_archive(unversioned)
    assert unversioned.read_text(encoding="utf-8") == original


def test_entity_archive_rejects_unreleased_future_version(tmp_path):
    source = tmp_path / "future.json"
    source.write_text(
        '{"kind":"alavette.material_package","version":6,"profiles":[]}',
        encoding="utf-8",
    )

    try:
        load_entity_archive(source)
    except ValueError as exc:
        assert str(exc) == "material_package_version_unsupported:6"
    else:
        raise AssertionError("unreleased package versions must not be accepted")


def test_builtin_plan_persists_only_stable_template_and_master_ids():
    scene = load_scene_from_library("official", mode_id="official")

    assert scene.template_id == "official_gbt"
    assert scene.master_id == "official_gbt_standard"
    assert not hasattr(scene, "template_ref")
    assert not hasattr(scene, "master_ref")


def test_execution_snapshot_freezes_effective_letter_master_before_live_source_changes(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = load_scene_from_library("official", mode_id="official")
    builtin_letter = get_master("official_gbt_letter", "official")
    assert builtin_letter is not None
    live_letter_path = tmp_path / "live-letter.docx"
    copy2(builtin_letter.docx_path, live_letter_path)
    live_letter = replace(builtin_letter, docx_path=live_letter_path)
    real_get_master = execution_session_resolution.get_master

    def _get_master(master_id, mode_id=None, **kwargs):
        if str(master_id) == "official_gbt_letter":
            return live_letter
        return real_get_master(master_id, mode_id, **kwargs)

    monkeypatch.setattr(execution_session_resolution, "get_master", _get_master)
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=create_builtin_template("official_gbt"),
        material_context=MaterialExecutionContext(
            mode_id="official",
            profile_id="official:letter",
            entity_data={"document_type": "letter"},
        ),
        input_path=source,
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="official_gbt",
        document_type_id="letter",
        **_object_preflight_receipt(scene, source),
    )

    assert snapshot.ready is True
    assert snapshot.master_ref.resource_id == "official_gbt_letter"
    frozen_letter = execution_session_frozen_official_master(
        snapshot,
        document_type_id="letter",
    )
    frozen_bytes = Path(frozen_letter.docx_path).read_bytes()
    assert frozen_letter.execution_frozen is True
    assert Path(frozen_letter.docx_path).is_relative_to(
        Path(snapshot.output_namespace)
    )

    live_letter_path.write_bytes(b"live source changed after snapshot")

    assert Path(frozen_letter.docx_path).read_bytes() == frozen_bytes
    contract = get_official_document_assembly_contract("letter")
    assert contract is not None
    values = {
        binding.field_key: f"value-{binding.field_key}"
        for binding in contract.field_bindings
        if binding.required
    }
    monkeypatch.setattr(
        master_library,
        "get_master",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("execution assembly reopened the live master library")
        ),
    )
    monkeypatch.setattr(
        master_library,
        "default_master",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("execution assembly used a live default master")
        ),
    )

    result = assemble_official_document_docx(
        "letter",
        values,
        tmp_path / "assembled",
        master=frozen_letter,
    )

    assert result.ok is True
    assert result.master_id == "official_gbt_letter"


def test_official_batch_runtime_uses_each_frozen_master_after_all_live_sources_change(
    tmp_path,
    monkeypatch,
):
    scene = load_scene_from_library("official", mode_id="official")
    live_masters = {}
    for master_id in ("official_gbt_letter", "official_gbt_minutes"):
        builtin = get_master(master_id, "official")
        assert builtin is not None
        live_path = tmp_path / f"live-{master_id}.docx"
        copy2(builtin.docx_path, live_path)
        live_masters[master_id] = replace(builtin, docx_path=live_path)
    real_get_master = execution_session_resolution.get_master

    def _get_master(master_id, mode_id=None, **kwargs):
        return live_masters.get(str(master_id)) or real_get_master(
            master_id,
            mode_id,
            **kwargs,
        )

    monkeypatch.setattr(execution_session_resolution, "get_master", _get_master)
    template = create_builtin_template("official_gbt")
    base_context = MaterialExecutionContext(
        mode_id="official",
        entity_data={"document_type": "notice"},
    )
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=template,
        material_context=base_context,
        input_path="",
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="official_gbt",
        document_type_id="per_item",
        official_document_type_ids=("letter", "minutes"),
    )
    assert snapshot.ready is True
    with pytest.raises(
        RuntimeError,
        match="official_master:per_item:document_type_required",
    ):
        execution_runtime._selected_official_master(
            scene,
            execution_session=snapshot,
        )

    for live_master in live_masters.values():
        live_master.docx_path.write_bytes(b"live source changed after snapshot")
    monkeypatch.setattr(
        master_library,
        "get_master",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("batch assembly reopened the live master library")
        ),
    )
    monkeypatch.setattr(
        master_library,
        "default_master",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("batch assembly used a live default master")
        ),
    )
    profiles = []
    for profile_id in ("letter", "minutes"):
        contract = get_official_document_assembly_contract(profile_id)
        assert contract is not None
        values = {
            binding.field_key: f"value-{profile_id}-{binding.field_key}"
            for binding in contract.field_bindings
            if binding.required
        }
        values["document_type"] = "" if profile_id == "letter" else profile_id
        profiles.append(
            EntityProfile(
                profile_id=f"{profile_id}_row",
                profile_name=profile_id,
                fields=values,
            )
        )

    payload = WorkbenchBatchProductionRunner(
        doc_path="",
        template=template,
        scene=scene,
        archive=EntityArchive(profiles=profiles),
        base_output_dir=snapshot.output_namespace,
        output_dir_template="{profile_id}",
        base_context=base_context,
        source_kind="official_document_table",
        source_path=str(tmp_path / "batch.csv"),
        item_metadata={
            "letter_row": {"official_profile_id": "letter"},
        },
        execution_session=snapshot,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert {
        item["official_profile_id"]: item["official_document_assembly"][
            "master_id"
        ]
        for item in payload["items"]
    } == {
        "letter": "official_gbt_letter",
        "minutes": "official_gbt_minutes",
    }


def test_master_template_contract_blocks_incompatible_pair(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = load_scene_from_library("official", mode_id="official")

    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=create_builtin_template("default"),
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=source,
        plan_id="official",
        template_id="default",
        document_type_id="notice",
        **_object_preflight_receipt(scene, source),
    )

    assert "master_template_incompatible:official_gbt_standard:default" in snapshot.issues


def test_shared_template_impact_lists_all_official_dependents(monkeypatch, tmp_path):
    import src.config.library as library

    monkeypatch.setattr(
        library,
        "SCENE_LIBRARY_DIR",
        tmp_path / "config_library" / "plans",
    )
    monkeypatch.setattr(
        library,
        "TEMPLATE_LIBRARY_DIR",
        tmp_path / "config_library" / "templates",
    )
    library.ensure_config_library()

    dependents = template_dependent_scene_descriptors(
        "official_gbt",
        mode_id="official",
    )

    assert {
        item.scene_id for item in dependents if item.source_type == "builtin"
    } == {"official"}


def test_execution_worker_persists_frozen_session_receipt(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = load_scene_from_library("official", mode_id="official")
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=create_builtin_template("official_gbt"),
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=source,
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="official_gbt",
        document_type_id="notice",
        **_object_preflight_receipt(scene, source),
    )

    class Runner:
        execution_session_snapshot = snapshot

    payload = ExecutionWorker(Runner())._normalize_result(
        {"status": "success"},
        "success",
    )
    receipt = Path(payload["execution_session_path"])
    saved = json.loads(receipt.read_text(encoding="utf-8"))

    assert receipt == Path(snapshot.output_namespace) / "execution_session.json"
    assert saved["session_id"] == snapshot.session_id
    assert saved["plan_ref"]["resource_id"] == "official"
    assert saved["template_ref"]["revision"].startswith("sha256:")
    assert saved["master_ref"]["resource_id"] == "official_gbt_standard"


def test_repeated_runs_get_distinct_output_namespaces(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = load_scene_from_library("official", mode_id="official")
    kwargs = {
        "mode_id": "official",
        "scene": scene,
        "template": create_builtin_template("official_gbt"),
        "material_context": MaterialExecutionContext(mode_id="official"),
        "input_path": source,
        "output_root": tmp_path / "runs",
        "plan_id": "official",
        "template_id": "official_gbt",
        "document_type_id": "notice",
        **_object_preflight_receipt(scene, source),
    }

    first = build_execution_session_snapshot(**kwargs)
    second = build_execution_session_snapshot(**kwargs)

    assert first.session_id != second.session_id
    assert first.output_namespace != second.output_namespace


def test_self_contained_entity_bundle_copies_assets_and_uses_relative_manifest_paths(
    tmp_path,
):
    source = tmp_path / "source.png"
    source.write_bytes(b"image")
    archive = EntityArchive(
        archive_id="archive_a",
        mode_id="custom",
        profiles=[
            EntityProfile(
                profile_id="p1",
                asset_paths={"logo": str(source)},
            )
        ],
    )

    bundle_dir = tmp_path / "bundle"
    target = save_entity_archive_bundle(archive, bundle_dir)
    payload = target.read_text(encoding="utf-8")
    moved_dir = tmp_path / "moved_bundle"
    bundle_dir.rename(moved_dir)
    loaded = load_entity_archive(moved_dir / "package.json")

    assert target.name == "package.json"
    assert '"kind": "alavette.material_package"' in payload
    assert str(source) not in payload
    copied_asset = Path(loaded.profiles[0].asset_paths["logo"])
    assert copied_asset.read_bytes() == b"image"
    assert loaded.profiles[0].asset_metadata["logo"]["sha256"].startswith("sha256:")
    assert check_material_batch_preflight(loaded).ok is True

    copied_asset.write_bytes(b"tampered")
    preflight = check_material_batch_preflight(loaded)

    assert preflight.ok is False
    assert "asset_hash_mismatch:p1:logo" in preflight.issues
