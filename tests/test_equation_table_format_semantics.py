import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from lxml import etree


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.modules.special.equation_table_format import EquationTableFormatModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.font_resolver import resolve_font
from src.shared.engine.line_spacing_ops import apply_line_spacing, sync_spacing_ooxml
from src.shared.engine.ooxml_ops import qn


def test_equation_table_format_keeps_number_specific_font_and_size():
    doc, table, formula_para, number_para = _build_equation_table_doc()

    config = ResolvedConfig()
    config.formula_table.block_alignment = "left"
    config.formula_table.formula_font_name = "Cambria Math"
    config.formula_table.formula_font_size_pt = 16.0
    config.formula_table.formula_line_spacing = 1.75
    config.formula_table.formula_space_before_pt = 9.0
    config.formula_table.formula_space_after_pt = 11.0
    config.formula_style.unify_font = True
    config.formula_style.unify_size = True
    config.formula_style.unify_spacing = True

    EquationTableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    tbl_jc = table._element.find(qn("w:tblPr")).find(qn("w:jc"))
    assert tbl_jc is not None
    assert tbl_jc.get(qn("w:val")) == "left"

    assert formula_para.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert number_para.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT

    formula_spacing = _spacing(formula_para)
    assert formula_spacing.get(qn("w:before")) == "180"
    assert formula_spacing.get(qn("w:after")) == "220"
    assert formula_spacing.get(qn("w:line")) == "420"
    assert formula_spacing.get(qn("w:lineRule")) == "auto"

    number_spacing = _spacing(number_para)
    assert number_spacing.get(qn("w:before")) == "180"
    assert number_spacing.get(qn("w:after")) == "220"
    assert number_spacing.get(qn("w:line")) == "420"
    assert number_spacing.get(qn("w:lineRule")) == "auto"

    resolved_formula_font = resolve_font("Cambria Math", lang="en")
    math_r_pr = _first_math_word_rpr(formula_para)
    assert math_r_pr.find(qn("w:rFonts")).get(qn("w:ascii")) == resolved_formula_font
    assert math_r_pr.find(qn("w:rFonts")).get(qn("w:hAnsi")) == resolved_formula_font
    assert math_r_pr.find(qn("w:rFonts")).get(qn("w:cs")) == resolved_formula_font
    assert math_r_pr.find(qn("w:sz")).get(qn("w:val")) == "32"
    assert math_r_pr.find(qn("w:szCs")).get(qn("w:val")) == "32"

    number_r_pr = _word_rpr(number_para.runs[0])
    resolved_number_font = resolve_font(
        "Times New Roman", lang="en"
    )
    assert number_r_pr.find(qn("w:rFonts")).get(qn("w:ascii")) == resolved_number_font
    assert number_r_pr.find(qn("w:rFonts")).get(qn("w:hAnsi")) == resolved_number_font
    assert number_r_pr.find(qn("w:rFonts")).get(qn("w:cs")) == resolved_number_font
    assert number_r_pr.find(qn("w:sz")).get(qn("w:val")) == "21"
    assert number_r_pr.find(qn("w:szCs")).get(qn("w:val")) == "21"


def test_equation_table_format_applies_formula_spacing_units():
    doc, _table, formula_para, number_para = _build_equation_table_doc()

    config = ResolvedConfig()
    config.formula_table.formula_line_spacing = 1.5
    config.formula_table.formula_space_before_pt = 2.0
    config.formula_table.formula_space_before_unit = "lines"
    config.formula_table.formula_space_after_pt = 0.0
    config.formula_table.formula_space_after_unit = "auto"
    config.formula_style.unify_spacing = True

    EquationTableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    for paragraph in (formula_para, number_para):
        spacing = _spacing(paragraph)
        assert spacing.get(qn("w:before")) is None
        assert spacing.get(qn("w:beforeLines")) == "200"
        assert spacing.get(qn("w:after")) is None
        assert spacing.get(qn("w:afterAutospacing")) == "1"


def test_equation_table_format_uses_number_specific_font_size_and_preserves_spacing_when_not_unified():
    doc, _table, formula_para, number_para = _build_equation_table_doc()

    config = ResolvedConfig()
    config.formula_table.formula_font_name = "Times New Roman"
    config.formula_table.formula_font_size_pt = 15.0
    config.formula_table.formula_line_spacing = 2.0
    config.formula_table.formula_space_before_pt = 8.0
    config.formula_table.formula_space_after_pt = 6.0
    config.formula_table.number_font_name = "Courier New"
    config.formula_table.number_font_size_pt = 10.0
    config.formula_style.unify_font = False
    config.formula_style.unify_size = False
    config.formula_style.unify_spacing = False

    EquationTableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    formula_spacing = _spacing(formula_para)
    assert formula_spacing.get(qn("w:before")) == "40"
    assert formula_spacing.get(qn("w:after")) == "60"
    assert formula_spacing.get(qn("w:line")) == "240"
    assert formula_spacing.get(qn("w:lineRule")) == "auto"

    number_spacing = _spacing(number_para)
    assert number_spacing.get(qn("w:before")) == "60"
    assert number_spacing.get(qn("w:after")) == "80"
    assert number_spacing.get(qn("w:line")) == "360"
    assert number_spacing.get(qn("w:lineRule")) == "auto"

    math_r_pr = _first_math_word_rpr(formula_para)
    assert math_r_pr.find(qn("w:rFonts")).get(qn("w:ascii")) == "Arial"
    assert math_r_pr.find(qn("w:sz")).get(qn("w:val")) == "20"
    assert math_r_pr.find(qn("w:szCs")).get(qn("w:val")) == "20"

    resolved_number_font = resolve_font("Courier New", lang="en")
    number_r_pr = _word_rpr(number_para.runs[0])
    assert number_r_pr.find(qn("w:rFonts")).get(qn("w:ascii")) == resolved_number_font
    assert number_r_pr.find(qn("w:rFonts")).get(qn("w:hAnsi")) == resolved_number_font
    assert number_r_pr.find(qn("w:rFonts")).get(qn("w:cs")) == resolved_number_font
    assert number_r_pr.find(qn("w:sz")).get(qn("w:val")) == "20"
    assert number_r_pr.find(qn("w:szCs")).get(qn("w:val")) == "20"


def test_equation_table_format_normalizes_existing_numbering_by_chapter():
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    table_1 = _add_equation_table(doc, ["(9-9)", "(9-10)"])
    doc.add_heading("第二章 方法", level=1)
    table_2 = _add_equation_table(doc, ["（99）"])

    config = ResolvedConfig()
    config.equation_numbering.numbering_format = "chapter.seq"
    context = PipelineContext(heading_map={0: 1, 1: 1})

    EquationTableFormatModule().apply(doc, config, ChangeTracker(), context)

    assert table_1.cell(0, 1).text == "(1.1)"
    assert table_1.cell(1, 1).text == "(1.2)"
    assert table_2.cell(0, 1).text == "（2.1）"


def test_equation_table_format_normalizes_existing_numbering_globally_without_wrappers():
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    table_1 = _add_equation_table(doc, ["9-9"])
    doc.add_heading("第二章 方法", level=1)
    table_2 = _add_equation_table(doc, ["(8-8)"])

    config = ResolvedConfig()
    config.equation_numbering.numbering_format = "global"
    context = PipelineContext(heading_map={0: 1, 1: 1})

    EquationTableFormatModule().apply(doc, config, ChangeTracker(), context)

    assert table_1.cell(0, 1).text == "1"
    assert table_2.cell(0, 1).text == "(2)"


def test_equation_table_format_infers_chapter_from_existing_number_without_heading_context():
    doc = Document()
    table = _add_equation_table(doc, ["(9-9)"])

    config = ResolvedConfig()
    config.equation_numbering.numbering_format = "chapter.seq"
    tracker = ChangeTracker()

    EquationTableFormatModule().apply(doc, config, tracker, PipelineContext())

    assert table.cell(0, 1).text == "(9.1)"
    records = tracker.get_by_module("equation_table_format")
    assert not any(record.change_type == "skip" for record in records)


def test_equation_table_format_keeps_10_5_pt_number_when_unify_size_is_enabled():
    doc, _table, _formula_para, number_para = _build_equation_table_doc()
    config = ResolvedConfig()
    config.formula_table.formula_font_size_pt = 12.0
    config.formula_table.number_font_size_pt = 10.5
    config.formula_style.unify_size = True

    EquationTableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    number_r_pr = _word_rpr(number_para.runs[0])
    assert number_r_pr.find(qn("w:sz")).get(qn("w:val")) == "21"
    assert number_r_pr.find(qn("w:szCs")).get(qn("w:val")) == "21"


def test_equation_table_format_shrinks_number_column_and_updates_cell_widths():
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    _append_math_formula(table.cell(0, 1).paragraphs[0], [("x", "Arial", "20")])
    table.cell(0, 2).paragraphs[0].add_run("(1-1)")
    before = [
        int(col.get(qn("w:w")))
        for col in table._element.find(qn("w:tblGrid")).findall(qn("w:gridCol"))
    ]

    EquationTableFormatModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    after = [
        int(col.get(qn("w:w")))
        for col in table._element.find(qn("w:tblGrid")).findall(qn("w:gridCol"))
    ]
    assert sum(after) == sum(before)
    assert after[-1] < before[-1]
    assert after[0] == after[-1]
    for index, cell in enumerate(table.rows[0].cells):
        assert int(cell._tc.find(qn("w:tcPr")).find(qn("w:tcW")).get(qn("w:w"))) == after[index]


def test_equation_table_format_generates_missing_number_in_global_mode():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    _append_math_formula(table.cell(0, 0).paragraphs[0], [("x", "Arial", "20")])
    config = ResolvedConfig()
    config.equation_numbering.numbering_format = "global"

    EquationTableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert table.cell(0, 1).text == "(1)"


def test_equation_table_format_starts_at_chapter_one_without_heading_context():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    _append_math_formula(table.cell(0, 0).paragraphs[0], [("x", "Arial", "20")])

    EquationTableFormatModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    assert table.cell(0, 1).text == "(1.1)"


def test_equation_table_format_clears_inherited_and_cell_borders_and_centers_cells():
    doc, table, _formula_para, _number_para = _build_equation_table_doc()
    table_style = etree.SubElement(table._element.find(qn("w:tblPr")), qn("w:tblStyle"))
    table_style.set(qn("w:val"), "TableGrid")
    for cell in table.rows[0].cells:
        cell_properties = cell._tc.find(qn("w:tcPr"))
        cell_borders = etree.SubElement(cell_properties, qn("w:tcBorders"))
        top_border = etree.SubElement(cell_borders, qn("w:top"))
        top_border.set(qn("w:val"), "single")
        vertical_alignment = etree.SubElement(cell_properties, qn("w:vAlign"))
        vertical_alignment.set(qn("w:val"), "top")

    EquationTableFormatModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    table_properties = table._element.find(qn("w:tblPr"))
    assert table_properties.find(qn("w:tblStyle")) is None
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        assert (
            table_properties.find(qn("w:tblBorders")).find(qn(f"w:{side}")).get(qn("w:val"))
            == "none"
        )
    for cell in table.rows[0].cells:
        cell_properties = cell._tc.find(qn("w:tcPr"))
        assert cell_properties.find(qn("w:tcBorders")) is None
        assert cell_properties.find(qn("w:vAlign")).get(qn("w:val")) == "center"


def test_equation_table_format_moves_same_cell_number_to_the_reserved_number_column():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    formula_paragraph = table.cell(0, 0).paragraphs[0]
    _append_math_formula(formula_paragraph, [("x", "Arial", "20")])
    formula_paragraph.add_run(" (9-9)")
    config = ResolvedConfig()
    config.equation_numbering.numbering_format = "global"

    EquationTableFormatModule().apply(
        doc,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    assert formula_paragraph._p.find(f".//{qn('m:oMath')}") is not None
    assert formula_paragraph.text == ""
    assert table.cell(0, 1).text == "(1)"


def test_equation_table_format_recovers_text_formula_with_same_cell_number():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "E=mc^2 (7)"
    table.cell(0, 1).text = ""
    config = ResolvedConfig()
    config.equation_numbering.numbering_format = "global"

    assert _is_equation_table(table) is True
    EquationTableFormatModule().apply(
        doc,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    assert table.cell(0, 0).text == "E=mc^2"
    assert table.cell(0, 1).text == "(1)"


def test_marked_equation_table_does_not_overwrite_non_number_content():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    description = etree.SubElement(table._element.find(qn("w:tblPr")), qn("w:tblDescription"))
    description.set(qn("w:val"), "alavette-equation-table")
    _append_math_formula(table.cell(0, 0).paragraphs[0], [("x", "Arial", "20")])
    table.cell(0, 1).text = "manual review"
    manual_paragraph = table.cell(0, 1).paragraphs[0]
    manual_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _seed_word_run_rpr(
        manual_paragraph.runs[0],
        font_name="Courier New",
        half_points="16",
    )

    context = PipelineContext()
    EquationTableFormatModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        context,
    )

    assert table.cell(0, 1).text == "manual review"
    assert manual_paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
    manual_r_pr = _word_rpr(manual_paragraph.runs[0])
    assert manual_r_pr.find(qn("w:rFonts")).get(qn("w:ascii")) == "Courier New"
    assert manual_r_pr.find(qn("w:sz")).get(qn("w:val")) == "16"
    assert context.formula_runtime["skipped_numbers"] == 1


def test_equation_number_style_does_not_spread_to_number_cell_note():
    doc, _table, _formula_paragraph, number_paragraph = _build_equation_table_doc()
    description = etree.SubElement(
        _table._element.find(qn("w:tblPr")),
        qn("w:tblDescription"),
    )
    description.set(qn("w:val"), "alavette-equation-table")
    number_cell = _table.cell(0, 1)
    note = number_cell.add_paragraph("manual note")
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _seed_word_run_rpr(
        note.runs[0],
        font_name="Courier New",
        half_points="16",
    )

    EquationTableFormatModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    assert number_paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert note.alignment == WD_ALIGN_PARAGRAPH.LEFT
    note_r_pr = _word_rpr(note.runs[0])
    assert note_r_pr.find(qn("w:rFonts")).get(qn("w:ascii")) == "Courier New"
    assert note_r_pr.find(qn("w:sz")).get(qn("w:val")) == "16"


def test_unmarked_equation_table_does_not_number_unrelated_text_rows():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    _append_math_formula(table.cell(0, 0).paragraphs[0], [("x", "Arial", "20")])
    table.cell(1, 0).text = "metadata"
    table.cell(1, 1).text = "(2024)"
    table.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    drawing = etree.SubElement(
        table.cell(1, 0).paragraphs[0]._p,
        qn("w:drawing"),
    )
    assert drawing is not None
    config = ResolvedConfig()
    config.equation_numbering.numbering_format = "global"

    EquationTableFormatModule().apply(
        doc,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    assert table.cell(0, 1).text == "(1)"
    assert table.cell(1, 0).text == "metadata"
    assert table.cell(1, 1).text == "(2024)"
    assert (
        table.cell(1, 0).paragraphs[0].alignment
        == WD_ALIGN_PARAGRAPH.LEFT
    )


def test_equation_table_formula_style_does_not_spread_to_explanatory_paragraph():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    formula_paragraph = table.cell(0, 0).paragraphs[0]
    _append_math_formula(formula_paragraph, [("x", "Arial", "20")])
    explanation = table.cell(0, 0).add_paragraph("explanatory note")
    explanation_run = explanation.runs[0]
    _seed_word_run_rpr(
        explanation_run,
        font_name="Courier New",
        half_points="16",
    )
    _seed_spacing(
        explanation,
        before_pt=6.0,
        after_pt=7.0,
        line_spacing=1.25,
    )

    config = ResolvedConfig()
    config.formula_table.formula_font_name = "Cambria Math"
    config.formula_table.formula_font_size_pt = 16.0
    config.formula_table.formula_line_spacing = 2.0
    config.formula_table.formula_space_before_pt = 9.0
    config.formula_table.formula_space_after_pt = 11.0
    config.formula_style.unify_font = True
    config.formula_style.unify_size = True
    config.formula_style.unify_spacing = True

    EquationTableFormatModule().apply(
        doc,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    explanation_r_pr = _word_rpr(explanation_run)
    assert explanation_r_pr.find(qn("w:rFonts")).get(qn("w:ascii")) == "Courier New"
    assert explanation_r_pr.find(qn("w:sz")).get(qn("w:val")) == "16"
    explanation_spacing = _spacing(explanation)
    assert explanation_spacing.get(qn("w:before")) == "120"
    assert explanation_spacing.get(qn("w:after")) == "140"
    assert explanation_spacing.get(qn("w:line")) == "300"


def test_large_data_table_with_one_native_math_cell_is_not_an_equation_table():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=21, cols=2)
    _append_math_formula(table.cell(0, 0).paragraphs[0], [("x", "Arial", "20")])

    assert _is_equation_table(table) is False


def test_omml_prose_in_data_table_is_not_an_equation_anchor():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    _append_math_formula(
        table.cell(0, 0).paragraphs[0],
        [("其中，混合液中分别得到多个样品。", "Arial", "20")],
    )

    assert _is_equation_table(table) is False


def test_generic_formula_word_in_table_description_is_not_a_structural_marker():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    description = etree.SubElement(
        table._element.find(qn("w:tblPr")),
        qn("w:tblDescription"),
    )
    description.set(qn("w:val"), "Formula results summary")
    table.cell(0, 0).text = "sample"
    table.cell(0, 1).text = "value"

    assert _is_equation_table(table) is False


def test_equation_table_marker_must_match_the_owned_token_exactly():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    description = etree.SubElement(
        table._element.find(qn("w:tblPr")),
        qn("w:tblDescription"),
    )
    description.set(qn("w:val"), "not-alavette-equation-table")
    table.cell(0, 0).text = "sample"
    table.cell(0, 1).text = "value"

    assert _is_equation_table(table) is False


def test_unwrapped_latex_formula_with_empty_number_slot_is_preserved_and_numbered():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = r"\frac{x}{y}"

    assert _is_equation_table(table) is True

    EquationTableFormatModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    assert table.cell(0, 0).text == r"\frac{x}{y}"
    assert table.cell(0, 1).text == "(1.1)"


def test_unwrapped_power_expression_with_empty_number_slot_is_equation_table():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "x^2+y^2"

    assert _is_equation_table(table) is True


def test_plain_status_assignment_with_empty_slot_is_not_equation_table():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "status=ok"

    assert _is_equation_table(table) is False

    table.cell(0, 1).text = "(1)"
    assert _is_equation_table(table) is False


def test_compact_sum_with_same_cell_number_uses_reserved_number_column():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "a+b (7)"

    assert _is_equation_table(table) is True
    EquationTableFormatModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    assert table.cell(0, 0).text == "a+b"
    assert table.cell(0, 1).text == "(1.1)"


def test_equation_table_detection_handles_long_omml_table_but_rejects_data_equals():
    from src.modules.table.table_format import _is_equation_table

    doc = Document()
    equation_table = doc.add_table(rows=6, cols=2)
    _append_math_formula(equation_table.cell(0, 0).paragraphs[0], [("x", "Arial", "20")])
    assert _is_equation_table(equation_table) is True

    data_table = doc.add_table(rows=2, cols=2)
    data_table.cell(0, 0).text = "条件"
    data_table.cell(0, 1).text = "结果"
    data_table.cell(1, 0).text = "a=1"
    data_table.cell(1, 1).text = "通过"
    assert _is_equation_table(data_table) is False


def _build_equation_table_doc():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)

    formula_para = table.cell(0, 0).paragraphs[0]
    _append_math_formula(
        formula_para,
        [
            ("x", "Arial", "20"),
            ("=", "Arial", "20"),
            ("1", "Arial", "20"),
        ],
    )
    formula_para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _seed_spacing(formula_para, before_pt=2.0, after_pt=3.0, line_spacing=1.0)

    number_para = table.cell(0, 1).paragraphs[0]
    number_run = number_para.add_run("(1-1)")
    number_run.font.name = "Arial"
    number_run.font.size = Pt(9.0)
    _seed_word_run_rpr(number_run, font_name="Arial", half_points="18")
    number_para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _seed_spacing(number_para, before_pt=3.0, after_pt=4.0, line_spacing=1.5)

    return doc, table, formula_para, number_para


def _add_equation_table(doc: Document, numbers: list[str]):
    table = doc.add_table(rows=len(numbers), cols=2)
    for row_index, text in enumerate(numbers):
        formula_para = table.cell(row_index, 0).paragraphs[0]
        _append_math_formula(formula_para, [("x", "Arial", "20")])
        number_para = table.cell(row_index, 1).paragraphs[0]
        number_run = number_para.add_run(text)
        _seed_word_run_rpr(number_run, font_name="Arial", half_points="18")
    return table


def _append_math_formula(paragraph, fragments: list[tuple[str, str, str]]) -> None:
    math_para = etree.SubElement(paragraph._element, qn("m:oMathPara"))
    math = etree.SubElement(math_para, qn("m:oMath"))

    for text, font_name, half_points in fragments:
        math_run = etree.SubElement(math, qn("m:r"))
        math_r_pr = etree.SubElement(math_run, qn("m:rPr"))
        etree.SubElement(math_r_pr, qn("m:nor"))
        style = etree.SubElement(math_r_pr, qn("m:sty"))
        style.set(qn("m:val"), "p")

        word_r_pr = etree.SubElement(math_run, qn("w:rPr"))
        r_fonts = etree.SubElement(word_r_pr, qn("w:rFonts"))
        r_fonts.set(qn("w:ascii"), font_name)
        r_fonts.set(qn("w:hAnsi"), font_name)
        r_fonts.set(qn("w:cs"), font_name)
        sz = etree.SubElement(word_r_pr, qn("w:sz"))
        sz.set(qn("w:val"), half_points)
        sz_cs = etree.SubElement(word_r_pr, qn("w:szCs"))
        sz_cs.set(qn("w:val"), half_points)

        text_node = etree.SubElement(math_run, qn("m:t"))
        text_node.text = text


def _seed_word_run_rpr(run, *, font_name: str, half_points: str) -> None:
    r_pr = run._element.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = etree.Element(qn("w:rPr"))
        run._element.insert(0, r_pr)

    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = etree.SubElement(r_pr, qn("w:rFonts"))
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)

    sz = r_pr.find(qn("w:sz"))
    if sz is None:
        sz = etree.SubElement(r_pr, qn("w:sz"))
    sz.set(qn("w:val"), half_points)

    sz_cs = r_pr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = etree.SubElement(r_pr, qn("w:szCs"))
    sz_cs.set(qn("w:val"), half_points)


def _seed_spacing(paragraph, *, before_pt: float, after_pt: float, line_spacing: float) -> None:
    apply_line_spacing(paragraph.paragraph_format, "multiple", line_spacing)
    paragraph.paragraph_format.space_before = Pt(before_pt)
    paragraph.paragraph_format.space_after = Pt(after_pt)
    sync_spacing_ooxml(
        paragraph._element,
        space_before_pt=before_pt,
        space_after_pt=after_pt,
        line_spacing_type="multiple",
        line_spacing_value=line_spacing,
    )


def _spacing(paragraph):
    spacing = paragraph._element.find(qn("w:pPr")).find(qn("w:spacing"))
    assert spacing is not None
    return spacing


def _first_math_word_rpr(paragraph):
    math_run = paragraph._element.find(f".//{qn('m:r')}")
    assert math_run is not None
    word_r_pr = math_run.find(qn("w:rPr"))
    assert word_r_pr is not None
    return word_r_pr


def _word_rpr(run):
    r_pr = run._element.find(qn("w:rPr"))
    assert r_pr is not None
    return r_pr
