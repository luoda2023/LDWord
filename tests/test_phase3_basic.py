"""
Phase 3 冒烟测试 — 基础排版模块 (4 个)
"""

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Cm, Pt
from src.shared.engine.font_resolver import resolve_font
from src.shared.engine.ooxml_ops import qn, find_or_create

from src.config.template import TemplateConfig, StyleConfig
from src.config.scene import SceneWorkspace
from src.config.resolver import resolve_config
from src.pipeline.runner import Pipeline


def test_page_setup_module():
    """M01: page_setup 真实 docx 链路测试"""
    from src.modules.basic.page_setup import PageSetupModule

    doc = Document()
    doc.add_paragraph("测试段落")
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)

    try:
        template = TemplateConfig()
        template.page_setup.paper_size = "A4"
        template.page_setup.margin.top_cm = 2.54
        template.page_setup.margin.left_cm = 3.17

        scene = SceneWorkspace()
        config = resolve_config(template, scene)

        pipeline = Pipeline(modules=[PageSetupModule()], config=config)
        result = pipeline.execute(tmp.name)

        assert result.success, f"执行失败: {result.error}"
        assert result.tracker.total_count >= 1

        out_doc = Document(result.output_paths["final"])
        section = out_doc.sections[0]
        assert abs(section.top_margin - Cm(2.54)) < 10000
        assert abs(section.left_margin - Cm(3.17)) < 10000
        print("  ✅ M01 page_setup 通过 (pipeline 集成)")
    finally:
        Path(tmp.name).unlink(missing_ok=True)
        out = result.output_paths.get("final", "")
        if out:
            Path(out).unlink(missing_ok=True)


def test_section_format_import():
    """M02: section_format 导入+实例化"""
    from src.modules.basic.section_format import SectionFormatModule

    mod = SectionFormatModule()
    assert mod.meta.name == "section_format"
    assert mod.meta.category == "basic"
    assert hasattr(mod, "apply")
    print("  ✅ M02 section_format 导入正常")


def test_paragraph_style_import():
    """M03: paragraph_style 导入+实例化"""
    from src.modules.basic.paragraph_style import ParagraphStyleModule

    mod = ParagraphStyleModule()
    assert mod.meta.name == "paragraph_style"
    assert "heading_recognition" in mod.meta.soft_after
    assert hasattr(mod, "apply")
    print("  ✅ M03 paragraph_style 导入正常")


def test_paragraph_style_syncs_heading_style_definition_text_format():
    """M03: paragraph_style 应同步修正 Word 标题样式定义的字号和字体。"""
    from src.modules.basic.paragraph_style import ParagraphStyleModule

    doc = Document()
    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Arial"
    heading2.font.size = Pt(42)
    r_pr = find_or_create(heading2._element, "w:rPr")
    r_fonts = find_or_create(r_pr, "w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:cs"), "Arial")
    r_fonts.set(qn("w:eastAsia"), "Arial")
    find_or_create(r_pr, "w:sz").set(qn("w:val"), "84")
    find_or_create(r_pr, "w:szCs").set(qn("w:val"), "84")
    doc.add_heading("第二章 方法", level=2)

    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)

    try:
        template = TemplateConfig(
            styles={
                "heading2": StyleConfig(size_pt=14, font_cn="黑体", font_en="Times New Roman", bold=True),
            }
        )
        config = resolve_config(template, SceneWorkspace())
        pipeline = Pipeline(modules=[ParagraphStyleModule()], config=config)
        result = pipeline.execute(tmp.name)

        assert result.success, f"执行失败: {result.error}"

        out_doc = Document(result.output_paths["final"])
        out_style = out_doc.styles["Heading 2"]
        expected_en = resolve_font("Times New Roman", lang="en")
        expected_cn = resolve_font("黑体", lang="cn")

        assert out_style.font.name == expected_en
        assert out_style.font.size is not None
        assert round(out_style.font.size.pt, 1) == 14.0

        out_r_pr = find_or_create(out_style._element, "w:rPr")
        out_r_fonts = find_or_create(out_r_pr, "w:rFonts")
        assert out_r_fonts.get(qn("w:ascii")) == expected_en
        assert out_r_fonts.get(qn("w:hAnsi")) == expected_en
        assert out_r_fonts.get(qn("w:cs")) == expected_en
        assert out_r_fonts.get(qn("w:eastAsia")) == expected_cn
        assert out_r_pr.find(qn("w:sz")).get(qn("w:val")) == "28"
        assert out_r_pr.find(qn("w:szCs")).get(qn("w:val")) == "28"

        para_run = out_doc.paragraphs[0].runs[0]
        assert para_run.font.name == expected_en
        assert para_run.font.size is not None
        assert round(para_run.font.size.pt, 1) == 14.0
        print("  ✅ M03 paragraph_style 已同步 Heading 2 样式定义字号/字体")
    finally:
        Path(tmp.name).unlink(missing_ok=True)
        out = result.output_paths.get("final", "")
        if out:
            Path(out).unlink(missing_ok=True)


def test_header_footer_import():
    """M04: header_footer 导入+实例化"""
    from src.modules.basic.header_footer import HeaderFooterModule

    mod = HeaderFooterModule()
    assert mod.meta.name == "header_footer"
    assert "heading_recognition" in mod.meta.soft_after
    assert hasattr(mod, "apply")
    print("  ✅ M04 header_footer 导入正常")


def test_all_modules_meta():
    """验证 4 个模块的 ModuleMeta 完整性"""
    from src.modules.basic.page_setup import PageSetupModule
    from src.modules.basic.section_format import SectionFormatModule
    from src.modules.basic.paragraph_style import ParagraphStyleModule
    from src.modules.basic.header_footer import HeaderFooterModule

    modules = [PageSetupModule(), SectionFormatModule(),
               ParagraphStyleModule(), HeaderFooterModule()]

    for mod in modules:
        m = mod.meta
        assert m.name, f"{mod} 缺少 name"
        assert m.description, f"{mod} 缺少 description"
        assert m.category == "basic", f"{mod} category 应为 basic"
        assert len(m.requires_config) >= 1, f"{mod} 应声明 requires_config"

    names = [m.meta.name for m in modules]
    assert len(names) == len(set(names)), "模块名不唯一！"
    print(f"  ✅ 4 个模块 meta 完整性验证通过: {names}")


if __name__ == "__main__":
    print("Phase 3 冒烟测试 — 📐 基础排版")
    print("=" * 50)

    test_page_setup_module()
    test_section_format_import()
    test_paragraph_style_import()
    test_header_footer_import()
    test_all_modules_meta()

    print("=" * 50)
    print("✅ 基础排版 4 个模块全部通过！")

