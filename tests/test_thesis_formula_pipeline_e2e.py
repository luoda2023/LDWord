from pathlib import Path

from docx import Document
from lxml import etree

from src.config.resolved import ResolvedConfig
from src.modules.special.equation_table_format import EquationTableFormatModule
from src.modules.special.formula_convert import FormulaConvertModule
from src.modules.structure.heading_recognition import HeadingRecognitionModule
from src.modules.validate.md_cleanup import MdCleanupModule
from src.modules.validate.validation import ValidationModule
from src.pipeline.runner import Pipeline
from src.shared.engine.ooxml_ops import qn


def _append_math(paragraph, text: str = "x") -> None:
    math = etree.SubElement(paragraph._element, qn("m:oMath"))
    math_run = etree.SubElement(math, qn("m:r"))
    etree.SubElement(math_run, qn("m:t")).text = text


def test_thesis_formula_pipeline_is_lossless_geometry_aware_and_idempotent(tmp_path):
    source = tmp_path / "formula-chain-input.docx"
    document = Document()
    document.add_paragraph(r"支持公式：$\frac{a_1}{b^2}+\alpha$。")
    document.add_paragraph(r"低置信公式：$\unsupported{x}$。")

    equation_table = document.add_table(rows=1, cols=3)
    _append_math(equation_table.cell(0, 1).paragraphs[0])

    data_table = document.add_table(rows=1, cols=2)
    data_table.cell(0, 0).text = "条件 a=1"
    data_table.cell(0, 1).text = "通过"
    data_xml_before = data_table._element.xml
    document.save(source)

    config = ResolvedConfig(strict_mode=False)
    config.module_switches.update(
        {
            "formula_convert": True,
            "equation_table_format": True,
            "md_cleanup": True,
            "heading_recognition": True,
            "validation": True,
        }
    )
    config.equation_numbering.numbering_format = "global"
    config.formula_table.formula_font_size_pt = 12.0
    config.formula_table.number_font_size_pt = 10.5
    config.formula_table.auto_shrink_number_column = True
    config.formula_style.unify_size = True

    modules = (
        FormulaConvertModule(),
        MdCleanupModule(),
        HeadingRecognitionModule(),
        EquationTableFormatModule(),
        ValidationModule(),
    )
    first = Pipeline(modules=modules, config=config).execute(str(source))

    assert first.success is True
    assert first.doc.paragraphs[0]._element.find(f".//{qn('m:f')}") is not None
    assert "$" not in first.doc.paragraphs[0].text
    assert first.doc.paragraphs[1].text == r"低置信公式：$\unsupported{x}$。"

    formatted_equation = first.doc.tables[0]
    assert formatted_equation.cell(0, 2).text == "(1)"
    number_size = formatted_equation.cell(0, 2).paragraphs[0]._element.find(
        f".//{qn('w:rPr')}/{qn('w:sz')}"
    )
    assert number_size is not None
    assert number_size.get(qn("w:val")) == "21"
    widths = [
        int(col.get(qn("w:w")))
        for col in formatted_equation._element.find(qn("w:tblGrid")).findall(
            qn("w:gridCol")
        )
    ]
    assert widths[0] == widths[-1] < widths[1]
    assert first.doc.tables[1]._element.xml == data_xml_before

    issue_messages = [
        issue.message for issue in first.context.validation_issues
    ]
    assert any("公式源码未安全转换" in message for message in issue_messages)
    assert not any("缺少编号" in message for message in issue_messages)
    assert not any("编号字号" in message for message in issue_messages)

    first_output = tmp_path / "formula-chain-first.docx"
    first.doc.save(first_output)
    second = Pipeline(modules=modules, config=config).execute(str(first_output))

    assert second.success is True
    assert second.doc.tables[0].cell(0, 2).text == "(1)"
    second_widths = [
        int(col.get(qn("w:w")))
        for col in second.doc.tables[0]._element.find(qn("w:tblGrid")).findall(
            qn("w:gridCol")
        )
    ]
    assert second_widths == widths
    assert second.doc.paragraphs[1].text == first.doc.paragraphs[1].text

    for result in (first, second):
        for output_path in result.output_paths.values():
            Path(output_path).unlink(missing_ok=True)
