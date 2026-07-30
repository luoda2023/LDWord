import sys
import json
from pathlib import Path

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ReplacementRule, ResolvedConfig
from src.config.material_preview import scan_docx_placeholders
from src.modules.fill.entity_fill import EntityFillModule
from src.modules.fill.placeholder_replace import PlaceholderReplaceModule
from src.pipeline.context import PipelineContext
from src.pipeline.result import PipelineResult
from src.pipeline.tracker import ChangeTracker
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.fixed_layout_text import (
    iter_fixed_layout_text_blocks,
    replace_fixed_layout_mapped_fields,
    replace_fixed_layout_placeholders,
)
from src.shared.engine.exact_material_placeholders import (
    replace_document_exact_placeholders,
)
from src.shared.engine.material_field_consistency import (
    inspect_material_field_consistency,
)
from src.shared.engine.ooxml_ops import qn


def test_entity_fill_replaces_content_control_and_textbox_placeholders():
    doc = Document()
    _append_content_control(doc, "{{@text:company_name}}")
    _append_textbox(doc, "{{@text:legal_person}}")
    config = ResolvedConfig()
    config.entity_data = {
        "company_name": "测试公司",
        "legal_person": "张三",
    }
    context = PipelineContext()

    EntityFillModule().apply(doc, config, ChangeTracker(), context)

    text = _all_xml_text(doc)
    assert "测试公司" in text
    assert "张三" in text
    assert "{{@text:company_name}}" not in text
    assert "{{@text:legal_person}}" not in text
    assert context.entity_values == {
        "company_name": "测试公司",
        "legal_person": "张三",
    }


def test_entity_fill_replaces_content_control_tags_aliases_and_textbox_anchors():
    doc = Document()
    _append_content_control(doc, "待填申请人", tag="applicant_name")
    _append_content_control(doc, "待填记录", alias="Record id")
    _append_textbox(doc, "待填单位", doc_pr_name="organization")
    _append_textbox(doc, "待填日期", doc_pr_descr="field:issue_date")
    config = ResolvedConfig()
    config.input_source_profile.material_schema_id = "form_batch_fields_v1"
    config.entity_data = {
        "applicant_name": "李四",
        "record_id": "R-001",
        "organization": "测试单位",
        "issue_date": "2026-06-18",
    }
    context = PipelineContext()

    EntityFillModule().apply(doc, config, ChangeTracker(), context)

    text = _all_xml_text(doc)
    assert "李四" in text
    assert "R-001" in text
    assert "测试单位" in text
    assert "2026-06-18" in text
    assert "待填" not in text
    assert context.entity_values == {
        "organization": "测试单位",
        "issue_date": "2026-06-18",
        "applicant_name": "李四",
        "record_id": "R-001",
    }


def test_entity_fill_uses_profile_aliases_and_reports_fixed_layout_mapping(tmp_path):
    doc = Document()
    _append_content_control(doc, "待填申请人", tag="applicant")
    config = ResolvedConfig()
    config.entity_data = {"applicant_name": "赵六"}
    config.field_aliases = {"applicant": "applicant_name"}
    context = PipelineContext()
    tracker = ChangeTracker()

    EntityFillModule().apply(doc, config, tracker, context)

    assert _all_xml_text(doc) == "赵六"
    assert context.entity_values == {"applicant_name": "赵六"}
    mapping_records = [
        record
        for record in tracker.get_all()
        if record.change_type == "fixed_layout_field_mapping"
    ]
    assert len(mapping_records) == 1
    assert mapping_records[0].before == "applicant"
    assert mapping_records[0].after == "applicant_name"

    result = PipelineResult(success=True, tracker=tracker, context=context)
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"
    write_json_report(
        result,
        input_path=tmp_path / "template.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=1,
        modules_total=1,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "template.docx",
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=1,
        modules_total=1,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    assert report_data["changes"][0]["change_type"] == "fixed_layout_field_mapping"
    assert "applicant_name" in report_md.read_text(encoding="utf-8")


def test_entity_fill_replaces_table_placeholders_and_text_aliases():
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "{{@text:company_name}} / {{@text:company}}"
    table.cell(1, 0).text = "{{@text:record}}"
    config = ResolvedConfig()
    config.entity_data = {
        "company_name": "Company A",
        "record_id": "R-001",
    }
    config.field_aliases = {
        "company": "company_name",
        "record": "record_id",
    }
    context = PipelineContext()

    EntityFillModule().apply(doc, config, ChangeTracker(), context)

    assert table.cell(0, 0).text == "Company A / Company A"
    assert table.cell(1, 0).text == "R-001"
    assert context.entity_values == {
        "company_name": "Company A",
        "record_id": "R-001",
    }


def test_placeholder_replace_replaces_fixed_layout_surfaces():
    doc = Document()
    _append_content_control(doc, "项目：{{project_name}}")
    _append_textbox(doc, "编号：__NO__")
    config = ResolvedConfig()
    config.replacements = [
        ReplacementRule(old="{{project_name}}", new="固定版位项目"),
        ReplacementRule(old="__NO__", new="A-001"),
    ]

    PlaceholderReplaceModule().apply(doc, config, ChangeTracker(), PipelineContext())

    text = _all_xml_text(doc)
    assert "固定版位项目" in text
    assert "A-001" in text
    assert "{{project_name}}" not in text
    assert "__NO__" not in text


def test_placeholder_replace_replaces_table_placeholders():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Project: {{project}}"
    config = ResolvedConfig()
    config.replacements = [ReplacementRule(old="{{project}}", new="Project A")]

    PlaceholderReplaceModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert table.cell(0, 0).text == "Project: Project A"


def test_material_preview_scans_table_and_fixed_layout_placeholders(tmp_path):
    source = tmp_path / "template.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{@text:table_field}}"
    _append_content_control(doc, "{{@text:control_field}}")
    _append_textbox(doc, "{{@text:box_field}}")
    doc.save(source)

    tokens = set(scan_docx_placeholders(source))

    assert {"table_field", "control_field", "box_field"} <= tokens


def test_exact_replacement_uses_same_nested_ooxml_surface_as_preview() -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{@text:table_field}}"
    _append_content_control(doc, "{{@text:control_field}}")
    _append_textbox(doc, "{{@text:box_field}}")

    result = replace_document_exact_placeholders(
        doc,
        {
            "table_field": "表格值",
            "control_field": "控件值",
            "box_field": "文本框值",
        },
    )

    text = _all_xml_text(doc)
    assert result.total_replacements == 3
    assert set(result.replaced_keys) == {
        "table_field",
        "control_field",
        "box_field",
    }
    assert "表格值" in text
    assert "控件值" in text
    assert "文本框值" in text
    assert "{{@text:" not in text


def test_fixed_layout_text_helper_reports_surface_counts():
    doc = Document()
    _append_content_control(doc, "{{party_a}}")
    _append_textbox(doc, "{{party_b}}")

    result = replace_fixed_layout_placeholders(
        doc,
        {
            "{{party_a}}": "甲方公司",
            "{{party_b}}": "乙方公司",
        },
    )
    blocks = iter_fixed_layout_text_blocks(doc)

    assert result.total_replacements == 2
    assert result.content_control_replacements == 1
    assert result.textbox_replacements == 1
    assert result.replaced_tokens == ("{{party_b}}", "{{party_a}}")
    assert [block.surface for block in blocks] == ["textboxes", "content_controls"]
    assert [block.text for block in blocks] == ["乙方公司", "甲方公司"]


def test_fixed_layout_text_helper_replaces_mapped_field_surfaces():
    doc = Document()
    _append_content_control(doc, "待填记录", alias="Record id")
    _append_textbox(doc, "待填单位", doc_pr_descr="field=organization")

    result = replace_fixed_layout_mapped_fields(
        doc,
        {
            "record_id": "R-002",
            "organization": "锚点单位",
        },
        schema_ids=("form_batch_fields_v1",),
    )

    assert result.total_replacements == 2
    assert result.content_control_replacements == 1
    assert result.textbox_replacements == 1
    assert set(result.replaced_fields) == {"record_id", "organization"}
    assert set(result.matched_identifiers) == {"Record id", "field=organization"}
    assert _all_xml_text(doc).splitlines() == ["R-002", "锚点单位"]


def test_fixed_layout_text_helper_replaces_vml_shape_anchor_textbox():
    doc = Document()
    _append_vml_textbox(doc, "待填申请人", shape_alt="field:applicant_name")

    result = replace_fixed_layout_mapped_fields(
        doc,
        {"applicant_name": "王五"},
    )

    assert result.total_replacements == 1
    assert result.textbox_replacements == 1
    assert result.replaced_fields == ("applicant_name",)
    assert result.matched_identifiers == ("field:applicant_name",)
    assert _all_xml_text(doc) == "王五"


def test_material_field_consistency_detects_fixed_layout_placeholder_residue():
    doc = Document()
    _append_content_control(doc, "甲方：{{@text:party_a}}")

    result = inspect_material_field_consistency(
        doc,
        schema_id="contract_parties_v1",
        entity_data={"party_a": "甲方公司"},
    )

    item = result.items[0]
    issue_kinds = {issue.kind for issue in item.issues}
    assert result.status == "warning"
    assert item.field_key == "party_a"
    assert item.placeholders_remaining == ("{{@text:party_a}}",)
    assert "unresolved_placeholder" in issue_kinds
    assert "expected_value_missing" in issue_kinds


def _append_content_control(
    doc: Document,
    text: str,
    *,
    tag: str = "",
    alias: str = "",
) -> None:
    sdt = etree.Element(qn("w:sdt"))
    if tag or alias:
        properties = etree.SubElement(sdt, qn("w:sdtPr"))
        if alias:
            alias_element = etree.SubElement(properties, qn("w:alias"))
            alias_element.set(qn("w:val"), alias)
        if tag:
            tag_element = etree.SubElement(properties, qn("w:tag"))
            tag_element.set(qn("w:val"), tag)
    content = etree.SubElement(sdt, qn("w:sdtContent"))
    _append_paragraph_with_text(content, text)
    _insert_before_sectpr(doc, sdt)


def _append_textbox(
    doc: Document,
    text: str,
    *,
    doc_pr_name: str = "",
    doc_pr_descr: str = "",
) -> None:
    paragraph = etree.Element(qn("w:p"))
    run = etree.SubElement(paragraph, qn("w:r"))
    drawing = etree.SubElement(run, qn("w:drawing"))
    parent = drawing
    if doc_pr_name or doc_pr_descr:
        parent = etree.SubElement(
            drawing,
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline",
        )
        doc_pr = etree.SubElement(
            parent,
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr",
        )
        doc_pr.set("id", "1")
        if doc_pr_name:
            doc_pr.set("name", doc_pr_name)
        if doc_pr_descr:
            doc_pr.set("descr", doc_pr_descr)
    text_box = etree.SubElement(parent, qn("w:txbxContent"))
    _append_paragraph_with_text(text_box, text)
    _insert_before_sectpr(doc, paragraph)


def _append_vml_textbox(
    doc: Document,
    text: str,
    *,
    shape_alt: str = "",
    shape_title: str = "",
    shape_id: str = "",
) -> None:
    paragraph = etree.Element(qn("w:p"))
    run = etree.SubElement(paragraph, qn("w:r"))
    pict = etree.SubElement(run, qn("w:pict"))
    shape = etree.SubElement(pict, "{urn:schemas-microsoft-com:vml}shape")
    if shape_id:
        shape.set("id", shape_id)
    if shape_alt:
        shape.set("alt", shape_alt)
    if shape_title:
        shape.set("title", shape_title)
    text_box = etree.SubElement(shape, "{urn:schemas-microsoft-com:vml}textbox")
    _append_paragraph_with_text(text_box, text)
    _insert_before_sectpr(doc, paragraph)


def _append_paragraph_with_text(parent, text: str) -> None:
    paragraph = etree.SubElement(parent, qn("w:p"))
    run = etree.SubElement(paragraph, qn("w:r"))
    text_node = etree.SubElement(run, qn("w:t"))
    text_node.text = text


def _insert_before_sectpr(doc: Document, element) -> None:
    body = doc._element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        body.insert(list(body).index(sect_pr), element)
    else:
        body.append(element)


def _all_xml_text(doc: Document) -> str:
    return "\n".join(
        text_node.text or ""
        for text_node in doc._element.findall(f".//{qn('w:t')}")
    )
