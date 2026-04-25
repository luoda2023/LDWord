import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.modules.validate.whitespace_normalize import WhitespaceNormalizeModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker


def _apply_whitespace(doc: Document, config: ResolvedConfig | None = None):
    module = WhitespaceNormalizeModule()
    tracker = ChangeTracker()
    module.apply(doc, config or ResolvedConfig(), tracker, PipelineContext())
    return tracker


def test_whitespace_normalize_cleans_space_variants_tabs_zero_width_and_edges():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("\u3000  Alpha")
    tab_run = para.add_run()
    tab_run._element.append(OxmlElement("w:tab"))
    para.add_run("Beta\u200b  ")

    _apply_whitespace(doc)

    assert para.text == "Alpha Beta"


def test_whitespace_normalize_skips_heading_style_paragraphs():
    doc = Document()
    para = doc.add_paragraph("  Heading  1  ")
    para.style = doc.styles["Heading 1"]

    _apply_whitespace(doc)

    assert para.text == "  Heading  1  "


def test_whitespace_normalize_skips_field_paragraphs():
    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run()
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = " TOC \\o \"1-3\" "
    run._element.append(instr)
    text_run = para.add_run("\uff21\uff22\uff23")

    _apply_whitespace(doc)

    assert text_run.text == "\uff21\uff22\uff23"


def test_whitespace_normalize_applies_conservative_full_half_width_conversion():
    doc = Document()
    para = doc.add_paragraph("\uff21\uff22\uff23\uff0ctest(\u4e2d\u6587)")

    _apply_whitespace(doc)

    assert para.text == "ABC,test\uff08\u4e2d\u6587\uff09"
