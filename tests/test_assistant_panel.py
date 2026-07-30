from __future__ import annotations

from pathlib import Path
from threading import Event
import time

from docx import Document

from src.assistant.application.session_coordinator import AssistantSessionCoordinator
from src.assistant.contracts.document_plan import DocumentPlan
from src.assistant.contracts.messages import AssistantMessage, ROLE_ASSISTANT
from src.assistant.domain.docx_format_evidence import (
    STANDARD_FORMAT_REFERENCE_ROLE,
)
from src.assistant.contracts.task_plan import (
    ARTIFACT_KIND_EXAM,
    SOURCE_ROLE_PRODUCTION_INPUT,
    SOURCE_ROLE_REFERENCE_MATERIAL,
)
from src.assistant.runtime.mock_provider import MockModelGateway
from src.assistant.runtime.events import (
    EVENT_CONTEXT_READY,
    EVENT_TEXT_DELTA,
    AssistantEvent,
)
from src.assistant.runtime.provider_contract import (
    PROVIDER_DONE,
    PROVIDER_START,
    PROVIDER_TEXT_DELTA,
    ProviderStreamEvent,
)
from src.assistant.contracts.runtime import (
    TURN_WAITING_TOOL_PERMISSION,
    TURN_WAITING_USER_QUESTION,
)
from src.assistant.runtime.providers.profiles import ProviderProfile, ProviderProfileStore
from src.assistant.runtime.providers.router import ProviderRouter
from src.assistant.runtime.providers.secrets import MemorySecretStore
from src.assistant.runtime.turn_runner import AssistantTurnRunner
from src.assistant.storage.session_store import AssistantSessionStore
from src.assistant.ui.assistant_panel import AssistantPanel
from src.assistant.ui.conversation_view import (
    AssistantConversationMessage,
    AssistantConversationSurface,
    AssistantSourceStrip,
)
from src.assistant.ui.interaction_card import AssistantInteractionCard
from src.assistant.ui.message_components import AssistantAttachmentStrip
from src.config.material_context import MaterialExecutionContext
from src.qt_api import QInputDialog, QLabel, QMessageBox, QPushButton, QTextCursor, Qt
from PySide6.QtTest import QTest
from src.ui.bridge import PanelBridge
from src.ui.panel_registry import PANEL_SPECS, create_panel


def _wait_until(
    qapp,
    predicate,
    *,
    timeout_seconds: float,
    label: str,
) -> None:
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        qapp.processEvents()
        if predicate():
            return
        time.sleep(0.01)
    raise AssertionError(f"timed out waiting for {label}")


def _panel(tmp_path) -> AssistantPanel:
    coordinator = AssistantSessionCoordinator(AssistantSessionStore(tmp_path / "assistant"))
    runner = AssistantTurnRunner(MockModelGateway("计划已准备，请确认后再执行。", chunk_size=4))
    router = ProviderRouter(
        profiles=ProviderProfileStore(tmp_path / "providers.json"),
        secrets=MemorySecretStore(),
    )
    return AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=runner,
        provider_router=router,
    )


class _GatedStreamGateway:
    def __init__(self) -> None:
        self.release = Event()
        self.cancelled = False

    def stream(self, _request):
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text="第一段")
        self.release.wait(5)
        if self.cancelled:
            return
        yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text="第二段")
        yield ProviderStreamEvent(PROVIDER_DONE)

    def cancel(self) -> bool:
        self.cancelled = True
        self.release.set()
        return True


class _FormatReferenceGateway:
    def __init__(self) -> None:
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(
            PROVIDER_TEXT_DELTA,
            text="已基于标准样稿的 Word 格式证据完成核验。",
        )
        yield ProviderStreamEvent(PROVIDER_DONE)

    def cancel(self) -> bool:
        return True


class _WaitingSideChannelGateway:
    def stream(self, _request):
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text="需要读取权限。")
        yield ProviderStreamEvent(
            PROVIDER_DONE,
            metadata={
                "side_channel": {
                    "confirmation_requests": [
                        {
                            "kind": "tool_permission",
                            "title": "允许读取当前文档？",
                            "body": "只会读取已添加的 DOCX。",
                        }
                    ],
                    "continuation_ref": {"continuation_id": "permission-1"},
                    "turn_status": TURN_WAITING_TOOL_PERMISSION,
                }
            },
        )

    def cancel(self) -> bool:
        return True


class _ResumableSideChannelGateway:
    def __init__(self) -> None:
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        yield ProviderStreamEvent(PROVIDER_START)
        if len(self.requests) == 1:
            yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text="需要读取权限。")
            yield ProviderStreamEvent(
                PROVIDER_DONE,
                metadata={
                    "side_channel": {
                        "confirmation_requests": [
                            {"kind": "tool_permission", "title": "允许读取？"}
                        ],
                        "continuation_ref": {"continuation_id": "resume-1"},
                        "turn_status": TURN_WAITING_TOOL_PERMISSION,
                    }
                },
            )
            return
        yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text="已按确认继续。")
        yield ProviderStreamEvent(PROVIDER_DONE)

    def cancel(self) -> bool:
        return True


class _QuestionContinuationGateway:
    def __init__(self) -> None:
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        yield ProviderStreamEvent(PROVIDER_START)
        if len(self.requests) == 1:
            yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text="需要确认交付范围。")
            yield ProviderStreamEvent(
                PROVIDER_DONE,
                metadata={
                    "side_channel": {
                        "confirmation_requests": [
                            {
                                "kind": "question",
                                "title": "选择交付范围",
                                "body": "请选择本次需要的交付物。",
                                "options": [
                                    {"id": "student", "label": "学生卷"},
                                    {"id": "answer", "label": "答案卷"},
                                ],
                                "multiple": True,
                            }
                        ],
                        "continuation_ref": {"continuation_id": "question-1"},
                        "turn_status": TURN_WAITING_USER_QUESTION,
                    }
                },
            )
            return
        yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text="已按所选范围继续。")
        yield ProviderStreamEvent(PROVIDER_DONE)

    def cancel(self) -> bool:
        return True


_PANEL_EXAM_MARKDOWN = """# 七年级数学测试卷

> 科目：数学　年级：七年级　考试时间：60 分钟　满分：15 分

## 一、选择题（本大题共 2 小题，每小题 5 分，共 10 分）

1. `2 + 3` 的结果是（　　）（5 分）
   A. 4
   B. 5
   C. 6
   D. 7

2. 下列数中最小的是（　　）（5 分）
   A. -2
   B. 0
   C. 1
   D. 3

## 二、解答题（本大题共 1 小题，共 5 分）

1. 计算：`8 - 3 + 2`。（5 分）

   answer_area_kind: free
   answer_lines: 4

## 答案速查

一、选择题
1. B
2. A

二、解答题
1. 7。
"""


class _ExamGenerationGateway:
    def __init__(self) -> None:
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        yield ProviderStreamEvent(PROVIDER_START)
        text = (
            _PANEL_EXAM_MARKDOWN
            if request.metadata.get("purpose") == "content_generation"
            else "已识别为试卷创作任务，将在确认后生成并校验题稿。"
        )
        yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text=text)
        yield ProviderStreamEvent(PROVIDER_DONE)

    def cancel(self) -> bool:
        return True


def test_panel_registry_exposes_design_sidebar_first_level_assistant(qapp, monkeypatch, tmp_path):
    monkeypatch.setenv("ALAVETTE_FORM_ASSISTANT_HOME", str(tmp_path))
    ids = [spec.id for spec in PANEL_SPECS]
    assert ids == [
        "workbench",
        "scene",
        "template",
        "assistant",
        "theme",
        "preferences",
    ]
    panel = create_panel("assistant", PanelBridge())
    try:
        assert isinstance(panel, AssistantPanel)
        assert panel._embedded is False
        assert panel._first_level is True
        panel.resize(1280, 800)
        panel.show()
        qapp.processEvents()
        assert panel._session_rail.isVisible()
        assert panel._context_rail.isHidden()
        assert panel._embedded_session_combo.isHidden()
        labels = "\n".join(label.text() for label in panel.findChildren(QLabel))
        assert "置顶" in labels
        assert "最近任务" in labels
    finally:
        panel.close()


def test_standard_format_reference_turn_extracts_evidence_without_production_plan(
    qapp,
    tmp_path,
):
    path = tmp_path / "标准规划文件.docx"
    document = Document()
    document.add_heading("规划标题", level=1)
    document.add_paragraph("规划正文")
    document.save(path)
    gateway = _FormatReferenceGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-format-reference")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
    )
    try:
        panel._empty_input.set_document_path(str(path))

        assert panel._send_message(
            "这个是标准的规划文件，帮我看看确定对应的格式要求"
        )
        assert panel._active_session is not None
        assert gateway.requests == []
        assert panel._active_session.document_job["status"] == (
            "needs_data_disclosure"
        )
        disclosure_cards = [
            card
            for card in panel._message_host.findChildren(
                AssistantInteractionCard
            )
            if card.interaction_type == "disclosure"
        ]
        assert len(disclosure_cards) == 1
        disclosure_facts = dict(
            disclosure_cards[0].presentation.facts
        )
        assert disclosure_facts["服务"] == "本地演示"
        assert disclosure_facts["模型"] == "form-assistant-mock"
        disclosure_id = panel._active_session.pending_continuation[
            "disclosure_id"
        ]
        panel._handle_card_action(
            "approve_provider_disclosure",
            {"disclosure_id": disclosure_id},
        )
        _wait_until(
            qapp,
            lambda: not panel._turn_workers,
            timeout_seconds=5,
            label="format-reference analysis",
        )

        assert len(gateway.requests) == 1
        material_message = gateway.requests[0].messages[-1]["content"]
        assert "<document_format_evidence>" in material_message
        assert "规划正文" not in material_message
        assert '"schema_version":"docx-format-evidence-v1"' in (
            material_message
        )
        assert '"schema_version":"docx-format-evidence-v1"' not in (
            gateway.requests[0].system_prompt
        )
        assert panel._active_session is not None
        assert panel._active_session.context_refs[0]["semantic_role"] == (
            STANDARD_FORMAT_REFERENCE_ROLE
        )
        assert panel._active_session.document_job["status"] == (
            "response_ready"
        )
        assert "plan_candidate" not in panel._active_session.document_job
        assert panel._active_session.active_plan == {}
        assert any(
            message.visible_text() == "已基于标准样稿的 Word 格式证据完成核验。"
            for message in panel._active_session.messages
        )
        evidence_cards = [
            card
            for card in panel._message_host.findChildren(
                AssistantInteractionCard
            )
            if card.interaction_type == "format_evidence"
        ]
        assert len(evidence_cards) == 1
        assert evidence_cards[0].presentation.eyebrow == "格式证据"
        assert ("样稿", path.name) in evidence_cards[0].presentation.facts
        facts = dict(evidence_cards[0].presentation.facts)
        assert int(facts["段落样式"]) >= 2
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_reference_material_disclosure_denial_keeps_body_out_of_provider(
    qapp,
    tmp_path,
):
    path = tmp_path / "private-reference.docx"
    document = Document()
    document.add_paragraph("Confidential acquisition target: Example Corp.")
    document.save(path)
    gateway = _FormatReferenceGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-disclosure-denial")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._on_composer_document_selected(str(path))

        assert panel._send_message("这份材料里提到了什么？") is True

        waiting = panel._active_session
        assert waiting is not None
        assert gateway.requests == []
        assert waiting.document_job["status"] == "needs_data_disclosure"
        assert waiting.pending_continuation["fields"] == ["document_text"]
        assert waiting.context_refs[0]["semantic_role"] == (
            SOURCE_ROLE_REFERENCE_MATERIAL
        )

        panel._handle_card_action(
            "deny_provider_disclosure",
            {
                "disclosure_id": waiting.pending_continuation[
                    "disclosure_id"
                ]
            },
        )

        closed = panel._active_session
        assert closed is not None
        assert gateway.requests == []
        assert closed.document_job["status"] == "response_closed"
        assert closed.document_job["disclosure_decision"] == "denied"
        assert closed.pending_continuation == {}
        assert closed.messages[-1].blocks[0].data["interaction_type"] == (
            "boundary"
        )
    finally:
        panel.close()


def test_content_generation_material_requires_fresh_d1_and_only_sends_after_approval(
    qapp,
    tmp_path,
):
    path = tmp_path / "project-brief.docx"
    document = Document()
    document.add_paragraph("Approved launch date: October 18.")
    document.save(path)
    gateway = _FormatReferenceGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-content-disclosure")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._on_composer_document_selected(str(path))
        assert panel._send_message("根据材料生成一份项目报告") is True
        assert panel._active_session is not None
        assert panel._active_session.document_job["status"] == "plan_ready"
        assert gateway.requests == []

        panel._handle_card_action("generate_content_draft", {})
        waiting = panel._active_session
        assert waiting is not None
        assert waiting.document_job["status"] == "needs_data_disclosure"
        assert waiting.pending_continuation["kind"] == (
            "local_content_disclosure"
        )
        assert gateway.requests == []
        disclosure_cards = [
            card
            for card in panel._message_host.findChildren(
                AssistantInteractionCard
            )
            if card.interaction_type == "disclosure"
        ]
        disclosure_facts = dict(
            disclosure_cards[-1].presentation.facts
        )
        assert disclosure_facts["服务"] == "本地演示"
        assert disclosure_facts["模型"] == "form-assistant-mock"

        first_disclosure_id = waiting.pending_continuation["disclosure_id"]
        panel._handle_card_action(
            "deny_content_disclosure",
            {"disclosure_id": first_disclosure_id},
        )
        denied = panel._active_session
        assert denied is not None
        assert denied.document_job["status"] == "plan_ready"
        assert denied.document_job["disclosure_decision"] == "denied"
        assert gateway.requests == []

        panel._handle_card_action("generate_content_draft", {})
        second_disclosure_id = panel._active_session.pending_continuation[
            "disclosure_id"
        ]
        assert second_disclosure_id != first_disclosure_id
        panel._handle_card_action(
            "approve_content_disclosure",
            {"disclosure_id": second_disclosure_id},
        )
        _wait_until(
            qapp,
            lambda: panel._content_worker is None,
            timeout_seconds=20,
            label="approved material content generation",
        )

        assert len(gateway.requests) == 1
        request = gateway.requests[0]
        assert request.metadata["purpose"] == "content_generation"
        assert "Approved launch date: October 18." in (
            request.messages[0]["content"]
        )
        assert str(tmp_path) not in request.messages[0]["content"]
        assert panel._active_session.document_job["status"] == (
            "content_draft_ready"
        )
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_first_level_design_sidebar_only_exposes_closed_session_flows(qapp, tmp_path):
    coordinator = AssistantSessionCoordinator(AssistantSessionStore(tmp_path / "first-level"))
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(MockModelGateway("计划已准备。")),
        first_level=True,
    )
    try:
        panel.resize(1280, 800)
        panel.show()
        qapp.processEvents()

        assert panel._session_rail.isVisible()
        assert panel._context_rail.isHidden()
        visible_copy = {
            widget.text()
            for widget_type in (QLabel, QPushButton)
            for widget in panel._session_sidebar.findChildren(widget_type)
        }
        assert {"对话", "新任务", "置顶", "最近任务"} <= visible_copy
        assert {
            "委托",
            "项目",
            "计划",
            "长期记忆",
            "技能中心",
        }.isdisjoint(visible_copy)
        assert panel._conversation_stack.count() == 2

        panel.new_session()
        session = panel._active_session
        assert session is not None
        assert panel._session_sidebar.recent_list.count() == 1
        assert panel._set_session_pinned(session.session_id, True)
        assert panel._session_sidebar.pinned_list.count() == 1
        assert panel._session_sidebar.recent_list.count() == 0
    finally:
        panel.close()


def test_first_level_sidebar_remains_reachable_in_compact_drawer(qapp, tmp_path):
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=AssistantSessionCoordinator(
            AssistantSessionStore(tmp_path / "first-level-compact")
        ),
        first_level=True,
    )
    try:
        panel.resize(760, 700)
        panel.show()
        qapp.processEvents()
        assert panel._responsive_mode == "compact"
        assert panel._session_rail.isHidden()
        assert panel._header_widget.isVisible()
        assert panel._session_toggle.isVisible()

        panel._session_toggle.click()
        qapp.processEvents()
        assert panel._session_drawer is not None
        assert panel._session_drawer.isVisible()
        assert panel._session_rail.window() is panel._session_drawer

        panel.resize(1280, 800)
        qapp.processEvents()
        assert panel._responsive_mode == "first_level"
        assert panel._session_rail.isVisible()
        assert panel._context_rail.isHidden()
    finally:
        panel.close()


def test_assistant_embedded_mode_is_center_only(qapp, tmp_path):
    coordinator = AssistantSessionCoordinator(AssistantSessionStore(tmp_path / "embedded"))
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(MockModelGateway("计划已准备。")),
        embedded=True,
    )
    try:
        panel.show()
        qapp.processEvents()
        assert panel._embedded is True
        assert panel._responsive_mode == "embedded"
        assert panel._session_rail.isHidden()
        assert panel._context_rail.isHidden()
        assert panel._session_toggle.text() == "新对话"
        assert panel._context_toggle.isHidden()

        panel.new_session()
        first = panel._active_session
        panel.new_session()
        assert first is not None
        first_index = panel._embedded_session_combo.findData(first.session_id)
        assert first_index > 0
        panel._embedded_session_combo.setCurrentIndex(first_index)
        assert panel._active_session is not None
        assert panel._active_session.session_id == first.session_id
    finally:
        panel.close()


def test_assistant_empty_state_matches_document_task_language(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        labels = "\n".join(label.text() for label in panel.findChildren(QLabel))
        assert "今天要创作什么？" in labels
        assert "快速开始" in labels
        assert "智能建议" not in labels
        assert "常用任务" in labels
        assert "确认文档目标" in labels
        assert (
            panel._empty_input._text_edit.placeholderText()
            == "上传材料、描述想法，或让我先帮你确认文档目标…"
        )
        assert "只添加材料不会上传" in labels
        assert panel._session_rail.width() == 240
        assert panel._context_rail.width() == 280
    finally:
        panel.close()


def test_assistant_creative_home_only_exposes_meaningful_controls(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        home = panel._creative_home
        composer = home.composer

        assert composer._attachment_button.text() == "添加材料"
        assert composer._model_label.text() == "模型"
        assert not hasattr(composer, "_web_button")
        assert not hasattr(composer, "_preset_button")
        assert not hasattr(composer, "_status_ring")
        assert not hasattr(home, "_continue_button")
        assert not composer._send_btn.isEnabled()
        assert composer._send_btn.toolTip() == "输入任务内容后发送"
        assert home.suggestion_buttons[0].isEnabled()
        assert not home.suggestion_buttons[1].isEnabled()
        assert not home.suggestion_buttons[2].isEnabled()

        home.suggestion_buttons[0].click()
        qapp.processEvents()
        assert home.suggestion_buttons[0].property("selected") is True
        assert "文档初稿" in composer.get_text()
        assert composer._send_btn.isEnabled()
        assert panel._active_session is None

        document = tmp_path / "material.docx"
        document.write_bytes(b"placeholder")
        panel.bridge.set_current_document_path(str(document))
        qapp.processEvents()
        assert composer._attachment_button.text() == "更换材料"
        assert home.suggestion_buttons[1].isEnabled()
        assert home.suggestion_buttons[2].isEnabled()

        home.suggestion_buttons[1].click()
        assert composer._send_btn.isEnabled()
        panel.bridge.set_current_document_path("")
        qapp.processEvents()
        assert not composer._send_btn.isEnabled()
        assert "请先添加 DOCX" in composer._keyboard_hint.text()

        composer.set_text("改为不依赖现有文档的普通任务")
        assert composer._send_btn.isEnabled()
        assert home._selected_action_label == ""
    finally:
        panel.close()


def test_custom_task_create_edit_keyboard_select_delete_and_rollback(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel = _panel(tmp_path)
    try:
        panel.resize(1280, 800)
        panel.show()
        qapp.processEvents()
        home = panel._creative_home
        home._custom_tasks = []
        home._rebuild_tasks()
        monkeypatch.setattr(home, "_save_custom_tasks", lambda: True)

        text_answers = iter((("自定义摘要", True), ("自定义摘要（编辑）", True)))
        prompt_answers = iter((("请生成摘要", True), ("请生成结构化摘要", True)))
        monkeypatch.setattr(
            QInputDialog,
            "getText",
            lambda *_args, **_kwargs: next(text_answers),
        )
        monkeypatch.setattr(
            QInputDialog,
            "getMultiLineText",
            lambda *_args, **_kwargs: next(prompt_answers),
        )

        home._add_custom_task()
        assert home._custom_tasks[0][:2] == ("自定义摘要", "请生成摘要")
        assert home._manage_button.isEnabled()

        custom_row = home._task_rows[-1]
        custom_row.setFocus()
        QTest.keyClick(custom_row, Qt.Key_Return)
        qapp.processEvents()
        assert home.composer.get_text() == "请生成摘要"
        assert custom_row.property("selected") is True

        home._edit_custom_task(0)
        assert home._custom_tasks[0][:2] == (
            "自定义摘要（编辑）",
            "请生成结构化摘要",
        )

        monkeypatch.setattr(
            QMessageBox,
            "question",
            lambda *_args, **_kwargs: QMessageBox.StandardButton.Yes,
        )
        home._delete_custom_task(0)
        assert home._custom_tasks == []
        assert not home._manage_button.isEnabled()

        home._custom_tasks = [("需要回滚", "原提示词", "sparkles")]
        home._rebuild_tasks()
        monkeypatch.setattr(home, "_save_custom_tasks", lambda: False)
        monkeypatch.setattr(QMessageBox, "warning", lambda *_args, **_kwargs: None)
        home._delete_custom_task(0)
        assert home._custom_tasks == [("需要回滚", "原提示词", "sparkles")]
    finally:
        panel.close()


def test_assistant_creative_home_restores_design_geometry_and_grid(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel.resize(1600, 1000)
        panel.show()
        qapp.processEvents()

        home = panel._creative_home
        available = max(320, home.width() - 144)
        expected_width = min(min(1180, max(860, int(available * 0.78))), available)
        assert home._center.width() == expected_width
        assert home.composer.height() == 220
        assert home._composer_slot.height() == 240
        assert home._task_scroll.height() == 150
        assert {row.height() for row in home._task_rows} == {50}

        origin = home._grid_point(0, 0, home.width(), home.height(), 42)
        next_column = home._grid_point(1, 0, home.width(), home.height(), 42)
        next_row = home._grid_point(0, 1, home.width(), home.height(), 42)
        assert abs((next_column[0] - origin[0]) - 42.0) < 0.01
        assert abs((next_column[1] - origin[1]) + 3.36) < 0.01
        assert abs((next_row[0] - origin[0]) - 5.04) < 0.01
        assert abs((next_row[1] - origin[1]) - 42.0) < 0.01
        assert home._grid_timer.isActive()
        background = home.grab().toImage()
        sampled_colors = {
            background.pixelColor(x, y).rgba()
            for x in range(0, min(240, background.width()), 6)
            for y in range(0, min(180, background.height()), 6)
        }
        assert len(sampled_colors) >= 4

        home._custom_tasks = [
            ("自定义任务一", "提示词一", "sparkles"),
            ("自定义任务二", "提示词二", "sparkles"),
        ]
        home._rebuild_tasks()
        qapp.processEvents()
        assert home._task_scroll.verticalScrollBar().maximum() > 0
        assert home._task_scroll.height() == 150

        home.hide()
        qapp.processEvents()
        assert not home._grid_timer.isActive()
    finally:
        panel.close()


def test_assistant_creative_home_quick_task_fills_without_auto_sending(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel._creative_home.suggestion_buttons[0].click()
        qapp.processEvents()

        assert "文档初稿" in panel._empty_input.get_text()
        assert panel._active_session is None
        assert panel._conversation_stack.currentWidget() is panel._empty_page
    finally:
        panel.close()


def test_assistant_creative_home_provider_applies_to_first_session(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        profile_id = panel._creative_home.composer.selected_provider_id()
        assert profile_id

        panel._send_message("生成一份项目报告")
        assert panel._active_session is not None
        assert panel._active_session.provider_profile_id == profile_id
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_assistant_home_allows_local_plan_without_provider_and_hot_updates_key(
    qapp,
    tmp_path,
):
    profiles = ProviderProfileStore(tmp_path / "providers.json")
    profiles.upsert(
        ProviderProfile(
            profile_id="cloud-main",
            label="公司模型",
            kind="openai_compatible",
            model_id="example-model",
            base_url="https://example.invalid/v1",
        )
    )
    secrets = MemorySecretStore()
    bridge = PanelBridge()
    panel = AssistantPanel(
        bridge,
        coordinator=AssistantSessionCoordinator(
            AssistantSessionStore(tmp_path / "assistant-readiness")
        ),
        turn_runner=AssistantTurnRunner(MockModelGateway("测试")),
        provider_router=ProviderRouter(profiles=profiles, secrets=secrets),
        first_level=True,
    )
    try:
        combo = panel._creative_home.composer._model_combo
        cloud_index = next(
            index
            for index in range(combo.count())
            if combo.itemData(index) == "cloud-main"
        )
        assert "未就绪" in combo.itemText(cloud_index)
        assert not combo.model().item(cloud_index).isEnabled()

        panel._creative_home.composer.select_provider("cloud-main")
        panel._creative_home.composer.set_text("生成项目报告")
        assert panel._creative_home.composer._send_btn.isEnabled()
        assert "API Key" in panel._creative_home.composer._model_combo.toolTip()

        secrets.set("cloud-main", "secret-value")
        bridge.assistant_provider_profiles_changed.emit()
        qapp.processEvents()

        cloud_index = next(
            index
            for index in range(combo.count())
            if combo.itemData(index) == "cloud-main"
        )
        assert combo.model().item(cloud_index).isEnabled()
        assert combo.currentData() == "cloud-main"
        assert panel._creative_home.composer._send_btn.isEnabled()
    finally:
        panel.close()


def test_assistant_provider_recovery_can_open_ai_settings(qapp, tmp_path):
    panel = _panel(tmp_path)
    destinations = []
    panel.bridge.navigate_to_panel.connect(destinations.append)
    try:
        panel._handle_card_action("open_ai_settings", {})
        target = next(
            index for index, spec in enumerate(PANEL_SPECS) if spec.id == "preferences"
        )
        assert destinations == [target]
        assert panel.bridge.preferred_preferences_page() == "ai"
    finally:
        panel.close()


def test_assistant_provider_recovery_can_retry_original_request(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel = _panel(tmp_path)
    retried = []
    monkeypatch.setattr(panel, "_send_message", retried.append)
    try:
        panel._handle_card_action(
            "retry_provider_request",
            {"retry_text": "重新生成项目报告"},
        )
        assert retried == ["重新生成项目报告"]
    finally:
        panel.close()


def test_active_session_model_identity_tracks_edited_profile(qapp, tmp_path):
    profiles = ProviderProfileStore(tmp_path / "providers.json")
    secrets = MemorySecretStore({"cloud-main": "secret-value"})
    original = ProviderProfile(
        profile_id="cloud-main",
        label="公司模型",
        kind="openai_compatible",
        model_id="model-v1",
        base_url="https://example.invalid/v1",
    )
    profiles.upsert(original)
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-model-sync")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        provider_router=ProviderRouter(profiles=profiles, secrets=secrets),
        first_level=True,
    )
    try:
        panel._active_session = coordinator.create_session(
            provider_profile_id="cloud-main",
            model_id="model-v1",
        )
        profiles.upsert(
            ProviderProfile(
                profile_id="cloud-main",
                label="公司模型",
                kind="openai_compatible",
                model_id="model-v2",
                base_url="https://example.invalid/v1",
            )
        )

        panel._refresh_provider_profiles()

        assert panel._active_session.model_id == "model-v2"
        stored = coordinator.load_session(panel._active_session.session_id)
        assert stored.model_id == "model-v2"
    finally:
        panel.close()


def test_assistant_creative_home_attachment_is_owned_by_session(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel = _panel(tmp_path)
    document = tmp_path / "source.docx"
    document.write_bytes(b"placeholder")
    monkeypatch.setattr(
        "src.assistant.ui.creative_home.QFileDialog.getOpenFileName",
        lambda *_args, **_kwargs: (str(document), "Word 文档 (*.docx)"),
    )
    try:
        panel._creative_home.composer._attachment_button.click()
        qapp.processEvents()

        assert panel.bridge.current_document_path() == ""
        assert panel._active_session is not None
        assert panel._active_session.context_refs[0]["path"] == str(document.resolve())
        assert panel._creative_home.composer._attachment_row.isVisible() is False
        panel.show()
        qapp.processEvents()
        assert panel._creative_home.composer._attachment_row.isVisible()
        assert panel._creative_home.composer._attachment_label.text() == "source.docx"
    finally:
        panel.close()


def test_switching_sessions_restores_each_session_attachment(qapp, tmp_path):
    panel = _panel(tmp_path)
    first_path = tmp_path / "first.docx"
    second_path = tmp_path / "second.docx"
    first_path.write_bytes(b"first")
    second_path.write_bytes(b"second")
    try:
        panel.new_session()
        panel._on_composer_document_selected(str(first_path))
        first = panel._active_session
        assert first is not None

        panel.new_session()
        panel._on_composer_document_selected(str(second_path))
        second = panel._active_session
        assert second is not None
        assert second.session_id != first.session_id

        panel._open_session_by_id(first.session_id)
        assert panel._composer.document_path() == str(first_path.resolve())
        assert panel._empty_input.document_path() == str(first_path.resolve())

        panel._open_session_by_id(second.session_id)
        assert panel._composer.document_path() == str(second_path.resolve())
        assert panel.bridge.current_document_path() == ""
    finally:
        panel.close()


def test_submitted_attachment_is_frozen_on_the_session_and_user_message(
    qapp,
    tmp_path,
):
    panel = _panel(tmp_path)
    document_path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("Material body")
    document.save(document_path)
    panel.bridge.set_current_document_path(str(document_path))
    try:
        assert panel._send_message("summarize this material") is True
        for _attempt in range(200):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)

        session = panel._active_session
        assert session is not None
        assert session.context_refs[0]["title"] == "brief.docx"
        assert session.context_refs[0]["path"] == str(document_path.resolve())
        assert session.messages[0].source_refs == session.context_refs
        user_rows = [
            row
            for row in panel._message_host.findChildren(AssistantConversationMessage)
            if row.property("role") == "user"
        ]
        assert user_rows
        strip = user_rows[0].findChild(AssistantAttachmentStrip)
        assert strip is not None
        assert len(strip.cards) == 1
        assert strip.cards[0].file.title == "brief.docx"
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_embedded_home_transitions_in_place_to_active_conversation(qapp, tmp_path):
    coordinator = AssistantSessionCoordinator(AssistantSessionStore(tmp_path / "embedded-transition"))
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(MockModelGateway("计划已准备。")),
        embedded=True,
    )
    try:
        panel.show()
        qapp.processEvents()
        assert panel._header_widget.isHidden()

        panel._empty_input.set_text("生成一份项目报告")
        panel._empty_input._send_btn.click()
        for _attempt in range(100):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)

        assert panel._conversation_stack.currentWidget() is panel._active_page
        assert panel._header_widget.isVisible()
        assert panel._session_rail.isHidden()
        assert panel._context_rail.isHidden()
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_assistant_local_plan_is_persisted_and_rendered_without_provider(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel._send_message("统一这份文档格式")
        for _attempt in range(100):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)

        session = panel._active_session
        assert session is not None
        assert session.title == "统一这份文档格式"
        assert len(session.messages) == 2
        assert session.messages[-1].blocks[0].data["interaction_type"] == "plan"
        assert session.active_plan["contract_kind"] == "document_plan"
        assert session.document_job["plan_source"] == "form_local_policy"
        assert panel._turn_worker is None
        assert panel._conversation_stack.currentWidget() is panel._active_page
        assert panel._session_list.count() == 1
        assert panel._coordinator.load_session(session.session_id) == session
    finally:
        panel.close()


def test_poc_path_a_executable_attachment_goes_directly_to_a_local_plan(
    qapp,
    tmp_path,
):
    source = tmp_path / "confidential-production-input.docx"
    document = Document()
    document.add_paragraph("This body must remain local during plan creation.")
    document.save(source)
    gateway = _FormatReferenceGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "poc-path-a")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._on_composer_document_selected(str(source))

        assert panel._send_message("统一这份 Word 文档格式") is True

        session = panel._active_session
        assert session is not None
        assert gateway.requests == []
        assert session.document_job["status"] == "plan_ready"
        assert session.document_job["plan_source"] == "form_local_policy"
        assert session.context_refs[0]["semantic_role"] == (
            SOURCE_ROLE_PRODUCTION_INPUT
        )
        assert session.pending_continuation == {}
        persisted_plan = DocumentPlan.from_dict(session.active_plan)
        assert persisted_plan.production_input_ref["path"] == str(
            source.resolve()
        )
    finally:
        panel.close()


def test_poc_path_b_ambiguous_attachment_recovers_to_the_selected_local_route(
    qapp,
    tmp_path,
):
    source = tmp_path / "bilingual-input.docx"
    document = Document()
    document.add_paragraph("中文 / English")
    document.save(source)
    gateway = _FormatReferenceGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "poc-path-b")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._on_composer_document_selected(str(source))
        assert panel._send_message("处理这份双语文档") is True

        waiting = panel._active_session
        assert waiting is not None
        assert gateway.requests == []
        assert waiting.document_job["status"] == "needs_route_clarification"
        assert waiting.pending_continuation["kind"] == (
            "local_route_clarification"
        )
        route_ids = {
            item["id"]
            for item in waiting.pending_continuation["route_choices"]
        }
        assert {
            "quick_bilingual_formatting",
            "bilingual_review_documents",
        } <= route_ids

        panel._handle_card_action(
            "submit_question_answer",
            {
                "clarification_id": waiting.pending_continuation[
                    "clarification_id"
                ],
                "selected_choice_ids": ["quick_bilingual_formatting"],
            },
        )

        resolved = panel._active_session
        assert resolved is not None
        assert gateway.requests == []
        assert resolved.document_job["status"] == "plan_ready"
        assert resolved.document_job["route_id"] == (
            "quick_bilingual_formatting"
        )
        assert resolved.active_plan["capability_ref"]["route_id"] == (
            "quick_bilingual_formatting"
        )
        assert resolved.pending_continuation == {}
    finally:
        panel.close()


def test_poc_path_c_professional_boundary_closes_without_disclosure_or_provider(
    qapp,
    tmp_path,
):
    source = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("Sensitive contract clauses.")
    document.save(source)
    gateway = _FormatReferenceGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "poc-path-c")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._on_composer_document_selected(str(source))

        assert panel._send_message("审查合同法律风险") is True

        session = panel._active_session
        assert session is not None
        assert gateway.requests == []
        assert session.document_job["status"] == "response_closed"
        assert session.document_job["capability_status"] == "gated"
        assert session.active_plan == {}
        assert session.pending_continuation == {}
        assert session.context_refs[0]["semantic_role"] == (
            SOURCE_ROLE_REFERENCE_MATERIAL
        )
        assert session.messages[-1].blocks[0].data["interaction_type"] == (
            "boundary"
        )
    finally:
        panel.close()


def test_general_chat_does_not_create_document_plan_candidate(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel._send_message("你好，请解释一下标题样式的作用")
        for _attempt in range(100):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)

        session = panel._active_session
        assert session is not None
        assert not session.active_plan
        assert "plan_candidate" not in session.document_job
        assert all(
            block.data.get("interaction_type") not in {"plan", "plan_candidate"}
            for message in session.messages
            for block in message.blocks
        )
    finally:
        panel.close()


def test_local_plan_persists_the_send_time_material_snapshot(qapp, tmp_path):
    panel = _panel(tmp_path)
    panel.bridge.set_current_material_context(
        MaterialExecutionContext(entity_data={"project": "turn-a"})
    )
    try:
        panel._send_message("生成一份项目报告")
        for _attempt in range(100):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)
        panel.bridge.set_current_material_context(
            MaterialExecutionContext(entity_data={"project": "live-b"})
        )

        assert panel._active_session is not None
        snapshot = panel._active_session.document_job["material_snapshot"]
        assert snapshot["context"]["frozen_field_values"] == {
            "project": "turn-a"
        }
        assert (
            panel._active_session.active_plan["material_snapshot_ref"]["digest"]
            == snapshot["digest"]
        )
    finally:
        panel.close()


def test_preflight_after_panel_restart_uses_the_persisted_plan_material_snapshot(
    qapp,
    tmp_path,
):
    source = tmp_path / "restart-source.docx"
    document = Document()
    document.add_paragraph("Stable input")
    document.save(source)
    store = AssistantSessionStore(tmp_path / "assistant-restart-material")
    first_bridge = PanelBridge()
    first_bridge.set_current_material_context(
        MaterialExecutionContext(entity_data={"project": "approved-alpha"})
    )
    first_panel = AssistantPanel(
        first_bridge,
        coordinator=AssistantSessionCoordinator(store),
        turn_runner=AssistantTurnRunner(_FormatReferenceGateway()),
        first_level=True,
    )
    first_panel._on_composer_document_selected(str(source))
    assert first_panel._send_message("统一这份 Word 文档格式") is True
    assert first_panel._active_session is not None
    session_id = first_panel._active_session.session_id
    approved_digest = first_panel._active_session.document_job[
        "material_snapshot"
    ]["digest"]
    first_panel.close()

    second_bridge = PanelBridge()
    second_bridge.set_current_material_context(
        MaterialExecutionContext(entity_data={"project": "live-beta"})
    )
    second_panel = AssistantPanel(
        second_bridge,
        coordinator=AssistantSessionCoordinator(store),
        turn_runner=AssistantTurnRunner(_FormatReferenceGateway()),
        first_level=True,
    )
    try:
        second_panel._open_session_by_id(session_id)
        assert second_panel._active_session is not None
        assert second_panel._active_session.session_id == session_id

        second_panel._handle_card_action("preflight", {})
        _wait_until(
            qapp,
            lambda: second_panel._preflight_worker is None,
            timeout_seconds=20,
            label="restart preflight",
        )

        reloaded = second_panel._active_session
        assert reloaded is not None
        assert reloaded.document_job["status"] == (
            "needs_execution_approval"
        )
        assert reloaded.document_job["preflight"][
            "material_context_digest"
        ] == approved_digest
        assert reloaded.document_job["preflight"]["ready"] is True
        approval_block = next(
            block
            for message in reversed(reloaded.messages)
            for block in message.blocks
            if block.data.get("interaction_type") == "approval"
        )
        approval_facts = {
            item["label"]: item["value"]
            for item in approval_block.data["facts"]
        }
        assert approval_facts["输入"] == source.name
        assert approval_facts["操作"] == "处理现有文档"
        assert approval_facts["原文件"] == "保留，不覆盖"
        assert {
            "input_hash",
            "material_context_digest",
            "plan_fingerprint",
            "resource_fingerprints",
        } <= set(approval_block.data["evidence"])
        assert approved_digest not in approval_block.text
        assert not [
            card
            for card in second_panel.findChildren(AssistantInteractionCard)
            if card.interaction_type == "progress" and card.isVisible()
        ]
    finally:
        second_panel.shutdown_active_execution(5000)
        second_panel.close()


def test_active_conversation_migrates_design_message_and_composer_contract(
    qapp,
    tmp_path,
):
    panel = _panel(tmp_path)
    try:
        panel.resize(1280, 800)
        panel.show()
        panel._send_message("请先分析结构，再给出修改建议")
        for _attempt in range(100):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)

        rows = [
            panel._message_layout.itemAt(index).widget()
            for index in range(panel._message_layout.count())
            if isinstance(
                panel._message_layout.itemAt(index).widget(),
                AssistantConversationMessage,
            )
        ]
        assert len(rows) == 2
        assert {row.property("role") for row in rows} == {"user", "assistant"}
        assert panel._composer.objectName() == "assistant_task_composer"
        assert panel._composer.height() == 154
        assert panel._composer._send_btn.width() == 32
        panel._sync_active_reading_widths()
        assert panel._composer.width() == panel._active_page.width() - 48
        assert panel._provider_combo is panel._composer._model_combo
        assert not panel._header_widget.isAncestorOf(panel._provider_combo)
        assert panel._header_icon.isVisible()
        assert panel._session_toggle.isHidden()
        assert panel._active_page.grab().toImage().width() > 0
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_active_conversation_projects_stream_events_before_persistence(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel.new_session()
        assert panel._active_session is not None
        panel._active_session = panel._coordinator.append_message(
            panel._active_session,
            AssistantMessage.text(role="user", text="生成一份摘要"),
            turn_status="provider_running",
        )
        turn_id = "turn-live-preview"
        panel._start_turn_preview(
            session_id=panel._active_session.session_id,
            turn_id=turn_id,
        )
        panel._render_active_session()

        panel._on_turn_event(AssistantEvent(EVENT_CONTEXT_READY, turn_id))
        panel._on_turn_event(
            AssistantEvent(EVENT_TEXT_DELTA, turn_id, text_delta="正在形成摘要")
        )
        qapp.processEvents()

        preview = panel._turn_preview_widget
        assert preview is not None
        assert preview._status_label.text() == "正在生成回复"
        assert preview._markdown.toPlainText() == "正在形成摘要"
        assert len(panel._active_session.messages) == 1

        cancelled = []
        panel._composer.cancel_requested.connect(lambda: cancelled.append(True))
        panel._composer.set_busy(True)
        assert panel._composer._send_btn.isEnabled()
        assert panel._composer._send_btn.accessibleName() == "停止"
        panel._composer._send_btn.click()
        assert cancelled == [True]
    finally:
        panel.close()


def test_worker_stream_reaches_active_conversation_before_turn_finishes(qapp, tmp_path):
    gateway = _GatedStreamGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-live-worker")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._send_message("请流式生成")
        for _attempt in range(100):
            qapp.processEvents()
            preview = panel._turn_preview_widget
            if preview is not None and preview._markdown.toPlainText() == "第一段":
                break
            QTest.qWait(10)

        assert panel._turn_worker is not None
        assert panel._turn_preview_widget is not None
        assert panel._turn_preview_widget._markdown.toPlainText() == "第一段"
        assert panel._composer._busy is True
        assert panel._active_session is not None
        assert len(panel._active_session.messages) == 1

        gateway.release.set()
        for _attempt in range(200):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)

        assert panel._turn_worker is None
        assert panel._composer._busy is False
        assert any(
            message.visible_text() == "第一段第二段"
            for message in panel._active_session.messages
        )
        assert panel._turn_preview_widget is None
    finally:
        gateway.release.set()
        panel.shutdown_active_execution(5000)
        panel.close()


def test_waiting_runtime_result_persists_continuation_and_renders_permission_card(
    qapp,
    tmp_path,
):
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-waiting-runtime")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(_WaitingSideChannelGateway()),
        first_level=True,
    )
    try:
        assert panel._send_message("analyze the document") is True
        for _attempt in range(200):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)

        session = panel._active_session
        assert session is not None
        assert session.turn_status == TURN_WAITING_TOOL_PERMISSION
        assert session.pending_continuation["continuation_id"] == "permission-1"
        assert not session.active_plan
        stored = coordinator.load_session(session.session_id)
        assert stored.document_job["runtime_results"][-1]["status"] == TURN_WAITING_TOOL_PERMISSION
        assert stored.document_job["runtime_results"][-1]["continuation_ref"] == {
            "continuation_id": "permission-1"
        }
        permission_blocks = [
            block
            for message in session.messages
            for block in message.blocks
            if block.data.get("interaction_type") == "permission"
        ]
        assert len(permission_blocks) == 1
        assert permission_blocks[0].data["actions"] == []
        assert "没有可验证的 Tool Call 恢复通道" in permission_blocks[0].text
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_permission_card_does_not_offer_unverifiable_cursor_resume(qapp, tmp_path):
    gateway = _ResumableSideChannelGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-resumable-runtime")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        assert panel._send_message("inspect material") is True
        for _attempt in range(200):
            qapp.processEvents()
            if not panel._turn_workers:
                break
            QTest.qWait(10)

        permission_card = next(
            card
            for card in panel._message_host.findChildren(AssistantInteractionCard)
            if card.interaction_type == "permission"
        )
        assert permission_card._buttons == []
        assert len(gateway.requests) == 1
        session = panel._active_session
        assert session is not None
        assert session.pending_continuation["continuation_id"] == "resume-1"
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_question_card_answer_resumes_the_existing_provider_continuation(
    qapp,
    tmp_path,
):
    gateway = _QuestionContinuationGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-question-continuation")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        assert panel._send_message("请比较学生卷和答案卷的区别") is True
        for _attempt in range(200):
            qapp.processEvents()
            if not panel._turn_workers:
                break
            QTest.qWait(10)

        question = next(
            card
            for card in panel._message_host.findChildren(AssistantInteractionCard)
            if card.interaction_type == "question"
        )
        question._choice_buttons[0].click()
        question._choice_buttons[1].click()
        question._question_submit.click()
        for _attempt in range(200):
            qapp.processEvents()
            if len(gateway.requests) == 2 and not panel._turn_workers:
                break
            QTest.qWait(10)

        assert len(gateway.requests) == 2
        assert gateway.requests[1].metadata["conversation_cursor"] == "question-1"
        assert gateway.requests[1].messages[-1]["content"] == "学生卷；答案卷"
        assert panel._active_session is not None
        assert panel._active_session.pending_continuation == {}
        assert panel._active_session.active_plan == {}
        assert "plan_candidate" not in panel._active_session.document_job
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_question_continuation_can_cancel_pending_document_request(
    qapp,
    tmp_path,
):
    gateway = _QuestionContinuationGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-question-cancel")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        assert panel._send_message("请比较学生卷和答案卷的区别") is True
        _wait_until(
            qapp,
            lambda: not panel._turn_workers,
            timeout_seconds=5,
            label="question before cancellation",
        )

        assert panel._send_message("取消") is True
        _wait_until(
            qapp,
            lambda: not panel._turn_workers,
            timeout_seconds=5,
            label="document request cancellation",
        )

        assert panel._active_session is not None
        assert panel._active_session.pending_continuation == {}
        assert "plan_candidate" not in panel._active_session.document_job
        assert "pending_document_request" not in panel._active_session.document_job
        assert panel._active_session.active_plan == {}
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_rejected_composer_submission_preserves_the_draft(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        composer = panel._empty_input
        emitted: list[str] = []
        composer.message_sent.connect(emitted.append)
        composer.set_submission_handler(lambda _text: False)
        composer.set_text("draft that must survive")

        composer._on_send()
        qapp.processEvents()

        assert composer.get_text() == "draft that must survive"
        assert emitted == []
    finally:
        panel.close()


def test_cross_session_turns_run_in_isolated_lanes_and_keep_origin_ownership(
    qapp,
    tmp_path,
):
    gateway = _GatedStreamGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-operation-owner")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        assert panel._send_message("origin request") is True
        for _attempt in range(100):
            qapp.processEvents()
            if panel._turn_preview_text:
                break
            QTest.qWait(10)

        assert panel._turn_worker is not None
        origin_id = panel._active_session.session_id
        panel.new_session()
        draft_session_id = panel._active_session.session_id
        assert draft_session_id != origin_id
        assert panel._empty_input._busy is False
        assert panel._empty_input._submission_gate == ""

        panel._empty_input.set_text("second session draft")
        panel._empty_input._on_send()
        qapp.processEvents()

        assert panel._empty_input.get_text() == ""
        assert len(coordinator.load_session(draft_session_id).messages) == 1
        assert set(panel._turn_workers) == {origin_id, draft_session_id}
        assert gateway.cancelled is False

        panel._open_session_by_id(origin_id)
        assert panel._composer._busy is True
        assert panel._composer._submission_gate == ""
        assert panel._turn_preview_widget is not None
        assert panel._turn_preview_widget._markdown.toPlainText() == "第一段"

        gateway.release.set()
        for _attempt in range(200):
            qapp.processEvents()
            if not panel._turn_workers:
                break
            QTest.qWait(10)
        assert not panel._turn_workers
        assert panel._active_session.session_id == origin_id

        panel._open_session_by_id(draft_session_id)
        assert panel._empty_input._submission_gate == ""
        assert panel._composer.get_text() == ""
        assert panel._composer._busy is False
    finally:
        gateway.release.set()
        panel.shutdown_active_execution(5000)
        panel.close()


def test_active_conversation_surface_retains_first_level_grid_identity():
    assert "paintEvent" in AssistantConversationSurface.__dict__


def test_message_sources_and_attachments_are_actionable(qapp):
    message = AssistantConversationMessage(
        text="Answer with [details](https://example.com/docs).",
        role="assistant",
        source_refs=(
            {"title": "Design notes", "url": "https://example.com/source"},
        ),
    )
    requested: list[dict[str, object]] = []
    message.source_requested.connect(requested.append)
    try:
        strip = message.findChild(AssistantSourceStrip)
        assert strip is not None
        assert strip._label.text() == "来源 · 1"
        strip._buttons[0].click()
        qapp.processEvents()
        assert requested == [
            {"title": "Design notes", "url": "https://example.com/source"}
        ]
        flags = message._markdown.textInteractionFlags()
        assert flags & Qt.LinksAccessibleByMouse
        assert flags & Qt.LinksAccessibleByKeyboard
    finally:
        message.close()


def test_assistant_message_quote_action_returns_text_to_composer(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel.new_session()
        assert panel._active_session is not None
        assistant = AssistantMessage.text(
            role=ROLE_ASSISTANT,
            text="第一条建议\n第二条建议",
        )
        panel._active_session = panel._coordinator.append_message(
            panel._active_session,
            assistant,
        )
        panel._render_active_session()

        panel._handle_message_action("quote", assistant.message_id)

        assert panel._composer.get_text() == "> 第一条建议\n> 第二条建议\n\n"
    finally:
        panel.close()


def test_live_message_appends_plain_text_without_reparsing_accumulated_markdown(
    qapp,
    monkeypatch,
):
    message = AssistantConversationMessage(
        text="first",
        role="assistant",
        live=True,
        status_text="模型正在生成",
    )
    try:
        monkeypatch.setattr(
            message._markdown,
            "set_markdown",
            lambda _text: (_ for _ in ()).throw(AssertionError("must not reparse")),
        )
        message.set_live_state(text="first second", status_text="正在生成回复")
        message.set_live_state(text="first second third", status_text="正在生成回复")
        qapp.processEvents()
        assert message._markdown.toPlainText() == "first second third"
    finally:
        message.close()


def test_message_reference_opens_supported_web_and_local_targets(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel = _panel(tmp_path)
    local_path = tmp_path / "evidence.txt"
    local_path.write_text("evidence", encoding="utf-8")
    opened: list[str] = []
    monkeypatch.setattr(
        "src.assistant.ui.assistant_panel.QDesktopServices.openUrl",
        lambda url: opened.append(url.toString()) or True,
    )
    try:
        panel._open_message_reference({"url": "https://example.com/evidence"})
        panel._open_message_reference({"path": str(local_path)})
        qapp.processEvents()
        assert opened[0] == "https://example.com/evidence"
        assert opened[1].startswith("file:")
        assert "evidence.txt" in opened[1]
    finally:
        panel.close()


def test_message_rerender_preserves_reader_scroll_when_not_following_latest(
    qapp,
    tmp_path,
):
    panel = _panel(tmp_path)
    panel.resize(900, 560)
    panel.show()
    try:
        panel.new_session()
        session = panel._active_session
        assert session is not None
        for index in range(24):
            role = ROLE_ASSISTANT if index % 2 else "user"
            session = panel._coordinator.append_message(
                session,
                AssistantMessage.text(
                    role=role,
                    text=f"message {index} " + ("long content " * 18),
                ),
            )
        panel._active_session = session
        panel._render_active_session()
        qapp.processEvents()
        QTest.qWait(20)
        qapp.processEvents()
        bar = panel._message_scroll.verticalScrollBar()
        assert bar.maximum() > 64
        bar.setValue(0)

        session = panel._coordinator.append_message(
            session,
            AssistantMessage.text(role=ROLE_ASSISTANT, text="new background update"),
        )
        panel._active_session = session
        panel._render_active_session()
        qapp.processEvents()
        QTest.qWait(20)
        qapp.processEvents()

        assert bar.value() == 0
        assert panel._jump_latest_button.isVisible()
    finally:
        panel.close()


def test_assistant_session_draft_and_management_round_trip(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel.new_session()
        first = panel._active_session
        assert first is not None
        panel._empty_input.set_text("尚未发送的要求")
        qapp.processEvents()
        assert panel._coordinator.load_session(first.session_id).draft_text == "尚未发送的要求"

        panel.new_session()
        second = panel._active_session
        assert second is not None
        first_item = next(
            panel._session_list.item(index)
            for index in range(panel._session_list.count())
            if panel._session_list.item(index).data(Qt.UserRole) == first.session_id
        )
        panel._open_session_item(first_item)
        assert panel._empty_input.get_text() == "尚未发送的要求"

        assert panel._rename_session(first.session_id, "固定的方案讨论")
        assert panel._set_session_pinned(first.session_id, True)
        restored = panel._coordinator.load_session(first.session_id)
        assert restored.title == "固定的方案讨论"
        assert restored.pinned is True
        assert panel._session_sidebar.pinned_list.item(0).data(Qt.UserRole) == first.session_id

        assert panel._delete_session(first.session_id)
        assert panel._active_session is not None
        assert panel._active_session.session_id == second.session_id
    finally:
        panel.close()


def test_assistant_composer_enter_sends_and_shift_enter_adds_newline(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel.new_session()
        editor = panel._empty_input._text_edit
        editor.setFocus()
        editor.setPlainText("第一行")
        editor.moveCursor(QTextCursor.End)
        QTest.keyClick(editor, Qt.Key_Return, Qt.ShiftModifier)
        editor.insertPlainText("第二行")

        assert panel._active_session is not None
        assert panel._active_session.messages == ()
        assert editor.toPlainText() == "第一行\n第二行"

        QTest.keyClick(editor, Qt.Key_Return)
        for _attempt in range(100):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)

        assert panel._active_session is not None
        assert panel._active_session.messages[0].visible_text() == "第一行\n第二行"
        assert editor.toPlainText() == ""
    finally:
        panel.close()


def test_assistant_renders_typed_interaction_card_in_message_stream(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel.new_session()
        assert panel._active_session is not None
        interaction = AssistantMessage.interaction(
            role=ROLE_ASSISTANT,
            interaction_type="disclosure",
            title="将发送以下数据",
            body="1 个文档，共 1200 个字符。",
            payload={"actions": [{"id": "allow_once", "label": "允许一次"}]},
        )
        panel._active_session = panel._coordinator.append_message(
            panel._active_session,
            interaction,
        )
        panel._render_active_session()
        qapp.processEvents()

        cards = panel.findChildren(AssistantInteractionCard)
        assert len(cards) == 1
        assert cards[0].interaction_type == "disclosure"
        assert cards[0]._buttons[0].text() == "允许一次"
    finally:
        panel.close()


def test_execution_completion_persists_and_renders_every_output_file(qapp, tmp_path):
    panel = _panel(tmp_path)
    student = tmp_path / "学生卷.docx"
    answer = tmp_path / "答案卷.docx"
    student.write_bytes(b"student")
    answer.write_bytes(b"answer")

    class _FinishedWorker:
        session_id = ""
        execution_id = "exam-execution"

        def deleteLater(self):
            return None

    try:
        panel.new_session()
        assert panel._active_session is not None
        worker = _FinishedWorker()
        worker.session_id = panel._active_session.session_id
        panel._execution_worker = worker

        panel._on_execution_finished(
            {
                "status": "success",
                "output_path": str(student),
                "output_paths": {
                    "student": str(student),
                    "answer_key": str(answer),
                },
                "assistant_input_hash_unchanged": True,
            }
        )
        qapp.processEvents()

        assert panel._active_session is not None
        result_message = panel._active_session.messages[-1]
        references = result_message.blocks[0].data["references"]
        assert [item["artifact_key"] for item in references] == [
            "student",
            "answer_key",
        ]
        artifact_cards = [
            card
            for card in panel.findChildren(AssistantInteractionCard)
            if card.interaction_type == "artifact"
        ]
        assert artifact_cards
        assert len(artifact_cards[-1]._file_cards) == 2
    finally:
        panel.close()


def test_assistant_responsive_layout_hides_secondary_rails(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel.show()
        assert panel.minimumSizeHint().width() < 900
        panel.resize(850, 600)
        qapp.processEvents()
        assert panel._responsive_mode == "compact"
        assert panel._session_rail.isHidden()
        assert panel._context_rail.isHidden()

        panel._toggle_session_rail()
        qapp.processEvents()
        assert panel._session_drawer is not None
        assert panel._session_drawer.isVisible()
        assert panel._session_rail.window() is panel._session_drawer
        panel._session_drawer.close()

        panel.resize(500, 600)
        qapp.processEvents()
        assert panel._responsive_mode == "compact"
        assert panel._conversation_title.isHidden()
        assert panel._provider_combo.isHidden()

        panel.resize(1200, 700)
        qapp.processEvents()
        assert panel._responsive_mode == "wide"
        assert panel._session_rail.isVisible()
        assert panel._context_rail.isVisible()
    finally:
        panel.close()


def test_existing_docx_path_a_runs_from_local_plan_to_new_output_without_provider(
    qapp,
    tmp_path,
):
    source = tmp_path / "local-format-source.docx"
    document = Document()
    document.add_heading("项目报告", level=1)
    document.add_paragraph("This input must remain unchanged.")
    document.save(source)
    original_bytes = source.read_bytes()
    gateway = _FormatReferenceGateway()
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=AssistantSessionCoordinator(
            AssistantSessionStore(tmp_path / "assistant-path-a-e2e")
        ),
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._on_composer_document_selected(str(source))
        assert panel._send_message("统一这份 Word 文档格式") is True
        assert panel._active_session is not None
        assert panel._active_session.document_job["status"] == "plan_ready"
        assert gateway.requests == []

        panel._handle_card_action("preflight", {})
        _wait_until(
            qapp,
            lambda: panel._preflight_worker is None,
            timeout_seconds=20,
            label="path A preflight",
        )
        assert panel._active_session.document_job["status"] == (
            "needs_execution_approval"
        )

        panel._handle_card_action("approve_execute", {})
        _wait_until(
            qapp,
            lambda: panel._execution_worker is None,
            timeout_seconds=60,
            label="path A local production",
        )

        session = panel._active_session
        assert session is not None
        assert session.document_job["status"] == "success", (
            session.document_job.get("result", {}).get("error_text")
        )
        output_path = Path(session.document_job["primary_output_path"])
        assert output_path.is_file()
        assert output_path.resolve() != source.resolve()
        assert output_path.suffix.casefold() == ".docx"
        assert source.read_bytes() == original_bytes
        assert gateway.requests == []
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_prompt_generation_flows_through_draft_preflight_approval_and_output(qapp, tmp_path):
    panel = _panel(tmp_path)
    try:
        panel._send_message("生成一份项目报告")
        for _attempt in range(100):
            qapp.processEvents()
            if panel._turn_worker is None:
                break
            QTest.qWait(10)
        assert panel._active_session is not None
        assert panel._active_session.active_plan["scene_ref"]["generation_mode"] == "from_prompt"

        panel._handle_card_action("generate_content_draft", {})
        for _attempt in range(1500):
            qapp.processEvents()
            if panel._content_worker is None:
                break
            QTest.qWait(10)

        session = panel._active_session
        assert session is not None
        assert session.document_job["status"] == "content_draft_ready"
        assert Path(session.active_plan["input_document_ref"]["path"]).is_file()
        assert session.active_plan["content_fragment_refs"][0]["fragment_digest"]

        panel._handle_card_action("preflight", {})
        assert panel._preflight_worker is not None
        for _attempt in range(1500):
            qapp.processEvents()
            if panel._preflight_worker is None:
                break
            QTest.qWait(10)

        assert panel._active_session is not None
        assert panel._active_session.document_job["status"] == "needs_execution_approval"
        assert panel._active_session.document_job["preflight"]["ready"] is True
        artifact_messages = [
            message
            for message in panel._active_session.messages
            if message.blocks[0].data.get("title") == "内容草稿已生成"
        ]
        assert artifact_messages
        assert {
            action["id"]
            for action in artifact_messages[-1].blocks[0].data["actions"]
        } == {"runtime_open_reference", "generate_content_draft", "preflight"}

        panel._handle_card_action("approve_execute", {})
        for _attempt in range(2000):
            qapp.processEvents()
            if panel._execution_worker is None:
                break
            QTest.qWait(10)

        assert panel._active_session is not None
        assert panel._active_session.document_job["status"] == "success", (
            panel._active_session.document_job.get("result", {}).get("error_text")
        )
        output_path = Path(panel._active_session.document_job["primary_output_path"])
        assert output_path.is_file()
        assert output_path.suffix.casefold() == ".docx"
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()


def test_exam_prompt_flows_through_ai_draft_and_assembly_to_both_deliveries(
    qapp,
    tmp_path,
):
    gateway = _ExamGenerationGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "assistant-exam-chain")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._send_message("生成一份七年级数学试卷，包含选择题和计算题")
        _wait_until(
            qapp,
            lambda: panel._turn_worker is None,
            timeout_seconds=5,
            label="exam request turn",
        )

        assert panel._active_session is not None
        assert panel._active_session.active_plan["generation_contract"][
            "artifact_kind"
        ] == ARTIFACT_KIND_EXAM

        panel._handle_card_action("generate_content_draft", {})
        _wait_until(
            qapp,
            lambda: panel._content_worker is None,
            timeout_seconds=20,
            label="exam content draft",
        )
        assert panel._active_session.document_job["status"] == "content_draft_ready"

        panel._handle_card_action("preflight", {})
        _wait_until(
            qapp,
            lambda: panel._preflight_worker is None,
            timeout_seconds=20,
            label="exam preflight",
        )
        assert (
            panel._active_session.document_job["status"]
            == "needs_execution_approval"
        )

        panel._handle_card_action("approve_execute", {})
        _wait_until(
            qapp,
            lambda: panel._execution_worker is None,
            timeout_seconds=60,
            label="exam assembly",
        )

        assert panel._active_session.document_job["status"] == "success", (
            panel._active_session.document_job.get("result", {}).get("error_text")
        )
        output_paths = panel._active_session.document_job["result"]["output_paths"]
        assert {"student", "answer_key"} <= set(output_paths)
        assert all(Path(output_paths[key]).is_file() for key in ("student", "answer_key"))
        references = panel._active_session.messages[-1].blocks[0].data["references"]
        assert {"student", "answer_key"} <= {
            item["artifact_key"] for item in references
        }
    finally:
        panel.shutdown_active_execution(5000)
        panel.close()
