from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from src.config.feature_configs import OutputConfig
from src.config.scene import DeliveryPreset, SceneWorkspace
from src.config.template import TemplateConfig
from src.pipeline.result import PipelineResult
import src.shared.engine.official_document_batch_history as history_module
import src.services.production_runtime.batch_reporting as batch_reporting
import src.services.production_runtime.delivery_reporting as delivery_reporting
import src.services.production_runtime.delivery_runtime as delivery_runtime
import src.services.production_runtime.execution_runtime as execution_runtime
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner


def _delivery_group_scene() -> SceneWorkspace:
    return SceneWorkspace(
        template_id="base_template",
        compatible_template_ids=["base_template", "target_template"],
        default_delivery_preset_id="first",
        delivery_presets=[
            DeliveryPreset(
                preset_id="first",
                target_template_id="target_template",
                filename_template="{stem}_first.docx",
                artifacts=OutputConfig(final_docx=True),
            ),
            DeliveryPreset(
                preset_id="second",
                target_template_id="base_template",
                filename_template="{stem}_second.docx",
                artifacts=OutputConfig(final_docx=True),
            ),
        ],
    )


def _install_successful_group_pipeline(monkeypatch, source: Path) -> None:
    class _PublishingPipeline:
        def __init__(self, **kwargs):
            self.config = kwargs["config"]
            self.output_dir = Path(kwargs["output_dir"])

        def execute(self, _doc_path):
            preset = self.config.delivery_presets[0]
            path = self.output_dir / f"{source.stem}_{preset.preset_id}.docx"
            path.parent.mkdir(parents=True, exist_ok=True)
            Document().save(path)
            return PipelineResult(
                success=True,
                status="success",
                output_paths={preset.preset_id: str(path)},
                config=self.config,
            )

    monkeypatch.setattr(execution_runtime, "Pipeline", _PublishingPipeline)
    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )


def _install_safe_group_artifact_writers(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(
        delivery_reporting,
        "_write_compare_docx_artifacts",
        lambda *_args, **_kwargs: {},
    )
    monkeypatch.setattr(
        delivery_reporting,
        "write_delivery_reports",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        delivery_reporting,
        "write_structured_intermediates",
        lambda *_args, **_kwargs: {},
    )
    monkeypatch.setattr(
        delivery_runtime,
        "write_material_manifest",
        lambda *_args, **_kwargs: {
            "material": str(tmp_path / "material-manifest.json")
        },
    )
    monkeypatch.setattr(
        delivery_runtime,
        "write_material_package_artifacts",
        lambda *_args, **_kwargs: {},
    )


@pytest.mark.parametrize(
    ("writer_owner", "writer_name", "failure_kind", "success_value"),
    [
        (delivery_reporting, "_write_compare_docx_artifacts", "compare_docx", {}),
        (delivery_reporting, "write_delivery_reports", "reports", []),
        (delivery_reporting, "write_structured_intermediates", "intermediates", {}),
        (delivery_runtime, "write_material_manifest", "material_manifest", {}),
        (delivery_runtime, "write_material_package_artifacts", "material_package", {}),
    ],
)
def test_delivery_group_auxiliary_failure_keeps_all_published_core_outputs(
    tmp_path,
    monkeypatch,
    writer_owner,
    writer_name,
    failure_kind,
    success_value,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    _install_successful_group_pipeline(monkeypatch, source)
    _install_safe_group_artifact_writers(monkeypatch, tmp_path)
    calls = {"count": 0}
    failure_path = tmp_path / f"{failure_kind}.artifact"

    def _fail_once(*_args, **_kwargs):
        calls["count"] += 1
        if calls["count"] == 1:
            raise OSError(5, "simulated auxiliary failure", str(failure_path))
        return success_value.copy() if isinstance(success_value, dict) else list(success_value)

    monkeypatch.setattr(writer_owner, writer_name, _fail_once)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_delivery_group_scene(),
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "partial_success"
    assert set(payload["output_paths"]) == {"first", "second"}
    assert all(Path(path).is_file() for path in payload["output_paths"].values())
    assert payload["failed_count"] == 0
    assert payload["artifact_failure_count"] == 1
    assert len(payload["artifact_failures"]) == 1
    failure = payload["artifact_failures"][0]
    assert failure["kind"] == failure_kind
    assert failure["path"] == str(failure_path)
    assert "simulated auxiliary failure" in failure["error"]
    assert failure_kind in payload["error_text"]


def test_delivery_group_cancellation_keeps_prior_failures_and_all_core_outputs(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    _install_safe_group_artifact_writers(monkeypatch, tmp_path)
    calls = {"pipeline": 0, "compare": 0}

    class _CancellingSecondPipeline:
        def __init__(self, **kwargs):
            self.config = kwargs["config"]
            self.output_dir = Path(kwargs["output_dir"])

        def execute(self, _doc_path):
            calls["pipeline"] += 1
            preset = self.config.delivery_presets[0]
            path = self.output_dir / f"{source.stem}_{preset.preset_id}.docx"
            path.parent.mkdir(parents=True, exist_ok=True)
            Document().save(path)
            cancelled = calls["pipeline"] == 2
            return PipelineResult(
                success=not cancelled,
                status="cancelled" if cancelled else "success",
                cancelled=cancelled,
                output_paths={preset.preset_id: str(path)},
                config=self.config,
            )

    def _fail_first_compare(*_args, **_kwargs):
        calls["compare"] += 1
        if calls["compare"] == 1:
            raise OSError(5, "simulated compare failure", str(tmp_path / "compare"))
        return {}

    monkeypatch.setattr(execution_runtime, "Pipeline", _CancellingSecondPipeline)
    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )
    monkeypatch.setattr(
        delivery_reporting,
        "_write_compare_docx_artifacts",
        _fail_first_compare,
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_delivery_group_scene(),
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "cancelled"
    assert set(payload["output_paths"]) == {"first", "second"}
    assert all(Path(path).is_file() for path in payload["output_paths"].values())
    assert payload["failed_count"] == 0
    assert payload["artifact_failure_count"] == 1
    assert payload["artifact_failures"][0]["kind"] == "compare_docx"


@pytest.mark.parametrize(
    ("cancel_on_check", "expected_pipeline_calls", "expected_evidence_groups"),
    [
        (1, 0, ["first"]),
        (3, 1, ["first", "second"]),
    ],
)
def test_delivery_group_pre_execution_cancel_publishes_current_group_evidence(
    tmp_path,
    monkeypatch,
    cancel_on_check,
    expected_pipeline_calls,
    expected_evidence_groups,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = _delivery_group_scene()
    for preset in scene.delivery_presets:
        preset.artifacts.report_json = True
        preset.include_structured_intermediate = True

    calls = {
        "cancel": 0,
        "pipeline": 0,
        "reports": [],
        "intermediates": [],
        "compare": [],
        "material": [],
    }

    class _SuccessfulPipeline:
        def __init__(self, **kwargs):
            self.config = kwargs["config"]
            self.output_dir = Path(kwargs["output_dir"])

        def execute(self, _doc_path):
            calls["pipeline"] += 1
            preset = self.config.delivery_presets[0]
            output_path = self.output_dir / f"{source.stem}_{preset.preset_id}.docx"
            output_path.parent.mkdir(parents=True, exist_ok=True)
            Document().save(output_path)
            return PipelineResult(
                success=True,
                status="success",
                output_paths={preset.preset_id: str(output_path)},
                config=self.config,
            )

    def _cancel_check():
        calls["cancel"] += 1
        return calls["cancel"] >= cancel_on_check

    def _write_reports(_result, **kwargs):
        preset_id = kwargs["config"].delivery_presets[0].preset_id
        calls["reports"].append(preset_id)
        return [str(tmp_path / f"{preset_id}_changes.json")]

    def _write_intermediates(_result, **kwargs):
        preset_id = kwargs["config"].delivery_presets[0].preset_id
        calls["intermediates"].append(preset_id)
        return {preset_id: str(tmp_path / f"{preset_id}_intermediate.json")}

    def _write_compare(_result, **kwargs):
        preset_id = kwargs["config"].delivery_presets[0].preset_id
        calls["compare"].append(preset_id)
        return {}

    def _write_material(**kwargs):
        preset_id = kwargs["config"].delivery_presets[0].preset_id
        calls["material"].append(preset_id)
        return {}

    monkeypatch.setattr(execution_runtime, "Pipeline", _SuccessfulPipeline)
    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )
    monkeypatch.setattr(delivery_reporting, "write_delivery_reports", _write_reports)
    monkeypatch.setattr(
        delivery_reporting,
        "write_structured_intermediates",
        _write_intermediates,
    )
    monkeypatch.setattr(
        delivery_reporting,
        "_write_compare_docx_artifacts",
        _write_compare,
    )
    monkeypatch.setattr(
        delivery_runtime,
        "write_material_manifest",
        _write_material,
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, _cancel_check)

    assert payload["status"] == "cancelled"
    assert calls["pipeline"] == expected_pipeline_calls
    assert calls["reports"] == expected_evidence_groups
    assert calls["intermediates"] == expected_evidence_groups
    assert calls["compare"] == expected_evidence_groups[:expected_pipeline_calls]
    assert calls["material"] == []
    assert [Path(path).stem for path in payload["report_paths"]] == [
        f"{preset_id}_changes" for preset_id in expected_evidence_groups
    ]
    assert list(payload["intermediate_paths"]) == expected_evidence_groups


def test_failed_delivery_groups_still_publish_configured_evidence(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = _delivery_group_scene()
    for preset in scene.delivery_presets:
        preset.artifacts.report_json = True
        preset.include_structured_intermediate = True

    class _FailingPipeline:
        def __init__(self, **kwargs):
            self.config = kwargs["config"]

        def execute(self, _doc_path):
            return PipelineResult(
                success=False,
                status="failed",
                error="simulated group failure",
                config=self.config,
            )

    report_calls: list[str] = []
    intermediate_calls: list[str] = []

    def _write_reports(_result, **kwargs):
        preset_id = kwargs["config"].delivery_presets[0].preset_id
        report_calls.append(preset_id)
        return [str(tmp_path / f"{preset_id}_failed.json")]

    def _write_intermediates(_result, **kwargs):
        preset_id = kwargs["config"].delivery_presets[0].preset_id
        intermediate_calls.append(preset_id)
        return {preset_id: str(tmp_path / f"{preset_id}_failed_intermediate.json")}

    monkeypatch.setattr(execution_runtime, "Pipeline", _FailingPipeline)
    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )
    monkeypatch.setattr(delivery_reporting, "write_delivery_reports", _write_reports)
    monkeypatch.setattr(
        delivery_reporting,
        "write_structured_intermediates",
        _write_intermediates,
    )
    monkeypatch.setattr(
        delivery_runtime,
        "write_material_manifest",
        lambda **_kwargs: {},
    )
    monkeypatch.setattr(
        delivery_runtime,
        "write_material_package_artifacts",
        lambda **_kwargs: {},
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert report_calls == ["first", "second"]
    assert intermediate_calls == ["first", "second"]
    assert {Path(path).name for path in payload["report_paths"]} == {
        "first_failed.json",
        "second_failed.json",
    }
    assert set(payload["intermediate_paths"]) == {"first", "second"}


def _batch_payload(tmp_path: Path, status: str) -> dict[str, object]:
    core_path = tmp_path / "core.docx"
    Document().save(core_path)
    return batch_reporting.build_batch_payload(
        status,
        [
            {
                "status": "success" if status == "success" else status,
                "profile_id": "profile-a",
                "profile_name": "Profile A",
                "output_path": str(core_path),
                "output_paths": {"final": str(core_path)},
                "failed_count": 0,
            }
        ],
        error_text="business failure" if status == "failed" else "",
    )


def _fail_staged_markdown(monkeypatch) -> None:
    original = batch_reporting.atomic_write_text

    def _write(path, text, *, encoding="utf-8"):
        if Path(path).suffix == ".md":
            raise OSError(5, "simulated markdown failure", str(path))
        return original(path, text, encoding=encoding)

    monkeypatch.setattr(batch_reporting, "atomic_write_text", _write)


@pytest.mark.parametrize(
    ("business_status", "expected_status"),
    [
        ("success", "partial_success"),
        ("failed", "failed"),
        ("cancelled", "cancelled"),
    ],
)
def test_generic_batch_report_failure_preserves_business_status_and_core_evidence(
    tmp_path,
    monkeypatch,
    business_status,
    expected_status,
):
    payload = _batch_payload(tmp_path, business_status)
    business_failed_count = payload["failed_count"]
    _fail_staged_markdown(monkeypatch)

    batch_reporting.attach_batch_reports(
        payload,
        tmp_path / "reports",
        tmp_path / "source.docx",
    )

    assert payload["status"] == expected_status
    assert payload["items"][0]["profile_id"] == "profile-a"
    assert Path(payload["output_paths"]["profile-a:final"]).is_file()
    assert payload["failed_count"] == business_failed_count
    assert payload["artifact_failure_count"] == 1
    assert payload["artifact_failures"][0]["kind"] == "batch_reports"
    assert "simulated markdown failure" in payload["artifact_failures"][0]["error"]
    assert not (tmp_path / "reports" / "source_batch_report.json").exists()
    assert not (tmp_path / "reports" / "source_batch_report.md").exists()


def test_batch_payload_aggregates_profile_outputs_and_artifact_failures(tmp_path):
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    Document().save(first)
    Document().save(second)

    payload = batch_reporting.build_batch_payload(
        "partial_success",
        [
            {
                "status": "partial_success",
                "profile_id": "profile-a",
                "profile_name": "Profile A",
                "output_path": str(first),
                "output_paths": {
                    "final": str(first),
                    "review": str(second),
                },
                "artifact_failure_count": 1,
                "artifact_failures": [
                    {
                        "kind": "reports",
                        "path": str(tmp_path / "failed-report.json"),
                        "error_type": "OSError",
                        "error": "simulated child report failure",
                    }
                ],
            }
        ],
        error_text="",
    )

    assert payload["output_paths"] == {
        "profile-a:final": str(first),
        "profile-a:review": str(second),
    }
    assert payload["artifact_failure_count"] == 1
    assert payload["artifact_failures"][0]["profile_id"] == "profile-a"
    assert "simulated child report failure" in payload["error_text"]


def test_official_batch_report_failure_commits_no_history(tmp_path, monkeypatch):
    payload = _batch_payload(tmp_path, "success")
    payload.update(
        {
            "batch_source_kind": "official_document_table",
            "batch_source_path": str(tmp_path / "source.csv"),
        }
    )
    _fail_staged_markdown(monkeypatch)

    batch_reporting.attach_batch_reports(
        payload,
        tmp_path / "reports",
        tmp_path / "source.csv",
    )

    assert payload["status"] == "partial_success"
    assert payload["items"][0]["profile_id"] == "profile-a"
    assert not (tmp_path / "reports" / "source_batch_report.json").exists()
    assert not (tmp_path / "reports" / "source_batch_report.md").exists()
    history_dir = tmp_path / "reports" / "batch_history"
    assert not history_dir.exists() or not list(history_dir.glob("*.json"))


def test_batch_report_keeps_business_and_artifact_failure_counts_separate(tmp_path):
    payload = _batch_payload(tmp_path, "success")
    payload.update(
        {
            "status": "partial_success",
            "artifact_failure_count": 1,
            "artifact_failures": [
                {
                    "kind": "reports",
                    "path": "earlier-report.json",
                    "error": "earlier report failed",
                }
            ],
        }
    )

    batch_reporting.attach_batch_reports(
        payload,
        tmp_path / "reports",
        tmp_path / "source.docx",
    )

    report = json.loads(
        (tmp_path / "reports" / "source_batch_report.json").read_text(
            encoding="utf-8"
        )
    )
    markdown = (
        tmp_path / "reports" / "source_batch_report.md"
    ).read_text(encoding="utf-8")
    assert report["failed_count"] == 0
    assert report["artifact_failure_count"] == 1
    assert "- Failed count: 0" in markdown
    assert "- Artifact failure count: 1" in markdown


def test_official_history_failure_occurs_after_both_main_reports_publish(
    tmp_path,
    monkeypatch,
):
    payload = _batch_payload(tmp_path, "success")
    payload.update(
        {
            "batch_source_kind": "official_document_table",
            "batch_source_path": str(tmp_path / "source.csv"),
        }
    )
    original_atomic_history_write = history_module._atomic_write_json

    def _fail_history(path, document):
        if Path(path).name != "batch_history_index.json":
            raise OSError(5, "simulated history failure", str(path))
        return original_atomic_history_write(path, document)

    monkeypatch.setattr(history_module, "_atomic_write_json", _fail_history)

    batch_reporting.attach_batch_reports(
        payload,
        tmp_path / "reports",
        tmp_path / "source.csv",
    )

    json_path = tmp_path / "reports" / "source_batch_report.json"
    markdown_path = tmp_path / "reports" / "source_batch_report.md"
    assert json_path.is_file()
    assert markdown_path.is_file()
    assert payload["status"] == "partial_success"
    assert payload["batch_report_paths"] == [str(json_path), str(markdown_path)]
    assert payload["items"][0]["profile_id"] == "profile-a"
    assert payload["artifact_failures"][0]["kind"] == "batch_reports"
    assert "simulated history failure" in payload["error_text"]
    report = json.loads(json_path.read_text(encoding="utf-8"))
    assert "batch_history" not in report
    assert "history_run_id" not in report.get("batch_isolation", {})
    assert "history_path" not in report.get("batch_isolation", {})
    markdown = markdown_path.read_text(encoding="utf-8")
    assert "Batch run id:" not in markdown
    assert "History path:" not in markdown
    assert payload["failed_count"] == 0
    assert payload["artifact_failure_count"] == 1
    history_dir = tmp_path / "reports" / "batch_history"
    assert not list(history_dir.glob("run_*.json"))
