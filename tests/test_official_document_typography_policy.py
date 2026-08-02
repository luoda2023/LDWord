from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import RGBColor

from src.shared.engine.font_resolver import canonicalize_font_name
from src.shared.engine.official_document_master import (
    write_official_master_family_docx,
)


def _east_asia_font(run) -> str:
    r_pr = run._element.rPr
    if r_pr is None or r_pr.rFonts is None:
        return ""
    return str(r_pr.rFonts.get(qn("w:eastAsia")) or "")


def _style_east_asia_font(style) -> str:
    r_pr = style.element.rPr
    if r_pr is None or r_pr.rFonts is None:
        return ""
    return str(r_pr.rFonts.get(qn("w:eastAsia")) or "")


def test_builtin_official_master_uses_semantic_non_founder_fonts(tmp_path):
    path = write_official_master_family_docx(
        tmp_path / "official.docx",
        family_id="common",
    )

    with ZipFile(path) as archive:
        xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in (
                "word/document.xml",
                "word/styles.xml",
                "word/fontTable.xml",
            )
        )
    assert "方正" not in xml
    assert "FZXiaoBiaoSong" not in xml

    document = Document(path)
    title = next(
        paragraph
        for paragraph in document.paragraphs
        if "{{@text:official_title}}" in paragraph.text
    )
    assert canonicalize_font_name(_east_asia_font(title.runs[0])) == "小标宋"
    assert (
        canonicalize_font_name(
            _style_east_asia_font(document.styles["Heading 1"])
        )
        == "黑体"
    )
    assert document.styles["Heading 1"].font.color.rgb == RGBColor(0, 0, 0)
    assert (
        canonicalize_font_name(
            _style_east_asia_font(document.styles["Heading 2"])
        )
        == "楷体"
    )
    for style_name in ("Heading 3", "Heading 4"):
        assert (
            canonicalize_font_name(
                _style_east_asia_font(document.styles[style_name])
            )
            == "仿宋"
        )


def test_large_redhead_mark_does_not_inherit_body_exact_line_height(tmp_path):
    path = write_official_master_family_docx(
        tmp_path / "official.docx",
        family_id="common",
    )

    document = Document(path)
    redhead = next(
        paragraph
        for paragraph in document.paragraphs
        if "{{@text:official_organization}}" in paragraph.text
    )
    run_size = redhead.runs[0].font.size.pt
    line_height = redhead.paragraph_format.line_spacing.pt

    assert run_size == 50
    assert redhead.paragraph_format.line_spacing_rule == WD_LINE_SPACING.AT_LEAST
    assert line_height >= run_size * 1.2


def test_common_master_anchors_the_atomic_imprint_to_the_text_area_bottom(
    tmp_path,
):
    path = write_official_master_family_docx(
        tmp_path / "official-common.docx",
        family_id="common",
    )

    document = Document(path)
    section = document.sections[0]
    title = next(
        paragraph
        for paragraph in document.paragraphs
        if "official_title" in paragraph.text
    )
    imprint = next(
        table
        for table in document.tables
        if "official_copy_scope" in "\n".join(
            cell.text for row in table.rows for cell in row.cells
        )
    )
    table_properties = imprint._tbl.tblPr
    position = table_properties.find(qn("w:tblpPr"))
    borders = table_properties.find(qn("w:tblBorders"))

    assert abs(section.footer_distance.mm - 23.5) < 0.1
    assert title.paragraph_format.space_before.pt == 56
    assert len(imprint.rows) == 1
    imprint_paragraphs = imprint.rows[0].cells[0].paragraphs
    assert imprint_paragraphs[0].text == "抄送：{{@text:official_copy_scope}}。"
    assert imprint_paragraphs[1].text == ""
    assert imprint_paragraphs[2].text == (
        "{{@text:official_printing_org}}\t"
        "{{@text:official_printing_date}}"
    )
    assert position is not None
    assert position.get(qn("w:vertAnchor")) == "margin"
    assert position.get(qn("w:tblpYSpec")) == "bottom"
    assert borders.find(qn("w:top")).get(qn("w:sz")) == "8"
    assert borders.find(qn("w:bottom")).get(qn("w:sz")) == "8"
    separator = imprint_paragraphs[1]._p.get_or_add_pPr().find(qn("w:pBdr"))
    assert separator.find(qn("w:bottom")).get(qn("w:sz")) == "6"
    assert imprint.rows[0]._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None


def test_upward_signer_uses_fangsong_label_and_kaiti_name(tmp_path):
    path = write_official_master_family_docx(
        tmp_path / "official-upward.docx",
        family_id="upward",
    )

    document = Document(path)
    signer = next(
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if "official_signer" in paragraph.text
    )

    text_runs = [run for run in signer.runs if run.text]
    assert text_runs[0].text == "签发人："
    assert text_runs[1].text == "{{@text:official_signer}}"
    assert canonicalize_font_name(_east_asia_font(text_runs[0])) == "仿宋"
    assert canonicalize_font_name(_east_asia_font(text_runs[1])) == "楷体"


def test_letter_and_order_masters_use_their_special_format_contracts(tmp_path):
    letter_path = write_official_master_family_docx(
        tmp_path / "official-letter.docx",
        family_id="letter",
    )
    order_path = write_official_master_family_docx(
        tmp_path / "official-order.docx",
        family_id="order",
    )

    letter = Document(letter_path)
    letter_mark = letter.paragraphs[0]
    mark_border = letter_mark._p.get_or_add_pPr().find(qn("w:pBdr"))
    footer_border = (
        letter.sections[0]
        .first_page_footer.paragraphs[0]
        ._p.get_or_add_pPr()
        .find(qn("w:pBdr"))
    )
    assert abs(letter.sections[0].footer_distance.mm - 20.0) < 0.1
    assert letter.sections[0].first_page_footer.paragraphs[0].text == ""
    assert mark_border.find(qn("w:bottom")).get(qn("w:val")) == (
        "thickThinSmallGap"
    )
    assert footer_border.find(qn("w:bottom")).get(qn("w:val")) == (
        "thinThickSmallGap"
    )
    assert any(
        "official_copy_scope" in cell.text
        for table in letter.tables
        for row in table.rows
        for cell in row.cells
    )

    order = Document(order_path)
    order_text = "\n".join(paragraph.text for paragraph in order.paragraphs)
    order_number = next(
        paragraph
        for paragraph in order.paragraphs
        if "official_document_no" in paragraph.text
    )
    signer = next(
        paragraph
        for paragraph in order.paragraphs
        if "official_signer" in paragraph.text
    )
    assert "签发人：" not in order_text
    assert order_number.paragraph_format.space_after.pt == 52
    assert canonicalize_font_name(_east_asia_font(signer.runs[0])) == "楷体"
