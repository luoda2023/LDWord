from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from src.config.builtin_scenes import create_builtin_scene
from src.config.builtin_templates import create_builtin_template
from src.services.execution_session import cleanup_execution_session_resources
from src.services import exam_markdown_source
from src.config.material_context import (
    MaterialExecutionContext,
    MaterialFieldResolution,
)
from src.config.object_preflight_evidence import build_object_preflight_evidence
from src.services.document_structure_evidence import (
    build_document_structure_evidence,
)
from src.services.production_runtime import exam_markdown_execution_input
from src.shared.engine.exam_question_schema import ExamMarkdownImportResult
from src.ui.panels.workbench import execution_thread_handle, execution_worker
from src.ui.panels.workbench.execution_session_controller import (
    WorkbenchExecutionSessionController,
)


COMPLETE_EXAM_MARKDOWN = """# 七年级数学单元测试

> 科目：数学　年级：七年级　考试时间：45 分钟　满分：5 分

## 一、选择题

1. 1 + 1 = （　　）（5 分）
   A. 1
   B. 2

## 答案速查

一、选择题
1. B
"""


def test_exam_markdown_import_to_dict_isolates_nested_payload():
    result = ExamMarkdownImportResult(
        payload={
            "sections": [
                {"questions": [{"prompt": "before", "answer": "A"}]}
            ]
        }
    )

    exported = result.to_dict()
    exported["payload"]["sections"][0]["questions"][0]["prompt"] = "after"

    assert result.payload["sections"][0]["questions"][0]["prompt"] == (
        "before"
    )
    assert exported["payload"] is not result.payload


def test_exam_source_projection_extends_an_approved_frozen_material_snapshot(
    tmp_path,
):
    source = tmp_path / "approved-generated-exam.md"
    source.write_text(COMPLETE_EXAM_MARKDOWN, encoding="utf-8")
    frozen_material = MaterialExecutionContext(mode_id="exam").with_frozen_field_resolution(
        MaterialFieldResolution(values={})
    )

    projected = exam_markdown_source.project_exam_markdown_source(
        source,
        frozen_material,
    ).material_context

    assert projected.field_values_frozen is True
    assert projected.resolved_entity_data()["exam_items"]
    assert projected.resolved_entity_data()["title"]
    assert frozen_material.resolved_entity_data() == {}


def _exam_scene():
    scene = create_builtin_scene("exam", mode_id="exam")
    scene.input_source_profile.failure_policy = "block"
    scene.input_source_profile.required_material_fields = [
        "title",
        "subject",
        "grade",
        "duration",
        "total_score",
    ]
    return scene


def _capture_runner(monkeypatch):
    class _Worker:
        def __init__(self, runner, parent=None):
            self.runner = runner

    class _Handle:
        def __init__(self, worker, parent=None):
            self.worker = worker

    monkeypatch.setattr(execution_worker, "ExecutionWorker", _Worker)
    monkeypatch.setattr(
        execution_thread_handle,
        "ThreadedExecutionHandle",
        _Handle,
    )


def test_controller_prepares_frozen_exam_markdown_once_before_block_gate(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "exam.md"
    source.write_text(COMPLETE_EXAM_MARKDOWN, encoding="utf-8")
    _capture_runner(monkeypatch)

    parse_calls: list[Path] = []
    original_parse = exam_markdown_source.parse_exam_markdown_file

    def _counted_parse(path):
        parse_calls.append(Path(path))
        return original_parse(path)

    monkeypatch.setattr(
        exam_markdown_source,
        "parse_exam_markdown_file",
        _counted_parse,
    )
    build = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="exam"),
        scene=_exam_scene(),
        material_context=MaterialExecutionContext(mode_id="exam"),
        mode_id="exam",
        plan_id="exam",
        template_id="default",
        output_root=str(tmp_path / "output"),
    )

    try:
        assert build.worker is not None
        assert build.session_snapshot is not None
        runner = build.worker.worker.runner
        prepared = runner._prepared_exam_markdown_input
        frozen_path = Path(build.session_snapshot.input_ref.frozen_path)
        assert parse_calls == [frozen_path]
        assert prepared.source_revision == build.session_snapshot.input_ref.frozen_revision
        assert prepared.material_context.entity_data["title"] == "七年级数学单元测试"
        assert prepared.material_context.entity_data["subject"] == "数学"

        source.write_text("# 已被替换的原始文件", encoding="utf-8")
        runner._run_exam_markdown_source = lambda **kwargs: {
            "status": "success",
            "title": kwargs["material_context"].entity_data["title"],
        }
        payload = runner.run(lambda *_args: None, lambda: False)

        assert payload["title"] == "七年级数学单元测试"
        assert parse_calls == [frozen_path]
        assert frozen_path.read_text(encoding="utf-8") == COMPLETE_EXAM_MARKDOWN
    finally:
        if build.session_snapshot is not None:
            cleanup_execution_session_resources(build.session_snapshot)


def test_controller_cleans_frozen_exam_input_when_prepared_gate_blocks(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "incomplete.md"
    source.write_text("# 数学练习\n", encoding="utf-8")
    _capture_runner(monkeypatch)

    build = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="exam"),
        scene=_exam_scene(),
        material_context=MaterialExecutionContext(mode_id="exam"),
        mode_id="exam",
        plan_id="exam",
        template_id="default",
        output_root=str(tmp_path / "output"),
    )

    assert build.worker is None
    assert build.session_snapshot is not None
    assert build.error_text
    assert not Path(build.session_snapshot.input_ref.frozen_path).exists()


def test_prepared_exam_markdown_rejects_source_revision_drift(tmp_path):
    source = tmp_path / "exam.md"
    source.write_text(COMPLETE_EXAM_MARKDOWN, encoding="utf-8")
    prepared = exam_markdown_execution_input.prepare_exam_markdown_input(
        scene=_exam_scene(),
        material_context=MaterialExecutionContext(mode_id="exam"),
        source_path=source,
    )
    source.write_text("# changed", encoding="utf-8")

    with pytest.raises(
        RuntimeError,
        match="execution_input_drift:exam_markdown:prepared_revision",
    ):
        exam_markdown_execution_input.validated_prepared_exam_markdown_input(
            prepared,
            source_path=source,
        )


def test_runner_rejects_prepared_exam_projection_from_another_scene(tmp_path):
    source = tmp_path / "exam.md"
    source.write_text(COMPLETE_EXAM_MARKDOWN, encoding="utf-8")
    scene_a = _exam_scene()
    prepared = exam_markdown_execution_input.prepare_exam_markdown_input(
        scene=scene_a,
        material_context=MaterialExecutionContext(mode_id="exam"),
        source_path=source,
    )
    scene_b = _exam_scene()
    scene_b.name = "another execution plan"

    from src.services.production_runtime.execution_runtime import (
        WorkbenchProductionRunner,
    )

    runner = WorkbenchProductionRunner(
        doc_path=str(source),
        template=create_builtin_template("default", mode_id="exam"),
        scene=scene_b,
        material_context=MaterialExecutionContext(mode_id="exam"),
        prepared_exam_markdown_input=prepared,
    )

    with pytest.raises(
        RuntimeError,
        match="execution_input_drift:exam_markdown:scene_revision",
    ):
        runner.run(lambda *_args: None, lambda: False)


def test_controller_rejects_stale_object_preflight_revision_and_cleans_input(
    tmp_path,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = create_builtin_scene("custom", mode_id="custom")
    evidence = build_object_preflight_evidence(scene, source)

    build = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="custom"),
        scene=scene,
        material_context=MaterialExecutionContext(mode_id="custom"),
        mode_id="custom",
        plan_id="custom",
        template_id="default",
        output_root=str(tmp_path / "output"),
        expected_input_revision="sha256:" + "0" * 64,
        object_preflight_confirmation_digest=evidence.evidence_digest,
    )

    assert build.worker is None
    assert build.session_snapshot is not None
    assert "object_preflight_confirmation_stale" in build.error_text
    assert build.session_snapshot.to_dict()[
        "object_preflight_confirmation"
    ]["status"] == "stale"
    assert build.session_snapshot.input_ref.frozen_path == ""


def test_controller_rejects_missing_object_preflight_receipt(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)

    build = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="custom"),
        scene=create_builtin_scene("custom", mode_id="custom"),
        material_context=MaterialExecutionContext(mode_id="custom"),
        mode_id="custom",
        plan_id="custom",
        template_id="default",
        output_root=str(tmp_path / "output"),
    )

    assert build.worker is None
    assert build.session_snapshot is not None
    assert "object_preflight_confirmation_revision_missing" in build.error_text
    assert "object_preflight_confirmation_digest_missing" in build.error_text
    assert build.session_snapshot.to_dict()[
        "object_preflight_confirmation"
    ]["status"] == "missing"
    assert build.session_snapshot.input_ref.frozen_path == ""


def test_controller_rejects_forged_object_preflight_digest(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = create_builtin_scene("custom", mode_id="custom")
    evidence = build_object_preflight_evidence(scene, source)

    build = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="custom"),
        scene=scene,
        material_context=MaterialExecutionContext(mode_id="custom"),
        mode_id="custom",
        plan_id="custom",
        template_id="default",
        output_root=str(tmp_path / "output"),
        expected_input_revision=evidence.source_revision,
        object_preflight_confirmation_digest="sha256:" + "0" * 64,
    )

    assert build.worker is None
    assert build.session_snapshot is not None
    assert "object_preflight_confirmation_digest_mismatch" in build.error_text
    assert build.session_snapshot.to_dict()[
        "object_preflight_confirmation"
    ]["status"] == "mismatch"
    assert build.session_snapshot.input_ref.frozen_path == ""


def test_controller_freezes_invalid_and_incomplete_preflight_receipt_statuses(
    tmp_path,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = create_builtin_scene("custom", mode_id="custom")
    evidence = build_object_preflight_evidence(scene, source)

    invalid = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="custom"),
        scene=scene,
        material_context=MaterialExecutionContext(mode_id="custom"),
        mode_id="custom",
        plan_id="custom",
        template_id="default",
        output_root=str(tmp_path / "invalid"),
        expected_input_revision=evidence.source_revision,
        object_preflight_confirmation_digest="not-a-digest",
    )
    incomplete = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="custom"),
        scene=scene,
        material_context=MaterialExecutionContext(mode_id="custom"),
        mode_id="custom",
        plan_id="custom",
        template_id="default",
        output_root=str(tmp_path / "incomplete"),
        expected_input_revision=evidence.source_revision,
    )

    assert invalid.session_snapshot is not None
    assert incomplete.session_snapshot is not None
    assert invalid.session_snapshot.to_dict()[
        "object_preflight_confirmation"
    ]["status"] == "invalid"
    assert incomplete.session_snapshot.to_dict()[
        "object_preflight_confirmation"
    ]["status"] == "missing"
    assert invalid.session_snapshot.input_ref.frozen_path == ""
    assert incomplete.session_snapshot.input_ref.frozen_path == ""


def test_controller_blocks_failed_docx_inspection_even_with_matching_receipt(
    tmp_path,
):
    source = tmp_path / "corrupt.docx"
    source.write_bytes(b"not a zip package")
    scene = create_builtin_scene("custom", mode_id="custom")
    scene.strict_mode = False
    scene.compliance_profile.failure_policy = "warn"
    evidence = build_object_preflight_evidence(scene, source)

    build = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="custom"),
        scene=scene,
        material_context=MaterialExecutionContext(mode_id="custom"),
        mode_id="custom",
        plan_id="custom",
        template_id="default",
        output_root=str(tmp_path / "output"),
        expected_input_revision=evidence.source_revision,
        object_preflight_confirmation_digest=evidence.evidence_digest,
    )

    assert evidence.blocked is True
    assert build.worker is None
    assert build.session_snapshot is not None
    assert "object_preflight_blocked" in build.error_text
    assert build.session_snapshot.to_dict()[
        "object_preflight_confirmation"
    ]["status"] == "blocked"
    assert build.session_snapshot.input_ref.frozen_path == ""


def test_controller_freezes_blocked_preflight_receipt_status(tmp_path):
    source = tmp_path / "strict.docx"
    Document().save(source)
    with ZipFile(source, "a") as package:
        package.writestr("word/embeddings/oleObject1.bin", b"ole")
    scene = create_builtin_scene("custom", mode_id="custom")
    policy = scene.compliance_profile.object_preflight
    policy.scan_targets = ["ole_objects"]
    policy.preservation_mode = "strict"
    policy.block_on = ["ole_objects"]
    evidence = build_object_preflight_evidence(scene, source)

    build = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="custom"),
        scene=scene,
        material_context=MaterialExecutionContext(mode_id="custom"),
        mode_id="custom",
        plan_id="custom",
        template_id="default",
        output_root=str(tmp_path / "output"),
        expected_input_revision=evidence.source_revision,
        object_preflight_confirmation_digest=evidence.evidence_digest,
    )

    assert build.worker is None
    assert build.session_snapshot is not None
    assert build.session_snapshot.to_dict()[
        "object_preflight_confirmation"
    ]["status"] == "blocked"
    assert build.session_snapshot.input_ref.frozen_path == ""


def test_execution_snapshot_preserves_object_preflight_confirmation_receipt(
    tmp_path,
):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("第一章 正文", level=1)
    document.add_paragraph("正文")
    document.save(source)
    scene = create_builtin_scene("custom", mode_id="custom")
    evidence = build_object_preflight_evidence(scene, source)
    structure_evidence = build_document_structure_evidence(source)
    build = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: source,
    ).build_worker(
        template=create_builtin_template("default", mode_id="custom"),
        scene=scene,
        material_context=MaterialExecutionContext(mode_id="custom"),
        mode_id="custom",
        plan_id="custom",
        template_id="default",
        output_root=str(tmp_path / "output"),
        expected_input_revision=evidence.source_revision,
        object_preflight_confirmation_digest=evidence.evidence_digest,
        document_structure_evidence=structure_evidence,
    )

    try:
        assert build.worker is not None
        assert build.session_snapshot is not None
        receipt = build.session_snapshot.to_dict()[
            "object_preflight_confirmation"
        ]
        assert receipt == {
            "status": "confirmed",
            "applicable": True,
            "blocked": False,
            "source_revision": build.session_snapshot.input_ref.source_revision,
            "evidence_digest": evidence.evidence_digest,
        }
    finally:
        if build.session_snapshot is not None:
            cleanup_execution_session_resources(build.session_snapshot)
