from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from src.assistant.adapters.content_generation_adapter import (
    AssistantContentGenerationAdapter,
    GeneratedExamDraft,
)
from src.assistant.application.content_generation_service import (
    AssistantContentGenerationService,
    ContentGenerationRequest,
    _blueprint_score_plan,
    _exam_section_heading,
    _normalize_exam_answer_numbering,
    _normalize_exam_header_metadata,
    _normalize_exam_score_metadata,
)
from src.assistant.contracts.permissions import DisclosureGrant
from src.assistant.domain.exam_authoring_contract import resolve_exam_blueprint
from src.assistant.runtime.mock_provider import MockModelGateway
from src.assistant.runtime.provider_contract import (
    PROVIDER_DONE,
    PROVIDER_START,
    PROVIDER_TEXT_DELTA,
    ProviderStreamEvent,
)
from src.assistant.ui.document_workflow_mixin import (
    _content_generation_failure_presentation,
)
from src.shared.engine.exam_question_schema import parse_exam_markdown_source


def test_generated_markdown_compiles_to_fragment_and_reviewable_docx(tmp_path):
    draft = AssistantContentGenerationAdapter(tmp_path).compile_and_compose(
        session_id="session-1",
        markdown="# 项目报告\n\n## 结论\n\n- 第一项\n- 第二项\n",
    )

    assert Path(draft.markdown_path).is_file()
    assert Path(draft.document_path).is_file()
    assert [block["kind"] for block in draft.fragment["blocks"]] == [
        "heading",
        "heading",
        "list",
    ]
    assert len(draft.fragment_digest) == 64
    assert len(draft.compose_receipt_id) == 64


def test_generated_content_rejects_direct_ooxml(tmp_path):
    adapter = AssistantContentGenerationAdapter(tmp_path)

    with pytest.raises(ValueError, match="direct_ooxml_is_forbidden"):
        adapter.compile_and_compose(
            session_id="session-1",
            markdown="<w:document>not allowed</w:document>",
        )


@pytest.mark.parametrize(
    ("markdown", "expected_code"),
    (
        (
            "# 投标文件\n\n"
            "{{@text:company_name}}\n"
            "{{@text:project_name}}\n"
            "{{@text:legal_person}}\n",
            "required_section_missing:project_understanding",
        ),
        (
            "# 投标文件\n\n"
            "{{@text:company_name}}\n"
            "{{@text:project_name}}\n"
            "{{@text:legal_person}}\n\n"
            "## 项目理解\n\n内容\n\n"
            "## 响应内容\n\n内容\n\n"
            "## 实施方案\n\n内容\n\n"
            "## 承诺事项\n\n{{@img:LOGO1}}\n",
            "provider_owned_image_token_forbidden",
        ),
    ),
)
def test_bidding_draft_fails_closed_on_domain_contract_violations(
    tmp_path,
    markdown,
    expected_code,
):
    adapter = AssistantContentGenerationAdapter(tmp_path)

    with pytest.raises(ValueError, match=expected_code):
        adapter.compile_and_compose(
            session_id="session-bidding-invalid",
            markdown=markdown,
            prompt_profile_id="assistant.bidding-markdown.v1",
        )


def test_bidding_local_profile_uses_canonical_logo_and_seal_roles(tmp_path):
    draft = AssistantContentGenerationAdapter(tmp_path).compile_and_compose(
        session_id="session-bidding-canonical-assets",
        markdown=(
            "# 投标文件\n\n"
            "{{@text:company_name}}\n"
            "{{@text:project_name}}\n"
            "{{@text:legal_person}}\n\n"
            "## 项目理解\n\n项目理解内容。\n\n"
            "## 响应内容\n\n响应内容。\n\n"
            "## 实施方案\n\n实施方案。\n\n"
            "## 承诺事项\n\n承诺事项。\n"
        ),
        prompt_profile_id="assistant.bidding-markdown.v1",
    )

    text = "\n".join(
        paragraph.text for paragraph in Document(draft.document_path).paragraphs
    )
    assert "{{@img:logo}}" in text
    assert "{{@img:seal}}" in text
    assert "{{@img:LOGO1}}" not in text
    assert "{{@img:公章1}}" not in text


def test_context_cannot_reach_provider_without_exact_disclosure_grant(tmp_path):
    with pytest.raises(PermissionError, match="not_authorized"):
        ContentGenerationRequest(
            session_id="session-1",
            turn_id="turn-1",
            prompt="起草报告",
            provider_id="cloud-main",
            model_id="model-a",
            context_text="confidential body",
            context_refs=("doc-1",),
            context_fields=("paragraphs",),
            context_fingerprints=(("doc-1", "hash-a"),),
        )

    wrong_model = DisclosureGrant(
        grant_id="grant-1",
        session_id="session-1",
        provider_id="cloud-main",
        model_id="model-b",
        allowed_refs=("doc-1",),
        allowed_fields=("paragraphs",),
        content_fingerprints={"doc-1": "hash-a"},
    )
    with pytest.raises(PermissionError, match="not_authorized"):
        ContentGenerationRequest(
            session_id="session-1",
            turn_id="turn-1",
            prompt="起草报告",
            provider_id="cloud-main",
            model_id="model-a",
            context_text="confidential body",
            context_refs=("doc-1",),
            context_fields=("paragraphs",),
            context_fingerprints=(("doc-1", "hash-a"),),
            disclosure_grant=wrong_model,
        )


def test_prompt_only_generation_uses_provider_then_local_compiler(tmp_path):
    service = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    )
    request = ContentGenerationRequest(
        session_id="session-1",
        turn_id="turn-1",
        prompt="生成一份项目报告",
        provider_id="mock-default",
        model_id="form-assistant-mock",
    )

    draft = service.generate(request, MockModelGateway())

    assert Path(draft.document_path).is_file()
    assert draft.fragment["blocks"]


def test_material_generation_extracts_authorized_docx_in_worker_context(tmp_path):
    material = tmp_path / "confidential-brief.docx"
    document = Document()
    document.add_paragraph("Approved material fact: launch is in October.")
    document.save(material)

    class CapturingGateway:
        def __init__(self) -> None:
            self.request = None

        def stream(self, request):
            self.request = request
            yield ProviderStreamEvent(PROVIDER_START)
            yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text="# Draft\n\nGenerated body")
            yield ProviderStreamEvent(PROVIDER_DONE)

        def cancel(self) -> bool:
            return True

    reference = str(material)
    grant = DisclosureGrant(
        grant_id="grant-material",
        session_id="session-material",
        provider_id="cloud-main",
        model_id="model-a",
        allowed_refs=(reference,),
        allowed_fields=("document_text",),
    )
    request = ContentGenerationRequest(
        session_id="session-material",
        turn_id="turn-material",
        prompt="Generate a launch report",
        provider_id="cloud-main",
        model_id="model-a",
        context_documents=({"title": material.name, "path": reference},),
        context_refs=(reference,),
        context_fields=("document_text",),
        disclosure_grant=grant,
    )
    gateway = CapturingGateway()

    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path / "outputs")
    ).generate(request, gateway)

    assert Path(draft.document_path).is_file()
    assert gateway.request is not None
    provider_text = gateway.request.messages[0]["content"]
    assert "launch is in October" in provider_text
    assert str(tmp_path) not in provider_text


def _standard_exam_question_phase() -> str:
    lines = [
        "# 高一信息技术 Python 基础单元测试",
        "> 科目：信息技术　年级：高一　考试时间：60 分钟　满分：100 分",
        "",
        "## 一、单项选择题",
        "",
    ]
    for number in range(1, 11):
        difficulty = "基础" if number <= 5 else "中等"
        lines.extend(
            (
                f"{number}. 第 {number} 道选择题的唯一题干。（3 分）",
                "   A. 选项 A",
                "   B. 选项 B",
                "   C. 选项 C",
                "   D. 选项 D",
                f"   difficulty: {difficulty}",
                f"   knowledge_points: 知识点{(number - 1) % 5 + 1}",
                "",
            )
        )
    lines.extend(("## 二、程序阅读与分析题", ""))
    for number in range(11, 16):
        lines.extend(
            (
                f"{number}. 分析第 {number} 段程序的运行结果。（6 分）",
                "   difficulty: 中等",
                f"   knowledge_points: 知识点{(number - 1) % 5 + 1}",
                "   answer_area_kind: lines",
                "   answer_lines: 3",
                "",
            )
        )
    lines.extend(("## 三、程序设计题", ""))
    for number, score in ((16, 12), (17, 13), (18, 15)):
        lines.extend(
            (
                f"{number}. 编写第 {number} 个完整程序。（{score} 分）",
                "   difficulty: 提高",
                f"   knowledge_points: 知识点{(number - 1) % 5 + 1}",
                "   answer_area_kind: free",
                "   answer_lines: 5",
                "",
            )
        )
    return "\n".join(lines)


def _standard_exam_answer_phase() -> str:
    return "## 答案速查\n" + "\n".join(
        f"{number}. {'A' if number <= 10 else f'第 {number} 题参考答案与评分标准。'}"
        for number in range(1, 19)
    )


def _standard_exam_answer_phase_with_section_restarts() -> str:
    groups = (
        ("一、单项选择题", range(1, 11)),
        ("二、程序阅读题", range(11, 16)),
        ("三、程序设计题", range(16, 19)),
    )
    lines = ["## 答案速查"]
    for title, global_numbers in groups:
        lines.append("### " + title)
        for local_number, global_number in enumerate(global_numbers, start=1):
            answer = "A" if global_number <= 10 else f"第 {global_number} 题参考答案。"
            lines.append(f"{local_number}. {answer}")
    return "\n".join(lines)


def test_answer_numbering_repairs_section_restarts_only_when_coverage_is_complete():
    repaired = _normalize_exam_answer_numbering(
        _standard_exam_answer_phase_with_section_restarts(),
        expected_count=18,
    )
    incomplete = _normalize_exam_answer_numbering(
        "## 答案速查\n1. A\n1. B\n",
        expected_count=18,
    )

    assert "11. 第 11 题参考答案。" in repaired
    assert "16. 第 16 题参考答案。" in repaired
    assert incomplete.count("1.") == 2


def test_exam_header_metadata_is_controller_owned_for_primary_school_request():
    intent = "帮我生成一份小学六年级的期中英语考试试卷"
    blueprint = resolve_exam_blueprint(intent, scene_id="exam_term")
    normalized = _normalize_exam_header_metadata(
        "# Wrong\n> 科目：数学　年级：初一　考试时间：10 分钟　满分：5 分\n"
        "## 一、选择题\n1. Test?（5 分）\n",
        intent=intent,
        blueprint=blueprint,
    )

    assert normalized.startswith("# 小学六年级英语期中试卷")
    assert "科目：英语　年级：小学六年级" in normalized
    assert "考试时间：90 分钟　满分：100 分" in normalized


def test_term_exam_score_normalization_repairs_2_point_4_decimal_averages():
    intent = "帮我生成一份小学六年级的期中英语考试试卷"
    blueprint = resolve_exam_blueprint(intent, scene_id="exam_term")
    lines = [
        "# 小学六年级英语期中试卷",
        "> 科目：英语　年级：小学六年级　考试时间：90 分钟　满分：100 分",
        "",
    ]
    number = 1
    for section in blueprint.section_blueprints:
        decimal_average = section.total_score / section.question_count
        lines.extend((f"## {section.label}", ""))
        for _ in range(section.question_count):
            lines.extend(
                (
                    f"{number}. 第 {number} 道完整题目。",
                    f"   score: {decimal_average}",
                    "",
                )
            )
            number += 1
    markdown = "\n".join(lines)

    desired_scores = _blueprint_score_plan(blueprint, markdown)
    normalized = _normalize_exam_score_metadata(markdown, desired_scores)
    result = parse_exam_markdown_source(normalized)
    scores = [
        float(question["score"])
        for section in result.payload["sections"]
        for question in section["questions"]
    ]

    assert "score: 2.4" not in normalized
    assert desired_scores[:10] == (
        2.0,
        2.0,
        2.0,
        2.0,
        2.0,
        2.0,
        3.0,
        3.0,
        3.0,
        3.0,
    )
    assert all(score.is_integer() for score in scores)
    assert sum(scores) == 100
    assert "每小题 2.4 分" not in _exam_section_heading(
        blueprint.section_blueprints[0]
    )


def test_rejected_exam_text_is_retained_for_diagnosis(tmp_path):
    adapter = AssistantContentGenerationAdapter(tmp_path)

    with pytest.raises(ValueError) as raised:
        adapter.compile_exam_markdown(
            session_id="session-rejected-exam",
            markdown="# incomplete exam\n",
            intent="生成一份六年级英语试卷",
            scene_id="exam",
        )

    error = str(raised.value)
    assert "|diagnostic=" in error
    diagnostic_path = Path(error.partition("|diagnostic=")[2])
    assert diagnostic_path.is_file()
    assert diagnostic_path.read_text(encoding="utf-8").startswith("# incomplete exam")


def test_exam_validation_failure_is_not_misreported_as_model_configuration_error():
    title, body = _content_generation_failure_presentation(
        "exam_content_compile_blocked:duplicate_answer_number"
        "|diagnostic=C:/drafts/rejected.exam.md"
    )

    assert title == "题稿未通过校验"
    assert "答案编号重复" in body
    assert "模型已经返回内容" in body
    assert "检查模型配置" not in body
    assert "C:/drafts/rejected.exam.md" in body


def _quiz_exam_question_phase() -> str:
    lines = [
        "# 高一信息技术基础随堂测验",
        "> 科目：信息技术　年级：高一　考试时间：25 分钟　满分：30 分",
        "",
    ]
    sections = (
        ("一、基础客观题", range(1, 5), 2),
        ("二、阅读/分析题", range(5, 8), 4),
        ("三、综合应用题", range(8, 9), 10),
    )
    for label, numbers, score in sections:
        lines.extend((f"## {label}", ""))
        for number in numbers:
            lines.extend(
                (
                    f"{number}. 第 {number} 道完整题目。（{score} 分）",
                    "   difficulty: 中等",
                    f"   knowledge_points: 知识点{(number - 1) % 3 + 1}",
                    "   answer_area_kind: lines",
                    "   answer_lines: 2",
                    "",
                )
            )
    return "\n".join(lines)


def _quiz_exam_answer_phase() -> str:
    return "## 答案速查\n" + "\n".join(
        f"{number}. 第 {number} 题参考答案与评分标准。"
        for number in range(1, 9)
    )


def test_user_scene_quiz_uses_authoritative_scale_and_staged_repair_path(
    tmp_path,
):
    class QuizGateway:
        def __init__(self) -> None:
            self.requests = []

        def stream(self, request):
            self.requests.append(request)
            phase = request.metadata["generation_phase"]
            text = (
                _quiz_exam_question_phase()
                if phase == "questions"
                else _quiz_exam_answer_phase()
            )
            yield ProviderStreamEvent(PROVIDER_START)
            yield ProviderStreamEvent(PROVIDER_DONE, text=text)

        def cancel(self) -> bool:
            return True

    request = ContentGenerationRequest(
        session_id="session-user-quiz",
        turn_id="turn-user-quiz",
        prompt="生成一份高一信息技术试卷",
        provider_id="cloud-main",
        model_id="model-a",
        artifact_kind="exam_paper_markdown",
        prompt_profile_id="assistant.exam-paper-markdown.v1",
        scene_id="my_school_exam",
        scale_profile_id="quiz",
    )
    gateway = QuizGateway()

    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    ).generate(request, gateway)

    assert isinstance(draft, GeneratedExamDraft)
    assert draft.validation_summary["question_count"] == 8
    assert draft.validation_summary["blueprint"]["profile_id"] == "quiz"
    assert [
        item.metadata["generation_phase"] for item in gateway.requests
    ] == ["questions", "answers"]


def test_large_exam_generation_uses_question_then_answer_phases(tmp_path):
    class StagedExamGateway:
        def __init__(self) -> None:
            self.requests = []

        def stream(self, request):
            self.requests.append(request)
            phase = request.metadata["generation_phase"]
            text = (
                _standard_exam_question_phase()
                if phase == "questions"
                else _standard_exam_answer_phase()
            )
            yield ProviderStreamEvent(PROVIDER_START)
            yield ProviderStreamEvent(PROVIDER_DONE, text=text)

        def cancel(self) -> bool:
            return True

    request = ContentGenerationRequest(
        session_id="session-exam-staged",
        turn_id="turn-exam-staged",
        prompt=(
            "生成一份高一信息技术 Python 基础单元测试，"
            "考试时间 60 分钟，满分 100 分"
        ),
        provider_id="cloud-main",
        model_id="model-a",
        artifact_kind="exam_paper_markdown",
        prompt_profile_id="assistant.exam-paper-markdown.v1",
        scene_id="exam",
    )
    gateway = StagedExamGateway()

    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    ).generate(request, gateway)

    assert isinstance(draft, GeneratedExamDraft)
    assert draft.validation_summary["question_count"] == 18
    assert draft.validation_summary["blueprint"]["student_page_range"] == [4, 8]
    assert [
        item.metadata["generation_phase"] for item in gateway.requests
    ] == ["questions", "answers"]


def test_staged_exam_locally_repairs_answer_numbers_restarted_by_section(tmp_path):
    class RestartedAnswerGateway:
        def stream(self, request):
            text = (
                _standard_exam_question_phase()
                if request.metadata["generation_phase"] == "questions"
                else _standard_exam_answer_phase_with_section_restarts()
            )
            yield ProviderStreamEvent(PROVIDER_START)
            yield ProviderStreamEvent(PROVIDER_DONE, text=text)

        def cancel(self) -> bool:
            return True

    request = ContentGenerationRequest(
        session_id="session-exam-answer-repair",
        turn_id="turn-exam-answer-repair",
        prompt=(
            "生成一份高一信息技术 Python 基础单元测试，"
            "考试时间 60 分钟，满分 100 分"
        ),
        provider_id="cloud-main",
        model_id="model-a",
        artifact_kind="exam_paper_markdown",
        prompt_profile_id="assistant.exam-paper-markdown.v1",
        scene_id="exam",
    )

    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    ).generate(request, RestartedAnswerGateway())
    markdown = Path(draft.markdown_path).read_text(encoding="utf-8")

    assert isinstance(draft, GeneratedExamDraft)
    assert "11. 第 11 题参考答案。" in markdown
    assert "16. 第 16 题参考答案。" in markdown


def test_large_exam_falls_back_to_individually_validated_sections(tmp_path):
    class SectionFallbackGateway:
        def __init__(self) -> None:
            self.requests = []
            self.sections = [
                "## " + part
                for part in _standard_exam_question_phase().split("## ")[1:]
            ]

        def stream(self, request):
            self.requests.append(request)
            phase = request.metadata["generation_phase"]
            if phase == "questions":
                text = "# 不完整试卷\n"
            elif phase.startswith("questions_section_"):
                section_index = int(phase.split("_")[2]) - 1
                text = self.sections[section_index]
            else:
                text = _standard_exam_answer_phase()
            yield ProviderStreamEvent(PROVIDER_START)
            yield ProviderStreamEvent(PROVIDER_DONE, text=text)

        def cancel(self) -> bool:
            return True

    request = ContentGenerationRequest(
        session_id="session-exam-section-fallback",
        turn_id="turn-exam-section-fallback",
        prompt=(
            "生成一份高一信息技术 Python 基础单元测试，"
            "考试时间 60 分钟，满分 100 分"
        ),
        provider_id="cloud-main",
        model_id="model-a",
        artifact_kind="exam_paper_markdown",
        prompt_profile_id="assistant.exam-paper-markdown.v1",
        scene_id="exam",
    )
    gateway = SectionFallbackGateway()

    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    ).generate(request, gateway)

    assert isinstance(draft, GeneratedExamDraft)
    assert draft.validation_summary["question_count"] == 18
    assert [
        item.metadata["generation_phase"] for item in gateway.requests
    ] == [
        "questions",
        "questions_section_1",
        "questions_section_2",
        "questions_section_3",
        "answers",
    ]


def test_section_fallback_repairs_score_arithmetic_without_another_model_call(
    tmp_path,
):
    class SectionScoreRepairGateway:
        def __init__(self) -> None:
            self.requests = []
            self.sections = [
                "## " + part
                for part in _standard_exam_question_phase().split("## ")[1:]
            ]
            self.sections[2] = _normalize_exam_score_metadata(
                self.sections[2],
                (13.0, 13.0, 13.0),
            )

        def stream(self, request):
            self.requests.append(request)
            phase = request.metadata["generation_phase"]
            if phase == "questions":
                text = "# incomplete\n"
            elif phase.startswith("questions_section_"):
                section_index = int(phase.split("_")[2]) - 1
                text = self.sections[section_index]
            else:
                text = _standard_exam_answer_phase()
            yield ProviderStreamEvent(PROVIDER_START)
            yield ProviderStreamEvent(PROVIDER_DONE, text=text)

        def cancel(self) -> bool:
            return True

    request = ContentGenerationRequest(
        session_id="session-exam-section-score-repair",
        turn_id="turn-exam-section-score-repair",
        prompt=(
            "生成一份高一信息技术 Python 基础单元测试，"
            "考试时间 60 分钟，满分 100 分。"
        ),
        provider_id="cloud-main",
        model_id="model-a",
        artifact_kind="exam_paper_markdown",
        prompt_profile_id="assistant.exam-paper-markdown.v1",
        scene_id="exam",
    )
    gateway = SectionScoreRepairGateway()

    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    ).generate(request, gateway)

    assert isinstance(draft, GeneratedExamDraft)
    assert draft.validation_summary["declared_total_score"] == 100
    assert [
        item.metadata["generation_phase"] for item in gateway.requests
    ] == [
        "questions",
        "questions_section_1",
        "questions_section_2",
        "questions_section_3",
        "answers",
    ]
