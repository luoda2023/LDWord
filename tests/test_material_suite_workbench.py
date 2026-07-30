from __future__ import annotations

import time

from docx import Document
from openpyxl import Workbook

from src.config.entity import EntityArchive, EntityProfile
from src.config.material_batch import MaterialBatchSelection
from src.config.material_import_draft import inspect_material_workbook
from src.qt_api import QApplication, Qt
from src.ui.bridge import PanelBridge
from src.ui.panels.workbench.material_suite_generation_detail import (
    _load_material_package,
)
from src.ui.panels.workbench.panel_v2 import WorkbenchPanel


def _app():
    return QApplication.instance() or QApplication([])


def _templates(root):
    word_root = root / "Word文档"
    word_root.mkdir(parents=True)
    document = Document()
    document.add_paragraph("{{项目名称}}")
    document.save(word_root / "计划书.docx")
    excel_root = root / "产品A_测试表"
    excel_root.mkdir()
    workbook = Workbook()
    workbook.active["A1"] = "{{项目编号}}"
    workbook.save(excel_root / "记录表.xlsx")
    workbook.close()


def _selection():
    archive = EntityArchive(
        archive_id="suite",
        archive_name="成套资料",
        profiles=[
            EntityProfile(
                profile_id="one",
                profile_name="项目一",
                fields={
                    "产品名称": "产品A",
                    "项目名称": "项目一",
                    "项目编号": "A-001",
                },
            )
        ],
    )
    return MaterialBatchSelection(
        package_id="suite",
        archive=archive,
        profile_ids=["one"],
    )


def test_workbench_exposes_independent_material_suite_fixed_card(tmp_path):
    _app()
    panel = WorkbenchPanel(PanelBridge())
    try:
        assert list(panel._navigation_cards)[:2] == [
            "quick_execute",
            "batch_generate",
        ]
        assert "material_suite_generate" in panel._navigation_cards
        assert (
            panel._detail_map["material_suite_generate"]
            is panel._suite_generation_detail
        )
        panel._nav_rail.select_card("material_suite_generate")
        assert panel._current_detail is panel._suite_generation_detail
    finally:
        panel.close()


def test_suite_detail_compiles_record_artifact_matrix_without_batch_document(
    tmp_path,
):
    _app()
    template_root = tmp_path / "模板文件"
    _templates(template_root)
    panel = WorkbenchPanel(PanelBridge())
    try:
        detail = panel._suite_generation_detail
        detail.set_material_batch_selection(_selection())
        detail.set_template_root(str(template_root))
        detail.set_output_root(str(tmp_path / "输出"))

        plan = detail.current_run_plan()

        assert plan.ok
        assert plan.artifact_count == 2
        assert detail.can_start_execution()
        assert "2 个文件" in detail._matrix_preview.toPlainText()
        # The existing batch page still owns the shared-DOCX workflow.
        assert panel._batch_generation_detail is not detail
    finally:
        panel.close()


def test_suite_workbench_flow_runs_through_shared_execution_lifecycle(tmp_path):
    app = _app()
    template_root = tmp_path / "模板文件"
    _templates(template_root)
    panel = WorkbenchPanel(PanelBridge())
    try:
        detail = panel._suite_generation_detail
        detail.set_material_batch_selection(_selection())
        detail.set_template_root(str(template_root))
        detail.set_output_root(str(tmp_path / "输出"))

        panel._suite_generation_flow.start()
        handle = panel._execution_worker
        deadline = time.monotonic() + 10
        while (
            handle is not None
            and (panel._execution_worker is not None or handle._thread.isRunning())
            and time.monotonic() < deadline
        ):
            app.processEvents()
            time.sleep(0.01)
        app.processEvents()

        assert panel._execution_worker is None
        assert detail._last_result_status == "success"
        assert (
            tmp_path / "输出" / "产品A" / "项目一" / "Word文档" / "计划书.docx"
        ).is_file()
    finally:
        panel.close()


def test_legacy_vba_workbook_import_becomes_stable_record_and_timeline(tmp_path):
    source = tmp_path / "ISO项目数据.xlsx"
    workbook = Workbook()
    primary = workbook.active
    primary.title = "项目基础数据"
    primary.append(
        [
            "产品名称",
            "预留",
            "项目名称",
            "项目编号",
            "项目来源",
            "项目开始日期",
            "项目结束日期",
            "目标成本",
        ]
    )
    primary.append(
        [
            "产品A",
            "",
            "项目一",
            "A-001",
            "客户",
            "2026-01-01",
            "2026-01-11",
            "100",
        ]
    )
    people = workbook.create_sheet("人员配置")
    people.append(["姓名1", "姓名2"])
    people.append(["张三", "李四"])
    custom = workbook.create_sheet("自定义文本")
    custom.append(["自定义文本1"])
    custom.append(["第一行\n第二行"])
    workbook.save(source)
    workbook.close()

    archive = _load_material_package(str(source))
    profile = archive.profiles[0]

    assert profile.profile_id == "record-0001"
    assert profile.profile_name == "项目一"
    assert profile.fields["姓名1"] == "张三"
    assert profile.fields["自定义文本1"] == "第一行\n第二行"
    assert list(profile.timeline_plans) == ["legacy_iso_timeline"]
    outputs = [
        node["outputs"][0]["field"]
        for node in profile.timeline_plans["legacy_iso_timeline"]["nodes"]
    ]
    assert outputs == [
        "节点_设计策划",
        "节点_编制计划书",
        "节点_设计输入",
        "节点_性能评审",
        "节点_设计输出",
        "节点_系统评审",
        "节点_设计验证",
        "节点_试产可行性",
        "节点_试产总结",
        "节点_设计确认",
        "节点_设计正稿",
    ]


def test_suite_record_matrix_keeps_drafts_visible_but_selects_only_active(
    tmp_path,
):
    _app()
    source = tmp_path / "项目矩阵.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["产品名称", "项目序号", "项目名称", "项目编号"])
    for product in ("产品A", "产品B"):
        for sequence in range(1, 4):
            active = product == "产品A"
            sheet.append(
                [
                    product,
                    sequence,
                    f"项目{sequence}" if active else "",
                    f"A-{sequence}" if active else "",
                ]
            )
    workbook.save(source)
    workbook.close()
    package = inspect_material_workbook(source).materialize()
    selection = MaterialBatchSelection(
        package_id=package.package_id,
        archive=package.to_entity_archive(),
        profile_ids=[
            item.record_id
            for item in package.records
            if item.lifecycle_state == "active"
        ],
        source_path=str(source),
        material_package=package,
    )
    panel = WorkbenchPanel(PanelBridge())
    try:
        detail = panel._suite_generation_detail
        detail.set_material_batch_selection(selection)

        assert detail._record_list.count() == 6
        checked = [
            detail._record_list.item(index).checkState() == Qt.Checked
            for index in range(detail._record_list.count())
        ]
        assert checked == [True, True, True, False, False, False]
        assert "6 个候选槽位" in detail._package_summary.text()
        assert "3 条活动记录" in detail._package_summary.text()
        assert "3 条草稿" in detail._package_summary.text()
    finally:
        panel.close()
