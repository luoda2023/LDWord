from __future__ import annotations

from pathlib import Path

from docx import Document
from PIL import Image

from src.config.content_artifacts import ContentMaterialBinding
from src.config.content_materials import (
    ContentInsertionRule,
    ImageBlock,
    ListBlock,
)
from src.services.material_content.artifact_repository import (
    ContentArtifactRepository,
)
from src.services.material_content.compiler import compile_content_material
from src.services.material_content.docx_renderer import render_document_fragment
from src.services.material_content.import_contract import ContentImportDisposition


TOKEN = "{{@file:body}}"


def _write_png(path: Path) -> bytes:
    Image.new("RGB", (24, 12), (20, 80, 160)).save(path, format="PNG")
    return path.read_bytes()


def test_markdown_compiles_once_and_runs_after_external_sources_are_deleted(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.md"
    image = tmp_path / "diagram.png"
    image_bytes = _write_png(image)
    source.write_text(
        "# Title\n\n3. third\n4. fourth\n\n![diagram](diagram.png)\n",
        encoding="utf-8",
    )
    repository = ContentArtifactRepository(tmp_path / "artifacts")

    result = compile_content_material(source, repository)

    assert result.blocked is False
    assert result.artifact_ref is not None
    binding = ContentMaterialBinding("body", "正文", result.artifact_ref)
    source.unlink()
    image.unlink()

    fragment = repository.load_fragment(binding.artifact_ref)
    ordered_list = next(block for block in fragment.blocks if isinstance(block, ListBlock))
    image_block = next(block for block in fragment.blocks if isinstance(block, ImageBlock))
    assert ordered_list.start == 3
    assert repository.resolve_resource(
        binding.artifact_ref,
        image_block.resource_id,
    ).read_bytes() == image_bytes

    target = Document()
    target.add_paragraph(TOKEN)
    render_receipt = render_document_fragment(
        target,
        fragment,
        ContentInsertionRule("body-rule", "body", TOKEN),
    )
    assert render_receipt.rendered_block_count == len(fragment.blocks)
    assert len(render_receipt.image_job_drafts) == 1


def test_same_markdown_bundle_in_another_directory_reuses_artifact_id(
    tmp_path: Path,
) -> None:
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    sources = []
    for folder_name in ("a", "b"):
        folder = tmp_path / folder_name
        folder.mkdir()
        source = folder / "source.md"
        source.write_text("# Same\n", encoding="utf-8")
        sources.append(source)

    first = compile_content_material(sources[0], repository)
    second = compile_content_material(sources[1], repository)

    assert first.blocked is False
    assert second.blocked is False
    assert first.artifact_ref == second.artifact_ref


def test_blocked_markdown_result_is_typed_compact_and_path_free(tmp_path: Path) -> None:
    source = tmp_path / "unsafe.md"
    source.write_text("<script>alert(1)</script>\n", encoding="utf-8")
    repository = ContentArtifactRepository(tmp_path / "artifacts")

    result = compile_content_material(source, repository)

    assert result.blocked is True
    assert result.artifact_ref is None
    assert len(result.findings) == 1
    finding = result.findings[0]
    assert finding.code == "raw_html_unsupported"
    assert str(tmp_path) not in finding.user_message
    assert finding.technical_context == ()
    assert not list(repository.artifacts_root.iterdir())


def test_common_markdown_normalizations_publish_typed_findings(tmp_path: Path) -> None:
    source = tmp_path / "normalized.md"
    source.write_text(
        "> 引用\n\n```text\ncode\n```\n\n- [x] 完成\n\n---\n\n"
        "[相对文档](guide.md)",
        encoding="utf-8",
    )
    result = compile_content_material(
        source, ContentArtifactRepository(tmp_path / "artifacts")
    )
    assert result.blocked is False
    by_code = {item.code: item for item in result.findings}
    assert by_code["markdown_blockquote_normalized"].disposition is (
        ContentImportDisposition.NORMALIZED
    )
    assert by_code["markdown_code_block_normalized"].count == 1
    assert by_code["markdown_task_list_normalized"].count == 1
    assert by_code["markdown_horizontal_rule_normalized"].count == 1
    assert by_code["markdown_relative_link_normalized"].disposition is (
        ContentImportDisposition.WARNING
    )


def test_equivalent_markdown_and_docx_converge_on_the_same_ir(tmp_path: Path) -> None:
    markdown = tmp_path / "same.md"
    markdown.write_text("# Title\n\nBody text\n", encoding="utf-8")
    docx = tmp_path / "same.docx"
    document = Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("Body text")
    document.save(docx)
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    md_result = compile_content_material(markdown, repository)
    docx_result = compile_content_material(docx, repository)
    assert md_result.blocked is False
    assert docx_result.blocked is False
    assert md_result.fragment == docx_result.fragment
    assert md_result.fragment.digest == docx_result.fragment.digest
