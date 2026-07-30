from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
import pytest

from src.assistant.domain.docx_format_evidence import (
    FORMAT_EVIDENCE_DISCLOSURE_FIELD,
    STANDARD_FORMAT_REFERENCE_ROLE,
    attachment_disclosure_fields,
    bind_attachment_semantic_roles,
    extract_docx_format_evidence,
    is_format_requirements_request,
)


FORMAT_QUERY = "这个是标准的规划文件，帮我看看确定对应的格式要求"


def _reference_docx(path):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.header_distance = Cm(1.3)
    section.footer_distance = Cm(1.4)
    section.header.paragraphs[0].text = "规划文件标准页眉"
    section.footer.paragraphs[0].text = "第 1 页"

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"),
        "宋体",
    )
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    heading = document.styles["Heading 1"]
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.paragraph_format.space_before = Pt(18)
    heading.paragraph_format.space_after = Pt(10)
    heading.paragraph_format.keep_with_next = True

    document.add_heading("第一章 项目规划", level=1)
    document.add_paragraph("这是正文。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "阶段"
    table.cell(0, 1).text = "目标"
    document.save(path)


def test_format_requirements_intent_binds_reference_role_and_disclosure_field():
    refs = ({"path": "C:/sample.docx", "title": "sample.docx"},)

    bound = bind_attachment_semantic_roles(refs, FORMAT_QUERY)

    assert is_format_requirements_request(FORMAT_QUERY) is True
    assert bound[0]["semantic_role"] == STANDARD_FORMAT_REFERENCE_ROLE
    assert attachment_disclosure_fields(bound) == (
        FORMAT_EVIDENCE_DISCLOSURE_FIELD,
    )
    cleared = bind_attachment_semantic_roles(bound, "请总结正文")
    assert "semantic_role" not in cleared[0]
    assert attachment_disclosure_fields(cleared) == ("document_text",)
    assert (
        is_format_requirements_request(
            "先分析这份标准文件的格式要求，再套用到另一份文档"
        )
        is True
    )
    assert is_format_requirements_request("什么是 Word 格式规范？") is False


def test_format_requirements_intent_does_not_guess_among_multiple_attachments():
    refs = (
        {"path": "C:/standard.docx", "title": "standard.docx"},
        {"path": "C:/target.docx", "title": "target.docx"},
    )

    bound = bind_attachment_semantic_roles(refs, FORMAT_QUERY)

    assert all("semantic_role" not in item for item in bound)
    assert attachment_disclosure_fields(bound) == ("document_text",)


def test_docx_format_evidence_extracts_geometry_styles_and_headers_without_path(
    tmp_path,
):
    path = tmp_path / "标准规划文件.docx"
    _reference_docx(path)

    evidence = extract_docx_format_evidence(path)

    assert evidence["schema_version"] == "docx-format-evidence-v1"
    assert evidence["source"]["name"] == path.name
    assert evidence["source"]["semantic_role"] == STANDARD_FORMAT_REFERENCE_ROLE
    assert str(tmp_path) not in str(evidence)
    assert len(evidence["source"]["sha256"]) == 64
    assert evidence["inventory"]["table_count"] == 1
    assert evidence["sections"][0]["page_width_cm"] == pytest.approx(21.0, abs=0.002)
    assert evidence["sections"][0]["page_height_cm"] == pytest.approx(29.7, abs=0.002)
    margins = evidence["sections"][0]["margins_cm"]
    assert margins["top"] == pytest.approx(2.6, abs=0.002)
    assert margins["right"] == pytest.approx(2.4, abs=0.002)
    assert margins["bottom"] == pytest.approx(2.5, abs=0.002)
    assert margins["left"] == pytest.approx(2.8, abs=0.002)
    assert margins["gutter"] == 0.0
    assert evidence["sections"][0]["headers"]["default"]["text_sample"] == (
        "规划文件标准页眉"
    )
    assert evidence["sections"][0]["footers"]["default"]["text_sample"] == "第 1 页"

    by_name = {row["name"]: row for row in evidence["paragraph_styles"]}
    assert by_name["Normal"]["font"]["east_asia_font"] == "宋体"
    assert by_name["Normal"]["font"]["size_pt"] == 10.5
    assert by_name["Normal"]["paragraph"]["line_spacing"]["value"] == {
        "kind": "multiple",
        "value": 1.5,
    }
    assert by_name["Heading 1"]["semantic_level"] == 1
    assert by_name["Heading 1"]["font"]["size_pt"] == 16.0
    assert by_name["Heading 1"]["font"]["bold"] is True
