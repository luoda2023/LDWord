import json
import os
from pathlib import Path

import pytest
from docx import Document

from src.config.resolved import ResolvedConfig
from src.config.scene import DeliveryPreset, InputSourceProfile
from src.shared.engine import exam_question_schema as exam_runtime_module
from src.shared.engine.exam_master_visual_verification import (
    ExamDocumentVisualQualityResult,
)
from src.shared.engine.exam_question_schema import build_exam_delivery_runtime
from src.shared.engine.fixed_layout_tables import row_height_state
from src.shared.io.artifact_publication import (
    publish_staged_artifacts as publish_staged_artifacts_transactionally,
)


def _config(
    *,
    answered: bool = True,
    include_answer_key: bool = False,
) -> ResolvedConfig:
    question = {
        "stem": "1 + 1 = ?",
        "score": 1,
    }
    if answered:
        question["answer"] = "2"
    payload = {
        "paper_title": "Atomic exam",
        "total_score": 1,
        "sections": [
            {
                "title": "Questions",
                "questions": [question],
            }
        ],
    }
    return ResolvedConfig(
        input_source_profile=InputSourceProfile(
            accepted_formats=["json"],
            structured_formats=["json"],
            material_schema_id="exam_items_v1",
        ),
        entity_data={"exam_items": payload},
        delivery_presets=[
            DeliveryPreset(preset_id="student", label="Student"),
            *(
                [
                    DeliveryPreset(
                        preset_id="answer_key",
                        label="Answer key",
                    )
                ]
                if include_answer_key
                else []
            ),
        ],
    )


def _old_finals(output_root: Path) -> tuple[Path, Path]:
    artifact_dir = output_root / "exam_runtime"
    artifact_dir.mkdir(parents=True)
    preview = artifact_dir / "atomic_preview.md"
    docx = artifact_dir / "atomic_student.docx"
    preview.write_bytes(b"old-preview")
    docx.write_bytes(b"old-docx")
    return preview, docx


def _assert_no_exam_staging(output_root: Path) -> None:
    assert not list(output_root.glob(".exam-runtime-*"))


def test_exam_runtime_publishes_preview_and_docx_as_one_final_set(
    tmp_path: Path,
) -> None:
    runtime = build_exam_delivery_runtime(
        _config(),
        output_dir=tmp_path,
        source_stem="atomic",
    )

    preview = (tmp_path / "exam_runtime" / "atomic_preview.md").resolve()
    docx = (tmp_path / "exam_runtime" / "atomic_student.docx").resolve()
    assert runtime.status == "ok"
    assert Path(runtime.markdown_preview_path) == preview
    assert Path(runtime.rendered_versions[0].docx_path) == docx
    assert preview.is_file()
    assert docx.is_file()
    _assert_no_exam_staging(tmp_path)


def test_exam_runtime_compiles_markdown_before_atomic_docx_publish(
    tmp_path: Path,
) -> None:
    config = _config()
    question = config.entity_data["exam_items"]["sections"][0]["questions"][0]
    question["stem"] = (
        "Read **important** data.\n"
        "| Item | Value |\n"
        "|---|---|\n"
        "| A | 1 |"
    )

    runtime = build_exam_delivery_runtime(
        config,
        output_dir=tmp_path,
        source_stem="markdown",
    )

    assert runtime.status == "ok"
    document = Document(runtime.rendered_versions[0].docx_path)
    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "**" not in body_text
    assert "|---|" not in body_text
    assert any(
        run.text == "important" and run.bold is True
        for paragraph in document.paragraphs
        for run in paragraph.runs
    )
    assert any(
        [[cell.text for cell in row.cells] for row in table.rows]
        == [["Item", "Value"], ["A", "1"]]
        for table in document.tables
    )


def test_exam_answer_sheet_delivery_preserves_fixed_layout_evidence(
    tmp_path: Path,
) -> None:
    config = _config()
    config.delivery_presets = [
        DeliveryPreset(preset_id="answer_sheet", label="Answer sheet")
    ]

    runtime = build_exam_delivery_runtime(
        config,
        output_dir=tmp_path,
        source_stem="answer-sheet",
    )

    versions = {version.preset_id: version for version in runtime.rendered_versions}
    answer_sheet = versions["answer_sheet"]
    sheet_doc = Document(answer_sheet.docx_path)
    assert row_height_state(sheet_doc.tables[0].rows[1]).height_twips == 440
    assert answer_sheet.fixed_layout_kind == "answer_sheet"
    assert answer_sheet.fixed_layout_row_height_twips == 440


def test_exam_visual_gate_checks_every_version_before_atomic_publish(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls = []

    def quality_ok(docx_path, _output_dir, **kwargs):
        calls.append((Path(docx_path).name, dict(kwargs)))
        minimum = int(kwargs["target_page_min"])
        maximum = int(kwargs["target_page_max"])
        return ExamDocumentVisualQualityResult(
            docx_path=Path(docx_path),
            status="quality_ok",
            target_page_min=minimum,
            target_page_max=maximum,
            actual_page_count=minimum,
            substantive_page_count=minimum,
            page_ink_ratios=(0.02,) * minimum,
            structured_response_page_numbers=(minimum,),
            renderer="fake",
        )

    monkeypatch.setattr(
        exam_runtime_module,
        "verify_exam_document_visual",
        quality_ok,
    )

    runtime = build_exam_delivery_runtime(
        _config(include_answer_key=True),
        output_dir=tmp_path,
        source_stem="quality",
        exam_scene_id="my_school_exam",
        exam_scale_profile_id="term",
        verify_visual_quality=True,
    )

    assert runtime.status == "ok"
    assert runtime.quality_status == "quality_ok"
    assert len(calls) == 2
    assert calls[0][1] == {
        "target_page_min": 6,
        "target_page_max": 10,
    }
    assert calls[1][1] == {
        "target_page_min": 1,
        "target_page_max": 10,
    }
    assert all(
        version.visual_quality["status"] == "quality_ok"
        for version in runtime.rendered_versions
    )
    manifest = Path(runtime.quality_manifest_path)
    assert manifest.is_file()
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    assert payload["scale_profile_id"] == "term"
    assert payload["status"] == "quality_ok"
    assert payload["versions"]["student"][
        "structured_response_page_numbers"
    ] == [6]
    assert payload["semantic_review_status"] == "manual_review_required"
    assert (
        payload["material_grounding_review_status"]
        == "manual_review_required"
    )
    _assert_no_exam_staging(tmp_path)


def test_exam_visual_warning_publishes_review_candidate_without_replacing_final_set(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    preview, docx = _old_finals(tmp_path)

    def quality_failed(docx_path, _output_dir, **kwargs):
        return ExamDocumentVisualQualityResult(
            docx_path=Path(docx_path),
            status="quality_failed",
            target_page_min=int(kwargs["target_page_min"]),
            target_page_max=int(kwargs["target_page_max"]),
            actual_page_count=9,
            substantive_page_count=9,
            issues=("page_count_above_target:9>8",),
            renderer="fake",
        )

    monkeypatch.setattr(
        exam_runtime_module,
        "verify_exam_document_visual",
        quality_failed,
    )

    runtime = build_exam_delivery_runtime(
        _config(),
        output_dir=tmp_path,
        source_stem="atomic",
        exam_scale_profile_id="standard",
        verify_visual_quality=True,
    )

    assert runtime.status == "warning"
    assert runtime.skipped_reason == ""
    assert runtime.quality_status == "quality_review_required"
    assert runtime.release_tier == "review"
    assert runtime.quality_issues == (
        "student:page_count_above_target:9>8",
    )
    assert preview.read_bytes() == b"old-preview"
    assert docx.read_bytes() == b"old-docx"
    review_dir = tmp_path / "exam_runtime" / "review_candidates"
    assert Path(runtime.markdown_preview_path) == (
        review_dir / "atomic_preview.md"
    ).resolve()
    assert Path(runtime.rendered_versions[0].docx_path) == (
        review_dir / "atomic_student.docx"
    ).resolve()
    assert Path(runtime.markdown_preview_path).is_file()
    assert Path(runtime.rendered_versions[0].docx_path).is_file()
    assert Path(runtime.quality_manifest_path).parent == review_dir.resolve()
    payload = json.loads(
        Path(runtime.quality_manifest_path).read_text(encoding="utf-8")
    )
    assert payload["release_tier"] == "review"
    _assert_no_exam_staging(tmp_path)


def test_exam_renderer_unavailable_publishes_as_quality_unverified(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def unavailable(docx_path, _output_dir, **kwargs):
        return ExamDocumentVisualQualityResult(
            docx_path=Path(docx_path),
            status="renderer_unavailable",
            target_page_min=int(kwargs["target_page_min"]),
            target_page_max=int(kwargs["target_page_max"]),
            issues=("docx_to_pdf_renderer_unavailable",),
        )

    monkeypatch.setattr(
        exam_runtime_module,
        "verify_exam_document_visual",
        unavailable,
    )

    runtime = build_exam_delivery_runtime(
        _config(),
        output_dir=tmp_path,
        source_stem="unverified",
        exam_scale_profile_id="quiz",
        verify_visual_quality=True,
    )

    assert runtime.status == "warning"
    assert runtime.quality_status == "quality_unverified"
    assert runtime.release_tier == "review"
    assert Path(runtime.rendered_versions[0].docx_path).is_file()
    assert Path(runtime.quality_manifest_path).is_file()
    assert Path(runtime.rendered_versions[0].docx_path).parent.name == (
        "review_candidates"
    )
    _assert_no_exam_staging(tmp_path)


@pytest.mark.parametrize(
    ("result_status", "result_issue"),
    [
        ("docx_missing", "docx_missing"),
        ("docx_corrupt", "docx_package_corrupt"),
    ],
)
def test_exam_missing_or_corrupt_docx_remains_a_fatal_publication_block(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    result_status: str,
    result_issue: str,
) -> None:
    preview, docx = _old_finals(tmp_path)

    def fatal(docx_path, _output_dir, **kwargs):
        return ExamDocumentVisualQualityResult(
            docx_path=Path(docx_path),
            status=result_status,
            target_page_min=int(kwargs["target_page_min"]),
            target_page_max=int(kwargs["target_page_max"]),
            issues=(result_issue,),
        )

    monkeypatch.setattr(
        exam_runtime_module,
        "verify_exam_document_visual",
        fatal,
    )

    runtime = build_exam_delivery_runtime(
        _config(),
        output_dir=tmp_path,
        source_stem="atomic",
        verify_visual_quality=True,
    )

    assert runtime.status == "blocked"
    assert runtime.skipped_reason == "visual_quality_fatal"
    assert runtime.quality_status == "quality_fatal"
    assert runtime.release_tier == "none"
    assert preview.read_bytes() == b"old-preview"
    assert docx.read_bytes() == b"old-docx"
    assert not (tmp_path / "exam_runtime" / "review_candidates").exists()
    _assert_no_exam_staging(tmp_path)


def test_exam_validation_error_does_not_publish_or_replace_old_finals(
    tmp_path: Path,
) -> None:
    preview, docx = _old_finals(tmp_path)

    runtime = build_exam_delivery_runtime(
        _config(answered=False),
        output_dir=tmp_path,
        source_stem="atomic",
    )

    assert runtime.status == "blocked"
    assert runtime.skipped_reason == "schema_validation_errors"
    assert runtime.markdown_preview_path == ""
    assert runtime.rendered_versions == ()
    assert preview.read_bytes() == b"old-preview"
    assert docx.read_bytes() == b"old-docx"
    _assert_no_exam_staging(tmp_path)


def test_exam_render_failure_leaves_no_partial_final_and_preserves_old_finals(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    preview, docx = _old_finals(tmp_path)

    def fail_after_partial_stage(_payload, _preset, docx_path, *, label):
        del label
        Path(docx_path).write_bytes(b"partial-staging-docx")
        raise RuntimeError("injected render failure")

    monkeypatch.setattr(
        exam_runtime_module,
        "_render_exam_version_docx",
        fail_after_partial_stage,
    )

    with pytest.raises(RuntimeError, match="injected render failure"):
        build_exam_delivery_runtime(
            _config(),
            output_dir=tmp_path,
            source_stem="atomic",
        )

    assert preview.read_bytes() == b"old-preview"
    assert docx.read_bytes() == b"old-docx"
    _assert_no_exam_staging(tmp_path)


def test_exam_publish_failure_rolls_back_every_old_final(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    preview, docx = _old_finals(tmp_path)

    class FailSecondCandidateReplace:
        def __init__(self) -> None:
            self.candidate_replaces = 0

        def __call__(self, source: str | Path, target: str | Path) -> None:
            if ".stage" in Path(source).name:
                self.candidate_replaces += 1
                if self.candidate_replaces == 2:
                    raise PermissionError("injected second publish failure")
            os.replace(source, target)

    failing_replace = FailSecondCandidateReplace()

    def publish_with_failure(artifacts, **kwargs):
        return publish_staged_artifacts_transactionally(
            artifacts,
            **kwargs,
            atomic_replace=failing_replace,
        )

    monkeypatch.setattr(
        exam_runtime_module,
        "publish_staged_artifacts",
        publish_with_failure,
    )

    with pytest.raises(PermissionError, match="injected second publish failure"):
        build_exam_delivery_runtime(
            _config(),
            output_dir=tmp_path,
            source_stem="atomic",
        )

    assert failing_replace.candidate_replaces == 2
    assert preview.read_bytes() == b"old-preview"
    assert docx.read_bytes() == b"old-docx"
    _assert_no_exam_staging(tmp_path)
