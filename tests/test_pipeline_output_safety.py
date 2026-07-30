from __future__ import annotations

import os
from pathlib import Path

import pytest
from docx import Document

from src.config.feature_configs import OutputConfig
from src.config.material_context import MaterialExecutionContext
from src.config.resolver import resolve_config
from src.config.scene import DeliveryPreset, SceneWorkspace
from src.config.template import TemplateConfig
from src.modules.base import BaseModule, ModuleMeta
from src.modules.fill.entity_fill import EntityFillModule
import src.pipeline.runner as runner_module
from src.pipeline.runner import Pipeline
import src.services.production_runtime.execution_runtime as execution_runtime
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner


class _MutationProbeModule(BaseModule):
    meta = ModuleMeta(
        name="output_safety_mutation_probe",
        description="Output safety mutation probe",
        category="test",
        execution_phase="fill",
        scope_behavior="document_level",
    )

    def __init__(self, applied: list[str]) -> None:
        self._applied = applied

    def apply(self, doc, config, tracker, context) -> None:
        self._applied.append("applied")
        doc.add_paragraph("MUTATED")


def _save_text_docx(path: Path, text: str) -> bytes:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    return path.read_bytes()


def _delivery_scene() -> SceneWorkspace:
    scene = SceneWorkspace(
        default_delivery_preset_id="first",
        delivery_presets=[
            DeliveryPreset(
                preset_id="first",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
            ),
            DeliveryPreset(
                preset_id="second",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    return scene


def _assert_no_transaction_payloads(root: Path) -> None:
    names = [path.name for path in root.rglob("*")]
    assert not any(name.endswith(".stage.docx") for name in names)
    assert not any(name.endswith(".backup") for name in names)
    assert not any(name.startswith(".material-exec-") for name in names)


def test_pipeline_blocks_lexical_source_alias_before_module_mutation(tmp_path):
    source = tmp_path / "source.docx"
    original = _save_text_docx(source, "ORIGINAL")
    alias_parent = tmp_path / "alias-parent"
    alias_parent.mkdir()
    source_alias = alias_parent / ".." / source.name
    applied: list[str] = []
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.enabled = False

    result = Pipeline(
        modules=[_MutationProbeModule(applied)],
        config=resolve_config(TemplateConfig(), scene),
        output_path_overrides={"final": source_alias},
    ).execute(str(source))

    assert result.success is False
    assert result.output_paths == {}
    assert applied == []
    assert source.read_bytes() == original
    assert result.context.output_target_preflight.has_errors is True
    assert {
        issue.kind
        for item in result.context.output_target_preflight.items
        for issue in item.issues
    } >= {"source_overwrite", "target_exists"}
    failures = result.tracker.get_failures()
    assert [item.change_type for item in failures] == ["output_target_blocked"]


def test_pipeline_blocks_existing_hardlink_alias_of_source(tmp_path):
    source = tmp_path / "source.docx"
    original = _save_text_docx(source, "ORIGINAL")
    hardlink = tmp_path / "source-hardlink.docx"
    try:
        os.link(source, hardlink)
    except OSError as exc:  # pragma: no cover - filesystem capability
        pytest.skip(f"hard links unavailable: {exc}")
    applied: list[str] = []
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.enabled = False

    result = Pipeline(
        modules=[_MutationProbeModule(applied)],
        config=resolve_config(TemplateConfig(), scene),
        output_path_overrides={"final": hardlink},
    ).execute(str(source))

    assert result.success is False
    assert applied == []
    assert source.read_bytes() == original
    assert hardlink.read_bytes() == original
    assert any(
        issue.kind == "source_overwrite"
        for item in result.context.output_target_preflight.items
        for issue in item.issues
    )


def test_existing_single_target_is_advisory_and_replaced_after_staging(tmp_path):
    source = tmp_path / "source.docx"
    _save_text_docx(source, "NEW")
    output_dir = tmp_path / "out"
    final = output_dir / "source_formatted.docx"
    old_bytes = _save_text_docx(final, "OLD")
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.enabled = False

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(output_dir),
        output_suffix="_formatted",
    ).execute(str(source))

    assert result.success is True
    assert result.output_paths == {"final": str(final)}
    assert final.read_bytes() != old_bytes
    assert [paragraph.text for paragraph in Document(final).paragraphs] == ["NEW"]
    issues = result.context.output_target_preflight.items[0].issues
    assert [(issue.kind, issue.severity) for issue in issues] == [
        ("target_exists", "warning")
    ]
    warnings = [
        item
        for item in result.tracker.get_all()
        if item.rule_name == "output_target_preflight"
    ]
    assert len(warnings) == 1
    assert warnings[0].change_type == "output_target_warning"
    assert warnings[0].success is True
    _assert_no_transaction_payloads(output_dir)


def test_multi_output_staging_failure_preserves_all_old_finals(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    _save_text_docx(source, "NEW")
    output_dir = tmp_path / "out"
    first = output_dir / "source_first.docx"
    second = output_dir / "source_second.docx"
    old_first = _save_text_docx(first, "OLD FIRST")
    old_second = _save_text_docx(second, "OLD SECOND")

    original_save = Pipeline._save_staged_document
    calls = 0

    def _fail_second_stage(self, doc, stage_path):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("simulated second-stage failure")
        return original_save(self, doc, stage_path)

    monkeypatch.setattr(Pipeline, "_save_staged_document", _fail_second_stage)
    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), _delivery_scene()),
        output_dir=str(output_dir),
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is False
    assert result.output_paths == {}
    assert "simulated second-stage failure" in result.error
    assert first.read_bytes() == old_first
    assert second.read_bytes() == old_second
    _assert_no_transaction_payloads(output_dir)


def test_cancel_during_output_staging_preserves_old_final_and_returns_cancelled(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    _save_text_docx(source, "NEW")
    output_dir = tmp_path / "out"
    final = output_dir / "source_formatted.docx"
    old_bytes = _save_text_docx(final, "OLD")
    cancelled = False
    original_save = Pipeline._save_staged_document

    def _save_then_cancel(self, doc, stage_path):
        nonlocal cancelled
        original_save(self, doc, stage_path)
        cancelled = True

    monkeypatch.setattr(Pipeline, "_save_staged_document", _save_then_cancel)
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.enabled = False

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(output_dir),
        output_suffix="_formatted",
        cancel_check=lambda: cancelled,
    ).execute(str(source))

    assert result.status == "cancelled"
    assert result.cancelled is True
    assert result.output_paths == {}
    assert final.read_bytes() == old_bytes
    _assert_no_transaction_payloads(output_dir)


def test_cancel_after_first_delivery_stage_preserves_complete_old_final_set(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    _save_text_docx(source, "NEW")
    output_dir = tmp_path / "out"
    first = output_dir / "source_first.docx"
    second = output_dir / "source_second.docx"
    old_first = _save_text_docx(first, "OLD FIRST")
    old_second = _save_text_docx(second, "OLD SECOND")
    cancelled = False
    original_save = Pipeline._save_staged_document

    def _save_first_then_cancel(self, doc, stage_path):
        nonlocal cancelled
        original_save(self, doc, stage_path)
        cancelled = True

    monkeypatch.setattr(Pipeline, "_save_staged_document", _save_first_then_cancel)

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), _delivery_scene()),
        output_dir=str(output_dir),
        force_delivery_presets=True,
        cancel_check=lambda: cancelled,
    ).execute(str(source))

    assert result.status == "cancelled"
    assert result.output_paths == {}
    assert first.read_bytes() == old_first
    assert second.read_bytes() == old_second
    _assert_no_transaction_payloads(output_dir)


def test_cancel_observed_after_publish_rolls_back_before_backup_release(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    _save_text_docx(source, "NEW")
    output_dir = tmp_path / "out"
    final = output_dir / "source_formatted.docx"
    old_bytes = _save_text_docx(final, "OLD")
    cancelled = False
    original_publish = runner_module.OwnedAssemblyTransaction.publish

    def _publish_then_cancel(self, candidates):
        nonlocal cancelled
        original_publish(self, candidates)
        cancelled = True

    monkeypatch.setattr(
        runner_module.OwnedAssemblyTransaction,
        "publish",
        _publish_then_cancel,
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.enabled = False

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(output_dir),
        output_suffix="_formatted",
        cancel_check=lambda: cancelled,
    ).execute(str(source))

    assert result.status == "cancelled"
    assert result.cancelled is True
    assert result.output_paths == {}
    assert final.read_bytes() == old_bytes
    _assert_no_transaction_payloads(output_dir)


def test_cancel_after_durable_commit_keeps_published_output_as_evidence(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    _save_text_docx(source, "NEW")
    output_dir = tmp_path / "out"
    final = output_dir / "source_formatted.docx"
    old_bytes = _save_text_docx(final, "OLD")
    cancelled = False
    original_release = runner_module.OwnedAssemblyTransaction.release_backups

    def _release_then_cancel(self):
        nonlocal cancelled
        original_release(self)
        cancelled = True

    monkeypatch.setattr(
        runner_module.OwnedAssemblyTransaction,
        "release_backups",
        _release_then_cancel,
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.enabled = False

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(output_dir),
        output_suffix="_formatted",
        cancel_check=lambda: cancelled,
    ).execute(str(source))

    assert result.status == "cancelled"
    assert result.cancelled is True
    assert result.output_paths == {"final": str(final)}
    assert final.read_bytes() != old_bytes
    assert [paragraph.text for paragraph in Document(final).paragraphs] == ["NEW"]
    _assert_no_transaction_payloads(output_dir)


def test_multi_output_publish_failure_rolls_back_complete_final_set(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    _save_text_docx(source, "NEW")
    output_dir = tmp_path / "out"
    first = output_dir / "source_first.docx"
    second = output_dir / "source_second.docx"
    old_first = _save_text_docx(first, "OLD FIRST")
    old_second = _save_text_docx(second, "OLD SECOND")

    real_transaction = runner_module.OwnedAssemblyTransaction
    publish_moves = 0

    def _replace_with_second_publish_failure(source_path, destination_path):
        nonlocal publish_moves
        source_candidate = Path(source_path)
        if source_candidate.name.endswith(".stage.docx"):
            publish_moves += 1
            if publish_moves == 2:
                raise OSError("simulated second-publish failure")
        return os.replace(source_path, destination_path)

    def _transaction_factory(**kwargs):
        return real_transaction(
            **kwargs,
            atomic_replace=_replace_with_second_publish_failure,
        )

    monkeypatch.setattr(
        runner_module,
        "OwnedAssemblyTransaction",
        _transaction_factory,
    )
    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), _delivery_scene()),
        output_dir=str(output_dir),
        force_delivery_presets=True,
    ).execute(str(source))

    assert publish_moves == 2
    assert result.success is False
    assert result.output_paths == {}
    assert "simulated second-publish failure" in result.error
    assert first.read_bytes() == old_first
    assert second.read_bytes() == old_second
    _assert_no_transaction_payloads(output_dir)


def test_field_only_workbench_route_blocks_source_overwrite_before_fill(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    original = _save_text_docx(source, "Company={{@text:company}}")
    scene = SceneWorkspace(
        default_delivery_preset_id="final",
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                output_dir_template="{document_dir}",
                filename_template="{stem}",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
            ),
            DeliveryPreset(
                preset_id="review",
                output_dir_template="{document_dir}",
                filename_template="{stem}",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    scene.module_switches["entity_fill"] = True
    monkeypatch.setattr(
        execution_runtime,
        "create_all_modules",
        lambda: [EntityFillModule()],
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={"company": "ACME"},
            exact_material_placeholders=True,
        ),
        output_dir=tmp_path,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert payload["material_assembly"]["status"] == "not_run"
    assert source.read_bytes() == original
    assert [paragraph.text for paragraph in Document(source).paragraphs] == [
        "Company={{@text:company}}"
    ]
    issue_kinds = {
        issue["kind"]
        for item in payload["output_target_preflight"]["items"]
        for issue in item["issues"]
    }
    assert issue_kinds >= {"duplicate_target", "source_overwrite"}
