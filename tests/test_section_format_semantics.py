import sys
from pathlib import Path
from types import SimpleNamespace

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.feature_configs import PageNumberPhaseConfig
from src.config.resolved import ResolvedConfig
from src.modules.basic.section_format import SectionFormatModule
from src.modules.structure.heading_recognition import DocSection
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn
from src.shared.engine.section_layout_planner import collect_section_inventory


def _fake_doc_tree(sections):
    return SimpleNamespace(
        sections=sections,
        get_section=lambda name: next((section for section in sections if section.section_type == name), None),
    )


def _section_break_type(para):
    ppr = para._element.find(qn("w:pPr"))
    if ppr is None:
        return None
    sect_pr = ppr.find(qn("w:sectPr"))
    if sect_pr is None:
        return None
    type_elem = sect_pr.find(qn("w:type"))
    if type_elem is None:
        return "nextPage"
    return type_elem.get(qn("w:val"))


def test_section_format_inserts_semantic_boundaries_without_splitting_body_chapters():
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_heading("Chapter 1 Intro", level=1)
    doc.add_paragraph("Body paragraph")
    doc.add_heading("Chapter 2 Method", level=1)
    doc.add_paragraph("More body")
    doc.add_paragraph("References")
    doc.add_paragraph("[1] Ref entry")

    sections = [
        DocSection("cover", 0, 1, confidence=10.0),
        DocSection("body", 1, 5, confidence=10.0),
        DocSection("references", 5, 7, confidence=10.0),
    ]
    context = PipelineContext(
        doc_tree=_fake_doc_tree(sections),
        heading_map={1: 1, 3: 1},
    )

    SectionFormatModule().apply(doc, ResolvedConfig(), ChangeTracker(), context)

    assert _section_break_type(doc.paragraphs[0]) == "nextPage"
    assert _section_break_type(doc.paragraphs[2]) is None
    assert _section_break_type(doc.paragraphs[4]) == "nextPage"


def test_section_format_inserts_minimal_carrier_when_table_sits_before_semantic_start():
    doc = Document()
    doc.add_paragraph("Cover")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "value"
    doc.add_paragraph("Chapter 1 Intro")
    doc.sections[0].header.paragraphs[0].text = "source header"

    sections = [
        DocSection("cover", 0, 1, confidence=10.0),
        DocSection("body", 1, 2, confidence=10.0),
    ]
    context = PipelineContext(
        doc_tree=_fake_doc_tree(sections),
        heading_map={1: 1},
    )

    SectionFormatModule().apply(doc, ResolvedConfig(), ChangeTracker(), context)

    assert _section_break_type(doc.paragraphs[0]) is None
    assert len(doc.sections) == 2
    assert doc.paragraphs[1].text == ""
    assert _section_break_type(doc.paragraphs[1]) == "nextPage"
    body_children = list(doc.element.body)
    assert body_children.index(table._element) < body_children.index(doc.paragraphs[1]._element)
    assert body_children.index(doc.paragraphs[1]._element) < body_children.index(
        doc.paragraphs[2]._element
    )
    assert context.section_execution_receipt.applied_operations[0].action == (
        "insert_carrier_break"
    )
    first_boundary, final_boundary = collect_section_inventory(doc).boundaries
    assert not any(ref[0] == "header" for ref in first_boundary.header_footer_refs)
    assert any(ref[0] == "header" for ref in final_boundary.header_footer_refs)

    context.doc_tree = _fake_doc_tree(
        [
            DocSection("cover", 0, 2, confidence=10.0),
            DocSection("body", 2, 3, confidence=10.0),
        ]
    )
    SectionFormatModule().apply(doc, ResolvedConfig(), ChangeTracker(), context)
    assert len(doc.sections) == 2
    assert context.section_execution_receipt.applied_operations == ()


def test_page_number_transition_keeps_next_page_break_over_global_continuous_setting():
    doc = Document()
    doc.add_paragraph("目录")
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文")

    sections = [
        DocSection("toc", 0, 1, confidence=10.0),
        DocSection("body", 1, 3, confidence=10.0),
    ]
    context = PipelineContext(
        doc_tree=_fake_doc_tree(sections),
        heading_map={1: 1},
    )
    config = ResolvedConfig()
    config.section.section_break_type = "continuous"
    config.header_footer.page_number_plan.phases = [
        PageNumberPhaseConfig(
            phase_id="toc",
            selectors=["toc"],
            visible=True,
            number_format="upperRoman",
            start_mode="restart",
            start_value=1,
        ),
        PageNumberPhaseConfig(
            phase_id="body",
            selectors=["body"],
            visible=True,
            number_format="decimal",
            start_mode="restart",
            start_value=1,
        ),
    ]

    SectionFormatModule().apply(doc, config, ChangeTracker(), context)

    assert _section_break_type(doc.paragraphs[0]) == "nextPage"
