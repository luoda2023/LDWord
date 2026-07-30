from __future__ import annotations

from pathlib import Path

from docx import Document

from src.config.feature_configs import OutputConfig
from src.config.resolver import resolve_config
from src.config.resolved import ResolvedConfig
from src.config.scene import DeliveryPreset, SceneWorkspace
from src.config.template import TemplateConfig
from src.pipeline.runner import Pipeline, plan_pipeline_output_paths


def _final_output_config() -> ResolvedConfig:
    return ResolvedConfig(
        output=OutputConfig(
            final_docx=True,
            compare_docx=False,
            report_json=False,
            report_markdown=False,
        )
    )


def test_output_planning_is_read_only_and_uses_logical_source(tmp_path: Path):
    source = tmp_path / "项目技术标.docx"
    planned = plan_pipeline_output_paths(
        source,
        _final_output_config(),
        output_dir=tmp_path / "delivery",
        output_suffix="_formatted",
    )

    assert planned == {
        "final": str(tmp_path / "delivery" / "项目技术标_formatted.docx")
    }
    assert not (tmp_path / "delivery").exists()


def test_delivery_output_planning_preserves_docx_suffix_for_dotted_stem(
    tmp_path: Path,
):
    source = tmp_path / "generated.bidding.docx"
    scene = SceneWorkspace(
        default_delivery_preset_id="original",
        delivery_presets=[
            DeliveryPreset(
                preset_id="original",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True),
            )
        ],
    )
    config = resolve_config(TemplateConfig(), scene)

    planned = plan_pipeline_output_paths(
        source,
        config,
        output_dir=tmp_path / "delivery",
        force_delivery_presets=True,
    )

    assert planned == {
        "original": str(
            tmp_path / "delivery" / "generated.bidding_original.docx"
        )
    }


def test_pipeline_can_write_owned_stage_while_preserving_logical_source(tmp_path: Path):
    logical_source = tmp_path / "投标文件.docx"
    physical_source = tmp_path / ".material-work" / logical_source.name
    physical_source.parent.mkdir()
    document = Document()
    document.add_paragraph("正文")
    document.save(physical_source)

    stage = tmp_path / ".投标文件_formatted.docx.lark-stage.docx"
    result = Pipeline(
        [],
        _final_output_config(),
        output_dir=str(tmp_path),
        output_suffix="_formatted",
        logical_source_path=logical_source,
        output_path_overrides={"final": stage},
    ).execute(str(physical_source))

    assert result.success
    assert result.output_paths == {"final": str(stage)}
    assert stage.is_file()
    assert not (tmp_path / "投标文件_formatted.docx").exists()
    assert result.context is not None
    assert result.context.source_doc_path == str(logical_source)
    assert result.context.source_doc_dir == str(logical_source.parent)
    assert result.context.working_doc_path == str(physical_source)
    assert result.context.working_doc_dir == str(physical_source.parent)


def test_pipeline_output_override_creates_its_parent(tmp_path: Path):
    source = tmp_path / "source.docx"
    Document().save(source)
    stage = tmp_path / "nested" / ".owned-stage.docx"

    result = Pipeline(
        [],
        _final_output_config(),
        output_path_overrides={"final": stage},
    ).execute(str(source))

    assert result.success
    assert stage.is_file()
