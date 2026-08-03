import sys
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.dataclass_utils import dict_to_dataclass
from src.config.feature_configs import HeaderFooterConfig, PageNumberPhaseConfig
from src.config.migration import normalize_template_payload
from src.config.resolved import ResolvedConfig
from src.modules.basic.header_footer import (
    HeaderFooterModule,
    _paragraph_has_field,
    _set_header_border,
    _set_page_number,
)
from src.modules.basic.section_format import _ensure_section_break_before_paragraph
from src.modules.structure.heading_recognition import DocSection, DocTree
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.field_builder import iter_field_instructions
from src.shared.engine.font_resolver import resolve_font
from src.shared.engine.ooxml_ops import qn


def test_header_footer_none_clears_existing_header_and_border():
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].add_run("Legacy Header")
    _set_header_border(section, True)

    config = ResolvedConfig()
    config.header_footer.header_mode = "none"
    config.header_footer.header_border = True

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    para = section.header.paragraphs[0]
    bottom = para._element.find(qn("w:pPr")).find(qn("w:pBdr")).find(qn("w:bottom"))

    assert para.text == ""
    assert bottom is not None
    assert bottom.get(qn("w:val")) == "none"


def test_header_footer_page_number_disabled_removes_existing_page_field():
    doc = Document()
    section = doc.sections[0]

    config = ResolvedConfig()
    _set_page_number(section, config.header_footer)
    assert _paragraph_has_field(section.footer.paragraphs[0], "PAGE") is True

    config.header_footer.page_number_enabled = False
    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert _paragraph_has_field(section.footer.paragraphs[0], "PAGE") is False
    assert section.footer.paragraphs[0].text == ""


def test_header_footer_fixed_footer_text_and_alignment_without_page_number():
    doc = Document()
    config = ResolvedConfig()
    config.header_footer.footer.content_mode = "fixed"
    config.header_footer.footer_text = "Confidential"
    config.header_footer.footer_alignment = "right"
    config.header_footer.page_number_enabled = False

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    para = doc.sections[0].footer.paragraphs[0]
    assert para.text == "Confidential"
    assert _paragraph_has_field(para, "PAGE") is False
    assert para._element.find(qn("w:pPr")).find(qn("w:jc")).get(qn("w:val")) == "right"


def test_header_footer_page_number_can_append_footer_text():
    doc = Document()
    config = ResolvedConfig()
    config.header_footer.footer.content_mode = "page_number_with_text"
    config.header_footer.footer_text = "Confidential"
    config.header_footer.footer_alignment = "left"
    config.header_footer.page_number_alignment = "left"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    para = doc.sections[0].footer.paragraphs[0]
    assert _paragraph_has_field(para, "PAGE") is True
    assert "Confidential" in para.text
    assert para._element.find(qn("w:pPr")).find(qn("w:jc")).get(qn("w:val")) == "left"


def test_footer_text_switch_does_not_disable_independent_page_number():
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].add_run("Legacy Header")
    section.first_page_header.paragraphs[0].add_run("Legacy First Header")
    section.even_page_footer.paragraphs[0].add_run("Legacy Even Footer")
    _set_header_border(section, True)
    _set_page_number(section, ResolvedConfig().header_footer)

    config = ResolvedConfig()
    hf = config.header_footer
    hf.behavior.different_first_page = True
    hf.behavior.different_odd_even_pages = True
    hf.header.enabled = False
    hf.header_mode = "fixed"
    hf.header_text = "Hidden Header"
    hf.header_border = True
    hf.footer.enabled = False
    hf.footer.content_mode = "page_number_with_text"
    hf.footer_text = "Hidden Footer"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    header_para = section.header.paragraphs[0]
    footer_para = section.footer.paragraphs[0]
    bottom = header_para._element.find(qn("w:pPr")).find(qn("w:pBdr")).find(qn("w:bottom"))
    assert header_para.text == ""
    assert bottom is not None
    assert bottom.get(qn("w:val")) == "none"
    assert "Hidden Footer" not in footer_para.text
    assert section.first_page_header.paragraphs[0].text == ""
    assert "Hidden Footer" not in section.even_page_footer.paragraphs[0].text
    assert _paragraph_has_field(footer_para, "PAGE") is True
    assert hf.header_mode == "fixed"
    assert hf.header_text == "Hidden Header"
    assert hf.footer.content_mode == "page_number_with_text"
    assert hf.footer_text == "Hidden Footer"


def test_header_footer_fixed_header_applies_emphasis_typography():
    doc = Document()
    config = ResolvedConfig()
    config.header_footer.header_mode = "fixed"
    config.header_footer.header_text = "Fixed Header"
    config.header_footer.header_alignment = "right"
    config.header_footer.font_cn = "SimSun"
    config.header_footer.font_en = "Times New Roman"
    config.header_footer.size_pt = 9
    config.header_footer.bold = True
    config.header_footer.italic = True

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    run = doc.sections[0].header.paragraphs[0].runs[0]
    assert _paragraph_alignment(doc.sections[0].header.paragraphs[0]) == "right"
    assert run.font.size.pt == 9
    assert run.font.bold is True
    assert run.font.italic is True


def test_header_footer_header_and_footer_typography_are_independent():
    doc = Document()
    config = ResolvedConfig()
    hf = config.header_footer
    hf.header_mode = "fixed"
    hf.header_text = "Fixed Header"
    hf.footer.content_mode = "fixed"
    hf.footer_text = "Fixed Footer"
    hf.page_number_enabled = False
    hf.header.typography.size_pt = 9
    hf.header.typography.bold = True
    hf.footer.typography.size_pt = 11
    hf.footer.typography.italic = True

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    header_run = doc.sections[0].header.paragraphs[0].runs[0]
    footer_run = doc.sections[0].footer.paragraphs[0].runs[0]
    assert header_run.font.size.pt == 9
    assert header_run.font.bold is True
    assert header_run.font.italic is False
    assert footer_run.font.size.pt == 11
    assert footer_run.font.bold is False
    assert footer_run.font.italic is True


def test_default_header_footer_typography_is_written_to_document():
    doc = Document()
    config = ResolvedConfig()
    config.header_footer.header_mode = "fixed"
    config.header_footer.header_text = "默认页眉"
    config.header_footer.footer.content_mode = "fixed"
    config.header_footer.footer_text = "Default Footer"
    config.header_footer.page_number_enabled = False

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    runs = (
        doc.sections[0].header.paragraphs[0].runs[0],
        doc.sections[0].footer.paragraphs[0].runs[0],
    )
    for run in runs:
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        assert run.font.name == resolve_font("Times New Roman", lang="en")
        assert fonts.get(qn("w:eastAsia")) == resolve_font("宋体", lang="cn")
        assert run.font.size.pt == 10.5
        assert run.font.bold is False
        assert run.font.italic is False


def test_header_footer_config_exposes_word_variant_model():
    cfg = HeaderFooterConfig()

    assert cfg.behavior.different_first_page is False
    assert cfg.behavior.different_odd_even_pages is False
    assert cfg.behavior.link_to_previous == "never"
    assert cfg.variants.default.header.mode == "inherit"
    assert cfg.variants.first.header.mode == "none"
    assert cfg.variants.first.footer.mode == "none"
    assert cfg.variants.even.footer.mode == "inherit"
    assert cfg.page_number_plan.first.visibility == "inherit"
    assert cfg.page_number_plan.first.template == ""
    assert cfg.page_number_plan.first.alignment == "inherit"
    assert cfg.page_number_plan.even.visibility == "inherit"
    assert cfg.header.border_style.width_pt == 0.5
    assert cfg.page_number_template == "{page}"
    assert cfg.header.typography is not cfg.footer.typography
    assert cfg.header.typography.font_cn == "宋体"
    assert cfg.header.typography.font_en == "Times New Roman"
    assert cfg.header.typography.size_pt == 10.5
    assert cfg.footer.typography.font_cn == "宋体"
    assert cfg.footer.typography.font_en == "Times New Roman"
    assert cfg.footer.typography.size_pt == 10.5


def test_header_footer_preserve_strategy_keeps_source_section_links():
    doc = Document()
    doc.add_paragraph("第一节")
    doc.add_section()
    doc.add_paragraph("第二节")
    assert doc.sections[1].header.is_linked_to_previous is True
    assert doc.sections[1].footer.is_linked_to_previous is True

    config = ResolvedConfig()
    config.header_footer.behavior.link_to_previous = "preserve"
    config.header_footer.header_mode = "fixed"
    config.header_footer.header_text = "共享页眉"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert doc.sections[1].header.is_linked_to_previous is True
    assert doc.sections[1].footer.is_linked_to_previous is True
    assert doc.sections[1].header.paragraphs[0].text == "共享页眉"


def test_template_normalization_moves_feature_policies_out_of_section():
    payload = normalize_template_payload(
        {
            "section": {
                "boundary_mode": "preserve_source",
                "section_break_type": None,
                "empty_break_policy": "preserve",
                "caption_table_break_policy": "remove_proven_redundant",
                "header_footer_link_mode": "preserve_source",
            },
            "caption": {},
            "header_footer": {"behavior": {"link_to_previous": "never"}},
        }
    )

    assert set(payload["section"]) == {
        "boundary_mode",
        "section_break_type",
        "empty_break_policy",
    }
    assert payload["caption"]["table_break_policy"] == "remove_proven_redundant"
    assert payload["header_footer"]["behavior"]["link_to_previous"] == "preserve"


def test_header_footer_template_normalization_preserves_word_variant_model():
    payload = normalize_template_payload(
        {
            "header_footer": {
                "behavior": {
                    "different_first_page": True,
                    "different_odd_even_pages": True,
                    "link_to_previous": "preserve",
                },
                "variants": {
                    "first": {
                        "header": {"mode": "fixed", "fixed_text": "First Header"},
                        "footer": {"mode": "none"},
                    },
                    "even": {
                        "footer": {
                            "mode": "template",
                            "template": "第 {page} 页 / 共 {pages} 页",
                        }
                    },
                },
                "footer": {"page_number_template": "第 {page} 页"},
            }
        }
    )

    cfg = dict_to_dataclass(HeaderFooterConfig, payload["header_footer"])

    assert cfg.behavior.different_first_page is True
    assert cfg.behavior.different_odd_even_pages is True
    assert cfg.behavior.link_to_previous == "preserve"
    assert cfg.variants.first.header.mode == "fixed"
    assert cfg.variants.first.header.fixed_text == "First Header"
    assert cfg.variants.even.footer.mode == "none"
    assert cfg.variants.even.footer.template == ""
    assert cfg.page_number_plan.even.visibility == "show"
    assert cfg.page_number_plan.even.template == "第 {page} 页 / 共 {pages} 页"
    assert cfg.page_number_template == "第 {page} 页"


def test_header_footer_applies_first_and_even_page_variants():
    doc = Document()
    config = ResolvedConfig()
    hf = config.header_footer
    hf.behavior.different_first_page = True
    hf.behavior.different_odd_even_pages = True
    hf.variants.first.header.mode = "fixed"
    hf.variants.first.header.fixed_text = "First Header"
    hf.variants.first.header.alignment = "left"
    hf.variants.first.footer.mode = "fixed"
    hf.variants.first.footer.fixed_text = "First Footer"
    hf.page_number_plan.first.visibility = "hide"
    hf.variants.even.header.mode = "fixed"
    hf.variants.even.header.fixed_text = "Even Header"
    hf.variants.even.header.alignment = "right"
    hf.page_number_plan.even.visibility = "show"
    hf.page_number_plan.even.template = "第 {page} 页 / 共 {pages} 页"
    hf.page_number_plan.even.alignment = "right"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    section = doc.sections[0]
    assert section.different_first_page_header_footer is True
    assert doc.settings.odd_and_even_pages_header_footer is True
    assert section.first_page_header.paragraphs[0].text == "First Header"
    assert _paragraph_alignment(section.first_page_header.paragraphs[0]) == "left"
    assert "First Footer" in section.first_page_footer.paragraphs[0].text
    assert _paragraph_has_field(section.first_page_footer.paragraphs[0], "PAGE") is False
    assert section.even_page_header.paragraphs[0].text == "Even Header"
    assert _paragraph_alignment(section.even_page_header.paragraphs[0]) == "right"
    even_instrs = _field_instr_texts(section.even_page_footer.paragraphs[0])
    assert any(instr.startswith("PAGE") for instr in even_instrs)
    assert any(instr.startswith("NUMPAGES") for instr in even_instrs)
    assert _paragraph_alignment(section.even_page_footer.paragraphs[0]) == "right"


def test_first_even_page_number_strategy_respects_global_master_gate():
    doc = Document()
    config = ResolvedConfig()
    hf = config.header_footer
    hf.behavior.different_first_page = True
    hf.behavior.different_odd_even_pages = True
    hf.page_number_enabled = False
    hf.page_number_plan.first.visibility = "show"
    hf.page_number_plan.even.visibility = "show"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    section = doc.sections[0]
    assert _paragraph_has_field(section.footer.paragraphs[0], "PAGE") is False
    assert _paragraph_has_field(section.first_page_footer.paragraphs[0], "PAGE") is False
    assert _paragraph_has_field(section.even_page_footer.paragraphs[0], "PAGE") is False


def test_first_page_force_show_overrides_hidden_numbering_phase():
    doc = Document()
    config = ResolvedConfig()
    hf = config.header_footer
    hf.behavior.different_first_page = True
    hf.page_number_plan.phases = [
        PageNumberPhaseConfig(
            phase_id="hidden_body",
            selectors=["body"],
            visible=False,
        )
    ]
    hf.page_number_plan.first.visibility = "show"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    section = doc.sections[0]
    assert _paragraph_has_field(section.footer.paragraphs[0], "PAGE") is False
    assert _paragraph_has_field(section.first_page_footer.paragraphs[0], "PAGE") is True


def test_footer_text_and_page_number_keep_independent_alignment_with_tab_stops():
    doc = Document()
    config = ResolvedConfig()
    hf = config.header_footer
    hf.footer.content_mode = "fixed"
    hf.footer_text = "Confidential"
    hf.footer_alignment = "left"
    hf.page_number_alignment = "right"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    para = doc.sections[0].footer.paragraphs[0]
    assert "Confidential" in para.text
    assert _paragraph_has_field(para, "PAGE") is True
    tabs = para._element.find(qn("w:pPr")).find(qn("w:tabs"))
    assert tabs is not None
    tab_specs = [(tab.get(qn("w:val")), int(tab.get(qn("w:pos")))) for tab in tabs]
    assert len(tab_specs) == 1
    assert tab_specs[0][0] == "right"
    assert tab_specs[0][1] > 0


def test_header_footer_page_number_template_can_include_total_pages():
    doc = Document()
    config = ResolvedConfig()
    config.header_footer.page_number_template = "第 {page} 页 / 共 {pages} 页"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    para = doc.sections[0].footer.paragraphs[0]
    instrs = _field_instr_texts(para)
    assert any(instr.startswith("PAGE") for instr in instrs)
    assert any(instr.startswith("NUMPAGES") for instr in instrs)
    assert "第 " in para.text
    assert " / 共 " in para.text
    assert " 页" in para.text


def test_header_footer_border_style_controls_line_width_color_and_spacing():
    doc = Document()
    config = ResolvedConfig()
    border = config.header_footer.header.border_style
    border.line_style = "double"
    border.width_pt = 1.5
    border.spacing_pt = 3
    border.color = "FF0000"

    HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())

    para = doc.sections[0].header.paragraphs[0]
    bottom = para._element.find(qn("w:pPr")).find(qn("w:pBdr")).find(qn("w:bottom"))
    assert bottom.get(qn("w:val")) == "double"
    assert bottom.get(qn("w:sz")) == "12"
    assert bottom.get(qn("w:space")) == "3"
    assert bottom.get(qn("w:color")) == "FF0000"


def _build_sectioned_doc():
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("目录")
    doc.add_paragraph("第一章 绪论")
    doc.add_paragraph("正文内容")
    doc.add_paragraph("参考文献")

    _ensure_section_break_before_paragraph(doc, 1, "nextPage")
    _ensure_section_break_before_paragraph(doc, 2, "nextPage")
    _ensure_section_break_before_paragraph(doc, 4, "nextPage")
    return doc


def _build_doc_tree():
    return DocTree(
        sections=[
            DocSection("cover", 0, 1),
            DocSection("toc", 1, 2),
            DocSection("body", 2, 4),
            DocSection("references", 4, 5),
        ]
    )


def _apply_thesis_page_plan(config):
    config.header_footer.page_number_plan.phases = [
        PageNumberPhaseConfig(
            phase_id="front",
            selectors=["front_matter"],
            visible=True,
            number_format="upperRoman",
            start_mode="restart",
            start_value=1,
        ),
        PageNumberPhaseConfig(
            phase_id="body",
            selectors=["body", "back_matter"],
            visible=True,
            number_format="decimal",
            start_mode="restart",
            start_value=1,
        ),
    ]


def _page_num_type(section):
    elem = section._sectPr.find(qn("w:pgNumType"))
    assert elem is not None
    return elem.get(qn("w:fmt")), elem.get(qn("w:start"))


def _paragraph_alignment(para):
    p_pr = para._element.find(qn("w:pPr"))
    assert p_pr is not None
    jc = p_pr.find(qn("w:jc"))
    assert jc is not None
    return jc.get(qn("w:val"))


def _field_instr_texts(para):
    return [" ".join((instr or "").split()) for _kind, _elem, instr in iter_field_instructions(para._element)]


def test_header_footer_doc_tree_strategy_hides_cover_and_restarts_front_and_body_numbering():
    doc = _build_sectioned_doc()
    config = ResolvedConfig()
    _apply_thesis_page_plan(config)
    context = PipelineContext(doc_tree=_build_doc_tree())

    HeaderFooterModule().apply(doc, config, ChangeTracker(), context)

    cover, toc, body, references = list(doc.sections)

    assert cover.header.paragraphs[0].text == ""
    assert cover.footer.paragraphs[0].text == ""

    assert _paragraph_has_field(toc.footer.paragraphs[0], "PAGE") is True
    assert _page_num_type(toc) == ("upperRoman", "1")

    assert _paragraph_has_field(body.footer.paragraphs[0], "PAGE") is True
    assert _page_num_type(body) == ("decimal", "1")
    body_instrs = _field_instr_texts(body.header.paragraphs[0])
    assert body_instrs == ["STYLEREF 1"]

    assert _paragraph_has_field(references.footer.paragraphs[0], "PAGE") is True
    assert _page_num_type(references) == ("decimal", None)

    ref_instrs = _field_instr_texts(references.header.paragraphs[0])
    assert any('STYLEREF "Heading 1 Unnumbered"' in instr for instr in ref_instrs)
    assert all("\\n" not in instr for instr in ref_instrs)


def test_header_footer_page_number_strategy_supports_lower_roman_front_matter_and_optional_body_restart():
    doc = _build_sectioned_doc()
    config = ResolvedConfig()
    config.header_footer.front_matter_page_number_format = "lowerRoman"
    config.header_footer.restart_body_page_number = False
    context = PipelineContext(doc_tree=_build_doc_tree())

    HeaderFooterModule().apply(doc, config, ChangeTracker(), context)

    _cover, toc, body, references = list(doc.sections)

    assert _page_num_type(toc) == ("lowerRoman", "1")
    assert _page_num_type(body) == ("decimal", None)
    assert _page_num_type(references) == ("decimal", None)
