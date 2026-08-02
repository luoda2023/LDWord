import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from src.config.feature_configs import OutputConfig
from src.config.library import load_scene_from_library, load_template_from_library
from src.config.resolver import resolve_config
from src.config.scene import ContentVisibilityRule, DeliveryPreset, SceneWorkspace
from src.config.template import TemplateConfig
from src.pipeline.result import PipelineResult
from src.services.execution_session import resolution as session_resolution
from src.services.execution_session import validation as session_validation
from src.services.production_runtime import delivery_reporting, execution_runtime
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner


def _source(path: Path) -> Path:
    document = Document()
    document.add_paragraph("Before text")
    document.save(path)
    return path


def test_exam_runtime_names_final_documents_from_the_paper_title(tmp_path):
    source = tmp_path / "7fe860ea293f4771ac6f7304f09b3182.exam.md"
    source.write_text(
        """# 小学六年级数学期中考试试卷
> 科目：数学　年级：六年级　考试时间：90 分钟　满分：100 分
## 一、选择题
1. 1 + 1 = ?（100 分）
   A. 1
   B. 2
## 答案速查
1. B
""",
        encoding="utf-8",
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=load_template_from_library("default", mode_id="exam"),
        scene=load_scene_from_library("exam", mode_id="exam"),
        output_dir=tmp_path / "delivery",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert {Path(path).name for path in payload["output_paths"].values()} == {
        "小学六年级数学期中考试试卷_学生卷.docx",
        "小学六年级数学期中考试试卷_答案卷.docx",
    }
    assert all("7fe860ea" not in path for path in payload["output_paths"].values())


def test_workbench_runner_publishes_delivery_artifact_chain(tmp_path, monkeypatch):
    source = _source(tmp_path / "source.docx")
    output_root = tmp_path / "delivery"
    scene = SceneWorkspace(
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                label="Review",
                output_dir_template="{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(
                    final_docx=True,
                    compare_docx=True,
                    report_json=True,
                    report_markdown=True,
                    material_manifest=True,
                    material_package=True,
                ),
                include_structured_intermediate=True,
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    monkeypatch.setattr(execution_runtime, "create_all_modules", list)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_root,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert Path(payload["output_paths"]["review"]).is_file()
    assert Path(payload["compare_paths"]["review"]).is_file()
    assert {Path(path).suffix for path in payload["report_paths"]} == {
        ".json",
        ".md",
    }
    intermediate_path = Path(payload["intermediate_paths"]["review"])
    intermediate = json.loads(intermediate_path.read_text(encoding="utf-8"))
    assert intermediate["kind"] == "delivery_structured_intermediate"
    assert intermediate["preset"]["preset_id"] == "review"
    assert intermediate["visibility"] == "local_diagnostic"
    manifest_path = Path(payload["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["delivery"]["compare_paths"] == payload["compare_paths"]
    assert manifest["delivery"]["intermediate_paths"] == payload[
        "intermediate_paths"
    ]
    with zipfile.ZipFile(payload["material_package_paths"]["zip"]) as archive:
        assert "manifest/material_manifest.json" in archive.namelist()


def _resolved_artifact_config(
    artifacts: OutputConfig,
    *,
    include_structured_intermediate: bool = False,
):
    scene = SceneWorkspace(
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                label="Review",
                artifacts=artifacts,
                include_structured_intermediate=include_structured_intermediate,
            )
        ],
    )
    return resolve_config(TemplateConfig(), scene, {})


def test_compare_docx_artifact_is_atomically_published(tmp_path):
    source = _source(tmp_path / "source.docx")
    revised = tmp_path / "revised.docx"
    revised_document = Document()
    revised_document.add_paragraph("After text")
    revised_document.save(revised)
    config = _resolved_artifact_config(
        OutputConfig(final_docx=True, compare_docx=True)
    )

    paths = delivery_reporting._write_compare_docx_artifacts(
        input_path=source,
        output_dir=tmp_path / "delivery",
        output_paths={"review": str(revised)},
        fallback_output_path=str(revised),
        config=config,
    )

    compare_path = Path(paths["review"])
    assert compare_path.is_file()
    assert not list(compare_path.parent.glob(".alavette-compare-*"))


def test_delivery_reports_follow_preset_artifact_toggles(tmp_path):
    source = _source(tmp_path / "source.docx")
    config = _resolved_artifact_config(
        OutputConfig(report_json=True, report_markdown=True)
    )

    paths = delivery_reporting.write_delivery_reports(
        PipelineResult(success=True),
        input_path=source,
        output_dir=tmp_path / "delivery",
        output_paths={},
        config=config,
        elapsed=0.25,
        modules_enabled=0,
        modules_total=0,
    )

    assert {Path(path).suffix for path in paths} == {".json", ".md"}
    assert all(Path(path).is_file() for path in paths)


def test_structured_intermediate_contains_delivery_contract(tmp_path):
    source = _source(tmp_path / "source.docx")
    config = _resolved_artifact_config(
        OutputConfig(),
        include_structured_intermediate=True,
    )

    paths = delivery_reporting.write_structured_intermediates(
        PipelineResult(success=False, status="failed", error="broken"),
        input_path=source,
        output_dir=tmp_path / "delivery",
        output_paths={},
        fallback_output_path="",
        config=config,
        elapsed=0.25,
        modules_enabled=0,
        modules_total=0,
    )

    intermediate = json.loads(Path(paths["review"]).read_text(encoding="utf-8"))
    assert intermediate["kind"] == "delivery_structured_intermediate"
    assert intermediate["preset"]["preset_id"] == "review"
    assert intermediate["status"] == "failed"


def test_pipeline_publishes_each_delivery_variant_with_visibility_receipts(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Visible")
    document.add_paragraph("{{#visibility:answer}}")
    document.add_paragraph("Teacher answer")
    document.add_paragraph("{{/visibility:answer}}")
    document.save(source)
    output_root = tmp_path / "delivery"
    scene = SceneWorkspace(
        default_delivery_preset_id="student",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                label="Student",
                artifacts=OutputConfig(final_docx=True),
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove")
                ],
            ),
            DeliveryPreset(
                preset_id="teacher",
                label="Teacher",
                artifacts=OutputConfig(final_docx=True),
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    monkeypatch.setattr(execution_runtime, "create_all_modules", list)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_root,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert set(payload["output_paths"]) == {"student", "teacher"}
    student_text = "\n".join(
        paragraph.text
        for paragraph in Document(payload["output_paths"]["student"]).paragraphs
    )
    teacher_text = "\n".join(
        paragraph.text
        for paragraph in Document(payload["output_paths"]["teacher"]).paragraphs
    )
    assert "Teacher answer" not in student_text
    assert "Teacher answer" in teacher_text
    assert payload["content_visibility_preview"]
    assert len(payload["content_visibility_receipts"]["student"]["removed_ranges"]) == 1


def test_failed_run_still_publishes_reports_and_material_package(
    tmp_path,
    monkeypatch,
):
    source = _source(tmp_path / "source.docx")
    output_root = tmp_path / "delivery"
    scene = SceneWorkspace(
        default_delivery_preset_id="archive",
        delivery_presets=[
            DeliveryPreset(
                preset_id="archive",
                label="Archive",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=True,
                    report_markdown=True,
                    material_manifest=True,
                    material_package=True,
                ),
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False

    class _FailingPipeline:
        def __init__(self, **_kwargs):
            pass

        def execute(self, _path):
            return PipelineResult(
                success=False,
                status="failed",
                error="broken delivery",
                failed_items=[{"preset_id": "archive", "reason": "broken"}],
            )

    monkeypatch.setattr(execution_runtime, "create_all_modules", list)
    monkeypatch.setattr(execution_runtime, "Pipeline", _FailingPipeline)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_root,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["error_text"] == "broken delivery"
    assert all(Path(path).is_file() for path in payload["report_paths"])
    assert Path(payload["material_manifest_paths"]["material"]).is_file()
    assert Path(payload["material_package_paths"]["zip"]).is_file()


def test_auxiliary_report_failure_preserves_primary_output(tmp_path, monkeypatch):
    source = _source(tmp_path / "source.docx")
    output_root = tmp_path / "delivery"
    scene = SceneWorkspace(
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                label="Review",
                artifacts=OutputConfig(final_docx=True, report_json=True),
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    monkeypatch.setattr(execution_runtime, "create_all_modules", list)
    monkeypatch.setattr(
        delivery_reporting,
        "write_delivery_reports",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("disk full")),
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_root,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "partial_success"
    assert Path(payload["output_path"]).is_file()
    assert payload["artifact_failure_count"] == 1
    assert payload["artifact_failures"][0]["kind"] == "reports"


def test_delivery_target_templates_are_captured_as_dependent_resources(monkeypatch):
    scene = SceneWorkspace(
        delivery_presets=[
            DeliveryPreset(preset_id="primary", target_template_id="primary"),
            DeliveryPreset(preset_id="review", target_template_id="review-template"),
        ]
    )
    monkeypatch.setattr(
        session_resolution,
        "get_template_entry",
        lambda *_args, **_kwargs: SimpleNamespace(path="", source_type="builtin"),
    )
    monkeypatch.setattr(
        session_resolution,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )

    refs, issues = session_resolution._delivery_target_template_refs(
        scene,
        mode_id="custom",
        primary_template_id="primary",
    )

    assert issues == []
    assert [(item.kind, item.resource_id, item.status) for item in refs] == [
        ("delivery_target_template", "review-template", "ok")
    ]


def test_v1_runner_fails_closed_for_unimplemented_alternate_target_template(
    tmp_path,
):
    source = _source(tmp_path / "source.docx")
    scene = SceneWorkspace(
        template_id="primary-template",
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                target_template_id="alternate-template",
            )
        ],
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=tmp_path / "delivery",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert payload["error_text"] == (
        "delivery_target_template_not_supported:review:alternate-template"
    )


def test_execution_session_rejects_alternate_delivery_target_before_worker():
    scene = SceneWorkspace(
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                target_template_id="alternate-template",
            )
        ]
    )

    issues = session_validation._validate_delivery_target_templates(
        SimpleNamespace(scene=scene),
        SimpleNamespace(template_id="primary-template"),
    )

    assert issues == [
        "delivery_target_template_not_supported:review:alternate-template"
    ]
