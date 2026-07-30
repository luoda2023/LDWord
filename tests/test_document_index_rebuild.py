from __future__ import annotations

from docx import Document

from src.modules.structure.heading_recognition import rebuild_document_index
from src.pipeline.context import PipelineContext


def test_rebuild_document_index_refreshes_paragraph_positions_after_insertion() -> None:
    document = Document()
    document.add_heading("Chapter", level=1)
    document.add_paragraph("Body")
    document.add_heading("Section", level=2)
    context = PipelineContext()

    first = rebuild_document_index(document, context)
    assert first.heading_map == {0: 1, 2: 2}
    assert context.heading_map == {0: 1, 2: 2}

    document.paragraphs[0].insert_paragraph_before("Preface")
    rebuilt = rebuild_document_index(document, context)

    assert rebuilt.heading_map == {1: 1, 3: 2}
    assert context.doc_tree is rebuilt
    assert context.heading_map == {1: 1, 3: 2}
