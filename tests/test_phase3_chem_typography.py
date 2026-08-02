from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.formula_policy import ChemTypographyOptions
from src.config.resolver import resolve_config
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.modules.special.chem_typography import (
    ChemTypographyModule,
    _apply_chem_typography,
)
from src.modules.structure.heading_recognition import DocTree
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn


def _marks_to_pattern(marks: list[str | None]) -> str:
    return ''.join('^' if m == 'superscript' else '_' if m == 'subscript' else '.' for m in marks)


def _run_vert_align(run) -> str | None:
    rpr = run._element.find(qn('w:rPr'))
    if rpr is None:
        return None
    va = rpr.find(qn('w:vertAlign'))
    if va is None:
        return None
    return va.get(qn('w:val')) or va.get('w:val')


def _paragraph_marks(para) -> str:
    buf: list[str] = []
    for run in para.runs:
        style = _run_vert_align(run)
        ch = '^' if style == 'superscript' else '_' if style == 'subscript' else '.'
        buf.extend(ch for _ in (run.text or ''))
    return ''.join(buf)


def test_build_marks_keeps_roman_oxidation_states_on_baseline():
    from src.shared.engine.chem_marks import build_chem_style_marks

    cfg = ChemTypographyOptions()
    assert _marks_to_pattern(build_chem_style_marks('Fe(IV)', chem_cfg=cfg)) == '......'
    assert _marks_to_pattern(build_chem_style_marks('Cr(VI)', chem_cfg=cfg)) == '......'
    assert _marks_to_pattern(build_chem_style_marks('Cr(III)', chem_cfg=cfg)) == '.......'



def test_build_marks_restores_common_formula_variants():
    from src.shared.engine.chem_marks import build_chem_style_marks

    cfg = ChemTypographyOptions()
    assert _marks_to_pattern(build_chem_style_marks('g-C3N4', chem_cfg=cfg)) == '..._._'
    assert _marks_to_pattern(build_chem_style_marks('\u03b1-Fe3O4', chem_cfg=cfg)) == '...._._'
    assert _marks_to_pattern(build_chem_style_marks('TiO2/In2S3', chem_cfg=cfg)) == '..._..._._'
    assert _marks_to_pattern(build_chem_style_marks('AgInS2/In2S3', chem_cfg=cfg)) == '....._..._._'
    assert _marks_to_pattern(build_chem_style_marks('1 O2', chem_cfg=cfg)) == '^.._'



def test_build_font_mask_keeps_formula_tokens_including_roman_state_forms():
    from src.shared.engine.chem_marks import build_chem_font_mask

    cfg = ChemTypographyOptions()
    assert ''.join('F' if x else '.' for x in build_chem_font_mask('Cr(VI)', chem_cfg=cfg)) == 'FFFFFF'
    assert ''.join('F' if x else '.' for x in build_chem_font_mask('TiO2/In2S3', chem_cfg=cfg)) == 'FFFFFFFFFF'
    assert ''.join('F' if x else '.' for x in build_chem_font_mask('Fe(IV)=O', chem_cfg=cfg)) == 'FFFFFFFF'
    assert ''.join('F' if x else '.' for x in build_chem_font_mask('Au/BiVO4', chem_cfg=cfg)) == 'FFFFFFFF'
    assert ''.join('F' if x else '.' for x in build_chem_font_mask('WO3/N-CDs', chem_cfg=cfg)) == 'FFFFFFFFF'



def test_apply_chem_typography_splits_runs_and_applies_marks_across_run_boundaries():
    doc = Document()
    para = doc.add_paragraph()
    para.add_run('TiO2/')
    para.add_run('In2S3')

    _apply_chem_typography(para, ChemTypographyOptions())

    assert para.text == 'TiO2/In2S3'
    assert _paragraph_marks(para) == '..._..._._'



def test_apply_chem_typography_does_not_superscript_roman_oxidation_state():
    doc = Document()
    para = doc.add_paragraph('Fe(IV) and Cr(VI)')

    _apply_chem_typography(para, ChemTypographyOptions())

    assert para.text == 'Fe(IV) and Cr(VI)'
    assert _paragraph_marks(para) == '.................'


def test_module_treats_all_disabled_scopes_as_no_operation():
    doc = Document()
    body_para = doc.add_paragraph('Body TiO2/In2S3')
    ref_para = doc.add_paragraph('[1] TiO2/In2S3 photocatalyst')

    scene = SceneWorkspace(mode_id="thesis")
    scene.ensure_thesis_formula_rules().chem_typography.enabled = True
    scene.ensure_thesis_formula_rules().chem_typography.scopes = {
        'references': False,
        'body': False,
        'headings': False,
        'abstract_cn': False,
        'abstract_en': False,
        'captions': False,
        'tables': False,
    }
    config = resolve_config(TemplateConfig(), scene)
    tracker = ChangeTracker()
    context = PipelineContext(
        doc_tree=DocTree(
            section_ranges={'references': (1, 2)},
        )
    )

    ChemTypographyModule().apply(doc, config, tracker, context)

    assert _paragraph_marks(body_para) == '.' * len(body_para.text)
    assert _paragraph_marks(ref_para) == '.' * len(ref_para.text)


def test_module_explicit_scopes_limit_body_headings_and_references():
    doc = Document()
    heading_para = doc.add_paragraph('TiO2/In2S3 heading')
    heading_para.style = doc.styles['Heading 1']
    body_para = doc.add_paragraph('Body TiO2/In2S3')
    ref_para = doc.add_paragraph('[1] TiO2/In2S3 photocatalyst')

    scene = SceneWorkspace(mode_id="thesis")
    scene.ensure_thesis_formula_rules().chem_typography.enabled = True
    scene.ensure_thesis_formula_rules().chem_typography.scopes = {
        'references': False,
        'body': True,
        'headings': False,
        'abstract_cn': False,
        'abstract_en': False,
        'captions': False,
        'tables': False,
    }
    config = resolve_config(TemplateConfig(), scene)
    tracker = ChangeTracker()
    context = PipelineContext(
        doc_tree=DocTree(
            heading_map={0: 1},
            section_ranges={'references': (2, 3)},
        )
    )

    ChemTypographyModule().apply(doc, config, tracker, context)

    assert _paragraph_marks(heading_para) == '..................'
    assert _paragraph_marks(body_para) == '........_..._._'
    assert _paragraph_marks(ref_para) == '............................'


def test_module_explicit_tables_scope_formats_table_cell_paragraphs():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell_para = table.cell(0, 0).paragraphs[0]
    cell_para.text = 'TiO2/In2S3'

    scene = SceneWorkspace(mode_id="thesis")
    scene.ensure_thesis_formula_rules().chem_typography.enabled = True
    scene.ensure_thesis_formula_rules().chem_typography.scopes = {
        'references': False,
        'body': False,
        'headings': False,
        'abstract_cn': False,
        'abstract_en': False,
        'captions': False,
        'tables': True,
    }
    config = resolve_config(TemplateConfig(), scene)
    tracker = ChangeTracker()
    context = PipelineContext()

    ChemTypographyModule().apply(doc, config, tracker, context)

    assert _paragraph_marks(cell_para) == '..._..._._'
