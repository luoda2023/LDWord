"""
Phase 3 冒烟测试 — 📊 表格与题注模块 (3 个)
"""

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def test_caption_module():
    """M08: caption 完整管线测试"""
    from src.modules.table.caption import CaptionModule, _scan_captions
    from src.modules.structure.heading_recognition import HeadingRecognitionModule
    from src.config.template import TemplateConfig
    from src.config.scene import SceneWorkspace
    from src.config.resolver import resolve_config
    from src.pipeline.runner import Pipeline

    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容。")
    doc.add_paragraph("图 1 实验装置图")
    doc.add_paragraph("正文内容。")
    doc.add_paragraph("表 1 实验数据")
    doc.add_heading("第二章 方法", level=1)
    doc.add_paragraph("图 2 流程图")
    doc.add_paragraph("表 2 结果统计")

    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)

    try:
        config = resolve_config(TemplateConfig(), SceneWorkspace())
        pipeline = Pipeline(
            modules=[HeadingRecognitionModule(), CaptionModule()],
            config=config,
        )
        result = pipeline.execute(tmp.name)

        assert result.success, f"执行失败: {result.error}"
        assert result.context.caption_counters is not None
        print(f"  ✅ M08 caption 通过 (counters={result.context.caption_counters})")
    finally:
        Path(tmp.name).unlink(missing_ok=True)
        out = result.output_paths.get("final", "")
        if out:
            Path(out).unlink(missing_ok=True)


def test_caption_scan_logic():
    """M08: 题注扫描逻辑测试"""
    from src.modules.table.caption import _scan_captions

    doc = Document()
    doc.add_paragraph("图 1.1 实验装置")
    doc.add_paragraph("正文段落")
    doc.add_paragraph("表 2.3 结果统计表")
    doc.add_paragraph("Figure 5 Overview")
    doc.add_paragraph("Table 3 Results")

    captions = _scan_captions(doc)
    assert len(captions) == 4
    assert captions[0].kind == "figure"
    assert captions[0].prefix == "图"
    assert captions[1].kind == "table"
    assert captions[2].kind == "figure"
    assert captions[3].kind == "table"
    print("  ✅ M08 题注扫描逻辑通过")


def test_caption_without_heading_recognition():
    """M08: 关闭 heading_recognition 时，caption 仍可回退到全局编号。"""
    from src.modules.table.caption import CaptionModule
    from src.config.template import TemplateConfig
    from src.config.scene import SceneWorkspace
    from src.config.resolver import resolve_config
    from src.pipeline.runner import Pipeline

    doc = Document()
    doc.add_paragraph("图 7 实验装置")
    doc.add_paragraph("正文段落")
    doc.add_paragraph("表 9 实验结果")

    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)

    try:
        scene = SceneWorkspace()
        scene.module_switches["heading_recognition"] = False
        config = resolve_config(TemplateConfig(), scene)

        pipeline = Pipeline(modules=[CaptionModule()], config=config)
        result = pipeline.execute(tmp.name)

        assert result.success, f"执行失败: {result.error}"
        assert result.doc.paragraphs[0].text.startswith("图1")
        assert result.doc.paragraphs[2].text.startswith("表1")
        print("  ✅ M08 caption 可在无 heading_map 时回退到全局编号")
    finally:
        Path(tmp.name).unlink(missing_ok=True)
        out = result.output_paths.get("final", "")
        if out:
            Path(out).unlink(missing_ok=True)


def test_table_format_import():
    """M09: table_format 导入"""
    from src.modules.table.table_format import TableFormatModule

    mod = TableFormatModule()
    assert mod.meta.name == "table_format"
    assert mod.meta.category == "table"
    print("  ✅ M09 table_format 导入正常")


def test_table_smart_levels_are_normalized_to_engine_supported_range():
    from src.config.feature_configs import (
        TABLE_SMART_LEVEL_OPTIONS,
        normalize_table_smart_levels,
    )

    assert TABLE_SMART_LEVEL_OPTIONS == (3, 4, 5, 6)
    assert normalize_table_smart_levels(2) == 4
    assert normalize_table_smart_levels(3) == 3
    assert normalize_table_smart_levels("6") == 6
    assert normalize_table_smart_levels("bad") == 4


def test_figure_table_center_import():
    """M10: figure_table_center 导入"""
    from src.modules.table.figure_table_center import FigureTableCenterModule

    mod = FigureTableCenterModule()
    assert mod.meta.name == "figure_table_center"
    assert "caption" in mod.meta.soft_after
    print("  ✅ M10 figure_table_center 导入正常")


def test_three_line_border_helper_uses_outer_outer_inner_semantics():
    """M09: helper 级三线表语义应为上=下粗，中间细。"""
    from src.shared.engine.table_builder import set_table_borders
    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    set_table_borders(
        table,
        mode="three_line",
        width_pt=2.0,
        header_width_pt=1.0,
        bottom_width_pt=0.5,
    )

    tbl_borders = table._element.find(qn("w:tblPr")).find(qn("w:tblBorders"))
    top = tbl_borders.find(qn("w:top"))
    bottom = tbl_borders.find(qn("w:bottom"))
    inside_h = tbl_borders.find(qn("w:insideH"))
    inside_v = tbl_borders.find(qn("w:insideV"))
    header_cell_bottom = (
        table.rows[0].cells[0]._element
        .find(qn("w:tcPr"))
        .find(qn("w:tcBorders"))
        .find(qn("w:bottom"))
    )

    assert top is not None and top.get(qn("w:sz")) == "8"
    assert bottom is not None and bottom.get(qn("w:sz")) == "8"
    assert inside_h is not None and inside_h.get(qn("w:val")) == "none"
    assert inside_v is not None and inside_v.get(qn("w:val")) == "none"
    assert header_cell_bottom is not None and header_cell_bottom.get(qn("w:sz")) == "4"
    print("  ✅ M09 helper 级三线表语义正确")


def test_no_border_helper_clears_cell_level_overrides():
    from src.shared.engine.table_builder import set_table_borders
    from src.shared.engine.ooxml_ops import find_or_create, qn

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    tc_pr = find_or_create(table.cell(0, 0)._element, "w:tcPr")
    tc_borders = find_or_create(tc_pr, "w:tcBorders")
    right = find_or_create(tc_borders, "w:right")
    right.set(qn("w:val"), "single")
    right.set(qn("w:sz"), "8")

    set_table_borders(table, mode="none")

    tc_borders = table.cell(0, 0)._element.find(qn("w:tcPr")).find(qn("w:tcBorders"))
    assert tc_borders.find(qn("w:right")).get(qn("w:val")) == "none"
    assert tc_borders.find(qn("w:right")).get(qn("w:sz")) == "0"


def test_three_line_helper_clears_cell_level_overrides():
    from src.shared.engine.table_builder import set_table_borders
    from src.shared.engine.ooxml_ops import find_or_create, qn

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    tc_pr = find_or_create(table.cell(0, 0)._element, "w:tcPr")
    tc_borders = find_or_create(tc_pr, "w:tcBorders")
    right = find_or_create(tc_borders, "w:right")
    right.set(qn("w:val"), "single")
    right.set(qn("w:sz"), "8")
    right.set(qn("w:color"), "FF0000")

    set_table_borders(table, mode="three_line")

    tc_borders_after = table.cell(0, 0)._element.find(qn("w:tcPr")).find(qn("w:tcBorders"))
    right_after = tc_borders_after.find(qn("w:right")) if tc_borders_after is not None else None
    bottom_after = tc_borders_after.find(qn("w:bottom")) if tc_borders_after is not None else None

    assert right_after is None
    assert bottom_after is not None and bottom_after.get(qn("w:val")) == "single"


def test_table_format_uses_outer_outer_inner_config_semantics():
    """M09: table_format 模块应复用 0.2.1 语义：outer 来自 three_line_header，inner 来自 three_line_bottom。"""
    from src.modules.table.table_format import _apply_three_line_borders
    from src.shared.engine.ooxml_ops import qn
    from src.config.template import TableConfig

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    cfg = TableConfig(
        border_mode="three_line",
        border_width_pt=2.0,
        three_line_header_width_pt=1.0,
        three_line_bottom_width_pt=0.5,
    )

    _apply_three_line_borders(table._element, cfg)

    tbl_borders = table._element.find(qn("w:tblPr")).find(qn("w:tblBorders"))
    top = tbl_borders.find(qn("w:top"))
    bottom = tbl_borders.find(qn("w:bottom"))
    inside_h = tbl_borders.find(qn("w:insideH"))
    inside_v = tbl_borders.find(qn("w:insideV"))
    header_cell_bottom = (
        table.rows[0].cells[0]._element
        .find(qn("w:tcPr"))
        .find(qn("w:tcBorders"))
        .find(qn("w:bottom"))
    )

    assert top is not None and top.get(qn("w:sz")) == "8"
    assert bottom is not None and bottom.get(qn("w:sz")) == "8"
    assert inside_h is not None and inside_h.get(qn("w:val")) == "none"
    assert inside_v is not None and inside_v.get(qn("w:val")) == "none"
    assert header_cell_bottom is not None and header_cell_bottom.get(qn("w:sz")) == "4"
    print("  ✅ M09 table_format 三线表语义正确")


def test_table_format_three_line_clears_existing_cell_border_overrides():
    from src.config.resolved import ResolvedConfig
    from src.modules.table.table_format import TableFormatModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.ooxml_ops import find_or_create, qn

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "H"
    table.cell(1, 1).text = "Data"
    tc_pr = find_or_create(table.cell(1, 1)._element, "w:tcPr")
    tc_borders = find_or_create(tc_pr, "w:tcBorders")
    left = find_or_create(tc_borders, "w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:color"), "FF0000")

    config = ResolvedConfig()
    config.table.border_mode = "three_line"

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    tc_borders_after = table.cell(1, 1)._element.find(qn("w:tcPr")).find(qn("w:tcBorders"))
    left_after = tc_borders_after.find(qn("w:left")) if tc_borders_after is not None else None
    tbl_left = table._element.find(qn("w:tblPr")).find(qn("w:tblBorders")).find(qn("w:left"))

    assert left_after is None
    assert tbl_left is not None and tbl_left.get(qn("w:val")) == "none"


def test_table_format_applies_full_grid_width_right_alignment_and_double_spacing():
    from src.config.resolved import ResolvedConfig
    from src.modules.table.table_format import TableFormatModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Data"

    config = ResolvedConfig()
    config.table.border_mode = "full_grid"
    config.table.border_width_pt = 1.5
    config.table.cell_alignment = "right"
    config.table.line_spacing_mode = "double"

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    tbl_borders = table._element.find(qn("w:tblPr")).find(qn("w:tblBorders"))
    top = tbl_borders.find(qn("w:top"))
    para = table.cell(0, 0).paragraphs[0]

    assert top is not None and top.get(qn("w:sz")) == "12"
    assert para.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert para.paragraph_format.line_spacing == 2.0


def test_table_format_applies_table_block_alignment_independently_from_cell_alignment():
    from src.config.resolved import ResolvedConfig
    from src.modules.table.table_format import TableFormatModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Data"

    config = ResolvedConfig()
    config.table.layout_mode = "compact"
    config.table.border_mode = "keep"
    config.table.table_alignment = "right"
    config.table.cell_alignment = "center"

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    tbl_jc = table._element.find(qn("w:tblPr")).find(qn("w:jc"))
    para = table.cell(0, 0).paragraphs[0]

    assert tbl_jc is not None and tbl_jc.get(qn("w:val")) == "right"
    assert para.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_table_format_can_leave_existing_table_block_alignment_unchanged():
    from src.config.resolved import ResolvedConfig
    from src.modules.table.table_format import TableFormatModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.ooxml_ops import find_or_create, qn

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    tbl_pr = find_or_create(table._element, "w:tblPr")
    jc = find_or_create(tbl_pr, "w:jc")
    jc.set(qn("w:val"), "right")

    config = ResolvedConfig()
    config.table.layout_mode = "keep"
    config.table.border_mode = "keep"
    config.table.table_alignment = None

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    tbl_jc = table._element.find(qn("w:tblPr")).find(qn("w:jc"))
    assert tbl_jc is not None and tbl_jc.get(qn("w:val")) == "right"


def test_table_format_can_bold_first_row_only():
    from src.config.resolved import ResolvedConfig
    from src.modules.table.table_format import TableFormatModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker

    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "表头"
    table.cell(1, 0).text = "正文"

    config = ResolvedConfig()
    config.table.first_row_bold = True

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert table.cell(0, 0).paragraphs[0].runs[0].font.bold is True
    assert table.cell(1, 0).paragraphs[0].runs[0].font.bold in (None, False)


def test_table_format_can_disable_existing_repeat_header():
    from src.config.resolved import ResolvedConfig
    from src.modules.table.table_format import TableFormatModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.ooxml_ops import find_or_create, qn

    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Header"
    table.cell(1, 0).text = "Body"
    tr_pr = find_or_create(table.rows[0]._element, "w:trPr")
    find_or_create(tr_pr, "w:tblHeader")

    config = ResolvedConfig()
    config.table.layout_mode = "keep"
    config.table.border_mode = "keep"
    config.table.table_alignment = None
    config.table.cell_alignment = None
    config.table.repeat_header = False

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert table.rows[0]._element.find(qn("w:trPr")).find(qn("w:tblHeader")) is None

    config.table.repeat_header = True
    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    header = table.rows[0]._element.find(qn("w:trPr")).find(qn("w:tblHeader"))
    assert header is not None and header.get(qn("w:val")) == "1"


def test_table_format_applies_color_table_header_and_zebra_semantics():
    from src.config.resolved import ResolvedConfig
    from src.modules.table.table_format import TableFormatModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "表头"
    table.cell(1, 0).text = "正文 A"
    table.cell(2, 0).text = "正文 B"

    config = ResolvedConfig()
    config.table.border_mode = "color_table"
    config.table.color_table_accent = "blue"
    config.table.color_table_variant = "header_grid_zebra"

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    header_tc_pr = table.cell(0, 0)._element.find(qn("w:tcPr"))
    body_tc_pr = table.cell(1, 0)._element.find(qn("w:tcPr"))
    tbl_borders = table._element.find(qn("w:tblPr")).find(qn("w:tblBorders"))

    assert header_tc_pr.find(qn("w:shd")).get(qn("w:fill")) == "4472C4"
    assert body_tc_pr.find(qn("w:shd")).get(qn("w:fill")) == "EAF1FB"
    assert tbl_borders.find(qn("w:insideV")).get(qn("w:val")) == "single"
    assert tbl_borders.find(qn("w:insideH")).get(qn("w:val")) == "single"

    run_pr = table.cell(0, 0).paragraphs[0].runs[0]._element.find(qn("w:rPr"))
    assert run_pr.find(qn("w:b")).get(qn("w:val")) == "1"
    assert run_pr.find(qn("w:color")).get(qn("w:val")) == "FFFFFF"


def test_table_format_none_strongly_clears_table_and_cell_borders():
    from src.config.resolved import ResolvedConfig
    from src.modules.table.table_format import TableFormatModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.ooxml_ops import find_or_create, qn

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Layout"
    tc_pr = find_or_create(table.cell(0, 0)._element, "w:tcPr")
    tc_borders = find_or_create(tc_pr, "w:tcBorders")
    left = find_or_create(tc_borders, "w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "8")

    config = ResolvedConfig()
    config.table.border_mode = "none"

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    tbl_borders = table._element.find(qn("w:tblPr")).find(qn("w:tblBorders"))
    tc_borders = table.cell(0, 0)._element.find(qn("w:tcPr")).find(qn("w:tcBorders"))

    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        assert tbl_borders.find(qn(f"w:{side}")).get(qn("w:val")) == "none"
        assert tc_borders.find(qn(f"w:{side}")).get(qn("w:val")) == "none"


def test_all_table_meta():
    """验证 3 个模块 meta 完整性"""
    from src.modules.table.caption import CaptionModule
    from src.modules.table.table_format import TableFormatModule
    from src.modules.table.figure_table_center import FigureTableCenterModule

    modules = [CaptionModule(), TableFormatModule(), FigureTableCenterModule()]
    for mod in modules:
        assert mod.meta.name
        assert mod.meta.description
        assert mod.meta.category == "table"

    names = [m.meta.name for m in modules]
    assert len(names) == len(set(names))
    print(f"  ✅ 3 个模块 meta 验证通过: {names}")


if __name__ == "__main__":
    print("Phase 3 冒烟测试 — 📊 表格与题注")
    print("=" * 50)

    test_caption_module()
    test_caption_scan_logic()
    test_table_format_import()
    test_figure_table_center_import()
    test_all_table_meta()

    print("=" * 50)
    print("✅ 表格与题注 3 个模块全部通过！")
