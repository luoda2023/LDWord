import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from src.config.resolver import resolve_config
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.execution_diagnostics import build_execution_diagnostics
from src.modules.base import BaseModule, ModuleMeta
from src.modules.basic.section_format import SectionFormatModule
from src.pipeline.runner import Pipeline
from src.pipeline.tracker import ChangeTracker
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.object_preflight import (
    inspect_docx_package,
    object_preflight_targets_for_touchpoints,
)
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner


def test_object_preflight_detects_fragile_docx_parts(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "fragile.docx",
        {
            "word/embeddings/oleObject1.bin": b"ole",
            "word/embeddings/workbook.xlsx": b"not-a-real-xlsx",
            "word/vbaProject.bin": b"macro",
            "word/comments.xml": (
                b'<w:comments xmlns:w="http://schemas.openxmlformats.org/'
                b'wordprocessingml/2006/main" />'
            ),
            "word/header1.xml": (
                b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b"<w:body><w:p><w:r><w:instrText>PAGE</w:instrText></w:r></w:p></w:body>"
                b"</w:document>"
            ),
            "word/footer1.xml": (
                b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b"<w:body><w:p><w:r><w:rPr><w:vanish /></w:rPr><w:t>Hidden</w:t></w:r></w:p></w:body>"
                b"</w:document>"
            ),
            "word/_rels/header1.xml.rels": (
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="embeddings/oleObject1.bin"/>'
                b"</Relationships>"
            ),
        },
    )

    result = inspect_docx_package(source)
    kinds = {finding.kind for finding in result.findings}

    assert "ole_objects" in kinds
    assert "embedded_workbooks" in kinds
    assert "macros" in kinds
    assert "comments" in kinds
    assert "fields" in kinds
    assert "hidden_text" in kinds


def test_object_preflight_detects_content_controls(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "content_controls.docx",
        {
            "word/custom-content-controls.xml": (
                b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b"<w:body><w:sdt><w:sdtContent><w:p><w:r><w:t>Name</w:t></w:r></w:p></w:sdtContent></w:sdt></w:body>"
                b"</w:document>"
            ),
        },
    )

    result = inspect_docx_package(source)
    kinds = {finding.kind for finding in result.findings}

    assert "content_controls" in kinds


def test_object_preflight_does_not_treat_revision_style_metadata_as_visio(
    tmp_path,
):
    source = tmp_path / "plain-python-docx.docx"
    document = Document()
    document.add_paragraph("Plain content")
    document.save(source)

    result = inspect_docx_package(source)
    kinds = {finding.kind for finding in result.findings}

    assert "visio_drawings" not in kinds


def test_object_preflight_detects_explicit_visio_markup(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "visio.docx",
        {
            "word/visio-marker.xml": (
                b'<v:document xmlns:v="http://schemas.microsoft.com/'
                b'visio/2012/main" />'
            ),
        },
    )

    result = inspect_docx_package(source)
    kinds = {finding.kind for finding in result.findings}

    assert "visio_drawings" in kinds


def test_object_preflight_respects_policy_scan_targets(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "filtered.docx",
        {
            "word/embeddings/oleObject1.bin": b"ole",
            "word/vbaProject.bin": b"macro",
            "word/comments.xml": (
                b'<w:comments xmlns:w="http://schemas.openxmlformats.org/'
                b'wordprocessingml/2006/main" />'
            ),
            "word/header1.xml": (
                b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b"<w:body><w:p><w:r><w:instrText>PAGE</w:instrText></w:r></w:p></w:body>"
                b"</w:document>"
            ),
        },
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.scan_targets = ["comments", "fields"]

    result = inspect_docx_package(source, scene.compliance_profile.object_preflight)
    kinds = {finding.kind for finding in result.findings}

    assert kinds == {"comments", "fields"}


def test_ooxml_touchpoints_map_to_object_preflight_targets():
    targets = object_preflight_targets_for_touchpoints(
        [
            "field",
            "comment",
            "revision",
            "hidden_text",
            "relationship",
            "content_control",
            "paragraph",
        ]
    )

    assert targets == (
        "fields",
        "comments",
        "tracked_changes",
        "hidden_text",
        "ole_objects",
        "embedded_workbooks",
        "embedded_packages",
        "visio_drawings",
        "macros",
        "content_controls",
    )


def test_pipeline_does_not_block_filtered_object_preflight_target(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "filtered-block.docx",
        {"word/embeddings/oleObject1.bin": b"ole"},
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.preservation_mode = "strict"
    scene.compliance_profile.object_preflight.block_on = ["ole_objects"]
    scene.compliance_profile.object_preflight.scan_targets = ["fields"]

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    diagnostics = build_execution_diagnostics(result)

    assert result.success is True
    assert result.context is not None
    assert result.context.object_preflight is not None
    assert result.context.object_preflight.findings == []
    assert not any(item["target"] == "ole_objects" for item in diagnostics["items"])
    assert (tmp_path / "out" / "filtered-block_new.docx").exists() is True


def test_object_preflight_report_includes_policy_and_planning_evidence(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "report-evidence.docx",
        {
            "word/comments.xml": (
                b'<w:comments xmlns:w="http://schemas.openxmlformats.org/'
                b'wordprocessingml/2006/main" />'
            ),
            "word/header1.xml": (
                b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b"<w:body><w:p><w:r><w:instrText>PAGE</w:instrText></w:r></w:p></w:body>"
                b"</w:document>"
            ),
        },
    )
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.compliance_profile.object_preflight.scan_targets = ["comments", "fields"]

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    json_report = tmp_path / "changes.json"
    markdown_report = tmp_path / "changes.md"
    output_path = Path(result.output_paths["final"])
    write_json_report(
        result,
        input_path=source,
        output_path=output_path,
        report_path=json_report,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=markdown_report,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )

    payload = json.loads(json_report.read_text(encoding="utf-8"))
    evidence = payload["object_preflight"]

    assert result.success is True
    assert evidence["planning_family_id"] == "contract_delivery"
    assert evidence["material_schema_id"] == "contract_parties_v1"
    assert "revision" in evidence["planning_ooxml_touchpoints"]
    assert "tracked_changes" in evidence["recommended_scan_targets"]
    assert evidence["scan_targets"] == ["comments", "fields"]
    assert {item["kind"] for item in evidence["findings"]} == {"comments", "fields"}
    assert evidence["findings_count"] == 2

    markdown = markdown_report.read_text(encoding="utf-8")
    assert "## 对象预检证据" in markdown
    assert "Planning family: contract_delivery" in markdown
    assert "Scan targets: comments, fields" in markdown


def test_workbench_runner_payload_includes_object_preflight_evidence(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "workbench-preflight.docx",
        {
            "word/comments.xml": (
                b'<w:comments xmlns:w="http://schemas.openxmlformats.org/'
                b'wordprocessingml/2006/main" />'
            ),
            "word/header1.xml": (
                b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b"<w:body><w:p><w:r><w:instrText>PAGE</w:instrText></w:r></w:p></w:body>"
                b"</w:document>"
            ),
        },
    )
    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.compliance_profile.object_preflight.scan_targets = ["comments", "fields"]

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    evidence = payload["object_preflight"]

    assert payload["status"] == "success"
    assert evidence["planning_family_id"] == "contract_delivery"
    assert evidence["material_schema_id"] == "contract_parties_v1"
    assert evidence["scan_targets"] == ["comments", "fields"]
    assert evidence["findings_count"] == 2
    assert {finding["kind"] for finding in evidence["findings"]} == {
        "comments",
        "fields",
    }


def test_pipeline_blocks_strict_object_preflight_findings(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "blocked.docx",
        {"word/embeddings/oleObject1.bin": b"ole"},
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.preservation_mode = "strict"
    scene.compliance_profile.object_preflight.block_on = ["ole_objects"]

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    diagnostics = build_execution_diagnostics(result)

    assert result.success is False
    assert result.status == "failed"
    assert result.context is not None
    assert result.context.object_preflight is not None
    assert result.context.object_preflight.blocking_findings
    assert diagnostics["count"] >= 1
    assert any(item["target"] == "ole_objects" for item in diagnostics["items"])
    assert (tmp_path / "out" / "blocked_new.docx").exists() is False


def test_pipeline_warns_for_nonblocking_object_preflight_findings(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "warn.docx",
        {"word/embeddings/oleObject1.bin": b"ole"},
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.preservation_mode = "warn"
    scene.compliance_profile.object_preflight.block_on = ["ole_objects"]

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    diagnostics = build_execution_diagnostics(result)

    assert result.success is True
    assert result.status == "success"
    assert result.context is not None
    assert result.context.object_preflight is not None
    assert result.context.object_preflight.findings
    assert result.context.object_preflight.blocking_findings == []
    assert diagnostics["count"] >= 1
    assert any(item["target"] == "ole_objects" for item in diagnostics["items"])
    assert (tmp_path / "out" / "warn_new.docx").exists() is True


def test_pipeline_skips_configured_high_risk_modules_after_preflight(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "skip.docx",
        {"word/embeddings/oleObject1.bin": b"ole"},
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.preservation_mode = "warn"
    scene.compliance_profile.object_preflight.block_on = []
    scene.compliance_profile.object_preflight.skip_high_risk_modules = True
    scene.compliance_profile.object_preflight.skip_modules_by_finding = {
        "ole_objects": ["risky_rewriter"],
    }

    result = Pipeline(
        modules=[_RiskyRewriteModule()],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    records = result.tracker.get_all() if result.tracker else []
    diagnostics = build_execution_diagnostics(result)
    json_report = tmp_path / "skip-changes.json"
    markdown_report = tmp_path / "skip-changes.md"
    write_json_report(
        result,
        input_path=source,
        output_path=Path(result.output_paths["final"]),
        report_path=json_report,
        elapsed=0.1,
        modules_enabled=1,
        modules_total=1,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=markdown_report,
        elapsed=0.1,
        modules_enabled=1,
        modules_total=1,
    )
    evidence = json.loads(json_report.read_text(encoding="utf-8"))["object_preflight"]

    assert result.success is True
    assert any(record.rule_name == "risky_rewriter" and record.change_type == "skip" for record in records)
    assert not any(record.rule_name == "risky_rewriter" and record.change_type == "format" for record in records)
    assert any(item["rule_name"] == "risky_rewriter" for item in diagnostics["items"])
    assert evidence["module_skips"] == [
        {
            "module_name": "risky_rewriter",
            "finding_kinds": ["ole_objects"],
            "reason": "Skipped because object preflight found ole_objects.",
        }
    ]
    markdown = markdown_report.read_text(encoding="utf-8")
    assert "Skipped modules:" in markdown
    assert "risky_rewriter: ole_objects" in markdown


def test_pipeline_runs_high_risk_modules_when_preflight_skips_are_disabled(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "noskip.docx",
        {"word/embeddings/oleObject1.bin": b"ole"},
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.preservation_mode = "warn"
    scene.compliance_profile.object_preflight.block_on = []
    scene.compliance_profile.object_preflight.skip_high_risk_modules = False
    scene.compliance_profile.object_preflight.skip_modules_by_finding = {
        "ole_objects": ["risky_rewriter"],
    }

    result = Pipeline(
        modules=[_RiskyRewriteModule()],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    records = result.tracker.get_all() if result.tracker else []

    assert result.success is True
    assert any(record.rule_name == "risky_rewriter" and record.change_type == "format" for record in records)
    assert not any(record.rule_name == "risky_rewriter" and record.change_type == "skip" for record in records)


def test_ole_finding_does_not_skip_section_inventory_and_safe_planner(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "section-safe-ole.docx",
        {"word/embeddings/oleObject1.bin": b"ole"},
    )
    scene = SceneWorkspace()
    scene.compliance_profile.object_preflight.preservation_mode = "warn"
    scene.compliance_profile.object_preflight.block_on = []
    scene.compliance_profile.object_preflight.skip_modules_by_finding = {
        "ole_objects": ["section_format"],
    }
    template = TemplateConfig()
    template.section.boundary_mode = "preserve_source"

    result = Pipeline(
        modules=[SectionFormatModule()],
        config=resolve_config(template, scene),
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    records = result.tracker.get_all() if result.tracker else []
    assert result.success is True
    assert result.context.section_inventory is not None
    assert result.context.section_execution_receipt is not None
    assert result.context.object_preflight_module_skips == []
    assert any(
        record.rule_name == "object_preflight"
        and record.change_type == "granular_risk_clearance"
        for record in records
    )


def _docx_with_parts(
    tmp_path: Path,
    filename: str,
    parts: dict[str, bytes],
) -> Path:
    source = tmp_path / filename
    Document().save(source)
    with ZipFile(source, "a") as package:
        for name, data in parts.items():
            package.writestr(name, data)
    return source


class _RiskyRewriteModule(BaseModule):
    meta = ModuleMeta(
        name="risky_rewriter",
        description="Risky package rewriter",
        category="test",
        execution_phase="format",
        scope_behavior="document_level",
    )

    def apply(self, doc, config, tracker: ChangeTracker, context) -> None:
        tracker.record(
            rule_name=self.meta.name,
            target="document",
            section="global",
            change_type="format",
            after="module ran",
        )
