from __future__ import annotations

from docx import Document

from src.config.attachment_materials import AttachmentBinding
from src.config.execution_target import ExecutionTarget, resolve_execution_target
from src.config.library import ensure_config_library, load_scene_from_library
from src.config.material_context import MaterialExecutionContext
from src.config.scene import SceneWorkspace
from src.qt_api import QAbstractItemView, QApplication
from src.ui.bridge import PanelBridge
from src.ui.panels.scene_content_detail import _ContentDetail
from src.ui.panels.scene_token_audit_projection import (
    build_scene_token_audit_projection,
)


def test_token_audit_requires_a_scannable_template_source():
    projection = build_scene_token_audit_projection(
        SceneWorkspace(scene_id="custom", mode_id="custom"),
        ExecutionTarget(mode_id="custom"),
    )

    assert projection.scan_status == "missing_source"
    assert projection.rows == ()
    assert projection.conclusion == "尚未校验"


def test_token_audit_marks_structured_assembly_as_not_applicable():
    projection = build_scene_token_audit_projection(
        SceneWorkspace(scene_id="exam", mode_id="exam"),
        ExecutionTarget(
            mode_id="exam",
            placeholder_source_kind="structured",
            placeholder_source_label="试卷结构化装配",
        ),
    )

    assert projection.scan_status == "not_applicable"
    assert projection.conclusion == "当前装配方式不使用模板 Token"


def test_token_audit_compares_template_and_package_existence_only(tmp_path):
    template_path = tmp_path / "token-audit.docx"
    document = Document()
    document.add_paragraph("标题：{{@text:title}}")
    document.add_paragraph("缺失字段：{{@text:missing}}")
    document.add_paragraph("错误类型：{{@img:title}}")
    document.add_paragraph("旧语法：{{title}}")
    document.add_paragraph("附件一：{{@attach:supporting}}")
    document.add_paragraph("附件二：{{@attach:supporting}}")
    document.add_paragraph("{{#visibility:answer}}")
    document.add_paragraph("答案内容")
    document.add_paragraph("{{/visibility:answer}}")
    document.save(template_path)

    context = MaterialExecutionContext(
        entity_data={"title": ""},
        attachment_bindings={
            "supporting": AttachmentBinding(role="supporting"),
            "archive": AttachmentBinding(role="archive"),
        },
    )
    projection = build_scene_token_audit_projection(
        SceneWorkspace(scene_id="custom", mode_id="custom"),
        ExecutionTarget(
            mode_id="custom",
            placeholder_source_kind="document",
            placeholder_source_path=str(template_path),
            placeholder_source_label="待校对模板",
        ),
        context,
    )
    rows = {row.token: row for row in projection.rows}

    assert projection.scan_status == "ready"
    assert rows["{{@text:title}}"].status == "ok"
    assert rows["{{@text:title}}"].package_present is True
    assert rows["{{@text:title}}"].content_text == "未填写"
    assert rows["{{@text:missing}}"].issue_code == "mapping_missing"
    assert rows["{{@text:missing}}"].package_present is False
    assert rows["{{@img:title}}"].issue_code == "type_conflict"
    assert rows["{{@img:title}}"].package_present is False
    assert rows["{{title}}"].issue_code == "invalid_token"
    assert rows["{{@attach:supporting}}"].issue_code == ""
    assert rows["{{@attach:supporting}}"].status == "ok"
    assert rows["{{@attach:supporting}}"].occurrence_count == 2
    assert "{{@attach:archive}}" not in rows
    assert "{{#visibility:answer}}" not in rows
    assert "{{/visibility:answer}}" not in rows
    assert projection.error_count == 3
    assert projection.warning_count == 0


def test_builtin_official_tokens_resolve_through_assembly_field_bindings():
    ensure_config_library()
    scene = load_scene_from_library("official", mode_id="official")
    target = resolve_execution_target(
        mode_id="official",
        scene=scene,
        official_document_type_id="notice",
    )
    projection = build_scene_token_audit_projection(
        scene,
        target,
        MaterialExecutionContext(
            mode_id="official",
            material_schema_ids=("official_document_v1",),
            entity_data={"title": "测试标题", "body": "测试正文"},
        ),
    )

    assert projection.scan_status == "ready"
    assert projection.token_count == 13
    assert projection.error_count == 0
    assert projection.warning_count == 0
    title = next(row for row in projection.rows if row.identifier == "official_title")
    assert title.token == "{{@text:official_title}}"
    assert title.package_present is True
    assert title.content_text == "测试标题"


def test_content_detail_uses_a_read_only_existence_table(tmp_path):
    app = QApplication.instance() or QApplication([])
    template_path = tmp_path / "detail-token-audit.docx"
    document = Document()
    document.add_paragraph("{{@text:title}}")
    document.add_paragraph("{{@text:missing}}")
    document.save(template_path)

    scene = SceneWorkspace(scene_id="custom", mode_id="custom")
    bridge = PanelBridge()
    bridge.set_current_scene(scene, emit_signal=False)
    bridge.set_current_document_path(str(template_path), emit_signal=False)
    bridge.set_current_material_context(
        MaterialExecutionContext(entity_data={"title": "测试标题"}),
        emit_signal=False,
    )
    detail = _ContentDetail(bridge)
    try:
        detail.set_scene(scene)
        app.processEvents()

        assert detail._title == "模板 Token 校验"
        assert detail._desc_label.isHidden() is True
        assert detail._detail_summary.isHidden() is True
        assert detail._advanced_section.is_expanded is False
        assert not hasattr(detail, "_token_document_preview")
        assert detail._token_audit_table.rowCount() == 1
        assert detail._token_view_control.segment_text(0) == "不一致 1"
        assert detail._token_view_control.segment_text(1) == "全部 Token 2"
        assert [
            detail._token_audit_table.horizontalHeaderItem(index).text()
            for index in range(detail._token_audit_table.columnCount())
        ] == ["Token", "模板", "资料包", "校验结果", "填写内容"]
        assert not hasattr(detail, "_token_source_meta")
        assert (
            detail._token_audit_table.selectionMode()
            == QAbstractItemView.NoSelection
        )
        assert detail._token_audit_table.item(0, 1).text() == "1 处"
        assert detail._token_audit_table.item(0, 2).text() == "缺少"
        assert detail._token_audit_table.item(0, 3).text() == "● 资料包缺少"
        assert detail._token_audit_table.item(0, 4).text() == "—"

        detail._token_view_control.set_current_index(1)
        app.processEvents()
        assert detail._token_audit_table.rowCount() == 2
        title_index = next(
            index
            for index in range(detail._token_audit_table.rowCount())
            if detail._token_audit_table.item(index, 0).text() == "{{@text:title}}"
        )
        assert detail._token_audit_table.item(title_index, 4).text() == "测试标题"
    finally:
        detail.close()
        app.processEvents()


def test_token_audit_refreshes_filled_content_without_changing_existence(tmp_path):
    app = QApplication.instance() or QApplication([])
    template_path = tmp_path / "return-token-audit.docx"
    document = Document()
    document.add_paragraph("{{@text:missing}}")
    document.save(template_path)

    scene = SceneWorkspace(scene_id="custom", mode_id="custom")
    bridge = PanelBridge()
    bridge.set_current_scene(scene, emit_signal=False)
    bridge.set_current_document_path(str(template_path), emit_signal=False)
    bridge.set_current_material_context(
        MaterialExecutionContext(entity_data={"missing": ""}),
        emit_signal=False,
    )
    detail = _ContentDetail(bridge)
    try:
        detail.set_scene(scene)
        app.processEvents()
        assert detail._token_audit_projection.error_count == 0
        assert detail._token_audit_table.rowCount() == 1
        assert detail._token_audit_table.item(0, 4).text() == "未填写"

        bridge.set_current_material_context(
            MaterialExecutionContext(entity_data={"missing": "已填写"}),
            emit_signal=False,
        )
        detail._refresh_token_audit()
        app.processEvents()

        assert detail._token_audit_projection.error_count == 0
        assert detail._token_audit_table.rowCount() == 1
        assert detail._token_audit_table.item(0, 4).text() == "已填写"
    finally:
        detail.close()
        app.processEvents()
