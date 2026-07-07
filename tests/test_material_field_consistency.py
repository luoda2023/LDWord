import json
from pathlib import Path

from docx import Document

from src.config.resolver import resolve_config
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.modules.fill.entity_fill import EntityFillModule
from src.pipeline.runner import Pipeline
from src.report_writer import write_json_report, write_markdown_report


def test_contract_material_field_consistency_enters_reports_after_fill(tmp_path):
    source = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("甲方：旧公司")
    doc.add_paragraph("甲方签约主体：{{party_a}}")
    doc.add_paragraph("乙方：{{party_b}}")
    doc.add_paragraph("合同金额：{{contract_amount}}")
    doc.add_paragraph("签署日期：{{signing_date}}")
    doc.save(source)

    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    config = resolve_config(
        TemplateConfig(),
        scene,
        entity_data={
            "party_a": "甲方公司",
            "party_b": "乙方公司",
            "contract_amount": "100000",
            "signing_date": "2026-06-16",
        },
    )

    result = Pipeline(
        modules=[EntityFillModule()],
        config=config,
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    json_report = tmp_path / "contract_changes.json"
    markdown_report = tmp_path / "contract_changes.md"
    write_json_report(
        result,
        input_path=source,
        output_path=Path(result.output_paths["final"]),
        report_path=json_report,
        elapsed=0.1,
        modules_enabled=1,
        modules_total=1,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=markdown_report,
        elapsed=0.1,
        modules_enabled=1,
        modules_total=1,
    )

    payload = json.loads(json_report.read_text(encoding="utf-8"))
    consistency = payload["material_field_consistency"]
    party_a = next(
        item for item in consistency["items"] if item["field_key"] == "party_a"
    )
    issue_kinds = {issue["kind"] for issue in consistency["issues"]}
    markdown = markdown_report.read_text(encoding="utf-8")

    assert result.success is True
    assert consistency["schema_id"] == "contract_parties_v1"
    assert consistency["family_id"] == "contract_delivery"
    assert consistency["status"] == "warning"
    assert party_a["occurrence_count"] == 1
    assert party_a["placeholders_remaining"] == []
    assert "label_value_conflict" in issue_kinds
    assert "Material Field Consistency" in markdown
    assert "party_a: warning" in markdown
    assert "label_value_conflict" in markdown
