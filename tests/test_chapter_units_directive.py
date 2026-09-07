# -*- coding: utf-8 -*-
"""Lock the authoritative chapter-unit convention.

All long-Markdown authoring profiles (narrative / bidding / engineering) embed
the SAME ``CHAPTER_UNITS_DIRECTIVE`` so the AI's chapter headings use exactly
the unit words the downstream chapter recognizer accepts.  A past edit had
dropped 单元 / 部分 / 卷 (only listing 章/篇/部), which made AI output mismatch
the recognizer.

The wording and the unit set are no longer hand-typed here or in the prompts:
both come from ``src/shared/engine.chapter_units`` (the same single source the
recognition tests consume).  These tests pin that ``CHAPTER_UNITS_DIRECTIVE`` is
*derived* from that shared source and covers every top-level unit, so a change
to the shared unit set (or an over-trimmed rewrite) fails loudly.
"""

from src.assistant.application.capability_registry import (
    BIDDING_PROMPT_PROFILE_ID,
    CHAPTER_UNITS_DIRECTIVE,
    ENGINEERING_PROMPT_PROFILE_ID,
    NARRATIVE_PROMPT_PROFILE_ID,
    system_prompt_for_profile,
)
from src.assistant.ui.chapter_workbench_mixin import _H1, _parse_outline_blocks
from src.assistant.application.directory_authoring_parser import _is_top_chapter
from src.modules.structure import heading_recognition as hr
from src.shared.engine.chapter_units import (
    SECTION_UNIT,
    TOP_LEVEL_UNITS,
    TOP_LEVEL_UNIT_EXAMPLES,
    render_chapter_units_directive,
)


def _workbench_kind(line: str):
    """_parse_outline_blocks 对该行判定的块类型（忽略纯图行）。"""
    for kind, _text in _parse_outline_blocks(line):
        if kind != 4:  # _IMAGE
            return kind
    return None


def test_directive_is_derived_from_shared_unit_set():
    # Guard against someone hand-editing CHAPTER_UNITS_DIRECTIVE back to a
    # trimmed prose block that no longer matches the shared recognizer units.
    assert CHAPTER_UNITS_DIRECTIVE == render_chapter_units_directive()


def test_directive_lists_every_top_level_unit():
    for unit in TOP_LEVEL_UNITS:
        assert f"第X{unit}" in CHAPTER_UNITS_DIRECTIVE, (
            f"CHAPTER_UNITS_DIRECTIVE no longer mentions 第X{unit}"
        )


def test_directive_explicitly_covers_unit_part_volume():
    # The regression target: 单元/部分/卷 must appear both as bare unit words
    # and in their 第X… forms (识别器把它们当一级标题/拆分边界).
    for unit in ("单元", "部分", "卷"):
        assert unit in CHAPTER_UNITS_DIRECTIVE
        assert f"第X{unit}" in CHAPTER_UNITS_DIRECTIVE


def test_directive_describes_subsections_and_items():
    assert f"第X{SECTION_UNIT}" in CHAPTER_UNITS_DIRECTIVE
    assert "x.x" in CHAPTER_UNITS_DIRECTIVE
    assert "1）2）" in CHAPTER_UNITS_DIRECTIVE


def test_directive_forbids_unrecognized_words():
    # Guards the trailing prohibition so made-up units cannot slip back in.
    assert "回" in CHAPTER_UNITS_DIRECTIVE
    assert "自造单位" in CHAPTER_UNITS_DIRECTIVE


def test_every_long_markdown_profile_embeds_the_constant():
    profiles = (
        NARRATIVE_PROMPT_PROFILE_ID,
        BIDDING_PROMPT_PROFILE_ID,
        ENGINEERING_PROMPT_PROFILE_ID,
    )
    for profile_id in profiles:
        prompt = system_prompt_for_profile(profile_id)
        assert CHAPTER_UNITS_DIRECTIVE in prompt, (
            f"{profile_id} no longer embeds CHAPTER_UNITS_DIRECTIVE"
        )
        # Every top-level unit reaches the actual prompt text.
        for unit in TOP_LEVEL_UNITS:
            assert f"第X{unit}" in prompt, (
                f"{profile_id} prompt is missing 第X{unit}"
            )


# ---------------------------------------------------------------------------
# 同步校验：单一共用样例清单 —— 每单位既被四处识别点接受，又被能力提示点名。
# ---------------------------------------------------------------------------
def test_shared_unit_examples_sync_prompt_and_pure_recognition_points():
    """对每个顶级单位：提示里的“第X<单位>”与代表标题一一对应，且工作台/大纲
    解析/docx 标题识别三处都判为 H1/顶级。

    这保证“只测试却提示漏掉单位”或“提示列了单位却识别不了”都不会静默发生。
    """
    for unit, sample in TOP_LEVEL_UNIT_EXAMPLES:
        assert f"第X{unit}" in CHAPTER_UNITS_DIRECTIVE, (
            f"directive 未提及 {unit}，但识别测试覆盖了它"
        )
        assert _workbench_kind(sample) == _H1, f"工作台未把 {sample!r} 判为 H1"
        assert _is_top_chapter(sample) is True, f"大纲解析未把 {sample!r} 当顶级"
        assert hr._detect_by_pattern(sample) == 1, f"docx 标题识别未把 {sample!r} 判为 level 1"


def test_shared_unit_examples_recognized_in_real_docx(tmp_path):
    """把共用样例清单逐条写成真实 docx 的 Heading-1，端到端拆分应得到同样顺序的
    顶级章节，印证四处识别点与提示同源。"""
    from docx import Document

    from src.assistant.application.chapter_document_editor import (
        detect_chapter_outline,
    )

    doc = Document()
    for unit, sample in TOP_LEVEL_UNIT_EXAMPLES:
        doc.add_heading(sample, level=1)
        doc.add_paragraph(f"{unit}的正文占位内容，足够长。")
    path = tmp_path / "units_sync.docx"
    doc.save(str(path))

    chapters = detect_chapter_outline(path)
    titles = [c.title for c in chapters]
    expected = [sample for _unit, sample in TOP_LEVEL_UNIT_EXAMPLES]
    assert titles == expected, f"真实 docx 拆分未得到样例全序：{titles}"
    assert all(c.level == 1 for c in chapters)
