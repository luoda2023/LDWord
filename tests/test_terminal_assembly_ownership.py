from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path
from shutil import copy2

from docx import Document
import pytest

from src.config.feature_configs import OutputConfig
from src.config.material_context import MaterialExecutionContext
from src.config.master_library import get_master
from src.config.resolver import resolve_config
from src.config.scene import DeliveryPreset, InputSourceProfile, SceneWorkspace
from src.config.template import TemplateConfig
from src.modules.base import BaseModule, ModuleMeta
from src.pipeline.runner import Pipeline, pipeline_terminal_assembly_owner
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner


class _MutationSentinel(BaseModule):
    meta = ModuleMeta(
        name="terminal_mutation_sentinel",
        description="Must never run after a terminal assembler claims ownership",
        category="test",
        execution_phase="format",
        scope_behavior="document_level",
    )

    def __init__(self, applied: list[str]) -> None:
        self._applied = applied

    def apply(self, doc, config, tracker, context) -> None:
        self._applied.append("applied")
        doc.add_paragraph("GENERIC MODULE MUTATION")


def _source_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("SOURCE MUST REMAIN UNCHANGED")
    doc.save(path)


def _frozen_master(tmp_path: Path, master_id: str, mode_id: str):
    master = get_master(master_id, mode_id=mode_id)
    assert master is not None
    frozen_path = tmp_path / f"frozen-{master.docx_path.name}"
    copy2(master.docx_path, frozen_path)
    return replace(
        master,
        docx_path=frozen_path,
        execution_frozen=True,
    )


def _exam_config():
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        default_delivery_preset_id="student",
        input_source_profile=InputSourceProfile(
            material_schema_id="exam_items_v1",
        ),
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                label="Student",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
            )
        ],
        exam_paper=None,
    )
    payload = {
        "paper_title": "Terminal ownership exam",
        "sections": [
            {
                "title": "Questions",
                "questions": [
                    {
                        "stem": "1 + 1 = ?",
                        "answer": "2",
                        "score": 1,
                    }
                ],
            }
        ],
    }
    return resolve_config(
        TemplateConfig(),
        scene,
        entity_data={"exam_items": json.dumps(payload)},
    )


def _official_config(*, complete: bool = True):
    scene = SceneWorkspace(
        scene_id="official",
        mode_id="official",
        input_source_profile=InputSourceProfile(
            material_schema_id="official_document_v1",
        ),
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                artifacts=OutputConfig(
                    final_docx=False,
                    report_json=False,
                    report_markdown=False,
                ),
            )
        ],
    )
    entity_data = {
        "document_type": "notice",
        "title": "Terminal ownership notice",
        "organization": "Example organization",
        "document_no": "EX-2026-01",
        "issue_date": "2026-07-14",
    }
    if complete:
        entity_data["body"] = "The official assembler owns this artifact set."
    return resolve_config(TemplateConfig(), scene, entity_data=entity_data)


def test_exam_terminal_assembler_skips_generic_modules_and_ordinary_output(tmp_path):
    source = tmp_path / "exam.docx"
    output_dir = tmp_path / "out"
    _source_docx(source)
    applied: list[str] = []

    result = Pipeline(
        [_MutationSentinel(applied)],
        _exam_config(),
        output_dir=str(output_dir),
        output_suffix="_ordinary",
        exam_master=_frozen_master(tmp_path, "default_exam", "exam"),
    ).execute(str(source))

    assert result.success is True
    assert result.context.terminal_assembly_owner == "exam"
    assert applied == []
    assert set(result.output_paths) == {"student", "answer_key"}
    assert not (output_dir / "exam_ordinary.docx").exists()
    assert Document(source).paragraphs[0].text == "SOURCE MUST REMAIN UNCHANGED"
    ownership = result.tracker.get_by_module("terminal_assembly_owner")
    assert len(ownership) == 1
    assert ownership[0].success is True
    assert "generic modules" in ownership[0].after


def test_official_terminal_assembler_skips_generic_modules_and_ordinary_output(tmp_path):
    source = tmp_path / "official.docx"
    output_dir = tmp_path / "out"
    _source_docx(source)
    applied: list[str] = []

    result = Pipeline(
        [_MutationSentinel(applied)],
        _official_config(),
        output_dir=str(output_dir),
        output_suffix="_ordinary",
        official_master=_frozen_master(
            tmp_path,
            "official_gbt_standard",
            "official",
        ),
        official_document_type_id="notice",
    ).execute(str(source))

    assert result.success is True
    assert result.context.terminal_assembly_owner == "official"
    assert applied == []
    assert set(result.output_paths) == {
        "official_docx",
        "internal_review_docx",
        "archive_manifest",
        "archive_manifest_md",
    }
    assert not (output_dir / "official_ordinary.docx").exists()
    assert Document(source).paragraphs[0].text == "SOURCE MUST REMAIN UNCHANGED"


def test_blocked_terminal_assembler_never_falls_back_to_ordinary_output(tmp_path):
    source = tmp_path / "official.docx"
    output_dir = tmp_path / "out"
    _source_docx(source)

    result = Pipeline(
        [],
        _official_config(complete=False),
        output_dir=str(output_dir),
        output_suffix="_ordinary",
        official_master=_frozen_master(
            tmp_path,
            "official_gbt_standard",
            "official",
        ),
        official_document_type_id="notice",
    ).execute(str(source))

    assert result.success is False
    assert result.output_paths == {}
    assert result.context.official_document_assembly.status == "missing_required_fields"
    assert not (output_dir / "official_ordinary.docx").exists()


def test_conflicting_terminal_schemas_fail_before_any_assembler_publishes(tmp_path):
    source = tmp_path / "conflict.docx"
    output_dir = tmp_path / "out"
    _source_docx(source)
    config = _exam_config()
    config.input_source_profile.material_schema_ids = [
        "exam_items_v1",
        "official_document_v1",
    ]

    result = Pipeline([], config, output_dir=str(output_dir)).execute(str(source))

    assert pipeline_terminal_assembly_owner(config) == "conflict"
    assert result.success is False
    assert result.output_paths == {}
    assert result.context.exam_delivery_runtime is None
    assert result.context.official_document_assembly is None
    assert not output_dir.exists()
    assert "ownership is ambiguous" in str(result.error)


def test_unknown_material_schema_never_falls_back_to_generic_pipeline(tmp_path):
    source = tmp_path / "unknown-schema.docx"
    output_dir = tmp_path / "out"
    _source_docx(source)
    applied: list[str] = []
    config = _official_config()
    config.input_source_profile.material_schema_id = "official_document_v1_typo"

    result = Pipeline(
        [_MutationSentinel(applied)],
        config,
        output_dir=str(output_dir),
    ).execute(str(source))

    assert pipeline_terminal_assembly_owner(config) == "invalid"
    assert result.success is False
    assert result.status == "failed"
    assert applied == []
    assert result.output_paths == {}
    assert not output_dir.exists()
    assert (
        "pipeline_configuration_invalid:unknown_material_schema:"
        "official_document_v1_typo"
    ) in str(result.error)


@pytest.mark.parametrize(
    "profile_name",
    ("input_source_profile", "compliance_profile"),
)
def test_unknown_failure_policy_never_falls_back_to_generic_pipeline(
    tmp_path,
    profile_name,
):
    source = tmp_path / f"unknown-policy-{profile_name}.docx"
    output_dir = tmp_path / f"out-{profile_name}"
    _source_docx(source)
    applied: list[str] = []
    config = _official_config()
    getattr(config, profile_name).failure_policy = "blok"

    result = Pipeline(
        [_MutationSentinel(applied)],
        config,
        output_dir=str(output_dir),
    ).execute(str(source))

    assert pipeline_terminal_assembly_owner(config) == "invalid"
    assert result.success is False
    assert result.status == "failed"
    assert applied == []
    assert result.output_paths == {}
    assert not output_dir.exists()
    assert f"{profile_name}.failure_policy" in str(result.error)
    assert "execution_failure_policy_invalid:blok" in str(result.error)


def test_terminal_assembler_rejects_generic_caller_owned_stage_transaction(tmp_path):
    source = tmp_path / "official.docx"
    output_dir = tmp_path / "out"
    _source_docx(source)

    result = Pipeline(
        [],
        _official_config(),
        output_dir=str(output_dir),
        output_path_overrides={"final": tmp_path / "owned-stage.docx"},
        output_paths_are_owned_stages=True,
    ).execute(str(source))

    assert result.success is False
    assert result.output_paths == {}
    assert result.context.official_document_assembly is None
    assert not output_dir.exists()
    assert "cannot run inside the generic caller-owned stage transaction" in str(
        result.error
    )


def test_workbench_blocks_active_generic_material_rules_for_terminal_owner(
    tmp_path,
    monkeypatch,
):
    import src.services.production_runtime.execution_runtime as runtime_module

    source = tmp_path / "official.docx"
    _source_docx(source)
    config = _official_config()
    runner = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(scene_id="official", mode_id="official"),
    )
    monkeypatch.setattr(runtime_module, "material_assembly_is_active", lambda _ctx: True)

    def forbidden_material_assembly(**_kwargs):
        raise AssertionError("generic material assembly must not own terminal output")

    monkeypatch.setattr(
        runtime_module,
        "run_workbench_material_assembly",
        forbidden_material_assembly,
    )

    result, _elapsed, modules_enabled, modules_total = runner._execute_config(
        input_path=source,
        output_dir=tmp_path / "out",
        config=config,
        progress_cb=lambda *_args: None,
        cancel_check=lambda: False,
        material_context=MaterialExecutionContext(),
    )

    assert result.success is False
    assert result.context.terminal_assembly_owner == "official"
    assert "cannot ignore or delegate active generic material assembly rules" in str(
        result.error
    )
    assert modules_enabled == 0
    assert modules_total > 0


def test_workbench_allows_field_only_terminal_context_and_reports_zero_modules(tmp_path):
    source = tmp_path / "official.docx"
    _source_docx(source)
    config = _official_config()
    runner = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(scene_id="official", mode_id="official"),
    )

    result, _elapsed, modules_enabled, modules_total = runner._execute_config(
        input_path=source,
        output_dir=tmp_path / "out",
        config=config,
        progress_cb=lambda *_args: None,
        cancel_check=lambda: False,
        material_context=MaterialExecutionContext(
            package_id="package-1",
            profile_id="profile-1",
            entity_data={"title": "Field-only projection"},
        ),
        official_master=_frozen_master(
            tmp_path,
            "official_gbt_standard",
            "official",
        ),
        official_document_type_id="notice",
    )

    assert result.success is True
    assert result.context.terminal_assembly_owner == "official"
    assert modules_enabled == 0
    assert modules_total > 0
