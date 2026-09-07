# -*- coding: utf-8 -*-
"""Regression tests for the 「润色整篇」 part-polish flow.

From a dock 篇/部分 group the panel launches a background worker that rewrites
every chapter under the part sequentially (each chapter a separate model call
so its result can be written back reliably), then writes the polished part into
a *copy* of the bound .docx preserving headings + layout.

These tests cover the pure/core pieces that need no full panel:
* :class:`_PartPolishWorker` — streams deltas and emits one ``chapter_finished``
  per target (in order), then ``finished``; a provider failure stops the run;
* the dock 「润色整篇」button — clicking it emits ``polish_part_requested`` and
  goes busy while a run is active (ignoring re-entry);
* ``_apply_part_polish_docx`` — applies every polished chapter cumulatively
  onto a nested 篇>章 docx, replacing only the target bodies while keeping all
  headings (count + text) intact and leaving other chapters untouched.

Run with ``QT_QPA_PLATFORM=offscreen`` and no ``pytest-qt``: a session-scoped
``QApplication`` fixture is provided here.
"""

import os
import time

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PySide6.QtWidgets import QApplication

from src.assistant.application.chapter_document_editor import (
    detect_chapter_outline,
)
from src.assistant.runtime.mock_provider import MockModelGateway
from src.assistant.ui.chapter_outline_dock import ChapterOutlineDock
from src.assistant.ui.chapter_rewrite_mixin import (
    _PartPolishWorker,
    _apply_part_polish_docx,
)


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    yield app


def _spin_until(qapp, predicate, timeout_s=5.0):
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        qapp.processEvents()
        if predicate():
            return True
        time.sleep(0.01)
    return predicate()


def _run_worker(qapp, gateway, **kwargs):
    worker = _PartPolishWorker(gateway, **kwargs)
    state = {"deltas": [], "done": [], "finished": False, "failed": None}
    worker.delta.connect(lambda i, d: state["deltas"].append((i, d)))
    worker.chapter_finished.connect(
        lambda i, t, text: state["done"].append((i, t, text))
    )
    worker.finished.connect(lambda: state.__setitem__("finished", True))
    worker.failed.connect(lambda e: state.__setitem__("failed", e))
    worker.start()
    assert _spin_until(
        qapp, lambda: state["finished"] or state["failed"] is not None
    ), state
    worker.shutdown()
    return state


PART_TARGETS = [
    {"index": 1, "title": "第一章 项目背景", "body": "旧正文一。"},
    {"index": 2, "title": "第二章 建设目标", "body": "旧正文二。"},
]


# ---------------------------------------------------------------------------
# Dock polish button
# ---------------------------------------------------------------------------
def test_dock_part_polish_button_emits(qapp):
    dock = ChapterOutlineDock()
    dock.show()
    _spin_until(qapp, lambda: dock.isVisible())
    titles = ["第一章 项目背景", "第二章 建设目标", "第三章 总体设计", "第四章 实施方案"]
    parts = ["第一篇 总论", "第一篇 总论", "第二篇 工程设计", ""]
    dock.set_outline(titles, parts=parts)
    _spin_until(qapp, lambda: not dock._rows[4].isHidden())

    part1 = dock.part_widget("第一篇 总论")
    assert part1 is not None
    emitted = []
    dock.polish_part_requested.connect(emitted.append)
    part1._polish.click()
    qapp.processEvents()
    assert emitted == ["第一篇 总论"]

    # Busy disables the button and ignores further clicks.
    part1.set_busy(True)
    part1._polish.click()
    qapp.processEvents()
    assert emitted == ["第一篇 总论"]
    assert not part1._polish.isEnabled()
    part1.set_busy(False)
    assert part1._polish.isEnabled()
    dock.hide()
    dock.deleteLater()


# ---------------------------------------------------------------------------
# Worker: sequential polish
# ---------------------------------------------------------------------------
def test_part_polish_worker_rewrites_every_chapter_in_order(qapp):
    gateway = MockModelGateway(reply="润色后的正文内容，更专业精炼。")
    state = _run_worker(
        qapp,
        gateway,
        part_title="第一篇 总论",
        targets=list(PART_TARGETS),
    )
    assert state["failed"] is None, state
    assert state["finished"] is True
    # one chapter_finished per target, in ascending index order
    assert [d[0] for d in state["done"]] == [1, 2]
    assert all(len(str(d[2] or "").strip()) > 0 for d in state["done"])
    # streamed at least one delta for each chapter
    assert len(state["deltas"]) >= 2


def test_part_polish_worker_failure_stops_and_reports(qapp):
    class _FailGateway:
        model = "mock"

        def stream(self, request):
            raise RuntimeError("boom_provider")

    state = _run_worker(
        qapp,
        _FailGateway(),
        part_title="第一篇 总论",
        targets=list(PART_TARGETS),
    )
    assert state["failed"] is not None
    assert "boom_provider" in state["failed"]


# ---------------------------------------------------------------------------
# Cumulative .docx write-back (layout preserved)
# ---------------------------------------------------------------------------
def _make_nested_docx(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("第一篇 总论", level=1)
    doc.add_paragraph("（篇头说明。）")
    doc.add_heading("第一章 项目背景", level=2)
    doc.add_paragraph("旧正文甲，待润色。")
    doc.add_heading("第二章 建设目标", level=2)
    doc.add_paragraph("旧正文乙，待润色。")
    doc.add_heading("第二篇 工程设计", level=1)
    doc.add_paragraph("（篇头说明二。）")
    doc.add_heading("第三章 总体设计", level=2)
    doc.add_paragraph("旧正文丙。")
    path = tmp_path / "report.docx"
    doc.save(str(path))
    return str(path)


def _headings(path):
    from docx import Document

    doc = Document(str(path))
    return [
        p.text.strip()
        for p in doc.paragraphs
        if (getattr(p.style, "name", "") or "").startswith("Heading")
    ]


def test_apply_part_polish_docx_writes_all_targets_preserving_headings(tmp_path):
    src = _make_nested_docx(tmp_path)
    # Verify part-1 membership first (chapter 1 & 2 belong to 第一篇).
    chapters = detect_chapter_outline(src)
    assert [c.title for c in chapters if c.part_title == "第一篇 总论"] == [
        "第一章 项目背景",
        "第二章 建设目标",
    ]
    done = {1: "新正文甲。更精炼。", 2: "新正文乙。更通顺。"}
    out = _apply_part_polish_docx(src, done)
    assert os.path.isfile(out)

    headings = _headings(out)
    assert "第一篇 总论" in headings
    assert "第一章 项目背景" in headings
    assert "第二章 建设目标" in headings
    assert "第三章 总体设计" in headings
    # same number of heading paragraphs as the source (no structure change)
    assert len(headings) == len(_headings(src)) == 5

    texts = [p.text.strip() for p in __import__("docx").Document(out).paragraphs]
    joined = "\n".join(texts)
    assert "新正文甲" in joined and "旧正文甲" not in joined
    assert "新正文乙" in joined and "旧正文乙" not in joined
    # chapter 3 (other part) untouched
    assert "旧正文丙" in joined


def test_apply_part_polish_docx_output_is_separate_copy(tmp_path):
    src = _make_nested_docx(tmp_path)
    out = _apply_part_polish_docx(src, {1: "只润色第一章。"})
    assert os.path.abspath(out) != os.path.abspath(src)
    # original source unchanged (no bodies replaced in it)
    texts = [p.text.strip() for p in __import__("docx").Document(src).paragraphs]
    assert "旧正文甲" in "\n".join(texts)
