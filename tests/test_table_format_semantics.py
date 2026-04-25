import sys
from pathlib import Path

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.config.table_style_presets import COLOR_TABLE_VARIANTS, color_variant
from src.modules.structure.heading_recognition import DocSection, DocTree
from src.modules.table.table_format import (
    TableFormatModule,
    _build_smart_width_plan,
    _extract_cell_lines,
    _normalize_body_paren_breaks,
    _normalize_first_row_unit_breaks,
)
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn


def _build_weighted_table(doc: Document):
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Normalized concentration (mg/mL)"
    table.cell(0, 1).text = "Light"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "This body column is intentionally much longer than the first one"
    return table


def _build_compactable_table(doc: Document):
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(0, 2).text = "Remark"
    table.cell(1, 0).text = "Density"
    table.cell(1, 1).text = "1.23"
    table.cell(1, 2).text = "Lab note"
    return table


def _table_layout_type(table) -> str | None:
    layout = table._element.find(qn("w:tblPr"))
    if layout is None:
        return None
    tbl_layout = layout.find(qn("w:tblLayout"))
    if tbl_layout is None:
        return None
    return tbl_layout.get(qn("w:type"))


def _table_width(table) -> int | None:
    tbl_pr = table._element.find(qn("w:tblPr"))
    if tbl_pr is None:
        return None
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        return None
    return int(tbl_w.get(qn("w:w"), "0"))


def _grid_widths(table) -> list[int]:
    grid = table._element.find(qn("w:tblGrid"))
    if grid is None:
        return []
    return [int(col.get(qn("w:w"), "0")) for col in grid.findall(qn("w:gridCol"))]


def _table_border(table, side: str) -> tuple[str | None, str | None]:
    tbl_pr = table._element.find(qn("w:tblPr"))
    if tbl_pr is None:
        return None, None
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        return None, None
    border = tbl_borders.find(qn(f"w:{side}"))
    if border is None:
        return None, None
    return border.get(qn("w:val")), border.get(qn("w:color"))


def _cell_border(cell, side: str) -> tuple[str | None, str | None]:
    tc_pr = cell._element.find(qn("w:tcPr"))
    if tc_pr is None:
        return None, None
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        return None, None
    border = tc_borders.find(qn(f"w:{side}"))
    if border is None:
        return None, None
    return border.get(qn("w:val")), border.get(qn("w:color"))


def test_table_format_only_formats_body_tables_when_doc_tree_is_present():
    doc = Document()
    doc.add_paragraph("Cover")
    cover_table = _build_weighted_table(doc)
    doc.add_paragraph("Body")
    body_table = _build_weighted_table(doc)

    cover_before = _grid_widths(cover_table)
    body_before = _grid_widths(body_table)

    config = ResolvedConfig()
    config.table.border_mode = "keep"
    config.table.layout_mode = "compact"

    context = PipelineContext(
        doc_tree=DocTree(
            sections=[
                DocSection("cover", 0, 1),
                DocSection("body", 1, 2),
            ]
        )
    )

    TableFormatModule().apply(doc, config, ChangeTracker(), context)

    assert _table_layout_type(cover_table) is None
    assert _table_layout_type(body_table) == "fixed"
    assert _grid_widths(cover_table) == cover_before
    assert _grid_widths(body_table) != body_before


def test_table_format_keep_layout_preserves_body_table_widths():
    doc = Document()
    doc.add_paragraph("Body")
    table = _build_weighted_table(doc)

    before_layout = _table_layout_type(table)
    before_grid = _grid_widths(table)
    before_width = _table_width(table)

    config = ResolvedConfig()
    config.table.border_mode = "keep"
    config.table.layout_mode = "keep"

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert _table_layout_type(table) == before_layout
    assert _grid_widths(table) == before_grid
    assert _table_width(table) == before_width


def test_table_format_skips_rich_content_tables():
    doc = Document()
    doc.add_paragraph("Body A")
    rich_table = _build_weighted_table(doc)
    doc.add_paragraph("Body B")
    plain_table = _build_weighted_table(doc)

    run = rich_table.cell(0, 0).paragraphs[0].add_run()
    etree.SubElement(run._element, qn("w:drawing"))

    rich_before = _grid_widths(rich_table)
    plain_before = _grid_widths(plain_table)

    config = ResolvedConfig()
    config.table.border_mode = "keep"
    config.table.layout_mode = "compact"

    context = PipelineContext(
        doc_tree=DocTree(sections=[DocSection("body", 0, 2)])
    )

    TableFormatModule().apply(doc, config, ChangeTracker(), context)

    assert _table_layout_type(rich_table) is None
    assert _grid_widths(rich_table) == rich_before
    assert _table_layout_type(plain_table) == "fixed"
    assert _grid_widths(plain_table) != plain_before


def test_table_format_smart_levels_change_clustering_plan():
    raw_targets = [(0, 2600), (1, 3400), (2, 4700), (3, 6200)]

    plan_2 = _build_smart_width_plan(raw_targets, 7000, max_levels=2)
    plan_4 = _build_smart_width_plan(raw_targets, 7000, max_levels=4)

    assert len(set(plan_2.values())) <= 2
    assert len(set(plan_4.values())) >= 3


def test_table_format_compact_mode_is_narrower_than_full_mode():
    compact_doc = Document()
    compact_table = _build_compactable_table(compact_doc)

    full_doc = Document()
    full_table = _build_compactable_table(full_doc)

    compact_config = ResolvedConfig()
    compact_config.table.border_mode = "keep"
    compact_config.table.layout_mode = "compact"

    full_config = ResolvedConfig()
    full_config.table.border_mode = "keep"
    full_config.table.layout_mode = "full"

    TableFormatModule().apply(compact_doc, compact_config, ChangeTracker(), PipelineContext())
    TableFormatModule().apply(full_doc, full_config, ChangeTracker(), PipelineContext())

    assert _table_width(compact_table) is not None
    assert _table_width(full_table) is not None
    assert _table_width(compact_table) < _table_width(full_table)


def test_table_format_ignores_legacy_row_height_setting():
    doc = Document()
    table = _build_compactable_table(doc)

    config = ResolvedConfig()
    config.table.border_mode = "keep"
    config.table.layout_mode = "keep"
    config.table.row_height_pt = 24

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert table._element.find(f".//{qn('w:trHeight')}") is None


def test_table_format_normalizes_header_unit_and_body_parenthesis_breaks():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Normalized concentration (mg/mL)"
    table.cell(0, 1).text = "Light"
    table.cell(1, 0).text = "Sample"
    table.cell(1, 1).text = "Commercial source (Shanghai)"

    header_changes = _normalize_first_row_unit_breaks(table._element, [1600, 1200])
    body_changes = _normalize_body_paren_breaks(table._element, [1200, 1400])

    assert header_changes == 1
    assert body_changes == 1
    assert _extract_cell_lines(table.rows[0].cells[0]._element)[:2] == [
        "Normalized concentration",
        "(mg/mL)",
    ]
    assert _extract_cell_lines(table.rows[1].cells[1]._element)[:2] == [
        "Commercial source",
        "(Shanghai)",
    ]


def test_color_table_header_rule_variant_replaces_low_value_column_only_style():
    keys = [variant.key for variant in COLOR_TABLE_VARIANTS]
    assert "header_rule" in keys
    assert "header_columns" not in keys
    assert color_variant("header_columns").key == "header_rule"


def test_table_format_color_header_rule_adds_table_and_header_rules():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Density"
    table.cell(1, 1).text = "1.23"

    config = ResolvedConfig()
    config.table.border_mode = "color_table"
    config.table.color_table_accent = "green"
    config.table.color_table_variant = "header_rule"
    config.table.three_line_header_width_pt = 1.5
    config.table.three_line_bottom_width_pt = 0.75

    TableFormatModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert _table_border(table, "top") == ("single", "70AD47")
    assert _table_border(table, "bottom") == ("single", "70AD47")
    tbl_pr = table._element.find(qn("w:tblPr"))
    borders = tbl_pr.find(qn("w:tblBorders"))
    assert borders.find(qn("w:top")).get(qn("w:sz")) == "12"
    assert borders.find(qn("w:bottom")).get(qn("w:sz")) == "12"
    assert _table_border(table, "insideV")[0] == "none"
    assert _table_border(table, "insideH")[0] == "none"
    assert _cell_border(table.cell(0, 0), "bottom") == ("single", "70AD47")
    assert _cell_border(table.cell(0, 1), "bottom") == ("single", "70AD47")
    assert table.cell(0, 0)._element.find(qn("w:tcPr")).find(qn("w:tcBorders")).find(qn("w:bottom")).get(qn("w:sz")) == "6"
