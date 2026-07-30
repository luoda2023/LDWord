import sys
import json
import pytest
import threading
import zipfile
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from xml.etree import ElementTree as ET
from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.entity import EntityArchive, EntityProfile
from src.config.material_batch import build_material_batch_items
from src.config.material_context import MaterialExecutionContext
from src.config.material_mappings import load_material_mapping
from src.config.materials import (
    AssetInsertionRule,
    AssetItem,
    build_image_insertions,
    parse_asset_insertion_rules,
    scan_asset_collection,
)
from src.config.feature_configs import OutputConfig
from src.config.library import load_template_from_library
from src.config.object_preflight_evidence import build_object_preflight_evidence
from src.config.resolved import ImageInsertionItem
from src.config.resolver import resolve_config
from src.config.scene import DeliveryPreset, SceneWorkspace
from src.config.scene_family_application import apply_planned_scene_family_defaults
from src.config.template import TemplateConfig
from src.shared.engine.media_ops import paragraph_has_image
from src.shared.engine.material_timeline import (
    default_timeline_plan,
    default_timeline_segment,
)
from src.qt_api import QApplication
from src.ui.bridge import PanelBridge
from src.ui.adapters.workbench_execution_adapter import batch_execution_issue_items
from src.ui.panels.assets_panel import AssetsPanel
from src.ui.panels.workbench.feature_detail_panes import ContentDataDetailPane
from src.services.production_runtime.execution_runtime import (
    WorkbenchBatchProductionRunner,
    WorkbenchProductionRunner,
)
from src.services.execution_session import (
    build_execution_session_snapshot,
    cleanup_execution_session_resources,
)


class _QuietSimpleHTTPRequestHandler(SimpleHTTPRequestHandler):
    def log_message(self, format, *args):  # noqa: A002
        return


def _run_exam_material_context(
    *,
    tmp_path: Path,
    source: Path,
    scene: SceneWorkspace,
    material_context: MaterialExecutionContext,
) -> dict[str, object]:
    """Run an exam fixture through the same frozen-master boundary as the product."""

    scene.scene_id = "exam"
    scene.mode_id = "exam"
    scene.template_id = "default"
    scene.compatible_template_ids = ["default"]
    scene.master_id = "default_exam"
    material_context.mode_id = "exam"
    template = load_template_from_library("default", mode_id="exam")
    preflight = build_object_preflight_evidence(scene, source)
    snapshot = build_execution_session_snapshot(
        mode_id="exam",
        scene=scene,
        template=template,
        material_context=material_context,
        input_path=source,
        output_root=tmp_path / "out",
        plan_id="exam",
        template_id="default",
        object_preflight_confirmation_revision=preflight.source_revision,
        object_preflight_confirmation_digest=preflight.evidence_digest,
    )
    try:
        return WorkbenchProductionRunner(
            doc_path=str(source),
            template=template,
            scene=scene,
            output_dir=Path(snapshot.output_namespace),
            material_context=material_context,
            execution_session=snapshot,
        ).run(lambda *_args: None, lambda: False)
    finally:
        cleanup_execution_session_resources(snapshot)


def _run_official_material_context(
    *,
    tmp_path: Path,
    source: Path,
    scene: SceneWorkspace,
    material_context: MaterialExecutionContext,
    document_type_id: str,
) -> tuple[dict[str, object], Path]:
    """Run an official fixture through the frozen type/master session boundary."""

    scene.scene_id = "official"
    scene.mode_id = "official"
    scene.template_id = "official_gbt"
    scene.compatible_template_ids = ["official_gbt"]
    scene.master_id = "official_gbt_standard"
    template = load_template_from_library("official_gbt", mode_id="official")
    material_context.mode_id = "official"
    preflight = build_object_preflight_evidence(scene, source)
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=template,
        material_context=material_context,
        input_path=source,
        output_root=tmp_path / "out",
        plan_id="official",
        template_id="official_gbt",
        document_type_id=document_type_id,
        object_preflight_confirmation_revision=preflight.source_revision,
        object_preflight_confirmation_digest=preflight.evidence_digest,
    )
    output_dir = Path(snapshot.output_namespace)
    try:
        result = WorkbenchProductionRunner(
            doc_path=str(source),
            template=template,
            scene=scene,
            output_dir=output_dir,
            material_context=material_context,
            execution_session=snapshot,
        ).run(lambda *_args: None, lambda: False)
        return result, output_dir
    finally:
        cleanup_execution_session_resources(snapshot)


class _HeaderAuthSimpleHTTPRequestHandler(_QuietSimpleHTTPRequestHandler):
    expected_header_name = "Authorization"
    expected_header_value = "Bearer exam-secret"

    def do_GET(self):
        if self.headers.get(self.expected_header_name) != self.expected_header_value:
            self.send_response(401)
            self.end_headers()
            self.wfile.write(b"unauthorized")
            return
        super().do_GET()


class _ConditionalNotModifiedSimpleHTTPRequestHandler(_QuietSimpleHTTPRequestHandler):
    expected_etag = '"exam-question-figure-v1"'
    expected_last_modified = "Wed, 01 Jan 2020 00:00:00 GMT"
    request_headers_seen: list[dict[str, str]] = []

    def do_GET(self):
        self.__class__.request_headers_seen.append(
            {
                "If-None-Match": self.headers.get("If-None-Match", ""),
                "If-Modified-Since": self.headers.get("If-Modified-Since", ""),
            }
        )
        if (
            self.headers.get("If-None-Match") == self.expected_etag
            or self.headers.get("If-Modified-Since") == self.expected_last_modified
        ):
            self.send_response(304)
            self.send_header("ETag", self.expected_etag)
            self.send_header("Last-Modified", self.expected_last_modified)
            self.end_headers()
            return
        super().do_GET()


class _FlakyRemoteAssetSimpleHTTPRequestHandler(_QuietSimpleHTTPRequestHandler):
    failed_response_count = 2
    request_count = 0

    def do_GET(self):
        self.__class__.request_count += 1
        if self.__class__.request_count <= self.__class__.failed_response_count:
            self.send_response(503)
            self.end_headers()
            self.wfile.write(b"temporary unavailable")
            return
        super().do_GET()


def _app():
    return QApplication.instance() or QApplication([])


def _start_http_file_server(
    root: Path,
    handler_cls=_QuietSimpleHTTPRequestHandler,
) -> tuple[ThreadingHTTPServer, str]:
    handler = partial(handler_cls, directory=str(root))
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    return server, f"http://127.0.0.1:{server.server_address[1]}"


def test_material_execution_context_normalizes_payload_and_clones():
    context = MaterialExecutionContext.from_payload(
        {
            "archive_id": "archive",
            "profile_id": "main",
            "profile_name": "投标人主体",
            "entity_data": {"company_name": "测试公司"},
            "entity_assets_dir": "/tmp/assets",
            "images": [{"path": "logo.png", "position": 2, "width_cm": "6.5"}],
            "asset_items": [
                {
                    "role": "Question Figure",
                    "label": "题目图片",
                    "path": "question.png",
                    "tags": ["exam", 1],
                    "width_cm": "7.25",
                    "metadata": {
                        "alt_text": "加法示意图",
                        "empty": "",
                        "caption": 12,
                    },
                }
            ],
        }
    )

    cloned = context.clone()
    cloned.entity_data["company_name"] = "被修改"
    cloned.images[0].path = "changed.png"
    cloned.asset_items[0].metadata["alt_text"] = "被修改说明"

    assert context.archive_id == "archive"
    assert context.profile_id == "main"
    assert context.profile_name == "投标人主体"
    assert context.entity_data["company_name"] == "测试公司"
    assert context.entity_assets_dir == "/tmp/assets"
    assert isinstance(context.images[0], ImageInsertionItem)
    assert context.images[0].width_cm == 6.5
    assert context.images[0].path == "logo.png"
    assert not hasattr(context, "replacements")
    assert context.asset_items[0].role == "question_figure"
    assert context.asset_items[0].tags == ["exam", "1"]
    assert context.asset_items[0].width_cm == 7.25
    assert context.asset_items[0].metadata == {
        "alt_text": "加法示意图",
        "caption": "12",
    }


def test_panel_bridge_stores_material_context_as_independent_runtime_state():
    bridge = PanelBridge()
    context = MaterialExecutionContext(entity_data={"company_name": "测试公司"})

    bridge.set_current_material_context(context, emit_signal=False)
    context.entity_data["company_name"] = "外部修改"

    current = bridge.current_material_context()
    current.entity_data["company_name"] = "再次修改"

    assert bridge.current_material_context().entity_data["company_name"] == "测试公司"


def test_asset_collection_scan_and_rules_build_image_insertions(tmp_path):
    logo_path = tmp_path / "logo.png"
    Image.new("RGB", (24, 24), color="red").save(logo_path)

    collection = scan_asset_collection(tmp_path)
    rules = parse_asset_insertion_rules("logo={{@img:logo}}")
    insertions = build_image_insertions(collection.items, rules)

    assert collection.items[0].role == "logo"
    assert collection.items[0].path == str(logo_path)
    assert isinstance(rules[0], AssetInsertionRule)
    assert insertions == [
        ImageInsertionItem(path=str(logo_path), position="{{@img:logo}}", width_cm=6.0)
    ]


def test_material_context_expands_asset_rules_into_resolve_images(tmp_path):
    logo_path = tmp_path / "logo.png"
    Image.new("RGB", (24, 24), color="red").save(logo_path)
    collection = scan_asset_collection(tmp_path)

    context = MaterialExecutionContext(
        asset_items=collection.items,
        image_rules=[AssetInsertionRule(asset_role="logo", target="{{@img:logo}}", width_cm=2.5)],
    )

    kwargs = context.to_resolve_kwargs()

    assert kwargs["images"] == [
        ImageInsertionItem(path=str(logo_path), position="{{@img:logo}}", width_cm=2.5)
    ]
    assert context.missing_required_asset_roles() == []


def test_material_context_carries_profile_field_aliases_to_resolver_kwargs():
    context = MaterialExecutionContext(
        entity_data={"company_name": "测试公司"},
        field_scopes={"company_name": "fixed", "project_name": "floating"},
        field_aliases={"company": "company_name"},
    )

    cloned = context.clone()
    cloned.field_aliases["company"] = "other_field"
    kwargs = context.to_resolve_kwargs()
    resolved = resolve_config(TemplateConfig(), SceneWorkspace(), **kwargs)

    assert context.field_aliases == {"company": "company_name"}
    assert kwargs["field_aliases"] == {"company": "company_name"}
    assert kwargs["field_scopes"] == {
        "company_name": "fixed",
        "project_name": "floating",
    }
    assert resolved.field_scopes == kwargs["field_scopes"]
    assert resolved.field_aliases == {"company": "company_name"}
    assert MaterialExecutionContext.from_payload(
        {"field_aliases": {"company": "company_name"}}
    ).field_aliases == {"company": "company_name"}


def test_material_mapping_loads_flat_json_as_entity_data(tmp_path):
    mapping_path = tmp_path / "mapping.json"
    mapping_path.write_text(
        json.dumps({"company_name": "测试公司", "legal_person": "张三"}, ensure_ascii=False),
        encoding="utf-8",
    )

    payload = load_material_mapping(mapping_path)

    assert payload.entity_data == {"company_name": "测试公司", "legal_person": "张三"}
    assert not hasattr(payload, "replacements")


def test_material_mapping_rejects_json_replacements(tmp_path):
    mapping_path = tmp_path / "replacements.json"
    mapping_path.write_text(
        json.dumps(
            {"replacements": [{"old": "{{@text:project}}", "new": "Project"}]},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="不再支持替换规则"):
        load_material_mapping(mapping_path)


def test_material_mapping_rejects_csv_replacements(tmp_path):
    mapping_path = tmp_path / "replacements.csv"
    mapping_path.write_text("old,new\n{{@text:project_name}},测试项目\n", encoding="utf-8")

    with pytest.raises(ValueError, match="不再支持 old/new 替换规则"):
        load_material_mapping(mapping_path)


def test_material_mapping_rejects_excel_replacements(tmp_path):
    from openpyxl import Workbook

    mapping_path = tmp_path / "replacements.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["old", "new"])
    sheet.append(["{{@text:project_name}}", "Project A"])
    workbook.save(mapping_path)

    with pytest.raises(ValueError, match="不再支持 old/new 替换规则"):
        load_material_mapping(mapping_path)


def test_material_batch_items_build_profile_contexts_and_output_dirs(tmp_path):
    archive = EntityArchive(
        archive_id="archive",
        archive_name="测试档案",
        profiles=[
            EntityProfile(
                profile_id="a",
                profile_name="主体 A",
                fields={"company_name": "测试公司A"},
            ),
            EntityProfile(
                profile_id="b",
                profile_name="主体 B",
                fields={"company_name": "测试公司B"},
            ),
        ],
    )

    items = build_material_batch_items(
        archive,
        profile_ids=["b"],
        base_output_dir=tmp_path,
        output_dir_template="{entity_name}",
    )

    assert len(items) == 1
    assert items[0].profile_id == "b"
    assert items[0].context.entity_data == {"company_name": "测试公司B"}
    assert items[0].context.field_aliases == {}
    assert Path(items[0].output_dir).name == "测试公司B"


def test_material_batch_items_inherit_image_bindings_and_scan_profile_assets(tmp_path):
    assets_dir = tmp_path / "entity_assets"
    assets_dir.mkdir()
    logo_path = assets_dir / "logo.png"
    Image.new("RGB", (24, 24), color="red").save(logo_path)

    archive = EntityArchive(
        archive_id="archive",
        profiles=[
            EntityProfile(
                profile_id="a",
                profile_name="Entity A",
                fields={"company_name": "Company A"},
                field_aliases={"company": "company_name"},
                assets_dir=str(assets_dir),
            )
        ],
    )

    items = build_material_batch_items(
        archive,
        base_output_dir=tmp_path,
        base_context=MaterialExecutionContext(
            image_rules=[AssetInsertionRule(asset_role="logo", target="{{@img:logo}}")],
        ),
    )

    assert len(items) == 1
    assert not hasattr(items[0].context, "replacements")
    assert items[0].context.field_aliases == {"company": "company_name"}
    assert items[0].context.asset_items[0].path == str(logo_path)
    assert items[0].context.missing_required_asset_roles() == []


def test_content_data_detail_exports_simple_entity_material_context():
    _app()
    detail = ContentDataDetailPane()
    try:
        detail._profile_name_edit.setText("投标人主体")
        detail._entity_fields_edit.set_text("company_name=测试公司\nlegal_person: 张三")
        detail._assets_picker.set_path("D:/assets/company")

        context = detail.material_context()

        assert context.profile_name == "投标人主体"
        assert context.entity_assets_dir == "D:/assets/company"
        assert context.entity_data == {
            "company_name": "测试公司",
            "legal_person": "张三",
        }
        assert "资料已填 2 项" in detail._summary.text()
    finally:
        detail.close()


def test_content_data_detail_preserves_full_material_runtime_metadata():
    _app()
    detail = ContentDataDetailPane()
    source = MaterialExecutionContext(
        archive_id="archive-a",
        profile_id="profile-a",
        profile_name="资料 A",
        entity_data={"公司": "测试公司"},
        field_aliases={"company": "公司"},
        image_rules=[AssetInsertionRule(asset_role="logo", target="{{@img:logo}}")],
        exact_material_placeholders=True,
    )
    try:
        detail.set_material_context(source, emit_signal=False)

        exported = detail.material_context()

        assert exported.archive_id == "archive-a"
        assert exported.profile_id == "profile-a"
        assert exported.field_aliases == {"company": "公司"}
        assert not hasattr(exported, "replacements")
        assert exported.image_rules == source.image_rules
        assert exported.exact_material_placeholders is True
        assert detail._confidence_row.isHidden()
        assert "精确 {{@text:字段}} 匹配" in detail._summary.text()
    finally:
        detail.close()


def test_assets_panel_material_context_builds_images_from_role_binding(tmp_path):
    _app()
    logo_path = tmp_path / "logo.png"
    Image.new("RGB", (24, 24), color="red").save(logo_path)
    panel = AssetsPanel(PanelBridge())
    try:
        panel._archive_name_edit.setText("测试资料包")
        panel._asset_paths["logo"] = str(logo_path)

        context = panel.material_context()
        images = context.to_resolve_kwargs()["images"]

        assert context.archive_name == "测试资料包"
        assert context.asset_items[0].role == "logo"
        assert context.image_material_rules["image:logo"].source_role == "logo"
        assert context.image_material_rules["image:logo"].anchor_token == "{{@img:LOGO1}}"
        # The new rule owns this role.  Sending the same source into the legacy
        # ImageInsertionModule would create an unreceipted duplicate insertion.
        assert images == []
    finally:
        panel.close()


def test_assets_panel_loads_json_mapping_into_token_fields(tmp_path):
    _app()
    mapping_path = tmp_path / "mapping.json"
    mapping_path.write_text(
        json.dumps(
            {
                "entity_data": {"company_name": "测试公司"},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    panel = AssetsPanel(PanelBridge())
    try:
        panel.load_mapping_from_path(mapping_path)
        context = panel.material_context()

        assert context.entity_data == {"company_name": "测试公司"}
        assert not hasattr(context, "replacements")
        assert "字段资料 1 项 · 已填写 1/1" in panel._summary.text()
        assert not hasattr(panel, "_replacement_rules_edit")
        assert not hasattr(panel, "_image_rules_edit")
    finally:
        panel.close()


def test_workbench_runner_applies_material_context_entity_fields(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("公司：{{@text:company_name}}")
    doc.save(source)

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            profile_name="投标人主体",
            entity_data={"company_name": "测试公司"},
        ),
    ).run(lambda *_args: None, lambda: False)

    output_doc = Document(payload["output_path"])
    output_text = "\n".join(paragraph.text for paragraph in output_doc.paragraphs)
    report_json = source.parent / "output" / "source_changes.json"

    assert payload["status"] == "success"
    assert "测试公司" in output_text
    assert "{{@text:company_name}}" not in output_text
    assert report_json.exists()
    assert '"rule_name": "entity_fill"' in report_json.read_text(encoding="utf-8")


def test_workbench_runner_surfaces_official_assembly_from_material_context(tmp_path):
    source = tmp_path / "official_source.docx"
    Document().save(source)

    scene = SceneWorkspace(
        scene_id="official",
        mode_id="official",
        template_id="official_gbt",
        master_id="official_gbt_standard",
    )
    scene.input_source_profile.material_schema_id = "official_document_v1"
    scene.input_source_profile.material_schema_ids = []
    scene.input_source_profile.failure_policy = "warn"
    scene.compliance_profile.rule_family = "official_document"
    scene.compliance_profile.profile_id = "official_document"
    scene.default_material_profile_id = "official:minutes"
    scene.compliance_profile.object_preflight.enabled = False
    scene.default_delivery_preset().artifacts.final_docx = False

    payload, output_dir = _run_official_material_context(
        tmp_path=tmp_path,
        source=source,
        scene=scene,
        document_type_id="notice",
        material_context=MaterialExecutionContext(
            profile_id="official:notice",
            profile_name="公文资料",
            entity_data={
                "document_type": "notice",
                "title": "关于开展资料归档检查的通知",
                "body": "请各部门按要求完成自查并提交材料。",
                "organization": "示例市档案局",
                "document_no": "示档发〔2026〕1号",
                "issue_date": "2026年7月9日",
            },
        ),
    )

    official_path = output_dir / "official_source_official.docx"
    internal_review_path = (
        output_dir / "official_source_official_internal_review.docx"
    )
    archive_manifest_path = (
        output_dir / "official_source_official_archive_manifest.json"
    )
    archive_manifest_md_path = (
        output_dir / "official_source_official_archive_manifest.md"
    )
    report_json = output_dir / "official_source_changes.json"

    assert payload["status"] == "success"
    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    assert payload["output_path"] == str(official_path)
    assert payload["output_paths"] == {
        "official_docx": str(official_path),
        "internal_review_docx": str(internal_review_path),
        "archive_manifest": str(archive_manifest_path),
        "archive_manifest_md": str(archive_manifest_md_path),
    }
    assert payload["official_document_assembly"]["status"] == "ok"
    assert payload["official_document_assembly"]["profile_id"] == "notice"
    assert official_path.is_file()
    assert internal_review_path.is_file()
    assert archive_manifest_path.is_file()
    assert archive_manifest_md_path.is_file()
    assert report_data["official_document_assembly"]["status"] == "ok"
    assert report_data["official_document_assembly"]["docx_path"] == str(official_path)
    assert report_data["official_document_assembly"]["output_paths"][
        "internal_review_docx"
    ] == str(internal_review_path)


def test_workbench_runner_inserts_scanned_asset_at_placeholder_anchor(tmp_path):
    source = tmp_path / "source.docx"
    logo_path = tmp_path / "logo.png"
    Image.new("RGB", (24, 24), color="red").save(logo_path)

    doc = Document()
    doc.add_paragraph("Logo: {{@img:logo}}")
    doc.save(source)

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["image_insertion"] = True

    collection = scan_asset_collection(tmp_path)
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            asset_items=collection.items,
            image_rules=[AssetInsertionRule(asset_role="logo", target="{{@img:logo}}", width_cm=2.0)],
        ),
    ).run(lambda *_args: None, lambda: False)

    output_doc = Document(payload["output_path"])
    output_para = output_doc.paragraphs[0]

    assert payload["status"] == "success"
    assert "{{@img:logo}}" not in output_para.text
    assert paragraph_has_image(output_para)


def test_workbench_runner_inserts_scanned_asset_at_table_anchor(tmp_path):
    source = tmp_path / "source.docx"
    seal_path = tmp_path / "公章.png"
    Image.new("RGB", (24, 24), color="red").save(seal_path)

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "盖章：{{@img:seal}}"
    doc.save(source)

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["image_insertion"] = True

    collection = scan_asset_collection(tmp_path)
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            asset_items=collection.items,
            image_rules=[AssetInsertionRule(asset_role="seal", target="{{@img:seal}}", width_cm=2.0)],
        ),
    ).run(lambda *_args: None, lambda: False)

    output_doc = Document(payload["output_path"])
    output_para = output_doc.tables[0].cell(0, 0).paragraphs[0]

    assert payload["status"] == "success"
    assert "{{@img:seal}}" not in output_para.text
    assert paragraph_has_image(output_para)


def test_workbench_runner_blocks_missing_required_asset(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Logo: {{@img:logo}}")
    doc.save(source)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(),
        material_context=MaterialExecutionContext(
            image_rules=[AssetInsertionRule(asset_role="logo", target="{{@img:logo}}")]
        ),
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert "logo" in payload["error_text"]


def test_workbench_runner_warns_material_schema_missing_fields_in_reports(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Contract")
    doc.save(source)

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.input_source_profile.failure_policy = "warn"

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={
                "party_a": "甲方",
                "party_b": "乙方",
                "contract_amount": "100000",
            },
        ),
    ).run(lambda *_args: None, lambda: False)

    report_json = source.parent / "output" / "source_changes.json"
    report_md = source.parent / "output" / "source_changes.md"
    report_json_text = report_json.read_text(encoding="utf-8")
    report_data = json.loads(report_json_text)

    assert payload["status"] == "success"
    assert payload["diagnostics_count"] >= 1
    assert "signing_date" in payload["diagnostics_summary"]
    assert payload["material_diagnostics"][0]["missing_field_keys"] == ["signing_date"]
    assert report_data["diagnostics"]["count"] >= 1
    material_items = [
        item
        for item in report_data["diagnostics"]["items"]
        if item["rule_name"] == "material_schema"
    ]
    assert material_items[0]["missing_field_keys"] == ["signing_date"]
    assert "signing_date" in report_md.read_text(encoding="utf-8")


def test_workbench_runner_does_not_block_timeline_warning_under_block_policy(tmp_path):
    source = tmp_path / "timeline_warning.docx"
    doc = Document()
    doc.add_paragraph("Date: {{@time:timeline_1}}")
    doc.save(source)

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True
    scene.input_source_profile.failure_policy = "block"
    plan = default_timeline_plan()
    plan["start_field"] = "timeline_start"
    plan["end_field"] = "timeline_end"
    for index, node in enumerate(plan["nodes"], start=1):
        node["outputs"] = [{"field": f"timeline_{index}", "format": "yyyy-MM-dd"}]

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={
                "timeline_start": "2025-01-01",
                "timeline_end": "2025-01-01",
            },
            timeline_plans={"primary": plan},
        ),
    ).run(lambda *_args: None, lambda: False)

    timeline_diagnostics = [
        diagnostic
        for diagnostic in payload["material_diagnostics"]
        if str(diagnostic.get("change_type") or "").startswith("timeline_")
    ]
    output_text = "\n".join(
        paragraph.text for paragraph in Document(payload["output_path"]).paragraphs
    )

    assert payload["status"] == "success"
    assert "2025-01-01" in output_text
    assert timeline_diagnostics
    assert {diagnostic["level"] for diagnostic in timeline_diagnostics} == {"warning"}


def test_workbench_runner_replaces_generated_chinese_timeline_tokens(tmp_path):
    source = tmp_path / "timeline_segment_tokens.docx"
    doc = Document()
    doc.add_paragraph(
        "{{@time:时间节点1-1}}｜{{@time:时间节点1-2}}｜{{@time:时间节点1-3}}"
    )
    doc.save(source)

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True
    plan = default_timeline_segment(
        1,
        node_count=3,
        start_value="2025-10-01",
        end_value="2025-10-11",
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            timeline_plans={"segment_1": plan},
        ),
    ).run(lambda *_args: None, lambda: False)
    output_text = "\n".join(
        paragraph.text for paragraph in Document(payload["output_path"]).paragraphs
    )

    assert payload["status"] == "success"
    assert output_text == "2025-10-01｜2025-10-06｜2025-10-11"


def test_workbench_runner_payload_includes_material_field_consistency(tmp_path):
    source = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("甲方：旧公司")
    doc.add_paragraph("甲方签约主体：{{@text:party_a}}")
    doc.add_paragraph("乙方：{{@text:party_b}}")
    doc.add_paragraph("合同金额：{{@text:contract_amount}}")
    doc.add_paragraph("签署日期：{{@text:signing_date}}")
    doc.save(source)

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True
    scene.input_source_profile.material_schema_id = "contract_parties_v1"

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={
                "party_a": "甲方公司",
                "party_b": "乙方公司",
                "contract_amount": "100000",
                "signing_date": "2026-06-16",
            },
        ),
    ).run(lambda *_args: None, lambda: False)

    consistency = payload["material_field_consistency"]
    issue_kinds = {issue["kind"] for issue in consistency["issues"]}

    assert payload["status"] == "success"
    assert consistency["schema_id"] == "contract_parties_v1"
    assert consistency["family_id"] == "contract_delivery"
    assert consistency["status"] == "warning"
    assert consistency["issue_count"] == 1
    assert "label_value_conflict" in issue_kinds


def test_workbench_runner_writes_material_manifest_for_attachment_inventory(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Qualification package")
    doc.save(source)
    certificate_path = tmp_path / "certificate.pdf"
    certificate_path.write_bytes(b"%PDF-1.4\n%certificate\n")

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.input_source_profile.material_schema_id = "qualification_archive_assets_v1"
    scene.input_source_profile.failure_policy = "warn"
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.default_delivery_preset().artifacts.material_package = True

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            profile_name="资质包",
            entity_data={
                "organization": "测试公司",
                "package_name": "投标资质包",
            },
            asset_items=[
                AssetItem(
                    item_id="certificate",
                    label="资质证书",
                    role="certificate",
                    path=str(certificate_path),
                    mime_type="application/pdf",
                )
            ],
        ),
    ).run(lambda *_args: None, lambda: False)

    manifest_path = Path(payload["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    certificate_role = next(
        item for item in manifest["asset_roles"] if item["role"] == "certificate"
    )
    business_license_role = next(
        item for item in manifest["asset_roles"] if item["role"] == "business_license"
    )

    assert payload["status"] == "success"
    assert manifest_path.name == "source_material_manifest.json"
    assert manifest["kind"] == "material_attachment_manifest"
    assert manifest["material_schema"]["schema_id"] == "qualification_archive_assets_v1"
    assert {
        item["role"]: item["package_subdir"]
        for item in manifest["material_schema"]["archive_directory_rules"]
    }["certificate"] == "assets/01_certificates"
    assert manifest["summary"]["filled_field_count"] == 2
    assert certificate_role["present"] is True
    assert certificate_role["archive_dir"] == "01_certificates"
    assert certificate_role["package_subdir"] == "assets/01_certificates"
    assert certificate_role["items"][0]["kind"] == "attachment"
    assert certificate_role["items"][0]["exists"] is True
    assert certificate_role["items"][0]["package_subdir"] == "assets/01_certificates"
    assert business_license_role["missing"] is True
    assert business_license_role["package_subdir"] == "assets/02_business_license"
    assert manifest["missing"]["asset_roles"] == ["business_license"]
    assert manifest["delivery"]["output_paths"]
    package_paths = payload["material_package_paths"]
    package_dir = Path(package_paths["directory"])
    package_zip = Path(package_paths["zip"])
    package_report = Path(package_paths["report"])
    package_manifest = json.loads(Path(package_paths["package_manifest"]).read_text(encoding="utf-8"))

    assert package_dir.is_dir()
    assert package_zip.exists()
    assert package_manifest["kind"] == "material_delivery_package"
    assert package_manifest["status"] == "incomplete"
    assert payload["material_package_receipt"]["status"] == "incomplete"
    assert any(
        item["path"] == "assets/01_certificates/certificate.pdf"
        for item in package_manifest["files"]
    )
    assert "business_license" in package_report.read_text(encoding="utf-8")
    with zipfile.ZipFile(package_zip, "r") as archive:
        names = set(archive.namelist())
    assert "manifest/material_manifest.json" in names
    assert "assets/01_certificates/certificate.pdf" in names
    assert "archive_report.md" in names
    assert "package_manifest.json" in names


def test_workbench_runner_writes_project_application_attachment_inventory(tmp_path):
    source = tmp_path / "application.docx"
    doc = Document()
    doc.add_paragraph("Project: {{@text:project_name}}")
    doc.save(source)
    application_form_path = tmp_path / "application_form.pdf"
    application_form_path.write_bytes(b"%PDF-1.4\n%application\n")

    scene = SceneWorkspace(scene_id="project_application", category="project_application")
    apply_planned_scene_family_defaults(scene)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            profile_name="Project application",
            entity_data={
                "project_name": "Project A",
                "applicant_unit": "Example Unit",
                "principal_investigator": "Jane Smith",
            },
            asset_items=[
                AssetItem(
                    item_id="application_form",
                    label="Application form",
                    role="application_form",
                    path=str(application_form_path),
                    mime_type="application/pdf",
                )
            ],
        ),
    ).run(lambda *_args: None, lambda: False)

    manifest_path = Path(payload["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    application_form_role = next(
        item for item in manifest["asset_roles"] if item["role"] == "application_form"
    )
    budget_sheet_role = next(
        item for item in manifest["asset_roles"] if item["role"] == "budget_sheet"
    )

    assert payload["status"] == "success"
    assert manifest["material_schema"]["schema_id"] == "project_application_materials_v1"
    assert manifest["material_schema"]["family"] == "project_application"
    assert manifest["summary"]["filled_field_count"] == 3
    assert application_form_role["present"] is True
    assert application_form_role["items"][0]["kind"] == "attachment"
    assert budget_sheet_role["missing"] is True
    assert manifest["missing"]["asset_roles"] == ["budget_sheet"]
    assert payload["material_package_paths"]["package_manifest"].endswith(
        "package_manifest.json"
    )


def test_workbench_runner_does_not_write_material_artifacts_without_output_switches(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Qualification package")
    doc.save(source)
    certificate_path = tmp_path / "certificate.pdf"
    certificate_path.write_bytes(b"%PDF-1.4\n%certificate\n")

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.input_source_profile.material_schema_id = "qualification_archive_assets_v1"
    scene.input_source_profile.failure_policy = "warn"

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            profile_name="资质包",
            entity_data={
                "organization": "测试公司",
                "package_name": "投标资质包",
            },
            asset_items=[
                AssetItem(
                    item_id="certificate",
                    label="资质证书",
                    role="certificate",
                    path=str(certificate_path),
                    mime_type="application/pdf",
                )
            ],
        ),
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert payload["material_manifest_paths"] == {}
    assert payload["material_package_paths"] == {}
    assert not (source.parent / "output" / "source_material_manifest.json").exists()
    assert not (source.parent / "output" / "source_material_package.zip").exists()


def test_workbench_runner_uses_delivery_preset_material_artifacts(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Qualification package")
    doc.save(source)
    certificate_path = tmp_path / "certificate.pdf"
    certificate_path.write_bytes(b"%PDF-1.4\n%certificate\n")

    scene = SceneWorkspace(
        default_delivery_preset_id="package_report",
        delivery_presets=[
            DeliveryPreset(
                preset_id="package_report",
                label="Package report",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(
                    final_docx=False,
                    report_json=False,
                    report_markdown=False,
                    material_manifest=True,
                    material_package=True,
                ),
            )
        ],
    )
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.input_source_profile.material_schema_id = "qualification_archive_assets_v1"
    scene.input_source_profile.failure_policy = "warn"

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            profile_name="资质包",
            entity_data={
                "organization": "测试公司",
                "package_name": "投标资质包",
            },
            asset_items=[
                AssetItem(
                    item_id="certificate",
                    label="资质证书",
                    role="certificate",
                    path=str(certificate_path),
                    mime_type="application/pdf",
                )
            ],
        ),
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert payload["output_paths"] == {}
    assert Path(payload["material_manifest_paths"]["material"]).exists()
    assert Path(payload["material_package_paths"]["zip"]).exists()


def test_contract_family_signing_copy_writes_material_package_with_missing_seal(tmp_path):
    source = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("甲方：甲方公司")
    doc.add_paragraph("乙方：乙方公司")
    doc.add_paragraph("合同金额：100000")
    doc.add_paragraph("签署日期：2026-06-16")
    doc.save(source)

    scene = SceneWorkspace(
        scene_id="contract_delivery",
        category="contract_delivery",
        template_id="default",
    )
    result = apply_planned_scene_family_defaults(scene)
    scene.module_switches = {name: False for name in scene.module_switches}

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={
                "party_a": "甲方公司",
                "party_b": "乙方公司",
                "contract_amount": "100000",
                "signing_date": "2026-06-16",
            },
        ),
    ).run(lambda *_args: None, lambda: False)

    manifest_path = Path(payload["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    seal_role = next(item for item in manifest["asset_roles"] if item["role"] == "seal")
    legal_signature_role = next(
        item for item in manifest["asset_roles"] if item["role"] == "legal_signature"
    )
    package_paths = payload["material_package_paths"]
    package_manifest_path = Path(package_paths["package_manifest"])
    package_manifest = json.loads(package_manifest_path.read_text(encoding="utf-8"))

    assert result.applied is True
    assert "signing_copy" in payload["output_paths"]
    assert payload["status"] == "success"
    assert payload["material_diagnostics"][0]["schema_ids"] == [
        "contract_parties_v1",
        "signature_assets_v1",
    ]
    assert manifest["material_schema"]["schema_id"] == "contract_parties_v1"
    assert manifest["material_schema"]["schema_ids"] == [
        "contract_parties_v1",
        "signature_assets_v1",
    ]
    assert "Signature and seal assets" in manifest["material_schema"]["labels"]
    assert "seal" in manifest["material_schema"]["required_asset_roles"]
    assert seal_role["required"] is True
    assert seal_role["present"] is False
    assert seal_role["missing"] is True
    assert legal_signature_role["required"] is False
    assert legal_signature_role["missing"] is False
    assert manifest["missing"]["asset_roles"] == ["seal"]
    assert Path(package_paths["directory"]).is_dir()
    assert Path(package_paths["zip"]).exists()
    assert package_manifest["status"] == "incomplete"
    assert package_manifest["missing"]["asset_roles"] == ["seal"]
    assert "seal" in Path(package_paths["report"]).read_text(encoding="utf-8")
    with zipfile.ZipFile(package_paths["zip"], "r") as archive:
        names = set(archive.namelist())
    assert "manifest/material_manifest.json" in names
    assert "package_manifest.json" in names
    assert "archive_report.md" in names


def test_contract_family_signing_copy_places_seal_and_packages_manifest(tmp_path):
    source = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("甲方：{{@text:party_a}}")
    doc.add_paragraph("乙方：{{@text:party_b}}")
    doc.add_paragraph("合同金额：{{@text:contract_amount}}")
    doc.add_paragraph("签署日期：{{@text:signing_date}}")
    doc.add_paragraph("盖章：{{@img:seal}}")
    doc.save(source)

    seal_path = tmp_path / "seal.png"
    Image.new("RGB", (320, 320), color="blue").save(seal_path)

    scene = SceneWorkspace(
        scene_id="contract_delivery",
        category="contract_delivery",
        template_id="default",
    )
    apply_planned_scene_family_defaults(scene)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={
                "party_a": "甲方公司",
                "party_b": "乙方公司",
                "contract_amount": "100000",
                "signing_date": "2026-06-16",
            },
            asset_items=[
                AssetItem(
                    role="seal",
                    label="公司公章",
                    path=str(seal_path),
                    width_cm=2.0,
                )
            ],
            image_rules=[
                AssetInsertionRule(
                    asset_role="seal",
                    target="{{@img:seal}}",
                    width_cm=2.0,
                    required=True,
                )
            ],
        ),
    ).run(lambda *_args: None, lambda: False)

    output_doc = Document(payload["output_paths"]["signing_copy"])
    output_text = "\n".join(paragraph.text for paragraph in output_doc.paragraphs)
    manifest_path = Path(payload["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    package_paths = payload["material_package_paths"]
    package_manifest = json.loads(
        Path(package_paths["package_manifest"]).read_text(encoding="utf-8")
    )
    seal_role = next(item for item in manifest["asset_roles"] if item["role"] == "seal")

    assert scene.module_switches["image_insertion"] is True
    assert payload["status"] == "success"
    assert "{{@img:seal}}" not in output_text
    assert output_doc.inline_shapes
    assert paragraph_has_image(output_doc.paragraphs[-1])
    assert seal_role["present"] is True
    assert seal_role["missing"] is False
    assert manifest["missing"]["asset_roles"] == []
    assert manifest["asset_items"][0]["role"] == "seal"
    assert package_manifest["status"] == "complete"
    with zipfile.ZipFile(package_paths["zip"], "r") as archive:
        names = set(archive.namelist())
    assert any(name.startswith("assets/seal/") and name.endswith("seal.png") for name in names)


def test_workbench_runner_injects_question_figure_asset_into_exam_runtime(tmp_path):
    source = tmp_path / "exam.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    figure_path = tmp_path / "question_figure.png"
    Image.new("RGB", (320, 240), color="green").save(figure_path)

    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.structured_formats = ["json"]
    scene.default_delivery_preset().artifacts.report_json = True
    scene.default_delivery_preset().artifacts.report_markdown = True
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.default_delivery_preset().artifacts.material_package = True

    payload = {
        "paper_title": "期末测试",
        "total_score": 5,
        "sections": [
            {
                "title": "选择题",
                "type": "single_choice",
                "questions": [
                    {
                        "stem": "1 + 1 = ?",
                        "options": ["A.1", "B.2"],
                        "answer": "B",
                        "score": 5,
                    }
                ],
            }
        ],
    }
    result = _run_exam_material_context(
        tmp_path=tmp_path,
        source=source,
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={"exam_items": json.dumps(payload, ensure_ascii=False)},
            asset_items=[
                AssetItem(
                    role="question_figure",
                    label="题目图片",
                    path=str(figure_path),
                    metadata={"alt_text": "加法示意图"},
                )
            ],
        ),
    )

    report_json = next(Path(path) for path in result["report_paths"] if str(path).endswith(".json"))
    report_md = next(Path(path) for path in result["report_paths"] if str(path).endswith(".md"))
    report_json_text = report_json.read_text(encoding="utf-8")
    report_data = json.loads(report_json_text)
    version = report_data["exam_delivery_runtime"]["rendered_versions"][0]
    version_doc = Document(version["docx_path"])

    assert result["status"] == "success"
    assert version["question_asset_count"] == 1
    assert version["rendered_question_asset_count"] == 1
    assert version["missing_question_asset_count"] == 0
    assert version["rendered_question_asset_alt_text_count"] == 1
    assert version["missing_question_asset_alt_text_count"] == 0
    assert len(version_doc.inline_shapes) == 1
    assert _first_doc_pr_attrs(version["docx_path"]).get("descr") == "加法示意图"
    markdown = report_md.read_text(encoding="utf-8")
    assert "assets=1/1" in markdown
    assert "altText=1/1" in markdown
    manifest_path = next(Path(path) for path in result["material_manifest_paths"].values())
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["asset_items"][0]["metadata"] == {"alt_text": "加法示意图"}


def test_workbench_runner_treats_remote_question_figure_url_as_missing_local_file(tmp_path):
    source = tmp_path / "exam.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.structured_formats = ["json"]
    scene.default_delivery_preset().artifacts.report_json = True
    scene.default_delivery_preset().artifacts.report_markdown = True
    scene.default_delivery_preset().artifacts.material_manifest = True

    payload = {
        "paper_title": "期末测试",
        "total_score": 5,
        "sections": [
            {
                "title": "选择题",
                "type": "single_choice",
                "questions": [
                    {
                        "stem": "1 + 1 = ?",
                        "options": ["A.1", "B.2"],
                        "answer": "B",
                        "score": 5,
                    }
                ],
            }
        ],
    }
    result = _run_exam_material_context(
        tmp_path=tmp_path,
        source=source,
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={"exam_items": json.dumps(payload, ensure_ascii=False)},
            asset_items=[
                AssetItem(
                    item_id="fig_q1",
                    role="question_figure",
                    label="第一题图片",
                    path="https://example.invalid/question_1.png",
                    metadata={"question_index": "1", "alt_text": "远程图不会下载"},
                )
            ],
        ),
    )

    manifest_path = next(Path(path) for path in result["material_manifest_paths"].values())
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    diagnostic = next(
        item
        for item in result["material_diagnostics"]
        if item["change_type"] == "preflight_missing_question_figure_file"
    )
    diagnostic_text = json.dumps(result["material_diagnostics"], ensure_ascii=False)

    assert result["status"] == "success"
    assert diagnostic["before"] == "https://example.invalid/question_1.png"
    assert not (tmp_path / "out" / "_remote_asset_cache").exists()
    assert "remote_asset_cache" not in manifest
    assert "remote_asset_items" not in diagnostic_text
    assert "download_url" not in diagnostic_text

def test_workbench_runner_maps_multiple_question_figure_assets_by_question_metadata(tmp_path):
    source = tmp_path / "exam.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    figure_one = tmp_path / "question_one.png"
    figure_two = tmp_path / "question_two.png"
    Image.new("RGB", (320, 240), color="green").save(figure_one)
    Image.new("RGB", (320, 240), color="blue").save(figure_two)

    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.structured_formats = ["json"]
    scene.default_delivery_preset().artifacts.report_json = True
    scene.default_delivery_preset().artifacts.report_markdown = True
    scene.default_delivery_preset().artifacts.material_manifest = True

    payload = {
        "paper_title": "期末测试",
        "total_score": 10,
        "sections": [
            {
                "title": "选择题",
                "type": "single_choice",
                "questions": [
                    {
                        "id": "q1",
                        "stem": "1 + 1 = ?",
                        "options": ["A.1", "B.2"],
                        "answer": "B",
                        "score": 5,
                    },
                    {
                        "id": "q2",
                        "stem": "3 - 1 = ?",
                        "options": ["A.1", "B.2"],
                        "answer": "B",
                        "score": 5,
                    },
                ],
            }
        ],
    }
    result = _run_exam_material_context(
        tmp_path=tmp_path,
        source=source,
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={"exam_items": json.dumps(payload, ensure_ascii=False)},
            asset_items=[
                AssetItem(
                    item_id="fig_q2",
                    role="question_figure",
                    label="第二题图片",
                    path=str(figure_two),
                    metadata={"question_id": "q2", "alt_text": "减法示意图"},
                ),
                AssetItem(
                    item_id="fig_q1",
                    role="question_figure",
                    label="第一题图片",
                    path=str(figure_one),
                    metadata={"question_index": "1", "alt_text": "加法示意图"},
                ),
            ],
        ),
    )

    report_json = next(Path(path) for path in result["report_paths"] if str(path).endswith(".json"))
    report_md = next(Path(path) for path in result["report_paths"] if str(path).endswith(".md"))
    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    version = report_data["exam_delivery_runtime"]["rendered_versions"][0]
    version_doc = Document(version["docx_path"])

    assert result["status"] == "success"
    assert version["question_asset_count"] == 2
    assert version["rendered_question_asset_count"] == 2
    assert version["missing_question_asset_count"] == 0
    assert version["rendered_question_asset_alt_text_count"] == 2
    assert version["missing_question_asset_alt_text_count"] == 0
    assert len(version_doc.inline_shapes) == 2
    assert [attrs.get("descr") for attrs in _doc_pr_attrs(version["docx_path"])] == [
        "加法示意图",
        "减法示意图",
    ]
    markdown = report_md.read_text(encoding="utf-8")
    assert "assets=2/2" in markdown
    assert "altText=2/2" in markdown
    manifest_path = next(Path(path) for path in result["material_manifest_paths"].values())
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert [item["metadata"] for item in manifest["asset_items"]] == [
        {"question_id": "q2", "alt_text": "减法示意图"},
        {"question_index": "1", "alt_text": "加法示意图"},
    ]


def test_workbench_runner_appends_same_question_figure_group_by_order(tmp_path):
    source = tmp_path / "exam.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    figure_one = tmp_path / "question_1_figure_1.png"
    figure_two = tmp_path / "question_1_figure_2.png"
    Image.new("RGB", (320, 240), color="green").save(figure_one)
    Image.new("RGB", (320, 240), color="blue").save(figure_two)

    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.structured_formats = ["json"]
    scene.default_delivery_preset().artifacts.report_json = True
    scene.default_delivery_preset().artifacts.report_markdown = True
    scene.default_delivery_preset().artifacts.material_manifest = True

    payload = {
        "paper_title": "期末测试",
        "total_score": 5,
        "sections": [
            {
                "title": "综合题",
                "type": "open",
                "questions": [
                    {
                        "id": "q1",
                        "stem": "观察两张图，回答问题。",
                        "answer": "略",
                        "score": 5,
                    }
                ],
            }
        ],
    }
    result = _run_exam_material_context(
        tmp_path=tmp_path,
        source=source,
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={"exam_items": json.dumps(payload, ensure_ascii=False)},
            asset_items=[
                AssetItem(
                    item_id="fig_q1_2",
                    role="question_figure",
                    label="第一题第二图",
                    path=str(figure_two),
                    metadata={
                        "question_id": "q1",
                        "figure_order": "2",
                        "alt_text": "第二张示意图",
                    },
                ),
                AssetItem(
                    item_id="fig_q1_1",
                    role="question_figure",
                    label="第一题第一图",
                    path=str(figure_one),
                    metadata={
                        "question_id": "q1",
                        "figure_order": "1",
                        "alt_text": "第一张示意图",
                    },
                ),
            ],
        ),
    )

    report_json = next(Path(path) for path in result["report_paths"] if str(path).endswith(".json"))
    report_md = next(Path(path) for path in result["report_paths"] if str(path).endswith(".md"))
    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    version = report_data["exam_delivery_runtime"]["rendered_versions"][0]
    version_doc = Document(version["docx_path"])

    assert result["status"] == "success"
    assert version["question_asset_count"] == 2
    assert version["rendered_question_asset_count"] == 2
    assert version["missing_question_asset_count"] == 0
    assert version["rendered_question_asset_alt_text_count"] == 2
    assert len(version_doc.inline_shapes) == 2
    assert [attrs.get("descr") for attrs in _doc_pr_attrs(version["docx_path"])] == [
        "第一张示意图",
        "第二张示意图",
    ]
    markdown = report_md.read_text(encoding="utf-8")
    assert "assets=2/2" in markdown
    assert "altText=2/2" in markdown


def _first_doc_pr_attrs(docx_path: str) -> dict[str, str]:
    attrs = _doc_pr_attrs(docx_path)
    assert attrs
    return attrs[0]


def _doc_pr_attrs(docx_path: str) -> list[dict[str, str]]:
    with zipfile.ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    ns = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
    return [dict(doc_pr.attrib) for doc_pr in root.findall(".//wp:docPr", ns)]


def test_workbench_runner_blocks_material_schema_missing_fields_with_preflight_report(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Contract")
    doc.save(source)

    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.input_source_profile.failure_policy = "block"
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.default_delivery_preset().artifacts.material_package = True

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={
                "party_a": "甲方",
                "party_b": "乙方",
                "contract_amount": "100000",
            },
        ),
    ).run(lambda *_args: None, lambda: False)

    report_json = source.parent / "output" / "source_changes.json"
    report_md = source.parent / "output" / "source_changes.md"
    report_data = json.loads(report_json.read_text(encoding="utf-8"))

    assert payload["status"] == "failed"
    assert payload["output_path"] == ""
    assert "signing_date" in payload["error_text"]
    assert payload["diagnostics_count"] == 1
    assert payload["report_paths"] == [str(report_json), str(report_md)]
    assert report_data["status"] == "failed"
    assert report_data["diagnostics"]["items"][0]["change_type"] == "preflight_missing_material_fields"
    assert report_data["diagnostics"]["items"][0]["missing_field_keys"] == ["signing_date"]
    assert "signing_date" in report_md.read_text(encoding="utf-8")
    manifest_path = Path(payload["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["missing"]["field_keys"] == ["signing_date"]
    assert manifest["diagnostics"][0]["change_type"] == "preflight_missing_material_fields"
    package_paths = payload["material_package_paths"]
    package_manifest = json.loads(Path(package_paths["package_manifest"]).read_text(encoding="utf-8"))
    assert Path(package_paths["zip"]).exists()
    assert package_manifest["status"] == "incomplete"
    assert package_manifest["missing"]["field_keys"] == ["signing_date"]
    assert "signing_date" in Path(package_paths["report"]).read_text(encoding="utf-8")


def test_blocked_preflight_reports_package_writer_failure_without_masking_reason(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "contract_parties_v1"
    scene.input_source_profile.failure_policy = "block"
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.default_delivery_preset().artifacts.material_package = True

    def _raise_package_failure(**_kwargs):
        raise OSError("package disk unavailable")

    monkeypatch.setattr(
        "src.services.production_runtime.material_preflight_reporting.write_material_package_artifacts",
        _raise_package_failure,
    )
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={
                "party_a": "Party A",
                "party_b": "Party B",
                "contract_amount": "100000",
            },
        ),
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert "signing_date" in payload["error_text"]
    assert "package disk unavailable" in payload["error_text"]
    assert payload["material_package_paths"] == {}
    assert payload["artifact_failure_count"] == 1
    assert payload["artifact_failures"][0]["kind"] == "material_package"


def test_unknown_material_schema_blocks_even_when_missing_material_policy_is_warn(
    tmp_path,
):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Unknown schema")
    doc.save(source)

    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "missing_schema_v1"
    scene.input_source_profile.failure_policy = "warn"
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.default_delivery_preset().artifacts.material_package = True

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(),
    ).run(lambda *_args: None, lambda: False)

    report_json = source.parent / "output" / "source_changes.json"
    report_md = source.parent / "output" / "source_changes.md"
    report_data = json.loads(report_json.read_text(encoding="utf-8"))

    assert payload["status"] == "failed"
    assert payload["output_path"] == ""
    assert "missing_schema_v1" in payload["error_text"]
    assert payload["diagnostics_count"] == 1
    assert payload["material_diagnostics"][0]["change_type"] == "preflight_unknown_material_schema"
    assert payload["material_diagnostics"][0]["missing_schema_ids"] == ["missing_schema_v1"]
    assert report_data["diagnostics"]["items"][0]["change_type"] == "preflight_unknown_material_schema"
    assert report_data["diagnostics"]["items"][0]["missing_schema_ids"] == ["missing_schema_v1"]
    assert "missing_schema_v1" in report_md.read_text(encoding="utf-8")

    manifest_path = Path(payload["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["material_schema"]["schema_ids"] == ["missing_schema_v1"]
    assert manifest["missing"]["schema_ids"] == ["missing_schema_v1"]
    assert manifest["summary"]["missing_schema_count"] == 1
    assert manifest["diagnostics"][0]["change_type"] == "preflight_unknown_material_schema"

    package_paths = payload["material_package_paths"]
    package_manifest = json.loads(Path(package_paths["package_manifest"]).read_text(encoding="utf-8"))
    assert Path(package_paths["zip"]).exists()
    assert package_manifest["status"] == "incomplete"
    assert package_manifest["missing"]["schema_ids"] == ["missing_schema_v1"]
    assert "missing_schema_v1" in Path(package_paths["report"]).read_text(encoding="utf-8")


def test_workbench_batch_runner_outputs_one_document_per_profile(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("公司：{{@text:company_name}}")
    doc.save(source)

    archive = EntityArchive(
        archive_id="archive",
        archive_name="测试档案",
        profiles=[
            EntityProfile(
                profile_id="a",
                profile_name="主体 A",
                fields={"company_name": "测试公司A"},
            ),
            EntityProfile(
                profile_id="b",
                profile_name="主体 B",
                fields={"company_name": "测试公司B"},
            ),
        ],
    )
    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "batch_output",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert len(payload["items"]) == 2
    assert len(payload["output_paths"]) == 2
    output_texts = []
    assert set(payload["output_paths"]) == {"a:final", "b:final"}
    for output_path in payload["output_paths"].values():
        output_doc = Document(output_path)
        output_texts.append("\n".join(paragraph.text for paragraph in output_doc.paragraphs))

    assert any("测试公司A" in text for text in output_texts)
    assert any("测试公司B" in text for text in output_texts)
    assert any(path.endswith("_batch_report.json") for path in payload["report_paths"])
    assert any(path.endswith("_batch_report.md") for path in payload["report_paths"])


def test_workbench_batch_runner_blocks_missing_assets(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Project: {{@text:project}}")
    doc.save(source)

    archive = EntityArchive(
        archive_id="archive",
        profiles=[
            EntityProfile(
                profile_id="a",
                profile_name="Entity A",
                fields={"company_name": "Company A"},
            )
        ],
    )
    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}

    failed_payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "missing_output",
        base_context=MaterialExecutionContext(
            image_rules=[AssetInsertionRule(asset_role="logo", target="{{@img:logo}}")]
        ),
    ).run(lambda *_args: None, lambda: False)

    assert failed_payload["status"] == "failed"
    assert "logo" in failed_payload["error_text"]
    assert failed_payload["batch_isolation"]["total_count"] == 1
    assert failed_payload["batch_isolation"]["failed_count"] == 1
    assert failed_payload["batch_isolation"]["failed_profiles"][0]["profile_id"] == "a"
    assert failed_payload["batch_isolation"]["failed_profiles"][0]["missing_asset_roles"] == [
        "logo"
    ]
    assert failed_payload["batch_issue_items"][0]["profile_id"] == "a"
    assert failed_payload["batch_issue_items"][0]["repair_target_type"] == "asset"
    assert failed_payload["batch_issue_items"][0]["repair_target_key"] == "logo"
    assert any(path.endswith("_batch_report.json") for path in failed_payload["report_paths"])


def test_workbench_batch_runner_isolates_missing_assets_per_profile(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Company: {{@text:company_name}}")
    doc.save(source)
    logo_path = tmp_path / "logo.png"
    logo_path.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
        b"\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01"
        b"\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    valid_profile_ids = [f"valid-{index:02d}" for index in range(1, 10)]
    archive = EntityArchive(
        archive_id="bid",
        profiles=[
            EntityProfile(
                profile_id=profile_id,
                profile_name=f"Company {index:02d}",
                fields={"company_name": f"Company {index:02d}"},
                asset_paths={"logo": str(logo_path)},
            )
            for index, profile_id in enumerate(valid_profile_ids, start=1)
        ]
        + [
            EntityProfile(
                profile_id="missing",
                profile_name="Company 10",
                fields={"company_name": "Company 10"},
            ),
        ],
    )
    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "batch_output",
        base_context=MaterialExecutionContext(
            image_rules=[AssetInsertionRule(asset_role="logo", target="{{@img:logo}}")]
        ),
    ).run(lambda *_args: None, lambda: False)

    batch_json = next(
        Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.json")
    )
    batch_md = next(
        Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.md")
    )
    report_data = json.loads(batch_json.read_text(encoding="utf-8"))

    assert payload["status"] == "partial_success"
    assert payload["batch_isolation"]["total_count"] == 10
    assert payload["batch_isolation"]["success_count"] == 9
    assert payload["batch_isolation"]["failed_count"] == 1
    assert {
        item["profile_id"]
        for item in payload["batch_isolation"]["successful_profiles"]
    } == set(valid_profile_ids)
    assert payload["batch_isolation"]["failed_profiles"][0]["profile_id"] == "missing"
    assert payload["batch_isolation"]["failed_profiles"][0]["missing_asset_roles"] == [
        "logo"
    ]
    assert payload["batch_issue_items"][0]["profile_id"] == "missing"
    assert payload["batch_issue_items"][0]["repair_target_type"] == "asset"
    assert set(payload["output_paths"]) == {
        f"{profile_id}:final" for profile_id in valid_profile_ids
    }
    for index, profile_id in enumerate(valid_profile_ids, start=1):
        assert f"Company {index:02d}" in "\n".join(
            paragraph.text
            for paragraph in Document(
                payload["output_paths"][f"{profile_id}:final"]
            ).paragraphs
        )
    assert report_data["batch_isolation"]["failed_profiles"][0]["profile_id"] == "missing"
    batch_markdown = batch_md.read_text(encoding="utf-8")
    assert "## Batch Isolation" in batch_markdown
    assert "success: 9" in batch_markdown
    assert "failed: 1" in batch_markdown
    assert "logo" in batch_markdown


def test_workbench_batch_runner_aggregates_material_schema_diagnostics(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Employee: {{@text:employee_name}}")
    doc.save(source)

    archive = EntityArchive(
        archive_id="hr",
        profiles=[
            EntityProfile(
                profile_id="ok",
                profile_name="完整员工",
                fields={"employee_name": "张三", "employee_id": "E001"},
            ),
            EntityProfile(
                profile_id="missing",
                profile_name="缺编号员工",
                fields={"employee_name": "李四"},
            ),
        ],
    )
    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True
    scene.input_source_profile.material_schema_id = "personnel_records_v1"
    scene.input_source_profile.failure_policy = "block"
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.default_delivery_preset().artifacts.material_package = True

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "batch_output",
    ).run(lambda *_args: None, lambda: False)

    batch_json = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.json"))
    batch_md = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.md"))
    report_data = json.loads(batch_json.read_text(encoding="utf-8"))

    assert payload["status"] == "partial_success"
    assert len(payload["material_diagnostics"]) == 1
    assert len(payload["material_manifest_paths"]) == 2
    assert len(payload["material_package_paths"]) == 8
    assert payload["material_package_receipts"]["ok"]["status"] == "complete"
    assert payload["material_package_receipts"]["missing"]["status"] == "incomplete"
    assert payload["material_diagnostics"][0]["missing_field_keys"] == ["employee_id"]
    assert payload["batch_issue_items"][0]["profile_id"] == "missing"
    assert payload["batch_issue_items"][0]["profile_name"] == "缺编号员工"
    assert payload["batch_issue_items"][0]["kind"] == "preflight_missing_material_fields"
    assert payload["batch_issue_items"][0]["repair_target_type"] == "field"
    assert payload["batch_issue_items"][0]["repair_target_key"] == "employee_id"
    assert report_data["material_manifest_paths"] == payload["material_manifest_paths"]
    assert report_data["material_package_paths"] == payload["material_package_paths"]
    assert report_data["material_package_receipts"] == payload["material_package_receipts"]
    assert report_data["material_diagnostics"][0]["missing_field_keys"] == ["employee_id"]
    assert report_data["batch_issue_items"][0]["profile_id"] == "missing"
    batch_markdown = batch_md.read_text(encoding="utf-8")
    assert "## Batch Issue Center" in batch_markdown
    assert "employee_id" in batch_markdown


def test_workbench_batch_runner_reports_missing_question_figure_file_repair_target(tmp_path):
    source = tmp_path / "exam_source.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    missing_question_two = tmp_path / "missing_question_2.png"
    archive = EntityArchive(
        archive_id="exam",
        profiles=[
            EntityProfile(
                profile_id="exam_a",
                profile_name="Exam A",
                fields={
                    "paper_title": "Final Exam",
                    "subject": "Math",
                    "grade": "Grade 7",
                    "duration": "90",
                    "total_score": "100",
                },
                asset_items=[
                    {
                        "item_id": "question_figure_2",
                        "label": "Question 2 figure",
                        "role": "question_figure",
                        "path": str(missing_question_two),
                        "metadata": {
                            "question_index": "2",
                            "alt_text": "Subtraction diagram",
                        },
                    }
                ],
            )
        ],
    )
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.failure_policy = "warn"

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "batch_output",
    ).run(lambda *_args: None, lambda: False)

    batch_json = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.json"))
    report_data = json.loads(batch_json.read_text(encoding="utf-8"))
    issue = next(
        item
        for item in payload["batch_issue_items"]
        if item["kind"] == "preflight_missing_question_figure_file"
    )
    target = json.loads(issue["repair_target_key"])

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert "Terminal assembler 'exam' blocked delivery" in payload["error_text"]
    assert payload["material_diagnostics"][0]["change_type"] == (
        "preflight_missing_question_figure_file"
    )
    assert issue["profile_id"] == "exam_a"
    assert issue["repair_target_type"] == "question_figure_item"
    assert target["item_id"] == "question_figure_2"
    assert target["question_index"] == "2"
    assert target["path"] == str(missing_question_two)
    assert report_data["batch_issue_items"][0]["repair_target_type"] == (
        "question_figure_item"
    )


def test_workbench_batch_runner_reports_missing_local_metadata_question_figure_repair_target(tmp_path):
    source = tmp_path / "exam_source.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    missing_local_path = tmp_path / "missing_question_2.png"
    archive = EntityArchive(
        archive_id="exam",
        profiles=[
            EntityProfile(
                profile_id="exam_a",
                profile_name="Exam A",
                fields={
                    "paper_title": "Final Exam",
                    "subject": "Math",
                    "grade": "Grade 7",
                    "duration": "90",
                    "total_score": "100",
                },
                asset_items=[
                    {
                        "item_id": "question_figure_2",
                        "label": "Question 2 figure",
                        "role": "question_figure",
                        "path": "",
                        "metadata": {
                            "question_index": "2",
                            "asset_id": "library-question-2",
                            "source": "local_library",
                            "local_path": str(missing_local_path),
                            "alt_text": "Subtraction diagram",
                        },
                    }
                ],
            )
        ],
    )
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.failure_policy = "warn"

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "batch_output",
    ).run(lambda *_args: None, lambda: False)

    batch_json = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.json"))
    report_data = json.loads(batch_json.read_text(encoding="utf-8"))
    issue = next(
        item
        for item in payload["batch_issue_items"]
        if item["kind"] == "preflight_missing_question_figure_file"
    )
    target = json.loads(issue["repair_target_key"])

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert "Terminal assembler 'exam' blocked delivery" in payload["error_text"]
    assert payload["material_diagnostics"][0]["before"] == str(missing_local_path)
    assert payload["material_diagnostics"][0]["missing_asset_items"][0]["path"] == (
        str(missing_local_path)
    )
    assert issue["profile_id"] == "exam_a"
    assert issue["repair_target_type"] == "question_figure_item"
    assert issue["missing_asset_items"][0]["metadata"]["local_path"] == str(missing_local_path)
    assert target["item_id"] == "question_figure_2"
    assert target["question_index"] == "2"
    assert target["path"] == str(missing_local_path)
    assert target["cache_path"] == str(missing_local_path)
    assert target["asset_id"] == "library-question-2"
    assert target["metadata"]["local_path"] == str(missing_local_path)
    assert report_data["batch_issue_items"][0]["repair_target_key"] == issue[
        "repair_target_key"
    ]

def test_workbench_batch_runner_reports_suspicious_question_figure_filename_mismatch(tmp_path):
    source = tmp_path / "exam_source.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    figure_path = tmp_path / "question_3.png"
    Image.new("RGB", (320, 240), color="green").save(figure_path)
    archive = EntityArchive(
        archive_id="exam",
        profiles=[
            EntityProfile(
                profile_id="exam_a",
                profile_name="Exam A",
                fields={
                    "paper_title": "Final Exam",
                    "subject": "Math",
                    "grade": "Grade 7",
                    "duration": "90",
                    "total_score": "100",
                },
                asset_items=[
                    {
                        "item_id": "question_figure_2",
                        "label": "Question 2 figure",
                        "role": "question_figure",
                        "path": str(figure_path),
                        "metadata": {
                            "question_index": "2",
                            "alt_text": "Subtraction diagram",
                        },
                    }
                ],
            )
        ],
    )
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.failure_policy = "warn"

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "batch_output",
    ).run(lambda *_args: None, lambda: False)

    batch_json = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.json"))
    batch_md = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.md"))
    report_data = json.loads(batch_json.read_text(encoding="utf-8"))
    issue = next(
        item
        for item in payload["batch_issue_items"]
        if item["kind"] == "preflight_suspicious_question_figure_mismatch"
    )
    target = json.loads(issue["repair_target_key"])
    workbench_issue = batch_execution_issue_items([issue])[0]

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert "Terminal assembler 'exam' blocked delivery" in payload["error_text"]
    assert payload["material_diagnostics"][0]["change_type"] == (
        "preflight_suspicious_question_figure_mismatch"
    )
    assert payload["material_diagnostics"][0]["expected_question_index"] == "2"
    assert payload["material_diagnostics"][0]["detected_question_index"] == "3"
    assert issue["profile_id"] == "exam_a"
    assert issue["repair_target_type"] == "question_figure_item"
    assert issue["suspicious_asset_items"][0]["detected_question_index"] == "3"
    assert target["item_id"] == "question_figure_2"
    assert target["question_index"] == "2"
    assert target["path"] == str(figure_path)
    assert report_data["batch_issue_items"][0]["suspicious_asset_items"][0][
        "detected_question_index"
    ] == "3"
    assert "looks like question 3" in batch_md.read_text(encoding="utf-8")
    assert "疑似错图" in workbench_issue.details[-1]
    assert "文件名题号 3" in workbench_issue.details[-1]


def test_workbench_batch_runner_reports_manual_question_figure_comparison_issue(tmp_path):
    source = tmp_path / "exam_source.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    figure_path = tmp_path / "question_2.png"
    Image.new("RGB", (320, 240), color="green").save(figure_path)
    archive = EntityArchive(
        archive_id="exam",
        profiles=[
            EntityProfile(
                profile_id="exam_a",
                profile_name="Exam A",
                fields={
                    "paper_title": "Final Exam",
                    "subject": "Math",
                    "grade": "Grade 7",
                    "duration": "90",
                    "total_score": "100",
                },
                asset_items=[
                    {
                        "item_id": "question_figure_2",
                        "label": "Question 2 figure",
                        "role": "question_figure",
                        "path": str(figure_path),
                        "metadata": {
                            "question_index": "2",
                            "alt_text": "Subtraction diagram",
                            "comparison_issue_status": "flagged",
                            "comparison_issue_type": "manual_compare",
                            "comparison_issue_reference": "question_1.png",
                            "comparison_issue_display_name": "question_1.png",
                            "comparison_issue_kind": "题图对比",
                            "comparison_issue_marked_at": "2026-06-21T00:00:00+00:00",
                            "comparison_issue_summary": "题2 题图对比: question_1.png",
                        },
                    }
                ],
            )
        ],
    )
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.failure_policy = "warn"

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "batch_output",
    ).run(lambda *_args: None, lambda: False)

    batch_json = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.json"))
    batch_md = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.md"))
    report_data = json.loads(batch_json.read_text(encoding="utf-8"))
    issue = next(
        item
        for item in payload["batch_issue_items"]
        if item["kind"] == "preflight_question_figure_manual_comparison_issue"
    )
    target = json.loads(issue["repair_target_key"])
    workbench_issue = batch_execution_issue_items([issue])[0]

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert "Terminal assembler 'exam' blocked delivery" in payload["error_text"]
    assert payload["material_diagnostics"][0]["change_type"] == (
        "preflight_question_figure_manual_comparison_issue"
    )
    assert issue["profile_id"] == "exam_a"
    assert issue["repair_target_type"] == "question_figure_item"
    assert issue["comparison_issue_items"][0]["comparison_issue_display_name"] == (
        "question_1.png"
    )
    assert target["item_id"] == "question_figure_2"
    assert target["question_index"] == "2"
    assert target["metadata"]["comparison_issue_status"] == "flagged"
    assert report_data["batch_issue_items"][0]["comparison_issue_items"][0][
        "comparison_issue_kind"
    ] == "题图对比"
    assert "Manual comparison issue for question 2" in batch_md.read_text(
        encoding="utf-8"
    )
    assert "人工对比问题" in workbench_issue.details[-1]
    assert "question_1.png" in workbench_issue.details[-1]


def test_workbench_batch_runner_reports_question_figure_comparison_matrix(tmp_path):
    source = tmp_path / "exam_source.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    figure_a = tmp_path / "question_2_a.png"
    figure_b = tmp_path / "question_2_b.png"
    figure_c = tmp_path / "question_5_b.png"
    expected_q2 = tmp_path / "question_2_expected.png"
    Image.new("RGB", (320, 240), color="green").save(figure_a)
    Image.new("RGB", (320, 240), color="blue").save(figure_b)
    Image.new("RGB", (320, 240), color="purple").save(figure_c)
    Image.new("RGB", (320, 240), color="yellow").save(expected_q2)
    archive = EntityArchive(
        archive_id="exam",
        profiles=[
            EntityProfile(
                profile_id="exam_a",
                profile_name="Exam A",
                fields={
                    "paper_title": "Final Exam A",
                    "subject": "Math",
                    "grade": "Grade 7",
                    "duration": "90",
                    "total_score": "100",
                },
                asset_items=[
                    {
                        "item_id": "question_figure_2_a",
                        "label": "Question 2 figure A",
                        "role": "question_figure",
                        "path": str(figure_a),
                        "metadata": {
                            "question_index": "2",
                            "comparison_issue_status": "flagged",
                            "comparison_issue_type": "manual_compare",
                            "comparison_issue_reference": str(expected_q2),
                            "comparison_issue_display_name": "question_2_expected.png",
                            "comparison_issue_kind": "题图对比",
                            "comparison_issue_marked_at": "2026-06-21T00:00:00+00:00",
                            "comparison_issue_summary": "题2 题图对比",
                            "comparison_issue_region_type": "current_view",
                            "comparison_issue_region_summary": (
                                "current_view(x=0, y=0, w=320, h=240, zoom=100%, image=320x240)"
                            ),
                            "comparison_issue_region_json": (
                                '{"schema_version":1,"type":"current_view","unit":"px"}'
                            ),
                        },
                    }
                ],
            ),
            EntityProfile(
                profile_id="exam_b",
                profile_name="Exam B",
                fields={
                    "paper_title": "Final Exam B",
                    "subject": "Math",
                    "grade": "Grade 7",
                    "duration": "90",
                    "total_score": "100",
                },
                asset_items=[
                    {
                        "item_id": "question_figure_2_b",
                        "label": "Question 2 figure B",
                        "role": "question_figure",
                        "path": str(figure_b),
                        "metadata": {
                            "question_index": "2",
                            "comparison_issue_status": "flagged",
                            "comparison_issue_type": "manual_compare",
                            "comparison_issue_reference": "question_2_expected.png",
                            "comparison_issue_display_name": "question_2_expected.png",
                            "comparison_issue_kind": "题图对比",
                            "comparison_issue_summary": "题2 题图对比",
                        },
                    },
                    {
                        "item_id": "question_figure_5_b",
                        "label": "Question 5 figure B",
                        "role": "question_figure",
                        "path": str(figure_c),
                        "metadata": {
                            "question_index": "5",
                            "comparison_issue_status": "flagged",
                            "comparison_issue_type": "manual_compare",
                            "comparison_issue_reference": "question_5_expected.png",
                            "comparison_issue_display_name": "question_5_expected.png",
                            "comparison_issue_kind": "题图对比",
                            "comparison_issue_summary": "题5 题图对比",
                        },
                    },
                ],
            ),
        ],
    )
    scene = SceneWorkspace()
    scene.input_source_profile.material_schema_id = "exam_items_v1"
    scene.input_source_profile.failure_policy = "warn"

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        archive=archive,
        base_output_dir=tmp_path / "batch_output",
    ).run(lambda *_args: None, lambda: False)

    batch_json = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.json"))
    batch_md = next(Path(path) for path in payload["report_paths"] if path.endswith("_batch_report.md"))
    report_data = json.loads(batch_json.read_text(encoding="utf-8"))
    matrix = report_data["question_figure_comparison_matrix"]
    repair_queue = report_data["question_figure_repair_queue"]
    q2 = next(item for item in matrix["questions"] if item["question_index"] == "2")
    profile_a = next(
        item for item in matrix["profiles"] if item["profile_id"] == "exam_a"
    )
    profile_b = next(
        item for item in matrix["profiles"] if item["profile_id"] == "exam_b"
    )

    assert payload["question_figure_comparison_matrix"]["total_issue_count"] == 3
    assert payload["question_figure_repair_queue"]["queue_count"] == 3
    assert matrix["kind"] == "question_figure_comparison_matrix"
    assert matrix["total_issue_count"] == 3
    assert matrix["profile_count"] == 2
    assert matrix["question_count"] == 2
    assert q2["issue_count"] == 2
    assert {item["profile_id"] for item in q2["profiles"]} == {"exam_a", "exam_b"}
    assert profile_b["issue_count"] == 2
    assert [item["question_index"] for item in profile_b["questions"]] == ["2", "5"]
    assert profile_b["questions"][0]["items"][0]["repair_target_type"] == (
        "question_figure_item"
    )
    assert profile_a["questions"][0]["items"][0]["region_type"] == "current_view"
    assert "current_view(" in profile_a["questions"][0]["items"][0]["region_summary"]
    region_issue = next(
        item
        for item in report_data["batch_issue_items"]
        if item["comparison_issue_items"]
        and item["comparison_issue_items"][0].get("comparison_issue_region_type")
        == "current_view"
    )
    assert region_issue["comparison_issue_items"][0][
        "comparison_issue_region_type"
    ] == "current_view"
    assert repair_queue["kind"] == "question_figure_repair_queue"
    assert repair_queue["status"] == "queued"
    assert repair_queue["queue_count"] == 3
    region_queue_entry = next(
        item for item in repair_queue["entries"] if item["region_type"] == "current_view"
    )
    assert region_queue_entry["action"] == "review_question_figure_replacement"
    assert region_queue_entry["status"] == "candidate"
    assert region_queue_entry["repair_target_type"] == "question_figure_item"
    assert region_queue_entry["requires_user_confirmation"] is True
    assert region_queue_entry["auto_apply_supported"] is False
    assert region_queue_entry["confirmation_action"] == (
        "confirm_question_figure_replacement"
    )
    assert region_queue_entry["confirmation_apply_supported"] is True
    assert region_queue_entry["confirmation_status"] == "ready"
    assert region_queue_entry["replacement_source_path"] == str(expected_q2)
    assert region_queue_entry["replacement_source_kind"] == "local_file"
    assert region_queue_entry["apply_blockers"] == []
    assert "current_view(" in region_queue_entry["region_summary"]
    blocked_queue_entry = next(
        item for item in repair_queue["entries"] if item["question_index"] == "5"
    )
    assert blocked_queue_entry["confirmation_apply_supported"] is False
    assert blocked_queue_entry["confirmation_status"] == "blocked"
    assert "comparison_reference_not_local_file" in blocked_queue_entry["apply_blockers"]
    markdown = batch_md.read_text(encoding="utf-8")
    assert "## Question Figure Comparison Matrix" in markdown
    assert "## Question Figure Repair Queue" in markdown
    assert "Batch confirmation plan: status=partial" in markdown
    assert "batch_apply=false" in markdown
    assert "review_question_figure_replacement" in markdown
    assert "confirm=ready" in markdown
    assert "confirm=blocked" in markdown
    assert "Exam B / Q5: 1 issue(s) -> question_5_expected.png" in markdown
    assert "current_view(" in markdown
