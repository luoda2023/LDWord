from __future__ import annotations

from pathlib import Path
from io import BytesIO

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from src.config.content_materials import (
    ContentInsertionRule,
    ImageBlock,
    InlineKind,
    ListBlock,
    ParagraphBlock,
    TableBlock,
)
from src.services.material_content.artifact_repository import (
    ContentArtifactRepository,
)
from src.services.material_content.compiler import compile_content_material
from src.services.material_content.docx_importer import DocxImportDiagnosticCode
from src.services.material_content.docx_normalizer import _diagnostic_message
from src.services.material_content.docx_renderer import render_document_fragment
from tests.content_docx_fixture_factory import build_anonymous_wps_compat_docx


def _append_equation(paragraph) -> None:
    paragraph._p.append(OxmlElement("m:oMath"))


def test_every_docx_blocker_code_has_a_named_object() -> None:
    for code in DocxImportDiagnosticCode:
        message = _diagnostic_message(code.value)
        assert "对象：" in message
        assert "未识别对象" not in message


def test_docx_blocker_message_names_the_unsupported_object(tmp_path: Path) -> None:
    source = tmp_path / "equation.docx"
    document = Document()
    paragraph = document.add_paragraph("公式：")
    _append_equation(paragraph)
    document.save(source)

    result = compile_content_material(
        source,
        ContentArtifactRepository(tmp_path / "artifacts"),
    )

    assert result.blocked is True
    equation = next(item for item in result.findings if item.code == "omml")
    assert equation.user_message == "DOCX 中有 1 个暂不支持的对象：公式。"


def test_diagnostic_limit_is_applied_after_semantic_object_aggregation(
    tmp_path: Path,
) -> None:
    source = tmp_path / "one-object-many-nodes.docx"
    document = Document()
    paragraph = document.add_paragraph("一个公式对象：")
    for _index in range(120):
        _append_equation(paragraph)
    document.save(source)

    result = compile_content_material(
        source,
        ContentArtifactRepository(tmp_path / "artifacts"),
    )

    blockers = tuple(
        item
        for item in result.findings
        if item.disposition.value == "blocker"
    )
    assert len(blockers) == 1
    assert blockers[0].code == "omml"
    assert all(item.code != "diagnostic_limit_exceeded" for item in blockers)


def test_diagnostic_limit_reports_omitted_semantic_objects(tmp_path: Path) -> None:
    source = tmp_path / "many-equations.docx"
    document = Document()
    for index in range(101):
        paragraph = document.add_paragraph(f"公式 {index + 1}：")
        _append_equation(paragraph)
    document.save(source)

    result = compile_content_material(
        source,
        ContentArtifactRepository(tmp_path / "artifacts"),
    )

    overflow = next(
        item
        for item in result.findings
        if item.code == "diagnostic_limit_exceeded"
    )
    assert overflow.count == 2
    assert overflow.user_message == "另有 2 个暂不支持的对象未展开。"
    assert result.findings[0].code == "omml"


def test_simple_docx_compiles_with_tabs_and_ignores_non_body_stories(
    tmp_path: Path,
) -> None:
    source = tmp_path / "simple.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("before").add_tab()
    paragraph.add_run("after")
    document.sections[0].header.paragraphs[0].text = "header text"
    document.save(source)
    repository = ContentArtifactRepository(tmp_path / "artifacts")

    result = compile_content_material(source, repository)

    assert result.blocked is False
    assert result.artifact_ref is not None
    assert {item.code for item in result.findings} == {
        "header_footer_ignored",
        "source_section_layout_ignored",
    }
    fragment = repository.load_fragment(result.artifact_ref)
    block = fragment.blocks[0]
    assert isinstance(block, ParagraphBlock)
    assert [item.kind for item in block.inlines] == [
        InlineKind.TEXT,
        InlineKind.TAB,
        InlineKind.TEXT,
    ]


def test_docx_artifact_keeps_embedded_image_after_source_is_deleted(
    tmp_path: Path,
) -> None:
    image = tmp_path / "image.png"
    Image.new("RGB", (20, 10), (60, 120, 180)).save(image, format="PNG")
    source = tmp_path / "image.docx"
    document = Document()
    document.add_picture(str(image))
    document.save(source)
    repository = ContentArtifactRepository(tmp_path / "artifacts")

    result = compile_content_material(source, repository)

    assert result.blocked is False
    assert result.artifact_ref is not None
    source.unlink()
    image.unlink()
    fragment = repository.load_fragment(result.artifact_ref)
    image_block = next(block for block in fragment.blocks if isinstance(block, ImageBlock))
    assert repository.resolve_resource(
        result.artifact_ref,
        image_block.resource_id,
    ).is_file()


def test_page_break_remains_semantic_not_a_section_blocker(tmp_path: Path) -> None:
    source = tmp_path / "page-break.docx"
    document = Document()
    run = document.add_paragraph("first").add_run()
    run.add_break(WD_BREAK.PAGE)
    document.add_paragraph("second")
    document.save(source)

    result = compile_content_material(
        source,
        ContentArtifactRepository(tmp_path / "artifacts"),
    )

    assert result.blocked is False


def test_wps_compatibility_fixture_compiles_with_only_typed_normalizations(
    tmp_path: Path,
) -> None:
    source = build_anonymous_wps_compat_docx(tmp_path / "wps-compat.docx")

    result = compile_content_material(
        source,
        ContentArtifactRepository(tmp_path / "artifacts"),
    )

    assert result.blocked is False
    codes = {item.code for item in result.findings}
    assert "disallowed_external_relationship" not in codes
    assert "alternate_content" not in codes
    assert "complex_section" not in codes
    assert "vml" not in codes
    assert "floating_image" not in codes
    assert "merged_table_cell" not in codes
    assert "vml_image_normalized" in codes
    assert "floating_image_inlined" in codes
    assert all(item.disposition.value != "blocker" for item in result.findings)
    assert result.artifact_ref is not None
    source.unlink()
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    fragment = repository.load_fragment(result.artifact_ref)
    images = [block for block in fragment.blocks if isinstance(block, ImageBlock)]
    assert len(images) == 4
    assert len({item.resource_id for item in images}) == 1
    by_code = {item.code: item for item in result.findings}
    assert by_code["vml_image_normalized"].count == 3
    assert by_code["floating_image_inlined"].count == 1
    assert repository.resolve_resource(
        result.artifact_ref,
        images[0].resource_id,
    ).is_file()

    target = Document()
    token = "{{@file:wps}}"
    target.add_paragraph(token)
    receipt = render_document_fragment(
        target,
        fragment,
        ContentInsertionRule("wps-rule", "wps", token),
    )
    assert len(receipt.image_job_drafts) == 4
    rendered = tmp_path / "wps-rendered.docx"
    target.save(rendered)
    reopened = Document(rendered)
    assert len(reopened.tables) == 1


def test_content_controls_revisions_and_cached_fields_are_normalized(
    tmp_path: Path,
) -> None:
    source = tmp_path / "normalized.docx"
    document = Document()

    control_paragraph = document.add_paragraph()
    control = OxmlElement("w:sdt")
    control.append(OxmlElement("w:sdtPr"))
    control_content = OxmlElement("w:sdtContent")
    control_run = OxmlElement("w:r")
    control_text = OxmlElement("w:t")
    control_text.text = "controlled"
    control_run.append(control_text)
    control_content.append(control_run)
    control.append(control_content)
    control_paragraph._p.append(control)

    revision_paragraph = document.add_paragraph("base-")
    inserted = OxmlElement("w:ins")
    inserted_run = OxmlElement("w:r")
    inserted_text = OxmlElement("w:t")
    inserted_text.text = "kept"
    inserted_run.append(inserted_text)
    inserted.append(inserted_run)
    deleted = OxmlElement("w:del")
    deleted_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "removed"
    deleted_run.append(deleted_text)
    deleted.append(deleted_run)
    revision_paragraph._p.extend((inserted, deleted))

    simple_field_paragraph = document.add_paragraph()
    simple_field = OxmlElement("w:fldSimple")
    simple_field.set(qn("w:instr"), " DATE ")
    simple_result = OxmlElement("w:r")
    simple_text = OxmlElement("w:t")
    simple_text.text = "2026-07-13"
    simple_result.append(simple_text)
    simple_field.append(simple_result)
    simple_field_paragraph._p.append(simple_field)

    complex = document.add_paragraph()
    for field_type, text in (
        ("begin", ""),
        ("", " PAGE "),
        ("separate", ""),
        ("", "7"),
        ("end", ""),
    ):
        run = OxmlElement("w:r")
        if field_type:
            marker = OxmlElement("w:fldChar")
            marker.set(qn("w:fldCharType"), field_type)
            run.append(marker)
        else:
            node = OxmlElement("w:instrText" if text.strip() == "PAGE" else "w:t")
            node.text = text
            run.append(node)
        complex._p.append(run)
    document.save(source)

    result = compile_content_material(
        source,
        ContentArtifactRepository(tmp_path / "artifacts"),
    )

    assert result.blocked is False
    visible = [
        "".join(inline.text for inline in block.inlines)
        for block in result.fragment.blocks
        if isinstance(block, ParagraphBlock)
    ]
    assert visible == ["controlled", "base-kept", "2026-07-13", "7"]
    codes = {item.code for item in result.findings}
    assert "content_control_unwrapped" in codes
    assert "revision_final_view" in codes
    assert "field_cached_result" in codes


def test_merged_and_multi_block_table_compiles_and_reopens_after_render(
    tmp_path: Path,
) -> None:
    source = tmp_path / "merged-table.docx"
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "wide"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 2).text = "right"
    table.cell(1, 0).text = "tall"
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 1).text = "first paragraph"
    table.cell(1, 1).add_paragraph("second paragraph")
    table.cell(1, 2).text = "item one"
    table.cell(1, 2).paragraphs[0].style = "List Number"
    table.cell(1, 2).add_paragraph("item two", style="List Number")
    table.cell(2, 1).text = "bottom middle"
    table.cell(2, 2).text = "bottom right"
    document.save(source)
    repository = ContentArtifactRepository(tmp_path / "artifacts")

    result = compile_content_material(source, repository)

    assert result.blocked is False
    table_block = next(
        block for block in result.fragment.blocks if isinstance(block, TableBlock)
    )
    assert [len(row.cells) for row in table_block.rows] == [2, 3, 2]
    assert table_block.rows[0].cells[0].colspan == 2
    assert table_block.rows[1].cells[0].rowspan == 2
    assert len(table_block.rows[1].cells[1].blocks) == 2
    assert isinstance(table_block.rows[1].cells[2].blocks[0], ListBlock)

    target = Document()
    token = "{{@file:table}}"
    target.add_paragraph(token)
    render_document_fragment(
        target,
        result.fragment,
        ContentInsertionRule("table-rule", "table", token),
    )
    stream = BytesIO()
    target.save(stream)
    stream.seek(0)
    reopened = Document(stream)
    rendered_table = reopened.tables[0]._tbl
    assert rendered_table.find(".//" + qn("w:gridSpan")) is not None
    assert len(rendered_table.findall(".//" + qn("w:vMerge"))) == 2


def test_table_cell_image_is_artifact_backed_and_emits_deferred_job(
    tmp_path: Path,
) -> None:
    image = tmp_path / "cell.png"
    Image.new("RGB", (18, 9), (180, 90, 30)).save(image, format="PNG")
    source = tmp_path / "table-image.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run().add_picture(str(image))
    document.save(source)
    repository = ContentArtifactRepository(tmp_path / "artifacts")

    result = compile_content_material(source, repository)

    assert result.blocked is False
    table_block = next(
        block for block in result.fragment.blocks if isinstance(block, TableBlock)
    )
    cell_image = table_block.rows[0].cells[0].blocks[0]
    assert isinstance(cell_image, ImageBlock)
    assert repository.resolve_resource(
        result.artifact_ref,
        cell_image.resource_id,
    ).is_file()

    target = Document()
    token = "{{@file:cell_image}}"
    target.add_paragraph(token)
    receipt = render_document_fragment(
        target,
        result.fragment,
        ContentInsertionRule("cell-image-rule", "cell_image", token),
    )
    assert len(receipt.image_job_drafts) == 1
    assert receipt.image_job_drafts[0].resource_id == cell_image.resource_id
