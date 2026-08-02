from types import SimpleNamespace
from pathlib import Path

from docx import Document
import pytest

from src.modules.base import BaseModule, ModuleMeta
from src.modules.registry import ALL_MODULES
from src.pipeline.module_phases import (
    MODULE_MUTATION_OVERRIDES,
    MODULE_PHASES,
    MODULE_SCOPE_BEHAVIORS,
    ModuleScopeBehavior,
    ModulePhaseContractError,
    build_module_phase_graph,
    module_scope_behavior,
    order_modules_by_phase,
)
from src.pipeline.phases import Phase, PhaseGraphValidationError
from src.pipeline.runner import Pipeline


EXPECTED_PHASES = {
    Phase.FILL: {
        "entity_fill",
        "source_fill",
        "placeholder_replace",
        "formula_convert",
        "md_cleanup",
    },
    Phase.SEMANTICS: {
        "heading_recognition",
        "whitespace_normalize",
        "heading_numbering",
        "caption",
        "reference_format",
        "citation_link",
    },
    Phase.FORMAT: {
        "section_format",
        "page_setup",
        "paragraph_style",
        "table_format",
        "figure_table_center",
        "chem_typography",
        "equation_table_format",
        "header_footer",
        "toc",
        "watermark",
    },
    Phase.STAGE_ARTIFACT: {"image_insertion"},
    Phase.FINAL_VALIDATION: {"validation"},
}

EXPECTED_SCOPE_BEHAVIORS = {
    ModuleScopeBehavior.STRUCTURE_DISCOVERY: {"heading_recognition"},
    ModuleScopeBehavior.REGION_FILTERED: {
        "whitespace_normalize",
        "heading_numbering",
        "caption",
        "reference_format",
        "citation_link",
        "paragraph_style",
        "table_format",
        "figure_table_center",
        "formula_convert",
        "chem_typography",
        "equation_table_format",
        "toc",
    },
    ModuleScopeBehavior.DOCUMENT_LEVEL: {
        "entity_fill",
        "source_fill",
        "placeholder_replace",
        "md_cleanup",
        "page_setup",
        "section_format",
        "header_footer",
        "watermark",
        "image_insertion",
        "validation",
    },
}


def _names(modules):
    return [module.meta.name for module in modules]


def test_explicit_phase_map_is_complete_and_contains_no_unregistered_names():
    registry_names = {module.meta.name for module in ALL_MODULES}

    assert set(MODULE_PHASES) == registry_names
    assert set().union(*EXPECTED_PHASES.values()) == registry_names
    for phase, expected_names in EXPECTED_PHASES.items():
        assert {
            name for name, assigned_phase in MODULE_PHASES.items()
            if assigned_phase is phase
        } == expected_names


def test_scope_behavior_map_is_complete_and_explicit():
    registry_names = {module.meta.name for module in ALL_MODULES}

    assert set(MODULE_SCOPE_BEHAVIORS) == registry_names
    assert set().union(*EXPECTED_SCOPE_BEHAVIORS.values()) == registry_names
    for behavior, expected_names in EXPECTED_SCOPE_BEHAVIORS.items():
        assert {
            name
            for name, assigned_behavior in MODULE_SCOPE_BEHAVIORS.items()
            if assigned_behavior is behavior
        } == expected_names
        assert {
            module.meta.name
            for module in ALL_MODULES
            if module_scope_behavior(module) is behavior
        } == expected_names


def test_full_registry_graph_projects_hard_and_present_soft_dependencies():
    graph = build_module_phase_graph(ALL_MODULES)
    present_names = {module.meta.name for module in ALL_MODULES}
    steps = {step.name: step for step in graph.steps}

    for module in ALL_MODULES:
        expected = tuple(
            dict.fromkeys(
                (*module.meta.depends_on, *(
                    dependency
                    for dependency in module.meta.soft_after
                    if dependency in present_names
                ))
            )
        )
        assert steps[module.meta.name].depends_on == expected


def test_missing_soft_hint_is_ignored_but_missing_hard_dependency_is_diagnostic():
    by_name = {module.meta.name: module for module in ALL_MODULES}

    graph = build_module_phase_graph([by_name["paragraph_style"]])
    assert graph.steps[0].depends_on == ()

    with pytest.raises(PhaseGraphValidationError, match="missing dependency"):
        build_module_phase_graph([by_name["heading_numbering"]])


def test_unknown_module_is_rejected_instead_of_inferred_from_category():
    unknown = SimpleNamespace(
        meta=SimpleNamespace(
            name="future_magic",
            category="format",
            depends_on=(),
            soft_after=(),
            modifies_structure=False,
        )
    )

    with pytest.raises(
        ModulePhaseContractError,
        match="no explicit phase assignment.*future_magic",
    ):
        build_module_phase_graph([unknown])


def test_unknown_module_with_explicit_phase_is_schedulable():
    unknown = SimpleNamespace(
        meta=SimpleNamespace(
            name="future_magic",
            category="anything",
            execution_phase="compose",
            scope_behavior="document_level",
            depends_on=(),
            soft_after=(),
            modifies_structure=False,
        )
    )

    graph = build_module_phase_graph([unknown])
    assert graph.steps[0].phase is Phase.COMPOSE
    assert order_modules_by_phase([unknown]) == (unknown,)


def test_registered_module_explicit_phase_must_match_compatibility_fact():
    conflict = SimpleNamespace(
        meta=SimpleNamespace(
            name="page_setup",
            category="anything",
            execution_phase="fill",
            scope_behavior="document_level",
            depends_on=(),
            soft_after=(),
            modifies_structure=False,
        )
    )

    with pytest.raises(ModulePhaseContractError, match="conflicts.*format"):
        build_module_phase_graph([conflict])


def test_unknown_module_requires_explicit_scope_behavior():
    unknown = SimpleNamespace(
        meta=SimpleNamespace(
            name="future_magic",
            category="anything",
            execution_phase="compose",
            depends_on=(),
            soft_after=(),
            modifies_structure=False,
        )
    )

    with pytest.raises(ModulePhaseContractError, match="no explicit scope_behavior"):
        build_module_phase_graph([unknown])


def test_structural_mutation_overrides_cover_body_index_invalidators():
    expected_invalidators = {
        "caption",
        "equation_table_format",
        "formula_convert",
        "section_format",
        "toc",
        "image_insertion",
    }
    graph = build_module_phase_graph(ALL_MODULES)
    invalidators = {
        step.name
        for step in graph.steps
        if step.invalidates_document_index
    }

    assert invalidators == expected_invalidators
    for name in expected_invalidators:
        contract = MODULE_MUTATION_OVERRIDES[name]
        assert contract.changes_document_structure is True
        assert contract.invalidates_document_index is True

    for module in ALL_MODULES:
        if module.meta.modifies_structure:
            assert module.meta.name in expected_invalidators


def test_default_phase_order_is_deterministic_and_uses_registry_ties():
    expected = [
        # Fill: independent modules retain registry order.
        "entity_fill",
        "source_fill",
        "placeholder_replace",
        "formula_convert",
        "md_cleanup",
        # Semantics: reference_format moves before its citation_link consumer.
        "heading_recognition",
        "heading_numbering",
        "caption",
        "whitespace_normalize",
        "reference_format",
        "citation_link",
        # Format: section topology is finalized before page properties are applied.
        "section_format",
        "page_setup",
        "paragraph_style",
        "header_footer",
        "toc",
        "table_format",
        "figure_table_center",
        "watermark",
        "chem_typography",
        "equation_table_format",
        # Explicit late barriers.
        "image_insertion",
        "validation",
    ]

    first = _names(order_modules_by_phase(ALL_MODULES))
    second = _names(order_modules_by_phase(tuple(ALL_MODULES)))
    assert first == expected
    assert second == expected


def test_image_is_after_every_format_step_and_validation_is_last():
    ordered_steps = build_module_phase_graph(ALL_MODULES).topological_order()
    positions = {step.name: index for index, step in enumerate(ordered_steps)}
    image_position = positions["image_insertion"]

    assert all(
        positions[step.name] < image_position
        for step in ordered_steps
        if step.phase is Phase.FORMAT
    )
    assert ordered_steps[-1].name == "validation"
    assert ordered_steps[-1].phase is Phase.FINAL_VALIDATION


class _RecordingModule(BaseModule):
    def __init__(
        self,
        name: str,
        phase: str,
        events: list[tuple[str, int]],
        *,
        fail: bool = False,
    ) -> None:
        self.meta = ModuleMeta(
            name=name,
            description=name,
            category="test",
            execution_phase=phase,
            scope_behavior=(
                "region_filtered"
                if name in MODULE_SCOPE_BEHAVIORS
                and MODULE_SCOPE_BEHAVIORS[name]
                is ModuleScopeBehavior.REGION_FILTERED
                else "document_level"
            ),
        )
        self._events = events
        self._fail = fail

    def apply(self, doc, config, tracker, context):
        self._events.append((self.meta.name, len(context.heading_map or {})))
        doc.add_heading(self.meta.name, level=1)
        if self._fail:
            raise RuntimeError("intentional failure")


def test_pipeline_uses_phase_barriers_and_registration_ties_not_alphabetical_scheduler():
    events: list[tuple[str, int]] = []
    format_z = _RecordingModule("z_format_plugin", "format", events)
    fill_z = _RecordingModule("z_fill_plugin", "fill", events)
    fill_a = _RecordingModule("a_fill_plugin", "fill", events)

    pipeline = Pipeline(
        modules=[format_z, fill_z, fill_a],
        config=SimpleNamespace(strict_mode=False),
    )

    assert _names(pipeline._modules) == [
        "z_fill_plugin",
        "a_fill_plugin",
        "z_format_plugin",
    ]


def test_pipeline_rebuilds_index_after_every_successful_graph_invalidator(
    tmp_path: Path,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    events: list[tuple[str, int]] = []
    modules = [
        _RecordingModule("image_insertion", "stage_artifact", events),
        _RecordingModule("toc", "format", events),
        _RecordingModule("section_format", "format", events),
        _RecordingModule("caption", "semantics", events),
    ]
    pipeline = Pipeline(
        modules=modules,
        config=SimpleNamespace(strict_mode=False),
        output_dir=str(tmp_path),
    )

    result = pipeline.execute(str(source))

    assert result.success
    assert _names(pipeline._modules) == [
        "caption",
        "toc",
        "section_format",
        "image_insertion",
    ]
    assert events == [
        ("caption", 0),
        ("toc", 1),
        ("section_format", 2),
        ("image_insertion", 3),
    ]
    rebuilds = [
        record
        for record in result.tracker.get_all()
        if record.change_type == "document_index_rebuild"
    ]
    assert [record.rule_name for record in rebuilds] == [
        "caption",
        "toc",
        "section_format",
        "image_insertion",
    ]
    assert result.context.heading_map is not None
    assert len(result.context.heading_map) == 4


def test_failed_or_preflight_skipped_invalidator_does_not_rebuild(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)

    failed_events: list[tuple[str, int]] = []
    failed = Pipeline(
        modules=[
            _RecordingModule(
                "caption", "semantics", failed_events, fail=True
            )
        ],
        config=SimpleNamespace(strict_mode=False),
        output_dir=str(tmp_path / "failed"),
    ).execute(str(source))
    assert failed.status == "partial_success"
    assert not any(
        record.change_type == "document_index_rebuild"
        for record in failed.tracker.get_all()
    )

    skipped_events: list[tuple[str, int]] = []
    skipped_pipeline = Pipeline(
        modules=[_RecordingModule("caption", "semantics", skipped_events)],
        config=SimpleNamespace(strict_mode=False),
        output_dir=str(tmp_path / "skipped"),
    )
    monkeypatch.setattr(
        skipped_pipeline,
        "_run_object_preflight",
        lambda _path, _ctx, *, safe_package: (
            False,
            {"caption": {"reason": "fixture skip"}},
        ),
    )
    skipped = skipped_pipeline.execute(str(source))
    assert skipped.success
    assert skipped_events == []
    assert not any(
        record.change_type == "document_index_rebuild"
        for record in skipped.tracker.get_all()
    )
