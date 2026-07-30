"""Valid in-process Office-layout fake shared by image-chain tests."""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from shutil import copy2

from docx import Document

from src.config.image_materials import ImagePlacementMode
from src.shared.engine.office_image_layout import (
    LAYOUT_SCHEMA_VERSION,
    LayoutStatus,
    OfficeImageJobReceipt,
    OfficeImageLayoutReceipt,
    PROVIDER_SPECS,
    StabilizationRoundReceipt,
    build_controlled_layout_shadow_path,
    image_inventory_sha256,
    image_preservation_sha256,
    inventory_ooxml_images,
)


class FakeSuccessfulLayout:
    def __init__(self, *, foreign_shadow: Path | None = None) -> None:
        self.foreign_shadow = foreign_shadow
        self.requests = []

    def __call__(self, request):
        self.requests.append(request)
        source = Path(request.source_docx_path)
        shadow = self.foreign_shadow or build_controlled_layout_shadow_path(
            source,
            nonce="a" * 32,
        )
        copy2(source, shadow)
        document = Document(shadow)
        expected_shapes: set[tuple[str, str]] = set()
        for request_job in request.jobs:
            shape_id = layout_shape_id(request_job.plan.job_id)
            inline = document.add_paragraph().add_run().add_picture(
                request_job.prepared_image.output_path
            )
            inline._inline.docPr.set("descr", shape_id)
            inline._inline.docPr.set("title", request_job.plan.job_id)
            expected_shapes.add((shape_id, request_job.plan.job_id))
        document.save(shadow)
        jobs = tuple(layout_job_receipt(job) for job in request.jobs)
        source_inventory = inventory_ooxml_images(source)
        shadow_inventory = inventory_ooxml_images(shadow)
        inserted_inventory = tuple(
            item
            for item in shadow_inventory
            if (item.alternative_text, item.title) in expected_shapes
        )
        inserted_set = set(inserted_inventory)
        retained_inventory = tuple(
            item for item in shadow_inventory if item not in inserted_set
        )
        spec = PROVIDER_SPECS[request.provider]
        return OfficeImageLayoutReceipt(
            schema_version=LAYOUT_SCHEMA_VERSION,
            transaction_id=request.transaction_id,
            provider=request.provider,
            adapter_name=spec.adapter_name,
            prog_id=spec.prog_id,
            status=LayoutStatus.SUCCESS,
            application_pid=321,
            source_docx_path=str(source),
            source_docx_sha256_before=request.source_docx_sha256,
            source_docx_sha256_after=request.source_docx_sha256,
            shadow_path=str(shadow),
            shadow_sha256=_hash(shadow),
            shadow_retained=True,
            source_images_unchanged=True,
            prepared_images_unchanged=True,
            document_open_count=1,
            document_save_count=1,
            repaginate_count=2,
            pdf_export_count=0,
            field_update_rounds=2,
            page_break_count_before=0,
            page_break_count_after=0,
            jobs=jobs,
            stabilization_rounds=(
                StabilizationRoundReceipt(
                    round_index=1,
                    page_count=1,
                    field_update_count=1,
                    job_pages=tuple(
                        (
                            item.job_id,
                            item.final_guard_page,
                            item.final_anchor_page,
                            item.final_image_page,
                        )
                        for item in jobs
                    ),
                    violation_job_ids=(),
                    corrected_job_ids=(),
                    revalidate_from_job_id="",
                    state_sha256=sha256(b"fake-stable-round").hexdigest(),
                    stable=True,
                ),
            ),
            warnings=(),
            failures=(),
            elapsed_ms=1.0,
            source_image_inventory_sha256=image_inventory_sha256(
                source_inventory
            ),
            shadow_image_inventory_sha256=image_inventory_sha256(
                shadow_inventory
            ),
            preexisting_image_semantic_sha256_before=(
                image_preservation_sha256(source_inventory)
            ),
            preexisting_image_semantic_sha256_after=(
                image_preservation_sha256(retained_inventory)
            ),
            preexisting_images_unchanged=True,
            inserted_images_job_owned=True,
            inserted_images_visible=True,
            sentinel_texts_hidden=True,
            source_image_inventory=source_inventory,
            shadow_image_inventory=shadow_inventory,
            inserted_image_inventory=inserted_inventory,
        )


def layout_shape_id(job_id: str) -> str:
    return "lark-layout-image-" + sha256(job_id.encode("utf-8")).hexdigest()[:20]


def layout_job_receipt(job) -> OfficeImageJobReceipt:
    plan = job.plan
    prepared = job.prepared_image
    aspect = float(prepared.width_px) / float(prepared.height_px)
    strict = plan.placement.mode is ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE
    guard_hash = "a" * 64 if strict else ""
    return OfficeImageJobReceipt(
        job_id=plan.job_id,
        mode=plan.placement.mode,
        sequence=plan.sequence,
        shape_id=layout_shape_id(plan.job_id),
        source_image_sha256=plan.image_ref.content_sha256,
        prepared_image_sha256=prepared.output_sha256,
        measurement_guard_page=1 if strict else None,
        measurement_anchor_page=1,
        measurement_anchor_start=1,
        anchor_y_pt=72.0,
        section_index=1,
        column_index=1,
        column_count=1,
        container_width_pt=450.0,
        flow_bottom_pt=720.0,
        paragraph_reserve_pt=12.0,
        safety_margin_pt=8.0,
        available_height_pt=None,
        width_cap_pt=450.0,
        initial_target_width_pt=220.0,
        initial_target_height_pt=220.0 / aspect,
        initial_scale_ratio=1.0,
        readability_warning=False,
        co_location_claimed=strict,
        final_guard_page=1 if strict else None,
        final_anchor_page=1,
        final_image_page=1,
        final_image_y_pt=72.0,
        final_width_pt=220.0,
        final_height_pt=220.0 / aspect,
        proportion_error=0.0,
        boundary_ok=True,
        guard_xml_sha256_before=guard_hash,
        guard_xml_sha256_after=guard_hash,
        guard_keep_sha256_before=guard_hash,
        guard_keep_sha256_after=guard_hash,
        initial_actual_width_pt=220.0,
        initial_actual_height_pt=220.0 / aspect,
        initial_proportion_error=0.0,
    )


def _hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


__all__ = ["FakeSuccessfulLayout", "layout_job_receipt", "layout_shape_id"]
