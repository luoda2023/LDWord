from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement

from src.config.resolved import ResolvedConfig
from src.modules.special.chem_typography import ChemTypographyModule
from src.modules.special.formula_to_table import _mark_equation_table
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn


def test_chem_typography_consumes_configured_western_font_and_restores_subscript():
    document = Document()
    paragraph = document.add_paragraph("H2O")
    config = ResolvedConfig()
    config.chem_typography.enabled = True
    config.chem_typography.western_font = "Arial"
    config.chem_typography.scopes["body"] = True

    ChemTypographyModule().apply(
        document,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    assert "".join(run.text for run in paragraph.runs) == "H2O"
    subscript_runs = [run for run in paragraph.runs if run.font.subscript]
    assert [run.text for run in subscript_runs] == ["2"]
    for run in paragraph.runs:
        fonts = run._element.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
        assert fonts is not None
        assert fonts.get(qn("w:ascii")) == "Arial"
        assert fonts.get(qn("w:hAnsi")) == "Arial"


def test_chem_typography_normalizes_stale_east_asia_font_hint():
    document = Document()
    run = document.add_paragraph("H2O").runs[0]
    run_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:hint"), "eastAsia")
    fonts.set(qn("w:asciiTheme"), "minorHAnsi")
    run_pr.append(fonts)
    run._element.insert(0, run_pr)

    config = ResolvedConfig()
    config.chem_typography.enabled = True
    config.chem_typography.scopes["body"] = True
    ChemTypographyModule().apply(
        document,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    normalized = run._element.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
    assert normalized is not None
    assert normalized.get(qn("w:hint")) == "default"
    assert normalized.get(qn("w:asciiTheme")) is None


def test_chem_typography_does_not_normalize_plain_text_font_hints():
    document = Document()
    run = document.add_paragraph("ordinary status note¹").runs[0]
    run_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:hint"), "eastAsia")
    fonts.set(qn("w:asciiTheme"), "minorHAnsi")
    run_pr.append(fonts)
    run._element.insert(0, run_pr)

    config = ResolvedConfig()
    config.chem_typography.enabled = True
    config.chem_typography.scopes["body"] = True
    ChemTypographyModule().apply(
        document,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    preserved = run._element.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
    assert preserved is not None
    assert preserved.get(qn("w:hint")) == "eastAsia"
    assert preserved.get(qn("w:asciiTheme")) == "minorHAnsi"
    assert len(document.paragraphs[0].runs) == 1


def test_chem_typography_does_not_rewrite_equation_table_scripts():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    _mark_equation_table(table)
    formula_run = table.cell(0, 0).paragraphs[0].add_run("H2O")
    config = ResolvedConfig()
    config.chem_typography.enabled = True
    config.chem_typography.scopes["tables"] = True

    ChemTypographyModule().apply(
        document,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    assert len(table.cell(0, 0).paragraphs[0].runs) == 1
    assert formula_run.font.subscript is not True
