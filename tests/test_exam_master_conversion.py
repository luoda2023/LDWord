from __future__ import annotations

import hashlib

from docx import Document
from docx.oxml import OxmlElement

from src.shared.engine.exam_master_conversion import (
    convert_exam_docx_to_user_master,
)


def _sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _complete_exam(path):
    document = Document()
    title = document.add_heading("七年级语文期中测试卷", level=1)
    graphic_run = title.add_run()
    graphic_run._r.append(OxmlElement("w:pict"))
    document.add_paragraph("考试时间：90分钟 满分：100分")
    score = document.add_table(rows=2, cols=3)
    score.cell(0, 0).text = "题号"
    score.cell(0, 1).text = "一"
    score.cell(0, 2).text = "总分"
    score.cell(1, 0).text = "评分"
    document.add_paragraph("一、积累与运用")
    document.add_paragraph("1. 选择正确的一项。")
    document.add_paragraph("A. 甲")
    document.add_paragraph("B. 乙")
    document.add_paragraph("答案解析部分")
    document.add_paragraph("1. A")
    document.save(path)


def test_complete_exam_is_distilled_without_mutating_source(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "master.docx"
    _complete_exam(source)
    before = _sha256(source)

    result = convert_exam_docx_to_user_master(source, output)

    assert _sha256(source) == before
    assert result.answer_section_detected is True
    assert result.preserved_graphic_run_count == 1
    converted = Document(output)
    body_text = "\n".join(paragraph.text for paragraph in converted.paragraphs)
    assert "{{af_title}}" in body_text
    assert "{{af_questions}}" in body_text
    assert "{{af_subject}}" in body_text
    assert "选择正确的一项" not in body_text
    assert "答案解析部分" not in body_text
    assert "<w:pict" in converted.paragraphs[0]._p.xml
    assert converted.tables[0].cell(1, 0).text == "评分"
    for style_name in (
        "Exam Section Heading",
        "Exam Question",
        "Exam Question Stem",
        "Exam Option",
        "Exam Answer Space",
    ):
        assert style_name in {style.name for style in converted.styles}


def test_conversion_rejects_document_without_detectable_exam_section(tmp_path):
    source = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("普通笔记")
    document.save(source)

    try:
        convert_exam_docx_to_user_master(source, tmp_path / "master.docx")
    except ValueError as exc:
        assert str(exc) == "first_exam_section_not_detected"
    else:  # pragma: no cover
        raise AssertionError("unsafe conversion should have been rejected")
