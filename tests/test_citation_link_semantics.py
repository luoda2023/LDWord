import sys
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.modules.special.citation_link import CitationLinkModule
from src.modules.structure.heading_recognition import HeadingRecognitionModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn


def _apply_citation_link(doc: Document, *, auto_number_reference_entries: bool):
    context = PipelineContext()
    HeadingRecognitionModule().apply(doc, ResolvedConfig(), ChangeTracker(), context)

    config = ResolvedConfig()
    config.citation_link.auto_number_reference_entries = auto_number_reference_entries
    tracker = ChangeTracker()
    CitationLinkModule().apply(doc, config, tracker, context)
    return tracker


def test_citation_link_builds_reference_entry_bookmarks_and_body_ref_fields():
    doc = Document()
    doc.add_heading("Chapter 1 Intro", level=1)
    body_para = doc.add_paragraph("See [1, 2] for details.")
    doc.add_paragraph("References")
    ref1 = doc.add_paragraph("[1] Author. Title[J]. Journal, 2024, 12(3): 1-8.")
    ref2 = doc.add_paragraph("[2] Author. Another title[J]. Journal, 2023, 11(2): 9-12.")

    _apply_citation_link(doc, auto_number_reference_entries=False)

    bookmark_names = [b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))]
    instr_texts = [t.text or "" for t in body_para._element.iter(qn("w:instrText"))]

    assert "_RefEntry_1" in bookmark_names
    assert "_RefEntry_2" in bookmark_names
    assert any("REF _RefEntry_1" in text for text in instr_texts)
    assert any("REF _RefEntry_2" in text for text in instr_texts)
    assert ref1.text.startswith("[1]")
    assert ref2.text.startswith("[2]")


def test_citation_link_uses_number_bookmark_targets_when_auto_numbering_is_enabled():
    doc = Document()
    doc.add_heading("Chapter 1 Intro", level=1)
    body_para = doc.add_paragraph("See [1] for details.")
    doc.add_paragraph("References")
    doc.add_paragraph("[1] Author. Title[J]. Journal, 2024, 12(3): 1-8.")

    _apply_citation_link(doc, auto_number_reference_entries=True)

    bookmark_names = [b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))]
    instr_texts = [t.text or "" for t in body_para._element.iter(qn("w:instrText"))]

    assert "_RefNum_1" in bookmark_names
    assert any("REF _RefNum_1" in text for text in instr_texts)
