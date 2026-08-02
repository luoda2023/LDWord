from __future__ import annotations

from dataclasses import asdict

from docx import Document

from src.config.resolved import ResolvedConfig
from src.config.scene import SceneWorkspace
from src.modules.validate.md_cleanup import MdCleanupModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.ui.panels.scene_card_definitions import CARD_DEFINITIONS
from src.ui.panels.scene_input_cleanup_rules import SceneInputCleanupRulesCard


def test_input_cleanup_card_projects_scene_without_mutating_it(qapp):
    scene = SceneWorkspace(scene_id="markdown_plan", mode_id="custom")
    scene.input_source_profile.accepted_formats = ["docx", "markdown"]
    scene.input_source_profile.markdown_policy = "preview_and_cleanup"
    scene.module_switches["md_cleanup"] = True
    scene.module_switches["whitespace_normalize"] = False
    before = asdict(scene)

    card = SceneInputCleanupRulesCard()
    try:
        card.set_scene(scene)
        qapp.processEvents()

        assert asdict(scene) == before
        assert not card._markdown_row.isHidden()
        assert card._markdown_policy.currentData() == "preview_and_cleanup"
        assert card._whitespace_enabled.isChecked() is False
        assert card._whitespace_options_row.isHidden()
    finally:
        card.close()


def test_input_cleanup_card_writes_real_plan_fields(qapp):
    scene = SceneWorkspace(scene_id="editable_plan", mode_id="custom")
    scene.input_source_profile.accepted_formats = ["docx", "markdown"]
    scene.input_source_profile.markdown_policy = "cleanup_only"
    scene.module_switches["whitespace_normalize"] = False
    card = SceneInputCleanupRulesCard()
    try:
        card.set_scene(scene)

        card._markdown_policy.setCurrentIndex(
            card._markdown_policy.findData("preview_and_cleanup")
        )
        card._whitespace_enabled.click()
        card._whitespace_option_checks["remove_zero_width"].setChecked(False)
        qapp.processEvents()

        assert scene.input_source_profile.markdown_policy == "preview_and_cleanup"
        assert scene.module_switches["md_cleanup"] is True
        assert scene.module_switches["whitespace_normalize"] is True
        assert scene.whitespace.remove_zero_width is False
    finally:
        card.close()


def test_markdown_control_is_hidden_when_plan_does_not_accept_markdown(qapp):
    scene = SceneWorkspace(scene_id="word_only", mode_id="official")
    scene.input_source_profile.accepted_formats = ["docx"]
    scene.input_source_profile.markdown_policy = "disabled"
    card = SceneInputCleanupRulesCard()
    try:
        card.set_scene(scene)
        qapp.processEvents()

        assert card._markdown_row.isHidden()
        assert card._whitespace_row.isHidden() is False
    finally:
        card.close()


def test_markup_cleanup_uses_actual_source_provenance():
    config = ResolvedConfig()
    config.input_source_profile.markdown_policy = "cleanup_only"
    config.md_cleanup.formula_copy_noise_cleanup = False
    config.md_cleanup.suppress_formula_fake_lists = False
    module = MdCleanupModule()

    word_doc = Document()
    word_doc.add_paragraph("**Word 中的字面 Markdown 示例**")
    module.apply(
        word_doc,
        config,
        ChangeTracker(),
        PipelineContext(source_doc_path="source.docx"),
    )
    assert word_doc.paragraphs[0].text == "**Word 中的字面 Markdown 示例**"

    markdown_doc = Document()
    markdown_doc.add_paragraph("**Markdown 来源**")
    module.apply(
        markdown_doc,
        config,
        ChangeTracker(),
        PipelineContext(source_doc_path="source.md"),
    )
    assert markdown_doc.paragraphs[0].text == "Markdown 来源"


def test_risk_check_has_no_visible_scene_navigation_contract():
    assert "scn_cleanup" not in CARD_DEFINITIONS
    assert all(title != "风险检查" for title, _icon in CARD_DEFINITIONS.values())
