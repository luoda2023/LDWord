"""
Phase 3 冒烟测试 — 🔢 结构与编号模块 (3 个)
"""

import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE


def test_heading_recognition_pipeline():
    """M05: heading_recognition 真实 docx 管线测试"""
    from src.modules.structure.heading_recognition import (
        HeadingRecognitionModule, DocTree, HeadingInfo,
    )
    from src.config.template import TemplateConfig
    from src.config.scene import SceneWorkspace
    from src.config.resolver import resolve_config
    from src.pipeline.runner import Pipeline

    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("这是正文内容。")
    doc.add_heading("1.1 研究背景", level=2)
    doc.add_paragraph("详细背景说明。")
    doc.add_heading("1.1.1 国内现状", level=3)
    doc.add_paragraph("国内发展情况。")
    doc.add_heading("第二章 方法", level=1)
    doc.add_paragraph("方法描述。")

    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)

    try:
        template = TemplateConfig()
        scene = SceneWorkspace()
        config = resolve_config(template, scene)
        pipeline = Pipeline(modules=[HeadingRecognitionModule()], config=config)
        result = pipeline.execute(tmp.name)
        assert result.success, f"执行失败: {result.error}"
        ctx = result.context
        assert ctx.doc_tree is not None
        assert len(ctx.doc_tree.headings) >= 4
        assert ctx.heading_map is not None
        levels = [h.level for h in ctx.doc_tree.headings]
        assert 1 in levels and 2 in levels and 3 in levels
        print(f"  ✅ M05 heading_recognition 通过 ({len(ctx.doc_tree.headings)} 个标题)")
    finally:
        Path(tmp.name).unlink(missing_ok=True)
        out = result.output_paths.get("final", "")
        if out:
            Path(out).unlink(missing_ok=True)


def test_heading_numbering_import():
    """M06: heading_numbering 导入+实例化"""
    from src.modules.structure.heading_numbering import HeadingNumberingModule
    mod = HeadingNumberingModule()
    assert mod.meta.name == "heading_numbering"
    assert "heading_recognition" in mod.meta.depends_on
    assert "heading_map" in mod.meta.consumes
    print("  ✅ M06 heading_numbering 导入正常")


def test_heading_numbering_logic():
    """M06: 模板化编号生成逻辑测试"""
    from src.modules.structure.heading_numbering import _format_level_number
    from src.config.template import HeadingLevelBindingConfig

    # --- 单级 chain ---
    # 第一章 (display_core_style auto-expand)
    b = HeadingLevelBindingConfig(
        enabled=True, display_core_style="chinese_chapter",
        chain="current_only", title_separator="\u3000",
    )
    counters = [0, 1, 0, 0, 0, 0, 0, 0, 0, 0]
    result = _format_level_number(1, counters, b)
    assert result == "第一章\u3000", f"Expected '第一章\\u3000', got {result!r}"

    # 第1章 (display_template override)
    b = HeadingLevelBindingConfig(
        enabled=True, display_template="第{nn}章",
        chain="current_only", title_separator="\u3000",
    )
    counters = [0, 3, 0, 0, 0, 0, 0, 0, 0, 0]
    result = _format_level_number(1, counters, b)
    assert result == "第3章\u3000", f"Expected '第3章\\u3000', got {result!r}"

    # 第 1 章 (spaces around number)
    b = HeadingLevelBindingConfig(
        enabled=True, display_template="第 {nn} 章",
        chain="current_only", title_separator=" ",
    )
    counters = [0, 5, 0, 0, 0, 0, 0, 0, 0, 0]
    result = _format_level_number(1, counters, b)
    assert result == "第 5 章 ", f"Got {result!r}"

    # (一) (括号中文)
    b = HeadingLevelBindingConfig(
        enabled=True, display_template="({cn})",
        chain="current_only", title_separator=" ",
    )
    counters = [0, 0, 3, 0, 0, 0, 0, 0, 0, 0]
    result = _format_level_number(2, counters, b)
    assert result == "(三) ", f"Got {result!r}"

    # ① 带圈 (auto-expand)
    b = HeadingLevelBindingConfig(
        enabled=True, display_core_style="circled",
        chain="current_only", title_separator=" ",
    )
    counters = [0, 0, 2, 0, 0, 0, 0, 0, 0, 0]
    result = _format_level_number(2, counters, b)
    assert "②" in result, f"Got {result!r}"

    # --- 多级 chain ---
    # 3.1 (parent.current, arabic)
    b = HeadingLevelBindingConfig(
        enabled=True, chain="parent.current",
        chain_separator=".", title_separator="\u3000",
    )
    counters = [0, 3, 1, 0, 0, 0, 0, 0, 0, 0]
    result = _format_level_number(2, counters, b)
    assert result == "3.1\u3000", f"Got {result!r}"

    # 第1-1节 (chain with template wrap)
    b = HeadingLevelBindingConfig(
        enabled=True, display_template="第{chain}节",
        chain="parent.current", chain_separator="-",
        title_separator=" ",
    )
    counters = [0, 1, 1, 0, 0, 0, 0, 0, 0, 0]
    result = _format_level_number(2, counters, b)
    assert result == "第1-1节 ", f"Got {result!r}"

    # 2.3.1 (parent.parent.current)
    b = HeadingLevelBindingConfig(
        enabled=True, chain="parent.parent.current",
        chain_separator=".", title_separator=" ",
    )
    counters = [0, 2, 3, 1, 0, 0, 0, 0, 0, 0]
    result = _format_level_number(3, counters, b)
    assert result == "2.3.1 ", f"Got {result!r}"

    print("  ✅ M06 模板化编号 — 所有格式通过")


def test_heading_numbering_logic_uses_reference_core_style_for_parent_segments():
    """M06: 多级 chain 中 parent 段应使用 source level 的 reference_core_style。"""
    from src.modules.structure.heading_numbering import _format_level_number
    from src.config.template import HeadingLevelBindingConfig

    parent = HeadingLevelBindingConfig(
        enabled=True,
        display_template="第{cn}章",
        display_core_style="chinese_lower",
        reference_core_style="roman_lower",
        chain="current_only",
        title_separator=" ",
    )
    child = HeadingLevelBindingConfig(
        enabled=True,
        display_template="{nn}",
        display_core_style="arabic",
        chain="parent.current",
        chain_separator=".",
        title_separator=" ",
    )
    counters = [0, 1, 1, 0, 0, 0, 0, 0, 0, 0]
    level_bindings = {
        "heading1": parent,
        "heading2": child,
    }

    result = _format_level_number(2, counters, child, level_bindings)

    assert result == "i.1 ", f"Expected 'i.1 ', got {result!r}"


def test_should_skip_numbering_edge_cases():
    """M06: _should_skip_numbering 应覆盖常见跳过边界。"""
    from src.modules.structure.heading_numbering import _should_skip_numbering

    doc = Document()
    unnumbered = doc.styles.add_style("Heading 1 Unnumbered", WD_STYLE_TYPE.PARAGRAPH)
    unnumbered.base_style = doc.styles["Heading 1"]
    non_numbered = {"参考文献", "致谢", "摘要", "Abstract"}
    prefixes = ["附录", "附件", "Appendix"]

    # 第1层: Unnumbered 样式
    para = doc.add_paragraph("绪论")
    para.style = unnumbered
    assert _should_skip_numbering(para, non_numbered, prefixes) is True

    # 第2层: 精确匹配
    para = doc.add_paragraph("参考文献")
    para.style = doc.styles["Heading 1"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is True

    para = doc.add_paragraph("1. 参考文献")
    para.style = doc.styles["Heading 1"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is True

    para = doc.add_paragraph("第一章 参考文献")
    para.style = doc.styles["Heading 1"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is True

    para = doc.add_paragraph("A. 致谢")
    para.style = doc.styles["Heading 1"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is True

    # 第3层: 前缀匹配
    para = doc.add_paragraph("附录A 数据表")
    para.style = doc.styles["Heading 1"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is True

    para = doc.add_paragraph("附件1")
    para.style = doc.styles["Heading 1"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is True

    para = doc.add_paragraph("Appendix B")
    para.style = doc.styles["Heading 1"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is True

    # 不应跳过
    para = doc.add_paragraph("1.1 研究背景")
    para.style = doc.styles["Heading 2"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is False

    para = doc.add_paragraph("1.")
    para.style = doc.styles["Heading 1"]
    assert _should_skip_numbering(para, non_numbered, prefixes) is False

    print("  ✅ M06 _should_skip_numbering 双词典边界覆盖通过")


def test_toc_import():
    """M07: toc 导入+实例化"""
    from src.modules.structure.toc import TocModule
    mod = TocModule()
    assert mod.meta.name == "toc"
    assert "heading_numbering" in mod.meta.depends_on
    assert "heading_map" in mod.meta.consumes
    print("  ✅ M07 toc 导入正常")


def test_toc_dirty_marks_only_toc_field():
    """M07: 仅将 TOC 域标记为 dirty，不污染其他域。"""
    from src.modules.structure.toc import _has_existing_toc, _mark_toc_for_update
    from src.shared.engine.field_builder import build_complex_field
    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    toc_para = doc.add_paragraph()
    for elem in build_complex_field(
        ' TOC \\o "1-3" \\h \\z \\u ', result_text="目录",
    ):
        toc_para._element.append(elem)

    page_para = doc.add_paragraph()
    for elem in build_complex_field(" PAGE ", result_text="1"):
        page_para._element.append(elem)

    assert _has_existing_toc(doc) is True
    toc_begin = next(
        elem for elem in toc_para._element.iter(qn("w:fldChar"))
        if elem.get(qn("w:fldCharType")) == "begin"
    )
    page_begin = next(
        elem for elem in page_para._element.iter(qn("w:fldChar"))
        if elem.get(qn("w:fldCharType")) == "begin"
    )
    _mark_toc_for_update(doc)
    assert toc_begin.get(qn("w:dirty")) == "true"
    assert page_begin.get(qn("w:dirty")) is None
    print("  ✅ M07 TOC dirty 标记范围正确")


def test_toc_marks_document_settings_for_update_on_open():
    """M07: TOC 应同时标记 settings/updateFields，提升打开后刷新成功率。"""
    from src.shared.engine.field_refresh import ensure_update_fields_on_open
    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    ensure_update_fields_on_open(doc)

    update_fields = doc.settings.element.find(qn("w:updateFields"))
    assert update_fields is not None
    assert update_fields.get(qn("w:val")) == "true"
    print("  ✅ M07 settings/updateFields 标记正确")


def test_pipeline_requests_post_save_field_refresh_for_toc_output():
    """M07: 含 TOC 输出时，Pipeline 应在保存后尝试字段刷新。"""
    from src.config.template import TemplateConfig
    from src.config.scene import SceneWorkspace
    from src.config.resolver import resolve_config
    from src.pipeline.runner import Pipeline
    from src.shared.engine.field_builder import build_complex_field

    doc = Document()
    para = doc.add_paragraph()
    for elem in build_complex_field(' TOC \\o "1-3" \\h \\z \\u ', result_text="目录"):
        para._element.append(elem)

    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)

    refresh_calls: list[tuple[str, int]] = []

    def _fake_refresh(path: str, timeout_sec: int = 30):
        refresh_calls.append((path, timeout_sec))
        return True, "ok(test)"

    try:
        config = resolve_config(TemplateConfig(), SceneWorkspace())
        pipeline = Pipeline(modules=[], config=config)
        with patch("src.pipeline.runner.refresh_doc_fields_with_word", side_effect=_fake_refresh):
            result = pipeline.execute(tmp.name)

        assert result.success, f"执行失败: {result.error}"
        assert refresh_calls, "保存后未触发字段刷新"
        assert refresh_calls[0][1] == 30
        assert Path(refresh_calls[0][0]).exists()
        print("  ✅ M07 Pipeline 保存后字段刷新调用正确")
    finally:
        Path(tmp.name).unlink(missing_ok=True)
        if "result" in locals():
            out = result.output_paths.get("final", "")
            if out:
                Path(out).unlink(missing_ok=True)


def test_word_field_refresh_can_be_disabled_by_env(monkeypatch):
    from src.shared.engine.field_refresh import refresh_doc_fields_with_word

    monkeypatch.setenv("LARK_DISABLE_WORD_COM_REFRESH", "1")

    ok, detail = refresh_doc_fields_with_word("__missing_output__.docx")

    assert not ok
    assert "disabled" in detail


def test_word_refresh_cleanup_only_terminates_new_automation_processes(monkeypatch):
    from src.shared.engine import field_refresh

    killed: list[set[int]] = []
    monkeypatch.setattr(field_refresh, "_automation_word_process_ids", lambda: {10, 20, 30})
    monkeypatch.setattr(
        field_refresh,
        "_terminate_process_ids",
        lambda process_ids: killed.append(set(process_ids)),
    )

    field_refresh._cleanup_new_automation_word_processes({10, 20})

    assert killed == [{30}]


def test_powershell_refresh_cleans_new_automation_word_on_timeout(monkeypatch, tmp_path):
    from src.shared.engine import field_refresh

    doc_path = tmp_path / "refresh.docx"
    doc_path.write_bytes(b"stub")
    snapshots = iter(({100}, {100, 101}))
    killed: list[set[int]] = []

    monkeypatch.setattr(field_refresh, "_automation_word_process_ids", lambda: next(snapshots))
    monkeypatch.setattr(
        field_refresh,
        "_terminate_process_ids",
        lambda process_ids: killed.append(set(process_ids)),
    )

    def _timeout_run(*args, **kwargs):
        raise field_refresh.subprocess.TimeoutExpired(
            cmd=args[0] if args else (),
            timeout=kwargs.get("timeout"),
        )

    monkeypatch.setattr(field_refresh.subprocess, "run", _timeout_run)

    ok, detail = field_refresh._refresh_via_powershell(doc_path, timeout_sec=1)

    assert not ok
    assert "timed out" in detail
    assert killed == [{101}]


def test_all_structure_meta():
    """验证 3 个模块的 ModuleMeta 完整性"""
    from src.modules.structure.heading_recognition import HeadingRecognitionModule
    from src.modules.structure.heading_numbering import HeadingNumberingModule
    from src.modules.structure.toc import TocModule

    modules = [HeadingRecognitionModule(), HeadingNumberingModule(), TocModule()]
    for mod in modules:
        m = mod.meta
        assert m.name and m.description
        assert m.category == "structure"

    assert "heading_recognition" in HeadingNumberingModule().meta.depends_on
    assert "heading_numbering" in TocModule().meta.depends_on

    names = [m.meta.name for m in modules]
    assert len(names) == len(set(names))
    print(f"  ✅ 3 个模块 meta + 依赖链验证通过: {names}")


if __name__ == "__main__":
    print("Phase 3 冒烟测试 — 🔢 结构与编号")
    print("=" * 50)
    test_heading_recognition_pipeline()
    test_heading_numbering_import()
    test_heading_numbering_logic()
    test_toc_import()
    test_all_structure_meta()
    print("=" * 50)
    print("✅ 结构与编号 3 个模块全部通过！")
