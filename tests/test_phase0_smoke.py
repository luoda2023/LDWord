"""Phase 0 smoke tests for pipeline skeleton behavior."""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from src.config.document_scope import DocumentScopePolicy
from src.modules.base import BaseModule, ModuleMeta
from src.modules.registry import create_all_modules
from src.pipeline.context import PipelineContext
from src.pipeline.runner import Pipeline
from src.pipeline.module_selection import (
    ModuleDisposition,
    build_module_selection_plan,
)
from src.pipeline.scheduler import (
    compute_dirty_modules,
    topological_sort,
    validate_data_flow,
    validate_schema_contract,
)


class StubPageSetup(BaseModule):
    meta = ModuleMeta(
        name="page_setup",
        description="Page setup",
        category="basic",
        execution_phase="format",
        requires_config=("page_setup",),
    )

    def apply(self, doc, config, tracker, context):
        tracker.record(
            rule_name=self.meta.name,
            target="page_margins",
            section="global",
            change_type="format",
            before="default",
            after="A4 standard",
        )


class StubHeadingRecognition(BaseModule):
    meta = ModuleMeta(
        name="heading_recognition",
        description="Heading recognition",
        category="structure",
        execution_phase="semantics",
        provides=("doc_tree", "heading_map"),
    )

    def apply(self, doc, config, tracker, context):
        context.doc_tree = {"mock": True}
        context.heading_map = {0: 1}
        tracker.record(
            rule_name=self.meta.name,
            target="headings",
            section="global",
            change_type="detect",
            after="found headings",
        )


class StubHeadingNumbering(BaseModule):
    meta = ModuleMeta(
        name="heading_numbering",
        description="Heading numbering",
        category="structure",
        execution_phase="semantics",
        depends_on=("heading_recognition",),
        consumes=("doc_tree", "heading_map"),
    )

    def apply(self, doc, config, tracker, context):
        assert context.doc_tree is not None
        assert context.heading_map == {0: 1}
        tracker.record(
            rule_name=self.meta.name,
            target="numbering",
            section="global",
            change_type="format",
            after="applied numbering",
        )


class StubCyclicA(BaseModule):
    meta = ModuleMeta(
        name="cyclic_a",
        description="Cycle A",
        category="test",
        execution_phase="fill",
        scope_behavior="document_level",
        depends_on=("cyclic_b",),
    )

    def apply(self, doc, config, tracker, context):
        return None


class StubCyclicB(BaseModule):
    meta = ModuleMeta(
        name="cyclic_b",
        description="Cycle B",
        category="test",
        execution_phase="fill",
        scope_behavior="document_level",
        depends_on=("cyclic_a",),
    )

    def apply(self, doc, config, tracker, context):
        return None


class StubFailingModule(BaseModule):
    meta = ModuleMeta(
        name="zz_fail",
        description="Intentional failure",
        category="test",
        execution_phase="format",
        scope_behavior="document_level",
    )

    def apply(self, doc, config, tracker, context):
        raise RuntimeError("boom")


def test_topological_sort_orders_dependencies():
    modules = [StubHeadingNumbering(), StubPageSetup(), StubHeadingRecognition()]
    names = [mod.meta.name for mod in topological_sort(modules)]
    assert names.index("heading_recognition") < names.index("heading_numbering")


def test_topological_sort_detects_cycle():
    try:
        topological_sort([StubCyclicA(), StubCyclicB()])
        assert False, "expected ValueError"
    except ValueError:
        pass


def test_validate_data_flow_success_and_failure():
    ordered = topological_sort([StubHeadingRecognition(), StubHeadingNumbering()])
    assert not validate_data_flow(ordered)
    assert validate_data_flow([StubHeadingNumbering()])


def test_compute_dirty_modules_uses_requires_config():
    modules = topological_sort([
        StubPageSetup(),
        StubHeadingRecognition(),
        StubHeadingNumbering(),
    ])
    dirty = compute_dirty_modules({"page_setup"}, modules)
    assert dirty == {"page_setup"}


def test_schema_contract_validation_for_registry_modules():
    errors = validate_schema_contract(create_all_modules())
    assert not errors, f"schema contract errors: {errors}"


def test_pipeline_execute_runs_basic_chain():
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    Document().save(tmp_path)

    try:
        pipeline = Pipeline(
            modules=[StubHeadingNumbering(), StubPageSetup(), StubHeadingRecognition()],
            config=SimpleNamespace(strict_mode=False),
        )
        result = pipeline.execute(tmp_path)

        assert result.success
        assert result.status == "success"
        assert result.tracker is not None
        assert result.tracker.total_count == 3
        assert result.tracker.failure_count == 0
        assert "final" in result.output_paths
    finally:
        Path(tmp_path).unlink(missing_ok=True)
        if 'result' in locals():
            Path(result.output_paths.get("final", "")).unlink(missing_ok=True)


def test_pipeline_progress_emits_user_flow_preflight_and_report_steps():
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    Document().save(tmp_path)
    progress: list[tuple[int, int, str]] = []

    try:
        pipeline = Pipeline(
            modules=[],
            config=SimpleNamespace(strict_mode=False),
            progress_callback=lambda cur, total, stage: progress.append(
                (cur, total, stage)
            ),
        )
        result = pipeline.execute(tmp_path)

        assert result.success
        stages = [stage for _cur, _total, stage in progress]
        assert stages == [
            "Loading document",
            "Object preflight",
            "Writing report",
            "Saving output",
            "Completed",
        ]
    finally:
        Path(tmp_path).unlink(missing_ok=True)
        if "result" in locals():
            Path(result.output_paths.get("final", "")).unlink(missing_ok=True)


def test_pipeline_cancelled_result():
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    Document().save(tmp_path)

    try:
        pipeline = Pipeline(
            modules=[StubPageSetup()],
            config=SimpleNamespace(strict_mode=False),
            cancel_check=lambda: True,
        )
        result = pipeline.execute(tmp_path)
        assert result.cancelled
        assert result.status == "cancelled"
    finally:
        Path(tmp_path).unlink(missing_ok=True)


def test_pipeline_custom_output_path():
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    Document().save(tmp_path)

    out_dir = Path(tempfile.mkdtemp())

    try:
        pipeline = Pipeline(
            modules=[StubPageSetup()],
            config=SimpleNamespace(strict_mode=False),
            output_dir=str(out_dir),
            output_suffix="_formatted",
        )
        result = pipeline.execute(tmp_path)
        final_path = Path(result.output_paths["final"])

        assert result.success
        assert final_path.parent == out_dir
        assert final_path.name.endswith("_formatted.docx")
    finally:
        Path(tmp_path).unlink(missing_ok=True)
        if 'result' in locals():
            Path(result.output_paths.get("final", "")).unlink(missing_ok=True)


def test_pipeline_fails_closed_before_running_an_invalid_document_scope():
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    Document().save(tmp_path)

    try:
        pipeline = Pipeline(
            modules=[],
            config=SimpleNamespace(
                strict_mode=True,
                mode_id="custom",
                document_scope=DocumentScopePolicy(mode="future"),
            ),
        )
        result = pipeline.execute(tmp_path)

        assert result.success is False
        assert result.status == "failed"
        assert result.context is None
        assert result.output_paths == {}
        assert result.error == "document_scope_mode_invalid:future"
    finally:
        Path(tmp_path).unlink(missing_ok=True)
        if "result" in locals() and result.output_paths.get("final"):
            Path(result.output_paths["final"]).unlink(missing_ok=True)


def test_module_selection_plan_prunes_unmet_hard_dependencies():
    from src.config.template import TemplateConfig
    from src.config.scene import SceneWorkspace
    from src.config.resolver import resolve_config

    scene = SceneWorkspace()
    scene.module_switches["heading_recognition"] = False
    config = resolve_config(TemplateConfig(), scene)

    selection = build_module_selection_plan(
        create_all_modules(),
        config.is_module_enabled,
    )
    enabled = list(selection.select_modules(create_all_modules()))
    names = [mod.meta.name for mod in enabled]

    assert "heading_recognition" not in names
    assert "heading_numbering" not in names
    assert "toc" not in names
    toc_decision = selection.decision_for("toc")
    assert toc_decision.disposition is ModuleDisposition.AUTO_PRUNED
    assert toc_decision.unmet_dependencies == ("heading_recognition",)

    pipeline = Pipeline(modules=enabled, config=config)
    assert set(mod.meta.name for mod in pipeline._modules) == set(names)


def test_toc_stays_enabled_when_heading_numbering_is_disabled():
    from src.config.template import TemplateConfig
    from src.config.scene import SceneWorkspace
    from src.config.resolver import resolve_config

    scene = SceneWorkspace()
    scene.module_switches["heading_numbering"] = False
    config = resolve_config(TemplateConfig(), scene)
    modules = create_all_modules()

    selection = build_module_selection_plan(modules, config.is_module_enabled)

    assert selection.is_effectively_enabled("heading_recognition") is True
    assert selection.is_effectively_enabled("heading_numbering") is False
    assert selection.is_effectively_enabled("toc") is True


def test_pipeline_context_slots_are_enforced():
    context = PipelineContext()
    try:
        context.chapter_counter = {"1": 1}
        assert False, "PipelineContext should reject undeclared attributes"
    except AttributeError:
        pass


def test_non_strict_failure_degrades_to_partial_success():
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    Document().save(tmp_path)

    try:
        pipeline = Pipeline(
            modules=[StubPageSetup(), StubFailingModule()],
            config=SimpleNamespace(strict_mode=False),
        )
        result = pipeline.execute(tmp_path)

        assert result.success is True
        assert result.status == "partial_success"
        assert len(result.failed_items) == 1
    finally:
        Path(tmp_path).unlink(missing_ok=True)
        if 'result' in locals():
            Path(result.output_paths.get("final", "")).unlink(missing_ok=True)


def test_strict_mode_failure_stops_pipeline():
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    Document().save(tmp_path)

    try:
        pipeline = Pipeline(
            modules=[StubPageSetup(), StubFailingModule()],
            config=SimpleNamespace(strict_mode=True),
        )
        result = pipeline.execute(tmp_path)

        assert result.success is False
        assert result.status == "failed"
        assert result.failed_items
        assert result.failed_items[0]["rule_name"] == "zz_fail"
    finally:
        Path(tmp_path).unlink(missing_ok=True)
