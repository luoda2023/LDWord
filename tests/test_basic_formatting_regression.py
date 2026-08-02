from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.builtin_templates import create_builtin_template
from src.config.heading_style_semantics import resolve_heading_style
from src.config.resolver import resolve_config
from src.config.scene import SceneWorkspace
from src.config.style_semantics import resolve_style_size_pt
from src.modules.basic.paragraph_style import ParagraphStyleModule
from src.modules.structure.heading_numbering import HeadingNumberingModule
from src.modules.validate.validation import ValidationModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.heading_numbering_ooxml import (
    effective_numbering,
    link_style_to_numbering,
    managed_abstract_count,
    managed_level_text,
)


def _default_config():
    return resolve_config(create_builtin_template("default"), SceneWorkspace())


def test_default_template_restores_v02_body_and_heading_size_ladder():
    config = _default_config()

    assert resolve_style_size_pt(config.styles["body"]) == 12.0
    assert [
        resolve_style_size_pt(
            resolve_heading_style(config.styles, level, include_body_fallback=True)
        )
        for level in range(1, 5)
    ] == [16.0, 14.0, 13.0, 12.0]


def test_paragraph_style_uses_heading_map_and_body_contract():
    document = Document()
    body = document.add_paragraph("正文内容")
    title = document.add_paragraph("章节标题")
    context = PipelineContext(heading_map={1: 1})

    ParagraphStyleModule().apply(
        document,
        _default_config(),
        ChangeTracker(),
        context,
    )

    assert body.style.name == "Normal"
    assert body.runs[0].font.size.pt == 12.0
    assert document.styles["Normal"].font.size.pt == 12.0
    assert title.style.name == "Heading 1"
    assert title.runs[0].font.size.pt == 16.0


def test_native_heading_numbering_has_one_owner_and_is_idempotent():
    document = Document()
    empty_heading = document.add_paragraph("")
    empty_heading.style = document.styles["Heading 1"]
    first = document.add_paragraph("第一章 概述")
    second = document.add_paragraph("2.1 规划背景")
    body = document.add_paragraph("正文内容")
    first.style = document.styles["Heading 1"]
    second.style = document.styles["Heading 2"]

    # Simulate an input document whose Heading styles already carry numbering.
    link_style_to_numbering(document.styles["Heading 1"], 42, 1)
    link_style_to_numbering(document.styles["Heading 2"], 42, 2)

    config = _default_config()
    context = PipelineContext(heading_map={1: 1, 2: 2})
    numbering = HeadingNumberingModule()
    paragraph_style = ParagraphStyleModule()

    for _ in range(2):
        numbering.apply(document, config, ChangeTracker(), context)
        paragraph_style.apply(document, config, ChangeTracker(), context)

    assert first.text == "概述"
    assert second.text == "规划背景"
    assert body.text == "正文内容"
    assert effective_numbering(empty_heading) is None
    first_numbering = effective_numbering(first)
    second_numbering = effective_numbering(second)
    assert first_numbering is not None
    assert second_numbering is not None
    assert first_numbering.source == "style:Heading 1"
    assert second_numbering.source == "style:Heading 2"
    assert first_numbering.num_id == second_numbering.num_id
    assert (first_numbering.level, second_numbering.level) == (0, 1)
    assert managed_abstract_count(document) == 1
    assert managed_level_text(document, 1) == "第%1章　"
    assert managed_level_text(document, 2) == "%1.%2　"
    assert first.runs[0].font.size.pt == 16.0
    assert second.runs[0].font.size.pt == 14.0
    assert body.runs[0].font.size.pt == 12.0


def test_validation_rejects_literal_and_native_double_numbering():
    document = Document()
    paragraph = document.add_paragraph("第一章 概述")
    paragraph.style = document.styles["Heading 1"]
    config = _default_config()
    context = PipelineContext(heading_map={0: 1})

    # Deliberately keep the literal prefix while enabling native numbering.
    link_style_to_numbering(document.styles["Heading 1"], 42, 1)
    issues = ValidationModule().validate(document, config, context)

    assert any(
        issue.level == "error" and "重复编号" in issue.message
        for issue in issues
    )
