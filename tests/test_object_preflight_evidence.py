import json
from dataclasses import FrozenInstanceError
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pytest

from src.config import object_preflight_evidence as evidence_module
from src.config.object_preflight_evidence import (
    build_object_preflight_evidence,
    object_preflight_evidence_applicable,
    object_preflight_evidence_is_current,
)
from src.config.scene import SceneWorkspace
from src.shared.engine import object_preflight as object_preflight_module


def test_non_docx_or_disabled_policy_is_not_applicable(tmp_path, monkeypatch):
    text_source = tmp_path / "source.txt"
    text_source.write_text("plain text", encoding="utf-8")
    docx_source = _docx_with_parts(tmp_path, "disabled.docx", {})
    scene = SceneWorkspace(scene_id="custom")

    def _unexpected_call(*_args, **_kwargs):
        raise AssertionError("non-applicable evidence must not read or inspect")

    monkeypatch.setattr(evidence_module, "file_content_revision", _unexpected_call)
    monkeypatch.setattr(evidence_module, "inspect_docx_package", _unexpected_call)

    text_evidence = build_object_preflight_evidence(scene, text_source)
    scene.compliance_profile.object_preflight.enabled = False
    disabled_evidence = build_object_preflight_evidence(scene, docx_source)

    assert object_preflight_evidence_applicable(scene, text_source) is False
    assert object_preflight_evidence_applicable(scene, docx_source) is False
    for evidence in (text_evidence, disabled_evidence):
        assert evidence.applicable is False
        assert evidence.payload == {}
        assert evidence.source_revision == ""
        assert evidence.evidence_digest == ""


def test_no_findings_builds_ui_equivalent_immutable_evidence(tmp_path):
    source = _docx_with_parts(tmp_path, "clean.docx", {})
    scene = SceneWorkspace(scene_id="custom")
    scene.compliance_profile.object_preflight.scan_targets = ["comments"]

    evidence = build_object_preflight_evidence(scene, source)
    payload = evidence.payload
    expected_key = json.dumps(
        {
            "doc_path": str(source.resolve()),
            "scene_id": "custom",
            "object_preflight": payload,
        },
        ensure_ascii=False,
        sort_keys=True,
    )

    assert evidence.applicable is True
    assert evidence.blocked is False
    assert payload["enabled"] is True
    assert payload["findings_count"] == 0
    assert payload["module_skips_count"] == 0
    assert evidence.source_revision.startswith("sha256:")
    assert evidence.canonical_key == expected_key
    assert evidence.evidence_digest == (
        "sha256:" + sha256(expected_key.encode("utf-8")).hexdigest()
    )

    payload["findings"].append({"kind": "caller_mutation"})
    assert evidence.payload["findings"] == []
    with pytest.raises(FrozenInstanceError):
        evidence.scene_id = "mutated"


def test_embedded_part_warning_does_not_disable_safe_section_planner(tmp_path):
    source = _docx_with_parts(
        tmp_path,
        "warning.docx",
        {"word/embeddings/oleObject1.bin": b"ole"},
    )
    scene = SceneWorkspace(scene_id="custom")
    scene.strict_mode = False
    policy = scene.compliance_profile.object_preflight
    policy.preservation_mode = "warn"
    policy.scan_targets = ["ole_objects"]
    policy.block_on = ["ole_objects"]
    policy.skip_modules_by_finding = {"ole_objects": ["section_format"]}

    evidence = build_object_preflight_evidence(scene, source)
    payload = evidence.payload

    assert payload["findings_count"] == 1
    assert payload["findings"][0]["kind"] == "ole_objects"
    assert payload["findings"][0]["severity"] == "warning"
    assert payload["blocking_findings_count"] == 0
    assert payload["blocked"] is False
    assert payload["module_skips"] == []


@pytest.mark.parametrize(
    ("strict_mode", "failure_policy"),
    [(True, "warn"), (False, "block")],
)
def test_strict_finding_blocks_for_strict_scene_or_block_policy(
    tmp_path,
    strict_mode,
    failure_policy,
):
    source = _docx_with_parts(
        tmp_path,
        f"blocked-{strict_mode}-{failure_policy}.docx",
        {"word/embeddings/oleObject1.bin": b"ole"},
    )
    scene = SceneWorkspace(scene_id="custom")
    scene.strict_mode = strict_mode
    scene.compliance_profile.failure_policy = failure_policy
    policy = scene.compliance_profile.object_preflight
    policy.preservation_mode = "strict"
    policy.scan_targets = ["ole_objects"]
    policy.block_on = ["ole_objects"]

    payload = build_object_preflight_evidence(scene, source).payload

    assert payload["findings"][0]["severity"] == "error"
    assert payload["blocking_findings_count"] == 1
    assert payload["blocked"] is True


def test_source_revision_drift_replaces_scan_result_with_blocker(
    tmp_path,
    monkeypatch,
):
    source = _docx_with_parts(tmp_path, "drift.docx", {})
    scene = SceneWorkspace(scene_id="custom")
    revisions = iter(["sha256:" + "a" * 64, "sha256:" + "b" * 64])
    monkeypatch.setattr(
        evidence_module,
        "file_content_revision",
        lambda _path: next(revisions),
    )

    evidence = build_object_preflight_evidence(scene, source)
    payload = evidence.payload

    assert evidence.source_revision == "sha256:" + "b" * 64
    assert payload["findings_count"] == 1
    assert payload["findings"][0]["kind"] == "source_changed_during_preflight"
    assert payload["findings"][0]["severity"] == "error"
    assert payload["blocked"] is True
    assert payload["module_skips"] == []


def test_source_read_failure_is_blocked(tmp_path, monkeypatch):
    source = _docx_with_parts(tmp_path, "unreadable.docx", {})
    scene = SceneWorkspace(scene_id="custom")

    def _permission_denied(_path):
        raise PermissionError("denied")

    monkeypatch.setattr(
        evidence_module,
        "file_content_revision",
        _permission_denied,
    )

    evidence = build_object_preflight_evidence(scene, source)
    payload = evidence.payload

    assert evidence.applicable is True
    assert evidence.source_revision == ""
    assert evidence.evidence_digest.startswith("sha256:")
    assert payload["findings"][0]["kind"] == "source_read_failed_during_preflight"
    assert payload["blocking_findings_count"] == 1
    assert payload["blocked"] is True


def test_invalid_docx_package_is_an_inspection_blocker(tmp_path):
    source = tmp_path / "invalid.docx"
    source.write_bytes(b"not a zip package")
    scene = SceneWorkspace(scene_id="custom")
    scene.strict_mode = False
    scene.compliance_profile.failure_policy = "warn"

    evidence = build_object_preflight_evidence(scene, source)
    payload = evidence.to_payload()

    assert evidence.applicable is True
    assert evidence.blocked is True
    assert payload["inspection_status"] == "invalid_docx_package"
    assert payload["blocking_findings_count"] == 1
    assert payload["findings"] == [
        {
            "kind": "object_preflight_inspection_failed",
            "location": str(source.resolve()),
            "message": (
                "Object preflight could not inspect a valid DOCX package "
                "(invalid_docx_package)."
            ),
            "severity": "error",
        }
    ]


def test_zip_without_required_docx_parts_is_an_inspection_blocker(tmp_path):
    source = tmp_path / "not-docx.docx"
    with ZipFile(source, "w") as package:
        package.writestr("arbitrary.txt", b"not a Word package")
    scene = SceneWorkspace(scene_id="custom")

    payload = build_object_preflight_evidence(scene, source).to_payload()

    assert payload["inspection_status"] == "invalid_docx_package"
    assert payload["blocked"] is True
    assert payload["findings"][0]["kind"] == "object_preflight_inspection_failed"


def test_malformed_core_xml_is_an_inspection_blocker(tmp_path):
    source = tmp_path / "malformed.docx"
    with ZipFile(source, "w") as package:
        package.writestr(
            "[Content_Types].xml",
            (
                b'<Types xmlns="http://schemas.openxmlformats.org/package/'
                b'2006/content-types" />'
            ),
        )
        package.writestr(
            "_rels/.rels",
            (
                b'<Relationships xmlns="http://schemas.openxmlformats.org/'
                b'package/2006/relationships" />'
            ),
        )
        package.writestr("word/document.xml", b"<w:document")
    scene = SceneWorkspace(scene_id="custom")

    evidence = build_object_preflight_evidence(scene, source)

    assert evidence.blocked is True
    assert evidence.to_payload()["inspection_status"] == "invalid_docx_package"


def test_package_read_runtime_error_becomes_inspection_blocker(
    tmp_path,
    monkeypatch,
):
    source = _docx_with_parts(tmp_path, "unreadable-package.docx", {})

    def _raise_unreadable(*_args, **_kwargs):
        raise RuntimeError("encrypted or unsupported package")

    monkeypatch.setattr(object_preflight_module, "ZipFile", _raise_unreadable)

    result = object_preflight_module.inspect_docx_package(source)

    assert result.inspection_status == "source_unreadable"
    assert result.inspection_succeeded is False
    assert result.blocking_findings[0].kind == "object_preflight_inspection_failed"


def test_cached_evidence_validity_binds_source_and_policy(tmp_path):
    source = _docx_with_parts(tmp_path, "cached.docx", {})
    scene = SceneWorkspace(scene_id="custom")
    scene.compliance_profile.object_preflight.scan_targets = ["comments"]
    evidence = build_object_preflight_evidence(scene, source)

    assert object_preflight_evidence_is_current(evidence, scene, source) is True

    scene.compliance_profile.object_preflight.scan_targets = ["fields"]
    assert object_preflight_evidence_is_current(evidence, scene, source) is False

    scene.compliance_profile.object_preflight.scan_targets = ["comments"]
    document = Document(source)
    document.add_paragraph("changed")
    document.save(source)
    assert object_preflight_evidence_is_current(evidence, scene, source) is False


def test_digest_is_stable_and_changes_with_bound_evidence(tmp_path):
    source = _docx_with_parts(tmp_path, "stable.docx", {})
    scene = SceneWorkspace(scene_id="custom")

    first = build_object_preflight_evidence(scene, source)
    second = build_object_preflight_evidence(scene, str(source.resolve()))

    assert second.payload == first.payload
    assert second.canonical_key == first.canonical_key
    assert second.evidence_digest == first.evidence_digest

    scene.scene_id = "another-scene"
    changed_scene = build_object_preflight_evidence(scene, source)
    assert changed_scene.payload == first.payload
    assert changed_scene.evidence_digest != first.evidence_digest

    scene.scene_id = "custom"
    document = Document(source)
    document.add_paragraph("content revision changed")
    document.save(source)
    changed_source = build_object_preflight_evidence(scene, source)
    assert changed_source.source_revision != first.source_revision
    assert changed_source.evidence_digest != first.evidence_digest


def _docx_with_parts(
    tmp_path: Path,
    filename: str,
    parts: dict[str, bytes],
) -> Path:
    source = tmp_path / filename
    Document().save(source)
    if parts:
        with ZipFile(source, "a") as package:
            for name, data in parts.items():
                package.writestr(name, data)
    return source
