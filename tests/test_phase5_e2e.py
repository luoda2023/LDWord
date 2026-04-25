"""
Phase 5 集成测试 — 全链 E2E

1. 注册表完整性 (22 模块)
2. 拓扑排序验证
3. 数据流验证
4. 全链 Pipeline 执行（真实 docx）
5. 增量执行（脏模块计算）
"""

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document


def test_registry_completeness():
    """注册表包含 22 个模块且名称唯一"""
    from src.modules.registry import ALL_MODULES, create_all_modules, list_module_names

    assert len(ALL_MODULES) == 22, f"期望 22 个模块, 实际 {len(ALL_MODULES)}"

    modules = create_all_modules()
    assert len(modules) == 22

    names = list_module_names()
    assert len(names) == 22
    assert len(set(names)) == 22, f"名称不唯一: {names}"

    # 检查分类覆盖
    categories = set(m.meta.category for m in modules)
    expected = {"basic", "structure", "table", "fill", "insert", "special", "validate"}
    assert categories == expected, f"分类不完整: {categories} vs {expected}"
    print("  registry complete: 22 modules, 7 categories")


def test_topological_sort():
    """拓扑排序不抛异常 + 依赖顺序正确"""
    from src.modules.registry import create_all_modules
    from src.pipeline.scheduler import topological_sort

    modules = create_all_modules()
    sorted_mods = topological_sort(modules)
    sorted_names = [m.meta.name for m in sorted_mods]

    # heading_recognition 必须在 heading_numbering 之前
    assert sorted_names.index("heading_recognition") < sorted_names.index("heading_numbering")
    # heading_numbering 必须在 toc 之前
    assert sorted_names.index("heading_numbering") < sorted_names.index("toc")

    print(f"  ✅ 拓扑排序通过: {len(sorted_names)} 模块")
    print(f"     排序: {' → '.join(sorted_names[:5])} → ... → {sorted_names[-1]}")


def test_data_flow_validation():
    """数据流校验: consumes 的 key 都有 provides 覆盖"""
    from src.modules.registry import create_all_modules
    from src.pipeline.scheduler import topological_sort, validate_data_flow

    modules = create_all_modules()
    sorted_mods = topological_sort(modules)
    errors = validate_data_flow(sorted_mods)

    assert not errors, f"数据流校验失败:\n" + "\n".join(f"  - {e}" for e in errors)
    print("  ✅ 数据流校验通过 (所有 consumes 都有 provides 覆盖)")


def test_dirty_module_propagation():
    """增量执行: 配置变化 → 脏模块传播"""
    from src.modules.registry import create_all_modules
    from src.pipeline.scheduler import topological_sort, compute_dirty_modules

    modules = create_all_modules()
    sorted_mods = topological_sort(modules)

    # 改变 heading_numbering 配置 → heading_numbering 直接变脏
    dirty = compute_dirty_modules({"heading_numbering"}, sorted_mods)
    assert "heading_numbering" in dirty
    print(f"  ✅ 脏模块传播: heading_numbering → {sorted(dirty)}")


def test_full_pipeline_e2e():
    """全链 E2E: 22 模块管线执行"""
    from src.modules.registry import create_all_modules
    from src.config.template import TemplateConfig
    from src.config.scene import SceneWorkspace
    from src.config.resolver import resolve_config
    from src.pipeline.runner import Pipeline

    # 创建测试文档
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容段落。这是第一段。")
    doc.add_paragraph("图 1 实验装置示意图")
    doc.add_paragraph("正文内容段落。这是第二段。")
    doc.add_heading("第二章 方法", level=2)
    doc.add_paragraph("表 1 实验数据统计")
    doc.add_paragraph("正文内容段落。这是第三段。")
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[1] 某某. 某某研究[J]. 某期刊, 2024.")
    doc.add_paragraph("[2] 某某. 某某综述[J]. 某期刊, 2023.")

    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)

    try:
        config = resolve_config(TemplateConfig(), SceneWorkspace())
        modules = create_all_modules()

        pipeline = Pipeline(modules=modules, config=config)
        result = pipeline.execute(tmp.name)

        assert result.success, f"管线执行失败: {result.error}"
        assert result.context is not None, "context 未传透"
        assert result.doc is not None, "doc 为空"

        # 验证 context 中有模块产出
        ctx = result.context
        assert ctx.doc_tree is not None, "doc_tree 未生成"
        assert ctx.heading_map is not None, "heading_map 未生成"

        print(f"  ✅ 全链 E2E 通过: {result.status}")
        print(f"     doc_tree 节点: {len(ctx.doc_tree) if isinstance(ctx.doc_tree, list) else 'dict'}")
        print(f"     heading_map: {len(ctx.heading_map)} 个标题")
        if result.failed_items:
            print(f"     ⚠️ {len(result.failed_items)} 个非关键失败")
    finally:
        Path(tmp.name).unlink(missing_ok=True)
        out = result.output_paths.get("final", "")
        if out:
            Path(out).unlink(missing_ok=True)


if __name__ == "__main__":
    print("Phase 5 集成测试 — 🔗 全链 E2E")
    print("=" * 50)

    test_registry_completeness()
    test_topological_sort()
    test_data_flow_validation()
    test_dirty_module_propagation()
    test_full_pipeline_e2e()

    print("=" * 50)
    print("✅ Phase 5 E2E 全链集成测试通过！")
