from __future__ import annotations

from docx import Document

from src.application import exam_user_plan_creation as creation
from src.config.library import ConfigLibraryEntry


def _source_exam(path):
    document = Document()
    document.add_heading("七年级语文期中测试卷", level=1)
    document.add_paragraph("考试时间：90分钟 满分：100分")
    document.add_paragraph("一、积累与运用")
    document.add_paragraph("1. 选择正确的一项。")
    document.add_paragraph("A. 甲")
    document.add_paragraph("答案解析部分")
    document.add_paragraph("1. A")
    document.save(path)


def test_create_user_plan_publishes_master_then_create_only_scene(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    _source_exam(source)
    master_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    monkeypatch.setattr(creation, "USER_EXAM_MASTER_DIR", master_dir)
    monkeypatch.setattr(creation, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(creation, "allocate_scene_id", lambda **_kwargs: "my-plan")
    saved = {}

    def fake_save(scene, scene_id, *, mode_id, expected_absent):
        saved["scene"] = scene
        saved["expected_absent"] = expected_absent
        saved["master_exists_at_commit"] = next(master_dir.glob("*.docx")).is_file()
        return ConfigLibraryEntry(
            "scene",
            scene_id,
            scene.name,
            tmp_path / f"{scene_id}.json",
            mode_id=mode_id,
            source_type="user",
        )

    monkeypatch.setattr(creation, "save_scene_to_library", fake_save)

    result = creation.create_exam_user_plan_from_docx(
        source,
        plan_name="学校七年级语文期中方案",
        preview_output_dir=tmp_path / "preview",
    )

    assert result.plan_id == "my-plan"
    assert result.preflight_status == "ok"
    assert result.master_path.is_file()
    assert result.sample_path is not None and result.sample_path.is_file()
    assert saved["expected_absent"] is True
    assert saved["master_exists_at_commit"] is True
    scene = saved["scene"]
    assert scene.master_id == result.master_id
    assert scene.template_id == "default"
    assert scene.exam_paper.custom_blank_styles[-1].style_id == result.master_id
    assert scene.exam_paper.custom_blank_styles[-1].master_docx_path.startswith(
        "config_library"
    )


def test_failed_scene_commit_removes_only_new_master(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    _source_exam(source)
    master_dir = tmp_path / "masters"
    monkeypatch.setattr(creation, "USER_EXAM_MASTER_DIR", master_dir)
    monkeypatch.setattr(creation, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(creation, "allocate_scene_id", lambda **_kwargs: "failed-plan")

    def fail_save(*_args, **_kwargs):
        raise RuntimeError("scene commit failed")

    monkeypatch.setattr(creation, "save_scene_to_library", fail_save)

    try:
        creation.create_exam_user_plan_from_docx(source, plan_name="失败方案")
    except RuntimeError as exc:
        assert str(exc) == "scene commit failed"
    else:  # pragma: no cover
        raise AssertionError("scene commit should have failed")

    assert list(master_dir.glob("*.docx")) == []
