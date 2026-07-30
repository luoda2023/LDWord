from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from docx import Document
import pytest

from src.config.content_materials import (
    ContentInsertionRule,
    content_anchor_token,
)
from src.config.entity import EntityArchive, EntityProfile
from src.config.image_materials import (
    ImageCoLocationGuard,
    ImageMaterialRule,
    ImagePlacementMode,
    ImagePlacementPolicy,
    ImageWatermarkPolicy,
    ImageWatermarkTextSource,
)
from src.config.material_schema_registry import MaterialAssetRoleSpec
from src.config.materials import AssetItem
from src.config.library import CONFIG_LIBRARY_ROOT
from src.qt_api import QApplication, QFileDialog
from src.services.material_attachments import build_single_attachment_binding
from content_artifact_test_utils import compile_content_binding
from src.services.material_content.artifact_repository import ContentArtifactRepository
from src.ui.bridge import PanelBridge
import src.ui.panels.assets.content_materials_presenter as content_materials_presenter
from src.ui.panels.assets.specs import AttachmentRoleSpec
from src.ui.panels.assets_panel import AssetsPanel


def _app() -> QApplication:
    return QApplication.instance() or QApplication([])


def _repository_snapshot(root: Path) -> tuple[tuple[str, str], ...]:
    if not root.exists():
        return ()
    return tuple(
        (path.relative_to(root).as_posix(), sha256(path.read_bytes()).hexdigest())
        for path in sorted(root.rglob("*"))
        if path.is_file()
    )


@pytest.fixture(autouse=True)
def _application_content_repository_stays_unchanged():
    root = CONFIG_LIBRARY_ROOT / "content_artifacts"
    before = _repository_snapshot(root)
    yield
    assert _repository_snapshot(root) == before


def _context_image_rule() -> ImageMaterialRule:
    return ImageMaterialRule(
        rule_id="image:logo",
        source_role="logo",
        anchor_token="{{@img:LOGO1}}",
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIXED_BOX,
            fixed_width_cm=6.0,
        ),
        required=False,
        watermark=ImageWatermarkPolicy(
            enabled=True,
            text_template="context watermark",
        ),
    )


def test_material_domains_have_separate_pages_and_attachment_is_not_in_image_card():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        primary_ids = tuple(panel._section_nav_cards)
        assert primary_ids == (
            "generate",
            "fields",
            "content",
            "timeline",
            "images",
            "attachments",
        )
        assert panel._content_card.parent() is panel._section_contents["content"]
        assert panel._image_rules_card.parent() is panel._section_contents["images"]
        assert panel._image_card.parent() is panel._section_contents["images"]
        assert panel._image_material_rules_container.parent() is panel._image_rules_card
        assert not panel._image_card.isAncestorOf(
            panel._image_material_rules_container
        )
        assert panel._attachment_card.parent() is panel._section_contents["attachments"]
        assert not hasattr(panel, "_attachment_inventory_label")
        assert panel._independent_attachments_title.text() == "独立附件"
        assert panel._attachment_folders_title.text() == "附件文件夹"
        assert panel._add_independent_attachment_btn.text() == "＋ 新增独立附件"
        assert panel._add_attachment_folder_btn.text() == "＋ 新增附件文件夹"
        assert panel._independent_attachments_body.parent() is panel._attachment_card
        assert panel._attachment_folders_body.parent() is panel._attachment_card
    finally:
        panel.close()


def test_material_repair_focus_opens_the_domain_that_owns_the_target():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        assert panel._focus_asset_slot("logo") is True
        assert panel._active_section_id == "images"

        panel._attachment_role_specs = (
            AttachmentRoleSpec(
                role="evidence",
                label="证明材料",
                accepted_types=("pdf",),
            ),
        )
        panel._sync_attachment_role_rows()

        assert panel._focus_attachment_role("evidence") is True
        assert panel._active_section_id == "attachments"
    finally:
        panel.close()


def test_schema_role_domain_is_explicit_and_never_inferred_from_file_formats():
    attachment = MaterialAssetRoleSpec(
        "evidence",
        "证明材料",
        accepted_types=("image", "pdf", "docx"),
        material_domain="attachment",
    )

    assert attachment.is_attachment is True
    assert attachment.accepted_types[0] == "image"
    try:
        MaterialAssetRoleSpec(
            "ambiguous",
            "歧义角色",
            accepted_types=("image", "pdf"),
        )
    except ValueError as exc:
        assert "image material roles" in str(exc)
    else:
        raise AssertionError("mixed formats require material_domain='attachment'")


def test_dynamic_attachment_row_is_themed_when_created_after_panel_setup():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._attachment_role_specs = (
            AttachmentRoleSpec(
                role="evidence",
                label="证明材料",
                accepted_types=("pdf",),
            ),
        )
        panel._sync_attachment_role_rows()

        row = panel._attachment_role_rows["evidence"]
        token = panel._attachment_role_token_edits["evidence"]
        name = panel._attachment_role_name_edits["evidence"]
        path = panel._attachment_role_path_edits["evidence"]
        assert row.styleSheet()
        assert token.text() == "{{@attach:evidence}}"
        assert name.text() == "证明材料"
        assert path.placeholderText() == "尚未选择附件"
        assert not panel._attachment_role_remove_buttons["evidence"].isEnabled()
    finally:
        panel.close()


def test_attachment_inventory_groups_single_and_package_roles():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._attachment_role_specs = (
            AttachmentRoleSpec(role="contract", label="合同"),
            AttachmentRoleSpec(
                role="certificates",
                label="证书集合",
                cardinality="multiple",
                source_kind="file_set",
                max_items=None,
            ),
            AttachmentRoleSpec(
                role="evidence_package",
                label="证明材料包",
                cardinality="multiple",
                source_kind="directory_package",
                recursive=True,
                max_items=None,
            ),
        )
        panel._sync_attachment_role_rows()

        assert panel._independent_attachments_count.text() == "1 项"
        assert panel._attachment_folders_count.text() == "2 项"
        assert panel._attachment_role_rows["contract"].parent() is (
            panel._independent_attachments_container
        )
        assert panel._attachment_role_rows["certificates"].parent() is (
            panel._attachment_folders_container
        )
        assert panel._attachment_role_rows["evidence_package"].parent() is (
            panel._attachment_folders_container
        )
    finally:
        panel.close()


def test_attachment_add_buttons_create_unbound_rows_before_source_selection(
    tmp_path: Path,
    monkeypatch,
):
    _app()
    pdf = tmp_path / "evidence.pdf"
    pdf.write_bytes(b"%PDF-1.4\nevidence")
    package = tmp_path / "package"
    package.mkdir()
    (package / "certificate.pdf").write_bytes(b"%PDF-1.4\ncertificate")
    panel = AssetsPanel(PanelBridge())
    try:
        monkeypatch.setattr(
            QFileDialog,
            "getOpenFileName",
            lambda *_args, **_kwargs: (_ for _ in ()).throw(
                AssertionError("header add must not open a file dialog")
            ),
        )
        panel._add_independent_attachment_btn.click()

        single_roles = [
            spec.role
            for spec in panel._attachment_role_specs
            if spec.source_kind == "single_file"
            and spec.deletable
        ]
        assert single_roles == ["附件1"]
        assert "附件1" not in panel._attachment_bindings
        assert panel._attachment_role_token_edits["附件1"].text() == (
            "{{@attach:附件1}}"
        )
        assert panel._independent_attachments_count.text() == "1 项"

        monkeypatch.setattr(
            QFileDialog,
            "getOpenFileName",
            lambda *_args, **_kwargs: (str(pdf), ""),
        )
        panel._attachment_role_choose_buttons["附件1"].click()
        assert panel._attachment_bindings["附件1"].source_kind.value == (
            "single_file"
        )
        assert panel._attachment_role_name_edits["附件1"].text() == "evidence"
        assert not panel._attachment_role_status_labels["附件1"].isHidden()
        assert "原样交付" in panel._attachment_role_status_labels["附件1"].text()
        assert panel._attachment_role_rows["附件1"].toolTip() == ""
        assert panel._attachment_role_status_labels["附件1"].toolTip() == ""

        monkeypatch.setattr(
            QFileDialog,
            "getExistingDirectory",
            lambda *_args, **_kwargs: str(package),
        )
        panel._add_attachment_folder_btn.click()

        folder_role = "附件文件夹1"
        assert folder_role not in panel._attachment_bindings
        panel._attachment_role_choose_buttons[folder_role].click()
        assert panel._attachment_bindings[folder_role].source_kind.value == (
            "directory_package"
        )
        assert panel._attachment_role_rows[folder_role].parent() is (
            panel._attachment_folders_container
        )
        assert panel._attachment_folders_count.text() == "1 项"
        assert panel._attachment_role_choose_buttons[folder_role].property(
            "pathAction"
        ) == "choose_directory"
        summary = panel._attachment_preparation_summaries[folder_role]
        assert summary.report().token_count == 0
        assert summary.isHidden()
        assert "border-radius" not in panel._attachment_role_rows[
            folder_role
        ].styleSheet()

        panel._add_independent_attachment_btn.click()
        before = tuple(spec.role for spec in panel._attachment_role_specs)
        monkeypatch.setattr(
            QFileDialog,
            "getOpenFileName",
            lambda *_args, **_kwargs: ("", ""),
        )
        panel._attachment_role_choose_buttons["附件2"].click()
        assert tuple(spec.role for spec in panel._attachment_role_specs) == before
        assert "附件2" not in panel._attachment_bindings

        panel._clear_attachment_file(folder_role)
        panel._attachment_role_remove_buttons[folder_role].click()
        assert folder_role not in {
            spec.role for spec in panel._attachment_role_specs
        }
        assert folder_role not in panel._attachment_role_rows
    finally:
        panel.close()


def test_attachment_docx_projects_token_requirements_and_enables_copy_mode(
    tmp_path: Path,
):
    _app()
    source = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("项目：{{@text:项目名称}}")
    document.save(source)
    panel = AssetsPanel(PanelBridge())
    try:
        panel._add_independent_attachment_btn.click()
        panel._apply_attachment_paths("附件1", (str(source),))

        binding = panel._attachment_bindings["附件1"]
        summary = panel._attachment_preparation_summaries["附件1"]
        assert binding.processing_mode.value == "substitute_copy"
        assert summary.objectName() == "attachment_preparation_summary"
        assert summary.report().token_count == 1
        assert summary.summary_text() == "1 个文件 · 1 项待补充"
        assert summary.status_text() == "1 个文件 · 1 项待补充"
        assert summary.prepare_button.text() == "去准备"
        assert not summary.report().requirements[0].is_ready
        assert "Token 1" in panel._attachment_role_status_labels["附件1"].text()
    finally:
        panel.close()


def test_attachment_legacy_tokens_are_explained_without_mutating_source(
    tmp_path: Path,
):
    _app()
    source = tmp_path / "legacy.docx"
    document = Document()
    document.add_paragraph("项目：{{项目名称}}")
    document.save(source)
    before = source.read_bytes()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._add_independent_attachment_btn.click()
        panel._apply_attachment_paths("附件1", (str(source),))

        summary = panel._attachment_preparation_summaries["附件1"]
        assert source.read_bytes() == before
        assert panel._attachment_bindings["附件1"].processing_mode.value == (
            "passthrough"
        )
        assert summary.report().token_count == 1
        assert summary.summary_text() == "1 个文件 · 1 项待补充"
        assert summary.status_text() == "1 个文件 · 1 项待补充"
    finally:
        panel.close()


def test_custom_attachment_token_rename_rekeys_definition_and_binding(tmp_path: Path):
    _app()
    pdf = tmp_path / "contract.pdf"
    pdf.write_bytes(b"%PDF-1.4\ncontract")
    panel = AssetsPanel(PanelBridge())
    try:
        panel._add_independent_attachment_btn.click()
        panel._apply_attachment_paths("附件1", (str(pdf),))

        panel._commit_attachment_role_token("附件1", "{{@attach:合同附件}}")

        assert [spec.role for spec in panel._attachment_role_specs] == ["合同附件"]
        assert panel._attachment_bindings["合同附件"].anchor_token == (
            "{{@attach:合同附件}}"
        )
        assert panel._attachment_role_token_edits["合同附件"].text() == (
            "{{@attach:合同附件}}"
        )
        assert panel._selected_profile().attachment_role_specs[0].role == "合同附件"
    finally:
        panel.close()


def test_archive_editor_round_trip_preserves_content_and_attachment_domains(
    tmp_path: Path,
    monkeypatch,
):
    _app()
    config_root = tmp_path / "config_library"
    monkeypatch.setattr(
        content_materials_presenter,
        "CONFIG_LIBRARY_ROOT",
        config_root,
    )
    markdown = tmp_path / "route.md"
    markdown.write_text("# 技术路线\n\n正文", encoding="utf-8")
    content = compile_content_binding(
        markdown,
        ContentArtifactRepository(config_root / "content_artifacts"),
        content_id="technical_route",
        label="技术路线",
    )
    rule = ContentInsertionRule(
        rule_id="content:technical_route",
        content_id="technical_route",
        anchor_token=content_anchor_token("technical_route"),
    )
    pdf = tmp_path / "evidence.pdf"
    pdf.write_bytes(b"%PDF-1.4\nevidence")
    attachment = build_single_attachment_binding(
        role="evidence",
        source_path=pdf,
        accepted_types=("pdf",),
    )
    archive = EntityArchive(
        archive_name="测试资料包",
        profiles=[
            EntityProfile(
                profile_name="第一份",
                image_material_rules={"image:logo": _context_image_rule()},
                content_bindings={content.content_id: content},
                content_rules=[rule],
                attachment_bindings={attachment.role: attachment},
            )
        ],
    )

    panel = AssetsPanel(PanelBridge())
    try:
        panel.set_archive(archive)
        exported = panel.current_archive().profiles[0]
        context = panel.material_context()

        assert exported.content_bindings == {content.content_id: content}
        assert exported.content_rules == [rule]
        assert exported.attachment_bindings == {attachment.role: attachment}
        assert panel._attachment_role_rows["evidence"].parent() is (
            panel._independent_attachments_container
        )
        logo_rule = exported.image_material_rules["image:logo"]
        assert logo_rule.placement.mode is ImagePlacementMode.NATURAL_SIZE
        assert logo_rule.watermark == _context_image_rule().watermark
        assert "evidence" not in exported.asset_paths
        assert "evidence" not in exported.asset_bindings
        assert context.content_bindings == {content.content_id: content}
        assert context.content_rules == [rule]
        assert context.attachment_bindings == {attachment.role: attachment}
        assert context.image_material_rules["image:logo"] == logo_rule
        assert context.to_resolve_kwargs()["images"] == []
    finally:
        panel.close()


def test_set_archive_hydrates_selected_profile_once():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        original = panel._set_editor_values
        calls: list[str] = []

        def counted_set_editor_values(**kwargs):
            calls.append(str(kwargs.get("profile_id", "")))
            return original(**kwargs)

        panel._set_editor_values = counted_set_editor_values
        panel.set_archive(
            EntityArchive(profiles=[EntityProfile(profile_id="selected")])
        )

        assert calls == ["selected"]
    finally:
        panel.close()


def test_legacy_attachment_path_migrates_out_of_shared_image_paths(tmp_path: Path):
    _app()
    pdf = tmp_path / "legacy.pdf"
    pdf.write_bytes(b"%PDF-1.4\nlegacy")
    panel = AssetsPanel(PanelBridge())
    try:
        panel._attachment_role_specs = (
            AttachmentRoleSpec(
                role="evidence",
                label="证明材料",
                accepted_types=("pdf",),
            ),
        )
        panel._set_editor_values(
            archive_id="",
            archive_name="",
            profile_id="",
            profile_name="",
            fields={},
            assets_dir="",
            asset_paths={"evidence": str(pdf)},
            attachment_bindings={},
        )

        assert "evidence" not in panel._asset_paths
        assert panel._attachment_bindings["evidence"].items[0].file_ref.source_path == str(
            pdf.resolve()
        )
    finally:
        panel.close()


def test_received_material_context_writes_image_rules_back_to_profile_and_editor():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.set_archive(
            EntityArchive(
                archive_id="material-package",
                profiles=[EntityProfile(profile_id="profile-1")],
            )
        )
        incoming = panel.material_context().clone()
        incoming.image_material_rules = {"image:logo": _context_image_rule()}

        panel._on_material_context_changed(incoming)

        profile = panel._selected_profile()
        expected = panel._image_material_rules["image:logo"]
        assert expected.placement.mode is ImagePlacementMode.NATURAL_SIZE
        assert expected.watermark == _context_image_rule().watermark
        assert profile.image_material_rules["image:logo"] == expected
        exported = panel.material_context()
        assert exported.image_material_rules["image:logo"] == expected
        assert exported.to_resolve_kwargs()["images"] == []
    finally:
        panel.close()


def test_received_execution_asset_projection_does_not_create_profile_sources(tmp_path):
    _app()
    logo = tmp_path / "logo.png"
    logo.write_bytes(b"logo")
    seal = tmp_path / "seal.png"
    seal.write_bytes(b"seal")
    external_dir = tmp_path / "external"
    external_dir.mkdir()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.set_archive(
            EntityArchive(
                archive_id="material-package",
                profiles=[
                    EntityProfile(
                        profile_id="profile-1",
                        asset_paths={"logo": str(logo)},
                    )
                ],
            )
        )
        incoming = panel.material_context().clone()
        incoming.entity_assets_dir = str(external_dir)
        incoming.asset_items = [AssetItem(role="seal", path=str(seal))]

        panel._on_material_context_changed(incoming)

        profile = panel._selected_profile()
        assert profile.assets_dir == ""
        assert profile.asset_paths == {"logo": str(logo)}
        assert profile.asset_items == []
        assert panel._assets_picker.path() == ""
        assert panel._asset_paths == {"logo": str(logo)}
        assert panel._asset_item_payloads == []
    finally:
        panel.close()


def test_content_token_preview_uses_only_content_binding():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._ensure_content_material_rule("technical_route")
        missing = panel._placeholder_preview_rows(
            ["@file:technical_route"], {}, [], []
        )
        attachment_like = panel._placeholder_preview_rows(["evidence"], {}, [], [])

        assert missing[0]["action_type"] == "content"
        assert missing[0]["source"] == "文件资料"
        assert attachment_like[0]["source"] == "文档占位符"
        assert "附件" not in attachment_like[0]["source"]
    finally:
        panel.close()


def test_image_rule_editor_applies_one_strategy_and_watermark_to_all_roles():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.set_archive(EntityArchive(profiles=[EntityProfile()]))
        assert panel._image_rule_adaptive_check.isChecked()
        assert panel._image_rules_card._description_label is None
        assert not hasattr(panel, "_image_rule_width_spins")

        panel._image_rule_adaptive_check.setChecked(False)
        assert all(
            rule.placement.mode is ImagePlacementMode.NATURAL_SIZE
            for rule in panel._image_material_rules.values()
        )

        panel._image_rule_adaptive_check.setChecked(True)
        panel._image_rule_watermark_check.setChecked(True)
        assert panel._image_rule_watermark_fixed_radio.isEnabled()
        assert panel._image_rule_watermark_fixed_radio.isChecked()
        panel._image_rule_watermark_edit.setText("示例水印{{credit_code}}")
        panel._commit_global_image_policy()

        rule = panel.current_archive().profiles[0].image_material_rules[
            "image:logo"
        ]
        qualification = panel._image_material_rules["image:qualification"]

        assert rule.placement.mode is ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE
        assert (
            rule.placement.co_location_guard
            is ImageCoLocationGuard.PRECEDING_NONEMPTY_PARAGRAPH
        )
        assert rule.watermark.enabled is True
        assert (
            rule.watermark.text_source
            is ImageWatermarkTextSource.FIXED_FIELD
        )
        assert rule.watermark.text_template == "示例水印{{credit_code}}"
        assert (
            qualification.placement.mode
            is ImagePlacementMode.FIT_CONTAINER_FLOW
        )
        assert all(
            item.watermark == rule.watermark
            for item in panel._image_material_rules.values()
        )

        panel._image_rule_watermark_free_radio.setChecked(True)
        free_rule = panel._image_material_rules["image:logo"]
        assert (
            free_rule.watermark.text_source
            is ImageWatermarkTextSource.WORKBENCH_FREE_FIELD
        )
        assert free_rule.watermark.text_template == ""
        assert not panel._image_rule_watermark_edit.isEnabled()
        assert panel._image_rule_watermark_edit.placeholderText() == "在工作台填写"

        panel._image_rule_adaptive_check.setChecked(False)
        assert panel._image_material_rules["image:logo"].watermark == free_rule.watermark
        panel._image_rule_watermark_check.setChecked(False)
        assert (
            panel._image_material_rules["image:logo"].placement.mode
            is ImagePlacementMode.NATURAL_SIZE
        )
    finally:
        panel.close()
