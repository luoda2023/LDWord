import sys
from pathlib import Path

import pytest
from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


def test_builtin_heading_numbering_schemes_define_all_eight_levels():
    from src.config.heading_presets import PRESET_CATALOG, SUPPORTED_LEVELS, get_preset_max_levels

    expected_keys = {f"heading{level}" for level in range(1, SUPPORTED_LEVELS + 1)}

    assert PRESET_CATALOG
    for key, entry in PRESET_CATALOG.items():
        assert set(entry["bindings"]) == expected_keys, key
        assert 1 <= int(entry["max_levels"]) <= SUPPORTED_LEVELS

    assert get_preset_max_levels("thesis_standard") == 4
    assert get_preset_max_levels("none") == 8


def test_adapter_expands_deep_levels_from_current_scheme_not_generic_fallback():
    from src.config.template import TemplateConfig
    from src.ui.adapters.heading_numbering_adapter import HeadingNumberingAdapter

    adapter = HeadingNumberingAdapter()
    adapter.set_template(TemplateConfig())
    adapter.apply_preset("thesis_standard")

    adapter.set_max_levels(6)

    heading5 = adapter.get_binding(5)
    heading6 = adapter.get_binding(6)
    assert heading5.display_template == "（{nn}）"
    assert heading5.chain == "current_only"
    assert heading6.display_core_style == "circled"
    assert adapter.preview_number(6).startswith("①")


def test_user_heading_numbering_scheme_can_be_saved_and_deleted(tmp_path, monkeypatch):
    import src.config.heading_presets as heading_presets
    from src.config.template import TemplateConfig
    from src.ui.adapters.heading_numbering_adapter import HeadingNumberingAdapter

    monkeypatch.setattr(heading_presets, "USER_SCHEME_DIR", tmp_path)

    adapter = HeadingNumberingAdapter()
    adapter.set_template(TemplateConfig())
    adapter.apply_preset("arabic_dot")
    adapter.set_max_levels(8)

    scheme_id = adapter.save_current_as_user_scheme("我的数字方案")

    assert scheme_id.startswith("user.")
    assert adapter.detect_active_preset() == scheme_id
    assert heading_presets.is_user_preset(scheme_id)
    saved = heading_presets.get_preset_bindings(scheme_id)
    assert saved is not None
    assert set(saved) == {f"heading{level}" for level in range(1, 9)}
    assert saved["heading8"].chain.count("parent") == 7

    assert adapter.delete_active_user_scheme() is True
    assert heading_presets.is_user_preset(scheme_id) is False


def test_user_heading_numbering_scheme_rejects_duplicate_display_names(tmp_path, monkeypatch):
    import src.config.heading_presets as heading_presets
    from src.config.template import TemplateConfig
    from src.ui.adapters.heading_numbering_adapter import HeadingNumberingAdapter

    monkeypatch.setattr(heading_presets, "USER_SCHEME_DIR", tmp_path)

    adapter = HeadingNumberingAdapter()
    adapter.set_template(TemplateConfig())

    with pytest.raises(ValueError, match="名称已存在"):
        adapter.save_current_as_user_scheme("论文标准")

    adapter.save_current_as_user_scheme("我的编号方案")
    adapter.apply_preset("arabic_dot")
    with pytest.raises(ValueError, match="名称已存在"):
        adapter.save_current_as_user_scheme("我的编号方案")


def test_user_heading_numbering_scheme_catalog_skips_duplicate_display_names(tmp_path, monkeypatch):
    import json

    import src.config.heading_presets as heading_presets

    monkeypatch.setattr(heading_presets, "USER_SCHEME_DIR", tmp_path)
    (tmp_path / "user.none_copy.json").write_text(
        json.dumps(
            {
                "scheme_id": "user.none_copy",
                "name": "无编号",
                "max_levels": 8,
                "level_bindings": {},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    catalog = heading_presets.get_scheme_catalog()

    assert "none" in catalog
    assert "user.none_copy" not in catalog


def test_heading_numbering_apply_honors_start_at_and_child_restart():
    from src.config.template import HeadingLevelBindingConfig, TemplateConfig
    from src.modules.structure.heading_numbering import HeadingNumberingModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.heading_numbering_ooxml import (
        effective_numbering,
        managed_level_start,
        managed_level_text,
    )

    doc = Document()
    doc.add_paragraph("Chapter A")
    doc.add_paragraph("Section A")
    doc.add_paragraph("Section B")
    doc.add_paragraph("Chapter B")
    doc.add_paragraph("Section C")

    cfg = TemplateConfig()
    cfg.heading_model.max_heading_levels = 2
    cfg.heading_numbering.level_bindings = {
        "heading1": HeadingLevelBindingConfig(
            enabled=True,
            display_template="{nn}",
            chain="current_only",
            title_separator=" ",
            start_at=3,
        ),
        "heading2": HeadingLevelBindingConfig(
            enabled=True,
            display_template="{chain}",
            chain="parent.current",
            chain_separator=".",
            title_separator=" ",
            start_at=0,
        ),
    }
    context = PipelineContext(heading_map={0: 1, 1: 2, 2: 2, 3: 1, 4: 2})

    HeadingNumberingModule().apply(doc, cfg, ChangeTracker(), context)

    assert [paragraph.text for paragraph in doc.paragraphs] == [
        "Chapter A",
        "Section A",
        "Section B",
        "Chapter B",
        "Section C",
    ]
    assert [effective_numbering(paragraph).level for paragraph in doc.paragraphs] == [
        0,
        1,
        1,
        0,
        1,
    ]
    assert managed_level_start(doc, 1) == 3
    assert managed_level_start(doc, 2) == 0
    assert managed_level_text(doc, 1) == "%1 "
    assert managed_level_text(doc, 2) == "%1.%2 "


def test_heading_numbering_restart_on_document_keeps_child_counter_continuous():
    from src.config.template import HeadingLevelBindingConfig, TemplateConfig
    from src.modules.structure.heading_numbering import HeadingNumberingModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.heading_numbering_ooxml import (
        effective_numbering,
        managed_level_restart,
    )

    doc = Document()
    doc.add_paragraph("Chapter A")
    doc.add_paragraph("Section A")
    doc.add_paragraph("Chapter B")
    doc.add_paragraph("Section B")

    cfg = TemplateConfig()
    cfg.heading_model.max_heading_levels = 2
    cfg.heading_numbering.level_bindings = {
        "heading1": HeadingLevelBindingConfig(
            enabled=True,
            display_template="{nn}",
            chain="current_only",
            title_separator=" ",
        ),
        "heading2": HeadingLevelBindingConfig(
            enabled=True,
            display_template="{nn}",
            chain="current_only",
            title_separator=" ",
            restart_on="document",
        ),
    }
    context = PipelineContext(heading_map={0: 1, 1: 2, 2: 1, 3: 2})

    HeadingNumberingModule().apply(doc, cfg, ChangeTracker(), context)

    assert [paragraph.text for paragraph in doc.paragraphs] == [
        "Chapter A",
        "Section A",
        "Chapter B",
        "Section B",
    ]
    assert all(effective_numbering(paragraph) is not None for paragraph in doc.paragraphs)
    assert managed_level_restart(doc, 2) == 0


def test_heading_numbering_restart_on_specific_heading_only_resets_on_that_level():
    from src.config.template import HeadingLevelBindingConfig, TemplateConfig
    from src.modules.structure.heading_numbering import HeadingNumberingModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.heading_numbering_ooxml import (
        effective_numbering,
        managed_level_restart,
    )

    doc = Document()
    doc.add_paragraph("Chapter A")
    doc.add_paragraph("Section A")
    doc.add_paragraph("Item A")
    doc.add_paragraph("Section B")
    doc.add_paragraph("Item B")
    doc.add_paragraph("Chapter B")
    doc.add_paragraph("Section C")
    doc.add_paragraph("Item C")

    cfg = TemplateConfig()
    cfg.heading_model.max_heading_levels = 3
    cfg.heading_numbering.level_bindings = {
        "heading1": HeadingLevelBindingConfig(enabled=True, display_template="{nn}", title_separator=" "),
        "heading2": HeadingLevelBindingConfig(enabled=True, display_template="{nn}", title_separator=" "),
        "heading3": HeadingLevelBindingConfig(
            enabled=True,
            display_template="{nn}",
            title_separator=" ",
            restart_on="heading1",
        ),
    }
    context = PipelineContext(heading_map={0: 1, 1: 2, 2: 3, 3: 2, 4: 3, 5: 1, 6: 2, 7: 3})

    HeadingNumberingModule().apply(doc, cfg, ChangeTracker(), context)

    assert [paragraph.text for paragraph in doc.paragraphs] == [
        "Chapter A",
        "Section A",
        "Item A",
        "Section B",
        "Item B",
        "Chapter B",
        "Section C",
        "Item C",
    ]
    assert all(effective_numbering(paragraph) is not None for paragraph in doc.paragraphs)
    assert managed_level_restart(doc, 3) == 1


def test_heading_numbering_excludes_level_from_native_toc_outline():
    from src.config.template import HeadingLevelBindingConfig, TemplateConfig
    from src.modules.structure.heading_numbering import HeadingNumberingModule
    from src.pipeline.context import PipelineContext
    from src.pipeline.tracker import ChangeTracker
    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    doc.add_paragraph("Included root")
    doc.add_paragraph("Hidden child")

    cfg = TemplateConfig()
    cfg.heading_model.max_heading_levels = 2
    cfg.heading_numbering.level_bindings = {
        "heading1": HeadingLevelBindingConfig(
            enabled=True,
            display_template="{nn}",
            chain="current_only",
            title_separator=" ",
            include_in_toc=True,
        ),
        "heading2": HeadingLevelBindingConfig(
            enabled=True,
            display_template="{nn}",
            chain="current_only",
            title_separator=" ",
            include_in_toc=False,
        ),
    }
    context = PipelineContext(heading_map={0: 1, 1: 2})

    HeadingNumberingModule().apply(doc, cfg, ChangeTracker(), context)

    root_outline = doc.paragraphs[0]._element.find(qn("w:pPr")).find(qn("w:outlineLvl"))
    child_outline = doc.paragraphs[1]._element.find(qn("w:pPr")).find(qn("w:outlineLvl"))
    assert root_outline.get(qn("w:val")) == "0"
    assert child_outline.get(qn("w:val")) == "9"


def test_plain_toc_collection_respects_level_include_in_toc_flag():
    from src.config.template import HeadingLevelBindingConfig, TemplateConfig
    from src.modules.structure.toc import _collect_plain_toc_entries
    from src.pipeline.context import PipelineContext

    doc = Document()
    doc.add_paragraph("Included root")
    doc.add_paragraph("Hidden child")

    cfg = TemplateConfig()
    cfg.heading_numbering.level_bindings = {
        "heading1": HeadingLevelBindingConfig(include_in_toc=True),
        "heading2": HeadingLevelBindingConfig(include_in_toc=False),
    }
    context = PipelineContext(heading_map={0: 1, 1: 2})

    entries = _collect_plain_toc_entries(doc, cfg, context, max_level=2)

    assert [entry["title"] for entry in entries] == ["Included root"]
