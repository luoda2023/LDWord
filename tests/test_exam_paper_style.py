from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from src.config.scene import ExamBlankStyleConfig, ExamPaperConfig
from src.shared.engine.exam_paper_style import (
    BUILTIN_EXAM_BLANK_STYLE_IDS,
    BUILTIN_EXAM_MASTER_DIR,
    EXAM_MARKDOWN_AUTHORING_PROMPT,
    EXAM_MASTER_CONVERSION_PROMPT,
    USER_EXAM_MASTER_DIR,
    builtin_exam_blank_style_options,
    create_exam_blank_master_copy,
    create_exam_blank_style_copy,
    ensure_builtin_exam_master_docx,
    ensure_current_exam_blank_master_docx,
    exam_blank_style_label,
    exam_blank_style_preview_lines,
    import_exam_blank_master_docx,
    inventory_user_exam_master_files,
    sync_user_exam_blank_master_files,
    write_exam_blank_master_docx,
    write_exam_blank_style_sample_docx,
    write_exam_paper_docx,
    write_exam_paper_docx_files,
)


def _all_docx_text(document: Document) -> str:
    paragraphs = []

    def collect(parent) -> None:
        paragraphs.extend(paragraph.text for paragraph in getattr(parent, "paragraphs", []))
        for table in getattr(parent, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    collect(cell)

    collect(document)
    return "\n".join(paragraphs)


def _sealed_header_xml_from_package(path) -> str:
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/header") or not name.endswith(".xml"):
                continue
            xml = archive.read(name).decode("utf-8")
            if "ExamSeal" in xml and "rotation:-5898240f" in xml:
                return xml
    raise AssertionError("sealed first-page header XML not found")


def test_exam_blank_style_copy_is_scene_owned_and_explicitly_resolvable():
    config = ExamPaperConfig()

    copied = create_exam_blank_style_copy(config, "default_exam")
    config.custom_blank_styles.append(copied)
    normalized = ExamPaperConfig(
        custom_blank_styles=[
            {
                "style_id": copied.style_id,
                "label": copied.label,
                "base_style_id": copied.base_style_id,
            }
        ],
    )

    assert not hasattr(normalized, "blank_style_id")
    assert normalized.custom_blank_styles == [
        ExamBlankStyleConfig(
            style_id=copied.style_id,
            label=copied.label,
            base_style_id="default_exam",
        )
    ]
    assert exam_blank_style_label(copied.style_id, normalized) == copied.label


def test_exam_blank_master_copy_creates_scene_owned_docx(tmp_path):
    config = ExamPaperConfig()

    copied = create_exam_blank_master_copy(
        config,
        "default_exam",
        output_dir=tmp_path / "masters",
    )
    config.custom_blank_styles.append(copied)

    path = ensure_current_exam_blank_master_docx(config, copied.style_id)
    assert path.exists()
    assert path.parent == tmp_path / "masters"
    assert copied.master_docx_path
    text = _all_docx_text(Document(path))
    assert "{{af_title}}" in text
    assert "{{af_questions}}" in text


def test_import_exam_blank_master_docx_copies_existing_word_master(tmp_path):
    source_path = tmp_path / "学校常用试卷母版.docx"
    source = Document()
    source.add_paragraph("学校常用母版标识")
    source.save(source_path)

    config = ExamPaperConfig()
    imported = import_exam_blank_master_docx(
        config,
        source_path,
        output_dir=tmp_path / "masters",
    )
    config.custom_blank_styles.append(imported)

    stored_path = ensure_current_exam_blank_master_docx(config, imported.style_id)
    assert imported.style_id == "user_imported_exam"
    assert imported.label == "学校常用试卷母版"
    assert stored_path.parent == tmp_path / "masters"
    assert stored_path.exists()
    assert source_path.exists()
    assert "学校常用母版标识" in _all_docx_text(Document(stored_path))

    sample_path = write_exam_blank_style_sample_docx(
        imported.style_id,
        tmp_path / "samples",
        config=config,
    )
    assert "学校常用母版标识" in _all_docx_text(Document(sample_path))


def test_import_exam_blank_master_docx_uses_exam_shell_error_copy(tmp_path):
    config = ExamPaperConfig()

    with pytest.raises(ValueError, match="请选择 .docx 格式的试卷卷面"):
        import_exam_blank_master_docx(config, tmp_path / "not_a_doc.txt")

    broken_docx = tmp_path / "broken.docx"
    broken_docx.write_text("not a real docx", encoding="utf-8")

    with pytest.raises(ValueError, match="选择的试卷卷面无法打开"):
        import_exam_blank_master_docx(config, broken_docx)


def test_exam_user_master_default_write_target_is_mode_scoped():
    assert USER_EXAM_MASTER_DIR.parts[-4:] == (
        "config_library",
        "masters",
        "exam",
        "user",
    )


def test_exam_blank_master_copy_defaults_to_mode_scoped_user_dir(
    tmp_path,
    monkeypatch,
):
    active_user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    monkeypatch.setattr(
        "src.shared.engine.exam_paper_style.USER_EXAM_MASTER_DIR",
        active_user_dir,
    )

    config = ExamPaperConfig()
    copied = create_exam_blank_master_copy(config, "default_exam")
    config.custom_blank_styles.append(copied)
    stored_path = ensure_current_exam_blank_master_docx(config, copied.style_id)

    assert stored_path.parent == active_user_dir
    assert active_user_dir.exists()


def test_missing_custom_exam_master_is_reported_without_silent_default_copy(
    tmp_path,
    monkeypatch,
):
    active_user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    missing_path = active_user_dir / "user_default_exam_copy.docx"
    monkeypatch.setattr(
        "src.shared.engine.exam_paper_style.USER_EXAM_MASTER_DIR",
        active_user_dir,
    )

    config = ExamPaperConfig(
        custom_blank_styles=[
            ExamBlankStyleConfig(
                style_id="user_default_exam_copy",
                label="Missing copy",
                master_docx_path=str(missing_path),
            )
        ],
    )
    with pytest.raises(FileNotFoundError, match="自定义卷面 DOCX 不存在"):
        ensure_current_exam_blank_master_docx(config, "user_default_exam_copy")

    assert not (active_user_dir / "user_default_exam_copy.docx").exists()
    assert not missing_path.exists()
    assert Path(config.custom_blank_styles[0].master_docx_path).resolve() == missing_path.resolve()


def test_user_master_folder_discovery_registers_valid_manual_docx_only(tmp_path):
    user_dir = tmp_path / "user_masters"

    manual_path = user_dir / "AI生成期中母版.docx"
    manual_path.parent.mkdir(parents=True)
    manual = Document()
    manual.add_paragraph("{{af_title}}")
    manual.add_paragraph("科目：{{af_subject}}")
    manual.add_paragraph("{{af_questions}}")
    manual.add_paragraph("{{af_answer_area}}")
    manual.save(manual_path)

    invalid_path = user_dir / "普通试卷.docx"
    invalid = Document()
    invalid.add_paragraph("这只是普通试卷，没有 Alavette 母版占位符。")
    invalid.save(invalid_path)

    generated_path = user_dir / "user_default_exam_copy_999.docx"
    generated = Document()
    generated.add_paragraph("{{af_title}}")
    generated.add_paragraph("{{af_questions}}")
    generated.save(generated_path)

    config = ExamPaperConfig()
    added = sync_user_exam_blank_master_files(config, user_dir)

    assert [style.label for style in added] == [
        "AI生成期中母版",
        "user_default_exam_copy_999",
    ]
    assert [style.label for style in config.custom_blank_styles] == [
        "AI生成期中母版",
        "user_default_exam_copy_999",
    ]
    discovered = config.custom_blank_styles[0]
    assert discovered.style_id.startswith("user_file_")
    assert discovered.master_docx_path.endswith("AI生成期中母版.docx")

    added_again = sync_user_exam_blank_master_files(config, user_dir)
    assert added_again == ()
    assert [style.label for style in config.custom_blank_styles] == [
        "AI生成期中母版",
        "user_default_exam_copy_999",
    ]


def test_user_master_inventory_classifies_references_and_stale_pool_files(tmp_path):
    user_dir = tmp_path / "user_masters"
    user_dir.mkdir()

    def write_docx(path: Path, *, valid_master: bool) -> None:
        document = Document()
        if valid_master:
            document.add_paragraph("{{af_title}}")
            document.add_paragraph("{{af_questions}}")
            document.add_paragraph("{{af_answer_area}}")
        else:
            document.add_paragraph("普通 Word 文件")
        document.save(path)

    referenced_path = user_dir / "user_default_exam_copy.docx"
    stale_copy_path = user_dir / "user_default_exam_copy_2.docx"
    stale_import_path = user_dir / "user_imported_exam.docx"
    manual_path = user_dir / "学校手工母版.docx"
    invalid_path = user_dir / "普通试卷.docx"
    for path in (referenced_path, stale_copy_path, stale_import_path, manual_path):
        write_docx(path, valid_master=True)
    write_docx(invalid_path, valid_master=False)

    config = ExamPaperConfig(
        custom_blank_styles=[
            ExamBlankStyleConfig(
                style_id="user_default_exam_copy",
                label="当前试卷母版副本",
                master_docx_path=str(referenced_path),
            )
        ]
    )

    items = inventory_user_exam_master_files(config, user_dir)
    by_name = {item.file_name: item for item in items}

    assert len(config.custom_blank_styles) == 1
    assert by_name["user_default_exam_copy.docx"].category == "referenced"
    assert by_name["user_default_exam_copy.docx"].referenced is True
    assert by_name["user_default_exam_copy.docx"].style_id == "user_default_exam_copy"
    assert by_name["user_default_exam_copy_2.docx"].category == "program_copy_unreferenced"
    assert by_name["user_imported_exam.docx"].category == "program_import_unreferenced"
    assert by_name["学校手工母版.docx"].category == "manual_discoverable"
    assert by_name["普通试卷.docx"].category == "unknown_or_invalid"

    cross_plan_config = ExamPaperConfig(
        custom_blank_styles=[
            ExamBlankStyleConfig(
                style_id="user_default_exam_copy_2",
                label="其他方案引用的母版",
                master_docx_path=str(stale_copy_path),
            )
        ]
    )
    cross_items = inventory_user_exam_master_files(
        ExamPaperConfig(),
        user_dir,
        reference_configs=[cross_plan_config],
    )
    cross_by_name = {item.file_name: item for item in cross_items}

    assert cross_by_name["user_default_exam_copy_2.docx"].category == "referenced"
    assert cross_by_name["user_default_exam_copy_2.docx"].referenced is True


def test_exam_blank_style_preview_switches_student_and_answer_versions():
    student_lines = "\n".join(exam_blank_style_preview_lines("default_exam"))
    answer_lines = "\n".join(
        exam_blank_style_preview_lines("default_exam", answer_version=True)
    )

    assert "学生卷：答案与解析不显示" not in student_lines
    assert "答案速查" not in student_lines
    assert "下列词语中加点字读音" in student_lines
    assert "答案速查" in answer_lines
    assert "1. C" in answer_lines
    assert "2. 按表达评分" in answer_lines
    assert "注意事项" not in answer_lines
    assert "下列词语中加点字读音" not in answer_lines
    assert "阅读材料，概括文章" not in answer_lines


def test_exam_ai_prompts_separate_content_and_master_responsibilities():
    assert "只输出 Markdown 正文" in EXAM_MARKDOWN_AUTHORING_PROMPT
    assert "不要编写页眉、页脚、页码、密封线" in EXAM_MARKDOWN_AUTHORING_PROMPT
    assert "## 答案速查" in EXAM_MARKDOWN_AUTHORING_PROMPT
    assert "answer_area_kind" in EXAM_MARKDOWN_AUTHORING_PROMPT
    assert "数学计算/解答/证明" in EXAM_MARKDOWN_AUTHORING_PROMPT
    assert "{{af_questions}}" not in EXAM_MARKDOWN_AUTHORING_PROMPT

    assert "Word 试卷卷面工程师" in EXAM_MASTER_CONVERSION_PROMPT
    assert "{{af_title}}" in EXAM_MASTER_CONVERSION_PROMPT
    assert "{{af_questions}}" in EXAM_MASTER_CONVERSION_PROMPT
    assert "{{af_answer_area}}" in EXAM_MASTER_CONVERSION_PROMPT
    assert "页眉页脚只作为卷面版式存在" in EXAM_MASTER_CONVERSION_PROMPT


def test_builtin_exam_blank_styles_expose_current_master_choices():
    assert BUILTIN_EXAM_BLANK_STYLE_IDS == ("default_exam",)
    options = dict(builtin_exam_blank_style_options())

    assert options == {"default_exam": "A4 标准卷面"}
    assert "compact_exam" not in options


def test_builtin_exam_master_directory_keeps_only_current_assets_prompts_and_manifests():
    expected_files = {
        "default_exam_v20.docx",
        "default_exam.master.json",
        "exam_content_generation_prompt.md",
        "exam_master_conversion_prompt.md",
    }
    actual_files = {path.name for path in Path(BUILTIN_EXAM_MASTER_DIR).iterdir() if path.is_file()}

    assert actual_files == expected_files
    assert EXAM_MARKDOWN_AUTHORING_PROMPT in (
        BUILTIN_EXAM_MASTER_DIR / "exam_content_generation_prompt.md"
    ).read_text(encoding="utf-8")
    assert EXAM_MASTER_CONVERSION_PROMPT in (
        BUILTIN_EXAM_MASTER_DIR / "exam_master_conversion_prompt.md"
    ).read_text(encoding="utf-8")


def test_exam_default_style_sample_docx_is_generated(tmp_path):
    path = write_exam_blank_style_sample_docx("default_exam", tmp_path)

    assert path.exists()
    document = Document(path)
    text = _all_docx_text(document)
    assert "七年级语文期中测试样张" in text
    assert "学生卷" in text
    assert "科目：语文" in text
    assert "年级：七年级" in text
    assert "考试时间：90 分钟" in text
    assert "满分：100 分" in text
    assert "{{af_title}}" not in text
    assert "{{af_questions}}" not in text
    assert "题目插入点" not in text
    assert "作答区" not in text
    assert "答案速查" not in text
    assert "答案：C" not in text
    assert "1. C" not in text


def test_exam_paper_writer_inserts_payload_into_current_master(tmp_path):
    payload = {
        "title": "五年级数学期中测试卷",
        "subject": "数学",
        "grade": "五年级",
        "duration": "90 分钟",
        "total_score": "100 分",
        "sections": [
            {
                "title": "一、填空题",
                "questions": [
                    {
                        "stem": "7 个 0.9 是（    ）。",
                        "answer": "6.3",
                        "score": "2",
                    },
                    {
                        "stem": "选择正确答案。",
                        "options": {"A": "3.14", "B": "6.28"},
                        "answer": "B",
                        "analysis": "按小数乘法计算。",
                        "score": "3",
                    },
                ],
            }
        ],
    }

    path = write_exam_paper_docx("default_exam", tmp_path, payload=payload)

    assert path.exists()
    assert path.name == "五年级数学期中测试卷.docx"
    text = _all_docx_text(Document(path))
    assert "五年级数学期中测试卷" in text
    assert "学生卷" in text
    assert "一、填空题" in text
    assert "一、题目区" not in text
    assert "1. 7 个 0.9 是（    ）。（2 分）" in text
    assert "A. 3.14" in text
    assert "答案速查" not in text
    assert "1. 6.3" not in text
    assert "2. B" not in text
    assert "答案：6.3" not in text
    assert "解析：按小数乘法计算。" not in text
    assert "{{af_title}}" not in text
    assert "{{af_questions}}" not in text
    assert "题目插入点" not in text
    assert "作答区" not in text

    first_page_header_xml = _sealed_header_xml_from_package(path)
    assert first_page_header_xml.count("AlternateContent") == 4
    assert first_page_header_xml.count("rotation:-5898240f") == 2
    assert "ExamSealDottedLine" in first_page_header_xml


def test_exam_paper_writer_preserves_user_docx_layout_and_existing_styles(tmp_path):
    user_master = tmp_path / "学校自定义卷面.docx"
    document = Document(ensure_builtin_exam_master_docx("default_exam"))
    section = document.sections[0]
    section.top_margin = Cm(3.1)
    section.bottom_margin = Cm(2.7)
    normal_style = document.styles["Normal"]
    normal_style.font.size = Pt(16)
    normal_style.paragraph_format.line_spacing = 2.0
    question_style = document.styles["Exam Question Stem"]
    question_style.font.size = Pt(15)
    question_style.paragraph_format.line_spacing = 1.8
    document.save(user_master)

    config = ExamPaperConfig(
        custom_blank_styles=[
            ExamBlankStyleConfig(
                style_id="user_file_school",
                label="学校自定义卷面",
                master_docx_path=str(user_master),
            )
        ],
    )
    output = write_exam_paper_docx(
        "user_file_school",
        tmp_path / "output",
        config=config,
    )

    rendered = Document(output)
    rendered_section = rendered.sections[0]
    assert round(rendered_section.top_margin.cm, 1) == 3.1
    assert round(rendered_section.bottom_margin.cm, 1) == 2.7
    assert rendered.styles["Normal"].font.size.pt == 16
    assert rendered.styles["Normal"].paragraph_format.line_spacing == 2.0
    assert rendered.styles["Exam Question Stem"].font.size.pt == 15
    assert rendered.styles["Exam Question Stem"].paragraph_format.line_spacing == 1.8


def test_dropped_minimal_user_docx_is_discovered_and_generates_normally(tmp_path):
    user_dir = tmp_path / "user"
    user_dir.mkdir()
    dropped = user_dir / "教师自制卷面.docx"
    document = Document()
    document.styles["Normal"].font.size = Pt(14)
    document.add_paragraph("{{af_title}}")
    document.add_paragraph("{{af_questions}}")
    document.save(dropped)

    config = ExamPaperConfig()
    discovered = sync_user_exam_blank_master_files(config, user_dir)
    assert len(discovered) == 1

    output = write_exam_paper_docx(
        discovered[0].style_id,
        tmp_path / "output",
        config=config,
    )

    rendered = Document(output)
    assert "七年级语文期中测试样张" in _all_docx_text(rendered)
    assert "1. 下列词语中加点字读音完全正确的一项是" in _all_docx_text(rendered)
    assert rendered.styles["Normal"].font.size.pt == 14
    assert "Exam Question Stem" in {style.name for style in rendered.styles}


def test_exam_paper_writer_outputs_separate_student_and_answer_key_files(tmp_path):
    payload = {
        "title": "五年级数学期中测试卷",
        "subject": "数学",
        "grade": "五年级",
        "duration": "90 分钟",
        "total_score": "100 分",
        "sections": [
            {
                "title": "一、填空题",
                "questions": [
                    {"stem": "7 个 0.9 是（    ）。", "answer": "6.3", "score": "2"},
                    {
                        "stem": "选择正确答案。",
                        "options": {"A": "3.14", "B": "6.28"},
                        "answer": "B",
                        "analysis": "按小数乘法计算。",
                        "score": "3",
                    },
                ],
            }
        ],
    }

    outputs = write_exam_paper_docx_files("default_exam", tmp_path, payload=payload)

    assert outputs.student_docx.exists()
    assert outputs.answer_key_docx is not None
    assert outputs.answer_key_docx.exists()
    assert outputs.student_docx.name == "五年级数学期中测试卷_学生卷.docx"
    assert outputs.answer_key_docx.name == "五年级数学期中测试卷_答案速查.docx"

    student_text = _all_docx_text(Document(outputs.student_docx))
    assert "学生卷" in student_text
    assert "答案速查" not in student_text
    assert "1. 7 个 0.9 是（    ）。（2 分）" in student_text
    assert "1. 6.3" not in student_text

    answer_text = _all_docx_text(Document(outputs.answer_key_docx))
    assert "答案速查" in answer_text
    assert "1. 6.3" in answer_text
    assert "2. B" in answer_text
    assert "7 个 0.9 是" not in answer_text
    assert "A. 3.14" not in answer_text
    assert "解析：按小数乘法计算。" in answer_text


def test_exam_questions_restart_numbering_by_section_and_use_hanging_layout(tmp_path):
    payload = {
        "title": "Restart Numbering Exam",
        "subject": "Language",
        "grade": "Grade 7",
        "duration": "45 min",
        "total_score": "30",
        "sections": [
            {
                "title": "I. Basics",
                "questions": [
                    {"stem": "Alpha question", "answer": "A1", "score": "3"},
                    {"stem": "Beta question", "answer": "B2", "score": "4"},
                ],
            },
            {
                "title": "II. Reading",
                "questions": [
                    {"stem": "Gamma question", "answer": "C3", "score": "5"},
                    {
                        "stem": "Delta question with a longer prompt that should wrap under the stem text instead of the number marker",
                        "answer": "D4",
                        "score": "6",
                    },
                ],
            },
        ],
    }

    outputs = write_exam_paper_docx_files("default_exam", tmp_path, payload=payload)

    student = Document(outputs.student_docx)
    student_text = _all_docx_text(student)
    assert "1. Alpha question" in student_text
    assert "2. Beta question" in student_text
    assert "1. Gamma question" in student_text
    assert "2. Delta question" in student_text
    assert "3. Gamma question" not in student_text
    question_paragraph = next(
        paragraph for paragraph in student.paragraphs if paragraph.text.startswith("1. Alpha question")
    )
    assert question_paragraph.style.name == "Exam Question Stem"
    question_format = student.styles["Exam Question Stem"].paragraph_format
    assert question_format.space_before.pt == 4
    assert question_format.space_after.pt == 1
    assert round(question_format.left_indent.cm, 2) == 0.72
    assert round(question_format.first_line_indent.cm, 2) == -0.72

    answer = Document(outputs.answer_key_docx)
    answer_text = _all_docx_text(answer)
    assert "1. A1" in answer_text
    assert "2. B2" in answer_text
    assert "1. C3" in answer_text
    assert "2. D4" in answer_text
    assert "3. C3" not in answer_text


def test_exam_text_response_questions_reserve_answer_space(tmp_path):
    payload = {
        "title": "Answer Space Exam",
        "subject": "Language",
        "grade": "Grade 7",
        "duration": "45 min",
        "total_score": "30",
        "sections": [
            {
                "title": "I. Basics",
                "questions": [
                    {
                        "type": "single_choice",
                        "stem": "Choose the right answer.",
                        "options": ["A. One", "B. Two"],
                        "answer": "B",
                        "score": "3",
                    },
                    {
                        "stem": "Complete the sentence: ____.",
                        "answer": "sample",
                        "score": "2",
                    },
                ],
            },
            {
                "title": "II. Reading",
                "questions": [
                    {
                        "type": "short_answer",
                        "stem": "Explain the character's two actions.",
                        "answer": "Helped others; kept learning",
                        "score": "9",
                        "answer_lines": 3,
                    },
                    {
                        "type": "writing",
                        "stem": "Write a short paragraph of no fewer than 200 words.",
                        "answer": "Rubric answer",
                        "score": "12",
                    },
                ],
            },
        ],
    }

    outputs = write_exam_paper_docx_files("default_exam", tmp_path, payload=payload)

    student = Document(outputs.student_docx)
    answer_spaces = [
        paragraph
        for paragraph in student.paragraphs
        if paragraph.style.name == "Exam Answer Space"
    ]
    assert len(answer_spaces) == 8
    assert all("<w:bottom" in paragraph._p.xml for paragraph in answer_spaces)
    assert student.styles["Exam Answer Space"].paragraph_format.space_before.pt == 2
    assert round(student.styles["Exam Answer Space"].paragraph_format.left_indent.cm, 2) == 0.72

    student_text = _all_docx_text(student)
    assert "1. Choose the right answer." in student_text
    assert "2. Complete the sentence: ____." in student_text
    assert "1. Explain the character's two actions." in student_text
    assert "2. Write a short paragraph" in student_text

    answer = Document(outputs.answer_key_docx)
    assert not [
        paragraph
        for paragraph in answer.paragraphs
        if paragraph.style.name == "Exam Answer Space"
    ]


def test_exam_math_solution_questions_use_free_answer_area(tmp_path):
    payload = {
        "title": "Math Free Area Exam",
        "subject": "数学",
        "grade": "Grade 5",
        "duration": "45 min",
        "total_score": "20",
        "sections": [
            {
                "title": "I. Calculation",
                "type": "calculation",
                "questions": [
                    {
                        "type": "single_choice",
                        "stem": "1 + 1 = ?",
                        "options": ["A. 1", "B. 2"],
                        "answer": "B",
                        "score": "2",
                    },
                    {
                        "type": "fill_blank",
                        "stem": "3 x 4 = ____.",
                        "answer": "12",
                        "score": "2",
                    },
                    {
                        "type": "calculation",
                        "stem": "Calculate 36 x 24 and show your process.",
                        "answer": "864",
                        "score": "8",
                    },
                    {
                        "stem": "Prove the two angles are equal.",
                        "answer": "Use parallel-line angle relationships.",
                        "score": "8",
                        "answer_area": {"kind": "free", "lines": 6},
                    },
                ],
            }
        ],
    }

    outputs = write_exam_paper_docx_files("default_exam", tmp_path, payload=payload)

    student = Document(outputs.student_docx)
    line_spaces = [
        paragraph
        for paragraph in student.paragraphs
        if paragraph.style.name == "Exam Answer Space"
    ]
    free_tables = [
        table
        for table in student.tables
        if table.cell(0, 0).paragraphs
        and table.cell(0, 0).paragraphs[0].style.name == "Exam Free Answer Area"
    ]
    assert not line_spaces
    assert len(free_tables) == 2
    assert all(table.cell(0, 0).text == "" for table in free_tables)
    assert [
        round(table.rows[0].height.cm, 2)
        for table in free_tables
    ] == [2.88, 4.32]
    table_indent = free_tables[0]._tbl.tblPr.first_child_found_in("w:tblInd")
    assert table_indent is not None
    assert table_indent.get(qn("w:w")) == str(int(0.72 * 567))

    answer = Document(outputs.answer_key_docx)
    assert not [
        table
        for table in answer.tables
        if table.cell(0, 0).paragraphs
        and table.cell(0, 0).paragraphs[0].style.name == "Exam Free Answer Area"
    ]


def test_exam_style_sample_docx_uses_current_master_file(tmp_path):
    config = ExamPaperConfig()
    copied = create_exam_blank_master_copy(
        config,
        "default_exam",
        output_dir=tmp_path / "masters",
    )
    config.custom_blank_styles.append(copied)
    master_path = ensure_current_exam_blank_master_docx(config, copied.style_id)

    master = Document(master_path)
    master.add_paragraph("自定义母版标识")
    master.save(master_path)

    sample_path = write_exam_blank_style_sample_docx(
        copied.style_id,
        tmp_path / "samples",
        config=config,
    )
    text = _all_docx_text(Document(sample_path))
    assert "自定义母版标识" in text
    assert "学生卷" in text


def test_exam_default_master_file_is_stable_builtin_docx():
    path = ensure_builtin_exam_master_docx("default_exam")

    assert path.exists()
    assert path.parent == BUILTIN_EXAM_MASTER_DIR
    assert path.name == "default_exam_v20.docx"
    document = Document(path)
    text = _all_docx_text(document)
    assert "{{af_title}}" in text
    assert "{{af_version}}" in text
    assert "科目：{{af_subject}}" in text
    assert "年级：{{af_grade}}" in text
    assert "考试时间：{{af_duration}}" in text
    assert "满分：{{af_total_score}}" in text
    assert "占位符保持原样；页眉页脚和密封线由试卷卷面决定。" in text
    assert "工作台字段：" not in text
    assert "生成结果字段：" not in text
    assert "题目从 {{af_questions}} 处生成。" not in text
    assert "{{af_questions}}" in text
    assert "{{af_answer_area}}" in text
    assert "默认试卷 · 空白母版" not in text
    assert document.paragraphs[0].text == "{{af_title}}"
    assert document.tables[0].rows[0].cells[0].text == "题号"
    assert len(document.tables[0].columns) == 8
    section_row = document.tables[1]
    assert len(section_row.rows) == 1
    assert len(section_row.columns) == 2
    grid_widths = [
        column.get(qn("w:w"))
        for column in section_row._tbl.tblGrid.iterchildren(tag=qn("w:gridCol"))
    ]
    assert grid_widths == ["1846", "7083"]
    score_cell = section_row.cell(0, 0)
    title_cell = section_row.cell(0, 1)
    assert score_cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    assert title_cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    assert len(score_cell.tables) == 1
    assert score_cell.tables[0].rows[0].cells[0].text == "得分"
    assert "一、题目区" in title_cell.text
    title_paragraph = title_cell.paragraphs[0]
    assert title_paragraph.paragraph_format.space_before.pt == 0
    assert title_paragraph.paragraph_format.space_after.pt == 0
    assert title_paragraph.paragraph_format.line_spacing == 1.0


def test_exam_default_master_uses_header_sealed_zone_not_outer_layout_table():
    path = ensure_builtin_exam_master_docx("default_exam")
    document = Document(path)

    assert document.paragraphs[0].text == "{{af_title}}"
    assert len(document.tables[0].columns) == 8
    section = document.sections[0]
    assert section.different_first_page_header_footer is True
    first_page_header_xml = _sealed_header_xml_from_package(path)
    regular_header_xml = section.header._element.xml
    assert "_x0000_s1026" in first_page_header_xml
    assert first_page_header_xml.count("<w:pict") == 7
    assert first_page_header_xml.count("AlternateContent") == 4
    assert "xmlns:w15=" in first_page_header_xml
    assert "xmlns:wpsCustomData=" in first_page_header_xml
    assert "ExamSealDottedLine" in first_page_header_xml
    assert "ExamSealDottedTextGuide" in first_page_header_xml
    assert "ExamSealSolidLineTop" in first_page_header_xml
    assert "ExamSealSolidLineBottom" in first_page_header_xml
    assert "ExamSealLabelBox" in first_page_header_xml
    assert first_page_header_xml.count("rotation:-5898240f") == 2
    assert "margin-left:-451.1pt" in first_page_header_xml
    assert "margin-left:-343.2pt" in first_page_header_xml
    assert "mso-width-relative:page" in first_page_header_xml
    assert "\u5b66\u6821\uff1a" in first_page_header_xml
    assert "\u73ed\u7ea7\uff1a" in first_page_header_xml
    assert "\u5b66\u53f7\uff1a" in first_page_header_xml
    assert "\u59d3\u540d\uff1a" in first_page_header_xml
    assert "margin-left:-102pt" not in first_page_header_xml
    assert "ExamSealDottedLine" not in regular_header_xml
    assert "ExamSealLabelBox" not in regular_header_xml
    assert "rotation:-5898240f" not in regular_header_xml


def test_exam_default_master_exposes_runtime_control_markers():
    path = ensure_builtin_exam_master_docx("default_exam")
    document = Document(path)

    marker_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Exam Control Marker"
    ]
    assert len(marker_paragraphs) == 2
    assert {paragraph.text for paragraph in marker_paragraphs} == {
        "{{af_questions}}",
        "{{af_answer_area}}",
    }
    assert all(not paragraph.runs[0].font.hidden for paragraph in marker_paragraphs)
    assert {paragraph.style.name for paragraph in marker_paragraphs} == {"Exam Control Marker"}


def test_exam_default_master_docx_is_generated_without_sample_answers(tmp_path):
    path = write_exam_blank_master_docx("default_exam", tmp_path)

    assert path.exists()
    document = Document(path)
    text = _all_docx_text(document)
    assert "{{af_title}}" in text
    assert "{{af_version}}" in text
    assert "科目：{{af_subject}}" in text
    assert "年级：{{af_grade}}" in text
    assert "考试时间：{{af_duration}}" in text
    assert "满分：{{af_total_score}}" in text
    assert "占位符保持原样" in text
    assert "工作台字段：" not in text
    assert "题目从 {{af_questions}} 处生成。" not in text
    assert "默认试卷 · 空白母版" not in text
    assert "题目区" in text
    assert "{{af_questions}}" in text
    assert "{{af_answer_area}}" in text
    assert "答案：C" not in text
    assert "七年级语文期中测试样张" not in text
    assert "一、积累与运用" not in text
