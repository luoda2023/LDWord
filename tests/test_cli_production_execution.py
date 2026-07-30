from __future__ import annotations

from dataclasses import replace
from pathlib import Path
import sys

from docx import Document
import pytest

import main
import src.cli_runner as cli_runner
from src.cli_runner import (
    CLI_EXIT_CANCELLED,
    CLI_EXIT_FAILED,
    CLI_EXIT_PARTIAL,
    CLI_EXIT_SUCCESS,
    CliExecutionResources,
    _emit_execution_summary,
    _resolve_cli_resources,
    run as run_cli,
)
from src.services.execution_session import execution_session_frozen_input_path
from src.config.library import load_scene_from_library, load_template_from_library
from src.config.loader import save_scene
from src.config.material_context import MaterialExecutionContext
from src.config.scene import SceneWorkspace
from src.services.console_output import console_print
from src.services.production_execution import (
    ProductionExecutionRequest,
    execute_production_request,
)


class _StrictCp936Stream:
    encoding = "cp936"

    def __init__(self) -> None:
        self.payload = bytearray()

    def write(self, text: str) -> int:
        encoded = text.encode(self.encoding)
        self.payload.extend(encoded)
        return len(text)

    def flush(self) -> None:
        return None

    def text(self) -> str:
        return bytes(self.payload).decode(self.encoding)


def _default_request(input_path: Path, output_root: Path) -> ProductionExecutionRequest:
    resources = _resolve_cli_resources(template_path=None, scene_path=None)
    return ProductionExecutionRequest(
        input_path=input_path,
        output_root=output_root,
        mode_id=resources.mode_id,
        scene=resources.scene,
        template=resources.template,
        plan_id=resources.plan_id,
        plan_path=resources.plan_path,
        plan_source_type=resources.plan_source_type,
        template_id=resources.template_id,
        template_path=resources.template_path,
        template_source_type=resources.template_source_type,
        material_context=MaterialExecutionContext(),
    )


def _write_docx(path: Path, text: str = "source") -> bytes:
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path.read_bytes()


def _cli_resources(mode_id: str) -> CliExecutionResources:
    if mode_id == "official":
        plan_id = "official"
        template_id = "official_gbt"
    else:
        plan_id = "custom"
        template_id = "default"
    return CliExecutionResources(
        mode_id=mode_id,
        scene=load_scene_from_library(plan_id, mode_id=mode_id),
        template=load_template_from_library(template_id, mode_id=mode_id),
        plan_id=plan_id,
        plan_path="",
        plan_source_type="builtin",
        template_id=template_id,
        template_path="",
        template_source_type="builtin",
    )


def test_console_output_and_argparse_are_cp936_safe(monkeypatch) -> None:
    stdout = _StrictCp936Stream()
    stderr = _StrictCp936Stream()
    monkeypatch.setattr(sys, "stdout", stdout)
    monkeypatch.setattr(sys, "stderr", stderr)

    console_print("状态✅ 路径😀.docx")
    with pytest.raises(SystemExit) as exc_info:
        main.parse_args(["--font-engine", "😀"])

    assert exc_info.value.code == 2
    assert "状态? 路径?.docx" in stdout.text()
    assert "invalid choice" in stderr.text()


def test_cli_defaults_match_gui_custom_default_binding() -> None:
    resources = _resolve_cli_resources(template_path=None, scene_path=None)

    assert resources.mode_id == "custom"
    assert resources.plan_id == "custom"
    assert resources.template_id == "default"
    assert resources.scene.mode_id == "custom"
    assert resources.scene.template_id == "default"


def test_explicit_scene_resolves_its_mode_scoped_template(tmp_path) -> None:
    scene_path = tmp_path / "explicit_thesis.json"
    save_scene(
        SceneWorkspace(
            scene_id="explicit_thesis",
            mode_id="thesis",
            template_id="thesis_gbt",
            compatible_template_ids=["thesis_gbt"],
        ),
        scene_path,
    )

    resources = _resolve_cli_resources(
        template_path=None,
        scene_path=str(scene_path),
    )

    assert resources.mode_id == "thesis"
    assert resources.plan_id == "explicit_thesis"
    assert resources.plan_path == str(scene_path.resolve())
    assert resources.plan_source_type == "external"
    assert resources.template_id == "thesis_gbt"
    assert "thesis" in {part.casefold() for part in Path(resources.template_path).parts}


@pytest.mark.parametrize(
    ("scene_id", "mode_id", "error_code"),
    (
        ("", "custom", "plan_identity_missing"),
        ("explicit", "", "plan_mode_missing"),
    ),
)
def test_explicit_scene_requires_complete_identity(
    tmp_path,
    scene_id,
    mode_id,
    error_code,
) -> None:
    scene_path = tmp_path / "must_not_supply_identity.json"
    save_scene(
        SceneWorkspace(
            scene_id=scene_id,
            mode_id=mode_id,
            template_id="default",
            compatible_template_ids=["default"],
        ),
        scene_path,
    )

    with pytest.raises(ValueError, match=error_code):
        _resolve_cli_resources(
            template_path=None,
            scene_path=str(scene_path),
        )


def test_explicit_resource_paths_do_not_query_runtime_library(
    tmp_path,
    monkeypatch,
) -> None:
    from src.services.execution_session import resolution as execution_session_resolution
    import src.services.production_execution as production_execution

    scene_path = tmp_path / "external-plan.json"
    save_scene(
        SceneWorkspace(
            scene_id="external_plan",
            mode_id="custom",
            template_id="default",
            compatible_template_ids=["default"],
        ),
        scene_path,
    )
    resources = _resolve_cli_resources(
        template_path=None,
        scene_path=str(scene_path),
    )
    source = tmp_path / "source.docx"
    _write_docx(source)

    def _must_not_query_library(*_args, **_kwargs):
        raise AssertionError("explicit resources must not query the runtime library")

    monkeypatch.setattr(
        execution_session_resolution,
        "get_scene_entry",
        _must_not_query_library,
    )
    monkeypatch.setattr(
        execution_session_resolution,
        "get_template_entry",
        _must_not_query_library,
    )
    monkeypatch.setattr(
        production_execution,
        "_execute_frozen_session",
        lambda _request, _snapshot, **_kwargs: {"status": "success"},
    )

    payload = execute_production_request(
        ProductionExecutionRequest(
            input_path=source,
            output_root=tmp_path / "output",
            mode_id=resources.mode_id,
            scene=resources.scene,
            template=resources.template,
            plan_id=resources.plan_id,
            plan_path=resources.plan_path,
            plan_source_type=resources.plan_source_type,
            template_id=resources.template_id,
            template_path=resources.template_path,
            template_source_type=resources.template_source_type,
            material_context=MaterialExecutionContext(),
        )
    )

    assert payload["status"] == "success"


def test_production_facade_freezes_input_and_cleans_after_success(
    tmp_path,
    monkeypatch,
) -> None:
    import src.services.production_execution as production_execution

    source = tmp_path / "source.docx"
    source_v1 = _write_docx(source, "source-v1")
    frozen_paths: list[Path] = []

    def fake_execute(request, snapshot, **_kwargs):
        frozen = execution_session_frozen_input_path(snapshot)
        assert frozen is not None
        frozen_paths.append(frozen)
        assert frozen != request.input_path
        assert frozen.read_bytes() == source_v1
        _write_docx(source, "source-v2")
        return {
            "status": "success",
            "output_path": str(Path(snapshot.output_namespace) / "final.docx"),
            "output_paths": {
                "final": str(Path(snapshot.output_namespace) / "final.docx")
            },
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }

    monkeypatch.setattr(production_execution, "_execute_frozen_session", fake_execute)
    payload = execute_production_request(_default_request(source, tmp_path / "output"))

    assert payload["status"] == "success"
    assert source.read_bytes() != source_v1
    assert Document(source).paragraphs[0].text == "source-v2"
    assert frozen_paths and not frozen_paths[0].exists()


def test_production_facade_rejects_noncanonical_mode_before_snapshot_and_runner(
    tmp_path,
    monkeypatch,
) -> None:
    import src.services.production_execution as production_execution

    source = tmp_path / "source.docx"
    _write_docx(source)
    scene = load_scene_from_library("official", mode_id="official")
    scene.mode_id = "custom"
    template = load_template_from_library("official_gbt", mode_id="official")
    observed_modes: list[tuple[str, str]] = []

    def fake_execute(request, snapshot, **_kwargs):
        observed_modes.append((request.mode_id, snapshot.mode_id))
        return {"status": "success"}

    monkeypatch.setattr(
        production_execution,
        "_execute_frozen_session",
        fake_execute,
    )

    payload = execute_production_request(
        ProductionExecutionRequest(
            input_path=source,
            output_root=tmp_path / "output",
            mode_id="official_document",
            scene=scene,
            template=template,
            plan_id="official",
            template_id="official_gbt",
            document_type_id="notice",
            material_context=MaterialExecutionContext(),
        )
    )

    assert payload["status"] == "failed"
    assert payload["error_text"] == (
        "execution_mode_not_canonical:requested_mode_id:"
        "official_document:official"
    )
    assert observed_modes == []


def test_headless_production_facade_rejects_non_official_document_type(
    tmp_path,
    monkeypatch,
) -> None:
    import src.services.production_execution as production_execution

    source = tmp_path / "source.docx"
    _write_docx(source)
    monkeypatch.setattr(
        production_execution,
        "_execute_frozen_session",
        lambda *_args, **_kwargs: pytest.fail(
            "a non-official document type must not reach the runner"
        ),
    )

    payload = execute_production_request(
        replace(
            _default_request(source, tmp_path / "output"),
            document_type_id="notice",
        )
    )

    assert payload["status"] == "failed"
    assert (
        "official_document_type_not_applicable:custom:notice"
        in str(payload["error_text"])
    )


def test_production_facade_cleans_frozen_input_after_failure(
    tmp_path,
    monkeypatch,
) -> None:
    import src.services.production_execution as production_execution

    source = tmp_path / "source.docx"
    _write_docx(source)
    frozen_paths: list[Path] = []

    def fail_execute(_request, snapshot, **_kwargs):
        frozen = execution_session_frozen_input_path(snapshot)
        assert frozen is not None
        frozen_paths.append(frozen)
        raise RuntimeError("runner exploded 😀")

    monkeypatch.setattr(production_execution, "_execute_frozen_session", fail_execute)
    payload = execute_production_request(_default_request(source, tmp_path / "output"))

    assert payload["status"] == "failed"
    assert "runner exploded" in str(payload["error_text"])
    assert frozen_paths and not frozen_paths[0].exists()


def test_production_facade_cleans_frozen_input_after_keyboard_cancel(
    tmp_path,
    monkeypatch,
) -> None:
    import src.services.production_execution as production_execution

    source = tmp_path / "source.docx"
    _write_docx(source)
    frozen_paths: list[Path] = []

    def cancel_execute(_request, snapshot, **_kwargs):
        frozen = execution_session_frozen_input_path(snapshot)
        assert frozen is not None
        frozen_paths.append(frozen)
        raise KeyboardInterrupt

    monkeypatch.setattr(
        production_execution,
        "_execute_frozen_session",
        cancel_execute,
    )
    payload = execute_production_request(_default_request(source, tmp_path / "output"))

    assert payload["status"] == "cancelled"
    assert frozen_paths and not frozen_paths[0].exists()


def test_production_facade_cancel_check_reaches_runner_without_progress(
    tmp_path,
    monkeypatch,
) -> None:
    import src.services.production_runtime.execution_runtime as execution_runtime

    source = tmp_path / "source.docx"
    _write_docx(source)
    cancel_calls = 0
    runner_observations: list[bool] = []

    def cancel_check() -> bool:
        nonlocal cancel_calls
        cancel_calls += 1
        return cancel_calls >= 2

    class _PollingRunner:
        def __init__(self, **kwargs):
            self.execution_session_snapshot = kwargs["execution_session"]

        def run(self, _progress_cb, runner_cancel_check):
            cancelled = bool(runner_cancel_check())
            runner_observations.append(cancelled)
            return {
                "status": "cancelled" if cancelled else "success",
                "output_path": "",
                "report_paths": [],
                "failed_count": 0,
                "error_text": "",
            }

    monkeypatch.setattr(
        execution_runtime,
        "WorkbenchProductionRunner",
        _PollingRunner,
    )

    payload = execute_production_request(
        _default_request(source, tmp_path / "output"),
        cancel_check=cancel_check,
    )

    assert runner_observations == [True]
    assert payload["status"] == "cancelled"


def test_cleanup_failure_downgrades_success_to_partial(
    tmp_path,
    monkeypatch,
) -> None:
    import src.services.production_execution as production_execution

    source = tmp_path / "source.docx"
    _write_docx(source)
    frozen_paths: list[Path] = []
    real_cleanup = production_execution.cleanup_execution_session_resources

    def successful_execute(_request, snapshot, **_kwargs):
        frozen = execution_session_frozen_input_path(snapshot)
        assert frozen is not None
        frozen_paths.append(frozen)
        return {
            "status": "success",
            "output_path": "",
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }

    def cleanup_with_issue(snapshot):
        assert real_cleanup(snapshot) == ()
        return ("execution_resource_cleanup_failed:input_document:source.docx",)

    monkeypatch.setattr(
        production_execution,
        "_execute_frozen_session",
        successful_execute,
    )
    monkeypatch.setattr(
        production_execution,
        "cleanup_execution_session_resources",
        cleanup_with_issue,
    )
    payload = execute_production_request(_default_request(source, tmp_path / "output"))

    assert payload["status"] == "partial_success"
    assert payload["artifact_failure_count"] == 1
    assert frozen_paths and not frozen_paths[0].exists()


def test_production_facade_rejects_output_root_that_is_input(tmp_path) -> None:
    source = tmp_path / "source.docx"
    _write_docx(source)

    payload = execute_production_request(_default_request(source, source))

    assert payload["status"] == "failed"
    assert payload["error_text"] == "output_root_overwrites_input_document"


def test_cli_real_production_path_smoke_and_resource_cleanup(tmp_path) -> None:
    source = tmp_path / "input.docx"
    document = Document()
    document.add_heading("Smoke", level=1)
    document.add_paragraph("Body")
    document.save(source)
    source_bytes = source.read_bytes()
    output_root = tmp_path / "output"

    exit_code = run_cli(source, output_dir=output_root)

    assert exit_code == CLI_EXIT_SUCCESS
    namespaces = [path for path in output_root.iterdir() if path.is_dir()]
    assert len(namespaces) == 1
    namespace = namespaces[0]
    assert source.read_bytes() == source_bytes
    assert (namespace / "input_formatted.docx").is_file()
    assert (namespace / "execution_session.json").is_file()
    assert not (namespace / ".execution_resources").exists()


def test_cli_real_production_path_accepts_explicit_plan_and_template(
    tmp_path,
) -> None:
    root = Path(__file__).resolve().parents[1]
    scene_path = root / "config_library" / "plans" / "custom" / "builtin" / "custom.json"
    template_path = (
        root
        / "config_library"
        / "templates"
        / "custom"
        / "builtin"
        / "default.json"
    )
    source = tmp_path / "explicit-input.docx"
    source_bytes = _write_docx(source, "explicit resources")
    output_root = tmp_path / "explicit-output"

    exit_code = run_cli(
        source,
        scene_path=str(scene_path),
        template_path=str(template_path),
        output_dir=output_root,
    )

    assert exit_code == CLI_EXIT_SUCCESS
    namespaces = [path for path in output_root.iterdir() if path.is_dir()]
    assert len(namespaces) == 1
    namespace = namespaces[0]
    assert source.read_bytes() == source_bytes
    assert (namespace / "explicit-input_formatted.docx").is_file()
    assert (namespace / "execution_session.json").is_file()
    assert not (namespace / ".execution_resources").exists()


@pytest.mark.parametrize(
    ("filename", "create_file", "expected_message"),
    (
        ("missing.docx", False, "输入文件不存在"),
        ("unsupported.txt", True, "不支持的文件格式"),
    ),
)
def test_cli_rejects_invalid_input_before_production_execution(
    tmp_path,
    monkeypatch,
    capsys,
    filename,
    create_file,
    expected_message,
) -> None:
    source = tmp_path / filename
    if create_file:
        source.write_text("not a docx", encoding="utf-8")
    monkeypatch.setattr(
        cli_runner,
        "execute_production_request",
        lambda *_args, **_kwargs: pytest.fail(
            "invalid input must not reach production execution"
        ),
    )

    exit_code = run_cli(source)

    assert exit_code == CLI_EXIT_FAILED
    assert expected_message in capsys.readouterr().out


def test_cli_official_mode_requires_explicit_document_type(
    tmp_path,
    monkeypatch,
    capsys,
) -> None:
    source = tmp_path / "input.docx"
    _write_docx(source)
    monkeypatch.setattr(
        cli_runner,
        "_resolve_cli_resources",
        lambda **_kwargs: _cli_resources("official"),
    )
    monkeypatch.setattr(
        cli_runner,
        "execute_production_request",
        lambda *_args, **_kwargs: pytest.fail(
            "execution must not start without an explicit official document type"
        ),
    )

    exit_code = run_cli(source)

    assert exit_code == CLI_EXIT_FAILED
    assert "official_document_type_missing" in capsys.readouterr().out


def test_cli_rejects_document_type_outside_official_mode(
    tmp_path,
    monkeypatch,
    capsys,
) -> None:
    source = tmp_path / "input.docx"
    _write_docx(source)
    monkeypatch.setattr(
        cli_runner,
        "_resolve_cli_resources",
        lambda **_kwargs: _cli_resources("custom"),
    )
    monkeypatch.setattr(
        cli_runner,
        "execute_production_request",
        lambda *_args, **_kwargs: pytest.fail(
            "execution must not start with an inapplicable document type"
        ),
    )

    exit_code = run_cli(source, document_type_id="notice")

    assert exit_code == CLI_EXIT_FAILED
    assert (
        "official_document_type_not_applicable:custom:notice"
        in capsys.readouterr().out
    )


def test_cli_propagates_explicit_official_document_type(
    tmp_path,
    monkeypatch,
) -> None:
    source = tmp_path / "input.docx"
    _write_docx(source)
    captured_requests: list[ProductionExecutionRequest] = []
    monkeypatch.setattr(
        cli_runner,
        "_resolve_cli_resources",
        lambda **_kwargs: _cli_resources("official"),
    )

    def fake_execute(request, **_kwargs):
        captured_requests.append(request)
        return {"status": "success"}

    monkeypatch.setattr(cli_runner, "execute_production_request", fake_execute)

    exit_code = run_cli(source, document_type_id="letter")

    assert exit_code == CLI_EXIT_SUCCESS
    assert len(captured_requests) == 1
    assert captured_requests[0].document_type_id == "letter"


@pytest.mark.parametrize(
    ("payload", "expected_exit"),
    [
        ({"status": "success"}, CLI_EXIT_SUCCESS),
        ({"status": "partial_success"}, CLI_EXIT_PARTIAL),
        (
            {"status": "success", "artifact_failure_count": 1},
            CLI_EXIT_PARTIAL,
        ),
        ({"status": "failed", "error_text": "failed"}, CLI_EXIT_FAILED),
        ({"status": "cancelled"}, CLI_EXIT_CANCELLED),
    ],
)
def test_cli_exit_code_preserves_terminal_status(payload, expected_exit) -> None:
    assert _emit_execution_summary(payload) == expected_exit


def test_cli_artifact_summary_is_cp936_safe() -> None:
    stream = _StrictCp936Stream()
    previous = sys.stdout
    sys.stdout = stream
    try:
        exit_code = _emit_execution_summary(
            {
                "status": "partial_success",
                "output_paths": {"final": "C:/输出/😀.docx"},
                "report_paths": ["C:/输出/报告✅.json"],
                "artifact_failure_count": 1,
                "error_text": "辅助产物😀失败",
            }
        )
    finally:
        sys.stdout = previous

    assert exit_code == CLI_EXIT_PARTIAL
    output = stream.text()
    assert "[ARTIFACT output:final]" in output
    assert "[PARTIAL]" in output


def test_cli_and_service_do_not_own_a_second_pipeline_implementation() -> None:
    root = Path(__file__).resolve().parents[1]
    cli_source = (root / "src" / "cli_runner.py").read_text(encoding="utf-8")
    facade_source = (root / "src" / "services" / "production_execution.py").read_text(
        encoding="utf-8"
    )

    assert "Pipeline(" not in cli_source
    assert "Pipeline(" not in facade_source
    assert "WorkbenchProductionRunner" in facade_source
    assert "ExecutionWorker" not in facade_source
    assert "execution_result_status" in facade_source
    assert "normalize_execution_result" in facade_source
