from __future__ import annotations

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.config.scene import ContentVisibilityRule, DeliveryPreset
from src.pipeline.runner import _apply_content_visibility_rules
from src.shared.engine.block_visibility import (
    BLOCK_VISIBILITY_RECEIPT_VERSION,
    BlockVisibilityPlanError,
    BlockVisibilityReceipt,
    apply_block_visibility,
    build_block_visibility_plan,
)


def _removing_preset(selector: str = "answer") -> DeliveryPreset:
    return DeliveryPreset(
        preset_id="student",
        content_visibility_rules=[
            ContentVisibilityRule(selector=selector, action="remove")
        ],
    )


def _append_content_image_sentinel(document, marker: str) -> None:
    paragraph = document.add_paragraph()
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), "91")
    bookmark_start.set(qn("w:name"), marker)
    paragraph._p.append(bookmark_start)
    paragraph.add_run(marker)
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), "91")
    paragraph._p.append(bookmark_end)


def test_direct_body_range_removes_whole_table_and_returns_canonical_receipt():
    document = Document()
    document.add_paragraph("Before")
    document.add_paragraph("{{#visibility:answer}}")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table answer A"
    table.cell(0, 1).text = "Table answer B"
    document.add_paragraph("{{/visibility:answer}}")
    document.add_paragraph("After")

    receipt = _apply_content_visibility_rules(document, _removing_preset())

    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Before",
        "After",
    ]
    assert document.tables == []
    assert receipt["removed_paragraph_count"] == 4
    assert receipt["removed_table_count"] == 1
    assert receipt["removed_body_element_count"] == 3
    assert receipt["stripped_marker_paragraph_count"] == 0
    assert receipt["remove_selectors"] == ["answer"]
    assert receipt["contract_version"] == BLOCK_VISIBILITY_RECEIPT_VERSION
    assert receipt["removed_ranges"] == [
        {
            "selector": "answer",
            "start_body_index": 1,
            "end_body_index": 3,
            "body_element_count": 3,
            "paragraph_count": 4,
            "table_count": 1,
            "content_image_markers": [],
        }
    ]


def test_marker_inside_table_is_blocked_before_any_mutation():
    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{#visibility:answer}}"
    document.add_paragraph("{{/visibility:answer}}")
    before = document.element.body.xml

    with pytest.raises(BlockVisibilityPlanError) as raised:
        build_block_visibility_plan(document, remove_selectors={"answer"})

    assert document.element.body.xml == before
    assert any(
        item.code == "marker_not_direct_body"
        and item.location == "document.body[1]/nested"
        for item in raised.value.diagnostics
    )


def test_removed_range_receipt_collects_content_image_sentinel_ownership():
    marker = "LarkContentImage_0123456789abcdef"
    document = Document()
    document.add_paragraph("{{#content:answer}}")
    _append_content_image_sentinel(document, marker)
    document.add_paragraph("{{/content:answer}}")

    receipt = apply_block_visibility(document, remove_selectors={"answer"})

    assert document.paragraphs == []
    assert receipt.removed_content_image_markers == (marker,)
    assert receipt.removed_ranges[0].content_image_markers == (marker,)
    assert receipt.to_dict()["removed_content_image_markers"] == [marker]
    assert BlockVisibilityReceipt.from_dict(receipt.to_dict()) == receipt

    corrupted = receipt.to_dict()
    corrupted["removed_content_image_markers"] = []
    with pytest.raises(ValueError, match="canonical receipt payload"):
        BlockVisibilityReceipt.from_dict(corrupted)


def test_retained_block_strips_cross_run_markers_and_removes_empty_marker_paragraph():
    document = Document()
    start = document.add_paragraph()
    start.add_run("Before {{#visibility:")
    selector = start.add_run("answer")
    selector.bold = True
    start.add_run("}} after")
    document.add_paragraph("Retained answer")
    document.add_paragraph("{{/visibility:answer}}")

    receipt = apply_block_visibility(document, remove_selectors=())

    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Before  after",
        "Retained answer",
    ]
    assert start.runs[1].bold is True
    assert receipt.removed_paragraph_count == 0
    assert receipt.stripped_marker_paragraph_count == 2
    assert receipt.stripped_marker_token_count == 2
    assert receipt.removed_empty_marker_paragraph_count == 1
    assert receipt.removed_body_element_count == 1


@pytest.mark.parametrize(
    ("paragraphs", "expected_code"),
    [
        (["{{/visibility:answer}}"], "orphan_end_marker"),
        (["{{#visibility:answer}}"], "unclosed_start_marker"),
        (
            [
                "{{#visibility:answer}}",
                "{{#visibility:detail}}",
                "{{/visibility:detail}}",
                "{{/visibility:answer}}",
            ],
            "nested_marker_block",
        ),
    ],
)
def test_invalid_direct_body_marker_structures_are_blocked_atomically(
    paragraphs,
    expected_code,
):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    before = document.element.body.xml

    with pytest.raises(BlockVisibilityPlanError) as raised:
        build_block_visibility_plan(document, remove_selectors={"answer"})

    assert document.element.body.xml == before
    assert expected_code in {item.code for item in raised.value.diagnostics}
