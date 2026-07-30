"""Anonymous OOXML fixtures that model common Word/WPS compatibility surfaces."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import ctypes
import os
import struct
from ctypes import wintypes
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from PIL import Image


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
V_NS = "urn:schemas-microsoft-com:vml"
O_NS = "urn:schemas-microsoft-com:office:office"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def build_anonymous_wps_compat_docx(path: Path) -> Path:
    """Create a private-data-free DOCX with the structures from the WPS sample."""

    image_path = path.with_name("fixture-image.png")
    Image.new("RGB", (64, 40), "white").save(image_path, format="PNG")
    document = Document()
    document.add_heading("匿名兼容性样本", level=1)
    tabs = document.add_paragraph("选项 A")
    tabs.add_run()._r.append(OxmlElement("w:tab"))
    tabs.add_run("选项 B")
    tabs.add_run()._r.append(OxmlElement("w:tab"))
    tabs.add_run("选项 C")
    merged = document.add_table(rows=2, cols=2)
    merged.cell(0, 0).text = "合并"
    merged.cell(0, 0).merge(merged.cell(0, 1))
    merged.cell(1, 0).text = "左"
    merged.cell(1, 1).text = "右"
    document.add_picture(str(image_path))
    document.add_paragraph("VML 图片位置")
    section = document.sections[0]._sectPr
    section.insert(0, OxmlElement("w:type"))
    section.append(OxmlElement("w:pgNumType"))
    document.save(path)
    image_path.unlink()

    with ZipFile(path, "r") as source:
        parts = {name: source.read(name) for name in source.namelist()}

    document_root = etree.fromstring(parts["word/document.xml"])
    ns = {"w": W_NS, "wp": WP_NS, "r": R_NS}
    inline = document_root.find(".//wp:inline", namespaces=ns)
    assert inline is not None
    inline.tag = f"{{{WP_NS}}}anchor"
    anchor_attributes = {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "251659264",
        "behindDoc": "0",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }
    for name, value in anchor_attributes.items():
        inline.set(name, value)
    simple_pos = etree.Element(f"{{{WP_NS}}}simplePos", x="0", y="0")
    position_h = etree.Element(f"{{{WP_NS}}}positionH", relativeFrom="column")
    etree.SubElement(position_h, f"{{{WP_NS}}}posOffset").text = "0"
    position_v = etree.Element(f"{{{WP_NS}}}positionV", relativeFrom="paragraph")
    etree.SubElement(position_v, f"{{{WP_NS}}}posOffset").text = "0"
    wrap_none = etree.Element(f"{{{WP_NS}}}wrapNone")
    inline.insert(0, simple_pos)
    inline.insert(1, position_h)
    inline.insert(2, position_v)
    extent_index = next(
        index
        for index, child in enumerate(inline)
        if etree.QName(child).localname == "extent"
    )
    inline.insert(extent_index + 2, wrap_none)

    blip = document_root.find(".//wp:anchor//{*}blip", namespaces=ns)
    assert blip is not None
    image_relationship_id = blip.get(f"{{{R_NS}}}embed")
    assert image_relationship_id
    body = document_root.find(".//w:body", namespaces=ns)
    assert body is not None
    for ordinal in range(1, 4):
        paragraph = etree.Element(f"{{{W_NS}}}p")
        run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        pict = etree.SubElement(run, f"{{{W_NS}}}pict")
        shape = etree.SubElement(pict, f"{{{V_NS}}}shape")
        shape.set("type", "#_x0000_t75")
        shape.set("style", "height:40pt;width:64pt")
        image_data = etree.SubElement(shape, f"{{{V_NS}}}imagedata")
        image_data.set(f"{{{R_NS}}}id", image_relationship_id)
        image_data.set(f"{{{O_NS}}}title", f"匿名图片{ordinal}")
        body.insert(max(0, len(body) - 1), paragraph)
    parts["word/document.xml"] = etree.tostring(
        document_root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone="yes",
    )

    settings_root = etree.fromstring(parts["word/settings.xml"])
    attached = etree.Element(f"{{{W_NS}}}attachedTemplate")
    attached.set(f"{{{R_NS}}}id", "rIdFixtureTemplate")
    settings_root.insert(0, attached)
    shape_defaults = etree.SubElement(settings_root, f"{{{W_NS}}}shapeDefaults")
    defaults = etree.SubElement(shape_defaults, f"{{{O_NS}}}shapedefaults")
    etree.SubElement(defaults, f"{{{V_NS}}}fill")
    etree.SubElement(defaults, f"{{{V_NS}}}stroke")
    alternate = etree.SubElement(settings_root, f"{{{MC_NS}}}AlternateContent")
    fallback = etree.SubElement(alternate, f"{{{MC_NS}}}Fallback")
    etree.SubElement(fallback, f"{{{W_NS}}}compat")
    parts["word/settings.xml"] = etree.tostring(
        settings_root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone="yes",
    )
    relationships = etree.Element(f"{{{REL_NS}}}Relationships")
    relation = etree.SubElement(relationships, f"{{{REL_NS}}}Relationship")
    relation.set("Id", "rIdFixtureTemplate")
    relation.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate",
    )
    relation.set("Target", "file:///C:/anonymous/template.docx")
    relation.set("TargetMode", "External")
    parts["word/_rels/settings.xml.rels"] = etree.tostring(
        relationships,
        xml_declaration=True,
        encoding="UTF-8",
        standalone="yes",
    )

    with ZipFile(path, "w", ZIP_DEFLATED) as target:
        for name, payload in sorted(parts.items()):
            target.writestr(name, payload)
    return path


def png_bytes(size: tuple[int, int] = (32, 18)) -> bytes:
    stream = BytesIO()
    Image.new("RGB", size, "white").save(stream, format="PNG")
    return stream.getvalue()


def emf_bytes() -> bytes:
    """Create a tiny anonymous EMF using the Windows GDI recorder."""

    if os.name != "nt":
        raise RuntimeError("EMF fixture generation requires Windows")

    class RECT(ctypes.Structure):
        _fields_ = [
            ("left", ctypes.c_long),
            ("top", ctypes.c_long),
            ("right", ctypes.c_long),
            ("bottom", ctypes.c_long),
        ]

    gdi = ctypes.WinDLL("gdi32", use_last_error=True)
    gdi.CreateEnhMetaFileW.argtypes = [
        wintypes.HDC,
        wintypes.LPCWSTR,
        ctypes.POINTER(RECT),
        wintypes.LPCWSTR,
    ]
    gdi.CreateEnhMetaFileW.restype = wintypes.HDC
    gdi.Rectangle.argtypes = [
        wintypes.HDC,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
    ]
    gdi.Rectangle.restype = wintypes.BOOL
    gdi.CloseEnhMetaFile.argtypes = [wintypes.HDC]
    gdi.CloseEnhMetaFile.restype = wintypes.HANDLE
    gdi.GetEnhMetaFileBits.argtypes = [
        wintypes.HANDLE,
        wintypes.UINT,
        ctypes.c_void_p,
    ]
    gdi.GetEnhMetaFileBits.restype = wintypes.UINT
    gdi.DeleteEnhMetaFile.argtypes = [wintypes.HANDLE]
    gdi.DeleteEnhMetaFile.restype = wintypes.BOOL

    frame = RECT(0, 0, 2000, 1000)
    device = gdi.CreateEnhMetaFileW(
        None,
        None,
        ctypes.byref(frame),
        "Anonymous\0Fixture\0\0",
    )
    if not device:
        raise OSError(ctypes.get_last_error(), "CreateEnhMetaFileW failed")
    try:
        if not gdi.Rectangle(device, 100, 100, 1900, 900):
            raise OSError(ctypes.get_last_error(), "GDI Rectangle failed")
        handle = gdi.CloseEnhMetaFile(device)
        device = None
    finally:
        if device:
            gdi.CloseEnhMetaFile(device)
    if not handle:
        raise OSError(ctypes.get_last_error(), "CloseEnhMetaFile failed")
    try:
        size = gdi.GetEnhMetaFileBits(handle, 0, None)
        buffer = (ctypes.c_ubyte * size)()
        if gdi.GetEnhMetaFileBits(handle, size, buffer) != size:
            raise OSError(ctypes.get_last_error(), "GetEnhMetaFileBits failed")
        return bytes(buffer)
    finally:
        gdi.DeleteEnhMetaFile(handle)


def wmf_bytes() -> bytes:
    """Create a tiny placeable WMF using the Windows GDI recorder."""

    if os.name != "nt":
        raise RuntimeError("WMF fixture generation requires Windows")
    gdi = ctypes.WinDLL("gdi32", use_last_error=True)
    gdi.CreateMetaFileW.argtypes = [wintypes.LPCWSTR]
    gdi.CreateMetaFileW.restype = wintypes.HDC
    gdi.Rectangle.argtypes = [
        wintypes.HDC,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
    ]
    gdi.Rectangle.restype = wintypes.BOOL
    gdi.CloseMetaFile.argtypes = [wintypes.HDC]
    gdi.CloseMetaFile.restype = wintypes.HANDLE
    gdi.GetMetaFileBitsEx.argtypes = [
        wintypes.HANDLE,
        wintypes.UINT,
        ctypes.c_void_p,
    ]
    gdi.GetMetaFileBitsEx.restype = wintypes.UINT
    gdi.DeleteMetaFile.argtypes = [wintypes.HANDLE]
    gdi.DeleteMetaFile.restype = wintypes.BOOL
    device = gdi.CreateMetaFileW(None)
    if not device:
        raise OSError(ctypes.get_last_error(), "CreateMetaFileW failed")
    if not gdi.Rectangle(device, 0, 0, 200, 100):
        raise OSError(ctypes.get_last_error(), "GDI Rectangle failed")
    handle = gdi.CloseMetaFile(device)
    if not handle:
        raise OSError(ctypes.get_last_error(), "CloseMetaFile failed")
    try:
        size = gdi.GetMetaFileBitsEx(handle, 0, None)
        buffer = (ctypes.c_ubyte * size)()
        if gdi.GetMetaFileBitsEx(handle, size, buffer) != size:
            raise OSError(ctypes.get_last_error(), "GetMetaFileBitsEx failed")
        standard = bytes(buffer)
    finally:
        gdi.DeleteMetaFile(handle)
    header_without_checksum = struct.pack(
        "<IHhhhhHI",
        0x9AC6CDD7,
        0,
        0,
        0,
        200,
        100,
        1440,
        0,
    )
    checksum = 0
    for word in struct.unpack("<10H", header_without_checksum):
        checksum ^= word
    return header_without_checksum + struct.pack("<H", checksum) + standard


def build_docx_with_emf(path: Path) -> Path:
    return _build_docx_with_metafile(path, emf_bytes(), "emf", "image/x-emf")


def build_docx_with_wmf(path: Path) -> Path:
    return _build_docx_with_metafile(path, wmf_bytes(), "wmf", "image/x-wmf")


def _build_docx_with_metafile(
    path: Path,
    payload: bytes,
    extension: str,
    media_type: str,
) -> Path:
    raster_path = path.with_name("emf-placeholder.png")
    Image.new("RGB", (20, 10), "white").save(raster_path, format="PNG")
    document = Document()
    document.add_picture(str(raster_path))
    document.save(path)
    raster_path.unlink()
    with ZipFile(path, "r") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    image_part = next(
        name for name in parts if name.startswith("word/media/") and name.endswith(".png")
    )
    metafile_part = image_part[:-4] + "." + extension
    parts[metafile_part] = payload
    del parts[image_part]

    relationships = etree.fromstring(parts["word/_rels/document.xml.rels"])
    for relation in relationships.findall(f"{{{REL_NS}}}Relationship"):
        if str(relation.get("Target", "")).endswith(Path(image_part).name):
            relation.set("Target", f"media/{Path(metafile_part).name}")
    parts["word/_rels/document.xml.rels"] = etree.tostring(
        relationships,
        xml_declaration=True,
        encoding="UTF-8",
    )

    content_types = etree.fromstring(parts["[Content_Types].xml"])
    for default in content_types.findall(f"{{{CT_NS}}}Default"):
        if str(default.get("Extension", "")).casefold() == "png":
            default.set("Extension", extension)
            default.set("ContentType", media_type)
    parts["[Content_Types].xml"] = etree.tostring(
        content_types,
        xml_declaration=True,
        encoding="UTF-8",
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, payload in sorted(parts.items()):
            archive.writestr(name, payload)
    return path


__all__ = [
    "build_anonymous_wps_compat_docx",
    "build_docx_with_emf",
    "build_docx_with_wmf",
    "emf_bytes",
    "png_bytes",
    "wmf_bytes",
]
