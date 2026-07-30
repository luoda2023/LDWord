import json
import sys
import zipfile
from pathlib import Path

from docx import Document
import pytest


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.config.execution_config_integrity import execution_config_integrity_issue
from src.modules.structure.heading_recognition import HeadingRecognitionModule
from src.modules.validate.validation import ValidationModule
from src.pipeline.context import PipelineContext
from src.pipeline.runner import Pipeline
from src.pipeline.tracker import ChangeTracker
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.count_engine import (
    COUNT_PROFILE_DIR,
    CountProfileRegistryError,
    count_document,
    get_count_profile,
    load_count_profiles_from_directory,
    list_count_profiles,
)


def test_count_engine_returns_multi_metric_document_counts():
    doc = Document()
    doc.add_paragraph("\u4e2d\u6587 ABC words 123")
    doc.add_paragraph("\u53c2\u8003\u6587\u732e")
    doc.add_paragraph("[1] Zhang A. Example.")
    doc.add_paragraph("[2] Li B. \u793a\u4f8b.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "\u8868\u683c text"

    result = count_document(doc, profile_id="thesis_cn")

    assert result.profile_id == "thesis_cn"
    assert result.counts["cjk_characters"] >= 8
    assert result.counts["english_words"] >= 4
    assert result.counts["number_tokens"] >= 3
    assert result.counts["paragraph_count"] == 4
    assert result.counts["table_count"] == 1
    assert result.counts["reference_count"] == 2
    assert "thesis_cn" in result.summary()


def test_count_profile_registry_covers_scene_and_p1_profiles():
    profile_ids = {profile.profile_id for profile in list_count_profiles()}

    assert {
        "basic",
        "thesis_cn",
        "school_thesis",
        "journal_words",
        "journal_display_items",
        "administrative_sections",
        "bid_package",
        "application_word_limits",
        "attachment_inventory",
        "product_asset_inventory",
        "exam_items",
        "bilingual_parallel_text",
        "disclosure_section_inventory",
        "batch_item_inventory",
        "chapter_inventory",
        "finance_attachment_inventory",
        "patent_section_inventory",
        "word_xml_full",
    }.issubset(profile_ids)

    journal = get_count_profile("journal_words")
    assert journal.scope == "main_text"
    assert "references" in journal.excluded_scopes
    assert "tables" in journal.excluded_scopes
    assert "count_profiles" in journal.source

    application = get_count_profile("application_word_limits")
    assert application.scope == "application_body"
    assert {limit.section_id for limit in application.section_limits} >= {
        "summary",
        "basis",
        "objectives",
        "implementation",
        "budget",
    }
    summary_limit = next(
        limit for limit in application.section_limits if limit.section_id == "summary"
    )
    assert summary_limit.max_characters_no_spaces == 500
    assert summary_limit.required is True

    full = get_count_profile("word_xml_full")
    assert "headers" in full.included_scopes
    assert "footers" in full.included_scopes
    assert "footnotes" in full.included_scopes
    assert "textboxes" in full.included_scopes
    assert "tracked_changes" in full.included_scopes
    assert "field_codes" in full.included_scopes

    administrative = get_count_profile("administrative_sections")
    assert administrative.scope == "official_policy_package"
    assert "headers" in administrative.included_scopes
    assert "footers" in administrative.included_scopes
    assert "field_codes" in administrative.included_scopes

    product = get_count_profile("product_asset_inventory")
    assert product.scope == "product_sales_package"
    assert "field_codes" in product.included_scopes
    assert "figure_count" in product.primary_metrics
    assert product.count_references is False

    disclosure = get_count_profile("disclosure_section_inventory")
    assert disclosure.scope == "disclosure_package"
    assert "comments" in disclosure.included_scopes
    assert "table_count" in disclosure.primary_metrics
    assert disclosure.count_references is False

    batch = get_count_profile("batch_item_inventory")
    assert batch.scope == "batch_document_package"
    assert "textboxes" in batch.included_scopes
    assert "field_codes_count" in batch.primary_metrics
    assert batch.count_references is False

    chapter = get_count_profile("chapter_inventory")
    assert chapter.scope == "long_document_package"
    assert "field_codes" in chapter.included_scopes
    assert "figure_count" in chapter.primary_metrics
    assert any("chapter hierarchy" in note for note in chapter.notes)

    finance = get_count_profile("finance_attachment_inventory")
    assert finance.scope == "finance_quote_package"
    assert "comments" in finance.included_scopes
    assert "field_codes_count" in finance.primary_metrics
    assert any("financial audit judgment" in note for note in finance.notes)

    patent = get_count_profile("patent_section_inventory")
    assert patent.scope == "patent_document_package"
    assert "comments" in patent.included_scopes
    assert "figure_count" in patent.primary_metrics
    assert any("patentability" in note for note in patent.notes)


def test_count_profiles_load_from_external_json_files(tmp_path):
    profile_dir = tmp_path / "count_profiles"
    profile_dir.mkdir()
    (profile_dir / "custom.json").write_text(
        json.dumps(
            {
                "profiles": [
                    {
                        "profile_id": "custom_words",
                        "label": "Custom words",
                        "scope": "custom_scope",
                        "included_scopes": ["body"],
                        "excluded_scopes": ["references"],
                        "primary_metrics": ["english_words"],
                        "count_references": False,
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    profiles = load_count_profiles_from_directory(profile_dir)

    assert len(profiles) == 1
    assert profiles[0].profile_id == "custom_words"
    assert profiles[0].scope == "custom_scope"
    assert profiles[0].count_references is False
    assert profiles[0].source.endswith("custom.json")
    assert (COUNT_PROFILE_DIR / "builtin.json").exists()

    with pytest.raises(
        CountProfileRegistryError,
        match="count_profile_registry_required_profiles_missing:basic",
    ):
        load_count_profiles_from_directory(
            profile_dir,
            required_profile_ids={"basic", "custom_words"},
        )


def test_count_profile_registry_missing_or_empty_fails_closed(tmp_path):
    missing = tmp_path / "missing"
    with pytest.raises(
        CountProfileRegistryError,
        match="count_profile_registry_missing:",
    ):
        load_count_profiles_from_directory(missing)

    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(
        CountProfileRegistryError,
        match="count_profile_registry_empty:",
    ):
        load_count_profiles_from_directory(empty)

    (empty / "empty.json").write_text('{"profiles": []}', encoding="utf-8")
    with pytest.raises(
        CountProfileRegistryError,
        match="count_profile_registry_empty:",
    ):
        load_count_profiles_from_directory(empty)


def test_unknown_count_profile_cannot_emit_a_successful_count_record():
    config = ResolvedConfig(strict_mode=False)
    config.compliance_profile.count_profile_id = "not_registered"
    context = PipelineContext()
    tracker = ChangeTracker()

    with pytest.raises(KeyError, match="unknown_count_profile:not_registered"):
        ValidationModule().apply(Document(), config, tracker, context)

    assert context.count_result is None
    assert tracker.get_by_module("count_engine") == []


def test_pipeline_uses_shared_integrity_gate_for_unknown_count_profile(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    config = ResolvedConfig(strict_mode=False)
    config.output.final_docx = False
    config.compliance_profile.count_profile_id = "not_registered"

    assert (
        execution_config_integrity_issue(config)
        == "unknown_count_profile:not_registered"
    )
    result = Pipeline(
        modules=[HeadingRecognitionModule(), ValidationModule()],
        config=config,
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    assert result.success is False
    assert result.error == (
        "pipeline_configuration_invalid:unknown_count_profile:not_registered"
    )
    assert result.context is None
    assert result.tracker is None


def test_count_profile_include_exclude_scopes_affect_text_metrics():
    doc = Document()
    doc.add_paragraph("Body Alpha Beta")
    doc.add_paragraph("\u53c2\u8003\u6587\u732e")
    doc.add_paragraph("[1] Reference Words.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table Words"

    broad = count_document(doc, profile_id="basic")
    journal = count_document(doc, profile_id="journal_words")

    assert journal.profile_name == "Journal main-text words"
    assert journal.scope == "main_text"
    assert journal.counts["english_words"] == 3
    assert broad.counts["english_words"] > journal.counts["english_words"]
    assert journal.counts["reference_count"] == 1
    assert journal.counts["table_count"] == 1
    assert "references" in journal.excluded_scopes


def test_count_engine_can_include_word_package_xml_ranges(tmp_path):
    source = tmp_path / "xml_ranges.docx"
    doc = Document()
    doc.add_paragraph("Body Alpha")
    doc.sections[0].header.paragraphs[0].text = "Header Scope"
    doc.sections[0].footer.paragraphs[0].text = "Footer Scope"
    doc.save(source)
    _inject_word_xml_count_ranges(source)

    counted = count_document(Document(source), profile_id="word_xml_full")
    basic = count_document(Document(source), profile_id="basic")

    assert counted.counts["headers_count"] == 1
    assert counted.counts["footers_count"] == 1
    assert counted.counts["footnotes_count"] == 1
    assert counted.counts["endnotes_count"] == 1
    assert counted.counts["comments_count"] == 1
    assert counted.counts["textboxes_count"] == 1
    assert counted.counts["hidden_text_count"] == 1
    assert counted.counts["tracked_changes_count"] == 2
    assert counted.counts["field_codes_count"] == 1
    assert counted.counts["field_results_count"] == 1
    assert counted.counts["english_words"] > basic.counts["english_words"]
    assert basic.counts["english_words"] == 4
    assert "word_xml_package" == counted.scope


def test_validation_module_runs_count_engine_when_count_profile_is_configured(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("\u7b2c\u4e00\u7ae0 \u7eea\u8bba", level=1)
    doc.add_paragraph("\u6b63\u6587\u5185\u5bb9 Alpha beta.")
    doc.add_paragraph("\u53c2\u8003\u6587\u732e")
    doc.add_paragraph("[1] Example.")
    doc.save(source)

    config = ResolvedConfig(strict_mode=False)
    config.output.final_docx = False
    config.compliance_profile.count_profile_id = "thesis_cn"

    result = Pipeline(
        modules=[HeadingRecognitionModule(), ValidationModule()],
        config=config,
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    assert result.success is True
    assert result.context.count_result.profile_id == "thesis_cn"
    assert result.context.count_result.counts["reference_count"] == 1
    assert any(
        record.rule_name == "count_engine" and record.change_type == "count_summary"
        for record in result.tracker.get_all()
    )

    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"
    write_json_report(
        result,
        input_path=source,
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=2,
        modules_total=2,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=2,
        modules_total=2,
    )

    payload = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")
    assert payload["counts"]["profile_id"] == "thesis_cn"
    assert payload["counts"]["profile_name"] == "Chinese thesis count"
    assert "count_profiles" in payload["counts"]["profile_source"]
    assert payload["counts"]["scope"] == "academic_document"
    assert "included_scopes" in payload["counts"]
    assert payload["counts"]["counts"]["reference_count"] == 1
    assert "## \u8ba1\u6570\u53e3\u5f84" in markdown
    assert "- Profile: thesis_cn" in markdown
    assert "- Name: Chinese thesis count" in markdown
    assert "- Included scopes:" in markdown


def _inject_word_xml_count_ranges(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}

    document_xml = entries["word/document.xml"].decode("utf-8")
    extra_paragraph = """
    <w:p>
      <w:r><w:t>Visible Xml</w:t></w:r>
      <w:r><w:rPr><w:vanish/></w:rPr><w:t>Hidden Scope</w:t></w:r>
      <w:ins w:id="1" w:author="tester" w:date="2026-06-16T00:00:00Z">
        <w:r><w:t>Inserted Scope</w:t></w:r>
      </w:ins>
      <w:del w:id="2" w:author="tester" w:date="2026-06-16T00:00:00Z">
        <w:r><w:delText>Deleted Scope</w:delText></w:r>
      </w:del>
      <w:r><w:fldChar w:fldCharType="begin"/></w:r>
      <w:r><w:instrText>PAGE \\* MERGEFORMAT</w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r>
      <w:r><w:t>7</w:t></w:r>
      <w:r><w:fldChar w:fldCharType="end"/></w:r>
      <w:r>
        <w:txbxContent>
          <w:p><w:r><w:t>Textbox Scope</w:t></w:r></w:p>
        </w:txbxContent>
      </w:r>
    </w:p>
    """
    entries["word/document.xml"] = document_xml.replace(
        "</w:body>",
        f"{extra_paragraph}</w:body>",
    ).encode("utf-8")

    entries["word/footnotes.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="1"><w:p><w:r><w:t>Footnote Scope</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    ).encode("utf-8")
    entries["word/endnotes.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:endnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:endnote w:id="1"><w:p><w:r><w:t>Endnote Scope</w:t></w:r></w:p></w:endnote>'
        "</w:endnotes>"
    ).encode("utf-8")
    entries["word/comments.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:comment w:id="0"><w:p><w:r><w:t>Comment Scope</w:t></w:r></w:p></w:comment>'
        "</w:comments>"
    ).encode("utf-8")

    rels = entries["word/_rels/document.xml.rels"].decode("utf-8")
    rels = rels.replace(
        "</Relationships>",
        (
            '<Relationship Id="rIdCountFootnotes" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
            'Target="footnotes.xml"/>'
            '<Relationship Id="rIdCountEndnotes" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes" '
            'Target="endnotes.xml"/>'
            '<Relationship Id="rIdCountComments" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" '
            'Target="comments.xml"/>'
            "</Relationships>"
        ),
    )
    entries["word/_rels/document.xml.rels"] = rels.encode("utf-8")

    content_types = entries["[Content_Types].xml"].decode("utf-8")
    content_types = content_types.replace(
        "</Types>",
        (
            '<Override PartName="/word/footnotes.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
            '<Override PartName="/word/endnotes.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"/>'
            '<Override PartName="/word/comments.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
            "</Types>"
        ),
    )
    entries["[Content_Types].xml"] = content_types.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
