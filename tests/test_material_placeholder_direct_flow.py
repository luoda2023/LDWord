import pytest
from docx import Document
from PySide6.QtTest import QSignalSpy, QTest

from src.config.entity import (
    EntityArchive,
    EntityProfile,
    load_entity_archive,
    save_entity_archive,
)
from src.config.execution_target import resolve_execution_target
from src.config.material_context import MaterialExecutionContext
from src.config.material_mappings import MaterialMappingPayload, load_material_mapping
from src.config.material_preview import (
    scan_docx_placeholder_inventory,
    scan_docx_placeholders,
)
from src.config.materials import AssetItem
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.qt_api import QApplication, QLabel, QLineEdit, QPoint, QPushButton, Qt
from src.ui.bridge import PanelBridge
from src.ui.panels.assets.batch_import import _load_batch_profiles_from_path
from src.ui.panels.assets_panel import AssetsPanel
from src.ui.panels.assets.material_context_application_presenter import (
    _material_execution_field_values,
)
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner
from src.services.material_assets import (
    asset_item_payload,
    question_figure_library_metadata_history_record,
)
from src.shared.ui.sizing import resolved_control_height
from src.shared.ui.compact_row_actions import CompactRowActions
from src.shared.ui.theme import get_theme


def _app() -> QApplication:
    return QApplication.instance() or QApplication([])


_OFFICIAL_TEST_FIELD_SCOPES = {
    "organization": "fixed",
    "copy_scope": "fixed",
    "title": "floating",
    "body": "floating",
    "document_no": "floating",
    "issue_date": "floating",
    "recipient": "floating",
    "attachment_note": "floating",
    "issuer": "floating",
    "printing_org": "floating",
    "printing_date": "floating",
    "security_level": "floating",
    "urgency": "floating",
}


def test_execution_field_projection_has_one_owner_per_scope():
    values = _material_execution_field_values(
        package_fields={
            "organization": "归档机关",
            "title": "不应回退的旧标题",
        },
        task_values={"body": "本次正文"},
        fixed_keys=["organization"],
        floating_keys=["title", "body"],
    )

    assert values == {
        "organization": "归档机关",
        "body": "本次正文",
    }


def _load_material_package(
    panel: AssetsPanel,
    *,
    field_scopes: dict[str, str] | None = None,
    fields: dict[str, str] | None = None,
    profile_id: str = "test-profile",
) -> None:
    panel._material_persistence.replace_identity(None, "")
    panel.set_archive(
        EntityArchive(
            archive_id="test-package",
            package_id="test-package",
            archive_name="测试资料包",
            profiles=[
                EntityProfile(
                    profile_id=profile_id,
                    profile_name="测试资料",
                    fields=dict(fields or {}),
                    field_scopes=dict(field_scopes or {}),
                )
            ],
        )
    )
    panel._refresh_summary()
    _app().processEvents()


def _load_official_test_package(panel: AssetsPanel) -> None:
    _load_material_package(
        panel,
        field_scopes=_OFFICIAL_TEST_FIELD_SCOPES,
        profile_id="official:test",
    )


def test_placeholder_inventory_keeps_user_defined_names_and_counts_occurrences(tmp_path):
    source = tmp_path / "custom-fields.docx"
    document = Document()
    document.add_paragraph(
        "{{@text:公司}} / {{@text:company}} / {{@text:公司}} / {{@text:公司全称（盖章）}} / {{@text:字段=自定义}}"
    )
    document.add_paragraph("{{@text: 公司 }}")
    document.save(source)

    inventory = scan_docx_placeholder_inventory(source)

    assert [item.key for item in inventory] == [
        "公司",
        "company",
        "公司全称（盖章）",
        "字段=自定义",
    ]
    assert [item.occurrence_count for item in inventory] == [2, 1, 1, 1]
    assert scan_docx_placeholders(source) == [
        "公司",
        "company",
        "公司全称（盖章）",
        "字段=自定义",
    ]


def test_placeholder_inventory_and_exact_runtime_share_split_run_matching(tmp_path):
    source = tmp_path / "split-placeholder.docx"
    document = Document()
    paragraph = document.add_paragraph("公司：")
    paragraph.add_run("{{@text:公")
    tail = paragraph.add_run("司}}")
    tail.bold = True
    paragraph.add_run("；旧语法：${公司}；别名：{{@text:company}}")
    document.save(source)

    inventory = scan_docx_placeholder_inventory(source)
    assert [(item.key, item.occurrence_count) for item in inventory] == [
        ("公司", 1),
        ("company", 1),
    ]

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(
            entity_data={"公司": "阿拉维特科技"},
            field_aliases={"company": "公司"},
            exact_material_placeholders=True,
        ),
    ).run(lambda *_args: None, lambda: False)

    output = Document(payload["output_path"])
    text = "\n".join(paragraph.text for paragraph in output.paragraphs)
    assert "公司：阿拉维特科技" in text
    assert "${公司}" in text
    assert "别名：阿拉维特科技" in text


def test_execution_target_uses_current_docx_without_guessing_format_template(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    bridge = PanelBridge()

    bridge.set_current_document_path(str(source))

    target = bridge.current_execution_target()
    assert target.placeholder_source_kind == "document"
    assert target.placeholder_source_path == str(source)
    assert target.placeholder_source_label == source.name


def test_official_execution_target_uses_explicit_master_contract(monkeypatch, tmp_path):
    master_path = tmp_path / "official.docx"
    Document().save(master_path)

    class _Master:
        master_id = "official_gbt_standard"
        label = "公文执行母版"
        docx_path = master_path

    monkeypatch.setattr(
        "src.config.execution_target.get_master",
        lambda master_id, mode_id: (
            _Master()
            if (master_id, mode_id) == ("official_gbt_standard", "official")
            else None
        ),
    )
    scene = SceneWorkspace(master_id="official_gbt_standard")

    target = resolve_execution_target(
        mode_id="official",
        scene=scene,
        document_path=str(tmp_path / "input.docx"),
        official_document_type_id="notice",
    )

    assert target.placeholder_source_kind == "master"
    assert target.placeholder_source_path == str(master_path)
    assert target.binding_field_key("official_title") == "title"
    assert target.binding_required("official_title") is True
    assert target.binding_required("official_recipient") is False


def test_official_execution_target_does_not_infer_document_type_from_other_sources():
    scene = SceneWorkspace(
        master_id="official_gbt_standard",
        default_material_profile_id="official:notice",
    )

    target = resolve_execution_target(
        mode_id="official",
        scene=scene,
        document_path="input.docx",
    )

    assert target.placeholder_source_kind == "unresolved_master"
    assert target.master_id == ""
    assert target.issues == ("official_document_type_missing",)


def test_official_execution_target_rejects_unknown_explicit_document_type():
    target = resolve_execution_target(
        mode_id="official",
        scene=SceneWorkspace(master_id="official_gbt_standard"),
        official_document_type_id="not_registered",
    )

    assert target.placeholder_source_kind == "unresolved_master"
    assert target.issues == (
        "official_document_type_unknown:not_registered",
    )


def test_official_execution_target_does_not_replace_incompatible_explicit_master(
    monkeypatch,
    tmp_path,
):
    master_path = tmp_path / "notice-only.docx"
    Document().save(master_path)

    class _Master:
        master_id = "notice_only"
        label = "Notice only"
        docx_path = master_path
        source_type = "builtin"
        supported_assembly_types = ("notice",)

    calls = []

    def get_master(master_id, mode_id):
        calls.append((master_id, mode_id))
        return _Master() if master_id == "notice_only" else None

    monkeypatch.setattr("src.config.execution_target.get_master", get_master)

    target = resolve_execution_target(
        mode_id="official",
        scene=SceneWorkspace(master_id="notice_only"),
        official_document_type_id="letter",
    )

    assert calls == [("notice_only", "official")]
    assert target.placeholder_source_kind == "unresolved_master"
    assert target.master_id == "notice_only"
    assert target.master_path == str(master_path)
    assert target.issues == (
        "master_ref_incompatible:notice_only:letter",
    )


def test_official_assets_owns_token_inventory_without_reading_plan_placeholders(
    monkeypatch,
):
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    bridge.set_current_official_document_type_id(
        "notice",
        source="test",
        emit_signal=False,
    )
    bridge.set_current_scene(
        SceneWorkspace(master_id="official_gbt_standard"),
        config_id="official",
        emit_signal=False,
    )
    scan_calls: list[int] = []

    def forbidden_scan(_path):
        scan_calls.append(1)
        raise AssertionError("official form must not discover fields by scanning DOCX")

    monkeypatch.setattr(
        "src.ui.panels.assets.preview_table_presenter.scan_docx_placeholders",
        forbidden_scan,
    )
    panel = AssetsPanel(bridge)
    try:
        # The panel intentionally activates an available library package on
        # startup.  This test owns a blank package explicitly so its contract
        # does not depend on which builtin/user packages exist on the host.
        _load_material_package(panel)
        tokens = panel._scanned_placeholders()

        assert scan_calls == []
        assert tokens == []
        assert list(panel._official_fixed_field_inputs) == []
        assert panel._official_floating_field_keys == []
        assert panel._generation_status_label.text() == "字段资料 0 项"
        assert panel._generation_detail_label.text() == (
            "固定字段 0 项 · 自由字段 0 项 · 当前已填写 0/0"
        )
        assert panel._overview_preview_table.rowCount() == 0

        _load_official_test_package(panel)

        assert list(panel._official_fixed_field_inputs) == [
            "organization",
            "copy_scope",
        ]
        assert panel._official_floating_field_keys == [
            "title",
            "body",
            "document_no",
            "issue_date",
            "recipient",
            "attachment_note",
            "issuer",
            "printing_org",
            "printing_date",
            "security_level",
            "urgency",
        ]
        assert panel._generate_card._title_label.text() == "资料 Token"
        assert panel._profile_card._title_label.text() == "字段资料"
        assert panel._generation_status_label.text() == "字段资料 13 项"
        assert panel._generation_detail_label.text() == (
            "固定字段 2 项 · 自由字段 11 项 · 当前已填写 0/13"
        )
        assert panel._overview_preview_table.item(0, 0).text() == "{{@text:发文机关1}}"
        assert panel._overview_preview_table.item(0, 1).text() == "—"
        assert panel._overview_preview_table.item(0, 2).text() == "未填写"
        assert panel._overview_preview_table.item(2, 0).text() == "{{@text:标题1}}"
        assert panel._overview_preview_table.item(2, 2).text() == "工作台填写"
        assert panel._preview_auto_match_btn.text() == "校验母版"
        assert (
            panel._official_field_name_labels["organization"].text()
            == "{{@text:发文机关1}}"
        )
        organization_edit = panel._official_fixed_field_inputs["organization"]
        organization_line_edit = organization_edit.findChild(QLineEdit)
        expected_height = resolved_control_height(get_theme(), "md")
        assert organization_edit.minimumHeight() >= expected_height
        assert organization_line_edit is not None
        assert organization_line_edit.minimumHeight() == expected_height
        assert organization_line_edit.maximumHeight() == expected_height
        assert organization_edit.parentWidget().minimumHeight() >= 56
        panel.resize(1300, 900)
        panel.show()
        panel._on_section_selected("fields")
        app.processEvents()
        assert panel._official_fixed_title.parentWidget().height() == expected_height
        assert panel._add_official_fixed_btn.height() == expected_height

        organization_edit.setText("综合办公室")
        QTest.qWait(350)
        app.processEvents()
        assert panel._official_fixed_field_inputs["organization"] is organization_edit
        overview_values = {
            panel._overview_preview_table.item(row, 0).text():
            panel._overview_preview_table.item(row, 1).text()
            for row in range(panel._overview_preview_table.rowCount())
        }
        assert overview_values["{{@text:发文机关1}}"] == "综合办公室"
        assert panel._generation_status_label.text() == "字段资料 13 项"
        assert panel._generation_detail_label.text() == (
            "固定字段 2 项 · 自由字段 11 项 · 当前已填写 1/13"
        )
        assert bridge.current_material_context().entity_data["organization"] == "综合办公室"
        assert (
            panel._official_field_name_labels["organization"].text()
            == "{{@text:发文机关1}}"
        )

        panel._preview_auto_match_btn.click()
        app.processEvents()
        assert panel._official_master_preflight.status == "ok"
    finally:
        panel.close()
        app.processEvents()


def test_document_scan_is_diagnostic_until_user_adds_exact_package_fields(tmp_path):
    app = _app()
    source = tmp_path / "custom-fields.docx"
    document = Document()
    document.add_paragraph("公司：{{@text:公司}}；项目：{{@text:项目名称}}；备注：{{@text:字段=自定义}}")
    document.save(source)

    bridge = PanelBridge()
    panel = AssetsPanel(bridge)
    try:
        _load_material_package(panel)
        bridge.document_loaded.emit(str(source))
        app.processEvents()
        assert panel._generate_card._title_label.text() == "资料 Token"
        assert not hasattr(panel, "_fill_missing_btn")
        assert not hasattr(panel, "_overview_preview_refresh_btn")
        assert not hasattr(panel, "_apply_btn")
        assert panel._scanned_placeholders() == ["公司", "项目名称", "字段=自定义"]
        assert panel._selected_profile().field_scopes == {}
        assert panel._official_floating_field_keys == []
        assert panel._overview_preview_table.rowCount() == 0

        for key in ("公司", "项目名称", "字段=自定义"):
            panel._add_unknown_placeholder_field(key)
        assert panel._official_floating_field_keys == ["公司", "项目名称", "字段=自定义"]
        assert panel._official_field_name_labels["字段=自定义"].text() == (
            "{{@text:字段=自定义}}"
        )

        runtime = panel.material_context()
        runtime.entity_data = {
            "公司": "阿拉维特科技",
            "项目名称": "XX建设项目",
            "字段=自定义": "保留等号字段",
        }
        bridge.set_current_material_context(runtime)
        app.processEvents()

        context = panel.material_context()
        assert context.exact_material_placeholders is True
        assert context.field_aliases == {}
        assert context.entity_data["公司"] == "阿拉维特科技"
        assert context.entity_data["项目名称"] == "XX建设项目"
        assert context.entity_data["字段=自定义"] == "保留等号字段"
        overview_values = [
            panel._overview_preview_table.item(row, 1).text()
            for row in range(panel._overview_preview_table.rowCount())
        ]
        assert overview_values == ["阿拉维特科技", "XX建设项目", "保留等号字段"]
    finally:
        panel.close()
        app.processEvents()


def test_official_assets_support_fixed_and_floating_custom_series():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        fixed_indexes = [
            panel._official_fixed_layout.itemAt(index).widget().findChild(
                QLabel, "official_material_field_index"
            )
            for index in range(panel._official_fixed_layout.count())
        ]
        floating_indexes = [
            panel._official_floating_layout.itemAt(index).widget().findChild(
                QLabel, "official_material_field_index"
            )
            for index in range(panel._official_floating_layout.count())
        ]
        assert [label.text() for label in fixed_indexes] == ["1", "2"]
        assert [label.text() for label in floating_indexes] == [
            str(index) for index in range(1, 12)
        ]

        panel._add_official_series_field("organization", "fixed")
        assert panel._official_fixed_field_keys[:3] == [
            "organization",
            "发文机关2",
            "copy_scope",
        ]
        assert (
            panel._official_field_name_labels["发文机关2"].text()
            == "{{@text:发文机关2}}"
        )
        panel._remove_official_field("organization")
        assert "organization" not in panel._selected_profile().field_scopes
        panel._undo_official_fixed_remove_btn.click()
        assert panel._selected_profile().field_scopes["organization"] == "fixed"

        panel._request_official_field("fixed")
        assert "固定字段1" in panel._official_fixed_field_inputs
        panel._add_official_series_field("固定字段1", "fixed")
        assert "固定字段2" in panel._official_fixed_field_inputs

        panel._request_official_field("floating")
        assert "自由字段1" in panel._official_floating_field_keys
        panel._add_official_series_field("自由字段1", "floating")
        assert "自由字段2" in panel._official_floating_field_keys

        fixed_indexes = [
            panel._official_fixed_layout.itemAt(index).widget().findChild(
                QLabel, "official_material_field_index"
            )
            for index in range(panel._official_fixed_layout.count())
        ]
        floating_indexes = [
            panel._official_floating_layout.itemAt(index).widget().findChild(
                QLabel, "official_material_field_index"
            )
            for index in range(panel._official_floating_layout.count())
        ]
        assert [label.text() for label in fixed_indexes] == [
            str(index) for index in range(1, 6)
        ]
        assert [label.text() for label in floating_indexes] == [
            str(index) for index in range(1, 14)
        ]

        panel._remove_official_field("固定字段2")
        assert "固定字段2" not in panel._selected_profile().field_scopes
        panel._remove_official_field("title")
        assert "title" not in panel._selected_profile().field_scopes
        assert "title" not in panel._official_floating_field_keys
        assert panel._undo_official_floating_remove_btn.isHidden() is False
        panel._undo_official_floating_remove_btn.click()
        assert panel._selected_profile().field_scopes["title"] == "floating"
        assert "title" in panel._official_floating_field_keys
        archive = panel.current_archive()
        assert archive.profiles[0].field_scopes["固定字段1"] == "fixed"
        assert archive.profiles[0].field_scopes["自由字段1"] == "floating"
    finally:
        panel.close()
        app.processEvents()


def test_official_field_column_guide_aligns_with_interactive_rows():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        panel.resize(1280, 900)
        panel.show()
        panel._section_nav.select_card("fields")
        QTest.qWait(30)
        app.processEvents()

        guide = panel._official_fixed_column_guide
        metrics = guide.metrics()
        key = panel._official_fixed_field_keys[0]
        row = panel._official_field_rows[key]
        token = panel._official_field_name_labels[key]
        index = panel._official_field_index_labels[key]
        value = panel._official_fixed_field_inputs[key]
        actions = row.findChild(CompactRowActions)

        assert guide.maximumHeight() == 32
        assert token.width() == metrics.token_width
        assert index.width() == metrics.index_width
        assert actions is not None
        assert actions.width() == metrics.actions_width
        assert token.mapTo(panel, QPoint(0, 0)).x() == (
            guide.token_label.mapTo(panel, QPoint(0, 0)).x()
        )
        assert value.mapTo(panel, QPoint(0, 0)).x() == (
            guide.value_label.mapTo(panel, QPoint(0, 0)).x()
        )
        assert actions.mapTo(panel, QPoint(0, 0)).x() == (
            guide.actions_label.mapTo(panel, QPoint(0, 0)).x()
        )
        fixed_rows = [
            panel._official_field_rows[field_key]
            for field_key in panel._official_fixed_field_keys
        ]
        assert {row.height() for row in fixed_rows} == {get_theme().token_row_min_height}
        assert all(get_theme().divider in row.styleSheet() for row in fixed_rows)
        fixed_last = fixed_rows[-1]
        floating_header = panel._official_floating_title.parentWidget()
        gap = floating_header.mapTo(panel, QPoint(0, 0)).y() - (
            fixed_last.mapTo(panel, QPoint(0, fixed_last.height())).y()
        )
        assert gap == panel._official_fields_panel.layout().spacing()
    finally:
        panel.close()
        app.processEvents()


def test_official_package_registry_does_not_import_tokens_from_active_plan():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        # Exercise a profile whose only registry source is ``field_scopes``;
        # do not inherit declarations from the auto-activated library package.
        _load_material_package(panel)
        profile = panel._selected_profile()
        profile.field_scopes = {"项目通知标题": "floating"}
        profile.fields = {}

        fixed, floating = panel._official_field_scope_projection()

        assert fixed == []
        assert floating == ["项目通知标题"]
        assert profile.field_scopes == {"项目通知标题": "floating"}
        assert "title" not in profile.field_scopes
        assert "body" not in profile.field_scopes
    finally:
        panel.close()
        app.processEvents()


def test_official_fixed_field_click_copies_and_double_click_edits_name_and_value(
    monkeypatch,
):
    app = _app()
    copy_toasts = []
    monkeypatch.setattr(
        "src.ui.panels.assets.field_status_presenter.Toast.show_success",
        lambda message: copy_toasts.append(message),
    )
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        name_edit = panel._official_field_name_labels["organization"]
        value_edit = panel._official_fixed_field_inputs["organization"]
        value_line_edit = value_edit.findChild(QLineEdit)
        copy_feedback_timeout_ms = max(1500, app.doubleClickInterval() * 3)

        name_copy_spy = QSignalSpy(name_edit.copied)
        QTest.mouseClick(name_edit, Qt.LeftButton)
        assert app.clipboard().text() == "{{@text:发文机关1}}"
        assert name_copy_spy.count() or name_copy_spy.wait(copy_feedback_timeout_ms)
        assert copy_toasts == ["复制成功"]

        QTest.mouseDClick(name_edit, Qt.LeftButton)
        assert name_edit.isReadOnly() is False
        name_edit.innerEditor().setText("签发单位")
        QTest.keyClick(name_edit.innerEditor(), Qt.Key_Return)
        app.processEvents()

        assert name_edit.text() == "{{@text:签发单位}}"
        assert panel._selected_profile().field_aliases["签发单位"] == "organization"
        assert panel.material_context().field_aliases["签发单位"] == "organization"

        value_edit.setText("综合办公室")
        value_copy_spy = QSignalSpy(value_edit.copied)
        QTest.mouseClick(value_line_edit, Qt.LeftButton)
        assert app.clipboard().text() == "综合办公室"
        assert value_copy_spy.count() or value_copy_spy.wait(copy_feedback_timeout_ms)
        assert copy_toasts == ["复制成功", "复制成功"]

        value_double_click_spy = QSignalSpy(value_edit.copied)
        QTest.mouseDClick(value_line_edit, Qt.LeftButton)
        assert value_edit.isEditing() is True
        assert value_line_edit.isReadOnly() is False
        value_line_edit.setText("市委办公室")
        QTest.keyClick(value_line_edit, Qt.Key_Return)
        app.processEvents()
        assert not value_double_click_spy.wait(copy_feedback_timeout_ms)

        assert value_edit.text() == "市委办公室"
        assert value_edit.isEditing() is False
        assert value_line_edit.isReadOnly() is True
        assert copy_toasts == ["复制成功", "复制成功"]

        row = name_edit.parentWidget()
        index = row.findChild(QLabel, "official_material_field_index")
        assert "background: transparent" in index.styleSheet()
        assert "color:" in index.styleSheet()
    finally:
        panel.close()
        app.processEvents()


def test_official_series_add_uses_renamed_display_number_without_duplicates():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        panel._commit_official_field_name("organization", "发文机关2")
        app.processEvents()

        first_value_edit = panel._official_fixed_field_inputs["organization"]
        first_line_edit = first_value_edit.findChild(QLineEdit)
        assert first_line_edit.placeholderText() == "填写发文机关2"

        panel._add_official_series_field("organization", "fixed")
        app.processEvents()

        assert panel._official_fixed_field_keys[:3] == [
            "organization",
            "发文机关3",
            "copy_scope",
        ]
        assert panel._official_field_name_labels["organization"].text() == "{{@text:发文机关2}}"
        assert panel._official_field_name_labels["发文机关3"].text() == "{{@text:发文机关3}}"
        assert panel._official_fixed_field_inputs[
            "发文机关3"
        ].findChild(QLineEdit).placeholderText() == "填写发文机关3"

        visible_names = [
            panel._official_field_name_labels[key].text()
            for key in panel._official_fixed_field_keys
        ]
        assert len(visible_names) == len(set(visible_names))
    finally:
        panel.close()
        app.processEvents()


def test_official_series_repairs_existing_duplicate_display_alias_without_losing_values():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        profile = panel._selected_profile()
        profile.field_aliases["发文机关2"] = "organization"
        items = list(profile.field_scopes.items())
        organization_index = next(
            index for index, (key, _scope) in enumerate(items) if key == "organization"
        )
        items.insert(organization_index + 1, ("发文机关2", "fixed"))
        profile.field_scopes = dict(items)
        profile.fields["organization"] = "第一机关"
        profile.fields["发文机关2"] = "第二机关"
        panel._template_field_values.update(profile.fields)
        panel._rendered_official_fixed_field_keys = ()

        panel._refresh_summary()
        app.processEvents()

        assert panel._official_field_name_labels["organization"].text() == "{{@text:发文机关2}}"
        assert panel._official_field_name_labels["发文机关2"].text() == "{{@text:发文机关3}}"
        assert profile.field_aliases["发文机关2"] == "organization"
        assert profile.field_aliases["发文机关3"] == "发文机关2"
        assert bridge.current_material_context().field_aliases["发文机关3"] == "发文机关2"
        assert panel._official_fixed_field_inputs["organization"].text() == "第一机关"
        assert panel._official_fixed_field_inputs["发文机关2"].text() == "第二机关"
    finally:
        panel.close()
        app.processEvents()


def test_renamed_official_series_round_trips_through_exact_docx_runtime(tmp_path):
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        panel._commit_official_field_name("organization", "发文机关2")
        panel._official_fixed_field_inputs["organization"].setText("第一机关")
        panel._add_official_series_field("organization", "fixed")
        panel._official_fixed_field_inputs["发文机关3"].setText("第二机关")
        app.processEvents()
        context = panel.material_context()

        assert context.entity_data["organization"] == "第一机关"
        assert context.entity_data["发文机关3"] == "第二机关"
        assert context.field_aliases["发文机关2"] == "organization"

        archive_path = tmp_path / "renamed-series-material.json"
        save_entity_archive(panel.current_archive(), archive_path)
        reloaded = load_entity_archive(archive_path)
        assert reloaded.profiles[0].field_aliases["发文机关2"] == "organization"
        assert reloaded.profiles[0].fields["发文机关3"] == "第二机关"

        source = tmp_path / "renamed-series.docx"
        document = Document()
        document.add_paragraph(
            "别名一：{{@text:发文机关2}}；新增项：{{@text:发文机关3}}；底层项：{{@text:organization}}"
        )
        document.save(source)

        scene = SceneWorkspace()
        scene.module_switches = {name: False for name in scene.module_switches}
        scene.module_switches["entity_fill"] = True
        payload = WorkbenchProductionRunner(
            doc_path=str(source),
            template=TemplateConfig(),
            scene=scene,
            material_context=context,
        ).run(lambda *_args: None, lambda: False)

        output = Document(payload["output_path"])
        text = "\n".join(paragraph.text for paragraph in output.paragraphs)
        assert text == "别名一：第一机关；新增项：第二机关；底层项：第一机关"
    finally:
        panel.close()
        app.processEvents()


def test_removing_and_restoring_renamed_official_field_cleans_and_restores_alias():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        panel._commit_official_field_name("organization", "发文机关2")
        assert panel._selected_profile().field_aliases["发文机关2"] == "organization"

        panel._remove_official_field("organization")
        assert "发文机关2" not in panel._selected_profile().field_aliases

        panel._undo_last_official_field_removal()
        assert panel._selected_profile().field_aliases["发文机关2"] == "organization"
        assert panel._official_field_name_labels["organization"].text() == "{{@text:发文机关2}}"
    finally:
        panel.close()
        app.processEvents()


def test_official_assets_function_action_binds_and_resolves_field(monkeypatch):
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        monkeypatch.setattr(
            "src.ui.panels.assets.field_status_presenter.choose_official_field_function",
            lambda **_kwargs: (
                "apply",
                {"function": "realtime_date"},
            ),
        )

        panel._open_official_field_function("organization")

        profile = panel._selected_profile()
        assert profile.field_functions["organization"]["function"] == "realtime_date"
        assert profile.fields["organization"]
        assert panel._official_fixed_field_inputs["organization"].isReadOnly()
        function_button = panel._official_function_buttons["organization"]
        assert function_button.toolTip() == "字段函数"
        actions = function_button.parentWidget()
        assert actions.button("add").property(
            "compactActionVariant"
        ) == "outlined-primary"
        assert actions.button("remove").property(
            "compactActionVariant"
        ) == "outlined-danger"
        assert actions.button("add").property("iconButtonVariant") == "secondary"
        assert actions.button("remove").property("iconButtonVariant") == "secondary"
        assert panel.material_context().field_functions == profile.field_functions
    finally:
        panel.close()
        app.processEvents()


def test_official_series_add_and_delete_preserve_rows_scroll_and_viewport_anchor():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        panel.resize(1280, 700)
        panel.show()
        panel._select_section("fields")
        QTest.qWait(30)
        app.processEvents()
        app.processEvents()

        source_key = "recipient"
        source_label = panel._official_field_name_labels[source_key]
        source_row = panel._official_field_rows[source_key]
        scroll = panel._detail_scroll
        scroll_bar = scroll.verticalScrollBar()
        scroll_bar.setValue(min(180, scroll_bar.maximum()))
        app.processEvents()
        assert scroll_bar.value() > 0

        original_rows = dict(panel._official_field_rows)
        before_y = source_label.mapTo(scroll.viewport(), QPoint(0, 0)).y()
        observed_scroll_values: list[int] = []
        scroll_bar.valueChanged.connect(observed_scroll_values.append)

        panel._add_official_series_field(source_key, "floating")
        for _ in range(3):
            app.processEvents()

        added_keys = set(panel._official_field_rows) - set(original_rows)
        assert added_keys == {"主送机关2"}
        assert panel._official_field_rows[source_key] is source_row
        assert panel._official_field_name_labels[source_key] is source_label
        assert all(
            panel._official_field_rows[key] is row
            for key, row in original_rows.items()
        )
        after_y = source_label.mapTo(scroll.viewport(), QPoint(0, 0)).y()
        assert abs(after_y - before_y) <= 1
        assert scroll_bar.value() > 0
        assert 0 not in observed_scroll_values
        assert panel._official_floating_field_previews["主送机关2"].hasFocus() is False

        added_key = added_keys.pop()
        section_keys = list(panel._official_floating_field_keys)
        added_index = section_keys.index(added_key)
        delete_anchor_key = (
            section_keys[added_index + 1]
            if added_index + 1 < len(section_keys)
            else section_keys[added_index - 1]
        )
        delete_anchor_label = panel._official_field_name_labels[delete_anchor_key]
        delete_anchor_row = panel._official_field_rows[delete_anchor_key]
        delete_before_y = delete_anchor_label.mapTo(
            scroll.viewport(), QPoint(0, 0)
        ).y()
        rows_before_delete = dict(panel._official_field_rows)
        observed_scroll_values.clear()

        panel._remove_official_field(added_key)
        for _ in range(3):
            app.processEvents()

        assert added_key not in panel._official_field_rows
        assert panel._official_field_rows[delete_anchor_key] is delete_anchor_row
        assert panel._official_field_name_labels[delete_anchor_key] is delete_anchor_label
        assert all(
            panel._official_field_rows[key] is row
            for key, row in rows_before_delete.items()
            if key != added_key
        )
        delete_after_y = delete_anchor_label.mapTo(
            scroll.viewport(), QPoint(0, 0)
        ).y()
        assert abs(delete_after_y - delete_before_y) <= 1
        assert scroll_bar.value() > 0
        assert 0 not in observed_scroll_values
    finally:
        panel.close()
        app.processEvents()


def test_official_keyed_controllers_cleanly_transfer_a_field_between_scopes():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        _load_official_test_package(panel)
        profile = panel._selected_profile()
        original_row = panel._official_field_rows["organization"]

        profile.field_scopes["organization"] = "floating"
        panel._refresh_official_field_editor()
        floating_row = panel._official_field_rows["organization"]
        assert floating_row is not original_row
        assert panel._official_field_row_scopes["organization"] == "floating"
        assert "organization" not in panel._official_fixed_field_inputs
        assert panel._official_floating_field_previews["organization"].parentWidget() is floating_row

        profile.field_scopes["organization"] = "fixed"
        panel._refresh_official_field_editor()
        fixed_row = panel._official_field_rows["organization"]
        assert fixed_row is not floating_row
        assert panel._official_field_row_scopes["organization"] == "fixed"
        assert "organization" not in panel._official_floating_field_previews
        assert panel._official_fixed_field_inputs["organization"].parentWidget() is fixed_row
        assert panel._template_field_inputs["organization"] is panel._official_fixed_field_inputs["organization"]
    finally:
        panel.close()
        app.processEvents()


def test_profile_add_copy_delete_mutate_list_without_reloading_or_replacing_survivors(
    monkeypatch,
):
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        assert panel._profile_list.count() == 1
        original_item = panel._profile_list.item(0)
        reload_calls: list[tuple[tuple[object, ...], dict[str, object]]] = []
        monkeypatch.setattr(
            panel,
            "_reload_profile_list",
            lambda *args, **kwargs: reload_calls.append((args, kwargs)),
        )

        panel._add_profile()
        added_item = panel._profile_list.item(1)
        assert panel._profile_list.count() == 2
        assert panel._profile_list.item(0) is original_item

        panel._copy_current_profile()
        copied_item = panel._profile_list.item(2)
        assert panel._profile_list.count() == 3
        assert panel._profile_list.item(0) is original_item
        assert panel._profile_list.item(1) is added_item

        panel._remove_current_profile()
        assert panel._profile_list.count() == 2
        assert panel._profile_list.item(0) is original_item
        assert panel._profile_list.item(1) is added_item
        assert copied_item is not panel._profile_list.item(1)
        assert reload_calls == []
    finally:
        panel.close()
        app.processEvents()


def test_batch_profile_append_preserves_existing_items_and_check_states(
    tmp_path,
    monkeypatch,
):
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._profiles = [
            EntityProfile(profile_id="existing_1", profile_name="既有一"),
            EntityProfile(profile_id="existing_2", profile_name="既有二"),
        ]
        panel._reload_profile_list(select_index=0)
        first_item = panel._profile_list.item(0)
        second_item = panel._profile_list.item(1)
        first_item.setCheckState(Qt.Unchecked)
        app.processEvents()

        source = tmp_path / "append-profiles.json"
        source.write_text(
            '{"profiles":[{"profile_id":"imported","profile_name":"导入资料",'
            '"fields":{"company_name":"新公司"}}]}',
            encoding="utf-8",
        )
        reload_calls: list[int] = []
        monkeypatch.setattr(
            panel,
            "_reload_profile_list",
            lambda **_kwargs: reload_calls.append(1),
        )

        imported = panel.load_batch_profiles_from_path(source)

        assert [profile.profile_id for profile in imported] == ["imported"]
        assert panel._profile_list.count() == 3
        assert panel._profile_list.item(0) is first_item
        assert panel._profile_list.item(1) is second_item
        assert first_item.checkState() == Qt.Unchecked
        assert second_item.checkState() == Qt.Checked
        assert panel._profile_list.item(2).checkState() == Qt.Checked
        assert panel._current_profile_index == 2
        assert panel.selected_batch_profile_ids() == ["existing_2", "imported"]
        assert reload_calls == []
    finally:
        panel.close()
        app.processEvents()


def test_material_archive_round_trips_fixed_and_floating_field_scopes(tmp_path):
    path = tmp_path / "scoped-material.json"
    save_entity_archive(
        EntityArchive(
            profiles=[
                EntityProfile(
                    fields={"organization": "综合办公室"},
                    field_scopes={
                        "organization": "fixed",
                        "title": "floating",
                    },
                )
            ]
        ),
        path,
    )

    loaded = load_entity_archive(path)

    assert loaded.profiles[0].field_scopes == {
        "organization": "fixed",
        "title": "floating",
    }


def test_text_field_typing_coalesces_expensive_assets_summary_refresh(monkeypatch):
    app = _app()
    panel = AssetsPanel(PanelBridge())
    calls = {"full": 0, "fields": 0}
    original_refresh = panel._refresh_summary
    original_fields_refresh = panel._refresh_field_only_projection

    def counted_refresh() -> None:
        calls["full"] += 1
        original_refresh()

    def counted_fields_refresh() -> None:
        calls["fields"] += 1
        original_fields_refresh()

    monkeypatch.setattr(panel, "_refresh_summary", counted_refresh)
    monkeypatch.setattr(
        panel,
        "_refresh_field_only_projection",
        counted_fields_refresh,
    )
    try:
        panel._manual_field_keys.append("自定义字段")
        panel._refresh_unknown_field_suggestions([])
        edit = panel._template_field_inputs["自定义字段"]
        for length in range(1, 21):
            edit.setText("x" * length)

        assert calls == {"full": 0, "fields": 0}
        assert edit.text() == "x" * 20

        QTest.qWait(320)
        app.processEvents()

        assert calls == {"full": 0, "fields": 1}
    finally:
        panel.close()
        app.processEvents()


def test_summary_refresh_scheduler_keeps_full_projection_priority(monkeypatch):
    app = _app()
    panel = AssetsPanel(PanelBridge())
    calls: list[str] = []
    monkeypatch.setattr(panel, "_refresh_summary", lambda: calls.append("full"))
    monkeypatch.setattr(
        panel,
        "_refresh_field_only_projection",
        lambda: calls.append("fields"),
    )
    try:
        panel._schedule_summary_refresh(10_000, scope="fields")
        panel._schedule_summary_refresh(10_000, scope="full")
        panel._schedule_summary_refresh(10_000, scope="fields")
        panel._summary_refresh_timer.stop()
        panel._run_scheduled_summary_refresh()

        assert calls == ["full"]
        assert panel._pending_summary_refresh_scope == ""
    finally:
        panel.close()
        app.processEvents()


def test_material_context_update_does_not_rebuild_profile_or_batch_state(monkeypatch):
    app = _app()
    bridge = PanelBridge()
    panel = AssetsPanel(bridge)
    _load_material_package(
        panel,
        field_scopes={"organization": "fixed", "title": "floating"},
        fields={"organization": "包内固定值"},
        profile_id="official:notice",
    )
    calls = {"refresh": 0, "light": 0, "reload": 0, "set_values": 0, "batch": 0}
    original_refresh = panel._refresh_summary
    original_light = panel._refresh_field_only_projection
    original_set_values = panel._set_editor_values

    def counted_refresh():
        calls["refresh"] += 1
        return original_refresh()

    def counted_set_values(*args, **kwargs):
        calls["set_values"] += 1
        return original_set_values(*args, **kwargs)

    def counted_light():
        calls["light"] += 1
        return original_light()

    monkeypatch.setattr(panel, "_refresh_summary", counted_refresh)
    monkeypatch.setattr(
        panel,
        "_reload_profile_list",
        lambda *args, **kwargs: calls.__setitem__("reload", calls["reload"] + 1),
    )
    monkeypatch.setattr(panel, "_set_editor_values", counted_set_values)
    monkeypatch.setattr(panel, "_refresh_field_only_projection", counted_light)
    monkeypatch.setattr(
        panel,
        "_sync_material_batch_selection",
        lambda *args, **kwargs: calls.__setitem__("batch", calls["batch"] + 1),
    )
    try:
        changed = bridge.set_current_material_context(
            MaterialExecutionContext(
                package_id="test-package",
                profile_id="official:notice",
                profile_name="公文资料",
                entity_data={"organization": "不应覆盖", "title": "测试"},
            )
        )

        assert changed is True
        assert calls == {
            "refresh": 0,
            "light": 1,
            "reload": 0,
            "set_values": 0,
            "batch": 0,
        }
        assert panel._selected_profile().fields == {"organization": "包内固定值"}
        assert panel._official_floating_field_previews["title"].text() == "测试"

        updated_context = bridge.current_material_context()
        updated_context.entity_data["title"] = "测试二"
        changed_field = bridge.set_current_material_context(updated_context)
        assert changed_field is True
        assert calls == {
            "refresh": 0,
            "light": 2,
            "reload": 0,
            "set_values": 0,
            "batch": 0,
        }
        assert panel._selected_profile().fields == {"organization": "包内固定值"}

        changed_again = bridge.set_current_material_context(
            bridge.current_material_context()
        )
        assert changed_again is False
        assert calls == {
            "refresh": 0,
            "light": 2,
            "reload": 0,
            "set_values": 0,
            "batch": 0,
        }
    finally:
        panel.close()
        app.processEvents()


def test_mismatched_runtime_context_cannot_mutate_selected_package():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        _load_material_package(
            panel,
            field_scopes={"organization": "fixed", "title": "floating"},
            fields={"organization": "包内固定值"},
            profile_id="owned-profile",
        )
        profile = panel._selected_profile()
        profile.assets_dir = "package-assets"
        panel._assets_picker.set_path("package-assets")

        panel._on_material_context_changed(
            MaterialExecutionContext(
                package_id="another-package",
                profile_id="another-profile",
                entity_data={"organization": "外部覆盖", "title": "外部任务值"},
                entity_assets_dir="external-assets",
            )
        )

        assert profile.fields == {"organization": "包内固定值"}
        assert profile.assets_dir == "package-assets"
        assert panel._task_field_values == {}
    finally:
        panel.close()
        app.processEvents()


def test_runtime_context_without_identity_cannot_mutate_selected_package():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        _load_material_package(
            panel,
            field_scopes={"organization": "fixed", "title": "floating"},
            fields={"organization": "包内固定值"},
            profile_id="owned-profile",
        )
        profile = panel._selected_profile()
        profile.assets_dir = "package-assets"
        profile.timeline_plans = {"owned": {"enabled": False}}
        panel._assets_picker.set_path("package-assets")

        panel._on_material_context_changed(
            MaterialExecutionContext(
                entity_data={"title": "无身份任务值"},
                entity_assets_dir="external-assets",
                timeline_plans={"foreign": {"enabled": False}},
            )
        )

        assert profile.fields == {"organization": "包内固定值"}
        assert profile.assets_dir == "package-assets"
        assert set(profile.timeline_plans) == {"owned"}
        assert panel._task_field_values == {}
    finally:
        panel.close()
        app.processEvents()


def test_text_fields_do_not_render_redundant_inline_status_explanations():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._manual_field_keys.extend(["信用代码", "电话"])
        panel._refresh_unknown_field_suggestions([])
        panel._template_field_inputs["信用代码"].setText("not-a-credit-code")
        panel._template_field_inputs["电话"].setText("not-a-phone")
        app.processEvents()

        assert panel.findChildren(QLabel, "asset_field_status") == []
        assert panel._fields_hint_label.isHidden()
    finally:
        panel.close()
        app.processEvents()


def test_blank_material_pack_has_no_scenario_specific_fields_or_default_requirement():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        _load_material_package(panel)
        assert panel._field_inputs == {}
        assert panel._template_field_inputs == {}
        assert not hasattr(panel, "_required_fields_edit")
        assert panel._empty_fields_label.isHidden()
        assert panel._official_fields_panel.isHidden() is False
        assert panel._official_fixed_count.text() == "0 项"
        assert panel._official_floating_count.text() == "0 项"
        assert panel._generation_detail_label.text() == (
            "固定字段 0 项 · 自由字段 0 项 · 当前已填写 0/0"
        )
        assert panel._batch_output_template() == "{profile_name}"
        assert all(
            "公司" not in panel._batch_output_naming_combo.itemText(index)
            for index in range(panel._batch_output_naming_combo.count())
        )
    finally:
        panel.close()
        app.processEvents()


def test_loaded_profile_fields_become_exact_dynamic_rows_without_fixed_schema():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._set_structured_fields(
            {"公司名称": "A 公司", "信用代码": "123", "custom_field": "value"}
        )

        assert list(panel._template_field_inputs) == ["公司名称", "信用代码", "custom_field"]
        assert panel._editor_fields() == {
            "公司名称": "A 公司",
            "信用代码": "123",
            "custom_field": "value",
        }
        assert panel._exact_material_field_key("{{@text:公司}}") == "公司"
        assert panel._exact_material_field_key("公司 名称") == ""
        assert panel._exact_material_field_key("{{@text: 公司 }}") == ""
    finally:
        panel.close()
        app.processEvents()


def test_add_material_field_uses_two_inline_inputs_and_wraps_code_automatically():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.show()
        app.processEvents()
        assert "Word 中写 {{@text:公司名称}}" in panel._field_mapping_example.text()
        assert panel._field_code_column_label.text() == "占位符"
        assert panel._field_value_column_label.text() == "字段内容"

        panel._add_material_field_btn.click()
        app.processEvents()

        draft_key = panel._manual_field_keys[-1]
        code_edit = panel._template_field_key_inputs[draft_key]
        value_edit = panel._template_field_inputs[draft_key]
        assert code_edit.objectName() == "material_field_code_input"
        assert value_edit.objectName() == "material_field_value_input"
        assert code_edit.text() == ""

        code_edit.setFocus()
        app.processEvents()
        code_edit.setText("公司名称")
        QTest.keyClick(code_edit, Qt.Key_Tab)
        QTest.qWait(20)
        app.processEvents()

        assert panel._manual_field_keys == ["公司名称"]
        assert panel._template_field_key_inputs["公司名称"].text() == "{{@text:公司名称}}"
        panel._template_field_inputs["公司名称"].setText("阿拉维特科技")
        app.processEvents()
        assert panel._editor_fields() == {"公司名称": "阿拉维特科技"}
    finally:
        panel.close()
        app.processEvents()


def test_renamed_numeric_field_rebinds_controller_delete_and_exact_runtime(tmp_path):
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._manual_field_keys.append("保留字段")
        panel._refresh_unknown_field_suggestions([])
        retained_row = panel._unknown_field_rows["保留字段"]

        panel._request_add_material_field()
        draft_key = panel._manual_field_keys[-1]
        draft_row = panel._unknown_field_rows[draft_key]
        panel._commit_material_field_key(draft_key, "12345678")
        app.processEvents()

        assert panel._manual_field_keys == ["保留字段", "12345678"]
        assert panel._unknown_fields_controller.keys() == ("保留字段", "12345678")
        assert draft_key not in panel._unknown_field_rows
        assert draft_key not in panel._template_field_inputs
        assert panel._unknown_field_rows["保留字段"] is retained_row
        assert panel._unknown_field_rows["12345678"] is not draft_row
        assert panel._unknown_field_rows["12345678"].property("materialFieldKey") == "12345678"

        panel._template_field_inputs["12345678"].setText("替换成功")
        app.processEvents()
        context = panel.material_context()
        assert context.entity_data["12345678"] == "替换成功"

        source = tmp_path / "numeric-field.docx"
        document = Document()
        document.add_paragraph("结果：{{@text:12345678}}")
        document.save(source)
        scene = SceneWorkspace()
        scene.module_switches = {name: False for name in scene.module_switches}
        scene.module_switches["entity_fill"] = True
        payload = WorkbenchProductionRunner(
            doc_path=str(source),
            template=TemplateConfig(),
            scene=scene,
            material_context=context,
        ).run(lambda *_args: None, lambda: False)
        output_text = "\n".join(
            paragraph.text for paragraph in Document(payload["output_path"]).paragraphs
        )
        assert output_text == "结果：替换成功"

        remove_button = panel._unknown_field_remove_buttons["12345678"]
        remove_button.click()
        assert "12345678" in panel._manual_field_keys
        app.processEvents()
        assert "12345678" not in panel._manual_field_keys
        assert "12345678" not in panel._unknown_fields_controller.keys()
        assert panel._unknown_field_rows["保留字段"] is retained_row
    finally:
        panel.close()
        app.processEvents()


def test_duplicate_inline_material_field_code_is_not_committed():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._manual_field_keys.append("公司名称")
        panel._refresh_unknown_field_suggestions([])
        panel._request_add_material_field()
        draft_key = panel._manual_field_keys[-1]
        code_edit = panel._template_field_key_inputs[draft_key]

        code_edit.setText("{{@text:公司名称}}")
        code_edit.editingFinished.emit()
        app.processEvents()

        assert panel._manual_field_keys.count("公司名称") == 1
        assert draft_key in panel._manual_field_keys
        assert panel._template_field_key_inputs[draft_key].text() == ""
    finally:
        panel.close()
        app.processEvents()


def test_material_field_remove_button_defers_row_rebuild_until_click_returns():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.show()
        panel._manual_field_keys.append("公司名称")
        panel._refresh_unknown_field_suggestions([])
        remove_button = panel.findChild(QPushButton, "material_field_remove_btn")
        assert remove_button is not None

        remove_button.click()
        assert "公司名称" in panel._manual_field_keys

        app.processEvents()
        assert "公司名称" not in panel._manual_field_keys
        assert "公司名称" not in panel._template_field_inputs
    finally:
        panel.close()
        app.processEvents()


def test_free_material_field_delete_preserves_nonzero_scroll_anchor_and_rows():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.resize(1280, 700)
        panel.show()
        panel._select_section("fields")
        panel._manual_field_keys.extend([f"field{index}" for index in range(1, 21)])
        panel._refresh_unknown_field_suggestions([])
        app.processEvents()
        app.processEvents()

        scroll = panel._detail_scroll
        bar = scroll.verticalScrollBar()
        bar.setValue(min(220, bar.maximum()))
        app.processEvents()
        assert bar.value() > 0
        anchor_key = "field6"
        anchor_row = panel._unknown_field_rows[anchor_key]
        before_y = anchor_row.mapTo(scroll.viewport(), QPoint(0, 0)).y()
        rows_before = dict(panel._unknown_field_rows)
        observed_values: list[int] = []
        bar.valueChanged.connect(observed_values.append)

        panel._remove_manual_field("field5")
        for _ in range(4):
            app.processEvents()

        after_y = anchor_row.mapTo(scroll.viewport(), QPoint(0, 0)).y()
        assert abs(after_y - before_y) <= 1
        assert panel._unknown_field_rows[anchor_key] is anchor_row
        assert all(
            panel._unknown_field_rows[key] is row
            for key, row in rows_before.items()
            if key != "field5"
        )
        assert "field5" not in panel._unknown_field_rows
        assert bar.value() > 0
        assert 0 not in observed_values
    finally:
        panel.close()
        app.processEvents()


def test_preview_table_action_button_defers_table_rebuild_until_click_returns():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.show()
        panel._refresh_preview_table(
            [
                {
                    "placeholder": "{{@text:用户自定义}}",
                    "value": "未填写",
                    "source": "",
                    "status": "未填写",
                    "action": "添加字段",
                    "action_type": "custom",
                    "action_key": "用户自定义",
                }
            ]
        )
        button = panel._preview_action_buttons["{{@text:用户自定义}}"]

        button.click()
        assert "用户自定义" not in panel._selected_profile().field_scopes

        app.processEvents()
        assert panel._selected_profile().field_scopes["用户自定义"] == "floating"
        assert "用户自定义" in panel._official_floating_field_previews
    finally:
        panel.close()
        app.processEvents()


def test_question_figure_replace_button_defers_table_rebuild(
    tmp_path,
    monkeypatch,
):
    app = _app()
    old_path = tmp_path / "old.png"
    new_path = tmp_path / "new.png"
    old_path.write_bytes(b"old")
    new_path.write_bytes(b"new")
    panel = AssetsPanel(PanelBridge())
    try:
        panel.show()
        panel._asset_item_payloads = [
            {
                "item_id": "q1",
                "label": "题目 1",
                "role": "question_figure",
                "path": str(old_path),
                "metadata": {"question_index": "1"},
            },
            {
                "item_id": "q2",
                "label": "题目 2",
                "role": "question_figure",
                "path": str(old_path),
                "metadata": {"question_index": "2"},
            },
        ]
        panel._refresh_summary()
        table = panel._question_figure_items_table
        button = table.cellWidget(0, 4)
        assert button is not None
        monkeypatch.setattr(
            "src.ui.panels.assets.question_figure_repair_actions_presenter.QFileDialog.getOpenFileName",
            lambda *args, **kwargs: (str(new_path), ""),
        )

        button.click()
        assert panel._asset_item_payloads[0]["path"] == str(old_path)

        app.processEvents()
        assert panel._asset_item_payloads[0]["path"] == str(new_path)
    finally:
        panel.close()
        app.processEvents()


def test_question_figure_history_rollback_button_defers_table_rebuild(tmp_path):
    app = _app()
    image_path = tmp_path / "q1.png"
    image_path.write_bytes(b"image")
    current = AssetItem(
        item_id="q1",
        label="题目 1",
        role="question_figure",
        path=str(image_path),
        metadata={
            "question_index": "1",
            "source": "remote-library",
            "asset_id": "new-q1",
            "alt_text": "new alt",
        },
    )
    history = question_figure_library_metadata_history_record(
        current,
        {
            "question_index": "1",
            "source": "local",
            "asset_id": "old-q1",
            "alt_text": "old alt",
        },
        dict(current.metadata),
    )
    assert history is not None
    panel = AssetsPanel(PanelBridge())
    try:
        panel.show()
        panel._asset_item_payloads = [asset_item_payload(current)]
        panel._selected_profile().asset_item_history = [history]
        panel._refresh_summary()
        table = panel._question_figure_library_version_history_table
        button = table.cellWidget(0, 4)
        assert button is not None

        button.click()
        assert panel._asset_item_payloads[0]["metadata"]["source"] == "remote-library"

        app.processEvents()
        assert panel._asset_item_payloads[0]["metadata"]["source"] == "local"
    finally:
        panel.close()
        app.processEvents()


def test_switching_dynamic_profiles_replaces_values_and_manual_delete_discards_value():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._set_structured_fields({"公司": "A 公司"})
        panel._set_structured_fields({"公司": "B 公司"})
        assert panel._template_field_inputs["公司"].text() == "B 公司"

        panel._remove_manual_field("公司")
        assert "公司" not in panel._template_field_inputs
        panel._manual_field_keys.append("公司")
        panel._refresh_unknown_field_suggestions([])
        assert panel._template_field_inputs["公司"].text() == ""
    finally:
        panel.close()
        app.processEvents()


def test_preview_repair_action_targets_dynamic_field_row_instead_of_hidden_legacy_editor():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._handle_preview_row_action("custom", "用户自定义")

        assert panel._selected_profile().field_scopes["用户自定义"] == "floating"
        assert "用户自定义" in panel._official_floating_field_previews
        assert panel._active_section_id == "fields"
        assert panel.focus_material_repair_target("field", "用户自定义") is True
    finally:
        panel.close()
        app.processEvents()


def test_conflicting_same_profile_fields_are_not_persisted():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._fields_edit.set_text("公司=A公司\n公司=B公司")
        app.processEvents()

        assert panel._field_conflict_keys() == ("公司",)
        assert panel._persist_current_profile_editor() is False
        assert panel._selected_profile().fields == {}
        with pytest.raises(ValueError, match="material_field_conflicts"):
            panel.material_context()
        assert not hasattr(panel, "_save_btn")
        assert not panel._fields_hint_label.isHidden()
        assert "无法写入" in panel._fields_hint_label.text()
    finally:
        panel.close()
        app.processEvents()


def test_profile_switch_aborts_and_preserves_editor_when_field_persist_fails():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.set_archive(
            EntityArchive(
                profiles=[
                    EntityProfile(profile_id="profile-a"),
                    EntityProfile(profile_id="profile-b"),
                ]
            )
        )
        panel._profile_list.setCurrentRow(0)
        conflicting_text = "公司=A 公司\n公司=B 公司"
        panel._fields_edit.set_text(conflicting_text)
        app.processEvents()

        panel._profile_list.setCurrentRow(1)
        app.processEvents()

        assert panel._current_profile_index == 0
        assert panel._profile_list.currentRow() == 0
        assert panel._fields_edit.get_text() == conflicting_text
        assert panel._profiles[0].fields == {}
    finally:
        panel.close()
        app.processEvents()


def test_mapping_import_preserves_field_inventory_order_through_persist():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    keys = [
        "zeta",
        "alpha",
        "middle",
        "beta",
        "omega",
        "gamma",
        "delta",
        "epsilon",
    ]
    try:
        panel._apply_mapping_payload(
            MaterialMappingPayload(
                entity_data={key: str(index) for index, key in enumerate(keys)}
            )
        )

        assert panel._persist_current_profile_editor() is True
        profile = panel._selected_profile()
        assert list(profile.field_scopes) == keys
        assert list(profile.fields) == keys
        assert profile.declared_field_keys == keys

        assert panel._persist_current_profile_editor() is True
        assert list(profile.field_scopes) == keys
        assert list(profile.fields) == keys
        assert profile.declared_field_keys == keys
    finally:
        panel.close()
        app.processEvents()


def test_persisted_manual_field_rename_migrates_package_owned_state():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        profile = panel._selected_profile()
        profile.field_scopes = {
            "第一项": "fixed",
            "旧字段": "fixed",
            "第三项": "floating",
        }
        profile.fields = {"第一项": "A", "旧字段": "B", "第三项": "C"}
        profile.declared_field_keys = ["第一项", "旧字段", "第三项"]
        profile.field_sources = {"旧字段": "imported_mapping"}
        profile.field_functions = {"旧字段": {"function": "realtime_date"}}
        profile.field_aliases = {"旧别名": "旧字段"}
        panel._manual_field_keys = ["第一项", "旧字段", "第三项"]
        panel._declared_field_keys = {"第一项", "旧字段", "第三项"}
        panel._imported_field_keys = {"旧字段"}
        panel._template_field_values = {
            "第一项": "A",
            "旧字段": "B",
            "第三项": "C",
        }

        panel._commit_material_field_key("旧字段", "新字段")

        assert list(profile.field_scopes) == ["第一项", "新字段", "第三项"]
        assert list(profile.fields) == ["第一项", "新字段", "第三项"]
        assert profile.declared_field_keys == ["第一项", "新字段", "第三项"]
        assert profile.field_functions == {
            "新字段": {"function": "realtime_date"}
        }
        assert profile.field_sources == {"新字段": "imported_mapping"}
        assert profile.field_aliases == {"旧别名": "新字段"}
        assert panel._manual_field_keys == ["第一项", "新字段", "第三项"]
        assert panel._persist_current_profile_editor() is True
        assert "旧字段" not in profile.field_scopes
        assert "旧字段" not in profile.declared_field_keys
    finally:
        panel.close()
        app.processEvents()


def test_batch_import_rejects_duplicate_exact_field_headers(tmp_path):
    source = tmp_path / "duplicate-fields.csv"
    source.write_text("公司,公司\nA公司,B公司\n", encoding="utf-8-sig")

    with pytest.raises(ValueError, match="重复字段.*公司"):
        _load_batch_profiles_from_path(source)


def test_mapping_import_rejects_replacement_rule_tables(tmp_path):
    source = tmp_path / "duplicate-replacements.csv"
    source.write_text(
        "old,new\n{{@text:公司}},A公司\n{{@text:公司}},B公司\n",
        encoding="utf-8-sig",
    )

    with pytest.raises(ValueError, match="不再支持 old/new 替换规则"):
        load_material_mapping(source)


def test_chinese_placeholder_name_runs_end_to_end(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("公司：{{@text:公司}}")
    document.save(source)

    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["entity_fill"] = True

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(entity_data={"公司": "阿拉维特科技"}),
    ).run(lambda *_args: None, lambda: False)

    output = Document(payload["output_path"])
    text = "\n".join(paragraph.text for paragraph in output.paragraphs)
    assert payload["status"] == "success"
    assert "公司：阿拉维特科技" in text
    assert "{{@text:公司}}" not in text
