from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from src.config.document_scope import DocumentScopePolicy
from src.config.document_structure_contract import (
    DetectedRegion,
    DocumentStructureEvidence,
    ParagraphAnchor,
    StructureReviewItem,
    pending_document_structure_review_roles,
)
from src.modules.structure.heading_recognition import analyze_document_tree
from src.pipeline.context import PipelineContext
from src.services.document_structure_evidence import (
    RegionDecision,
    build_document_structure_evidence,
    read_document_paragraph_anchors,
)
from src.shared.engine.document_scope_runtime import (
    bind_document_scope,
    document_scope_allows_paragraph,
    document_scope_allows_role,
    project_document_scope_tree,
)
from src.shared.engine.document_structure_model import DocSection, DocTree


def _save_document_with_appendix(path):
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容")
    doc.add_heading("附录A 补充材料", level=1)
    doc.add_paragraph("附录内容")
    doc.save(path)


def test_excluded_region_keeps_its_boundary_but_is_not_exposed_or_writable(tmp_path):
    source = tmp_path / "appendix.docx"
    _save_document_with_appendix(source)
    evidence = build_document_structure_evidence(source)
    appendix = next(region for region in evidence.regions if region.role_id == "appendix")
    doc = Document(source)
    policy = DocumentScopePolicy(mode="all")
    binding = bind_document_scope(
        doc,
        evidence,
        (RegionDecision("appendix", "exclude"),),
        policy,
        mode_id="thesis",
    )

    tree = project_document_scope_tree(
        doc,
        analyze_document_tree(doc),
        binding,
        policy,
        mode_id="thesis",
    )

    appendix_index = appendix.start_anchor.source_index
    assert tree.get_section("appendix") is None
    assert tree.is_role_writable("appendix") is False
    assert tree.is_paragraph_writable(appendix_index) is False
    assert tree.is_paragraph_writable(appendix_index + 1) is False
    assert tree.is_paragraph_writable(0) is True


def test_unclassified_gap_is_not_treated_as_writable_body():
    tree = DocTree(
        sections=[DocSection("body", 2, 4)],
        section_ranges={"body": (2, 4)},
        writable_roles=frozenset({"body"}),
    )

    assert tree.is_paragraph_writable(0) is False
    assert tree.is_paragraph_writable(2) is True


def test_body_scope_can_use_an_excluded_role_start_as_its_end_boundary(tmp_path):
    source = tmp_path / "manual-boundary.docx"
    original = Document()
    original.add_heading("第一章 正文", level=1)
    original.add_paragraph("正文内容")
    original.add_paragraph("未识别的后置区域")
    original.add_paragraph("不应写入")
    original.save(source)
    evidence = build_document_structure_evidence(source)
    anchors = read_document_paragraph_anchors(source)
    doc = Document(source)
    policy = DocumentScopePolicy(mode="body")

    binding = bind_document_scope(
        doc,
        evidence,
        (
            RegionDecision(
                "references",
                "set_start",
                start_anchor=anchors[2],
            ),
        ),
        policy,
        mode_id="custom",
    )
    tree = project_document_scope_tree(
        doc,
        analyze_document_tree(doc),
        binding,
        policy,
        mode_id="custom",
    )

    assert tree.get_section("body").end_index == 2
    assert tree.get_section("references") is not None
    assert tree.is_paragraph_writable(1) is True
    assert tree.is_paragraph_writable(2) is False
    assert tree.is_paragraph_writable(3) is False


def test_review_includes_uncertain_boundary_after_the_selected_region():
    def anchor(index: int) -> ParagraphAnchor:
        return ParagraphAnchor(index, str(index), "", "", 0, f"p{index}")

    evidence = DocumentStructureEvidence(
        source_path="source.docx",
        source_revision="sha256:" + ("1" * 64),
        detector_revision="test",
        evidence_digest="sha256:" + ("2" * 64),
        regions=(
            DetectedRegion("body", anchor(0), anchor(2), "accepted", ()),
            DetectedRegion("references", anchor(2), None, "review", ()),
        ),
        review_items=(
            StructureReviewItem("references", "region_requires_confirmation"),
        ),
    )

    assert pending_document_structure_review_roles(
        evidence,
        (),
        included_roles=("body",),
    ) == ("references",)
    assert pending_document_structure_review_roles(
        evidence,
        (RegionDecision("references", "accept"),),
        included_roles=("body",),
    ) == ()


def test_structure_analysis_does_not_mutate_outline_metadata():
    doc = Document()
    paragraph = doc.add_paragraph("[1] Author. Title. 2024.")
    paragraph_properties = paragraph._element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    paragraph_properties.append(outline)
    before = etree.tostring(doc.element, encoding="utf-8")

    analyze_document_tree(doc)

    assert etree.tostring(doc.element, encoding="utf-8") == before


def test_scope_helpers_are_unrestricted_only_outside_the_execution_gate():
    context = PipelineContext()
    assert context.doc_tree is None
    assert document_scope_allows_role(context, "references") is True
    assert document_scope_allows_paragraph(context, 1) is True

    context.document_scope_gate_active = True
    assert document_scope_allows_role(context, "references") is False
    assert document_scope_allows_paragraph(context, 1) is False

    context.doc_tree = DocTree(
        sections=[DocSection("references", 1, 3)],
        section_ranges={"references": (1, 3)},
        writable_roles=frozenset({"references"}),
    )
    assert document_scope_allows_role(context, "references") is True
    assert document_scope_allows_paragraph(context, 1) is True
