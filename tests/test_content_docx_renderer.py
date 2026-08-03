from base64 import b64decode
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from src.config.content_materials import (
    ContentInsertionRule,
    DocumentFragment,
    HeadingBlock,
    ImageBlock,
    InlineContent,
    InlineKind,
    InlineVerticalAlignment,
    ListBlock,
    ListItem,
    ListNumberFormat,
    PageBreakBlock,
    ParagraphBlock,
    TableBlock,
    TableCell,
    TableRow,
)
from src.services.material_content.docx_renderer import (
    ContentDocxRenderer,
    ContentRenderBlockedError,
    render_document_fragment,
)


TOKEN = "{{@file:route}}"


def _rule(**overrides) -> ContentInsertionRule:
    values = {
        "rule_id": "route-rule",
        "content_id": "route",
        "anchor_token": TOKEN,
    }
    values.update(overrides)
    return ContentInsertionRule(**values)


def _text(value: str, **formatting) -> InlineContent:
    return InlineContent(text=value, **formatting)


def _round_trip(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return Document(stream)


def test_semantic_blocks_preserve_order_styles_inline_formatting_and_field_tokens():
    document = Document()
    document.add_paragraph(TOKEN)
    fragment = DocumentFragment(
        (
            HeadingBlock(2, (_text("技术路线", bold=True),)),
            ParagraphBlock(
                (
                    _text("投标人：", italic=True),
                    InlineContent(
                        kind=InlineKind.FIELD_TOKEN,
                        field_key="@text:company_name",
                        underline=True,
                    ),
                    InlineContent(kind=InlineKind.SOFT_BREAK),
                    _text("下一行"),
                    InlineContent(kind=InlineKind.HARD_BREAK),
                    _text("结束"),
                )
            ),
        )
    )

    receipt = render_document_fragment(
        document,
        fragment,
        _rule(heading_level_offset=-1),
    )

    assert receipt.rendered_block_count == 2
    assert [paragraph.style.style_id for paragraph in document.paragraphs] == [
        "Heading1",
        "Normal",
    ]
    assert document.paragraphs[0].text == "技术路线"
    assert "{{@text:company_name}}" in document.paragraphs[1].text
    assert TOKEN not in document.element.xml
    assert document.paragraphs[0]._p.find(".//" + qn("w:b")) is not None
    assert document.paragraphs[1]._p.find(".//" + qn("w:i")) is not None
    assert document.paragraphs[1]._p.find(".//" + qn("w:u")) is not None
    assert len(document.paragraphs[1]._p.findall(".//" + qn("w:br"))) == 1

    reopened = _round_trip(document)
    assert [paragraph.style.style_id for paragraph in reopened.paragraphs] == [
        "Heading1",
        "Normal",
    ]
    assert "{{@text:company_name}}" in reopened.paragraphs[1].text


def test_plain_text_mode_flattens_structure_and_drops_formatting_and_images():
    document = Document()
    document.add_paragraph(TOKEN)
    fragment = DocumentFragment(
        (
            HeadingBlock(2, (_text("标题", bold=True),)),
            ListBlock(
                ordered=True,
                start=2,
                nesting=0,
                items=(ListItem((_text("事项", italic=True),)),),
            ),
            TableBlock(
                (
                    TableRow(
                        (
                            TableCell((ParagraphBlock((_text("甲"),)),)),
                            TableCell((ParagraphBlock((_text("乙"),)),)),
                        )
                    ),
                )
            ),
            ImageBlock("sha256/image.png", alt_text="示意图"),
        )
    )

    receipt = render_document_fragment(
        document,
        fragment,
        _rule(format_mode="plain_text"),
    )

    assert [paragraph.text for paragraph in document.paragraphs] == [
        "标题",
        "2. 事项",
        "甲\t乙",
    ]
    assert all(paragraph.style.style_id == "Normal" for paragraph in document.paragraphs)
    assert not document.tables
    assert not receipt.image_job_drafts
    assert document.element.body.find(".//" + qn("w:b")) is None
    assert document.element.body.find(".//" + qn("w:i")) is None


def test_tab_strikethrough_and_vertical_alignment_survive_docx_round_trip():
    document = Document()
    document.add_paragraph(TOKEN)
    fragment = DocumentFragment(
        (
            ParagraphBlock(
                (
                    _text("before"),
                    InlineContent(kind=InlineKind.TAB),
                    _text("strike", strikethrough=True),
                    _text(
                        "super",
                        vertical_alignment=InlineVerticalAlignment.SUPERSCRIPT,
                    ),
                    _text(
                        "sub",
                        vertical_alignment=InlineVerticalAlignment.SUBSCRIPT,
                    ),
                )
            ),
        )
    )

    render_document_fragment(document, fragment, _rule())
    reopened = _round_trip(document)
    paragraph = reopened.paragraphs[0]

    assert paragraph._p.find(".//" + qn("w:tab")) is not None
    assert paragraph._p.find(".//" + qn("w:strike")) is not None
    alignments = [
        node.get(qn("w:val"))
        for node in paragraph._p.findall(".//" + qn("w:vertAlign"))
    ]
    assert alignments == ["superscript", "subscript"]


def test_nine_level_chinese_numbering_start_and_marker_are_rendered():
    document = Document()
    document.add_paragraph(TOKEN)
    fragment = DocumentFragment(
        (
            ListBlock(
                True,
                5,
                8,
                (ListItem((_text("第九层"),)),),
                number_format=ListNumberFormat.CHINESE_COUNTING,
                marker_template="%9、",
            ),
        )
    )

    render_document_fragment(document, fragment, _rule())
    numbering = document.part.numbering_part.element

    level = numbering.find(
        ".//w:lvl[@w:ilvl='8']",
        namespaces={"w": qn("w:w").split("}", 1)[0][1:]},
    )
    assert level is not None
    assert level.find(qn("w:numFmt")).get(qn("w:val")) == "chineseCounting"
    assert level.find(qn("w:lvlText")).get(qn("w:val")) == "%9、"
    override = numbering.find(
        ".//w:lvlOverride[@w:ilvl='8']",
        namespaces={"w": qn("w:w").split("}", 1)[0][1:]},
    )
    assert override is not None
    assert override.find(qn("w:startOverride")).get(qn("w:val")) == "5"


def test_all_occurrences_render_in_document_order_and_exactly_one_blocks_atomically():
    fragment = DocumentFragment((ParagraphBlock((_text("复用内容"),)),))
    document = Document()
    document.add_paragraph(TOKEN)
    document.add_paragraph("中间")
    document.add_paragraph(TOKEN)

    receipt = render_document_fragment(
        document,
        fragment,
        _rule(occurrence_policy="all"),
    )
    assert receipt.occurrence_count == 2
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "复用内容",
        "中间",
        "复用内容",
    ]
    assert [item.original_body_index for item in receipt.occurrences] == [0, 2]

    blocked = Document()
    blocked.add_paragraph(TOKEN)
    blocked.add_paragraph(TOKEN)
    before = blocked.element.xml
    preflight = ContentDocxRenderer().preflight(blocked, fragment, _rule())
    assert preflight.ready is False
    assert any(item.code == "anchor_occurrence_mismatch" for item in preflight.diagnostics)
    assert blocked.element.xml == before
    with pytest.raises(ContentRenderBlockedError, match="anchor_occurrence_mismatch"):
        render_document_fragment(blocked, fragment, _rule())
    assert blocked.element.xml == before


@pytest.mark.parametrize("surface", ["mixed_body", "table", "header"])
def test_nonexclusive_or_forbidden_anchor_surfaces_block_before_writes(surface):
    document = Document()
    if surface == "mixed_body":
        document.add_paragraph("前缀 " + TOKEN)
    elif surface == "table":
        document.add_table(rows=1, cols=1).cell(0, 0).text = TOKEN
        document.add_paragraph(TOKEN)
    else:
        document.sections[0].header.paragraphs[0].text = TOKEN
        document.add_paragraph(TOKEN)
    before = document.element.xml
    fragment = DocumentFragment((ParagraphBlock((_text("不会写入"),)),))

    with pytest.raises(ContentRenderBlockedError):
        render_document_fragment(document, fragment, _rule())

    assert document.element.xml == before
    assert "不会写入" not in document.element.xml


def test_anchor_paragraph_cannot_silently_delete_a_section_boundary():
    document = Document()
    paragraph = document.add_paragraph(TOKEN)
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    before = document.element.xml

    with pytest.raises(ContentRenderBlockedError) as raised:
        render_document_fragment(
            document,
            DocumentFragment((ParagraphBlock((_text("正文"),)),)),
            _rule(),
        )

    assert any(
        item.code == "anchor_section_boundary_unsupported"
        for item in raised.value.diagnostics
    )
    assert document.element.xml == before


def test_anchor_paragraph_cannot_silently_delete_non_plain_ooxml():
    document = Document()
    paragraph = document.add_paragraph(TOKEN)
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "9")
    bookmark.set(qn("w:name"), "must-not-be-deleted")
    paragraph._p.append(bookmark)
    before = document.element.xml

    with pytest.raises(ContentRenderBlockedError) as raised:
        render_document_fragment(
            document,
            DocumentFragment((ParagraphBlock((_text("正文"),)),)),
            _rule(),
        )

    assert any(
        item.code == "anchor_not_isolated"
        for item in raised.value.diagnostics
    )
    assert document.element.xml == before


def test_renderer_emits_image_job_without_owning_artifact_resource_validation():
    document = Document()
    document.add_paragraph(TOKEN)
    fragment = DocumentFragment((ImageBlock("missing.png"),))

    receipt = render_document_fragment(document, fragment, _rule())

    assert len(receipt.image_job_drafts) == 1
    assert receipt.image_job_drafts[0].resource_id == "missing.png"


def test_relative_heading_maps_shallowest_below_preceding_heading_and_bounds_block():
    document = Document()
    document.add_paragraph("父标题", style="Heading 2")
    document.add_paragraph(TOKEN)
    fragment = DocumentFragment(
        (
            HeadingBlock(2, (_text("浅层"),)),
            HeadingBlock(3, (_text("深层"),)),
        )
    )
    render_document_fragment(
        document,
        fragment,
        _rule(heading_policy="relative_to_anchor"),
    )
    assert [paragraph.style.style_id for paragraph in document.paragraphs] == [
        "Heading2",
        "Heading3",
        "Heading4",
    ]

    blocked = Document()
    blocked.add_paragraph(TOKEN)
    before = blocked.element.xml
    with pytest.raises(ContentRenderBlockedError, match="heading_level_out_of_range"):
        render_document_fragment(
            blocked,
            DocumentFragment((HeadingBlock(9, (_text("越界"),)),)),
            _rule(heading_level_offset=1),
        )
    assert blocked.element.xml == before


def test_each_list_block_gets_real_independent_numbering_restarted_from_one():
    document = Document()
    document.add_paragraph(TOKEN)
    fragment = DocumentFragment(
        (
            ListBlock(True, 1, 0, (ListItem((_text("一"),)), ListItem((_text("二"),)))),
            ListBlock(True, 1, 0, (ListItem((_text("重新一"),)),)),
            ListBlock(False, 1, 1, (ListItem((_text("项目"),)),)),
        )
    )
    render_document_fragment(document, fragment, _rule())

    num_ids = [
        paragraph._p.find(".//" + qn("w:numId")).get(qn("w:val"))
        for paragraph in document.paragraphs
    ]
    assert num_ids[0] == num_ids[1]
    assert len({num_ids[0], num_ids[2], num_ids[3]}) == 3
    numbering = document.part.numbering_part.element
    numbering_tags = [item.tag for item in numbering]
    assert max(
        index for index, tag in enumerate(numbering_tags) if tag == qn("w:abstractNum")
    ) < min(index for index, tag in enumerate(numbering_tags) if tag == qn("w:num"))
    for num_id in {num_ids[0], num_ids[2], num_ids[3]}:
        num = next(
            item
            for item in numbering.findall(qn("w:num"))
            if item.get(qn("w:numId")) == num_id
        )
        abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        abstract = next(
            item
            for item in numbering.findall(qn("w:abstractNum"))
            if item.get(qn("w:abstractNumId")) == abstract_id
        )
        assert all(
            level.find(qn("w:start")).get(qn("w:val")) == "1"
            for level in abstract.findall(qn("w:lvl"))
        )
    bullet_num = next(
        item
        for item in numbering.findall(qn("w:num"))
        if item.get(qn("w:numId")) == num_ids[3]
    )
    bullet_abstract_id = bullet_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    bullet_abstract = next(
        item
        for item in numbering.findall(qn("w:abstractNum"))
        if item.get(qn("w:abstractNumId")) == bullet_abstract_id
    )
    assert bullet_abstract.find(qn("w:lvl")).find(qn("w:numFmt")).get(qn("w:val")) == "bullet"
    reopened = _round_trip(document)
    assert all(
        paragraph._p.find(".//" + qn("w:numId")) is not None
        for paragraph in reopened.paragraphs
    )


def test_external_links_create_target_relationships_and_unsupported_protocol_blocks():
    document = Document()
    document.add_paragraph(TOKEN)
    fragment = DocumentFragment(
        (
            ParagraphBlock(
                (
                    InlineContent(
                        kind=InlineKind.HYPERLINK,
                        text="网站",
                        href="https://example.com/path",
                    ),
                    _text(" / "),
                    InlineContent(
                        kind=InlineKind.HYPERLINK,
                        text="邮箱",
                        href="mailto:user@example.com",
                    ),
                )
            ),
        )
    )
    render_document_fragment(document, fragment, _rule())
    hyperlinks = document.paragraphs[0]._p.findall(qn("w:hyperlink"))
    assert len(hyperlinks) == 2
    targets = {
        document.part.rels[item.get(qn("r:id"))].target_ref for item in hyperlinks
    }
    assert targets == {"https://example.com/path", "mailto:user@example.com"}
    assert all(
        document.part.rels[item.get(qn("r:id"))].reltype == RT.HYPERLINK
        for item in hyperlinks
    )

    blocked = Document()
    blocked.add_paragraph(TOKEN)
    with pytest.raises(ContentRenderBlockedError, match="hyperlink_protocol_unsupported"):
        render_document_fragment(
            blocked,
            DocumentFragment(
                (
                    ParagraphBlock(
                        (
                            InlineContent(
                                kind=InlineKind.HYPERLINK,
                                text="本地",
                                href="file:///C:/secret.txt",
                            ),
                        )
                    ),
                )
            ),
            _rule(),
        )


def test_simple_table_has_rectangular_grid_consistent_cell_widths_and_no_fixed_height():
    document = Document()
    document.add_paragraph(TOKEN)
    block = TableBlock(
        (
            TableRow(
                (
                    TableCell((ParagraphBlock((_text("A"),)),)),
                    TableCell((ParagraphBlock((_text("B"),)),)),
                )
            ),
            TableRow(
                (
                    TableCell((ParagraphBlock((_text("C"),)),)),
                    TableCell((ParagraphBlock((_text("D"),)),)),
                )
            ),
        )
    )
    render_document_fragment(document, DocumentFragment((block,)), _rule())

    assert len(document.tables) == 1
    table = document.tables[0]._tbl
    table_properties = table.find(qn("w:tblPr"))
    assert table_properties.find(qn("w:tblStyle")).get(qn("w:val")) == "TableGrid"
    assert table_properties.find(qn("w:tblInd")).get(qn("w:w")) == "120"
    cell_margins = table_properties.find(qn("w:tblCellMar"))
    assert [
        int(cell_margins.find(qn(f"w:{side}")).get(qn("w:w")))
        for side in ("top", "left", "bottom", "right")
    ] == [80, 120, 80, 120]
    grid_widths = [
        int(item.get(qn("w:w")))
        for item in table.find(qn("w:tblGrid")).findall(qn("w:gridCol"))
    ]
    assert len(grid_widths) == 2
    assert max(grid_widths) - min(grid_widths) <= 1
    for row in table.findall(qn("w:tr")):
        assert row.find(qn("w:trPr")) is None or row.find(".//" + qn("w:trHeight")) is None
        widths = [
            int(cell.find(".//" + qn("w:tcW")).get(qn("w:w")))
            for cell in row.findall(qn("w:tc"))
        ]
        assert widths == grid_widths

    reopened = _round_trip(document)
    assert [[cell.text for cell in row.cells] for row in reopened.tables[0].rows] == [
        ["A", "B"],
        ["C", "D"],
    ]


def test_page_break_policy_drops_or_preserves_explicit_break_deterministically():
    fragment = DocumentFragment(
        (ParagraphBlock((_text("前"),)), PageBreakBlock(), ParagraphBlock((_text("后"),)))
    )
    dropped = Document()
    dropped.add_paragraph(TOKEN)
    drop_receipt = render_document_fragment(dropped, fragment, _rule())
    assert drop_receipt.dropped_page_break_count == 1
    assert dropped.element.body.find(".//" + qn("w:br")) is None

    preserved = Document()
    preserved.add_paragraph(TOKEN)
    keep_receipt = render_document_fragment(
        preserved,
        fragment,
        _rule(page_break_policy="preserve_explicit"),
    )
    assert keep_receipt.dropped_page_break_count == 0
    page_break = preserved.element.body.find(".//" + qn("w:br"))
    assert page_break is not None
    assert page_break.get(qn("w:type")) == "page"


def test_image_blocks_create_only_deterministic_deferred_sentinels_and_collision_safe_ids():
    def build_document():
        document = Document()
        document.add_picture(
            BytesIO(
                b64decode(
                    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwC"
                    "AAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
                )
            )
        )
        paragraph = document.add_paragraph()
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), "17")
        start.set(qn("w:name"), "ExistingMarker")
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "17")
        paragraph._p.append(start)
        paragraph._p.append(end)
        document.add_paragraph(TOKEN)
        return document

    fragment = DocumentFragment(
        (ImageBlock("diagram.png", alt_text="路线图", width_px=800, height_px=600),),
    )
    first = build_document()
    second = build_document()
    first_receipt = render_document_fragment(first, fragment, _rule())
    second_receipt = render_document_fragment(second, fragment, _rule())

    draft = first_receipt.image_job_drafts[0]
    existing_docpr_id = max(
        int(item.get("id"))
        for item in first.element.body.findall(".//" + qn("wp:docPr"))
    )
    assert draft.resource_id == "diagram.png"
    assert draft.sequence == 0
    assert draft.reserved_docpr_id > existing_docpr_id
    assert draft.stable_marker_id == second_receipt.image_job_drafts[0].stable_marker_id
    assert draft.occurrence_id == first_receipt.occurrences[0].occurrence_id
    assert len(first.element.body.findall(".//" + qn("w:drawing"))) == 1
    bookmarks = first.element.body.findall(".//" + qn("w:bookmarkStart"))
    assert {item.get(qn("w:name")) for item in bookmarks} == {
        "ExistingMarker",
        draft.stable_marker_id,
    }
    assert len({item.get(qn("w:id")) for item in bookmarks}) == 2
    marker = next(item for item in bookmarks if item.get(qn("w:name")) == draft.stable_marker_id)
    assert int(marker.get(qn("w:id"))) > 17
    assert first.element.body.find(".//" + qn("w:vanish")) is not None
    assert TOKEN not in first.element.xml

    reopened = _round_trip(first)
    assert any(
        item.get(qn("w:name")) == draft.stable_marker_id
        for item in reopened.element.body.findall(".//" + qn("w:bookmarkStart"))
    )
