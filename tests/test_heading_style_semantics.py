import sys
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.heading_style_semantics import (
    NON_NUMBERED_HEADING_STYLE_KEY,
    resolve_heading_style,
    resolve_heading_style_source,
    resolve_non_numbered_heading_style,
)
from src.config.resolver import resolve_config, resolve_template_baseline
from src.config.scene import SceneWorkspace
from src.config.style_semantics import resolve_style_size_pt
from src.config.template import StyleConfig, TemplateConfig
from src.modules.basic.paragraph_style import ParagraphStyleModule
from src.modules.registry import create_all_modules
from src.pipeline.context import PipelineContext
from src.pipeline.module_selection import build_module_selection_plan
from src.pipeline.runner import Pipeline
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn
from src.ui.adapters.heading_numbering_adapter import HeadingNumberingAdapter
from src.ui.panels.template_preview.model import PreviewBlockKind, TemplatePreviewMode
from src.ui.panels.template_preview.projector import build_template_preview_projection


def test_resolve_heading_style_uses_shared_fallback_chain():
    cfg = TemplateConfig()
    cfg.styles = {
        "normal": StyleConfig(size_pt=10.0, bold=False),
        "body": StyleConfig(size_pt=12.0, bold=False, alignment="justify"),
        "heading": StyleConfig(size_pt=14.0, bold=True, alignment="center"),
        "heading1": StyleConfig(size_pt=16.0, italic=True, alignment="right"),
    }

    level_one = resolve_heading_style(cfg, 1, include_body_fallback=True)
    level_two = resolve_heading_style(cfg, 2, include_body_fallback=True)

    assert level_one is not None
    assert level_two is not None
    assert resolve_style_size_pt(level_one) == 16.0
    assert level_one.italic is True
    assert level_one.alignment == "right"
    assert resolve_style_size_pt(level_two) == 14.0
    assert level_two.bold is True
    assert level_two.alignment == "center"


def test_heading_style_override_is_seeded_from_body_fallback():
    cfg = TemplateConfig()
    cfg.styles = {
        "body": StyleConfig(size_pt=13.0, bold=True, alignment="center"),
    }
    adapter = HeadingNumberingAdapter()
    adapter.set_template(cfg)

    adapter.set_heading_style_field(1, "size_pt", 18.0)

    assert "heading1" in cfg.styles
    assert cfg.styles["heading1"].size_pt == 18.0
    assert cfg.styles["heading1"].bold is True
    assert cfg.styles["heading1"].alignment == "center"


def test_resolve_heading_style_source_reports_effective_origin():
    cfg = TemplateConfig()
    cfg.styles = {
        "body": StyleConfig(size_pt=12.0),
        "heading": StyleConfig(size_pt=14.0, bold=True),
    }

    assert resolve_heading_style_source(cfg, 1, include_body_fallback=True) == "heading"

    cfg.styles["heading1"] = StyleConfig(size_pt=16.0, italic=True)
    assert resolve_heading_style_source(cfg, 1, include_body_fallback=True) == "heading1"


def test_non_numbered_heading_style_accepts_resolved_config():
    cfg = TemplateConfig()
    cfg.heading_model.non_numbered_heading_style_mode = "custom"
    cfg.styles = {
        "body": StyleConfig(size_pt=12.0),
        "heading1": StyleConfig(size_pt=16.0, bold=True),
        NON_NUMBERED_HEADING_STYLE_KEY: StyleConfig(size_pt=13.0, italic=True),
    }

    resolved = resolve_config(cfg, SceneWorkspace())
    style = resolve_non_numbered_heading_style(resolved, include_body_fallback=True)

    assert style is not None
    assert resolve_style_size_pt(style) == 13.0
    assert style.italic is True


def test_non_numbered_heading_gets_the_word_style_used_by_styleref():
    document = Document()
    paragraph = document.add_heading("参考文献", level=1)
    template = TemplateConfig()
    template.styles["heading1"] = StyleConfig(size_pt=16, bold=True)
    config = resolve_config(template, SceneWorkspace())
    context = PipelineContext(
        doc_tree=SimpleNamespace(
            get_special_title_match=lambda index: (
                "special-title:exact:test" if index == 0 else None
            ),
            get_section_for_paragraph=lambda _index: "references",
        )
    )

    ParagraphStyleModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    expected_style = config.heading_model.non_numbered_heading_style_name
    assert paragraph.style.name == expected_style
    assert document.styles[expected_style].base_style.name == "Heading 1"


def test_next_page_section_setting_applies_page_break_to_body_chapters_only():
    document = Document()
    chapter = document.add_heading("第一章 绪论", level=1)
    references = document.add_heading("参考文献", level=1)
    template = TemplateConfig()
    template.styles["heading1"] = StyleConfig(size_pt=16, bold=True)
    template.section.section_break_type = "nextPage"
    config = resolve_config(template, SceneWorkspace())
    context = PipelineContext(
        doc_tree=SimpleNamespace(
            get_special_title_match=lambda index: (
                "special-title:exact:test" if index == 1 else None
            ),
            get_section_for_paragraph=lambda index: (
                "body" if index == 0 else "references"
            ),
        )
    )

    ParagraphStyleModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    assert chapter.paragraph_format.page_break_before is True
    assert references.paragraph_format.page_break_before is not True


def test_paragraph_formatting_keeps_section_properties_as_the_last_ppr_child():
    document = Document()
    paragraph = document.add_paragraph("摘要正文")
    paragraph_properties = paragraph._p.get_or_add_pPr()
    paragraph_properties.append(OxmlElement("w:sectPr"))
    paragraph_properties.append(OxmlElement("w:pPrChange"))
    template = TemplateConfig()
    template.styles["normal"] = StyleConfig(
        size_pt=12,
        space_before_pt=6,
        space_after_pt=6,
    )
    config = resolve_config(template, SceneWorkspace())

    ParagraphStyleModule().apply(
        document,
        config,
        ChangeTracker(),
        PipelineContext(),
    )

    children = list(paragraph._p.get_or_add_pPr())
    child_tags = [child.tag for child in children]
    assert paragraph._p[0] is paragraph._p.pPr
    assert child_tags.index(qn("w:spacing")) < child_tags.index(qn("w:ind"))
    assert child_tags.index(qn("w:ind")) < child_tags.index(qn("w:jc"))
    assert child_tags.index(qn("w:jc")) < child_tags.index(qn("w:sectPr"))
    assert child_tags.index(qn("w:sectPr")) < child_tags.index(qn("w:pPrChange"))


def test_preview_adapter_and_runtime_share_heading_body_fallback(tmp_path):
    cfg = TemplateConfig()
    cfg.styles = {
        "body": StyleConfig(size_pt=13.0, bold=True, alignment="center"),
    }

    resolved = resolve_template_baseline(cfg)
    modules = create_all_modules()
    selection = build_module_selection_plan(modules, is_requested=lambda _name: True)
    projection = build_template_preview_projection(
        resolved,
        selection,
        mode=TemplatePreviewMode.TEMPLATE_BASELINE,
    )
    preview_style = next(
        block.style
        for block in projection.blocks
        if block.kind is PreviewBlockKind.HEADING and block.level == 1
    )
    adapter = HeadingNumberingAdapter()
    adapter.set_template(cfg)
    adapter_style = adapter.get_heading_style(1)

    assert preview_style.size_pt == 13.0
    assert preview_style.bold is True
    assert preview_style.alignment == "center"
    assert resolve_style_size_pt(adapter_style) == 13.0
    assert adapter_style.bold is True
    assert adapter_style.alignment == "center"

    source = tmp_path / "heading_body_fallback.docx"
    doc = Document()
    doc.add_heading("Shared fallback", level=1)
    doc.save(source)

    config = resolve_config(cfg, SceneWorkspace())
    pipeline = Pipeline(modules=[ParagraphStyleModule()], config=config)
    result = pipeline.execute(str(source))

    assert result.success, f"pipeline failed: {result.error}"

    out_doc = Document(result.output_paths["final"])
    para = out_doc.paragraphs[0]
    run = para.runs[0]

    assert run.font.size is not None
    assert round(run.font.size.pt, 1) == 13.0
    assert run.font.bold is True
    assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER
