import json

from docx import Document

from src.config.resolved import ResolvedConfig
from src.config.scene import ComplianceProfile, InputSourceProfile
from src.pipeline.runner import Pipeline
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.journal_citation_schema import inspect_journal_citations
from src.services.production_runtime.delivery_reporting import pipeline_context_payload


def _journal_config(entity_data: dict[str, object]) -> ResolvedConfig:
    return ResolvedConfig(
        input_source_profile=InputSourceProfile(
            accepted_formats=["docx", "bibtex", "csl_json"],
            structured_formats=["bibtex", "csl_json"],
            material_schema_id="journal_submission_materials_v1",
        ),
        compliance_profile=ComplianceProfile(
            profile_id="journal_submission",
            rule_family="journal_en",
            count_profile_id="journal_words",
        ),
        entity_data=dict(entity_data),
    )


def test_journal_citations_parse_bibtex_and_match_body_citation_keys():
    doc = Document()
    doc.add_paragraph(r"Prior work \cite{smith2020,missing2024} and @doe2021.")
    bibtex = """
@article{smith2020,
  title={A study},
  author={Smith, Jane},
  year={2020}
}
@article{doe2021,
  title={Another study},
  author={Doe, John},
  year={2021}
}
"""

    result = inspect_journal_citations(_journal_config({"bibtex": bibtex}), doc)
    issue_kinds = {issue.kind for issue in result.issues}

    assert result.status == "error"
    assert result.source_key == "entity_data.bibtex"
    assert result.summary.source_format == "bibtex"
    assert result.summary.reference_count == 2
    assert result.summary.citation_count == 3
    assert result.summary.matched_citation_count == 2
    assert result.summary.missing_reference_count == 1
    assert "missing_reference_for_citation" in issue_kinds
    assert "missing2024" in result.citation_keys


def test_journal_citations_parse_csl_json_and_report_duplicate_or_incomplete_items():
    csl_items = [
        {
            "id": "smith2020",
            "type": "article-journal",
            "title": "A study",
            "author": [{"family": "Smith", "given": "Jane"}],
            "issued": {"date-parts": [[2020]]},
        },
        {
            "id": "smith2020",
            "type": "article-journal",
            "author": [{"family": "Smith", "given": "Jane"}],
        },
    ]

    result = inspect_journal_citations(
        _journal_config(
            {
                "csl_json": json.dumps(csl_items, ensure_ascii=False),
                "citation_keys": ["smith2020"],
            }
        )
    )
    issue_kinds = {issue.kind for issue in result.issues}

    assert result.status == "error"
    assert result.summary.source_format == "csl_json"
    assert result.summary.reference_count == 1
    assert result.summary.duplicate_key_count == 1
    assert result.summary.incomplete_reference_count == 1
    assert "duplicate_reference_key" in issue_kinds
    assert "incomplete_reference_metadata" in issue_kinds


def test_pipeline_records_journal_citation_result_and_reports(tmp_path):
    source = tmp_path / "journal.docx"
    doc = Document()
    doc.add_paragraph(r"This manuscript cites \cite{smith2020}.")
    doc.save(source)

    bibtex = """
@article{smith2020,
  title={A study},
  author={Smith, Jane},
  year={2020}
}
"""
    result = Pipeline([], _journal_config({"bibtex": bibtex}), output_dir=str(tmp_path)).execute(
        str(source)
    )

    assert result.success is True
    assert result.context is not None
    assert result.context.journal_citations.status == "ok"
    assert pipeline_context_payload(result.context)["journal_citations"]["status"] == "ok"
    records = result.tracker.get_by_module("journal_citations")
    assert records
    assert records[0].change_type == "citation_source_validation"
    assert "references=1" in records[0].after
    assert "citations=1" in records[0].after

    report_json = tmp_path / "journal_changes.json"
    report_md = tmp_path / "journal_changes.md"
    write_json_report(
        result,
        input_path=source,
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )

    data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert data["journal_citations"]["status"] == "ok"
    assert data["journal_citations"]["summary"]["reference_count"] == 1
    assert data["journal_citations"]["summary"]["matched_citation_count"] == 1
    assert "## 英文期刊引用源校验" in markdown
    assert "smith2020" in markdown
