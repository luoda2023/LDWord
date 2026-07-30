from __future__ import annotations

import inspect
import json
import subprocess
import sys
from dataclasses import replace
from hashlib import sha256
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from src.config.content_materials import FileAssetRef
from src.config.image_materials import (
    ImageAnchorOrigin,
    ImageAnchorRef,
    ImageCoLocationGuard,
    ImagePlacementMode,
    ImagePlacementPolicy,
    ResolvedImageInsertionPlan,
    ResolvedImageWatermark,
)
from src.shared.engine import office_image_layout as layout_module
from src.shared.engine import office_image_layout_child_lifecycle as lifecycle_module
from src.shared.engine import office_image_layout_contracts as contract_module
from src.shared.engine import office_image_layout_coordinator as coordinator_module
from src.shared.engine import office_image_layout_inventory as inventory_module
from src.shared.engine import office_image_layout_preflight as preflight_module
from src.shared.engine import office_image_layout_worker as worker_module
from src.shared.engine.office_image_layout import (
    PROVIDER_SPECS,
    ImageGeometryError,
    ImageGeometryInput,
    LayoutFailure,
    LayoutFailureCode,
    LayoutStatus,
    OfficeImageJobReceipt,
    OfficeImageLayoutJob,
    OfficeImageLayoutReceipt,
    OfficeImageLayoutRequest,
    OfficeImageProvider,
    StabilizationRoundReceipt,
    compute_image_geometry,
    run_office_image_layout,
)
from src.shared.engine.prepared_image import PreparedImage
from src.shared.io import safe_docx_package


def _sha(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _png(path: Path, size: tuple[int, int] = (800, 500)) -> None:
    Image.new("RGB", size, (40, 100, 180)).save(path, format="PNG")


def _bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _docx(
    path: Path,
    *,
    anchor: str = "anchor_1",
    guard: str = "guard_1",
    duplicate_anchor: bool = False,
    anchor_in_table: bool = False,
    include_guard: bool = True,
) -> None:
    doc = Document()
    if include_guard:
        guard_paragraph = doc.add_paragraph("Arbitrary heading, any font height")
        _bookmark(guard_paragraph, guard, 1)
    if anchor_in_table:
        paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        paragraph.add_run(anchor)
    else:
        paragraph = doc.add_paragraph(anchor)
    _bookmark(paragraph, anchor, 2)
    if duplicate_anchor:
        duplicate = doc.add_paragraph("duplicate")
        _bookmark(duplicate, anchor, 3)
    for index in range(40):
        doc.add_paragraph(f"filler {index}")
    doc.save(path)


def _docx_with_existing_images(path: Path, image_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("body image").add_run().add_picture(str(image_path))
    doc.sections[0].header.paragraphs[0].add_run().add_picture(str(image_path))
    guard_paragraph = doc.add_paragraph("Arbitrary heading, any font height")
    _bookmark(guard_paragraph, "guard_1", 1)
    anchor_paragraph = doc.add_paragraph("anchor_1")
    _bookmark(anchor_paragraph, "anchor_1", 2)
    doc.save(path)


def _docx_with_hidden_anchor_sentinel(
    path: Path,
    *,
    anchor: str = "anchor_1",
    guard: str = "guard_1",
) -> None:
    doc = Document()
    guard_paragraph = doc.add_paragraph("Arbitrary heading, any font height")
    _bookmark(guard_paragraph, guard, 1)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(anchor)
    run._r.get_or_add_rPr().append(OxmlElement("w:vanish"))
    run._r.get_or_add_rPr().append(OxmlElement("w:noProof"))
    _bookmark(paragraph, anchor, 2)
    for index in range(10):
        doc.add_paragraph(f"filler {index}")
    doc.save(path)


def _rewrite_zip_member(path: Path, member: str, payload: bytes) -> None:
    temporary = path.with_suffix(".rewritten.docx")
    with ZipFile(path, "r") as source, ZipFile(
        temporary,
        "w",
        compression=ZIP_DEFLATED,
    ) as target:
        for info in source.infolist():
            target.writestr(info, payload if info.filename == member else source.read(info))
    temporary.replace(path)


def _image_identity(**changes: str):
    values = {
        "part_name": "word/document.xml",
        "drawing_kind": "inline",
        "doc_pr_id": "7",
        "doc_pr_name": "Picture 7",
        "alternative_text": "existing-alt",
        "title": "existing-title",
        "run_visibility": "visible",
        "layout_extent_cx_emu": 1_080_000,
        "layout_extent_cy_emu": 360_000,
        "drawing_semantic_sha256": "6" * 64,
        "relationship_reference": "embed",
        "relationship_id": "rId7",
        "relationship_type": (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
        ),
        "relationship_target": "media/image7.png",
        "target_mode": "internal",
        "media_part_name": "word/media/image7.png",
        "media_sha256": "7" * 64,
    }
    values.update(changes)
    return layout_module.OOXMLImageIdentity(**values)


def _job(
    tmp_path: Path,
    *,
    job_id: str = "job-1",
    anchor: str = "anchor_1",
    guard: str = "guard_1",
    mode: ImagePlacementMode = ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
    sequence: int | None = 0,
    fixed_width_cm: float | None = None,
) -> OfficeImageLayoutJob:
    source = tmp_path / f"{job_id}-source.png"
    prepared_path = tmp_path / f"{job_id}-prepared.png"
    _png(source)
    prepared_path.write_bytes(source.read_bytes())
    digest = _sha(source)
    if mode is ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE:
        placement = ImagePlacementPolicy(
            mode=mode,
            max_width_cm=None,
            co_location_guard=ImageCoLocationGuard.PRECEDING_NONEMPTY_PARAGRAPH,
        )
    elif mode in {
        ImagePlacementMode.NATURAL_SIZE,
        ImagePlacementMode.FIT_CONTAINER_FLOW,
    }:
        placement = ImagePlacementPolicy(mode=mode)
    else:
        placement = ImagePlacementPolicy(
            mode=mode,
            fixed_width_cm=fixed_width_cm or 10.0,
        )
    plan = ResolvedImageInsertionPlan(
        job_id=job_id,
        image_ref=FileAssetRef(
            source_path=str(source),
            original_name=source.name,
            media_type="image/png",
            content_sha256=digest,
            byte_size=source.stat().st_size,
        ),
        anchor=ImageAnchorRef(
            origin=ImageAnchorOrigin.CONTENT_RESOURCE,
            stable_marker_id=anchor,
            guard_marker_id=(
                guard
                if mode is ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE
                else ""
            ),
        ),
        occurrence_id=f"occ-{job_id}",
        watermark=ResolvedImageWatermark.disabled(),
        placement=placement,
        source_role="test",
        sequence=sequence,
    )
    prepared = PreparedImage(
        cache_key=digest,
        output_path=str(prepared_path),
        output_sha256=_sha(prepared_path),
        media_type="image/png",
        width_px=800,
        height_px=500,
        source_sha256=digest,
        watermark_text_sha256="",
    )
    return OfficeImageLayoutJob(plan, prepared)


def _request(
    docx: Path,
    jobs: tuple[OfficeImageLayoutJob, ...],
    *,
    rounds: int = 3,
) -> OfficeImageLayoutRequest:
    return OfficeImageLayoutRequest(
        transaction_id="tx-test",
        source_docx_path=str(docx),
        source_docx_sha256=_sha(docx),
        source_docx_size=docx.stat().st_size,
        provider=OfficeImageProvider.WORD,
        jobs=jobs,
        timeout_seconds=5,
        max_stabilization_rounds=rounds,
    )


def _job_receipt(
    job_id: str = "job-1",
    *,
    claimed: bool = True,
) -> OfficeImageJobReceipt:
    return OfficeImageJobReceipt(
        job_id=job_id,
        mode=(
            ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE
            if claimed
            else ImagePlacementMode.FIXED_BOX_FLOW
        ),
        sequence=0,
        shape_id=f"shape-{job_id}",
        source_image_sha256="d" * 64,
        prepared_image_sha256="e" * 64,
        measurement_guard_page=1 if claimed else None,
        measurement_anchor_page=1,
        measurement_anchor_start=20,
        anchor_y_pt=100,
        section_index=1,
        column_index=1,
        column_count=1,
        container_width_pt=450,
        flow_bottom_pt=760,
        paragraph_reserve_pt=16,
        safety_margin_pt=8,
        available_height_pt=636 if claimed else None,
        width_cap_pt=450,
        initial_target_width_pt=400,
        initial_target_height_pt=250,
        initial_scale_ratio=0.667,
        readability_warning=False,
        co_location_claimed=claimed,
        final_guard_page=1 if claimed else None,
        final_anchor_page=1,
        final_image_page=1,
        final_image_y_pt=120,
        final_width_pt=400,
        final_height_pt=250,
        proportion_error=0,
        boundary_ok=True,
        guard_xml_sha256_before="a" * 64 if claimed else "",
        guard_xml_sha256_after="a" * 64 if claimed else "",
        guard_keep_sha256_before="b" * 64 if claimed else "",
        guard_keep_sha256_after="b" * 64 if claimed else "",
        correction_count=1,
        initial_actual_width_pt=400,
        initial_actual_height_pt=250,
        initial_proportion_error=0.002273,
        dimension_candidate_count=3,
        selected_width_offset_pt=0.75,
    )


def test_fit_geometry_uses_real_anchor_height_and_only_shrinks_image() -> None:
    short_heading = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
            1200,
            800,
            450,
            anchor_y_pt=120,
            flow_bottom_pt=760,
            paragraph_reserve_pt=20,
            safety_margin_pt=10,
        )
    )
    tall_heading = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
            1200,
            800,
            450,
            anchor_y_pt=650,
            flow_bottom_pt=760,
            paragraph_reserve_pt=20,
            safety_margin_pt=10,
        )
    )

    assert short_heading.target_width_pt == 450
    assert short_heading.target_height_pt == 300
    assert tall_heading.available_height_pt == 80
    assert tall_heading.target_height_pt == 80
    assert tall_heading.target_width_pt == 120
    assert tall_heading.aspect_ratio == short_heading.aspect_ratio == 1.5
    # The pure contract has no heading-move/page-break decision: measured Y only
    # changes the contained image dimensions.
    assert not hasattr(tall_heading, "move_preceding_text")


def test_all_placement_modes_and_none_max_width_respect_container() -> None:
    natural = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.NATURAL_SIZE,
            200,
            100,
            300,
            100,
            700,
            20,
        )
    )
    container_flow = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIT_CONTAINER_FLOW,
            200,
            100,
            300,
            100,
            700,
            20,
        )
    )
    fixed = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIXED_BOX,
            800,
            400,
            300,
            100,
            700,
            20,
            fixed_width_pt=400,
            max_width_pt=None,
        )
    )
    flow = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIXED_BOX_FLOW,
            800,
            400,
            300,
            100,
            700,
            20,
            fixed_width_pt=250,
            max_width_pt=200,
        )
    )
    strict = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
            800,
            400,
            180,
            100,
            700,
            20,
            max_width_pt=None,
        )
    )

    assert (natural.target_width_pt, natural.target_height_pt) == (150, 75)
    assert natural.scale_ratio_96dpi == 1.0
    assert (container_flow.target_width_pt, container_flow.target_height_pt) == (
        300,
        150,
    )
    assert container_flow.scale_ratio_96dpi == 2.0
    assert (fixed.target_width_pt, fixed.target_height_pt) == (300, 150)
    assert fixed.width_cap_pt == 300
    assert (flow.target_width_pt, flow.target_height_pt) == (200, 100)
    assert strict.width_cap_pt == 180
    assert strict.target_width_pt == 180


def test_no_remaining_space_and_word_minimum_are_hard_blocks() -> None:
    with pytest.raises(ImageGeometryError) as no_space:
        compute_image_geometry(
            ImageGeometryInput(
                ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
                800,
                400,
                300,
                anchor_y_pt=690,
                flow_bottom_pt=700,
                paragraph_reserve_pt=8,
                safety_margin_pt=3,
            )
        )
    assert no_space.value.code is LayoutFailureCode.NO_REMAINING_SPACE

    with pytest.raises(ImageGeometryError) as too_small:
        compute_image_geometry(
            ImageGeometryInput(
                ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
                1000,
                1000,
                300,
                anchor_y_pt=698.8,
                flow_bottom_pt=700,
                paragraph_reserve_pt=0,
                safety_margin_pt=1,
                word_min_dimension_pt=0.5,
            )
        )
    assert too_small.value.code is LayoutFailureCode.BELOW_WORD_MINIMUM

    readable_but_warned = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
            1000,
            1000,
            300,
            anchor_y_pt=640,
            flow_bottom_pt=700,
            paragraph_reserve_pt=5,
            safety_margin_pt=5,
        )
    )
    assert readable_but_warned.target_height_pt == 50
    assert readable_but_warned.readability_warning is True


def test_current_column_width_is_the_geometry_cap() -> None:
    first_column = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIXED_BOX,
            1000,
            500,
            container_width_pt=220,
            anchor_y_pt=100,
            flow_bottom_pt=700,
            paragraph_reserve_pt=20,
            fixed_width_pt=400,
        )
    )
    wide_single_column = compute_image_geometry(
        ImageGeometryInput(
            ImagePlacementMode.FIXED_BOX,
            1000,
            500,
            container_width_pt=460,
            anchor_y_pt=100,
            flow_bottom_pt=700,
            paragraph_reserve_pt=20,
            fixed_width_pt=400,
        )
    )
    assert first_column.target_width_pt == 220
    assert wide_single_column.target_width_pt == 400


def test_explicit_width_and_height_avoids_wps_auto_dimension_ratio_drift() -> None:
    class QuantizedShape:
        def __init__(self) -> None:
            self._lock = -1
            self._width = 100.0
            self._height = 140.0
            self.events: list[tuple[str, float | int]] = []

        @property
        def LockAspectRatio(self):
            return self._lock

        @LockAspectRatio.setter
        def LockAspectRatio(self, value):
            self.events.append(("lock", value))
            self._lock = value

        @property
        def Width(self):
            return self._width

        @Width.setter
        def Width(self, value):
            self.events.append(("width", value))
            self._width = round(float(value) * 20) / 20
            if self._lock == -1:
                # Reproduce the measured WPS auto-derived 198.0pt height that
                # yielded ratio_error ~= 0.002273 for a 500x700 source.
                self._height = 198.0

        @property
        def Height(self):
            return self._height

        @Height.setter
        def Height(self, value):
            self.events.append(("height", value))
            self._height = round(float(value) * 20) / 20

    aspect = 500 / 700
    target_width = 5 * layout_module.POINTS_PER_CM
    target_height = target_width / aspect
    old_shape = QuantizedShape()
    old_shape.Width = target_width
    old_error = abs(old_shape.Width / old_shape.Height - aspect) / aspect
    assert old_error > layout_module.MAX_IMAGE_PROPORTION_ERROR

    shape = QuantizedShape()
    decision = worker_module._set_inline_shape_dimensions(
        shape,
        target_width,
        target_height,
        aspect,
        "job-quantized",
    )
    actual_width = decision.actual_width_pt
    actual_height = decision.actual_height_pt
    explicit_error = abs(actual_width / actual_height - aspect) / aspect

    assert [event[0] for event in shape.events] == [
        "lock",
        "width",
        "height",
        "lock",
    ]
    assert (actual_width, actual_height) == (141.75, 198.45)
    assert explicit_error < layout_module.MAX_IMAGE_PROPORTION_ERROR
    assert decision.corrected is False
    assert decision.candidate_count == 1
    assert decision.initial_width_pt == actual_width
    assert decision.initial_height_pt == actual_height
    assert decision.initial_proportion_error == decision.proportion_error


@pytest.mark.parametrize(
    "mode",
    (ImagePlacementMode.FIXED_BOX, ImagePlacementMode.FIXED_BOX_FLOW),
)
def test_bounded_wps_point_bucket_search_selects_minimum_viable_candidate(
    mode: ImagePlacementMode,
) -> None:
    aspect = 500 / 700
    geometry = compute_image_geometry(
        ImageGeometryInput(
            mode,
            500,
            700,
            container_width_pt=450,
            anchor_y_pt=100,
            flow_bottom_pt=700,
            paragraph_reserve_pt=20,
            fixed_width_pt=5 * layout_module.POINTS_PER_CM,
        )
    )
    target_width = geometry.target_width_pt
    target_height = geometry.target_height_pt

    class WpsBucketShape:
        def __init__(self) -> None:
            self._lock = -1
            self._requested_width = target_width
            self._width = 141.75
            self._height = 198.0

        @property
        def LockAspectRatio(self):
            return self._lock

        @LockAspectRatio.setter
        def LockAspectRatio(self, value):
            self._lock = value
            if value != -1:
                return
            offset = self._requested_width - target_width
            if offset >= 0.75 - 1e-9:
                self._width, self._height = 142.5, 199.5
            elif offset <= -0.25 + 1e-9:
                self._width, self._height = 141.0, 197.25
            else:
                self._width, self._height = 141.75, 198.0

        @property
        def Width(self):
            return self._width if self._lock == -1 else self._requested_width

        @Width.setter
        def Width(self, value):
            self._requested_width = float(value)

        @property
        def Height(self):
            return self._height

        @Height.setter
        def Height(self, _value):
            # Reproduce WPS deriving its own point bucket on relock.
            pass

    decision = worker_module._set_inline_shape_dimensions(
        WpsBucketShape(),
        target_width,
        target_height,
        aspect,
        "job-wps-bucket",
    )

    assert decision.corrected is True
    assert decision.candidate_count == 3
    assert decision.selected_width_offset_pt == 0.75
    assert (decision.initial_width_pt, decision.initial_height_pt) == (
        141.75,
        198.0,
    )
    assert decision.initial_proportion_error > (
        layout_module.MAX_IMAGE_PROPORTION_ERROR
    )
    assert (decision.actual_width_pt, decision.actual_height_pt) == (142.5, 199.5)
    assert decision.proportion_error == 0


def test_dimension_search_blocks_ratio_valid_candidate_outside_size_bound() -> None:
    class FarShape:
        @property
        def LockAspectRatio(self):
            return -1

        @LockAspectRatio.setter
        def LockAspectRatio(self, _value):
            pass

        @property
        def Width(self):
            return 144.0

        @Width.setter
        def Width(self, _value):
            pass

        @property
        def Height(self):
            return 201.6

        @Height.setter
        def Height(self, _value):
            pass

    target_width = 5 * layout_module.POINTS_PER_CM
    aspect = 500 / 700
    with pytest.raises(lifecycle_module.ChildBlocked) as captured:
        worker_module._set_inline_shape_dimensions(
            FarShape(),
            target_width,
            target_width / aspect,
            aspect,
            "job-outside-bound",
        )

    assert captured.value.failure.code is LayoutFailureCode.PROPORTION_CHANGED
    assert "no bounded image dimension solution" in captured.value.failure.message
    assert "search=±1pt" in captured.value.failure.message


def test_live_column_geometry_selects_the_second_text_column() -> None:
    class Column:
        def __init__(self, width: float, space_after: float):
            self.Width = width
            self.SpaceAfter = space_after

    class Columns:
        Count = 2

        def __init__(self):
            self.values = (Column(200, 20), Column(180, 0))

        def Item(self, index: int):
            return self.values[index - 1]

    class Setup:
        LeftMargin = 72
        TextColumns = Columns()

    class Section:
        PageSetup = Setup()

    class SectionsCollection:
        def Item(self, _index: int):
            return Section()

    class Target:
        Sections = SectionsCollection()

        def Information(self, constant: int):
            assert constant == 5
            return 300.0

    _section, index, count, width = worker_module._column_geometry(Target())
    assert (index, count, width) == (2, 2, 180)


def test_bookmark_preflight_blocks_table_missing_guard_and_duplicate_marker(
    tmp_path: Path,
) -> None:
    table_docx = tmp_path / "table.docx"
    _docx(table_docx, anchor_in_table=True)
    table_request = _request(table_docx, (_job(tmp_path, job_id="table"),))
    failures, _identities = preflight_module.parent_preflight(table_request)
    assert LayoutFailureCode.UNSUPPORTED_SURFACE in {item.code for item in failures}

    missing_guard_docx = tmp_path / "missing-guard.docx"
    _docx(missing_guard_docx, include_guard=False)
    missing_request = _request(
        missing_guard_docx,
        (_job(tmp_path, job_id="missing"),),
    )
    failures, _identities = preflight_module.parent_preflight(missing_request)
    assert LayoutFailureCode.BOOKMARK_MISSING in {item.code for item in failures}

    duplicate_docx = tmp_path / "duplicate.docx"
    _docx(duplicate_docx, duplicate_anchor=True)
    duplicate_request = _request(
        duplicate_docx,
        (_job(tmp_path, job_id="duplicate"),),
    )
    failures, _identities = preflight_module.parent_preflight(duplicate_request)
    assert LayoutFailureCode.BOOKMARK_NOT_UNIQUE in {item.code for item in failures}


def test_controlled_shadow_name_is_short_source_bound_and_stem_independent(
    tmp_path: Path,
) -> None:
    long_source = tmp_path / (("very-long-stage-name-" * 7) + ".docx")
    _docx(long_source)
    shadow = layout_module.build_controlled_layout_shadow_path(
        long_source,
        nonce="a" * 32,
    )
    old_style = long_source.parent / (
        f".{long_source.stem}.lark-layout-{'a' * 32}.docx"
    )

    assert shadow.parent == long_source.parent
    assert long_source.stem not in shadow.name
    assert len(shadow.name) <= 72
    assert len(str(shadow)) < len(str(old_style))
    assert layout_module.is_controlled_layout_shadow(long_source, shadow)
    assert not layout_module.is_controlled_layout_shadow(
        long_source,
        shadow.with_name(
            f"{layout_module.LAYOUT_SHADOW_PREFIX}{'0' * 16}-{'a' * 32}.docx"
        ),
    )
    other_source = long_source.with_name("other-source.docx")
    assert not layout_module.is_controlled_layout_shadow(other_source, shadow)
    with pytest.raises(ValueError):
        layout_module.build_controlled_layout_shadow_path(long_source, nonce="A" * 32)


def test_long_stage_parent_preflight_uses_bounded_shadow_and_cleans_it(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    long_source = tmp_path / (("pipeline-stage-" * 9) + ".docx")
    _docx(long_source)
    job = _job(tmp_path)
    request = _request(long_source, (job,))
    source_before = long_source.read_bytes()
    observed: list[Path] = []
    monkeypatch.setattr(coordinator_module, "_registration_status", lambda _prog: (True, "ok"))

    def blocked(_request, spec, shadow):
        observed.append(shadow)
        assert shadow.is_file()
        assert layout_module.is_controlled_layout_shadow(long_source, shadow)
        return contract_module.build_layout_receipt(
            _request,
            spec,
            LayoutStatus.BLOCKED,
            failures=(LayoutFailure(LayoutFailureCode.OFFICE_ERROR, "fake"),),
        )

    monkeypatch.setattr(coordinator_module, "_run_layout_child", blocked)

    receipt = run_office_image_layout(request)

    assert receipt.status is LayoutStatus.BLOCKED
    assert len(observed) == 1
    old_style = long_source.parent / (
        f".{long_source.stem}.lark-layout-{'a' * 32}.docx"
    )
    assert len(str(observed[0])) < len(str(old_style))
    assert len(observed[0].name) <= 72
    assert not observed[0].exists()
    assert long_source.read_bytes() == source_before


def test_ooxml_image_inventory_covers_body_and_header_without_zip_metadata(
    tmp_path: Path,
) -> None:
    image = tmp_path / "existing.png"
    _png(image, (320, 200))
    docx = tmp_path / "existing.docx"
    _docx_with_existing_images(docx, image)

    inventory = layout_module.inventory_ooxml_images(docx)

    assert len(inventory) == 2
    assert {item.part_name for item in inventory} == {
        "word/document.xml",
        "word/header1.xml",
    }
    assert {item.drawing_kind for item in inventory} == {"inline"}
    assert all(item.doc_pr_id and item.doc_pr_name for item in inventory)
    assert all(item.relationship_id.startswith("rId") for item in inventory)
    assert all(item.relationship_target for item in inventory)
    assert all(item.media_sha256 == _sha(image) for item in inventory)
    assert all(item.layout_extent_cx_emu > 0 for item in inventory)
    assert all(item.layout_extent_cy_emu > 0 for item in inventory)
    first_hash = layout_module.image_inventory_sha256(inventory)

    # Repacking changes ZIP metadata/compression but not any image identity.
    repacked = tmp_path / "repacked.docx"
    with ZipFile(docx, "r") as source, ZipFile(
        repacked,
        "w",
        compression=ZIP_DEFLATED,
        compresslevel=1,
    ) as target:
        for name in reversed(source.namelist()):
            target.writestr(name, source.read(name))
    repacked_inventory = layout_module.inventory_ooxml_images(repacked)
    assert repacked_inventory == inventory
    assert layout_module.image_inventory_sha256(repacked_inventory) == first_hash


def test_office_layout_blocks_compression_bomb_before_provider_probe(
    tmp_path: Path,
) -> None:
    docx = tmp_path / "bomb.docx"
    _docx(docx)
    _rewrite_zip_member(docx, "word/document.xml", b"A" * 1_000_000)
    job = _job(tmp_path)

    receipt = run_office_image_layout(_request(docx, (job,)))

    assert receipt.status is LayoutStatus.BLOCKED
    assert LayoutFailureCode.IMAGE_INVENTORY_INVALID in {
        item.code for item in receipt.failures
    }
    assert "zip_compression_ratio_exceeded" in " ".join(
        item.message for item in receipt.failures
    )


def test_office_layout_rejects_casefold_duplicate_zip_members(
    tmp_path: Path,
) -> None:
    docx = tmp_path / "duplicate.docx"
    _docx(docx)
    with ZipFile(docx, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("WORD/DOCUMENT.XML", b"<document/>")
    job = _job(tmp_path)

    receipt = run_office_image_layout(_request(docx, (job,)))

    assert receipt.status is LayoutStatus.BLOCKED
    assert "zip_duplicate_member" in " ".join(
        item.message for item in receipt.failures
    )


def test_office_layout_enforces_member_and_xml_limits(
    tmp_path: Path,
    monkeypatch,
) -> None:
    docx = tmp_path / "bounded.docx"
    _docx(docx)
    job = _job(tmp_path)
    monkeypatch.setattr(
        coordinator_module,
        "OFFICE_DOCX_PACKAGE_LIMITS",
        safe_docx_package.DocxPackageLimits(max_members=5),
    )

    member_receipt = run_office_image_layout(_request(docx, (job,)))

    assert member_receipt.status is LayoutStatus.BLOCKED
    assert "zip_member_limit_exceeded" in " ".join(
        item.message for item in member_receipt.failures
    )

    monkeypatch.setattr(
        coordinator_module,
        "OFFICE_DOCX_PACKAGE_LIMITS",
        safe_docx_package.DocxPackageLimits(max_xml_nodes=20),
    )
    xml_receipt = run_office_image_layout(_request(docx, (job,)))

    assert xml_receipt.status is LayoutStatus.BLOCKED
    assert "xml_node_limit_exceeded" in " ".join(
        item.message for item in xml_receipt.failures
    )


def test_request_builder_rejects_oversized_docx_without_read_bytes(
    tmp_path: Path,
    monkeypatch,
) -> None:
    docx = tmp_path / "oversized.docx"
    _docx(docx)
    job = _job(tmp_path)
    monkeypatch.setattr(
        coordinator_module,
        "OFFICE_DOCX_PACKAGE_LIMITS",
        safe_docx_package.DocxPackageLimits(max_source_bytes=32),
    )

    with pytest.raises(safe_docx_package.DocxPackageError) as raised:
        layout_module.build_office_image_layout_request(
            docx,
            (job.plan,),
            {job.job_id: job.prepared_image},
            provider=OfficeImageProvider.WORD,
        )

    assert raised.value.code == "source_too_large"


def test_drawing_hash_normalizes_wp_extent_but_keeps_xfrm_extent_exact(
    tmp_path: Path,
) -> None:
    image = tmp_path / "wps-defaults.png"
    _png(image, (320, 200))
    docx = tmp_path / "wps-defaults.docx"
    _docx_with_existing_images(docx, image)
    package = safe_docx_package.SafeDocxPackage.open_path(docx)
    names = set(package.part_names)
    root = package.parse_xml("word/document.xml")
    relationships = inventory_module._image_relationships_for_part(
        package,
        names,
        "word/document.xml",
    )
    inline = next(root.iter(inventory_module._WP_INLINE))
    baseline = inventory_module._drawing_semantic_sha256(inline, relationships)

    inline.set("distT", "0")
    inline.set("distB", "0")
    inline.set("distL", "114300")
    inline.set("distR", "114300")
    effect_extent = inventory_module.etree.Element(
        f"{{{inventory_module._WP_NS}}}effectExtent",
        l="0",
        t="0",
        r="5715",
        b="7620",
    )
    inline.insert(1, effect_extent)
    c_nv_pr = next(
        inline.iter(f"{{{inventory_module._PIC_NS}}}cNvPr")
    )
    c_nv_pr.set("id", "91")
    c_nv_pr.set("name", "Picture 1")
    c_nv_pic_pr = next(
        inline.iter(f"{{{inventory_module._PIC_NS}}}cNvPicPr")
    )
    c_nv_pic_pr.append(
        inventory_module.etree.Element(
            f"{{{inventory_module._A_NS}}}picLocks",
            noChangeAspect="1",
        )
    )

    assert inventory_module._drawing_semantic_sha256(inline, relationships) == baseline
    extent = next(inline.iter(f"{{{inventory_module._WP_NS}}}extent"))
    extent.set("cx", str(int(extent.get("cx")) + 1))
    assert inventory_module._drawing_semantic_sha256(inline, relationships) == baseline
    xfrm_extent = next(inline.iter(f"{{{inventory_module._A_NS}}}ext"))
    xfrm_extent.set("cx", str(int(xfrm_extent.get("cx")) + 1))
    assert inventory_module._drawing_semantic_sha256(inline, relationships) != baseline


def test_image_inventory_verification_requires_exact_existing_and_job_owned_delta(
    tmp_path: Path,
) -> None:
    job = _job(tmp_path)
    before = (_image_identity(),)
    shape_id = contract_module.job_owned_shape_id(job.job_id)
    inserted = _image_identity(
        doc_pr_id="8",
        doc_pr_name="Picture 8",
        alternative_text=shape_id,
        title=job.job_id,
        relationship_id="rId8",
        relationship_target="media/image8.png",
        media_part_name="word/media/image8.png",
        media_sha256=job.prepared_image.output_sha256,
    )
    job_receipt = replace(
        _job_receipt(),
        shape_id=shape_id,
        prepared_image_sha256=job.prepared_image.output_sha256,
    )

    verified = inventory_module.verify_ooxml_image_invariants(
        before,
        before + (inserted,),
        (job,),
        (job_receipt,),
    )
    assert verified.preexisting_unchanged is True
    assert verified.inserted_images_job_owned is True
    assert verified.inserted_images_visible is True
    assert verified.inserted == (inserted,)

    provider_renumbered = replace(
        before[0],
        doc_pr_id="19033724",
        relationship_id="rId42",
    )
    renumbered = inventory_module.verify_ooxml_image_invariants(
        before,
        (provider_renumbered, inserted),
        (job,),
        (job_receipt,),
    )
    assert renumbered.preexisting_unchanged is True
    assert renumbered.inserted_images_job_owned is True
    assert renumbered.inserted_images_visible is True
    assert renumbered.inserted == (inserted,)

    wps_quantized = replace(
        before[0],
        layout_extent_cx_emu=1_079_500,
        layout_extent_cy_emu=359_410,
    )
    tolerated = inventory_module.verify_ooxml_image_invariants(
        before,
        (wps_quantized, inserted),
        (job,),
        (job_receipt,),
    )
    assert tolerated.preexisting_unchanged is True
    assert tolerated.inserted_images_job_owned is True
    assert layout_module.image_preservation_sha256(before) == (
        layout_module.image_preservation_sha256((wps_quantized,))
    )

    over_tolerance = replace(
        before[0],
        layout_extent_cx_emu=(
            before[0].layout_extent_cx_emu
            - layout_module.PREEXISTING_WP_EXTENT_TOLERANCE_EMU
            - 1
        ),
    )
    rejected_extent = inventory_module.verify_ooxml_image_invariants(
        before,
        (over_tolerance, inserted),
        (job,),
        (job_receipt,),
    )
    assert rejected_extent.preexisting_unchanged is False
    assert rejected_extent.inserted_images_job_owned is False

    existing_drift = replace(before[0], doc_pr_name="Office rewrote identity")
    drifted = inventory_module.verify_ooxml_image_invariants(
        before,
        (existing_drift, inserted),
        (job,),
        (job_receipt,),
    )
    assert drifted.preexisting_unchanged is False
    assert drifted.inserted_images_job_owned is False

    unowned = replace(inserted, alternative_text="not-this-transaction")
    rejected = inventory_module.verify_ooxml_image_invariants(
        before,
        before + (unowned,),
        (job,),
        (job_receipt,),
    )
    assert rejected.preexisting_unchanged is True
    assert rejected.inserted_images_job_owned is False
    assert rejected.inserted_images_visible is False

    hidden_inserted = replace(inserted, run_visibility="vanish")
    hidden = inventory_module.verify_ooxml_image_invariants(
        before,
        before + (hidden_inserted,),
        (job,),
        (job_receipt,),
    )
    assert hidden.preexisting_unchanged is True
    assert hidden.inserted_images_job_owned is True
    assert hidden.inserted_images_visible is False

    hidden_existing = replace(before[0], run_visibility="vanish")
    unhidden_existing = inventory_module.verify_ooxml_image_invariants(
        (hidden_existing,),
        before + (inserted,),
        (job,),
        (job_receipt,),
    )
    assert unhidden_existing.preexisting_unchanged is False


def test_preflight_rejects_page_break_policy_and_prepared_image_drift(
    tmp_path: Path,
) -> None:
    docx = tmp_path / "preflight.docx"
    _docx(docx)
    job = _job(tmp_path, mode=ImagePlacementMode.FIXED_BOX)
    unsafe_placement = replace(job.plan.placement, allow_page_break=True)
    unsafe_job = replace(job, plan=replace(job.plan, placement=unsafe_placement))

    failures, _ = preflight_module.parent_preflight(_request(docx, (unsafe_job,)))
    assert LayoutFailureCode.INVALID_REQUEST in {item.code for item in failures}

    Path(job.prepared_image.output_path).write_bytes(b"prepared-image-drift")
    failures, _ = preflight_module.parent_preflight(_request(docx, (job,)))
    assert LayoutFailureCode.PREPARED_IMAGE_IDENTITY_MISMATCH in {
        item.code for item in failures
    }


def test_multi_image_strict_is_blocked_but_flow_requires_strict_sequence(
    tmp_path: Path,
) -> None:
    docx = tmp_path / "multi.docx"
    _docx(docx)
    strict_jobs = (
        _job(tmp_path, job_id="strict-1", sequence=0),
        _job(tmp_path, job_id="strict-2", sequence=1),
    )
    failures, _ = preflight_module.parent_preflight(_request(docx, strict_jobs))
    assert LayoutFailureCode.DUPLICATE_STRICT_ANCHOR in {item.code for item in failures}

    flow_jobs = (
        _job(
            tmp_path,
            job_id="flow-2",
            mode=ImagePlacementMode.FIXED_BOX_FLOW,
            sequence=2,
        ),
        _job(
            tmp_path,
            job_id="flow-1",
            mode=ImagePlacementMode.FIXED_BOX_FLOW,
            sequence=1,
        ),
    )
    failures, _ = preflight_module.parent_preflight(_request(docx, flow_jobs))
    assert LayoutFailureCode.INVALID_FLOW_SEQUENCE in {item.code for item in failures}

    valid_flow_jobs = tuple(reversed(flow_jobs))
    failures, _ = preflight_module.parent_preflight(_request(docx, valid_flow_jobs))
    assert LayoutFailureCode.INVALID_FLOW_SEQUENCE not in {item.code for item in failures}


@pytest.mark.parametrize(
    "mode",
    (
        ImagePlacementMode.NATURAL_SIZE,
        ImagePlacementMode.FIT_CONTAINER_FLOW,
    ),
)
def test_global_multi_image_flow_modes_share_anchor(
    tmp_path: Path,
    mode: ImagePlacementMode,
) -> None:
    docx = tmp_path / f"multi-{mode.value}.docx"
    _docx(docx)
    jobs = (
        _job(tmp_path, job_id=f"{mode.value}-1", mode=mode, sequence=0),
        _job(tmp_path, job_id=f"{mode.value}-2", mode=mode, sequence=1),
    )

    failures, _ = preflight_module.parent_preflight(_request(docx, jobs))
    codes = {item.code for item in failures}

    assert LayoutFailureCode.DUPLICATE_STRICT_ANCHOR not in codes
    assert LayoutFailureCode.INVALID_FLOW_SEQUENCE not in codes


def test_request_and_receipt_roundtrip_and_identity_excludes_elapsed(tmp_path: Path) -> None:
    docx = tmp_path / "roundtrip.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    assert OfficeImageLayoutRequest.from_dict(request.to_dict()) == request

    round_receipt = StabilizationRoundReceipt(
        1,
        2,
        0,
        (("job-1", 1, 1, 1),),
        (),
        (),
        "",
        "c" * 64,
        True,
    )
    existing_inventory = (_image_identity(),)
    inserted_inventory = (
        _image_identity(
            doc_pr_id="8",
            relationship_id="rId8",
            relationship_target="media/image8.png",
            media_part_name="word/media/image8.png",
            media_sha256="e" * 64,
        ),
    )
    shadow_inventory = existing_inventory + inserted_inventory
    receipt = OfficeImageLayoutReceipt(
        schema_version=1,
        transaction_id="tx-test",
        provider=OfficeImageProvider.WORD,
        adapter_name="word_adapter_v1",
        prog_id="Word.Application",
        status=LayoutStatus.SUCCESS,
        application_pid=1234,
        source_docx_path=str(docx),
        source_docx_sha256_before=_sha(docx),
        source_docx_sha256_after=_sha(docx),
        shadow_path=str(tmp_path / ".shadow.docx"),
        shadow_sha256="d" * 64,
        shadow_retained=True,
        source_images_unchanged=True,
        prepared_images_unchanged=True,
        document_open_count=1,
        document_save_count=1,
        repaginate_count=3,
        pdf_export_count=0,
        field_update_rounds=3,
        page_break_count_before=0,
        page_break_count_after=0,
        jobs=(_job_receipt(),),
        stabilization_rounds=(round_receipt,),
        warnings=(),
        failures=(),
        elapsed_ms=10,
        source_image_inventory_sha256=layout_module.image_inventory_sha256(
            existing_inventory
        ),
        shadow_image_inventory_sha256=layout_module.image_inventory_sha256(
            shadow_inventory
        ),
        preexisting_image_semantic_sha256_before=(
            layout_module.image_preservation_sha256(existing_inventory)
        ),
        preexisting_image_semantic_sha256_after=(
            layout_module.image_preservation_sha256(existing_inventory)
        ),
        preexisting_images_unchanged=True,
        inserted_images_job_owned=True,
        inserted_images_visible=True,
        sentinel_texts_hidden=True,
        source_image_inventory=existing_inventory,
        shadow_image_inventory=shadow_inventory,
        inserted_image_inventory=inserted_inventory,
    )
    restored = OfficeImageLayoutReceipt.from_dict(receipt.to_dict())
    assert restored == receipt
    assert replace(receipt, elapsed_ms=999).identity_sha256 == receipt.identity_sha256
    assert receipt.to_dict()["identity_sha256"] == receipt.identity_sha256

    legacy_v3_payload = json.loads(json.dumps(receipt.to_dict()))
    legacy_v3_payload["schema_version"] = 3
    for job_payload in legacy_v3_payload["jobs"]:
        job_payload.pop("initial_actual_width_pt", None)
        job_payload.pop("initial_actual_height_pt", None)
        job_payload.pop("initial_proportion_error", None)
        job_payload.pop("dimension_candidate_count", None)
        job_payload.pop("selected_width_offset_pt", None)
    legacy_v3 = OfficeImageLayoutReceipt.from_dict(legacy_v3_payload)
    assert legacy_v3.schema_version == 3
    assert legacy_v3.jobs[0].initial_actual_width_pt is None
    assert legacy_v3.jobs[0].dimension_candidate_count == 1
    assert legacy_v3.jobs[0].selected_width_offset_pt == 0

    legacy_v2_payload = json.loads(json.dumps(legacy_v3_payload))
    legacy_v2_payload["schema_version"] = 2
    for key in (
        "source_image_inventory",
        "shadow_image_inventory",
        "inserted_image_inventory",
    ):
        for item in legacy_v2_payload[key]:
            item.pop("layout_extent_cx_emu", None)
            item.pop("layout_extent_cy_emu", None)
    legacy_v2 = OfficeImageLayoutReceipt.from_dict(legacy_v2_payload)
    assert legacy_v2.schema_version == 2
    assert all(
        item.layout_extent_cx_emu == item.layout_extent_cy_emu == 0
        for item in legacy_v2.source_image_inventory
    )


class _FakeContent:
    Text = "stable text"


class _FakeDocument:
    def __init__(self) -> None:
        self.Content = _FakeContent()
        self.repaginate_calls = 0

    def Repaginate(self) -> None:
        self.repaginate_calls += 1

    def ComputeStatistics(self, _kind: int) -> int:
        return 4


def test_stabilization_is_global_and_bounded_not_per_image(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    document = _FakeDocument()
    request_job = _job(tmp_path)
    state = worker_module._ChildJobState(
        request_job,
        object(),
        _job_receipt(),
        "a" * 64,
        "b" * 64,
    )
    field_calls: list[int] = []
    verify_calls: list[str] = []
    monkeypatch.setattr(
        worker_module,
        "_update_fields_and_toc",
        lambda _document: field_calls.append(1) or 0,
    )

    def verify(_document, current, _shapes_by_id):
        verify_calls.append(current.receipt.job_id)
        return current.receipt, False

    monkeypatch.setattr(worker_module, "_index_inline_shapes", lambda _document: {})
    monkeypatch.setattr(worker_module, "_verify_job", verify)
    request = OfficeImageLayoutRequest(
        "tx",
        "x.docx",
        "a" * 64,
        1,
        OfficeImageProvider.WORD,
        (request_job,),
        max_stabilization_rounds=3,
    )

    rounds, stable, repaginates, field_rounds = worker_module._stabilize(
        document,
        [state],
        request,
    )

    assert stable is True
    assert len(rounds) == repaginates == field_rounds == 2
    assert document.repaginate_calls == 2
    assert verify_calls == ["job-1", "job-1"]
    assert len(field_calls) == 2


def test_stabilization_indexes_inline_shapes_once_per_round(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    class Shape:
        def __init__(self, alternative_text: str) -> None:
            self.AlternativeText = alternative_text

    class CountingInlineShapes:
        def __init__(self, shapes) -> None:
            self.shapes = tuple(shapes)
            self.Count = len(self.shapes)
            self.item_calls = 0

        def Item(self, index: int):
            self.item_calls += 1
            return self.shapes[index - 1]

    jobs = tuple(
        _job(
            tmp_path,
            job_id=f"indexed-{index}",
            mode=ImagePlacementMode.FIXED_BOX_FLOW,
            sequence=index,
        )
        for index in range(3)
    )
    states = [
        worker_module._ChildJobState(
            job,
            object(),
            _job_receipt(job.job_id, claimed=False),
            "",
            "",
        )
        for job in jobs
    ]
    document = _FakeDocument()
    document.InlineShapes = CountingInlineShapes(
        Shape(state.receipt.shape_id) for state in states
    )
    index_ids_seen: list[tuple[str, ...]] = []
    original_verify_job = worker_module._verify_job
    monkeypatch.setattr(worker_module, "_update_fields_and_toc", lambda _document: 0)

    def verify(_document, current, shapes_by_id):
        index_ids_seen.append(tuple(sorted(shapes_by_id)))
        assert shapes_by_id[current.receipt.shape_id]
        return current.receipt, False

    monkeypatch.setattr(worker_module, "_verify_job", verify)
    request = OfficeImageLayoutRequest(
        "tx-indexed",
        "x.docx",
        "a" * 64,
        1,
        OfficeImageProvider.WORD,
        jobs,
        max_stabilization_rounds=3,
    )

    rounds, stable, _repaginates, _field_rounds = worker_module._stabilize(
        document,
        states,
        request,
    )

    assert stable is True
    assert len(rounds) == 2
    assert document.InlineShapes.item_calls == len(states) * len(rounds)
    assert len(index_ids_seen) == len(states) * len(rounds)

    with pytest.raises(lifecycle_module.ChildBlocked) as duplicate:
        original_verify_job(
            object(),
            states[0],
            {states[0].receipt.shape_id: (object(), object())},
        )
    assert duplicate.value.failure.code is LayoutFailureCode.SHAPE_OWNERSHIP_INVALID


def test_persistent_violation_stops_at_configured_round_limit(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    document = _FakeDocument()
    request_job = _job(tmp_path)
    receipt = replace(
        _job_receipt(),
        final_guard_page=1,
        final_anchor_page=2,
        final_image_page=2,
        boundary_ok=False,
    )
    state = worker_module._ChildJobState(
        request_job,
        object(),
        receipt,
        "a" * 64,
        "b" * 64,
    )
    shrinks: list[str] = []
    monkeypatch.setattr(worker_module, "_update_fields_and_toc", lambda _document: 0)
    monkeypatch.setattr(worker_module, "_index_inline_shapes", lambda _document: {})
    monkeypatch.setattr(
        worker_module,
        "_verify_job",
        lambda _document, current, _shapes_by_id: (current.receipt, True),
    )
    monkeypatch.setattr(
        worker_module,
        "_shrink_violating_job",
        lambda current, _factor, _threshold: shrinks.append(current.receipt.job_id),
    )
    request = OfficeImageLayoutRequest(
        "tx",
        "x.docx",
        "a" * 64,
        1,
        OfficeImageProvider.WORD,
        (request_job,),
        max_stabilization_rounds=2,
    )

    rounds, stable, repaginates, _field_rounds = worker_module._stabilize(
        document,
        [state],
        request,
    )

    assert stable is False
    assert len(rounds) == repaginates == 2
    assert shrinks == ["job-1"]
    assert rounds[0].revalidate_from_job_id == "job-1"


def test_failed_child_cleans_controlled_shadow_and_preserves_all_inputs(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "source.docx"
    _docx(docx)
    job = _job(tmp_path)
    request = _request(docx, (job,))
    source_before = docx.read_bytes()
    image_before = Path(job.plan.image_ref.source_path).read_bytes()
    prepared_before = Path(job.prepared_image.output_path).read_bytes()
    monkeypatch.setattr(coordinator_module, "_registration_status", lambda _prog: (True, "ok"))

    def blocked(_request, spec, _shadow):
        return contract_module.build_layout_receipt(
            _request,
            spec,
            LayoutStatus.BLOCKED,
            failures=(LayoutFailure(LayoutFailureCode.OFFICE_ERROR, "fake"),),
        )

    monkeypatch.setattr(coordinator_module, "_run_layout_child", blocked)

    receipt = run_office_image_layout(request)

    assert receipt.status is LayoutStatus.BLOCKED
    assert receipt.shadow_retained is False
    assert not list(tmp_path.glob(".lark-layout-*.docx"))
    assert docx.read_bytes() == source_before
    assert Path(job.plan.image_ref.source_path).read_bytes() == image_before
    assert Path(job.prepared_image.output_path).read_bytes() == prepared_before


def test_unexpected_child_runner_error_is_structured_and_cleans_shadow(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "unexpected-child-error.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    monkeypatch.setattr(coordinator_module, "_registration_status", lambda _prog: (True, "ok"))
    monkeypatch.setattr(
        coordinator_module,
        "_run_layout_child",
        lambda *_args: (_ for _ in ()).throw(RuntimeError("broker launch exploded")),
    )

    receipt = run_office_image_layout(request)

    assert receipt.status is LayoutStatus.CHILD_FAILED
    assert receipt.failures[0].code is LayoutFailureCode.CHILD_PROTOCOL_ERROR
    assert "RuntimeError: broker launch exploded" in receipt.failures[0].message
    assert receipt.shadow_retained is False
    assert receipt.shadow_path == ""
    assert not list(tmp_path.glob(".lark-layout-*.docx"))


def test_parent_finalization_error_is_structured_and_still_cleans_shadow(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "finalization-error.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    monkeypatch.setattr(coordinator_module, "_registration_status", lambda _prog: (True, "ok"))
    monkeypatch.setattr(
        coordinator_module,
        "_run_layout_child",
        lambda child_request, spec, _shadow: contract_module.build_layout_receipt(
            child_request,
            spec,
            LayoutStatus.BLOCKED,
            failures=(LayoutFailure(LayoutFailureCode.OFFICE_ERROR, "child blocked"),),
        ),
    )
    monkeypatch.setattr(
        coordinator_module,
        "_finalize_layout_receipt",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("verify failed")),
    )

    receipt = run_office_image_layout(request)

    assert receipt.status is LayoutStatus.CHILD_FAILED
    assert [failure.code for failure in receipt.failures] == [
        LayoutFailureCode.OFFICE_ERROR,
        LayoutFailureCode.CHILD_PROTOCOL_ERROR,
    ]
    assert "parent finalization failed: RuntimeError: verify failed" in (
        receipt.failures[-1].message
    )
    assert receipt.shadow_retained is False
    assert not list(tmp_path.glob(".lark-layout-*.docx"))


def test_phase_k_blocks_and_cleans_shadow_when_office_mutates_existing_image(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    existing = tmp_path / "existing.png"
    _png(existing, (240, 160))
    docx = tmp_path / "phase-k.docx"
    _docx_with_existing_images(docx, existing)
    job = _job(tmp_path)
    request = _request(docx, (job,))
    source_before = docx.read_bytes()
    monkeypatch.setattr(coordinator_module, "_registration_status", lambda _prog: (True, "ok"))

    def fake_saved_shadow(_request, spec, shadow):
        inventory = layout_module.inventory_ooxml_images(shadow)
        _rewrite_zip_member(shadow, inventory[0].media_part_name, b"mutated-by-office")
        owned_receipt = replace(
            _job_receipt(),
            shape_id=contract_module.job_owned_shape_id(job.job_id),
            prepared_image_sha256=job.prepared_image.output_sha256,
        )
        return contract_module.build_layout_receipt(
            _request,
            spec,
            LayoutStatus.SUCCESS,
            open_count=1,
            save_count=1,
            jobs=(owned_receipt,),
        )

    monkeypatch.setattr(coordinator_module, "_run_layout_child", fake_saved_shadow)

    receipt = run_office_image_layout(request)

    assert receipt.status is LayoutStatus.BLOCKED
    assert receipt.shadow_retained is False
    assert receipt.preexisting_images_unchanged is False
    assert receipt.inserted_images_job_owned is False
    assert receipt.source_image_inventory_sha256
    assert receipt.shadow_image_inventory_sha256
    assert (
        receipt.preexisting_image_semantic_sha256_before
        != receipt.preexisting_image_semantic_sha256_after
    )
    assert receipt.source_image_inventory
    assert receipt.shadow_image_inventory
    assert LayoutFailureCode.PREEXISTING_IMAGE_MUTATED in {
        item.code for item in receipt.failures
    }
    assert LayoutFailureCode.UNOWNED_IMAGE_ADDED in {
        item.code for item in receipt.failures
    }
    assert not list(tmp_path.glob(".lark-layout-*.docx"))
    assert docx.read_bytes() == source_before


def test_phase_k_retains_only_exact_job_owned_image_delta(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    existing = tmp_path / "existing-success.png"
    _png(existing, (240, 160))
    docx = tmp_path / "phase-k-success.docx"
    _docx_with_existing_images(docx, existing)
    job = _job(tmp_path)
    request = _request(docx, (job,))
    source_before = docx.read_bytes()
    monkeypatch.setattr(coordinator_module, "_registration_status", lambda _prog: (True, "ok"))

    def fake_saved_shadow(_request, spec, shadow):
        document = Document(shadow)
        shape = document.add_paragraph().add_run().add_picture(
            job.prepared_image.output_path
        )
        shape_id = contract_module.job_owned_shape_id(job.job_id)
        shape._inline.docPr.set("descr", shape_id)
        shape._inline.docPr.set("title", job.job_id)
        document.save(shadow)
        owned_receipt = replace(
            _job_receipt(),
            shape_id=shape_id,
            prepared_image_sha256=job.prepared_image.output_sha256,
        )
        return contract_module.build_layout_receipt(
            _request,
            spec,
            LayoutStatus.SUCCESS,
            open_count=1,
            save_count=1,
            jobs=(owned_receipt,),
        )

    monkeypatch.setattr(coordinator_module, "_run_layout_child", fake_saved_shadow)

    receipt = run_office_image_layout(request)

    assert receipt.succeeded is True
    assert receipt.shadow_retained is True
    assert receipt.preexisting_images_unchanged is True
    assert receipt.inserted_images_job_owned is True
    assert receipt.inserted_images_visible is True
    assert receipt.sentinel_texts_hidden is True
    assert (
        receipt.preexisting_image_semantic_sha256_before
        == receipt.preexisting_image_semantic_sha256_after
    )
    assert len(receipt.source_image_inventory) == 2
    assert len(receipt.shadow_image_inventory) == 3
    assert len(receipt.inserted_image_inventory) == 1
    assert receipt.inserted_image_inventory[0].alternative_text == (
        contract_module.job_owned_shape_id(job.job_id)
    )
    assert receipt.inserted_image_inventory[0].title == job.job_id
    assert docx.read_bytes() == source_before
    Path(receipt.shadow_path).unlink()


def test_phase_k_blocks_job_owned_image_when_parent_run_is_hidden(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "hidden-job-image.docx"
    _docx(docx)
    job = _job(tmp_path)
    request = _request(docx, (job,))
    monkeypatch.setattr(coordinator_module, "_registration_status", lambda _prog: (True, "ok"))

    def fake_hidden_shadow(_request, spec, shadow):
        document = Document(shadow)
        shape = document.add_paragraph().add_run().add_picture(
            job.prepared_image.output_path
        )
        shape_id = contract_module.job_owned_shape_id(job.job_id)
        shape._inline.docPr.set("descr", shape_id)
        shape._inline.docPr.set("title", job.job_id)
        image_run = shape._inline.getparent().getparent()
        rpr = image_run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            image_run.insert(0, rpr)
        rpr.append(OxmlElement("w:vanish"))
        document.save(shadow)
        owned_receipt = replace(
            _job_receipt(),
            shape_id=shape_id,
            prepared_image_sha256=job.prepared_image.output_sha256,
        )
        return contract_module.build_layout_receipt(
            _request,
            spec,
            LayoutStatus.SUCCESS,
            open_count=1,
            save_count=1,
            jobs=(owned_receipt,),
        )

    monkeypatch.setattr(coordinator_module, "_run_layout_child", fake_hidden_shadow)

    receipt = run_office_image_layout(request)

    assert receipt.status is LayoutStatus.BLOCKED
    assert receipt.shadow_retained is False
    assert receipt.inserted_images_job_owned is True
    assert receipt.inserted_images_visible is False
    assert receipt.inserted_image_inventory[0].run_visibility == "vanish"
    assert LayoutFailureCode.JOB_IMAGE_HIDDEN in {
        item.code for item in receipt.failures
    }
    assert not list(tmp_path.glob(".lark-layout-*.docx"))


def test_phase_k_blocks_if_image_visibility_change_reveals_hidden_sentinel(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "revealed-sentinel.docx"
    _docx_with_hidden_anchor_sentinel(docx)
    job = _job(tmp_path)
    request = _request(docx, (job,))
    monkeypatch.setattr(coordinator_module, "_registration_status", lambda _prog: (True, "ok"))

    def fake_revealed_shadow(_request, spec, shadow):
        document = Document(shadow)
        shape = document.add_paragraph().add_run().add_picture(
            job.prepared_image.output_path
        )
        shape_id = contract_module.job_owned_shape_id(job.job_id)
        shape._inline.docPr.set("descr", shape_id)
        shape._inline.docPr.set("title", job.job_id)
        for text in document.element.iter(qn("w:t")):
            if text.text != "anchor_1":
                continue
            run = text.getparent()
            rpr = run.find(qn("w:rPr"))
            vanish = None if rpr is None else rpr.find(qn("w:vanish"))
            if vanish is not None:
                rpr.remove(vanish)
        document.save(shadow)
        owned_receipt = replace(
            _job_receipt(),
            shape_id=shape_id,
            prepared_image_sha256=job.prepared_image.output_sha256,
        )
        return contract_module.build_layout_receipt(
            _request,
            spec,
            LayoutStatus.SUCCESS,
            open_count=1,
            save_count=1,
            jobs=(owned_receipt,),
        )

    monkeypatch.setattr(coordinator_module, "_run_layout_child", fake_revealed_shadow)

    receipt = run_office_image_layout(request)

    assert receipt.status is LayoutStatus.BLOCKED
    assert receipt.inserted_images_job_owned is True
    assert receipt.inserted_images_visible is True
    assert receipt.sentinel_texts_hidden is False
    assert LayoutFailureCode.SENTINEL_REVEALED in {
        item.code for item in receipt.failures
    }
    assert not list(tmp_path.glob(".lark-layout-*.docx"))


def test_long_prepared_cache_path_is_staged_to_verified_short_office_copy(
    tmp_path: Path,
) -> None:
    docx = tmp_path / "long-prepared.docx"
    _docx(docx)
    job = _job(tmp_path)
    long_parent = tmp_path
    while coordinator_module._office_path_utf16_units(long_parent / "prepared.png") <= 270:
        long_parent = long_parent / ("cache-segment-" + "x" * 36)
    long_parent.mkdir(parents=True)
    long_prepared = long_parent / "prepared.png"
    original_prepared = Path(job.prepared_image.output_path)
    long_prepared.write_bytes(original_prepared.read_bytes())
    long_job = replace(
        job,
        prepared_image=replace(
            job.prepared_image,
            output_path=str(long_prepared),
            output_sha256=_sha(long_prepared),
        ),
    )
    request = _request(docx, (long_job,))
    ipc_dir = tmp_path / "ipc"
    ipc_dir.mkdir()

    staged = coordinator_module._stage_prepared_images_for_office(request, ipc_dir)
    staged_path = Path(staged.jobs[0].prepared_image.output_path)

    assert coordinator_module._office_path_utf16_units(long_prepared) > 260
    assert staged_path == (ipc_dir / "i0000.png").resolve()
    assert coordinator_module._office_path_utf16_units(staged_path) <= (
        layout_module.OFFICE_COM_PATH_MAX_UTF16_UNITS
    )
    assert _sha(staged_path) == long_job.prepared_image.output_sha256
    assert _sha(long_prepared) == long_job.prepared_image.output_sha256
    assert staged.jobs[0].plan == long_job.plan


def test_over_budget_office_shadow_path_is_structured_block_before_process(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "path-budget.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    overlong = tmp_path
    while coordinator_module._office_path_utf16_units(overlong / "shadow.docx") <= 245:
        overlong = overlong / ("shadow-segment-" + "y" * 40)
    shadow = overlong / "shadow.docx"

    def process_must_not_start(*_args, **_kwargs):
        raise AssertionError("Office broker process must not start")

    monkeypatch.setattr(coordinator_module.subprocess, "Popen", process_must_not_start)

    receipt = coordinator_module._run_layout_child(
        request,
        PROVIDER_SPECS[OfficeImageProvider.WORD],
        shadow,
    )

    assert receipt.status is LayoutStatus.BLOCKED
    assert receipt.document_open_count == 0
    assert receipt.failures[0].code is LayoutFailureCode.PATH_TOO_LONG
    assert "UTF-16 units" in receipt.failures[0].message


def test_child_timeout_targets_only_reported_office_pid(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "timeout.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    shadow = tmp_path / ".timeout.lark-layout-owned.docx"
    shadow.write_bytes(docx.read_bytes())

    class FakeProcess:
        pid = 111
        returncode = None
        calls = 0

        def communicate(self, timeout: float):
            self.calls += 1
            if self.calls == 1:
                raise subprocess.TimeoutExpired(["python"], timeout)
            self.returncode = -9
            return "", ""

        def poll(self):
            return self.returncode

    fake = FakeProcess()
    broker_prepared_paths: list[Path] = []
    monkeypatch.setattr(coordinator_module, "_process_ids_for_names", lambda _names: set())

    def fake_popen(command, *_args, **_kwargs):
        envelope = json.loads(Path(command[-1]).read_text(encoding="utf-8"))
        child_request = OfficeImageLayoutRequest.from_dict(envelope["request"])
        broker_prepared_paths.extend(
            Path(item.prepared_image.output_path) for item in child_request.jobs
        )
        return fake

    monkeypatch.setattr(coordinator_module.subprocess, "Popen", fake_popen)
    monkeypatch.setattr(coordinator_module, "_terminate_child", lambda _process: None)
    monkeypatch.setattr(coordinator_module, "_read_pid_file", lambda _path: 222)
    cleaned: list[int | None] = []

    def cleanup(pid):
        cleaned.append(pid)
        return "owned Office PID required forced cleanup"

    monkeypatch.setattr(coordinator_module, "_ensure_owned_pid_stopped", cleanup)

    receipt = coordinator_module._run_layout_child(
        request,
        PROVIDER_SPECS[OfficeImageProvider.WORD],
        shadow,
    )

    assert receipt.status is LayoutStatus.TIMED_OUT
    assert receipt.application_pid == 222
    assert cleaned == [222]
    assert receipt.failures[0].code is LayoutFailureCode.OFFICE_TIMEOUT
    assert [path.name for path in broker_prepared_paths] == ["i0000.png"]
    assert broker_prepared_paths[0].parent.name.startswith("lk-ol-")
    assert broker_prepared_paths[0] != Path(request.jobs[0].prepared_image.output_path)


def test_child_pipe_error_terminates_broker_and_cleans_only_owned_office_pid(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "pipe-error.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    shadow = tmp_path / ".pipe-error.lark-layout-owned.docx"
    shadow.write_bytes(docx.read_bytes())

    class FakeProcess:
        pid = 111
        returncode = None
        calls = 0

        def communicate(self, timeout: float):
            self.calls += 1
            if self.calls == 1:
                raise OSError("pipe broke")
            self.returncode = -9
            return "", ""

        def poll(self):
            return self.returncode

    process = FakeProcess()
    terminated: list[int] = []
    cleaned: list[int | None] = []
    monkeypatch.setattr(coordinator_module, "_process_ids_for_names", lambda _names: set())
    monkeypatch.setattr(coordinator_module.subprocess, "Popen", lambda *_a, **_k: process)
    monkeypatch.setattr(
        coordinator_module,
        "_terminate_child",
        lambda child: terminated.append(child.pid),
    )
    monkeypatch.setattr(coordinator_module, "_read_pid_file", lambda _path: 222)
    monkeypatch.setattr(
        coordinator_module,
        "_ensure_owned_pid_stopped",
        lambda pid: cleaned.append(pid) or "owned Office PID required forced cleanup",
    )

    receipt = coordinator_module._run_layout_child(
        request,
        PROVIDER_SPECS[OfficeImageProvider.WORD],
        shadow,
    )

    assert receipt.status is LayoutStatus.CHILD_FAILED
    assert receipt.application_pid == 222
    assert process.calls == 2
    assert terminated == [111]
    assert cleaned == [222]
    assert receipt.failures[0].code is LayoutFailureCode.CHILD_PROTOCOL_ERROR
    assert "pipe broke" in receipt.failures[0].message
    assert receipt.warnings[0].code is (
        contract_module.LayoutWarningCode.OWNED_PROCESS_FORCED_CLEANUP
    )


def test_child_receipt_pid_recovers_cleanup_when_pid_record_is_missing(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "receipt-pid.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    shadow = tmp_path / ".receipt-pid.lark-layout-owned.docx"
    shadow.write_bytes(docx.read_bytes())

    class FakeProcess:
        pid = 111
        returncode = 0

        def communicate(self, timeout: float):
            return "", ""

        def poll(self):
            return self.returncode

    def fake_popen(command, *_args, **_kwargs):
        envelope = json.loads(Path(command[-1]).read_text(encoding="utf-8"))
        receipt = contract_module.build_layout_receipt(
            request,
            PROVIDER_SPECS[OfficeImageProvider.WORD],
            LayoutStatus.BLOCKED,
            application_pid=222,
            failures=(LayoutFailure(LayoutFailureCode.OFFICE_ERROR, "pid write failed"),),
        )
        Path(envelope["result_path"]).write_text(
            json.dumps(receipt.to_dict(), ensure_ascii=False),
            encoding="utf-8",
        )
        return FakeProcess()

    cleaned: list[int | None] = []
    monkeypatch.setattr(coordinator_module, "_process_ids_for_names", lambda _names: set())
    monkeypatch.setattr(coordinator_module.subprocess, "Popen", fake_popen)
    monkeypatch.setattr(
        coordinator_module,
        "_ensure_owned_pid_stopped",
        lambda pid: cleaned.append(pid) or (
            "owned Office PID required forced cleanup" if pid is not None else ""
        ),
    )

    receipt = coordinator_module._run_layout_child(
        request,
        PROVIDER_SPECS[OfficeImageProvider.WORD],
        shadow,
    )

    assert receipt.status is LayoutStatus.BLOCKED
    assert receipt.application_pid == 222
    assert cleaned == [222]
    assert receipt.warnings[-1].code is (
        contract_module.LayoutWarningCode.OWNED_PROCESS_FORCED_CLEANUP
    )


def test_successful_child_receipt_without_owned_pid_fails_protocol(
    tmp_path: Path,
) -> None:
    docx = tmp_path / "missing-success-pid.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    result_path = tmp_path / "result.json"
    child_receipt = contract_module.build_layout_receipt(
        request,
        PROVIDER_SPECS[OfficeImageProvider.WORD],
        LayoutStatus.SUCCESS,
    )
    result_path.write_text(
        json.dumps(child_receipt.to_dict(), ensure_ascii=False),
        encoding="utf-8",
    )

    receipt = coordinator_module._load_child_receipt(
        request,
        PROVIDER_SPECS[OfficeImageProvider.WORD],
        result_path,
        owned_pid=None,
        before_pids=set(),
        cleanup_warnings=(),
    )

    assert receipt.status is LayoutStatus.CHILD_FAILED
    assert receipt.failures[0].code is LayoutFailureCode.CHILD_PROTOCOL_ERROR
    assert "did not report an owned Office PID" in receipt.failures[0].message


def test_child_execute_releases_partial_office_session_on_runtime_failure(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    docx = tmp_path / "worker-cleanup.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    events: list[object] = []

    class PythonCom:
        @staticmethod
        def CoInitialize() -> None:
            events.append("coinitialize")

        @staticmethod
        def CoUninitialize() -> None:
            events.append("couninitialize")

    class View:
        Type = 0

    class Document:
        ActiveWindow = type("Window", (), {"View": View()})()

        def Close(self, save: bool) -> None:
            events.append(("close", save))

    class Application:
        Visible = True
        DisplayAlerts = 1

        def Quit(self) -> None:
            events.append("quit")

    document = Document()
    application = Application()

    class Adapter:
        def create_application(self):
            events.append("create")
            return application

        def open_shadow(self, _application, _path):
            events.append("open")
            return document

    monkeypatch.setitem(sys.modules, "pythoncom", PythonCom)
    monkeypatch.setitem(worker_module._ADAPTERS, OfficeImageProvider.WORD, Adapter)
    monkeypatch.setattr(lifecycle_module, "_application_process_id", lambda *_args: 4321)
    monkeypatch.setattr(
        worker_module,
        "_page_break_count",
        lambda _document: (_ for _ in ()).throw(RuntimeError("layout read failed")),
    )
    envelope = {
        "request": request.to_dict(),
        "shadow_path": str(docx),
        "pid_path": str(tmp_path / "office-pid.json"),
        "before_pids": [],
    }

    receipt = worker_module._child_execute(envelope)

    assert receipt.status is LayoutStatus.BLOCKED
    assert receipt.application_pid == 4321
    assert receipt.document_open_count == 1
    assert receipt.document_save_count == 0
    assert receipt.failures[0].code is LayoutFailureCode.OFFICE_ERROR
    assert events == [
        "coinitialize",
        "create",
        "open",
        ("close", False),
        "quit",
        "couninitialize",
    ]


def test_source_contract_has_no_dispatch_fallback_pdf_or_per_image_repaginate() -> None:
    facade_source = inspect.getsource(layout_module)
    source = "\n".join(
        (
            inspect.getsource(coordinator_module),
            inspect.getsource(preflight_module),
            inspect.getsource(lifecycle_module),
            inspect.getsource(worker_module),
        )
    )
    insert_source = inspect.getsource(worker_module._insert_one_job)
    assert "DispatchEx" not in facade_source
    assert "def _insert_one_job" not in facade_source
    assert "win32com.client.Dispatch(" not in source
    assert "DispatchEx" in source
    assert "ExportAsFixedFormat" not in source
    assert "SaveAs" not in source
    assert "Repaginate" not in insert_source
    assert "InsertBreak" not in source
    assert source.count("application.Documents.Open(") == 1
    assert "class _WordOfficeAdapter" in source
    assert "class _WpsOfficeAdapter" in source
    assert 'prog_id = "Word.Application"' in source
    assert 'prog_id = "KWPS.Application"' in source


def test_cli_reports_preflight_failure_without_starting_office(tmp_path: Path) -> None:
    docx = tmp_path / "cli.docx"
    _docx(docx)
    request = _request(docx, (_job(tmp_path),))
    request_path = tmp_path / "request.json"
    request_path.write_text(
        json.dumps(request.to_dict(), ensure_ascii=False),
        encoding="utf-8",
    )
    # Drift after the immutable request snapshot: the CLI must fail in the
    # parent preflight, before registration checks, shadow creation, or COM.
    docx.write_bytes(docx.read_bytes() + b"drift")
    repository_root = Path(layout_module.__file__).resolve().parents[3]
    script = repository_root / "scripts" / "run_office_image_layout.py"

    completed = subprocess.run(
        [sys.executable, str(script), str(request_path), "--compact"],
        cwd=repository_root,
        capture_output=True,
        text=True,
        timeout=10,
        check=False,
    )

    assert completed.returncode == 2
    payload = json.loads(completed.stdout)
    assert payload["status"] == LayoutStatus.BLOCKED.value
    assert payload["document_open_count"] == 0
    assert payload["shadow_retained"] is False
    assert payload["failures"][0]["code"] == (
        LayoutFailureCode.SOURCE_DOCX_IDENTITY_MISMATCH.value
    )
    assert not list(tmp_path.glob(".lark-layout-*.docx"))
