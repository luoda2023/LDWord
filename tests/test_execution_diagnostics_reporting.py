import json
import sys
from dataclasses import dataclass, replace
from pathlib import Path
from types import SimpleNamespace

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.execution_diagnostics import build_execution_diagnostics
from src.config.scene import SceneWorkspace
from src.config.master_library import default_master
from src.config.resolver import resolve_config
from src.config.style_source_report_summary import build_style_source_report_summary
from src.config.template import StyleConfig, TemplateConfig
from src.modules.base import BaseModule, ModuleMeta
from src.pipeline.context import PipelineContext
from src.pipeline.result import PipelineResult
from src.pipeline.runner import Pipeline
from src.pipeline.runner import build_coverage_boundary_report_items
from src.pipeline.tracker import ChangeTracker
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.official_numbering_preservation import (
    inspect_official_numbering_preservation,
)
from src.shared.engine.technical_chapter_inventory import (
    inspect_technical_chapter_inventory,
)
from src.shared.engine.application_section_word_limits import (
    inspect_application_section_word_limits,
)


class RuntimeTraceTableModule(BaseModule):
    meta = ModuleMeta(
        name="table_runtime_trace",
        description="Table runtime trace",
        category="table",
        requires_config=("table",),
        execution_phase="format",
        scope_behavior="region_filtered",
    )

    def apply(self, doc, config, tracker, context):
        return None


class RuntimeTraceStylesModule(BaseModule):
    meta = ModuleMeta(
        name="styles_runtime_trace",
        description="Styles runtime trace",
        category="style",
        requires_config=("styles",),
        execution_phase="format",
        scope_behavior="region_filtered",
    )

    def apply(self, doc, config, tracker, context):
        return None


def _build_result_with_diagnostics() -> PipelineResult:
    tracker = ChangeTracker()
    tracker.record(
        rule_name="equation_table_format",
        target="3 个公式编号",
        section="global",
        change_type="skip",
        before="chapter-aware numbering normalization",
        after="skipped due to missing chapter context",
    )
    tracker.record(
        rule_name="table_format",
        target="pipeline",
        section="global",
        change_type="error",
        success=False,
        failure_reason="table layout explosion",
    )
    return PipelineResult(
        success=True,
        status="partial_success",
        tracker=tracker,
        failed_items=[
            {
                "rule_name": "table_format",
                "target": "pipeline",
                "section": "global",
                "change_type": "error",
                "paragraph_index": -1,
                "reason": "table layout explosion",
            }
        ],
    )


def test_build_execution_diagnostics_collects_skip_and_failure_records():
    result = _build_result_with_diagnostics()

    diagnostics = build_execution_diagnostics(result, summary_limit=2)

    assert diagnostics["count"] == 2
    assert diagnostics["items"][0]["change_type"] == "skip"
    assert diagnostics["items"][0]["reason"] == "skipped due to missing chapter context"
    assert diagnostics["items"][1]["reason"] == "table layout explosion"
    assert "诊断提示（2）" in diagnostics["summary"]


def test_report_writer_keeps_document_family_sections_in_dedicated_module():
    writer_source = (ROOT / "src/report_writer.py").read_text(encoding="utf-8")
    section_source = (
        ROOT / "src/reporting/document_sections.py"
    ).read_text(encoding="utf-8")

    assert "from src.reporting.document_sections import" in writer_source
    for function_name in (
        "_extract_official_numbering_preservation",
        "_format_official_numbering_preservation_markdown",
        "_extract_technical_chapter_inventory",
        "_format_technical_chapter_inventory_markdown",
        "_extract_application_section_word_limits",
        "_format_application_section_word_limits_markdown",
    ):
        assert f"def {function_name}" not in writer_source
        assert f"def {function_name}" in section_source


def test_report_writer_keeps_exam_sections_in_dedicated_module():
    writer_source = (ROOT / "src/report_writer.py").read_text(encoding="utf-8")
    section_source = (ROOT / "src/reporting/exam_sections.py").read_text(
        encoding="utf-8"
    )

    assert "from src.reporting.exam_sections import" in writer_source
    for function_name in (
        "_extract_exam_question_schema",
        "_extract_exam_delivery_runtime",
        "_format_exam_question_schema_markdown",
        "_format_exam_delivery_runtime_markdown",
    ):
        assert f"def {function_name}" not in writer_source
        assert f"def {function_name}" in section_source


def test_report_writer_keeps_material_sections_in_dedicated_module():
    writer_source = (ROOT / "src/report_writer.py").read_text(encoding="utf-8")
    section_source = (ROOT / "src/reporting/material_sections.py").read_text(
        encoding="utf-8"
    )

    assert "from src.reporting.material_sections import" in writer_source
    for function_name in (
        "_extract_material_field_consistency",
        "_format_material_field_consistency_markdown",
        "_extract_object_preflight",
        "_format_object_preflight_markdown",
        "_extract_coverage_boundaries",
        "_format_coverage_boundaries_markdown",
    ):
        assert f"def {function_name}" not in writer_source
        assert f"def {function_name}" in section_source


def test_report_writer_emits_diagnostics_into_json_and_markdown(tmp_path):
    result = _build_result_with_diagnostics()
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"

    write_json_report(
        result,
        input_path=tmp_path / "sample.docx",
        output_path=None,
        report_path=report_json,
        elapsed=1.23,
        modules_enabled=3,
        modules_total=5,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "sample.docx",
        report_path=report_md,
        elapsed=1.23,
        modules_enabled=3,
        modules_total=5,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert report_data["diagnostics"]["count"] == 2
    assert report_data["diagnostics"]["items"][0]["reason"] == "skipped due to missing chapter context"
    assert "## 诊断提示 (2 项)" in markdown
    assert "[equation_table_format] 3 个公式编号: skipped due to missing chapter context" in markdown


def test_report_writer_emits_style_source_summary_into_json_and_markdown(tmp_path):
    result = _build_result_with_diagnostics()
    template = TemplateConfig(name="默认格式")
    template.styles["body"] = StyleConfig(line_spacing_type="exact", line_spacing_pt=20)
    scene = SceneWorkspace(scene_id="custom", template_id="default")
    style_source_summary = build_style_source_report_summary(scene, template)
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"

    write_json_report(
        result,
        input_path=tmp_path / "sample.docx",
        output_path=None,
        report_path=report_json,
        elapsed=1.23,
        modules_enabled=3,
        modules_total=5,
        style_source_summary=style_source_summary,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "sample.docx",
        report_path=report_md,
        elapsed=1.23,
        modules_enabled=3,
        modules_total=5,
        style_source_summary=style_source_summary,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert report_data["style_source"]["template_label"] == "默认格式"
    assert report_data["style_source"] == {
        "template_label": "默认格式",
        "summary": "本次使用模板“默认格式”。",
    }
    assert "## 样式来源" in markdown
    assert "- 摘要: 本次使用模板“默认格式”。" in markdown
    assert "- 模板: 默认格式" in markdown


def test_official_numbering_preservation_detects_preserve_mode_rewrites():
    original = Document()
    original.add_heading("一、总体要求", level=1)
    current = Document()
    current.add_heading("第一章 总体要求", level=1)
    tracker = ChangeTracker()
    tracker.record(
        rule_name="heading_numbering",
        target="1 个标题",
        section="global",
        change_type="format",
        before="(mixed/无编号)",
        after="已添加编号",
    )
    config = SimpleNamespace(
        strict_mode=False,
        module_switches={"heading_numbering": True},
        compliance_profile=SimpleNamespace(
            rule_family="official_document",
            profile_id="official_document",
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="official_document_v1",
            material_schema_ids=[],
        ),
    )

    evidence = inspect_official_numbering_preservation(
        config,
        original_doc=original,
        current_doc=current,
        heading_map={0: 1},
        tracker_records=tracker.get_all(),
    )

    assert evidence.status == "needs_review"
    assert evidence.strategy == "preserve"
    assert evidence.manual_confirmation_required is True
    assert evidence.summary.heading_count == 1
    assert evidence.summary.changed_heading_count == 1
    assert evidence.summary.heading_numbering_changed_count == 1
    assert evidence.changed_headings[0].before == "一、总体要求"
    assert evidence.changed_headings[0].after == "第一章 总体要求"
    assert any(
        issue.kind == "heading_numbering_recorded_in_preserve_mode"
        for issue in evidence.issues
    )


def test_report_writer_emits_official_numbering_preservation_evidence(tmp_path):
    context = PipelineContext(
        official_numbering_preservation={
            "family_id": "official_document",
            "status": "needs_review",
            "strategy": "preserve",
            "rule_family": "official_document",
            "profile_id": "official_document",
            "material_schema_ids": ["official_document_v1"],
            "word_surfaces": ["numbering.xml", "w:pPr/w:numPr"],
            "heading_numbering_enabled": True,
            "manual_confirmation_required": True,
            "summary": {
                "heading_count": 1,
                "changed_heading_count": 1,
                "heading_numbering_record_count": 1,
                "heading_numbering_changed_count": 1,
            },
            "changed_headings": [
                {
                    "paragraph_index": 0,
                    "level": 1,
                    "before": "一、总体要求",
                    "after": "第一章 总体要求",
                }
            ],
            "issues": [
                {
                    "kind": "heading_text_changed_in_preserve_mode",
                    "severity": "warning",
                    "message": "Heading text changed while preservation was requested.",
                }
            ],
        }
    )
    result = PipelineResult(success=True, context=context)
    report_json = tmp_path / "official_numbering.json"
    report_md = tmp_path / "official_numbering.md"

    write_json_report(
        result,
        input_path=tmp_path / "official.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "official.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    evidence = report_data["official_numbering_preservation"]
    assert evidence["status"] == "needs_review"
    assert evidence["strategy"] == "preserve"
    assert evidence["summary"]["changed_heading_count"] == 1
    assert "## 公文编号保留证据" in markdown
    assert "- Strategy: preserve" in markdown
    assert "一、总体要求 -> 第一章 总体要求" in markdown


def test_pipeline_records_official_numbering_preservation_context(tmp_path):
    source = tmp_path / "official.docx"
    doc = Document()
    doc.add_heading("一、总体要求", level=1)
    doc.save(source)

    class _MarkOfficialHeadingMap(BaseModule):
        meta = ModuleMeta(
            name="mark_official_heading_map",
            description="Mark official heading map",
            category="structure",
            provides=("heading_map",),
            execution_phase="semantics",
            scope_behavior="structure_discovery",
        )

        def apply(self, doc, config, tracker, context):
            context.heading_map = {0: 1}

    class _RewriteOfficialHeadingNumber(BaseModule):
        meta = ModuleMeta(
            name="rewrite_official_heading_number",
            description="Rewrite official heading number",
            category="structure",
            consumes=("heading_map",),
            execution_phase="semantics",
            scope_behavior="region_filtered",
        )

        def apply(self, doc, config, tracker, context):
            before = doc.paragraphs[0].text
            doc.paragraphs[0].text = "第一章 总体要求"
            tracker.record(
                rule_name="heading_numbering",
                target="1 个标题",
                section="global",
                change_type="format",
                before=before,
                after=doc.paragraphs[0].text,
            )

    config = SimpleNamespace(
        strict_mode=False,
        module_switches={"heading_numbering": True},
        compliance_profile=SimpleNamespace(
            rule_family="official_document",
            profile_id="official_document",
            object_preflight=SimpleNamespace(enabled=False),
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="",
            material_schema_ids=[],
        ),
        entity_data={},
        delivery_presets=[],
        output=SimpleNamespace(final_docx=False),
    )

    result = Pipeline(
        modules=[_MarkOfficialHeadingMap(), _RewriteOfficialHeadingNumber()],
        config=config,
    ).execute(str(source))

    assert result.success is True
    evidence = result.context.official_numbering_preservation
    assert evidence.status == "needs_review"
    assert evidence.summary.changed_heading_count == 1
    assert evidence.summary.heading_numbering_record_count == 1
    diagnostics = build_execution_diagnostics(result)
    assert diagnostics["count"] == 1
    assert diagnostics["items"][0]["rule_name"] == "official_numbering_preservation"


def test_report_writer_emits_official_document_assembly_evidence(tmp_path):
    context = PipelineContext(
        official_document_assembly={
            "status": "ok",
            "profile_id": "notice",
            "master_id": "official_gbt_standard",
            "material_schema_ids": ["official_document_v1"],
            "docx_path": "dist/official_notice.docx",
            "internal_review_docx_path": "dist/official_notice_internal_review.docx",
            "archive_manifest_path": "dist/official_notice_archive_manifest.json",
            "archive_manifest_markdown_path": "dist/official_notice_archive_manifest.md",
            "review_pdf_path": "dist/official_notice_review.pdf",
            "review_pdf_status": "generated",
            "review_pdf_renderer": "word_com",
            "review_pdf_issue": "",
            "output_paths": {
                "official_docx": "dist/official_notice.docx",
                "internal_review_docx": "dist/official_notice_internal_review.docx",
                "archive_manifest": "dist/official_notice_archive_manifest.json",
                "archive_manifest_md": "dist/official_notice_archive_manifest.md",
                "review_pdf": "dist/official_notice_review.pdf",
            },
            "replaced_placeholders": [
                "official_title",
                "official_body",
                "official_issue_date",
            ],
            "missing_required_fields": [],
            "unresolved_placeholders": [],
        }
    )
    result = PipelineResult(success=True, context=context)
    report_json = tmp_path / "official_assembly.json"
    report_md = tmp_path / "official_assembly.md"

    write_json_report(
        result,
        input_path=tmp_path / "official.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "official.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    evidence = report_data["official_document_assembly"]
    assert evidence["status"] == "ok"
    assert evidence["profile_id"] == "notice"
    assert evidence["master_id"] == "official_gbt_standard"
    assert "## 公文版式装配证据" in markdown
    assert evidence["output_paths"]["internal_review_docx"].endswith(
        "_internal_review.docx"
    )
    assert "Formal DOCX: `dist/official_notice.docx`" in markdown
    assert (
        "Internal review DOCX: `dist/official_notice_internal_review.docx`"
        in markdown
    )
    assert "Archive manifest: `dist/official_notice_archive_manifest.json`" in markdown
    assert evidence["review_pdf_status"] == "generated"
    assert evidence["review_pdf_renderer"] == "word_com"
    assert "Review PDF: `dist/official_notice_review.pdf`" in markdown
    assert "Review PDF status: generated (word_com)" in markdown
    assert "archive_manifest_md" in markdown
    assert "official_title" in markdown


def test_pipeline_records_official_document_assembly_context_and_output(tmp_path):
    source = tmp_path / "official.docx"
    Document().save(source)

    output_dir = tmp_path / "out"
    config = SimpleNamespace(
        strict_mode=False,
        compliance_profile=SimpleNamespace(
            rule_family="official_document",
            profile_id="official_document",
            object_preflight=SimpleNamespace(enabled=False),
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="official_document_v1",
            material_schema_ids=[],
        ),
        entity_data={
            "document_type": "notice",
            "title": "关于开展资料归档检查的通知",
            "body": "请各部门按要求完成自查并提交材料。",
            "organization": "示例市档案局",
            "document_no": "示档发〔2026〕1号",
            "issue_date": "2026年7月9日",
        },
        delivery_presets=[],
        output=SimpleNamespace(final_docx=False),
    )

    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        official_document_type_id="notice",
    ).execute(str(source))

    evidence = result.context.official_document_assembly
    docx_path = output_dir / "official_official.docx"
    internal_review_path = output_dir / "official_official_internal_review.docx"
    archive_manifest_path = output_dir / "official_official_archive_manifest.json"
    archive_manifest_md_path = output_dir / "official_official_archive_manifest.md"
    records = result.tracker.get_by_module("official_document_assembly")

    assert result.success is True
    assert evidence.status == "ok"
    assert evidence.profile_id == "notice"
    assert evidence.master_id == "official_gbt_standard"
    assert evidence.docx_path == docx_path
    assert evidence.internal_review_docx_path == internal_review_path
    assert evidence.archive_manifest_path == archive_manifest_path
    assert evidence.archive_manifest_markdown_path == archive_manifest_md_path
    assert evidence.review_pdf_status == "not_requested"
    assert evidence.review_pdf_path is None
    assert result.output_paths == {
        "official_docx": str(docx_path),
        "internal_review_docx": str(internal_review_path),
        "archive_manifest": str(archive_manifest_path),
        "archive_manifest_md": str(archive_manifest_md_path),
    }
    assert docx_path.is_file()
    assert internal_review_path.is_file()
    assert archive_manifest_path.is_file()
    assert archive_manifest_md_path.is_file()
    assert len(records) == 1
    assert records[0].success is True
    assert "review_pdf=not_requested" in records[0].after


def test_pipeline_uses_explicit_official_master_for_assembly(tmp_path, monkeypatch):
    import src.shared.engine.official_document_assembly as assembly_module

    source = tmp_path / "official.docx"
    Document().save(source)
    default = default_master("official")
    assert default is not None
    selected = replace(default, execution_frozen=True)
    resolver_calls = []
    real_resolver = assembly_module.resolve_official_master_for_contract

    def _resolve_master(contract, *, requested=None):
        resolver_calls.append(requested)
        return real_resolver(contract, requested=requested)

    monkeypatch.setattr(
        assembly_module,
        "resolve_official_master_for_contract",
        _resolve_master,
    )
    config = SimpleNamespace(
        strict_mode=False,
        compliance_profile=SimpleNamespace(
            rule_family="official_document",
            profile_id="official_document",
            object_preflight=SimpleNamespace(enabled=False),
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="official_document_v1",
            material_schema_ids=[],
        ),
        entity_data={
            "document_type": "notice",
            "title": "关于所选公文母版链路的通知",
            "body": "验证运行时使用显式选择的公文母版。",
            "organization": "示例单位",
            "document_no": "示发〔2026〕2号",
            "issue_date": "2026年7月11日",
        },
        delivery_presets=[],
        output=SimpleNamespace(final_docx=False),
    )

    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(tmp_path / "out"),
        official_master=selected,
        official_document_type_id="notice",
    ).execute(str(source))

    assert result.success is True
    assert resolver_calls == [selected]
    assert result.context.official_document_assembly.master_id == selected.master_id


def test_pipeline_generates_opted_in_official_review_pdf(tmp_path, monkeypatch):
    import src.shared.engine.official_document_assembly as assembly_module

    source = tmp_path / "official.docx"
    Document().save(source)

    def fake_renderer(_docx_path, pdf_path):
        Path(pdf_path).write_bytes(b"%PDF-1.4\n% test review pdf\n")

    monkeypatch.setattr(
        assembly_module,
        "_default_review_pdf_renderer_resolver",
        lambda: ("fake_pdf", fake_renderer),
    )
    output_dir = tmp_path / "out"
    config = SimpleNamespace(
        strict_mode=False,
        compliance_profile=SimpleNamespace(
            rule_family="official_document",
            profile_id="official_document",
            object_preflight=SimpleNamespace(enabled=False),
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="official_document_v1",
            material_schema_ids=[],
        ),
        entity_data={
            "document_type": "notice",
            "title": "关于开展资料归档检查的通知",
            "body": "请各部门按要求完成自查并提交材料。",
            "organization": "示例市档案局",
            "document_no": "示档发〔2026〕1号",
            "issue_date": "2026年7月9日",
        },
        delivery_presets=[],
        output=SimpleNamespace(final_docx=False, review_pdf=True),
    )

    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        official_document_type_id="notice",
    ).execute(str(source))

    evidence = result.context.official_document_assembly
    records = result.tracker.get_by_module("official_document_assembly")
    assert result.success is True
    assert evidence.status == "ok"
    assert evidence.review_pdf_status == "generated"
    assert evidence.review_pdf_renderer == "fake_pdf"
    assert evidence.review_pdf_path == output_dir / "official_official_review.pdf"
    assert evidence.review_pdf_path.is_file()
    assert result.output_paths["review_pdf"] == str(evidence.review_pdf_path)
    assert records[0].success is True
    assert "review_pdf=generated" in records[0].after


def test_pipeline_degrades_when_official_review_pdf_renderer_is_unavailable(
    tmp_path,
    monkeypatch,
):
    import src.shared.engine.official_document_assembly as assembly_module

    source = tmp_path / "official.docx"
    Document().save(source)
    monkeypatch.setattr(
        assembly_module,
        "_default_review_pdf_renderer_resolver",
        lambda: None,
    )
    output_dir = tmp_path / "out"
    config = SimpleNamespace(
        strict_mode=False,
        compliance_profile=SimpleNamespace(
            rule_family="official_document",
            profile_id="official_document",
            object_preflight=SimpleNamespace(enabled=False),
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="official_document_v1",
            material_schema_ids=[],
        ),
        entity_data={
            "document_type": "notice",
            "title": "关于开展资料归档检查的通知",
            "body": "请各部门按要求完成自查并提交材料。",
            "organization": "示例市档案局",
            "document_no": "示档发〔2026〕1号",
            "issue_date": "2026年7月9日",
        },
        delivery_presets=[],
        output=SimpleNamespace(final_docx=False, review_pdf=True),
    )

    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        official_document_type_id="notice",
    ).execute(str(source))

    evidence = result.context.official_document_assembly
    records = result.tracker.get_by_module("official_document_assembly")
    assert result.success is True
    assert evidence.status == "ok"
    assert evidence.review_pdf_status == "renderer_unavailable"
    assert evidence.review_pdf_path is None
    assert evidence.review_pdf_issue == "docx_to_pdf_renderer_unavailable"
    assert "review_pdf" not in result.output_paths
    assert records[0].success is True
    assert "review_pdf=renderer_unavailable" in records[0].after


def test_pipeline_reports_official_document_assembly_missing_fields(tmp_path):
    source = tmp_path / "official.docx"
    Document().save(source)

    config = SimpleNamespace(
        strict_mode=False,
        compliance_profile=SimpleNamespace(
            rule_family="official_document",
            profile_id="official_document",
            object_preflight=SimpleNamespace(enabled=False),
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="official_document_v1",
            material_schema_ids=[],
        ),
        entity_data={
            "document_type": "notice",
            "title": "缺正文的通知",
            "organization": "示例单位",
            "document_no": "示发〔2026〕2号",
            "issue_date": "2026年7月9日",
        },
        delivery_presets=[],
        output=SimpleNamespace(final_docx=False),
    )

    result = Pipeline(
        modules=[],
        config=config,
        official_document_type_id="notice",
    ).execute(str(source))
    evidence = result.context.official_document_assembly
    diagnostics = build_execution_diagnostics(result)

    assert result.success is False
    assert "Terminal assembler 'official' blocked delivery" in str(result.error)
    assert evidence.status == "missing_required_fields"
    assert evidence.missing_required_fields == ("body",)
    assert result.output_paths == {}
    assert any(
        item["rule_name"] == "official_document_assembly"
        and item["success"] is False
        for item in diagnostics["items"]
    )


def test_technical_chapter_inventory_counts_chapters_appendices_and_objects():
    doc = Document()
    doc.add_paragraph("Chapter 1 Overview")
    doc.add_paragraph("1.1 Scope")
    doc.add_paragraph("Appendix A Assets")
    doc.add_table(rows=1, cols=1)
    config = SimpleNamespace(
        compliance_profile=SimpleNamespace(
            rule_family="technical_review",
            profile_id="technical_review",
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="technical_document_v1",
            material_schema_ids=[],
        ),
    )

    evidence = inspect_technical_chapter_inventory(
        config,
        doc,
        heading_map={0: 1, 1: 2, 2: 1},
    )

    assert evidence.status == "ok"
    assert evidence.family_id == "technical_long_docs"
    assert evidence.summary.heading_count == 3
    assert evidence.summary.chapter_count == 1
    assert evidence.summary.appendix_count == 1
    assert evidence.summary.max_heading_level == 2
    assert evidence.summary.table_count == 1
    assert evidence.chapters[0].title == "Chapter 1 Overview"
    assert "w:tbl" in evidence.word_surfaces


def test_report_writer_emits_technical_chapter_inventory_evidence(tmp_path):
    context = PipelineContext(
        technical_chapter_inventory={
            "family_id": "technical_long_docs",
            "status": "ok",
            "rule_family": "technical_review",
            "profile_id": "technical_review",
            "material_schema_ids": ["technical_document_v1"],
            "word_surfaces": ["w:p/w:pPr", "numbering.xml", "w:tbl"],
            "manual_confirmation_required": False,
            "summary": {
                "heading_count": 3,
                "chapter_count": 1,
                "appendix_count": 1,
                "max_heading_level": 2,
                "table_count": 1,
                "figure_count": 0,
                "toc_present": False,
            },
            "chapters": [
                {
                    "paragraph_index": 0,
                    "level": 1,
                    "title": "Chapter 1 Overview",
                    "section_type": "body",
                    "start_index": 0,
                    "end_index": 3,
                }
            ],
            "headings": [
                {
                    "paragraph_index": 0,
                    "level": 1,
                    "title": "Chapter 1 Overview",
                    "section_type": "body",
                    "start_index": 0,
                    "end_index": 3,
                }
            ],
            "issues": [],
        }
    )
    result = PipelineResult(success=True, context=context)
    report_json = tmp_path / "technical_inventory.json"
    report_md = tmp_path / "technical_inventory.md"

    write_json_report(
        result,
        input_path=tmp_path / "manual.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "manual.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    evidence = report_data["technical_chapter_inventory"]
    assert evidence["status"] == "ok"
    assert evidence["summary"]["chapter_count"] == 1
    assert evidence["summary"]["appendix_count"] == 1
    assert "## 技术长文档章节清单" in markdown
    assert "- Chapters: 1" in markdown
    assert "Chapter 1 Overview" in markdown


def test_pipeline_records_technical_chapter_inventory_context(tmp_path):
    source = tmp_path / "manual.docx"
    doc = Document()
    doc.add_paragraph("Chapter 1 Overview")
    doc.add_paragraph("1.1 Scope")
    doc.add_paragraph("Appendix A Assets")
    doc.save(source)

    class _MarkTechnicalHeadingMap(BaseModule):
        meta = ModuleMeta(
            name="mark_technical_heading_map",
            description="Mark technical heading map",
            category="structure",
            provides=("heading_map",),
            execution_phase="semantics",
            scope_behavior="structure_discovery",
        )

        def apply(self, doc, config, tracker, context):
            context.heading_map = {0: 1, 1: 2, 2: 1}

    config = SimpleNamespace(
        compliance_profile=SimpleNamespace(
            rule_family="technical_review",
            profile_id="technical_review",
            object_preflight=SimpleNamespace(enabled=False),
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="technical_document_v1",
            material_schema_ids=[],
        ),
        entity_data={},
        delivery_presets=[],
        output=SimpleNamespace(final_docx=False),
    )

    result = Pipeline(
        modules=[_MarkTechnicalHeadingMap()],
        config=config,
    ).execute(str(source))

    assert result.success is True
    evidence = result.context.technical_chapter_inventory
    assert evidence.status == "ok"
    assert evidence.summary.chapter_count == 1
    assert evidence.summary.appendix_count == 1
    assert result.tracker.get_by_module("technical_chapter_inventory") == []


def test_application_section_word_limits_detects_exceeded_and_missing_sections():
    doc = Document()
    doc.add_paragraph("Project summary")
    doc.add_paragraph("A" * 520)
    doc.add_paragraph("Budget")
    doc.add_paragraph("Short budget.")
    config = SimpleNamespace(
        compliance_profile=SimpleNamespace(
            rule_family="project_application",
            profile_id="project_application",
            count_profile_id="application_word_limits",
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="project_application_materials_v1",
            material_schema_ids=[],
        ),
    )

    evidence = inspect_application_section_word_limits(
        config,
        doc,
        heading_map={0: 1, 2: 1},
    )

    assert evidence.status == "warning"
    assert evidence.profile_id == "application_word_limits"
    assert evidence.summary.matched_section_count == 2
    assert evidence.summary.exceeded_section_count == 1
    assert evidence.summary.required_missing_count == 0
    summary = next(section for section in evidence.sections if section.section_id == "summary")
    assert summary.status == "exceeded"
    assert summary.counts["characters_no_spaces"] == 520
    assert any(
        issue.kind == "section_word_limit_exceeded"
        and issue.section_id == "summary"
        for issue in evidence.issues
    )


def test_report_writer_emits_application_section_word_limit_evidence(tmp_path):
    context = PipelineContext(
        application_section_word_limits={
            "family_id": "application_reports",
            "status": "warning",
            "profile_id": "application_word_limits",
            "profile_name": "Project application word limits",
            "rule_family": "project_application",
            "material_schema_ids": ["project_application_materials_v1"],
            "word_surfaces": ["w:p/w:pPr", "w:r/w:t"],
            "manual_confirmation_required": True,
            "summary": {
                "section_count": 2,
                "matched_section_count": 2,
                "exceeded_section_count": 1,
                "required_missing_count": 0,
                "table_count": 0,
            },
            "sections": [
                {
                    "section_id": "summary",
                    "label": "Project summary",
                    "heading_title": "Project summary",
                    "paragraph_index": 0,
                    "start_index": 1,
                    "end_index": 2,
                    "status": "exceeded",
                    "counts": {
                        "characters_no_spaces": 520,
                        "cjk_characters": 0,
                        "english_words": 1,
                    },
                    "limits": {
                        "max_characters_no_spaces": 500,
                        "max_cjk_characters": 0,
                        "max_english_words": 0,
                    },
                }
            ],
            "issues": [
                {
                    "kind": "section_word_limit_exceeded",
                    "section_id": "summary",
                    "severity": "warning",
                    "message": "Project summary exceeds the configured section limit.",
                }
            ],
        }
    )
    result = PipelineResult(success=True, context=context)
    report_json = tmp_path / "application_limits.json"
    report_md = tmp_path / "application_limits.md"

    write_json_report(
        result,
        input_path=tmp_path / "application.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "application.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    evidence = report_data["application_section_word_limits"]
    assert evidence["status"] == "warning"
    assert evidence["summary"]["exceeded_section_count"] == 1
    assert evidence["sections"][0]["section_id"] == "summary"
    assert "## 应用材料分节限字证据" in markdown
    assert "summary: Project summary [exceeded] chars=520/500" in markdown


def test_pipeline_records_application_section_word_limit_context(tmp_path):
    source = tmp_path / "application.docx"
    doc = Document()
    doc.add_paragraph("Project summary")
    doc.add_paragraph("A" * 520)
    doc.add_paragraph("Budget")
    doc.add_paragraph("Short budget.")
    doc.save(source)

    class _MarkApplicationHeadingMap(BaseModule):
        meta = ModuleMeta(
            name="mark_application_heading_map",
            description="Mark application heading map",
            category="structure",
            provides=("heading_map",),
            execution_phase="semantics",
            scope_behavior="structure_discovery",
        )

        def apply(self, doc, config, tracker, context):
            context.heading_map = {0: 1, 2: 1}

    config = SimpleNamespace(
        compliance_profile=SimpleNamespace(
            rule_family="project_application",
            profile_id="project_application",
            count_profile_id="application_word_limits",
            object_preflight=SimpleNamespace(enabled=False),
        ),
        input_source_profile=SimpleNamespace(
            material_schema_id="project_application_materials_v1",
            material_schema_ids=[],
        ),
        entity_data={},
        delivery_presets=[],
        output=SimpleNamespace(final_docx=False),
    )

    result = Pipeline(
        modules=[_MarkApplicationHeadingMap()],
        config=config,
    ).execute(str(source))

    assert result.success is True
    evidence = result.context.application_section_word_limits
    assert evidence.status == "warning"
    assert evidence.summary.exceeded_section_count == 1
    diagnostics = build_execution_diagnostics(result)
    assert any(
        item["rule_name"] == "application_section_word_limits"
        for item in diagnostics["items"]
    )


def test_report_writer_emits_academic_citation_and_formula_confidence(tmp_path):
    tracker = ChangeTracker()
    tracker.record(
        rule_name="citation_link",
        target="summary",
        section="global",
        change_type="field",
        before="entry_bookmarks=2, number_bookmarks=0, references=2",
        after="linked=1, unresolved=1, duplicates=0",
    )
    tracker.record(
        rule_name="equation_table_format",
        target="1 个公式表格",
        section="global",
        change_type="format",
        before="(mixed)",
        after="table=center, formula=center, number=right, numbering=chapter.seq",
    )
    tracker.record(
        rule_name="equation_table_format",
        target="1 个公式编号",
        section="global",
        change_type="skip",
        before="chapter-aware numbering normalization",
        after="skipped due to missing chapter context",
    )
    context = PipelineContext(
        count_result={
            "profile_id": "thesis_cn",
            "profile_name": "Chinese thesis count",
            "scope": "academic_document",
            "counts": {
                "reference_count": 2,
                "equation_count": 1,
            },
        }
    )
    result = PipelineResult(success=True, tracker=tracker, context=context)
    report_json = tmp_path / "academic_changes.json"
    report_md = tmp_path / "academic_changes.md"

    write_json_report(
        result,
        input_path=tmp_path / "thesis.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=2,
        modules_total=2,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "thesis.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=2,
        modules_total=2,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")
    evidence = report_data["academic_confidence"]

    assert evidence["status"] == "needs_review"
    assert evidence["confidence_level"] == "low"
    assert evidence["manual_confirmation_required"] is True
    assert evidence["profile_id"] == "thesis_cn"
    assert evidence["citation"]["reference_count"] == 2
    assert evidence["citation"]["linked_citation_count"] == 1
    assert evidence["citation"]["unresolved_citation_count"] == 1
    assert evidence["formula"]["equation_count"] == 1
    assert evidence["formula"]["formatted_equation_table_count"] == 1
    assert evidence["formula"]["skipped_number_count"] == 1
    assert any(issue["domain"] == "formula" for issue in evidence["issues"])
    assert any(issue["domain"] == "citation" for issue in evidence["issues"])
    assert "## 学术引用与公式置信度" in markdown
    assert "Citation: status=needs_review, references=2, linked=1, unresolved=1" in markdown
    assert "Formula: status=needs_review, equations=1, tables=1" in markdown
    assert "skipped due to missing chapter context" in markdown


def test_pipeline_builds_coverage_boundary_report_items_from_rule_family():
    config = SimpleNamespace(
        compliance_profile=SimpleNamespace(rule_family="exam_teaching"),
        input_source_profile=SimpleNamespace(
            material_schema_id="",
            material_schema_ids=[],
        ),
    )

    items = build_coverage_boundary_report_items(config)

    assert [item["pack_id"] for item in items] == ["exam_education"]
    assert items[0]["boundary"] == "does not guarantee AI content quality or complex diagram generation"
    assert items[0]["plugin_manual_gate"]["plugin_entry_id"] == (
        "exam_ai_quality_diagram_plugin"
    )
    assert items[0]["plugin_manual_gate"]["manual_confirmation_required"] is True
    assert items[0]["plugin_manual_gate"]["confidence_report_required"] is True
    assert items[0]["plugin_manual_gate"]["risk_domain_ids"] == [
        "ai_content_quality",
        "complex_diagram_generation",
    ]
    assert items[0]["plugin_manual_gate"]["confirmation_decision_states"] == [
        "accepted",
        "rejected",
        "needs_plugin_handoff",
    ]
    assert "AI quality and complex diagram plugin/manual gate entry is registered" in (
        items[0]["implemented_closures"]
    )
    assert items[0]["closure_tasks"] == []


def test_pipeline_does_not_report_plugin_boundary_for_normal_report_scene():
    config = SimpleNamespace(
        scene_id="report",
        category="report",
        compliance_profile=SimpleNamespace(rule_family="basic_format"),
        input_source_profile=SimpleNamespace(
            material_schema_id="",
            material_schema_ids=[],
        ),
    )

    assert build_coverage_boundary_report_items(config) == []


def test_pipeline_reports_direct_import_ai_plugin_manual_gate():
    config = SimpleNamespace(
        scene_id="import_ai_boundary",
        category="import_ai_boundary",
        compliance_profile=SimpleNamespace(rule_family="import_ai_boundary"),
        input_source_profile=SimpleNamespace(
            material_schema_id="",
            material_schema_ids=[],
        ),
    )

    items = build_coverage_boundary_report_items(config)

    assert [item["pack_id"] for item in items] == ["import_ai_boundary"]
    gate = items[0]["plugin_manual_gate"]
    assert gate["gate_id"] == "import_ai_conversion_gate"
    assert gate["plugin_entry_id"] == "import_ai_assistant_plugin"
    assert gate["confidence_report_required"] is True
    assert gate["manual_confirmation_required"] is True
    assert gate["risk_domain_ids"] == [
        "ocr_confidence",
        "pdf_conversion",
        "full_latex_conversion",
        "ai_content_quality",
        "complex_diagram_generation",
    ]
    assert gate["report_fields"] == [
        "source_type",
        "confidence_level",
        "low_confidence_regions",
        "manual_decision",
        "handoff_pack_id",
        "handoff_family_id",
        "handoff_status",
        "fallback_strategy",
    ]
    assert "lossless_pdf_to_word" in gate["unsupported_core_inputs"]


def test_pipeline_builds_contract_delivery_legal_boundary_report_items():
    config = SimpleNamespace(
        scene_id="contract_delivery",
        category="contract_delivery",
        compliance_profile=SimpleNamespace(rule_family="contract_delivery"),
        input_source_profile=SimpleNamespace(
            material_schema_id="contract_parties_v1",
            material_schema_ids=["signature_assets_v1"],
        ),
    )

    items = build_coverage_boundary_report_items(config)

    assert [item["pack_id"] for item in items] == ["contract_delivery"]
    assert items[0]["boundary"] == "does not provide legal advice or judge clause validity"
    assert "contract field consistency has report and Workbench summary evidence" in (
        items[0]["implemented_closures"]
    )


def test_report_writer_emits_coverage_boundary_evidence(tmp_path):
    context = PipelineContext(
        coverage_boundaries=[
            {
                "pack_id": "exam_education",
                "label": "Exam and teaching materials",
                "boundary": "does not guarantee AI content quality or complex diagram generation",
                "primary_landings": ["exam_teaching"],
                "secondary_landings": ["DeliveryPreset"],
                "capability_axis_ids": ["input_fact_source", "plugin_boundary"],
                "missing_closures": ["AI quality and complex diagram plugin entry"],
                "closure_tasks": [
                    {
                        "summary": "AI quality and complex diagram plugin entry",
                        "priority": "P2",
                        "owner": "plugin",
                        "target_phase": "L4",
                        "validation_commands": [
                            "python -m pytest tests/test_scene_coverage_manifest.py -q"
                        ],
                    }
                ],
                "plugin_manual_gate": {
                    "gate_id": "exam_ai_complex_diagram_gate",
                    "pack_id": "exam_education",
                    "label": "Exam AI quality and complex diagram gate",
                    "plugin_entry_id": "exam_ai_quality_diagram_plugin",
                    "plugin_entry_label": "Exam AI quality / complex diagram plugin",
                    "manual_confirmation_required": True,
                    "confidence_report_required": True,
                    "boundary_report_required": True,
                    "blocks_core_execution_until_confirmed": True,
                    "professional_review_required": False,
                    "risk_domain_ids": [
                        "ai_content_quality",
                        "complex_diagram_generation",
                    ],
                    "accepted_inputs": ["ai_generated_questions"],
                    "unsupported_core_inputs": ["ai_quality_judgment"],
                    "confirmation_scope": ["question correctness"],
                    "confirmation_decision_states": [
                        "accepted",
                        "rejected",
                        "needs_plugin_handoff",
                    ],
                    "report_fields": ["confidence_level", "manual_decision"],
                },
            }
        ]
    )
    result = PipelineResult(success=True, context=context)
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"

    write_json_report(
        result,
        input_path=tmp_path / "exam.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "exam.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert report_data["coverage_boundaries"][0]["pack_id"] == "exam_education"
    assert report_data["coverage_boundaries"][0]["closure_tasks"][0]["owner"] == "plugin"
    assert report_data["coverage_boundaries"][0]["plugin_manual_gate"][
        "plugin_entry_id"
    ] == "exam_ai_quality_diagram_plugin"
    assert "## 方案边界证据" in markdown
    assert "exam_education (Exam and teaching materials)" in markdown
    assert "Plugin entry: exam_ai_quality_diagram_plugin" in markdown
    assert "Manual confirmation: required" in markdown
    assert "Confidence report: required" in markdown
    assert "Risk domains: ai_content_quality, complex_diagram_generation" in markdown
    assert "Decision states: accepted, rejected, needs_plugin_handoff" in markdown
    assert "Report fields: confidence_level, manual_decision" in markdown
    assert "[P2/plugin/L4] AI quality and complex diagram plugin entry" in markdown


def test_report_writer_emits_exam_question_schema_evidence(tmp_path):
    context = PipelineContext(
        exam_question_schema={
            "schema_id": "exam_items_v1",
            "family_id": "exam_teaching",
            "status": "error",
            "source_key": "entity_data.exam_items",
            "manual_confirmation_required": True,
            "summary": {
                "section_count": 1,
                "question_count": 1,
                "answered_question_count": 0,
                "analysis_count": 0,
                "knowledge_point_count": 0,
                "figure_count": 0,
                "scored_question_count": 1,
                "declared_total_score": 10,
                "computed_total_score": 10,
                "question_types": ["short_answer"],
            },
            "issue_count": 1,
            "error_count": 1,
            "warning_count": 0,
            "issues": [
                {
                    "path": "sections.0.questions.0.answer",
                    "kind": "missing_answer",
                    "severity": "error",
                    "message": "Answer must remain in the structured source.",
                }
            ],
        }
    )
    result = PipelineResult(success=True, context=context)
    report_json = tmp_path / "exam_schema.json"
    report_md = tmp_path / "exam_schema.md"

    write_json_report(
        result,
        input_path=tmp_path / "exam.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "exam.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert report_data["exam_question_schema"]["schema_id"] == "exam_items_v1"
    assert report_data["exam_question_schema"]["error_count"] == 1
    assert "## 试卷题源结构校验" in markdown
    assert "missing_answer" in markdown


def test_report_writer_emits_exam_delivery_runtime_evidence(tmp_path):
    context = PipelineContext(
        exam_delivery_runtime={
            "schema_id": "exam_items_v1",
            "family_id": "exam_teaching",
            "status": "ok",
            "source_key": "entity_data.exam_items",
            "markdown_preview_path": "preview/exam.md",
            "markdown_preview_excerpt": "一、选择题",
            "version_count": 1,
            "master_evidence": {
                "mode_id": "exam",
                "master_id": "default_exam",
                "master_label": "A4 标准卷面",
                "master_source_type": "builtin",
                "master_docx_path": "config_library/masters/exam/builtin/default_exam_v20.docx",
                "master_version": "exam-master-v20-free-answer-area-2026-07-06",
                "manifest_path": (
                    "config_library/masters/exam/builtin/default_exam.master.json"
                ),
                "placeholder_contract_status": "ok",
                "required_placeholders": ["af_title", "af_questions"],
                "optional_placeholders": ["af_answer_area"],
                "runtime_fields": ["title", "subject"],
                "generated_fields": ["version"],
            },
            "rendered_versions": [
                {
                    "preset_id": "student",
                    "label": "学生版",
                    "docx_path": "dist/exam_student.docx",
                    "hidden_selectors": ["answers", "analysis"],
                    "visible_question_count": 2,
                    "visible_answer_count": 0,
                    "visible_analysis_count": 0,
                    "fixed_layout_kind": "answer_sheet",
                    "fixed_layout_row_count": 4,
                    "fixed_layout_column_count": 2,
                    "fixed_layout_row_height_twips": 720,
                    "question_asset_count": 1,
                    "rendered_question_asset_count": 1,
                    "question_asset_alt_text_count": 1,
                    "rendered_question_asset_alt_text_count": 1,
                }
            ],
        }
    )
    result = PipelineResult(success=True, context=context)
    report_json = tmp_path / "exam_delivery.json"
    report_md = tmp_path / "exam_delivery.md"

    write_json_report(
        result,
        input_path=tmp_path / "exam.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "exam.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")
    rendered_version = report_data["exam_delivery_runtime"]["rendered_versions"][0]
    master = report_data["exam_delivery_runtime"]["master_evidence"]

    assert rendered_version["preset_id"] == "student"
    assert rendered_version["fixed_layout_row_height_twips"] == 720
    assert master["mode_id"] == "exam"
    assert master["master_id"] == "default_exam"
    assert master["placeholder_contract_status"] == "ok"
    assert "## 试卷多版本运行时渲染" in markdown
    assert "Work mode: exam" in markdown
    assert "Master: A4 标准卷面 / default_exam" in markdown
    assert "contract=ok" in markdown
    assert "Master DOCX: `config_library/masters/exam/builtin/default_exam_v20.docx`" in markdown
    assert "student: `dist/exam_student.docx`" in markdown
    assert "assets=1/1" in markdown
    assert "Preview excerpt: 一、选择题" in markdown


def test_report_writer_emits_contract_delivery_legal_boundary_evidence(tmp_path):
    config = SimpleNamespace(
        scene_id="contract_delivery",
        category="contract_delivery",
        compliance_profile=SimpleNamespace(rule_family="contract_delivery"),
        input_source_profile=SimpleNamespace(
            material_schema_id="contract_parties_v1",
            material_schema_ids=["signature_assets_v1"],
        ),
    )
    context = PipelineContext(
        coverage_boundaries=build_coverage_boundary_report_items(config)
    )
    result = PipelineResult(success=True, context=context)
    report_json = tmp_path / "contract_changes.json"
    report_md = tmp_path / "contract_changes.md"

    write_json_report(
        result,
        input_path=tmp_path / "contract.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "contract.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert report_data["coverage_boundaries"][0]["pack_id"] == "contract_delivery"
    assert report_data["coverage_boundaries"][0]["boundary"] == (
        "does not provide legal advice or judge clause validity"
    )
    assert "contract_delivery (Contract delivery)" in markdown
    assert "does not provide legal advice or judge clause validity" in markdown


def test_report_writer_emits_scene_product_readiness_for_green_contract_delivery(
    tmp_path,
):
    config = SimpleNamespace(
        scene_id="contract_delivery",
        category="contract_delivery",
        compliance_profile=SimpleNamespace(rule_family="contract_delivery"),
        input_source_profile=SimpleNamespace(
            material_schema_id="contract_parties_v1",
            material_schema_ids=[],
        ),
    )
    context = PipelineContext(
        coverage_boundaries=build_coverage_boundary_report_items(config)
    )
    result = PipelineResult(success=True, config=config, context=context)
    report_json = tmp_path / "contract_readiness.json"
    report_md = tmp_path / "contract_readiness.md"

    write_json_report(
        result,
        input_path=tmp_path / "contract.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "contract.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    evidence = report_data["scene_product_readiness"]
    markdown = report_md.read_text(encoding="utf-8")
    subjects = {
        (subject["subject_type"], subject["subject_id"]): subject
        for subject in evidence["subjects"]
    }

    assert evidence["status"] == "green_l5"
    assert evidence["registry_audit_status"] == "clean"
    assert evidence["static_closed_but_not_green_count"] == 0
    assert evidence["level_counts"]["green_l5"] == 2
    assert evidence["remaining_gap_count"] == 0
    assert subjects[("pack", "contract_delivery")]["product_readiness_level"] == "green_l5"
    assert subjects[("family", "contract_delivery")]["product_readiness_level"] == "green_l5"
    assert subjects[("pack", "contract_delivery")]["is_green_l5"] is True
    assert "pre-execution legal-boundary banner" in (
        subjects[("pack", "contract_delivery")]["evidence_surfaces"]
    )
    assert "## 方案产品成熟度证据" in markdown
    assert "Static closed != Green/L5: 0" in markdown
    assert "field evidence UI" in markdown


def test_report_writer_emits_parameter_ownership_evidence(tmp_path):
    @dataclass
    class FutureSceneWorkspace(SceneWorkspace):
        experimental_knob: str = ""

    result = PipelineResult(
        success=True,
        config=FutureSceneWorkspace(
            scene_id="contract_delivery",
            category="contract_delivery",
            template_id="contract_template",
        ),
    )
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"

    write_json_report(
        result,
        input_path=tmp_path / "contract.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "contract.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    evidence = report_data["parameter_ownership"]
    markdown = report_md.read_text(encoding="utf-8")

    assert evidence["status"] == "has_gaps"
    assert evidence["config_identity"]["scene_id"] == "contract_delivery"
    assert evidence["layer_counts"]["material"] > 0
    assert evidence["layer_counts"]["output"] > 0
    assert evidence["execution_consumer_anchor_status"] == "clean"
    assert evidence["execution_consumer_anchor_audit"]["gap_count"] == 0
    assert any(
        sample["source_path"] == "src/report_writer.py"
        for sample in evidence["execution_consumer_anchor_samples"]
    )
    assert evidence["audit"]["missing_top_level_paths"] == ["experimental_knob"]
    assert any(
        sample["path"] == "input_source_profile.material_schema_id"
        and sample["owner_layer"] == "material"
        and sample["runtime_available"]
        for sample in evidence["samples"]
    )
    assert "## 参数归属证据" in markdown
    assert "Consumer anchors: clean" in markdown
    assert "experimental_knob" in markdown
    assert "input_source_profile.material_schema_id" in markdown


def test_report_writer_emits_control_contract_evidence(tmp_path):
    result = PipelineResult(success=True, config=SceneWorkspace(scene_id="report"))
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"

    write_json_report(
        result,
        input_path=tmp_path / "report.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "report.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    evidence = report_data["control_contracts"]
    markdown = report_md.read_text(encoding="utf-8")
    special_sample = next(
        sample
        for sample in evidence["samples"]
        if sample["contract_id"] == "body.special_indent"
    )
    row_height_sample = next(
        sample
        for sample in evidence["samples"]
        if sample["contract_id"] == "fixed_layout.table_row_height"
    )
    delivery_sample = next(
        sample
        for sample in evidence["samples"]
        if sample["contract_id"] == "output.delivery_preset"
    )
    plugin_sample = next(
        sample
        for sample in evidence["samples"]
        if sample["contract_id"] == "plugin.manual_gate"
    )

    assert evidence["status"] == "clean"
    assert evidence["contract_count"] == 16
    assert evidence["owner_counts"]["template"] == 9
    assert evidence["owner_counts"]["scene"] == 3
    assert evidence["owner_counts"]["material"] == 1
    assert evidence["owner_counts"]["output"] == 2
    assert evidence["owner_counts"]["plugin"] == 1
    assert evidence["audit"]["gap_count"] == 0
    assert "SpecialIndentInput" in evidence["canonical_controls"]
    assert "StyledComboBox + preset editor + artifact toggles" in evidence[
        "canonical_controls"
    ]
    assert special_sample["canonical_control"] == "SpecialIndentInput"
    assert special_sample["unit_set"] == ["chars", "pt", "cm"]
    assert "mode=none" in special_sample["disabled_state_rule"]
    special_locations = special_sample["evidence_locations"]
    assert any(
        location["source_path"] == "src/shared/ui/paragraph_style_inputs.py"
        and location["marker"] == "class SpecialIndentInput"
        and location["line_number"] > 0
        for location in special_locations
    )
    assert row_height_sample["owner_layer"] == "scene"
    assert row_height_sample["workbench_surface"] == "fixed-layout policy/report"
    assert "word.w:trHeight" in row_height_sample["parameter_paths"]
    assert delivery_sample["owner_layer"] == "output"
    assert "scene.default_delivery_preset_id" in delivery_sample["parameter_paths"]
    assert plugin_sample["owner_layer"] == "plugin"
    assert "plugin_manual_gate.*" in plugin_sample["parameter_paths"]
    assert "## 控件契约证据" in markdown
    assert "Status: clean" in markdown
    assert "body.special_indent: SpecialIndentInput" in markdown
    assert "src/shared/ui/paragraph_style_inputs.py:" in markdown
    assert "#class SpecialIndentInput" in markdown
    assert "fixed_layout.table_row_height" in markdown


def test_report_writer_emits_parameter_effective_value_evidence(tmp_path):
    scene = SceneWorkspace(scene_id="report", category="report")
    scene.template_overrides["table.layout_mode"] = "full"
    config = resolve_config(TemplateConfig(), scene)
    result = PipelineResult(success=True, config=config)
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"

    write_json_report(
        result,
        input_path=tmp_path / "report.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "report.docx",
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    evidence = report_data["parameter_ownership"]
    markdown = report_md.read_text(encoding="utf-8")
    table_sample = next(
        sample
        for sample in evidence["effective_value_samples"]
        if sample["path"] == "table.layout_mode"
    )

    assert evidence["effective_value_status"] == "available"
    assert evidence["effective_value_source_counts"]["scene"] > 0
    assert table_sample["runtime_path"] == "table.layout_mode"
    assert table_sample["source"] == "scene"
    assert table_sample["source_kind"] == "provenance"
    assert table_sample["value"] == "full"
    assert table_sample["template_default"] == "smart"
    assert table_sample["overridden"] is True
    assert "Effective values: available" in markdown
    assert "table.layout_mode: source=scene" in markdown


def test_pipeline_report_emits_parameter_runtime_consumption_trace(tmp_path):
    scene = SceneWorkspace(scene_id="report", category="report")
    scene.template_overrides["table.layout_mode"] = "full"
    config = resolve_config(TemplateConfig(), scene)
    input_doc = tmp_path / "runtime.docx"
    Document().save(str(input_doc))
    pipeline = Pipeline(
        [RuntimeTraceTableModule()],
        config,
        output_dir=str(tmp_path),
    )

    result = pipeline.execute(str(input_doc))
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"
    write_json_report(
        result,
        input_path=input_doc,
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=1,
        modules_total=1,
        include_internal_evidence=True,
    )
    write_markdown_report(
        result,
        input_path=input_doc,
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=1,
        modules_total=1,
        include_internal_evidence=True,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    evidence = report_data["parameter_ownership"]
    markdown = report_md.read_text(encoding="utf-8")
    sample = evidence["runtime_consumption_samples"][0]

    assert result.success is True
    assert result.context.parameter_runtime_consumption[0]["module_name"] == "table_runtime_trace"
    assert evidence["runtime_consumption_status"] == "available"
    assert evidence["runtime_consumption_event_count"] == 1
    assert evidence["runtime_consumption_module_count"] == 1
    assert sample["module_name"] == "table_runtime_trace"
    assert sample["status"] == "executed"
    assert sample["config_paths"] == ["table"]
    assert sample["runtime_available_paths"] == ["table"]
    assert "table.layout_mode" in sample["owned_spec_paths"]
    assert sample["source_counts"]["scene"] > 0
    assert "Runtime consumption: available" in markdown
    assert "table_runtime_trace: executed, config=table" in markdown


def test_pipeline_report_links_control_contracts_to_effective_values_and_runtime(tmp_path):
    template = TemplateConfig()
    template.styles["body"] = StyleConfig(
        special_indent_mode="none",
        special_indent_value=0,
        special_indent_unit="chars",
    )
    scene = SceneWorkspace(
        scene_id="report",
        category="report",
        template_overrides={
            "styles.body.special_indent_mode": "first_line",
            "styles.body.special_indent_value": 2.0,
            "styles.body.special_indent_unit": "chars",
        },
    )
    config = resolve_config(template, scene)
    input_doc = tmp_path / "style_runtime.docx"
    Document().save(str(input_doc))
    pipeline = Pipeline(
        [RuntimeTraceStylesModule()],
        config,
        output_dir=str(tmp_path),
    )

    result = pipeline.execute(str(input_doc))
    report_json = tmp_path / "changes.json"
    report_md = tmp_path / "changes.md"
    write_json_report(
        result,
        input_path=input_doc,
        output_path=None,
        report_path=report_json,
        elapsed=0.5,
        modules_enabled=1,
        modules_total=1,
        include_internal_evidence=True,
    )
    write_markdown_report(
        result,
        input_path=input_doc,
        report_path=report_md,
        elapsed=0.5,
        modules_enabled=1,
        modules_total=1,
        include_internal_evidence=True,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    evidence = report_data["control_contracts"]
    markdown = report_md.read_text(encoding="utf-8")
    special_trace = next(
        sample
        for sample in evidence["field_trace_samples"]
        if sample["contract_id"] == "body.special_indent"
    )
    mode_trace = next(
        trace
        for trace in special_trace["parameter_traces"]
        if trace["runtime_path"] == "styles.body.special_indent_mode"
    )

    assert result.success is True
    assert evidence["field_trace_status_counts"]["effective_and_consumed"] >= 1
    assert special_trace["trace_status"] == "effective_and_consumed"
    assert mode_trace["runtime_available"] is True
    assert mode_trace["source"] == "scene"
    assert mode_trace["source_kind"] == "provenance"
    assert mode_trace["overridden"] is True
    assert mode_trace["value"] == "first_line"
    assert "template.styles.body.special_indent_mode" in mode_trace["declared_paths"]
    assert mode_trace["declared_paths"] == [
        "template.styles.body.special_indent_mode"
    ]
    event = special_trace["runtime_consumption_events"][0]
    assert event["module_name"] == "styles_runtime_trace"
    assert event["status"] == "executed"
    assert event["config_paths"] == ["styles"]
    assert "styles.body.special_indent_mode" in event["matched_runtime_paths"]
    assert "Field trace:" in markdown
    assert "body.special_indent: effective_and_consumed" in markdown
    assert "styles.body.special_indent_mode[available=yes, source=scene]" in markdown
    assert "styles_runtime_trace:executed@styles.body.special_indent_mode" in markdown
