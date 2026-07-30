import json
import subprocess
import sys
from pathlib import Path
from zipfile import is_zipfile

from docx import Document

from src.config.fixed_layout import FixedLayoutRowHeightPolicy
from src.config.resolver import resolve_config
from src.config.scene import SceneWorkspace
from src.config.scene_coverage_manifest import SCENE_COVERAGE_PACK_MAP
from src.config.scene_high_frequency_request_samples import (
    list_high_frequency_request_samples,
)
from src.config.scene_request_cell_fixture_registry import (
    audit_scene_request_cell_fixtures,
    build_scene_request_cell_fixture_summary,
    list_scene_request_cell_fixtures,
    request_cell_fixtures_for_pack,
)
from src.config.scene_sample_fixture_registry import (
    REQUIRED_SAMPLE_PACK_IDS,
    REQUIRED_SAMPLE_SURFACES,
    audit_scene_sample_fixtures,
    build_scene_sample_coverage_summary,
    build_scene_sample_fixture_summary,
    get_scene_sample_fixture,
    list_scene_sample_fixtures,
    sample_surface_coverage,
    scene_sample_fixtures_for_pack,
)
from src.config.template import TemplateConfig
from src.modules.base import BaseModule, ModuleMeta
from src.pipeline.context import PipelineContext
from src.pipeline.result import PipelineResult
from src.pipeline.runner import Pipeline
from src.pipeline.tracker import ChangeTracker
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.fixed_layout_tables import (
    apply_fixed_layout_row_height_policy,
    row_height_state,
)
from src.shared.engine.scene_sample_docx_builder import (
    build_scene_sample_docx,
    build_scene_sample_docx_library,
    scene_sample_fixture_manifest_paths,
)
from src.shared.engine.object_preflight import inspect_docx_package


def test_scene_sample_fixture_registry_covers_n2_134_required_packs_and_surfaces():
    assert audit_scene_sample_fixtures() == ()

    fixtures = {fixture.fixture_id: fixture for fixture in list_scene_sample_fixtures()}
    coverage = sample_surface_coverage()

    assert REQUIRED_SAMPLE_PACK_IDS == tuple(SCENE_COVERAGE_PACK_MAP)
    assert len(fixtures) == len(REQUIRED_SAMPLE_PACK_IDS) + 30
    assert set(REQUIRED_SAMPLE_PACK_IDS) == {
        fixture.pack_id for fixture in fixtures.values()
    }
    assert fixtures["application_reports_product_sales_assets"].family_id == (
        "product_sales_documents"
    )
    assert fixtures["batch_forms_hr_batch_offer"].family_id == "hr_batch_documents"
    assert fixtures["quick_formatting_preserve_fields_degraded"].request_sample_ids == (
        "quick_template_cleanup",
    )
    assert "preserve_layout" in (
        fixtures["quick_formatting_preserve_fields_degraded"].expected_behaviors
    )
    assert fixtures[
        "quick_formatting_business_template_cleanup"
    ].request_sample_ids == (
        "quick_word_cleanup",
        "personal_resume_formatting",
    )
    assert "table_grid" in (
        fixtures["quick_formatting_business_template_cleanup"].docx_surfaces
    )
    assert "preserve_layout" in (
        fixtures["quick_formatting_business_template_cleanup"].expected_behaviors
    )
    assert "template_field_repair" in (
        fixtures["quick_formatting_business_template_cleanup"].report_expectations
    )
    assert fixtures["exam_education_controls_textboxes"].request_sample_ids == (
        "exam_multi_version",
    )
    assert fixtures["exam_education_structured_multiversion"].request_sample_ids == (
        "exam_answer_sheet",
        "exam_answer_analysis",
    )
    assert fixtures["english_journal_bibtex_csl_degraded"].request_sample_ids == (
        "english_journal_submission_package",
    )
    assert "reviewed_journal_profile_update" in (
        fixtures["english_journal_revision_comments"].report_expectations
    )
    assert "submission_artifact_manifest" in (
        fixtures["english_journal_revision_comments"].report_expectations
    )
    assert fixtures[
        "chinese_academic_school_rule_section_confirmation"
    ].family_id == "thesis_cn"
    assert fixtures[
        "chinese_academic_school_rule_section_confirmation"
    ].request_sample_ids == ("chinese_thesis_count",)
    assert "manual_confirmation" in (
        fixtures[
            "chinese_academic_school_rule_section_confirmation"
        ].expected_behaviors
    )
    assert "school_rule_source_selection" in (
        fixtures[
            "chinese_academic_school_rule_section_confirmation"
        ].report_expectations
    )
    assert "section_classifier_confirmation" in (
        fixtures[
            "chinese_academic_school_rule_section_confirmation"
        ].report_expectations
    )
    assert "skip_report" in (
        fixtures["english_journal_bibtex_csl_degraded"].expected_behaviors
    )
    assert "reviewed_journal_profile_update" in (
        fixtures["english_journal_bibtex_csl_degraded"].report_expectations
    )
    assert "submission_artifact_manifest" in (
        fixtures["english_journal_bibtex_csl_degraded"].report_expectations
    )
    assert fixtures[
        "english_journal_publisher_rule_manual_boundary"
    ].manual_gate_id == "journal_publisher_rule_review_gate"
    assert fixtures[
        "english_journal_publisher_rule_manual_boundary"
    ].request_sample_ids == ("english_journal_response_letter",)
    assert "reviewed_journal_profile_update" in (
        fixtures[
            "english_journal_publisher_rule_manual_boundary"
        ].report_expectations
    )
    assert "submission_artifact_manifest" in (
        fixtures[
            "english_journal_publisher_rule_manual_boundary"
        ].report_expectations
    )
    assert fixtures["contract_delivery_signature_fields_degraded"].request_sample_ids == (
        "contract_signing_consistency",
        "contract_signature_package_fields",
    )
    assert "skip_report" in (
        fixtures["contract_delivery_signature_fields_degraded"].expected_behaviors
    )
    assert fixtures["batch_forms_hr_missing_fields_manual_boundary"].request_sample_ids == (
        "batch_employee_certificate",
    )
    assert "manual_confirmation" in (
        fixtures["batch_forms_hr_missing_fields_manual_boundary"].expected_behaviors
    )
    assert "profile_specific_preview" in (
        fixtures["batch_forms_hr_missing_fields_manual_boundary"].report_expectations
    )
    assert "profile_specific_preview" in (
        fixtures["batch_forms_hr_batch_offer"].report_expectations
    )
    assert "fixed_layout_profile_browser" in (
        fixtures["batch_forms_fixed_layout"].report_expectations
    )
    assert "answer_sheet_reuse_path" in (
        fixtures["batch_forms_fixed_layout"].report_expectations
    )
    assert fixtures[
        "batch_forms_fixed_layout_placeholder_manual_boundary"
    ].request_sample_ids == ("batch_certificate_printing",)
    assert "placeholder_residue_report" in (
        fixtures[
            "batch_forms_fixed_layout_placeholder_manual_boundary"
        ].report_expectations
    )
    assert "fixed_layout_profile_browser" in (
        fixtures[
            "batch_forms_fixed_layout_placeholder_manual_boundary"
        ].report_expectations
    )
    assert "answer_sheet_reuse_path" in (
        fixtures[
            "batch_forms_fixed_layout_placeholder_manual_boundary"
        ].report_expectations
    )
    assert fixtures[
        "official_policy_formal_archive_manual_boundary"
    ].family_id == "meeting_policy_documents"
    assert fixtures[
        "official_policy_formal_archive_manual_boundary"
    ].request_sample_ids == ("official_notice_formal_archive",)
    assert "manual_confirmation" in (
        fixtures[
            "official_policy_formal_archive_manual_boundary"
        ].expected_behaviors
    )
    assert "watermark_status_report" in (
        fixtures[
            "official_policy_formal_archive_manual_boundary"
        ].report_expectations
    )
    assert fixtures[
        "official_policy_metadata_archive_report"
    ].family_id == "meeting_policy_documents"
    assert fixtures[
        "official_policy_metadata_archive_report"
    ].request_sample_ids == ("official_policy_collection",)
    assert fixtures[
        "official_policy_metadata_archive_report"
    ].expected_behaviors == ("detect",)
    assert "headers_footers" in (
        fixtures["official_policy_metadata_archive_report"].docx_surfaces
    )
    assert "package_relationships" in (
        fixtures["official_policy_metadata_archive_report"].docx_surfaces
    )
    assert "official_metadata_report" in (
        fixtures["official_policy_metadata_archive_report"].report_expectations
    )
    assert "formal_internal_archive_manifest" in (
        fixtures["official_policy_metadata_archive_report"].report_expectations
    )
    assert fixtures[
        "technical_long_docs_chapter_inventory_manual_boundary"
    ].family_id == "long_document_publishing"
    assert fixtures[
        "technical_long_docs_chapter_inventory_manual_boundary"
    ].request_sample_ids == ("technical_sop_chapter_inventory",)
    assert "manual_confirmation" in (
        fixtures[
            "technical_long_docs_chapter_inventory_manual_boundary"
        ].expected_behaviors
    )
    assert "cross_reference_status_report" in (
        fixtures[
            "technical_long_docs_chapter_inventory_manual_boundary"
        ].report_expectations
    )
    assert fixtures[
        "technical_long_docs_index_appendix_merge_boundary"
    ].family_id == "long_document_publishing"
    assert fixtures[
        "technical_long_docs_index_appendix_merge_boundary"
    ].request_sample_ids == ("technical_product_manual",)
    assert "manual_confirmation" in (
        fixtures[
            "technical_long_docs_index_appendix_merge_boundary"
        ].expected_behaviors
    )
    assert "headers_footers" in (
        fixtures["technical_long_docs_index_appendix_merge_boundary"].docx_surfaces
    )
    assert "package_relationships" in (
        fixtures["technical_long_docs_index_appendix_merge_boundary"].docx_surfaces
    )
    assert "index_appendix_inventory" in (
        fixtures[
            "technical_long_docs_index_appendix_merge_boundary"
        ].report_expectations
    )
    assert "multi_file_merge_boundary_report" in (
        fixtures[
            "technical_long_docs_index_appendix_merge_boundary"
        ].report_expectations
    )
    assert fixtures["bidding_materials_missing_attachment_degraded"].family_id == (
        "qualification_archive_packages"
    )
    assert fixtures[
        "bidding_materials_missing_attachment_degraded"
    ].request_sample_ids == ("bidding_license_archive",)
    assert "skip_report" in (
        fixtures["bidding_materials_missing_attachment_degraded"].expected_behaviors
    )
    assert "missing_items_report" in (
        fixtures["bidding_materials_missing_attachment_degraded"].report_expectations
    )
    assert fixtures[
        "bidding_materials_original_copy_manual_boundary"
    ].request_sample_ids == ("bidding_original_copy",)
    assert "manual_confirmation" in (
        fixtures[
            "bidding_materials_original_copy_manual_boundary"
        ].expected_behaviors
    )
    assert "qualification_authenticity_boundary_report" in (
        fixtures[
            "bidding_materials_original_copy_manual_boundary"
        ].report_expectations
    )
    assert fixtures[
        "bidding_materials_consortium_seal_residue_degraded"
    ].request_sample_ids == ("bidding_consortium_seal_archive",)
    assert "skip_report" in (
        fixtures[
            "bidding_materials_consortium_seal_residue_degraded"
        ].expected_behaviors
    )
    assert "manual_confirmation" in (
        fixtures[
            "bidding_materials_consortium_seal_residue_degraded"
        ].expected_behaviors
    )
    assert "expiry_metadata_report" in (
        fixtures[
            "bidding_materials_consortium_seal_residue_degraded"
        ].report_expectations
    )
    assert "consortium_archive_manifest" in (
        fixtures[
            "bidding_materials_consortium_seal_residue_degraded"
        ].report_expectations
    )
    assert "seal_position_residue_report" in (
        fixtures[
            "bidding_materials_consortium_seal_residue_degraded"
        ].report_expectations
    )
    assert fixtures[
        "application_reports_project_attachment_degraded"
    ].request_sample_ids == ("application_project_limits",)
    assert "skip_report" in (
        fixtures[
            "application_reports_project_attachment_degraded"
        ].expected_behaviors
    )
    assert "submission_system_boundary_report" in (
        fixtures[
            "application_reports_project_attachment_degraded"
        ].report_expectations
    )
    assert fixtures[
        "application_reports_project_budget_manual_boundary"
    ].request_sample_ids == ("application_review_budget",)
    assert "manual_confirmation" in (
        fixtures[
            "application_reports_project_budget_manual_boundary"
        ].expected_behaviors
    )
    assert "rule_source_governance" in (
        fixtures[
            "application_reports_project_budget_manual_boundary"
        ].report_expectations
    )
    assert "submission_system_boundary_report" in (
        fixtures[
            "application_reports_project_budget_manual_boundary"
        ].report_expectations
    )
    assert "asset_consistency_report" in (
        fixtures["application_reports_product_sales_assets"].report_expectations
    )
    assert "quote_body_disambiguation" in (
        fixtures["application_reports_product_sales_assets"].report_expectations
    )
    assert fixtures[
        "application_reports_product_asset_degraded"
    ].request_sample_ids == ("application_customer_product_manual",)
    assert "asset_consistency_report" in (
        fixtures["application_reports_product_asset_degraded"].report_expectations
    )
    assert "quote_body_disambiguation" in (
        fixtures["application_reports_product_asset_degraded"].report_expectations
    )
    assert fixtures[
        "application_reports_product_version_manual_boundary"
    ].request_sample_ids == ("application_presales_plan",)
    assert "customer_internal_version_report" in (
        fixtures[
            "application_reports_product_version_manual_boundary"
        ].report_expectations
    )
    assert "asset_consistency_report" in (
        fixtures[
            "application_reports_product_version_manual_boundary"
        ].report_expectations
    )
    assert "quote_body_disambiguation" in (
        fixtures[
            "application_reports_product_version_manual_boundary"
        ].report_expectations
    )
    assert fixtures[
        "professional_disclosure_source_quality_degraded"
    ].request_sample_ids == ("professional_finance_quote",)
    assert fixtures[
        "professional_disclosure_source_quality_degraded"
    ].manual_gate_id == "professional_disclosure_review_gate"
    assert "finance_spreadsheet_mapping_report" in (
        fixtures[
            "professional_disclosure_finance_table_mapping_boundary"
        ].report_expectations
    )
    assert fixtures[
        "professional_disclosure_finance_table_mapping_boundary"
    ].request_sample_ids == ("professional_finance_quote",)
    assert "claim_quality_boundary_ui" in (
        fixtures[
            "professional_disclosure_patent_claim_quality_boundary"
        ].report_expectations
    )
    assert fixtures[
        "professional_disclosure_patent_claim_quality_boundary"
    ].request_sample_ids == ("professional_patent_claims",)
    assert "termbase_ui" in (
        fixtures[
            "professional_disclosure_bilingual_termbase_boundary"
        ].report_expectations
    )
    assert fixtures[
        "professional_disclosure_bilingual_termbase_boundary"
    ].request_sample_ids == ("professional_bilingual_terms",)
    assert "regulated_rule_source_governance" in (
        fixtures[
            "professional_disclosure_regulated_assurance_boundary"
        ].report_expectations
    )
    assert "assurance_boundary_ui" in (
        fixtures[
            "professional_disclosure_regulated_assurance_boundary"
        ].report_expectations
    )
    assert fixtures[
        "professional_disclosure_regulated_assurance_boundary"
    ].request_sample_ids == ("professional_esg_archive",)
    assert fixtures[
        "import_ai_boundary_conversion_confidence_degraded"
    ].request_sample_ids == ("import_ocr_pdf",)
    assert fixtures[
        "import_ai_boundary_conversion_confidence_degraded"
    ].manual_gate_id == "import_ai_conversion_gate"
    assert "confidence_artifact_manifest" in (
        fixtures["import_ai_boundary_latex_handoff_confidence"].report_expectations
    )
    assert "ocr_pdf_latex_plugin_boundary" in (
        fixtures["import_ai_boundary_latex_handoff_confidence"].report_expectations
    )
    assert fixtures["import_ai_boundary_latex_handoff_confidence"].request_sample_ids == (
        "import_latex_project",
        "import_ai_diagrams",
    )
    for pack_id in REQUIRED_SAMPLE_PACK_IDS:
        assert scene_sample_fixtures_for_pack(pack_id)
    for surface in REQUIRED_SAMPLE_SURFACES:
        assert coverage[surface], f"missing sample coverage for {surface}"

    summary = build_scene_sample_fixture_summary(
        get_scene_sample_fixture("batch_forms_fixed_layout")
    )
    assert "fixed_row_height" in summary
    assert "preserve_layout" in summary


def test_scene_sample_fixture_summary_is_pack_level_gate_for_ui_surfaces():
    contract = build_scene_sample_coverage_summary(["contract_delivery"])
    report = build_scene_sample_coverage_summary(["quick_formatting", "application_reports"])

    assert contract.is_clean is True
    assert contract.fixture_ids == (
        "contract_delivery_revisions",
        "contract_delivery_signature_fields_degraded",
    )
    assert "tracked_changes" in contract.docx_surfaces
    assert "hidden_text" in contract.docx_surfaces
    assert contract.boundary_notes == (
        "Contract samples do not imply legal review.",
        (
            "Missing or inconsistent contract party, amount, date, or signature "
            "assets should downgrade to a field/signature boundary report."
        ),
    )

    assert report.is_clean is True
    assert report.fixture_ids == (
        "quick_formatting_fields_comments",
        "quick_formatting_preserve_fields_degraded",
        "quick_formatting_business_template_cleanup",
        "application_reports_hidden_comments",
        "application_reports_project_attachment_degraded",
        "application_reports_project_budget_manual_boundary",
        "application_reports_product_sales_assets",
        "application_reports_product_asset_degraded",
        "application_reports_product_version_manual_boundary",
    )
    assert "comments" in report.docx_surfaces
    assert "hidden_text" in report.docx_surfaces


def test_scene_request_cell_fixture_registry_covers_high_frequency_samples():
    cells = {cell.sample_id: cell for cell in list_scene_request_cell_fixtures()}
    request_samples = {sample.sample_id: sample for sample in list_high_frequency_request_samples()}
    summary = build_scene_request_cell_fixture_summary()

    assert audit_scene_request_cell_fixtures() == ()
    assert set(cells) == set(request_samples)
    assert summary.is_clean is True
    assert summary.cell_count == len(request_samples)
    assert summary.fixture_cell_count == len(request_samples) - 3
    assert summary.negative_control_count == 3
    assert summary.family_proxy_count == 0
    assert summary.manual_boundary_count == 12
    assert summary.ambiguous_count == 7
    assert set(summary.pack_ids) == set(REQUIRED_SAMPLE_PACK_IDS)
    assert summary.family_proxy_sample_ids == ()

    legal = cells["professional_legal_opinion"]
    unknown = cells["unknown_stays_unmatched"]
    ppt = cells["negative_ppt_poster_design"]
    product_manual = cells["application_customer_product_manual"]
    hr_offer = cells["batch_hr_offer"]
    project_form = cells["project_application_form"]
    batch_notice = cells["ambiguous_batch_notice"]
    exam_versions = cells["exam_multi_version"]
    exam_sheet = cells["exam_answer_sheet"]
    exam_analysis = cells["exam_answer_analysis"]
    quick_cleanup = cells["quick_word_cleanup"]
    personal_resume = cells["personal_resume_formatting"]
    journal_submission = cells["english_journal_submission_package"]
    journal_response = cells["english_journal_response_letter"]
    contract_signing = cells["contract_signing_consistency"]
    contract_legal = cells["ambiguous_contract_legal_review"]
    official_archive = cells["official_notice_formal_archive"]
    technical_inventory = cells["technical_sop_chapter_inventory"]

    assert legal.coverage_level == "manual_boundary_fixture"
    assert legal.manual_gate_ids == ("professional_disclosure_review_gate",)
    assert unknown.coverage_level == "negative_control"
    assert unknown.fixture_ids == ()
    assert ppt.coverage_level == "negative_control"
    assert ppt.fixture_ids == ()
    assert product_manual.coverage_level == "direct_family_fixture"
    assert "application_reports_product_sales_assets" in product_manual.fixture_ids
    assert product_manual.family_proxy_gap_ids == ()
    assert hr_offer.coverage_level == "direct_family_fixture"
    assert "batch_forms_hr_batch_offer" in hr_offer.fixture_ids
    assert project_form.coverage_level == "ambiguous_fixture_set"
    assert set(project_form.expected_pack_ids) == {"application_reports", "batch_forms"}
    assert batch_notice.coverage_level == "ambiguous_fixture_set"
    assert set(batch_notice.expected_pack_ids) == {"batch_forms", "official_policy"}
    assert exam_versions.coverage_level == "manual_boundary_fixture"
    assert exam_versions.fixture_ids == ("exam_education_controls_textboxes",)
    assert exam_versions.manual_gate_ids == ("exam_ai_complex_diagram_gate",)
    assert exam_sheet.coverage_level == "direct_family_fixture"
    assert exam_sheet.fixture_ids == ("exam_education_structured_multiversion",)
    assert exam_sheet.manual_gate_ids == ()
    assert exam_analysis.coverage_level == "direct_family_fixture"
    assert exam_analysis.fixture_ids == ("exam_education_structured_multiversion",)
    assert exam_analysis.manual_gate_ids == ()
    assert quick_cleanup.coverage_level == "direct_family_fixture"
    assert quick_cleanup.fixture_ids == (
        "quick_formatting_fields_comments",
        "quick_formatting_business_template_cleanup",
    )
    assert quick_cleanup.manual_gate_ids == ()
    assert personal_resume.coverage_level == "direct_family_fixture"
    assert personal_resume.fixture_ids == (
        "quick_formatting_fields_comments",
        "quick_formatting_business_template_cleanup",
    )
    assert personal_resume.manual_gate_ids == ()
    assert journal_submission.coverage_level == "direct_family_fixture"
    assert journal_submission.fixture_ids == (
        "english_journal_revision_comments",
        "english_journal_bibtex_csl_degraded",
    )
    assert journal_response.coverage_level == "manual_boundary_fixture"
    assert journal_response.fixture_ids == (
        "english_journal_revision_comments",
        "english_journal_publisher_rule_manual_boundary",
    )
    assert journal_response.manual_gate_ids == (
        "journal_publisher_rule_review_gate",
    )
    assert contract_signing.coverage_level == "direct_family_fixture"
    assert contract_signing.fixture_ids == (
        "contract_delivery_revisions",
        "contract_delivery_signature_fields_degraded",
    )
    assert contract_legal.coverage_level == "ambiguous_fixture_set"
    assert contract_legal.expected_pack_ids == (
        "contract_delivery",
        "professional_disclosure",
    )
    assert contract_legal.manual_gate_ids == (
        "professional_disclosure_review_gate",
    )
    assert official_archive.coverage_level == "direct_family_fixture"
    assert official_archive.fixture_ids == (
        "official_policy_fields_comments",
        "official_policy_formal_archive_manual_boundary",
    )
    assert official_archive.manual_gate_ids == ()
    assert "formal/internal" in official_archive.boundary_notes[0]
    assert technical_inventory.coverage_level == "direct_family_fixture"
    assert technical_inventory.fixture_ids == (
        "technical_long_docs_skip_objects",
        "technical_long_docs_chapter_inventory_manual_boundary",
    )
    assert technical_inventory.manual_gate_ids == ()
    assert "Chapter inventory" in technical_inventory.boundary_notes[1]

    professional_cells = request_cell_fixtures_for_pack("professional_disclosure")
    assert {cell.sample_id for cell in professional_cells} >= {
        "professional_legal_opinion",
        "ambiguous_quote_plan",
    }


def test_scene_sample_docx_fixtures_are_real_packages_and_preflight_detects_surfaces(tmp_path):
    for fixture in list_scene_sample_fixtures():
        source = build_scene_sample_docx(tmp_path, fixture)

        assert is_zipfile(source)
        result = inspect_docx_package(source)
        finding_kinds = {finding.kind for finding in result.findings}

        assert set(fixture.expected_preflight_findings) <= finding_kinds
        if "fixed_row_height" in fixture.docx_surfaces:
            doc = Document(source)
            state = row_height_state(doc.tables[0].rows[0])
            assert state.height_twips == 400
            assert state.rule == "atLeast"


def test_scene_sample_docx_library_generator_writes_openable_manifest(tmp_path):
    library = build_scene_sample_docx_library(tmp_path / "scene_samples")
    payload = json.loads(Path(library.manifest_path).read_text(encoding="utf-8"))
    request_cells = {item["sample_id"]: item for item in payload["request_cells"]}
    request_cell_report = Path(payload["request_cell_report_path"])

    assert payload["artifact_count"] == len(list_scene_sample_fixtures())
    assert len(payload["artifacts"]) == len(list_scene_sample_fixtures())
    assert payload["request_cell_count"] == len(list_high_frequency_request_samples())
    assert len(payload["request_cells"]) == len(list_high_frequency_request_samples())
    assert request_cell_report.exists()
    assert '<a id="request-cell-contract-signing-consistency"></a>' in (
        request_cell_report.read_text(encoding="utf-8")
    )
    assert payload["request_cell_summary"]["family_proxy_count"] == 0
    assert payload["request_cell_summary"]["negative_control_count"] == 3
    assert request_cells["application_customer_product_manual"]["coverage_level"] == (
        "direct_family_fixture"
    )
    assert request_cells["application_customer_product_manual"]["report_path"] == str(
        request_cell_report
    )
    assert request_cells["application_customer_product_manual"]["report_anchor"] == (
        "request-cell-application-customer-product-manual"
    )
    assert request_cells["application_customer_product_manual"]["family_proxy_gap_ids"] == []
    assert request_cells["batch_hr_offer"]["coverage_level"] == "direct_family_fixture"
    assert request_cells["professional_medical_regulatory"]["coverage_level"] == (
        "manual_boundary_fixture"
    )
    assert request_cells["unknown_stays_unmatched"]["fixture_ids"] == []
    assert request_cells["negative_ppt_poster_design"]["fixture_ids"] == []
    assert request_cells["project_application_form"]["coverage_level"] == (
        "ambiguous_fixture_set"
    )
    assert request_cells["ambiguous_batch_notice"]["coverage_level"] == (
        "ambiguous_fixture_set"
    )
    assert {item["pack_id"] for item in payload["artifacts"]} == set(
        REQUIRED_SAMPLE_PACK_IDS
    )
    for artifact in library.artifacts:
        source = Path(artifact.path)
        assert source.exists()
        assert source.suffix == ".docx"
        assert is_zipfile(source)
    batch = next(
        artifact
        for artifact in library.artifacts
        if artifact.fixture_id == "batch_forms_fixed_layout"
    )
    doc = Document(batch.path)
    assert row_height_state(doc.tables[0].rows[0]).height_twips == 400
    assert scene_sample_fixture_manifest_paths(tmp_path / "scene_samples") == {
        "fixture_manifest": library.manifest_path
    }


def test_scene_sample_fixture_generation_script_writes_openable_library(tmp_path):
    output_dir = tmp_path / "cli_scene_samples"
    result = subprocess.run(
        [
            sys.executable,
            "scripts/generate_scene_sample_fixtures.py",
            str(output_dir),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    payload = json.loads(result.stdout)

    assert payload["artifact_count"] == len(list_scene_sample_fixtures())
    assert payload["request_cell_count"] == len(list_high_frequency_request_samples())
    assert payload["request_cell_summary"]["fixture_cell_count"] == (
        len(list_high_frequency_request_samples()) - 3
    )
    assert Path(payload["manifest_path"]).exists()
    assert Path(payload["request_cell_report_path"]).exists()
    assert len(list(output_dir.glob("*.docx"))) == len(list_scene_sample_fixtures())


def test_scene_matrix_release_gate_script_covers_samples_readiness_and_cells(tmp_path):
    output_dir = tmp_path / "release_gate_samples"
    result = subprocess.run(
        [
            sys.executable,
            "scripts/verify_scene_matrix_release_gate.py",
            "--json",
            "--output-dir",
            str(output_dir),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    payload = json.loads(result.stdout)

    assert payload["status"] == "passed"
    assert payload["failed_check_ids"] == []
    assert payload["counts"]["coverage_pack_count"] == len(REQUIRED_SAMPLE_PACK_IDS)
    assert payload["counts"]["sample_fixture_count"] == len(list_scene_sample_fixtures())
    assert payload["counts"]["high_frequency_request_sample_count"] == len(
        list_high_frequency_request_samples()
    )
    assert payload["counts"]["request_cell_count"] == len(
        list_high_frequency_request_samples()
    )
    assert payload["counts"]["request_cell_family_proxy_count"] == 0
    assert payload["counts"]["product_readiness_subject_count"] == 27
    assert payload["request_cell_summary"]["family_proxy_sample_ids"] == []
    assert payload["checks"]["scene_product_readiness"]["status"] == "passed"
    assert payload["checks"]["scene_request_cell_release_threshold"]["status"] == (
        "passed"
    )
    assert (output_dir / "manifest.json").exists()


def test_report_writer_emits_scene_sample_fixture_manifest_evidence(tmp_path, monkeypatch):
    library = build_scene_sample_docx_library(tmp_path / "scene_samples")
    report_json = tmp_path / "sample_manifest_report.json"
    report_md = tmp_path / "sample_manifest_report.md"
    result = PipelineResult(success=True)

    monkeypatch.setattr(
        "src.report_writer.scene_sample_fixture_manifest_paths",
        lambda: {"fixture_manifest": library.manifest_path},
    )

    write_json_report(
        result,
        input_path=tmp_path / "sample.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "sample.docx",
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )

    payload = json.loads(report_json.read_text(encoding="utf-8"))
    evidence = payload["scene_sample_fixtures"]
    markdown = report_md.read_text(encoding="utf-8")

    assert evidence["status"] == "available"
    assert evidence["manifest_path"] == library.manifest_path
    assert evidence["artifact_count"] == len(list_scene_sample_fixtures())
    assert evidence["request_cell_count"] == len(list_high_frequency_request_samples())
    assert evidence["request_cell_report_path"] == library.request_cell_report_path
    assert evidence["request_cell_report_available"] is True
    assert {
        item["sample_id"]: item["anchor"]
        for item in evidence["request_cell_report_anchors"]
    }["contract_signing_consistency"] == "request-cell-contract-signing-consistency"
    assert evidence["request_cell_summary"]["family_proxy_count"] == 0
    assert evidence["request_cell_family_proxy_sample_ids"] == []
    assert "contract_delivery" in evidence["pack_ids"]
    assert "contract_delivery_revisions" in evidence["fixture_ids"]
    assert "application_reports_product_sales_assets" in evidence["fixture_ids"]
    assert "batch_forms_hr_batch_offer" in evidence["fixture_ids"]
    assert "manual_boundary_fixture" in evidence["request_cell_coverage_levels"]
    assert "tracked_changes" in evidence["docx_surfaces"]
    assert "manual_confirmation" in evidence["expected_behaviors"]
    assert "professional_disclosure_review_gate" in evidence["manual_gate_ids"]
    assert "Contract samples do not imply legal review." in evidence["boundary_notes"]
    assert "## 样本库证据" in markdown
    assert f"- Manifest: `{library.manifest_path}`" in markdown
    assert f"- Request-cell report: `{library.request_cell_report_path}` (yes)" in (
        markdown
    )
    assert f"- Artifacts: {len(list_scene_sample_fixtures())}" in markdown
    assert "- Request cells: 53" in markdown
    assert "quick_template_cleanup#request-cell-quick-template-cleanup" in markdown
    assert "family_proxy=0" in markdown
    assert "contract_delivery_revisions" in markdown


def test_report_writer_omits_scene_sample_fixture_manifest_when_missing(
    tmp_path,
    monkeypatch,
):
    report_json = tmp_path / "missing_sample_manifest_report.json"
    report_md = tmp_path / "missing_sample_manifest_report.md"
    result = PipelineResult(success=True)

    monkeypatch.setattr(
        "src.report_writer.scene_sample_fixture_manifest_paths",
        lambda: {},
    )

    write_json_report(
        result,
        input_path=tmp_path / "sample.docx",
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )
    write_markdown_report(
        result,
        input_path=tmp_path / "sample.docx",
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
        include_internal_evidence=True,
    )

    payload = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert payload["scene_sample_fixtures"] is None
    assert "## 样本库证据" not in markdown


def test_scene_sample_skip_report_uses_real_docx_object_preflight(tmp_path):
    fixture = get_scene_sample_fixture("technical_long_docs_skip_objects")
    source = build_scene_sample_docx(tmp_path, fixture)
    scene = SceneWorkspace(scene_id="technical", category="technical")
    scene.compliance_profile.object_preflight.preservation_mode = "warn"
    scene.compliance_profile.object_preflight.block_on = []
    scene.compliance_profile.object_preflight.skip_high_risk_modules = True
    scene.compliance_profile.object_preflight.skip_modules_by_finding = {
        "ole_objects": ["risky_rewriter"],
        "embedded_workbooks": ["risky_rewriter"],
        "visio_drawings": ["risky_rewriter"],
    }

    result = Pipeline(
        modules=[_RiskyRewriteModule()],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "out"),
    ).execute(str(source))

    report_json = tmp_path / "skip_report.json"
    report_md = tmp_path / "skip_report.md"
    write_json_report(
        result,
        input_path=source,
        output_path=Path(result.output_paths["final"]),
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=1,
        modules_total=1,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=1,
        modules_total=1,
    )

    payload = json.loads(report_json.read_text(encoding="utf-8"))
    evidence = payload["object_preflight"]
    markdown = report_md.read_text(encoding="utf-8")

    assert result.success is True
    assert {finding["kind"] for finding in evidence["findings"]} >= {
        "ole_objects",
        "embedded_workbooks",
        "visio_drawings",
    }
    assert evidence["module_skips_count"] == 1
    assert evidence["module_skips"][0]["module_name"] == "risky_rewriter"
    assert "risky_rewriter: embedded_workbooks, ole_objects, visio_drawings" in markdown


def test_scene_sample_block_and_manual_confirmation_report_use_real_docx(tmp_path):
    fixture = get_scene_sample_fixture("import_ai_boundary_blocking_macro")
    source = build_scene_sample_docx(tmp_path, fixture)
    scene = SceneWorkspace(scene_id="import_ai_boundary", category="import_ai_boundary")
    scene.compliance_profile.rule_family = "import_ai_boundary"
    scene.compliance_profile.object_preflight.preservation_mode = "strict"
    scene.compliance_profile.object_preflight.block_on = ["macros"]

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "blocked_out"),
    ).execute(str(source))

    report_json = tmp_path / "blocked_report.json"
    report_md = tmp_path / "blocked_report.md"
    write_json_report(
        result,
        input_path=source,
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )

    payload = json.loads(report_json.read_text(encoding="utf-8"))
    findings = payload["object_preflight"]["findings"]
    gate = payload["coverage_boundaries"][0]["plugin_manual_gate"]
    markdown = report_md.read_text(encoding="utf-8")

    assert result.success is False
    assert any(
        finding["kind"] == "macros" and finding["severity"] == "error"
        for finding in findings
    )
    assert gate["gate_id"] == "import_ai_conversion_gate"
    assert gate["manual_confirmation_required"] is True
    assert "Manual confirmation: required" in markdown
    assert "VBA macro project is present" in markdown


def test_scene_sample_fixed_layout_preserves_existing_row_height(tmp_path):
    fixture = get_scene_sample_fixture("batch_forms_fixed_layout")
    source = build_scene_sample_docx(tmp_path, fixture)
    doc = Document(source)
    policy = FixedLayoutRowHeightPolicy(
        policy_id="form_batch_documents.preserve_sample",
        family_id="form_batch_documents",
        label="Preserve fixed row height sample",
        mode="preserve_existing",
        row_height_pt=None,
        parameter_path="form_batch_documents.table.row_height_pt",
    )

    result = apply_fixed_layout_row_height_policy(doc.tables[0], policy)

    assert result.preserved_count == 1
    assert result.changed_count == 0
    assert row_height_state(doc.tables[0].rows[0]).height_twips == 400
    assert "preserve_layout" in fixture.expected_behaviors

class _RiskyRewriteModule(BaseModule):
    meta = ModuleMeta(
        name="risky_rewriter",
        description="Risky package rewriter",
        category="test",
        execution_phase="format",
        scope_behavior="document_level",
    )

    def apply(self, doc, config, tracker: ChangeTracker, context: PipelineContext) -> None:
        tracker.record(
            rule_name=self.meta.name,
            target="document",
            section="global",
            change_type="format",
            after="module ran",
        )
