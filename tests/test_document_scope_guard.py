from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.config.document_scope import DocumentScopePolicy
from src.modules.structure.heading_recognition import rebuild_document_index
from src.pipeline.context import PipelineContext
from src.services.document_structure_evidence import build_document_structure_evidence
from src.shared.engine.document_scope_guard import (
    capture_document_scope_guard,
    validate_document_scope_guard,
)
from src.shared.engine.document_scope_runtime import bind_document_scope


def _cover_document(path):
    doc = Document()
    title = doc.add_paragraph("\u901a\u7528\u5c01\u9762\u6807\u9898")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _index in range(8):
        doc.add_paragraph("")
    authority = doc.add_paragraph("\u67d0\u5e02\u4eba\u6c11\u653f\u5e9c")
    authority.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date = doc.add_paragraph("\u4e8c\u25cb\u4e8c\u516d\u5e74\u4e8c\u6708")
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("\u7b2c\u4e00\u7ae0 \u6b63\u6587", level=1)
    for index in range(30):
        doc.add_paragraph(f"\u6b63\u6587 {index}\u3002")
    doc.save(path)


def _bound_body_context(path, doc):
    evidence = build_document_structure_evidence(path)
    policy = DocumentScopePolicy(mode="body")
    context = PipelineContext(
        mode_id="custom",
        document_scope=policy,
        document_scope_gate_active=True,
    )
    context.document_scope_binding = bind_document_scope(
        doc,
        evidence,
        (),
        policy,
        mode_id="custom",
    )
    rebuild_document_index(doc, context)
    return context


def test_guard_allows_body_formatting_but_blocks_cover_mutation(tmp_path):
    source = tmp_path / "cover.docx"
    _cover_document(source)
    doc = Document(source)
    context = _bound_body_context(source, doc)
    guard = capture_document_scope_guard(doc, context)
    body = context.doc_tree.get_section("body")

    doc.paragraphs[body.start_index].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    assert validate_document_scope_guard(guard) == ()

    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    violations = validate_document_scope_guard(guard)
    assert any(item.startswith("paragraph_changed:cover@0") for item in violations)


def test_guard_blocks_shared_style_changes_that_affect_cover(tmp_path):
    source = tmp_path / "cover-style.docx"
    _cover_document(source)
    doc = Document(source)
    context = _bound_body_context(source, doc)
    guard = capture_document_scope_guard(doc, context)

    doc.styles["Normal"].font.bold = True

    assert any(
        item.startswith("style_changed:")
        for item in validate_document_scope_guard(guard)
    )
