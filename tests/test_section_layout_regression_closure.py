from __future__ import annotations

import json
from dataclasses import asdict
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.shared import Cm

from src.config.builtin_templates import create_builtin_template
from src.config.library import _load_template_entry_config
from src.config.loader import (
    ConfigLoadError,
    load_compatible_user_template,
    load_template,
)
from src.config.resolved import ResolvedConfig
from src.config.template import SectionMarginConfig, TemplateConfig
from src.modules.basic.header_footer import HeaderFooterModule
from src.modules.basic.page_setup import (
    PageSetupModule,
    _section_orientation_contradiction,
)
from src.modules.basic.section_format import SectionFormatModule
from src.modules.structure.heading_recognition import DocSection
from src.modules.validate.validation import ValidationModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.reporting.section_layout import extract_section_layout_evidence
from src.shared.engine.field_builder import iter_field_instructions
from src.shared.engine.ooxml_ops import qn
from src.shared.engine.section_layout_planner import (
    build_section_execution_plan,
    collect_section_inventory,
    execute_section_execution_plan,
)


def _tree(*sections):
    return SimpleNamespace(
        sections=list(sections),
        get_section=lambda name: next(
            (section for section in sections if section.section_type == name),
            None,
        ),
        get_section_for_paragraph=lambda index: next(
            (
                section.section_type
                for section in sections
                if section.start_index <= index < section.end_index
            ),
            "body",
        ),
    )


def _cover_body_context() -> PipelineContext:
    return PipelineContext(
        doc_tree=_tree(
            DocSection("cover", 0, 1, confidence=10),
            DocSection("body", 1, 2, confidence=10),
        )
    )


def _field_instructions(paragraph) -> list[str]:
    return [
        " ".join((instruction or "").split())
        for _kind, _element, instruction in iter_field_instructions(
            paragraph._element
        )
    ]


def test_inserted_semantic_boundary_never_aliases_the_existing_header_part():
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_paragraph("Body")
    doc.sections[0].header.paragraphs[0].text = "body source header"
    config = ResolvedConfig()
    config.section.boundary_mode = "semantic_rebuild"
    config.section.header_footer_link_mode = "preserve_source"
    context = _cover_body_context()

    SectionFormatModule().apply(doc, config, ChangeTracker(), context)

    inventory = collect_section_inventory(doc)
    assert inventory.section_count == 2
    assert not any(ref[0] == "header" for ref in inventory.boundaries[0].header_footer_refs)
    assert any(ref[0] == "header" for ref in inventory.boundaries[1].header_footer_refs)

    HeaderFooterModule().apply(doc, config, ChangeTracker(), context)
    cover, body = list(doc.sections)
    assert cover.header.part is not body.header.part
    assert cover.header.paragraphs[0].text == ""
    assert _field_instructions(cover.header.paragraphs[0]) == []
    assert _field_instructions(body.header.paragraphs[0]) == ["STYLEREF 1"]
    assert context.final_section_inventory.digest == collect_section_inventory(doc).digest


def test_semantic_rebuild_accepts_existing_odd_page_boundary():
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_paragraph("Body")
    config = ResolvedConfig()
    config.section.boundary_mode = "normalize_all"
    config.section.section_break_type = "oddPage"
    context = _cover_body_context()
    SectionFormatModule().apply(doc, config, ChangeTracker(), context)

    first = collect_section_inventory(doc)
    assert first.boundaries[0].break_type == "oddPage"

    config.section.boundary_mode = "semantic_rebuild"
    config.section.section_break_type = None
    SectionFormatModule().apply(doc, config, ChangeTracker(), context)
    second = collect_section_inventory(doc)

    assert second.boundaries[0].break_type == "oddPage"
    assert context.section_execution_receipt.applied_operations == ()


def test_tracked_section_properties_block_type_mutation_without_partial_write():
    doc = Document()
    sect_pr = doc.element.body.sectPr
    sect_pr.append(OxmlElement("w:sectPrChange"))
    before_signature = collect_section_inventory(doc).boundaries[0].full_signature
    config = ResolvedConfig()
    config.section.boundary_mode = "normalize_all"
    config.section.section_break_type = "continuous"

    plan = build_section_execution_plan(doc, config, PipelineContext())
    receipt = execute_section_execution_plan(doc, plan)

    assert plan.blocked_operations
    assert "sectPrChange" in plan.blocked_operations[0].blocked_reason
    assert receipt.applied_operations == ()
    assert receipt.validation_errors
    assert collect_section_inventory(doc).boundaries[0].full_signature == before_signature


def test_nested_content_control_section_is_inventoried_and_mutation_is_blocked():
    doc = Document()
    paragraph = doc.add_paragraph("nested boundary")
    body = doc.element.body
    body.remove(paragraph._element)
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    sdt.append(content)
    content.append(paragraph._element)
    body.insert(0, sdt)
    ppr = OxmlElement("w:pPr")
    paragraph._element.insert(0, ppr)
    nested_sect_pr = OxmlElement("w:sectPr")
    break_type = OxmlElement("w:type")
    break_type.set(qn("w:val"), "continuous")
    nested_sect_pr.append(break_type)
    ppr.append(nested_sect_pr)

    inventory = collect_section_inventory(doc)
    assert inventory.section_count == 2
    assert len(doc.sections) == 1
    assert inventory.boundaries[0].anchor_kind == "nested_paragraph"
    assert "sdt" in inventory.boundaries[0].protected_markup

    config = ResolvedConfig()
    config.section.boundary_mode = "normalize_all"
    config.section.section_break_type = "nextPage"
    plan = build_section_execution_plan(doc, config, PipelineContext())
    receipt = execute_section_execution_plan(doc, plan)

    assert any("nested" in item.blocked_reason for item in plan.blocked_operations)
    assert receipt.applied_operations == ()
    assert collect_section_inventory(doc).boundaries[0].break_type == "continuous"
    header_issues = HeaderFooterModule().validate(doc, config, PipelineContext())
    assert any(issue.location == "document.sections.protected" for issue in header_issues)
    with pytest.raises(RuntimeError, match="protected section properties"):
        HeaderFooterModule().apply(doc, config, ChangeTracker(), PipelineContext())
    page_issues = PageSetupModule().validate(doc, config, PipelineContext())
    assert any("受保护分节" in issue.message for issue in page_issues)
    nested_page_size = collect_section_inventory(doc).boundaries[0].page_size
    PageSetupModule().apply(doc, config, ChangeTracker(), PipelineContext())
    assert collect_section_inventory(doc).boundaries[0].page_size == nested_page_size


def test_non_table_body_content_still_blocks_semantic_carrier_recovery():
    doc = Document()
    doc.add_paragraph("Cover")
    target = doc.add_paragraph("Body")
    alt_chunk = OxmlElement("w:altChunk")
    alt_chunk.set(qn("r:id"), "rIdMissing")
    target._element.addprevious(alt_chunk)
    config = ResolvedConfig()
    config.section.boundary_mode = "semantic_rebuild"

    plan = build_section_execution_plan(doc, config, _cover_body_context())
    receipt = execute_section_execution_plan(doc, plan)

    assert plan.blocked_operations
    assert plan.blocked_operations[0].action == "insert_break"
    assert "altChunk" in plan.blocked_operations[0].blocked_reason
    assert receipt.applied_operations == ()
    assert collect_section_inventory(doc).section_count == 1


def test_page_setup_repairs_explicit_orientation_geometry_contradiction():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section._sectPr.find(qn("w:pgSz")).set(qn("w:orient"), "portrait")
    config = ResolvedConfig()
    config.page_setup.paper_size_mode = "force_template"
    config.page_setup.orientation_mode = "preserve_source"
    context = PipelineContext()

    assert _section_orientation_contradiction(section)
    assert any(
        issue.level == "warning"
        for issue in PageSetupModule().validate(doc, config, context)
    )
    PageSetupModule().apply(doc, config, ChangeTracker(), context)

    assert not _section_orientation_contradiction(section)
    assert section.orientation == WD_ORIENT.PORTRAIT
    assert section.page_width < section.page_height


def test_per_section_orientation_is_applied_and_finally_verified():
    doc = Document()
    doc.add_paragraph("first")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("second")
    config = ResolvedConfig()
    config.page_setup.orientation_mode = "per_section"
    config.page_setup.orientation_by_section = {"2": "landscape"}
    context = PipelineContext()

    assert PageSetupModule().validate(doc, config, context) == []
    PageSetupModule().apply(doc, config, ChangeTracker(), context)
    assert doc.sections[1].page_width > doc.sections[1].page_height
    assert not any(
        "页面方向不符" in issue.message
        for issue in ValidationModule().validate(doc, config, context)
    )

    doc.sections[1].page_width = Cm(21)
    doc.sections[1].page_height = Cm(29.7)
    doc.sections[1].orientation = WD_ORIENT.PORTRAIT
    issues = ValidationModule().validate(doc, config, context)
    assert any("页面方向不符" in issue.message for issue in issues)


def test_per_section_orientation_rejects_empty_and_unknown_mapping():
    doc = Document()
    config = ResolvedConfig()
    config.page_setup.orientation_mode = "per_section"
    context = PipelineContext()

    issues = PageSetupModule().validate(doc, config, context)
    assert any("requires at least one" in issue.message for issue in issues)

    config.page_setup.orientation_by_section = {"99": "landscape"}
    issues = PageSetupModule().validate(doc, config, context)
    assert any("does not match" in issue.message for issue in issues)


def test_per_section_paper_and_complete_margins_are_applied_and_verified():
    doc = Document()
    doc.add_paragraph("first")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("second")
    first_source = collect_section_inventory(doc).boundaries[0]
    config = ResolvedConfig()
    config.page_setup.paper_size_mode = "per_section"
    config.page_setup.paper_size_by_section = {"2": "A3"}
    config.page_setup.orientation_mode = "per_section"
    config.page_setup.orientation_by_section = {"2": "landscape"}
    config.page_setup.margin_mode = "per_section"
    config.page_setup.margin_by_section = {
        "2": SectionMarginConfig(
            top_cm=2.0,
            bottom_cm=2.1,
            left_cm=2.2,
            right_cm=2.3,
            gutter_cm=0.4,
            header_distance_cm=0.8,
            footer_distance_cm=0.9,
        )
    }
    context = PipelineContext()

    assert PageSetupModule().validate(doc, config, context) == []
    PageSetupModule().apply(doc, config, ChangeTracker(), context)

    final = collect_section_inventory(doc)
    assert final.boundaries[0].page_size == first_source.page_size
    assert final.boundaries[0].page_margins == first_source.page_margins
    second = doc.sections[1]
    assert second.page_width > second.page_height
    assert second.page_width.cm == pytest.approx(42.0, abs=0.02)
    assert second.page_height.cm == pytest.approx(29.7, abs=0.02)
    assert second.top_margin.cm == pytest.approx(2.0, abs=0.02)
    assert second.bottom_margin.cm == pytest.approx(2.1, abs=0.02)
    assert second.left_margin.cm == pytest.approx(2.2, abs=0.02)
    assert second.right_margin.cm == pytest.approx(2.3, abs=0.02)
    assert second.gutter.cm == pytest.approx(0.4, abs=0.02)
    assert second.header_distance.cm == pytest.approx(0.8, abs=0.02)
    assert second.footer_distance.cm == pytest.approx(0.9, abs=0.02)
    assert not any(
        issue.level == "error"
        for issue in ValidationModule().validate(doc, config, context)
    )

    second.right_margin = Cm(1.0)
    assert any(
        "右边距" in issue.message
        for issue in ValidationModule().validate(doc, config, context)
    )


def test_per_section_layout_uses_ordinal_before_semantic_role():
    doc = Document()
    doc.add_paragraph("first")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("appendix")
    context = PipelineContext(
        doc_tree=SimpleNamespace(
            get_section_for_paragraph=lambda index: (
                "appendix" if index > 0 else "body"
            )
        )
    )
    config = ResolvedConfig()
    config.page_setup.paper_size_mode = "per_section"
    config.page_setup.paper_size_by_section = {
        "appendix": "B5",
        "2": "A3",
    }
    config.page_setup.orientation_mode = "per_section"
    config.page_setup.orientation_by_section = {"appendix": "landscape"}
    config.page_setup.margin_mode = "per_section"
    config.page_setup.margin_by_section = {
        "appendix": SectionMarginConfig(
            top_cm=1.5,
            bottom_cm=1.5,
            left_cm=1.5,
            right_cm=1.5,
            gutter_cm=0,
            header_distance_cm=0.7,
            footer_distance_cm=0.7,
        )
    }

    assert PageSetupModule().validate(doc, config, context) == []
    PageSetupModule().apply(doc, config, ChangeTracker(), context)

    second = doc.sections[1]
    assert second.page_width.cm == pytest.approx(42.0, abs=0.02)
    assert second.page_height.cm == pytest.approx(29.7, abs=0.02)
    assert second.top_margin.cm == pytest.approx(1.5, abs=0.02)


def test_per_section_paper_and_margin_reject_empty_unknown_and_invalid_maps():
    doc = Document()
    config = ResolvedConfig()
    config.page_setup.paper_size_mode = "per_section"
    config.page_setup.margin_mode = "per_section"

    issues = PageSetupModule().validate(doc, config, PipelineContext())
    assert any("per_section paper size requires" in issue.message for issue in issues)
    assert any("per_section margin requires" in issue.message for issue in issues)

    config.page_setup.paper_size_by_section = {"99": "TABLOID"}
    config.page_setup.margin_by_section = {
        "99": SectionMarginConfig(left_cm=-1)
    }
    issues = PageSetupModule().validate(doc, config, PipelineContext())
    assert any("unsupported per-section paper size" in issue.message for issue in issues)
    assert any("does not match" in issue.message for issue in issues)
    assert any("non-negative" in issue.message for issue in issues)


def test_final_validation_honors_disabled_page_setup_module():
    doc = Document()
    doc.sections[0].top_margin = Cm(0.1)
    config = ResolvedConfig(
        module_switches={"page_setup": False, "section_format": False}
    )

    issues = ValidationModule().validate(doc, config, PipelineContext())

    assert not any(
        any(token in issue.message for token in ("边距", "纸张", "页面方向"))
        for issue in issues
    )


def test_final_validation_detects_preserved_paper_and_margin_mutation():
    doc = Document()
    config = ResolvedConfig()
    config.page_setup.paper_size_mode = "preserve_source"
    config.page_setup.orientation_mode = "preserve_source"
    config.page_setup.margin_mode = "preserve_source"
    context = PipelineContext()
    PageSetupModule().apply(doc, config, ChangeTracker(), context)

    doc.sections[0].page_width = Cm(10)
    doc.sections[0].top_margin = Cm(1)
    issues = ValidationModule().validate(doc, config, context)
    messages = [issue.message for issue in issues]

    assert any("纸张规格改变" in message for message in messages)
    assert any("页边距或页眉页脚距离改变" in message for message in messages)


def test_final_validation_checks_every_forced_margin_and_distance():
    doc = Document()
    config = ResolvedConfig()
    context = PipelineContext()
    PageSetupModule().apply(doc, config, ChangeTracker(), context)
    section = doc.sections[0]
    section.bottom_margin = Cm(1)
    section.left_margin = Cm(1)
    section.right_margin = Cm(1)
    section.gutter = Cm(1)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)

    messages = [
        issue.message
        for issue in ValidationModule().validate(doc, config, context)
    ]
    for label in (
        "下边距",
        "左边距",
        "右边距",
        "装订线",
        "页眉距边界",
        "页脚距边界",
    ):
        assert any(label in message for message in messages)


def test_report_final_digest_is_captured_after_page_setup():
    doc = Document()
    config = ResolvedConfig()
    config.page_setup.paper_size_by_section = {"body": "A4"}
    config.page_setup.orientation_by_section = {"body": "portrait"}
    config.page_setup.margin_by_section = {
        "body": SectionMarginConfig(top_cm=2.5)
    }
    context = PipelineContext()
    SectionFormatModule().apply(doc, config, ChangeTracker(), context)
    section_stage_digest = context.section_execution_receipt.final_digest
    PageSetupModule().apply(doc, config, ChangeTracker(), context)
    result = SimpleNamespace(context=context, config=config)

    evidence = extract_section_layout_evidence(result)

    assert evidence is not None
    assert evidence["after_section_format_digest"] == section_stage_digest
    assert evidence["final_digest"] == context.page_setup_final_inventory.digest
    assert evidence["final_digest"] != section_stage_digest
    assert len(evidence["final_sections"]) == evidence["final_section_count"]
    assert evidence["paper_size_by_section"] == {"body": "A4"}
    assert evidence["orientation_by_section"] == {"body": "portrait"}
    assert evidence["margin_by_section"]["body"]["top_cm"] == 2.5


def test_user_library_compatibility_is_bounded_to_known_v02_policy_fields(tmp_path):
    payload = asdict(TemplateConfig())
    payload["name"] = "legacy user template"
    for key in (
        "paper_size_mode",
        "orientation_mode",
        "margin_mode",
        "paper_size_by_section",
        "orientation_by_section",
        "margin_by_section",
    ):
        del payload["page_setup"][key]
    payload["section"] = {"section_break_type": "nextPage"}
    payload["header_footer"]["behavior"]["link_to_previous"] = "never"
    del payload["header_footer"]["page_number_plan"]["first"]
    del payload["header_footer"]["page_number_plan"]["even"]
    target = tmp_path / "legacy.json"
    target.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(ConfigLoadError):
        load_template(target)

    compatible = load_compatible_user_template(target)
    routed = _load_template_entry_config(target, target.stem, "user")

    assert compatible.page_setup.paper_size_mode == "force_template"
    assert compatible.page_setup.orientation_mode == "preserve_source"
    assert compatible.page_setup.paper_size_by_section == {}
    assert compatible.page_setup.margin_by_section == {}
    assert compatible.section.boundary_mode == "normalize_all"
    assert compatible.section.header_footer_link_mode == "semantic_rebuild"
    assert routed == compatible


def test_custom_quick_formatting_preserves_source_page_geometry():
    template = create_builtin_template("default", mode_id="custom")

    assert template.page_setup.paper_size_mode == "preserve_source"
    assert template.page_setup.orientation_mode == "preserve_source"
    assert template.page_setup.margin_mode == "preserve_source"
    assert template.page_setup.paper_size_by_section == {}
    assert template.page_setup.margin_by_section == {}
    assert template.section.boundary_mode == "preserve_source"
