import json
from pathlib import Path

from docx import Document

from src.config.exam_master_inventory_audit import (
    audit_active_exam_user_master_pool,
    audit_exam_user_master_pool,
    build_exam_user_master_archive_plan,
    format_exam_user_master_archive_plan_markdown,
    format_exam_user_master_pool_audit_markdown,
    write_exam_user_master_archive_plan_manifest,
    write_exam_user_master_archive_plan_markdown,
    write_exam_user_master_pool_audit_manifest,
    write_exam_user_master_pool_audit_markdown,
)
from src.config.loader import save_scene
from src.config.scene import (
    ExamBlankStyleConfig,
    ExamPaperConfig,
    SceneWorkspace,
)


def _write_exam_master_docx(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("{{af_title}}")
    document.add_paragraph("{{af_questions}}")
    document.add_paragraph("{{af_answer_area}}")
    document.save(path)


def _write_plain_docx(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("plain document")
    document.save(path)


def test_exam_user_master_pool_audit_collects_references_before_cleanup(tmp_path):
    user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    referenced_path = user_dir / "user_default_exam_copy.docx"
    builtin_referenced_path = user_dir / "user_imported_exam.docx"
    stale_path = user_dir / "user_default_exam_copy_2.docx"
    manual_path = user_dir / "school_manual_master.docx"
    invalid_path = user_dir / "plain.docx"
    for path in (referenced_path, builtin_referenced_path, stale_path, manual_path):
        _write_exam_master_docx(path)
    _write_plain_docx(invalid_path)

    plan_root = tmp_path / "config_library" / "plans"
    save_scene(
        SceneWorkspace(
            scene_id="exam",
            template_id="default",
            exam_paper=ExamPaperConfig(
                custom_blank_styles=[
                    ExamBlankStyleConfig(
                        style_id="user_default_exam_copy",
                        label="Current exam copy",
                        master_docx_path=str(referenced_path),
                    )
                ]
            ),
        ),
        plan_root / "exam" / "user" / "exam_user.json",
    )
    save_scene(
        SceneWorkspace(
            scene_id="exam_term",
            template_id="default",
            exam_paper=ExamPaperConfig(
                custom_blank_styles=[
                    ExamBlankStyleConfig(
                        style_id="user_imported_exam",
                        label="Built-in imported copy",
                        master_docx_path=str(builtin_referenced_path),
                    )
                ]
            ),
        ),
        plan_root / "exam" / "builtin" / "exam_term.json",
    )

    audit = audit_exam_user_master_pool(
        user_master_dir=user_dir,
        scene_roots=(plan_root,),
    )
    by_name = {item.file_name: item for item in audit.inventory_items}

    assert audit.scanned_scene_count == 2
    assert audit.referenced_file_count == 2
    assert audit.cleanup_candidate_count == 1
    assert by_name["user_default_exam_copy.docx"].category == "referenced"
    assert by_name["user_imported_exam.docx"].category == "referenced"
    assert by_name["user_default_exam_copy_2.docx"].category == "program_copy_unreferenced"
    assert by_name["school_manual_master.docx"].category == "manual_discoverable"
    assert by_name["plain.docx"].category == "unknown_or_invalid"

    referenced_sources = audit.references_for_item(by_name["user_imported_exam.docx"])
    assert len(referenced_sources) == 1
    assert referenced_sources[0].config_id == "exam_term"
    assert referenced_sources[0].source_type == "builtin"

    payload = audit.to_payload()
    assert payload["kind"] == "alavette.exam_master.user_pool_audit"
    assert payload["pool_id"] == "custom"
    assert payload["summary"]["cleanup_candidate_count"] == 1
    cleanup_items = [
        item for item in payload["items"] if item["cleanup_candidate"]
    ]
    assert [item["file_name"] for item in cleanup_items] == [
        "user_default_exam_copy_2.docx"
    ]


def test_exam_user_master_pool_audit_writes_json_manifest_and_markdown(tmp_path):
    user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    _write_exam_master_docx(user_dir / "user_default_exam_copy_2.docx")
    plan_root = tmp_path / "config_library" / "plans"
    plan_root.mkdir(parents=True)

    manifest_path = write_exam_user_master_pool_audit_manifest(
        tmp_path / "audit" / "exam_master_pool.json",
        user_master_dir=user_dir,
        scene_roots=(plan_root,),
    )
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))

    assert payload["summary"]["total_file_count"] == 1
    assert payload["summary"]["cleanup_candidate_count"] == 1
    assert payload["items"][0]["category"] == "program_copy_unreferenced"

    audit = audit_exam_user_master_pool(
        user_master_dir=user_dir,
        scene_roots=(plan_root,),
    )
    markdown = format_exam_user_master_pool_audit_markdown(audit)
    assert "Exam Master User Pool Audit" in markdown
    assert "- Pool: `custom`" in markdown
    assert "`program_copy_unreferenced`: 1" in markdown

    markdown_path = write_exam_user_master_pool_audit_markdown(
        tmp_path / "audit" / "exam_master_pool.md",
        user_master_dir=user_dir,
        scene_roots=(plan_root,),
    )
    markdown_text = markdown_path.read_text(encoding="utf-8")
    assert "- Pool: `custom`" in markdown_text
    assert "`program_copy_unreferenced`: 1" in markdown_text


def test_active_exam_user_master_pool_reads_only_mode_scoped_directory(
    tmp_path,
    monkeypatch,
):
    import src.config.exam_master_inventory_audit as audit_module

    active_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    unrelated_dir = tmp_path / "unrelated" / "exam" / "user"
    active_path = active_dir / "user_default_exam_copy_2.docx"
    unrelated_path = unrelated_dir / "user_imported_exam_2.docx"
    _write_exam_master_docx(active_path)
    _write_exam_master_docx(unrelated_path)
    monkeypatch.setattr(audit_module, "USER_EXAM_MASTER_DIR", active_dir)

    active_audit = audit_active_exam_user_master_pool(scene_roots=())

    assert active_audit.pool_id == "active"
    assert [item.file_name for item in active_audit.inventory_items] == [
        "user_default_exam_copy_2.docx"
    ]
    assert active_audit.cleanup_candidate_count == 1
    assert active_audit.to_payload()["pool_id"] == "active"
    assert unrelated_path.exists()


def test_exam_user_master_pool_audit_protects_current_unsaved_config(tmp_path):
    user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    current_path = user_dir / "user_default_exam_copy_7.docx"
    _write_exam_master_docx(current_path)

    audit = audit_exam_user_master_pool(
        user_master_dir=user_dir,
        scene_roots=(),
        current_config=ExamPaperConfig(
            custom_blank_styles=[
                ExamBlankStyleConfig(
                    style_id="user_default_exam_copy_7",
                    label="Unsaved copy",
                    master_docx_path=str(current_path),
                )
            ]
        ),
    )
    item = audit.inventory_items[0]

    assert item.category == "referenced"
    assert item.referenced is True
    assert audit.cleanup_candidate_count == 0
    references = audit.references_for_item(item)
    assert references[0].source_type == "current_session"


def test_exam_user_master_archive_plan_is_dry_run_and_protects_non_candidates(tmp_path):
    user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    referenced_path = user_dir / "user_default_exam_copy.docx"
    stale_path = user_dir / "user_default_exam_copy_2.docx"
    manual_path = user_dir / "school_manual_master.docx"
    invalid_path = user_dir / "plain.docx"
    for path in (referenced_path, stale_path, manual_path):
        _write_exam_master_docx(path)
    _write_plain_docx(invalid_path)

    audit = audit_exam_user_master_pool(
        user_master_dir=user_dir,
        scene_roots=(),
        current_config=ExamPaperConfig(
            custom_blank_styles=[
                ExamBlankStyleConfig(
                    style_id="user_default_exam_copy",
                    label="Referenced copy",
                    master_docx_path=str(referenced_path),
                )
            ]
        ),
    )
    archive_root = tmp_path / "archive" / "user_legacy"
    plan = build_exam_user_master_archive_plan(audit, archive_root=archive_root)

    assert plan.dry_run is True
    assert [entry.file_name for entry in plan.entries] == [
        "user_default_exam_copy_2.docx"
    ]
    entry = plan.entries[0]
    assert entry.category == "program_copy_unreferenced"
    assert entry.source_path == stale_path.resolve()
    assert entry.archive_path == archive_root / stale_path.name
    assert len(entry.sha1) == 40
    assert entry.size_bytes > 0
    assert stale_path.exists()
    assert not entry.archive_path.exists()

    payload = plan.to_payload()
    assert payload["kind"] == "alavette.exam_master.user_pool_archive_plan"
    assert payload["dry_run"] is True
    assert payload["summary"]["planned_file_count"] == 1
    assert payload["summary"]["protected_file_count"] == 3
    protected = {
        item["file_name"]: item["archive_protected_reason"]
        for item in payload["protected_items"]
    }
    assert protected["user_default_exam_copy.docx"] == "referenced_by_plan"
    assert protected["school_manual_master.docx"] == (
        "category_not_cleanup_candidate:manual_discoverable"
    )
    assert protected["plain.docx"] == "category_not_cleanup_candidate:unknown_or_invalid"


def test_exam_user_master_archive_plan_writes_json_and_markdown_without_moving_files(
    tmp_path,
):
    user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    stale_path = user_dir / "user_imported_exam_2.docx"
    _write_exam_master_docx(stale_path)
    archive_root = tmp_path / "archive" / "exam_user_masters"

    manifest_path = write_exam_user_master_archive_plan_manifest(
        tmp_path / "audit" / "archive_plan.json",
        user_master_dir=user_dir,
        scene_roots=(),
        pool_id="legacy",
        archive_root=archive_root,
    )
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))

    assert payload["pool_id"] == "legacy"
    assert payload["dry_run"] is True
    assert payload["summary"]["planned_file_count"] == 1
    assert payload["planned_entries"][0]["file_name"] == stale_path.name
    assert payload["planned_entries"][0]["archive_path"] == str(
        archive_root / stale_path.name
    )
    assert stale_path.exists()
    assert not (archive_root / stale_path.name).exists()

    audit = audit_exam_user_master_pool(
        user_master_dir=user_dir,
        scene_roots=(),
        pool_id="legacy",
    )
    markdown = format_exam_user_master_archive_plan_markdown(
        build_exam_user_master_archive_plan(audit, archive_root=archive_root)
    )
    assert "Exam Master User Pool Archive Plan" in markdown
    assert "- Dry run: `true`" in markdown
    assert stale_path.name in markdown
    assert str(archive_root / stale_path.name) in markdown

    markdown_path = write_exam_user_master_archive_plan_markdown(
        tmp_path / "audit" / "archive_plan.md",
        user_master_dir=user_dir,
        scene_roots=(),
        pool_id="legacy",
        archive_root=archive_root,
    )
    markdown_text = markdown_path.read_text(encoding="utf-8")
    assert "- Pool: `legacy`" in markdown_text
    assert stale_path.name in markdown_text
    assert stale_path.exists()
    assert not (archive_root / stale_path.name).exists()
