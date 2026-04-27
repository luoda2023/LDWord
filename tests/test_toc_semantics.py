import sys
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.enum.style import WD_STYLE_TYPE


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.config.template import StyleConfig
from src.modules.structure.heading_recognition import HeadingRecognitionModule
from src.modules.structure.toc import (
    TocModule,
    _fallback_toc_section,
    _find_insert_position,
    _format_existing_toc_paragraphs,
    _insert_toc,
    _resolve_existing_toc_range,
)
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.field_builder import iter_field_instructions
from src.shared.engine.ooxml_ops import qn
from src.shared.engine.toc_style_ops import sync_toc_styles


def test_toc_style_sync_updates_heading_and_entry_styles():
    doc = Document()
    if "TOC Heading" not in [style.name for style in doc.styles]:
        doc.styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)

    config_styles = {
        "toc_title": StyleConfig(
            font_cn="Heiti",
            font_en="Arial",
            size_pt=16,
            alignment="center",
            bold=True,
            italic=True,
        ),
        "toc_chapter": StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=14, alignment="left"),
        "toc_level1": StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=12, alignment="left"),
        "toc_level2": StyleConfig(
            font_cn="FangSong",
            font_en="Calibri",
            size_pt=11,
            alignment="left",
            italic=True,
        ),
    }

    changed = sync_toc_styles(doc, config_styles)

    assert changed == 4
    toc_heading = doc.styles["TOC Heading"]
    toc_1 = doc.styles["TOC 1"]
    toc_2 = doc.styles["TOC 2"]
    toc_3 = doc.styles["TOC 3"]

    assert toc_heading.font.size.pt == 16
    assert toc_1.font.size.pt == 14
    assert toc_2.font.size.pt == 12
    assert toc_3.font.size.pt == 11
    assert toc_heading.font.bold is True
    assert toc_heading.font.italic is True
    assert toc_3.font.italic is True

    toc_heading_outline = toc_heading.element.find(qn("w:pPr")).find(qn("w:outlineLvl"))
    assert toc_heading_outline is not None
    assert toc_heading_outline.get(qn("w:val")) == "9"


def test_toc_style_sync_falls_back_to_shared_toc_style_when_role_styles_are_missing():
    doc = Document()
    config_styles = {
        "toc": StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=12, alignment="left"),
    }

    changed = sync_toc_styles(doc, config_styles)

    assert changed == 4
    assert doc.styles["TOC Heading"].font.size.pt == 12
    assert doc.styles["TOC 1"].font.size.pt == 12
    assert doc.styles["TOC 2"].font.size.pt == 12
    assert doc.styles["TOC 3"].font.size.pt == 12


def test_toc_insert_position_prefers_existing_toc_section_from_doc_tree():
    doc = Document()
    context = SimpleNamespace(
        heading_map={5: 1},
        doc_tree=SimpleNamespace(
            get_section=lambda section: SimpleNamespace(start_index=2, end_index=8) if section == "toc" else None
        ),
    )

    assert _find_insert_position(doc, "auto", context) == 2


def test_toc_insert_position_uses_cover_end_before_first_heading():
    doc = Document()
    for _ in range(6):
        doc.add_paragraph("x")

    context = SimpleNamespace(
        heading_map={4: 1},
        doc_tree=SimpleNamespace(
            get_section=lambda section: (
                SimpleNamespace(start_index=0, end_index=3)
                if section == "cover"
                else None
            )
        ),
    )

    assert _find_insert_position(doc, "auto", context) == 3


def test_toc_native_mode_updates_existing_field_depth():
    doc = Document()
    _insert_toc(doc, 0, 2)

    config = ResolvedConfig()
    config.toc.max_level = 4

    TocModule().apply(doc, config, ChangeTracker(), PipelineContext())

    instructions = [
        " ".join((instr or "").split())
        for _kind, _elem, instr in iter_field_instructions(doc.element.body)
    ]
    assert any('TOC \\o "1-4"' in instr for instr in instructions)


def test_toc_formats_existing_paragraphs_using_doc_tree_range():
    doc = Document()
    if "TOC Heading" not in [style.name for style in doc.styles]:
        doc.styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
    title = doc.add_paragraph("Contents")
    entry = doc.add_paragraph("1.1 Intro\t1")

    config = ResolvedConfig()
    config.styles["toc_title"] = StyleConfig(
        font_cn="Heiti",
        font_en="Arial",
        size_pt=16,
        alignment="center",
        bold=True,
        italic=True,
    )
    config.styles["toc_chapter"] = StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=14, alignment="left")
    config.styles["toc_level1"] = StyleConfig(
        font_cn="Songti",
        font_en="Times New Roman",
        size_pt=12,
        alignment="left",
        italic=True,
    )
    sync_toc_styles(doc, config.styles)

    context = SimpleNamespace(
        doc_tree=SimpleNamespace(get_section=lambda section: SimpleNamespace(start_index=0, end_index=2) if section == "toc" else None)
    )

    changed = _format_existing_toc_paragraphs(doc, config, context)

    assert changed == 2
    assert title.style.name == "TOC Heading"
    assert entry.style.name == "TOC 2"
    assert title.runs[0].font.bold is True
    assert title.runs[0].font.italic is True
    assert entry.runs[0].font.italic is True
    title_outline = title._element.find(qn("w:pPr")).find(qn("w:outlineLvl"))
    assert title_outline is not None
    assert title_outline.get(qn("w:val")) == "9"


def test_toc_fallback_section_uses_title_and_entry_block():
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_paragraph("Contents")
    doc.add_paragraph("1 Intro\t1")
    doc.add_paragraph("1.1 Background\t2")
    doc.add_paragraph("Chapter 1 Intro")

    assert _fallback_toc_section(doc) == (1, 4)


def test_toc_resolve_existing_range_rejects_suspicious_doc_tree_section():
    doc = Document()
    if "TOC Heading" not in [style.name for style in doc.styles]:
        doc.styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
    title = doc.add_paragraph("Contents")
    entry = doc.add_paragraph("1 Intro\t1")
    body_heading = doc.add_heading("Chapter 1 Intro", level=1)
    doc.add_paragraph("Body content")

    config = ResolvedConfig()
    config.styles["toc_title"] = StyleConfig(font_cn="Heiti", font_en="Arial", size_pt=16, alignment="center")
    config.styles["toc_chapter"] = StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=14, alignment="left")
    sync_toc_styles(doc, config.styles)

    context = SimpleNamespace(
        doc_tree=SimpleNamespace(get_section=lambda section: SimpleNamespace(start_index=0, end_index=4) if section == "toc" else None)
    )

    assert _resolve_existing_toc_range(doc, context) == (0, 2)

    changed = _format_existing_toc_paragraphs(doc, config, context)
    assert changed == 2
    assert title.style.name == "TOC Heading"
    assert entry.style.name == "TOC 1"
    assert body_heading.style.name != "TOC 1"


def test_toc_plain_mode_inserts_plain_entries_from_front_body_and_back_matter():
    doc = Document()
    doc.add_paragraph("硕士学位论文")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("Abstract body")
    doc.add_heading("Chapter 1 Intro", level=1)
    doc.add_heading("1.1 Background", level=2)
    doc.add_paragraph("References")
    doc.add_paragraph("[1] Author. Title[J]. Journal, 2024, 12(3): 1-8.")
    doc.add_paragraph("[2] Author. Another title[J]. Journal, 2023, 11(2): 9-12.")

    context = PipelineContext()
    HeadingRecognitionModule().apply(doc, ResolvedConfig(), ChangeTracker(), context)

    config = ResolvedConfig()
    config.toc.mode = "plain"
    config.styles["toc_title"] = StyleConfig(font_cn="Heiti", font_en="Arial", size_pt=16, alignment="center")
    config.styles["toc_chapter"] = StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=14, alignment="left")
    config.styles["toc_level1"] = StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=12, alignment="left")
    config.styles["toc_level2"] = StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=10.5, alignment="left")

    TocModule().apply(doc, config, ChangeTracker(), context)

    texts = [para.text for para in doc.paragraphs[:7]]
    assert texts == ["硕士学位论文", "目录", "Abstract", "Chapter 1 Intro", "1.1 Background", "References", "Abstract"]
    assert doc.paragraphs[2].style.name == "TOC 1"
    assert doc.paragraphs[3].style.name == "TOC 1"
    assert doc.paragraphs[4].style.name == "TOC 2"
    assert doc.paragraphs[5].style.name == "TOC 1"


def test_toc_plain_mode_rebuilds_existing_plain_toc_block():
    doc = Document()
    if "TOC Heading" not in [style.name for style in doc.styles]:
        doc.styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
    if "TOC 1" not in [style.name for style in doc.styles]:
        doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    old_title = doc.add_paragraph("Contents")
    old_entry = doc.add_paragraph("Old entry")
    old_entry.style = doc.styles["TOC 1"]
    doc.add_paragraph("硕士学位论文")
    doc.add_heading("Chapter 1 Intro", level=1)
    doc.add_paragraph("Body content")

    context = PipelineContext()
    HeadingRecognitionModule().apply(doc, ResolvedConfig(), ChangeTracker(), context)

    config = ResolvedConfig()
    config.toc.mode = "plain"
    config.styles["toc_title"] = StyleConfig(font_cn="Heiti", font_en="Arial", size_pt=16, alignment="center")
    config.styles["toc_chapter"] = StyleConfig(font_cn="Songti", font_en="Times New Roman", size_pt=14, alignment="left")

    TocModule().apply(doc, config, ChangeTracker(), context)

    assert doc.paragraphs[0].text == "目录"
    assert doc.paragraphs[1].text == "Chapter 1 Intro"
    assert doc.paragraphs[1].style.name == "TOC 1"
    assert old_title.text != doc.paragraphs[0].text
