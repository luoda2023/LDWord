import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.ui.panels.workbench.execution_runtime import _render_batch_markdown_report
from src.ui.panels.workbench.question_figure_repair_runtime import (
    _apply_question_figure_repair_batch_execution_to_docx,
    _batch_question_figure_repair_queue,
    _dry_run_question_figure_repair_batch_apply_guard,
    _freeze_question_figure_repair_batch_confirmation_plan,
    _plan_question_figure_repair_batch_apply_execution,
    _resolve_question_figure_repair_queue_conflict,
    _rollback_question_figure_repair_batch_apply_from_audit,
)


def test_question_figure_repair_queue_blocks_conflicting_ready_candidates(tmp_path):
    current = tmp_path / "question_2.png"
    replacement_a = tmp_path / "question_2_expected_a.png"
    replacement_b = tmp_path / "question_2_expected_b.png"
    Image.new("RGB", (320, 240), color="blue").save(current)
    Image.new("RGB", (320, 240), color="yellow").save(replacement_a)
    Image.new("RGB", (320, 240), color="red").save(replacement_b)
    target_key = json.dumps(
        {
            "role": "question_figure",
            "item_id": "question_figure_2",
            "question_index": "2",
            "path": str(current),
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )

    queue = _batch_question_figure_repair_queue(
        [
            {
                "issue_id": "issue-a",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_a",
                "profile_name": "Exam A",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key,
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_2",
                        "label": "Question figure 2",
                        "path": str(current),
                        "question_index": "2",
                        "comparison_issue_reference": str(replacement_a),
                        "comparison_issue_display_name": replacement_a.name,
                        "comparison_issue_summary": "候选 A",
                    },
                    {
                        "item_id": "question_figure_2",
                        "label": "Question figure 2",
                        "path": str(current),
                        "question_index": "2",
                        "comparison_issue_reference": str(replacement_b),
                        "comparison_issue_display_name": replacement_b.name,
                        "comparison_issue_summary": "候选 B",
                    },
                ],
            }
        ]
    )

    assert queue["kind"] == "question_figure_repair_queue"
    assert queue["status"] == "conflict"
    assert queue["queue_count"] == 2
    assert queue["conflict_group_count"] == 1
    assert queue["conflict_count"] == 2
    entries = queue["entries"]
    assert {entry["confirmation_status"] for entry in entries} == {"conflict"}
    assert {entry["confirmation_apply_supported"] for entry in entries} == {False}
    assert {
        entry["confirmation_action"] for entry in entries
    } == {"resolve_question_figure_replacement_conflict"}
    conflict_group_ids = {entry["conflict_group_id"] for entry in entries}
    assert len(conflict_group_ids) == 1
    for entry in entries:
        assert entry["conflict_candidate_count"] == 2
        assert len(entry["conflict_candidate_queue_ids"]) == 2
        assert "candidate_conflict_same_repair_target" in entry["apply_blockers"]
        assert entry["conflict_resolution_select_supported"] is True
        assert entry["conflict_resolution_action"] == (
            "select_question_figure_replacement_conflict_candidate"
        )
        assert entry["replacement_source_kind"] == "local_file"
        assert entry["replacement_source_path"] in {
            str(replacement_a),
            str(replacement_b),
        }


def test_question_figure_repair_batch_confirmation_plan_summarizes_queue(tmp_path):
    current_ready = tmp_path / "question_2.png"
    replacement_ready = tmp_path / "question_2_expected.png"
    current_blocked = tmp_path / "question_5.png"
    current_conflict = tmp_path / "question_7.png"
    replacement_conflict_a = tmp_path / "question_7_expected_a.png"
    replacement_conflict_b = tmp_path / "question_7_expected_b.png"
    for image_path, color in (
        (current_ready, "blue"),
        (replacement_ready, "yellow"),
        (current_blocked, "green"),
        (current_conflict, "purple"),
        (replacement_conflict_a, "red"),
        (replacement_conflict_b, "orange"),
    ):
        Image.new("RGB", (320, 240), color=color).save(image_path)

    def target_key(item_id: str, question_index: str, path: Path) -> str:
        return json.dumps(
            {
                "role": "question_figure",
                "item_id": item_id,
                "question_index": question_index,
                "path": str(path),
            },
            ensure_ascii=False,
            separators=(",", ":"),
        )

    queue = _batch_question_figure_repair_queue(
        [
            {
                "issue_id": "issue-ready",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_ready",
                "profile_name": "Exam Ready",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key(
                    "question_figure_2", "2", current_ready
                ),
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_2",
                        "label": "Question figure 2",
                        "path": str(current_ready),
                        "question_index": "2",
                        "comparison_issue_reference": str(replacement_ready),
                        "comparison_issue_display_name": replacement_ready.name,
                    }
                ],
            },
            {
                "issue_id": "issue-blocked",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_blocked",
                "profile_name": "Exam Blocked",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key(
                    "question_figure_5", "5", current_blocked
                ),
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_5",
                        "label": "Question figure 5",
                        "path": str(current_blocked),
                        "question_index": "5",
                        "comparison_issue_reference": "asset://expected-q5",
                        "comparison_issue_display_name": "expected-q5",
                    }
                ],
            },
            {
                "issue_id": "issue-conflict",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_conflict",
                "profile_name": "Exam Conflict",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key(
                    "question_figure_7", "7", current_conflict
                ),
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_7",
                        "label": "Question figure 7",
                        "path": str(current_conflict),
                        "question_index": "7",
                        "comparison_issue_reference": str(replacement_conflict_a),
                        "comparison_issue_display_name": replacement_conflict_a.name,
                    },
                    {
                        "item_id": "question_figure_7",
                        "label": "Question figure 7",
                        "path": str(current_conflict),
                        "question_index": "7",
                        "comparison_issue_reference": str(replacement_conflict_b),
                        "comparison_issue_display_name": replacement_conflict_b.name,
                    },
                ],
            },
        ]
    )

    plan = queue["batch_confirmation_plan"]

    assert queue["status"] == "conflict"
    assert plan["kind"] == "question_figure_repair_batch_confirmation_plan"
    assert plan["status"] == "partial"
    assert plan["action"] == "review_question_figure_repair_batch_confirmation_plan"
    assert plan["entry_count"] == 4
    assert plan["target_count"] == 3
    assert plan["eligible_count"] == 1
    assert plan["blocked_count"] == 3
    assert plan["conflict_count"] == 2
    assert plan["duplicate_target_count"] == 0
    assert plan["batch_confirmation_supported"] is True
    assert plan["batch_apply_supported"] is False
    assert plan["plan_fingerprint"]
    assert len(plan["eligible_queue_ids"]) == 1
    assert len(plan["blocked_queue_ids"]) == 3
    assert plan["eligible_entries"][0]["question_index"] == "2"
    assert any(
        "comparison_reference_not_local_file" in entry["blockers"]
        for entry in plan["blocked_entries"]
    )
    assert sum(
        "candidate_conflict_same_repair_target" in entry["blockers"]
        for entry in plan["blocked_entries"]
    ) == 2


def test_question_figure_repair_batch_confirmation_freeze_requires_confirmation(
    tmp_path,
):
    current_ready = tmp_path / "question_2.png"
    replacement_ready = tmp_path / "question_2_expected.png"
    current_blocked = tmp_path / "question_5.png"
    for image_path, color in (
        (current_ready, "blue"),
        (replacement_ready, "yellow"),
        (current_blocked, "green"),
    ):
        Image.new("RGB", (320, 240), color=color).save(image_path)

    def target_key(item_id: str, question_index: str, path: Path) -> str:
        return json.dumps(
            {
                "role": "question_figure",
                "item_id": item_id,
                "question_index": question_index,
                "path": str(path),
            },
            ensure_ascii=False,
            separators=(",", ":"),
        )

    queue = _batch_question_figure_repair_queue(
        [
            {
                "issue_id": "issue-ready",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_ready",
                "profile_name": "Exam Ready",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key(
                    "question_figure_2", "2", current_ready
                ),
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_2",
                        "label": "Question figure 2",
                        "path": str(current_ready),
                        "question_index": "2",
                        "comparison_issue_reference": str(replacement_ready),
                        "comparison_issue_display_name": replacement_ready.name,
                    }
                ],
            },
            {
                "issue_id": "issue-blocked",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_blocked",
                "profile_name": "Exam Blocked",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key(
                    "question_figure_5", "5", current_blocked
                ),
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_5",
                        "label": "Question figure 5",
                        "path": str(current_blocked),
                        "question_index": "5",
                        "comparison_issue_reference": "asset://expected-q5",
                        "comparison_issue_display_name": "expected-q5",
                    }
                ],
            },
        ]
    )
    plan = queue["batch_confirmation_plan"]
    eligible_queue_id = plan["eligible_queue_ids"][0]

    not_confirmed = _freeze_question_figure_repair_batch_confirmation_plan(queue)
    not_confirmed_freeze = not_confirmed["batch_confirmation_freeze"]
    assert not_confirmed_freeze["status"] == "not_confirmed"
    assert not_confirmed_freeze["confirmed"] is False
    assert "manual_confirmation_required" in not_confirmed_freeze["blockers"]
    assert not_confirmed_freeze["frozen_queue_ids"] == []

    stale = _freeze_question_figure_repair_batch_confirmation_plan(
        queue,
        confirmed=True,
        expected_plan_fingerprint="stale-plan",
    )
    stale_freeze = stale["batch_confirmation_freeze"]
    assert stale_freeze["status"] == "blocked"
    assert stale_freeze["fingerprint_matched"] is False
    assert "plan_fingerprint_mismatch" in stale_freeze["blockers"]

    frozen = _freeze_question_figure_repair_batch_confirmation_plan(
        queue,
        confirmed=True,
        confirmed_by="batch-reviewer",
        confirmation_note="freeze current ready candidates",
        expected_plan_fingerprint=plan["plan_fingerprint"],
    )
    freeze = frozen["batch_confirmation_freeze"]
    assert freeze["kind"] == "question_figure_repair_batch_confirmation_freeze"
    assert freeze["status"] == "partial_frozen"
    assert freeze["confirmed"] is True
    assert freeze["confirmed_by"] == "batch-reviewer"
    assert freeze["confirmation_note"] == "freeze current ready candidates"
    assert freeze["plan_fingerprint"] == plan["plan_fingerprint"]
    assert freeze["fingerprint_matched"] is True
    assert freeze["frozen_queue_ids"] == [eligible_queue_id]
    assert freeze["frozen_candidate_count"] == 1
    assert freeze["frozen_target_count"] == 1
    assert freeze["blocked_count"] == 1
    assert freeze["batch_apply_supported"] is False
    assert freeze["execution_plan_ready"] is True
    assert freeze["frozen_entries"][0]["queue_id"] == eligible_queue_id
    assert freeze["frozen_entries"][0]["replacement_source_path"] == str(
        replacement_ready
    )
    markdown = _render_batch_markdown_report(
        {
            "status": "partial_success",
            "input_path": "source.docx",
            "failed_count": 0,
            "items": [],
            "question_figure_repair_queue": frozen,
        }
    )
    assert "Batch confirmation freeze: status=partial_frozen" in markdown
    assert f"fingerprint={plan['plan_fingerprint']}" in markdown
    assert "batch_apply=false" in markdown

    invalid = _freeze_question_figure_repair_batch_confirmation_plan(
        queue,
        selected_queue_ids=["not-in-plan"],
        confirmed=True,
        expected_plan_fingerprint=plan["plan_fingerprint"],
    )
    invalid_freeze = invalid["batch_confirmation_freeze"]
    assert invalid_freeze["status"] == "blocked"
    assert invalid_freeze["invalid_queue_ids"] == ["not-in-plan"]
    assert "selected_queue_id_not_eligible" in invalid_freeze["blockers"]


def test_question_figure_repair_batch_apply_dry_run_guards_frozen_plan(tmp_path):
    current_ready = tmp_path / "question_2.png"
    replacement_ready = tmp_path / "question_2_expected.png"
    Image.new("RGB", (320, 240), color="blue").save(current_ready)
    Image.new("RGB", (320, 240), color="yellow").save(replacement_ready)
    target_key = json.dumps(
        {
            "role": "question_figure",
            "item_id": "question_figure_2",
            "question_index": "2",
            "path": str(current_ready),
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    queue = _batch_question_figure_repair_queue(
        [
            {
                "issue_id": "issue-ready",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_ready",
                "profile_name": "Exam Ready",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key,
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_2",
                        "label": "Question figure 2",
                        "path": str(current_ready),
                        "question_index": "2",
                        "comparison_issue_reference": str(replacement_ready),
                        "comparison_issue_display_name": replacement_ready.name,
                    }
                ],
            }
        ]
    )
    plan = queue["batch_confirmation_plan"]
    frozen = _freeze_question_figure_repair_batch_confirmation_plan(
        queue,
        confirmed=True,
        confirmed_by="batch-reviewer",
        expected_plan_fingerprint=plan["plan_fingerprint"],
    )
    freeze = frozen["batch_confirmation_freeze"]
    guarded = _dry_run_question_figure_repair_batch_apply_guard(
        frozen,
        expected_freeze_id=freeze["freeze_id"],
        expected_plan_fingerprint=freeze["plan_fingerprint"],
    )
    dry_run = guarded["batch_apply_dry_run"]

    assert dry_run["kind"] == "question_figure_repair_batch_apply_dry_run"
    assert dry_run["status"] == "ready"
    assert dry_run["execution_guard_passed"] is True
    assert dry_run["batch_apply_supported"] is False
    assert dry_run["freeze_id"] == freeze["freeze_id"]
    assert dry_run["fingerprint_matched"] is True
    assert dry_run["checked_candidate_count"] == 1
    assert dry_run["checked_target_count"] == 1
    assert dry_run["ready_queue_ids"] == freeze["frozen_queue_ids"]
    assert dry_run["blocked_queue_ids"] == []
    assert dry_run["candidate_results"][0]["status"] == "ready"

    markdown = _render_batch_markdown_report(
        {
            "status": "partial_success",
            "input_path": "source.docx",
            "failed_count": 0,
            "items": [],
            "question_figure_repair_queue": guarded,
        }
    )
    assert "Batch apply dry-run: status=ready" in markdown
    assert "guard_passed=true" in markdown
    assert "batch_apply=false" in markdown

    replacement_ready.unlink()
    blocked = _dry_run_question_figure_repair_batch_apply_guard(
        frozen,
        expected_freeze_id=freeze["freeze_id"],
        expected_plan_fingerprint=freeze["plan_fingerprint"],
    )
    blocked_dry_run = blocked["batch_apply_dry_run"]

    assert blocked_dry_run["status"] == "blocked"
    assert blocked_dry_run["execution_guard_passed"] is False
    assert blocked_dry_run["batch_apply_supported"] is False
    assert "plan_fingerprint_mismatch" in blocked_dry_run["blockers"]
    assert "frozen_candidate_guard_failed" in blocked_dry_run["blockers"]
    assert blocked_dry_run["blocked_queue_ids"] == freeze["frozen_queue_ids"]
    assert blocked_dry_run["candidate_results"][0]["status"] == "blocked"
    assert "frozen_candidate_not_currently_eligible" in blocked_dry_run[
        "candidate_results"
    ][0]["blockers"]


def test_question_figure_repair_batch_apply_execution_plan_drafts_guarded_work(
    tmp_path,
):
    current_ready = tmp_path / "question_3.png"
    replacement_ready = tmp_path / "question_3_expected.png"
    Image.new("RGB", (320, 240), color="green").save(current_ready)
    Image.new("RGB", (320, 240), color="orange").save(replacement_ready)
    target_key = json.dumps(
        {
            "role": "question_figure",
            "item_id": "question_figure_3",
            "question_index": "3",
            "path": str(current_ready),
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    queue = _batch_question_figure_repair_queue(
        [
            {
                "issue_id": "issue-ready-plan",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_ready",
                "profile_name": "Exam Ready",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key,
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_3",
                        "label": "Question figure 3",
                        "path": str(current_ready),
                        "question_index": "3",
                        "comparison_issue_reference": str(replacement_ready),
                        "comparison_issue_display_name": replacement_ready.name,
                    }
                ],
            }
        ]
    )
    plan = queue["batch_confirmation_plan"]
    frozen = _freeze_question_figure_repair_batch_confirmation_plan(
        queue,
        confirmed=True,
        confirmed_by="batch-reviewer",
        expected_plan_fingerprint=plan["plan_fingerprint"],
    )
    freeze = frozen["batch_confirmation_freeze"]
    guarded = _dry_run_question_figure_repair_batch_apply_guard(
        frozen,
        expected_freeze_id=freeze["freeze_id"],
        expected_plan_fingerprint=freeze["plan_fingerprint"],
    )

    planned = _plan_question_figure_repair_batch_apply_execution(
        guarded,
        confirmed=True,
        confirmed_by="final-reviewer",
        confirmation_note="ready to draft",
        expected_freeze_id=freeze["freeze_id"],
        expected_plan_fingerprint=freeze["plan_fingerprint"],
    )
    execution_plan = planned["batch_apply_execution_plan"]

    assert execution_plan["kind"] == (
        "question_figure_repair_batch_apply_execution_plan"
    )
    assert execution_plan["status"] == "planned"
    assert execution_plan["execution_plan_ready"] is True
    assert execution_plan["batch_apply_supported"] is False
    assert execution_plan["word_write_enabled"] is False
    assert execution_plan["final_confirmation_provided"] is True
    assert execution_plan["freeze_id"] == freeze["freeze_id"]
    assert execution_plan["dry_run_status"] == "ready"
    assert execution_plan["fingerprint_matched"] is True
    assert execution_plan["planned_candidate_count"] == 1
    assert execution_plan["planned_target_count"] == 1
    assert execution_plan["planned_queue_ids"] == freeze["frozen_queue_ids"]
    assert execution_plan["blocked_queue_ids"] == []
    assert execution_plan["plan_id"].startswith(
        "question-figure-batch-apply-plan:"
    )
    assert execution_plan["execution_plan_fingerprint"]
    entry = execution_plan["planned_entries"][0]
    assert entry["status"] == "planned"
    assert entry["operation"] == "replace_question_figure_asset"
    assert entry["queue_id"] == freeze["frozen_queue_ids"][0]
    assert entry["replacement_source_path"] == str(replacement_ready)
    assert entry["repair_target_key"] == target_key
    assert entry["word_write_enabled"] is False

    markdown = _render_batch_markdown_report(
        {
            "status": "partial_success",
            "input_path": "source.docx",
            "failed_count": 0,
            "items": [],
            "question_figure_repair_queue": planned,
        }
    )
    assert "Batch apply execution plan: status=planned" in markdown
    assert "planned=1" in markdown
    assert "final_confirmation=provided" in markdown
    assert "batch_apply=false" in markdown

    replacement_ready.unlink()
    blocked = _plan_question_figure_repair_batch_apply_execution(
        guarded,
        confirmed=True,
        confirmed_by="final-reviewer",
        expected_freeze_id=freeze["freeze_id"],
        expected_plan_fingerprint=freeze["plan_fingerprint"],
    )
    blocked_plan = blocked["batch_apply_execution_plan"]

    assert blocked_plan["status"] == "blocked"
    assert blocked_plan["execution_plan_ready"] is False
    assert blocked_plan["batch_apply_supported"] is False
    assert "batch_apply_dry_run_not_ready" in blocked_plan["blockers"]
    assert "batch_apply_guard_not_passed" in blocked_plan["blockers"]
    assert "batch_apply_dry_run_has_blockers" in blocked_plan["blockers"]
    assert "plan_fingerprint_mismatch" in blocked_plan["blockers"]
    assert blocked_plan["planned_candidate_count"] == 0
    assert blocked_plan["dry_run_blocked_queue_ids"] == freeze["frozen_queue_ids"]


def test_question_figure_repair_batch_apply_execution_writes_docx_media_copy(
    tmp_path,
):
    current_ready = tmp_path / "question_4.png"
    replacement_ready = tmp_path / "question_4_expected.png"
    Image.new("RGB", (320, 240), color="purple").save(current_ready)
    Image.new("RGB", (320, 240), color="white").save(replacement_ready)
    source_docx = tmp_path / "source_exam.docx"
    output_docx = tmp_path / "source_exam_repaired.docx"
    document = Document()
    document.add_paragraph("Question 4")
    document.add_picture(str(current_ready))
    document.save(source_docx)

    target_key = json.dumps(
        {
            "role": "question_figure",
            "item_id": "question_figure_4",
            "question_index": "4",
            "path": str(current_ready),
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    queue = _batch_question_figure_repair_queue(
        [
            {
                "issue_id": "issue-ready-docx",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_ready",
                "profile_name": "Exam Ready",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key,
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_4",
                        "label": "Question figure 4",
                        "path": str(current_ready),
                        "question_index": "4",
                        "comparison_issue_reference": str(replacement_ready),
                        "comparison_issue_display_name": replacement_ready.name,
                    }
                ],
            }
        ]
    )
    plan = queue["batch_confirmation_plan"]
    frozen = _freeze_question_figure_repair_batch_confirmation_plan(
        queue,
        confirmed=True,
        confirmed_by="batch-reviewer",
        expected_plan_fingerprint=plan["plan_fingerprint"],
    )
    freeze = frozen["batch_confirmation_freeze"]
    guarded = _dry_run_question_figure_repair_batch_apply_guard(
        frozen,
        expected_freeze_id=freeze["freeze_id"],
        expected_plan_fingerprint=freeze["plan_fingerprint"],
    )
    planned = _plan_question_figure_repair_batch_apply_execution(
        guarded,
        confirmed=True,
        confirmed_by="final-reviewer",
        expected_freeze_id=freeze["freeze_id"],
        expected_plan_fingerprint=freeze["plan_fingerprint"],
    )
    execution_plan = planned["batch_apply_execution_plan"]
    applied = _apply_question_figure_repair_batch_execution_to_docx(
        planned,
        input_docx_path=source_docx,
        output_docx_path=output_docx,
        confirmed=True,
        confirmed_by="docx-reviewer",
        expected_plan_id=execution_plan["plan_id"],
        expected_execution_plan_fingerprint=execution_plan[
            "execution_plan_fingerprint"
        ],
    )
    result = applied["batch_apply_execution_result"]

    assert result["kind"] == "question_figure_repair_batch_apply_execution_result"
    assert result["status"] == "applied"
    assert result["batch_apply_supported"] is True
    assert result["word_write_enabled"] is True
    assert result["applied_count"] == 1
    assert result["applied_target_count"] == 1
    assert result["output_path"] == str(output_docx)
    assert result["output_sha1"]
    assert result["audit_supported"] is True
    assert result["audit_written"] is True
    assert result["audit_path"].endswith("question_figure_batch_apply_audit.json")
    assert result["audit_record_id"].startswith(
        "question-figure-batch-apply-audit:"
    )
    assert result["entry_results"][0]["status"] == "applied"
    assert result["entry_results"][0]["media_path"].startswith("word/media/")
    assert result["entry_results"][0]["original_path"] == str(current_ready)
    assert result["entry_results"][0]["replacement_source_path"] == str(
        replacement_ready
    )
    apply_manifest = applied["batch_apply_transaction_manifest"]
    assert (
        apply_manifest["kind"]
        == "question_figure_repair_batch_apply_transaction_manifest"
    )
    assert apply_manifest["status"] == "tracked"
    assert apply_manifest["expected_apply_audit_found"] is True
    assert apply_manifest["transaction_count"] == 1
    assert apply_manifest["active_count"] == 1
    assert apply_manifest["rolled_back_count"] == 0
    assert apply_manifest["orphan_rollback_count"] == 0
    assert apply_manifest["transactions"][0]["transaction_status"] == "applied"
    assert apply_manifest["transactions"][0]["apply_audit_id"] == (
        result["audit_record_id"]
    )
    assert apply_manifest["transactions"][0]["rollback_available"] is True
    apply_task_summary = apply_manifest["task_summary"]
    assert (
        apply_task_summary["kind"]
        == "question_figure_repair_batch_apply_transaction_task_summary"
    )
    assert apply_task_summary["status"] == "active"
    assert apply_task_summary["next_action"] == "review_active_transaction"
    assert apply_task_summary["transaction_count"] == 1
    assert apply_task_summary["active_count"] == 1
    assert apply_task_summary["rolled_back_count"] == 0
    assert apply_task_summary["rollback_available_count"] == 1
    assert apply_task_summary["latest_apply_audit_id"] == result["audit_record_id"]
    assert apply_manifest["artifact_supported"] is True
    assert apply_manifest["artifact_written"] is True
    assert apply_manifest["artifact_path"].endswith(
        "question_figure_batch_apply_transaction_manifest.json"
    )
    assert apply_manifest["report_supported"] is True
    assert apply_manifest["report_written"] is True
    assert apply_manifest["report_path"].endswith(
        "question_figure_batch_apply_transaction_manifest.md"
    )
    apply_manifest_path = Path(apply_manifest["artifact_path"])
    apply_report_path = Path(apply_manifest["report_path"])
    apply_manifest_payload = json.loads(
        apply_manifest_path.read_text(encoding="utf-8")
    )
    apply_report_text = apply_report_path.read_text(encoding="utf-8")
    assert apply_manifest_payload["transaction_count"] == 1
    assert apply_manifest_payload["active_count"] == 1
    assert apply_manifest_payload["rolled_back_count"] == 0
    assert apply_manifest_payload["report_written"] is True
    assert apply_manifest_payload["task_summary"]["status"] == "active"
    assert (
        apply_manifest_payload["task_summary"]["rollback_available_count"]
        == 1
    )
    assert "# Question Figure Batch Apply Transaction Manifest" in apply_report_text
    assert "## Task Summary" in apply_report_text
    assert "- transactions: 1" in apply_report_text
    assert "- active: 1" in apply_report_text
    assert "- rollback_available: 1" in apply_report_text
    assert result["audit_record_id"] in apply_report_text
    assert output_docx.is_file()
    assert len(Document(output_docx).inline_shapes) == 1

    with zipfile.ZipFile(source_docx) as archive:
        source_media = [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]
    with zipfile.ZipFile(output_docx) as archive:
        output_media = [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]
    assert current_ready.read_bytes() in source_media
    assert replacement_ready.read_bytes() in output_media
    assert current_ready.read_bytes() not in output_media
    audit_path = Path(result["audit_path"])
    audit_payload = json.loads(audit_path.read_text(encoding="utf-8"))
    assert audit_payload["kind"] == "question_figure_repair_batch_apply_audit"
    assert audit_payload["entry_count"] == 1
    audit_record = audit_payload["records"][0]
    assert audit_record["audit_id"] == result["audit_record_id"]
    assert audit_record["plan_id"] == execution_plan["plan_id"]
    assert audit_record["output_path"] == str(output_docx)
    assert audit_record["output_sha1"] == result["output_sha1"]
    assert audit_record["applied_count"] == 1
    assert audit_record["entry_results"][0]["queue_id"] == (
        result["entry_results"][0]["queue_id"]
    )

    markdown = _render_batch_markdown_report(
        {
            "status": "partial_success",
            "input_path": "source.docx",
            "failed_count": 0,
            "items": [],
            "question_figure_repair_queue": applied,
        }
    )
    assert "Batch apply execution result: status=applied" in markdown
    assert "applied=1" in markdown
    assert "word_write=true" in markdown
    assert "audit=written" in markdown
    assert "Batch apply transaction manifest: status=tracked" in markdown
    assert "task=active" in markdown
    assert "rollback_available=1" in markdown
    assert "artifact=written" in markdown
    assert "report=written" in markdown

    rollback_docx = tmp_path / "source_exam_rollback.docx"
    rolled_back = _rollback_question_figure_repair_batch_apply_from_audit(
        applied,
        applied_docx_path=output_docx,
        rollback_docx_path=rollback_docx,
        confirmed=True,
        confirmed_by="rollback-reviewer",
        expected_audit_id=result["audit_record_id"],
        expected_output_sha1=result["output_sha1"],
    )
    rollback_result = rolled_back["batch_apply_rollback_result"]
    rollback_manifest = rolled_back["batch_apply_transaction_manifest"]
    assert (
        rollback_result["kind"]
        == "question_figure_repair_batch_apply_rollback_result"
    )
    assert rollback_result["status"] == "rolled_back"
    assert rollback_result["linked_apply_audit_id"] == result["audit_record_id"]
    assert rollback_result["rollback_supported"] is True
    assert rollback_result["word_write_enabled"] is True
    assert rollback_result["restored_count"] == 1
    assert rollback_result["restored_target_count"] == 1
    assert rollback_result["output_path"] == str(rollback_docx)
    assert rollback_result["output_sha1"]
    assert rollback_result["audit_supported"] is True
    assert rollback_result["audit_written"] is True
    assert rollback_result["audit_path"] == str(audit_path)
    assert rollback_result["audit_record_id"].startswith(
        "question-figure-batch-rollback-audit:"
    )
    assert rollback_result["entry_results"][0]["status"] == "restored"
    assert rollback_result["entry_results"][0]["media_path"] == (
        result["entry_results"][0]["media_path"]
    )
    assert rollback_manifest["status"] == "tracked"
    assert rollback_manifest["transaction_count"] == 1
    assert rollback_manifest["active_count"] == 0
    assert rollback_manifest["rolled_back_count"] == 1
    assert rollback_manifest["orphan_rollback_count"] == 0
    rollback_transaction = rollback_manifest["transactions"][0]
    assert rollback_transaction["transaction_status"] == "rolled_back"
    assert rollback_transaction["apply_audit_id"] == result["audit_record_id"]
    assert rollback_transaction["rollback_count"] == 1
    assert rollback_transaction["latest_rollback_audit_id"] == (
        rollback_result["audit_record_id"]
    )
    assert rollback_transaction["latest_rollback_output_path"] == str(rollback_docx)
    assert rollback_transaction["rollback_available"] is False
    rollback_task_summary = rollback_manifest["task_summary"]
    assert rollback_task_summary["status"] == "completed"
    assert rollback_task_summary["next_action"] == "review_transaction_history"
    assert rollback_task_summary["transaction_count"] == 1
    assert rollback_task_summary["active_count"] == 0
    assert rollback_task_summary["rolled_back_count"] == 1
    assert rollback_task_summary["rollback_available_count"] == 0
    assert rollback_task_summary["latest_apply_audit_id"] == result["audit_record_id"]
    assert rollback_task_summary["latest_rollback_audit_id"] == (
        rollback_result["audit_record_id"]
    )
    assert rollback_manifest["artifact_supported"] is True
    assert rollback_manifest["artifact_written"] is True
    assert rollback_manifest["artifact_path"] == str(apply_manifest_path)
    assert rollback_manifest["report_supported"] is True
    assert rollback_manifest["report_written"] is True
    assert rollback_manifest["report_path"] == str(apply_report_path)
    rollback_manifest_payload = json.loads(
        apply_manifest_path.read_text(encoding="utf-8")
    )
    rollback_report_text = apply_report_path.read_text(encoding="utf-8")
    assert rollback_manifest_payload["transaction_count"] == 1
    assert rollback_manifest_payload["active_count"] == 0
    assert rollback_manifest_payload["rolled_back_count"] == 1
    assert rollback_manifest_payload["report_written"] is True
    assert rollback_manifest_payload["task_summary"]["status"] == "completed"
    assert (
        rollback_manifest_payload["task_summary"]["rollback_available_count"]
        == 0
    )
    assert rollback_manifest_payload["transactions"][0][
        "latest_rollback_audit_id"
    ] == rollback_result["audit_record_id"]
    assert "- rolled_back: 1" in rollback_report_text
    assert "- rollback_available: 0" in rollback_report_text
    assert rollback_result["audit_record_id"] in rollback_report_text
    assert "- rollback_available: false" in rollback_report_text
    assert rollback_docx.is_file()
    assert len(Document(rollback_docx).inline_shapes) == 1
    with zipfile.ZipFile(rollback_docx) as archive:
        rollback_media = [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]
    assert current_ready.read_bytes() in rollback_media
    assert replacement_ready.read_bytes() not in rollback_media

    rollback_audit_payload = json.loads(audit_path.read_text(encoding="utf-8"))
    assert rollback_audit_payload["entry_count"] == 2
    assert [
        record["action"] for record in rollback_audit_payload["records"]
    ] == [
        "question_figure_repair_batch_apply_word_media_write",
        "question_figure_repair_batch_apply_word_media_rollback",
    ]
    rollback_audit_record = rollback_audit_payload["records"][1]
    assert rollback_audit_record["audit_id"] == rollback_result["audit_record_id"]
    assert rollback_audit_record["linked_apply_audit_id"] == result[
        "audit_record_id"
    ]
    assert rollback_audit_record["output_path"] == str(rollback_docx)
    assert rollback_audit_record["restored_count"] == 1

    rollback_markdown = _render_batch_markdown_report(
        {
            "status": "partial_success",
            "input_path": "source.docx",
            "failed_count": 0,
            "items": [],
            "question_figure_repair_queue": rolled_back,
        }
    )
    assert "Batch apply rollback result: status=rolled_back" in rollback_markdown
    assert "restored=1" in rollback_markdown
    assert "word_write=true" in rollback_markdown
    assert "audit=written" in rollback_markdown
    assert "Batch apply transaction manifest: status=tracked" in rollback_markdown
    assert "transactions=1" in rollback_markdown
    assert "task=completed" in rollback_markdown
    assert "rollback_available=0" in rollback_markdown
    assert "artifact=written" in rollback_markdown
    assert "report=written" in rollback_markdown
    assert "rolled_back=1" in rollback_markdown
    assert "artifact=written" in rollback_markdown

    blocked_rollback_docx = tmp_path / "blocked_rollback.docx"
    blocked_rollback = _rollback_question_figure_repair_batch_apply_from_audit(
        applied,
        applied_docx_path=output_docx,
        rollback_docx_path=blocked_rollback_docx,
        confirmed=True,
        expected_audit_id="stale-audit",
        expected_output_sha1=result["output_sha1"],
    )
    blocked_rollback_result = blocked_rollback["batch_apply_rollback_result"]
    assert blocked_rollback_result["status"] == "blocked"
    assert blocked_rollback_result["word_write_enabled"] is False
    assert blocked_rollback_result["audit_written"] is False
    assert "audit_id_mismatch" in blocked_rollback_result["blockers"]
    assert not blocked_rollback_docx.exists()

    blocked = _apply_question_figure_repair_batch_execution_to_docx(
        planned,
        input_docx_path=source_docx,
        output_docx_path=tmp_path / "blocked.docx",
        confirmed=True,
        expected_plan_id="stale-plan",
        expected_execution_plan_fingerprint=execution_plan[
            "execution_plan_fingerprint"
        ],
    )
    blocked_result = blocked["batch_apply_execution_result"]
    assert blocked_result["status"] == "blocked"
    assert blocked_result["word_write_enabled"] is False
    assert blocked_result["audit_written"] is False
    assert "plan_id_mismatch" in blocked_result["blockers"]
    assert not (tmp_path / "blocked.docx").exists()


def test_question_figure_repair_queue_resolves_selected_conflict_candidate(tmp_path):
    current = tmp_path / "question_2.png"
    replacement_a = tmp_path / "question_2_expected_a.png"
    replacement_b = tmp_path / "question_2_expected_b.png"
    Image.new("RGB", (320, 240), color="blue").save(current)
    Image.new("RGB", (320, 240), color="yellow").save(replacement_a)
    Image.new("RGB", (320, 240), color="red").save(replacement_b)
    target_key = json.dumps(
        {
            "role": "question_figure",
            "item_id": "question_figure_2",
            "question_index": "2",
            "path": str(current),
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )

    queue = _batch_question_figure_repair_queue(
        [
            {
                "issue_id": "issue-a",
                "kind": "manual_comparison_issue",
                "profile_id": "exam_a",
                "profile_name": "Exam A",
                "repair_target_type": "question_figure_item",
                "repair_target_key": target_key,
                "comparison_issue_items": [
                    {
                        "item_id": "question_figure_2",
                        "label": "Question figure 2",
                        "path": str(current),
                        "question_index": "2",
                        "comparison_issue_reference": str(replacement_a),
                        "comparison_issue_display_name": replacement_a.name,
                    },
                    {
                        "item_id": "question_figure_2",
                        "label": "Question figure 2",
                        "path": str(current),
                        "question_index": "2",
                        "comparison_issue_reference": str(replacement_b),
                        "comparison_issue_display_name": replacement_b.name,
                    },
                ],
            }
        ]
    )
    selected_queue_id = next(
        entry["queue_id"]
        for entry in queue["entries"]
        if entry["replacement_source_path"] == str(replacement_b)
    )

    not_confirmed = _resolve_question_figure_repair_queue_conflict(
        queue,
        selected_queue_id,
    )
    assert not_confirmed["resolution_status"] == "not_confirmed"
    assert not_confirmed["status"] == "conflict"
    assert {entry["confirmation_status"] for entry in queue["entries"]} == {
        "conflict"
    }

    resolved = _resolve_question_figure_repair_queue_conflict(
        queue,
        selected_queue_id,
        confirmed=True,
        resolved_by="manual-reviewer",
        resolution_note="choose clearer reference",
    )

    assert resolved["kind"] == "question_figure_repair_queue"
    assert resolved["status"] == "queued"
    assert resolved["queue_count"] == 2
    assert resolved["conflict_group_count"] == 0
    assert resolved["conflict_count"] == 0
    assert resolved["resolution_status"] == "resolved"
    assert resolved["resolved_conflict_group_count"] == 1
    assert resolved["resolved_conflict_count"] == 2
    assert len(resolved["resolution_records"]) == 1
    resolution_record = resolved["resolution_records"][0]
    assert resolution_record["action"] == "resolve_question_figure_repair_conflict"
    assert resolution_record["status"] == "resolved"
    assert resolution_record["selected_queue_id"] == selected_queue_id
    assert resolution_record["resolved_by"] == "manual-reviewer"
    assert resolution_record["resolution_note"] == "choose clearer reference"
    assert len(resolution_record["candidate_queue_ids"]) == 2
    assert len(resolution_record["rejected_queue_ids"]) == 1

    selected = next(
        entry
        for entry in resolved["entries"]
        if entry["queue_id"] == selected_queue_id
    )
    rejected = next(
        entry
        for entry in resolved["entries"]
        if entry["queue_id"] != selected_queue_id
    )
    assert selected["confirmation_status"] == "ready"
    assert selected["confirmation_action"] == "confirm_question_figure_replacement"
    assert selected["confirmation_apply_supported"] is True
    assert selected["conflict_resolution_status"] == "selected"
    assert selected["apply_blockers"] == []
    assert rejected["status"] == "rejected"
    assert rejected["confirmation_status"] == "rejected"
    assert rejected["confirmation_apply_supported"] is False
    assert rejected["conflict_resolution_status"] == "rejected"
    assert "candidate_rejected_by_conflict_resolution" in rejected["apply_blockers"]
    plan = resolved["batch_confirmation_plan"]
    assert plan["status"] == "partial"
    assert plan["plan_fingerprint"]
    assert plan["eligible_queue_ids"] == [selected_queue_id]
    assert rejected["queue_id"] in plan["blocked_queue_ids"]
