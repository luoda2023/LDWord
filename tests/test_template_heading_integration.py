import sys
from dataclasses import asdict
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.builtin_templates import create_builtin_template
from src.config.resolver import resolve_config
from src.config.scene import SceneWorkspace
from src.config.template import HeadingLevelBindingConfig
from src.qt_api import QApplication
from src.pipeline.runner import Pipeline
from src.ui.bridge import PanelBridge
from src.ui.panels.template_panel import TemplatePanel
from src.ui.panels.template_preview.model import PreviewBlockKind
from src.config.loader import load_template
from src.modules.basic.paragraph_style import ParagraphStyleModule
from src.modules.structure.heading_numbering import HeadingNumberingModule
from src.modules.structure.heading_recognition import HeadingRecognitionModule


def _app():
    return QApplication.instance() or QApplication([])


def test_heading_detail_updates_heading_preview_summary_when_levels_change():
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        before = panel._nav_cards["tpl_heading"]._full_subtitle

        panel._heading_detail._levels_slider.setValue(3)
        app.processEvents()

        after = panel._nav_cards["tpl_heading"]._full_subtitle

        assert panel._current_template.heading_model.max_heading_levels == 3
        assert before != after
        assert "3" in after
        assert bridge.is_template_dirty() is True
    finally:
        panel.close()
        app.processEvents()


def test_projected_overview_matches_heading_panel_on_initial_load():
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        panel.show()
        app.processEvents()

        overview_preview = panel._overview_detail._preview.projection
        heading_adapter = panel._heading_detail._adapter

        heading1_text = next(
            block.text
            for block in overview_preview.blocks
            if block.kind is PreviewBlockKind.HEADING and block.level == 1
        )
        expected_text = heading_adapter.preview_heading_text(1, "绪论")

        assert heading1_text == expected_text
    finally:
        panel.close()
        app.processEvents()


def test_special_title_editor_links_one_scope_option_through_its_lifecycle():
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        from src.config.special_title_rules import (
            special_title_selector,
            special_title_selector_options,
        )

        heading = panel._heading_detail
        heading._nn_texts_edit.setText("摘要, 目录, 参考文献, 缩略语表, 鸣谢")
        app.processEvents()

        assert "鸣谢" in panel._current_template.heading_model.non_numbered_title_texts
        selector = special_title_selector("exact", "鸣谢")
        dynamic_selectors = {
            option_selector
            for option_selector, _label in special_title_selector_options(
                panel._current_template.heading_model
            )
        }
        assert set(
            panel._header_footer_detail
            ._header_scope_editor
            ._selector_buttons
        ) == {"cover", "body", *dynamic_selectors}
        assert set(
            panel._header_footer_detail
            ._footer_scope_editor
            ._selector_buttons
        ) == {"cover", "body", *dynamic_selectors}
        expected_page_selectors = {
            "all_numbered_content",
            "front_matter",
            "body",
            "back_matter",
            "cover",
            *dynamic_selectors,
        }
        assert all(
            set(row.selector_editor._selector_buttons) == expected_page_selectors
            for row in panel._header_footer_detail._page_phase_rows
        )
        assert not {
            "statement",
            "authorization",
            "front_note",
        }.intersection(expected_page_selectors)
        assert (
            selector
            in panel._header_footer_detail._header_scope_editor._selector_buttons
        )
        assert (
            selector
            in panel._header_footer_detail._footer_scope_editor._selector_buttons
        )
        assert all(
            selector in row.selector_editor._selector_buttons
            for row in panel._header_footer_detail._page_phase_rows
        )
        assert (
            panel._header_footer_detail
            ._header_scope_editor
            ._selector_buttons[selector]
            .isChecked()
            is False
        )
        assert all(
            row.selector_editor._selector_buttons[selector].isChecked() is False
            for row in panel._header_footer_detail._page_phase_rows
        )

        panel._header_footer_detail._header_scope_editor._selector_buttons[
            selector
        ].click()
        footer_mode_combo = panel._header_footer_detail._footer_content_combo
        footer_mode_index = footer_mode_combo.findData("fixed")
        footer_mode_combo.setCurrentIndex(
            footer_mode_index
        )
        app.processEvents()
        panel._header_footer_detail._footer_scope_editor._selector_buttons[
            selector
        ].click()
        panel._header_footer_detail._page_phase_rows[0].selector_editor._selector_buttons[
            selector
        ].click()
        app.processEvents()

        renamed_selector = special_title_selector("exact", "鸣谢说明")
        heading._nn_texts_edit.setText(
            "摘要, 目录, 参考文献, 缩略语表, 鸣谢说明"
        )
        app.processEvents()

        assert (
            selector
            not in panel._header_footer_detail._header_scope_editor._selector_buttons
        )
        assert (
            panel._header_footer_detail
            ._header_scope_editor
            ._selector_buttons[renamed_selector]
            .isChecked()
            is True
        )
        assert (
            panel._header_footer_detail
            ._footer_scope_editor
            ._selector_buttons[renamed_selector]
            .isChecked()
            is True
        )
        assert (
            panel._header_footer_detail
            ._page_phase_rows[0]
            .selector_editor
            ._selector_buttons[renamed_selector]
            .isChecked()
            is True
        )

        heading._nn_texts_edit.setText("摘要, 目录, 参考文献, 缩略语表")
        app.processEvents()

        assert (
            renamed_selector
            not in panel._header_footer_detail._header_scope_editor._selector_buttons
        )
        assert (
            renamed_selector
            not in panel._current_template.header_footer.header.hidden_selectors
        )
        assert (
            renamed_selector
            not in panel._current_template.header_footer.footer.hidden_selectors
        )
        assert all(
            renamed_selector not in phase.selectors
            for phase in panel._current_template.header_footer.page_number_plan.phases
        )
        assert not hasattr(heading, "_nn_role_edits")
        assert bridge.is_template_dirty() is True
    finally:
        panel.close()
        app.processEvents()


def test_heading_detail_restore_recomputes_template_dirty_state():
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        heading = panel._heading_detail
        panel.show()
        app.processEvents()
        heading._adv_list.setCurrentRow(0)
        app.processEvents()

        heading._hd_size_combo.setEditText("18")
        app.processEvents()

        assert bridge.is_template_dirty() is True
        assert heading._adapter.has_unsaved_changes is True

        heading._on_restore()
        app.processEvents()

        assert bridge.is_template_dirty() is False
        assert heading._adapter.has_unsaved_changes is False
    finally:
        panel.close()
        app.processEvents()


def test_heading_detail_preset_change_enters_save_restore_state_machine():
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        heading = panel._heading_detail
        original_bindings = {
            key: asdict(value)
            for key, value in panel._current_template.heading_numbering.level_bindings.items()
        }

        target_index = None
        active_preset = heading._adapter.detect_active_preset()
        for index in range(heading._preset_cb.count()):
            key = heading._preset_cb.itemData(index)
            if key and key != active_preset:
                target_index = index
                break

        assert target_index is not None
        assert bridge.is_template_dirty() is False
        assert heading._restore_btn.isEnabled() is False
        assert heading._save_btn.isEnabled() is False

        heading._preset_cb.setCurrentIndex(target_index)
        app.processEvents()

        assert bridge.is_template_dirty() is True
        assert heading._adapter.has_unsaved_changes is True
        assert heading._restore_btn.isEnabled() is True
        assert heading._save_btn.isEnabled() is True

        heading._restore_btn.click()
        app.processEvents()

        restored_bindings = {
            key: asdict(value)
            for key, value in panel._current_template.heading_numbering.level_bindings.items()
        }
        assert restored_bindings == original_bindings
        assert bridge.is_template_dirty() is False
        assert heading._adapter.has_unsaved_changes is False
        assert heading._restore_btn.isEnabled() is False
        assert heading._save_btn.isEnabled() is False
    finally:
        panel.close()
        app.processEvents()


def test_heading_detail_save_button_persists_template_via_template_panel(tmp_path):
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        target = tmp_path / "heading_saved_template.json"
        panel._current_template_id = ""
        panel._current_template_path = str(target)
        panel._current_template_source = "file"
        panel._sync_template_file_status()

        heading = panel._heading_detail
        heading._adv_list.setCurrentRow(0)
        app.processEvents()

        heading._hd_size_combo.setEditText("18")
        app.processEvents()

        assert bridge.is_template_dirty() is True
        assert heading._adapter.has_unsaved_changes is True

        heading._save_btn.click()
        app.processEvents()

        assert target.exists()
        assert bridge.is_template_dirty() is False
        assert heading._adapter.has_unsaved_changes is False

        reloaded = load_template(target)
        assert reloaded.styles["heading1"].size_pt == 18.0
    finally:
        panel.close()
        app.processEvents()


def test_heading_detail_updates_overview_preview_on_every_style_edit():
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        panel.show()
        app.processEvents()
        preview = panel._overview_detail._preview

        def heading_size_pt():
            for block in preview.projection.blocks:
                if block.kind is PreviewBlockKind.HEADING and block.level == 1:
                    return block.style.size_pt
            raise AssertionError("heading1 preview line not found")

        heading = panel._heading_detail
        heading._adv_list.setCurrentRow(0)
        app.processEvents()

        px0 = heading_size_pt()
        heading._hd_size_combo.setEditText("18")
        app.processEvents()
        px18 = heading_size_pt()
        heading._hd_size_combo.setEditText("24")
        app.processEvents()
        px24 = heading_size_pt()

        assert px18 > px0
        assert px24 > px18
    finally:
        panel.close()
        app.processEvents()


def test_heading_detail_uses_preview_first_inspector_without_explanation_text_or_action_button():
    """Heading detail should expose preview first, with heading style always available."""
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        heading = panel._heading_detail
        heading._adv_list.setCurrentRow(0)
        app.processEvents()

        # Preview-first inspector should be available immediately without explanatory text blocks.
        assert not heading._result_heading_label.isHidden()
        assert not hasattr(heading, "_result_meta_label")
        assert not hasattr(heading, "_detail_hint")
        assert not hasattr(heading, "_style_summary_primary")
        assert not hasattr(heading, "_expert_summary_primary")
        assert not hasattr(heading, "_ref_style_hint_label")
        assert not hasattr(heading, "_style_toggle_btn")
        assert not heading._style_editor_container.isHidden()
        assert heading._expert_editor_container.isHidden()

        heading._expert_toggle_btn.click()
        app.processEvents()
        assert not heading._expert_editor_container.isHidden()

        # Font combo and numbering controls should still exist and be accessible
        assert heading._hd_font_cn is not None
        assert heading._hd_size_combo is not None
        assert heading._core_style_cb is not None

        # No override action buttons should exist
        assert not hasattr(heading, '_style_action_btn')
        assert not hasattr(heading, '_numbering_action_btn')
    finally:
        panel.close()
        app.processEvents()


def test_heading_preview_text_matches_runtime_number_separator_once(tmp_path):
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        template = panel._current_template
        binding = HeadingLevelBindingConfig(enabled=True)
        template.heading_numbering.level_bindings["heading1"] = binding
        binding.display_core_style = "arabic"
        binding.display_template = "{nn}"
        binding.title_separator = "::"

        panel._on_template_edited(template)
        heading_text = next(
            block.text
            for block in panel._overview_detail._preview.projection.blocks
            if block.kind is PreviewBlockKind.HEADING and block.level == 1
        )

        assert heading_text == "1::绪论"

        # Verify adapter preview_heading_text matches
        adapter = panel._heading_detail._adapter
        built = adapter.preview_heading_text(1, "绪论")
        assert built == "1::绪论"

        source = tmp_path / "heading_separator_once.docx"
        doc = Document()
        doc.add_heading("绪论", level=1)
        doc.save(source)

        config = resolve_config(template, SceneWorkspace())
        pipeline = Pipeline(
            modules=[HeadingRecognitionModule(), HeadingNumberingModule(), ParagraphStyleModule()],
            config=config,
        )
        result = pipeline.execute(str(source))

        assert result.success, f"执行失败: {result.error}"

        out_doc = Document(result.output_paths["final"])
        from src.shared.engine.heading_numbering_ooxml import (
            effective_numbering,
            managed_level_text,
        )

        assert out_doc.paragraphs[0].text == "绪论"
        assert effective_numbering(out_doc.paragraphs[0]) is not None
        assert managed_level_text(out_doc, 1) == "%1::"
    finally:
        panel.close()
        app.processEvents()


def test_heading_detail_style_edit_flows_into_final_document_pipeline(tmp_path):
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        heading = panel._heading_detail
        heading._adv_list.setCurrentRow(0)
        app.processEvents()
        heading._hd_size_combo.setEditText("18")
        app.processEvents()

        source = tmp_path / "heading_input.docx"
        doc = Document()
        doc.add_heading("绪论", level=1)
        doc.save(source)

        config = resolve_config(panel._current_template, SceneWorkspace())
        pipeline = Pipeline(modules=[ParagraphStyleModule()], config=config)
        result = pipeline.execute(str(source))

        assert result.success, f"执行失败: {result.error}"

        out_doc = Document(result.output_paths["final"])
        run = out_doc.paragraphs[0].runs[0]
        assert run.font.size is not None
        assert round(run.font.size.pt, 1) == 18.0
    finally:
        panel.close()
        app.processEvents()


def test_default_builtin_template_applies_heading_style_in_pipeline(tmp_path):
    source = tmp_path / "default_heading_input.docx"
    doc = Document()
    doc.add_heading("示例一级标题", level=1)
    doc.save(source)

    config = resolve_config(create_builtin_template("default"), SceneWorkspace())
    pipeline = Pipeline(modules=[ParagraphStyleModule()], config=config)
    result = pipeline.execute(str(source))

    assert result.success, f"执行失败: {result.error}"

    out_doc = Document(result.output_paths["final"])
    para = out_doc.paragraphs[0]
    run = para.runs[0]

    assert run.font.size is not None
    assert round(run.font.size.pt, 1) == 16.0
    assert run.font.bold is True
    assert para.paragraph_format.first_line_indent is not None
    assert round(para.paragraph_format.first_line_indent.pt, 1) == 0.0


def test_heading_detail_spacing_edit_writes_back_to_template_style():
    app = _app()
    bridge = PanelBridge()
    panel = TemplatePanel(bridge)

    try:
        heading = panel._heading_detail
        heading._adv_list.setCurrentRow(0)
        app.processEvents()

        heading._hd_space_before.set_value(12, "pt")
        heading._hd_space_after.set_value(6, "pt")
        app.processEvents()

        style = panel._current_template.styles["heading1"]
        assert style.space_before_pt == 12
        assert style.space_after_pt == 6
    finally:
        panel.close()
        app.processEvents()
