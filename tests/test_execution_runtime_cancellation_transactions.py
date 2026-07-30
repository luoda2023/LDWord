from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image

from src.config.entity import EntityArchive
from src.config.feature_configs import OutputConfig
from src.config.material_context import MaterialExecutionContext
from src.config.materials import AssetItem
from src.config.object_preflight_evidence import build_object_preflight_evidence
from src.config.resolver import resolve_config
from src.config.scene import (
    DeliveryPreset,
    ExamPaperConfig,
    InputSourceProfile,
    SceneWorkspace,
)
from src.config.template import TemplateConfig
from src.pipeline.result import PipelineResult
from src.pipeline.runner import Pipeline
from src.shared.engine.official_document_assembly import (
    OfficialDocumentAssemblyResult,
)
from src.shared.engine.official_document_batch_history import (
    load_official_document_batch_history,
)
import src.pipeline.runner as pipeline_runner_module
import src.config.material_context as material_context_module
import src.services.production_runtime.execution_runtime as execution_runtime
from src.services.production_runtime import batch_reporting, delivery_reporting
from src.services.production_runtime import execution_preflight
from src.services.production_runtime.execution_runtime import (
    WorkbenchBatchProductionRunner,
    WorkbenchProductionRunner,
)
from src.services.execution_session import (
    build_execution_session_snapshot,
    cleanup_execution_session_resources,
)


def _official_config():
    scene = SceneWorkspace(
        scene_id="official",
        mode_id="official",
        input_source_profile=InputSourceProfile(
            material_schema_id="official_document_v1",
        ),
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                artifacts=OutputConfig(
                    final_docx=False,
                    report_json=False,
                    report_markdown=False,
                ),
            )
        ],
    )
    return resolve_config(
        TemplateConfig(),
        scene,
        entity_data={
            "document_type": "notice",
            "title": "Cancellation evidence",
            "body": "Published before cancellation was observed.",
            "organization": "Example organization",
            "document_no": "EX-2026-01",
            "issue_date": "2026-07-14",
        },
    )


def _cancelled_evidence_config():
    scene = SceneWorkspace(
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                filename_template="{stem}_review.docx",
                artifacts=OutputConfig(
                    final_docx=True,
                    compare_docx=True,
                    report_json=True,
                    report_markdown=False,
                ),
                include_structured_intermediate=True,
            )
        ],
    )
    return resolve_config(TemplateConfig(), scene)


def _exam_execution_session(
    *,
    scene: SceneWorkspace,
    template: TemplateConfig,
    input_path: Path,
    output_root: Path,
    material_context: MaterialExecutionContext,
):
    evidence = build_object_preflight_evidence(scene, input_path)
    return build_execution_session_snapshot(
        mode_id="exam",
        scene=scene,
        template=template,
        material_context=material_context,
        input_path=input_path,
        output_root=output_root,
        plan_id="exam",
        template_id="default",
        object_preflight_confirmation_revision=evidence.source_revision,
        object_preflight_confirmation_digest=evidence.evidence_digest,
    )


def test_finalize_cancelled_empty_outputs_still_invokes_configured_evidence_only(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    config = _cancelled_evidence_config()
    result = PipelineResult(
        success=False,
        status="cancelled",
        cancelled=True,
        output_paths={},
        config=config,
    )
    calls: list[str] = []

    monkeypatch.setattr(
        delivery_reporting,
        "write_enabled_result_reports",
        lambda *_args, **_kwargs: calls.append("reports") or ["cancelled.json"],
    )
    monkeypatch.setattr(
        delivery_reporting,
        "write_structured_intermediates",
        lambda *_args, **_kwargs: calls.append("intermediates")
        or {"review": "cancelled_intermediate.json"},
    )
    monkeypatch.setattr(
        delivery_reporting,
        "_write_compare_docx_artifacts",
        lambda *_args, **_kwargs: calls.append("compare") or {},
    )
    monkeypatch.setattr(
        delivery_reporting,
        "write_material_manifest",
        lambda *_args, **_kwargs: calls.append("material") or {},
    )

    outcome = delivery_reporting.finalize_result_artifacts(
        result,
        input_path=source,
        output_dir=tmp_path / "out",
        output_paths={},
        fallback_output_path="",
        config=config,
        elapsed=0.0,
        modules_enabled=0,
        modules_total=0,
        material_diagnostics=[],
        material_context=MaterialExecutionContext(),
        style_source_summary=None,
        include_material_artifacts=True,
    )

    assert calls == ["reports", "intermediates"]
    assert outcome.report_paths == ["cancelled.json"]
    assert outcome.intermediate_paths == {
        "review": "cancelled_intermediate.json"
    }
    assert outcome.compare_paths == {}
    assert outcome.material_manifest_paths == {}
    assert outcome.material_package_paths == {}


def test_single_cancelled_payload_with_no_outputs_returns_report_and_intermediate(
    tmp_path,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "out"
    config = _cancelled_evidence_config()
    result = PipelineResult(
        success=False,
        status="cancelled",
        cancelled=True,
        output_paths={},
        config=config,
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(),
        output_dir=output_dir,
    )._payload_from_result(
        result,
        input_path=source,
        output_dir=output_dir,
        config=config,
        elapsed=0.0,
        modules_enabled=0,
        modules_total=0,
    )

    assert payload["status"] == "cancelled"
    assert payload["output_paths"] == {}
    assert [Path(path).name for path in payload["report_paths"]] == [
        "source_review_changes.json"
    ]
    assert payload["intermediate_paths"] == {
        "review": str(output_dir / "source_review_intermediate.json")
    }
    assert all(Path(path).is_file() for path in payload["report_paths"])
    intermediate_path = Path(payload["intermediate_paths"]["review"])
    assert intermediate_path.is_file()
    assert json.loads(intermediate_path.read_text(encoding="utf-8"))["status"] == (
        "cancelled"
    )
    assert payload["compare_paths"] == {}
    assert payload["material_manifest_paths"] == {}
    assert payload["material_package_paths"] == {}


def test_generic_runner_immediate_cancel_precedes_material_preflight(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = SceneWorkspace(
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                artifacts=OutputConfig(
                    final_docx=False,
                    report_json=False,
                    report_markdown=False,
                ),
            )
        ]
    )

    def _unexpected_material_evaluation(*_args, **_kwargs):
        pytest.fail("immediate cancellation must precede material evaluation")

    monkeypatch.setattr(
        execution_runtime,
        "material_context_with_exam_question_assets",
        _unexpected_material_evaluation,
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=MaterialExecutionContext(),
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: True)

    assert payload["status"] == "cancelled"
    assert payload["output_paths"] == {}
    assert payload["material_diagnostics"] == []


@pytest.mark.parametrize(
    "cancel_stage",
    ("before_material", "after_parse", "after_validation"),
)
def test_exam_markdown_early_cancel_always_writes_configured_evidence(
    tmp_path,
    cancel_stage,
):
    source = tmp_path / "exam.md"
    source.write_text(
        """# 数学测试

## 一、填空题

1. 1 + 1 = ______。（5 分）

## 答案速查

一、填空题
1. 2
""",
        encoding="utf-8",
    )
    output_dir = tmp_path / f"out-{cancel_stage}"
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        master_id="default_exam",
        category="exam_paper",
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                filename_template="{stem}_review.docx",
                artifacts=OutputConfig(
                    final_docx=False,
                    compare_docx=True,
                    report_json=True,
                    report_markdown=False,
                ),
                include_structured_intermediate=True,
            )
        ],
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown"],
            structured_formats=["json"],
        ),
        exam_paper=ExamPaperConfig(answer_policy="student_plus_answer"),
    )
    reached = {"parse": False, "validation": False}

    def _cancel_at_stage() -> bool:
        if cancel_stage == "before_material":
            return True
        if cancel_stage == "after_parse":
            return reached["parse"]
        return reached["validation"]

    def _record_progress(_current, _total, message) -> None:
        if message == "解析 Markdown 题稿":
            reached["parse"] = True
        elif message == "校验试卷题源":
            reached["validation"] = True

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_dir,
    ).run(_record_progress, _cancel_at_stage)

    assert payload["status"] == "cancelled"
    assert payload["output_paths"] == {}
    assert len(payload["report_paths"]) == 1
    assert all(Path(path).is_file() for path in payload["report_paths"])
    assert payload["intermediate_paths"] == {
        "review": str(output_dir / "exam_review_intermediate.json")
    }
    intermediate_path = Path(payload["intermediate_paths"]["review"])
    assert intermediate_path.is_file()
    assert json.loads(intermediate_path.read_text(encoding="utf-8"))["status"] == (
        "cancelled"
    )
    assert payload["compare_paths"] == {}
    assert payload["material_manifest_paths"] == {}
    assert payload["material_package_paths"] == {}
    if cancel_stage == "before_material":
        assert "exam_markdown_import" not in payload
    else:
        assert payload["exam_markdown_import"]["summary"]["question_count"] == 1


def test_exam_markdown_cancel_fallback_does_not_invent_artifact_authorization(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "cancel-before-config.md"
    source.write_text("# 已取消题稿\n", encoding="utf-8")
    output_dir = tmp_path / "cancel-before-config-out"

    def _reject_config(*_args, **_kwargs):
        raise ValueError("invalid session override")

    monkeypatch.setattr(execution_runtime, "resolve_config", _reject_config)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(
            scene_id="exam",
            mode_id="exam",
            category="exam_paper",
            input_source_profile=InputSourceProfile(
                accepted_formats=["markdown"],
            ),
            exam_paper=ExamPaperConfig(),
        ),
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: True)

    assert payload["status"] == "cancelled"
    assert payload["artifact_failure_count"] == 1
    assert payload["artifact_failures"][0]["kind"] == "report_config"
    assert payload["artifact_failures"][0]["error_type"] == "ValueError"
    assert payload["artifact_failures"][0]["fallback"] == "template_baseline"
    assert payload["report_paths"] == []
    assert payload["intermediate_paths"] == {}


def test_exam_markdown_cancel_config_retry_preserves_scene_evidence_policy(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "cancel-with-scene-policy.md"
    source.write_text("# 已取消题稿\n", encoding="utf-8")
    output_dir = tmp_path / "cancel-with-scene-policy-out"
    real_resolve_config = execution_runtime.resolve_config
    resolve_calls = 0

    def _reject_session_once(*args, **kwargs):
        nonlocal resolve_calls
        resolve_calls += 1
        if resolve_calls == 1:
            raise ValueError("invalid session override")
        return real_resolve_config(*args, **kwargs)

    monkeypatch.setattr(
        execution_runtime,
        "resolve_config",
        _reject_session_once,
    )
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        master_id="default_exam",
        category="exam_paper",
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                filename_template="{stem}_review.docx",
                artifacts=OutputConfig(
                    final_docx=False,
                    report_json=True,
                    report_markdown=False,
                ),
                include_structured_intermediate=True,
            )
        ],
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown"],
        ),
        exam_paper=ExamPaperConfig(),
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        session_overrides={"output.report_json": False},
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: True)

    assert payload["status"] == "cancelled"
    assert resolve_calls == 2
    assert payload["artifact_failures"][0]["fallback"] == (
        "scene_without_session_overrides"
    )
    assert [Path(path).name for path in payload["report_paths"]] == [
        "cancel-with-scene-policy_review_changes.json"
    ]
    assert payload["intermediate_paths"] == {
        "review": str(
            output_dir / "cancel-with-scene-policy_review_intermediate.json"
        )
    }


def test_exam_markdown_cancel_does_not_write_reports_without_scene_policy(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "cancel-with-two-failures.md"
    source.write_text("# 已取消题稿\n", encoding="utf-8")

    def _reject_config(*_args, **_kwargs):
        raise ValueError("invalid session override")

    monkeypatch.setattr(execution_runtime, "resolve_config", _reject_config)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(
            scene_id="exam",
            mode_id="exam",
            category="exam_paper",
            input_source_profile=InputSourceProfile(
                accepted_formats=["markdown"],
            ),
            exam_paper=ExamPaperConfig(),
        ),
        output_dir=tmp_path / "cancel-with-two-failures-out",
    ).run(lambda *_args: None, lambda: True)

    assert payload["status"] == "cancelled"
    assert payload["artifact_failure_count"] == 1
    assert payload["artifact_failures"][0]["kind"] == "report_config"
    assert payload["report_paths"] == []
    assert payload["error_text"].count("invalid session override") == 1


def test_exam_markdown_cancel_after_parse_takes_precedence_over_blocker(
    tmp_path,
):
    source = tmp_path / "blocked-exam.md"
    source.write_text("# 空题稿\n", encoding="utf-8")
    output_dir = tmp_path / "blocked-out"
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        category="exam_paper",
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                artifacts=OutputConfig(final_docx=False, report_json=True),
                include_structured_intermediate=True,
            )
        ],
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown"],
            required_material_fields=["required_but_missing"],
            failure_policy="block",
        ),
        exam_paper=ExamPaperConfig(),
    )
    parse_started = False

    def _cancel_after_parse() -> bool:
        return parse_started

    def _record_parse_start(_current, _total, message) -> None:
        nonlocal parse_started
        if message == "解析 Markdown 题稿":
            parse_started = True

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_dir,
    ).run(_record_parse_start, _cancel_after_parse)

    assert payload["status"] == "cancelled"
    assert any(
        item.get("target") == "required_fields"
        for item in payload["material_diagnostics"]
    )
    assert payload["report_paths"]
    assert payload["intermediate_paths"]


def test_exam_markdown_resolves_dynamic_material_fields_only_once(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "frozen-fields.md"
    source.write_text(
        """# 数学测试

## 一、填空题

1. 1 + 1 = ______。（5 分）

## 答案速查

一、填空题
1. 2
""",
        encoding="utf-8",
    )
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        category="exam_paper",
        template_id="default",
        compatible_template_ids=["default"],
        master_id="default_exam",
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown"],
            material_schema_id="exam_items_v1",
        ),
        exam_paper=ExamPaperConfig(),
    )
    figure_path = tmp_path / "question-figure.png"
    Image.new("RGB", (64, 48), color="green").save(figure_path)
    material_context = MaterialExecutionContext(
        field_functions={"generated_at": {"function": "current_datetime"}},
        asset_items=[
            AssetItem(
                role="question_figure",
                label="题图",
                path=str(figure_path),
            )
        ],
    )
    real_resolver = material_context_module.resolve_field_functions
    provider_calls = 0

    def _counted_resolver(*args, **kwargs):
        nonlocal provider_calls
        provider_calls += 1
        return real_resolver(*args, **kwargs)

    monkeypatch.setattr(
        material_context_module,
        "resolve_field_functions",
        _counted_resolver,
    )

    template = TemplateConfig()
    execution_session = _exam_execution_session(
        scene=scene,
        template=template,
        input_path=source,
        output_root=tmp_path / "frozen-out",
        material_context=material_context,
    )
    try:
        payload = WorkbenchProductionRunner(
            doc_path=str(source),
            template=template,
            scene=scene,
            material_context=material_context,
            output_dir=Path(execution_session.output_namespace),
            execution_session=execution_session,
        ).run(lambda *_args: None, lambda: False)
    finally:
        cleanup_execution_session_resources(execution_session)

    assert payload["status"] in {"success", "partial_success"}
    assert provider_calls == 1
    rendered_versions = payload["exam_delivery_runtime"]["rendered_versions"]
    assert rendered_versions[0]["question_asset_count"] == 1
    assert rendered_versions[0]["rendered_question_asset_count"] == 1


def test_official_terminal_cancellation_after_assembly_preserves_published_evidence(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "official.docx"
    Document().save(source)
    output_dir = tmp_path / "out"
    published = output_dir / "official_official.docx"
    state = {"assembled": False}

    def _assemble(*_args, **_kwargs):
        published.parent.mkdir(parents=True, exist_ok=True)
        Document().save(published)
        state["assembled"] = True
        return OfficialDocumentAssemblyResult(
            status="ok",
            profile_id="notice",
            docx_path=published,
        )

    monkeypatch.setattr(
        pipeline_runner_module,
        "assemble_official_document_docx",
        _assemble,
    )
    config = _official_config()
    result = Pipeline(
        modules=[],
        config=config,
        output_dir=output_dir,
        cancel_check=lambda: state["assembled"],
        official_document_type_id="notice",
    ).execute(source)

    assert result.status == "cancelled"
    assert result.cancelled is True
    assert result.output_paths == {"official_docx": str(published)}
    assert result.context.official_document_assembly.status == "ok"
    assert published.exists()

    runner = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(),
        output_dir=output_dir,
    )
    payload = runner._payload_from_result(
        result,
        input_path=source,
        output_dir=output_dir,
        config=config,
        elapsed=0.0,
        modules_enabled=0,
        modules_total=0,
    )

    assert payload["status"] == "cancelled"
    assert payload["output_path"] == str(published)
    assert payload["output_paths"] == {"official_docx": str(published)}
    assert payload["official_document_assembly"]["status"] == "ok"


def test_exam_markdown_cancellation_after_assembly_preserves_runtime_outputs(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "exam.md"
    source.write_text(
        """# 数学测试

## 一、填空题

1. 1 + 1 = ______。（5 分）

## 答案速查

一、填空题
1. 2
""",
        encoding="utf-8",
    )
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        category="exam_paper",
        template_id="default",
        compatible_template_ids=["default"],
        master_id="default_exam",
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown"],
            structured_formats=["json"],
        ),
        exam_paper=ExamPaperConfig(answer_policy="student_plus_answer"),
    )
    original_builder = execution_runtime.build_exam_delivery_runtime
    state = {"assembled": False}

    def _build_and_cancel(*args, **kwargs):
        result = original_builder(*args, **kwargs)
        state["assembled"] = True
        return result

    monkeypatch.setattr(
        execution_runtime,
        "build_exam_delivery_runtime",
        _build_and_cancel,
    )

    template = TemplateConfig()
    material_context = MaterialExecutionContext(mode_id="exam")
    execution_session = _exam_execution_session(
        scene=scene,
        template=template,
        input_path=source,
        output_root=tmp_path / "out",
        material_context=material_context,
    )
    try:
        payload = WorkbenchProductionRunner(
            doc_path=str(source),
            template=template,
            scene=scene,
            material_context=material_context,
            output_dir=Path(execution_session.output_namespace),
            execution_session=execution_session,
        ).run(lambda *_args: None, lambda: state["assembled"])
    finally:
        cleanup_execution_session_resources(execution_session)

    assert payload["status"] == "cancelled"
    assert set(payload["output_paths"]) == {"student", "answer_key"}
    assert all(Path(path).exists() for path in payload["output_paths"].values())
    assert payload["report_paths"]
    assert all(Path(path).exists() for path in payload["report_paths"])
    assert payload["exam_delivery_runtime"]["status"] in {"ok", "warning"}
    assert payload["exam_markdown_import"]["summary"]["question_count"] == 1


def test_generic_batch_last_item_cancelled_is_not_aggregated_as_failed(
    tmp_path,
    monkeypatch,
):
    item = SimpleNamespace(
        profile_id="only",
        profile_name="Only profile",
        output_dir=str(tmp_path / "out" / "only"),
        context=MaterialExecutionContext(),
    )

    class _CancelledProductionRunner:
        def __init__(self, **_kwargs):
            pass

        def run(self, _progress_cb, _cancel_check):
            return {
                "status": "cancelled",
                "output_path": str(tmp_path / "already-published.docx"),
                "output_paths": {"final": str(tmp_path / "already-published.docx")},
            }

    monkeypatch.setattr(
        execution_runtime,
        "build_material_batch_items",
        lambda *_args, **_kwargs: [item],
    )
    monkeypatch.setattr(
        execution_runtime,
        "WorkbenchProductionRunner",
        _CancelledProductionRunner,
    )
    monkeypatch.setattr(execution_runtime, "attach_batch_reports", lambda *_args: None)

    payload = WorkbenchBatchProductionRunner(
        doc_path=str(tmp_path / "source.docx"),
        template=TemplateConfig(),
        scene=SceneWorkspace(),
        archive=EntityArchive(),
        base_output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "cancelled"
    assert payload["items"][0]["status"] == "cancelled"
    assert payload["pending_profile_ids"] == []
    assert payload["error_text"] == ""


def test_cancelled_official_batch_projects_pending_ids_into_report_and_history(tmp_path):
    payload = batch_reporting.build_batch_payload("cancelled", [], error_text="")
    payload.update(
        {
            "batch_source_kind": "official_document_table",
            "batch_source_path": str(tmp_path / "source.csv"),
            "pending_profile_ids": ["pending-a", "pending-b"],
        }
    )

    batch_reporting.attach_batch_reports(
        payload,
        tmp_path / "reports",
        tmp_path / "source.csv",
    )

    report_path = tmp_path / "reports" / "source_batch_report.json"
    report_payload = json.loads(report_path.read_text(encoding="utf-8"))
    history = load_official_document_batch_history(payload["batch_history_path"])
    assert report_payload["pending_profile_ids"] == ["pending-a", "pending-b"]
    assert history is not None
    assert history.failed_profile_ids == ("pending-a", "pending-b")


@pytest.mark.parametrize(
    ("profile_ids", "expected_pending"),
    [
        (["only"], []),
        (["first", "pending"], ["pending"]),
    ],
)
def test_official_batch_post_assembly_cancellation_keeps_items_and_pending_ids(
    tmp_path,
    monkeypatch,
    profile_ids,
    expected_pending,
):
    items = [
        SimpleNamespace(
                profile_id=profile_id,
                profile_name=profile_id.title(),
                output_dir=str(tmp_path / "out" / profile_id),
            context=MaterialExecutionContext(
                entity_data={"document_type": "notice"},
            ),
        )
        for profile_id in profile_ids
    ]
    state = {"assembled": False}

    def _assemble(*_args, output_dir, **_kwargs):
        published = Path(output_dir) / "published.docx"
        published.parent.mkdir(parents=True, exist_ok=True)
        Document().save(published)
        state["assembled"] = True
        return OfficialDocumentAssemblyResult(
            status="ok",
            profile_id="notice",
            docx_path=published,
        )

    monkeypatch.setattr(execution_runtime, "assemble_official_document_docx", _assemble)
    monkeypatch.setattr(
        execution_preflight,
        "material_requirement_diagnostics",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        execution_preflight,
        "image_anchor_diagnostics",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        execution_preflight,
        "question_figure_file_diagnostics",
        lambda *_args, **_kwargs: [],
    )
    monkeypatch.setattr(
        execution_preflight,
        "missing_asset_rule_diagnostics",
        lambda *_args, **_kwargs: [],
    )
    runner = WorkbenchBatchProductionRunner(
        doc_path=str(tmp_path / "source.docx"),
        template=TemplateConfig(),
        scene=SceneWorkspace(scene_id="official", mode_id="official"),
        archive=EntityArchive(),
        base_output_dir=tmp_path / "out",
    )
    monkeypatch.setattr(runner, "_attach_official_batch_reports", lambda *_args: None)

    payload = runner._run_official_document_batch(
        items,
        input_path=tmp_path / "source.docx",
        base_output_dir=tmp_path / "out",
        progress_cb=lambda *_args: None,
        cancel_check=lambda: state["assembled"],
    )

    assert payload["status"] == "cancelled"
    assert [item["profile_id"] for item in payload["items"]] == [profile_ids[0]]
    assert payload["items"][0]["status"] == "success"
    assert Path(payload["items"][0]["output_path"]).exists()
    assert payload["pending_profile_ids"] == expected_pending
