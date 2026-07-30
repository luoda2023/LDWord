from __future__ import annotations

from pathlib import Path
import shutil
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from PIL import Image
import pytest

import src.services.material_attachments.intake as intake_module
from src.services.material_attachments import (
    AttachmentIntakeError,
    build_attachment_binding,
    build_directory_attachment_binding,
    build_single_attachment_binding,
    inspect_attachment_file,
)
from src.config.attachment_materials import (
    AttachmentProcessingMode,
    AttachmentSourceKind,
    normalize_attachment_relative_path,
)


def test_image_file_may_be_an_attachment_without_becoming_an_inline_image(tmp_path: Path):
    source = tmp_path / "evidence.png"
    Image.new("RGB", (12, 12), "white").save(source)

    binding = build_single_attachment_binding(
        role="supporting_proof",
        source_path=source,
        accepted_types=("image", "pdf"),
    )

    assert binding.role == "supporting_proof"
    assert binding.items[0].file_ref.media_type == "image/png"
    assert binding.accepted_media_types == ("application/pdf", "image/*")


def test_docx_attachment_is_validated_as_openxml_word_package(tmp_path: Path):
    source = tmp_path / "appendix.docx"
    document = Document()
    document.add_paragraph("附件")
    document.save(source)

    file_ref = inspect_attachment_file(source, accepted_types=("docx",))

    assert file_ref.original_name == "appendix.docx"
    assert "wordprocessingml.document" in file_ref.media_type


def test_fake_docx_bytes_are_rejected(tmp_path: Path):
    source = tmp_path / "pollution.docx"
    source.write_bytes(b"not a docx")

    with pytest.raises(AttachmentIntakeError, match="openxml_package_invalid"):
        inspect_attachment_file(source, accepted_types=("docx",))


def test_attachment_intake_rejects_source_size_limit_before_full_read(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "large.pdf"
    source.write_bytes(b"%PDF-1.4\n" + b"x" * 64)
    monkeypatch.setattr(intake_module, "_MAX_ATTACHMENT_SOURCE_BYTES", 16)

    with pytest.raises(AttachmentIntakeError, match="source_too_large"):
        inspect_attachment_file(source, accepted_types=("pdf",))


def test_attachment_intake_rejects_openxml_compression_bomb(tmp_path: Path) -> None:
    source = tmp_path / "bomb.docx"
    with ZipFile(source, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("word/document.xml", b"A" * 1_000_000)

    with pytest.raises(AttachmentIntakeError, match="openxml_package_invalid"):
        inspect_attachment_file(source, accepted_types=("docx",))


def test_disallowed_extension_is_rejected_before_binding(tmp_path: Path):
    source = tmp_path / "sheet.xlsx"
    source.write_bytes(b"not relevant")

    with pytest.raises(AttachmentIntakeError, match="extension_not_allowed"):
        inspect_attachment_file(source, accepted_types=("pdf",))


def test_multiple_attachment_binding_is_ordered_and_bounded(tmp_path: Path):
    first = tmp_path / "a.pdf"
    second = tmp_path / "b.pdf"
    first.write_bytes(b"%PDF-1.4\nfirst")
    second.write_bytes(b"%PDF-1.4\nsecond")

    binding = build_attachment_binding(
        role="evidence",
        source_paths=(second, first),
        accepted_types=("pdf",),
        cardinality="multiple",
        min_items=1,
        max_items=3,
    )

    assert [item.file_ref.original_name for item in binding.items] == ["a.pdf", "b.pdf"]
    assert binding.max_items == 3


def test_directory_package_preserves_tree_and_natural_relative_order(tmp_path: Path):
    root = tmp_path / "qualification-package"
    (root / "section10").mkdir(parents=True)
    (root / "section2").mkdir()
    (root / "cover.pdf").write_bytes(b"%PDF-1.4\ncover")
    (root / "section10" / "proof.pdf").write_bytes(b"%PDF-1.4\nten")
    (root / "section2" / "proof.pdf").write_bytes(b"%PDF-1.4\ntwo")

    binding = build_directory_attachment_binding(
        role="qualification_package",
        source_directory=root,
        accepted_types=("pdf",),
        recursive=True,
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
    )

    assert binding.source_kind is AttachmentSourceKind.DIRECTORY_PACKAGE
    assert binding.processing_mode is AttachmentProcessingMode.SUBSTITUTE_COPY
    assert binding.source_path == str(root.resolve())
    assert [item.relative_path for item in binding.items] == [
        "cover.pdf",
        "section2/proof.pdf",
        "section10/proof.pdf",
    ]
    assert [item.sequence for item in binding.items] == [0, 1, 2]


def test_directory_package_ignores_office_lock_files(tmp_path: Path) -> None:
    package = tmp_path / "package"
    package.mkdir()
    document = Document()
    document.add_paragraph("{{@text:project}}")
    document.save(package / "template.docx")
    (package / "~$template.docx").write_bytes(b"not-an-openxml-package")
    (package / ".~lock.template.docx#").write_bytes(b"lock")

    binding = build_directory_attachment_binding(
        role="templates",
        source_directory=package,
        accepted_types=("docx",),
        recursive=True,
        max_items=None,
    )

    assert [item.relative_path for item in binding.items] == ["template.docx"]


def test_nonrecursive_directory_package_only_collects_root_files(tmp_path: Path):
    root = tmp_path / "package"
    (root / "nested").mkdir(parents=True)
    (root / "root.pdf").write_bytes(b"%PDF-1.4\nroot")
    (root / "nested" / "nested.pdf").write_bytes(b"%PDF-1.4\nnested")

    binding = build_directory_attachment_binding(
        role="package",
        source_directory=root,
        accepted_types=("pdf",),
        recursive=False,
    )

    assert [item.relative_path for item in binding.items] == ["root.pdf"]


def test_item_identity_is_path_stable_while_revision_tracks_content(tmp_path: Path):
    root = tmp_path / "package"
    root.mkdir()
    source = root / "evidence.pdf"
    source.write_bytes(b"%PDF-1.4\nfirst")
    first = build_directory_attachment_binding(
        role="evidence",
        source_directory=root,
        accepted_types=("pdf",),
    )
    source.write_bytes(b"%PDF-1.4\nchanged")
    changed = build_directory_attachment_binding(
        role="evidence",
        source_directory=root,
        accepted_types=("pdf",),
    )

    assert changed.items[0].item_id == first.items[0].item_id
    assert changed.binding_revision != first.binding_revision


def test_binding_revision_is_portable_when_directory_moves(tmp_path: Path):
    first_root = tmp_path / "first" / "package"
    first_root.mkdir(parents=True)
    (first_root / "evidence.pdf").write_bytes(b"%PDF-1.4\nsame")
    second_root = tmp_path / "second" / "package"
    shutil.copytree(first_root, second_root)

    first = build_directory_attachment_binding(
        role="evidence",
        source_directory=first_root,
        accepted_types=("pdf",),
    )
    second = build_directory_attachment_binding(
        role="evidence",
        source_directory=second_root,
        accepted_types=("pdf",),
    )

    assert first.source_path != second.source_path
    assert first.binding_revision == second.binding_revision


def test_file_set_rejects_casefold_output_collision(tmp_path: Path):
    first_dir = tmp_path / "first"
    second_dir = tmp_path / "second"
    first_dir.mkdir()
    second_dir.mkdir()
    first = first_dir / "Proof.pdf"
    second = second_dir / "proof.PDF"
    first.write_bytes(b"%PDF-1.4\nfirst")
    second.write_bytes(b"%PDF-1.4\nsecond")

    with pytest.raises(AttachmentIntakeError, match="relative_path.*unique"):
        build_attachment_binding(
            role="evidence",
            source_paths=(first, second),
            accepted_types=("pdf",),
            source_kind=AttachmentSourceKind.FILE_SET,
        )


@pytest.mark.parametrize(
    "relative_path",
    ["../escape.pdf", "a/../escape.pdf", "/absolute.pdf", "C:/drive.pdf", "a//b.pdf"],
)
def test_relative_path_rejects_escape_and_noncanonical_forms(relative_path: str):
    with pytest.raises(ValueError, match="attachment_relative_path"):
        normalize_attachment_relative_path(relative_path)


def test_directory_package_rejects_symbolic_links(tmp_path: Path):
    root = tmp_path / "package"
    root.mkdir()
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(b"%PDF-1.4\noutside")
    link = root / "linked.pdf"
    try:
        link.symlink_to(outside)
    except OSError:
        pytest.skip("symbolic links are not available for this Windows account")

    with pytest.raises(AttachmentIntakeError, match="link_forbidden"):
        build_directory_attachment_binding(
            role="evidence",
            source_directory=root,
            accepted_types=("pdf",),
        )
