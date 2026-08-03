from types import SimpleNamespace

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.shared import Cm

from src.config.resolved import ResolvedConfig
from src.modules.basic.page_setup import PageSetupModule
from src.modules.basic.section_format import SectionFormatModule
from src.modules.structure.heading_recognition import DocSection
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn
from src.shared.engine.section_layout_planner import (
    build_section_execution_plan,
    collect_section_inventory,
)


def _tree(*sections):
    return SimpleNamespace(
        sections=list(sections),
        get_section=lambda name: next(
            (section for section in sections if section.section_type == name),
            None,
        ),
        get_section_for_paragraph=lambda index: next(
            (
                section.section_type
                for section in sections
                if section.start_index <= index < section.end_index
            ),
            "body",
        ),
    )


def test_preserve_source_keeps_existing_empty_section_boundaries():
    doc = Document()
    doc.add_paragraph("content")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("tail")
    before = collect_section_inventory(doc)

    config = ResolvedConfig()
    config.section.boundary_mode = "preserve_source"
    context = PipelineContext()
    SectionFormatModule().apply(doc, config, ChangeTracker(), context)

    after = collect_section_inventory(doc)
    assert before.section_count == 3
    assert after.section_count == before.section_count
    assert context.section_execution_receipt.applied_operations == ()


def test_explicit_cleanup_removes_only_a_proven_empty_redundant_section():
    doc = Document()
    doc.add_paragraph("content")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("tail")

    config = ResolvedConfig()
    config.section.boundary_mode = "preserve_source"
    config.section.empty_break_policy = "remove_proven_redundant"
    context = PipelineContext()
    SectionFormatModule().apply(doc, config, ChangeTracker(), context)

    receipt = context.section_execution_receipt
    assert receipt.source_section_count == 3
    assert receipt.final_section_count == 2
    assert receipt.removed_count == 1


def test_caption_owned_policy_removes_only_a_proven_caption_table_break():
    doc = Document()
    caption = doc.add_paragraph("表 1  测试数据")
    caption.style = doc.styles["Caption"]
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_table(rows=1, cols=1)

    config = ResolvedConfig()
    config.section.boundary_mode = "preserve_source"
    config.caption.table_break_policy = "remove_proven_redundant"
    context = PipelineContext()

    SectionFormatModule().apply(doc, config, ChangeTracker(), context)

    receipt = context.section_execution_receipt
    assert receipt.source_section_count == 2
    assert receipt.final_section_count == 1
    assert receipt.removed_count == 1


def test_cleanup_preserves_empty_boundary_that_contains_a_field():
    doc = Document()
    doc.add_paragraph("content")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("tail")
    boundary_paragraph = doc.paragraphs[2]
    run = OxmlElement("w:r")
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    run.append(field)
    boundary_paragraph._element.append(run)

    config = ResolvedConfig()
    config.section.boundary_mode = "preserve_source"
    config.section.empty_break_policy = "remove_proven_redundant"
    plan = build_section_execution_plan(doc, config, PipelineContext())

    assert not any(operation.action == "remove_break" for operation in plan.operations)


def test_semantic_insertion_clones_the_local_not_final_section_geometry():
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_paragraph("Body starts")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Appendix")
    doc.sections[0].page_width = Cm(21)
    doc.sections[0].page_height = Cm(29.7)
    doc.sections[1].page_width = Cm(42)
    doc.sections[1].page_height = Cm(29.7)

    config = ResolvedConfig()
    config.section.boundary_mode = "semantic_rebuild"
    context = PipelineContext(
        doc_tree=_tree(
            DocSection("cover", 0, 1, confidence=10),
            DocSection("body", 1, 3, confidence=10),
            DocSection("appendix", 3, 4, confidence=10),
        )
    )
    SectionFormatModule().apply(doc, config, ChangeTracker(), context)

    ppr = doc.paragraphs[0]._element.find(qn("w:pPr"))
    inserted = ppr.find(qn("w:sectPr"))
    pg_size = inserted.find(qn("w:pgSz"))
    assert int(pg_size.get(qn("w:w"))) < int(pg_size.get(qn("w:h")))
    assert len(doc.sections) == 3


def test_semantic_section_execution_is_idempotent():
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_paragraph("Body")
    context = PipelineContext(
        doc_tree=_tree(
            DocSection("cover", 0, 1, confidence=10),
            DocSection("body", 1, 2, confidence=10),
        )
    )
    config = ResolvedConfig()
    config.section.boundary_mode = "semantic_rebuild"
    module = SectionFormatModule()

    module.apply(doc, config, ChangeTracker(), context)
    first = collect_section_inventory(doc)
    module.apply(doc, config, ChangeTracker(), context)
    second = collect_section_inventory(doc)

    assert first.section_count == 2
    assert second.digest == first.digest
    assert context.section_execution_receipt.applied_operations == ()


def test_page_setup_preserves_source_orientation_while_forcing_paper_size():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(42)
    section.page_height = Cm(29.7)
    config = ResolvedConfig()
    config.page_setup.paper_size = "A4"
    config.page_setup.paper_size_mode = "force_template"
    config.page_setup.orientation_mode = "preserve_source"

    PageSetupModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert section.page_width > section.page_height
    assert abs(section.page_width.cm - 29.7) < 0.02
    assert abs(section.page_height.cm - 21.0) < 0.02


def test_page_setup_can_force_orientation_without_replacing_source_paper():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(42)
    section.page_height = Cm(29.7)
    config = ResolvedConfig()
    config.page_setup.paper_size_mode = "preserve_source"
    config.page_setup.orientation_mode = "force_template"
    config.page_setup.orientation = "portrait"

    PageSetupModule().apply(doc, config, ChangeTracker(), PipelineContext())

    assert section.page_width < section.page_height
    assert abs(section.page_width.cm - 29.7) < 0.02
    assert abs(section.page_height.cm - 42.0) < 0.02
