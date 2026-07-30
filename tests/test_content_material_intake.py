from __future__ import annotations

from pathlib import Path

from docx import Document
from PIL import Image

from src.services.material_content.artifact_repository import ContentArtifactRepository
from src.services.material_content.compiler import compile_content_material


def test_markdown_intake_discovers_exact_local_image_manifest(tmp_path: Path):
    image_path = tmp_path / "images" / "route.png"
    image_path.parent.mkdir()
    Image.new("RGB", (20, 10), "navy").save(image_path)
    source = tmp_path / "route.md"
    source.write_text("# 技术路线\n\n![路线图](images/route.png)\n", encoding="utf-8")

    repository = ContentArtifactRepository(tmp_path / "content-artifacts")
    result = compile_content_material(source, repository)

    assert not result.blocked, result.findings
    assert result.artifact_ref is not None
    artifact = repository.validate(result.artifact_ref)
    assert artifact.manifest.source.format == "markdown"
    assert artifact.manifest.source.original_name == "route.md"
    assert len(artifact.manifest.resources) == 1
    assert artifact.manifest.resources[0].media_type == "image/png"
    assert result.inventory.image_count == 1


def test_docx_intake_validates_semantics_and_keeps_resources_embedded(tmp_path: Path):
    source = tmp_path / "background.docx"
    document = Document()
    document.add_heading("项目背景", level=1)
    document.add_paragraph("背景正文")
    document.save(source)

    repository = ContentArtifactRepository(tmp_path / "content-artifacts")
    result = compile_content_material(source, repository)

    assert not result.blocked, result.findings
    assert result.artifact_ref is not None
    artifact = repository.validate(result.artifact_ref)
    assert artifact.manifest.source.format == "docx"
    assert artifact.manifest.resources == ()
    assert result.inventory.heading_count == 1


def test_intake_rejects_non_content_extension(tmp_path: Path):
    source = tmp_path / "attachment.pdf"
    source.write_bytes(b"%PDF-1.4")

    result = compile_content_material(
        source,
        ContentArtifactRepository(tmp_path / "content-artifacts"),
    )

    assert result.blocked
    assert result.artifact_ref is None
    assert {item.code for item in result.findings} == {"source_extension_not_allowed"}


def test_markdown_intake_rejects_resource_escape(tmp_path: Path):
    outside = tmp_path / "outside.png"
    Image.new("RGB", (10, 10), "white").save(outside)
    source_dir = tmp_path / "content"
    source_dir.mkdir()
    source = source_dir / "bad.md"
    source.write_text("![bad](../outside.png)", encoding="utf-8")

    result = compile_content_material(
        source,
        ContentArtifactRepository(tmp_path / "content-artifacts"),
    )

    assert result.blocked
    assert result.artifact_ref is None
    assert {item.code for item in result.findings} == {"image_path_escape"}
