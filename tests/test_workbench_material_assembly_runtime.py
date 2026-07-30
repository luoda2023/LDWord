from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
import json
from pathlib import Path
from shutil import copy2
from types import SimpleNamespace

from docx import Document
from PIL import Image
import pytest

from src.config.attachment_materials import AttachmentProcessingMode
from src.config.content_artifacts import ContentMaterialBinding
from src.config.content_materials import ContentInsertionRule, content_anchor_token
from src.config.library import CONFIG_LIBRARY_ROOT
from src.config.image_materials import (
    ImageMaterialRule,
    ImagePlacementMode,
    ImagePlacementPolicy,
)
from src.config.material_context import MaterialExecutionContext
from src.config.materials import AssetItem
from src.config.resolver import resolve_config
from src.config.scene import ContentVisibilityRule, DeliveryPreset, SceneWorkspace
from src.config.template import TemplateConfig
from src.pipeline.context import PipelineContext
from src.pipeline.result import PipelineResult
import src.pipeline.runner as pipeline_runner
import src.services.material_content.composer as content_composer
import src.services.material_execution.dependency_indexer as dependency_indexer
from src.services.material_content.artifact_repository import ContentArtifactRepository
from src.services.material_content.compiler import compile_content_material
from src.services.material_attachments import build_single_attachment_binding
from src.services.material_attachments import (
    AttachmentBundleProcessingError,
    AttachmentProcessingDiagnostic,
)
from src.services.material_execution import (
    FileEvidence,
    MaterialAssemblyDependencies,
    MaterialAssemblyService as RealMaterialAssemblyService,
    MaterialPipelineRequest,
    PipelineVariantTarget,
    VisibilityEvidenceKind,
)
from src.shared.engine.office_image_layout import OfficeImageProvider
from tests.material_image_layout_fake import FakeSuccessfulLayout
from src.services.material_delivery import DeliveryPackageBuildError
from src.shared.engine.block_visibility import apply_block_visibility
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner
from src.services.production_runtime import material_assembly_runtime as assembly_runtime
from src.services.production_runtime import material_artifacts


def _repository_snapshot(root: Path) -> tuple[tuple[str, str], ...]:
    if not root.exists():
        return ()
    return tuple(
        (path.relative_to(root).as_posix(), sha256(path.read_bytes()).hexdigest())
        for path in sorted(root.rglob("*"))
        if path.is_file()
    )


@pytest.fixture(autouse=True)
def _application_content_repository_stays_unchanged():
    root = CONFIG_LIBRARY_ROOT / "content_artifacts"
    before = _repository_snapshot(root)
    yield
    assert _repository_snapshot(root) == before


def _isolated_content_repository(
    tmp_path: Path,
    monkeypatch,
) -> ContentArtifactRepository:
    config_root = tmp_path / "config_library"
    monkeypatch.setattr(assembly_runtime, "CONFIG_LIBRARY_ROOT", config_root)
    monkeypatch.setattr(content_composer, "CONFIG_LIBRARY_ROOT", config_root)
    monkeypatch.setattr(dependency_indexer, "CONFIG_LIBRARY_ROOT", config_root)
    return ContentArtifactRepository(config_root / "content_artifacts")


def _content_context(
    tmp_path: Path,
    monkeypatch,
    *,
    anchor_present: bool = True,
):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph(
        content_anchor_token("technical_route") if anchor_present else "No anchor"
    )
    document.save(source)
    markdown = tmp_path / "route.md"
    markdown.write_text("Reusable technical route.\n", encoding="utf-8")
    result = compile_content_material(
        markdown,
        _isolated_content_repository(tmp_path, monkeypatch),
    )
    assert not result.blocked, result.findings
    assert result.artifact_ref is not None
    binding = ContentMaterialBinding(
        content_id="technical_route",
        label="Technical route",
        artifact_ref=result.artifact_ref,
    )
    rule = ContentInsertionRule(
        rule_id="technical-route-rule",
        content_id="technical_route",
        anchor_token=content_anchor_token("technical_route"),
    )
    context = MaterialExecutionContext(
        profile_id="p1",
        profile_name="P1",
        content_bindings={binding.content_id: binding},
        content_rules=[rule],
    )
    return source, context


def _minimal_scene() -> SceneWorkspace:
    scene = SceneWorkspace()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.compliance_profile.object_preflight.enabled = False
    scene.default_delivery_preset().artifacts.report_json = False
    scene.default_delivery_preset().artifacts.report_markdown = False
    return scene


def _optional_image_rule(*, required: bool = False) -> ImageMaterialRule:
    return ImageMaterialRule(
        rule_id="image:logo",
        source_role="logo",
        anchor_token="{{@img:LOGO1}}",
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIXED_BOX,
            fixed_width_cm=6.0,
        ),
        required=required,
    )


def test_active_material_route_rejects_source_final_before_service_ownership(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source, context = _content_context(tmp_path, monkeypatch)
    original = source.read_bytes()
    alias_parent = tmp_path / "alias-parent"
    alias_parent.mkdir()
    scene = _minimal_scene()
    scene.delivery_presets = [
        DeliveryPreset(
            preset_id="final",
            output_dir_template="{document_dir}/alias-parent/..",
            filename_template="{stem}",
        )
    ]
    service_calls: list[str] = []

    class _UnexpectedMaterialAssemblyService:
        def __init__(self, *_args, **_kwargs) -> None:
            service_calls.append("constructed")
            raise AssertionError("service must not acquire transaction ownership")

    monkeypatch.setattr(
        assembly_runtime,
        "MaterialAssemblyService",
        _UnexpectedMaterialAssemblyService,
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=context,
        output_dir=tmp_path,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert "source overwrite is forbidden" in payload["error_text"]
    assert service_calls == []
    assert source.read_bytes() == original


def test_dormant_optional_image_rules_do_not_activate_assembly(tmp_path: Path):
    rule = _optional_image_rule()
    dormant = MaterialExecutionContext(image_material_rules={rule.rule_id: rule})
    selected = MaterialExecutionContext(
        image_material_rules={rule.rule_id: rule},
        asset_items=[AssetItem(role="logo", source_path=str(tmp_path / "logo.png"))],
    )
    required_rule = _optional_image_rule(required=True)
    required = MaterialExecutionContext(
        image_material_rules={required_rule.rule_id: required_rule}
    )

    assert assembly_runtime.material_assembly_is_active(dormant) is False
    assert assembly_runtime.material_assembly_is_active(selected) is True
    assert assembly_runtime.material_assembly_is_active(required) is True


def test_workbench_executes_canonicalized_image_role_without_silent_token_loss(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("{{@img:LOGO1}}")
    document.save(source)
    image_path = tmp_path / "logo.png"
    Image.new("RGB", (80, 60), "red").save(image_path)
    rule = ImageMaterialRule(
        rule_id="image:logo",
        source_role="品牌标志",
        anchor_token="{{@img:LOGO1}}",
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.NATURAL_SIZE,
        ),
    )
    context = MaterialExecutionContext(
        asset_items=[
            AssetItem(
                item_id="logo-1",
                role="LOGO",
                path=str(image_path),
            )
        ],
        image_material_rules={rule.rule_id: rule},
    )
    fake_layout = FakeSuccessfulLayout()
    dependencies = replace(
        MaterialAssemblyDependencies(),
        layout_runner=fake_layout,
    )
    monkeypatch.setattr(
        assembly_runtime,
        "_qualified_office_provider",
        lambda _path: OfficeImageProvider.WORD,
    )
    monkeypatch.setattr(
        assembly_runtime,
        "_material_cache_root",
        lambda: tmp_path / "material-cache",
    )
    monkeypatch.setattr(
        assembly_runtime,
        "MaterialAssemblyService",
        lambda callback: RealMaterialAssemblyService(
            callback,
            dependencies=dependencies,
        ),
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_minimal_scene(),
        material_context=context,
        output_dir=tmp_path / "output",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success", payload["error_text"]
    output = Document(payload["output_path"])
    assert "{{@img:LOGO1}}" not in "\n".join(
        paragraph.text for paragraph in output.paragraphs
    )
    assert len(output.inline_shapes) == 1
    assert len(fake_layout.requests) == 1


def test_attachment_binding_activates_assembly_and_projects_bundle_receipt(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Main document")
    document.save(source)
    attachment = tmp_path / "evidence.pdf"
    attachment.write_bytes(b"%PDF-1.7\nattachment\n")
    binding = build_single_attachment_binding(
        role="evidence",
        source_path=attachment,
        accepted_types=("pdf",),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
    )
    document = Document(source)
    document.add_paragraph(binding.anchor_token)
    document.save(source)
    context = MaterialExecutionContext(
        attachment_bindings={binding.role: binding}
    )
    scene = _minimal_scene()
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.default_delivery_preset().artifacts.material_package = True
    output_dir = tmp_path / "out"

    assert assembly_runtime.material_assembly_is_active(context) is True
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=context,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    bundles = payload["attachment_bundles"]
    receipt = bundles["receipts"]["evidence"]
    processed = Path(receipt["output_directory"]) / "evidence.pdf"
    manifest = json.loads(
        Path(payload["material_manifest_paths"]["material"]).read_text(
            encoding="utf-8"
        )
    )
    package_manifest = json.loads(
        Path(payload["material_package_paths"]["package_manifest"]).read_text(
            encoding="utf-8"
        )
    )
    package_paths = {
        (item["category"], item["path"])
        for item in package_manifest["files"]
    }
    final_text = "\n".join(
        paragraph.text
        for paragraph in Document(payload["output_path"]).paragraphs
    )

    assert payload["status"] == "success"
    assert bundles["status"] == "applied"
    assert bundles["file_count"] == 1
    assert receipt["files"][0]["status"] == "not_applicable"
    assert processed.read_bytes() == attachment.read_bytes()
    assert "附件：evidence.pdf" in final_text
    assert binding.anchor_token not in final_text
    assert (
        payload["material_assembly_receipt"]["content_compose_receipt"]
        ["attachment_render_receipts"][0]["rendered_item_count"]
        == 1
    )
    assert manifest["material_domains"]["attachment"]["execution"] == bundles
    assert manifest["delivery"]["attachment_bundle_paths"]["evidence"] == str(
        processed.parent
    )
    assert ("attachment", "attachments/evidence/evidence.pdf") in package_paths
    assert (
        "attachment_receipt",
        "attachments/evidence/attachment_bundle_receipt.json",
    ) in package_paths


def test_attachment_bundle_failure_keeps_main_output_and_marks_partial_package(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "source.docx"
    Document().save(source)
    attachment = tmp_path / "evidence.pdf"
    attachment.write_bytes(b"%PDF-1.7\nattachment\n")
    binding = build_single_attachment_binding(
        role="evidence",
        source_path=attachment,
        accepted_types=("pdf",),
    )
    context = MaterialExecutionContext(
        attachment_bindings={binding.role: binding}
    )
    scene = _minimal_scene()
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.default_delivery_preset().artifacts.material_package = True

    class _FailingAttachmentService:
        def process(self, *_args, **_kwargs):
            raise AttachmentBundleProcessingError(
                (
                    AttachmentProcessingDiagnostic(
                        code="simulated_attachment_failure",
                        message="simulated attachment failure",
                    ),
                )
            )

    monkeypatch.setattr(
        assembly_runtime,
        "AttachmentBundleService",
        _FailingAttachmentService,
    )
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=context,
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: False)

    package_manifest = json.loads(
        Path(payload["material_package_paths"]["package_manifest"]).read_text(
            encoding="utf-8"
        )
    )
    assert payload["status"] == "partial_success"
    assert Path(payload["output_path"]).is_file()
    assert payload["attachment_bundles"]["status"] == "failed"
    assert payload["attachment_bundles"]["failed_binding_count"] == 1
    assert "evidence" in payload["attachment_bundles"]["errors"]
    assert package_manifest["status"] == "partial_success"
    assert not any(
        item["category"] == "attachment"
        for item in package_manifest["files"]
    )


def test_content_only_uses_transactional_assembly_without_office_probe(
    tmp_path: Path,
    monkeypatch,
):
    source, context = _content_context(tmp_path, monkeypatch)
    dormant_rule = _optional_image_rule()
    context.image_material_rules = {dormant_rule.rule_id: dormant_rule}
    output_dir = tmp_path / "out"
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "local"))

    assert assembly_runtime._has_image_layout_risk(context) is False

    def unexpected_probe(*_args, **_kwargs):
        raise AssertionError("Markdown content without images must not probe Office")

    def unexpected_inner_transaction(**_kwargs):
        raise AssertionError(
            "assembler-owned stages must not open a nested transaction"
        )

    monkeypatch.setattr(
        assembly_runtime,
        "resolve_office_layout_capability",
        unexpected_probe,
    )
    monkeypatch.setattr(
        pipeline_runner,
        "OwnedAssemblyTransaction",
        unexpected_inner_transaction,
    )
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_minimal_scene(),
        material_context=context,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    final = Path(payload["output_path"])
    text = "\n".join(paragraph.text for paragraph in Document(final).paragraphs)
    assert payload["status"] == "success"
    assert "Reusable technical route." in text
    assert content_anchor_token("technical_route") not in text
    assert final.parent == output_dir
    assert not (output_dir / "material-execution").exists()
    assert payload["material_assembly"]["status"] == "applied"
    assert payload["material_assembly_receipt"] is not None
    assert "material_assembly_error" not in payload
    assert (
        payload["material_assembly"]["receipt"]
        == payload["material_assembly_receipt"]
    )


def test_published_assembly_is_retained_when_report_writer_fails(
    tmp_path: Path,
    monkeypatch,
):
    source, context = _content_context(tmp_path, monkeypatch)
    output_dir = tmp_path / "out"

    def fail_reports(*_args, **_kwargs):
        raise OSError(5, "simulated report write failure", str(output_dir / "report.json"))

    monkeypatch.setattr(
        "src.services.production_runtime.delivery_reporting.write_enabled_result_reports",
        fail_reports,
    )
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_minimal_scene(),
        material_context=context,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "partial_success"
    assert Path(payload["output_path"]).is_file()
    assert payload["output_paths"]
    assert payload["material_assembly_receipt"] is not None
    assert "material_assembly_error" not in payload
    assert len(payload["artifact_failures"]) == 1
    failure = payload["artifact_failures"][0]
    assert failure["kind"] == "reports"
    assert failure["path"] == str(output_dir / "report.json")
    assert failure["error_type"] == "OSError"
    assert "simulated report write failure" in failure["error"]


def test_requested_material_package_failure_is_visible_without_hiding_core_output(
    tmp_path: Path,
    monkeypatch,
):
    source, context = _content_context(tmp_path, monkeypatch)
    scene = _minimal_scene()
    scene.default_delivery_preset().artifacts.material_package = True

    def fail_package(*_args, **_kwargs):
        raise DeliveryPackageBuildError("simulated package publication failure")

    monkeypatch.setattr(
        material_artifacts,
        "build_material_delivery_package",
        fail_package,
    )
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=context,
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "partial_success"
    assert Path(payload["output_path"]).is_file()
    assert payload["material_manifest_paths"]
    assert payload["material_package_paths"] == {}
    assert payload["artifact_failures"] == [
        {
            "kind": "material_package",
            "path": "",
            "error_type": "DeliveryPackageBuildError",
            "error": "simulated package publication failure",
        }
    ]


def test_failed_assembly_error_survives_failure_report_writer_error(
    tmp_path: Path,
    monkeypatch,
):
    source, context = _content_context(
        tmp_path,
        monkeypatch,
        anchor_present=False,
    )

    def fail_reports(*_args, **_kwargs):
        raise OSError("simulated failure-report error")

    monkeypatch.setattr(
        "src.services.production_runtime.delivery_reporting.write_enabled_result_reports",
        fail_reports,
    )
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_minimal_scene(),
        material_context=context,
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert "material_assembly_receipt" not in payload
    assert payload["material_assembly_error"] is not None
    assert payload["artifact_failures"][0]["kind"] == "reports"
    assert "simulated failure-report error" in payload["error_text"]


def test_assembly_failure_preserves_existing_final(tmp_path: Path, monkeypatch):
    source, context = _content_context(
        tmp_path,
        monkeypatch,
        anchor_present=False,
    )
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    old_final = output_dir / "source_formatted.docx"
    old_document = Document()
    old_document.add_paragraph("old final")
    old_document.save(old_final)
    old_hash = sha256(old_final.read_bytes()).hexdigest()
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "local"))
    scene = _minimal_scene()
    scene.default_delivery_preset().artifacts.report_json = True
    scene.default_delivery_preset().artifacts.report_markdown = True
    scene.default_delivery_preset().artifacts.material_manifest = True
    scene.delivery_presets[0].include_structured_intermediate = True

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        material_context=context,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert sha256(old_final.read_bytes()).hexdigest() == old_hash
    assert "old final" in "\n".join(
        paragraph.text for paragraph in Document(old_final).paragraphs
    )
    assert payload["material_assembly"]["status"] == "failed"
    assert "material_assembly_receipt" not in payload
    assert payload["material_assembly_error"] is not None
    assert (
        payload["material_assembly"]["error"]
        == payload["material_assembly_error"]
    )
    assert len(payload["report_paths"]) == 2
    json_report = next(
        Path(path) for path in payload["report_paths"] if path.endswith(".json")
    )
    markdown_report = next(
        Path(path) for path in payload["report_paths"] if path.endswith(".md")
    )
    report_payload = json.loads(json_report.read_text(encoding="utf-8"))
    assert report_payload["output"] == ""
    assert report_payload["material_assembly"]["status"] == "failed"
    assert (
        report_payload["material_assembly"]["error"]
        == payload["material_assembly_error"]
    )
    assert "content_compose_failed" in markdown_report.read_text(encoding="utf-8")
    assert set(payload["intermediate_paths"]) == {"final"}
    intermediate = json.loads(
        Path(payload["intermediate_paths"]["final"]).read_text(encoding="utf-8")
    )
    assert intermediate["output"] == ""
    assert intermediate["material_assembly"]["status"] == "failed"
    assert (
        intermediate["result"]["material_assembly_error"]
        == payload["material_assembly_error"]
    )
    manifest_path = Path(payload["material_manifest_paths"]["material"])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["execution"]["status"] == "failed"
    assert manifest["execution"]["applied"] is None
    assert manifest["execution"]["failure"] == payload["material_assembly_error"]


def test_non_active_context_keeps_legacy_runner_path(tmp_path: Path, monkeypatch):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("plain")
    document.save(source)

    def unexpected_new_chain(**_kwargs):
        raise AssertionError("non-active execution entered material assembly")

    monkeypatch.setattr(
        "src.services.production_runtime.execution_runtime.run_workbench_material_assembly",
        unexpected_new_chain,
    )
    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_minimal_scene(),
        material_context=MaterialExecutionContext(entity_data={"name": "A"}),
        output_dir=tmp_path / "out",
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert Path(payload["output_path"]).is_file()


def test_pipeline_adapter_excludes_legacy_image_module_and_returns_visibility(
    tmp_path: Path,
    monkeypatch,
):
    source = tmp_path / "logical.docx"
    prepared = tmp_path / "prepared.docx"
    final = tmp_path / "final.docx"
    stage = tmp_path / ".final.material-stage.docx"
    Document().save(source)
    document = Document()
    document.add_paragraph("staged")
    document.save(prepared)
    digest = sha256(prepared.read_bytes()).hexdigest()
    captured: dict[str, object] = {}
    visibility_receipt = apply_block_visibility(
        Document(),
        remove_selectors=(),
    )
    receipt_id = visibility_receipt.receipt_id

    class FakePipeline:
        def __init__(self, *, modules, **kwargs):
            captured["module_names"] = [module.meta.name for module in modules]
            captured["kwargs"] = kwargs

        def execute(self, input_path):
            kwargs = captured["kwargs"]
            target = Path(kwargs["output_path_overrides"]["final"])
            copy2(input_path, target)
            context = PipelineContext(
                content_visibility_receipts={
                    "final": visibility_receipt.to_dict()
                }
            )
            return PipelineResult(
                success=True,
                output_paths={"final": str(target)},
                context=context,
            )

    monkeypatch.setattr(assembly_runtime, "Pipeline", FakePipeline)
    scene = _minimal_scene()
    scene.module_switches["image_insertion"] = True
    config = resolve_config(TemplateConfig(), scene)
    adapter = assembly_runtime._WorkbenchMaterialPipelineAdapter(
        config=config,
        output_dir=tmp_path,
        output_suffix="_formatted",
        progress_cb=lambda *_args: None,
        cancel_check=lambda: False,
        force_delivery_presets=False,
        official_master=None,
        image_rule_tokens=("{{@img:qualification1}}",),
    )
    request = MaterialPipelineRequest(
        execution_id="e1",
        logical_source_path=str(source),
        prepared_input=FileEvidence(str(prepared), digest, prepared.stat().st_size),
        intake_snapshot_id=sha256(b"snapshot").hexdigest(),
        targets=(
            PipelineVariantTarget(
                variant_id="final",
                variant_version="1",
                final_output_path=str(final),
                owned_stage_output_path=str(stage),
            ),
        ),
    )

    outcome = adapter(request, cancel_check=lambda: False)

    kwargs = captured["kwargs"]
    assert "image_insertion" not in captured["module_names"]
    assert kwargs["logical_source_path"] == str(source)
    assert kwargs["output_path_overrides"] == {"final": str(stage)}
    assert kwargs["defer_field_refresh"] is True
    assert outcome.visibility_evidence[0].kind is VisibilityEvidenceKind.RECEIPT
    assert outcome.visibility_evidence[0].visibility_receipt_id == receipt_id


@pytest.mark.parametrize("tamper", ("canonical", "missing_identity", "nested_legacy"))
def test_visibility_handoff_rejects_noncanonical_receipts(
    tmp_path: Path,
    tamper: str,
):
    raw = apply_block_visibility(Document(), remove_selectors=()).to_dict()
    if tamper == "canonical":
        raw["body_sha256_after"] = "0" * 64
    elif tamper == "missing_identity":
        raw.pop("receipt_id")
    else:
        raw = {"receipt": raw}
    result = PipelineResult(
        success=True,
        context=PipelineContext(
            content_visibility_receipts={"final": raw}
        ),
    )
    target = PipelineVariantTarget(
        variant_id="final",
        variant_version="1",
        final_output_path=str(tmp_path / "final.docx"),
        owned_stage_output_path=str(tmp_path / ".final.stage.docx"),
    )

    with pytest.raises(RuntimeError, match="visibility"):
        assembly_runtime._visibility_evidence(
            result,
            (target,),
            config=SimpleNamespace(delivery_presets=[]),
        )


def test_pipeline_adapter_never_fabricates_required_visibility_receipt(
    tmp_path: Path,
    monkeypatch,
):
    source = tmp_path / "logical.docx"
    prepared = tmp_path / "prepared.docx"
    final = tmp_path / "final.docx"
    stage = tmp_path / ".final.material-stage.docx"
    Document().save(source)
    Document().save(prepared)

    class FakePipeline:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

        def execute(self, input_path):
            target = Path(self.kwargs["output_path_overrides"]["final"])
            copy2(input_path, target)
            return PipelineResult(
                success=True,
                output_paths={"final": str(target)},
                context=PipelineContext(),
            )

    monkeypatch.setattr(assembly_runtime, "Pipeline", FakePipeline)
    scene = _minimal_scene()
    scene.delivery_presets = [
        DeliveryPreset(
            preset_id="final",
            content_visibility_rules=[
                ContentVisibilityRule(
                    rule_id="hide-draft",
                    selector="draft-block",
                )
            ],
        )
    ]
    config = resolve_config(TemplateConfig(), scene)
    adapter = assembly_runtime._WorkbenchMaterialPipelineAdapter(
        config=config,
        output_dir=tmp_path,
        output_suffix="_formatted",
        progress_cb=lambda *_args: None,
        cancel_check=lambda: False,
        force_delivery_presets=False,
        official_master=None,
        image_rule_tokens=(),
    )
    digest = sha256(prepared.read_bytes()).hexdigest()
    request = MaterialPipelineRequest(
        execution_id="e-no-receipt",
        logical_source_path=str(source),
        prepared_input=FileEvidence(
            str(prepared),
            digest,
            prepared.stat().st_size,
        ),
        intake_snapshot_id=sha256(b"snapshot-no-receipt").hexdigest(),
        targets=(
            PipelineVariantTarget(
                variant_id="final",
                variant_version="1",
                final_output_path=str(final),
                owned_stage_output_path=str(stage),
            ),
        ),
    )

    with pytest.raises(RuntimeError, match="returned no visibility receipt"):
        adapter(request)


def test_compiled_docx_without_images_has_no_office_layout_risk(
    tmp_path: Path,
    monkeypatch,
):
    content = tmp_path / "content.docx"
    Document().save(content)
    result = compile_content_material(
        content,
        _isolated_content_repository(tmp_path, monkeypatch),
    )
    assert not result.blocked, result.findings
    assert result.artifact_ref is not None
    binding = ContentMaterialBinding(
        content_id="content",
        label="Content",
        artifact_ref=result.artifact_ref,
    )
    context = MaterialExecutionContext(
        content_bindings={"content": binding},
        content_rules=[
            ContentInsertionRule(
                rule_id="content-rule",
                content_id="content",
                anchor_token=content_anchor_token("content"),
            )
        ],
    )
    assert assembly_runtime._has_image_layout_risk(context) is False


def test_pipeline_adapter_refreshes_only_variant_without_pending_image_work(
    tmp_path: Path,
    monkeypatch,
):
    source = tmp_path / "logical.docx"
    prepared = tmp_path / "prepared.docx"
    final = tmp_path / "final.docx"
    stage = tmp_path / ".final.material-stage.docx"
    Document().save(source)
    document = Document()
    document.add_paragraph("content only")
    document.save(prepared)
    digest = sha256(prepared.read_bytes()).hexdigest()
    refreshed: list[Path] = []

    class FakePipeline:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

        def execute(self, input_path):
            target = Path(self.kwargs["output_path_overrides"]["final"])
            copy2(input_path, target)
            return PipelineResult(
                success=True,
                output_paths={"final": str(target)},
                context=PipelineContext(),
            )

    monkeypatch.setattr(assembly_runtime, "Pipeline", FakePipeline)
    monkeypatch.setattr(assembly_runtime, "document_has_toc", lambda _doc: True)
    monkeypatch.setattr(
        assembly_runtime,
        "refresh_doc_fields_with_word",
        lambda path, timeout_sec: refreshed.append(Path(path)) or (True, "ok"),
    )
    config = resolve_config(TemplateConfig(), _minimal_scene())
    request = MaterialPipelineRequest(
        execution_id="e1",
        logical_source_path=str(source),
        prepared_input=FileEvidence(str(prepared), digest, prepared.stat().st_size),
        intake_snapshot_id=sha256(b"snapshot").hexdigest(),
        targets=(
            PipelineVariantTarget(
                variant_id="final",
                variant_version="1",
                final_output_path=str(final),
                owned_stage_output_path=str(stage),
            ),
        ),
    )
    adapter = assembly_runtime._WorkbenchMaterialPipelineAdapter(
        config=config,
        output_dir=tmp_path,
        output_suffix="_formatted",
        progress_cb=lambda *_args: None,
        cancel_check=lambda: False,
        force_delivery_presets=False,
        official_master=None,
        image_rule_tokens=("{{@img:qualification1}}",),
    )

    adapter(request)
    assert refreshed == [stage]

    document = Document()
    document.add_paragraph("{{@img:qualification1}}")
    document.save(prepared)
    request = MaterialPipelineRequest(
        execution_id="e2",
        logical_source_path=str(source),
        prepared_input=FileEvidence(
            str(prepared),
            sha256(prepared.read_bytes()).hexdigest(),
            prepared.stat().st_size,
        ),
        intake_snapshot_id=sha256(b"snapshot-2").hexdigest(),
        targets=(
            PipelineVariantTarget(
                variant_id="final",
                variant_version="1",
                final_output_path=str(final),
                owned_stage_output_path=str(stage),
            ),
        ),
    )
    adapter(request)
    assert refreshed == [stage]
