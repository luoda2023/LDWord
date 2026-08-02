from __future__ import annotations

import ast
from io import BytesIO
from pathlib import Path

from docx import Document
from lxml import etree

from src.config.resolved import ResolvedConfig
from src.formula_core import (
    convert_formula_node,
    normalize_formula_node,
    parse_document_formulas,
)
from src.modules.special.equation_table_format import EquationTableFormatModule
from src.modules.special.formula_convert import FormulaConvertModule
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn
from src.utils.ooxml_paragraph import replace_paragraph_payload_with_omml

ROOT = Path(__file__).resolve().parents[1]
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def test_migrated_formula_core_parses_complex_latex_and_builds_structured_omml():
    document = Document()
    document.add_paragraph(
        r"$$\begin{bmatrix}a_1&\frac{b}{c}\\d&\sum_{i=1}^{n}i\end{bmatrix}$$"
    )

    result = parse_document_formulas(document)
    assert result.total == 1
    occurrence = result.occurrences[0]
    normalized = normalize_formula_node(occurrence.node)
    outcome = convert_formula_node(normalized, "word_native", block=False)

    assert occurrence.source_type == "latex"
    assert outcome.success is True
    assert outcome.omml_element is not None
    assert outcome.omml_element.find(f".//{qn('m:m')}") is not None
    assert outcome.omml_element.find(f".//{qn('m:f')}") is not None
    assert outcome.omml_element.find(f".//{qn('m:nary')}") is not None


def test_thesis_formula_chain_converts_then_containerizes_standalone_formula():
    document = Document()
    document.add_paragraph(r"$$\frac{x_1}{y^2}$$")
    config = ResolvedConfig(mode_id="thesis")
    config.equation_numbering.numbering_format = "global"
    tracker = ChangeTracker()
    context = PipelineContext(mode_id="thesis")

    FormulaConvertModule().apply(document, config, tracker, context)
    EquationTableFormatModule().apply(document, config, tracker, context)

    assert len(document.paragraphs) == 0
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(0, 0)._tc.find(f".//{qn('m:f')}") is not None
    assert table.cell(0, 1).text == "(1)"
    assert tracker.get_by_module("formula_to_table")
    assert tracker.get_by_module("formula_style")


def test_formula_convert_nested_enablement_is_honored_by_the_module():
    document = Document()
    paragraph = document.add_paragraph(r"$$x^2$$")
    config = ResolvedConfig(mode_id="thesis")
    config.formula_convert.enabled = False
    context = PipelineContext(mode_id="thesis")

    FormulaConvertModule().apply(
        document,
        config,
        ChangeTracker(),
        context,
    )

    assert paragraph.text == r"$$x^2$$"
    assert paragraph._element.find(f".//{qn('m:oMath')}") is None
    assert context.formula_runtime["convert_enabled"] is False


def test_thesis_formula_subpasses_keep_their_v02_independent_enablement():
    document = Document()
    document.add_paragraph(r"$$x^2$$")
    config = ResolvedConfig(mode_id="thesis")
    config.formula_to_table.enabled = False
    config.equation_numbering.enabled = True
    config.formula_style.enabled = True
    tracker = ChangeTracker()
    context = PipelineContext(mode_id="thesis")

    FormulaConvertModule().apply(document, config, tracker, context)
    EquationTableFormatModule().apply(document, config, tracker, context)

    assert len(document.paragraphs) == 1
    assert len(document.tables) == 0
    assert not tracker.get_by_module("formula_to_table")
    style_record_count = len(tracker.get_by_module("formula_style"))

    config.formula_to_table.enabled = True
    config.equation_numbering.enabled = False
    config.formula_style.enabled = False
    EquationTableFormatModule().apply(document, config, tracker, context)

    assert len(document.paragraphs) == 0
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 1).text == ""
    assert tracker.get_by_module("formula_to_table")
    assert len(tracker.get_by_module("formula_style")) == style_record_count


def test_thesis_formula_to_table_retains_the_v02_block_only_strategy():
    document = Document()
    document.add_paragraph(r"$x_1$")
    config = ResolvedConfig(mode_id="thesis")
    config.formula_to_table.block_only = False
    config.equation_numbering.enabled = False
    config.formula_style.enabled = False
    tracker = ChangeTracker()
    context = PipelineContext(mode_id="thesis")

    FormulaConvertModule().apply(document, config, tracker, context)
    EquationTableFormatModule().apply(document, config, tracker, context)

    assert len(document.paragraphs) == 0
    assert len(document.tables) == 1
    assert tracker.get_by_module("formula_to_table")


def test_formula_conversion_does_not_reclassify_plain_data_table_cells():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "条件 a=1"
    table.cell(0, 1).text = "通过"
    before = table._element.xml

    FormulaConvertModule().apply(
        document,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    assert table._element.xml == before
    assert table._element.find(f".//{qn('m:oMath')}") is None


def test_paragraph_replacement_preserves_cross_paragraph_bookmarks_on_roundtrip():
    document = Document()
    first = document.add_paragraph("x=y")
    second = document.add_paragraph("after")
    start = etree.Element(f"{{{W_NS}}}bookmarkStart")
    start.set(f"{{{W_NS}}}id", "191")
    start.set(f"{{{W_NS}}}name", "OLE_LINK107")
    end = etree.Element(f"{{{W_NS}}}bookmarkEnd")
    end.set(f"{{{W_NS}}}id", "191")
    first._element.insert(0, start)
    second._element.insert(0, end)

    omath = etree.Element(f"{{{M_NS}}}oMath")
    run = etree.SubElement(omath, f"{{{M_NS}}}r")
    etree.SubElement(run, f"{{{M_NS}}}t").text = "x=y"
    assert replace_paragraph_payload_with_omml(first._element, omath).applied is True
    assert first._element.find(qn("m:oMath")) is not None
    assert first._element.find(f"{qn('w:r')}/{qn('m:oMath')}") is None

    payload = BytesIO()
    document.save(payload)
    payload.seek(0)
    reloaded = Document(payload)
    starts = reloaded.element.body.findall(f".//{{{W_NS}}}bookmarkStart")
    ends = reloaded.element.body.findall(f".//{{{W_NS}}}bookmarkEnd")
    assert [item.get(f"{{{W_NS}}}id") for item in starts].count("191") == 1
    assert [item.get(f"{{{W_NS}}}id") for item in ends].count("191") == 1


def test_mathtype_ole_core_remains_python_311_compatible():
    source_path = ROOT / "src" / "formula_core" / "mathtype_ole.py"
    tree = ast.parse(
        source_path.read_text(encoding="utf-8"),
        filename=str(source_path),
        feature_version=11,
    )
    assert tree is not None
