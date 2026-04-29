import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.modules.structure.heading_recognition import (
    HeadingRecognitionModule,
    _detect_heading,
)
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker


def _apply_heading_recognition(doc: Document):
    module = HeadingRecognitionModule()
    config = ResolvedConfig()
    context = PipelineContext()
    tracker = ChangeTracker()
    module.apply(doc, config, tracker, context)
    return context


def test_heading_recognition_ignores_toc_like_lines_even_with_heading_visual_traits():
    doc = Document()
    para = doc.add_paragraph("第一章 绪论\t1")
    para.runs[0].bold = True
    para.runs[0].font.size = Pt(16)

    assert _detect_heading(para, para.text.strip()) is None


def test_heading_recognition_ignores_reference_entries_even_with_heading_visual_traits():
    doc = Document()
    para = doc.add_paragraph("[1] Author. Title[J]. Journal, 2024, 12(3): 1-8.")
    para.runs[0].bold = True
    para.runs[0].font.size = Pt(16)

    assert _detect_heading(para, para.text.strip()) is None


def test_heading_recognition_ignores_pageref_toc_entries_even_with_heading_style():
    from src.shared.engine.field_builder import build_pageref_field

    doc = Document()
    para = doc.add_paragraph("Abstract")
    para.style = doc.styles["Heading 1"]
    para._element.append(build_pageref_field("_Toc123"))

    assert _detect_heading(para, para.text.strip()) is None


def test_heading_recognition_ignores_structured_table_rows_with_heading_visual_traits():
    doc = Document()
    para = doc.add_paragraph("41.9 mW cm-2\t1562\tthis work")
    para.runs[0].bold = True
    para.runs[0].font.size = Pt(16)

    assert _detect_heading(para, para.text.strip()) is None


def test_heading_recognition_detects_references_section_from_reference_entry_cluster():
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容")
    doc.add_paragraph("[1] Author. Title[J]. Journal, 2024, 12(3): 1-8.")
    doc.add_paragraph("[2] Author. Another title[J]. Journal, 2023, 11(2): 9-12.")

    context = _apply_heading_recognition(doc)
    doc_tree = context.doc_tree

    assert doc_tree.get_section("references") is not None
    assert doc_tree.get_section("references").start_index == 2
    assert doc_tree.get_section_for_paragraph(2) == "references"
    assert doc_tree.get_section_for_paragraph(3) == "references"


def test_heading_recognition_detects_toc_section_from_toc_styles():
    doc = Document()
    doc.add_paragraph("目录")
    if "TOC 1" not in [style.name for style in doc.styles]:
        doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    if "TOC 2" not in [style.name for style in doc.styles]:
        doc.styles.add_style("TOC 2", WD_STYLE_TYPE.PARAGRAPH)

    toc_1 = doc.add_paragraph("第一章 绪论\t1")
    toc_1.style = doc.styles["TOC 1"]
    toc_2 = doc.add_paragraph("1.1 研究背景\t2")
    toc_2.style = doc.styles["TOC 2"]
    doc.add_heading("第一章 绪论", level=1)

    context = _apply_heading_recognition(doc)
    doc_tree = context.doc_tree

    assert doc_tree.get_section("toc") is not None
    assert doc_tree.get_section("toc").start_index == 0
    assert doc_tree.get_section_for_paragraph(1) == "toc"
    assert doc_tree.get_section_for_paragraph(2) == "toc"


def test_heading_recognition_detects_cover_section_from_content():
    doc = Document()
    doc.add_paragraph("硕士学位论文")
    doc.add_paragraph("题目")
    doc.add_paragraph("摘要")
    doc.add_paragraph("摘要内容")

    context = _apply_heading_recognition(doc)
    doc_tree = context.doc_tree

    assert doc_tree.get_section("cover") is not None
    assert doc_tree.get_section("cover").start_index == 0


def test_heading_recognition_exposes_only_body_headings_for_downstream_numbering():
    doc = Document()
    doc.add_heading("硕士学位论文", level=1)
    doc.add_heading("摘要", level=1)
    doc.add_paragraph("摘要正文。")
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容。")

    context = _apply_heading_recognition(doc)

    assert context.doc_tree.get_section_for_paragraph(0) == "cover"
    assert context.doc_tree.get_section_for_paragraph(1) == "abstract_cn"
    assert context.doc_tree.get_section_for_paragraph(3) == "body"
    assert context.heading_map == {3: 1}
    assert [heading.para_index for heading in context.doc_tree.headings] == [3]


def test_heading_recognition_detects_pre_numbering_statement_pages():
    doc = Document()
    doc.add_paragraph("硕士学位论文")
    doc.add_paragraph("原创性声明")
    doc.add_paragraph("本人郑重声明。")
    doc.add_paragraph("学位论文版权使用授权书")
    doc.add_paragraph("授权说明。")
    doc.add_paragraph("摘要")
    doc.add_paragraph("摘要内容")
    doc.add_heading("第一章 绪论", level=1)

    context = _apply_heading_recognition(doc)
    doc_tree = context.doc_tree

    assert doc_tree.get_section("statement") is not None
    assert doc_tree.get_section("authorization") is not None
    assert doc_tree.get_section_for_paragraph(1) == "statement"
    assert doc_tree.get_section_for_paragraph(3) == "authorization"
    assert doc_tree.get_section_for_paragraph(5) == "abstract_cn"
    assert doc_tree.get_section_for_paragraph(7) == "body"


def test_heading_recognition_infers_body_between_front_and_back_matter():
    doc = Document()
    doc.add_paragraph("硕士学位论文")
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] Author. Title[J]. Journal, 2024, 12(3): 1-8.")
    doc.add_paragraph("[2] Author. Another title[J]. Journal, 2023, 11(2): 9-12.")

    context = _apply_heading_recognition(doc)
    doc_tree = context.doc_tree

    assert doc_tree.get_section("cover") is not None
    assert doc_tree.get_section("body") is not None
    assert doc_tree.get_section("body").start_index == 1
    assert doc_tree.get_section_for_paragraph(1) == "body"
    assert doc_tree.get_section_for_paragraph(2) == "body"
    assert doc_tree.get_section_for_paragraph(3) == "references"


def test_heading_recognition_detects_abstract_and_acknowledgment_sections_without_heading_styles():
    doc = Document()
    doc.add_paragraph("摘要")
    doc.add_paragraph("这里是摘要内容。")
    doc.add_paragraph("第一章 绪论")
    doc.add_paragraph("正文内容。")
    doc.add_paragraph("致谢")
    doc.add_paragraph("感谢老师。")

    context = _apply_heading_recognition(doc)
    doc_tree = context.doc_tree

    assert doc_tree.get_section("abstract_cn") is not None
    assert doc_tree.get_section("acknowledgment") is not None
    assert doc_tree.get_section_for_paragraph(0) == "abstract_cn"
    assert doc_tree.get_section_for_paragraph(1) == "abstract_cn"
    assert doc_tree.get_section_for_paragraph(2) == "body"
    assert doc_tree.get_section_for_paragraph(4) == "acknowledgment"
    assert doc_tree.get_section_for_paragraph(5) == "acknowledgment"


def test_heading_recognition_detects_inline_abstract_anchor():
    doc = Document()
    doc.add_paragraph("摘要：本文研究毕业论文格式自动修订中的结构识别问题。")
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容。")

    context = _apply_heading_recognition(doc)
    doc_tree = context.doc_tree

    assert doc_tree.get_section("abstract_cn") is not None
    assert doc_tree.get_section("abstract_cn").start_index == 0
    assert doc_tree.get_section_for_paragraph(0) == "abstract_cn"
    assert context.heading_map == {1: 1}


def test_heading_recognition_keeps_appendix_section_at_first_appendix_title():
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容。")
    doc.add_paragraph("附录A 数据表")
    doc.add_paragraph("附录A 正文。")
    doc.add_paragraph("附录B 访谈提纲")
    doc.add_paragraph("附录B 正文。")

    context = _apply_heading_recognition(doc)
    appendix = context.doc_tree.get_section("appendix")

    assert appendix is not None
    assert appendix.start_index == 2
    assert context.doc_tree.get_section_for_paragraph(2) == "appendix"
    assert context.doc_tree.get_section_for_paragraph(4) == "appendix"
