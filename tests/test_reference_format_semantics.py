import sys
from pathlib import Path
from types import SimpleNamespace

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.config.template import StyleConfig
from src.modules.special.reference_format import (
    ReferenceFormatModule,
    _format_reference_entry,
    _resolve_reference_entry_style,
)
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn


def test_reference_format_uses_reference_variant_style_with_rule_overrides():
    config = ResolvedConfig()
    config.styles["body"] = StyleConfig(
        font_cn="宋体",
        font_en="Times New Roman",
        size_pt=12,
        alignment="justify",
        line_spacing_type="multiple",
        line_spacing_pt=1.5,
        space_before_pt=6,
        space_after_pt=6,
    )
    config.styles["references_body"] = StyleConfig(
        font_cn="黑体",
        font_en="Arial",
        size_pt=11,
        alignment="left",
        line_spacing_type="exact",
        line_spacing_pt=18,
        space_before_pt=4,
        space_after_pt=8,
    )
    config.reference_style.hanging_indent_cm = 0.74
    config.reference_style.space_after_pt = 2
    config.reference_style.space_after_unit = "lines"
    config.reference_style.font_en = "Calibri"

    style = _resolve_reference_entry_style(config)

    assert style.font_cn == "黑体"
    assert style.font_en == "Calibri"
    assert style.size_pt == 11
    assert style.alignment == "left"
    assert style.line_spacing_type == "exact"
    assert style.line_spacing_pt == 18
    assert style.space_before_pt == 0
    assert style.space_after_pt == 2
    assert style.space_after_unit == "lines"
    assert style.special_indent_mode == "hanging"
    assert style.special_indent_value == 0.74
    assert style.special_indent_unit == "cm"


def test_reference_format_syncs_spacing_indent_and_removes_numbering():
    document = Document()
    para = document.add_paragraph("[1] 参考文献条目")
    p_pr = para._element.get_or_add_pPr()
    p_pr.append(document._part._element.makeelement(qn("w:numPr")))

    style = StyleConfig(
        font_cn="宋体",
        font_en="Times New Roman",
        size_pt=10.5,
        alignment="justify",
        line_spacing_type="exact",
        line_spacing_pt=16,
        space_before_pt=0,
        space_after_pt=6,
        special_indent_mode="hanging",
        special_indent_value=0.74,
        special_indent_unit="cm",
        hanging_indent_chars=0.74,
        hanging_indent_unit="cm",
    )

    _format_reference_entry(para, style)

    spacing = para._element.find(qn("w:pPr")).find(qn("w:spacing"))
    ind = para._element.find(qn("w:pPr")).find(qn("w:ind"))
    num_pr = para._element.find(qn("w:pPr")).find(qn("w:numPr"))

    assert spacing.get(qn("w:before")) == "0"
    assert spacing.get(qn("w:after")) == "120"
    assert spacing.get(qn("w:line")) == "320"
    assert spacing.get(qn("w:lineRule")) == "exact"
    assert ind.get(qn("w:left")) == "420"
    assert ind.get(qn("w:hanging")) == "420"
    assert ind.get(qn("w:leftChars")) is None
    assert ind.get(qn("w:hangingChars")) is None
    assert num_pr is None
    assert para.runs[0].font.size.pt == 10.5


def test_reference_format_syncs_rule_spacing_unit_to_ooxml():
    document = Document()
    para = document.add_paragraph("[1] Reference")

    config = ResolvedConfig()
    config.styles["body"] = StyleConfig(line_spacing_type="exact", line_spacing_pt=16)
    config.reference_style.space_after_pt = 1.5
    config.reference_style.space_after_unit = "lines"

    style = _resolve_reference_entry_style(config)
    _format_reference_entry(para, style)

    spacing = para._element.find(qn("w:pPr")).find(qn("w:spacing"))
    assert spacing.get(qn("w:after")) is None
    assert spacing.get(qn("w:afterLines")) == "150"


def test_reference_format_uses_doc_tree_reference_range_when_title_is_missing():
    document = Document()
    document.add_paragraph("正文段落")
    ref1 = document.add_paragraph("[1] Author. Title[J]. Journal, 2024, 12(3): 1-8.")
    ref2 = document.add_paragraph("[2] Another. Title[J]. Journal, 2023, 11(2): 9-12.")
    document.add_paragraph("附录")

    config = ResolvedConfig()
    config.styles["body"] = StyleConfig(font_cn="宋体", font_en="Times New Roman", size_pt=12)
    config.styles["references_body"] = StyleConfig(
        font_cn="黑体",
        font_en="Arial",
        size_pt=10.5,
        alignment="justify",
        line_spacing_type="exact",
        line_spacing_pt=16,
    )
    config.reference_style.hanging_indent_cm = 0.74
    config.reference_style.space_after_pt = 4

    context = PipelineContext()
    context.doc_tree = SimpleNamespace(
        get_section=lambda section: SimpleNamespace(start_index=1, end_index=3)
        if section == "references"
        else None
    )

    ReferenceFormatModule().apply(document, config, ChangeTracker(), context)

    ref1_ind = ref1._element.find(qn("w:pPr")).find(qn("w:ind"))
    ref2_ind = ref2._element.find(qn("w:pPr")).find(qn("w:ind"))

    assert ref1.runs[0].font.size.pt == 10.5
    assert ref2.runs[0].font.size.pt == 10.5
    assert ref1_ind.get(qn("w:hanging")) == "420"
    assert ref2_ind.get(qn("w:hanging")) == "420"
