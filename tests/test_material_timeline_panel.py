from docx import Document
from PySide6.QtTest import QTest

from src.config.entity import EntityProfile
from src.qt_api import QApplication, Qt
from src.shared.engine.material_timeline import (
    default_timeline_plan,
    default_timeline_segment,
)
from src.shared.ui.sizing import resolved_control_height
from src.shared.ui.theme import get_theme
from src.ui.bridge import PanelBridge
from src.ui.panels.assets_panel import AssetsPanel
from src.ui.panels.workbench.quick_execution_detail import QuickExecutionDetail


def _app():
    return QApplication.instance() or QApplication([])


def _add_segment(panel: AssetsPanel) -> str:
    panel._add_timeline_segment()
    return next(
        plan_id
        for plan_id, plan in panel._selected_profile().timeline_plans.items()
        if not bool(plan.get("deleted", False))
    )


def _set_segment_dates(
    panel: AssetsPanel,
    plan_id: str,
    start: str,
    end: str,
) -> None:
    panel._timeline_start_edits[plan_id].setText(start)
    panel._timeline_end_edits[plan_id].setText(end)


def test_timeline_is_a_permanent_page_without_empty_placeholder_or_summary_cards():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        assert not panel._section_nav_cards["timeline"].isHidden()
        assert "timeline" in panel._section_pages
        assert not hasattr(panel, "_timeline_entry_card")
        assert not hasattr(panel, "_timeline_back_btn")
        assert panel._section_summary_cards["timeline"].isHidden()
        assert len(panel._timeline_segments_controller) == 0
        assert not hasattr(panel, "_timeline_empty_label")
        assert panel._timeline_segments_container.isHidden()
        assert panel._timeline_summary_items() == []
    finally:
        panel.close()


def test_adding_first_segment_generates_stable_copyable_tokens():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        plan = panel._selected_profile().timeline_plans[plan_id]
        tokens = [node["outputs"][0]["field"] for node in plan["nodes"] if node["active"]]

        assert plan["segment_no"] == 1
        assert tokens == ["时间节点1-1", "时间节点1-2", "时间节点1-3"]
        assert len(panel._timeline_segments_controller) == 1
        assert len(panel._timeline_node_controllers[plan_id]) == 3
        assert panel._timeline_node_token_edits[(plan_id, "node_2")].text() == "{{@time:时间节点1-2}}"
    finally:
        panel.close()


def test_direct_segment_dates_resolve_tokens_and_inherit_start_date_format():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        _set_segment_dates(panel, plan_id, "2025年10月1日", "2025年10月11日")

        values = panel.material_context().resolved_entity_data()

        assert values["时间节点1-1"] == "2025年10月1日"
        assert values["时间节点1-2"] == "2025年10月6日"
        assert values["时间节点1-3"] == "2025年10月11日"
        assert panel._timeline_node_result_edits[(plan_id, "node_2")].text() == "2025年10月6日"
    finally:
        panel.close()


def test_mixed_input_date_styles_warn_and_still_follow_the_start_style():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        _set_segment_dates(panel, plan_id, "2025-10-01", "2025年10月11日")

        values = panel.material_context().resolved_entity_data()

        assert values["时间节点1-2"] == "2025-10-06"
        assert panel._timeline_segment_status_labels[plan_id].text() == "1 项提醒"
    finally:
        panel.close()


def test_explicit_date_format_controls_generated_dates_and_persists():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        _set_segment_dates(panel, plan_id, "2025-05-31", "2025-06-10")
        combo = panel._timeline_date_format_combos[plan_id]

        combo.setCurrentIndex(combo.findData("d/M/yyyy"))

        plan = panel._selected_profile().timeline_plans[plan_id]
        assert plan["format_mode"] == "d/M/yyyy"
        assert plan["output_format"] == "d/M/yyyy"
        assert panel.material_context().resolved_entity_data()["时间节点1-1"] == "31/5/2025"
        assert panel._timeline_segment_status_labels[plan_id].text() == "长期复用"
    finally:
        panel.close()


def test_date_format_menu_covers_common_chinese_and_numeric_spellings():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        combo = panel._timeline_date_format_combos[plan_id]
        formats = {combo.itemData(index) for index in range(combo.count())}

        assert {
            "auto",
            "yyyy-M-d",
            "yyyy年M月d日",
            "M月d日",
            "yyyy.M.d",
            "yyyy/M/d",
            "d/M/yyyy",
            "M/d/yyyy",
        } <= formats
    finally:
        panel.close()


def test_free_time_input_scope_is_editable_in_workbench_and_drives_timeline():
    _app()
    panel = AssetsPanel(PanelBridge())
    detail = QuickExecutionDetail()
    try:
        panel._archive_id_edit.setText("timeline-package")
        panel._profile_id_edit.setText("timeline-profile")
        plan_id = _add_segment(panel)
        combo = panel._timeline_input_scope_combos[plan_id]
        _set_segment_dates(panel, plan_id, "2025-05-31", "2025-06-10")

        combo.setCurrentIndex(combo.findData("floating"))

        plan = panel._selected_profile().timeline_plans[plan_id]
        assert plan["input_scope"] == "floating"
        assert plan["start_value"] == ""
        assert plan["end_value"] == ""
        assert plan["start_field"] == "时间段1开始日期"
        assert plan["end_field"] == "时间段1结束日期"
        context = panel.material_context()
        assert context.field_scopes["时间段1开始日期"] == "floating"
        assert context.field_scopes["时间段1结束日期"] == "floating"
        assert panel._timeline_segment_status_labels[plan_id].text() == "每次填写"
        assert not panel._timeline_start_edits[plan_id].isEnabled()
        assert not panel._timeline_end_edits[plan_id].isEnabled()

        detail.set_work_mode("official")
        detail.set_material_context(context)
        assert "时间段1开始日期" in detail._floating_field_inputs
        assert "时间段1结束日期" in detail._floating_field_inputs
        assert detail._floating_field_inputs["时间段1开始日期"]._editor_kind == "date"
        assert (
            detail._floating_field_name_labels["时间段1开始日期"].text()
            == "{{@time:时间段1开始日期}}"
        )
        changed_contexts = []
        detail.material_context_changed.connect(changed_contexts.append)
        detail._floating_field_inputs["时间段1开始日期"].setText("2025-06-01")
        detail._floating_field_inputs["时间段1结束日期"].setText("2025-06-11")
        resolved = changed_contexts[-1].resolved_entity_data()
        assert resolved["时间节点1-1"] == "2025-06-01"
        assert resolved["时间节点1-3"] == "2025-06-11"

        panel._last_received_material_context = context.clone()
        panel._on_material_context_changed(changed_contexts[-1])
        assert panel._timeline_start_edits[plan_id].text() == "2025-06-01"
        assert panel._timeline_end_edits[plan_id].text() == "2025-06-11"

        weekend = panel._timeline_weekend_combos[plan_id]
        weekend.setCurrentIndex(weekend.findData("forward"))
        profile = panel._selected_profile()
        assert "时间段1开始日期" not in profile.fields
        assert "时间段1结束日期" not in profile.fields

        QTest.mouseClick(
            panel._timeline_node_token_edits[(plan_id, "node_2")],
            Qt.LeftButton,
        )
        panel._duplicate_timeline_segment(plan_id)
        profile = panel._selected_profile()
        assert "时间段1开始日期" not in profile.fields
        assert "时间段1结束日期" not in profile.fields
        assert "时间段2开始日期" not in profile.fields
        assert "时间段2结束日期" not in profile.fields
        next_context = panel.material_context()
        assert next_context.entity_data["时间段1开始日期"] == "2025-06-01"
        assert next_context.entity_data["时间段1结束日期"] == "2025-06-11"
        assert next_context.entity_data["时间段2开始日期"] == "2025-06-01"
        assert next_context.entity_data["时间段2结束日期"] == "2025-06-11"
    finally:
        detail.close()
        panel.close()


def test_deleting_floating_segment_does_not_resurrect_its_anchor_fields(
    monkeypatch,
):
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        _set_segment_dates(panel, plan_id, "2025-05-31", "2025-06-10")
        scope_combo = panel._timeline_input_scope_combos[plan_id]
        scope_combo.setCurrentIndex(scope_combo.findData("floating"))
        assert panel._persist_current_profile_editor() is True

        profile = panel._selected_profile()
        anchor_keys = {"时间段1开始日期", "时间段1结束日期"}
        assert anchor_keys <= set(profile.field_scopes)
        assert anchor_keys <= set(profile.declared_field_keys)
        monkeypatch.setattr(
            "src.ui.panels.assets.timeline_presenter.confirm",
            lambda *args, **kwargs: True,
        )

        panel._remove_timeline_segment(plan_id)
        panel._refresh_timeline_ui()

        deleted = profile.timeline_plans[plan_id]
        assert deleted["deleted"] is True
        assert deleted["start_field"] == "时间段1开始日期"
        assert deleted["end_field"] == "时间段1结束日期"
        assert anchor_keys.isdisjoint(profile.field_scopes)
        assert anchor_keys.isdisjoint(profile.fields)
        assert anchor_keys.isdisjoint(profile.declared_field_keys)
        assert anchor_keys.isdisjoint(panel.material_context().field_scopes)
    finally:
        panel.close()


def test_non_deleted_timeline_references_cannot_be_renamed_or_deleted_as_fields():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan = default_timeline_segment(1)
        plan.update(
            {
                "enabled": False,
                "input_scope": "floating",
                "start_field": "自定义开始日",
                "end_field": "自定义结束日",
            }
        )
        profile = panel._selected_profile()
        profile.timeline_plans = {"custom": plan}
        profile.field_scopes = {
            "自定义开始日": "floating",
            "自定义结束日": "floating",
        }
        profile.fields = {"自定义开始日": "2026-01-01"}
        profile.declared_field_keys = ["自定义开始日", "自定义结束日"]
        panel._manual_field_keys = ["自定义开始日", "自定义结束日"]
        panel._declared_field_keys = {"自定义开始日", "自定义结束日"}

        panel._remove_official_field("自定义开始日")
        panel._remove_manual_field("自定义开始日")
        panel._commit_material_field_key("自定义开始日", "改名后的开始日")

        assert profile.timeline_plans["custom"]["start_field"] == "自定义开始日"
        assert profile.field_scopes["自定义开始日"] == "floating"
        assert profile.fields["自定义开始日"] == "2026-01-01"
        assert "自定义开始日" in profile.declared_field_keys
        assert "改名后的开始日" not in profile.field_scopes
    finally:
        panel.close()


def test_timeline_segment_can_collapse_and_uses_material_field_control_height():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel.resize(1440, 900)
        panel.show()
        panel._section_nav.select_card("timeline")
        plan_id = _add_segment(panel)
        app.processEvents()

        assert panel._timeline_segment_ranges[plan_id].text() == "3 项"
        assert "font-weight: 700" in panel._timeline_segment_ranges[plan_id].styleSheet()
        expected_height = resolved_control_height(get_theme(), "md")
        assert panel._timeline_start_edits[plan_id].height() == expected_height
        assert panel._timeline_date_format_combos[plan_id].height() == expected_height
        assert panel._timeline_node_rows[(plan_id, "node_1")].minimumHeight() == 56
        guide = panel._timeline_node_column_guides[plan_id]
        metrics = guide.metrics()
        assert metrics.guide_visible is True
        assert guide.position_label.text() == "位置（%）"
        assert panel._timeline_node_token_edits[(plan_id, "node_1")].width() == guide.token_label.width()
        assert panel._timeline_node_ratio_edits[(plan_id, "node_1")].width() == (
            metrics.position_width
        )

        QTest.mouseClick(
            panel._timeline_segment_collapse_buttons[plan_id],
            Qt.LeftButton,
        )
        assert panel._timeline_segment_bodies[plan_id].isHidden()
        QTest.mouseClick(
            panel._timeline_segment_collapse_buttons[plan_id],
            Qt.LeftButton,
        )
        assert not panel._timeline_segment_bodies[plan_id].isHidden()
    finally:
        panel.close()


def test_legacy_free_time_anchor_names_migrate_without_losing_dates():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan = default_timeline_segment(1)
        plan.update(
            {
                "input_scope": "floating",
                "start_field": "时间段开始日期1",
                "end_field": "时间段结束日期1",
                "start_value": "",
                "end_value": "",
                "enabled": True,
            }
        )
        profile = panel._selected_profile()
        profile.timeline_plans = {"segment_1": plan}
        profile.fields = {
            "时间段开始日期1": "2025-05-31",
            "时间段结束日期1": "2025-06-10",
        }
        profile.field_scopes = {
            "时间段开始日期1": "floating",
            "时间段结束日期1": "floating",
        }
        panel._template_field_values.update(profile.fields)

        panel._refresh_timeline_ui()

        migrated = profile.timeline_plans["segment_1"]
        assert migrated["start_field"] == "时间段1开始日期"
        assert migrated["end_field"] == "时间段1结束日期"
        assert profile.fields["时间段1开始日期"] == "2025-05-31"
        assert profile.fields["时间段1结束日期"] == "2025-06-10"
        assert "时间段开始日期1" not in profile.fields
        assert panel._resolve_timeline_preview_fields(profile.fields)["时间节点1-3"] == "2025-06-10"
    finally:
        panel.close()


def test_clicking_timeline_token_copies_the_exact_braced_token():
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        token_edit = panel._timeline_node_token_edits[(plan_id, "node_2")]

        QTest.mouseClick(token_edit, Qt.LeftButton)

        assert app.clipboard().text() == "{{@time:时间节点1-2}}"
        assert panel._selected_profile().timeline_plans[plan_id]["token_copied"] is True
    finally:
        panel.close()


def test_node_count_is_the_only_node_add_remove_control_and_restores_stable_slots():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        spin = panel._timeline_node_count_spins[plan_id]

        spin.setValue(5)
        five_node_plan = panel._selected_profile().timeline_plans[plan_id]
        active = [node for node in five_node_plan["nodes"] if node["active"]]
        node_5_id = active[-1]["node_id"]
        assert [node["rule"]["value"] for node in active] == ["0", "0.25", "0.5", "0.75", "1"]
        assert [node["outputs"][0]["field"] for node in active] == [
            "时间节点1-1",
            "时间节点1-2",
            "时间节点1-3",
            "时间节点1-4",
            "时间节点1-5",
        ]
        assert panel._timeline_node_ratio_edits[(plan_id, "node_1")].isReadOnly()
        assert not panel._timeline_node_ratio_edits[(plan_id, "node_2")].isReadOnly()
        assert panel._timeline_node_ratio_edits[(plan_id, "node_5")].isReadOnly()

        spin.setValue(3)
        assert sum(
            bool(node["active"])
            for node in panel._selected_profile().timeline_plans[plan_id]["nodes"]
        ) == 3

        spin.setValue(5)
        restored = [
            node
            for node in panel._selected_profile().timeline_plans[plan_id]["nodes"]
            if node["active"]
        ]
        assert restored[-1]["node_id"] == node_5_id
        assert restored[-1]["outputs"][0]["field"] == "时间节点1-5"
    finally:
        panel.close()


def test_middle_node_ratio_fine_tunes_its_token_date_and_context():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        _set_segment_dates(panel, plan_id, "2025-01-01", "2025-01-11")
        ratio = panel._timeline_node_ratio_edits[(plan_id, "node_2")]

        ratio.setText("20")

        assert panel._timeline_node_result_edits[(plan_id, "node_2")].text() == "2025-01-03"
        assert panel.material_context().resolved_entity_data()["时间节点1-2"] == "2025-01-03"
        assert panel._selected_profile().timeline_plans[plan_id]["nodes"][1]["rule"]["value"] == "0.2"
    finally:
        panel.close()


def test_weekend_strategy_is_scoped_to_one_segment():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        first_id = _add_segment(panel)
        _set_segment_dates(panel, first_id, "2026-07-10", "2026-07-11")
        first_weekend = panel._timeline_weekend_combos[first_id]
        first_weekend.setCurrentIndex(first_weekend.findData("forward"))
        panel._duplicate_timeline_segment(first_id)
        second_id = panel._active_timeline_segment_items()[-1][0]
        second_weekend = panel._timeline_weekend_combos[second_id]
        second_weekend.setCurrentIndex(second_weekend.findData("none"))

        values = panel.material_context().resolved_entity_data()

        assert values["时间节点1-3"] == "2026-07-13"
        assert values["时间节点2-3"] == "2026-07-11"
        assert panel._selected_profile().timeline_plans[first_id]["calendar"]["weekend_adjust"] == "forward"
        assert panel._selected_profile().timeline_plans[second_id]["calendar"]["weekend_adjust"] == "none"
    finally:
        panel.close()


def test_copying_whole_segment_deep_copies_rules_but_generates_new_tokens():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        first_id = _add_segment(panel)
        _set_segment_dates(panel, first_id, "2025-01-01", "2025-01-11")
        panel._timeline_node_ratio_edits[(first_id, "node_2")].setText("20")

        panel._duplicate_timeline_segment(first_id)
        active_items = panel._active_timeline_segment_items()
        second_id, second = active_items[-1]

        assert second["segment_no"] == 2
        assert second["nodes"][1]["rule"]["value"] == "0.2"
        assert second["nodes"][1]["outputs"] == [
            {"field": "时间节点2-2", "format": "yyyy-MM-dd"}
        ]

        panel._timeline_node_ratio_edits[(second_id, "node_2")].setText("70")
        assert panel._selected_profile().timeline_plans[first_id]["nodes"][1]["rule"]["value"] == "0.2"
        assert panel._selected_profile().timeline_plans[second_id]["nodes"][1]["rule"]["value"] == "0.7"
    finally:
        panel.close()


def test_unreferenced_deleted_segment_reuses_its_plan_slot_without_renumbering_others(
    monkeypatch,
):
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        first_id = _add_segment(panel)
        panel._timeline_node_count_spins[first_id].setValue(5)
        panel._duplicate_timeline_segment(first_id)
        second_id = panel._active_timeline_segment_items()[-1][0]
        monkeypatch.setattr(
            "src.ui.panels.assets.timeline_presenter.confirm",
            lambda *args, **kwargs: True,
        )

        panel._remove_timeline_segment(first_id)
        deleted = panel._selected_profile().timeline_plans[first_id]
        assert deleted["deleted"] is True
        assert deleted["number_state"] == "reusable"
        assert panel._selected_profile().timeline_plans[second_id]["segment_no"] == 2
        assert panel._timeline_node_token_edits[(second_id, "node_1")].text() == "{{@time:时间节点2-1}}"

        panel._add_timeline_segment()
        active_items = panel._active_timeline_segment_items()
        assert [plan_id for plan_id, _plan in active_items] == [first_id, second_id]
        assert [plan["segment_no"] for _plan_id, plan in active_items] == [1, 2]
        assert panel._timeline_node_token_edits[(first_id, "node_1")].text() == "{{@time:时间节点1-1}}"
        assert panel._timeline_node_token_edits[(second_id, "node_1")].text() == "{{@time:时间节点2-1}}"
        assert {"时间节点1-4", "时间节点1-5"}.issubset(
            panel._selected_profile().timeline_plans[first_id]["retired_outputs"]
        )
    finally:
        panel.close()


def test_copied_segment_defaults_to_reserved_and_new_segment_uses_next_number(
    monkeypatch,
):
    app = _app()
    panel = AssetsPanel(PanelBridge())
    try:
        first_id = _add_segment(panel)
        panel._duplicate_timeline_segment(first_id)
        second_id = panel._active_timeline_segment_items()[-1][0]
        QTest.mouseClick(
            panel._timeline_node_token_edits[(first_id, "node_1")],
            Qt.LeftButton,
        )
        chooser_calls = []

        def choose_state(**kwargs):
            chooser_calls.append(kwargs)
            return "reserved"

        monkeypatch.setattr(
            panel,
            "_choose_timeline_deletion_number_state",
            choose_state,
        )

        panel._remove_timeline_segment(first_id)

        assert app.clipboard().text() == "{{@time:时间节点1-1}}"
        assert chooser_calls == [
            {
                "segment_no": 1,
                "referenced_fields": set(),
                "token_copied": True,
                "reference_occurrences": 0,
                "scan_failed": False,
            }
        ]
        assert panel._selected_profile().timeline_plans[first_id]["number_state"] == "reserved"

        panel._add_timeline_segment()
        active_items = panel._active_timeline_segment_items()
        assert [plan["segment_no"] for _plan_id, plan in active_items] == [2, 3]
        assert panel._selected_profile().timeline_plans[second_id]["segment_no"] == 2
        assert panel._timeline_node_token_edits[(second_id, "node_1")].text() == "{{@time:时间节点2-1}}"
    finally:
        panel.close()


def test_user_can_release_a_copied_segment_number_for_reuse(monkeypatch):
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        first_id = _add_segment(panel)
        panel._duplicate_timeline_segment(first_id)
        second_id = panel._active_timeline_segment_items()[-1][0]
        QTest.mouseClick(
            panel._timeline_node_token_edits[(first_id, "node_2")],
            Qt.LeftButton,
        )
        monkeypatch.setattr(
            panel,
            "_choose_timeline_deletion_number_state",
            lambda **_kwargs: "reusable",
        )

        panel._remove_timeline_segment(first_id)
        panel._add_timeline_segment()

        active_items = panel._active_timeline_segment_items()
        assert [plan_id for plan_id, _plan in active_items] == [first_id, second_id]
        assert [plan["segment_no"] for _plan_id, plan in active_items] == [1, 2]
        reused = panel._selected_profile().timeline_plans[first_id]
        assert reused["number_state"] == "active"
        assert reused["token_copied"] is False
    finally:
        panel.close()


def test_docx_reference_evidence_counts_occurrences_and_inactive_node_tokens(
    monkeypatch,
    tmp_path,
):
    source = tmp_path / "timeline-references.docx"
    document = Document()
    document.add_paragraph(
        "{{@time:时间节点1-5}} / {{@time:时间节点1-5}} / {{@time:时间节点1-2}}"
    )
    document.add_paragraph("{{@time:时间节点2-1}} / {{@time:时间节点10-1}} / {{@text:其他字段}}")
    document.save(source)
    bridge = PanelBridge()
    bridge.set_current_document_path(str(source))
    _app()
    panel = AssetsPanel(bridge)
    try:
        plan_id = _add_segment(panel)
        panel._timeline_node_count_spins[plan_id].setValue(5)
        panel._timeline_node_count_spins[plan_id].setValue(3)
        plan = panel._selected_profile().timeline_plans[plan_id]

        evidence = panel._timeline_segment_reference_evidence(plan)

        assert evidence.scan_status == "ok"
        assert set(evidence.token_keys) == {"时间节点1-2", "时间节点1-5"}
        assert evidence.occurrence_count == 3
        chooser_calls = []

        def choose_state(**kwargs):
            chooser_calls.append(kwargs)
            return "reserved"

        monkeypatch.setattr(
            panel,
            "_choose_timeline_deletion_number_state",
            choose_state,
        )

        panel._remove_timeline_segment(plan_id)

        assert len(chooser_calls) == 1
        assert chooser_calls[0]["referenced_fields"] == {
            "时间节点1-2",
            "时间节点1-5",
        }
        assert chooser_calls[0]["reference_occurrences"] == 3
        assert chooser_calls[0]["scan_failed"] is False
        assert panel._selected_profile().timeline_plans[plan_id]["number_state"] == "reserved"
    finally:
        panel.close()


def test_docx_scan_failure_requires_an_explicit_number_choice(monkeypatch, tmp_path):
    source = tmp_path / "timeline-scan-failure.docx"
    Document().save(source)
    bridge = PanelBridge()
    bridge.set_current_document_path(str(source))
    _app()
    panel = AssetsPanel(bridge)
    try:
        plan_id = _add_segment(panel)
        chooser_calls = []

        def fail_scan(_path):
            raise OSError("scan failed")

        def choose_state(**kwargs):
            chooser_calls.append(kwargs)
            return "reserved"

        monkeypatch.setattr(
            "src.ui.panels.assets.timeline_presenter.scan_docx_placeholder_inventory",
            fail_scan,
        )
        monkeypatch.setattr(
            panel,
            "_choose_timeline_deletion_number_state",
            choose_state,
        )

        panel._remove_timeline_segment(plan_id)

        assert chooser_calls == [
            {
                "segment_no": 1,
                "referenced_fields": set(),
                "token_copied": False,
                "reference_occurrences": 0,
                "scan_failed": True,
            }
        ]
        assert panel._selected_profile().timeline_plans[plan_id]["number_state"] == "reserved"
    finally:
        panel.close()


def test_copying_whole_segment_reuses_the_smallest_reusable_plan_slot(monkeypatch):
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        first_id = _add_segment(panel)
        panel._duplicate_timeline_segment(first_id)
        second_id = panel._active_timeline_segment_items()[-1][0]
        panel._timeline_node_ratio_edits[(second_id, "node_2")].setText("70")
        panel._duplicate_timeline_segment(second_id)
        third_id = panel._active_timeline_segment_items()[-1][0]
        monkeypatch.setattr(
            "src.ui.panels.assets.timeline_presenter.confirm",
            lambda *args, **kwargs: True,
        )
        panel._remove_timeline_segment(second_id)
        panel._remove_timeline_segment(first_id)

        panel._duplicate_timeline_segment(third_id)

        active_items = panel._active_timeline_segment_items()
        assert [plan_id for plan_id, _plan in active_items] == [first_id, third_id]
        assert [plan["segment_no"] for _plan_id, plan in active_items] == [1, 3]
        reused = panel._selected_profile().timeline_plans[first_id]
        assert reused["nodes"][1]["rule"]["value"] == "0.7"
        assert reused["nodes"][1]["outputs"] == [
            {"field": "时间节点1-2", "format": "yyyy-MM-dd"}
        ]
        assert panel._selected_profile().timeline_plans[second_id]["number_state"] == "reusable"
        assert panel._selected_profile().timeline_plans[third_id]["segment_no"] == 3
        assert panel._timeline_node_token_edits[(third_id, "node_1")].text() == "{{@time:时间节点3-1}}"
    finally:
        panel.close()


def test_empty_draft_segment_reports_locally_but_does_not_block_execution():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        resolution = panel.material_context().resolve_material_fields()

        assert resolution.timeline_issues == ()
        assert not hasattr(panel, "_timeline_segment_issue_labels")
        assert panel._timeline_segment_status_labels[plan_id].text() == "缺 2"
        assert not any(key.startswith("时间节点") for key in resolution.values)
    finally:
        panel.close()


def test_profile_switch_reconciles_outer_segments_and_inner_nodes():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        _add_segment(panel)
        second_plan = default_timeline_segment(4, node_count=2)
        panel._insert_profile_incrementally(
            1,
            EntityProfile(
                profile_id="profile_2",
                profile_name="第 2 份",
                timeline_plans={"segment_4": second_plan},
            ),
        )

        assert tuple(panel._timeline_segments_controller.keys()) == ("segment_4",)
        assert len(panel._timeline_node_controllers["segment_4"]) == 2
        assert panel._timeline_node_token_edits[("segment_4", "node_2")].text() == "{{@time:时间节点4-2}}"
    finally:
        panel.close()


def test_loading_material_context_reconciles_multiple_segments():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        panel._archive_id_edit.setText("timeline-package")
        panel._profile_id_edit.setText("timeline-profile")
        context = panel.material_context()
        context.timeline_plans = {
            "segment_2": default_timeline_segment(2, node_count=2),
            "segment_5": default_timeline_segment(5, node_count=4),
        }

        panel._on_material_context_changed(context)

        assert tuple(panel._timeline_segments_controller.keys()) == (
            "segment_2",
            "segment_5",
        )
        assert len(panel._timeline_node_controllers["segment_5"]) == 4
    finally:
        panel.close()


def test_legacy_plan_keeps_old_output_as_alias_and_adds_canonical_token():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        legacy = default_timeline_plan()
        legacy["start_field"] = "legacy_start"
        legacy["end_field"] = "legacy_end"
        profile = panel._selected_profile()
        profile.fields.update(
            {"legacy_start": "2025-01-01", "legacy_end": "2025-01-11"}
        )
        panel._set_structured_fields(profile.fields)
        profile.timeline_plans = {"primary": legacy}

        panel._refresh_timeline_ui()
        migrated = profile.timeline_plans["primary"]
        values = panel.material_context().resolved_entity_data()

        assert migrated["segment_no"] == 1
        assert migrated["nodes"][0]["outputs"][0]["field"] == "时间节点1-1"
        assert migrated["nodes"][0]["outputs"][1]["field"] == "节点_开始节点"
        assert values["时间节点1-2"] == "2025-01-06"
        assert values["节点_节点 2"] == "2025-01-06"
        panel._refresh_unknown_field_suggestions(["节点_开始节点"])
        legacy_edit = panel._template_field_inputs["节点_开始节点"]
        assert legacy_edit.isReadOnly()
        assert legacy_edit.text() == "2025-01-01"
    finally:
        panel.close()


def test_resizing_migrated_legacy_plan_keeps_old_token_aliases():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        legacy = default_timeline_plan()
        legacy["start_field"] = "legacy_start"
        legacy["end_field"] = "legacy_end"
        profile = panel._selected_profile()
        profile.fields.update(
            {"legacy_start": "2025-01-01", "legacy_end": "2025-01-11"}
        )
        panel._set_structured_fields(profile.fields)
        profile.timeline_plans = {"primary": legacy}
        panel._refresh_timeline_ui()

        panel._timeline_node_count_spins["primary"].setValue(4)
        outputs = profile.timeline_plans["primary"]["nodes"][0]["outputs"]
        values = panel.material_context().resolved_entity_data()

        assert outputs == [
            {"field": "时间节点1-1", "format": "yyyy-MM-dd"},
            {"field": "节点_开始节点", "format": "yyyy-MM-dd"},
        ]
        assert values["节点_开始节点"] == "2025-01-01"
        assert values["时间节点1-4"] == "2025-01-11"
    finally:
        panel.close()


def test_paused_legacy_plan_with_complete_dates_is_reactivated_by_v2_migration():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        legacy = default_timeline_plan()
        legacy["enabled"] = False
        legacy["start_field"] = "legacy_start"
        legacy["end_field"] = "legacy_end"
        profile = panel._selected_profile()
        profile.fields.update(
            {"legacy_start": "2025-01-01", "legacy_end": "2025-01-11"}
        )
        panel._set_structured_fields(profile.fields)
        profile.timeline_plans = {"primary": legacy}

        panel._refresh_timeline_ui()
        values = panel.material_context().resolved_entity_data()

        assert profile.timeline_plans["primary"]["enabled"] is True
        assert values["时间节点1-2"] == "2025-01-06"
        assert panel._timeline_node_result_edits[("primary", "node_2")].text() == "2025-01-06"
    finally:
        panel.close()


def test_document_token_owned_by_segment_is_read_only_and_shows_result():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        _set_segment_dates(panel, plan_id, "2025-10-01", "2025-10-11")
        panel._refresh_unknown_field_suggestions(["时间节点1-1"])

        edit = panel._template_field_inputs["时间节点1-1"]

        assert edit.isReadOnly()
        assert edit.text() == "2025-10-01"
        assert edit.placeholderText() == "由时间计划计算"
    finally:
        panel.close()


def test_ratio_order_error_is_visible_and_segment_outputs_are_atomic():
    _app()
    panel = AssetsPanel(PanelBridge())
    try:
        plan_id = _add_segment(panel)
        panel._timeline_node_count_spins[plan_id].setValue(4)
        _set_segment_dates(panel, plan_id, "2025-01-01", "2025-01-11")
        panel._timeline_node_ratio_edits[(plan_id, "node_2")].setText("90")
        panel._timeline_node_ratio_edits[(plan_id, "node_3")].setText("10")

        resolution = panel.material_context().resolve_material_fields()

        assert any(
            issue.code == "node_ratio_order_invalid"
            for issue in resolution.timeline_issues
        )
        assert "时间节点1-2" not in resolution.values
        assert panel._timeline_segment_status_labels[plan_id].text() == "缺 1"
    finally:
        panel.close()
