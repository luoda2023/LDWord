import base64
import json
from dataclasses import replace
from shutil import copy2
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

from src.config.resolved import ResolvedConfig
from src.config.resolver import resolve_config
from src.config.library import load_scene_from_library, load_template_from_library
from src.config.master_library import get_master
from src.config.material_context import MaterialExecutionContext
from src.config.object_preflight_evidence import build_object_preflight_evidence
from src.config.scene import (
    ContentVisibilityRule,
    DeliveryPreset,
    ExamPaperConfig,
    InputSourceProfile,
    SceneWorkspace,
)
from src.config.template import TemplateConfig
from src.pipeline.runner import Pipeline
from src.services.execution_session import (
    build_execution_session_snapshot,
    cleanup_execution_session_resources,
    execution_session_frozen_master,
)
from src.services.execution_session import resolution as execution_session_resolution
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.exam_question_schema import (
    _preset_hidden_selectors,
    build_exam_delivery_runtime,
    exam_markdown_import_entity_data,
    inspect_exam_question_schema,
    parse_exam_markdown_source,
    render_exam_markdown_preview,
)
from src.shared.engine.fixed_layout_tables import row_height_state


_SAMPLE_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+ip1s"
    "AAAAASUVORK5CYII="
)


def _write_sample_png(tmp_path) -> str:
    path = tmp_path / "question_figure.png"
    path.write_bytes(base64.b64decode(_SAMPLE_PNG))
    return str(path)


def _exam_config(payload, *, delivery_presets=None) -> ResolvedConfig:
    presets = list(delivery_presets or [])
    return ResolvedConfig(
        input_source_profile=InputSourceProfile(
            accepted_formats=["docx", "json"],
            structured_formats=["json"],
            material_schema_id="exam_items_v1",
        ),
        entity_data={"exam_items": json.dumps(payload, ensure_ascii=False)},
        delivery_presets=presets,
        default_delivery_preset_id=(presets[0].preset_id if presets else ""),
    )


def _resolved_exam_master():
    master = get_master("default_exam", mode_id="exam")
    assert master is not None
    return master


def _exam_delivery_presets() -> list[DeliveryPreset]:
    return [
        DeliveryPreset(
            preset_id="student_version",
            label="Student version",
            content_visibility_rules=[
                ContentVisibilityRule(selector="answer", action="remove"),
                ContentVisibilityRule(selector="analysis", action="remove"),
                ContentVisibilityRule(selector="solution", action="remove"),
                ContentVisibilityRule(selector="knowledge_points", action="remove"),
            ],
        ),
        DeliveryPreset(
            preset_id="teacher_version",
            label="Teacher version",
        ),
        DeliveryPreset(
            preset_id="answer_key",
            label="Answer key",
            content_visibility_rules=[
                ContentVisibilityRule(selector="question_only", action="remove"),
                ContentVisibilityRule(selector="analysis", action="remove"),
                ContentVisibilityRule(selector="solution", action="remove"),
                ContentVisibilityRule(selector="knowledge_points", action="remove"),
            ],
        ),
        DeliveryPreset(
            preset_id="analysis_version",
            label="Analysis version",
            content_visibility_rules=[
                ContentVisibilityRule(selector="question_only", action="remove"),
            ],
        ),
        DeliveryPreset(
            preset_id="answer_sheet",
            label="Answer sheet",
            content_visibility_rules=[
                ContentVisibilityRule(selector="question_body", action="remove"),
                ContentVisibilityRule(selector="answer", action="remove"),
                ContentVisibilityRule(selector="analysis", action="remove"),
                ContentVisibilityRule(selector="knowledge_points", action="remove"),
            ],
        ),
    ]


def test_exam_visibility_projection_uses_the_canonical_remove_only_contract():
    assert _preset_hidden_selectors(
        {
            "content_visibility_rules": [
                {
                    "selector_type": "marker_block",
                    "selector": "answer",
                    "action": "remove",
                }
            ]
        }
    ) == ("answer",)
    for retired_action in ("hide", "exclude"):
        assert _preset_hidden_selectors(
            {
                "content_visibility_rules": [
                    {
                        "selector_type": "marker_block",
                        "selector": "answer",
                        "action": retired_action,
                    }
                ]
            }
        ) == ()


def test_parse_exam_markdown_source_imports_prompt_contract():
    markdown = """# 七年级数学期中模拟试卷

> 科目：数学　年级：七年级　考试时间：90 分钟　满分：15 分

## 一、选择题（本大题共 2 小题，每小题 5 分，共 10 分）

1. 下列各数中，最小的是（　　）（5 分）
   A. -3
   B. 0
   C. 2
   D. -1

2. 计算 `-4 + 7` 的结果是（　　）（5 分）
   A. -11
   B. -3
   C. 3
   D. 11

## 二、解答题（本大题共 1 小题，共 5 分）

1. 计算：`(-6) + 9 - (-4) - 7`。（5 分）

   answer_area_kind: free
   answer_lines: 4

## 答案速查

一、选择题
1. A
2. C

二、解答题
1. 0。解析：`(-6) + 9 + 4 - 7 = 0`。
"""

    result = parse_exam_markdown_source(markdown)
    entity_data = exam_markdown_import_entity_data(result)
    payload = result.payload

    assert result.status == "ok"
    assert result.summary.section_count == 2
    assert result.summary.question_count == 3
    assert result.summary.answered_question_count == 3
    assert result.summary.analysis_count == 1
    assert result.summary.declared_total_score == 15
    assert result.summary.computed_total_score == 15
    assert payload["paper_title"] == "七年级数学期中模拟试卷"
    assert payload["subject"] == "数学"
    assert payload["grade"] == "七年级"
    assert payload["duration"] == "90 分钟"
    assert payload["total_score"] == "15 分"
    assert payload["sections"][0]["questions"][0]["options"] == [
        "A. -3",
        "B. 0",
        "C. 2",
        "D. -1",
    ]
    assert payload["sections"][1]["questions"][0]["answer_area_kind"] == "free"
    assert payload["sections"][1]["questions"][0]["answer_lines"] == 4
    assert payload["sections"][1]["questions"][0]["answer"] == "0"
    assert "解析" not in payload["sections"][1]["questions"][0]["answer"]
    assert "exam_items" in entity_data
    assert json.loads(entity_data["exam_items"])["sections"][0]["questions"][1]["answer"] == "C"


def test_parse_exam_markdown_source_reports_missing_answer_block():
    result = parse_exam_markdown_source(
        """# 单元测试

> 科目：数学　满分：5 分

## 一、选择题

1. 1 + 1 = （　　）（5 分）
   A. 1
   B. 2
"""
    )

    issue_kinds = {issue.kind for issue in result.issues}

    assert result.status == "error"
    assert "missing_answer_block" in issue_kinds
    assert "missing_answer" in issue_kinds


def test_exam_question_schema_accepts_single_structured_source():
    payload = {
        "paper_title": "期末测试",
        "subject": "数学",
        "grade": "八年级",
        "duration": "90 分钟",
        "total_score": 15,
        "sections": [
            {
                "title": "选择题",
                "type": "single_choice",
                "questions": [
                    {
                        "stem": "1 + 1 = ?",
                        "options": ["A.1", "B.2"],
                        "answer": "B",
                        "analysis": "基础加法。",
                        "score": 5,
                        "knowledge_points": ["整数加法"],
                    },
                    {
                        "stem": "2 + 3 = ?",
                        "options": ["A.5", "B.6"],
                        "answer": "A",
                        "score": 10,
                    },
                ],
            }
        ],
    }

    result = inspect_exam_question_schema(_exam_config(payload))

    assert result.status == "ok"
    assert result.source_key == "entity_data.exam_items"
    assert result.summary.section_count == 1
    assert result.summary.question_count == 2
    assert result.summary.answered_question_count == 2
    assert result.summary.computed_total_score == 15
    assert "single_choice" in result.summary.question_types


def test_exam_question_schema_reports_missing_answer_and_score_mismatch():
    payload = {
        "paper_title": "期末测试",
        "total_score": 20,
        "sections": [
            {
                "title": "选择题",
                "type": "single_choice",
                "questions": [
                    {
                        "stem": "1 + 1 = ?",
                        "options": ["A.1", "B.2"],
                        "score": 5,
                    },
                    {
                        "answer": "A",
                        "score": 5,
                    },
                ],
            }
        ],
    }

    result = inspect_exam_question_schema(_exam_config(payload))
    issue_kinds = {issue.kind for issue in result.issues}

    assert result.status == "error"
    assert result.error_count >= 3
    assert "missing_answer" in issue_kinds
    assert "missing_stem" in issue_kinds
    assert "total_score_mismatch" in issue_kinds
    assert result.manual_confirmation_required is True


def test_exam_delivery_runtime_renders_markdown_preview_and_word_versions(tmp_path):
    figure_path = _write_sample_png(tmp_path)
    payload = {
        "paper_title": "期末测试",
        "subject": "数学",
        "total_score": 5,
        "sections": [
            {
                "title": "选择题",
                "type": "single_choice",
                "questions": [
                    {
                        "stem": "1 + 1 = ?",
                        "options": ["A.1", "B.2"],
                        "answer": "B",
                        "analysis": "基础加法。",
                        "score": 5,
                        "knowledge_points": ["整数加法"],
                        "figure": {
                            "source": "file",
                            "path": figure_path,
                            "alt": "加法示意图",
                        },
                    }
                ],
            }
        ],
    }
    config = _exam_config(payload, delivery_presets=_exam_delivery_presets())

    markdown = render_exam_markdown_preview(payload)
    runtime = build_exam_delivery_runtime(
        config,
        output_dir=tmp_path,
        source_stem="exam_source",
    )
    versions = {version.preset_id: version for version in runtime.rendered_versions}

    assert "# 期末测试" in markdown
    assert "Answer: B" in markdown
    assert runtime.status == "ok"
    assert runtime.version_count == 5
    assert runtime.markdown_preview_path.endswith("exam_source_preview.md")
    assert {
        "student_version",
        "teacher_version",
        "answer_key",
        "analysis_version",
        "answer_sheet",
    } == set(versions)

    student_doc = Document(versions["student_version"].docx_path)
    student_text = "\n".join(paragraph.text for paragraph in student_doc.paragraphs)
    assert "1 + 1 = ?" in student_text
    assert "Answer: B" not in student_text
    assert "Analysis:" not in student_text
    assert len(student_doc.inline_shapes) == 1
    assert versions["student_version"].visible_answer_count == 0
    assert versions["student_version"].question_asset_count == 1
    assert versions["student_version"].rendered_question_asset_count == 1
    assert versions["student_version"].missing_question_asset_count == 0
    assert versions["student_version"].question_asset_alt_text_count == 1
    assert versions["student_version"].rendered_question_asset_alt_text_count == 1
    assert versions["student_version"].missing_question_asset_alt_text_count == 0
    student_doc_pr = _first_doc_pr_attrs(versions["student_version"].docx_path)
    assert student_doc_pr.get("descr") == "加法示意图"
    assert student_doc_pr.get("title") == "file"

    teacher_doc = Document(versions["teacher_version"].docx_path)
    teacher_text = "\n".join(paragraph.text for paragraph in teacher_doc.paragraphs)
    assert "1 + 1 = ?" in teacher_text
    assert "Answer: B" in teacher_text
    assert "Analysis: 基础加法。" in teacher_text
    assert len(teacher_doc.inline_shapes) == 1
    assert versions["teacher_version"].rendered_question_asset_count == 1
    assert versions["teacher_version"].rendered_question_asset_alt_text_count == 1

    answer_doc = Document(versions["answer_key"].docx_path)
    answer_text = "\n".join(paragraph.text for paragraph in answer_doc.paragraphs)
    assert "1 + 1 = ?" not in answer_text
    assert "Answer: B" in answer_text
    assert len(answer_doc.inline_shapes) == 0
    assert versions["answer_key"].visible_question_count == 0
    assert versions["answer_key"].question_asset_count == 0
    assert versions["answer_key"].rendered_question_asset_alt_text_count == 0

    sheet_doc = Document(versions["answer_sheet"].docx_path)
    sheet_text = "\n".join(paragraph.text for paragraph in sheet_doc.paragraphs)
    assert "1 + 1 = ?" not in sheet_text
    assert "Answer: B" not in sheet_text
    assert len(sheet_doc.inline_shapes) == 0
    assert len(sheet_doc.tables) == 1
    assert sheet_doc.tables[0].cell(1, 2).text == "____________________________"
    assert row_height_state(sheet_doc.tables[0].rows[1]).height_twips == 440
    assert row_height_state(sheet_doc.tables[0].rows[1]).rule == "exact"
    assert versions["answer_sheet"].fixed_layout_kind == "answer_sheet"
    assert versions["answer_sheet"].fixed_layout_row_count == 2
    assert versions["answer_sheet"].fixed_layout_column_count == 4
    assert versions["answer_sheet"].fixed_layout_row_height_twips == 440


def test_scene_exam_paper_config_renders_structured_source_into_master_docx(tmp_path):
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        category="exam_paper",
        master_id="default_exam",
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown", "docx"],
            structured_formats=["json"],
            material_schema_id="exam_items_v1",
        ),
        exam_paper=ExamPaperConfig(
            question_structure_mode="markdown_headings",
            answer_policy="student_only",
            runtime_fields=["title", "subject", "grade", "duration", "total_score"],
        ),
    )
    payload = {
        "paper_title": "题源里的旧标题",
        "sections": [
            {
                "title": "一、填空题",
                "questions": [
                    {
                        "stem": "7 个 0.9 是（    ）。",
                        "answer": "6.3",
                        "score": 2,
                    }
                ],
            }
        ],
    }
    config = resolve_config(
        TemplateConfig(),
        scene,
        entity_data={
            "exam_items": json.dumps(payload, ensure_ascii=False),
            "title": "工作台填写的期中试卷",
            "subject": "数学",
            "grade": "五年级",
            "duration": "90 分钟",
            "total_score": "2",
        },
    )

    validation = inspect_exam_question_schema(config)
    runtime = build_exam_delivery_runtime(
        config,
        output_dir=tmp_path,
        source_stem="scene_exam",
        validation=validation,
        master=_resolved_exam_master(),
    )

    assert config.exam_paper is not None
    assert runtime.status == "ok"
    assert runtime.version_count == 1
    version = runtime.rendered_versions[0]
    assert version.preset_id == "student"
    assert version.hidden_selectors == ("answer", "analysis", "solution", "knowledge_points")
    text = "\n".join(paragraph.text for paragraph in Document(version.docx_path).paragraphs)
    assert "工作台填写的期中试卷" in text
    assert "题源里的旧标题" not in text
    assert "学生卷" in text
    assert "答案版" not in text
    assert "答案速查" not in text
    assert "1. 7 个 0.9 是（    ）。" in text


def test_scene_exam_paper_config_renders_student_and_answer_key_files(tmp_path):
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        category="exam_paper",
        master_id="default_exam",
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown", "docx"],
            structured_formats=["json"],
            material_schema_id="exam_items_v1",
        ),
        exam_paper=ExamPaperConfig(
            question_structure_mode="markdown_headings",
            answer_policy="student_plus_answer",
            runtime_fields=["title", "subject", "grade", "duration", "total_score"],
        ),
    )
    payload = {
        "paper_title": "题源里的旧标题",
        "sections": [
            {
                "title": "一、填空题",
                "questions": [
                    {
                        "stem": "7 个 0.9 是（    ）。",
                        "answer": "6.3",
                        "score": 2,
                    }
                ],
            }
        ],
    }
    config = resolve_config(
        TemplateConfig(),
        scene,
        entity_data={
            "exam_items": json.dumps(payload, ensure_ascii=False),
            "title": "工作台填写的期中试卷",
            "subject": "数学",
            "grade": "五年级",
            "duration": "90 分钟",
            "total_score": "2",
        },
    )

    validation = inspect_exam_question_schema(config)
    runtime = build_exam_delivery_runtime(
        config,
        output_dir=tmp_path,
        source_stem="scene_exam",
        validation=validation,
        master=_resolved_exam_master(),
    )
    versions = {version.preset_id: version for version in runtime.rendered_versions}

    assert runtime.status == "ok"
    assert runtime.version_count == 2
    assert runtime.master_evidence is not None
    assert runtime.master_evidence.mode_id == "exam"
    assert runtime.master_evidence.master_id == "default_exam"
    assert runtime.master_evidence.master_source_type == "builtin"
    assert runtime.master_evidence.placeholder_contract_status == "ok"
    assert runtime.master_evidence.master_docx_path.endswith(
        "config_library\\masters\\exam\\builtin\\default_exam_v20.docx"
    ) or runtime.master_evidence.master_docx_path.endswith(
        "config_library/masters/exam/builtin/default_exam_v20.docx"
    )
    assert runtime.master_evidence.manifest_path.endswith(
        "config_library\\masters\\exam\\builtin\\default_exam.master.json"
    ) or runtime.master_evidence.manifest_path.endswith(
        "config_library/masters/exam/builtin/default_exam.master.json"
    )
    assert set(versions) == {"student", "answer_key"}
    assert versions["student"].label == "学生卷"
    assert versions["answer_key"].label == "答案速查"
    assert versions["student"].docx_path.endswith("scene_exam_学生卷.docx")
    assert versions["answer_key"].docx_path.endswith("scene_exam_答案速查.docx")
    assert versions["student"].visible_question_count == 1
    assert versions["student"].visible_answer_count == 0
    assert versions["answer_key"].visible_question_count == 0
    assert versions["answer_key"].visible_answer_count == 1

    student_text = "\n".join(paragraph.text for paragraph in Document(versions["student"].docx_path).paragraphs)
    answer_text = "\n".join(paragraph.text for paragraph in Document(versions["answer_key"].docx_path).paragraphs)
    assert "工作台填写的期中试卷" in student_text
    assert "学生卷" in student_text
    assert "答案速查" not in student_text
    assert "7 个 0.9 是" in student_text
    assert "答案速查" in answer_text
    assert "1. 6.3" in answer_text
    assert "7 个 0.9 是" not in answer_text


def test_scene_exam_paper_config_can_render_answer_key_only(tmp_path):
    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        category="exam_paper",
        master_id="default_exam",
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown", "docx"],
            structured_formats=["json"],
            material_schema_id="exam_items_v1",
        ),
        exam_paper=ExamPaperConfig(
            answer_policy="answer_only",
        ),
    )
    payload = {
        "paper_title": "题源里的旧标题",
        "sections": [
            {
                "title": "一、填空题",
                "questions": [
                    {
                        "stem": "7 个 0.9 是（    ）。",
                        "answer": "6.3",
                        "score": 2,
                    }
                ],
            }
        ],
    }
    config = resolve_config(
        TemplateConfig(),
        scene,
        entity_data={
            "exam_items": json.dumps(payload, ensure_ascii=False),
            "title": "工作台填写的期中试卷",
        },
    )

    validation = inspect_exam_question_schema(config)
    runtime = build_exam_delivery_runtime(
        config,
        output_dir=tmp_path,
        source_stem="scene_exam",
        validation=validation,
        master=_resolved_exam_master(),
    )
    versions = {version.preset_id: version for version in runtime.rendered_versions}

    assert runtime.status == "ok"
    assert runtime.version_count == 1
    assert set(versions) == {"answer_key"}
    assert versions["answer_key"].docx_path.endswith("scene_exam_答案速查.docx")
    assert not (tmp_path / "scene_exam_学生卷.docx").exists()


def test_pipeline_without_execution_master_blocks_exam_runtime(tmp_path):
    source = tmp_path / "scene_exam.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        category="exam_paper",
        master_id="default_exam",
        input_source_profile=InputSourceProfile(
            accepted_formats=["markdown", "docx"],
            structured_formats=["json"],
            material_schema_id="exam_items_v1",
        ),
        exam_paper=ExamPaperConfig(
            answer_policy="student_plus_answer",
        ),
    )
    payload = {
        "paper_title": "题源标题",
        "sections": [
            {
                "title": "一、填空题",
                "questions": [
                    {"stem": "7 个 0.9 是（    ）。", "answer": "6.3", "score": 2}
                ],
            }
        ],
    }
    config = resolve_config(
        TemplateConfig(),
        scene,
        entity_data={"exam_items": json.dumps(payload, ensure_ascii=False)},
    )

    result = Pipeline([], config, output_dir=str(tmp_path)).execute(str(source))

    assert result.success is False
    assert result.output_paths == {}
    assert "execution_master_required" in str(result.error)
    assert result.context.exam_delivery_runtime.status == "blocked"
    assert result.context.exam_delivery_runtime.master_evidence is None


def test_exam_runtime_consumes_frozen_master_after_live_source_is_deleted(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "exam_source.docx"
    Document().save(source)
    scene = load_scene_from_library("exam", mode_id="exam")
    template = load_template_from_library("default", mode_id="exam")

    builtin_master = _resolved_exam_master()
    live_master_path = tmp_path / "live_exam_master.docx"
    copy2(builtin_master.docx_path, live_master_path)
    live_master = replace(
        builtin_master,
        source_type="user",
        readonly=False,
        docx_path=live_master_path,
    )
    monkeypatch.setattr(
        execution_session_resolution,
        "get_master",
        lambda *_args, **_kwargs: live_master,
    )

    exam_payload = {
        "paper_title": "Frozen master exam",
        "total_score": 2,
        "sections": [
            {
                "title": "Questions",
                "questions": [
                    {
                        "stem": "1 + 1 = ?",
                        "answer": "2",
                        "score": 2,
                    }
                ],
            }
        ],
    }
    material_context = MaterialExecutionContext(
        mode_id="exam",
        entity_data={
            "exam_items": json.dumps(exam_payload, ensure_ascii=False),
        },
    )
    preflight = build_object_preflight_evidence(scene, source)
    snapshot = build_execution_session_snapshot(
        mode_id="exam",
        scene=scene,
        template=template,
        material_context=material_context,
        input_path=source,
        output_root=tmp_path / "runs",
        plan_id="exam",
        template_id="default",
        object_preflight_confirmation_revision=preflight.source_revision,
        object_preflight_confirmation_digest=preflight.evidence_digest,
    )
    try:
        frozen_master = execution_session_frozen_master(snapshot)
        assert frozen_master is not None
        assert frozen_master.execution_frozen is True
        assert frozen_master.docx_path != live_master_path
        assert frozen_master.docx_path.is_file()

        live_master_path.unlink()

        def _reject_live_master_lookup(*_args, **_kwargs):
            raise AssertionError("live exam master lookup must not run after snapshot")

        monkeypatch.setattr(
            "src.shared.engine.exam_paper_style.ensure_current_exam_blank_master_docx",
            _reject_live_master_lookup,
        )
        config = resolve_config(
            template,
            scene,
            **material_context.to_resolve_kwargs(),
        )
        runtime = build_exam_delivery_runtime(
            config,
            output_dir=tmp_path / "output",
            source_stem="frozen_exam",
            master=frozen_master,
        )

        assert runtime.status == "ok"
        assert runtime.rendered_versions
        assert all(
            Path(version.docx_path).is_file()
            for version in runtime.rendered_versions
        )
        assert runtime.master_evidence is not None
        assert Path(runtime.master_evidence.master_docx_path) == frozen_master.docx_path
        assert not live_master_path.exists()

        runner_payload = WorkbenchProductionRunner(
            doc_path=str(source),
            template=template,
            scene=scene,
            material_context=material_context,
            output_dir=Path(snapshot.output_namespace),
            execution_session=snapshot,
        ).run(lambda *_args: None, lambda: False)
        runner_master = runner_payload["exam_delivery_runtime"]["master_evidence"]

        assert runner_payload["status"] == "success"
        assert runner_master["master_id"] == frozen_master.master_id
        assert Path(runner_master["master_docx_path"]) == frozen_master.docx_path
        assert runner_payload["output_paths"]
    finally:
        cleanup_execution_session_resources(snapshot)


def _first_doc_pr_attrs(docx_path: str) -> dict[str, str]:
    with zipfile.ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    ns = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
    doc_pr = root.find(".//wp:docPr", ns)
    assert doc_pr is not None
    return dict(doc_pr.attrib)


def test_exam_delivery_runtime_reports_missing_question_asset_alt_text(tmp_path):
    figure_path = _write_sample_png(tmp_path)
    payload = {
        "paper_title": "期末测试",
        "sections": [
            {
                "title": "选择题",
                "questions": [
                    {
                        "stem": "观察图片，回答问题。",
                        "answer": "B",
                        "score": 5,
                        "figure": {"source": "file", "path": figure_path},
                    }
                ],
            }
        ],
    }
    runtime = build_exam_delivery_runtime(
        _exam_config(payload, delivery_presets=_exam_delivery_presets()),
        output_dir=tmp_path,
        source_stem="exam_without_alt",
    )
    versions = {version.preset_id: version for version in runtime.rendered_versions}

    assert versions["student_version"].question_asset_count == 1
    assert versions["student_version"].rendered_question_asset_count == 1
    assert versions["student_version"].question_asset_alt_text_count == 0
    assert versions["student_version"].rendered_question_asset_alt_text_count == 0
    assert versions["student_version"].missing_question_asset_alt_text_count == 1


def test_pipeline_records_exam_question_schema_result_and_reports(tmp_path):
    source = tmp_path / "exam.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)

    payload = {
        "paper_title": "期末测试",
        "total_score": 10,
        "sections": [
            {
                "title": "填空题",
                "type": "short_answer",
                "questions": [
                    {
                        "stem": "水的化学式是？",
                        "score": 10,
                    }
                ],
            }
        ],
    }
    config = _exam_config(payload)
    result = Pipeline([], config, output_dir=str(tmp_path)).execute(str(source))

    assert result.success is False
    assert result.output_paths == {}
    assert "Terminal assembler 'exam' blocked delivery" in str(result.error)
    assert result.context is not None
    assert result.context.exam_question_schema.status == "error"
    records = result.tracker.get_by_module("exam_question_schema")
    assert records
    assert records[0].change_type == "schema_validation_warning"
    assert "questions=1" in records[0].after

    report_json = tmp_path / "exam_changes.json"
    report_md = tmp_path / "exam_changes.md"
    write_json_report(
        result,
        input_path=source,
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )

    data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert data["exam_question_schema"]["status"] == "error"
    assert data["exam_question_schema"]["summary"]["question_count"] == 1
    assert any(
        issue["kind"] == "missing_answer"
        for issue in data["exam_question_schema"]["issues"]
    )
    assert "## 试卷题源结构校验" in markdown
    assert "missing_answer" in markdown


def test_pipeline_records_exam_delivery_runtime_and_reports(tmp_path):
    source = tmp_path / "exam.docx"
    doc = Document()
    doc.add_paragraph("Exam source")
    doc.save(source)
    figure_path = _write_sample_png(tmp_path)

    payload = {
        "paper_title": "期末测试",
        "total_score": 5,
        "sections": [
            {
                "title": "选择题",
                "type": "single_choice",
                "questions": [
                    {
                        "stem": "1 + 1 = ?",
                        "options": ["A.1", "B.2"],
                        "answer": "B",
                        "analysis": "基础加法。",
                        "score": 5,
                        "figure": {"source": "file", "path": figure_path},
                    }
                ],
            }
        ],
    }
    config = _exam_config(payload, delivery_presets=_exam_delivery_presets())
    result = Pipeline([], config, output_dir=str(tmp_path)).execute(str(source))

    assert result.success is True
    assert result.context is not None
    assert result.context.exam_delivery_runtime.status == "ok"
    assert result.context.exam_delivery_runtime.version_count == 5
    records = result.tracker.get_by_module("exam_delivery_runtime")
    assert records
    assert "versions=5" in records[0].after

    report_json = tmp_path / "exam_runtime_changes.json"
    report_md = tmp_path / "exam_runtime_changes.md"
    write_json_report(
        result,
        input_path=source,
        output_path=None,
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )

    data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert data["exam_delivery_runtime"]["status"] == "ok"
    assert data["exam_delivery_runtime"]["version_count"] == 5
    assert data["exam_delivery_runtime"]["markdown_preview_path"].endswith(
        "exam_preview.md"
    )
    assert len(data["exam_delivery_runtime"]["rendered_versions"]) == 5
    assert data["exam_delivery_runtime"]["master_evidence"] == {}
    versions = {
        version["preset_id"]: version
        for version in data["exam_delivery_runtime"]["rendered_versions"]
    }
    assert versions["answer_sheet"]["fixed_layout_kind"] == "answer_sheet"
    assert versions["answer_sheet"]["fixed_layout_row_height_twips"] == 440
    assert versions["student_version"]["question_asset_count"] == 1
    assert versions["student_version"]["rendered_question_asset_count"] == 1
    assert versions["student_version"]["question_asset_alt_text_count"] == 0
    assert versions["student_version"]["rendered_question_asset_alt_text_count"] == 0
    assert versions["student_version"]["missing_question_asset_alt_text_count"] == 1
    assert versions["answer_key"]["question_asset_count"] == 0
    assert "## 试卷多版本运行时渲染" in markdown
    assert "student_version" in markdown
    assert "fixed_layout=answer_sheet" in markdown
    assert "assets=1/1" in markdown
    assert "altText=0/1" in markdown
    assert "missing_altText=1" in markdown
