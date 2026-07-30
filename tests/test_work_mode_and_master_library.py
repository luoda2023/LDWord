import json

from docx import Document
import pytest

import src.config.master_library as master_library
import src.config.library as config_library
from src.config.builtin_scenes import create_builtin_scene
from src.services.execution_session import (
    build_execution_session_snapshot,
    cleanup_execution_session_resources,
)
from src.config.library import load_scene_from_library
from src.config.master_library import (
    MASTER_MANIFEST_ROOT,
    MasterManifestError,
    create_official_master_copy,
    default_master,
    get_master,
    list_masters,
)
from src.config.master_preflight import (
    check_master_preflight,
    check_placeholder_contract,
)
from src.config.official_document_profiles import (
    get_official_document_assembly_contract,
    get_official_document_profile,
    list_official_document_assembly_contracts,
    list_official_document_profiles,
    resolve_official_batch_document_type_id,
)
from src.config.material_schema_registry import get_material_schema
from src.config.material_context import MaterialExecutionContext
from src.config.object_preflight_evidence import build_object_preflight_evidence
from src.config.scene import ExamBlankStyleConfig, ExamPaperConfig
from src.config.work_mode import (
    default_work_mode,
    get_work_mode,
    list_work_modes,
    work_mode_for_scene_id,
)
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.shared.engine.exam_paper_style import (
    BUILTIN_EXAM_MASTER_DIR,
)
from src.shared.engine.official_document_master import (
    OFFICIAL_MASTER_VERSION,
    write_official_gbt_master_docx,
)
from src.shared.engine.official_document_assembly import assemble_official_document_docx
from src.ui.bridge import PanelBridge
from src.ui.adapters.config_selector_models import master_selector_options
from src.ui.panel_registry import PANEL_SPECS


def test_work_mode_registry_lists_core_modes_without_replacing_scene_ids():
    modes = list_work_modes()
    mode_ids = [mode.mode_id for mode in modes]

    assert mode_ids[:5] == ["custom", "exam", "thesis", "bidding", "official"]
    assert {"technical", "report"}.issubset(mode_ids)

    exam = get_work_mode("试卷版")
    assert exam is not None
    assert exam.mode_id == "exam"
    assert exam.default_scene_id == "exam"
    assert exam.default_template_id == "default"
    assert exam.default_master_id == "default_exam"
    assert exam.master_domain == "exam_blank_master"
    assert exam.material_schema_ids == ("exam_items_v1",)
    assert "masters" not in exam.visible_panels

    assert default_work_mode().mode_id == "custom"
    assert work_mode_for_scene_id("official").mode_id == "official"
    assert work_mode_for_scene_id("exam_quiz").mode_id == "exam"
    assert work_mode_for_scene_id("exam_term").mode_id == "exam"
    assert work_mode_for_scene_id("unknown_scene").mode_id == "custom"


def test_master_library_remains_internal_without_a_sidebar_management_page():
    assert "masters" not in {spec.id for spec in PANEL_SPECS}


def test_sidebar_has_no_unimplemented_pipeline_page():
    assert "pipeline" not in {spec.id for spec in PANEL_SPECS}


def test_panel_bridge_tracks_current_work_mode_without_touching_scene_template():
    bridge = PanelBridge()
    scene = SceneWorkspace(scene_id="custom", template_id="default")
    template = TemplateConfig(name="Default Template")
    seen_modes = []
    seen_scenes = []
    seen_templates = []
    bridge.work_mode_changed.connect(seen_modes.append)
    bridge.scene_changed.connect(seen_scenes.append)
    bridge.template_changed.connect(seen_templates.append)

    bridge.set_current_scene(scene, config_id="custom", emit_signal=False)
    bridge.set_current_template(template, config_id="default", emit_signal=False)

    assert bridge.current_work_mode_id() == "custom"
    assert bridge.current_scene_id() == "custom"
    assert bridge.current_template_id() == "default"

    bridge.set_current_work_mode("exam")

    assert bridge.current_work_mode_id() == "exam"
    assert bridge.current_work_mode().label == "试卷版"
    assert [mode.mode_id for mode in seen_modes] == ["exam"]
    assert bridge.current_scene_id() == "custom"
    assert bridge.current_template_id() == "default"
    assert seen_scenes == []
    assert seen_templates == []

    bridge.set_current_work_mode("missing-mode", emit_signal=False)

    assert bridge.current_work_mode_id() == "custom"
    assert [mode.mode_id for mode in seen_modes] == ["exam"]


def test_master_library_wraps_mode_scoped_exam_builtin_master():
    masters = list_masters("exam")
    default = default_master("exam")

    assert default is not None
    assert default in masters
    assert [master.master_id for master in masters] == ["default_exam"]
    assert {master.label for master in masters} == {"A4 标准卷面"}
    assert default.master_id == "default_exam"
    assert default.qualified_id == "exam/default_exam"
    assert default.source_type == "builtin"
    assert default.readonly is True
    assert default.docx_path == BUILTIN_EXAM_MASTER_DIR / "default_exam_v20.docx"
    assert default.docx_path.exists()
    assert default.manifest_path == (
        MASTER_MANIFEST_ROOT / "exam" / "builtin" / "default_exam.master.json"
    )
    assert default.template_config_id == "default"
    assert default.placeholder_contract.required == ("af_title", "af_questions")
    assert "af_answer_area" in default.placeholder_contract.optional
    assert "title" in default.placeholder_contract.runtime_fields
    assert "legacy_exam_questions_marker" in default.placeholder_contract.legacy_aliases

    assert get_master("default_exam") is None
    assert get_master("exam/default_exam").qualified_id == "exam/default_exam"
    assert get_master("exam/default_exam").master_id == "default_exam"


def test_master_library_reads_exam_manifest_as_the_canonical_builtin_source():
    default = default_master("exam")

    assert default is not None
    assert default.manifest_path is not None
    assert default.manifest_path.name == "default_exam.master.json"
    assert default.label == "A4 标准卷面"
    assert default.summary.startswith("A4 标准卷面")
    assert default.master_version == "exam-master-v20-free-answer-area-2026-07-06"
    assert default.placeholder_contract.generated_fields == ("version",)
    assert "does not judge exam content quality" in default.boundaries


def test_exam_builtin_library_exposes_one_manifest_per_real_docx():
    masters = {
        master.master_id: master
        for master in list_masters("exam")
        if master.source_type == "builtin"
    }

    assert set(masters) == {"default_exam"}
    assert masters["default_exam"].manifest_path == (
        MASTER_MANIFEST_ROOT / "exam" / "builtin" / "default_exam.master.json"
    )
    assert masters["default_exam"].docx_path == (
        BUILTIN_EXAM_MASTER_DIR / "default_exam_v20.docx"
    )


def test_exam_builtin_manifest_cannot_redirect_canonical_docx(
    monkeypatch,
    tmp_path,
):
    manifest_dir = tmp_path / "exam" / "builtin"
    manifest_dir.mkdir(parents=True)
    payload = json.loads(
        (
            MASTER_MANIFEST_ROOT
            / "exam"
            / "builtin"
            / "default_exam.master.json"
        ).read_text(encoding="utf-8")
    )
    payload["docx_path"] = "outside/attempted_override.docx"
    (manifest_dir / "default_exam.master.json").write_text(
        json.dumps(payload),
        encoding="utf-8",
    )
    monkeypatch.setattr(master_library, "MASTER_MANIFEST_ROOT", tmp_path)

    with pytest.raises(MasterManifestError, match="master_manifest_asset_invalid"):
        master_library.default_master("exam")


def test_exam_builtin_manifest_is_required_not_an_optional_overlay(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(master_library, "MASTER_MANIFEST_ROOT", tmp_path)

    with pytest.raises(MasterManifestError, match="master_manifest_missing"):
        master_library.default_master("exam")


def test_exam_builtin_manifest_rejects_corrupt_json_without_fallback(
    monkeypatch,
    tmp_path,
):
    manifest_dir = tmp_path / "exam" / "builtin"
    manifest_dir.mkdir(parents=True)
    (manifest_dir / "default_exam.master.json").write_text(
        "{not-json",
        encoding="utf-8",
    )
    monkeypatch.setattr(master_library, "MASTER_MANIFEST_ROOT", tmp_path)

    with pytest.raises(MasterManifestError, match="master_manifest_unreadable"):
        master_library.default_master("exam")


def test_exam_builtin_manifest_rejects_missing_current_schema_field(
    monkeypatch,
    tmp_path,
):
    manifest_dir = tmp_path / "exam" / "builtin"
    manifest_dir.mkdir(parents=True)
    payload = json.loads(
        (
            MASTER_MANIFEST_ROOT
            / "exam"
            / "builtin"
            / "default_exam.master.json"
        ).read_text(encoding="utf-8")
    )
    del payload["label"]
    (manifest_dir / "default_exam.master.json").write_text(
        json.dumps(payload),
        encoding="utf-8",
    )
    monkeypatch.setattr(master_library, "MASTER_MANIFEST_ROOT", tmp_path)

    with pytest.raises(MasterManifestError, match=r"label must be non-empty text"):
        master_library.default_master("exam")


def test_exam_builtin_manifest_rejects_stale_placeholder_index(
    monkeypatch,
    tmp_path,
):
    manifest_dir = tmp_path / "exam" / "builtin"
    manifest_dir.mkdir(parents=True)
    payload = json.loads(
        (
            MASTER_MANIFEST_ROOT
            / "exam"
            / "builtin"
            / "default_exam.master.json"
        ).read_text(encoding="utf-8")
    )
    payload["placeholder_index"]["sha256"] = "0" * 64
    (manifest_dir / "default_exam.master.json").write_text(
        json.dumps(payload),
        encoding="utf-8",
    )
    monkeypatch.setattr(master_library, "MASTER_MANIFEST_ROOT", tmp_path)

    with pytest.raises(MasterManifestError, match="master_manifest_index_stale"):
        master_library.default_master("exam")


def test_master_library_discovers_manual_exam_user_masters_without_mutating_input(
    tmp_path,
):
    user_dir = tmp_path / "user_masters"
    user_dir.mkdir()
    manual_path = user_dir / "AI生成期中母版.docx"
    doc = Document()
    doc.add_paragraph("{{af_title}}")
    doc.add_paragraph("科目：{{af_subject}}")
    doc.add_paragraph("{{af_questions}}")
    doc.add_paragraph("{{af_answer_area}}")
    doc.save(manual_path)

    invalid_path = user_dir / "普通试卷.docx"
    invalid = Document()
    invalid.add_paragraph("这只是普通试卷，没有 Alavette 母版占位符。")
    invalid.save(invalid_path)

    config = ExamPaperConfig()
    masters = list_masters("exam", exam_config=config, user_master_dir=user_dir)
    discovered = [master for master in masters if master.source_type == "discovered"]

    assert config.custom_blank_styles == []
    assert len(discovered) == 1
    assert discovered[0].master_id.startswith("user_file_")
    assert discovered[0].label == "AI生成期中母版"
    assert discovered[0].docx_path == manual_path
    assert discovered[0].readonly is False
    assert discovered[0].base_master_id == "default_exam"
    assert all(master.docx_path != invalid_path for master in masters)


def test_exam_user_master_ids_are_normalized_away_from_builtin_ids(tmp_path):
    config = ExamPaperConfig(
        custom_blank_styles=[
            ExamBlankStyleConfig(
                style_id="default_exam",
                label="User Default Exam Master",
                base_style_id="default_exam",
            )
        ],
    )

    assert [style.style_id for style in config.custom_blank_styles] == [
        "user_default_exam_copy"
    ]

    masters = list_masters("exam", exam_config=config, user_master_dir=tmp_path)
    options = master_selector_options(
        "exam",
        exam_config=config,
        user_master_dir=tmp_path,
    )

    assert [master.master_id for master in masters].count("default_exam") == 1
    assert [master.master_id for master in masters].count("user_default_exam_copy") == 1
    assert [option.value for option in options].count("default_exam") == 1
    assert [option.value for option in options].count("user_default_exam_copy") == 1
    user_master = next(
        master for master in masters if master.master_id == "user_default_exam_copy"
    )
    user_option = next(
        option for option in options if option.value == "user_default_exam_copy"
    )
    assert user_master.source_type == "user"
    assert user_master.label == "User Default Exam Master"
    assert user_option.source_type == "user"
    assert user_option.label == "User Default Exam Master"

    reloaded = ExamPaperConfig(
        custom_blank_styles=config.custom_blank_styles,
    )
    assert [style.style_id for style in reloaded.custom_blank_styles] == [
        "user_default_exam_copy"
    ]


def test_master_library_lists_official_builtin_master_asset():
    official = default_master("official")

    assert official is not None
    assert official.master_id == "official_gbt_standard"
    assert official.qualified_id == "official/official_gbt_standard"
    assert official.source_type == "builtin"
    assert official.template_config_id == "official_gbt"
    assert official.manifest_path == (
        MASTER_MANIFEST_ROOT
        / "official"
        / "builtin"
        / "official_gbt_standard.master.json"
    )
    assert official.docx_path.name == "official_gbt_standard.docx"
    assert official.docx_path.exists() is True
    assert official.label == "GB/T 9704 通用红头公文版式"
    assert official.status == "内置，已接入文种版式路由和 runtime 公文装配"
    assert "尚未接入" not in official.status
    assert official.master_version == OFFICIAL_MASTER_VERSION
    assert official.placeholder_contract.required == (
        "official_title",
        "official_body",
    )
    assert "official_document_no" in official.placeholder_contract.optional
    assert "notice" in official.supported_assembly_types
    assert "minutes" not in official.supported_assembly_types
    assert "does not verify seal legality" in official.boundaries


def test_official_master_docx_generator_writes_required_placeholders(tmp_path):
    generated = write_official_gbt_master_docx(tmp_path / "official_master.docx")
    doc = Document(generated)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs) + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )

    assert generated.is_file()
    assert len(doc.sections) == 1
    first_section = doc.sections[0]
    assert round(first_section.page_width.cm, 1) == 21.0
    assert round(first_section.page_height.cm, 1) == 29.7
    assert round(first_section.top_margin.cm, 1) == 3.7
    assert round(first_section.left_margin.cm, 1) == 2.8
    assert "{{@text:official_title}}" in text
    assert "{{@text:official_body}}" in text
    assert "{{@text:official_document_no}}" in text
    assert "{{@text:official_meeting_attendees}}" not in text
    assert "版式占位符说明" not in text
    assert "母版占位符说明" not in text
    assert doc.core_properties.comments == OFFICIAL_MASTER_VERSION


def test_master_preflight_checks_docx_required_placeholders(tmp_path):
    default = default_master("exam")
    official = default_master("official")

    assert default is not None
    default_result = check_master_preflight(default)
    assert default_result.status == "ok"
    assert default_result.missing_required_placeholders == ()

    assert official is not None
    official_result = check_master_preflight(official)
    assert official_result.status == "ok"
    assert official_result.missing_required_placeholders == ()

    incomplete_path = tmp_path / "incomplete_master.docx"
    doc = Document()
    doc.add_paragraph("{{af_title}}")
    doc.save(incomplete_path)

    status, missing = check_placeholder_contract(
        incomplete_path,
        default.placeholder_contract,
    )

    assert status == "missing_required_placeholders"
    assert missing == ("af_questions",)


def test_official_user_docx_is_discovered_without_user_json(tmp_path):
    user_dir = tmp_path / "config_library" / "masters" / "official" / "user"
    user_dir.mkdir(parents=True)
    valid_path = user_dir / "agency_letter.docx"
    valid = Document()
    valid.add_paragraph("{{@text:official_title}}")
    valid.add_paragraph("{{@text:official_body}}")
    valid.save(valid_path)
    invalid = Document()
    invalid.add_paragraph("plain Word file")
    invalid.save(user_dir / "invalid.docx")

    masters = list_masters("official", user_master_dir=user_dir)

    assert [master.source_type for master in masters].count("builtin") == 5
    assert [master.source_type for master in masters].count("user") == 1
    user = next(master for master in masters if master.source_type == "user")
    assert user.label == "agency_letter"
    assert user.docx_path == valid_path
    assert user.manifest_path is None
    assert user.readonly is False
    assert check_master_preflight(user).status == "ok"
    assert user.supported_assembly_types == default_master(
        "official"
    ).supported_assembly_types


def test_official_master_copy_creates_discoverable_user_docx(tmp_path):
    builtin = default_master("official")
    assert builtin is not None

    copied = create_official_master_copy(builtin, output_dir=tmp_path)
    discovered = list_masters("official", user_master_dir=tmp_path)

    assert copied.docx_path.is_file()
    assert copied.source_type == "user"
    assert copied.readonly is False
    assert [
        master.master_id for master in discovered if master.source_type == "user"
    ] == [copied.master_id]


def test_official_letter_copy_keeps_layout_family_without_user_json(tmp_path):
    builtin = get_master("official_gbt_letter", "official")
    assert builtin is not None

    copied = create_official_master_copy(builtin, output_dir=tmp_path)
    discovered = next(
        master
        for master in list_masters("official", user_master_dir=tmp_path)
        if master.source_type == "user"
    )

    assert copied.base_master_id == "official_gbt_letter"
    assert discovered.base_master_id == "official_gbt_letter"
    assert "letter" in discovered.supported_assembly_types
    assert "minutes" not in discovered.supported_assembly_types


def test_official_scene_validation_preserves_user_master_id_without_path(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(master_library, "OFFICIAL_USER_MASTER_DIR", tmp_path)
    builtin = default_master("official")
    assert builtin is not None
    copied = create_official_master_copy(builtin, output_dir=tmp_path)
    scene = create_builtin_scene("official", mode_id="official")
    scene.scene_id = "official_custom"
    scene.master_id = copied.master_id

    normalized = config_library.validate_scene_resource_ids(
        scene,
        mode_id="official",
    )

    assert normalized.master_id == copied.master_id
    assert not hasattr(normalized, "master_ref")
    assert not hasattr(normalized, "template_ref")


def test_discovered_official_user_docx_assembles_letter(tmp_path):
    user_dir = tmp_path / "user_masters"
    user_dir.mkdir()
    source = user_dir / "letter_layout.docx"
    doc = Document()
    doc.core_properties.keywords = "alavette:official-layout=letter"
    doc.add_paragraph("{{@text:official_title}}")
    doc.add_paragraph("{{@text:official_body}}")
    doc.save(source)
    user = next(
        master
        for master in list_masters("official", user_master_dir=user_dir)
        if master.source_type == "user"
    )

    result = assemble_official_document_docx(
        "letter",
        {
            "title": "测试函",
            "body": "函件正文",
            "organization": "测试单位",
            "document_no": "测函〔2026〕1号",
            "recipient": "有关单位",
            "issue_date": "2026年7月11日",
        },
        tmp_path / "output",
        master=user,
    )

    assert result.ok is True
    assert result.master_id == user.master_id
    assert result.docx_path is not None and result.docx_path.is_file()


def test_execution_snapshot_preserves_compatible_user_official_master(
    monkeypatch,
    tmp_path,
):
    user_dir = tmp_path / "user_masters"
    monkeypatch.setattr(master_library, "OFFICIAL_USER_MASTER_DIR", user_dir)
    builtin = default_master("official")
    assert builtin is not None
    user = create_official_master_copy(builtin, output_dir=user_dir)
    source = tmp_path / "source.docx"
    Document().save(source)
    scene = load_scene_from_library("official", mode_id="official")
    scene.master_id = user.master_id
    evidence = build_object_preflight_evidence(scene, source)
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=config_library.load_template_from_library(
            "official_gbt",
            mode_id="official",
        ),
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=source,
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="official_gbt",
        document_type_id="notice",
        object_preflight_confirmation_revision=evidence.source_revision,
        object_preflight_confirmation_digest=evidence.evidence_digest,
    )

    try:
        assert snapshot.ready
        assert snapshot.master_ref.requested_id == user.master_id
        assert snapshot.master_ref.effective_id == user.master_id
        assert snapshot.master_ref.source_type == "user"
        assert dict(snapshot.official_master_ids_by_document_type) == {
            "notice": user.master_id
        }
        assert snapshot.frozen_master is not None
        assert snapshot.frozen_master.master_id == user.master_id
        assert snapshot.frozen_master.execution_frozen is True
        assert snapshot.frozen_master.docx_path.read_bytes() == user.docx_path.read_bytes()
    finally:
        cleanup_execution_session_resources(snapshot)


def test_official_document_profiles_cover_core_document_types():
    profiles = list_official_document_profiles()
    profile_ids = {profile.profile_id for profile in profiles}

    assert len(profiles) == 15
    assert {
        "resolution",
        "decision",
        "order",
        "bulletin",
        "announcement",
        "notice_public",
        "opinion",
        "notice",
        "circular",
        "report",
        "request",
        "approval",
        "proposal",
        "letter",
        "minutes",
    } <= profile_ids
    assert get_official_document_profile("notice").label == "通知"
    assert get_official_document_profile("minutes").category == "meeting"
    assert get_official_document_profile("missing") is None


def test_official_batch_document_type_prefers_resolved_context_over_metadata():
    assert resolve_official_batch_document_type_id(
        {"document_type": "  official:minutes  "},
        {"official_profile_id": "letter"},
    ) == "minutes"


def test_official_batch_document_type_uses_metadata_only_for_empty_context():
    assert resolve_official_batch_document_type_id(
        {"document_type": "  "},
        {"official_profile_id": " official:letter "},
    ) == "letter"
    assert resolve_official_batch_document_type_id({}, {}) == ""


def test_official_batch_document_type_preserves_unknown_for_boundary_validation():
    assert resolve_official_batch_document_type_id(
        {"document_type": " future_official_type "},
        {"official_profile_id": "notice"},
    ) == "future_official_type"


def test_official_document_profiles_bind_material_schemas_to_master_placeholders():
    official = default_master("official")
    contracts = list_official_document_assembly_contracts()
    notice = get_official_document_assembly_contract("notice")
    minutes = get_official_document_assembly_contract("minutes")
    profiles_by_id = {
        profile.profile_id: profile for profile in list_official_document_profiles()
    }

    assert official is not None
    assert len(contracts) == 15
    assert notice is not None
    assert minutes is not None
    assert notice.master_id == "official_gbt_standard"
    assert notice.material_schema_ids == ("official_document_v1",)
    assert minutes.material_schema_ids == (
        "official_document_v1",
        "administrative_meeting_fields_v1",
    )
    assert "official_docx" in notice.delivery_versions
    assert "internal_review_docx" in notice.delivery_versions
    assert "archive_manifest" in notice.delivery_versions
    assert "review_pdf" in notice.delivery_versions
    assert "official_title" in notice.placeholder_ids
    assert "official_body" in notice.placeholder_ids
    assert "official_meeting_time" in minutes.placeholder_ids
    assert "official_meeting_attendees" in minutes.placeholder_ids

    for contract in contracts:
        profile = profiles_by_id[contract.profile_id]
        profile_placeholders = set(profile.required_placeholders)
        profile_placeholders.update(profile.optional_placeholders)
        contract_master = get_master(contract.master_id, "official")
        assert contract_master is not None
        placeholder_ids = set(contract_master.placeholder_contract.required)
        placeholder_ids.update(contract_master.placeholder_contract.optional)
        assert set(contract.material_schema_ids)
        for schema_id in contract.material_schema_ids:
            schema = get_material_schema(schema_id)
            schema_field_keys = {field.key for field in schema.fields}
            for binding in contract.field_bindings:
                if binding.material_schema_id != schema_id:
                    continue
                assert binding.field_key in schema_field_keys
                if not binding.applicable:
                    continue
                assert binding.placeholder_id in placeholder_ids
                assert binding.placeholder_id in profile_placeholders
        assert {
            binding.material_schema_id for binding in contract.field_bindings
        } <= set(contract.material_schema_ids)
    supported_by_family = {
        profile_id
        for master in list_masters("official")
        if master.source_type == "builtin"
        for profile_id in master.supported_assembly_types
    }
    assert supported_by_family == set(profiles_by_id)


def test_builtin_official_plan_requests_optional_review_pdf_only_for_official_mode():
    plan = load_scene_from_library("official", mode_id="official")
    assert plan.default_delivery_preset().artifacts.review_pdf is True

    exam = load_scene_from_library("exam", mode_id="exam")
    assert exam.default_delivery_preset().artifacts.review_pdf is False


def test_master_library_returns_empty_for_modes_without_master_contracts_yet():
    assert list_masters("bidding") == ()
    assert default_master("bidding") is None


def test_master_library_does_not_fallback_between_work_modes():
    for mode_id in ("custom", "thesis", "bidding", "technical", "report", "unknown"):
        assert list_masters(mode_id) == ()
        assert default_master(mode_id) is None
        assert get_master("default_exam", mode_id) is None
        assert get_master("official_gbt_standard", mode_id) is None

    exam_ids = {master.master_id for master in list_masters("exam")}
    official_ids = {master.master_id for master in list_masters("official")}

    assert "official_gbt_standard" not in exam_ids
    assert "default_exam" not in official_ids
