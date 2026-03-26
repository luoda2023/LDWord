from docx import Document

from src.docx_io.field_refresh import refresh_doc_fields_with_word
from src.engine.pipeline import Pipeline
from src.scene.manager import load_default_scene


def test_refresh_write_back_replace_failure_keeps_original_file(tmp_path, monkeypatch):
    target_path = tmp_path / "output.docx"
    target_path.write_bytes(b"old-bytes")

    def _fake_refresh(shadow_path, _timeout_sec):
        shadow_path.write_bytes(b"new-bytes")
        return True, "ok(pywin32)"

    def _fail_replace(_src, _dst):
        raise PermissionError("replace blocked")

    monkeypatch.setattr("src.docx_io.field_refresh._refresh_via_pywin32", _fake_refresh)
    monkeypatch.setattr("src.docx_io.field_refresh.os.replace", _fail_replace)

    ok, detail = refresh_doc_fields_with_word(str(target_path), timeout_sec=10)

    assert ok is False
    assert "failed to write back" in detail
    assert target_path.read_bytes() == b"old-bytes"


def test_pipeline_uses_30s_default_field_refresh_timeout_when_env_missing(tmp_path, monkeypatch):
    source = tmp_path / "input.docx"
    doc = Document()
    doc.add_paragraph("目录")
    doc.add_paragraph("旧目录条目\t1")
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    doc.add_paragraph("正文内容")
    doc.save(source)

    captured: dict[str, int] = {}

    def _fake_refresh(doc_path, timeout_sec=10):
        captured["timeout_sec"] = timeout_sec
        return False, "refresh skipped in test"

    cfg = load_default_scene()
    cfg.output.final_docx = True
    cfg.output.compare_docx = False
    cfg.output.report_json = False
    cfg.output.report_markdown = False

    monkeypatch.delenv("DOCX_DISABLE_FIELD_REFRESH", raising=False)
    monkeypatch.delenv("DOCX_FIELD_REFRESH_TIMEOUT_SEC", raising=False)
    monkeypatch.setattr("src.engine.pipeline.refresh_doc_fields_with_word", _fake_refresh)

    result = Pipeline(cfg).run(str(source))

    assert result.success is True
    assert captured["timeout_sec"] == 30
