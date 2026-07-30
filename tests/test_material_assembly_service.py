from __future__ import annotations

import ast
from dataclasses import replace
from hashlib import sha256
import os
from pathlib import Path
from shutil import copy2
from zipfile import ZipFile

from docx import Document
from PIL import Image
import pytest

from docx.oxml.ns import qn

from src.config.attachment_materials import AttachmentBinding, AttachmentItem
from src.config.content_artifacts import ContentMaterialBinding
from src.config.content_materials import (
    ContentInsertionRule,
    FileAssetRef,
    content_anchor_token,
)
from src.config.entity import EntityProfile
from src.config.image_materials import (
    ImageCardinality,
    ImageCoLocationGuard,
    ImageMaterialRule,
    ImageOccurrencePolicy,
    ImagePlacementMode,
    ImagePlacementPolicy,
    ImageSourceItem,
    ImageWatermarkPolicy,
)
from src.services.material_execution import (
    AssemblyFailureCode,
    AssemblyStage,
    MaterialAssemblyDependencies,
    MaterialAssemblyError,
    MaterialAssemblyRequest,
    MaterialAssemblyService,
    MaterialPipelineOutcome,
    MaterialVariantRequest,
    PipelineVisibilityEvidence,
    VisibilityEvidenceKind,
)
from src.services.material_execution import assembly as assembly_module
from src.services.material_assets.image_transform_batch import (
    ImageTransformBatch,
    ImageTransformBatchError,
)
from src.services.material_content.composer import ContentMaterialComposer
from src.services.material_content.artifact_repository import ContentArtifactRepository
from src.services.material_content.compiler import compile_content_material
from src.services.material_execution.material_snapshot_builder import MaterialSnapshotBuilder
from src.services.material_execution.dependency_indexer import MaterialDependencyIndexer
from src.shared.engine.office_image_layout import (
    OfficeImageProvider,
    image_inventory_sha256,
    image_preservation_sha256,
)
from tests.material_image_layout_fake import (
    FakeSuccessfulLayout as _FakeSuccessfulLayout,
)


def _hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _save_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def _visibility(variant_id: str) -> PipelineVisibilityEvidence:
    return PipelineVisibilityEvidence(
        variant_id=variant_id,
        kind=VisibilityEvidenceKind.RECEIPT,
        visibility_receipt_id=sha256(
            f"visibility:{variant_id}".encode("utf-8")
        ).hexdigest(),
    )


class _CopyPipeline:
    def __init__(self, *, mutate_source: bool = False) -> None:
        self.mutate_source = mutate_source
        self.requests = []

    def __call__(self, request, *, cancel_check=None):
        self.requests.append(request)
        for target in request.targets:
            copy2(request.prepared_input.path, target.owned_stage_output_path)
            document = Document(target.owned_stage_output_path)
            document.add_paragraph(f"variant={target.variant_id}")
            document.save(target.owned_stage_output_path)
        if self.mutate_source:
            _save_docx(Path(request.logical_source_path), "source drift")
        return MaterialPipelineOutcome(
            tuple(_visibility(target.variant_id) for target in request.targets)
        )


def _request(
    tmp_path: Path,
    *,
    source: Path,
    variants: tuple[MaterialVariantRequest, ...],
    profile: EntityProfile | None = None,
    image_rules=(),
    image_source_items=(),
) -> MaterialAssemblyRequest:
    return MaterialAssemblyRequest(
        source_docx_path=str(source),
        profile=profile or EntityProfile(profile_id="p1", profile_name="P1"),
        frozen_field_values={"project": "Alpha"},
        material_schema_id="schema-a",
        material_schema_version="1",
        rule_versions={"fields": "1", "content": "1", "images": "1"},
        variants=variants,
        office_provider=OfficeImageProvider.WORD,
        image_rules=tuple(image_rules),
        image_source_items=tuple(image_source_items),
        image_cache_dir=str(tmp_path / "image-cache"),
        work_root=str(tmp_path / "work"),
    )


def _variant(variant_id: str, path: Path) -> MaterialVariantRequest:
    return MaterialVariantRequest(variant_id, "1", str(path))


def _image_fixture(tmp_path: Path):
    image_path = tmp_path / "qualification.png"
    Image.new("RGB", (180, 120), "white").save(image_path)
    image_ref = FileAssetRef(
        source_path=str(image_path),
        original_name=image_path.name,
        media_type="image/png",
        content_sha256=_hash(image_path),
        byte_size=image_path.stat().st_size,
    )
    rule = ImageMaterialRule(
        rule_id="qualification-rule",
        source_role="qualification",
        anchor_token="{{@img:qualification1}}",
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIXED_BOX_FLOW,
            fixed_width_cm=8.0,
        ),
        required=True,
        occurrence_policy=ImageOccurrencePolicy.EXACTLY_ONE,
        cardinality=ImageCardinality.SINGLE,
        watermark=ImageWatermarkPolicy(enabled=False),
    )
    source_item = ImageSourceItem(
        role="qualification",
        item_id="qualification-1",
        sequence=0,
        image_ref=image_ref,
    )
    return rule, source_item


def _content_image_profile(
    tmp_path: Path,
) -> tuple[EntityProfile, str, ContentArtifactRepository]:
    content_dir = tmp_path / "content"
    content_dir.mkdir()
    image_path = content_dir / "diagram.png"
    Image.new("RGB", (160, 100), "white").save(image_path)
    markdown = content_dir / "route.md"
    markdown.write_text("# Route\n\n![diagram](diagram.png)\n", encoding="utf-8")
    repository = ContentArtifactRepository(tmp_path / "content-artifacts")
    compiled = compile_content_material(markdown, repository)
    assert not compiled.blocked, compiled.findings
    assert compiled.artifact_ref is not None
    binding = ContentMaterialBinding(
        content_id="route",
        label="Route",
        artifact_ref=compiled.artifact_ref,
    )
    rule = ContentInsertionRule(
        rule_id="route-rule",
        content_id="route",
        anchor_token=content_anchor_token("route"),
    )
    return (
        EntityProfile(
            profile_id="p-content",
            profile_name="Content",
            content_bindings={binding.content_id: binding},
            content_rules=[rule],
        ),
        rule.anchor_token,
        repository,
    )


class _VisibilityFilteringPipeline:
    def __call__(self, request, *, cancel_check=None):
        prepared = Document(request.prepared_input.path)
        marker_names = tuple(
            item.get(qn("w:name"))
            for item in prepared.element.body.findall(".//" + qn("w:bookmarkStart"))
            if item.get(qn("w:name"))
        )
        assert len(marker_names) == 1
        marker = marker_names[0]
        evidence = []
        for target in request.targets:
            copy2(request.prepared_input.path, target.owned_stage_output_path)
            removed = ()
            if target.variant_id == "suppressed":
                document = Document(target.owned_stage_output_path)
                for paragraph in document.paragraphs:
                    names = {
                        item.get(qn("w:name"))
                        for item in paragraph._p.findall(".//" + qn("w:bookmarkStart"))
                    }
                    if marker in names:
                        paragraph._p.getparent().remove(paragraph._p)
                        break
                document.save(target.owned_stage_output_path)
                removed = (marker,)
            evidence.append(
                PipelineVisibilityEvidence(
                    variant_id=target.variant_id,
                    kind=VisibilityEvidenceKind.RECEIPT,
                    visibility_receipt_id=sha256(
                        f"visibility:{target.variant_id}:{marker}".encode("utf-8")
                    ).hexdigest(),
                    removed_content_image_markers=removed,
                )
            )
        return MaterialPipelineOutcome(tuple(evidence))


class _CancelBlindComposer:
    """Delegate composition without consuming the service boundary counter."""

    def __init__(self) -> None:
        self._delegate = ContentMaterialComposer()

    def compose(self, request, *, cancel_check=None):
        return self._delegate.compose(request, cancel_check=None)


def _cancel_blind_transform_factory(plans, **kwargs):
    delegate = ImageTransformBatch(plans, **kwargs)

    class _Batch:
        def run(self, *, cancel_check=None):
            return delegate.run(cancel_check=None)

    return _Batch()


def _text_content_profile(
    tmp_path: Path,
) -> tuple[EntityProfile, Path, str, ContentArtifactRepository]:
    markdown = tmp_path / "reusable.md"
    markdown.write_text("Reusable technical route.\n", encoding="utf-8")
    repository = ContentArtifactRepository(tmp_path / "content-artifacts")
    compiled = compile_content_material(markdown, repository)
    assert not compiled.blocked, compiled.findings
    assert compiled.artifact_ref is not None
    binding = ContentMaterialBinding(
        content_id="route",
        label="Route",
        artifact_ref=compiled.artifact_ref,
    )
    rule = ContentInsertionRule(
        rule_id="route-rule",
        content_id=binding.content_id,
        anchor_token=content_anchor_token(binding.content_id),
    )
    return (
        EntityProfile(
            profile_id="p-text-content",
            profile_name="Text content",
            content_bindings={binding.content_id: binding},
            content_rules=[rule],
        ),
        markdown,
        rule.anchor_token,
        repository,
    )


class _RepositorySnapshotBuilder:
    def __init__(self, repository: ContentArtifactRepository) -> None:
        self._repository = repository
        self._delegate = MaterialSnapshotBuilder()

    def build(self, request):
        return self._delegate.build(
            replace(request, content_artifact_root=self._repository.root)
        )


def _content_dependencies(
    repository: ContentArtifactRepository,
    **overrides,
) -> MaterialAssemblyDependencies:
    return replace(
        MaterialAssemblyDependencies(),
        snapshot_builder=_RepositorySnapshotBuilder(repository),
        dependency_indexer=MaterialDependencyIndexer(
            content_repository=repository,
        ),
        content_composer=ContentMaterialComposer(repository=repository),
        **overrides,
    )


def test_no_image_success_keeps_intake_and_delivery_identity_separate(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "source")
    _save_docx(final, "old final")
    pipeline = _CopyPipeline()

    receipt = MaterialAssemblyService(pipeline).execute(
        _request(tmp_path, source=source, variants=(_variant("final", final),))
    )

    assert receipt.intake_snapshot.snapshot_id
    assert "resolved_image_plans" not in receipt.intake_snapshot.to_dict()
    variant = receipt.variants[0]
    assert variant.delivery_plan.snapshot_id == receipt.intake_snapshot.snapshot_id
    assert variant.delivery_plan.resolved_image_plans == ()
    assert variant.office_layout_receipt is None
    assert variant.image_transform_receipt.jobs == ()
    assert variant.final_output.sha256 == _hash(final)
    assert "variant=final" in "\n".join(p.text for p in Document(final).paragraphs)
    assert receipt.to_dict()["receipt_id"] == receipt.receipt_id
    assert not list(tmp_path.glob(".*.material-*"))


def test_two_variants_are_both_staged_before_publish(tmp_path: Path):
    source = tmp_path / "source.docx"
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    _save_docx(source, "source")
    pipeline = _CopyPipeline()

    receipt = MaterialAssemblyService(pipeline).execute(
        _request(
            tmp_path,
            source=source,
            variants=(_variant("first", first), _variant("second", second)),
        )
    )

    assert [item.variant.variant_id for item in receipt.variants] == ["first", "second"]
    assert first.is_file() and second.is_file()
    assert "variant=first" in "\n".join(p.text for p in Document(first).paragraphs)
    assert "variant=second" in "\n".join(p.text for p in Document(second).paragraphs)
    callback_request = pipeline.requests[0]
    assert all(
        Path(target.owned_stage_output_path).parent
        == Path(target.final_output_path).parent
        for target in callback_request.targets
    )


def test_second_variant_publish_failure_restores_every_old_final(tmp_path: Path):
    source = tmp_path / "source.docx"
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    _save_docx(source, "source")
    _save_docx(first, "old first")
    _save_docx(second, "old second")
    old_hashes = (_hash(first), _hash(second))
    failed = False

    def fail_second_publish(src, dst):
        nonlocal failed
        if (
            not failed
            and Path(dst).resolve() == second.resolve()
            and Path(src).name.endswith(".stage.docx")
        ):
            failed = True
            raise OSError("simulated second publish failure")
        return os.replace(src, dst)

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            atomic_replace=fail_second_publish,
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("first", first), _variant("second", second)),
            )
        )

    assert captured.value.primary.code is AssemblyFailureCode.ATOMIC_PUBLISH_FAILED
    assert captured.value.primary.stage is AssemblyStage.PUBLISH
    assert captured.value.primary.variant_id == "second"
    assert (_hash(first), _hash(second)) == old_hashes
    assert not list(tmp_path.glob(".*.material-*"))


def test_receipt_failure_after_publish_restores_every_old_final_and_backups(
    tmp_path: Path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    _save_docx(source, "source")
    _save_docx(first, "old first")
    _save_docx(second, "old second")
    old_hashes = (_hash(first), _hash(second))

    def fail_receipt_construction(*_args, **_kwargs):
        raise RuntimeError("simulated receipt construction failure")

    monkeypatch.setattr(
        assembly_module,
        "MaterialAssemblyReceipt",
        fail_receipt_construction,
    )
    service = MaterialAssemblyService(_CopyPipeline())

    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("first", first), _variant("second", second)),
            )
        )

    assert captured.value.primary.code is AssemblyFailureCode.ATOMIC_PUBLISH_FAILED
    assert captured.value.primary.stage is AssemblyStage.PUBLISH
    assert captured.value.primary.variant_id == "second"
    assert (_hash(first), _hash(second)) == old_hashes
    assert not list(tmp_path.glob(".*.material-*"))
    assert not list(tmp_path.glob(".*.backup"))


def test_image_layout_failure_rolls_back_old_final(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "{{@img:qualification1}}")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    rule, source_item = _image_fixture(tmp_path)

    def fail_layout(_request):
        raise RuntimeError("layout failed")

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            layout_runner=fail_layout,
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
                image_rules=(rule,),
                image_source_items=(source_item,),
            )
        )

    assert captured.value.primary.code is AssemblyFailureCode.OFFICE_LAYOUT_FAILED
    assert _hash(final) == old_hash
    assert not list(tmp_path.glob(".*.material-*"))
    assert not list(tmp_path.glob(".lark-layout-*.docx"))


def test_image_policy_freeze_failure_has_intake_code_and_preserves_final(
    tmp_path: Path,
):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "source")
    _save_docx(final, "old final")
    old_hash = _hash(final)

    class BrokenFreezer:
        def freeze(self, **_kwargs):
            raise RuntimeError("cannot freeze image policy")

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            image_policy_freezer=BrokenFreezer(),
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(tmp_path, source=source, variants=(_variant("final", final),))
        )

    assert captured.value.primary.code is AssemblyFailureCode.INTAKE_FREEZE_FAILED
    assert captured.value.primary.stage is AssemblyStage.INTAKE_FREEZE
    assert _hash(final) == old_hash


def test_final_drift_during_verify_has_exact_code_and_is_rolled_back(
    tmp_path: Path,
):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "source")
    _save_docx(final, "old final")
    old_hash = _hash(final)

    def mutate_final_when_image_plan_loads(path: str):
        _save_docx(final, "foreign writer changed final")
        return Document(path)

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            document_loader=mutate_final_when_image_plan_loads,
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(tmp_path, source=source, variants=(_variant("final", final),))
        )

    assert captured.value.primary.code is AssemblyFailureCode.FINAL_DRIFT
    assert captured.value.primary.stage is AssemblyStage.VERIFY
    assert _hash(final) == old_hash


def test_unclassified_verify_failure_has_final_verify_code() -> None:
    error = assembly_module._map_exception(
        RuntimeError("candidate verification failed"),
        AssemblyStage.VERIFY,
        variant_id="review",
    )

    assert error.primary.code is AssemblyFailureCode.FINAL_VERIFY_FAILED
    assert error.primary.stage is AssemblyStage.VERIFY
    assert error.primary.variant_id == "review"


def test_source_drift_is_rejected_and_old_final_survives(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "source")
    _save_docx(final, "old final")
    old_hash = _hash(final)

    with pytest.raises(MaterialAssemblyError) as captured:
        MaterialAssemblyService(_CopyPipeline(mutate_source=True)).execute(
            _request(tmp_path, source=source, variants=(_variant("final", final),))
        )

    assert captured.value.primary.code is AssemblyFailureCode.SOURCE_DRIFT
    assert _hash(final) == old_hash


def test_foreign_layout_shadow_is_rejected_without_deleting_it(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    foreign = tmp_path / "foreign.docx"
    _save_docx(source, "{{@img:qualification1}}")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    rule, source_item = _image_fixture(tmp_path)
    fake_layout = _FakeSuccessfulLayout(foreign_shadow=foreign)
    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            layout_runner=fake_layout,
        ),
    )

    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
                image_rules=(rule,),
                image_source_items=(source_item,),
            )
        )

    assert captured.value.primary.code is AssemblyFailureCode.FOREIGN_SHADOW
    assert _hash(final) == old_hash
    assert foreign.is_file()


def test_image_success_uses_layout_shadow_as_only_publish_candidate(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "{{@img:qualification1}}")
    rule, source_item = _image_fixture(tmp_path)
    fake_layout = _FakeSuccessfulLayout()
    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            layout_runner=fake_layout,
        ),
    )

    receipt = service.execute(
        _request(
            tmp_path,
            source=source,
            variants=(_variant("final", final),),
            image_rules=(rule,),
            image_source_items=(source_item,),
        )
    )

    variant = receipt.variants[0]
    assert len(variant.delivery_plan.resolved_image_plans) == 1
    assert variant.office_layout_receipt is not None
    assert (
        variant.office_layout_receipt.source_image_inventory_sha256
        == image_inventory_sha256(())
    )
    assert (
        variant.office_layout_receipt.preexisting_image_semantic_sha256_before
        == image_preservation_sha256(())
    )
    assert variant.publish_candidate.sha256 == variant.final_output.sha256
    assert _hash(source) == receipt.source.sha256
    assert not list(tmp_path.glob(".lark-layout-*.docx"))


def test_phase_k_accepts_preserved_preexisting_image_inventory(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    existing = tmp_path / "existing.png"
    Image.new("RGB", (80, 60), "blue").save(existing)
    document = Document()
    document.add_paragraph().add_run().add_picture(str(existing))
    document.add_paragraph("{{@img:qualification1}}")
    document.save(source)
    rule, source_item = _image_fixture(tmp_path)
    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            layout_runner=_FakeSuccessfulLayout(),
        ),
    )

    receipt = service.execute(
        _request(
            tmp_path,
            source=source,
            variants=(_variant("final", final),),
            image_rules=(rule,),
            image_source_items=(source_item,),
        )
    )

    layout = receipt.variants[0].office_layout_receipt
    assert layout is not None
    assert layout.preexisting_images_unchanged is True
    assert layout.inserted_images_job_owned is True
    assert layout.inserted_images_visible is True
    assert layout.sentinel_texts_hidden is True
    assert len(layout.source_image_inventory) == 1
    assert len(layout.inserted_image_inventory) == 1
    assert (
        layout.preexisting_image_semantic_sha256_before
        == layout.preexisting_image_semantic_sha256_after
    )


def test_variant_visibility_receipt_is_the_only_content_draft_suppression_owner(
    tmp_path: Path,
):
    source = tmp_path / "source.docx"
    suppressed = tmp_path / "suppressed.docx"
    retained = tmp_path / "retained.docx"
    profile, anchor_token, repository = _content_image_profile(tmp_path)
    _save_docx(source, anchor_token)
    fake_layout = _FakeSuccessfulLayout()
    service = MaterialAssemblyService(
        _VisibilityFilteringPipeline(),
        dependencies=_content_dependencies(
            repository,
            layout_runner=fake_layout,
        ),
    )

    receipt = service.execute(
        _request(
            tmp_path,
            source=source,
            profile=profile,
            variants=(
                _variant("suppressed", suppressed),
                _variant("retained", retained),
            ),
        )
    )

    by_variant = {item.variant.variant_id: item for item in receipt.variants}
    hidden = by_variant["suppressed"]
    visible = by_variant["retained"]
    assert hidden.delivery_plan.resolved_image_plans == ()
    assert len(hidden.not_applicable_content_image_job_ids) == 1
    assert hidden.not_applicable_content_image_markers == (
        hidden.visibility_evidence.removed_content_image_markers
    )
    assert len(visible.delivery_plan.resolved_image_plans) == 1
    assert visible.not_applicable_content_image_job_ids == ()
    assert visible.office_layout_receipt is not None


@pytest.mark.parametrize(
    ("cancel_at", "expected_stage"),
        (
            (1, "preflight"),
            (2, "intake_freeze"),
            (3, "dependency_index"),
            (4, "content_compose"),
            (5, "pipeline"),
            (6, "image_plan"),
            (7, "verify"),
            (8, "publish"),
    ),
)
def test_cancellation_at_each_service_boundary_preserves_old_final(
    tmp_path: Path,
    cancel_at: int,
    expected_stage: str,
):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "source")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    calls = 0

    def cancel_check() -> bool:
        nonlocal calls
        calls += 1
        return calls == cancel_at

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            content_composer=_CancelBlindComposer(),
            transform_batch_factory=_cancel_blind_transform_factory,
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(tmp_path, source=source, variants=(_variant("final", final),)),
            cancel_check=cancel_check,
        )

    assert captured.value.primary.code is AssemblyFailureCode.CANCELLED
    assert expected_stage in captured.value.primary.message
    assert _hash(final) == old_hash
    assert not list(tmp_path.glob(".*.material-*"))


def test_second_variant_cancellation_preserves_complete_old_final_set(
    tmp_path: Path,
):
    source = tmp_path / "source.docx"
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    _save_docx(source, "source")
    _save_docx(first, "old first")
    _save_docx(second, "old second")
    old_hashes = (_hash(first), _hash(second))
    calls = 0

    def cancel_check() -> bool:
        nonlocal calls
        calls += 1
        return calls == 7

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            content_composer=_CancelBlindComposer(),
            transform_batch_factory=_cancel_blind_transform_factory,
        ),
    )

    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("first", first), _variant("second", second)),
            ),
            cancel_check=cancel_check,
        )

    assert captured.value.primary.code is AssemblyFailureCode.CANCELLED
    assert captured.value.primary.stage is AssemblyStage.CANCEL
    assert "image_plan" in captured.value.primary.message
    assert captured.value.primary.variant_id == "second"
    assert (_hash(first), _hash(second)) == old_hashes
    assert not list(tmp_path.glob(".*.material-*"))


def test_cancellation_before_office_layout_preserves_old_final(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "{{@img:qualification1}}")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    rule, source_item = _image_fixture(tmp_path)
    calls = 0

    def cancel_check() -> bool:
        nonlocal calls
        calls += 1
        return calls == 7

    def unexpected_layout(_request):
        raise AssertionError("cancelled assembly must not open Office")

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            content_composer=_CancelBlindComposer(),
            transform_batch_factory=_cancel_blind_transform_factory,
            layout_runner=unexpected_layout,
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
                image_rules=(rule,),
                image_source_items=(source_item,),
            ),
            cancel_check=cancel_check,
        )

    assert captured.value.primary.code is AssemblyFailureCode.CANCELLED
    assert "office_layout" in captured.value.primary.message
    assert _hash(final) == old_hash


def test_transform_cancellation_preserves_old_final(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "{{@img:qualification1}}")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    rule, source_item = _image_fixture(tmp_path)

    class _CancelledBatch:
        def run(self, *, cancel_check=None):
            raise ImageTransformBatchError(
                "cancelled",
                "cancelled in transform",
                job_id="qualification-rule:qualification-1:0",
            )

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            transform_batch_factory=lambda _plans, **_kwargs: _CancelledBatch(),
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
                image_rules=(rule,),
                image_source_items=(source_item,),
            )
        )

    assert captured.value.primary.code is AssemblyFailureCode.CANCELLED
    assert "image_transform" in captured.value.primary.message
    assert _hash(final) == old_hash


def test_image_transform_failure_preserves_old_final(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "{{@img:qualification1}}")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    rule, source_item = _image_fixture(tmp_path)

    class _FailedBatch:
        def run(self, *, cancel_check=None):
            raise RuntimeError("transform failed")

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            transform_batch_factory=lambda _plans, **_kwargs: _FailedBatch(),
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
                image_rules=(rule,),
                image_source_items=(source_item,),
            )
        )

    assert captured.value.primary.code is AssemblyFailureCode.IMAGE_TRANSFORM_FAILED
    assert _hash(final) == old_hash
    assert not list(tmp_path.glob(".*.material-*"))


def test_original_content_source_mutation_does_not_change_compiled_artifact(
    tmp_path: Path,
):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    profile, markdown, anchor, repository = _text_content_profile(tmp_path)
    _save_docx(source, anchor)
    _save_docx(final, "old final")
    old_hash = _hash(final)
    delegate = _CopyPipeline()

    def mutate_content_after_pipeline(request, *, cancel_check=None):
        outcome = delegate(request, cancel_check=cancel_check)
        markdown.write_text("drifted technical route.\n", encoding="utf-8")
        return outcome

    receipt = MaterialAssemblyService(
        mutate_content_after_pipeline,
        dependencies=_content_dependencies(repository),
    ).execute(
        _request(
            tmp_path,
            source=source,
            profile=profile,
            variants=(_variant("final", final),),
        )
    )

    assert receipt.variants[0].final_output.sha256 == _hash(final)
    assert _hash(final) != old_hash
    text = "\n".join(paragraph.text for paragraph in Document(final).paragraphs)
    assert "Reusable technical route." in text
    assert "drifted technical route." not in text


def test_image_source_drift_after_layout_is_rejected(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "{{@img:qualification1}}")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    rule, source_item = _image_fixture(tmp_path)
    image_path = Path(source_item.image_ref.source_path)
    fake_layout = _FakeSuccessfulLayout()

    def mutate_image_after_layout(request):
        receipt = fake_layout(request)
        Image.new("RGB", (180, 120), "black").save(image_path)
        return receipt

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            layout_runner=mutate_image_after_layout,
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
                image_rules=(rule,),
                image_source_items=(source_item,),
            )
        )

    assert captured.value.primary.code is AssemblyFailureCode.SOURCE_DRIFT
    assert _hash(final) == old_hash


def test_attachment_is_frozen_but_never_inserted_into_docx(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    attachment = tmp_path / "supporting-evidence.docx"
    sentinel = "ATTACHMENT-MUST-NEVER-BE-INLINED-7f1b"
    _save_docx(source, "source body")
    _save_docx(attachment, sentinel)
    attachment_ref = FileAssetRef(
        source_path=str(attachment),
        original_name=attachment.name,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        content_sha256=_hash(attachment),
        byte_size=attachment.stat().st_size,
    )
    profile = EntityProfile(
        profile_id="p-attachment",
        profile_name="Attachment boundary",
        attachment_bindings={
            "evidence": AttachmentBinding(
                role="evidence",
                items=(
                    AttachmentItem(
                        item_id="evidence-1",
                        file_ref=attachment_ref,
                    ),
                ),
                required=True,
                min_items=1,
            )
        },
    )

    receipt = MaterialAssemblyService(_CopyPipeline()).execute(
        _request(
            tmp_path,
            source=source,
            profile=profile,
            variants=(_variant("final", final),),
        )
    )

    assert {
        binding.role for binding in receipt.intake_snapshot.attachment_bindings
    } == {"evidence"}
    with ZipFile(final) as archive:
        names = tuple(archive.namelist())
        xml_bytes = b"".join(
            archive.read(name) for name in names if name.endswith(".xml")
        )
    assert not any(name.startswith("word/embeddings/") for name in names)
    assert sentinel.encode("utf-8") not in xml_bytes
    assert attachment.name.encode("utf-8") not in xml_bytes


def test_unknown_main_material_token_is_blocked_by_dependency_preflight(
    tmp_path: Path,
):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "{{@text:unknown_field}}")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    pipeline = _CopyPipeline()

    with pytest.raises(MaterialAssemblyError) as captured:
        MaterialAssemblyService(pipeline).execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
            )
        )

    assert captured.value.primary.code is AssemblyFailureCode.DEPENDENCY_INDEX_INVALID
    assert "unknown_field_token" in captured.value.primary.message
    assert not pipeline.requests
    assert _hash(final) == old_hash


def test_pipeline_cannot_mutate_final_instead_of_its_owned_stage(tmp_path: Path):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "source")
    _save_docx(final, "old final")
    old_hash = _hash(final)

    def malicious_pipeline(request, *, cancel_check=None):
        for target in request.targets:
            copy2(request.prepared_input.path, target.owned_stage_output_path)
            _save_docx(Path(target.final_output_path), "illicit direct publish")
        return MaterialPipelineOutcome(
            tuple(_visibility(target.variant_id) for target in request.targets)
        )

    with pytest.raises(MaterialAssemblyError) as captured:
        MaterialAssemblyService(malicious_pipeline).execute(
            _request(tmp_path, source=source, variants=(_variant("final", final),))
        )

    assert captured.value.primary.code is AssemblyFailureCode.FINAL_DRIFT
    assert _hash(final) == old_hash
    assert not list(tmp_path.glob(".*.material-*"))


@pytest.mark.parametrize(
    "invalid_field",
    (
        "preexisting_images_unchanged",
        "inserted_images_job_owned",
        "inserted_images_visible",
        "sentinel_texts_hidden",
    ),
)
def test_phase_k_false_success_evidence_is_rejected(
    tmp_path: Path,
    invalid_field: str,
):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    _save_docx(source, "{{@img:qualification1}}")
    _save_docx(final, "old final")
    old_hash = _hash(final)
    rule, source_item = _image_fixture(tmp_path)
    fake_layout = _FakeSuccessfulLayout()

    def forged_success(request):
        return replace(fake_layout(request), **{invalid_field: False})

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            layout_runner=forged_success,
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
                image_rules=(rule,),
                image_source_items=(source_item,),
            )
        )

    assert (
        captured.value.primary.code
        is AssemblyFailureCode.OFFICE_LAYOUT_RECEIPT_INVALID
    )
    assert _hash(final) == old_hash


@pytest.mark.parametrize(
    "forgery",
    ("boundary", "strict_pages", "unstable_round"),
)
def test_phase_k_forged_geometry_evidence_is_rejected(
    tmp_path: Path,
    forgery: str,
):
    source = tmp_path / "source.docx"
    final = tmp_path / "final.docx"
    document = Document()
    document.add_paragraph("Qualification certificate")
    document.add_paragraph("{{@img:qualification1}}")
    document.save(source)
    _save_docx(final, "old final")
    old_hash = _hash(final)
    rule, source_item = _image_fixture(tmp_path)
    strict_rule = replace(
        rule,
        placement=ImagePlacementPolicy(
            mode=ImagePlacementMode.FIT_REMAINING_ANCHOR_PAGE,
            max_width_cm=16.0,
            co_location_guard=(
                ImageCoLocationGuard.PRECEDING_NONEMPTY_PARAGRAPH
            ),
        ),
    )
    fake_layout = _FakeSuccessfulLayout()

    def forged_success(request):
        receipt = fake_layout(request)
        if forgery == "unstable_round":
            rounds = tuple(
                replace(item, stable=False)
                if item is receipt.stabilization_rounds[-1]
                else item
                for item in receipt.stabilization_rounds
            )
            return replace(receipt, stabilization_rounds=rounds)
        job = receipt.jobs[0]
        if forgery == "boundary":
            job = replace(job, boundary_ok=False)
        elif forgery == "strict_pages":
            job = replace(job, final_image_page=2)
        return replace(receipt, jobs=(job,))

    service = MaterialAssemblyService(
        _CopyPipeline(),
        dependencies=replace(
            MaterialAssemblyDependencies(),
            layout_runner=forged_success,
        ),
    )
    with pytest.raises(MaterialAssemblyError) as captured:
        service.execute(
            _request(
                tmp_path,
                source=source,
                variants=(_variant("final", final),),
                image_rules=(strict_rule,),
                image_source_items=(source_item,),
            )
        )

    assert (
        captured.value.primary.code
        is AssemblyFailureCode.OFFICE_LAYOUT_RECEIPT_INVALID
    )
    assert _hash(final) == old_hash


def test_material_assembly_execute_remains_a_bounded_phase_coordinator():
    source = Path(assembly_module.__file__).read_text(encoding="utf-8")
    tree = ast.parse(source)
    service = next(
        node
        for node in tree.body
        if isinstance(node, ast.ClassDef)
        and node.name == "MaterialAssemblyService"
    )
    methods = {
        node.name: node
        for node in service.body
        if isinstance(node, ast.FunctionDef)
    }

    assert methods["execute"].end_lineno - methods["execute"].lineno + 1 <= 100
    assert {
        "_run_preflight",
        "_freeze_intake",
        "_index_dependencies",
        "_compose_content",
        "_run_pipeline",
        "_prepare_variants",
        "_prepare_variant",
        "_verify_publish_readiness",
        "_publish",
    } <= set(methods)
    assert all(
        method.end_lineno - method.lineno + 1 <= 100
        for name, method in methods.items()
        if name != "__init__"
    )
