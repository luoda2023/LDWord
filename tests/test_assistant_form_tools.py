from __future__ import annotations

from docx import Document

from src.assistant.adapters.form_tools import build_form_tool_registry
from src.assistant.adapters.workspace_state_adapter import WorkspaceSnapshot
from src.assistant.contracts.permissions import PermissionDecision
from src.assistant.domain.docx_format_evidence import (
    DOCX_FORMAT_EVIDENCE_SCHEMA_VERSION,
)
from src.assistant.tools.gateway import FormToolGateway
from src.assistant.tools.registry import ToolCall


def _snapshot(path="") -> WorkspaceSnapshot:
    return WorkspaceSnapshot(
        mode_id="custom",
        mode_label="通用版",
        scene_id="custom",
        scene_source_type="builtin",
        template_id="default",
        template_source_type="builtin",
        input_path=path,
        input_name="example.docx" if path else "",
        input_exists=bool(path),
        material_summary={"field_count": 2, "image_count": 0},
    )


def _call(name, arguments=None):
    return ToolCall(
        call_id="call-1",
        session_id="session-1",
        turn_id="turn-1",
        tool_name=name,
        arguments=arguments or {},
    )


def test_workspace_tool_never_exposes_absolute_path(tmp_path):
    path = tmp_path / "example.docx"
    registry = build_form_tool_registry(_snapshot(str(path)))

    result = FormToolGateway(registry).invoke(_call("get_current_workspace_state"))

    assert result.status == "success"
    assert "input_path" not in result.output
    assert result.output["input_name"] == "example.docx"


def test_document_inspection_requires_permission_and_returns_bounded_text(tmp_path):
    path = tmp_path / "example.docx"
    document = Document()
    document.add_heading("标题", level=1)
    document.add_paragraph("正文" * 100)
    document.add_table(rows=1, cols=2)
    document.save(path)
    registry = build_form_tool_registry(_snapshot(str(path)))
    gateway = FormToolGateway(registry)
    call = _call(
        "inspect_input_document",
        {"path": str(path), "include_text": True, "max_characters": 20},
    )

    assert gateway.invoke(call).status == "needs_permission"
    result = gateway.invoke(
        call,
        decision=PermissionDecision(request_id="call-1", allowed=True),
    )

    assert result.status == "success"
    assert result.output["name"] == "example.docx"
    assert result.output["heading_count"] == 1
    assert result.output["table_count"] == 1
    assert len(result.output["text"]) == 20
    assert result.output["text_truncated"] is True


def test_document_inspection_can_return_structured_format_evidence(tmp_path):
    path = tmp_path / "standard-reference.docx"
    document = Document()
    document.add_heading("标准标题", level=1)
    document.add_paragraph("标准正文")
    document.save(path)
    registry = build_form_tool_registry(_snapshot(str(path)))
    gateway = FormToolGateway(registry)
    call = _call(
        "inspect_input_document",
        {"path": str(path), "include_format_evidence": True},
    )

    result = gateway.invoke(
        call,
        decision=PermissionDecision(request_id="call-1", allowed=True),
    )

    assert result.status == "success"
    evidence = result.output["format_evidence"]
    assert evidence["schema_version"] == DOCX_FORMAT_EVIDENCE_SCHEMA_VERSION
    assert evidence["source"]["name"] == path.name
    assert evidence["inventory"]["used_paragraph_style_count"] >= 2


def test_scene_router_tool_uses_deterministic_form_router():
    registry = build_form_tool_registry(_snapshot())
    result = FormToolGateway(registry).invoke(
        _call("match_scene_request", {"query": "帮我统一Word格式并修目录"})
    )

    assert result.status == "success"
    assert result.output["selected_route_id"] == "quick_formatting_general"


def test_form_tool_catalog_contains_read_and_draft_boundaries():
    names = {item["name"] for item in build_form_tool_registry(_snapshot()).public_catalog()}
    assert {
        "get_current_workspace_state",
        "list_work_modes",
        "match_scene_request",
        "list_scene_candidates",
        "list_template_candidates",
        "list_material_candidates",
        "list_output_presets",
        "inspect_input_document",
        "create_exam_user_plan_from_docx",
        "select_work_mode_draft",
        "bind_scene_draft",
        "bind_template_draft",
        "set_output_policy_draft",
        "generate_document_fragments",
    }.issubset(names)


def test_exam_user_plan_creation_requires_durable_write_permission(tmp_path):
    registry = build_form_tool_registry(_snapshot())
    result = FormToolGateway(registry).invoke(
        _call(
            "create_exam_user_plan_from_docx",
            {"source_path": str(tmp_path / "current-exam.docx")},
        )
    )

    assert result.status == "needs_permission"
