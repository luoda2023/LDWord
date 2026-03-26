from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from src.engine.change_tracker import ChangeTracker
from src.engine.doc_tree import DocTree
from src.engine.pipeline import Pipeline
from src.engine.rules import table_format as table_helpers
from src.engine.rules.equation_table_format import EquationTableFormatRule
from src.engine.rules.formula_convert import FormulaConvertRule
from src.engine.rules.heading_detect import HeadingDetectRule
from src.engine.rules.formula_style import FormulaStyleRule, _semantic_segments_for_text
from src.engine.rules.formula_to_table import FormulaToTableRule
from src.formula_core.ast import FormulaNode
from src.formula_core.convert import convert_formula_node
from src.formula_core.parse import parse_document_formulas
from src.scene.manager import load_scene_from_data
from src.scene.schema import SceneConfig

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _formula_stats(context: dict, rule_name: str) -> dict:
    return context.get("formula_runtime", {}).get("stats", {}).get(rule_name, {})


def _stub_word_page_spans(monkeypatch, spans):
    monkeypatch.setattr(
        "src.engine.page_scope._resolve_word_paragraph_page_spans",
        lambda doc, source_doc_path=None, timeout_sec=None: list(spans),
    )


def _set_native_formula(paragraph, latex: str, *, block: bool = True) -> None:
    outcome = convert_formula_node(
        FormulaNode(
            kind="equation",
            payload={"latex": latex},
            source_type="latex",
            confidence=0.95,
        ),
        "word_native",
        block=block,
    )
    assert outcome.success is True
    assert outcome.omml_element is not None

    p = paragraph._p
    for child in list(p):
        p.remove(child)

    omml = deepcopy(outcome.omml_element)
    local_name = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if local_name == "oMathPara":
        p.append(omml)
        return

    run = p.makeelement(f"{{{_W_NS}}}r")
    run.append(omml)
    p.append(run)


def _add_mathtype_formula(paragraph, text: str) -> None:
    paragraph.add_run(text)
    run = OxmlElement("w:r")
    obj = OxmlElement("w:object")
    control = OxmlElement("w:control")
    control.set(qn("w:name"), "MathType")
    obj.append(control)
    run.append(obj)
    paragraph._p.append(run)


def _build_heading_context(doc: Document, cfg: SceneConfig) -> dict:
    doc_tree = DocTree()
    doc_tree.build(doc)
    context = {"doc_tree": doc_tree}
    HeadingDetectRule().apply(doc, cfg, ChangeTracker(), context)
    return context


def _table_grid_widths(tbl) -> list[int]:
    return [
        int(col.get(f"{{{_W_NS}}}w"))
        for col in tbl._tbl.findall(f".//{{{_W_NS}}}gridCol")
    ]


def _table_alignment_values(tbl) -> list[str]:
    return tbl._tbl.xpath(
        "./*[namespace-uri()='%s' and local-name()='tblPr']"
        "/*[namespace-uri()='%s' and local-name()='jc']"
        "/@*[local-name()='val']"
        % (_W_NS, _W_NS)
    )


def _cell_alignment_values(cell) -> list[str]:
    return cell._tc.xpath(
        ".//*[namespace-uri()='%s' and local-name()='pPr']"
        "/*[namespace-uri()='%s' and local-name()='jc']"
        "/@*[local-name()='val']"
        % (_W_NS, _W_NS)
    )


def _ensure_run_rpr(run):
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        rpr = etree.SubElement(run._element, qn("w:rPr"))
    return rpr


def _realcase_c3n4_docx_path() -> Path:
    root = Path(__file__).resolve().parent / "test-equation"
    for path in root.glob("*C3N4.docx"):
        if not path.name.startswith("~$"):
            return path
    raise FileNotFoundError("missing C3N4 realcase docx fixture")


def test_scene_manager_syncs_equation_table_flag_and_pipeline_step():
    legacy_cfg = load_scene_from_data(
        {
            "pipeline": ["page_setup", "table_format", "equation_table_format", "section_format"],
        }
    )
    assert legacy_cfg.equation_table_format.enabled is True
    assert "equation_table_format" in legacy_cfg.pipeline

    cfg = load_scene_from_data(
        {
            "pipeline": ["page_setup", "table_format", "section_format"],
            "equation_table_format": {"enabled": True, "numbering_format": "chapter-seq"},
        }
    )
    assert cfg.equation_table_format.enabled is True
    assert cfg.equation_table_format.numbering_format == "chapter-seq"
    assert "equation_table_format" in cfg.pipeline

    cfg2 = load_scene_from_data(
        {
            "pipeline": ["page_setup", "table_format", "equation_table_format", "section_format"],
            "equation_table_format": {"enabled": False},
        }
    )
    assert cfg2.equation_table_format.enabled is False
    assert "equation_table_format" not in cfg2.pipeline


def test_scene_manager_syncs_citation_link_flag_and_pipeline_step():
    legacy_cfg = load_scene_from_data(
        {
            "pipeline": ["page_setup", "table_format", "citation_link", "validation"],
        }
    )
    assert legacy_cfg.citation_link.enabled is True
    assert "citation_link" in legacy_cfg.pipeline

    cfg = load_scene_from_data(
        {
            "pipeline": ["page_setup", "table_format", "validation"],
            "citation_link": {"enabled": True},
        }
    )
    assert cfg.citation_link.enabled is True
    assert "citation_link" in cfg.pipeline
    assert cfg.pipeline.index("citation_link") == cfg.pipeline.index("table_format") + 1

    cfg2 = load_scene_from_data(
        {
            "pipeline": ["page_setup", "table_format", "citation_link", "validation"],
            "citation_link": {"enabled": False},
        }
    )
    assert cfg2.citation_link.enabled is False
    assert "citation_link" not in cfg2.pipeline


def test_scene_manager_auto_upgrades_stale_pipeline_critical_rules():
    stale_rules = list(SceneConfig().pipeline_critical_rules)
    cfg = load_scene_from_data(
        {
            "pipeline": [
                "page_setup",
                "table_format",
                "formula_convert",
                "formula_to_table",
                "section_format",
                "citation_link",
                "validation",
            ],
            "formula_convert": {"enabled": True},
            "formula_to_table": {"enabled": True},
            "citation_link": {"enabled": True},
            "pipeline_critical_rules": stale_rules,
        }
    )

    assert cfg.pipeline_critical_rules == cfg.pipeline


def test_scene_manager_loads_formula_table_visual_config_payload():
    cfg = load_scene_from_data(
        {
            "formula_table": {
                "formula_font_name": "XITS Math",
                "formula_font_size_pt": 11.5,
                "formula_font_size_display": "11.5pt",
                "formula_line_spacing": 1.25,
                "formula_space_before_pt": 2.0,
                "formula_space_after_pt": 3.0,
                "block_alignment": "right",
                "table_alignment": "left",
                "formula_cell_alignment": "left",
                "number_alignment": "center",
                "number_font_name": "Arial",
                "number_font_size_pt": 9.0,
                "number_font_size_display": "小五",
                "auto_shrink_number_column": False,
            }
        }
    )

    assert cfg.formula_table.formula_font_name == "XITS Math"
    assert cfg.formula_table.formula_font_size_pt == 11.5
    assert cfg.formula_table.formula_font_size_display == "11.5pt"
    assert cfg.formula_table.formula_line_spacing == 1.25
    assert cfg.formula_table.formula_space_before_pt == 2.0
    assert cfg.formula_table.formula_space_after_pt == 3.0
    assert cfg.formula_table.block_alignment == "right"
    assert cfg.formula_table.table_alignment == "left"
    assert cfg.formula_table.formula_cell_alignment == "left"
    assert cfg.formula_table.number_alignment == "center"
    assert cfg.formula_table.number_font_name == "Arial"
    assert cfg.formula_table.number_font_size_pt == 9.0
    assert cfg.formula_table.number_font_size_display == "小五"
    assert cfg.formula_table.auto_shrink_number_column is False


def test_equation_table_numbering_does_not_advance_on_failed_row_write():
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "A = B"
    tbl.cell(0, 1).text = "occupied"
    tbl.cell(1, 0).text = "C = D"
    tbl.cell(1, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(1, 1).text.strip() == "(1.1)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("auto_number") == 1


def test_equation_table_format_skips_explanatory_equals_table():
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "A = B"
    tbl.cell(0, 1).text = "变量关系说明"
    tbl.cell(1, 0).text = "C = D"
    tbl.cell(1, 1).text = "进一步解释"

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text == "变量关系说明"
    assert tbl.cell(1, 1).text == "进一步解释"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 0
    assert stats.get("auto_number") == 0


def test_equation_table_format_skips_single_row_reaction_explanation_table():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "A + B -> C"
    tbl.cell(0, 1).text = "反应条件"

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text == "反应条件"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 0
    assert stats.get("auto_number") == 0


def test_equation_table_format_skips_english_explanatory_equals_table():
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "F = ma"
    tbl.cell(0, 1).text = "force relation"
    tbl.cell(1, 0).text = "E = mc^2"
    tbl.cell(1, 1).text = "energy relation"

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text == "force relation"
    assert tbl.cell(1, 1).text == "energy relation"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 0
    assert stats.get("auto_number") == 0


def test_equation_table_format_skips_three_column_formula_explanation_table():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=3)
    tbl.cell(0, 0).text = "x^2+y^2=z^2"
    tbl.cell(0, 1).text = ""
    tbl.cell(0, 2).text = "where z is hypotenuse"

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 2).text == "where z is hypotenuse"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 0
    assert stats.get("auto_number") == 0


def test_equation_table_format_keeps_blank_right_slot_formula_table():
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "F = ma"
    tbl.cell(0, 1).text = ""
    tbl.cell(1, 0).text = "E = mc^2"
    tbl.cell(1, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text.strip() == "(1.1)"
    assert tbl.cell(1, 1).text.strip() == "(1.2)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("auto_number") == 2


def test_equation_table_format_uses_configured_hyphen_separator():
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "F = ma"
    tbl.cell(0, 1).text = ""
    tbl.cell(1, 0).text = "E = mc^2"
    tbl.cell(1, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    cfg.equation_table_format.numbering_format = "chapter-seq"
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text.strip() == "(1-1)"
    assert tbl.cell(1, 1).text.strip() == "(1-2)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("auto_number") == 2
    assert stats.get("renumber") == 0


def test_equation_table_format_uses_configured_plain_sequence():
    doc = Document()
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "F = ma"
    tbl.cell(0, 1).text = ""
    tbl.cell(1, 0).text = "E = mc^2"
    tbl.cell(1, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    cfg.equation_table_format.numbering_format = "seq"
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text.strip() == "(1)"
    assert tbl.cell(1, 1).text.strip() == "(2)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("auto_number") == 2


def test_equation_table_format_keeps_three_column_manual_number_slot():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=3)
    tbl.cell(0, 0).text = "x^2+y^2=z^2"
    tbl.cell(0, 1).text = ""
    tbl.cell(0, 2).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 2).text.strip() == "(1.1)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("auto_number") == 1


def test_equation_table_format_detects_bare_latex_fraction_table():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = r"\frac{a}{b}"
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text.strip() == "(1.1)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("auto_number") == 1


def test_equation_table_format_detects_plain_text_script_formula_table():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "x_i^2+y_i^2"
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text.strip() == "(1.1)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("auto_number") == 1


def test_formula_style_couples_with_hand_authored_latex_equation_table():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = r"\sqrt{x}"
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    left_para = tbl.cell(0, 0).paragraphs[0]
    right_para = tbl.cell(0, 1).paragraphs[0]
    assert right_para.text.strip() == "(1.1)"
    assert left_para.runs[0].font.name == "Cambria Math"
    assert left_para.runs[0].font.size.pt == 12.0

    style_stats = _formula_stats(context, "formula_style")
    assert style_stats.get("converted") == 2
    assert style_stats.get("styled_tables") == 1
    assert style_stats.get("styled_cells", 0) >= 1


def test_formula_style_only_touches_formula_paragraph_inside_equation_table_cell():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    left_para = tbl.cell(0, 0).paragraphs[0]
    left_para.text = r"\sqrt{x}"
    note_para = tbl.cell(0, 0).add_paragraph("note text")
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    assert left_para.runs[0].font.name == "Cambria Math"
    assert left_para.runs[0].font.size.pt == 12.0
    assert note_para.runs[0].font.name is None
    assert note_para.runs[0].font.size is None


def test_formula_pipeline_end_to_end_recovers_math_family_formula_to_numbered_styled_table(
        tmp_path, monkeypatch):
    monkeypatch.setenv("DOCX_DISABLE_FIELD_REFRESH", "1")

    doc_path = tmp_path / "escaped_math_family_formula.docx"
    doc = Document()
    doc.add_paragraph("ESC_11111111athcal{F}(x)=ESC_22222222athtt{x}_0")
    doc.save(str(doc_path))

    cfg = SceneConfig()
    cfg.pipeline = [
        "formula_convert",
        "formula_to_table",
        "equation_table_format",
        "formula_style",
    ]
    cfg.formula_convert.enabled = True
    cfg.formula_to_table.enabled = True
    cfg.equation_table_format.enabled = True
    cfg.formula_style.enabled = True
    cfg.output.final_docx = False
    cfg.output.compare_docx = False
    cfg.output.report_json = False
    cfg.output.report_markdown = False

    result = Pipeline(cfg).run(str(doc_path))

    assert result.success is True
    assert len(result.doc.tables) == 1

    tbl = result.doc.tables[0]
    left_para = tbl.cell(0, 0).paragraphs[0]
    right_para = tbl.cell(0, 1).paragraphs[0]

    assert right_para.text.strip() == "(1.1)"
    assert "oMath" in left_para._p.xml
    assert "ESC_" not in left_para._p.xml
    assert r"\\" not in left_para._p.xml

    font_nodes = left_para._p.xpath(
        ".//*[namespace-uri()='%s' and local-name()='rFonts']" % _W_NS
    )
    size_nodes = left_para._p.xpath(
        ".//*[namespace-uri()='%s' and (local-name()='sz' or local-name()='szCs')]"
        % _W_NS
    )
    assert any(
        node.get(f"{{{_W_NS}}}ascii") == "Cambria Math"
        for node in font_nodes
    )
    assert size_nodes
    assert all(node.get(f"{{{_W_NS}}}val") == "24" for node in size_nodes)


def test_pipeline_range_mode_skips_formula_rules(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCX_DISABLE_FIELD_REFRESH", "1")
    _stub_word_page_spans(monkeypatch, [(1, 1)])

    doc_path = tmp_path / "range_formula.docx"
    doc = Document()
    doc.add_paragraph("x = y")
    doc.save(str(doc_path))

    cfg = SceneConfig()
    cfg.pipeline = [
        "formula_convert",
        "formula_to_table",
        "equation_table_format",
        "formula_style",
    ]
    cfg.formula_convert.enabled = True
    cfg.formula_to_table.enabled = True
    cfg.equation_table_format.enabled = True
    cfg.formula_style.enabled = True
    cfg.format_scope.mode = "manual"
    cfg.format_scope.page_ranges_text = "1-1"
    cfg.output.final_docx = False
    cfg.output.compare_docx = False
    cfg.output.report_json = False
    cfg.output.report_markdown = False

    result = Pipeline(cfg).run(str(doc_path))
    assert result.success is True

    skipped_formula_rules = {
        rec.rule_name
        for rec in result.tracker.records
        if rec.target == "range_scope" and rec.change_type == "skip"
    }
    assert {
        "formula_convert",
        "formula_to_table",
        "equation_table_format",
        "formula_style",
    }.issubset(skipped_formula_rules)


def test_pipeline_runtime_syncs_optional_enabled_steps_for_raw_scene_config():
    cfg = SceneConfig()
    cfg.whitespace_normalize.enabled = True
    cfg.formula_convert.enabled = True
    cfg.formula_to_table.enabled = True

    step_names = [step.name for step in Pipeline(cfg).steps]

    assert "whitespace_normalize" in step_names
    assert step_names.index("whitespace_normalize") == step_names.index("md_cleanup") + 1
    assert "formula_convert" in step_names
    assert step_names.index("formula_convert") == step_names.index("table_format") + 1
    assert step_names.index("formula_to_table") == step_names.index("formula_convert") + 1
    assert "citation_link" in step_names
    assert step_names.index("citation_link") > step_names.index("formula_to_table")
    assert step_names.index("citation_link") < step_names.index("header_footer")


def test_pipeline_strict_mode_marks_auto_managed_formula_failures_as_failed():
    cfg = SceneConfig()
    cfg.formula_convert.enabled = True

    pipeline = Pipeline(cfg)
    pipeline.tracker.record(
        rule_name="formula_convert",
        target="summary",
        section="formula",
        change_type="error",
        before="",
        after="",
        paragraph_index=-1,
        success=False,
        failure_reason="boom",
    )

    result = pipeline._finalize_result(doc=None, original_doc=None, output_paths={})

    assert "formula_convert" in pipeline._critical_rules()
    assert result.success is False
    assert result.status == "failed"


def test_formula_convert_skips_fallback_literal_conversion():
    doc = Document()
    para = doc.add_paragraph("a=b@c")

    cfg = SceneConfig()
    cfg.formula_convert.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaConvertRule().apply(doc, cfg, tracker, context)

    assert "oMath" not in para._p.xml
    stats = _formula_stats(context, "formula_convert")
    assert stats.get("converted") == 0
    assert stats.get("skipped_low_confidence") == 1
    low_conf = context.get("formula_runtime", {}).get("low_confidence", [])
    assert any(item.get("reason") == "fallback_literal_conversion" for item in low_conf)


def test_formula_convert_escape_placeholder_noise_does_not_flow_into_table():
    doc = Document()
    para = doc.add_paragraph("x=ESC_badfrac{1}{2}")

    cfg = SceneConfig()
    cfg.formula_convert.enabled = True
    cfg.formula_to_table.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaConvertRule().apply(doc, cfg, tracker, context)
    FormulaToTableRule().apply(doc, cfg, tracker, context)

    assert "oMath" not in para._p.xml
    assert len(doc.tables) == 0

    convert_stats = _formula_stats(context, "formula_convert")
    assert convert_stats.get("converted") == 0
    assert convert_stats.get("skipped_low_confidence") == 1
    low_conf = context.get("formula_runtime", {}).get("low_confidence", [])
    assert any(item.get("reason") == "escape_placeholder_noise" for item in low_conf)

    to_table_stats = _formula_stats(context, "formula_to_table")
    assert to_table_stats.get("matched") == 0
    assert to_table_stats.get("converted") == 0


def test_formula_convert_allows_recovered_escape_placeholder_formula():
    doc = Document()
    para = doc.add_paragraph("a2+b2ESC_04e74865qrt{a^2+b^2}a2+b2\u200b")

    cfg = SceneConfig()
    cfg.formula_convert.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaConvertRule().apply(doc, cfg, tracker, context)

    assert "oMath" in para._p.xml
    stats = _formula_stats(context, "formula_convert")
    assert stats.get("converted") == 1
    assert stats.get("skipped_low_confidence") == 0


def test_formula_convert_allows_rendered_segment_extracted_formula():
    doc = Document()
    para = doc.add_paragraph("\u3000arcsin\u2061xESC_bd6e0b20rcsin xarcsinx")

    cfg = SceneConfig()
    cfg.formula_convert.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaConvertRule().apply(doc, cfg, tracker, context)

    assert "oMath" in para._p.xml
    stats = _formula_stats(context, "formula_convert")
    assert stats.get("converted") == 1
    assert stats.get("skipped_low_confidence") == 0


def test_formula_to_table_requires_word_native_when_convert_enabled():
    doc = Document()
    doc.add_paragraph("a+b=c")

    cfg = SceneConfig()
    cfg.formula_to_table.enabled = True

    tracker = ChangeTracker()
    context = {"formula_runtime": {"convert_enabled": True}}
    FormulaToTableRule().apply(doc, cfg, tracker, context)

    assert len(doc.tables) == 0
    stats = _formula_stats(context, "formula_to_table")
    assert stats.get("matched") == 0
    assert stats.get("converted") == 0


def test_formula_style_skips_plain_text_when_convert_enabled():
    doc = Document()
    para = doc.add_paragraph("a+b=c")

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {"formula_runtime": {"convert_enabled": True}}
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    stats = _formula_stats(context, "formula_style")
    assert stats.get("matched") == 0
    assert stats.get("converted") == 0
    assert para.runs[0].font.name is None
    assert not any(
        rec.rule_name == "formula_style"
        and rec.change_type == "style"
        and rec.paragraph_index == 0
        for rec in tracker.records
    )


def test_formula_style_applies_font_and_size_to_word_native_omml():
    doc = Document()
    para = doc.add_paragraph()
    node = FormulaNode(
        kind="equation",
        payload={"latex": r"\frac{\pi^2}{6}"},
        source_type="latex",
        confidence=0.95,
    )
    outcome = convert_formula_node(node, "word_native", block=True)
    assert outcome.success is True
    para._p.append(outcome.omml_element)

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    font_nodes = para._p.xpath(
        ".//*[namespace-uri()='%s' and local-name()='rFonts']" % _W_NS
    )
    size_nodes = para._p.xpath(
        ".//*[namespace-uri()='%s' and (local-name()='sz' or local-name()='szCs')]"
        % _W_NS
    )

    assert font_nodes
    assert any(
        node.get(f"{{{_W_NS}}}ascii") == "Cambria Math"
        for node in font_nodes
    )
    assert size_nodes
    assert all(node.get(f"{{{_W_NS}}}val") == "24" for node in size_nodes)

    stats = _formula_stats(context, "formula_style")
    assert stats.get("converted") == 1


def test_formula_style_uses_custom_formula_table_visual_config_for_formula_paragraph():
    doc = Document()
    para = doc.add_paragraph()
    _set_native_formula(para, r"\frac{\pi^2}{6}")

    cfg = SceneConfig()
    cfg.formula_style.enabled = True
    cfg.formula_table.formula_font_name = "XITS Math"
    cfg.formula_table.formula_font_size_pt = 11.5
    cfg.formula_table.formula_line_spacing = 1.5
    cfg.formula_table.formula_space_before_pt = 2.0
    cfg.formula_table.formula_space_after_pt = 3.0
    cfg.formula_table.block_alignment = "right"

    FormulaStyleRule().apply(doc, cfg, ChangeTracker(), {})

    font_nodes = para._p.xpath(
        ".//*[namespace-uri()='%s' and local-name()='rFonts']" % _W_NS
    )
    size_nodes = para._p.xpath(
        ".//*[namespace-uri()='%s' and (local-name()='sz' or local-name()='szCs')]"
        % _W_NS
    )

    assert para.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert abs(float(para.paragraph_format.line_spacing) - 1.5) < 1e-6
    assert abs(float(para.paragraph_format.space_before.pt) - 2.0) < 0.05
    assert abs(float(para.paragraph_format.space_after.pt) - 3.0) < 0.05
    assert font_nodes
    assert any(
        node.get(f"{{{_W_NS}}}ascii") == "XITS Math"
        for node in font_nodes
    )
    assert size_nodes
    assert all(node.get(f"{{{_W_NS}}}val") == "23" for node in size_nodes)


def test_formula_style_uses_custom_formula_table_visual_config_for_equation_tables():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    _set_native_formula(tbl.cell(0, 0).paragraphs[0], r"x=1")
    tbl.cell(0, 1).text = "(1.1)"

    cfg = SceneConfig()
    cfg.formula_style.enabled = True
    cfg.formula_table.table_alignment = "right"
    cfg.formula_table.formula_cell_alignment = "left"
    cfg.formula_table.number_alignment = "center"
    cfg.formula_table.number_font_name = "Arial"
    cfg.formula_table.number_font_size_pt = 9.0

    FormulaStyleRule().apply(doc, cfg, ChangeTracker(), {})

    assert _table_alignment_values(tbl) == ["right"]
    assert "left" in _cell_alignment_values(tbl.cell(0, 0))
    assert "center" in _cell_alignment_values(tbl.cell(0, 1))

    number_font_nodes = tbl.cell(0, 1)._tc.xpath(
        ".//*[namespace-uri()='%s' and local-name()='rFonts']" % _W_NS
    )
    number_size_nodes = tbl.cell(0, 1)._tc.xpath(
        ".//*[namespace-uri()='%s' and (local-name()='sz' or local-name()='szCs')]"
        % _W_NS
    )
    assert number_font_nodes
    assert any(
        node.get(f"{{{_W_NS}}}ascii") == "Arial"
        for node in number_font_nodes
    )
    assert number_size_nodes
    assert all(node.get(f"{{{_W_NS}}}val") == "18" for node in number_size_nodes)


def test_formula_style_marks_plain_text_formula_as_visual_only():
    doc = Document()
    para = doc.add_paragraph("a^2+b^2=c^2")

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    stats = _formula_stats(context, "formula_style")
    assert stats.get("visual_only") == 1
    assert stats.get("native_styled") == 0
    assert para.runs[0].font.name == "Cambria Math"
    assert any(
        rec.rule_name == "formula_style"
        and rec.change_type == "style"
        and rec.paragraph_index == 0
        and "视觉样式" in rec.after
        for rec in tracker.records
    )


def test_formula_style_normalizes_plain_text_formula_font_slots_and_hint():
    doc = Document()
    para = doc.add_paragraph("η=ΔG×nt×S×I (2.1)")
    run = para.runs[0]
    rpr = _ensure_run_rpr(run)
    rfonts = etree.SubElement(rpr, qn("w:rFonts"))
    rfonts.set(qn("w:eastAsia"), "等线")
    rfonts.set(qn("w:hint"), "eastAsia")
    rfonts.set(qn("w:eastAsiaTheme"), "minorEastAsia")

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    FormulaStyleRule().apply(doc, cfg, ChangeTracker(), {})

    rfonts = run._element.find(".//" + qn("w:rFonts"))
    assert rfonts is not None
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        assert rfonts.get(qn(f"w:{key}")) == "Cambria Math"
    assert rfonts.get(qn("w:hint")) == "default"
    assert rfonts.get(qn("w:eastAsiaTheme")) is None


def test_formula_style_normalizes_omml_run_font_slots_hint_and_theme():
    doc = Document()
    para = doc.add_paragraph()
    _set_native_formula(para, r"\eta=\Delta G\times nt\times S\times I")

    m_runs = para._p.xpath(
        ".//*[namespace-uri()='%s' and local-name()='r']" % _M_NS
    )
    assert m_runs
    for m_run in m_runs:
        m_rpr = m_run.find(f"{{{_M_NS}}}rPr")
        if m_rpr is None:
            m_rpr = etree.SubElement(m_run, f"{{{_M_NS}}}rPr")
        w_rpr = m_rpr.find(qn("w:rPr"))
        if w_rpr is None:
            w_rpr = etree.SubElement(m_rpr, qn("w:rPr"))
        rfonts = w_rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = etree.SubElement(w_rpr, qn("w:rFonts"))
        rfonts.set(qn("w:eastAsia"), "等线")
        rfonts.set(qn("w:hint"), "eastAsia")
        rfonts.set(qn("w:eastAsiaTheme"), "minorEastAsia")

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    FormulaStyleRule().apply(doc, cfg, ChangeTracker(), {})

    font_nodes = para._p.xpath(
        ".//*[namespace-uri()='%s' and local-name()='rFonts']" % _W_NS
    )
    assert font_nodes
    for node in font_nodes:
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            assert node.get(qn(f"w:{key}")) == "Cambria Math"
        assert node.get(qn("w:hint")) == "default"
        assert node.get(qn("w:eastAsiaTheme")) is None
        assert node.get(qn("w:asciiTheme")) is None
        assert node.get(qn("w:hAnsiTheme")) is None


def test_equation_table_format_refreshes_heading_context_after_formula_to_table():
    doc = Document()
    h1 = doc.add_paragraph("Chapter 1")
    h1.style = "Heading 1"
    p1 = doc.add_paragraph()
    _set_native_formula(p1, r"x=1")

    h2 = doc.add_paragraph("Chapter 2")
    h2.style = "Heading 1"
    p2 = doc.add_paragraph()
    _set_native_formula(p2, r"y=2")

    cfg = SceneConfig()
    cfg.formula_to_table.enabled = True
    cfg.equation_table_format.enabled = True

    tracker = ChangeTracker()
    context = _build_heading_context(doc, cfg)

    FormulaToTableRule().apply(doc, cfg, tracker, context)
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert len(doc.tables) == 2
    assert doc.tables[0].cell(0, 1).text.strip() == "(1.1)"
    assert doc.tables[1].cell(0, 1).text.strip() == "(2.1)"


def test_equation_table_format_refreshes_heading_context_after_formula_convert():
    doc = Document()
    h1 = doc.add_paragraph("Chapter 1")
    h1.style = "Heading 1"
    doc.add_paragraph(r"\[")
    doc.add_paragraph(r"x=1")
    doc.add_paragraph(r"\]")

    h2 = doc.add_paragraph("Chapter 2")
    h2.style = "Heading 1"

    tbl = doc.add_table(rows=1, cols=2)
    _set_native_formula(tbl.cell(0, 0).paragraphs[0], r"y=2")
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.formula_convert.enabled = True
    cfg.equation_table_format.enabled = True

    tracker = ChangeTracker()
    context = _build_heading_context(doc, cfg)

    FormulaConvertRule().apply(doc, cfg, tracker, context)
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 1).text.strip() == "(2.1)"


def test_equation_table_format_expands_number_column_without_formula_style():
    doc = Document()
    para = doc.add_paragraph()
    _set_native_formula(para, r"x=1")

    cfg = SceneConfig()
    cfg.formula_to_table.enabled = True
    cfg.equation_table_format.enabled = True

    tracker = ChangeTracker()
    context = {}

    FormulaToTableRule().apply(doc, cfg, tracker, context)
    before_widths = _table_grid_widths(doc.tables[0])

    EquationTableFormatRule().apply(doc, cfg, tracker, context)
    after_widths = _table_grid_widths(doc.tables[0])

    assert doc.tables[0].cell(0, 1).text.strip() == "(1.1)"
    assert before_widths[-1] < after_widths[-1]
    assert after_widths[-1] >= 850


def test_formula_style_styles_plain_text_formula_cell_when_convert_enabled():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "a^2+b^2=c^2"
    tbl.cell(0, 1).text = "(1.1)"

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {"formula_runtime": {"convert_enabled": True}}
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    para = tbl.cell(0, 0).paragraphs[0]
    assert para.runs[0].font.name == "Cambria Math"
    assert para.runs[0].font.size.pt == 12.0


def test_formula_style_styles_bare_latex_formula_cell_when_convert_enabled():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = r"\frac{a}{b}"
    tbl.cell(0, 1).text = "(1.1)"

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {"formula_runtime": {"convert_enabled": True}}
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    para = tbl.cell(0, 0).paragraphs[0]
    assert para.runs[0].font.name == "Cambria Math"
    assert para.runs[0].font.size.pt == 12.0


def test_formula_style_applies_omml_style_inside_mixed_paragraph_without_touching_text_runs():
    doc = Document()
    para = doc.add_paragraph("由此可得 ")
    outcome = convert_formula_node(
        FormulaNode(
            kind="equation",
            payload={"latex": r"x^2+y^2=z^2"},
            source_type="latex",
            confidence=0.95,
        ),
        "word_native",
        block=False,
    )
    assert outcome.success is True
    run = para.add_run()
    run._r.append(deepcopy(outcome.omml_element))
    para.add_run(" 成立")

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    assert para.alignment is None
    text_runs = [r for r in para.runs if "oMath" not in r._r.xml]
    assert any((r.text or "").strip() for r in text_runs)
    assert all(r.font.name is None for r in text_runs if (r.text or "").strip())

    font_nodes = para._p.xpath(
        ".//*[namespace-uri()='%s' and local-name()='rFonts']" % _W_NS
    )
    assert any(
        node.get(f"{{{_W_NS}}}ascii") == "Cambria Math"
        for node in font_nodes
    )
    assert not any(
        rec.rule_name == "formula_style"
        and rec.change_type == "skip"
        and rec.paragraph_index == 0
        for rec in tracker.records
    )


def test_formula_style_splits_multi_text_omml_run_for_semantic_style():
    xml = """
    <m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
             xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <m:r>
        <m:t>sin</m:t>
        <m:t>(x)</m:t>
      </m:r>
    </m:oMath>
    """

    doc = Document()
    para = doc.add_paragraph()
    para._p.append(etree.fromstring(xml.encode("utf-8")))

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    m_runs = para._p.xpath(
        ".//*[namespace-uri()='%s' and local-name()='r']" % _M_NS
    )
    assert len(m_runs) == 3

    texts = [
        run.xpath(
            "./*[namespace-uri()='%s' and local-name()='t']/text()" % _M_NS
        )
        for run in m_runs
    ]
    italics = [
        run.xpath(
            "./*[namespace-uri()='%s' and local-name()='rPr']"
            "/*[namespace-uri()='%s' and local-name()='rPr']"
            "/*[namespace-uri()='%s' and local-name()='i']/@*[local-name()='val']"
            % (_M_NS, _W_NS, _W_NS)
        )
        for run in m_runs
    ]

    assert texts[0] == ["sin("]
    assert texts[1] == ["x"]
    assert texts[2] == [")"]
    assert italics[0] == ["0"]
    assert italics[1] == ["1"]
    assert italics[2] == ["0"]


def test_formula_style_recognizes_lim_and_det_as_upright_functions():
    lim_segments = _semantic_segments_for_text(
        "lim x",
        in_sub_or_sup=False,
        base_of_superscript=False,
    )
    det_segments = _semantic_segments_for_text(
        "det A",
        in_sub_or_sup=False,
        base_of_superscript=False,
    )

    assert lim_segments[0] == {"text": "lim", "italic": False, "bold": False}
    assert lim_segments[2] == {"text": "x", "italic": True, "bold": False}
    assert det_segments[0] == {"text": "det", "italic": False, "bold": False}
    assert det_segments[2] == {"text": "A", "italic": True, "bold": False}


def test_formula_style_recognizes_math_alphabet_families_as_upright():
    samples = [
        (r"\mathbb R", {"text": "R", "italic": False, "bold": False}),
        (r"\mathcal F", {"text": "F", "italic": False, "bold": False}),
        (r"\mathtt x", {"text": "x", "italic": False, "bold": False}),
        (r"\mathsf A", {"text": "A", "italic": False, "bold": False}),
        (r"\mathfrak g", {"text": "g", "italic": False, "bold": False}),
        (r"\mathscr L", {"text": "L", "italic": False, "bold": False}),
    ]

    for text, expected_symbol in samples:
        segments = _semantic_segments_for_text(
            text,
            in_sub_or_sup=False,
            base_of_superscript=False,
        )
        assert expected_symbol in segments


def test_formula_style_recognizes_variant_greek_commands_as_italic():
    segments = _semantic_segments_for_text(
        r"\varepsilon+\varphi=\vartheta",
        in_sub_or_sup=False,
        base_of_superscript=False,
    )
    assert {"text": r"\varepsilon", "italic": True, "bold": False} in segments
    assert {"text": r"\varphi", "italic": True, "bold": False} in segments
    assert {"text": r"\vartheta", "italic": True, "bold": False} in segments


def test_formula_style_preserves_explicit_omml_msty_without_conflicting_w_toggles():
    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    for latex, expected_style, expected_scr in (
        (r"\mathbf A", "b", None),
        (r"\boldsymbol x", "bi", None),
        (r"\mathrm d", "p", None),
        (r"\mathsf A", "p", "sans-serif"),
    ):
        doc = Document()
        para = doc.add_paragraph()
        outcome = convert_formula_node(
            FormulaNode(
                kind="equation",
                payload={"latex": latex},
                source_type="latex",
                confidence=0.95,
            ),
            "word_native",
            block=False,
        )
        assert outcome.success is True
        assert outcome.omml_element is not None
        para._p.append(deepcopy(outcome.omml_element))

        FormulaStyleRule().apply(doc, cfg, ChangeTracker(), {})

        m_runs = para._p.xpath(
            ".//*[namespace-uri()='%s' and local-name()='r']" % _M_NS
        )
        assert len(m_runs) == 1
        m_run = m_runs[0]

        sty_vals = m_run.xpath(
            "./*[namespace-uri()='%s' and local-name()='rPr']"
            "/*[namespace-uri()='%s' and local-name()='sty']"
            "/@*[local-name()='val']"
            % (_M_NS, _M_NS)
        )
        assert sty_vals == [expected_style]

        scr_vals = m_run.xpath(
            "./*[namespace-uri()='%s' and local-name()='rPr']"
            "/*[namespace-uri()='%s' and local-name()='scr']"
            "/@*[local-name()='val']"
            % (_M_NS, _M_NS)
        )
        if expected_scr is None:
            assert scr_vals == []
        else:
            assert scr_vals == [expected_scr]

        toggle_nodes = m_run.xpath(
            "./*[namespace-uri()='%s' and local-name()='rPr']"
            "/*[namespace-uri()='%s' and local-name()='rPr']"
            "/*[local-name()='i' or local-name()='iCs' or local-name()='b' or local-name()='bCs']"
            % (_M_NS, _W_NS)
        )
        assert toggle_nodes == []


def test_parse_document_formulas_skips_figure_caption_with_chem_markers():
    doc = Document()
    doc.add_paragraph("图2.12CQDs和BD−CQDs体系中的DMPO−⋅OH信号\t(2.3)")

    parsed = parse_document_formulas(doc)

    assert parsed.total == 0


def test_parse_document_formulas_realcase_skips_prose_omml_equation_table_cell():
    doc = Document(str(_realcase_c3n4_docx_path()))

    parsed = parse_document_formulas(doc)

    assert "table[1] r1c1 p1" not in {occ.location for occ in parsed.occurrences}


def test_equation_table_format_rejects_caption_text_with_right_number_cell():
    doc = Document()
    doc.add_paragraph("anchor")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "图2.12CQDs和BD−CQDs体系中的DMPO−⋅OH信号"
    tbl.cell(0, 1).text = "(2.3)"

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 1).text == "(2.3)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 0
    assert stats.get("auto_number") == 0
    assert stats.get("renumber") == 0


def test_equation_table_format_detects_same_cell_tab_number_tail_for_renumber():
    doc = Document()
    doc.add_paragraph("anchor")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "Eads = E(*O2/*H+) – E(*) − E(O2/H+)\t(2-4)"
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 0).text.strip().endswith("(2.1)")
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("auto_number") == 0
    assert stats.get("renumber") == 1


def test_equation_table_format_detects_same_cell_tab_number_tail_for_plain_sequence():
    doc = Document()
    doc.add_paragraph("anchor")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "Eads = E(*O2/*H+) – E(*) − E(O2/H+)\t(4)"
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    cfg.equation_table_format.numbering_format = "seq"
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 0).text.strip().endswith("(1)")
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("auto_number") == 0
    assert stats.get("renumber") == 1


def test_equation_table_format_rejects_chemical_process_sentence_with_same_cell_number_tail():
    doc = Document()
    doc.add_paragraph("anchor")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(
        0, 0
    ).text = (
        "将0.2gBA、BrPE、NaBH4或DMP与20mgCQDs−起溶于20mLCCl4/H2O（体积比为1:1）混合液中。"
        "搅拌10h后，连续透析36h，并将透析后的溶液冷冻干燥后分别得到a−CQDs（BA−CQDs）、"
        "r−CQDs（BrPE−CQDs）、s−CQDs（NaBH4−CQDs）和d−CQDs（DMP−CQDs）粉末。\t(2.1)"
    )
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 0).text.strip().endswith("(2.1)")
    assert tbl.cell(0, 1).text == ""
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 0
    assert stats.get("auto_number") == 0
    assert stats.get("renumber") == 0


def test_equation_table_format_rejects_chemical_sample_name_with_same_cell_number_tail():
    doc = Document()
    doc.add_paragraph("anchor")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "BA-CQDs/NaBH4\t(2.1)"
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 0).text.strip().endswith("(2.1)")
    assert tbl.cell(0, 1).text == ""
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 0
    assert stats.get("auto_number") == 0
    assert stats.get("renumber") == 0


def test_formula_convert_converts_mathtype_formula_inside_equation_table():
    doc = Document()
    doc.add_paragraph("anchor")
    tbl = doc.add_table(rows=1, cols=2)
    _add_mathtype_formula(tbl.cell(0, 0).paragraphs[0], "x_i^2+y_i^2=z_i^2\t(2-4)")
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.formula_convert.enabled = True
    cfg.equation_table_format.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaConvertRule().apply(doc, cfg, tracker, context)
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert "oMath" in tbl.cell(0, 0)._tc.xml
    assert "(2-4)" not in tbl.cell(0, 0).text
    assert tbl.cell(0, 1).text.strip() == "(1.1)"

    convert_stats = context.get("formula_runtime", {}).get("stats", {}).get("formula_convert", {})
    eq_stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert convert_stats.get("converted") == 1
    assert convert_stats.get("skipped_dependency") == 0
    assert eq_stats.get("tables") == 1
    assert eq_stats.get("auto_number") == 1


def test_equation_table_format_handles_leading_table_without_body_anchor():
    doc = Document()
    tbl = doc.add_table(rows=1, cols=2)
    _add_mathtype_formula(tbl.cell(0, 0).paragraphs[0], "x_i^2+y_i^2=z_i^2\t(2-4)")
    tbl.cell(0, 1).text = ""

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert tbl.cell(0, 0).text == "x_i^2+y_i^2=z_i^2"
    assert tbl.cell(0, 1).text.strip() == "(2.1)"
    stats = context.get("formula_runtime", {}).get("stats", {}).get("equation_table_format", {})
    assert stats.get("tables") == 1
    assert stats.get("renumber") == 1


def test_formula_style_reports_styled_number_cells_for_equation_tables():
    doc = Document()
    doc.add_paragraph("anchor")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "x=1"
    tbl.cell(0, 1).text = "(1.1)"

    cfg = SceneConfig()
    cfg.formula_style.enabled = True

    tracker = ChangeTracker()
    context = {}
    FormulaStyleRule().apply(doc, cfg, tracker, context)

    stats = context.get("formula_runtime", {}).get("stats", {}).get("formula_style", {})
    assert stats.get("styled_tables") == 1
    assert stats.get("styled_number_cells") == 1


def test_equation_table_format_realcase_skips_false_prose_omml_table_and_unifies_numbers():
    doc = Document(str(_realcase_c3n4_docx_path()))

    cfg = SceneConfig()
    cfg.equation_table_format.enabled = True
    tracker = ChangeTracker()
    context = {}
    EquationTableFormatRule().apply(doc, cfg, tracker, context)

    assert table_helpers._is_equation_table(doc.tables[0]._tbl) is False
    assert doc.tables[0].cell(0, 1).text.strip() == "(2.1)"
    assert doc.tables[1].cell(0, 1).text.strip() == "(2.1)"
    assert doc.tables[2].cell(0, 1).text.strip() == "(2.2)"
    assert doc.tables[2].cell(1, 1).text.strip() == "(2.3)"
    assert doc.tables[3].cell(0, 1).text.strip() == "(2.4)"
    assert doc.tables[4].cell(0, 1).text.strip() == "(2.5)"

    stats = _formula_stats(context, "equation_table_format")
    assert stats == {"tables": 4, "auto_number": 0, "renumber": 5}
