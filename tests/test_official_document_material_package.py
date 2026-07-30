import json
from pathlib import Path
from shutil import copytree

from docx import Document

from src.config.library import ensure_config_library, load_scene_from_library
from src.config import material_package_library
from src.config.entity import ENTITY_PACKAGE_VERSION, load_entity_archive
from src.config.material_context import MaterialExecutionContext
from src.config.official_document_profiles import (
    OFFICIAL_PLAN_ENTRY_BUILTIN,
    get_official_document_assembly_contract,
    list_common_official_document_profiles,
    list_official_document_plan_entry_decisions,
)
import src.shared.engine.official_document_material_package as material_package_module
from src.shared.engine.official_document_material_package import (
    PACKAGE_VERSION,
    export_official_document_material_package,
    export_official_document_material_package_to_user_library,
    get_official_document_material_package_sample,
    list_official_document_material_package_samples,
    load_official_document_material_package,
    load_official_document_material_package_sample,
    load_official_document_material_batch,
    load_official_document_material_table,
)
from src.shared.engine.official_document_assembly import assemble_official_document_docx
from src.shared.engine.material_timeline import default_timeline_plan
from src.shared.engine.official_document_batch_history import (
    build_official_document_batch_retry_selection,
    list_official_document_batch_history,
)
from src.ui.adapters.config_selector_models import material_package_sample_selector_options
from src.ui.bridge import PanelBridge
from src.ui.panels.assets_panel import AssetsPanel
from src.services.production_runtime.execution_runtime import WorkbenchBatchProductionRunner


def _all_docx_text(path):
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_official_document_material_package_exports_and_loads_notice_context(tmp_path):
    context = MaterialExecutionContext(
        profile_id="official:notice",
        profile_name="Notice",
        entity_data={
            "title": "Notice title",
            "body": "Notice body",
            "organization": "Archive Office",
            "document_no": "A-2026-1",
            "issue_date": "2026-07-10",
            "recipient": "Departments",
        },
    )
    package_path = tmp_path / "notice_material.json"

    exported = export_official_document_material_package(
        context,
        package_path,
        profile_id="notice",
    )
    loaded = load_official_document_material_package(package_path)

    assert exported.status == "ok"
    assert exported.profile_id == "notice"
    assert exported.material_schema_ids == ("official_document_v1",)
    assert exported.missing_required_fields == ()
    assert loaded.status == "ok"
    assert loaded.context.profile_id == "official:notice"
    assert loaded.context.entity_data["document_type"] == "notice"
    assert loaded.context.entity_data["title"] == "Notice title"
    payload = json.loads(package_path.read_text(encoding="utf-8"))
    assert payload["kind"] == "alavette.material_package"
    assert payload["version"] == ENTITY_PACKAGE_VERSION
    assert payload["mode_id"] == "official"
    assert payload["profiles"][0]["profile_id"] == "official:notice"
    assert payload["profiles"][0]["fields"]["document_type"] == "notice"


def test_official_material_package_honors_removed_builtin_fields(tmp_path):
    context = MaterialExecutionContext(
        profile_id="official:notice",
        field_scopes={"title": "removed"},
        entity_data={
            "body": "通知正文",
            "organization": "档案处",
            "document_no": "档发〔2026〕12号",
            "issue_date": "2026-07-11",
        },
    )

    result = export_official_document_material_package(
        context,
        tmp_path / "removed-title-material.json",
        profile_id="notice",
    )

    assert result.status == "ok"
    assert result.missing_required_fields == ()


def test_official_material_package_requires_explicit_profile_owner_and_writes_nothing(
    tmp_path,
):
    target = tmp_path / "missing-owner.json"
    context = MaterialExecutionContext(
        profile_id="official:notice",
        entity_data={"document_type": "notice", "title": "Notice"},
    )

    result = export_official_document_material_package(context, target)

    assert result.status == "missing_profile"
    assert result.profile_id == ""
    assert result.issues == ("official_profile_id_missing",)
    assert context.profile_id == "official:notice"
    assert context.entity_data["document_type"] == "notice"
    assert not target.exists()


def test_official_material_package_rejects_identity_conflict_without_overwrite(
    tmp_path,
):
    target = tmp_path / "existing.json"
    target.write_bytes(b"existing-package-bytes")
    context = MaterialExecutionContext(
        profile_id="official:notice",
        entity_data={"document_type": "letter", "title": "Conflicting source"},
    )

    result = export_official_document_material_package(
        context,
        target,
        profile_id="notice",
    )

    assert result.status == "profile_conflict"
    assert result.profile_id == "notice"
    assert result.issues == (
        "official_profile_identity_conflict: selected=notice; document_type=letter",
    )
    assert context.entity_data["document_type"] == "letter"
    assert target.read_bytes() == b"existing-package-bytes"


def test_official_document_material_table_loads_horizontal_csv_with_chinese_headers(
    tmp_path,
):
    table_path = tmp_path / "letter_material.csv"
    table_path.write_text(
        "文种,标题,正文,发文机关,发文字号,成文日期,主送机关\n"
        "letter,关于补充材料的函,请补充提交材料。,档案处,档函〔2026〕1号,2026-07-10,项目单位\n",
        encoding="utf-8-sig",
    )

    loaded = load_official_document_material_table(
        table_path,
        profile_id="letter",
    )

    assert loaded.status == "ok"
    assert loaded.profile_id == "letter"
    assert loaded.context.profile_id == "official:letter"
    assert loaded.context.entity_data["document_type"] == "letter"
    assert loaded.context.entity_data["title"] == "关于补充材料的函"
    assert loaded.context.entity_data["document_no"] == "档函〔2026〕1号"
    assert loaded.context.entity_data["recipient"] == "项目单位"
    assert loaded.missing_required_fields == ()
    assert loaded.unknown_fields == ()


def test_official_document_material_table_uses_explicit_selection_only_when_missing(
    tmp_path,
):
    table_path = tmp_path / "selected_notice.csv"
    table_path.write_text(
        "title,body,organization,document_no,issue_date\n"
        "Notice,Body,Office,N-1,2026-07-10\n",
        encoding="utf-8",
    )
    source_bytes = table_path.read_bytes()

    loaded = load_official_document_material_table(
        table_path,
        profile_id="notice",
    )

    assert loaded.status == "ok"
    assert loaded.profile_id == "notice"
    assert loaded.context.entity_data["document_type"] == "notice"
    assert table_path.read_bytes() == source_bytes


def test_official_document_material_table_rejects_missing_and_conflicting_identity(
    tmp_path,
):
    missing_path = tmp_path / "missing_profile.csv"
    missing_path.write_text("title\nNo profile\n", encoding="utf-8")
    conflict_path = tmp_path / "conflicting_profile.csv"
    conflict_path.write_text(
        "document_type,title\nletter,Letter source\n",
        encoding="utf-8",
    )
    conflict_bytes = conflict_path.read_bytes()

    missing = load_official_document_material_table(missing_path)
    conflict = load_official_document_material_table(
        conflict_path,
        profile_id="notice",
    )

    assert missing.status == "missing_profile"
    assert missing.context.entity_data == {"title": "No profile"}
    assert missing.issues == ("official_profile_id_missing",)
    assert conflict.status == "profile_conflict"
    assert conflict.profile_id == "notice"
    assert conflict.context.profile_id == ""
    assert conflict.context.entity_data["document_type"] == "letter"
    assert conflict.issues == (
        "official_profile_identity_conflict: selected=notice; document_type=letter",
    )
    assert conflict_path.read_bytes() == conflict_bytes


def test_official_document_material_table_loads_vertical_minutes_excel(tmp_path):
    from openpyxl import Workbook

    table_path = tmp_path / "minutes_material.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["key", "value"])
    sheet.append(["文种", "minutes"])
    sheet.append(["标题", "档案工作协调会纪要"])
    sheet.append(["正文", "会议研究了归档安排。"])
    sheet.append(["发文机关", "综合办公室"])
    sheet.append(["发文字号", "办纪〔2026〕2号"])
    sheet.append(["成文日期", "2026-07-10"])
    sheet.append(["会议时间", "2026-07-10 09:00"])
    sheet.append(["参会人员", "张三、李四"])
    workbook.save(table_path)
    workbook.close()

    loaded = load_official_document_material_table(table_path)

    assert loaded.status == "ok"
    assert loaded.profile_id == "minutes"
    assert loaded.material_schema_ids == (
        "official_document_v1",
        "administrative_meeting_fields_v1",
    )
    assert loaded.context.entity_data["meeting_date"] == "2026-07-10 09:00"
    assert loaded.context.entity_data["participants"] == "张三、李四"
    assert loaded.unknown_fields == ()


def test_official_document_material_table_rejects_empty_or_unknown_profile_tables(
    tmp_path,
):
    empty_path = tmp_path / "empty.csv"
    empty_path.write_text("key,value\n", encoding="utf-8")
    unknown_path = tmp_path / "unknown.csv"
    unknown_path.write_text(
        "文种,标题\n自定义文种,测试标题\n",
        encoding="utf-8",
    )

    empty = load_official_document_material_table(
        empty_path,
        profile_id="notice",
    )
    unknown = load_official_document_material_table(unknown_path)

    assert empty.status == "invalid_table"
    assert empty.context.entity_data == {}
    assert empty.issues == ("official_material_table_entity_data_empty",)
    assert unknown.status == "unknown_profile"
    assert unknown.profile_id == "自定义文种"
    assert unknown.context.entity_data["title"] == "测试标题"


def test_official_document_material_table_does_not_silently_take_first_batch_row(
    tmp_path,
):
    table_path = tmp_path / "multiple_notices.csv"
    table_path.write_text(
        "document_type,title,body,organization,document_no,issue_date\n"
        "notice,Notice A,Body A,Office,A-1,2026-07-10\n"
        "notice,Notice B,Body B,Office,B-1,2026-07-11\n",
        encoding="utf-8",
    )

    loaded = load_official_document_material_table(table_path)

    assert loaded.status == "multiple_records"
    assert loaded.context.entity_data == {}
    assert loaded.issues == ("official_material_table_multiple_records: 2",)


def test_official_document_material_batch_keeps_task_and_document_type_ids_separate(
    tmp_path,
):
    table_path = tmp_path / "official_batch.csv"
    table_path.write_text(
        "task_id,document_type,title,body,organization,document_no,issue_date\n"
        "task,notice,Notice A,Body A,Office,A-1,2026-07-10\n"
        "task,letter,Letter B,Body B,Office,,2026-07-11\n"
        ",missing,Unknown C,Body C,Office,C-1,2026-07-12\n",
        encoding="utf-8",
    )

    loaded = load_official_document_material_batch(table_path)

    assert loaded.status == "ready_with_issues"
    assert loaded.invalid_count == 2
    assert [item.item_id for item in loaded.items] == [
        "task",
        "task_2",
        "row_0003",
    ]
    assert [item.material.profile_id for item in loaded.items] == [
        "notice",
        "letter",
        "missing",
    ]
    assert loaded.selection.profile_ids == ["task", "task_2", "row_0003"]
    assert loaded.selection.source_kind == "official_document_table"
    assert loaded.selection.source_path == str(table_path)
    assert loaded.selection.output_dir_template == "{profile_id}_{entity_name}"
    assert loaded.selection.archive.profiles[0].fields["document_type"] == "notice"
    assert loaded.selection.archive.profiles[1].fields["document_type"] == "letter"
    assert loaded.selection.item_metadata["task_2"]["import_status"] == (
        "missing_required_fields"
    )
    assert loaded.selection.item_metadata["row_0003"]["import_status"] == (
        "unknown_profile"
    )


def test_official_document_material_batch_requires_horizontal_records(tmp_path):
    table_path = tmp_path / "vertical.csv"
    table_path.write_text(
        "key,value\n"
        "document_type,notice\n"
        "title,One notice\n",
        encoding="utf-8",
    )

    loaded = load_official_document_material_batch(table_path)

    assert loaded.status == "invalid_table"
    assert loaded.items == ()
    assert loaded.selection.archive.profiles == []
    assert loaded.issues == (
        "official_material_batch_requires_horizontal_records",
    )


def test_official_document_batch_runner_assembles_without_source_docx_and_isolates_failures(
    tmp_path,
):
    table_path = tmp_path / "official_batch.csv"
    table_path.write_text(
        "task_id,document_type,title,body,organization,document_no,issue_date\n"
        "notice_ok,notice,Notice A,Body A,Office,A-1,2026-07-10\n"
        "letter_bad,letter,Letter B,Body B,Office,,2026-07-11\n",
        encoding="utf-8",
    )
    batch = load_official_document_material_batch(table_path)
    scene = load_scene_from_library("official", mode_id="official")
    selection = batch.selection

    payload = WorkbenchBatchProductionRunner(
        doc_path="",
        template=None,
        scene=scene,
        archive=selection.archive,
        profile_ids=selection.profile_ids,
        base_output_dir=tmp_path / "out",
        output_dir_template=selection.output_dir_template,
        base_context=selection.base_context,
        source_kind=selection.source_kind,
        source_path=selection.source_path,
        item_metadata=selection.item_metadata,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "partial_success"
    assert payload["batch_source_kind"] == "official_document_table"
    assert payload["batch_source_path"] == str(table_path)
    assert [item["status"] for item in payload["items"]] == [
        "success",
        "failed",
    ]
    assert payload["items"][0]["profile_id"] == "notice_ok"
    assert payload["items"][0]["official_profile_id"] == "notice"
    assert payload["items"][1]["profile_id"] == "letter_bad"
    assert payload["items"][1]["official_profile_id"] == "letter"
    assert payload["items"][0]["output_path"].endswith(".docx")
    assert Path(payload["items"][0]["output_path"]).is_file()
    assert payload["items"][1]["output_path"] == ""
    assert payload["items"][1]["failed_count"] == 1
    assert payload["batch_isolation"]["success_count"] == 1
    assert payload["batch_isolation"]["failed_count"] == 1
    assert len(payload["batch_report_paths"]) == 2
    assert all(Path(path).is_file() for path in payload["batch_report_paths"])
    report_json_path = next(
        Path(path)
        for path in payload["batch_report_paths"]
        if str(path).endswith(".json")
    )
    report_payload = json.loads(report_json_path.read_text(encoding="utf-8"))
    assert report_payload["batch_source_kind"] == "official_document_table"
    assert report_payload["batch_source_path"] == str(table_path)
    assert "batch_history" not in report_payload
    assert "history_run_id" not in report_payload.get("batch_isolation", {})
    assert payload["batch_isolation"]["retry_eligible_profile_ids"] == [
        "letter_bad"
    ]
    assert Path(payload["batch_history_path"]).is_file()

    retry_selection = build_official_document_batch_retry_selection(
        selection,
        payload["batch_isolation"]["retry_eligible_profile_ids"],
        retry_of_run_id=payload["batch_run_id"],
        attempt_number=2,
    )
    retry_profile = next(
        profile
        for profile in retry_selection.archive.profiles
        if profile.profile_id == "letter_bad"
    )
    retry_profile.fields["document_no"] = "B-2"
    retry_payload = WorkbenchBatchProductionRunner(
        doc_path="",
        template=None,
        scene=scene,
        archive=retry_selection.archive,
        profile_ids=retry_selection.profile_ids,
        base_output_dir=tmp_path / "out",
        output_dir_template=retry_selection.output_dir_template,
        base_context=retry_selection.base_context,
        source_kind=retry_selection.source_kind,
        source_path=retry_selection.source_path,
        item_metadata=retry_selection.item_metadata,
        retry_of_run_id=payload["batch_run_id"],
        attempt_number=2,
    ).run(lambda *_args: None, lambda: False)

    assert retry_payload["status"] == "success"
    assert [item["profile_id"] for item in retry_payload["items"]] == [
        "letter_bad"
    ]
    assert retry_payload["batch_isolation"]["retry_eligible_profile_ids"] == []
    assert retry_payload["batch_isolation"]["retry_of_run_id"] == payload[
        "batch_run_id"
    ]
    assert retry_payload["batch_isolation"]["attempt_number"] == 2
    history = list_official_document_batch_history(tmp_path / "out")
    assert len(history) == 2
    assert history[0].retry_of_run_id == payload["batch_run_id"]
    assert history[0].failed_profile_ids == ()


def test_official_document_batch_applies_timeline_preflight_failure_policy(tmp_path):
    table_path = tmp_path / "official_timeline_batch.csv"
    table_path.write_text(
        "task_id,document_type,title,body,organization,document_no,issue_date\n"
        "notice_timeline,notice,Notice A,Body A,Office,A-1,2026-07-10\n",
        encoding="utf-8",
    )
    batch = load_official_document_material_batch(table_path)
    selection = batch.selection
    plan = default_timeline_plan()
    plan["start_field"] = "missing_timeline_start"
    plan["end_field"] = "missing_timeline_end"
    selection.archive.profiles[0].timeline_plans = {"primary": plan}

    warn_scene = load_scene_from_library("official", mode_id="official")
    warn_scene.input_source_profile.failure_policy = "warn"
    warn_payload = WorkbenchBatchProductionRunner(
        doc_path="",
        template=None,
        scene=warn_scene,
        archive=selection.archive,
        profile_ids=selection.profile_ids,
        base_output_dir=tmp_path / "warn_out",
        output_dir_template=selection.output_dir_template,
        base_context=selection.base_context,
        source_kind=selection.source_kind,
        source_path=selection.source_path,
        item_metadata=selection.item_metadata,
    ).run(lambda *_args: None, lambda: False)
    warn_timeline_diagnostics = [
        diagnostic
        for diagnostic in warn_payload["items"][0]["material_diagnostics"]
        if str(diagnostic.get("change_type") or "").startswith("timeline_")
    ]

    block_scene = load_scene_from_library("official", mode_id="official")
    block_scene.input_source_profile.failure_policy = "block"
    block_payload = WorkbenchBatchProductionRunner(
        doc_path="",
        template=None,
        scene=block_scene,
        archive=selection.archive,
        profile_ids=selection.profile_ids,
        base_output_dir=tmp_path / "block_out",
        output_dir_template=selection.output_dir_template,
        base_context=selection.base_context,
        source_kind=selection.source_kind,
        source_path=selection.source_path,
        item_metadata=selection.item_metadata,
    ).run(lambda *_args: None, lambda: False)
    block_item = block_payload["items"][0]
    block_timeline_diagnostics = [
        diagnostic
        for diagnostic in block_item["material_diagnostics"]
        if str(diagnostic.get("change_type") or "").startswith("timeline_")
    ]

    assert warn_payload["status"] == "success"
    assert warn_timeline_diagnostics
    assert {diagnostic["level"] for diagnostic in warn_timeline_diagnostics} == {"warning"}
    assert Path(warn_payload["items"][0]["output_path"]).is_file()
    assert block_payload["status"] == "failed"
    assert block_item["status"] == "failed"
    assert block_item["output_path"] == ""
    assert block_item["official_document_assembly"] == {}
    assert block_timeline_diagnostics
    assert {diagnostic["level"] for diagnostic in block_timeline_diagnostics} == {"error"}


def test_official_document_material_package_lists_builtin_profile_samples():
    builtin_dir = material_package_library.material_package_builtin_dir("official")
    samples = list_official_document_material_package_samples(source_type="builtin")
    samples_by_id = {sample.sample_id: sample for sample in samples}

    expected_ids = {
        "approval_archive_system",
        "letter_material_request",
        "notice_archive_check",
        "minutes_coordination",
        "report_work_summary",
        "request_archive_system",
    }
    assert set(samples_by_id) == expected_ids
    assert not tuple(builtin_dir.glob("*.json"))
    assert {
        path.parent.name
        for path in builtin_dir.glob("*/package.json")
    } == expected_ids
    assert all(sample.path.name == "package.json" for sample in samples)
    assert all(
        sample.path.parent.parent == builtin_dir
        for sample in samples
    )
    archives = [load_entity_archive(sample.path) for sample in samples]
    assert all(archive.kind == "alavette.material_package" for archive in archives)
    assert all(archive.version == ENTITY_PACKAGE_VERSION for archive in archives)
    assert all(archive.mode_id == "official" for archive in archives)
    assert samples_by_id["letter_material_request"].profile_id == "letter"
    assert samples_by_id["letter_material_request"].source_type == "builtin"
    assert samples_by_id["letter_material_request"].material_schema_ids == (
        "official_document_v1",
    )
    assert samples_by_id["notice_archive_check"].profile_id == "notice"
    assert samples_by_id["notice_archive_check"].source_type == "builtin"
    assert samples_by_id["notice_archive_check"].material_schema_ids == (
        "official_document_v1",
    )
    assert samples_by_id["minutes_coordination"].profile_id == "minutes"
    assert samples_by_id["minutes_coordination"].material_schema_ids == (
        "official_document_v1",
        "administrative_meeting_fields_v1",
    )
    assert get_official_document_material_package_sample(
        "builtin/notice_archive_check"
    ) == samples_by_id["notice_archive_check"]
    assert [
        option.value
        for option in material_package_sample_selector_options(
            "official",
            profile_id="official:letter",
        )
    ] == ["builtin/letter_material_request"]
    assert [
        option.value
        for option in material_package_sample_selector_options(
            "official",
            profile_id="minutes",
        )
    ] == ["builtin/minutes_coordination"]
    assert material_package_sample_selector_options(
        "official",
        profile_id="missing_profile",
    ) == ()


def test_noncanonical_official_library_entry_is_visible_but_fail_closed(
    monkeypatch,
    tmp_path,
):
    library_root = tmp_path / "material_packages"
    monkeypatch.setattr(
        material_package_library,
        "MATERIAL_PACKAGE_LIBRARY_DIR",
        library_root,
    )
    package_path = (
        material_package_library.material_package_builtin_dir("official")
        / "legacy_notice"
        / "package.json"
    )
    package_path.parent.mkdir(parents=True)
    package_path.write_text(
        json.dumps(
                {
                    "kind": "alavette.official_document.material_package",
                    "version": 1,
                    "mode_id": "official",
                    "package_id": "legacy_notice",
                }
        ),
        encoding="utf-8",
    )

    entries = material_package_library.list_material_package_entries(
        mode_id="official"
    )
    samples = list_official_document_material_package_samples()
    options = material_package_sample_selector_options("official")
    loaded = load_official_document_material_package_sample(
        "builtin/legacy_notice"
    )

    assert [entry.qualified_id for entry in entries] == ["builtin/legacy_notice"]
    assert not entries[0].is_available
    assert entries[0].load_error == "material_package_version_unsupported:1"
    assert [sample.qualified_id for sample in samples] == ["builtin/legacy_notice"]
    assert samples[0].load_error == entries[0].load_error
    assert [option.value for option in options] == ["builtin/legacy_notice"]
    assert options[0].disabled is True
    assert loaded.status == "invalid_package"
    assert loaded.context.is_empty()


def test_assets_panel_and_official_scene_projection_share_canonical_library():
    bridge = PanelBridge()
    bridge.set_current_work_mode("official", emit_signal=False)
    panel = AssetsPanel(bridge)
    try:
        panel._refresh_archive_selector_from_library()
        asset_entries = {
            entry.qualified_id
            for entry in panel._archive_entries_by_path.values()
            if entry.source_type == "builtin"
        }
        scene_samples = {
            sample.qualified_id
            for sample in list_official_document_material_package_samples(
                source_type="builtin"
            )
        }

        assert asset_entries == scene_samples
        assert asset_entries == {
            "builtin/approval_archive_system",
            "builtin/letter_material_request",
            "builtin/minutes_coordination",
            "builtin/notice_archive_check",
            "builtin/report_work_summary",
            "builtin/request_archive_system",
        }
    finally:
        panel.close()


def test_common_official_document_types_have_samples_independently_from_plans():
    ensure_config_library()

    builtin_decisions = {
        decision.profile_id: decision.plan_id
        for decision in list_official_document_plan_entry_decisions()
        if decision.entry_kind == OFFICIAL_PLAN_ENTRY_BUILTIN
    }
    samples_by_profile = {
        sample.profile_id: sample.qualified_id
        for sample in list_official_document_material_package_samples(
            source_type="builtin",
        )
    }

    assert builtin_decisions == {"notice": "official"}
    assert samples_by_profile == {
        "notice": "builtin/notice_archive_check",
        "letter": "builtin/letter_material_request",
        "minutes": "builtin/minutes_coordination",
        "report": "builtin/report_work_summary",
        "request": "builtin/request_archive_system",
        "approval": "builtin/approval_archive_system",
    }

    for profile in list_common_official_document_profiles():
        profile_id = profile.profile_id
        contract = get_official_document_assembly_contract(profile_id)
        sample = load_official_document_material_package_sample(
            samples_by_profile[profile_id],
            source_type="builtin",
        )

        assert contract is not None
        assert sample.status == "ok"
        assert sample.profile_id == profile_id
        assert tuple(contract.material_schema_ids) == sample.material_schema_ids


def test_official_document_material_package_saves_user_library_sample(monkeypatch, tmp_path):
    library_root = tmp_path / "config_library" / "material_packages"
    monkeypatch.setattr(
        material_package_library,
        "MATERIAL_PACKAGE_LIBRARY_DIR",
        library_root,
    )
    user_dir = material_package_library.material_package_user_dir("official")
    user_dir.mkdir(parents=True)
    ignored_flat_path = user_dir / "school_notice.material.json"
    ignored_flat_path.write_text("{}", encoding="utf-8")
    context = MaterialExecutionContext(
        profile_id="official:notice",
        profile_name="Notice",
        entity_data={
            "title": "Notice title",
            "body": "Notice body",
            "organization": "Archive Office",
            "document_no": "A-2026-1",
            "issue_date": "2026-07-10",
        },
    )

    result = export_official_document_material_package_to_user_library(
        context,
        package_id="school_notice",
        label="School notice package",
        profile_id="notice",
    )

    assert result.status == "ok"
    assert result.package_path == user_dir / "school_notice" / "package.json"
    assert ignored_flat_path.exists()
    payload = json.loads(result.package_path.read_text(encoding="utf-8"))
    assert payload["kind"] == "alavette.material_package"
    assert payload["version"] == ENTITY_PACKAGE_VERSION
    assert payload["package_id"] == "school_notice"
    assert payload["archive_name"] == "School notice package"
    assert payload["profiles"][0]["profile_id"] == "official:notice"
    assert payload["profiles"][0]["fields"]["document_type"] == "notice"

    samples = list_official_document_material_package_samples(source_type="user")
    assert [sample.qualified_id for sample in samples] == ["user/school_notice"]
    assert samples[0].label == "School notice package"
    assert get_official_document_material_package_sample(
        "user/school_notice"
    ) == samples[0]
    selector_options = material_package_sample_selector_options("official")
    selector_by_value = {option.value: option for option in selector_options}
    assert "user/school_notice" in selector_by_value
    assert selector_by_value["user/school_notice"].label == "用户资料包：School notice package"


def test_official_material_package_library_export_conflict_creates_no_artifacts(
    monkeypatch,
    tmp_path,
):
    library_root = tmp_path / "config_library" / "material_packages"
    monkeypatch.setattr(
        material_package_library,
        "MATERIAL_PACKAGE_LIBRARY_DIR",
        library_root,
    )
    context = MaterialExecutionContext(
        profile_id="official:letter",
        entity_data={"document_type": "letter", "title": "Letter source"},
    )

    result = export_official_document_material_package_to_user_library(
        context,
        package_id="must_not_exist",
        profile_id="notice",
    )

    assert result.status == "profile_conflict"
    assert result.package_path is None
    assert context.profile_id == "official:letter"
    assert context.entity_data["document_type"] == "letter"
    assert not library_root.exists()


def test_same_id_builtin_and_user_material_packages_require_qualified_identity(
    monkeypatch,
    tmp_path,
):
    library_root = tmp_path / "config_library" / "material_packages"
    monkeypatch.setattr(
        material_package_library,
        "MATERIAL_PACKAGE_LIBRARY_DIR",
        library_root,
    )
    builtin_source = (
        Path(__file__).resolve().parents[1]
        / "config_library"
        / "material_packages"
        / "official"
        / "builtin"
    )
    builtin_target = material_package_library.material_package_builtin_dir("official")
    copytree(builtin_source, builtin_target)
    saved = export_official_document_material_package_to_user_library(
        MaterialExecutionContext(
            profile_id="official:notice",
            entity_data={
                "title": "User notice title",
                "body": "User notice body",
                "organization": "User Archive Office",
                "document_no": "U-2026-1",
                "issue_date": "2026-07-10",
            },
        ),
        package_id="notice_archive_check",
        label="User Notice Archive Check",
        profile_id="notice",
    )

    assert saved.status == "ok"
    matching = [
        sample
        for sample in list_official_document_material_package_samples()
        if sample.sample_id == "notice_archive_check"
    ]
    assert {sample.qualified_id for sample in matching} == {
        "builtin/notice_archive_check",
        "user/notice_archive_check",
    }

    builtin = get_official_document_material_package_sample(
        "builtin/notice_archive_check"
    )
    user = get_official_document_material_package_sample("user/notice_archive_check")
    assert builtin is not None
    assert user is not None
    assert builtin.source_type == "builtin"
    assert user.source_type == "user"
    assert get_official_document_material_package_sample("notice_archive_check") is None

    ambiguous = load_official_document_material_package_sample("notice_archive_check")
    loaded_user = load_official_document_material_package_sample(
        "user/notice_archive_check"
    )
    assert ambiguous.status == "sample_ambiguous"
    assert ambiguous.issues == (
        "official_material_package_sample_ambiguous: notice_archive_check; "
        "matches=builtin/notice_archive_check,user/notice_archive_check",
    )
    assert loaded_user.status == "ok"
    assert loaded_user.context.entity_data["title"] == "User notice title"

    selector_values = {
        option.value
        for option in material_package_sample_selector_options(
            "official",
            profile_id="notice",
        )
    }
    assert {
        "builtin/notice_archive_check",
        "user/notice_archive_check",
    } <= selector_values


def test_builtin_official_material_package_samples_load_cleanly():
    for sample_id in (
        "letter_material_request",
        "notice_archive_check",
        "minutes_coordination",
    ):
        loaded = load_official_document_material_package_sample(
            sample_id,
            source_type="builtin",
        )

        assert loaded.status == "ok"
        assert loaded.missing_required_fields == ()
        assert loaded.unknown_fields == ()
        assert loaded.context.profile_id == f"official:{loaded.profile_id}"
        assert loaded.context.entity_data["document_type"] == loaded.profile_id

    missing = load_official_document_material_package_sample("missing_sample")
    assert missing.status == "sample_not_found"
    assert missing.issues == (
        "official_material_package_sample_not_found: missing_sample",
    )


def test_builtin_official_material_package_samples_feed_official_assembly(tmp_path):
    for sample_id in (
        "letter_material_request",
        "notice_archive_check",
        "minutes_coordination",
    ):
        loaded = load_official_document_material_package_sample(
            sample_id,
            source_type="builtin",
        )
        assembly = assemble_official_document_docx(
            loaded.profile_id,
            loaded.context.entity_data,
            tmp_path / sample_id,
        )

        assert assembly.status == "ok"
        assert assembly.docx_path is not None
        assert assembly.docx_path.is_file()
        assert assembly.archive_manifest_path is not None
        assert assembly.archive_manifest_path.is_file()
        assert loaded.context.entity_data["title"] in _all_docx_text(assembly.docx_path)


def test_official_document_material_package_load_feeds_official_assembly(tmp_path):
    package_path = tmp_path / "notice_material.json"
    export_official_document_material_package(
        MaterialExecutionContext(
            profile_id="official:notice",
            entity_data={
                "title": "Notice title",
                "body": "Notice body",
                "organization": "Archive Office",
                "document_no": "A-2026-1",
                "issue_date": "2026-07-10",
            },
        ),
        package_path,
        profile_id="notice",
    )

    loaded = load_official_document_material_package(package_path)
    assembly = assemble_official_document_docx(
        loaded.profile_id,
        loaded.context.entity_data,
        tmp_path / "out",
    )

    assert loaded.status == "ok"
    assert assembly.status == "ok"
    assert assembly.output_paths["official_docx"].endswith(".docx")
    assert "internal_review_docx" in assembly.output_paths
    assert "archive_manifest" in assembly.output_paths


def test_official_document_material_package_uses_minutes_schema_ids(tmp_path):
    context = MaterialExecutionContext(
        profile_id="official:minutes",
        entity_data={
            "title": "Meeting minutes",
            "body": "Discussed archive process.",
            "organization": "Office",
            "document_no": "M-2026-1",
            "issue_date": "2026-07-10",
            "meeting_date": "2026-07-10 AM",
            "participants": "Alice, Bob",
        },
    )

    result = export_official_document_material_package(
        context,
        tmp_path / "minutes_material.json",
        profile_id="minutes",
    )
    loaded = load_official_document_material_package(result.package_path)

    assert result.status == "ok"
    assert result.material_schema_ids == (
        "official_document_v1",
        "administrative_meeting_fields_v1",
    )
    assert loaded.context.profile_id == "official:minutes"
    assert loaded.context.entity_data["meeting_date"] == "2026-07-10 AM"
    assert loaded.unknown_fields == ()


def test_official_document_material_package_preserves_draft_with_missing_fields(
    tmp_path,
):
    result = export_official_document_material_package(
        MaterialExecutionContext(
            profile_id="official:letter",
            entity_data={
                "title": "Letter title",
                "body": "Letter body",
            },
        ),
        tmp_path / "letter_draft.json",
        profile_id="letter",
    )
    loaded = load_official_document_material_package(result.package_path)

    assert result.status == "missing_required_fields"
    assert result.missing_required_fields == (
        "organization",
        "document_no",
        "issue_date",
    )
    assert loaded.status == "missing_required_fields"
    assert loaded.context.entity_data["document_type"] == "letter"
    assert loaded.issues == (
        "missing_required_field: organization",
        "missing_required_field: document_no",
        "missing_required_field: issue_date",
    )


def test_official_document_material_package_reports_unknown_profile(tmp_path):
    package_path = tmp_path / "unknown.json"
    result = export_official_document_material_package(
        MaterialExecutionContext(
            profile_id="official:missing",
            entity_data={"document_type": "missing", "title": "Title"},
        ),
        package_path,
        profile_id="missing",
    )

    assert result.status == "unknown_profile"
    assert result.profile_id == "missing"
    assert result.context.is_empty()
    assert result.issues == ("unknown_official_profile: missing",)
    assert not package_path.exists()


def test_official_document_material_package_rejects_unversioned_payload(tmp_path):
    source_path = tmp_path / "legacy_notice.json"
    source_payload = {
        "profile_id": "official:notice",
        "profile_name": "Legacy notice",
        "sample_id": "legacy_notice",
        "sample_label": "Legacy notice sample",
        "title": "Legacy title",
        "body": "Legacy body",
        "organization": "Archive Office",
        "document_no": "N-1",
        "issue_date": "2026-07-10",
    }
    source_path.write_text(
        json.dumps(source_payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    loaded = load_official_document_material_package(source_path)

    assert loaded.status == "invalid_package"
    assert loaded.context.is_empty()
    assert loaded.issues == ("material_package_version_unsupported:0",)
    assert not hasattr(
        material_package_module,
        "convert_official_document_material_package",
    )


def test_official_document_material_package_rejects_dedicated_v1_protocol(tmp_path):
    source_path = tmp_path / "official_v1_notice.json"
    source_path.write_text(
        json.dumps(
            {
                "kind": "alavette.official_document.material_package",
                "version": 1,
                "profile_id": "official:notice",
                "entity_data": {"title": "Legacy title"},
            }
        ),
        encoding="utf-8",
    )

    result = load_official_document_material_package(source_path)

    assert result.status == "invalid_package"
    assert result.context.is_empty()
    assert result.issues == ("material_package_version_unsupported:1",)


def test_official_document_material_package_blocks_future_versions_and_wrong_kinds(
    tmp_path,
):
    future_path = tmp_path / "future.json"
    wrong_kind_path = tmp_path / "wrong_kind.json"
    context = MaterialExecutionContext(
        profile_id="official:notice",
        entity_data={
            "document_type": "notice",
            "title": "Notice title",
            "body": "Notice body",
            "organization": "Office",
            "document_no": "N-1",
            "issue_date": "2026-07-10",
        },
    )
    export_official_document_material_package(
        context,
        future_path,
        profile_id="notice",
    )
    export_official_document_material_package(
        context,
        wrong_kind_path,
        profile_id="notice",
    )
    future_payload = json.loads(future_path.read_text(encoding="utf-8"))
    future_payload["version"] = PACKAGE_VERSION + 1
    future_path.write_text(json.dumps(future_payload), encoding="utf-8")
    wrong_kind_payload = json.loads(wrong_kind_path.read_text(encoding="utf-8"))
    wrong_kind_payload["kind"] = "alavette.other.package"
    wrong_kind_path.write_text(
        json.dumps(wrong_kind_payload),
        encoding="utf-8",
    )

    future_loaded = load_official_document_material_package(future_path)
    wrong_kind_loaded = load_official_document_material_package(
        wrong_kind_path
    )

    assert future_loaded.status == "invalid_package"
    assert future_loaded.context.is_empty()
    assert future_loaded.issues == (
        f"material_package_version_unsupported:{PACKAGE_VERSION + 1}",
    )
    assert wrong_kind_loaded.status == "invalid_package"
    assert wrong_kind_loaded.context.is_empty()
    assert wrong_kind_loaded.issues == (
        "material_package_kind_invalid:alavette.other.package",
    )


def test_official_document_material_package_tracks_unknown_fields(tmp_path):
    result = export_official_document_material_package(
        MaterialExecutionContext(
            profile_id="official:notice",
            entity_data={
                "title": "Notice title",
                "body": "Notice body",
                "organization": "Office",
                "document_no": "N-1",
                "issue_date": "2026-07-10",
                "unexpected": "kept for review",
            },
        ),
        tmp_path / "notice_extra.json",
        profile_id="notice",
    )

    assert result.status == "ok"
    assert result.unknown_fields == ("unexpected",)
    loaded = load_official_document_material_package(result.package_path)
    assert loaded.unknown_fields == ("unexpected",)
    assert loaded.context.entity_data["unexpected"] == "kept for review"
