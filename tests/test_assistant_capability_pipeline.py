from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path
import re
import zipfile

from docx import Document
from PIL import Image
import pytest

from src.assistant.adapters.content_generation_adapter import (
    AssistantContentGenerationAdapter,
    GeneratedContentDraft,
    GeneratedExamDraft,
)
from src.assistant.adapters.production_adapter import AssistantProductionAdapter
from src.assistant.adapters.workspace_state_adapter import WorkspaceSnapshot
from src.assistant.application.capability_registry import capability_route_coverage
from src.assistant.application.content_generation_service import (
    AssistantContentGenerationService,
    ContentGenerationRequest,
)
from src.assistant.application.generated_draft_binding import bind_generated_draft
from src.assistant.application.plan_builder import (
    FormDocumentPlanBuilder,
    requests_form_document_action,
)
from src.assistant.application.plan_presentation import present_document_plan
from src.assistant.contracts.document_plan import DocumentPlan, OutputPolicy
from src.assistant.contracts.execution import ExecutionApproval
from src.assistant.contracts.permissions import DisclosureGrant
from src.assistant.contracts.task_plan import (
    ARTIFACT_KIND_EXAM,
    CAPABILITY_EXECUTABLE,
    CAPABILITY_GATED,
    CAPABILITY_PLANNED,
    SOURCE_ROLE_STRUCTURED_SOURCE,
    TASK_OPERATION_AUTHOR,
)
from src.assistant.runtime.provider_contract import (
    PROVIDER_DONE,
    PROVIDER_START,
    PROVIDER_TEXT_DELTA,
    ProviderStreamEvent,
)
from src.assistant.runtime.turn_runner import attachment_fingerprints
from src.config.entity import EntityArchive, EntityProfile
from src.config.image_materials import (
    ImageMaterialRule,
    ImagePlacementMode,
    ImagePlacementPolicy,
)
from src.config.scene_natural_request_router import list_natural_request_routes
from src.config.material_context import MaterialExecutionContext
from src.config.materials import AssetItem
from src.services.material_execution import (
    MaterialAssemblyDependencies,
    MaterialAssemblyService as RealMaterialAssemblyService,
)
from src.services.production_runtime import (
    material_assembly_runtime as assembly_runtime,
)
from src.shared.engine.office_image_layout import OfficeImageProvider
from tests.material_image_layout_fake import FakeSuccessfulLayout


EXAM_MARKDOWN = """# 七年级数学期中模拟试卷

> 科目：数学　年级：七年级　考试时间：90 分钟　满分：15 分

## 一、选择题（本大题共 2 小题，每小题 5 分，共 10 分）

1. 下列各数中，最小的是（　　）（5 分）
   A. -3
   B. 0
   C. 2
   D. -1

2. 计算 `-4 + 7` 的结果是（　　）（5 分）
   A. -11
   B. -3
   C. 3
   D. 11

## 二、解答题（本大题共 1 小题，共 5 分）

1. 计算：`(-6) + 9 - (-4) - 7`。（5 分）

   answer_area_kind: free
   answer_lines: 4

## 答案速查

一、选择题
1. A
2. C

二、解答题
1. 0。解析：`(-6) + 9 + 4 - 7 = 0`。
"""


class _FixedGateway:
    def __init__(self, content: str) -> None:
        self.content = content
        self.request = None

    def stream(self, request):
        self.request = request
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text=self.content)
        yield ProviderStreamEvent(PROVIDER_DONE)

    def cancel(self) -> bool:
        return True


def _workspace() -> WorkspaceSnapshot:
    return WorkspaceSnapshot(
        mode_id="custom",
        mode_label="通用版",
        scene_id="custom",
        scene_source_type="builtin",
        template_id="default",
        template_source_type="builtin",
        input_path="",
        input_name="",
        input_exists=False,
        material_summary={},
    )


def _approval(preflight) -> ExecutionApproval:
    return ExecutionApproval(
        approval_id="approval-exam-e2e",
        session_id="session-exam",
        plan_id=preflight.plan_id,
        plan_revision=preflight.plan_revision,
        preflight_hash=preflight.evidence_hash,
        input_hash=preflight.input_hash,
        output_root=preflight.output_root,
        overwrite_policy="deny",
        approved_at=datetime.now(timezone.utc).isoformat(),
    )


def _docx_text(path: Path) -> str:
    document = Document(path)
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def test_every_natural_route_has_explicit_capability_status():
    coverage = capability_route_coverage()
    route_ids = {route.route_id for route in list_natural_request_routes()}

    assert set(coverage) == route_ids
    assert set(coverage.values()) <= {
        CAPABILITY_EXECUTABLE,
        CAPABILITY_PLANNED,
        CAPABILITY_GATED,
    }
    assert coverage["exam_teaching_versions"] == CAPABILITY_EXECUTABLE
    assert coverage["english_journal_submission"] == CAPABILITY_PLANNED
    assert coverage["project_application_package"] == CAPABILITY_PLANNED
    assert coverage["product_sales_document"] == CAPABILITY_PLANNED
    assert coverage["legal_document_manual_boundary"] == CAPABILITY_GATED


def test_assistant_panel_does_not_own_exam_branching():
    source = (
        Path("src/assistant/ui/assistant_panel.py")
        .read_text(encoding="utf-8")
    )

    assert 'mode_id == "exam"' not in source
    assert "compile_exam_markdown" not in source
    assert "exam_items_v1" not in source


def test_exam_authoring_plan_is_typed_and_does_not_request_docx():
    plan = FormDocumentPlanBuilder().build(
        query="生成一份试卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )

    assert requests_form_document_action("生成一份试卷")
    assert plan.work_mode_id == "exam"
    assert plan.operation == TASK_OPERATION_AUTHOR
    assert plan.capability_ref.status == CAPABILITY_EXECUTABLE
    assert plan.generation_contract.required
    assert plan.generation_contract.artifact_kind == ARTIFACT_KIND_EXAM
    assert plan.production_contract.input_role == SOURCE_ROLE_STRUCTURED_SOURCE
    assert plan.production_contract.terminal_assembler == "exam"
    assert plan.delivery_contract.required_artifact_keys == ("student", "answer_key")
    assert not plan.unresolved_questions


def test_exam_plan_presentation_uses_product_language_not_config_ids():
    plan = FormDocumentPlanBuilder().build(
        query="生成一份七年级数学试卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )

    presentation = present_document_plan(plan)

    assert presentation.title == "试卷生成计划"
    assert "模式：试卷版" in presentation.body
    assert "交付：学生卷、答案卷" in presentation.body
    assert "工作模式：exam" not in presentation.body
    assert "模板：default" not in presentation.body
    assert presentation.actions == (
        {"id": "generate_content_draft", "label": "生成并校验题稿"},
    )


def test_bidding_authoring_generates_tokenized_reviewable_docx(tmp_path):
    workspace = WorkspaceSnapshot(
        mode_id="custom",
        mode_label="通用版",
        scene_id="custom",
        scene_source_type="builtin",
        template_id="default",
        template_source_type="builtin",
        input_path="",
        input_name="",
        input_exists=False,
        material_summary={
            "package_id": "bidder-main",
            "profile_id": "company-a",
            "material_schema_ids": ["bid_materials_v1"],
            "field_count": 3,
            "asset_count": 2,
        },
    )
    plan = FormDocumentPlanBuilder().build(
        query="生成一份投标标书正文",
        workspace=workspace,
        turn_id="turn-bid-author",
    )
    markdown = """# 投标文件

投标人：{{@text:company_name}}

项目名称：{{@text:project_name}}

法定代表人：{{@text:legal_person}}

## 项目理解

根据招标要求形成项目理解。

## 响应内容

逐项响应招标要求。

## 实施方案

按阶段组织实施。

## 承诺事项

本文件中的企业标志与公章由本地资料包装配。
"""
    gateway = _FixedGateway(markdown)
    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path / "assistant")
    ).generate(
        ContentGenerationRequest(
            session_id="session-bid",
            turn_id="turn-bid-author",
            prompt=plan.intent,
            provider_id="fixed",
            model_id="fixed",
            capability_id=plan.capability_ref.capability_id,
            artifact_kind=plan.generation_contract.artifact_kind,
            prompt_profile_id=plan.generation_contract.prompt_profile_id,
        ),
        gateway,
    )
    bound = bind_generated_draft(plan, draft)
    text = _docx_text(Path(draft.document_path))

    assert plan.capability_ref.status == CAPABILITY_EXECUTABLE
    assert plan.generation_contract.prompt_profile_id == (
        "assistant.bidding-markdown.v1"
    )
    assert not plan.unresolved_questions
    assert "{{@text:company_name}}" in text
    assert "{{@img:LOGO1}}" in text
    assert "{{@img:公章1}}" in text
    assert draft.document_profile_id == "assistant.bidding-local-assembly.v1"
    assert bound.production_input_ref["path"] == draft.document_path
    assert "{{@text:company_name}}" in gateway.request.system_prompt
    assert "不要输出 @img Token" in gateway.request.system_prompt
    assert "不得虚构" in gateway.request.system_prompt


def test_bidding_generated_draft_consumes_package_fields_and_images_end_to_end(
    tmp_path,
    monkeypatch,
):
    logo_path = tmp_path / "logo.png"
    seal_path = tmp_path / "seal.png"
    Image.new("RGB", (80, 40), color="blue").save(logo_path)
    Image.new("RGB", (80, 80), color="red").save(seal_path)
    image_rules = {
        rule_id: ImageMaterialRule(
            rule_id=rule_id,
            source_role=role,
            anchor_token=token,
            placement=ImagePlacementPolicy(
                mode=ImagePlacementMode.FIXED_BOX,
                fixed_width_cm=6.0,
            ),
        )
        for rule_id, role, token in (
            ("image:logo", "logo", "{{@img:LOGO1}}"),
            ("image:seal", "seal", "{{@img:公章1}}"),
        )
    }
    package = EntityArchive(
        archive_id="bidder-main",
        archive_name="投标资料包",
        mode_id="bidding",
        package_id="bidder-main",
        material_schema_ids=["bid_materials_v1"],
        profiles=[
            EntityProfile(
                profile_id="company-a",
                profile_name="投标主体",
                image_material_rules=image_rules,
            )
        ],
    )
    profile = package.profiles[0]
    assert package.material_schema_ids == ["bid_materials_v1"]
    fake_layout = FakeSuccessfulLayout()
    dependencies = replace(
        MaterialAssemblyDependencies(),
        layout_runner=fake_layout,
    )
    monkeypatch.setattr(
        assembly_runtime,
        "_qualified_office_provider",
        lambda _path: OfficeImageProvider.WORD,
    )
    monkeypatch.setattr(
        assembly_runtime,
        "_material_cache_root",
        lambda: tmp_path / "material-cache",
    )
    monkeypatch.setattr(
        assembly_runtime,
        "MaterialAssemblyService",
        lambda callback: RealMaterialAssemblyService(
            callback,
            dependencies=dependencies,
        ),
    )
    workspace = WorkspaceSnapshot(
        mode_id="custom",
        mode_label="通用版",
        scene_id="custom",
        scene_source_type="builtin",
        template_id="default",
        template_source_type="builtin",
        input_path="",
        input_name="",
        input_exists=False,
        material_summary={
            "package_id": "bidder-main",
            "profile_id": "company-a",
            "material_schema_ids": ["bid_materials_v1"],
            "field_count": 3,
            "asset_count": 2,
        },
    )
    plan = FormDocumentPlanBuilder().build(
        query="生成一份投标标书正文",
        workspace=workspace,
        turn_id="turn-bid-package-e2e",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(
            output_root=str(tmp_path / "outputs"),
            filename_suffix="_bid",
        ),
    )
    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path / "assistant")
    ).generate(
        ContentGenerationRequest(
            session_id="session-bid-package",
            turn_id="turn-bid-package-e2e",
            prompt=plan.intent,
            provider_id="fixed",
            model_id="fixed",
            capability_id=plan.capability_ref.capability_id,
            artifact_kind=plan.generation_contract.artifact_kind,
            prompt_profile_id=plan.generation_contract.prompt_profile_id,
        ),
        _FixedGateway(
            """# 投标文件

投标人：{{@text:company_name}}

项目名称：{{@text:project_name}}

法定代表人：{{@text:legal_person}}

## 项目响应

本项目将按招标要求组织实施。

## 项目理解

已梳理项目目标和范围。

## 实施方案

按阶段组织实施。

## 承诺事项

未明确事项标记为待补充。
"""
        ),
    )
    plan = bind_generated_draft(plan, draft)
    adapter = AssistantProductionAdapter()
    material_context = MaterialExecutionContext(
            mode_id="bidding",
            scene_id="bidding",
            package_id="bidder-main",
            material_schema_ids=("bid_materials_v1",),
            archive_id="bidder-main",
            profile_id="company-a",
            profile_name="投标主体",
            entity_data={
                "company_name": "示例建设有限公司",
                "project_name": "示例工程项目",
                "legal_person": "张三",
            },
            asset_items=[
                AssetItem(
                    item_id="logo",
                    role="logo",
                    path=str(logo_path),
                    mime_type="image/png",
                ),
                AssetItem(
                    item_id="seal",
                    role="seal",
                    path=str(seal_path),
                    mime_type="image/png",
                ),
            ],
            image_material_rules=dict(profile.image_material_rules),
        )
    preflight = adapter.build_preflight(
        plan,
        material_context=material_context,
    )

    assert preflight.ready, preflight.issues
    result = adapter.execute_approved_plan(
        plan,
        preflight,
        _approval(preflight),
        material_context=material_context,
    )

    assert result["status"] == "success", result
    output_path = Path(result["output_path"])
    assert output_path.is_file()
    output = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in output.paragraphs)
    assert "示例建设有限公司" in text
    assert "示例工程项目" in text
    assert "张三" in text
    assert "{{@text:" not in text
    assert "{{@img:" not in text
    assert len(output.inline_shapes) == 2
    assert set(result["output_paths"]) == {"original", "copy"}
    assert len(fake_layout.requests) == 2


def test_generated_exam_draft_is_validated_without_generic_docx_downgrade(tmp_path):
    plan = FormDocumentPlanBuilder().build(
        query="生成一份七年级数学试卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    gateway = _FixedGateway(EXAM_MARKDOWN)
    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    ).generate(
        ContentGenerationRequest(
            session_id="session-exam",
            turn_id="turn-exam",
            prompt=plan.intent,
            provider_id="fixed",
            model_id="fixed",
            capability_id=plan.capability_ref.capability_id,
            artifact_kind=plan.generation_contract.artifact_kind,
            prompt_profile_id=plan.generation_contract.prompt_profile_id,
        ),
        gateway,
    )

    assert isinstance(draft, GeneratedExamDraft)
    assert not isinstance(draft, GeneratedContentDraft)
    assert Path(draft.production_input_path).suffix == ".md"
    assert draft.validation_summary["question_count"] == 3
    assert "答案速查" in gateway.request.system_prompt
    assert not hasattr(draft, "document_path")

    bound = bind_generated_draft(plan, draft)
    assert bound.production_input_artifact is not None
    assert bound.production_input_artifact.role == SOURCE_ROLE_STRUCTURED_SOURCE
    assert bound.production_input_ref["path"] == draft.markdown_path


def test_exam_generation_to_student_and_answer_docx_is_closed(tmp_path):
    plan = FormDocumentPlanBuilder().build(
        query="生成一份七年级数学试卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(
            output_root=str(tmp_path / "outputs"),
            filename_suffix="_ai",
        ),
    )
    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path / "assistant")
    ).generate(
        ContentGenerationRequest(
            session_id="session-exam",
            turn_id="turn-exam",
            prompt=plan.intent,
            provider_id="fixed",
            model_id="fixed",
            capability_id=plan.capability_ref.capability_id,
            artifact_kind=plan.generation_contract.artifact_kind,
            prompt_profile_id=plan.generation_contract.prompt_profile_id,
        ),
        _FixedGateway(EXAM_MARKDOWN),
    )
    plan = bind_generated_draft(plan, draft)
    adapter = AssistantProductionAdapter()

    preflight = adapter.build_preflight(plan)
    assert preflight.ready, preflight.issues
    assert set(preflight.resource_fingerprints) == {
        "scene",
        "template",
        "master",
        "materials",
    }

    result = adapter.execute_approved_plan(plan, preflight, _approval(preflight))

    assert result["status"] == "success", result.get("error_text")
    output_paths = result["output_paths"]
    assert {"student", "answer_key"} <= set(output_paths)
    for key in ("student", "answer_key"):
        path = Path(output_paths[key])
        assert path.is_file()
        Document(path)
    assert "答案速查" not in _docx_text(Path(output_paths["student"]))
    assert "答案速查" in _docx_text(Path(output_paths["answer_key"]))
    assert Path(output_paths["student"]).read_bytes() != Path(
        output_paths["answer_key"]
    ).read_bytes()
    runtime = result["exam_delivery_runtime"]
    versions = {
        item["preset_id"]: item for item in runtime["rendered_versions"]
    }
    assert versions["student"]["visible_question_count"] == 3
    assert versions["student"]["visible_answer_count"] == 0
    assert versions["answer_key"]["visible_question_count"] == 0
    assert versions["answer_key"]["visible_answer_count"] == 3
    assert versions["answer_key"]["visible_analysis_count"] == 1
    assert "解析：" in _docx_text(Path(output_paths["answer_key"]))


def test_assistant_qualification_route_delivers_complete_attachment_package(
    tmp_path,
):
    input_path = tmp_path / "qualification-index.docx"
    source = Document()
    source.add_paragraph("投标资质附件归档索引")
    source.save(input_path)
    certificate = tmp_path / "certificate.pdf"
    license_file = tmp_path / "business-license.pdf"
    certificate.write_bytes(b"%PDF-1.4\ncertificate")
    license_file.write_bytes(b"%PDF-1.4\nbusiness-license")
    workspace = WorkspaceSnapshot(
        mode_id="custom",
        mode_label="通用版",
        scene_id="custom",
        scene_source_type="builtin",
        template_id="default",
        template_source_type="builtin",
        input_path=str(input_path),
        input_name=input_path.name,
        input_exists=True,
        material_summary={
            "package_id": "qualification-main",
            "profile_id": "company-a",
            "field_count": 2,
            "asset_count": 2,
        },
    )
    plan = FormDocumentPlanBuilder().build(
        query="整理投标资质证书材料",
        workspace=workspace,
        turn_id="turn-qualification",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(
            output_root=str(tmp_path / "outputs"),
            filename_suffix="_archive",
        ),
    )
    adapter = AssistantProductionAdapter()
    material_context = MaterialExecutionContext(
            mode_id="bidding",
            scene_id="bidding",
            package_id="qualification-main",
            material_schema_ids=("qualification_archive_assets_v1",),
            archive_id="qualification-main",
            profile_id="company-a",
            profile_name="投标资质包",
            entity_data={
                "organization": "示例建设有限公司",
                "package_name": "示例项目投标资质包",
            },
            asset_items=[
                AssetItem(
                    item_id="certificate",
                    role="certificate",
                    path=str(certificate),
                    mime_type="application/pdf",
                ),
                AssetItem(
                    item_id="business-license",
                    role="business_license",
                    path=str(license_file),
                    mime_type="application/pdf",
                ),
            ],
        )
    preflight = adapter.build_preflight(
        plan,
        material_context=material_context,
    )

    assert preflight.ready, preflight.issues
    assert plan.scene_ref["family_id"] == "qualification_archive_packages"
    assert plan.delivery_contract.default_preset_id == "attachment_package"
    assert plan.delivery_contract.required_artifact_keys == (
        "material_manifest",
        "material_package",
    )

    result = adapter.execute_approved_plan(
        plan,
        preflight,
        _approval(preflight),
        material_context=material_context,
    )

    assert result["status"] == "success", result.get("error_text")
    assert result["output_paths"] == {}
    assert result["material_package_receipt"]["status"] == "complete"
    package_zip = Path(result["material_package_paths"]["zip"])
    assert package_zip.is_file()
    with zipfile.ZipFile(package_zip) as archive:
        names = set(archive.namelist())
    assert "assets/01_certificates/certificate.pdf" in names
    assert "assets/02_business_license/business-license.pdf" in names
    assert "manifest/material_manifest.json" in names


def test_exam_compiler_rejects_unstructured_markdown(tmp_path):
    adapter = AssistantContentGenerationAdapter(tmp_path)

    with pytest.raises(ValueError, match="exam_content_compile_blocked"):
        adapter.compile_exam_markdown(
            session_id="session-exam",
            markdown="# 只有标题，没有题目",
        )


@pytest.mark.parametrize(
    ("markdown", "expected_code"),
    (
        (
            re.sub(r"（\d+\s*分）", "", EXAM_MARKDOWN),
            "score_coverage_incomplete",
        ),
        (
            re.sub(
                r"(?m)^\s+[A-D]\.\s+.*\n",
                "",
                EXAM_MARKDOWN,
                count=4,
            ),
            "choice_options_incomplete",
        ),
        (
            EXAM_MARKDOWN.replace(
                "七年级数学期中模拟试卷",
                "{{试卷标题}}",
            ),
            "unresolved_template_placeholder",
        ),
        (
            EXAM_MARKDOWN.replace(
                "2. 计算 `-4 + 7`",
                "3. 计算 `-4 + 7`",
            ),
            "question_number_sequence_invalid",
        ),
        (
            EXAM_MARKDOWN.replace(
                "本大题共 2 小题",
                "本大题共 3 小题",
            ),
            "section_question_count_mismatch",
        ),
        (
            EXAM_MARKDOWN.replace(
                "一、选择题\n1. A",
                "二、解答题\n1. A",
            ),
            "answer_group_section_mismatch",
        ),
    ),
)
def test_exam_compiler_rejects_structurally_incomplete_ai_drafts(
    tmp_path,
    markdown,
    expected_code,
):
    adapter = AssistantContentGenerationAdapter(tmp_path)

    with pytest.raises(ValueError, match=expected_code):
        adapter.compile_exam_markdown(
            session_id="session-exam",
            markdown=markdown,
        )


def test_exam_generation_rejects_draft_that_violates_explicit_user_requirements(
    tmp_path,
):
    plan = FormDocumentPlanBuilder().build(
        query="生成七年级数学试卷，满分100分，共10道选择题",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    service = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    )

    with pytest.raises(
        ValueError,
        match=(
            "requested_total_score_mismatch"
            ".*requested_question_count_mismatch"
            ".*requested_question_type_count_mismatch:choice"
        ),
    ):
        service.generate(
            ContentGenerationRequest(
                session_id="session-exam",
                turn_id="turn-exam",
                prompt=plan.intent,
                provider_id="fixed",
                model_id="fixed",
                capability_id=plan.capability_ref.capability_id,
                artifact_kind=plan.generation_contract.artifact_kind,
                prompt_profile_id=plan.generation_contract.prompt_profile_id,
            ),
            _FixedGateway(EXAM_MARKDOWN),
        )


def test_exam_generation_requires_full_analysis_when_user_explicitly_requests_it(
    tmp_path,
):
    plan = FormDocumentPlanBuilder().build(
        query="生成七年级数学试卷，并附详细解析",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    service = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path)
    )

    with pytest.raises(
        ValueError,
        match="requested_analysis_coverage_incomplete",
    ):
        service.generate(
            ContentGenerationRequest(
                session_id="session-exam",
                turn_id="turn-exam",
                prompt=plan.intent,
                provider_id="fixed",
                model_id="fixed",
                capability_id=plan.capability_ref.capability_id,
                artifact_kind=plan.generation_contract.artifact_kind,
                prompt_profile_id=plan.generation_contract.prompt_profile_id,
            ),
            _FixedGateway(EXAM_MARKDOWN),
        )


@pytest.mark.parametrize(
    ("query", "required"),
    (
        ("生成一份试卷，只要学生卷", ("student",)),
        ("生成一份试卷，只要答案卷", ("answer_key",)),
        (
            "生成一份试卷，需要学生卷和答案卷",
            ("student", "answer_key"),
        ),
    ),
)
def test_exam_delivery_selection_is_a_typed_plan_contract(query, required):
    plan = FormDocumentPlanBuilder().build(
        query=query,
        workspace=_workspace(),
        turn_id="turn-exam",
    )

    assert plan.delivery_contract.required_artifact_keys == required


def test_student_only_exam_plan_drives_terminal_assembly(tmp_path):
    plan = FormDocumentPlanBuilder().build(
        query="生成一份七年级数学试卷，只要学生卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(output_root=str(tmp_path / "outputs")),
    )
    draft = AssistantContentGenerationAdapter(
        tmp_path / "assistant"
    ).compile_exam_markdown(
        session_id="session-exam",
        markdown=EXAM_MARKDOWN,
        intent=plan.intent,
    )
    plan = bind_generated_draft(plan, draft)
    adapter = AssistantProductionAdapter()
    preflight = adapter.build_preflight(plan)

    assert preflight.ready, preflight.issues
    result = adapter.execute_approved_plan(plan, preflight, _approval(preflight))

    assert result["status"] == "success", result.get("error_text")
    assert set(result["output_paths"]) == {"student"}
    assert Path(result["output_paths"]["student"]).is_file()


def test_answer_only_exam_plan_drives_terminal_assembly(tmp_path):
    plan = FormDocumentPlanBuilder().build(
        query="生成一份七年级数学试卷，只要答案卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(output_root=str(tmp_path / "outputs")),
    )
    draft = AssistantContentGenerationAdapter(
        tmp_path / "assistant"
    ).compile_exam_markdown(
        session_id="session-exam",
        markdown=EXAM_MARKDOWN,
        intent=plan.intent,
    )
    plan = bind_generated_draft(plan, draft)
    adapter = AssistantProductionAdapter()
    preflight = adapter.build_preflight(plan)

    assert preflight.ready, preflight.issues
    result = adapter.execute_approved_plan(plan, preflight, _approval(preflight))

    assert result["status"] == "success", result.get("error_text")
    assert set(result["output_paths"]) == {"answer_key"}
    assert "答案速查" in _docx_text(Path(result["output_paths"]["answer_key"]))


def test_authorized_material_exam_generation_reaches_assembly_without_mutating_source(
    tmp_path,
):
    material = tmp_path / "七年级有理数复习材料.docx"
    source_document = Document()
    source_document.add_heading("复习范围", level=1)
    source_document.add_paragraph("知识点：有理数的加减法、绝对值与数轴。")
    source_document.save(material)
    source_bytes = material.read_bytes()
    workspace = replace(
        _workspace(),
        input_path=str(material),
        input_name=material.name,
        input_exists=True,
    )
    plan = FormDocumentPlanBuilder().build(
        query="根据材料生成一份七年级数学试卷",
        workspace=workspace,
        turn_id="turn-material-exam",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(output_root=str(tmp_path / "outputs")),
    )
    refs = tuple(dict(item) for item in plan.material_refs)
    ref_ids = tuple(str(item["path"]) for item in refs)
    fingerprints = attachment_fingerprints(refs)
    grant = DisclosureGrant(
        grant_id="grant-material-exam",
        session_id="session-material-exam",
        provider_id="fixed",
        model_id="fixed",
        allowed_refs=ref_ids,
        allowed_fields=("document_text",),
        content_fingerprints=fingerprints,
    )
    gateway = _FixedGateway(EXAM_MARKDOWN)

    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path / "assistant")
    ).generate(
        ContentGenerationRequest(
            session_id="session-material-exam",
            turn_id="turn-material-exam",
            prompt=plan.intent,
            provider_id="fixed",
            model_id="fixed",
            context_documents=refs,
            context_refs=ref_ids,
            context_fields=("document_text",),
            context_fingerprints=tuple(fingerprints.items()),
            disclosure_grant=grant,
            capability_id=plan.capability_ref.capability_id,
            artifact_kind=plan.generation_contract.artifact_kind,
            prompt_profile_id=plan.generation_contract.prompt_profile_id,
        ),
        gateway,
    )
    plan = bind_generated_draft(plan, draft)
    adapter = AssistantProductionAdapter()
    preflight = adapter.build_preflight(plan)
    result = adapter.execute_approved_plan(plan, preflight, _approval(preflight))

    assert plan.scene_ref["generation_mode"] == "generated_draft"
    assert preflight.ready, preflight.issues
    assert result["status"] == "success", result.get("error_text")
    assert {"student", "answer_key"} == set(result["output_paths"])
    provider_text = gateway.request.messages[0]["content"]
    assert "有理数的加减法" in provider_text
    assert str(tmp_path) not in provider_text
    assert material.read_bytes() == source_bytes


def test_planned_scene_is_not_silently_downgraded_to_generic_authoring():
    plan = FormDocumentPlanBuilder().build(
        query="为客户生成产品手册售前方案",
        workspace=_workspace(),
        turn_id="turn-product",
    )

    assert plan.capability_ref.route_id == "product_sales_document"
    assert plan.capability_ref.status == CAPABILITY_PLANNED
    assert not plan.generation_contract.required
    assert plan.unresolved_questions
    assert any("不会按通用文档降级" in item for item in plan.unresolved_questions)
    assert present_document_plan(plan).actions == ()


def test_ambiguous_route_cannot_start_generic_generation():
    plan = FormDocumentPlanBuilder().build(
        query="生成一份产品手册",
        workspace=_workspace(),
        turn_id="turn-ambiguous",
    )

    assert plan.capability_ref.status == CAPABILITY_PLANNED
    assert not plan.generation_contract.required
    assert plan.unresolved_questions


def test_official_plan_carries_task_level_document_type():
    workspace = replace(
        _workspace(),
        mode_id="official",
        mode_label="公文版",
        scene_id="official",
        template_id="official_gbt",
        document_type_id="notice",
    )
    plan = FormDocumentPlanBuilder().build(
        query="将通知套用公文格式",
        workspace=workspace,
        turn_id="turn-official",
    )

    assert plan.work_mode_id == "official"
    assert plan.production_contract.terminal_assembler == "official"
    assert plan.production_contract.document_type_id == "notice"


def test_delivery_contract_rejects_partial_exam_success(tmp_path):
    plan = FormDocumentPlanBuilder().build(
        query="生成一份七年级数学试卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(output_root=str(tmp_path / "outputs")),
    )
    draft = AssistantContentGenerationAdapter(
        tmp_path / "assistant"
    ).compile_exam_markdown(
        session_id="session-exam",
        markdown=EXAM_MARKDOWN,
    )
    plan = bind_generated_draft(plan, draft)
    student = tmp_path / "student.docx"
    Document().save(student)

    def incomplete_executor(_request, **_kwargs):
        return {
            "status": "success",
            "output_path": "",
            "output_paths": {"student": str(student)},
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }

    adapter = AssistantProductionAdapter(executor=incomplete_executor)
    preflight = adapter.build_preflight(plan)
    assert preflight.ready

    result = adapter.execute_approved_plan(plan, preflight, _approval(preflight))

    assert result["status"] == "failed"
    assert result["error_text"] == "delivery_contract_incomplete:answer_key"


def test_delivery_contract_rejects_paths_that_were_not_written(tmp_path):
    plan = FormDocumentPlanBuilder().build(
        query="生成一份七年级数学试卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(output_root=str(tmp_path / "outputs")),
    )
    draft = AssistantContentGenerationAdapter(
        tmp_path / "assistant"
    ).compile_exam_markdown(
        session_id="session-exam",
        markdown=EXAM_MARKDOWN,
    )
    plan = bind_generated_draft(plan, draft)

    def phantom_success(_request, **_kwargs):
        return {
            "status": "success",
            "output_path": "",
            "output_paths": {
                "student": str(tmp_path / "missing-student.docx"),
                "answer_key": str(tmp_path / "missing-answer.docx"),
            },
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }

    adapter = AssistantProductionAdapter(executor=phantom_success)
    preflight = adapter.build_preflight(plan)
    assert preflight.ready

    result = adapter.execute_approved_plan(plan, preflight, _approval(preflight))

    assert result["status"] == "failed"
    assert result["error_text"] == (
        "delivery_contract_incomplete:student,answer_key"
    )


def test_delivery_contract_rejects_empty_or_non_docx_exam_artifacts(tmp_path):
    plan = FormDocumentPlanBuilder().build(
        query="生成一份七年级数学试卷",
        workspace=_workspace(),
        turn_id="turn-exam",
    )
    plan = replace(
        plan,
        output_policy=OutputPolicy(output_root=str(tmp_path / "outputs")),
    )
    draft = AssistantContentGenerationAdapter(
        tmp_path / "assistant"
    ).compile_exam_markdown(
        session_id="session-exam",
        markdown=EXAM_MARKDOWN,
    )
    plan = bind_generated_draft(plan, draft)
    empty_student = tmp_path / "empty-student.docx"
    fake_answer = tmp_path / "fake-answer.txt"
    empty_student.touch()
    fake_answer.write_text("not a Word document", encoding="utf-8")

    adapter = AssistantProductionAdapter(
        executor=lambda _request, **_kwargs: {
            "status": "success",
            "output_path": str(empty_student),
            "output_paths": {
                "student": str(empty_student),
                "answer_key": str(fake_answer),
            },
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }
    )
    preflight = adapter.build_preflight(plan)

    result = adapter.execute_approved_plan(plan, preflight, _approval(preflight))

    assert result["status"] == "failed"
    assert result["error_text"] == (
        "delivery_contract_incomplete:student,answer_key"
    )


def test_legacy_generated_exam_plan_fails_closed_to_structured_contract():
    plan = DocumentPlan.from_dict(
        {
            "plan_id": "legacy-exam",
            "revision": 2,
            "intent": "生成试卷",
            "created_by_turn_id": "turn-legacy",
            "work_mode_id": "exam",
            "scene_ref": {"id": "exam", "generation_mode": "generated_draft"},
            "template_ref": {"id": "default"},
            "input_document_ref": {
                "path": "legacy-generic-draft.docx",
                "name": "legacy-generic-draft.docx",
            },
        }
    )

    assert plan.generation_contract.required
    assert plan.generation_contract.artifact_kind == ARTIFACT_KIND_EXAM
    assert plan.production_contract.terminal_assembler == "exam"
    assert plan.production_contract.accepted_suffixes == (".md", ".markdown")
    assert plan.delivery_contract.required_artifact_keys == ("student", "answer_key")
