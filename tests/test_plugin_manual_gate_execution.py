from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.shared import Cm
import pytest

from src.config.execution_config_integrity import execution_config_integrity_issue
from src.config.material_context import MaterialExecutionContext
from src.config.object_preflight_evidence import build_object_preflight_evidence
from src.config.plugin_manual_gate import (
    get_plugin_manual_gate,
    plugin_manual_gate_execution_issue,
)
from src.config.resolved import ResolvedConfig
from src.config.scene import ComplianceProfile, ObjectPreflightPolicy, SceneWorkspace
from src.config.scene_family_application import apply_planned_scene_family_defaults
from src.config.template import TemplateConfig
from src.modules.basic.page_setup import PageSetupModule
from src.pipeline.runner import Pipeline
from src.services.execution_session import build_execution_session_snapshot
from src.services.execution_session import resolution as execution_session_resolution


PROFESSIONAL_FAMILY_IDS = (
    "finance_quote_documents",
    "ip_patent_documents",
    "bilingual_translation_documents",
    "regulated_disclosure_documents",
)
PROFESSIONAL_GATE_ID = "professional_disclosure_review_gate"


def _expected_issue(family_id: str) -> str:
    return (
        f"plugin_manual_gate_required:{PROFESSIONAL_GATE_ID}:{family_id}:"
        "verified_external_receipt_missing"
    )


@pytest.mark.parametrize("family_id", PROFESSIONAL_FAMILY_IDS)
def test_professional_families_require_verified_external_receipt(family_id):
    config = ResolvedConfig(
        compliance_profile=ComplianceProfile(rule_family=family_id)
    )

    assert plugin_manual_gate_execution_issue(config) == _expected_issue(family_id)
    assert execution_config_integrity_issue(config) == _expected_issue(family_id)


def test_professional_gate_does_not_infer_family_from_pack_schema_or_alias():
    config = ResolvedConfig(
        compliance_profile=ComplianceProfile(rule_family="basic_format")
    )
    config.scene_id = "professional_disclosure"
    config.category = "professional_disclosure"
    config.input_source_profile.material_schema_id = (
        "regulated_disclosure_materials_v1"
    )

    assert plugin_manual_gate_execution_issue(config) == ""

    config.compliance_profile.rule_family = "regulated-disclosure-documents"
    assert plugin_manual_gate_execution_issue(config) == ""


def test_regulated_family_makes_execution_session_not_ready(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    plan_path = tmp_path / "plan_a.json"
    template_path = tmp_path / "template_a.json"
    source = tmp_path / "source.docx"
    plan_path.write_text('{"name":"plan"}', encoding="utf-8")
    template_path.write_text('{"name":"template"}', encoding="utf-8")
    Document().save(source)
    monkeypatch.setattr(
        execution_session_resolution,
        "get_scene_entry",
        lambda resource_id, *, mode_id: SimpleNamespace(
            path=plan_path,
            source_type="user",
        ),
    )
    monkeypatch.setattr(
        execution_session_resolution,
        "get_template_entry",
        lambda resource_id, *, mode_id: SimpleNamespace(
            path=template_path,
            source_type="user",
        ),
    )

    scene = SceneWorkspace(
        scene_id="plan_a",
        mode_id="custom",
        template_id="template_a",
        compatible_template_ids=["template_a"],
    )
    applied = apply_planned_scene_family_defaults(
        scene,
        family_id="regulated_disclosure_documents",
    )
    assert applied.applied is True
    evidence = build_object_preflight_evidence(scene, source)

    snapshot = build_execution_session_snapshot(
        mode_id="custom",
        scene=scene,
        template=TemplateConfig(),
        material_context=MaterialExecutionContext(mode_id="custom"),
        input_path=source,
        output_root=tmp_path / "runs",
        plan_id="plan_a",
        template_id="template_a",
        object_preflight_confirmation_revision=evidence.source_revision,
        object_preflight_confirmation_digest=evidence.evidence_digest,
    )

    assert snapshot.ready is False
    assert _expected_issue("regulated_disclosure_documents") in snapshot.issues
    assert snapshot.input_ref.frozen_path == ""


@pytest.mark.parametrize("family_id", PROFESSIONAL_FAMILY_IDS)
def test_pipeline_blocks_professional_family_before_mutation_or_output(
    tmp_path: Path,
    family_id: str,
):
    source = tmp_path / f"{family_id}.docx"
    document = Document()
    document.sections[0].top_margin = Cm(1)
    document.save(source)
    before = source.read_bytes()
    output_dir = tmp_path / f"out-{family_id}"
    config = ResolvedConfig(
        compliance_profile=ComplianceProfile(rule_family=family_id)
    )

    result = Pipeline(
        [PageSetupModule()],
        config,
        output_dir=str(output_dir),
    ).execute(str(source))

    assert result.success is False
    assert result.status == "failed"
    assert result.error == (
        "pipeline_configuration_invalid:" + _expected_issue(family_id)
    )
    assert result.context is None
    assert result.tracker is None
    assert result.output_paths == {}
    assert source.read_bytes() == before
    assert round(Document(source).sections[0].top_margin.cm, 2) == 1.0
    assert not output_dir.exists()


@pytest.mark.parametrize("rule_family", ("exam_teaching", "journal_en"))
def test_normal_exam_and_journal_families_remain_executable(
    tmp_path: Path,
    rule_family: str,
):
    source = tmp_path / f"{rule_family}.docx"
    Document().save(source)
    config = ResolvedConfig(
        compliance_profile=ComplianceProfile(
            rule_family=rule_family,
            count_profile_id={
                "exam_teaching": "exam_items",
                "journal_en": "journal_words",
            }[rule_family],
            object_preflight=ObjectPreflightPolicy(enabled=False),
        )
    )

    assert plugin_manual_gate_execution_issue(config) == ""
    result = Pipeline([], config, output_dir=str(tmp_path / "out")).execute(
        str(source)
    )
    assert result.success is True


def test_import_ai_gate_remains_a_non_executable_pack_boundary():
    gate = get_plugin_manual_gate("import_ai_boundary")
    scene = SceneWorkspace(
        scene_id="import_ai_boundary",
        category="import_ai_boundary",
    )

    application = apply_planned_scene_family_defaults(
        scene,
        family_id="import_ai_boundary",
    )

    assert gate.blocking_family_ids == ()
    assert application.applied is False
    assert scene.compliance_profile.rule_family == "basic_format"
