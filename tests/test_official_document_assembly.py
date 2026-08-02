import json
from dataclasses import replace
from types import SimpleNamespace

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

from src.config import master_library
from src.shared.engine.official_document_assembly import (
    _placeholder_replacements,
    assemble_official_document_docx,
)
from src.config.official_document_profiles import (
    get_official_document_assembly_contract,
)
from src.config.master_library import default_master
from src.shared.engine import official_document_assembly
from src.pipeline.runner import _official_delivery_versions


def _all_docx_text(path):
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_official_replacements_support_numbered_and_custom_placeholders():
    contract = get_official_document_assembly_contract("notice")

    replacements = _placeholder_replacements(
        contract,
        {
            "organization": "示例市档案局",
            "发文机关2": "示例市财政局",
            "自定义字段1": "自定义内容",
        },
    )

    assert replacements["{{@text:发文机关1}}"] == "示例市档案局"
    assert replacements["{{@text:发文机关2}}"] == "示例市财政局"
    assert replacements["{{@text:自定义字段1}}"] == "自定义内容"

    renamed = _placeholder_replacements(
        contract,
        {
            "organization": "示例市档案局",
            "发文机关3": "示例市财政局",
        },
        field_aliases={"发文机关2": "organization"},
    )
    assert renamed["{{@text:发文机关2}}"] == "示例市档案局"
    assert renamed["{{@text:发文机关3}}"] == "示例市财政局"


def test_execution_frozen_master_never_falls_back_to_live_contract_master(
    tmp_path,
    monkeypatch,
):
    standard = default_master("official")
    assert standard is not None
    frozen_standard = replace(standard, execution_frozen=True)
    monkeypatch.setattr(
        master_library,
        "get_master",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("frozen execution master fell back to live library")
        ),
    )
    monkeypatch.setattr(
        master_library,
        "default_master",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("frozen execution master fell back to live default")
        ),
    )

    result = assemble_official_document_docx(
        "letter",
        {},
        tmp_path,
        master=frozen_standard,
    )

    assert result.status == "master_incompatible"
    assert result.master_id == "official_gbt_standard"


def test_official_assembly_applies_renamed_alias_and_added_series_to_user_master(
    tmp_path,
):
    builtin_master = default_master("official")
    assert builtin_master is not None
    master_path = tmp_path / "renamed-series-master.docx"
    document = Document(builtin_master.docx_path)
    document.add_paragraph("别名一：{{@text:发文机关2}}；新增项：{{@text:发文机关3}}")
    document.save(master_path)
    user_master = replace(
        builtin_master,
        docx_path=master_path,
        source_type="user",
        readonly=False,
        manifest_path=None,
    )

    result = assemble_official_document_docx(
        "notice",
        {
            "title": "别名链路验证通知",
            "body": "验证公文正式件装配链路。",
            "organization": "第一机关",
            "发文机关3": "第二机关",
            "document_no": "测发〔2026〕1号",
            "issue_date": "2026年7月12日",
        },
        tmp_path,
        field_aliases={"发文机关2": "organization"},
        filename="renamed-series-output.docx",
        master=user_master,
    )

    assert result.ok is True
    text = _all_docx_text(result.docx_path)
    assert "别名一：第一机关；新增项：第二机关" in text
    assert "{{@text:发文机关2}}" not in text
    assert "{{@text:发文机关3}}" not in text


def test_user_master_keeps_explicit_small_title_song_family(tmp_path):
    builtin_master = default_master("official")
    assert builtin_master is not None
    master_path = tmp_path / "explicit-small-title-song.docx"
    document = Document(builtin_master.docx_path)
    redhead = next(
        paragraph
        for paragraph in document.paragraphs
        if "{{@text:official_organization}}" in paragraph.text
    )
    run_fonts = redhead.runs[0]._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia"):
        run_fonts.set(qn(f"w:{attribute}"), "方正小标宋_GBK")
    redhead.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    redhead.paragraph_format.line_spacing = Pt(28)
    document.save(master_path)
    user_master = replace(
        builtin_master,
        docx_path=master_path,
        source_type="user",
        readonly=False,
        manifest_path=None,
    )

    result = assemble_official_document_docx(
        "notice",
        {
            "title": "用户模板字体保真验证",
            "body": "正文内容。",
            "organization": "用户指定机关",
            "document_no": "示发〔2026〕1号",
            "issue_date": "2026年7月31日",
        },
        tmp_path,
        filename="explicit-small-title-song-output.docx",
        master=user_master,
    )

    assert result.ok is True
    output = Document(result.docx_path)
    rendered_redhead = next(
        paragraph
        for paragraph in output.paragraphs
        if paragraph.text == "用户指定机关"
    )
    output_fonts = (
        rendered_redhead.runs[0]
        ._element.get_or_add_rPr()
        .get_or_add_rFonts()
    )
    assert output_fonts.get(qn("w:eastAsia")) == "方正小标宋_GBK"
    assert (
        rendered_redhead.paragraph_format.line_spacing_rule
        == WD_LINE_SPACING.AT_LEAST
    )
    assert rendered_redhead.paragraph_format.line_spacing.pt >= 60


def test_assemble_official_notice_docx_replaces_master_placeholders(tmp_path):
    result = assemble_official_document_docx(
        "notice",
        {
            "title": "关于开展档案检查的通知",
            "body": "请各部门按要求完成自查并提交材料。",
            "organization": "示例市档案局",
            "document_no": "示档发〔2026〕3号",
            "issue_date": "2026年7月9日",
            "issuer": "示例市档案局",
            "recipient": "各部门",
            "copy_scope": "办公室",
            "printing_org": "示例市档案局办公室",
            "printing_date": "2026年7月9日",
        },
        tmp_path,
    )

    assert result.ok is True
    assert result.docx_path is not None
    assert result.docx_path.is_file()
    assert result.profile_id == "notice"
    assert result.master_id == "official_gbt_standard"
    assert result.material_schema_ids == ("official_document_v1",)
    assert "official_title" in result.replaced_placeholders
    assert "official_body" in result.replaced_placeholders
    assert result.unresolved_placeholders == ()
    assert result.internal_review_docx_path is not None
    assert result.internal_review_docx_path.is_file()
    assert result.archive_manifest_path is not None
    assert result.archive_manifest_path.is_file()
    assert result.archive_manifest_markdown_path is not None
    assert result.archive_manifest_markdown_path.is_file()
    assert result.output_paths == {
        "official_docx": str(result.docx_path),
        "internal_review_docx": str(result.internal_review_docx_path),
        "archive_manifest": str(result.archive_manifest_path),
        "archive_manifest_md": str(result.archive_manifest_markdown_path),
    }

    text = _all_docx_text(result.docx_path)
    formal_document = Document(result.docx_path)
    review_text = _all_docx_text(result.internal_review_docx_path)
    archive_manifest = json.loads(
        result.archive_manifest_path.read_text(encoding="utf-8")
    )
    assert "关于开展档案检查的通知" in text
    assert "请各部门按要求完成自查并提交材料。" in text
    assert "{{@text:official_title}}" not in text
    assert "{{@text:official_body}}" not in text
    assert "\u5185\u90e8\u5ba1\u9605\u7a3f" not in text
    assert "\u5185\u90e8\u5ba1\u9605\u7a3f" in review_text
    assert archive_manifest["status"] == "ok"
    assert archive_manifest["profile_id"] == "notice"
    assert archive_manifest["archive_fields"]["title"]
    assert archive_manifest["output_paths"] == result.output_paths
    assert "official_docx" in archive_manifest["delivery_versions"]
    assert "internal_review_docx" in archive_manifest["delivery_versions"]
    assert "archive_manifest" in archive_manifest["delivery_versions"]
    assert (
        "material bindings are applied by the official document runtime assembly"
        in archive_manifest["boundaries"]
    )
    assert not any(
        "does not execute runtime assembly yet" in boundary
        for boundary in archive_manifest["boundaries"]
    )
    assert "版式占位符说明" not in text
    assert "母版占位符说明" not in text
    assert len(formal_document.sections) == 1
    assert not any(
        paragraph._p.xpath(".//w:sectPr")
        for paragraph in formal_document.paragraphs
    )


def test_official_assembly_publishes_only_the_selected_word_version(tmp_path):
    values = {
        "title": "关于开展档案检查的通知",
        "body": "请各部门按要求完成自查并提交材料。",
        "organization": "示例市档案局",
        "document_no": "示档发〔2026〕3号",
        "issue_date": "2026年7月9日",
    }

    formal = assemble_official_document_docx(
        "notice",
        values,
        tmp_path / "formal",
        delivery_versions=("official_docx",),
    )
    review = assemble_official_document_docx(
        "notice",
        values,
        tmp_path / "review",
        delivery_versions=("internal_review_docx",),
    )

    assert formal.ok is True
    assert set(formal.output_paths) == {"official_docx"}
    assert formal.internal_review_docx_path is None
    assert formal.archive_manifest_path is None
    assert review.ok is True
    assert set(review.output_paths) == {"internal_review_docx"}
    assert review.docx_path is None
    assert review.archive_manifest_path is None
    assert "内部审阅稿" in _all_docx_text(review.internal_review_docx_path)


def test_official_delivery_preset_controls_the_real_assembly_bundle():
    assert _official_delivery_versions(
        SimpleNamespace(default_delivery_preset_id="formal")
    ) == ("official_docx",)
    assert _official_delivery_versions(
        SimpleNamespace(default_delivery_preset_id="internal_review")
    ) == ("internal_review_docx",)
    assert _official_delivery_versions(
        SimpleNamespace(default_delivery_preset_id="formal_minutes")
    ) == (
        "official_docx",
        "internal_review_docx",
        "archive_manifest",
    )


def test_assemble_official_body_uses_real_semantic_heading_paragraphs(tmp_path):
    result = assemble_official_document_docx(
        "notice",
        {
            "title": "关于开展档案检查的通知",
            "body": (
                "各区县档案主管部门、市直各单位：\n"
                "一、工作安排\n"
                "（一）检查范围\n"
                "1. 完成材料核验\n"
                "（1）核对归档目录\n"
                "请于规定时间内报送材料。"
            ),
            "organization": "示例市档案局",
            "document_no": "示档发〔2026〕1号",
            "issuer": "示例市档案局",
            "issue_date": "2026年7月31日",
        },
        tmp_path,
        filename="semantic-body.docx",
    )

    assert result.ok is True
    document = Document(result.docx_path)
    by_text = {
        paragraph.text.strip(): paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    }
    assert by_text["各区县档案主管部门、市直各单位："].style.name == "Normal"
    assert by_text["一、工作安排"].style.name == "Heading 1"
    assert by_text["（一）检查范围"].style.name == "Heading 2"
    assert by_text["1. 完成材料核验"].style.name == "Heading 3"
    assert by_text["（1）核对归档目录"].style.name == "Heading 4"
    assert by_text["请于规定时间内报送材料。"].style.name == "Normal"
    for text in (
        "各区县档案主管部门、市直各单位：",
        "一、工作安排",
        "（一）检查范围",
        "1. 完成材料核验",
        "（1）核对归档目录",
        "请于规定时间内报送材料。",
    ):
        assert not by_text[text]._p.xpath(".//w:br")
    for text in (
        "一、工作安排",
        "（一）检查范围",
        "1. 完成材料核验",
        "（1）核对归档目录",
    ):
        paragraph = by_text[text]
        assert paragraph.paragraph_format.keep_with_next is True
        assert paragraph.paragraph_format.keep_together is True
        assert paragraph.paragraph_format.page_break_before is False
    assert by_text["请于规定时间内报送材料。"].paragraph_format.keep_with_next is True


def test_builtin_official_typography_uses_unemboldened_small_title_song_roles(
    tmp_path,
):
    result = assemble_official_document_docx(
        "notice",
        {
            "title": "关于开展档案检查的通知",
            "body": "一、开展材料检查。",
            "organization": "示例市档案局",
            "document_no": "示档发〔2026〕1号",
            "issuer": "示例市档案局",
            "issue_date": "2026年8月2日",
        },
        tmp_path,
        filename="gbt-typography.docx",
    )

    assert result.ok is True
    document = Document(result.docx_path)
    redhead = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "示例市档案局"
        and paragraph.alignment is not None
        and paragraph.runs
        and paragraph.runs[0].font.color.rgb is not None
    )
    title = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "关于开展档案检查的通知"
    )
    issuer = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "示例市档案局"
    ][-1]

    assert redhead.runs[0].bold is False
    assert title.runs[0].bold is False
    assert title.paragraph_format.space_before.pt == 56
    assert abs(issuer.paragraph_format.right_indent.mm - 11.0) < 0.1


def test_assemble_official_body_removes_recipient_owned_duplicate_line(tmp_path):
    result = assemble_official_document_docx(
        "notice",
        {
            "title": "部署会议通知",
            "recipient": "各区县档案主管部门、市直各单位：",
            "body": (
                "各区县档案主管部门，市直各单位：\n"
                "现将有关事项通知如下：\n"
                "一、会议安排"
            ),
            "organization": "示例市档案局",
            "document_no": "示档发〔2026〕1号",
            "issue_date": "2026年7月31日",
        },
        tmp_path,
        filename="deduplicated-recipient.docx",
    )

    assert result.ok is True
    document = Document(result.docx_path)
    recipient_lines = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if "各区县档案主管部门" in paragraph.text
    ]
    assert recipient_lines == ["各区县档案主管部门、市直各单位："]
    assert any(
        paragraph.text.strip() == "一、会议安排"
        and paragraph.style.name == "Heading 1"
        for paragraph in document.paragraphs
    )


def test_assemble_official_minutes_uses_meeting_schema_bindings(tmp_path):
    result = assemble_official_document_docx(
        "minutes",
        {
            "title": "专题协调会议纪要",
            "body": "会议研究了近期材料归档和流程优化事项。",
            "organization": "示例市办公室",
            "document_no": "会议纪要〔2026〕8号",
            "issue_date": "2026年7月9日",
            "meeting_date": "2026年7月9日上午",
            "participants": "张三、李四、王五",
        },
        tmp_path,
        filename="minutes.docx",
    )

    assert result.ok is True
    assert result.material_schema_ids == (
        "official_document_v1",
        "administrative_meeting_fields_v1",
    )
    assert "official_meeting_time" in result.replaced_placeholders
    assert "official_meeting_attendees" in result.replaced_placeholders
    assert result.unresolved_placeholders == ()
    assert result.docx_path is not None
    document = Document(result.docx_path)
    text = _all_docx_text(result.docx_path)
    assert "会议时间：2026年7月9日上午" in text
    assert "出席：张三、李四、王五" in text
    assert result.master_id == "official_gbt_minutes"
    assert "2026年7月9日" in text
    assert "印发机关" not in text


def test_builtin_assembly_normalizes_master_owned_suffixes_and_imprint_text(
    tmp_path,
):
    order_result = assemble_official_document_docx(
        "order",
        {
            "title": "关于公布管理办法的命令",
            "body": "现予公布。",
            "organization": "示例市人民政府令",
            "document_no": "第3号",
            "issue_date": "2026年8月2日",
            "signer": "市长 张明",
        },
        tmp_path / "order",
        filename="order.docx",
    )
    minutes_result = assemble_official_document_docx(
        "minutes",
        {
            "title": "专题会议纪要",
            "body": "会议研究了有关事项。",
            "organization": "示例市教育局专题会议纪要",
            "document_no": "第2期",
            "issue_date": "2026年8月2日",
        },
        tmp_path / "minutes",
        filename="minutes.docx",
    )
    notice_result = assemble_official_document_docx(
        "notice",
        {
            "title": "关于归档检查的通知",
            "body": "请按要求完成检查。",
            "organization": "示例市档案局",
            "document_no": "示档发〔2026〕3号",
            "issue_date": "2026年8月2日",
            "copy_scope": "办公室。",
            "printing_org": "示例市档案局办公室",
            "printing_date": "2026年8月2日印发",
        },
        tmp_path / "notice",
        filename="notice.docx",
    )

    assert order_result.ok and minutes_result.ok and notice_result.ok
    assert "示例市人民政府令令" not in _all_docx_text(order_result.docx_path)
    assert "示例市人民政府令" in _all_docx_text(order_result.docx_path)
    assert "专题会议纪要纪要" not in _all_docx_text(minutes_result.docx_path)
    assert "示例市教育局专题会议纪要" in _all_docx_text(
        minutes_result.docx_path
    )
    notice = Document(notice_result.docx_path)
    imprint = notice.tables[-1]
    assert len(imprint.rows) == 1
    paragraphs = imprint.rows[0].cells[0].paragraphs
    assert paragraphs[0].text == "抄送：办公室。"
    assert paragraphs[1].text == ""
    assert paragraphs[2].text == (
        "示例市档案局办公室\t2026年8月2日印发"
    )


def test_public_announcement_can_omit_a_document_number(tmp_path):
    result = assemble_official_document_docx(
        "announcement",
        {
            "title": "关于开放馆藏档案的公告",
            "body": "现依法向社会开放一批馆藏档案。",
            "organization": "示例市档案馆",
            "issue_date": "2026年8月2日",
        },
        tmp_path,
        filename="announcement.docx",
    )

    assert result.ok is True
    assert "document_no" not in result.missing_required_fields
    assert "official_document_no" in result.replaced_placeholders
    assert not any(
        "official_document_no" in paragraph.text
        for paragraph in Document(result.docx_path).paragraphs
    )


def test_assemble_official_document_omits_empty_optional_paragraphs(tmp_path):
    result = assemble_official_document_docx(
        "notice",
        {
            "title": "无附件通知",
            "body": "本通知没有主送、附件说明或署名机关。",
            "organization": "示例单位",
            "document_no": "示发〔2026〕9号",
            "issue_date": "2026年7月10日",
        },
        tmp_path,
        filename="without_optional_paragraphs.docx",
    )

    assert result.ok is True
    assert result.docx_path is not None
    document = Document(result.docx_path)
    paragraph_texts = [paragraph.text.strip() for paragraph in document.paragraphs]
    assert "附件：" not in paragraph_texts
    assert "：" not in paragraph_texts
    assert "official_attachment_note" not in result.replaced_placeholders
    assert "official_recipient" not in result.replaced_placeholders
    assert "official_issuer" in result.replaced_placeholders
    issuer = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "示例单位"
        and paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    )
    assert abs(issuer.paragraph_format.right_indent.mm - 11.0) < 0.1


def test_assemble_official_document_honors_removed_builtin_fields(tmp_path):
    result = assemble_official_document_docx(
        "notice",
        {
            "body": "正文保留，标题由资料包明确删除。",
            "organization": "示例单位",
            "document_no": "示发〔2026〕11号",
            "issue_date": "2026年7月11日",
        },
        tmp_path,
        filename="removed-title.docx",
        removed_field_keys={"title"},
    )

    assert result.ok is True
    assert result.missing_required_fields == ()
    assert result.docx_path is not None
    assert "{{@text:official_title}}" not in _all_docx_text(result.docx_path)


def test_assemble_official_document_generates_optional_review_pdf(tmp_path):
    rendered_sources = []

    def fake_renderer(source_path, target_path):
        rendered_sources.append(source_path)
        assert "内部审阅稿" in _all_docx_text(source_path)
        target_path.write_bytes(b"%PDF-1.4\n% review\n")

    result = assemble_official_document_docx(
        "notice",
        {
            "title": "带审阅 PDF 的通知",
            "body": "用于验证审阅 PDF 可选交付。",
            "organization": "示例单位",
            "document_no": "示发〔2026〕10号",
            "issue_date": "2026年7月10日",
        },
        tmp_path,
        filename="review_pdf.docx",
        generate_review_pdf=True,
        review_pdf_renderer_resolver=lambda: ("fake_pdf", fake_renderer),
    )

    assert result.ok is True
    assert result.review_pdf_status == "generated"
    assert result.review_pdf_renderer == "fake_pdf"
    assert result.review_pdf_issue == ""
    assert result.review_pdf_path == tmp_path / "review_pdf_review.pdf"
    assert result.review_pdf_path.is_file()
    assert len(rendered_sources) == 1
    assert rendered_sources[0].name == result.internal_review_docx_path.name
    assert rendered_sources[0] != result.internal_review_docx_path
    assert result.output_paths["review_pdf"] == str(result.review_pdf_path)
    archive_manifest = json.loads(result.archive_manifest_path.read_text(encoding="utf-8"))
    assert archive_manifest["review_pdf"] == {
        "status": "generated",
        "path": str(result.review_pdf_path),
        "renderer": "fake_pdf",
        "issue": "",
    }
    assert archive_manifest["output_paths"]["review_pdf"] == str(result.review_pdf_path)


def test_assemble_official_document_degrades_when_review_pdf_renderer_unavailable(tmp_path):
    result = assemble_official_document_docx(
        "notice",
        {
            "title": "无 PDF 渲染器通知",
            "body": "DOCX 交付仍应成功。",
            "organization": "示例单位",
            "document_no": "示发〔2026〕11号",
            "issue_date": "2026年7月10日",
        },
        tmp_path,
        filename="renderer_unavailable.docx",
        generate_review_pdf=True,
        review_pdf_renderer_resolver=lambda: None,
    )

    assert result.ok is True
    assert result.review_pdf_status == "renderer_unavailable"
    assert result.review_pdf_path is None
    assert result.review_pdf_renderer == ""
    assert result.review_pdf_issue == "docx_to_pdf_renderer_unavailable"
    assert "review_pdf" not in result.output_paths
    archive_manifest = json.loads(result.archive_manifest_path.read_text(encoding="utf-8"))
    assert archive_manifest["review_pdf"]["status"] == "renderer_unavailable"
    assert archive_manifest["review_pdf"]["path"] == ""
    assert "review_pdf" not in archive_manifest["output_paths"]


def test_assemble_official_document_degrades_and_removes_failed_review_pdf(tmp_path):
    def broken_renderer(_source_path, target_path):
        target_path.write_bytes(b"partial")
        raise RuntimeError("renderer crashed")

    result = assemble_official_document_docx(
        "notice",
        {
            "title": "PDF 失败通知",
            "body": "PDF 失败不能破坏 DOCX 交付。",
            "organization": "示例单位",
            "document_no": "示发〔2026〕12号",
            "issue_date": "2026年7月10日",
        },
        tmp_path,
        filename="renderer_failed.docx",
        generate_review_pdf=True,
        review_pdf_renderer_resolver=lambda: ("broken_pdf", broken_renderer),
    )

    assert result.ok is True
    assert result.review_pdf_status == "render_failed"
    assert result.review_pdf_path is None
    assert result.review_pdf_renderer == "broken_pdf"
    assert "renderer crashed" in result.review_pdf_issue
    assert "review_pdf" not in result.output_paths
    assert not (tmp_path / "renderer_failed_review.pdf").exists()


def test_assemble_official_document_blocks_missing_required_fields(tmp_path):
    result = assemble_official_document_docx(
        "notice",
        {
            "title": "缺正文的通知",
            "organization": "示例单位",
            "document_no": "示发〔2026〕1号",
            "issue_date": "2026年7月9日",
        },
        tmp_path,
    )

    assert result.status == "missing_required_fields"
    assert result.docx_path is None
    assert result.output_paths == {}
    assert result.missing_required_fields == ("body",)


def test_assemble_official_document_blocks_unknown_profile(tmp_path):
    result = assemble_official_document_docx(
        "unknown_official_type",
        {
            "title": "未知文种",
            "body": "正文",
        },
        tmp_path,
    )

    assert result.status == "unknown_profile"
    assert result.docx_path is None
    assert result.output_paths == {}


def test_assemble_official_document_auto_routes_stale_builtin_master(tmp_path):
    common = default_master("official")
    assert common is not None

    result = assemble_official_document_docx(
        "letter",
        {
            "title": "测试函",
            "body": "函件正文。",
            "organization": "测试单位",
            "document_no": "测函〔2026〕1号",
            "recipient": "有关单位",
            "issue_date": "2026年7月11日",
            "copy_scope": "办公室",
        },
        tmp_path,
        master=common,
    )

    assert result.ok is True
    assert result.master_id == "official_gbt_letter"
    document = Document(result.docx_path)
    section = document.sections[0]
    assert abs(section.top_margin.mm - 23.5) < 0.1
    assert abs(section.footer_distance.mm - 20.0) < 0.1
    assert section.different_first_page_header_footer is True
    assert section.first_page_footer.paragraphs[0].text == ""
    assert any(
        cell.text == "抄送：办公室"
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    mark = document.paragraphs[0]
    assert mark.text == "测试单位"
    assert mark.runs[0].bold is False
    assert abs(mark.paragraph_format.left_indent.mm + 7.0) < 0.1
    document_number = next(
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if paragraph.text == "测函〔2026〕1号"
    )
    assert document_number.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_assemble_official_document_blocks_incompatible_user_layout(tmp_path):
    common = default_master("official")
    assert common is not None
    user_common = replace(common, source_type="user", readonly=False)

    result = assemble_official_document_docx(
        "minutes",
        {
            "title": "会议纪要",
            "body": "会议正文。",
            "organization": "测试单位",
            "document_no": "第1期",
            "issue_date": "2026年7月11日",
        },
        tmp_path,
        master=user_common,
    )

    assert result.status == "master_incompatible"
