from hashlib import sha256
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import pytest

from content_artifact_test_utils import compile_content_binding
from src.config.attachment_materials import (
    AttachmentBinding,
    AttachmentItem,
    AttachmentSourceKind,
)
from src.config.content_materials import (
    ContentInsertionRule,
    FileAssetRef,
    content_anchor_token,
)
from src.services.material_content.artifact_repository import ContentArtifactRepository
from src.services.material_content.composer import (
    ContentComposeCancelledError,
    ContentComposeError,
    ContentComposeInputs,
    ContentComposeRequest,
    ContentMaterialComposer,
)


def _rule(content_id: str) -> ContentInsertionRule:
    return ContentInsertionRule(
        f"content:{content_id}", content_id, content_anchor_token(content_id)
    )


def _inputs(bindings, rules, attachment_bindings=()) -> ContentComposeInputs:
    return ContentComposeInputs(
        content_bindings=tuple(bindings),
        content_rules=tuple(rules),
        attachment_bindings=tuple(attachment_bindings),
    )


def _target(path: Path, *tokens: str) -> None:
    document = Document()
    for token in tokens:
        document.add_paragraph(token)
    document.save(path)


def _attachment_binding(tmp_path: Path, role: str) -> AttachmentBinding:
    source = tmp_path / f"{role}.pdf"
    source.write_bytes(role.encode("utf-8"))
    payload = source.read_bytes()
    return AttachmentBinding(
        role=role,
        source_kind=AttachmentSourceKind.SINGLE_FILE,
        items=(
            AttachmentItem(
                item_id=f"{role}-1",
                label=role,
                file_ref=FileAssetRef(
                    source_path=str(source),
                    original_name=source.name,
                    media_type="application/pdf",
                    content_sha256=sha256(payload).hexdigest(),
                    byte_size=len(payload),
                ),
            ),
        ),
    )


def test_docx_artifact_is_a_normal_composer_input(tmp_path: Path) -> None:
    source = tmp_path / "content.docx"
    content = Document()
    content.add_heading("技术路线", 1)
    content.add_paragraph("正文")
    content.save(source)
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    binding = compile_content_binding(source, repository, content_id="route")
    rule = _rule("route")
    source.unlink()
    target = tmp_path / "target.docx"
    output = tmp_path / "output.docx"
    _target(target, rule.anchor_token)
    receipt = ContentMaterialComposer(repository=repository).compose(
        ContentComposeRequest(str(target), str(output), _inputs((binding,), (rule,)))
    )
    assert receipt.content_artifacts[0].artifact_id == binding.artifact_ref.artifact_id
    assert [item.text for item in Document(output).paragraphs[:2]] == ["技术路线", "正文"]


def test_same_resource_id_is_namespaced_by_content_id(tmp_path: Path) -> None:
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    bindings = []
    rules = []
    for index, color in enumerate(("red", "blue"), start=1):
        folder = tmp_path / f"content-{index}"
        folder.mkdir()
        Image.new("RGB", (10, 10), color).save(folder / "diagram.png")
        source = folder / "content.md"
        source.write_text("![图](diagram.png)", encoding="utf-8")
        content_id = f"part{index}"
        bindings.append(
            compile_content_binding(source, repository, content_id=content_id)
        )
        rules.append(_rule(content_id))
    target = tmp_path / "target.docx"
    output = tmp_path / "output.docx"
    _target(target, *(rule.anchor_token for rule in rules))
    receipt = ContentMaterialComposer(repository=repository).compose(
        ContentComposeRequest(str(target), str(output), _inputs(bindings, rules))
    )
    keys = [item.key for item in receipt.resource_materialization.entries]
    assert {item.content_id for item in keys} == {"part1", "part2"}
    assert len(keys) == 2


def test_multiple_rules_follow_input_order_and_receipt_is_deterministic(tmp_path: Path) -> None:
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    bindings = []
    rules = []
    for content_id in ("first", "second"):
        source = tmp_path / f"{content_id}.md"
        source.write_text(content_id, encoding="utf-8")
        bindings.append(compile_content_binding(source, repository, content_id=content_id))
        rules.append(_rule(content_id))
    inputs = _inputs(bindings, rules)
    target = tmp_path / "target.docx"
    _target(target, rules[0].anchor_token, rules[1].anchor_token)
    composer = ContentMaterialComposer(repository=repository)
    first = composer.compose(ContentComposeRequest(str(target), str(tmp_path / "a.docx"), inputs))
    second = composer.compose(ContentComposeRequest(str(target), str(tmp_path / "b.docx"), inputs))
    assert [item.content_id for item in first.content_artifacts] == ["first", "second"]
    assert first.output_sha256 == second.output_sha256


def test_anchor_failure_preserves_existing_output(tmp_path: Path) -> None:
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    source = tmp_path / "content.md"
    source.write_text("正文", encoding="utf-8")
    binding = compile_content_binding(source, repository)
    rule = _rule(binding.content_id)
    target = tmp_path / "target.docx"
    _target(target, "missing anchor")
    output = tmp_path / "output.docx"
    output.write_bytes(b"existing")
    with pytest.raises(ContentComposeError):
        ContentMaterialComposer(repository=repository).compose(
            ContentComposeRequest(str(target), str(output), _inputs((binding,), (rule,)))
        )
    assert output.read_bytes() == b"existing"


def test_every_anchor_is_preflighted_before_rendering(tmp_path: Path) -> None:
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    bindings = []
    rules = []
    for content_id in ("valid", "missing"):
        source = tmp_path / f"{content_id}.md"
        source.write_text(content_id, encoding="utf-8")
        bindings.append(compile_content_binding(source, repository, content_id=content_id))
        rules.append(_rule(content_id))
    target = tmp_path / "target.docx"
    _target(target, rules[0].anchor_token)
    output = tmp_path / "output.docx"
    with pytest.raises(ContentComposeError):
        ContentMaterialComposer(repository=repository).compose(
            ContentComposeRequest(str(target), str(output), _inputs(bindings, rules))
        )
    assert not output.exists()
    assert Document(target).paragraphs[0].text == rules[0].anchor_token


@pytest.mark.parametrize(
    ("corruption", "expected_code"),
    (
        ("bookmark", "attachment_anchor_not_isolated_body"),
        ("field", "attachment_anchor_not_isolated_body"),
        ("section", "attachment_anchor_section_boundary"),
    ),
)
def test_attachment_anchor_preflight_preserves_source_and_creates_no_output(
    tmp_path: Path,
    corruption: str,
    expected_code: str,
) -> None:
    valid = _attachment_binding(tmp_path, "a_valid")
    invalid = _attachment_binding(tmp_path, "z_invalid")
    target = tmp_path / "target.docx"
    output = tmp_path / "output.docx"
    document = Document()
    document.add_paragraph(valid.anchor_token)
    paragraph = document.add_paragraph(invalid.anchor_token)
    if corruption == "bookmark":
        bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), "11")
        bookmark.set(qn("w:name"), "must-not-be-deleted")
        paragraph._p.append(bookmark)
    elif corruption == "field":
        run = OxmlElement("w:r")
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "begin")
        run.append(field)
        paragraph._p.append(run)
    else:
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    document.save(target)
    source_before = target.read_bytes()

    with pytest.raises(ContentComposeError) as raised:
        ContentMaterialComposer(
            repository=ContentArtifactRepository(tmp_path / "artifacts")
        ).compose(
            ContentComposeRequest(
                str(target),
                str(output),
                _inputs((), (), (valid, invalid)),
            )
        )

    assert {
        item.code for item in raised.value.diagnostics
    } == {f"attachment_renderer_{expected_code}"}
    assert target.read_bytes() == source_before
    assert not output.exists()
    assert not list(tmp_path.glob(".output.docx.compose-*"))


def test_output_equal_to_source_is_blocked(tmp_path: Path) -> None:
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    source = tmp_path / "content.md"
    source.write_text("正文", encoding="utf-8")
    binding = compile_content_binding(source, repository)
    rule = _rule(binding.content_id)
    target = tmp_path / "target.docx"
    _target(target, rule.anchor_token)
    with pytest.raises(ContentComposeError):
        ContentMaterialComposer(repository=repository).compose(
            ContentComposeRequest(str(target), str(target), _inputs((binding,), (rule,)))
        )


def test_cancellation_before_publish_is_atomic(tmp_path: Path) -> None:
    repository = ContentArtifactRepository(tmp_path / "artifacts")
    source = tmp_path / "content.md"
    source.write_text("正文", encoding="utf-8")
    binding = compile_content_binding(source, repository)
    rule = _rule(binding.content_id)
    target = tmp_path / "target.docx"
    output = tmp_path / "output.docx"
    _target(target, rule.anchor_token)
    output.write_bytes(b"existing")
    calls = 0
    def cancel() -> bool:
        nonlocal calls
        calls += 1
        return calls >= 5
    with pytest.raises(ContentComposeCancelledError):
        ContentMaterialComposer(repository=repository).compose(
            ContentComposeRequest(str(target), str(output), _inputs((binding,), (rule,))),
            cancel_check=cancel,
        )
    assert output.read_bytes() == b"existing"
    assert not list(tmp_path.glob("*.tmp"))


def test_composer_rejects_compression_bomb_before_python_docx_open(
    tmp_path: Path,
) -> None:
    target = tmp_path / "bomb.docx"
    with ZipFile(target, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"A" * 1_000_000)
    output = tmp_path / "output.docx"
    output.write_bytes(b"existing")

    with pytest.raises(ContentComposeError) as raised:
        ContentMaterialComposer(
            repository=ContentArtifactRepository(tmp_path / "artifacts")
        ).compose(
            ContentComposeRequest(
                str(target),
                str(output),
                _inputs((), ()),
            )
        )

    assert "zip_compression_ratio_exceeded" in str(raised.value)
    assert output.read_bytes() == b"existing"
