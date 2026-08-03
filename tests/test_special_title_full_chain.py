from docx import Document

from src.config.resolved import ResolvedConfig
from src.config.special_title_rules import special_title_selector
from src.config.template import (
    HeadingLevelBindingConfig,
    PageNumberPhaseConfig,
    StyleConfig,
)
from src.modules.basic.header_footer import HeaderFooterModule
from src.modules.basic.paragraph_style import ParagraphStyleModule
from src.modules.basic.section_format import SectionFormatModule
from src.modules.structure.heading_numbering import HeadingNumberingModule
from src.modules.structure.heading_recognition import (
    HeadingRecognitionModule,
    rebuild_document_index,
)
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.heading_numbering_ooxml import effective_numbering
from src.shared.engine.page_number_planner import build_page_number_execution_plan


def test_special_title_configuration_reaches_all_document_consumers():
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    document.add_paragraph("正文")
    special_title = document.add_paragraph("摘要")
    document.add_paragraph("摘要正文")
    document.add_heading("第二章 方法", level=1)

    config = ResolvedConfig()
    config.heading_model.non_numbered_title_texts = ["摘要"]
    config.heading_model.non_numbered_prefixes = []
    config.heading_model.non_numbered_heading_style_mode = "custom"
    config.heading_numbering.level_bindings = {
        "heading1": HeadingLevelBindingConfig(enabled=True)
    }
    config.styles = {
        "normal": StyleConfig(),
        "body": StyleConfig(),
        "heading1": StyleConfig(size_pt=16, bold=True),
        "non_numbered_heading": StyleConfig(size_pt=15, bold=True),
    }

    selector = special_title_selector("exact", "摘要")
    config.header_footer.header.mode = "fixed"
    config.header_footer.header.fixed_text = "HEADER"
    config.header_footer.header.hidden_selectors = [selector]
    config.header_footer.footer.content_mode = "fixed"
    config.header_footer.footer.fixed_text = "FOOTER"
    config.header_footer.footer.hidden_selectors = ["body"]
    config.header_footer.page_number_plan.phases = [
        PageNumberPhaseConfig(
            phase_id="body",
            selectors=["body"],
            visible=True,
        ),
        PageNumberPhaseConfig(
            phase_id="summary",
            selectors=[selector],
            visible=True,
            number_format="upperRoman",
            start_mode="restart",
        ),
    ]

    context = PipelineContext()
    tracker = ChangeTracker()
    HeadingRecognitionModule().apply(document, config, tracker, context)

    assert context.doc_tree.get_special_title_match(2) == selector

    HeadingNumberingModule().apply(document, config, tracker, context)
    assert effective_numbering(special_title) is None

    SectionFormatModule().apply(document, config, tracker, context)
    rebuild_document_index(document, context, config)
    ParagraphStyleModule().apply(document, config, tracker, context)

    assert (
        special_title.style.name
        == config.heading_model.non_numbered_heading_style_name
    )

    plan = build_page_number_execution_plan(
        document,
        context,
        config.header_footer,
    )

    assert [section.start_index for section in plan.sections] == [0, 2, 4]
    assert plan.sections[1].phase_id == "summary"
    assert plan.sections[1].number_format == "upperRoman"
    assert plan.sections[1].header_visible is False
    assert plan.sections[1].footer_text_visible is True

    HeaderFooterModule().apply(document, config, tracker, context)

    assert document.sections[1].header.paragraphs[0].text == ""
    assert "FOOTER" in document.sections[1].footer.paragraphs[0].text
