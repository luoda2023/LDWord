import json
import os
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

from docx import Document
import pytest


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolver import resolve_config
from src.config.resolved import ResolvedConfig
from src.config.feature_configs import OutputConfig
from src.config.scene import ContentVisibilityRule, DeliveryPreset, SceneWorkspace
from src.config.scene_family_application import apply_planned_scene_family_defaults
from src.config.template import StyleConfig, TemplateConfig
from src.modules.base import BaseModule, ModuleMeta
from src.modules.basic.header_footer import HeaderFooterModule
from src.modules.basic.section_format import SectionFormatModule
from src.modules.structure.heading_recognition import HeadingRecognitionModule
from src.pipeline.result import PipelineResult
from src.pipeline.tracker import ChangeTracker
from src.pipeline.runner import Pipeline
from src.report_writer import write_json_report, write_markdown_report
from src.shared.engine.field_builder import iter_field_instructions
import src.services.production_runtime.execution_runtime as execution_runtime
from src.services.production_runtime import delivery_reporting, delivery_runtime
from src.services.production_runtime.delivery_group_preflight import (
    PreparedDeliveryTargetGroup,
    build_delivery_group_output_preflight,
)
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner


def test_unresolved_runtime_config_authorizes_no_artifacts():
    config = ResolvedConfig()
    output = config.output

    assert config.default_delivery_preset_id == ""
    assert config.delivery_presets == []

    assert not any(
        (
            output.final_docx,
            output.compare_docx,
            output.compare_text,
            output.compare_formatting,
            output.report_json,
            output.report_markdown,
            output.material_manifest,
            output.material_package,
            output.review_pdf,
        )
    )


def test_explicit_zero_delivery_scene_resolves_without_artifact_authorization():
    scene = SceneWorkspace(
        default_delivery_preset_id="",
        delivery_presets=[],
    )

    resolved = resolve_config(TemplateConfig(), scene)

    assert resolved.default_delivery_preset_id == ""
    assert resolved.delivery_presets == []
    assert not any(vars(resolved.output).values())


def test_pipeline_skips_final_output_when_final_docx_is_disabled(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)

    template = TemplateConfig()
    scene = SceneWorkspace()
    scene.default_delivery_preset().artifacts.final_docx = False
    config = resolve_config(template, scene)

    output_dir = tmp_path / "out"
    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        output_suffix="_formatted",
    ).execute(str(source))

    assert result.success is True
    assert result.output_paths == {}
    assert (output_dir / "source_formatted.docx").exists() is False


def test_pipeline_writes_docx_for_each_delivery_preset_with_final_artifact(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="final",
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                label="Final",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=True, report_markdown=False),
            ),
            DeliveryPreset(
                preset_id="review",
                label="Review",
                output_dir_template="{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=True),
            ),
            DeliveryPreset(
                preset_id="compliance_report",
                label="Compliance",
                output_dir_template="reports/{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=False, report_json=True, report_markdown=True),
            ),
        ],
    )
    config = resolve_config(TemplateConfig(), scene)

    output_dir = tmp_path / "out"
    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        output_suffix="_formatted",
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is True
    assert result.output_paths == {
        "final": str(output_dir / "source_final.docx"),
        "review": str(output_dir / "review" / "source_review.docx"),
    }
    assert (output_dir / "source_final.docx").exists() is True
    assert (output_dir / "review" / "source_review.docx").exists() is True
    assert (output_dir / "source_compliance_report.docx").exists() is False


@pytest.mark.parametrize(
    ("corruption", "expected_issue"),
    (
        ("duplicate", "duplicate_delivery_preset_id:final"),
        ("empty_preset", "delivery_preset_id_empty:0"),
        ("empty_default", "default_delivery_preset_id_empty"),
        ("unknown_default", "default_delivery_preset_id_unknown:missing"),
    ),
)
def test_invalid_delivery_identity_blocks_before_mutation_or_output(
    tmp_path,
    corruption,
    expected_issue,
):
    source = tmp_path / "source.docx"
    source_doc = Document()
    source_doc.add_paragraph("SOURCE MUST REMAIN UNCHANGED")
    source_doc.save(source)
    source_bytes = source.read_bytes()
    output_dir = tmp_path / "out"
    applied: list[str] = []

    class _MutationProbeModule(BaseModule):
        meta = ModuleMeta(
            name="delivery_identity_mutation_probe",
            description="must not run for invalid delivery identity",
            category="test",
            execution_phase="format",
            scope_behavior="document_level",
        )

        def apply(self, doc, config, tracker, context):
            applied.append("applied")
            doc.add_paragraph("MUTATED")

    scene = SceneWorkspace(
        default_delivery_preset_id="final",
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                filename_template="{stem}_first",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
            ),
            DeliveryPreset(
                preset_id="review",
                filename_template="{stem}_second",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
            ),
        ],
    )
    config = resolve_config(TemplateConfig(), scene)
    if corruption == "duplicate":
        config.delivery_presets[1].preset_id = "final"
    elif corruption == "empty_preset":
        config.delivery_presets[0].preset_id = ""
    elif corruption == "empty_default":
        config.default_delivery_preset_id = ""
    else:
        config.default_delivery_preset_id = "missing"

    result = Pipeline(
        modules=[_MutationProbeModule()],
        config=config,
        output_dir=str(output_dir),
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is False
    assert result.status == "failed"
    assert expected_issue in str(result.error)
    assert applied == []
    assert result.output_paths == {}
    assert not output_dir.exists()
    assert source.read_bytes() == source_bytes


def test_publication_defense_rejects_duplicate_output_id_before_staging(tmp_path):
    config = resolve_config(TemplateConfig(), SceneWorkspace())
    pipeline = Pipeline(modules=[], config=config)
    document = Document()
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"

    with pytest.raises(ValueError, match="duplicate_output_id:same"):
        pipeline._publish_output_documents(
            [
                ("same", first, document),
                ("same", second, document),
            ]
        )

    assert not first.exists()
    assert not second.exists()


def test_pipeline_renders_preset_label_with_shared_display_name(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="student_version",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student_version",
                label="Student version",
                filename_template="{stem}_{preset_label}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            )
        ],
    )
    config = resolve_config(TemplateConfig(), scene)

    output_dir = tmp_path / "out"
    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        output_suffix="_formatted",
        force_delivery_presets=True,
    ).execute(str(source))

    expected = output_dir / "source_学生版.docx"
    assert result.success is True
    assert result.output_paths == {"student_version": str(expected)}
    assert expected.exists() is True


def test_pipeline_records_journal_submission_package_evidence_and_reports(tmp_path):
    source = tmp_path / "journal.docx"
    doc = Document()
    doc.add_paragraph(r"This manuscript cites \cite{smith2020}.")
    doc.save(source)

    scene = SceneWorkspace(scene_id="journal_en", category="journal_en")
    apply_planned_scene_family_defaults(scene)
    config = resolve_config(
        TemplateConfig(),
        scene,
        entity_data={
            "article_title": "Journal article",
            "author": "Jane Smith",
            "affiliation": "Example University",
            "corresponding_author": "Jane Smith",
            "bibtex": """
@article{smith2020,
  title={A study},
  author={Smith, Jane},
  year={2020}
}
""",
        },
    )

    output_dir = tmp_path / "out"
    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        output_suffix="_formatted",
    ).execute(str(source))

    assert result.success is True
    assert result.output_paths.keys() >= {
        "submission_manuscript",
        "review_copy",
        "cover_letter",
    }
    evidence = result.context.journal_submission_package
    assert evidence.status == "ok"
    assert evidence.summary.required_component_count == 3
    assert evidence.summary.satisfied_required_count == 3
    assert evidence.summary.output_count == 3
    assert evidence.summary.material_artifact_enabled_count == 1
    tracker_items = result.tracker.get_by_module("journal_submission_package")
    assert tracker_items
    assert tracker_items[0].change_type == "submission_package"
    assert "required=3/3" in tracker_items[0].after

    report_json = output_dir / "journal_submission_manuscript_changes.json"
    report_md = output_dir / "journal_submission_manuscript_changes.md"
    write_json_report(
        result,
        input_path=source,
        output_path=Path(result.output_paths["submission_manuscript"]),
        report_path=report_json,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )
    write_markdown_report(
        result,
        input_path=source,
        report_path=report_md,
        elapsed=0.1,
        modules_enabled=0,
        modules_total=0,
    )

    report_data = json.loads(report_json.read_text(encoding="utf-8"))
    markdown = report_md.read_text(encoding="utf-8")

    assert report_data["journal_submission_package"]["status"] == "ok"
    assert report_data["journal_submission_package"]["summary"]["satisfied_required_count"] == 3
    assert "## 英文期刊投稿包证据" in markdown
    assert "submission_manuscript" in markdown
    assert "declaration_package" in markdown


def test_pipeline_can_force_delivery_output_for_single_preset(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                output_dir_template="{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=True),
            ),
        ],
    )
    config = resolve_config(TemplateConfig(), scene)

    output_dir = tmp_path / "out"
    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        output_suffix="_formatted",
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is True
    assert result.output_paths == {
        "review": str(output_dir / "review" / "source_review.docx"),
    }
    assert (output_dir / "review" / "source_review.docx").exists() is True
    assert (output_dir / "source_formatted.docx").exists() is False


def test_pipeline_reports_delivery_output_target_preflight_issues(tmp_path):
    source = tmp_path / "source.docx"
    source_doc = Document()
    source_doc.add_paragraph("SOURCE")
    source_doc.save(source)
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    existing_target = output_dir / "source.docx"
    existing_doc = Document()
    existing_doc.add_paragraph("OLD TARGET")
    existing_doc.save(existing_target)

    scene = SceneWorkspace(
        default_delivery_preset_id="final",
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                label="Final",
                filename_template="{stem}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            ),
            DeliveryPreset(
                preset_id="review",
                label="Review",
                filename_template="{stem}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    config = resolve_config(TemplateConfig(), scene)

    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        output_suffix="_formatted",
    ).execute(str(source))

    preflight = result.context.output_target_preflight
    issue_kinds = [
        issue.kind
        for item in preflight.items
        for issue in item.issues
    ]
    tracker_items = [
        item for item in result.tracker.get_all()
        if item.rule_name == "output_target_preflight"
    ]

    assert result.success is False
    assert result.output_paths == {}
    assert [paragraph.text for paragraph in Document(existing_target).paragraphs] == [
        "OLD TARGET"
    ]
    assert [item.path for item in preflight.items] == [
        str(existing_target),
        str(existing_target),
    ]
    assert preflight.has_issues is True
    assert preflight.has_errors is True
    assert preflight.error_count == 2
    assert issue_kinds.count("duplicate_target") == 2
    assert issue_kinds.count("target_exists") == 2
    assert {
        issue.severity
        for item in preflight.items
        for issue in item.issues
        if issue.kind == "duplicate_target"
    } == {"error"}
    assert {
        issue.severity
        for item in preflight.items
        for issue in item.issues
        if issue.kind == "target_exists"
    } == {"warning"}
    issue_messages = [
        issue.message
        for item in preflight.items
        for issue in item.issues
    ]
    assert any(message.startswith("Final ") for message in issue_messages)
    assert any(message.startswith("Review ") for message in issue_messages)
    assert not any(message.startswith("final ") for message in issue_messages)
    assert not any(message.startswith("review ") for message in issue_messages)
    assert tracker_items
    assert {item.change_type for item in tracker_items} == {
        "output_target_warning",
        "output_target_blocked",
    }
    blocked = next(
        item for item in tracker_items if item.change_type == "output_target_blocked"
    )
    assert blocked.success is False
    assert blocked.failure_reason == result.error


def test_pipeline_output_target_preflight_uses_shared_delivery_display_label(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    existing_target = output_dir / "source.docx"
    Document().save(existing_target)

    scene = SceneWorkspace(
        default_delivery_preset_id="answer_key",
        delivery_presets=[
            DeliveryPreset(
                preset_id="answer_key",
                label="Answer key",
                filename_template="{stem}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    config = resolve_config(TemplateConfig(), scene)

    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        output_suffix="_formatted",
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is True
    preflight_item = result.context.output_target_preflight.items[0]
    assert preflight_item.label == "答案速查"
    assert preflight_item.issues[0].message.startswith("答案速查 ")
    assert "输出文件已存在" in preflight_item.issues[0].message
    assert "answer_key" not in preflight_item.issues[0].message


def test_pipeline_applies_delivery_visibility_marker_blocks_per_preset(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("1. What is 1 + 1?")
    doc.add_paragraph("{{#visibility:answer}}")
    doc.add_paragraph("Answer: 2")
    doc.add_paragraph("{{/visibility:answer}}")
    doc.add_paragraph("End")
    doc.save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="student",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                label="Student",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
                content_visibility_rules=[
                    ContentVisibilityRule(
                        rule_id="hide_answers",
                        label="Hide answers",
                        selector="answer",
                        action="remove",
                    )
                ],
            ),
            DeliveryPreset(
                preset_id="teacher",
                label="Teacher",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    config = resolve_config(TemplateConfig(), scene)

    output_dir = tmp_path / "out"
    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(output_dir),
        output_suffix="_formatted",
    ).execute(str(source))

    student_doc = Document(result.output_paths["student"])
    teacher_doc = Document(result.output_paths["teacher"])
    student_text = "\n".join(paragraph.text for paragraph in student_doc.paragraphs)
    teacher_text = "\n".join(paragraph.text for paragraph in teacher_doc.paragraphs)

    assert result.success is True
    assert "Answer: 2" not in student_text
    assert "{{#visibility:answer}}" not in student_text
    assert "{{/visibility:answer}}" not in student_text
    assert "Answer: 2" in teacher_text
    assert "{{#visibility:answer}}" not in teacher_text
    assert "{{/visibility:answer}}" not in teacher_text
    previews = result.context.content_visibility_preview
    assert [(item.preset_id, item.removed_paragraph_count) for item in previews] == [
        ("student", 3),
        ("teacher", 0),
    ]
    assert previews[0].removed_blocks[0].selector == "answer"
    assert previews[0].removed_blocks[0].context_before == "1. What is 1 + 1?"
    assert previews[0].removed_blocks[0].context_after == "End"
    visibility_receipts = result.context.content_visibility_receipts
    assert set(visibility_receipts) == {"student", "teacher"}
    assert visibility_receipts["student"]["removed_body_element_count"] == 3
    assert visibility_receipts["teacher"]["removed_body_element_count"] == 2
    assert visibility_receipts["student"]["receipt_id"]
    visibility_changes = [
        item
        for item in result.tracker.get_all()
        if item.rule_name == "delivery_visibility"
    ]
    preflight_changes = [
        item
        for item in result.tracker.get_all()
        if item.rule_name == "delivery_visibility_preflight"
    ]
    assert visibility_changes
    assert "answer" in visibility_changes[0].after
    assert (
        "student: remove body=3, paragraphs=3, tables=0 "
        "in 1 block(s) sample=Answer: 2"
    ) in preflight_changes[0].after


def test_pipeline_reports_visibility_preflight_without_blocking_output(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Question only")
    doc.save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="student",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove")
                ],
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    config = resolve_config(TemplateConfig(), scene)

    result = Pipeline(
        modules=[],
        config=config,
        output_dir=str(tmp_path / "out"),
        output_suffix="_formatted",
        force_delivery_presets=True,
    ).execute(str(source))

    preflight_changes = [
        item
        for item in result.tracker.get_all()
        if item.rule_name == "delivery_visibility_preflight"
    ]

    assert result.success is True
    assert result.output_paths["student"].endswith("source_student.docx")
    assert preflight_changes
    assert preflight_changes[0].change_type == "visibility_preflight_warning"
    assert "规则 selector 未在文档中找到: answer" in preflight_changes[0].after
    assert "student: remove 0" not in preflight_changes[0].after
    assert result.context.content_visibility_scan.missing_rule_selectors == ["answer"]
    assert result.context.content_visibility_preview[0].missing_selectors == ["answer"]


def test_pipeline_blocks_invalid_visibility_action_before_docx_mutation(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Question")
    doc.add_paragraph("{{#visibility:answer}}")
    doc.add_paragraph("SENSITIVE ANSWER")
    doc.add_paragraph("{{/visibility:answer}}")
    doc.save(source)
    source_before = source.read_bytes()
    applied: list[str] = []

    class _MutationProbeModule(BaseModule):
        meta = ModuleMeta(
            name="invalid_visibility_action_probe",
            description="Invalid visibility action probe",
            category="test",
            execution_phase="fill",
            scope_behavior="document_level",
        )

        def apply(self, doc, config, tracker, context):
            applied.append("applied")
            doc.add_paragraph("MUTATED")

    scene = SceneWorkspace(
        default_delivery_preset_id="student",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
                content_visibility_rules=[
                    ContentVisibilityRule(
                        rule_id="hide-answer",
                        selector="answer",
                        action="remvoe",
                    )
                ],
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    output_dir = tmp_path / "out"

    result = Pipeline(
        modules=[_MutationProbeModule()],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(output_dir),
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is False
    assert result.status == "failed"
    assert applied == []
    assert result.output_paths == {}
    assert not output_dir.exists()
    assert source.read_bytes() == source_before
    diagnostics = result.context.content_visibility_scan.blocking_diagnostics
    assert {item["code"] for item in diagnostics} == {
        "content_visibility_action_invalid"
    }
    assert "SENSITIVE ANSWER" in "\n".join(
        paragraph.text for paragraph in Document(source).paragraphs
    )


def test_pipeline_blocks_sensitive_visibility_selector_typo_before_marker_strip(
    tmp_path,
):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Question")
    doc.add_paragraph("{{#visibility:answer}}")
    doc.add_paragraph("SENSITIVE ANSWER")
    doc.add_paragraph("{{/visibility:answer}}")
    doc.save(source)
    source_before = source.read_bytes()
    scene = SceneWorkspace(
        default_delivery_preset_id="student",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
                content_visibility_rules=[
                    ContentVisibilityRule(
                        rule_id="hide-answer",
                        selector="anwser",
                        action="remove",
                    )
                ],
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    output_dir = tmp_path / "out"

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(output_dir),
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is False
    assert result.status == "failed"
    assert result.output_paths == {}
    assert not output_dir.exists()
    assert source.read_bytes() == source_before
    scan = result.context.content_visibility_scan
    assert scan.missing_rule_selectors == ["anwser"]
    assert scan.unused_document_selectors == ["answer"]
    assert {item["code"] for item in scan.blocking_diagnostics} == {
        "content_visibility_sensitive_selector_mismatch"
    }
    assert "Content visibility preflight blocked execution" in result.error


def test_pipeline_keeps_unrelated_optional_visibility_rule_nonblocking(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("{{#visibility:appendix}}")
    doc.add_paragraph("Optional appendix")
    doc.add_paragraph("{{/visibility:appendix}}")
    doc.save(source)
    scene = SceneWorkspace(
        default_delivery_preset_id="student",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove")
                ],
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False

    result = Pipeline(
        modules=[],
        config=resolve_config(TemplateConfig(), scene),
        output_dir=str(tmp_path / "out"),
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is True
    output_text = "\n".join(
        paragraph.text
        for paragraph in Document(result.output_paths["student"]).paragraphs
    )
    assert "Optional appendix" in output_text
    assert "visibility:appendix" not in output_text
    assert result.context.content_visibility_scan.has_blocking_issues is False


def test_pipeline_blocks_invalid_visibility_grammar_before_module_apply(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{#visibility:answer}}"
    doc.add_paragraph("{{/visibility:answer}}")
    doc.save(source)
    applied: list[str] = []

    class _MutationProbeModule(BaseModule):
        meta = ModuleMeta(
            name="visibility_mutation_probe",
            description="Visibility mutation probe",
            category="test",
            execution_phase="fill",
            scope_behavior="document_level",
        )

        def apply(self, doc, config, tracker, context):
            applied.append("applied")
            doc.add_paragraph("MUTATED")

    scene = SceneWorkspace(
        default_delivery_preset_id="student",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=False,
                    report_markdown=False,
                ),
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove")
                ],
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False
    config = resolve_config(TemplateConfig(), scene)

    result = Pipeline(
        modules=[_MutationProbeModule()],
        config=config,
        output_dir=str(tmp_path / "out"),
        force_delivery_presets=True,
    ).execute(str(source))

    assert result.success is False
    assert applied == []
    assert result.output_paths == {}
    assert result.context.content_visibility_preview == []
    assert result.context.content_visibility_scan.has_blocking_issues is True
    assert "Content visibility preflight blocked execution" in result.error
    failures = result.tracker.get_failures()
    assert len(failures) == 1
    assert failures[0].change_type == "visibility_preflight_blocked"


def test_workbench_runner_reports_delivery_visibility_changes(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Question")
    doc.add_paragraph("{{#visibility:answer}}")
    doc.add_paragraph("Answer")
    doc.add_paragraph("{{/visibility:answer}}")
    doc.save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="student",
        delivery_presets=[
            DeliveryPreset(
                preset_id="student",
                label="Student",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=True, report_markdown=False),
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove")
                ],
            ),
            DeliveryPreset(
                preset_id="teacher",
                label="Teacher",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    output_dir = source.parent / "output"
    student_doc = Document(output_dir / "source_student.docx")
    student_text = "\n".join(paragraph.text for paragraph in student_doc.paragraphs)
    report_data = json.loads((output_dir / "source_student_changes.json").read_text(encoding="utf-8"))

    assert payload["status"] == "success"
    assert payload["content_visibility_scan"]["selectors"] == ["answer"]
    assert payload["content_visibility_scan"]["used_rule_selectors"] == ["answer"]
    assert payload["content_visibility_scan"]["has_issues"] is False
    assert payload["content_visibility_preview"][0]["preset_id"] == "student"
    assert payload["content_visibility_preview"][0]["removed_paragraph_count"] == 3
    assert payload["content_visibility_preview"][0]["removed_text_samples"] == ["Answer"]
    assert payload["content_visibility_preview"][0]["removed_blocks"] == [
        {
            "selector": "answer",
            "start_paragraph_index": 1,
            "end_paragraph_index": 3,
            "paragraph_count": 3,
            "text_samples": ["Answer"],
            "context_before": "Question",
            "context_after": "",
            "start_body_index": 1,
            "end_body_index": 3,
            "body_element_count": 3,
            "table_count": 0,
            "content_image_markers": [],
        }
    ]
    assert "Answer" not in student_text
    assert any(
        item["rule_name"] == "delivery_visibility_preflight"
        for item in report_data["changes"]
    )
    assert any(item["rule_name"] == "delivery_visibility" for item in report_data["changes"])


def test_official_markdown_input_does_not_use_exam_paper_residual():
    official_scene = SceneWorkspace(
        scene_id="official",
        mode_id="official",
        category="government",
        master_id="official_default",
    )
    official_user_scene = SceneWorkspace(
        scene_id="official_custom",
        mode_id="official",
        category="government",
        master_id="official_default",
    )
    exam_scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        category="exam_paper",
        master_id="default_exam",
    )

    assert official_scene.exam_paper is not None
    assert official_user_scene.exam_paper is not None
    assert execution_runtime._is_exam_markdown_input(
        official_scene,
        Path("notice.md"),
    ) is False
    assert execution_runtime._is_exam_markdown_input(
        official_user_scene,
        Path("custom-official.md"),
    ) is False
    assert execution_runtime._is_exam_markdown_input(
        exam_scene,
        Path("exam_source.md"),
    ) is True
    assert execution_runtime._is_exam_markdown_input(
        exam_scene,
        Path("exam_source.docx"),
    ) is False


def test_workbench_runner_payload_includes_output_target_preflight(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = source.parent / "output"
    output_dir.mkdir()
    existing_target = output_dir / "source.docx"
    Document().save(existing_target)

    scene = SceneWorkspace(
        default_delivery_preset_id="final",
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                filename_template="{stem}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            ),
            DeliveryPreset(
                preset_id="review",
                filename_template="{stem}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    preflight = payload["output_target_preflight"]
    issue_kinds = [
        issue["kind"]
        for item in preflight["items"]
        for issue in item["issues"]
    ]

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert preflight["has_issues"] is True
    assert preflight["has_errors"] is True
    assert preflight["issue_count"] == 4
    assert preflight["error_count"] == 2
    assert issue_kinds.count("duplicate_target") == 2
    assert issue_kinds.count("target_exists") == 2
    assert preflight["items"][0]["path"] == str(existing_target)


def test_workbench_runner_does_not_attach_engineering_sample_manifest(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)

    template = TemplateConfig()
    scene = SceneWorkspace()
    scene.default_delivery_preset().artifacts.report_json = False
    scene.default_delivery_preset().artifacts.report_markdown = False

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=template,
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert "scene_sample_manifest_paths" not in payload


def test_workbench_runner_can_emit_json_report_without_final_docx(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)

    template = TemplateConfig()
    scene = SceneWorkspace()
    artifacts = scene.default_delivery_preset().artifacts
    artifacts.final_docx = False
    artifacts.report_json = True
    artifacts.report_markdown = False

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=template,
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    report_json = source.parent / "output" / "source_changes.json"
    report_md = source.parent / "output" / "source_changes.md"

    assert payload["status"] == "success"
    assert payload["output_path"] == ""
    assert payload["report_paths"] == [str(report_json)]
    assert report_json.exists() is True
    assert report_md.exists() is False
    assert '"output": ""' in report_json.read_text(encoding="utf-8")


def test_workbench_runner_resolves_target_template_for_delivery_presets(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)

    base_template = TemplateConfig()
    base_template.styles["body"] = StyleConfig(font_cn="BaseFont")

    target_template = TemplateConfig()
    target_template.styles["body"] = StyleConfig(font_cn="TargetFont")

    scene = SceneWorkspace(
        scene_id="exam",
        mode_id="exam",
        template_id="base_template",
        compatible_template_ids=["base_template", "target_template"],
        master_id="default_exam",
        default_delivery_preset_id="submission",
        delivery_presets=[
            DeliveryPreset(
                preset_id="submission",
                label="Submission",
                target_template_id="target_template",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=True, report_markdown=False),
            ),
            DeliveryPreset(
                preset_id="review",
                label="Review",
                target_template_id="base_template",
                output_dir_template="{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=True),
            ),
        ],
    )
    loaded_template_requests = []
    captured_runs = []

    class _StubPipeline:
        def __init__(self, **kwargs):
            captured_runs.append(kwargs)

        def execute(self, _doc_path: str) -> PipelineResult:
            config = captured_runs[-1]["config"]
            output_dir = Path(captured_runs[-1]["output_dir"])
            output_paths = {}
            for preset in config.delivery_presets:
                preset_id = str(preset.preset_id)
                output_paths[preset_id] = str(output_dir / f"{preset_id}.docx")
            return PipelineResult(
                success=True,
                status="success",
                output_paths=output_paths,
            )

    def _load_template(
        template_id: str,
        *,
        mode_id: str | None = None,
    ) -> TemplateConfig:
        loaded_template_requests.append((template_id, mode_id))
        assert template_id == "target_template"
        assert mode_id == "exam"
        return target_template

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])
    monkeypatch.setattr(execution_runtime, "Pipeline", _StubPipeline)
    monkeypatch.setattr(delivery_runtime, "load_template_from_library", _load_template)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=base_template,
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    fonts_by_preset = {
        run["config"].delivery_presets[0].preset_id: run["config"].styles["body"].font_cn
        for run in captured_runs
    }

    assert payload["status"] == "success"
    assert payload["output_path"] == str(source.parent / "output" / "submission.docx")
    assert payload["output_paths"] == {
        "submission": str(source.parent / "output" / "submission.docx"),
        "review": str(source.parent / "output" / "review.docx"),
    }
    assert loaded_template_requests == [("target_template", "exam")]
    assert fonts_by_preset == {
        "submission": "TargetFont",
        "review": "BaseFont",
    }
    assert all(run["force_delivery_presets"] is True for run in captured_runs)
    assert sorted(Path(path).name for path in payload["report_paths"]) == [
        "source_review_changes.md",
        "source_submission_changes.json",
    ]


def test_delivery_target_groups_preflight_cross_group_hardlink_alias_before_execution(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    first_target = output_dir / "shared-a.docx"
    second_target = output_dir / "shared-b.docx"
    first_target.write_bytes(b"existing-final")
    os.link(first_target, second_target)

    scene = SceneWorkspace(
        template_id="base_template",
        compatible_template_ids=["base_template", "target_template"],
        default_delivery_preset_id="submission",
        delivery_presets=[
            DeliveryPreset(
                preset_id="submission",
                target_template_id="target_template",
                filename_template=first_target.name,
                artifacts=OutputConfig(final_docx=True),
            ),
            DeliveryPreset(
                preset_id="review",
                target_template_id="base_template",
                filename_template=second_target.name,
                artifacts=OutputConfig(final_docx=True),
            ),
        ],
    )
    pipeline_calls: list[dict[str, object]] = []

    class _ForbiddenPipeline:
        def __init__(self, **kwargs):
            pipeline_calls.append(kwargs)
            raise AssertionError("no delivery group may start before global preflight")

    monkeypatch.setattr(execution_runtime, "Pipeline", _ForbiddenPipeline)
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert pipeline_calls == []
    assert first_target.read_bytes() == b"existing-final"
    assert second_target.read_bytes() == b"existing-final"
    issue_kinds = {
        issue["kind"]
        for item in payload["output_target_preflight"]["items"]
        for issue in item["issues"]
    }
    assert "duplicate_target" in issue_kinds


def test_delivery_group_preflight_rejects_each_same_group_auxiliary_final_collision(
    tmp_path,
):
    shared_path = str(tmp_path / "shared.docx")
    artifact_kinds = (
        "report_json",
        "report_markdown",
        "compare_docx",
        "structured_intermediate",
    )

    for artifact_kind in artifact_kinds:
        payload = build_delivery_group_output_preflight(
            [
                PreparedDeliveryTargetGroup(
                    target_template_id="target_template",
                    config=object(),
                    planned_output_paths={"final": shared_path},
                    terminal_owner="",
                    error="",
                    planned_artifact_paths={
                        f"final:{artifact_kind}": shared_path,
                    },
                )
            ]
        )

        assert payload["has_errors"] is True
        issue_kinds = {
            issue["kind"]
            for item in payload["items"]
            for issue in item["issues"]
        }
        assert "duplicate_artifact_target" in issue_kinds


def test_delivery_target_group_preflight_blocks_same_group_hardlink_artifact(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "out"
    output_dir.mkdir()
    final_path = output_dir / "shared.docx"
    report_path = output_dir / "shared_changes.json"
    final_path.write_bytes(b"existing-shared-artifact")
    os.link(final_path, report_path)
    scene = SceneWorkspace(
        template_id="base_template",
        compatible_template_ids=["base_template", "target_template"],
        default_delivery_preset_id="final",
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                target_template_id="target_template",
                filename_template="shared.docx",
                artifacts=OutputConfig(final_docx=True, report_json=True),
            )
        ],
    )
    pipeline_calls: list[dict[str, object]] = []

    class _ForbiddenPipeline:
        def __init__(self, **kwargs):
            pipeline_calls.append(kwargs)
            raise AssertionError("same-group artifact collision must fail preflight")

    monkeypatch.setattr(execution_runtime, "Pipeline", _ForbiddenPipeline)
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert pipeline_calls == []
    assert final_path.read_bytes() == b"existing-shared-artifact"
    assert report_path.read_bytes() == b"existing-shared-artifact"
    issue_kinds = {
        issue["kind"]
        for item in payload["output_target_preflight"]["items"]
        for issue in item["issues"]
    }
    assert "duplicate_artifact_target" in issue_kinds


def test_delivery_target_groups_conflict_does_not_create_output_root(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "not-created"
    scene = SceneWorkspace(
        template_id="base_template",
        default_delivery_preset_id="first",
        delivery_presets=[
            DeliveryPreset(
                preset_id="first",
                target_template_id="target_template",
                filename_template="same-final.docx",
                artifacts=OutputConfig(final_docx=True),
            ),
            DeliveryPreset(
                preset_id="second",
                target_template_id="base_template",
                filename_template="same-final.docx",
                artifacts=OutputConfig(final_docx=True),
            ),
        ],
    )

    class _ForbiddenPipeline:
        def __init__(self, **_kwargs):
            raise AssertionError("global preflight must run before any group pipeline")

    monkeypatch.setattr(execution_runtime, "Pipeline", _ForbiddenPipeline)
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert "同一 final" in payload["error_text"]
    assert not output_dir.exists()


def test_delivery_target_groups_preflight_auxiliary_collision_without_final_docx(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "not-created"
    scene = SceneWorkspace(
        template_id="base_template",
        default_delivery_preset_id="first",
        delivery_presets=[
            DeliveryPreset(
                preset_id="first",
                target_template_id="target_template",
                filename_template="same-report.docx",
                artifacts=OutputConfig(final_docx=False, report_json=True),
            ),
            DeliveryPreset(
                preset_id="second",
                target_template_id="base_template",
                filename_template="same-report.docx",
                artifacts=OutputConfig(final_docx=False, report_json=True),
            ),
        ],
    )
    pipeline_calls: list[dict[str, object]] = []

    class _ForbiddenPipeline:
        def __init__(self, **kwargs):
            pipeline_calls.append(kwargs)
            raise AssertionError("artifact collision must fail before execution")

    monkeypatch.setattr(execution_runtime, "Pipeline", _ForbiddenPipeline)
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(),
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert pipeline_calls == []
    assert not output_dir.exists()
    issue_kinds = {
        issue["kind"]
        for item in payload["output_target_preflight"]["items"]
        for issue in item["issues"]
    }
    assert "duplicate_artifact_target" in issue_kinds


def test_delivery_template_load_prefers_explicit_scene_mode(monkeypatch):
    current_template = TemplateConfig(name="Current")
    loaded_template = TemplateConfig(name="Loaded")
    calls = []

    def _load_template(
        template_id: str,
        *,
        mode_id: str | None = None,
    ) -> TemplateConfig:
        calls.append((template_id, mode_id))
        return loaded_template

    monkeypatch.setattr(delivery_runtime, "load_template_from_library", _load_template)
    scene = SceneWorkspace(
        scene_id="school_midterm_copy",
        mode_id="exam",
        template_id="base_template",
    )

    result = delivery_runtime.load_delivery_template(
        current_template,
        scene=scene,
        target_template_id="target_template",
    )

    assert result is loaded_template
    assert calls == [("target_template", "exam")]


def test_workbench_runner_uses_delivery_preset_artifacts_for_outputs_and_reports(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                label="Final",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=True, report_markdown=False),
            ),
            DeliveryPreset(
                preset_id="review",
                label="Review",
                output_dir_template="{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=True),
            ),
            DeliveryPreset(
                preset_id="compliance_report",
                label="Compliance",
                output_dir_template="reports/{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=False, report_json=True, report_markdown=True),
            ),
        ],
    )

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    output_dir = source.parent / "output"
    assert payload["status"] == "success"
    assert payload["output_path"] == str(output_dir / "review" / "source_review.docx")
    assert payload["output_paths"] == {
        "final": str(output_dir / "source_final.docx"),
        "review": str(output_dir / "review" / "source_review.docx"),
    }
    assert sorted(Path(path).name for path in payload["report_paths"]) == [
        "source_compliance_report_changes.json",
        "source_compliance_report_changes.md",
        "source_final_changes.json",
        "source_review_changes.md",
    ]
    assert (output_dir / "source_final.docx").exists() is True
    assert (output_dir / "review" / "source_review.docx").exists() is True
    assert (output_dir / "source_final_changes.json").exists() is True
    assert (output_dir / "review" / "source_review_changes.md").exists() is True
    assert (output_dir / "reports" / "compliance_report" / "source_compliance_report_changes.json").exists() is True
    assert (output_dir / "reports" / "compliance_report" / "source_compliance_report_changes.md").exists() is True
    assert '"output": "' in (output_dir / "source_final_changes.json").read_text(encoding="utf-8")
    assert '"output": ""' in (
        output_dir / "reports" / "compliance_report" / "source_compliance_report_changes.json"
    ).read_text(encoding="utf-8")


def test_workbench_runner_delivery_markdown_report_uses_display_label(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="answer_key",
        delivery_presets=[
            DeliveryPreset(
                preset_id="answer_key",
                label="Answer key",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(
                    final_docx=True,
                    report_json=True,
                    report_markdown=True,
                ),
            )
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    output_dir = source.parent / "output"
    report_md = output_dir / "source_answer_key_changes.md"
    report_json = output_dir / "source_answer_key_changes.json"
    markdown = report_md.read_text(encoding="utf-8")
    report_data = json.loads(report_json.read_text(encoding="utf-8"))

    assert payload["status"] == "success"
    assert report_md.exists() is True
    assert "## 交付" in markdown
    assert "- 版本: **答案速查**" in markdown
    assert f"- 输出文件: {output_dir / 'source_answer_key.docx'}" in markdown
    assert "- 版本 ID: `answer_key`" in markdown
    assert "- 原始标签: Answer key" in markdown
    assert "- 报告级别: summary" in markdown
    assert "- 版本: **answer_key**" not in markdown
    assert report_data["delivery_preset"] == {
        "preset_id": "answer_key",
        "label": "Answer key",
        "display_label": "答案速查",
        "target_template_id": "",
        "report_level": "summary",
    }


def test_workbench_runner_writes_technical_long_document_delivery_package(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Chapter 1 Overview")
    doc.add_paragraph("Appendix A Assets")
    doc.save(source)

    scene = SceneWorkspace(
        scene_id="technical",
        category="technical",
        template_id="tech_standard",
    )
    apply_planned_scene_family_defaults(scene)

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    output_dir = source.parent / "output"
    assert payload["status"] == "success"
    assert set(payload["output_paths"]) == {
        "review_copy",
        "proof_copy",
        "final_docx",
    }
    assert Path(payload["output_paths"]["review_copy"]) == (
        output_dir / "technical" / "review_copy" / "source_review_copy.docx"
    )
    assert Path(payload["output_paths"]["proof_copy"]) == (
        output_dir / "technical" / "proof_copy" / "source_proof_copy.docx"
    )
    assert Path(payload["output_paths"]["final_docx"]) == (
        output_dir / "technical" / "final_docx" / "source_final_docx.docx"
    )
    assert set(payload["compare_paths"]) == {"review_copy", "proof_copy"}
    assert set(payload["intermediate_paths"]) == {
        "review_copy",
        "proof_copy",
        "archive_package",
    }
    assert payload["material_manifest_paths"]["material"].endswith(
        "source_material_manifest.json"
    )
    assert payload["material_package_paths"]["package_manifest"].endswith(
        "package_manifest.json"
    )
    assert Path(payload["material_package_paths"]["zip"]).exists() is True

    manifest = json.loads(
        Path(payload["material_manifest_paths"]["material"]).read_text(
            encoding="utf-8"
        )
    )
    assert manifest["delivery"]["output_paths"].keys() >= {
        "review_copy",
        "proof_copy",
        "final_docx",
    }
    assert manifest["delivery"]["compare_paths"].keys() >= {
        "review_copy",
        "proof_copy",
    }
    assert "archive_package" in manifest["delivery"]["intermediate_paths"]

    package_manifest = json.loads(
        Path(payload["material_package_paths"]["package_manifest"]).read_text(
            encoding="utf-8"
        )
    )
    assert package_manifest["kind"] == "material_delivery_package"


def test_workbench_runner_writes_product_pre_sales_delivery_package(tmp_path, monkeypatch):
    source = tmp_path / "product.docx"
    doc = Document()
    doc.add_paragraph("Product overview")
    doc.add_paragraph("Customer case")
    doc.save(source)

    scene = SceneWorkspace(
        scene_id="product_sales_documents",
        category="product_sales_documents",
    )
    apply_planned_scene_family_defaults(scene)

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    output_dir = source.parent / "output"
    assert payload["status"] == "success"
    assert set(payload["output_paths"]) == {
        "customer_copy",
        "internal_review",
        "pre_sales_package",
    }
    assert Path(payload["output_paths"]["customer_copy"]) == (
        output_dir / "product" / "customer_copy" / "product_customer_copy.docx"
    )
    assert Path(payload["output_paths"]["internal_review"]) == (
        output_dir / "product" / "internal_review" / "product_internal_review.docx"
    )
    assert Path(payload["output_paths"]["pre_sales_package"]) == (
        output_dir / "product" / "pre_sales_package" / "product_pre_sales_package.docx"
    )
    assert set(payload["compare_paths"]) == {"internal_review", "pre_sales_package"}
    assert set(payload["intermediate_paths"]) == {
        "customer_copy",
        "internal_review",
        "asset_report",
        "pre_sales_package",
    }
    assert payload["material_manifest_paths"]["material"].endswith(
        "product_material_manifest.json"
    )
    assert payload["material_package_paths"]["package_manifest"].endswith(
        "package_manifest.json"
    )
    assert Path(payload["material_package_paths"]["zip"]).exists() is True

    manifest = json.loads(
        Path(payload["material_manifest_paths"]["material"]).read_text(
            encoding="utf-8"
        )
    )
    assert manifest["material_schema"]["schema_id"] == "product_assets_v1"
    assert manifest["material_schema"]["schema_ids"] == [
        "product_assets_v1",
        "case_study_assets_v1",
    ]
    assert manifest["delivery"]["output_paths"].keys() >= {
        "customer_copy",
        "internal_review",
        "pre_sales_package",
    }
    assert manifest["delivery"]["compare_paths"].keys() >= {
        "internal_review",
        "pre_sales_package",
    }
    assert "asset_report" in manifest["delivery"]["intermediate_paths"]

    package_manifest = json.loads(
        Path(payload["material_package_paths"]["package_manifest"]).read_text(
            encoding="utf-8"
        )
    )
    assert package_manifest["kind"] == "material_delivery_package"
    packaged_files = {
        (item.get("category"), item.get("key"))
        for item in package_manifest["files"]
    }
    assert ("output", "pre_sales_package") in packaged_files
    assert ("compare", "pre_sales_package") in packaged_files
    assert ("intermediate", "pre_sales_package") in packaged_files


def test_workbench_runner_blocks_regulated_disclosure_without_external_receipt(
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "disclosure.docx"
    doc = Document()
    doc.add_paragraph("Annual disclosure")
    doc.add_paragraph("ESG table notes")
    doc.save(source)

    scene = SceneWorkspace(
        scene_id="regulated_disclosure_documents",
        category="regulated_disclosure_documents",
    )
    apply_planned_scene_family_defaults(scene)

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    output_dir = source.parent / "output"
    assert payload["status"] == "failed"
    assert payload["error_text"] == (
        "pipeline_configuration_invalid:plugin_manual_gate_required:"
        "professional_disclosure_review_gate:regulated_disclosure_documents:"
        "verified_external_receipt_missing"
    )
    assert payload["output_paths"] == {}
    assert payload["compare_paths"] == {}
    assert list(output_dir.rglob("*.docx")) == []


def test_workbench_runner_writes_material_package_for_failed_delivery_run(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="archive",
        delivery_presets=[
            DeliveryPreset(
                preset_id="archive",
                label="Archive",
                artifacts=OutputConfig(
                    final_docx=True,
                    material_manifest=True,
                    material_package=True,
                ),
            ),
        ],
    )
    scene.input_source_profile.required_material_fields = ["employee_id"]

    class _FailingPipeline:
        def __init__(self, **_kwargs):
            pass

        def execute(self, _doc_path: str) -> PipelineResult:
            return PipelineResult(
                success=False,
                status="failed",
                error="broken delivery",
                failed_items=[{"preset_id": "archive", "reason": "broken delivery"}],
                output_paths={},
            )

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])
    monkeypatch.setattr(execution_runtime, "Pipeline", _FailingPipeline)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["failed_count"] == 1
    assert payload["error_text"] == "broken delivery"
    assert payload["output_paths"] == {}
    assert payload["compare_paths"] == {}
    assert payload["intermediate_paths"] == {}
    assert {Path(path).suffix for path in payload["report_paths"]} == {
        ".json",
        ".md",
    }
    assert all(Path(path).is_file() for path in payload["report_paths"])
    assert payload["material_manifest_paths"]["material"].endswith(
        "source_material_manifest.json"
    )
    assert payload["material_package_paths"]["package_manifest"].endswith(
        "package_manifest.json"
    )
    assert payload["material_package_paths"]["zip"]
    assert Path(payload["material_package_paths"]["zip"]).exists() is True

    manifest = json.loads(
        Path(payload["material_manifest_paths"]["material"]).read_text(
            encoding="utf-8"
        )
    )
    assert manifest["kind"] == "material_attachment_manifest"
    assert manifest["delivery"]["output_paths"] == {}
    assert manifest["delivery"]["report_paths"] == payload["report_paths"]
    assert manifest["material_schema"]["required_field_keys"] == ["employee_id"]


def test_workbench_runner_writes_structured_intermediate_for_delivery_preset(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("正文内容")
    doc.save(source)

    scene = SceneWorkspace(
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="final",
                label="Final",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=False),
            ),
            DeliveryPreset(
                preset_id="review",
                label="Review",
                output_dir_template="{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(final_docx=True, report_json=False, report_markdown=True),
                include_structured_intermediate=True,
                report_level="detailed",
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    output_dir = source.parent / "output"
    intermediate_path = output_dir / "review" / "source_review_intermediate.json"
    assert payload["status"] == "success"
    assert payload["intermediate_paths"] == {"review": str(intermediate_path)}
    assert intermediate_path.exists() is True

    intermediate = json.loads(intermediate_path.read_text(encoding="utf-8"))
    assert intermediate["kind"] == "delivery_structured_intermediate"
    assert intermediate["schema_version"] == 1
    assert intermediate["input"] == str(source)
    assert intermediate["output"] == str(output_dir / "review" / "source_review.docx")
    assert intermediate["preset"]["preset_id"] == "review"
    assert intermediate["preset"]["include_structured_intermediate"] is True
    assert intermediate["preset"]["artifacts"]["report_markdown"] is True
    assert intermediate["preset"]["content_visibility_rules"] == []
    assert intermediate["modules_enabled"] == 0
    assert intermediate["diagnostics"]["count"] == 0
    assert intermediate["changes"] == []
    assert intermediate["context"]["source_doc_path"] == str(source)


def test_pipeline_context_payload_includes_technical_chapter_inventory():
    context = SimpleNamespace(
        source_doc_path="source.docx",
        source_doc_dir="",
        entity_values={},
        source_values={},
        exam_question_schema=None,
        journal_citations=None,
        journal_submission_package=None,
        official_numbering_preservation=None,
        technical_chapter_inventory={
            "status": "ok",
            "summary": {"chapter_count": 1, "appendix_count": 1},
        },
        application_section_word_limits={
            "status": "warning",
            "summary": {"exceeded_section_count": 1},
        },
        inserted_images=[],
    )

    payload = delivery_reporting.pipeline_context_payload(context)

    assert payload["technical_chapter_inventory"]["status"] == "ok"
    assert payload["technical_chapter_inventory"]["summary"]["chapter_count"] == 1
    assert payload["application_section_word_limits"]["status"] == "warning"
    assert (
        payload["application_section_word_limits"]["summary"]["exceeded_section_count"]
        == 1
    )


def test_workbench_report_templates_render_preset_label_with_shared_display_name(tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    output_dir = tmp_path / "out"

    preset = DeliveryPreset(
        preset_id="answer_key",
        label="Answer key",
        output_dir_template="{preset_label}",
        filename_template="{stem}_{preset_label}",
    )
    assert delivery_reporting.delivery_report_stem(source, preset) == "source_答案速查"
    assert delivery_reporting.delivery_report_dir(output_dir, source, preset) == (
        output_dir / "答案速查"
    )

    custom = DeliveryPreset(
        preset_id="review",
        label="Review package",
        output_dir_template="{preset_label}",
        filename_template="{stem}_{preset_label}",
    )
    assert delivery_reporting.delivery_report_stem(source, custom) == "source_Review package"
    assert delivery_reporting.delivery_report_dir(output_dir, source, custom) == (
        output_dir / "Review package"
    )


def test_workbench_delivery_preset_payload_keeps_raw_and_display_labels():
    payload = delivery_reporting.delivery_preset_payload(
        DeliveryPreset(preset_id="answer_key", label="Answer key")
    )

    assert payload["label"] == "Answer key"
    assert payload["display_label"] == "答案速查"


def test_workbench_runner_writes_compare_docx_for_delivery_preset(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Before text")
    doc.save(source)

    class _ReplaceTextModule(BaseModule):
        meta = ModuleMeta(
            name="replace_text",
            description="Replace text",
            category="test",
            execution_phase="fill",
            scope_behavior="region_filtered",
        )

        def apply(self, doc, config, tracker, context):
            doc.paragraphs[0].text = "After text"
            tracker.record(
                rule_name=self.meta.name,
                target="paragraph 1",
                section="body",
                change_type="replace",
                before="Before text",
                after="After text",
            )

    scene = SceneWorkspace(
        default_delivery_preset_id="review",
        delivery_presets=[
            DeliveryPreset(
                preset_id="review",
                label="Review",
                output_dir_template="{preset_id}",
                filename_template="{stem}_{preset_id}",
                artifacts=OutputConfig(
                    final_docx=True,
                    compare_docx=True,
                    compare_text=True,
                    compare_formatting=False,
                    report_json=False,
                    report_markdown=False,
                ),
            ),
        ],
    )
    scene.compliance_profile.object_preflight.enabled = False

    monkeypatch.setattr(
        execution_runtime,
        "create_all_modules",
        lambda: [_ReplaceTextModule()],
    )
    monkeypatch.setattr(
        execution_runtime,
        "build_module_selection_plan",
        lambda modules, predicate: SimpleNamespace(
            select_modules=lambda available: tuple(available),
        ),
    )

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    output_dir = source.parent / "output"
    compare_path = output_dir / "review" / "source_review_compare.docx"
    assert payload["status"] == "success"
    assert payload["output_paths"] == {
        "review": str(output_dir / "review" / "source_review.docx"),
    }
    assert payload["compare_paths"] == {"review": str(compare_path)}
    assert compare_path.exists() is True

    with zipfile.ZipFile(compare_path) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "<w:del " in document_xml
    assert "<w:ins " in document_xml
    assert "Before" in document_xml
    assert "After" in document_xml


def test_workbench_runner_surfaces_diagnostics_summary(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)

    template = TemplateConfig()
    scene = SceneWorkspace()
    artifacts = scene.default_delivery_preset().artifacts
    artifacts.final_docx = False
    artifacts.report_json = False
    artifacts.report_markdown = False

    tracker = ChangeTracker()
    tracker.record(
        rule_name="equation_table_format",
        target="1 个公式编号",
        section="global",
        change_type="skip",
        before="chapter-aware numbering normalization",
        after="skipped due to missing chapter context",
    )

    class _StubPipeline:
        def __init__(self, **_kwargs):
            pass

        def execute(self, _doc_path: str) -> PipelineResult:
            return PipelineResult(
                success=True,
                status="success",
                tracker=tracker,
                output_paths={},
            )

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])
    monkeypatch.setattr(execution_runtime, "Pipeline", _StubPipeline)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=template,
        scene=scene,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert payload["diagnostics_count"] == 1
    assert "missing chapter context" in payload["diagnostics_summary"]


def test_workbench_runner_applies_session_overrides_before_pipeline(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)

    captured = {}

    class _StubPipeline:
        def __init__(self, **kwargs):
            captured["config"] = kwargs["config"]

        def execute(self, _doc_path: str) -> PipelineResult:
            return PipelineResult(
                success=True,
                status="success",
                output_paths={},
            )

    monkeypatch.setattr(execution_runtime, "create_all_modules", lambda: [])
    monkeypatch.setattr(execution_runtime, "Pipeline", _StubPipeline)

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(),
        session_overrides={"header_footer.suppress_header_footer_selectors": []},
    ).run(lambda *_args: None, lambda: False)

    config = captured["config"]
    assert payload["status"] == "success"
    assert config.header_footer.suppress_header_footer_selectors == []
    assert config.get_with_source("header_footer.suppress_header_footer_selectors").source == "session"


def test_workbench_runner_session_override_changes_cover_page_number_visibility(tmp_path, monkeypatch):
    source = tmp_path / "thesis.docx"
    doc = Document()
    for text in ["博士学位论文", "原创性声明", "摘要", "目录"]:
        doc.add_paragraph(text)
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容")
    doc.save(source)

    monkeypatch.setattr(
        execution_runtime,
        "create_all_modules",
        lambda: [
            HeadingRecognitionModule(),
            SectionFormatModule(),
            HeaderFooterModule(),
        ],
    )

    default_payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(),
    ).run(lambda *_args: None, lambda: False)
    default_doc = Document(default_payload["output_path"])

    override_payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=SceneWorkspace(),
        session_overrides={"header_footer.suppress_header_footer_selectors": []},
    ).run(lambda *_args: None, lambda: False)
    override_doc = Document(override_payload["output_path"])

    assert default_payload["status"] in {"success", "partial_success"}
    assert override_payload["status"] in {"success", "partial_success"}
    assert _section_has_page_field(default_doc.sections[0]) is False
    assert _section_has_page_field(override_doc.sections[0]) is True


def _section_has_page_field(section) -> bool:
    footer = section.footer
    for para in footer.paragraphs:
        for _kind, _elem, instr in iter_field_instructions(para._element):
            if "PAGE" in str(instr or "").upper():
                return True
    return False
