import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.scene import ContentVisibilityRule, DeliveryPreset
from src.shared.engine.content_visibility import (
    content_visibility_configuration_diagnostics,
    preview_content_visibility_effects,
    scan_content_visibility_markers,
)


def test_content_visibility_scan_matches_rules_and_reports_marker_issues():
    doc = Document()
    doc.add_paragraph("{{#visibility:answer}}")
    doc.add_paragraph("Answer")
    doc.add_paragraph("{{/visibility:answer}}")
    doc.add_paragraph("{{/content:solution}}")
    doc.add_paragraph("{{#visibility:analysis}}")
    doc.add_paragraph("{{#visibility:hint}}")

    scan = scan_content_visibility_markers(
        doc,
        [
            DeliveryPreset(
                preset_id="student",
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove"),
                    ContentVisibilityRule(selector="missing", action="remove"),
                ],
            )
        ],
    )

    assert scan.selectors == ["analysis", "answer", "hint", "solution"]
    assert scan.start_counts == {"analysis": 1, "answer": 1, "hint": 1}
    assert scan.end_counts == {"answer": 1, "solution": 1}
    assert scan.used_rule_selectors == ["answer", "missing"]
    assert scan.missing_rule_selectors == ["missing"]
    assert scan.unused_document_selectors == ["analysis", "hint", "solution"]
    assert scan.orphan_end_selectors == {"solution": 1}
    assert scan.nested_selectors == {"hint": 1}
    assert scan.unclosed_selectors == {"analysis": 1, "hint": 1}
    assert scan.has_issues is True
    assert "规则 selector 未在文档中找到: missing" in scan.issue_messages()


def test_content_visibility_scan_accepts_content_alias_and_dict_rules():
    doc = Document()
    doc.add_paragraph("{{#content:answer}}")
    doc.add_paragraph("Answer")
    doc.add_paragraph("{{/content:answer}}")

    scan = scan_content_visibility_markers(
        doc,
        [
            {
                "content_visibility_rules": [
                    {"selector": "answer", "action": "remove"},
                ]
            }
        ],
    )

    assert scan.selectors == ["answer"]
    assert scan.used_rule_selectors == ["answer"]
    assert scan.missing_rule_selectors == []
    assert scan.unused_document_selectors == []
    assert scan.has_issues is False


def test_content_visibility_scan_blocks_invalid_executable_rule_configuration():
    doc = Document()
    doc.add_paragraph("Question only")
    presets = [
        {
            "preset_id": "student",
            "content_visibility_rules": [
                {
                    "rule_id": "bad-type",
                    "selector_type": "marker_blocks",
                    "selector": "answer",
                    "action": "remove",
                },
                {
                    "rule_id": "bad-action",
                    "selector_type": "marker_block",
                    "selector": "analysis",
                    "action": "remvoe",
                },
                {
                    "rule_id": "retired-hide-action",
                    "selector_type": "marker_block",
                    "selector": "answer",
                    "action": "hide",
                },
                {
                    "rule_id": "retired-exclude-action",
                    "selector_type": "marker_block",
                    "selector": "solution",
                    "action": "exclude",
                },
            ],
        }
    ]

    diagnostics = content_visibility_configuration_diagnostics(presets)
    scan = scan_content_visibility_markers(doc, presets)

    expected_codes = {
        "content_visibility_selector_type_invalid",
        "content_visibility_action_invalid",
    }
    assert {item["code"] for item in diagnostics} == expected_codes
    invalid_action_messages = [
        str(item["message"])
        for item in diagnostics
        if item["code"] == "content_visibility_action_invalid"
    ]
    assert len(invalid_action_messages) == 3
    assert any("action 'hide'" in message for message in invalid_action_messages)
    assert any("action 'exclude'" in message for message in invalid_action_messages)
    assert expected_codes <= {
        item["code"] for item in scan.blocking_diagnostics
    }
    assert scan.has_blocking_issues is True
    assert scan.grammar_valid is True


def test_content_visibility_scan_blocks_sensitive_selector_near_match_only():
    answer_doc = Document()
    answer_doc.add_paragraph("{{#visibility:answer}}")
    answer_doc.add_paragraph("Sensitive answer")
    answer_doc.add_paragraph("{{/visibility:answer}}")
    typo_preset = DeliveryPreset(
        preset_id="student",
        content_visibility_rules=[
            ContentVisibilityRule(selector="anwser", action="remove")
        ],
    )

    typo_scan = scan_content_visibility_markers(answer_doc, [typo_preset])

    assert typo_scan.missing_rule_selectors == ["anwser"]
    assert typo_scan.unused_document_selectors == ["answer"]
    assert {
        item["code"] for item in typo_scan.blocking_diagnostics
    } == {"content_visibility_sensitive_selector_mismatch"}

    optional_doc = Document()
    optional_doc.add_paragraph("{{#visibility:appendix}}")
    optional_doc.add_paragraph("Optional appendix")
    optional_doc.add_paragraph("{{/visibility:appendix}}")
    optional_preset = DeliveryPreset(
        preset_id="student",
        content_visibility_rules=[
            ContentVisibilityRule(selector="answer", action="remove")
        ],
    )

    optional_scan = scan_content_visibility_markers(optional_doc, [optional_preset])

    assert optional_scan.missing_rule_selectors == ["answer"]
    assert optional_scan.unused_document_selectors == ["appendix"]
    assert optional_scan.has_blocking_issues is False


def test_content_visibility_preview_reports_per_preset_removed_blocks():
    doc = Document()
    doc.add_paragraph("Question")
    doc.add_paragraph("{{#visibility:answer}}")
    doc.add_paragraph("Answer")
    doc.add_paragraph("{{/visibility:answer}}")
    doc.add_paragraph("End")

    previews = preview_content_visibility_effects(
        doc,
        [
            DeliveryPreset(
                preset_id="student",
                label="Student",
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove"),
                    ContentVisibilityRule(selector="missing", action="remove"),
                ],
            ),
            DeliveryPreset(preset_id="teacher", label="Teacher"),
        ],
    )

    student, teacher = previews
    assert student.preset_id == "student"
    assert student.remove_selectors == ["answer", "missing"]
    assert student.matched_selectors == ["answer"]
    assert student.missing_selectors == ["missing"]
    assert student.removed_block_count == 1
    assert student.removed_paragraph_count == 3
    assert student.stripped_marker_paragraph_count == 0
    assert student.removed_text_samples == ["Answer"]
    assert len(student.removed_blocks) == 1
    block = student.removed_blocks[0]
    assert block.selector == "answer"
    assert block.start_paragraph_index == 1
    assert block.end_paragraph_index == 3
    assert block.paragraph_count == 3
    assert block.text_samples == ["Answer"]
    assert block.context_before == "Question"
    assert block.context_after == "End"

    assert teacher.preset_id == "teacher"
    assert teacher.remove_selectors == []
    assert teacher.matched_selectors == []
    assert teacher.removed_block_count == 0
    assert teacher.removed_paragraph_count == 0
    assert teacher.stripped_marker_paragraph_count == 2
    assert teacher.removed_blocks == []


def test_content_visibility_preview_uses_strict_body_plan_for_table_counts():
    doc = Document()
    doc.add_paragraph("Question")
    doc.add_paragraph("{{#visibility:answer}}")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Answer A"
    table.cell(0, 1).text = "Answer B"
    doc.add_paragraph("{{/visibility:answer}}")
    doc.add_paragraph("End")

    preview = preview_content_visibility_effects(
        doc,
        [
            DeliveryPreset(
                preset_id="student",
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove")
                ],
            )
        ],
    )[0]

    assert preview.matched_selectors == ["answer"]
    assert preview.missing_selectors == []
    assert preview.removed_body_element_count == 3
    assert preview.removed_paragraph_count == 4
    assert preview.removed_table_count == 1
    assert preview.removed_text_samples == ["Answer A", "Answer B"]
    assert preview.strict_plan_id
    assert len(preview.removed_blocks) == 1
    block = preview.removed_blocks[0]
    assert block.start_body_index == 1
    assert block.end_body_index == 3
    assert block.body_element_count == 3
    assert block.paragraph_count == 4
    assert block.table_count == 1
    assert block.context_before == "Question"
    assert block.context_after == "End"


def test_content_visibility_scan_projects_strict_nested_surface_diagnostic():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{#visibility:answer}}"
    doc.add_paragraph("{{/visibility:answer}}")

    scan = scan_content_visibility_markers(
        doc,
        [
            DeliveryPreset(
                preset_id="student",
                content_visibility_rules=[
                    ContentVisibilityRule(selector="answer", action="remove")
                ],
            )
        ],
    )

    assert scan.grammar_valid is False
    assert scan.has_blocking_issues is True
    assert scan.strict_plan_id == ""
    assert {
        item["code"] for item in scan.blocking_diagnostics
    } >= {"marker_not_direct_body", "orphan_end_marker"}


def test_content_visibility_scan_projects_mismatch_and_unsupported_range():
    mismatched = Document()
    mismatched.add_paragraph("{{#visibility:answer}}")
    mismatched.add_paragraph("{{/visibility:detail}}")
    mismatched.add_paragraph("{{/visibility:answer}}")

    mismatch_scan = scan_content_visibility_markers(mismatched)
    assert mismatch_scan.mismatched_end_selectors == {"detail": 1}
    assert "mismatched_end_marker" in {
        item["code"] for item in mismatch_scan.blocking_diagnostics
    }

    unsupported = Document()
    unsupported.add_paragraph("{{#visibility:answer}}")
    unsupported.add_paragraph("{{/visibility:answer}}")
    unsupported.element.body.insert(1, OxmlElement("w:sdt"))

    unsupported_scan = scan_content_visibility_markers(unsupported)
    assert unsupported_scan.has_blocking_issues is True
    assert "unsupported_body_element_in_range" in {
        item["code"] for item in unsupported_scan.blocking_diagnostics
    }
