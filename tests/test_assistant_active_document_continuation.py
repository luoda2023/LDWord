from __future__ import annotations

import hashlib
import time
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document

from src.assistant.adapters.content_generation_adapter import GeneratedExamDraft
from src.assistant.application.active_document_continuation import (
    ACTION_APPROVE_EXECUTE,
    ACTION_GENERATE_CONTENT_DRAFT,
    ACTION_OPEN_ARTIFACT,
    ACTION_OPEN_CONTENT_DRAFT,
    ACTION_OPEN_OUTPUT_FOLDER,
    ACTION_PREFLIGHT,
    ACTION_RETRY_PREFLIGHT,
    ACTION_REVISE_CONTENT_DRAFT,
    ACTION_RESTORE_PREVIOUS_DRAFT,
    is_active_document_revision_request,
    resolve_active_document_continuation,
    validate_active_document_action,
)
from src.assistant.application.content_generation_service import (
    ContentGenerationRequest,
)
from src.assistant.application.document_job_controller import (
    DocumentJobController,
)
from src.assistant.application.exam_plan_editing import ExamPlanEditValues
from src.assistant.application.session_coordinator import (
    AssistantSessionCoordinator,
)
from src.assistant.contracts.document_plan import DocumentPlan, OutputPolicy
from src.assistant.contracts.execution import PreflightReceipt
from src.assistant.contracts.material_snapshot import MaterialExecutionEnvelope
from src.assistant.contracts.messages import ROLE_ASSISTANT, ROLE_USER, AssistantMessage
from src.assistant.contracts.task_plan import (
    ARTIFACT_KIND_EXAM,
    SOURCE_ROLE_STRUCTURED_SOURCE,
    DeliveryContract,
    GenerationContract,
    ProductionContract,
    SourceArtifactRef,
)
from src.assistant.runtime.provider_contract import (
    PROVIDER_DONE,
    PROVIDER_START,
    ProviderStreamEvent,
)
from src.assistant.runtime.turn_runner import AssistantTurnRunner
from src.assistant.storage.session_store import AssistantSessionStore
from src.assistant.ui.assistant_panel import AssistantPanel
from src.assistant.ui.conversation_presentation import (
    build_interaction_action_scope,
)
from src.assistant.ui.exam_plan_editor import ExamPlanEditor
from src.ui.bridge import PanelBridge


@pytest.mark.parametrize(
    ("query", "action_id"),
    [
        ("把这个拼装好，最后给我word文档", ACTION_PREFLIGHT),
        ("继续，输出 Word", ACTION_PREFLIGHT),
        ("给我最终docx文档", ACTION_PREFLIGHT),
        ("打开内容草稿", ACTION_OPEN_CONTENT_DRAFT),
        ("重新生成试卷", ACTION_GENERATE_CONTENT_DRAFT),
    ],
)
def test_content_draft_followups_resolve_to_bounded_host_actions(
    query,
    action_id,
):
    decision = resolve_active_document_continuation(
        query,
        job_status="content_draft_ready",
        has_active_plan=True,
    )

    assert decision is not None
    assert decision.action_id == action_id


def test_revision_request_is_not_mistaken_for_production_permission():
    query = "修改第 3 题并调整题量，最后再生成 Word"

    draft_decision = resolve_active_document_continuation(
        query,
        job_status="content_draft_ready",
        has_active_plan=True,
    )
    approval_decision = resolve_active_document_continuation(
        query,
        job_status="needs_execution_approval",
        has_active_plan=True,
    )

    assert draft_decision is not None
    assert draft_decision.action_id == ACTION_REVISE_CONTENT_DRAFT
    assert approval_decision is not None
    assert approval_decision.action_id == ACTION_REVISE_CONTENT_DRAFT
    assert is_active_document_revision_request(
        query,
        job_status="content_draft_ready",
    )
    assert is_active_document_revision_request(
        query,
        job_status="needs_execution_approval",
    )


def test_failed_revision_can_resolve_restore_previous_draft():
    decision = resolve_active_document_continuation(
        "撤销这次修改，恢复上一版草稿",
        job_status="failed",
        has_active_plan=True,
    )

    assert decision is not None
    assert decision.action_id == ACTION_RESTORE_PREVIOUS_DRAFT


def test_approval_requires_approval_state_and_explicit_language():
    approved = resolve_active_document_continuation(
        "确认并生成 Word",
        job_status="needs_execution_approval",
        has_active_plan=True,
    )
    casual = resolve_active_document_continuation(
        "继续看看",
        job_status="needs_execution_approval",
        has_active_plan=True,
    )
    wrong_state = resolve_active_document_continuation(
        "确认并生成 Word",
        job_status="content_draft_ready",
        has_active_plan=True,
    )

    assert approved is not None
    assert approved.action_id == ACTION_APPROVE_EXECUTE
    assert casual is None
    assert wrong_state is not None
    assert wrong_state.action_id == ACTION_PREFLIGHT


def test_pending_provider_or_local_question_blocks_document_shortcuts():
    decision = resolve_active_document_continuation(
        "继续",
        job_status="content_draft_ready",
        has_active_plan=True,
        pending_kind="official_field_completion",
    )

    assert decision is None


@pytest.mark.parametrize(
    ("generation_pending", "action_id"),
    [
        (True, ACTION_GENERATE_CONTENT_DRAFT),
        (False, ACTION_PREFLIGHT),
    ],
)
def test_final_docx_request_advances_the_current_plan(
    generation_pending,
    action_id,
):
    decision = resolve_active_document_continuation(
        "给我最终docx文档",
        job_status="plan_ready",
        has_active_plan=True,
        generation_pending=generation_pending,
    )

    assert decision is not None
    assert decision.action_id == action_id


@pytest.mark.parametrize(
    "query",
    ("下一步", "按这个来", "就这样", "可以", "命题", "出题"),
)
def test_short_plan_followups_start_the_pending_generation(query):
    decision = resolve_active_document_continuation(
        query,
        job_status="plan_ready",
        has_active_plan=True,
        generation_pending=True,
    )

    assert decision is not None
    assert decision.action_id == ACTION_GENERATE_CONTENT_DRAFT


def test_authorized_but_not_started_generation_can_resume_locally():
    decision = resolve_active_document_continuation(
        "继续",
        job_status="content_generation_ready",
        has_active_plan=True,
        generation_pending=True,
    )

    assert decision is not None
    assert decision.action_id == ACTION_GENERATE_CONTENT_DRAFT


@pytest.mark.parametrize(
    ("status", "query", "action_id"),
    [
        ("preflight_failed", "重新检查", ACTION_RETRY_PREFLIGHT),
        ("needs_execution_approval", "重新检查", ACTION_RETRY_PREFLIGHT),
        ("needs_execution_approval", "给我最终docx文档", ACTION_APPROVE_EXECUTE),
        ("success", "打开生成的文档", ACTION_OPEN_ARTIFACT),
        ("success", "给我最终docx文档", ACTION_OPEN_ARTIFACT),
        ("partial_success", "查看结果", ACTION_OPEN_ARTIFACT),
    ],
)
def test_recovery_and_output_followups_are_state_specific(
    status,
    query,
    action_id,
):
    decision = resolve_active_document_continuation(
        query,
        job_status=status,
        has_active_plan=True,
    )

    assert decision is not None
    assert decision.action_id == action_id


def test_action_validation_rejects_stale_plan_and_stale_preflight(
    tmp_path,
):
    plan = _bound_exam_plan(tmp_path)
    job = _draft_job(plan)

    assert (
        validate_active_document_action(
            ACTION_PREFLIGHT,
            job=job,
            plan=plan,
        )
        == ""
    )
    assert validate_active_document_action(
        ACTION_PREFLIGHT,
        job={**job, "plan_revision": plan.revision - 1},
        plan=plan,
    ) == "assistant_document_action_plan_stale"

    stale = _preflight(plan, plan_fingerprint="stale")
    approval_job = {
        **job,
        "status": "needs_execution_approval",
        "preflight": stale.to_dict(),
    }
    assert validate_active_document_action(
        ACTION_APPROVE_EXECUTE,
        job=approval_job,
        plan=plan,
    ) == "assistant_preflight_stale"
    with pytest.raises(ValueError):
        DocumentJobController.approve(
            session_id="session",
            plan=plan,
            preflight=stale,
        )


@pytest.mark.parametrize("status", ("success", "partial_success"))
def test_output_folder_action_is_available_for_completed_artifacts(
    tmp_path,
    status,
):
    plan = _bound_exam_plan(tmp_path)

    assert validate_active_document_action(
        ACTION_OPEN_OUTPUT_FOLDER,
        job={"status": status},
        plan=plan,
    ) == ""


class _CaptureGateway:
    def __init__(self) -> None:
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(PROVIDER_DONE, text="ok")

    def cancel(self) -> bool:
        return True


def test_explicit_recovery_command_promotes_a_previous_chat_request(
    qapp,
    tmp_path,
    monkeypatch,
):
    gateway = _CaptureGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "sessions")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    session = coordinator.create_session()
    session = coordinator.append_message(
        session,
        AssistantMessage.text(role=ROLE_USER, text="给我整一份通知"),
    )
    session = coordinator.append_message(
        session,
        AssistantMessage.text(role=ROLE_ASSISTANT, text="旧版误入普通对话"),
    )
    panel._active_session = session
    captured = []
    monkeypatch.setattr(
        panel,
        "_handle_local_request_policy",
        lambda policy, **kwargs: captured.append((policy, kwargs)),
    )
    try:
        assert panel._send_message("转为文档任务") is True

        assert gateway.requests == []
        assert len(captured) == 1
        assert captured[0][0].kind == "document_action"
        assert captured[0][0].query == "给我整一份通知"
    finally:
        panel.close()


def test_typed_handoff_uses_current_plan_without_provider_or_replanning(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel, gateway, original_plan = _panel_with_bound_draft(tmp_path)
    dispatched = []
    monkeypatch.setattr(
        panel,
        "_run_preflight",
        lambda session, plan: dispatched.append((session, plan)),
    )
    try:
        assert panel._send_message("给我最终docx文档") is True

        assert len(dispatched) == 1
        assert dispatched[0][1].fingerprint == original_plan.fingerprint
        assert gateway.requests == []
        persisted = panel._coordinator.load_session(
            panel._active_session.session_id
        )
        persisted_plan = DocumentPlan.from_dict(persisted.active_plan)
        assert persisted_plan.fingerprint == original_plan.fingerprint
        assert persisted_plan.production_input_artifact is not None
        assert persisted_plan.production_input_artifact.path == (
            original_plan.production_input_artifact.path
        )
        assert persisted.messages[-1].role == ROLE_USER
    finally:
        panel.close()


def test_card_and_typed_handoff_share_current_plan_dispatch(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel, gateway, original_plan = _panel_with_bound_draft(tmp_path)
    dispatched = []
    monkeypatch.setattr(
        panel,
        "_run_preflight",
        lambda session, plan: dispatched.append((session, plan)),
    )
    try:
        session = panel._active_session
        assert session is not None
        action_scope = build_interaction_action_scope(
            pending_continuation=session.pending_continuation,
            active_plan=session.active_plan,
            document_job=session.document_job,
            turn_status=session.turn_status,
        )
        panel._handle_card_action(
            "preflight",
            {
                "action_scope": {
                    **action_scope,
                    "job_status": "plan_ready",
                }
            },
        )
        assert dispatched == []
        panel._handle_card_action(
            "preflight",
            {
                "action_scope": action_scope,
            },
        )

        assert len(dispatched) == 1
        assert dispatched[0][1].fingerprint == original_plan.fingerprint
        assert gateway.requests == []
    finally:
        panel.close()


def test_draft_revision_creates_a_new_generation_plan_and_rebinds_new_source(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel, gateway, original_plan = _panel_with_bound_draft(tmp_path)
    preflight_calls = []
    generation_calls = []
    monkeypatch.setattr(
        panel,
        "_run_preflight",
        lambda session, plan: preflight_calls.append((session, plan)),
    )
    monkeypatch.setattr(
        panel,
        "_start_content_generation",
        lambda session, plan, **kwargs: generation_calls.append(
            (session, plan, kwargs)
        ),
    )
    try:
        assert panel._send_message(
            "修改第3题并调整题量，最后再生成 Word"
        ) is True
        assert preflight_calls == []
        assert gateway.requests == []
        assert len(generation_calls) == 1
        persisted = panel._coordinator.load_session(
            panel._active_session.session_id
        )
        persisted_plan = DocumentPlan.from_dict(persisted.active_plan)
        assert persisted_plan.revision == original_plan.revision + 1
        assert persisted_plan.fingerprint != original_plan.fingerprint
        assert persisted.document_job["status"] == "plan_ready"
        assert persisted.document_job["invalidated_reason"] == (
            "assistant_content_revision_requested"
        )
        assert persisted_plan.production_input_artifact is None
        assert persisted.document_job["revision_source_path"].endswith(
            "current.exam.md"
        )
        assert generation_calls[0][2]["revision_context_text"].startswith(
            "# 小学数学试卷"
        )

        revised_path = tmp_path / "revised.exam.md"
        revised_path.write_text("# 修订后的小学数学试卷", encoding="utf-8")
        revised_draft = GeneratedExamDraft(
            draft_id="revised-draft",
            session_id=persisted.session_id,
            markdown_path=str(revised_path),
            source_digest=hashlib.sha256(revised_path.read_bytes()).hexdigest(),
            artifact_id="artifact-revised-draft",
            schema_id="exam_items_v1",
            validation_summary={},
        )
        running = panel._coordinator.update_state(
            persisted,
            document_job={
                **dict(persisted.document_job),
                "status": "content_generation_running",
            },
        )
        panel._active_session = running
        panel._publish_generated_draft(running, persisted_plan, revised_draft)
        qapp.processEvents()

        rebound = panel._coordinator.load_session(persisted.session_id)
        rebound_plan = DocumentPlan.from_dict(rebound.active_plan)
        assert rebound_plan.revision == original_plan.revision + 2
        assert rebound_plan.production_input_artifact is not None
        assert rebound_plan.production_input_artifact.path == str(revised_path)
        assert rebound.document_job["generated_content_draft_id"] == (
            "revised-draft"
        )
    finally:
        panel.shutdown_active_execution(1000)
        panel.close()


def test_failed_draft_revision_can_restore_previous_bound_source(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel, gateway, original_plan = _panel_with_bound_draft(tmp_path)
    monkeypatch.setattr(
        panel,
        "_start_content_generation",
        lambda *_args, **_kwargs: None,
    )
    try:
        panel._revise_content_draft(
            panel._active_session,
            original_plan,
            "修改第 3 题",
        )
        revised_session = panel._coordinator.load_session(
            panel._active_session.session_id
        )
        revised_plan = DocumentPlan.from_dict(revised_session.active_plan)
        failed_session = panel._coordinator.update_state(
            revised_session,
            document_job={
                **dict(revised_session.document_job),
                "status": "failed",
                "error_text": "provider_error",
            },
        )

        assert validate_active_document_action(
            ACTION_RESTORE_PREVIOUS_DRAFT,
            job=failed_session.document_job,
            plan=revised_plan,
        ) == ""
        panel._restore_previous_content_draft(failed_session, revised_plan)
        qapp.processEvents()

        restored_session = panel._coordinator.load_session(
            failed_session.session_id
        )
        restored_plan = DocumentPlan.from_dict(restored_session.active_plan)
        assert restored_plan.revision == revised_plan.revision + 1
        assert restored_plan.production_input_artifact is not None
        assert restored_plan.production_input_artifact.path == (
            original_plan.production_input_artifact.path
        )
        assert restored_session.document_job["status"] == "content_draft_ready"
        assert "previous_plan_snapshot" not in restored_session.document_job
        assert restored_session.messages[-1].blocks[0].data["title"] == (
            "已恢复上一版内容草稿"
        )
        assert gateway.requests == []
    finally:
        panel.close()


def test_typed_confirmation_rejects_stale_preflight_without_execution(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel, gateway, plan = _panel_with_bound_draft(tmp_path)
    session = panel._active_session
    session = panel._coordinator.update_state(
        session,
        document_job={
            **dict(session.document_job),
            "status": "preflight_running",
        },
    )
    stale = _preflight(plan, plan_fingerprint="stale")
    session = panel._coordinator.update_state(
        session,
        document_job={
            **dict(session.document_job),
            "status": "needs_execution_approval",
            "preflight": stale.to_dict(),
        },
    )
    panel._active_session = session
    executions = []
    monkeypatch.setattr(
        panel,
        "_start_execution",
        lambda current, active_plan: executions.append(
            (current, active_plan)
        ),
    )
    try:
        assert panel._send_message("确认并生成 Word") is True

        assert executions == []
        assert gateway.requests == []
        assert panel._active_session.messages[-1].blocks[0].data[
            "reason"
        ] == "assistant_preflight_stale"
    finally:
        panel.close()


def test_failed_preflight_keeps_blocker_visible_and_local_recovery_available(
    qapp,
    tmp_path,
):
    panel, _gateway, plan = _panel_with_bound_draft(tmp_path)
    session = panel._coordinator.update_state(
        panel._active_session,
        document_job={
            **dict(panel._active_session.document_job),
            "status": "preflight_running",
        },
    )
    panel._active_session = session

    class _FinishedPreflightWorker:
        session_id = session.session_id

        def deleteLater(self):
            return None

    panel._preflight_worker = _FinishedPreflightWorker()
    receipt = PreflightReceipt(
        preflight_id="preflight-blocked",
        plan_id=plan.plan_id,
        plan_revision=plan.revision,
        plan_fingerprint=plan.fingerprint,
        input_hash="a" * 64,
        material_snapshot_digest=MaterialExecutionEnvelope.capture(None).digest,
        output_root=plan.output_policy.output_root,
        ready=False,
        issues=("official_draft_field_missing:organization",),
        warnings=(
            "official_draft_warning:成文日期已暂按当天日期填写。",
        ),
    )
    try:
        panel._on_preflight_finished(receipt)
        qapp.processEvents()

        persisted = panel._coordinator.load_session(session.session_id)
        block = persisted.messages[-1].blocks[0]
        action_ids = [item["id"] for item in block.data["actions"]]
        assert persisted.document_job["status"] == "preflight_failed"
        assert "内容草稿缺少发文机关" in block.text
        assert "official_draft_field_missing" not in block.text
        assert block.data["notices"] == ["成文日期已暂按当天日期填写。"]
        assert action_ids == [
            "retry_preflight",
            "edit_official_plan_requirements",
        ]
    finally:
        panel.close()


def test_ready_preflight_presents_warnings_without_blocking_approval(
    qapp,
    tmp_path,
):
    panel, _gateway, plan = _panel_with_bound_draft(tmp_path)
    session = panel._coordinator.update_state(
        panel._active_session,
        document_job={
            **dict(panel._active_session.document_job),
            "status": "preflight_running",
        },
    )
    panel._active_session = session

    class _FinishedPreflightWorker:
        session_id = session.session_id

        def deleteLater(self):
            return None

    panel._preflight_worker = _FinishedPreflightWorker()
    receipt = PreflightReceipt(
        preflight_id="preflight-warning",
        plan_id=plan.plan_id,
        plan_revision=plan.revision,
        plan_fingerprint=plan.fingerprint,
        input_hash="a" * 64,
        material_snapshot_digest=MaterialExecutionEnvelope.capture(None).digest,
        output_root=plan.output_policy.output_root,
        ready=True,
        warnings=(
            "official_draft_warning:材料未提供发文字号，已使用“待编”。",
        ),
    )
    try:
        panel._on_preflight_finished(receipt)
        qapp.processEvents()

        persisted = panel._coordinator.load_session(session.session_id)
        block = persisted.messages[-1].blocks[0]
        assert persisted.document_job["status"] == "needs_execution_approval"
        assert block.data["title"] == "执行前检查已通过"
        assert block.text == ""
        assert block.data["notices"] == ["材料未提供发文字号，已使用“待编”。"]
        assert block.data["actions"] == [
            {"id": "approve_execute", "label": "确认并生成 Word"}
        ]
    finally:
        panel.close()


def test_revision_request_invalidates_existing_execution_approval(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel, gateway, plan = _panel_with_bound_draft(tmp_path)
    session = panel._active_session
    session = panel._coordinator.update_state(
        session,
        document_job={
            **dict(session.document_job),
            "status": "preflight_running",
        },
    )
    session = panel._coordinator.update_state(
        session,
        document_job={
            **dict(session.document_job),
            "status": "needs_execution_approval",
            "preflight": _preflight(plan).to_dict(),
        },
    )
    panel._active_session = session
    generation_calls = []
    monkeypatch.setattr(
        panel,
        "_start_content_generation",
        lambda current, active_plan, **kwargs: generation_calls.append(
            (current, active_plan, kwargs)
        ),
    )
    try:
        assert panel._send_message("修改第3题并优化题量") is True
        assert gateway.requests == []
        assert len(generation_calls) == 1
        persisted = panel._coordinator.load_session(session.session_id)
        assert persisted.document_job["status"] == "plan_ready"
        assert persisted.document_job["invalidated_reason"] == (
            "assistant_content_revision_requested"
        )
        assert "preflight" not in persisted.document_job
        revised = DocumentPlan.from_dict(persisted.active_plan)
        assert revised.revision == plan.revision + 1
        assert revised.fingerprint != plan.fingerprint
    finally:
        panel.shutdown_active_execution(1000)
        panel.close()


def test_typed_draft_handoff_completes_preflight_and_execution_loop(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel, gateway, plan = _panel_with_bound_draft(tmp_path)
    output = tmp_path / "outputs" / "小学数学试卷_学生卷.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    generated_document = Document()
    generated_document.add_paragraph("模拟生成的最终 DOCX 文档")
    generated_document.save(output)

    def build_ready_preflight(
        current_plan,
        *,
        material_snapshot=None,
    ):
        del material_snapshot
        return _preflight(current_plan)

    def execute_current_plan(**kwargs):
        assert kwargs["plan"].fingerprint == plan.fingerprint
        assert kwargs["preflight"].plan_fingerprint == plan.fingerprint
        assert kwargs["approval"].authorizes(kwargs["preflight"])
        return {
            "status": "success",
            "output_path": str(output),
            "output_paths": {"student": str(output)},
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "exam_delivery_runtime": {
                "quality_status": "quality_ok",
                "rendered_versions": [],
            },
        }

    monkeypatch.setattr(
        panel._document_jobs,
        "preflight",
        build_ready_preflight,
    )
    monkeypatch.setattr(
        panel._document_jobs,
        "execute",
        execute_current_plan,
    )
    try:
        assert panel._send_message("给我最终docx文档") is True
        _wait_for_local_worker(qapp, panel, "_preflight_worker")
        assert panel._active_session.document_job["status"] == (
            "needs_execution_approval"
        )

        assert panel._send_message("确认并生成 Word") is True
        _wait_for_local_worker(qapp, panel, "_execution_worker")

        persisted = panel._coordinator.load_session(
            panel._active_session.session_id
        )
        assert gateway.requests == []
        assert persisted.document_job["status"] == "success"
        assert persisted.document_job["primary_output_path"] == str(output)
        execution_blocks = [
            block
            for message in persisted.messages
            for block in message.blocks
            if block.data.get("interaction_type") == "artifact"
            or block.data.get("progress_kind") == "execution"
        ]
        assert execution_blocks
        assert all(block.text == "" for block in execution_blocks)
        assert "模拟生成的最终 DOCX 文档" in "\n".join(
            paragraph.text for paragraph in Document(output).paragraphs
        )
        assert DocumentPlan.from_dict(
            persisted.active_plan
        ).fingerprint == plan.fingerprint
        assert [
            message.visible_text()
            for message in persisted.messages
            if message.role == ROLE_USER
        ] == [
            "给我最终docx文档",
            "确认并生成 Word",
        ]
    finally:
        panel.shutdown_active_execution(1000)
        panel.close()


def _panel_with_bound_draft(
    tmp_path: Path,
) -> tuple[AssistantPanel, _CaptureGateway, DocumentPlan]:
    gateway = _CaptureGateway()
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "sessions")
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    plan = _bound_exam_plan(tmp_path)
    session = coordinator.create_session()
    session = coordinator.update_state(
        session,
        active_plan=plan.to_dict(),
        document_job=_draft_job(plan),
        turn_status="completed",
    )
    panel._active_session = session
    return panel, gateway, plan


def test_stale_content_generation_cannot_bind_to_a_newer_plan(
    qapp,
    tmp_path,
):
    panel, _gateway, original_plan = _panel_with_bound_draft(tmp_path)
    newer_plan = replace(original_plan, revision=original_plan.revision + 1)
    session = panel._coordinator.update_state(
        panel._active_session,
        active_plan=newer_plan.to_dict(),
        document_job={
            **_draft_job(original_plan),
            "status": "content_generation_running",
        },
    )
    panel._active_session = session

    class _FinishedWorker:
        request = ContentGenerationRequest(
            session_id=session.session_id,
            turn_id=original_plan.created_by_turn_id,
            prompt=original_plan.intent,
            provider_id="mock-default",
            model_id="form-assistant-mock",
            plan_id=original_plan.plan_id,
            plan_revision=original_plan.revision,
            plan_fingerprint=original_plan.fingerprint,
            artifact_kind=ARTIFACT_KIND_EXAM,
        )

        def deleteLater(self):
            return None

    panel._content_worker = _FinishedWorker()
    stale_output = tmp_path / "stale.exam.md"
    stale_output.write_text("# stale", encoding="utf-8")
    draft = GeneratedExamDraft(
        draft_id="stale-draft",
        session_id=session.session_id,
        markdown_path=str(stale_output),
        source_digest="sha256:stale",
        artifact_id="artifact-stale",
        schema_id="exam_items_v1",
        validation_summary={},
    )
    try:
        panel._on_content_generation_finished(draft)
        qapp.processEvents()

        restored = panel._coordinator.load_session(session.session_id)
        assert DocumentPlan.from_dict(restored.active_plan) == newer_plan
        assert restored.document_job["status"] == "failed"
        assert restored.document_job["error_text"] == (
            "assistant_content_generation_plan_stale"
        )
        assert restored.messages[-1].blocks[0].data["title"] == (
            "已丢弃过期内容草稿"
        )
    finally:
        panel.close()


def test_published_exam_draft_uses_a_readable_title_and_no_generic_body(
    qapp,
    tmp_path,
):
    panel, _gateway, plan = _panel_with_bound_draft(tmp_path)
    draft_id = "3fc8e0ed872e4633a3f156afced8c7ba"
    draft_path = tmp_path / f"{draft_id}.exam.md"
    draft_path.write_text("# 小学数学试卷", encoding="utf-8")
    draft = GeneratedExamDraft(
        draft_id=draft_id,
        session_id=panel._active_session.session_id,
        markdown_path=str(draft_path),
        source_digest="sha256:current",
        artifact_id=f"exam-{draft_id}",
        schema_id="exam_items_v1",
        validation_summary={},
    )
    try:
        panel._publish_generated_draft(panel._active_session, plan, draft)
        qapp.processEvents()

        persisted = panel._coordinator.load_session(
            panel._active_session.session_id
        )
        block = persisted.messages[-1].blocks[0]
        reference_title = block.data["reference"]["title"]
        assert block.text == ""
        assert reference_title.endswith("-内容草稿.md")
        assert draft_id not in reference_title
    finally:
        panel.close()


def test_exam_plan_edit_creates_a_new_revision_and_invalidates_old_artifacts(
    qapp,
    tmp_path,
):
    panel, _gateway, original_plan = _panel_with_bound_draft(tmp_path)
    old_source = Path(original_plan.production_input_artifact.path)
    old_job = {
        **panel._active_session.document_job,
        "status": "needs_execution_approval",
        "preflight": _preflight(original_plan).to_dict(),
        "approval_id": "approval-old",
    }
    panel._active_session = panel._coordinator.update_state(
        panel._active_session,
        document_job={
            **panel._active_session.document_job,
            "status": "preflight_running",
        },
    )
    panel._active_session = panel._coordinator.update_state(
        panel._active_session,
        document_job=old_job,
    )
    values = ExamPlanEditValues(
        school_stage="小学",
        grade="六年级",
        subject="英语",
        exam_period="期末考试",
        question_count=24,
        duration_minutes=90,
        total_score=100,
        textbook_edition="人教版",
        semester="上册",
        scope_hint="Unit 1-4",
        include_student=True,
        include_answer=False,
        output_root=str(tmp_path / "revised-output"),
        master_id="default_exam",
    )

    try:
        panel._apply_exam_plan_edit(
            values,
            expected_plan_id=original_plan.plan_id,
            expected_revision=original_plan.revision,
        )
        qapp.processEvents()

        persisted = panel._coordinator.load_session(
            panel._active_session.session_id
        )
        revised = DocumentPlan.from_dict(persisted.active_plan)
        revised_job = persisted.document_job

        assert revised.plan_id == original_plan.plan_id
        assert revised.revision == original_plan.revision + 1
        assert revised.production_input_artifact is None
        assert revised.output_policy.output_root == str(
            (tmp_path / "revised-output").resolve()
        )
        assert "小学六年级英语期末考试试卷" in revised.intent
        assert revised.delivery_contract.required_artifact_keys == ("student",)
        assert revised.scene_ref["master_id"] == "default_exam"
        assert revised.scene_ref["master_label"] == "A4 标准卷面"
        assert revised.template_ref["id"] == "default"
        assert revised_job["status"] == "plan_ready"
        assert revised_job["plan_revision"] == revised.revision
        assert revised_job["invalidated_plan_revision"] == original_plan.revision
        assert revised_job["material_snapshot"] == old_job["material_snapshot"]
        assert "preflight" not in revised_job
        assert "approval_id" not in revised_job
        assert "generated_content_preview_path" not in revised_job
        assert old_source.is_file()
        assert persisted.messages[-1].blocks[0].data["interaction_type"] == "plan"
    finally:
        panel.close()


def test_exam_plan_edit_drawer_opens_even_when_the_original_grade_is_missing(
    qapp,
    tmp_path,
):
    panel, _gateway, plan = _panel_with_bound_draft(tmp_path)
    try:
        panel._open_exam_plan_editor(
            {"plan_id": plan.plan_id, "revision": plan.revision}
        )
        qapp.processEvents()

        drawer = panel._exam_plan_editor_drawer
        editor = drawer._body_layout.itemAt(0).widget()
        assert drawer.isVisible()
        assert isinstance(editor, ExamPlanEditor)
        assert editor.grade.text() == ""
    finally:
        panel.close()


def test_exam_plan_edit_refuses_an_unverifiable_material_snapshot(qapp, tmp_path):
    panel, _gateway, plan = _panel_with_bound_draft(tmp_path)
    panel._active_session = panel._coordinator.update_state(
        panel._active_session,
        document_job={
            **panel._active_session.document_job,
            "material_snapshot": {},
        },
    )
    values = ExamPlanEditValues(
        school_stage="小学",
        grade="六年级",
        subject="数学",
        exam_period="单元测试",
        question_count=18,
        duration_minutes=60,
        total_score=100,
    )

    try:
        panel._apply_exam_plan_edit(
            values,
            expected_plan_id=plan.plan_id,
            expected_revision=plan.revision,
        )
        qapp.processEvents()

        persisted = panel._coordinator.load_session(
            panel._active_session.session_id
        )
        assert DocumentPlan.from_dict(persisted.active_plan) == plan
        assert persisted.messages[-1].blocks[0].data["reason"] == (
            "assistant_material_snapshot_unavailable"
        )
    finally:
        panel.close()


def test_exam_plan_edit_refuses_a_missing_master_without_replacing_the_plan(
    qapp,
    tmp_path,
):
    panel, _gateway, plan = _panel_with_bound_draft(tmp_path)
    values = ExamPlanEditValues(
        school_stage="小学",
        grade="六年级",
        subject="数学",
        exam_period="期中考试",
        question_count=24,
        duration_minutes=90,
        total_score=100,
        master_id="missing_exam_master",
    )

    try:
        panel._apply_exam_plan_edit(
            values,
            expected_plan_id=plan.plan_id,
            expected_revision=plan.revision,
        )
        qapp.processEvents()

        persisted = panel._coordinator.load_session(
            panel._active_session.session_id
        )
        assert DocumentPlan.from_dict(persisted.active_plan) == plan
        assert persisted.messages[-1].blocks[0].data["reason"] == (
            "assistant_exam_master_unavailable"
        )
    finally:
        panel.close()


def _bound_exam_plan(tmp_path: Path) -> DocumentPlan:
    source = tmp_path / "current.exam.md"
    source.write_text(
        "# 小学数学试卷\n\n## 一、选择题\n\n1. 1 + 1 = ?\n",
        encoding="utf-8",
    )
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    material = MaterialExecutionEnvelope.capture(None)
    return DocumentPlan(
        plan_id="plan-current-draft",
        revision=2,
        intent="生成一份小学数学试卷",
        created_by_turn_id="turn-current-draft",
        input_document_ref={
            "path": str(source),
            "name": source.name,
            "source": "assistant_generated_draft",
        },
        work_mode_id="exam",
        scene_ref={
            "id": "exam_education",
            "generation_mode": "generated_draft",
        },
        template_ref={"id": "default"},
        material_snapshot_ref=material.reference(),
        output_policy=OutputPolicy(
            output_root=str(tmp_path / "outputs"),
        ),
        operation="author",
        source_artifacts=(
            SourceArtifactRef(
                artifact_id="artifact-current-draft",
                role=SOURCE_ROLE_STRUCTURED_SOURCE,
                media_type="text/markdown",
                path=str(source),
                name=source.name,
                schema_id="exam_items_v1",
                digest=f"sha256:{digest}",
                source_kind="assistant_generated",
            ),
        ),
        generation_contract=GenerationContract(
            required=True,
            artifact_kind=ARTIFACT_KIND_EXAM,
            prompt_profile_id="exam_content_generation_v1",
            validator_id="exam_items_v1",
        ),
        production_contract=ProductionContract(
            input_role=SOURCE_ROLE_STRUCTURED_SOURCE,
            artifact_kind=ARTIFACT_KIND_EXAM,
            validator_id="exam_items_v1",
            terminal_assembler="exam",
            accepted_suffixes=(".md",),
        ),
        delivery_contract=DeliveryContract(
            required_artifact_keys=("student", "answer_key"),
        ),
    )


def _draft_job(
    plan: DocumentPlan,
) -> dict[str, object]:
    source = Path(plan.production_input_artifact.path)
    return {
        "job_id": "job-current-draft",
        "status": "content_draft_ready",
        "plan_id": plan.plan_id,
        "plan_revision": plan.revision,
        "material_snapshot": MaterialExecutionEnvelope.capture(None).to_dict(),
        "generated_content_preview_path": str(source),
        "generated_content_document_path": str(source),
        "generated_content_production_input_path": str(source),
        "generated_content_artifact_kind": ARTIFACT_KIND_EXAM,
        "generated_content_schema_id": "exam_items_v1",
    }


def _preflight(
    plan: DocumentPlan,
    *,
    plan_fingerprint: str | None = None,
) -> PreflightReceipt:
    return PreflightReceipt(
        preflight_id="preflight-current",
        plan_id=plan.plan_id,
        plan_revision=plan.revision,
        plan_fingerprint=plan_fingerprint or plan.fingerprint,
        input_hash="a" * 64,
        material_snapshot_digest=MaterialExecutionEnvelope.capture(None).digest,
        output_root=plan.output_policy.output_root,
        ready=True,
    )


def _wait_for_turn(qapp, panel: AssistantPanel) -> None:
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline and panel._turn_worker is not None:
        qapp.processEvents()
        time.sleep(0.01)
    assert panel._turn_worker is None


def _wait_for_local_worker(
    qapp,
    panel: AssistantPanel,
    attribute: str,
) -> None:
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline and getattr(panel, attribute) is not None:
        qapp.processEvents()
        time.sleep(0.01)
    assert getattr(panel, attribute) is None
