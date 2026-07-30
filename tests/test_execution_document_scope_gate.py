from __future__ import annotations

from pathlib import Path

from docx import Document

from src.config.builtin_scenes import list_builtin_scene_resources
from src.config.document_scope import DOCUMENT_SCOPE_MODES
from src.config.library import load_scene_from_library
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.services.document_structure_evidence import (
    build_document_structure_evidence,
)
from src.services.production_runtime.execution_runtime import WorkbenchProductionRunner
from src.ui.panels.workbench.execution_session_controller import (
    WorkbenchExecutionSessionController,
)


def _whitespace_scene(mode: str) -> SceneWorkspace:
    scene = SceneWorkspace(mode_id="custom")
    scene.document_scope.mode = mode
    scene.document_scope.__post_init__()
    scene.module_switches = {name: False for name in scene.module_switches}
    scene.module_switches["whitespace_normalize"] = True
    scene.compliance_profile.object_preflight.enabled = False
    scene.default_delivery_preset().artifacts.report_json = False
    scene.default_delivery_preset().artifacts.report_markdown = False
    return scene


def test_workbench_rejects_invalid_scope_before_mutating_input(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Must   remain unchanged")
    document.save(source)
    before = source.read_bytes()
    output_dir = tmp_path / "output"

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_whitespace_scene("future"),
        output_dir=output_dir,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert "document_scope_mode_invalid:future" in payload["error_text"]
    assert source.read_bytes() == before
    assert not output_dir.exists()


def test_session_controller_rejects_invalid_scope_before_selecting_input() -> None:
    path_requests: list[str] = []
    controller = WorkbenchExecutionSessionController(
        resolve_document_path=lambda: path_requests.append("requested") or None,
    )

    result = controller.build_worker(
        template=TemplateConfig(),
        scene=_whitespace_scene("future"),
    )

    assert result.worker is None
    assert result.cancelled is False
    assert "document_scope_mode_invalid:future" in result.error_text
    assert path_requests == []


def test_body_scope_uses_current_document_evidence_to_limit_writes(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("第一章 正文", level=1)
    document.add_paragraph("Inside   body")
    document.add_heading("参考文献", level=1)
    document.add_paragraph("[1] Outside   references")
    document.save(source)
    evidence = build_document_structure_evidence(source)

    assert evidence.ready is True
    assert [(region.role_id, region.start_anchor.source_index) for region in evidence.regions] == [
        ("body", 0),
        ("references", 2),
    ]

    payload = WorkbenchProductionRunner(
        doc_path=str(source),
        template=TemplateConfig(),
        scene=_whitespace_scene("body"),
        output_dir=tmp_path / "output",
        document_structure_evidence=evidence,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    output = Document(payload["output_paths"]["final"])
    assert output.paragraphs[1].text == "Inside body"
    assert output.paragraphs[3].text == "[1] Outside   references"


def test_builtin_scenes_use_only_the_new_document_scope_contract() -> None:
    for mode_id, scene_id, _path in list_builtin_scene_resources():
        scene = load_scene_from_library(scene_id, mode_id=mode_id)
        assert scene.document_scope.mode in DOCUMENT_SCOPE_MODES
        assert not hasattr(scene, "format_scope")
        assert not hasattr(scene, "application_boundary")
        assert not hasattr(scene, "section_styles")
