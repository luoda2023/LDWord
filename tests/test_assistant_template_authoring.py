from __future__ import annotations

import json
import re
import time
from pathlib import Path

from docx import Document

from src.assistant.application.session_coordinator import AssistantSessionCoordinator
from src.assistant.contracts.runtime import (
    TURN_FAILED,
    AssistantTurnRequest,
)
from src.assistant.domain.docx_format_evidence import (
    FORMAT_EVIDENCE_DISCLOSURE_FIELD,
    attachment_disclosure_fields,
    bind_attachment_semantic_roles,
    is_template_authoring_request,
)
from src.assistant.runtime.provider_contract import (
    PROVIDER_DONE,
    PROVIDER_START,
    ProviderStreamEvent,
)
from src.assistant.runtime.turn_runner import (
    AssistantTurnRunner,
    build_attachment_context,
)
from src.assistant.storage.session_store import AssistantSessionStore
from src.assistant.ui.assistant_panel import AssistantPanel
from src.assistant.ui.interaction_card import AssistantInteractionCard
from src.assistant.ui.workers import AssistantTurnWorker
from src.config import library
from src.config.library import list_template_entries
from src.config.template_authoring_contract import AUTHORING_RESULT_KIND
from src.qt_api import QDesktopServices
from src.ui.bridge import PanelBridge


_QUERY = "帮我按照要求制作论文的模板"


class _TemplateAuthoringGateway:
    def __init__(self) -> None:
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        bundle = request.messages[-1]["content"]
        match = re.search(
            r"<template_authoring_baseline>\s*(.*?)\s*</template_authoring_baseline>",
            bundle,
            re.DOTALL,
        )
        assert match is not None
        result = json.loads(match.group(1))
        result["kind"] = AUTHORING_RESULT_KIND
        result["template"]["name"] = "上传规范自动生成模板"
        result["template"]["description"] = "来自对话中上传的 DOCX 规范"
        result["observations"] = {
            "master": [],
            "scene": [],
            "material": [],
            "unsupported": [],
        }
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(
            PROVIDER_DONE,
            text=json.dumps(result, ensure_ascii=False),
        )

    def cancel(self) -> bool:
        return True


def _wait_for_turn(qapp, panel: AssistantPanel) -> None:
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline and panel._turn_worker is not None:
        qapp.processEvents()
        time.sleep(0.01)
    assert panel._turn_worker is None


def test_template_authoring_intent_discloses_body_and_format_evidence() -> None:
    refs = ({"path": "C:/sample.docx", "title": "sample.docx"},)

    bound = bind_attachment_semantic_roles(refs, _QUERY)

    assert is_template_authoring_request(_QUERY) is True
    assert bound[0]["template_authoring_source"] is True
    assert attachment_disclosure_fields(bound) == (
        "document_text",
        FORMAT_EVIDENCE_DISCLOSURE_FIELD,
    )
    assert is_template_authoring_request("按这个模板生成文档") is False
    assert is_template_authoring_request("不要生成模板，只分析格式") is False


def test_template_authoring_source_bypasses_generic_attachment_compaction(
    tmp_path: Path,
    monkeypatch,
) -> None:
    import src.assistant.runtime.turn_runner as turn_runner_module

    source = tmp_path / "论文规范.docx"
    document = Document()
    document.add_heading("格式要求", level=1)
    document.add_paragraph("正文使用宋体小四，固定值 20 磅行距。")
    document.save(source)
    refs = bind_attachment_semantic_roles(
        ({"path": str(source), "title": source.name},),
        _QUERY,
    )
    monkeypatch.setattr(
        turn_runner_module,
        "_MAX_ATTACHMENT_CONTEXT_CHARACTERS",
        1,
    )
    monkeypatch.setattr(
        turn_runner_module,
        "_MAX_ATTACHMENT_FORMAT_EVIDENCE_CHARACTERS",
        1,
    )

    _prompt, audit = build_attachment_context(refs)

    coverage = {item["kind"]: item for item in audit["attachment_coverage"]}
    assert coverage["document_text"]["mode"] == "full"
    assert coverage["document_text"]["omitted_characters"] == 0
    assert coverage["document_format_evidence"]["mode"] == "full"
    assert coverage["document_format_evidence"]["omitted_characters"] == 0
    assert audit["attachment_format_evidence_truncated"] is False


def test_uploaded_docx_template_authoring_writes_validated_user_template(
    qapp,
    tmp_path: Path,
    monkeypatch,
) -> None:
    config_root = tmp_path / "config_library"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", config_root / "templates")
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", config_root / "plans")
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda _url: True)

    source = tmp_path / "学校论文格式规范.docx"
    document = Document()
    document.add_heading("排版格式要求", level=1)
    document.add_paragraph("正文使用宋体小四，固定值 20 磅行距。")
    document.save(source)

    gateway = _TemplateAuthoringGateway()
    bridge = PanelBridge()
    bridge.set_current_work_mode("thesis")
    panel = AssistantPanel(
        bridge,
        coordinator=AssistantSessionCoordinator(
            AssistantSessionStore(tmp_path / "sessions")
        ),
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._on_composer_document_selected(str(source))
        assert panel._send_message(_QUERY) is True
        disclosure = panel._active_session.messages[-1].blocks[0].data
        assert disclosure["facts"][1]["value"] == "附件正文、格式结构证据"
        assert any(item["label"] == "写入目标" for item in disclosure["facts"])

        panel._resolve_provider_disclosure(disclosure, approved=True)
        _wait_for_turn(qapp, panel)

        entries = [
            entry
            for entry in list_template_entries(mode_id="thesis")
            if entry.source_type == "user"
        ]
        assert len(entries) == 1
        assert entries[0].name == "上传规范自动生成模板"
        assert entries[0].path.is_file()

        session = panel._active_session
        assert session is not None
        assert session.document_job["template_authoring"] == {
            "status": "completed",
            "mode_id": "thesis",
            "template_id": entries[0].config_id,
            "template_name": "上传规范自动生成模板",
            "template_path": str(entries[0].path),
            "warning_count": 0,
        }
        artifact_blocks = [
            block
            for message in session.messages
            for block in message.blocks
            if block.data.get("title") == "新模板已创建并写入模板库"
        ]
        assert len(artifact_blocks) == 1
        assert artifact_blocks[0].data["reference"]["path"] == str(entries[0].path)
        assert not any(
            message.visible_text().lstrip().startswith("{")
            for message in session.messages
        )
        provider_request = gateway.requests[0]
        assert provider_request.messages[-1]["content"].startswith(
            "<template_authoring_instructions>"
        )
        assert "只生成模板创作结果 JSON" in provider_request.system_prompt
        assert "只读的“格式规范分析”" not in provider_request.system_prompt
        assert not any(
            block.data.get("title") == "文档处理计划"
            for message in session.messages
            for block in message.blocks
        )
        assert session.document_job.get("content_generation_purpose") is None
    finally:
        panel.close()


def test_turn_worker_reports_workspace_failure_instead_of_dying_silently(
    qapp,
) -> None:
    class _ExplodingRunner:
        @staticmethod
        def run(*_args, **_kwargs):
            raise OSError("模板工作台不可读")

    request = AssistantTurnRequest(
        turn_id="turn-template-error",
        session_id="session-template-error",
        user_message=_QUERY,
        provider_profile_id="provider",
        model_id="model",
        template_authoring_mode_id="thesis",
    )
    worker = AssistantTurnWorker(_ExplodingRunner(), request)
    results = []
    worker.finished.connect(results.append)

    worker._run()
    qapp.processEvents()

    assert len(results) == 1
    assert results[0].status == TURN_FAILED
    assert results[0].error == {
        "category": "internal",
        "message": "模板工作台不可读",
    }


def test_template_authoring_rejects_single_non_docx_source(
    qapp,
    tmp_path: Path,
) -> None:
    source = tmp_path / "论文规范.md"
    source.write_text("# 格式要求\n正文使用宋体小四。", encoding="utf-8")
    gateway = _TemplateAuthoringGateway()
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=AssistantSessionCoordinator(
            AssistantSessionStore(tmp_path / "sessions")
        ),
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )
    try:
        panel._on_composer_document_selected(str(source))

        assert panel._send_message(_QUERY) is True
        qapp.processEvents()

        assert gateway.requests == []
        session = panel._active_session
        assert session is not None
        assert session.document_job["reason"] == (
            "template_authoring_requires_single_docx"
        )
        boundary = session.messages[-1].blocks[0].data
        assert boundary["title"] == "请选择一份模板来源 DOCX"
        assert "来源必须是 DOCX" in session.messages[-1].blocks[0].text
    finally:
        panel.close()


def test_host_owned_template_file_card_uses_local_template_action(
    qapp,
    tmp_path: Path,
) -> None:
    target = tmp_path / "template.json"
    target.write_text("{}", encoding="utf-8")
    card = AssistantInteractionCard(
        interaction_type="artifact",
        title="新模板已创建并写入模板库",
        body="",
        payload={
            "reference": {
                "path": str(target),
                "title": target.name,
                "kind": "template_config",
                "owner": "form",
            },
            "actions": [],
        },
    )
    emitted = []
    card.action_requested.connect(
        lambda action, payload: emitted.append((action, payload))
    )
    try:
        qapp.processEvents()
        card._file_cards[0]._open.click()
        assert emitted[0][0] == "open_template_artifact"
    finally:
        card.close()
