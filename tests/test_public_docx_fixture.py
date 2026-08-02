from pathlib import Path

from docx import Document

from scripts.check_public_release import validate_public_docx

ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "tests" / "TEST-1" / "中图分类号.docx"


def test_public_thesis_fixture_is_synthetic_and_reproducible():
    document = Document(FIXTURE)

    assert validate_public_docx(FIXTURE) == []
    assert document.core_properties.author == "Alavette Form"
    assert document.core_properties.last_modified_by == "Alavette Form"
    assert len(document.tables) == 2
    assert "合成学位论文版式测试样本" in {
        paragraph.text for paragraph in document.paragraphs
    }
    assert (ROOT / "scripts" / "generate_public_docx_fixture.py").is_file()


def test_public_docx_validator_rejects_contact_data_and_personal_metadata(tmp_path):
    path = tmp_path / "private.docx"
    document = Document()
    document.add_paragraph("contact: person@example.com")
    document.core_properties.author = "Private Author"
    document.core_properties.last_modified_by = "Private Editor"
    document.save(path)

    assert validate_public_docx(path) == [
        "Email address remains in public DOCX: private.docx",
        "Personal DOCX creator metadata remains in private.docx: 'Private Author'",
        "Personal DOCX lastModifiedBy metadata remains in private.docx: 'Private Editor'",
    ]
