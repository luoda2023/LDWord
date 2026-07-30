from __future__ import annotations

import copy
from dataclasses import replace
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from src.services.execution_session import freezing as execution_session_freezing
from src.services.execution_session import resolution as execution_session_resolution
from src.services.execution_session import (
    build_execution_session_snapshot,
    cleanup_execution_session_resources,
    execution_session_frozen_input_path,
    validate_ready_session_bindings,
)
from src.config.library import load_scene_from_library, load_template_from_library
from src.config.master_library import MasterSpec
from src.config.material_context import MaterialExecutionContext
from src.config.object_preflight_evidence import build_object_preflight_evidence
from src.services.execution_session.support import object_revision
from src.config.scene import DeliveryPreset, SceneWorkspace
from src.config.template import TemplateConfig
from src.shared.engine.official_document_assembly import (
    OfficialDocumentAssemblyResult,
)
from src.ui.adapters.workbench_execution_adapter import WorkbenchExecutionAdapter
import src.services.production_runtime.execution_runtime as execution_runtime
import src.services.production_runtime.delivery_runtime as delivery_runtime
import src.services.production_runtime.execution_preflight as execution_preflight
from src.ui.panels.workbench.execution_worker import ExecutionWorker


def _object_preflight_receipt(scene, input_path) -> dict[str, str]:
    evidence = build_object_preflight_evidence(scene, input_path)
    return {
        "object_preflight_confirmation_revision": evidence.source_revision,
        "object_preflight_confirmation_digest": evidence.evidence_digest,
    }


def _build_resource_graph(tmp_path: Path, monkeypatch):
    plan_path = tmp_path / "official_plan.json"
    primary_path = tmp_path / "primary_template.json"
    target_path = tmp_path / "target_template.json"
    target_b_path = tmp_path / "target_template_b.json"
    master_path = tmp_path / "official_master.docx"
    input_path = tmp_path / "input.docx"
    plan_path.write_text('{"name":"official plan"}', encoding="utf-8")
    primary_path.write_text('{"name":"primary"}', encoding="utf-8")
    target_path.write_text('{"name":"target"}', encoding="utf-8")
    target_b_path.write_text('{"name":"target b"}', encoding="utf-8")
    master_path.write_bytes(b"master-v1")
    Document().save(input_path)

    primary = TemplateConfig(name="primary effective")
    target = TemplateConfig(name="target effective")
    target_b = TemplateConfig(name="target b effective")
    master = MasterSpec(
        master_id="master_a",
        mode_id="official",
        label="Master A",
        family="official",
        source_type="user",
        docx_path=master_path,
        master_version="v1",
    )
    scene = SceneWorkspace(
        name="official effective",
        scene_id="official",
        mode_id="official",
        template_id="primary_template",
        compatible_template_ids=[
            "primary_template",
            "target_template",
            "target_template_b",
        ],
        default_delivery_preset_id="primary",
        delivery_presets=[
            DeliveryPreset(
                preset_id="primary",
                target_template_id="primary_template",
            ),
            DeliveryPreset(
                preset_id="target",
                target_template_id="target_template",
            ),
            DeliveryPreset(
                preset_id="target_b",
                target_template_id="target_template_b",
            ),
        ],
        master_id="master_a",
    )
    entries = {
        "primary_template": SimpleNamespace(
            path=primary_path,
            source_type="user",
        ),
        "target_template": SimpleNamespace(
            path=target_path,
            source_type="user",
        ),
        "target_template_b": SimpleNamespace(
            path=target_b_path,
            source_type="user",
        ),
    }
    targets = {
        "target_template": target,
        "target_template_b": target_b,
    }
    monkeypatch.setattr(
        execution_session_resolution,
        "get_scene_entry",
        lambda resource_id, *, mode_id: SimpleNamespace(
            path=plan_path,
            source_type="user",
        ),
    )
    monkeypatch.setattr(
        execution_session_resolution,
        "get_template_entry",
        lambda resource_id, *, mode_id: entries.get(resource_id),
    )
    monkeypatch.setattr(
        execution_session_resolution,
        "load_template_from_library",
        lambda resource_id, *, mode_id: targets[resource_id],
        raising=False,
    )
    monkeypatch.setattr(
        execution_session_resolution,
        "get_master",
        lambda *_args, **_kwargs: master,
    )
    material_context = MaterialExecutionContext(mode_id="official")
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=primary,
        material_context=material_context,
        input_path=input_path,
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="primary_template",
        document_type_id="notice",
        **_object_preflight_receipt(scene, input_path),
    )
    return SimpleNamespace(
        snapshot=snapshot,
        scene=scene,
        primary=primary,
        material_context=material_context,
        target=target,
        target_b=target_b,
        master=master,
        plan_path=plan_path,
        primary_path=primary_path,
        target_path=target_path,
        target_b_path=target_b_path,
        master_path=master_path,
        input_path=input_path,
    )


def _install_runtime_template_entry(monkeypatch, graph):
    monkeypatch.setattr(
        delivery_runtime,
        "get_template_entry",
        lambda resource_id, *, mode_id: SimpleNamespace(
            path=(
                graph.target_b_path
                if resource_id == "target_template_b"
                else graph.target_path
            ),
            source_type="user",
        ),
        raising=False,
    )


def test_snapshot_lists_master_and_delivery_target_effective_resources(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)

    assert graph.snapshot.ready
    frozen_input = execution_session_frozen_input_path(graph.snapshot)
    assert frozen_input is not None
    assert frozen_input.is_file()
    assert frozen_input.read_bytes() == graph.input_path.read_bytes()
    assert frozen_input.stat().st_mode & 0o222 == 0
    assert frozen_input.name == graph.input_path.name
    assert graph.snapshot.master_ref.source_revision.startswith("sha256:")
    assert graph.snapshot.master_ref.frozen_revision == (
        graph.snapshot.master_ref.source_revision
    )
    assert Path(graph.snapshot.master_ref.frozen_path).is_file()
    assert Path(graph.snapshot.master_ref.frozen_path).read_bytes() == b"master-v1"
    assert Path(graph.snapshot.master_ref.frozen_path).stat().st_mode & 0o222 == 0
    assert Path(graph.snapshot.master_ref.frozen_path).is_relative_to(
        Path(graph.snapshot.output_namespace)
    )
    assert graph.snapshot.master_ref.effective_revision == (
        object_revision(graph.master)
    )
    assert len(graph.snapshot.dependent_resource_refs) == 2
    assert [
        resource_ref.resource_id
        for resource_ref in graph.snapshot.dependent_resource_refs
    ] == ["target_template", "target_template_b"]
    target_ref = graph.snapshot.dependent_resource_refs[0]
    assert target_ref.kind == "delivery_target_template"
    assert target_ref.resource_id == "target_template"
    assert target_ref.source_revision.startswith("sha256:")
    assert target_ref.effective_revision == object_revision(
        graph.target
    )
    assert graph.snapshot.session_overrides_revision == object_revision({})
    assert validate_ready_session_bindings(
        graph.snapshot,
        scene=graph.scene,
        template=graph.primary,
        material_context=graph.material_context,
        output_namespace=graph.snapshot.output_namespace,
    ) == ()


@pytest.mark.parametrize(
    "drift_kind",
    (
        "scene",
        "template",
        "material_context",
        "output_namespace",
        "session_overrides",
    ),
)
def test_runner_blocks_ready_snapshot_bound_to_different_runtime_inputs(
    tmp_path,
    monkeypatch,
    drift_kind,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    scene = copy.deepcopy(graph.scene)
    template = copy.deepcopy(graph.primary)
    material_context = graph.material_context.clone()
    output_dir = Path(graph.snapshot.output_namespace)
    session_overrides: dict[str, object] = {}

    if drift_kind == "scene":
        scene.name = "drifted scene"
    elif drift_kind == "template":
        template.name = "drifted template"
    elif drift_kind == "material_context":
        material_context.entity_data["title"] = "drifted material"
    elif drift_kind == "output_namespace":
        output_dir = tmp_path / "different-output"
    else:
        session_overrides["output.report_json"] = False

    pipeline_calls: list[str] = []

    class _PipelineMustNotRun:
        def __init__(self, **_kwargs):
            pipeline_calls.append("constructed")

    monkeypatch.setattr(execution_runtime, "Pipeline", _PipelineMustNotRun)
    payload = execution_runtime.WorkbenchProductionRunner(
        doc_path=str(graph.input_path),
        template=template,
        scene=scene,
        session_overrides=session_overrides,
        material_context=material_context,
        output_dir=output_dir,
        execution_session=graph.snapshot,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["error_text"] == (
        f"execution_session_binding_drift:{drift_kind}"
    )
    assert pipeline_calls == []


def test_normal_library_primary_template_and_master_validate(tmp_path):
    input_path = tmp_path / "input.docx"
    Document().save(input_path)
    scene = load_scene_from_library("official", mode_id="official")
    template = load_template_from_library("official_gbt", mode_id="official")
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=template,
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=input_path,
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="official_gbt",
        document_type_id="notice",
        **_object_preflight_receipt(scene, input_path),
    )

    selected_template = delivery_runtime.load_delivery_template(
        template,
        scene=scene,
        target_template_id="official_gbt",
        execution_session=snapshot,
    )
    selected_master = execution_runtime._selected_official_master(
        scene,
        execution_session=snapshot,
    )

    assert snapshot.ready
    assert selected_template is template
    assert selected_master.master_id == snapshot.master_ref.resource_id


@pytest.mark.parametrize(
    ("document_type_id", "expected_master_id", "material_document_type"),
    (
        ("letter", "official_gbt_letter", None),
        ("minutes", "official_gbt_minutes", "notice"),
    ),
)
def test_single_official_session_drives_runner_profile_and_master_identity(
    tmp_path,
    document_type_id,
    expected_master_id,
    material_document_type,
):
    source = tmp_path / f"{document_type_id}.docx"
    Document().save(source)
    scene = load_scene_from_library("official", mode_id="official")
    scene.default_delivery_preset().artifacts.review_pdf = False
    template = load_template_from_library("official_gbt", mode_id="official")
    entity_data = {
        "title": f"Frozen {document_type_id} profile",
        "body": "The session-owned profile selects the assembler contract.",
        "organization": "Example organization",
        "document_no": "EX-2026-07",
        "issue_date": "2026-07-15",
    }
    if material_document_type is not None:
        # Deliberately contradictory: the frozen session identity must win.
        entity_data["document_type"] = material_document_type
    material_context = MaterialExecutionContext(
        mode_id="official",
        entity_data=entity_data,
    )
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=template,
        material_context=material_context,
        input_path=source,
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="official_gbt",
        document_type_id=document_type_id,
        **_object_preflight_receipt(scene, source),
    )
    try:
        assert snapshot.ready is True
        assert snapshot.document_type_id == document_type_id
        assert snapshot.master_ref.resource_id == expected_master_id

        payload = execution_runtime.WorkbenchProductionRunner(
            doc_path=str(source),
            template=template,
            scene=scene,
            material_context=material_context,
            output_dir=Path(snapshot.output_namespace),
            execution_session=snapshot,
        ).run(lambda *_args: None, lambda: False)

        assembly = payload["official_document_assembly"]
        assert payload["status"] == "success"
        assert assembly["profile_id"] == document_type_id
        assert assembly["master_id"] == expected_master_id
        assert Path(assembly["docx_path"]).is_file()
    finally:
        cleanup_execution_session_resources(snapshot)


def test_snapshot_blocks_unresolved_delivery_target_without_fallback(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)

    def _missing_target(*_args, **_kwargs):
        raise FileNotFoundError("target missing")

    monkeypatch.setattr(
        execution_session_resolution,
        "load_template_from_library",
        _missing_target,
    )
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=graph.scene,
        template=graph.primary,
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=graph.input_path,
        output_root=tmp_path / "blocked-runs",
        plan_id="official",
        template_id="primary_template",
        document_type_id="notice",
        **_object_preflight_receipt(graph.scene, graph.input_path),
    )

    assert not snapshot.ready
    assert "dependent_template_ref_unresolved:target_template" in snapshot.issues
    assert snapshot.dependent_resource_refs[0].status == "missing"
    assert snapshot.dependent_resource_refs[0].effective_id == ""


def test_delivery_target_programming_error_is_not_downgraded_to_missing(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)

    def _broken_loader(*_args, **_kwargs):
        raise RuntimeError("template registry programming error")

    monkeypatch.setattr(
        execution_session_resolution,
        "load_template_from_library",
        _broken_loader,
    )

    with pytest.raises(RuntimeError, match="template registry programming error"):
        build_execution_session_snapshot(
            mode_id="official",
            scene=graph.scene,
            template=graph.primary,
            material_context=MaterialExecutionContext(mode_id="official"),
            input_path=graph.input_path,
            output_root=tmp_path / "broken-runs",
            plan_id="official",
            template_id="primary_template",
            document_type_id="notice",
            **_object_preflight_receipt(graph.scene, graph.input_path),
        )


def test_official_master_loader_drift_is_not_requeried_after_snapshot(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)

    def _must_not_reload(*_args, **_kwargs):
        raise AssertionError("master registry must not be queried after snapshot")

    monkeypatch.setattr(
        execution_runtime,
        "get_master",
        _must_not_reload,
    )

    selected = execution_runtime._selected_official_master(
        graph.scene,
        execution_session=graph.snapshot,
    )

    assert selected is graph.snapshot.frozen_master
    assert selected.docx_path == Path(graph.snapshot.master_ref.frozen_path)


def test_official_master_source_overwrite_keeps_using_frozen_bytes(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    graph.master_path.write_bytes(b"master-v2")

    selected = execution_runtime._selected_official_master(
        graph.scene,
        execution_session=graph.snapshot,
    )

    assert graph.master_path.read_bytes() == b"master-v2"
    assert selected.docx_path.read_bytes() == b"master-v1"


def test_official_master_source_deletion_keeps_using_frozen_bytes(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    graph.master_path.unlink()

    selected = execution_runtime._selected_official_master(
        graph.scene,
        execution_session=graph.snapshot,
    )

    assert not graph.master_path.exists()
    assert selected.docx_path.read_bytes() == b"master-v1"


def test_input_source_overwrite_keeps_runner_on_session_owned_bytes(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    original_bytes = graph.input_path.read_bytes()
    graph.input_path.write_bytes(b"replacement-input-v2")
    seen_input_bytes: list[bytes] = []

    runner = execution_runtime.WorkbenchProductionRunner(
        doc_path=str(graph.input_path),
        template=graph.primary,
        scene=graph.scene,
        material_context=graph.material_context,
        output_dir=graph.snapshot.output_namespace,
        execution_session=graph.snapshot,
    )

    def _run_document_input(**kwargs):
        seen_input_bytes.append(kwargs["input_path"].read_bytes())
        return {"status": "success"}

    monkeypatch.setattr(runner, "_run_document_input", _run_document_input)

    payload = runner.run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert graph.input_path.read_bytes() == b"replacement-input-v2"
    assert seen_input_bytes == [original_bytes]


def test_frozen_input_tamper_blocks_before_runtime_preflight(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    frozen_input = Path(graph.snapshot.input_ref.frozen_path)
    frozen_input.chmod(0o666)
    frozen_input.write_bytes(b"tampered-input")
    runner = execution_runtime.WorkbenchProductionRunner(
        doc_path=str(graph.input_path),
        template=graph.primary,
        scene=graph.scene,
        material_context=graph.material_context,
        output_dir=graph.snapshot.output_namespace,
        execution_session=graph.snapshot,
    )

    with pytest.raises(
        RuntimeError,
        match="execution_resource_drift:input_document:input.docx:frozen_revision",
    ):
        runner.run(lambda *_args: None, lambda: False)


def test_frozen_official_master_tamper_blocks_before_pipeline_construction(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    frozen_path = Path(graph.snapshot.master_ref.frozen_path)
    frozen_path.chmod(0o666)
    frozen_path.write_bytes(b"tampered")
    pipeline_calls = []

    class _PipelineMustNotRun:
        def __init__(self, **_kwargs):
            pipeline_calls.append("constructed")

    monkeypatch.setattr(execution_runtime, "Pipeline", _PipelineMustNotRun)
    runner = execution_runtime.WorkbenchProductionRunner(
        doc_path=str(graph.input_path),
        template=graph.primary,
        scene=graph.scene,
        material_context=graph.material_context,
        output_dir=graph.snapshot.output_namespace,
        execution_session=graph.snapshot,
    )

    with pytest.raises(
        RuntimeError,
        match="execution_resource_drift:master:master_a:frozen_revision",
    ):
        runner.run(lambda *_args: None, lambda: False)

    assert pipeline_calls == []
    assert not list((tmp_path / "out").glob("*.docx"))


def test_delivery_target_loader_drift_fails_closed(tmp_path, monkeypatch):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    _install_runtime_template_entry(monkeypatch, graph)
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(name="drifted target"),
    )

    with pytest.raises(
        RuntimeError,
        match=(
            "execution_resource_drift:delivery_target_template:"
            "target_template:effective_revision"
        ),
    ):
        delivery_runtime.load_delivery_template(
            graph.primary,
            scene=graph.scene,
            target_template_id="target_template",
            execution_session=graph.snapshot,
        )


def test_delivery_template_uses_session_mode_before_custom_scene_mode(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    runtime_scene = copy.deepcopy(graph.scene)
    runtime_scene.mode_id = "custom"
    observed_modes: list[tuple[str, str]] = []

    def load_target(resource_id, *, mode_id):
        observed_modes.append(("load", mode_id))
        assert resource_id == "target_template"
        return graph.target

    def target_entry(resource_id, *, mode_id):
        observed_modes.append(("entry", mode_id))
        assert resource_id == "target_template"
        return SimpleNamespace(
            path=graph.target_path,
            source_type="user",
        )

    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        load_target,
    )
    monkeypatch.setattr(
        delivery_runtime,
        "get_template_entry",
        target_entry,
    )

    selected = delivery_runtime.load_delivery_template(
        graph.primary,
        scene=runtime_scene,
        target_template_id="target_template",
        execution_session=graph.snapshot,
    )

    assert selected is graph.target
    assert observed_modes == [
        ("load", "official"),
        ("entry", "official"),
    ]


def test_delivery_target_disk_drift_fails_closed(tmp_path, monkeypatch):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    _install_runtime_template_entry(monkeypatch, graph)
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: graph.target,
    )
    graph.target_path.write_text('{"name":"target v2"}', encoding="utf-8")

    with pytest.raises(
        RuntimeError,
        match=(
            "execution_resource_drift:delivery_target_template:"
            "target_template:source_revision"
        ),
    ):
        delivery_runtime.load_delivery_template(
            graph.primary,
            scene=graph.scene,
            target_template_id="target_template",
            execution_session=graph.snapshot,
        )


def test_delivery_group_drift_blocks_before_any_pipeline_output(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    _install_runtime_template_entry(monkeypatch, graph)
    monkeypatch.setattr(
        execution_runtime,
        "get_master",
        lambda *_args, **_kwargs: graph.master,
    )
    monkeypatch.setattr(
        delivery_runtime,
        "load_template_from_library",
        lambda *_args, **_kwargs: TemplateConfig(name="drifted target"),
    )
    pipeline_calls = []

    class _PipelineMustNotRun:
        def __init__(self, **_kwargs):
            pipeline_calls.append("constructed")

        def execute(self, _path):
            pipeline_calls.append("executed")
            raise AssertionError("pipeline must not run after resource drift")

    monkeypatch.setattr(execution_runtime, "Pipeline", _PipelineMustNotRun)
    payload = execution_runtime.WorkbenchProductionRunner(
        doc_path=str(graph.input_path),
        template=graph.primary,
        scene=graph.scene,
        material_context=graph.material_context,
        output_dir=graph.snapshot.output_namespace,
        execution_session=graph.snapshot,
    ).run(lambda *_args: None, lambda: False)

    assert payload["status"] == "failed"
    assert payload["output_paths"] == {}
    assert "execution_resource_drift" in payload["error_text"]
    assert pipeline_calls == []


def test_official_batch_reuses_one_validated_master_for_every_item(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    batch_items = [
        SimpleNamespace(
            profile_id=f"entity-{index}",
            profile_name=f"Entity {index}",
            output_dir=(
                Path(graph.snapshot.output_namespace) / f"entity-{index}"
            ),
            context=MaterialExecutionContext(
                mode_id="official",
                entity_data={"document_type": "notice"},
            ),
        )
        for index in (1, 2)
    ]
    monkeypatch.setattr(
        execution_runtime,
        "build_material_batch_items",
        lambda *_args, **_kwargs: batch_items,
    )
    monkeypatch.setattr(
        execution_runtime,
        "check_material_batch_preflight",
        lambda *_args, **_kwargs: SimpleNamespace(ok=True, issues=()),
    )
    master_lookups = []

    def _master_lookup(*_args, **_kwargs):
        master_lookups.append("lookup")
        if len(master_lookups) == 1:
            return graph.master
        return replace(graph.master, label="Drifted after execution began")

    monkeypatch.setattr(execution_runtime, "get_master", _master_lookup)
    consumed_masters = []

    def _assemble(profile_id, _entity_data, *, output_dir, master, **_kwargs):
        consumed_masters.append(master)
        output_path = Path(output_dir) / f"{profile_id}.docx"
        return OfficialDocumentAssemblyResult(
            status="ok",
            profile_id=profile_id,
            master_id=master.master_id,
            docx_path=output_path,
        )

    monkeypatch.setattr(
        execution_runtime,
        "assemble_official_document_docx",
        _assemble,
    )
    for diagnostic_builder in (
        "material_requirement_diagnostics",
        "image_anchor_diagnostics",
        "question_figure_file_diagnostics",
        "missing_asset_rule_diagnostics",
    ):
        monkeypatch.setattr(
            execution_preflight,
            diagnostic_builder,
            lambda *_args, **_kwargs: [],
        )
    runner = execution_runtime.WorkbenchBatchProductionRunner(
        doc_path=str(graph.input_path),
        template=graph.primary,
        scene=graph.scene,
        archive=SimpleNamespace(),
        source_kind="official_document_table",
        source_path=str(graph.input_path),
        base_output_dir=graph.snapshot.output_namespace,
        base_context=graph.material_context,
        execution_session=graph.snapshot,
    )
    monkeypatch.setattr(runner, "_attach_official_batch_reports", lambda *_args: None)

    payload = runner.run(lambda *_args: None, lambda: False)

    assert payload["status"] == "success"
    assert master_lookups == []
    assert consumed_masters == [
        graph.snapshot.frozen_master,
        graph.snapshot.frozen_master,
    ]
    assert consumed_masters[0] is consumed_masters[1]
    assert consumed_masters[0].docx_path == Path(graph.snapshot.master_ref.frozen_path)


def test_frozen_master_lives_with_receipt_and_cleanup_is_namespace_scoped(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    frozen_path = Path(graph.snapshot.master_ref.frozen_path)
    frozen_input_path = Path(graph.snapshot.input_ref.frozen_path)

    class Runner:
        execution_session_snapshot = graph.snapshot

    payload = ExecutionWorker(Runner())._normalize_result(
        {"status": "success"},
        "success",
    )
    receipt_path = Path(payload["execution_session_path"])
    unrelated_path = Path(graph.snapshot.output_namespace) / "keep.txt"
    unrelated_path.write_text("keep", encoding="utf-8")

    assert receipt_path.is_file()
    assert frozen_path.is_file()
    assert frozen_input_path.is_file()
    assert cleanup_execution_session_resources(graph.snapshot) == ()
    assert not frozen_path.exists()
    assert not frozen_input_path.exists()
    assert receipt_path.is_file()
    assert unrelated_path.read_text(encoding="utf-8") == "keep"


def test_worker_cleans_session_owned_input_and_master_after_receipt(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    frozen_paths = {
        Path(graph.snapshot.input_ref.frozen_path),
        Path(graph.snapshot.master_ref.frozen_path),
    }

    class Runner:
        execution_session_snapshot = graph.snapshot

        @staticmethod
        def run(_progress, _cancelled):
            return {"status": "success"}

    succeeded: list[dict[str, object]] = []
    worker = ExecutionWorker(Runner())
    worker.execution_succeeded.connect(succeeded.append)
    worker.run()

    assert len(succeeded) == 1
    assert all(not path.exists() for path in frozen_paths)
    assert Path(succeeded[0]["execution_session_path"]).is_file()


def test_master_freeze_write_failure_blocks_and_cleans_session_staging(
    tmp_path,
    monkeypatch,
):
    input_path = tmp_path / "input.docx"
    Document().save(input_path)
    scene = load_scene_from_library("official", mode_id="official")
    template = load_template_from_library("official_gbt", mode_id="official")

    def _fail_after_write(path, payload):
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(payload)
        raise OSError("simulated freeze publication failure")

    monkeypatch.setattr(
        execution_session_freezing,
        "atomic_write_bytes",
        _fail_after_write,
    )
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=template,
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=input_path,
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="official_gbt",
        document_type_id="notice",
        **_object_preflight_receipt(scene, input_path),
    )

    assert not snapshot.ready
    assert snapshot.master_ref.status == "freeze_failed"
    assert snapshot.issues == (
        "master_freeze_failed:official_gbt_standard:OSError",
    )
    assert not (Path(snapshot.output_namespace) / ".execution_resources").exists()


def test_input_freeze_failure_rolls_back_already_frozen_master(
    tmp_path,
    monkeypatch,
):
    input_path = tmp_path / "input.docx"
    Document().save(input_path)
    scene = load_scene_from_library("official", mode_id="official")
    template = load_template_from_library("official_gbt", mode_id="official")

    def _fail_input_freeze(*_args, **_kwargs):
        raise OSError("simulated input freeze failure")

    monkeypatch.setattr(
        execution_session_freezing,
        "_materialize_frozen_input",
        _fail_input_freeze,
    )
    snapshot = build_execution_session_snapshot(
        mode_id="official",
        scene=scene,
        template=template,
        material_context=MaterialExecutionContext(mode_id="official"),
        input_path=input_path,
        output_root=tmp_path / "runs",
        plan_id="official",
        template_id="official_gbt",
        document_type_id="notice",
        **_object_preflight_receipt(scene, input_path),
    )

    assert not snapshot.ready
    assert snapshot.input_ref.status == "freeze_failed"
    assert snapshot.master_ref.frozen_path == ""
    assert snapshot.master_ref.frozen_revision == ""
    assert snapshot.frozen_master is None
    assert snapshot.issues == ("input_freeze_failed:input.docx:OSError",)
    assert not (Path(snapshot.output_namespace) / ".execution_resources").exists()


def test_primary_template_consumes_frozen_object_after_source_drift(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)
    graph.primary_path.write_text('{"name":"primary v2"}', encoding="utf-8")

    selected = delivery_runtime.load_delivery_template(
        graph.primary,
        scene=graph.scene,
        target_template_id="primary_template",
        execution_session=graph.snapshot,
    )

    assert selected is graph.primary


def test_dependent_resource_graph_roundtrips_through_worker_and_ui(
    tmp_path,
    monkeypatch,
):
    graph = _build_resource_graph(tmp_path, monkeypatch)

    class Runner:
        execution_session_snapshot = graph.snapshot

    payload = ExecutionWorker(Runner())._normalize_result(
        {"status": "success"},
        "success",
    )
    saved = json.loads(
        Path(payload["execution_session_path"]).read_text(encoding="utf-8")
    )
    state = WorkbenchExecutionAdapter().build_result_state(
        terminal_payload={
            "status": "success",
            "output_path": "",
            "report_paths": [],
            "failed_count": 0,
            "error_text": "",
            "execution_session": saved,
        }
    )

    assert saved == graph.snapshot.to_dict()
    assert saved["dependent_resource_refs"][0]["resource_id"] == "target_template"
    assert state.execution_session["dependent_resource_refs"] == saved[
        "dependent_resource_refs"
    ]
