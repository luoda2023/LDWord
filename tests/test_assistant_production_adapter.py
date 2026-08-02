from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from src.assistant.adapters import production_adapter as module
from src.assistant.adapters.production_adapter import (
    AUTHORING_ONLY_MATERIAL_WARNING,
    MATERIAL_USAGE_AUTHORING_ONLY,
    AssistantProductionAdapter,
)
from src.assistant.adapters.workspace_state_adapter import WorkspaceSnapshot
from src.assistant.application.plan_builder import FormDocumentPlanBuilder
from src.assistant.contracts.material_snapshot import MaterialExecutionEnvelope
from src.assistant.contracts.document_plan import DocumentPlan, OutputPolicy
from src.assistant.contracts.execution import ExecutionApproval
from src.assistant.contracts.task_plan import (
    SOURCE_ROLE_PRODUCTION_INPUT,
    GenerationContract,
    ProductionContract,
    SourceArtifactRef,
)
from src.application.materials import (
    ExecutionMaterialRecord,
    ExecutionMaterialSnapshot,
)
from src.application.materials import execution as material_execution
from src.domain.materials import MaterialPackageRef


def _plan(input_path: Path, output_root: Path) -> DocumentPlan:
    return DocumentPlan(
        plan_id="assistant-plan-1",
        revision=1,
        intent="统一排版",
        created_by_turn_id="turn-1",
        input_document_ref={"path": str(input_path), "name": input_path.name},
        work_mode_id="custom",
        scene_ref={"id": "custom"},
        template_ref={"id": "default"},
        output_policy=OutputPolicy(output_root=str(output_root)),
    )


def _patch_resources(monkeypatch, tmp_path):
    scene_path = tmp_path / "scene.json"
    template_path = tmp_path / "template.json"
    scene_path.write_text("scene-v1", encoding="utf-8")
    template_path.write_text("template-v1", encoding="utf-8")
    scene_entry = SimpleNamespace(path=scene_path, is_available=True, source_type="builtin")
    template_entry = SimpleNamespace(path=template_path, is_available=True, source_type="builtin")
    monkeypatch.setattr(module, "get_scene_entry", lambda *_args, **_kwargs: scene_entry)
    monkeypatch.setattr(module, "get_template_entry", lambda *_args, **_kwargs: template_entry)
    monkeypatch.setattr(module, "load_scene_from_library", lambda *_args, **_kwargs: object())
    monkeypatch.setattr(module, "load_template_from_library", lambda *_args, **_kwargs: object())
    return scene_path, template_path


def _approval(preflight):
    return ExecutionApproval(
        approval_id="approval-1",
        session_id="session-1",
        plan_id=preflight.plan_id,
        plan_revision=preflight.plan_revision,
        preflight_hash=preflight.evidence_hash,
        input_hash=preflight.input_hash,
        output_root=preflight.output_root,
        overwrite_policy="deny",
        approved_at="2026-07-16T00:00:00+00:00",
    )


def test_exam_master_selection_is_projected_onto_the_runtime_scene(tmp_path):
    input_path = tmp_path / "exam.md"
    input_path.write_text("# 试卷\n\n1. 题目", encoding="utf-8")
    plan = replace(
        _plan(input_path, tmp_path / "out"),
        work_mode_id="exam",
        scene_ref={"id": "exam", "master_id": "school_exam_master"},
        production_contract=ProductionContract(terminal_assembler="exam"),
    )
    scene = SimpleNamespace(
        master_id="default_exam",
        exam_paper=SimpleNamespace(answer_policy="student_plus_answer"),
        delivery_presets=(),
    )

    issue = module._apply_plan_scene_contract(scene, plan)

    assert issue == ""
    assert scene.master_id == "school_exam_master"


def _custom_material_snapshot() -> ExecutionMaterialSnapshot:
    snapshot = ExecutionMaterialSnapshot(
        snapshot_id="",
        run_id="run-cross-mode",
        package_ref=MaterialPackageRef(
            package_id=f"pkg_{'a' * 32}",
            revision=f"sha256:{'b' * 64}",
        ),
        work_mode_id="custom",
        material_contract_id="generic_document_v1",
        recipe_id="document_batch",
        scene_id="custom",
        document_type="",
        package_display_name="用户材料",
        package_field_values={},
        package_field_owners={},
        groups=(),
        records=(
            ExecutionMaterialRecord(
                record_id=f"rec_{'c' * 32}",
                display_name="材料记录",
                group_id="",
                field_values={"subject": "材料中的主题"},
                field_owners={"subject": "record"},
                resources={},
                resource_owners={},
            ),
        ),
    )
    return replace(
        snapshot,
        snapshot_id=material_execution._snapshot_content_id(snapshot),
    )


def _generated_cross_mode_plan(
    input_path: Path,
    output_root: Path,
    material_snapshot: ExecutionMaterialSnapshot,
) -> DocumentPlan:
    base = _plan(input_path, output_root)
    return replace(
        base,
        work_mode_id="exam",
        scene_ref={"id": "custom", "generation_mode": "generated_draft"},
        source_artifacts=(
            SourceArtifactRef(
                artifact_id="artifact-generated-cross-mode",
                role=SOURCE_ROLE_PRODUCTION_INPUT,
                media_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                path=str(input_path),
                name=input_path.name,
                digest=f"sha256:{module.file_sha256(input_path)}",
                source_kind="assistant_generated",
            ),
        ),
        generation_contract=GenerationContract(
            required=True,
            artifact_kind="narrative_document",
            prompt_profile_id="narrative_generation_v1",
            validator_id="",
        ),
        material_snapshot_ref=MaterialExecutionEnvelope.capture(
            material_snapshot
        ).reference(),
    )


def test_preflight_binds_input_plan_and_resource_hashes(tmp_path, monkeypatch):
    input_path = tmp_path / "input.docx"
    Document().save(input_path)
    _patch_resources(monkeypatch, tmp_path)
    plan = _plan(input_path, tmp_path / "out")

    preflight = AssistantProductionAdapter().build_preflight(plan)

    assert preflight.ready
    assert preflight.plan_fingerprint == plan.fingerprint
    assert preflight.input_hash == module.file_sha256(input_path)
    assert set(preflight.resource_fingerprints) == {
        "materials",
        "scene",
        "template",
    }
    assert preflight.resource_fingerprints["materials"] == (
        preflight.material_snapshot_digest
    )


def test_approved_plan_is_the_only_path_to_form_execution(tmp_path, monkeypatch):
    input_path = tmp_path / "input.docx"
    Document().save(input_path)
    _patch_resources(monkeypatch, tmp_path)
    calls = []

    def executor(request, **kwargs):
        calls.append((request, kwargs))
        return {
            "status": "success",
            "output_path": str(tmp_path / "out" / "result.docx"),
            "output_paths": {},
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }

    adapter = AssistantProductionAdapter(executor=executor)
    plan = _plan(input_path, tmp_path / "out")
    preflight = adapter.build_preflight(plan)
    result = adapter.execute_approved_plan(plan, preflight, _approval(preflight))

    assert result["status"] == "success"
    assert result["assistant_input_hash_unchanged"] is True
    assert calls[0][0].input_path == input_path
    assert calls[0][0].plan_id == "custom"


def test_generated_draft_uses_cross_mode_material_for_authoring_only(
    tmp_path,
    monkeypatch,
):
    input_path = tmp_path / "generated.docx"
    document = Document()
    document.add_paragraph("根据材料生成的结构化草稿")
    document.save(input_path)
    _patch_resources(monkeypatch, tmp_path)
    snapshot = _custom_material_snapshot()
    plan = _generated_cross_mode_plan(
        input_path,
        tmp_path / "out",
        snapshot,
    )
    calls = []

    def executor(request, **_kwargs):
        calls.append(request)
        return {
            "status": "success",
            "output_path": str(tmp_path / "out" / "result.docx"),
            "output_paths": {},
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }

    adapter = AssistantProductionAdapter(executor=executor)
    preflight = adapter.build_preflight(
        plan,
        material_snapshot=snapshot,
    )
    result = adapter.execute_approved_plan(
        plan,
        preflight,
        _approval(preflight),
        material_snapshot=snapshot,
    )

    assert preflight.ready, preflight.issues
    assert AUTHORING_ONLY_MATERIAL_WARNING in preflight.warnings
    assert (
        preflight.resource_fingerprints["material_usage"]
        == MATERIAL_USAGE_AUTHORING_ONLY
    )
    assert result["status"] == "success"
    assert calls and calls[0].material_snapshot is None


def test_existing_document_rejects_cross_mode_execution_material(
    tmp_path,
    monkeypatch,
):
    input_path = tmp_path / "existing.docx"
    Document().save(input_path)
    _patch_resources(monkeypatch, tmp_path)
    snapshot = _custom_material_snapshot()
    plan = replace(
        _plan(input_path, tmp_path / "out"),
        work_mode_id="exam",
        material_snapshot_ref=MaterialExecutionEnvelope.capture(
            snapshot
        ).reference(),
    )

    preflight = AssistantProductionAdapter().build_preflight(
        plan,
        material_snapshot=snapshot,
    )

    assert not preflight.ready
    assert "execution_material_mode_mismatch" in preflight.issues


def test_exam_execution_receives_authoritative_scale_and_visual_gate(
    tmp_path,
    monkeypatch,
):
    input_path = tmp_path / "input.docx"
    Document().save(input_path)
    _patch_resources(monkeypatch, tmp_path)
    calls = []

    def executor(request, **kwargs):
        calls.append((request, kwargs))
        return {
            "status": "success",
            "output_path": str(tmp_path / "out" / "result.docx"),
            "output_paths": {},
            "report_paths": [],
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }

    plan = _plan(input_path, tmp_path / "out")
    plan = replace(
        plan,
        scene_ref={"id": "my_school_exam", "scale_profile_id": "term"},
        production_contract=replace(
            plan.production_contract,
            terminal_assembler="exam",
        ),
    )
    adapter = AssistantProductionAdapter(executor=executor)
    preflight = adapter.build_preflight(plan)

    result = adapter.execute_approved_plan(
        plan,
        preflight,
        _approval(preflight),
    )

    assert result["status"] == "success"
    assert calls
    request = calls[0][0]
    assert request.exam_scale_profile_id == "term"
    assert request.exam_visual_quality_required is True


def test_execution_rejects_changed_input_or_resource(tmp_path, monkeypatch):
    input_path = tmp_path / "input.docx"
    Document().save(input_path)
    scene_path, _template_path = _patch_resources(monkeypatch, tmp_path)
    calls = []
    adapter = AssistantProductionAdapter(executor=lambda *_args, **_kwargs: calls.append(1) or {})
    plan = _plan(input_path, tmp_path / "out")
    preflight = adapter.build_preflight(plan)
    approval = _approval(preflight)

    input_path.write_bytes(input_path.read_bytes() + b"changed")
    assert adapter.execute_approved_plan(plan, preflight, approval)["error_text"] == "input_changed_after_preflight"
    assert calls == []

    Document().save(input_path)
    fresh = adapter.build_preflight(plan)
    fresh_approval = _approval(fresh)
    scene_path.write_text("scene-v2", encoding="utf-8")
    assert adapter.execute_approved_plan(plan, fresh, fresh_approval)["error_text"] == "approved_resource_changed"
    assert calls == []


def test_qualification_archive_plan_projects_family_and_package_delivery(tmp_path):
    input_path = tmp_path / "qualification-index.docx"
    Document().save(input_path)
    workspace = WorkspaceSnapshot(
        mode_id="custom",
        mode_label="通用版",
        scene_id="custom",
        scene_source_type="builtin",
        template_id="default",
        template_source_type="builtin",
        input_path=str(input_path),
        input_name=input_path.name,
        input_exists=True,
        material_summary={
            "package_id": "qualification-main",
            "profile_id": "company-a",
            "field_count": 2,
            "asset_count": 2,
        },
    )
    plan = FormDocumentPlanBuilder().build(
        query="整理投标资质证书材料",
        workspace=workspace,
        turn_id="turn-qualification",
    )
    plan = DocumentPlan.from_dict(
        {
            **plan.to_dict(),
            "output_policy": {
                "output_root": str(tmp_path / "out"),
                "filename_suffix": "_archive",
                "overwrite": False,
            },
        }
    )
    captured = []

    def executor(request, **_kwargs):
        captured.append(request)
        output_root = Path(request.output_root)
        output_root.mkdir(parents=True, exist_ok=True)
        manifest = output_root / "material_manifest.json"
        package = output_root / "qualification.zip"
        manifest.write_text("{}", encoding="utf-8")
        package.write_bytes(b"PK-package")
        return {
            "status": "success",
            "output_path": "",
            "output_paths": {},
            "report_paths": [],
            "material_manifest_paths": {"material": str(manifest)},
            "material_package_paths": {"zip": str(package)},
            "failed_count": 0,
            "artifact_failure_count": 0,
            "error_text": "",
        }

    adapter = AssistantProductionAdapter(executor=executor)
    preflight = adapter.build_preflight(plan)
    result = adapter.execute_approved_plan(
        plan,
        preflight,
        _approval(preflight),
    )

    assert preflight.ready, preflight.issues
    assert result["status"] == "success"
    assert captured
    request = captured[0]
    assert request.scene.input_source_profile.material_schema_id == (
        "qualification_archive_assets_v1"
    )
    assert request.scene.default_delivery_preset_id == "attachment_package"
    assert {
        item.preset_id for item in request.scene.delivery_presets
    } == {
        "attachment_package",
        "missing_items_report",
        "archive_manifest",
    }
