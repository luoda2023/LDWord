from __future__ import annotations

import json
import time
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from src.assistant.application.capability_registry import (
    classify_task_operation,
)
from src.assistant.application.request_policy import (
    POLICY_DOCUMENT_ADVISORY,
    POLICY_RESPONSE_CLOSED,
    evaluate_request_policy,
)
from src.assistant.application.session_coordinator import (
    AssistantSessionCoordinator,
)
from src.assistant.contracts.messages import (
    ROLE_ASSISTANT,
    ROLE_USER,
    AssistantMessage,
)
from src.assistant.contracts.runtime import (
    MAX_ASSISTANT_HISTORY_CHARACTERS,
    MAX_ASSISTANT_USER_MESSAGE_CHARACTERS,
    TURN_FAILED,
    TURN_WAITING_USER_QUESTION,
    AssistantTurnRequest,
)
from src.assistant.contracts.task_plan import TASK_OPERATION_REVIEW
from src.assistant.domain.docx_format_evidence import (
    STANDARD_FORMAT_REFERENCE_ROLE,
    extract_docx_format_evidence,
)
from src.assistant.runtime.provider_contract import (
    PROVIDER_DONE,
    PROVIDER_ERROR,
    PROVIDER_START,
    ProviderStreamEvent,
)
from src.assistant.runtime.providers.profiles import (
    ProviderProfile,
    ProviderProfileStore,
)
from src.assistant.runtime.providers.router import ProviderRouter
from src.assistant.runtime.providers.secrets import MemorySecretStore
from src.assistant.runtime.turn_runner import (
    AssistantTurnRunner,
    build_attachment_context,
)
from src.assistant.storage.session_store import AssistantSessionStore
from src.assistant.ui.assistant_panel import AssistantPanel
from src.assistant.ui.conversation_presentation import (
    build_interaction_action_scope,
)
from src.assistant.ui.creative_home import AssistantHeroComposer
from src.assistant.ui.interaction_card import AssistantInteractionCard
from src.config.scene_natural_request_router import (
    route_natural_scene_request,
)
from src.ui.bridge import PanelBridge


def _wait_for_turn(qapp, panel: AssistantPanel) -> None:
    deadline = time.monotonic() + 8
    while time.monotonic() < deadline and panel._turn_worker is not None:
        qapp.processEvents()
        time.sleep(0.01)
    assert panel._turn_worker is None


class _CaptureGateway:
    def __init__(self) -> None:
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(PROVIDER_DONE, text="ok")

    def cancel(self) -> bool:
        return True


def _panel(tmp_path: Path, gateway) -> AssistantPanel:
    return AssistantPanel(
        PanelBridge(),
        coordinator=AssistantSessionCoordinator(
            AssistantSessionStore(tmp_path / "sessions")
        ),
        turn_runner=AssistantTurnRunner(gateway),
        first_level=True,
    )


def test_missing_selected_attachment_blocks_submit_and_preserves_draft(
    qapp,
    tmp_path,
):
    source = tmp_path / "selected.docx"
    Document().save(source)
    gateway = _CaptureGateway()
    panel = _panel(tmp_path, gateway)
    try:
        panel._on_composer_document_selected(str(source))
        source.unlink()
        composer = panel._empty_input
        composer.set_text("请深入分析这份文档")

        composer._on_send()
        qapp.processEvents()

        assert gateway.requests == []
        assert composer.get_text() == "请深入分析这份文档"
        assert "附件已失效" in composer._submission_gate
        assert "selected.docx" in composer._submission_gate
    finally:
        panel.close()


def test_composer_material_set_stays_bound_to_one_session(qapp, tmp_path):
    docx_path = tmp_path / "primary.docx"
    Document().save(docx_path)
    markdown_path = tmp_path / "notes.md"
    markdown_path.write_text("# Notes", encoding="utf-8")
    panel = _panel(tmp_path, _CaptureGateway())
    try:
        panel._empty_input._add_document_paths(
            (str(docx_path), str(markdown_path))
        )
        qapp.processEvents()

        assert panel._active_session is not None
        assert tuple(
            Path(str(item["path"])).name
            for item in panel._active_session.context_refs
        ) == (docx_path.name, markdown_path.name)
        assert panel._empty_input.document_paths() == (
            str(docx_path.resolve()),
            str(markdown_path.resolve()),
        )
        assert panel._composer.document_paths() == (
            str(docx_path.resolve()),
            str(markdown_path.resolve()),
        )

        panel._empty_input._clear_attachment(str(docx_path.resolve()))
        qapp.processEvents()

        assert tuple(
            Path(str(item["path"])).name
            for item in panel._active_session.context_refs
        ) == (markdown_path.name,)
        assert panel._empty_input.document_paths() == (
            str(markdown_path.resolve()),
        )
    finally:
        panel.close()


def test_successful_send_consumes_composer_materials_exactly_once(
    qapp,
    tmp_path,
):
    source = tmp_path / "analysis.docx"
    Document().save(source)
    panel = _panel(tmp_path, _CaptureGateway())
    try:
        panel._on_composer_document_selected(str(source))
        composer = panel._empty_input
        composer.set_text("请分析这份文档")

        composer._on_send()
        qapp.processEvents()

        session = panel._active_session
        assert session is not None
        assert session.context_refs == ()
        assert panel._empty_input.document_paths() == ()
        assert panel._composer.document_paths() == ()
        assert panel._context_refs_for_submission() == ()
        submitted = next(
            message
            for message in reversed(session.messages)
            if message.role == ROLE_USER
        )
        assert tuple(
            str(item.get("path") or "") for item in submitted.source_refs
        ) == (str(source.resolve()),)
        assert tuple(
            str(item.get("path") or "")
            for item in session.pending_continuation.get(
                "context_refs",
                (),
            )
        ) == (str(source.resolve()),)
    finally:
        panel.close()


def test_retry_action_recovers_attachments_from_the_failed_user_message(
    qapp,
    tmp_path,
    monkeypatch,
):
    source = tmp_path / "retry.docx"
    Document().save(source)
    panel = _panel(tmp_path, _CaptureGateway())
    try:
        session = panel._coordinator.create_session()
        source_refs = (
            {
                "type": "file",
                "name": source.name,
                "path": str(source.resolve()),
            },
        )
        session = panel._coordinator.append_message(
            session,
            AssistantMessage.text(
                role=ROLE_USER,
                text="请分析附件",
                source_refs=source_refs,
            ),
        )
        session = panel._coordinator.update_state(
            session,
            document_job={"status": "failed"},
            turn_status=TURN_FAILED,
        )
        panel._active_session = session
        captured: dict[str, object] = {}

        def _capture_retry(text: str, **kwargs) -> bool:
            captured["text"] = text
            captured.update(kwargs)
            return True

        monkeypatch.setattr(panel, "_send_message", _capture_retry)
        panel._handle_card_action(
            "retry_provider_request",
            {
                "retry_text": "请分析附件",
                "action_scope": build_interaction_action_scope(
                    pending_continuation=session.pending_continuation,
                    active_plan=session.active_plan,
                    document_job=session.document_job,
                    turn_status=session.turn_status,
                ),
            },
        )

        assert captured["text"] == "请分析附件"
        assert captured["context_refs_override"] == source_refs
    finally:
        panel.close()


def test_truncated_attachment_text_declares_partial_coverage(tmp_path):
    source = tmp_path / "long.docx"
    document = Document()
    document.add_paragraph("x" * 50_000)
    document.save(source)

    prompt, audit = build_attachment_context(
        ({"path": str(source), "name": source.name},)
    )

    assert audit["attachment_text_truncated"] is True
    assert audit["attachment_text_character_count"] == 40_000
    assert audit["attachment_text_total_character_count"] == 50_000
    assert 'mode="partial"' in prompt
    assert "不得把分析表述为全文结论" in prompt
    assert audit["attachment_coverage"][0]["omitted_characters"] == 10_000


def test_oversized_format_evidence_is_compacted_not_dropped(
    monkeypatch,
    tmp_path,
):
    source = tmp_path / "complex.docx"
    Document().save(source)
    evidence = {
        "schema_version": "test",
        "source": {"name": source.name, "sha256": "a" * 64},
        "inventory": {
            "section_count": 1,
            "used_paragraph_style_count": 100,
        },
        "document_settings": {},
        "sections": [{"value": "s" * 5_000}] * 12,
        "paragraph_styles": [{"value": "p" * 5_000}] * 24,
        "table_styles": [],
        "table_format_profiles": [],
        "numbering_profiles": [],
        "direct_formatting": {},
        "evidence_policy": {"limitations": []},
    }
    monkeypatch.setattr(
        "src.assistant.runtime.turn_runner.extract_docx_format_evidence",
        lambda _path: evidence,
    )

    prompt, audit = build_attachment_context(
        (
            {
                "path": str(source),
                "name": source.name,
                "semantic_role": STANDARD_FORMAT_REFERENCE_ROLE,
            },
        )
    )

    assert "<document_format_evidence>" in prompt
    assert 'mode="summary"' in prompt
    assert audit["attachment_format_evidence_count"] == 1
    assert audit["attachment_format_evidence_truncated"] is True
    assert audit["attachment_format_evidence_character_count"] > 0
    assert audit["attachment_coverage"][0]["mode"] == "summary"


def test_real_complex_docx_format_evidence_is_compacted_with_coverage(
    tmp_path,
):
    source = tmp_path / "real-complex-reference.docx"
    document = Document()
    for index in range(24):
        style = document.styles.add_style(
            f"审计样式-{index:02d}-" + ("层级" * 20),
            WD_STYLE_TYPE.PARAGRAPH,
        )
        style.font.name = f"Audit Font {index:02d}"
        style.font.size = Pt(9 + index / 2)
        style.font.bold = bool(index % 2)
        style.font.italic = bool(index % 3)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH(index % 4)
        style.paragraph_format.left_indent = Cm(index / 10)
        style.paragraph_format.right_indent = Cm(index / 20)
        style.paragraph_format.first_line_indent = Cm(0.5)
        style.paragraph_format.space_before = Pt(index)
        style.paragraph_format.space_after = Pt(index + 1)
        style.paragraph_format.line_spacing = 1 + index / 20
        document.add_paragraph(f"复杂格式证据段落 {index}", style=style)
    for index in range(11):
        section = document.add_section()
        section.header.paragraphs[0].text = (
            f"第 {index + 2} 节页眉 " + ("格式证据" * 20)
        )
        section.footer.paragraphs[0].text = (
            f"第 {index + 2} 节页脚 " + ("审计范围" * 20)
        )
    document.save(source)

    evidence = extract_docx_format_evidence(source)
    serialized = json.dumps(evidence, ensure_ascii=False)
    assert len(serialized) > 20_000

    prompt, audit = build_attachment_context(
        (
            {
                "path": str(source),
                "name": source.name,
                "semantic_role": STANDARD_FORMAT_REFERENCE_ROLE,
            },
        )
    )

    assert "<document_format_evidence>" in prompt
    assert 'mode="summary"' in prompt
    assert audit["attachment_format_evidence_count"] == 1
    assert audit["attachment_format_evidence_truncated"] is True
    assert audit["attachment_format_evidence_original_character_count"] > 20_000
    assert audit["attachment_coverage"][0]["mode"] == "summary"


class _RetryContinuationGateway:
    def __init__(self) -> None:
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        yield ProviderStreamEvent(PROVIDER_START)
        if len(self.requests) == 1:
            yield ProviderStreamEvent(
                PROVIDER_DONE,
                text="need choice",
                metadata={
                    "side_channel": {
                        "confirmation_requests": [
                            {
                                "kind": "question",
                                "title": "Choose",
                                "options": [
                                    {"id": "student", "label": "Student"}
                                ],
                            }
                        ],
                        "continuation_ref": {
                            "continuation_id": "cursor-1"
                        },
                        "turn_status": TURN_WAITING_USER_QUESTION,
                    }
                },
            )
        elif len(self.requests) == 2:
            yield ProviderStreamEvent(
                PROVIDER_ERROR,
                text="temporary failure",
            )
        else:
            yield ProviderStreamEvent(
                PROVIDER_DONE,
                text="recovered",
            )

    def cancel(self) -> bool:
        return True


def test_failed_confirmation_retry_reuses_cursor_without_duplicate_user_message(
    qapp,
    tmp_path,
):
    gateway = _RetryContinuationGateway()
    panel = _panel(tmp_path, gateway)
    try:
        assert panel._send_message("start") is True
        _wait_for_turn(qapp, panel)
        session = panel._active_session
        assert session is not None
        panel._handle_card_action(
            "submit_question_answer",
            {
                "response": "Student",
                "selected_choice_ids": ["student"],
                "continuation_ref": dict(session.pending_continuation),
                "action_scope": build_interaction_action_scope(
                    pending_continuation=session.pending_continuation,
                    active_plan=session.active_plan,
                    document_job=session.document_job,
                    turn_status=session.turn_status,
                ),
            },
        )
        _wait_for_turn(qapp, panel)
        assert panel._active_session is not None
        assert panel._active_session.pending_continuation[
            "last_attempt_failed"
        ] is True

        session = panel._active_session
        panel._handle_card_action(
            "retry_provider_request",
            {
                "retry_text": "Student",
                "action_scope": build_interaction_action_scope(
                    pending_continuation=session.pending_continuation,
                    active_plan=session.active_plan,
                    document_job=session.document_job,
                    turn_status=session.turn_status,
                ),
            },
        )
        _wait_for_turn(qapp, panel)

        assert [
            request.metadata["conversation_cursor"]
            for request in gateway.requests
        ] == ["", "cursor-1", "cursor-1"]
        assert panel._active_session.pending_continuation == {}
        assert len(
            [
                message
                for message in panel._active_session.messages
                if message.role == ROLE_USER
            ]
        ) == 2
    finally:
        panel.shutdown_active_execution(1000)
        panel.close()


def test_new_local_document_request_supersedes_provider_question_cursor(
    qapp,
    tmp_path,
):
    gateway = _RetryContinuationGateway()
    panel = _panel(tmp_path, gateway)
    try:
        assert panel._send_message("start") is True
        _wait_for_turn(qapp, panel)
        session = panel._active_session
        assert session is not None
        assert session.pending_continuation["continuation_id"] == "cursor-1"
        assert len(gateway.requests) == 1

        assert panel._send_message("我要发个通知") is True
        qapp.processEvents()

        session = panel._active_session
        assert session is not None
        assert len(gateway.requests) == 1
        assert session.document_job["status"] == "plan_ready"
        assert session.pending_continuation == {}
        assert session.active_plan["operation"] == "author"
        assert session.active_plan["work_mode_id"] == "official"
    finally:
        panel.shutdown_active_execution(1000)
        panel.close()


def test_startup_recovery_closes_stale_provider_job_and_allows_local_plan(
    qapp,
    tmp_path,
):
    source = tmp_path / "input.docx"
    Document().save(source)
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "recovery")
    )
    session = coordinator.create_session()
    session = coordinator.update_state(
        session,
        context_refs=(
            {"path": str(source.resolve()), "name": source.name},
        ),
        document_job={"status": "provider_running"},
        turn_status="provider_running",
    )
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(_CaptureGateway()),
        first_level=True,
    )
    try:
        recovered = coordinator.load_session(session.session_id)
        assert recovered.turn_status == "failed"
        assert recovered.document_job["status"] == "failed"

        panel._open_session_by_id(session.session_id)
        assert panel._send_message("统一这份 Word 文档格式") is True
        assert panel._active_session is not None
        assert panel._active_session.document_job["status"] == "plan_ready"
    finally:
        panel.close()


def test_exam_make_request_clarifies_stage_then_renders_plan_without_provider(
    qapp,
    tmp_path,
):
    gateway = _CaptureGateway()
    panel = _panel(tmp_path, gateway)
    try:
        assert panel._send_message(
            "我需要制作一份初中六年级语文期中考试试卷"
        ) is True
        qapp.processEvents()

        session = panel._active_session
        assert session is not None
        assert gateway.requests == []
        assert session.document_job["status"] == "needs_route_clarification"
        assert session.pending_continuation["kind"] == (
            "local_exam_clarification"
        )
        question = session.messages[-1].blocks[0]
        assert question.data["interaction_type"] == "question"

        panel._handle_card_action(
            "submit_question_answer",
            {
                "clarification_id": session.pending_continuation[
                    "clarification_id"
                ],
                "selected_choice_ids": ["middle_preparatory"],
            },
        )
        qapp.processEvents()

        session = panel._active_session
        assert session is not None
        assert session.document_job["status"] == "plan_ready"
        assert session.active_plan["operation"] == "author"
        assert session.active_plan["scene_ref"]["id"] == "exam_term"
        assert session.active_plan["scene_ref"]["generation_mode"] == (
            "from_prompt"
        )
        block = session.messages[-1].blocks[0]
        assert block.type == "interaction"
        assert block.data["interaction_type"] == "plan"
        assert any(
            fact == {
                "label": "年级",
                "value": "初中预备班（六年级）",
            }
            for fact in block.data["facts"]
        )
        assert any(
            action["id"] == "generate_content_draft"
            for action in block.data["actions"]
        )
    finally:
        panel.close()


@pytest.mark.parametrize(
    "query",
    (
        "我需要写一个公文",
        "我要发一个公文",
        "我有写一篇公文",
    ),
)
def test_generic_official_request_uses_one_local_structured_intake(
    qapp,
    tmp_path,
    query,
):
    gateway = _CaptureGateway()
    panel = _panel(tmp_path, gateway)
    try:
        assert panel._send_message(query) is True
        qapp.processEvents()

        session = panel._active_session
        assert session is not None
        assert gateway.requests == []
        assert session.document_job["status"] == "needs_route_clarification"
        assert session.pending_continuation["kind"] == (
            "local_official_clarification"
        )
        question = session.messages[-1].blocks[0]
        assert question.data["interaction_type"] == "question"
        request = question.data["confirmation_request"]
        assert request["requires_choice"] is True
        assert request["choice_columns"] == 2
        assert request["initial_choice_count"] == 7
        assert request["expand_choices_label"] == "更多文种（8）"
        assert [item["id"] for item in request["options"][:7]] == [
            "notice",
            "notice_public",
            "circular",
            "report",
            "request",
            "letter",
            "minutes",
        ]
        assert len(request["options"]) == 15
        assert request["input_columns"] == 2
        assert request["compact_heading"] is True
        assert request["submit_label"] == "生成计划"
        assert not question.data.get("body")
        assert [item["id"] for item in request["inputs"]] == [
            "organization",
            "recipient",
            "purpose",
        ]

        panel._handle_card_action(
            "submit_question_answer",
            {
                "clarification_id": session.pending_continuation[
                    "clarification_id"
                ],
                "selected_choice_ids": ["notice"],
                "field_values": {
                    "organization": "示例市教育局",
                    "recipient": "各区教育局",
                    "purpose": "部署秋季校园安全检查工作",
                },
            },
        )
        qapp.processEvents()

        session = panel._active_session
        assert session is not None
        assert session.document_job["status"] == "plan_ready"
        assert session.pending_continuation == {}
        assert session.active_plan["operation"] == "author"
        assert session.active_plan["work_mode_id"] == "official"
        assert session.active_plan["production_contract"][
            "document_type_id"
        ] == "notice"
        assert session.active_plan["scene_ref"]["official_field_values"] == {
            "document_type": "notice",
            "organization": "示例市教育局",
            "recipient": "各区教育局",
        }
        assert session.active_plan["scene_ref"][
            "official_content_requirements"
        ] == "部署秋季校园安全检查工作"
        assert session.active_plan["blocking_issues"] == []
        plan_card = session.messages[-1].blocks[0]
        assert any(
            action["id"] == "generate_content_draft"
            for action in plan_card.data["actions"]
        )
        assert gateway.requests == []
    finally:
        panel.close()


@pytest.mark.parametrize(
    "query, expects_import_boundary",
    [
        ("PDF论文排版", True),
        ("分析PDF论文", True),
        ("扫描论文OCR导入", True),
        ("OCR毕业论文", True),
        ("PDF转Word论文", True),
        ("导入扫描论文", True),
        ("扫描件论文转换", True),
        ("把PDF论文转成Word", True),
        ("识别扫描版课程论文", True),
        ("OCR导入论文后再排版", True),
        ("分析这份论文", False),
        ("审阅这篇论文", False),
        ("评估论文结构", False),
        ("检查毕业论文", False),
        ("论文摘要分析", False),
        ("分析课程论文", False),
        ("总结这份论文", False),
        ("修改论文格式", False),
        ("论文排版", False),
        ("文献综述分析", False),
        ("开题报告分析", False),
        ("论文参考文献检查", False),
        ("论文写作建议", False),
        ("诊断这篇论文的问题", False),
    ],
)
def test_thesis_import_boundary_requires_media_or_conversion_intent(
    query,
    expects_import_boundary,
):
    routed = route_natural_scene_request(query)
    strongest_is_import = bool(
        routed.matches
        and routed.matches[0].route.route_type == "import_boundary"
    )

    assert strongest_is_import is expects_import_boundary


def test_plain_thesis_analysis_is_review_advisory_not_import_gate():
    decision = evaluate_request_policy(
        "分析这份论文",
        workspace_mode_id="custom",
        has_attachment=True,
    )

    assert classify_task_operation("分析这份论文") == TASK_OPERATION_REVIEW
    assert decision.kind != POLICY_RESPONSE_CLOSED
    assert decision.route_id != "import_pdf_thesis_boundary"


def test_reference_analysis_plus_apply_keeps_sample_role_and_blocks_production(
    qapp,
    tmp_path,
):
    source = tmp_path / "standard.docx"
    Document().save(source)
    gateway = _CaptureGateway()
    panel = _panel(tmp_path, gateway)
    query = "先分析这份标准文件的格式要求，再套用到另一份文档"
    try:
        panel._on_composer_document_selected(str(source))
        assert panel._send_message(query) is True

        session = panel._active_session
        assert session is not None
        assert session.document_job["status"] == "needs_data_disclosure"
        assert session.active_plan == {}
        assert session.context_refs == ()
        disclosed_ref = session.pending_continuation["context_refs"][0]
        assert disclosed_ref["semantic_role"] == (
            STANDARD_FORMAT_REFERENCE_ROLE
        )
        assert disclosed_ref["target_attachment_required"] is True
        assert "目标文档尚未提供" in session.messages[-1].blocks[0].text
        assert gateway.requests == []
        decision = evaluate_request_policy(
            query,
            workspace_mode_id="custom",
            has_attachment=True,
        )
        assert decision.kind == POLICY_DOCUMENT_ADVISORY
    finally:
        panel.close()


def test_oversized_single_requirement_is_blocked_before_session_or_provider(
    qapp,
    tmp_path,
):
    gateway = _CaptureGateway()
    panel = _panel(tmp_path, gateway)
    try:
        text = "x" * (MAX_ASSISTANT_USER_MESSAGE_CHARACTERS + 1)
        panel._empty_input.set_text(text)
        qapp.processEvents()

        assert panel._empty_input._send_btn.isEnabled() is False
        assert "单次最多" in panel._empty_input._input_limit_requirement
        assert panel._send_message(
            text,
            source=panel._empty_input,
        ) is False
        assert gateway.requests == []
        assert panel._active_session is None
        assert panel._empty_input.get_text() == text
    finally:
        panel.close()


def test_large_history_uses_recent_window_and_reports_coverage():
    history = tuple(
        AssistantMessage.text(
            role=ROLE_USER if index % 2 == 0 else ROLE_ASSISTANT,
            text=f"{index:02d}-" + ("x" * 10_000),
        )
        for index in range(40)
    )
    request = AssistantTurnRequest(
        turn_id="history-budget",
        session_id="session-history-budget",
        user_message="current",
        provider_profile_id="mock-default",
        model_id="form-assistant-mock",
        history=history,
    )
    gateway = _CaptureGateway()

    result = AssistantTurnRunner(gateway).run(request)

    assert result.status != TURN_FAILED
    assert len(gateway.requests) == 1
    sent = gateway.requests[0].messages
    assert "conversation_history_coverage" in sent[0]["content"]
    sent_history_characters = sum(
        len(item["content"])
        for item in sent
        if "conversation_history_coverage" not in item["content"]
    )
    assert sent_history_characters <= MAX_ASSISTANT_HISTORY_CHARACTERS
    audit = result.provider_audit["history"]
    assert audit["truncated"] is True
    assert audit["omitted_message_count"] > 0
    assert audit["sent_character_count"] <= MAX_ASSISTANT_HISTORY_CHARACTERS


def test_turn_runner_rejects_oversized_direct_request_before_gateway():
    request = AssistantTurnRequest(
        turn_id="oversized-direct",
        session_id="session-oversized-direct",
        user_message="x" * (MAX_ASSISTANT_USER_MESSAGE_CHARACTERS + 1),
        provider_profile_id="mock-default",
        model_id="form-assistant-mock",
    )
    gateway = _CaptureGateway()

    result = AssistantTurnRunner(gateway).run(request)

    assert result.status == TURN_FAILED
    assert result.error["message"] == "user_message_too_large"
    assert gateway.requests == []


def _provider_switch_panel(tmp_path, qapp):
    profiles = ProviderProfileStore(tmp_path / "providers.json")
    profiles.upsert(
        ProviderProfile(
            profile_id="cloud-a",
            label="Cloud A",
            kind="openai_compatible",
            model_id="model-a",
            base_url="https://a.example.invalid/v1",
        )
    )
    profiles.upsert(
        ProviderProfile(
            profile_id="cloud-a-same-domain",
            label="Cloud A2",
            kind="openai_compatible",
            model_id="model-a2",
            base_url="https://a.example.invalid/v1/",
        )
    )
    profiles.upsert(
        ProviderProfile(
            profile_id="cloud-b",
            label="Cloud B",
            kind="openai_compatible",
            model_id="model-b",
            base_url="https://b.example.invalid/v1",
        )
    )
    secrets = MemorySecretStore(
        {
            "cloud-a": "secret-a",
            "cloud-a-same-domain": "secret-a2",
            "cloud-b": "secret-b",
        }
    )
    coordinator = AssistantSessionCoordinator(
        AssistantSessionStore(tmp_path / "provider-switch-sessions")
    )
    session = coordinator.create_session(
        provider_profile_id="cloud-a",
        model_id="model-a",
    )
    session = coordinator.append_message(
        session,
        AssistantMessage.text(role=ROLE_USER, text="private history"),
    )
    session = coordinator.append_message(
        session,
        AssistantMessage.text(role=ROLE_ASSISTANT, text="private reply"),
    )
    gateway = _CaptureGateway()
    panel = AssistantPanel(
        PanelBridge(),
        coordinator=coordinator,
        turn_runner=AssistantTurnRunner(gateway),
        provider_router=ProviderRouter(
            profiles=profiles,
            secrets=secrets,
        ),
        first_level=True,
    )
    panel._open_session_by_id(session.session_id)
    qapp.processEvents()
    return panel, coordinator, gateway, session


def _select_profile(qapp, panel: AssistantPanel, profile_id: str) -> None:
    combo = panel._provider_combo
    index = next(
        row
        for row in range(combo.count())
        if combo.itemData(row) == profile_id
    )
    combo.setCurrentIndex(index)
    qapp.processEvents()


def test_same_provider_domain_switch_does_not_require_history_disclosure(
    qapp,
    tmp_path,
):
    panel, _coordinator, _gateway, _session = _provider_switch_panel(
        tmp_path,
        qapp,
    )
    panel._confirm_provider_history_transition = lambda *_args, **_kwargs: (
        pytest.fail("same data domain must not request cross-domain consent")
    )
    try:
        _select_profile(qapp, panel, "cloud-a-same-domain")

        assert panel._active_session is not None
        assert panel._active_session.provider_profile_id == (
            "cloud-a-same-domain"
        )
        assert panel._active_session.provider_history_grant == {}
    finally:
        panel.close()


def test_cross_provider_carry_persists_grant_and_audits_sent_history(
    qapp,
    tmp_path,
):
    panel, coordinator, gateway, _session = _provider_switch_panel(
        tmp_path,
        qapp,
    )
    panel._confirm_provider_history_transition = (
        lambda *_args, **_kwargs: "carry"
    )
    try:
        _select_profile(qapp, panel, "cloud-b")

        switched = panel._active_session
        assert switched is not None
        assert switched.provider_profile_id == "cloud-b"
        assert switched.provider_history_grant["target_profile_id"] == (
            "cloud-b"
        )
        stored = coordinator.load_session(switched.session_id)
        assert stored.provider_history_grant["source_profile_id"] == "cloud-a"

        assert panel._send_message("continue") is True
        _wait_for_turn(qapp, panel)
        assert gateway.requests[-1].model == "model-b"
        contents = "|".join(
            item["content"] for item in gateway.requests[-1].messages
        )
        assert "private history" in contents
        runtime_result = panel._active_session.document_job[
            "runtime_results"
        ][-1]
        grant = runtime_result["provider_audit"]["history"][
            "disclosure_grant"
        ]
        assert grant["target_profile_id"] == "cloud-b"
        assert grant["history_message_count"] == 2
    finally:
        panel.shutdown_active_execution(1000)
        panel.close()


def test_cross_provider_can_create_blank_session_without_history(
    qapp,
    tmp_path,
):
    panel, coordinator, _gateway, old_session = _provider_switch_panel(
        tmp_path,
        qapp,
    )
    panel._confirm_provider_history_transition = (
        lambda *_args, **_kwargs: "new"
    )
    try:
        _select_profile(qapp, panel, "cloud-b")

        assert panel._active_session is not None
        assert panel._active_session.session_id != old_session.session_id
        assert panel._active_session.provider_profile_id == "cloud-b"
        assert panel._active_session.messages == ()
        assert coordinator.load_session(old_session.session_id).messages
    finally:
        panel.close()


def test_cross_provider_cancel_restores_original_selection(
    qapp,
    tmp_path,
):
    panel, _coordinator, _gateway, old_session = _provider_switch_panel(
        tmp_path,
        qapp,
    )
    panel._confirm_provider_history_transition = (
        lambda *_args, **_kwargs: "cancel"
    )
    try:
        _select_profile(qapp, panel, "cloud-b")

        assert panel._active_session is not None
        assert panel._active_session.session_id == old_session.session_id
        assert panel._active_session.provider_profile_id == "cloud-a"
        assert panel._provider_combo.currentData() == "cloud-a"
    finally:
        panel.close()


def test_submission_exception_is_visible_retryable_and_preserves_draft(qapp):
    composer = AssistantHeroComposer(PanelBridge())
    composer.set_submission_handler(
        lambda _text: (_ for _ in ()).throw(OSError("disk full"))
    )
    try:
        composer.set_text("important requirement")
        composer._on_send()
        qapp.processEvents()

        assert composer.get_text() == "important requirement"
        assert "需求未写入" in composer._submission_error
        assert "OSError" in composer._submission_error
        assert composer._keyboard_hint.text() == composer._submission_error
        assert composer._send_btn.isEnabled() is True
    finally:
        composer.close()


def test_draft_persistence_failure_is_caught_and_shown(
    qapp,
    tmp_path,
    monkeypatch,
):
    panel = _panel(tmp_path, _CaptureGateway())
    try:
        panel.new_session()
        with monkeypatch.context() as scoped:
            scoped.setattr(
                panel._coordinator,
                "persist",
                lambda *_args, **_kwargs: (_ for _ in ()).throw(
                    OSError("disk full")
                ),
            )

            panel._empty_input.set_text("draft survives")
            qapp.processEvents()
            assert not panel._flush_active_draft()

            assert panel._empty_input.get_text() == "draft survives"
            assert "草稿暂未保存" in panel._empty_input._submission_error
            assert "OSError" in panel._empty_input._submission_error
    finally:
        panel.close()


def test_question_card_renders_every_provider_option(qapp):
    options = [
        {"id": str(index), "label": f"Option {index}"}
        for index in range(7)
    ]
    card = AssistantInteractionCard(
        interaction_type="question",
        title="Choose",
        body="Every option must remain reachable.",
        payload={
            "confirmation_request": {
                "options": options,
                "multiple": True,
            }
        },
    )
    try:
        qapp.processEvents()
        assert len(card.presentation.choices) == 7
        assert len(card._choice_buttons) == 7
        assert [
            button.choice_id for button in card._choice_buttons
        ] == [str(index) for index in range(7)]
    finally:
        card.close()


@pytest.mark.parametrize(
    "interaction_type",
    [
        "question",
        "disclosure",
        "permission",
        "plan_candidate",
        "plan",
        "approval",
        "progress",
    ],
)
def test_historical_interaction_cards_do_not_render_redundant_status_receipts(
    qapp,
    interaction_type,
):
    card = AssistantInteractionCard(
        interaction_type=interaction_type,
        title="历史卡片",
        body="",
        payload={"active": False},
    )
    try:
        qapp.processEvents()
        assert not hasattr(card, "_receipt")
        assert card.findChild(type(card._type_icon), "assistant_card_receipt") is None
    finally:
        card.close()


def test_corrupt_session_is_visible_as_recovery_entry_and_original_is_kept(
    qapp,
    tmp_path,
):
    store = AssistantSessionStore(tmp_path / "corrupt-sessions")
    store.sessions_dir.mkdir(parents=True)
    broken = store.sessions_dir / "broken.json"
    broken.write_text("{not-json", encoding="utf-8")

    summaries = store.list_summaries()

    assert len(summaries) == 1
    assert summaries[0].corrupt is True
    assert summaries[0].recovery_path == str(broken.resolve())
    assert broken.read_text(encoding="utf-8") == "{not-json"

    panel = AssistantPanel(
        PanelBridge(),
        coordinator=AssistantSessionCoordinator(store),
        turn_runner=AssistantTurnRunner(_CaptureGateway()),
        first_level=True,
    )
    opened = []
    panel._show_corrupt_session_recovery = opened.append
    try:
        panel._open_session_by_id(summaries[0].session_id)
        assert opened == [str(broken.resolve())]
        assert broken.is_file()
    finally:
        panel.close()
