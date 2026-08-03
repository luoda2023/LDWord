from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from docx import Document
from PIL import Image

from src.application.materials import (
    ExecutionMaterialRecord,
    ExecutionMaterialSnapshot,
    ExecutionResource,
    MaterialPreviewSnapshot,
)
from src.config.document_scope import DocumentScopePolicy
from src.config.execution_feature_state import (
    DISABLED_SELECTOR_VALUE,
    MATERIAL_CONTROLLED_MODULES,
    project_execution_scene,
)
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.domain.materials import (
    MaterialObjectRef,
    MaterialPackageRef,
    MaterialRunSelection,
    generate_package_id,
    generate_record_id,
)
from src.modules.fill.entity_fill import EntityFillModule
from src.services.document_structure_evidence import (
    build_document_structure_evidence,
)
from src.services.production_execution import (
    ProductionExecutionRequest,
    execute_production_request,
)
from src.ui.panels.workbench.quick_execution_detail import QuickExecutionDetail


def _material_selection() -> tuple[MaterialRunSelection, MaterialPreviewSnapshot]:
    selection = MaterialRunSelection(
        package_ref=MaterialPackageRef(
            generate_package_id(),
            "sha256:" + "a" * 64,
        ),
        selected_record_ids=(generate_record_id(),),
    )
    preview = MaterialPreviewSnapshot(
        package_id=selection.package_ref.package_id,
        package_revision=selection.package_ref.revision,
        package_name="独立资料包",
        source_type="user",
        current_record_id=selection.selected_record_ids[0],
        current_record_name="记录 1",
        record_count=1,
        resource_count=2,
    )
    return selection, preview


def _image_material_snapshot(tmp_path) -> ExecutionMaterialSnapshot:
    image_path = tmp_path / "photo.png"
    Image.new("RGB", (320, 180), "navy").save(image_path)
    image_payload = image_path.read_bytes()
    resource = ExecutionResource(
        MaterialObjectRef(
            object_id="sha256:" + sha256(image_payload).hexdigest(),
            media_type="image/png",
            original_name=image_path.name,
            size=len(image_payload),
        ),
        str(image_path.resolve()),
    )
    return _execution_material_snapshot(
        ExecutionMaterialRecord(
            record_id=generate_record_id(),
            display_name="Record 1",
            group_id="",
            field_values={},
            field_owners={},
            resources={"photo": (resource,)},
            resource_owners={"photo": ("record",)},
        ),
        resource_domains={"photo": "image"},
    )


def _execution_material_snapshot(
    record: ExecutionMaterialRecord,
    *,
    resource_domains: dict[str, str] | None = None,
) -> ExecutionMaterialSnapshot:
    return ExecutionMaterialSnapshot(
        snapshot_id="",
        run_id="run-independent-material",
        package_ref=MaterialPackageRef(
            generate_package_id(),
            "sha256:" + "b" * 64,
        ),
        work_mode_id="custom",
        material_contract_id="generic_document_v1",
        recipe_id="document_batch",
        scene_id="custom",
        document_type="",
        package_display_name="Materials",
        package_field_values={},
        package_field_owners={},
        groups=(),
        records=(record,),
        resource_domains=resource_domains or {},
    )


def test_plan_off_enables_material_runtime_modules_independently() -> None:
    scene = SceneWorkspace()
    scene.module_switches.update(
        {
            "paragraph_style": True,
            "placeholder_replace": False,
            "image_insertion": False,
        }
    )
    scene.input_source_profile.require_material_package = True
    scene.input_source_profile.material_schema_id = "generic_document_v1"
    scene.input_source_profile.material_schema_ids = ["generic_document_v1"]
    scene.input_source_profile.required_material_fields = ["title"]
    scene.input_source_profile.required_image_roles = ["logo"]

    projected = project_execution_scene(
        scene,
        plan_enabled=False,
        template_enabled=False,
        material_enabled=True,
    )

    assert projected.module_switches["paragraph_style"] is False
    assert all(
        projected.module_switches[module_name] is True
        for module_name in MATERIAL_CONTROLLED_MODULES
    )
    assert projected.input_source_profile.require_material_package is True
    assert projected.input_source_profile.material_schema_id == "generic_document_v1"
    assert projected.input_source_profile.material_schema_ids == [
        "generic_document_v1"
    ]
    assert projected.input_source_profile.required_material_fields == ["title"]
    assert projected.input_source_profile.required_image_roles == ["logo"]
    artifacts = projected.default_delivery_preset().artifacts
    assert artifacts.final_docx is True
    assert artifacts.compare_docx is False
    assert artifacts.report_json is False
    assert artifacts.report_markdown is False
    assert artifacts.material_manifest is False
    assert artifacts.material_package is False
    assert projected.default_delivery_preset().include_structured_intermediate is False


def test_plan_on_cannot_disable_enabled_material_runtime_modules() -> None:
    scene = SceneWorkspace()
    for module_name in MATERIAL_CONTROLLED_MODULES:
        scene.module_switches[module_name] = False

    projected = project_execution_scene(
        scene,
        plan_enabled=True,
        template_enabled=False,
        material_enabled=True,
    )

    assert all(
        projected.module_switches[module_name] is True
        for module_name in MATERIAL_CONTROLLED_MODULES
    )


def test_material_off_still_disables_material_runtime_modules() -> None:
    scene = SceneWorkspace()
    for module_name in MATERIAL_CONTROLLED_MODULES:
        scene.module_switches[module_name] = True
    scene.input_source_profile.require_material_package = True
    scene.input_source_profile.required_material_fields = ["title"]
    scene.input_source_profile.required_image_roles = ["logo"]

    projected = project_execution_scene(
        scene,
        plan_enabled=False,
        template_enabled=False,
        material_enabled=False,
    )

    assert all(
        projected.module_switches[module_name] is False
        for module_name in MATERIAL_CONTROLLED_MODULES
    )
    assert projected.input_source_profile.require_material_package is False
    assert projected.input_source_profile.required_material_fields == []
    assert projected.input_source_profile.required_image_roles == []


def test_quick_execution_keeps_material_selected_when_plan_is_off(qapp) -> None:
    selection, preview = _material_selection()
    detail = QuickExecutionDetail(include_shared_chrome=False)
    try:
        detail.set_material_selection(selection, preview_snapshot=preview)
        disabled_index = detail._scene_combo.findData(DISABLED_SELECTOR_VALUE)
        assert disabled_index >= 0

        detail._scene_combo.setCurrentIndex(disabled_index)
        qapp.processEvents()

        assert detail.plan_enabled() is False
        assert detail.material_package_enabled() is True
        assert detail.execution_material_selection() == selection
        assert detail.current_execution_gate_decision().can_run is True
        assert detail._material_preview.projection().package_label == "独立资料包"
        assert detail._material_preview.projection().summary_text != "本次不使用资料包"
        assert all(
            detail.execution_scene().module_switches[module_name] is True
            for module_name in MATERIAL_CONTROLLED_MODULES
        )
    finally:
        detail.close()
        qapp.processEvents()


def test_plan_off_material_image_generation_ignores_plan_scope_anchor(
    tmp_path,
    monkeypatch,
) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("{{@img:photo}}")
    document.save(source)
    evidence = build_document_structure_evidence(source)
    assert evidence.regions

    snapshot = _image_material_snapshot(tmp_path)
    scene = project_execution_scene(
        SceneWorkspace(mode_id="custom"),
        plan_enabled=False,
        template_enabled=True,
        material_enabled=True,
    )
    monkeypatch.setattr(
        "src.services.production_runtime.execution_runtime.create_all_modules",
        list,
    )

    result = execute_production_request(
        ProductionExecutionRequest(
            input_path=source,
            output_root=tmp_path / "outputs",
            mode_id="custom",
            scene=scene,
            template=TemplateConfig(),
            plan_id=DISABLED_SELECTOR_VALUE,
            material_snapshot=snapshot,
            document_structure_evidence=evidence,
        )
    )

    assert result["status"] == "success"
    assert result["materialized_resource_count"] == 1
    assert result["compare_paths"] == {}
    assert result["report_paths"] == []
    assert result["intermediate_paths"] == {}
    assert result.get("material_manifest_paths", {}) == {}
    assert result.get("material_package_paths", {}) == {}
    output = Path(str(result["output_path"]))
    assert output.is_file()
    assert {
        item.name
        for item in output.parent.iterdir()
        if item.name != ".lark-material-transactions"
    } == {output.name}
    generated = Document(output)
    assert len(generated.inline_shapes) == 1
    assert "{{@img:photo}}" not in "\n".join(
        paragraph.text for paragraph in generated.paragraphs
    )


def test_material_image_without_matching_placeholder_fails_preflight(
    tmp_path,
) -> None:
    source = tmp_path / "source-without-image-token.docx"
    document = Document()
    document.add_paragraph("No image token")
    document.save(source)
    scene = project_execution_scene(
        SceneWorkspace(mode_id="custom"),
        plan_enabled=False,
        template_enabled=True,
        material_enabled=True,
    )

    result = execute_production_request(
        ProductionExecutionRequest(
            input_path=source,
            output_root=tmp_path / "outputs",
            mode_id="custom",
            scene=scene,
            template=TemplateConfig(name="Template"),
            plan_id="",
            output_suffix="-formatted",
            material_snapshot=_image_material_snapshot(tmp_path),
        )
    )

    assert result["status"] == "failed"
    assert "material.bind.template_resources_unsupported:photo" in str(
        result["error_text"]
    )
    assert not list((tmp_path / "outputs").glob("*.docx"))


def test_repeated_material_generation_versions_the_output_path(tmp_path) -> None:
    source = tmp_path / "source-retry.docx"
    document = Document()
    document.add_paragraph("{{@img:photo}}")
    document.save(source)
    scene = project_execution_scene(
        SceneWorkspace(mode_id="custom"),
        plan_enabled=False,
        template_enabled=True,
        material_enabled=True,
    )
    request = ProductionExecutionRequest(
        input_path=source,
        output_root=tmp_path / "outputs",
        mode_id="custom",
        scene=scene,
        template=TemplateConfig(name="Template"),
        plan_id="",
        output_suffix="-formatted",
        material_snapshot=_image_material_snapshot(tmp_path),
    )

    first = execute_production_request(request)
    second = execute_production_request(request)

    assert first["status"] == "success"
    assert second["status"] == "success"
    assert first["output_path"] != second["output_path"]
    assert Path(str(first["output_path"])).name == "source-retry-formatted.docx"
    assert Path(str(second["output_path"])).name == (
        "source-retry-formatted (2).docx"
    )


def test_plan_on_rebinds_scope_after_material_image_expansion(
    tmp_path,
    monkeypatch,
) -> None:
    source = tmp_path / "source-plan-on.docx"
    document = Document()
    document.add_paragraph("{{@img:photo}}")
    document.save(source)
    evidence = build_document_structure_evidence(source)
    scene = SceneWorkspace(
        mode_id="custom",
        document_scope=DocumentScopePolicy(mode="body"),
    )
    scene = project_execution_scene(
        scene,
        plan_enabled=True,
        template_enabled=True,
        material_enabled=True,
    )
    monkeypatch.setattr(
        "src.services.production_runtime.execution_runtime.create_all_modules",
        list,
    )

    result = execute_production_request(
        ProductionExecutionRequest(
            input_path=source,
            output_root=tmp_path / "outputs-plan-on",
            mode_id="custom",
            scene=scene,
            template=TemplateConfig(),
            plan_id="custom",
            material_snapshot=_image_material_snapshot(tmp_path),
            document_structure_evidence=evidence,
        )
    )

    assert result["status"] == "success"
    generated = Document(Path(str(result["output_path"])))
    assert len(generated.inline_shapes) == 1


def test_material_text_outside_body_is_independent_from_plan_scope(
    tmp_path,
    monkeypatch,
) -> None:
    source = tmp_path / "outside-body.docx"
    document = Document()
    document.add_paragraph("\u6458\u8981")
    document.add_paragraph("{{@text:title}}")
    document.add_heading("\u7b2c\u4e00\u7ae0 \u7eea\u8bba", level=1)
    document.add_paragraph("\u6b63\u6587\u5185\u5bb9")
    document.save(source)
    evidence = build_document_structure_evidence(source)
    assert {region.role_id for region in evidence.regions} >= {
        "abstract_cn",
        "body",
    }
    record = ExecutionMaterialRecord(
        record_id=generate_record_id(),
        display_name="Record 1",
        group_id="",
        field_values={"title": "Material title"},
        field_owners={"title": "record"},
        resources={},
        resource_owners={},
    )
    scene = SceneWorkspace(
        mode_id="custom",
        document_scope=DocumentScopePolicy(mode="body"),
    )
    scene = project_execution_scene(
        scene,
        plan_enabled=True,
        template_enabled=True,
        material_enabled=True,
    )
    monkeypatch.setattr(
        "src.services.production_runtime.execution_runtime.create_all_modules",
        lambda: [EntityFillModule()],
    )

    result = execute_production_request(
        ProductionExecutionRequest(
            input_path=source,
            output_root=tmp_path / "outputs-outside-body",
            mode_id="custom",
            scene=scene,
            template=TemplateConfig(),
            plan_id="custom",
            material_snapshot=_execution_material_snapshot(record),
            document_structure_evidence=evidence,
        )
    )

    assert result["status"] == "success"
    generated = Document(Path(str(result["output_path"])))
    assert generated.paragraphs[1].text == "Material title"
