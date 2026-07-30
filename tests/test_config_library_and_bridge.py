import ast
import sys
import tempfile
import json
from pathlib import Path

from docx import Document
import pytest


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.library import (
    default_scene_entry,
    default_scene_descriptor,
    default_template_entry,
    ensure_config_library,
    get_scene_descriptor,
    get_template_entry,
    list_template_entries,
    list_scene_descriptors,
    load_template_from_library,
    load_scene_from_library,
)
from src.config.builtin_templates import create_builtin_template
from src.config.loader import load_scene, save_scene, save_template
from src.config.master_library import MasterSpec
from src.config.official_document_profiles import (
    OFFICIAL_PLAN_ENTRY_BUILTIN,
    OFFICIAL_PLAN_ENTRY_CANDIDATE,
    OFFICIAL_PLAN_ENTRY_PROFILE_ONLY,
    list_official_document_builtin_plan_profile_ids,
    list_common_official_document_profiles,
    list_official_document_plan_entry_decisions,
    list_official_document_profiles,
)
from src.config.scene import SceneWorkspace
from src.config.template import TemplateConfig
from src.config.work_mode import list_work_modes
from src.ui.adapters.config_selector_models import (
    master_display_label,
    master_selector_options,
    material_package_sample_selector_options,
    plan_selector_options,
    strip_source_prefix,
    template_selector_options,
)
from src.qt_api import QApplication
from src.shared.ui.styled_combo_box import SOURCE_BADGE_TEXT_ROLE
from src.ui.bridge import PanelBridge
from src.ui.panels.scene_panel import ScenePanel
from src.ui.panels.scene_state_projection import is_scene_selector_group
from src.ui.panels.workbench import WorkbenchPanel


def _app():
    return QApplication.instance() or QApplication([])


def _combo_values(combo, *, skip_scene_groups: bool = False) -> list[str]:
    values: list[str] = []
    for index in range(combo.count()):
        value = str(combo.itemData(index) or "").strip()
        if skip_scene_groups and is_scene_selector_group(value):
            continue
        values.append(value)
    return values


def _combo_labels(combo, *, skip_scene_groups: bool = False) -> list[str]:
    labels: list[str] = []
    for index in range(combo.count()):
        value = str(combo.itemData(index) or "").strip()
        if skip_scene_groups and is_scene_selector_group(value):
            continue
        labels.append(combo.itemText(index))
    return labels


def _call_name(node):
    if isinstance(node, ast.Name):
        return node.id
    if isinstance(node, ast.Attribute):
        return node.attr
    return ""


def _has_mode_context(name: str, node: ast.Call) -> bool:
    keyword_names = {keyword.arg for keyword in node.keywords if keyword.arg}
    if name in {
        "list_scene_descriptors",
        "load_scene_from_library",
        "get_template_entry",
        "load_template_from_library",
    }:
        return "mode_id" in keyword_names
    if name in {"master_selector_options", "list_masters", "default_master"}:
        return "mode_id" in keyword_names or bool(node.args)
    if name == "get_master":
        return "mode_id" in keyword_names or len(node.args) >= 2
    return True


def test_runtime_config_library_has_no_flat_or_legacy_resource_roots():
    forbidden_names = {
        "LEGACY_SCENE_LIBRARY_DIR",
        "LEGACY_TEMPLATE_LIBRARY_DIR",
    }
    forbidden_path_parts = {"scenes", "templates"}
    violations: list[str] = []

    for path in sorted((ROOT / "src").rglob("*.py")):
        source = path.read_text(encoding="utf-8-sig")
        tree = ast.parse(source, filename=str(path))
        relative = path.relative_to(ROOT)
        for node in ast.walk(tree):
            name = (
                node.id
                if isinstance(node, ast.Name)
                else node.attr if isinstance(node, ast.Attribute) else ""
            )
            if name in forbidden_names:
                violations.append(f"{relative}:{node.lineno} {name}")
            if (
                isinstance(node, ast.BinOp)
                and isinstance(node.op, ast.Div)
                and isinstance(node.left, ast.Name)
                and node.left.id == "_PROJECT_ROOT"
                and isinstance(node.right, ast.Constant)
                and node.right.value in forbidden_path_parts
            ):
                violations.append(
                    f"{relative}:{node.lineno} project-relative {node.right.value} root"
                )
            if (
                isinstance(node, ast.Call)
                and _call_name(node.func) == "Path"
                and node.args
                and isinstance(node.args[0], ast.Constant)
                and node.args[0].value in forbidden_path_parts
            ):
                violations.append(
                    f"{relative}:{node.lineno} relative {node.args[0].value} root"
                )

    assert violations == []
    assert list((ROOT / "scenes").glob("*.json")) == []
    assert list((ROOT / "templates").glob("*.json")) == []


def test_config_library_ignores_unscoped_flat_files(monkeypatch, tmp_path):
    import src.config.library as library

    template_dir = tmp_path / "config_library" / "templates"
    scene_dir = tmp_path / "config_library" / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)
    template_dir.mkdir(parents=True)
    scene_dir.mkdir(parents=True)
    save_template(create_builtin_template("default"), template_dir / "flat.json")
    save_scene(SceneWorkspace(scene_id="flat"), scene_dir / "flat.json")

    assert all(entry.config_id != "flat" for entry in list_template_entries())
    assert all(entry.config_id != "flat" for entry in list_scene_descriptors())
    assert library.is_template_library_path(template_dir / "flat.json") is False
    assert library.is_scene_library_path(scene_dir / "flat.json") is False
    assert library.template_source_type_for_path(template_dir / "flat.json") == "external"
    assert library.scene_source_type_for_path(scene_dir / "flat.json") == "external"


def test_config_library_missing_ids_never_fall_back_to_runtime_factories(
    monkeypatch,
    tmp_path,
):
    import pytest
    import src.config.library as library

    monkeypatch.setattr(
        library,
        "TEMPLATE_LIBRARY_DIR",
        tmp_path / "config_library" / "templates",
    )
    monkeypatch.setattr(
        library,
        "SCENE_LIBRARY_DIR",
        tmp_path / "config_library" / "plans",
    )
    ensure_config_library()

    with pytest.raises(FileNotFoundError, match="plan_ref_unresolved"):
        load_scene_from_library("not_registered", mode_id="exam")
    with pytest.raises(FileNotFoundError, match="template_ref_unresolved"):
        load_template_from_library(
            "not_registered",
            mode_id="exam",
        )


def test_ui_config_library_calls_are_mode_scoped():
    protected_names = {
        "list_scene_descriptors",
        "load_scene_from_library",
        "get_template_entry",
        "load_template_from_library",
        "master_selector_options",
        "list_masters",
        "get_master",
        "default_master",
    }
    violations: list[str] = []

    for path in sorted((ROOT / "src" / "ui").rglob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            name = _call_name(node.func)
            if name not in protected_names or _has_mode_context(name, node):
                continue
            violations.append(f"{path.relative_to(ROOT)}:{node.lineno} {name}")

    assert violations == []


def test_ui_exam_master_surface_does_not_write_master_files_directly():
    protected_names = {
        "create_exam_blank_master_copy",
        "import_exam_blank_master_docx",
    }
    checked_calls = 0
    violations: list[str] = []

    for relative_path in (
        Path("src/ui/panels/scene_panel.py"),
        Path("src/ui/panels/scene_exam_detail.py"),
    ):
        path = ROOT / relative_path
        tree = ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            name = _call_name(node.func)
            if name not in protected_names:
                continue
            checked_calls += 1
            output_keyword = next(
                (keyword for keyword in node.keywords if keyword.arg == "output_dir"),
                None,
            )
            output_value = getattr(output_keyword, "value", None)
            if not (
                isinstance(output_value, ast.Name)
                and output_value.id == "USER_EXAM_MASTER_DIR"
            ):
                violations.append(f"{relative_path}:{node.lineno} {name}")

    assert checked_calls == 0
    assert violations == []


def test_ui_exam_master_inventory_calls_use_only_explicit_active_pool_dir():
    protected_names = {
        "sync_user_exam_blank_master_files",
        "audit_exam_user_master_pool",
    }
    checked_calls = 0
    active_calls = 0
    violations: list[str] = []

    for relative_path in (
        Path("src/ui/panels/scene_panel.py"),
        Path("src/ui/panels/scene_exam_detail.py"),
    ):
        path = ROOT / relative_path
        tree = ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            name = _call_name(node.func)
            if name not in protected_names:
                continue
            checked_calls += 1
            if name == "sync_user_exam_blank_master_files":
                directory_keyword = next(
                    (keyword for keyword in node.keywords if keyword.arg == "directory"),
                    None,
                )
                directory_value = (
                    directory_keyword.value
                    if directory_keyword is not None
                    else (node.args[1] if len(node.args) >= 2 else None)
                )
            else:
                directory_keyword = next(
                    (keyword for keyword in node.keywords if keyword.arg == "user_master_dir"),
                    None,
                )
                directory_value = (
                    directory_keyword.value if directory_keyword is not None else None
                )
            directory_name = (
                directory_value.id if isinstance(directory_value, ast.Name) else ""
            )
            if directory_name == "USER_EXAM_MASTER_DIR":
                active_calls += 1
                continue
            violations.append(f"{relative_path}:{node.lineno} {name}")

    assert checked_calls >= 1
    assert active_calls >= 1
    assert violations == []


def test_config_library_seeds_default_scene_and_template_entries():
    ensure_config_library()

    template_entry = default_template_entry()
    scene_entry = default_scene_entry()

    assert template_entry is not None
    assert template_entry.path.exists()
    assert scene_entry is not None
    assert scene_entry.path.exists()

    scene = load_scene_from_library(scene_entry.config_id)
    assert scene.scene_id == scene_entry.config_id
    assert scene.template_id
    assert scene.template_id in scene.compatible_template_ids


def test_config_library_seeds_mode_scoped_templates_and_plans(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()

    assert (template_dir / "custom" / "builtin" / "default.json").exists()
    assert (template_dir / "exam" / "builtin" / "default.json").exists()
    assert (template_dir / "thesis" / "builtin" / "thesis_gbt.json").exists()
    assert (template_dir / "official" / "builtin" / "official_gbt.json").exists()
    assert (scene_dir / "exam" / "builtin" / "exam.json").exists()
    assert (scene_dir / "exam" / "builtin" / "exam_quiz.json").exists()
    assert (scene_dir / "exam" / "builtin" / "exam_term.json").exists()
    assert (scene_dir / "official" / "builtin" / "official.json").exists()


def test_template_entries_are_filtered_by_work_mode(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()

    assert {entry.config_id for entry in list_template_entries(mode_id="exam")} == {"default"}
    assert {entry.config_id for entry in list_template_entries(mode_id="thesis")} == {
        "thesis_custom",
        "thesis_gbt",
    }
    assert {entry.config_id for entry in list_template_entries(mode_id="official")} == {
        "official_gbt",
        "official_custom",
    }
    assert "official_gbt" not in {
        entry.config_id for entry in list_template_entries(mode_id="exam")
    }


def test_save_template_to_library_writes_mode_scoped_user_dir(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    template = create_builtin_template("default")
    template.name = "Exam Custom Format"

    entry = library.save_template_to_library(
        template,
        template_id="school_exam_format",
        mode_id="exam",
    )

    saved_path = template_dir / "exam" / "user" / "school_exam_format.json"
    assert entry.mode_id == "exam"
    assert entry.path == saved_path
    assert saved_path.exists()
    assert not (template_dir / "custom" / "user" / "school_exam_format.json").exists()
    assert json.loads(saved_path.read_text(encoding="utf-8"))["name"] == "Exam Custom Format"

    exam_entry = get_template_entry("school_exam_format", mode_id="exam")
    custom_entry = get_template_entry("school_exam_format", mode_id="custom")
    exam_template = load_template_from_library("school_exam_format", mode_id="exam")
    assert exam_entry is not None
    assert exam_entry.path == saved_path
    assert custom_entry is None
    assert exam_template.name == "Exam Custom Format"


def test_mode_scoped_template_load_does_not_cross_same_id_defaults(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    custom_default = create_builtin_template("default")
    custom_default.name = "通用用户默认格式"
    exam_default = create_builtin_template("default")
    exam_default.name = "试卷用户默认格式"
    save_template(
        custom_default,
        template_dir / "custom" / "user" / "school_default.json",
    )
    save_template(
        exam_default,
        template_dir / "exam" / "user" / "school_default.json",
    )

    custom_entry = get_template_entry("school_default", mode_id="custom")
    exam_entry = get_template_entry("school_default", mode_id="exam")
    custom_template = load_template_from_library(
        "school_default",
        mode_id="custom",
    )
    exam_template = load_template_from_library("school_default", mode_id="exam")

    assert custom_entry is not None
    assert exam_entry is not None
    assert custom_entry.path == template_dir / "custom" / "user" / "school_default.json"
    assert exam_entry.path == template_dir / "exam" / "user" / "school_default.json"
    assert custom_template.name == "通用用户默认格式"
    assert exam_template.name == "试卷用户默认格式"
    assert default_template_entry("exam").config_id == "default"


def test_mode_scoped_user_resources_cannot_shadow_same_id_builtins(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    user_template = create_builtin_template("default")
    user_template.name = "User Exam Default"
    save_template(user_template, template_dir / "exam" / "user" / "default.json")

    user_plan = SceneWorkspace(
        scene_id="exam",
        name="User Exam Plan",
        mode_id="exam",
        template_id="default",
        compatible_template_ids=["default"],
        master_id="default_exam",
    )
    save_scene(user_plan, scene_dir / "exam" / "user" / "exam.json")

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="template_id_ambiguous",
    ):
        list_template_entries(mode_id="exam")
    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="plan_id_ambiguous",
    ):
        list_scene_descriptors(mode_id="exam")
    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="config_id_ambiguous",
    ):
        load_template_from_library("default", mode_id="exam")
    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="config_id_ambiguous",
    ):
        load_scene_from_library("exam", mode_id="exam")


def test_save_apis_reject_user_ids_that_shadow_builtins(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", root / "templates")
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", root / "plans")
    ensure_config_library()

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="template_id_shadowed",
    ):
        library.save_template_to_library(
            create_builtin_template("default"),
            template_id="default",
            mode_id="exam",
        )

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="plan_id_shadowed",
    ):
        library.save_scene_to_library(
            SceneWorkspace(
                scene_id="exam",
                name="Shadowing exam plan",
                mode_id="exam",
                template_id="default",
                compatible_template_ids=["default"],
                master_id="default_exam",
            ),
            scene_id="exam",
            mode_id="exam",
        )


def test_user_same_id_resources_fail_before_selector_label_projection(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    user_template = create_builtin_template("default")
    user_template.name = ""
    save_template(user_template, template_dir / "exam" / "user" / "default.json")
    save_scene(
        SceneWorkspace(
            scene_id="exam",
            name="",
            description="",
            mode_id="exam",
            template_id="default",
            compatible_template_ids=["default"],
            master_id="default_exam",
        ),
        scene_dir / "exam" / "user" / "exam.json",
    )

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="plan_id_ambiguous",
    ):
        plan_selector_options("exam")
    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="config_id_ambiguous",
    ):
        template_selector_options(
            "exam",
            template_ids=["default"],
            fallback_labels={"default": "Built-in Exam Default"},
            include_source_prefix=False,
        )


def test_broken_user_same_id_template_is_rejected_as_ambiguous(
    monkeypatch,
):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    user_path = template_dir / "exam" / "user" / "default.json"
    user_path.parent.mkdir(parents=True, exist_ok=True)
    user_path.write_text("{broken", encoding="utf-8")

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="template_id_ambiguous",
    ):
        list_template_entries(mode_id="exam")
    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="config_id_ambiguous",
    ):
        get_template_entry("default", mode_id="exam")
    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="config_id_ambiguous",
    ):
        template_selector_options("exam", template_ids=["default"])
    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="config_id_ambiguous",
    ):
        load_template_from_library("default", mode_id="exam")


def test_broken_user_same_id_plan_is_rejected_as_ambiguous(
    monkeypatch,
):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    user_path = scene_dir / "exam" / "user" / "exam.json"
    user_path.parent.mkdir(parents=True, exist_ok=True)
    user_path.write_text("{broken", encoding="utf-8")

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="plan_id_ambiguous",
    ):
        list_scene_descriptors(mode_id="exam")
    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="plan_id_ambiguous",
    ):
        plan_selector_options("exam")


def test_unscoped_plan_descriptors_keep_same_id_in_each_work_mode(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    for mode_id, name, template_id, master_id in (
        ("exam", "Exam Shared", "default", "default_exam"),
        ("official", "Official Shared", "official_gbt", "official_default"),
    ):
        save_scene(
            SceneWorkspace(
                scene_id="shared",
                name=name,
                mode_id=mode_id,
                template_id=template_id,
                compatible_template_ids=[template_id],
                master_id=master_id,
            ),
            scene_dir / mode_id / "user" / "shared.json",
        )

    matching = {
        (item.mode_id, item.config_id, item.name)
        for item in list_scene_descriptors()
        if item.config_id == "shared"
    }

    assert matching == {
        ("exam", "shared", "Exam Shared"),
        ("official", "shared", "Official Shared"),
    }


def test_mode_scoped_plan_loading_rejects_cross_mode_builtin_template_ids(monkeypatch):
    import pytest
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    stale_plan = SceneWorkspace(
        scene_id="exam_cross_mode_template",
        name="Exam Cross Mode Template",
        mode_id="exam",
        template_id="official_gbt",
        compatible_template_ids=["official_gbt", "default"],
        master_id="default_exam",
    )
    stale_plan_path = scene_dir / "exam" / "user" / "exam_cross_mode_template.json"
    save_scene(stale_plan, stale_plan_path)

    descriptor = get_scene_descriptor(
        "exam_cross_mode_template",
        mode_id="exam",
    )
    assert descriptor is not None
    assert descriptor.is_available is False
    assert "plan_template_ref_unresolved" in descriptor.load_error

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="plan_template_ref_unresolved",
    ):
        load_scene_from_library(
            "exam_cross_mode_template",
            mode_id="exam",
        )
    with pytest.raises(FileNotFoundError, match="template_ref_unresolved"):
        load_template_from_library("official_gbt", mode_id="exam")
    loaded_template = load_template_from_library("default", mode_id="exam")
    official_loaded_template = load_template_from_library(
        "official_gbt",
        mode_id="official",
    )

    assert loaded_template.name
    assert official_loaded_template.name


def test_scene_descriptors_are_library_backed_and_keep_builtin_order():
    ensure_config_library()

    descriptors = list_scene_descriptors()
    builtin_ids = {"custom", "exam", "thesis", "bidding", "official", "technical", "report"}
    builtin_descriptors = [
        descriptor for descriptor in descriptors if descriptor.config_id in builtin_ids
    ]

    assert [descriptor.config_id for descriptor in builtin_descriptors[:4]] == ["custom", "exam", "thesis", "bidding"]
    assert builtin_descriptors[0].scene_id == "custom"
    assert builtin_descriptors[0].display_name == "自定义"
    assert builtin_descriptors[0].path.exists()
    exam = next(descriptor for descriptor in descriptors if descriptor.config_id == "exam")
    assert exam.display_name == "默认试卷"
    exam_descriptors = list_scene_descriptors(mode_id="exam")
    assert [descriptor.config_id for descriptor in exam_descriptors[:3]] == [
        "exam",
        "exam_quiz",
        "exam_term",
    ]
    assert [descriptor.display_name for descriptor in exam_descriptors[:3]] == [
        "默认试卷",
        "随堂测验",
        "期中期末",
    ]
    quiz = load_scene_from_library("exam_quiz", mode_id="exam")
    term = load_scene_from_library("exam_term", mode_id="exam")
    assert quiz.master_id == "default_exam"
    assert term.master_id == "default_exam"
    thesis = next(descriptor for descriptor in descriptors if descriptor.config_id == "thesis")
    assert thesis.description == "中文学位论文·课程论文·综述·开题"
    assert "期刊论文" not in thesis.description


def test_selector_projections_distinguish_plans_templates_and_masters(tmp_path):
    ensure_config_library()

    plan_options = plan_selector_options("exam", include_source_prefix=True)
    assert [option.value for option in plan_options[:3]] == [
        "exam",
        "exam_quiz",
        "exam_term",
    ]
    assert [option.label for option in plan_options[:3]] == [
        "内置方案：默认试卷方案",
        "内置方案：随堂测验方案",
        "内置方案：期中期末方案",
    ]

    template_options = template_selector_options("exam")
    assert [option.value for option in template_options] == ["default"]
    assert template_options[0].label.startswith("内置模板：")

    master_options = master_selector_options("exam", user_master_dir=tmp_path)
    assert [option.value for option in master_options] == ["default_exam"]
    assert [option.label for option in master_options] == ["A4 标准卷面"]
    assert {
        option.label.removeprefix("内置方案：")
        for option in plan_options[:3]
    }.isdisjoint({option.label for option in master_options[:3]})

    official_master_options = master_selector_options("official")
    assert [option.value for option in official_master_options] == [
        "official_gbt_standard",
        "official_gbt_upward",
        "official_gbt_letter",
        "official_gbt_minutes",
        "official_gbt_order",
    ]
    assert official_master_options[0].label == "GB/T 9704 通用红头公文版式"

    for unbacked_mode in ("custom", "thesis", "bidding", "technical", "report"):
        assert master_selector_options(unbacked_mode, user_master_dir=tmp_path) == ()


def test_master_display_label_uses_specific_or_neutral_nouns():
    official_master = MasterSpec(
        master_id="official_custom",
        mode_id="official",
        label="公文母版",
        family="official",
        source_type="builtin",
        docx_path=ROOT / "missing.docx",
    )
    generic_master = MasterSpec(
        master_id="future_shell",
        mode_id="future",
        label="默认",
        family="future",
        source_type="builtin",
        docx_path=ROOT / "missing.docx",
    )

    assert master_display_label(official_master) == "公文版式"
    assert master_display_label(generic_master) == "默认装配项"
    assert strip_source_prefix("内置装配项：默认装配项") == "默认装配项"


def test_selector_projections_keep_material_packages_mode_scoped(tmp_path):
    ensure_config_library()

    exam_plan_values = {
        option.value
        for option in plan_selector_options("exam", include_source_prefix=True)
    }
    official_plan_values = {
        option.value
        for option in plan_selector_options("official", include_source_prefix=True)
    }
    assert {"official"}.isdisjoint(exam_plan_values)
    assert {"exam", "exam_quiz", "exam_term"}.isdisjoint(official_plan_values)

    exam_template_values = {
        option.value
        for option in template_selector_options(
            "exam",
            template_ids=["default", "official_gbt", "thesis_gbt"],
        )
    }
    official_template_values = {
        option.value
        for option in template_selector_options(
            "official",
            template_ids=["official_gbt", "default", "thesis_gbt"],
        )
    }
    assert exam_template_values == {"default"}
    assert official_template_values == {"official_gbt"}

    exam_master_values = {
        option.value
        for option in master_selector_options("exam", user_master_dir=tmp_path)
    }
    official_master_values = {
        option.value
        for option in master_selector_options("official")
    }
    assert exam_master_values == {"default_exam"}
    assert official_master_values == {
        "official_gbt_standard",
        "official_gbt_upward",
        "official_gbt_letter",
        "official_gbt_minutes",
        "official_gbt_order",
    }
    assert official_master_values.isdisjoint(exam_master_values)

    official_material_options = material_package_sample_selector_options(
        "official",
        include_source_prefix=True,
    )
    official_materials = {option.value: option for option in official_material_options}
    assert {
        "builtin/letter_material_request",
        "builtin/notice_archive_check",
        "builtin/minutes_coordination",
    } <= set(official_materials)
    letter = official_materials["builtin/letter_material_request"]
    assert letter.label == "内置资料包：函件资料包样例"
    assert letter.source_type == "builtin"
    assert strip_source_prefix(letter.label) == "函件资料包样例"
    notice = official_materials["builtin/notice_archive_check"]
    assert notice.label == "内置资料包：通知资料包样例"
    assert notice.source_type == "builtin"
    assert "material_packages" in notice.tooltip
    assert "official" in notice.tooltip
    assert "builtin" in notice.tooltip
    assert strip_source_prefix(notice.label) == "通知资料包样例"

    for mode_id in ("custom", "exam", "thesis", "bidding", "technical", "report"):
        assert material_package_sample_selector_options(mode_id) == ()


def test_selector_projections_expose_only_mode_owned_builtin_resources(tmp_path):
    ensure_config_library()

    expected_plans = {
        "custom": {"custom"},
        "exam": {"exam", "exam_quiz", "exam_term"},
        "thesis": {"thesis"},
        "bidding": {"bidding"},
        "official": {"official"},
        "technical": {"technical"},
        "report": {"report"},
    }
    expected_templates = {
        "custom": {"default"},
        "exam": {"default"},
        "thesis": {"thesis_custom", "thesis_gbt"},
        "bidding": {"bid_custom", "bid_engineering", "bid_procurement"},
        "official": {"official_custom", "official_gbt"},
        "technical": {"tech_custom", "tech_standard"},
        "report": {"report_custom", "report_default"},
    }
    expected_masters = {
        "custom": set(),
        "exam": {"default_exam"},
        "thesis": set(),
        "bidding": set(),
        "official": {
            "official_gbt_standard",
            "official_gbt_upward",
            "official_gbt_letter",
            "official_gbt_minutes",
            "official_gbt_order",
        },
        "technical": set(),
        "report": set(),
    }
    expected_material_packages = {
        "custom": set(),
        "exam": set(),
        "thesis": set(),
        "bidding": set(),
        "official": {
            "builtin/approval_archive_system",
            "builtin/letter_material_request",
            "builtin/minutes_coordination",
            "builtin/notice_archive_check",
            "builtin/report_work_summary",
            "builtin/request_archive_system",
        },
        "technical": set(),
        "report": set(),
    }

    for mode in list_work_modes():
        mode_id = mode.mode_id
        master_user_dir = tmp_path / mode_id / "masters"

        plan_values = {
            option.value
            for option in plan_selector_options(mode_id)
            if option.source_type == "builtin"
        }
        template_values = {
            option.value
            for option in template_selector_options(mode_id)
            if option.source_type == "builtin"
        }
        master_values = {
            option.value
            for option in master_selector_options(mode_id, user_master_dir=master_user_dir)
            if option.source_type == "builtin"
        }
        material_values = {
            option.value
            for option in material_package_sample_selector_options(mode_id)
            if option.source_type == "builtin"
        }

        assert plan_values == expected_plans[mode_id]
        assert template_values == expected_templates[mode_id]
        assert master_values == expected_masters[mode_id]
        assert material_values == expected_material_packages[mode_id]


def test_template_selector_options_stay_within_work_mode_when_ids_are_stale():
    ensure_config_library()

    exam_values = [
        option.value
        for option in template_selector_options(
            "exam",
            template_ids=["default", "official_gbt", "thesis_gbt"],
        )
    ]
    official_values = [
        option.value
        for option in template_selector_options(
            "official",
            template_ids=["official_gbt", "default", "thesis_gbt"],
        )
    ]

    assert exam_values == ["default"]
    assert official_values == ["official_gbt"]


def test_exam_master_selector_exposes_valid_docx_from_active_user_folder(tmp_path):
    user_dir = tmp_path / "config_library" / "masters" / "exam" / "user"
    user_dir.mkdir(parents=True)

    stale_program_copy = user_dir / "user_default_exam_copy_999.docx"
    manual_copy = user_dir / "学校手工卷面.docx"
    for path in (stale_program_copy, manual_copy):
        document = Document()
        document.add_paragraph("{{af_title}}")
        document.add_paragraph("{{af_questions}}")
        document.add_paragraph("{{af_answer_area}}")
        document.save(path)

    options = master_selector_options("exam", user_master_dir=user_dir)
    values = [option.value for option in options]
    labels = [option.label for option in options]

    assert sum(value.startswith("user_file_") for value in values) == 2
    assert "user_default_exam_copy_999" in labels
    assert "学校手工卷面" in labels


def test_official_plan_selector_exposes_one_base_plan_and_six_document_types():
    ensure_config_library()

    official_seed = json.loads(
        (
            ROOT
            / "config_library"
            / "plans"
            / "official"
            / "builtin"
            / "official.json"
        ).read_text(encoding="utf-8")
    )
    assert "exam_paper" in official_seed

    plan_options = tuple(
        option
        for option in plan_selector_options(
            "official",
            include_source_prefix=True,
        )
        if option.source_type == "builtin"
    )
    values = [option.value for option in plan_options]
    labels = {option.value: option.label for option in plan_options}

    assert values == ["official"]
    assert labels["official"] == "内置方案：公文基础方案"
    assert [
        profile.profile_id for profile in list_common_official_document_profiles()
    ] == ["notice", "letter", "minutes", "report", "request", "approval"]

    default_scene = load_scene_from_library("official", mode_id="official")

    assert default_scene.default_material_profile_id == "official:notice"
    assert default_scene.input_source_profile.material_schema_ids == [
        "official_document_v1"
    ]
    assert default_scene.input_source_profile.required_material_fields == [
        "title",
        "body",
        "organization",
        "document_no",
        "issue_date",
    ]
def test_official_document_type_is_one_shared_task_state_across_panels():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official")
    workbench = WorkbenchPanel(bridge)
    scene_panel = ScenePanel(bridge)
    try:
        scene_panel._ensure_detail_loaded("scn_content")
        app.processEvents()

        bridge.set_current_official_document_type_id("letter")
        app.processEvents()

        assert bridge.current_official_document_type_id() == "letter"
        assert workbench._quick_execution_detail.official_document_type_id() == "letter"
        assert scene_panel._content._official_material_profile.currentData() == "letter"

        minutes_index = workbench._quick_execution_detail._document_type_combo.findData(
            "minutes"
        )
        workbench._quick_execution_detail._document_type_combo.setCurrentIndex(
            minutes_index
        )
        app.processEvents()

        assert bridge.current_official_document_type_id() == "minutes"
        assert scene_panel._content._official_material_profile.currentData() == "minutes"
        assert bridge.current_scene_id() == "official"
    finally:
        scene_panel.close()
        workbench.close()
        app.processEvents()


def test_scene_panel_user_plan_activation_broadcasts_to_workbench():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("official")
    workbench = WorkbenchPanel(bridge)
    scene_panel = ScenePanel(bridge)
    received: list[str] = []
    bridge.scene_changed.connect(
        lambda scene: received.append(str(getattr(scene, "scene_id", "") or ""))
    )
    try:
        scene = load_scene_from_library("official", mode_id="official")
        descriptor = get_scene_descriptor("official", mode_id="official")
        scene_panel._activate_scene(
            scene,
            scene_id="official",
            path=str(descriptor.path),
            source="library",
            source_type=descriptor.source_type,
            dirty=False,
        )
        app.processEvents()

        assert received[-1] == "official"
        assert workbench._current_scene is scene
        assert workbench._quick_execution_detail.current_scene() == scene
        assert workbench._quick_execution_detail.current_scene() is not scene
    finally:
        scene_panel.close()
        workbench.close()
        app.processEvents()
def test_official_plan_entry_decisions_gate_profile_promotion():
    ensure_config_library()

    profile_ids = {
        profile.profile_id for profile in list_official_document_profiles()
    }
    decisions = list_official_document_plan_entry_decisions()
    decisions_by_profile = {
        decision.profile_id: decision for decision in decisions
    }
    builtin_decisions = {
        decision.profile_id: decision.plan_id
        for decision in decisions
        if decision.entry_kind == OFFICIAL_PLAN_ENTRY_BUILTIN
    }
    candidate_profile_ids = {
        decision.profile_id
        for decision in decisions
        if decision.entry_kind == OFFICIAL_PLAN_ENTRY_CANDIDATE
    }
    profile_only_ids = {
        decision.profile_id
        for decision in decisions
        if decision.entry_kind == OFFICIAL_PLAN_ENTRY_PROFILE_ONLY
    }
    official_builtin_plan_ids = {
        path.stem
        for path in (
            ROOT / "config_library" / "plans" / "official" / "builtin"
        ).glob("*.json")
    }

    assert set(decisions_by_profile) == profile_ids
    assert len(decisions) == len(profile_ids)
    assert builtin_decisions == {
        "notice": "official",
    }
    assert list_official_document_builtin_plan_profile_ids() == ("notice",)
    assert candidate_profile_ids == set()
    assert profile_only_ids == profile_ids - set(builtin_decisions) - candidate_profile_ids
    assert official_builtin_plan_ids == {"official"}

    for profile_id, plan_id in builtin_decisions.items():
        scene = load_scene_from_library(plan_id, mode_id="official")
        assert scene.default_material_profile_id == f"official:{profile_id}"
    for decision in decisions:
        if decision.entry_kind != OFFICIAL_PLAN_ENTRY_BUILTIN:
            assert decision.plan_id == ""


def test_official_plan_selector_options_all_derive_material_schema_contracts():
    ensure_config_library()

    options = [
        option
        for option in plan_selector_options("official", include_source_prefix=True)
        if not option.disabled
    ]

    assert options
    for option in options:
        scene = load_scene_from_library(option.value, mode_id="official")
        profile = scene.input_source_profile
        assert profile.material_schema_ids, option.value
        assert profile.material_schema_id == profile.material_schema_ids[0]
        assert "title" in profile.required_material_fields
        assert "body" in profile.required_material_fields


def test_all_builtin_plan_sources_are_structurally_complete_for_exam_paper():
    ensure_config_library()

    plan_root = ROOT / "config_library" / "plans"
    exam_sources: set[str] = set()

    for path in sorted(plan_root.glob("*/builtin/*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        mode_id = path.relative_to(plan_root).parts[0]
        assert "exam_paper" in payload, str(path.relative_to(ROOT))
        if mode_id == "exam":
            exam_sources.add(path.name)

    assert {"exam.json", "exam_quiz.json", "exam_term.json"} <= exam_sources


def test_builtin_plan_sources_are_self_describing_and_reference_mode_owned_templates():
    ensure_config_library()

    plan_root = ROOT / "config_library" / "plans"
    template_root = ROOT / "config_library" / "templates"
    checked_plan_ids: set[str] = set()

    for path in sorted(plan_root.glob("*/builtin/*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        mode_id = path.relative_to(plan_root).parts[0]
        loaded = load_scene(path)
        template_ids = {
            str(template_id or "").strip()
            for template_id in (
                payload.get("template_id"),
                *list(payload.get("compatible_template_ids", []) or []),
            )
            if str(template_id or "").strip()
        }

        assert payload.get("mode_id") == mode_id, str(path.relative_to(ROOT))
        assert loaded.mode_id == mode_id, str(path.relative_to(ROOT))
        assert template_ids, str(path.relative_to(ROOT))
        for template_id in template_ids:
            assert (
                template_root / mode_id / "builtin" / f"{template_id}.json"
            ).is_file(), f"{path.relative_to(ROOT)} -> {template_id}"
        checked_plan_ids.add(path.stem)

    assert {
        "custom",
        "exam",
        "exam_quiz",
        "exam_term",
        "thesis",
        "bidding",
        "official",
        "technical",
        "report",
    } <= checked_plan_ids


def test_scene_descriptors_put_user_scenes_before_builtin_scenes(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    scene = SceneWorkspace(
        scene_id="exam_default_copy",
        name="试卷-期中",
        mode_id="exam",
        template_id="default",
        compatible_template_ids=["default"],
        master_id="default_exam",
    )
    library.save_scene_to_library(scene, scene_id="exam_default_copy", mode_id="exam")

    descriptors = list_scene_descriptors(mode_id="exam")
    builtin_ids = {"custom", "exam", "thesis", "bidding", "official", "technical", "report"}

    assert descriptors[0].config_id == "exam_default_copy"
    assert descriptors[0].display_name == "试卷-期中"
    assert descriptors[0].mode_id == "exam"
    assert next(
        index for index, descriptor in enumerate(descriptors)
        if descriptor.config_id in builtin_ids
    ) > 0
    saved_path = scene_dir / "exam" / "user" / "exam_default_copy.json"
    assert saved_path.exists()
    saved_payload = json.loads(saved_path.read_text(encoding="utf-8"))
    assert saved_payload["mode_id"] == "exam"

    loaded = load_scene_from_library("exam_default_copy", mode_id="exam")
    assert loaded.mode_id == "exam"
    assert loaded.scene_id == "exam_default_copy"


def test_save_scene_to_library_uses_scene_mode_when_argument_missing(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    scene = SceneWorkspace(
        scene_id="school_midterm_plan",
        name="School Midterm Plan",
        mode_id="exam",
        template_id="default",
        compatible_template_ids=["default"],
        master_id="default_exam",
    )

    entry = library.save_scene_to_library(scene, scene_id="school_midterm_plan")

    saved_path = scene_dir / "exam" / "user" / "school_midterm_plan.json"
    assert entry.mode_id == "exam"
    assert entry.path == saved_path
    assert saved_path.exists()
    assert not (scene_dir / "custom" / "user" / "school_midterm_plan.json").exists()
    saved_payload = json.loads(saved_path.read_text(encoding="utf-8"))
    assert saved_payload["mode_id"] == "exam"

    exam_descriptors = list_scene_descriptors(mode_id="exam")
    custom_descriptors = list_scene_descriptors(mode_id="custom")
    assert exam_descriptors[0].config_id == "school_midterm_plan"
    assert exam_descriptors[0].mode_id == "exam"
    assert "school_midterm_plan" not in {
        descriptor.config_id for descriptor in custom_descriptors
    }

    loaded = load_scene_from_library("school_midterm_plan", mode_id="exam")
    assert loaded.mode_id == "exam"
    assert loaded.scene_id == "school_midterm_plan"


def test_scene_descriptors_keep_broken_scene_visible_with_error(monkeypatch):
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    broken_path = scene_dir / "custom" / "user" / "broken_scene.json"
    broken_path.parent.mkdir(parents=True, exist_ok=True)
    broken_path.write_text("{not-valid-json", encoding="utf-8")

    descriptors = list_scene_descriptors()
    broken = next(descriptor for descriptor in descriptors if descriptor.config_id == "broken_scene")

    assert broken.is_available is False
    assert broken.load_error
    assert "Unavailable" in broken.display_name

    fetched = get_scene_descriptor("broken_scene")
    assert fetched is not None
    assert fetched.is_available is False

    default_descriptor = default_scene_descriptor()
    assert default_descriptor is not None
    assert default_descriptor.is_available is True


def test_load_scene_from_library_raises_for_broken_scene_file(monkeypatch):
    import pytest
    import src.config.library as library

    root = Path(tempfile.mkdtemp())
    template_dir = root / "templates"
    scene_dir = root / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)

    ensure_config_library()
    broken_path = scene_dir / "custom" / "user" / "broken_scene.json"
    broken_path.parent.mkdir(parents=True, exist_ok=True)
    broken_path.write_text("{not-valid-json", encoding="utf-8")

    with pytest.raises(Exception):
        load_scene_from_library("broken_scene")


def test_panel_bridge_tracks_scene_and_template_context_metadata():
    bridge = PanelBridge()
    scene = SceneWorkspace(scene_id="custom", template_id="default")
    template = TemplateConfig(name="Default Template")

    seen_scenes = []
    seen_templates = []
    bridge.scene_changed.connect(seen_scenes.append)
    bridge.template_changed.connect(seen_templates.append)

    bridge.set_current_scene(
        scene,
        config_id="custom",
        path="C:/configs/scenes/custom.json",
        source="library",
        source_type="user",
    )
    bridge.set_current_template(
        template,
        config_id="default",
        path="C:/configs/templates/default.json",
        source="library",
        source_type="builtin",
    )

    assert bridge.current_scene() is scene
    assert bridge.current_scene_id() == "custom"
    assert bridge.current_scene_path().endswith("custom.json")
    assert bridge.current_scene_source() == "library"
    assert bridge.current_scene_source_type() == "user"
    assert bridge.current_template() == template
    assert bridge.current_template() is not template
    assert bridge.current_template_id() == "default"
    assert bridge.current_template_path().endswith("default.json")
    assert bridge.current_template_source() == "library"
    assert bridge.current_template_source_type() == "builtin"
    assert seen_scenes == [scene]
    assert seen_templates == [template]


def test_panel_bridge_preserves_source_type_only_for_the_same_context_identity():
    bridge = PanelBridge()
    first_scene = SceneWorkspace(scene_id="exam", template_id="default")
    edited_scene = SceneWorkspace(scene_id="exam", template_id="default")

    bridge.set_current_scene(
        first_scene,
        config_id="exam",
        path="C:/configs/plans/exam/user/exam.json",
        source="library",
        source_type="user",
        emit_signal=False,
    )
    bridge.set_current_scene(
        edited_scene,
        config_id="exam",
        path="C:/configs/plans/exam/user/exam.json",
        source="library",
        emit_signal=False,
    )
    assert bridge.current_scene_source_type() == "user"

    bridge.set_current_scene(
        SceneWorkspace(scene_id="exam_quiz", template_id="default"),
        config_id="exam_quiz",
        path="C:/configs/plans/exam/builtin/exam_quiz.json",
        source="library",
        emit_signal=False,
    )
    assert bridge.current_scene_source_type() == ""

    bridge.set_current_template(
        TemplateConfig(name="External"),
        config_id="external",
        path="C:/tmp/external.json",
        source="file",
        source_type="external",
        emit_signal=False,
    )
    bridge.set_current_template(
        TemplateConfig(name="External Edited"),
        config_id="external",
        path="C:/tmp/external.json",
        source="file",
        emit_signal=False,
    )
    assert bridge.current_template_source_type() == "external"

    bridge.set_current_template(
        TemplateConfig(name="Default"),
        config_id="default",
        path="C:/configs/templates/exam/builtin/default.json",
        source="library",
        emit_signal=False,
    )
    assert bridge.current_template_source_type() == ""


def test_panel_bridge_protects_dirty_template_commits_from_navigation_reload():
    bridge = PanelBridge()
    path = "C:/configs/templates/custom/user/default.json"
    committed = TemplateConfig(name="Saved baseline")
    incoming_disk = TemplateConfig(name="Unexpected disk reload")

    bridge.set_current_work_mode("custom", emit_signal=False)
    bridge.set_current_template(
        committed,
        config_id="default",
        path=path,
        source="library",
        source_type="user",
        emit_signal=False,
    )
    bridge.replace_protected_template_commits(
        [("custom", "default", path, committed)]
    )
    bridge.mark_template_dirty()

    assert (
        bridge.set_current_template(
            incoming_disk,
            config_id="default",
            path=path,
            source="library",
            source_type="user",
            emit_signal=False,
        )
        is False
    )
    assert bridge.current_template().name == "Saved baseline"

    bridge.set_current_work_mode("exam", emit_signal=False)
    bridge.set_current_template(
        TemplateConfig(name="Exam"),
        config_id="default",
        path="C:/configs/templates/exam/builtin/default.json",
        source="library",
        source_type="builtin",
        emit_signal=False,
    )
    bridge.set_current_work_mode("custom", emit_signal=False)

    assert bridge.set_current_template(
        incoming_disk,
        config_id="default",
        path=path,
        source="library",
        source_type="user",
        emit_signal=False,
    )
    assert bridge.current_template().name == "Saved baseline"


def test_workbench_quick_context_does_not_alias_bridge_scene_and_keeps_dirty_state():
    app = _app()
    bridge = PanelBridge()
    panel = WorkbenchPanel(bridge)
    try:
        quick_scene = panel._quick_execution_detail.current_scene()
        bridge_scene = bridge.current_scene()
        assert quick_scene is not bridge_scene

        quick_scene.module_switches["heading_numbering"] = False
        assert bridge.current_scene().module_switches.get("heading_numbering") is not False

        bridge.update_current_scene_module_switches({"heading_numbering": False})
        app.processEvents()
        assert bridge.is_scene_dirty() is True

        panel._on_quick_binding_changed(
            panel._quick_execution_detail.current_scene(),
            bridge.current_template_id(),
        )
        assert bridge.is_scene_dirty() is True
        assert bridge.current_scene().module_switches["heading_numbering"] is False
    finally:
        panel.close()
        app.processEvents()


def test_workbench_bootstraps_and_updates_bridge_binding_from_quick_execution():
    app = _app()
    bridge = PanelBridge()
    panel = WorkbenchPanel(bridge)
    try:
        assert bridge.current_scene() is not None
        assert bridge.current_template() is not None
        assert bridge.current_scene_id() == panel._quick_execution_detail.current_scene_id()
        assert bridge.current_template_id() == panel._quick_execution_detail.current_template_id()

        assert panel._quick_execution_detail._scene_combo.findData("thesis") < 0

        bridge.set_current_work_mode("thesis")
        app.processEvents()
        thesis_index = panel._quick_execution_detail._scene_combo.findData("thesis")
        assert thesis_index >= 0
        assert (
            panel._quick_execution_detail._scene_combo.itemText(thesis_index)
            == "论文排版方案"
        )
        assert panel._quick_execution_detail._scene_combo.itemData(
            thesis_index,
            SOURCE_BADGE_TEXT_ROLE,
        ) == "内置"
        thesis_template_index = panel._quick_execution_detail._template_combo.findData(
            "thesis_gbt"
        )
        assert thesis_template_index >= 0
        assert not panel._quick_execution_detail._template_combo.itemText(
            thesis_template_index
        ).startswith("内置模板：")
        assert panel._quick_execution_detail._template_combo.itemData(
            thesis_template_index,
            SOURCE_BADGE_TEXT_ROLE,
        ) == "内置"
        panel._on_quick_binding_changed(
            panel._quick_execution_detail.current_scene(),
            panel._quick_execution_detail.current_template_id(),
        )

        assert bridge.current_scene_id() == "thesis"
        assert bridge.current_template_id() == "thesis_gbt"
    finally:
        panel.close()
        app.processEvents()


def test_workbench_quick_execution_plan_selector_matches_exam_mode():
    app = _app()
    bridge = PanelBridge()
    panel = WorkbenchPanel(bridge)
    try:
        bridge.set_current_work_mode("exam")
        app.processEvents()

        combo = panel._quick_execution_detail._scene_combo
        values = [
            str(combo.itemData(index) or "").strip()
            for index in range(combo.count())
        ]
        labels = [combo.itemText(index) for index in range(combo.count())]

        assert values[:3] == ["exam", "exam_quiz", "exam_term"]
        assert labels[:3] == ["默认试卷方案", "随堂测验方案", "期中期末方案"]
        assert [
            combo.itemData(index, SOURCE_BADGE_TEXT_ROLE)
            for index in range(3)
        ] == ["内置", "内置", "内置"]
        assert "official" not in values
    finally:
        panel.close()
        app.processEvents()


def test_workbench_quick_execution_rejects_same_id_user_plan(
    tmp_path,
    monkeypatch,
):
    import src.config.library as library

    scene_dir = tmp_path / "plans"
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)
    ensure_config_library()
    save_scene(
        SceneWorkspace(
            scene_id="exam",
            name="校级期中试卷",
            mode_id="exam",
            template_id="default",
            compatible_template_ids=["default"],
            master_id="default_exam",
        ),
        scene_dir / "exam" / "user" / "exam.json",
    )

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="plan_id_ambiguous",
    ):
        list_scene_descriptors(mode_id="exam")


def test_workbench_quick_execution_rejects_broken_same_id_user_template(
    tmp_path,
    monkeypatch,
):
    import src.config.library as library

    template_dir = tmp_path / "templates"
    scene_dir = tmp_path / "plans"
    monkeypatch.setattr(library, "TEMPLATE_LIBRARY_DIR", template_dir)
    monkeypatch.setattr(library, "SCENE_LIBRARY_DIR", scene_dir)
    ensure_config_library()
    user_path = template_dir / "exam" / "user" / "default.json"
    user_path.parent.mkdir(parents=True, exist_ok=True)
    user_path.write_text("{broken", encoding="utf-8")

    with pytest.raises(
        library.ConfigReferenceResolutionError,
        match="template_id_ambiguous",
    ):
        list_template_entries(mode_id="exam")


def test_workbench_and_scene_overview_plan_selectors_share_mode_options():
    app = _app()
    bridge = PanelBridge()
    workbench = WorkbenchPanel(bridge)
    scene_panel = ScenePanel(bridge)
    try:
        for mode_id, forbidden_values in (
            (
                "exam",
                {"official"},
            ),
            ("official", {"exam", "exam_quiz", "exam_term"}),
        ):
            bridge.set_current_work_mode(mode_id)
            app.processEvents()

            expected = [
                option.value
                for option in plan_selector_options(mode_id, include_source_prefix=True)
            ]
            quick_values = _combo_values(workbench._quick_execution_detail._scene_combo)
            overview_values = _combo_values(
                scene_panel._overview._combo,
                skip_scene_groups=True,
            )

            assert quick_values == expected
            assert overview_values == expected
            assert forbidden_values.isdisjoint(quick_values)
            assert forbidden_values.isdisjoint(overview_values)
    finally:
        scene_panel.close()
        workbench.close()
        app.processEvents()


def test_workbench_plan_selector_and_plan_preview_master_selector_stay_distinct(
    monkeypatch,
    tmp_path,
):
    monkeypatch.setattr(
        "src.ui.panels.scene_exam_detail.USER_EXAM_MASTER_DIR",
        tmp_path / "exam_user_masters",
    )
    app = _app()
    bridge = PanelBridge()
    workbench = WorkbenchPanel(bridge)
    scene_panel = ScenePanel(bridge)
    try:
        bridge.set_current_work_mode("exam")
        app.processEvents()
        scene_panel._ensure_detail_loaded("scn_exam_paper")
        app.processEvents()

        exam_plan_options = plan_selector_options("exam", include_source_prefix=False)
        exam_master_options = master_selector_options(
            "exam",
            user_master_dir=tmp_path / "exam_user_masters",
        )
        quick_plan_values = _combo_values(workbench._quick_execution_detail._scene_combo)
        quick_plan_labels = _combo_labels(workbench._quick_execution_detail._scene_combo)
        preview_master_values = _combo_values(scene_panel._exam_paper._blank_style)
        preview_master_labels = _combo_labels(scene_panel._exam_paper._blank_style)

        assert quick_plan_values == [option.value for option in exam_plan_options]
        assert quick_plan_labels == [option.label for option in exam_plan_options]
        assert preview_master_values == [option.value for option in exam_master_options]
        assert preview_master_labels == [option.label for option in exam_master_options]
        assert set(quick_plan_values).isdisjoint(preview_master_values)
        assert set(quick_plan_labels).isdisjoint(preview_master_labels)
        assert all("方案" in label for label in quick_plan_labels[:3])
        assert all("卷面" in label for label in preview_master_labels[:3])

        bridge.set_current_work_mode("official")
        app.processEvents()
        scene_panel._ensure_detail_loaded("scn_exam_paper")
        app.processEvents()

        official_plan_options = plan_selector_options(
            "official",
            include_source_prefix=False,
        )
        official_master_options = master_selector_options("official")
        quick_plan_values = _combo_values(workbench._quick_execution_detail._scene_combo)
        quick_plan_labels = _combo_labels(workbench._quick_execution_detail._scene_combo)
        official_master_values = _combo_values(scene_panel._exam_paper._official_master)
        official_master_labels = _combo_labels(scene_panel._exam_paper._official_master)

        assert quick_plan_values == [option.value for option in official_plan_options]
        assert quick_plan_labels == [option.label for option in official_plan_options]
        assert official_master_values == [option.value for option in official_master_options]
        assert official_master_labels == [option.label for option in official_master_options]
        assert set(quick_plan_values).isdisjoint(official_master_values)
        assert set(quick_plan_labels).isdisjoint(official_master_labels)
        assert all("方案" in label for label in quick_plan_labels)
        assert all(
            any(noun in label for noun in ("版式", "格式"))
            for label in official_master_labels
        )
    finally:
        scene_panel.close()
        workbench.close()
        app.processEvents()


def test_scene_panel_plan_selection_updates_loaded_rules_detail():
    app = _app()
    bridge = PanelBridge()
    bridge.set_current_work_mode("exam", emit_signal=False)
    scene_descriptor = get_scene_descriptor("exam", mode_id="exam")
    template_entry = get_template_entry("default", mode_id="exam")
    scene = load_scene_from_library("exam", mode_id="exam")
    template = load_template_from_library("default", mode_id="exam")
    assert scene_descriptor is not None
    assert template_entry is not None
    bridge.set_current_scene(
        scene,
        config_id="exam",
        path=str(scene_descriptor.path),
        source="library",
        emit_signal=False,
    )
    bridge.set_current_template(
        template,
        config_id="default",
        path=str(template_entry.path),
        source="library",
        emit_signal=False,
    )
    panel = ScenePanel(bridge)
    try:
        panel._ensure_detail_loaded("scn_rules")
        app.processEvents()

        combo = panel._overview._combo
        assert _combo_values(combo, skip_scene_groups=True) == [
            option.value for option in plan_selector_options("exam")
        ]

        term_index = combo.findData("exam_term")
        assert term_index >= 0
        combo.setCurrentIndex(term_index)
        app.processEvents()

        assert bridge.current_scene_id() == "exam_term"
        assert panel._current_scene.scene_id == "exam_term"
        assert panel._rules._current_scene is panel._current_scene
        assert panel._scope._current_scene is panel._current_scene
        assert panel._output._current_scene is panel._current_scene
        assert panel._rules._output_rules._current_scene is panel._current_scene
    finally:
        panel.close()
        app.processEvents()


def test_workbench_quick_execution_template_selector_is_mode_scoped():
    app = _app()
    bridge = PanelBridge()
    workbench = WorkbenchPanel(bridge)
    try:
        for mode_id, forbidden_values in (
            ("exam", {"official_gbt", "official_custom", "thesis_gbt"}),
            ("official", {"default", "thesis_gbt"}),
        ):
            bridge.set_current_work_mode(mode_id)
            app.processEvents()

            detail = workbench._quick_execution_detail
            current_scene = detail.current_scene()
            expected = [
                option.value
                for option in template_selector_options(
                    mode_id,
                    template_ids=list(current_scene.compatible_template_ids or ()),
                )
            ]
            quick_values = _combo_values(detail._template_combo)

            assert quick_values == expected
            assert forbidden_values.isdisjoint(quick_values)
    finally:
        workbench.close()
        app.processEvents()


def test_workbench_quick_execution_official_plan_template_cascade_is_mode_scoped():
    app = _app()
    bridge = PanelBridge()
    workbench = WorkbenchPanel(bridge)
    try:
        bridge.set_current_work_mode("official")
        app.processEvents()

        detail = workbench._quick_execution_detail
        expected_plan_options = plan_selector_options(
            "official",
            include_source_prefix=False,
        )
        assert _combo_values(detail._scene_combo) == [
            option.value for option in expected_plan_options
        ]
        assert [detail._scene_combo.itemText(index) for index in range(detail._scene_combo.count())] == [
            option.label for option in expected_plan_options
        ]

        for scene_id in ("official",):
            scene_index = detail._scene_combo.findData(scene_id)
            assert scene_index >= 0
            detail._scene_combo.setCurrentIndex(scene_index)
            app.processEvents()

            current_scene = detail.current_scene()
            expected_template_values = [
                option.value
                for option in template_selector_options(
                    "official",
                    template_ids=list(current_scene.compatible_template_ids or ()),
                )
            ]

            assert current_scene.scene_id == scene_id
            assert _combo_values(detail._template_combo) == expected_template_values
            assert "default" not in expected_template_values
            assert "thesis_gbt" not in expected_template_values
            assert all(
                value in {"official_gbt", "official_custom"}
                for value in expected_template_values
            )
        assert _combo_values(detail._document_type_combo) == [
            "notice",
            "letter",
            "minutes",
            "report",
            "request",
            "approval",
        ]
    finally:
        workbench.close()
        app.processEvents()


def test_workbench_quick_execution_exam_plan_selection_keeps_real_bound_master():
    app = _app()
    bridge = PanelBridge()
    panel = WorkbenchPanel(bridge)
    try:
        bridge.set_current_work_mode("exam")
        app.processEvents()

        combo = panel._quick_execution_detail._scene_combo
        quiz_index = combo.findData("exam_quiz")
        assert quiz_index >= 0

        combo.setCurrentIndex(quiz_index)
        app.processEvents()

        quick_scene = panel._quick_execution_detail.current_scene()
        bridge_scene = bridge.current_scene()

        assert quick_scene.scene_id == "exam_quiz"
        assert quick_scene.master_id == "default_exam"
        assert bridge.current_scene_id() == "exam_quiz"
        assert bridge_scene.master_id == "default_exam"
    finally:
        panel.close()
        app.processEvents()


def test_workbench_does_not_write_external_template_id_into_plan():
    app = _app()
    bridge = PanelBridge()
    panel = WorkbenchPanel(bridge)
    try:
        detail = panel._quick_execution_detail
        before_scene = detail.current_scene()
        template = TemplateConfig(name="Custom Saved Template")
        bridge.set_current_template(
            template,
            config_id="custom_saved",
            path="C:/tmp/custom_saved.json",
            source="file",
            source_type="external",
        )
        app.processEvents()

        assert bridge.current_template_id() == "custom_saved"
        assert bridge.current_template_source() == "file"
        assert bridge.current_template_source_type() == "external"
        assert detail.current_template_id() == before_scene.template_id
        assert detail.current_scene().template_id == before_scene.template_id
        assert (
            detail.current_scene().compatible_template_ids
            == before_scene.compatible_template_ids
        )
        assert "custom_saved" not in detail.current_scene().compatible_template_ids
        assert detail._template_combo.findData("custom_saved") < 0
    finally:
        panel.close()
        app.processEvents()


def test_workbench_same_template_binding_does_not_reload_or_overwrite_bridge_template(
    monkeypatch,
    caplog,
):
    app = _app()
    bridge = PanelBridge()
    panel = WorkbenchPanel(bridge)
    try:
        import src.ui.panels.workbench.panel_v2 as panel_v2

        original_template = bridge.current_template()
        original_template_id = bridge.current_template_id()

        def _raise_template_load(_template_id: str, **_kwargs):
            raise AssertionError("same committed template must not be reloaded")

        monkeypatch.setattr(panel_v2, "load_template_from_library", _raise_template_load)

        panel._on_quick_binding_changed(panel._current_scene, original_template_id)

        assert bridge.current_template() is original_template
        assert bridge.current_template_id() == original_template_id
        assert not caplog.records
    finally:
        panel.close()
        app.processEvents()
