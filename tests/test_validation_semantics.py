import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.resolved import ResolvedConfig
from src.modules.basic.section_format import SectionFormatModule
from src.modules.structure.heading_recognition import HeadingRecognitionModule
from src.modules.validate.validation import ValidationModule
from src.pipeline.context import PipelineContext
from src.pipeline.runner import Pipeline
from src.shared.engine.document_structure_model import DocSection, DocTree


def test_validation_module_warns_on_heading_style_mismatch():
    doc = Document()
    para = doc.add_heading("第一章 绪论", level=1)
    para.style = doc.styles["Normal"]

    config = ResolvedConfig()
    module = ValidationModule()
    context = SimpleNamespace(heading_map={0: 1})

    issues = module.validate(doc, config, context)

    assert any("标题样式不符" in issue.message for issue in issues)


def test_validation_module_reports_broken_toc_bookmark_placeholder():
    doc = Document()
    doc.add_paragraph("第一章 绪论\tError! Bookmark not defined")

    issues = ValidationModule().validate(doc, ResolvedConfig(), SimpleNamespace(heading_map={}))

    assert any(issue.level == "error" for issue in issues)


def test_validation_reports_preserved_formula_source_and_missing_equation_number():
    from lxml import etree

    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    doc.add_paragraph(r"需要复核 $\unsupported{x}$")
    table = doc.add_table(rows=1, cols=2)
    math = etree.SubElement(table.cell(0, 0).paragraphs[0]._element, qn("m:oMath"))
    math_run = etree.SubElement(math, qn("m:r"))
    etree.SubElement(math_run, qn("m:t")).text = "x"

    config = ResolvedConfig()
    config.module_switches["formula_convert"] = True
    config.module_switches["equation_table_format"] = True

    issues = ValidationModule().validate(
        doc,
        config,
        SimpleNamespace(heading_map={}),
    )

    messages = [issue.message for issue in issues]
    assert any("公式源码未安全转换" in message for message in messages)
    assert any("公式行缺少编号" in message for message in messages)


def test_formula_validation_ignores_sources_and_tables_outside_reviewed_scope():
    from lxml import etree

    from src.shared.engine.ooxml_ops import qn

    doc = Document()
    doc.add_paragraph(r"cover $\unsupported{x}$")
    table = doc.add_table(rows=1, cols=2)
    math = etree.SubElement(
        table.cell(0, 0).paragraphs[0]._element,
        qn("m:oMath"),
    )
    math_run = etree.SubElement(math, qn("m:r"))
    etree.SubElement(math_run, qn("m:t")).text = "x"
    doc.add_paragraph("body")

    config = ResolvedConfig()
    config.module_switches["formula_convert"] = True
    config.module_switches["equation_table_format"] = True
    context = PipelineContext(
        document_scope_gate_active=True,
        doc_tree=DocTree(
            sections=[
                DocSection("cover", 0, 1),
                DocSection("body", 1, 2),
            ],
            writable_roles=frozenset({"body"}),
        ),
    )

    issues = ValidationModule().validate(doc, config, context)
    formula_messages = [
        issue.message
        for issue in issues
        if "公式" in issue.message or "OLE" in issue.message
    ]

    assert formula_messages == []


def test_section_format_validate_rejects_unknown_break_type():
    issues = SectionFormatModule().validate(
        Document(),
        SimpleNamespace(section=SimpleNamespace(section_break_type="mystery")),
        SimpleNamespace(),
    )

    assert issues
    assert issues[0].level == "error"


def test_pipeline_strict_mode_fails_on_validation_errors():
    doc = Document()
    doc.add_paragraph("第一章 绪论\tError! Bookmark not defined")

    tmp = tempfile.NamedTemporaryFile(
        suffix=".docx",
        delete=False,
        dir=str(ROOT),
    )
    tmp.close()
    target = Path(tmp.name)
    doc.save(target)

    try:
        pipeline = Pipeline(
            modules=[HeadingRecognitionModule(), ValidationModule()],
            config=ResolvedConfig(strict_mode=False),
        )

        result = pipeline.execute(str(target))
        assert result.success is True
        assert result.status == "partial_success"

        strict_cfg = ResolvedConfig(strict_mode=True)
        strict_pipeline = Pipeline(
            modules=[HeadingRecognitionModule(), ValidationModule()],
            config=strict_cfg,
        )

        strict_result = strict_pipeline.execute(str(target))
        assert strict_result.success is False
        assert strict_result.status == "failed"
    finally:
        target.unlink(missing_ok=True)
        if "result" in locals():
            final_path = result.output_paths.get("final", "")
            if final_path:
                Path(final_path).unlink(missing_ok=True)
        if "strict_result" in locals():
            final_path = strict_result.output_paths.get("final", "")
            if final_path:
                Path(final_path).unlink(missing_ok=True)
