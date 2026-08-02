from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.config.resolved import ResolvedConfig
from src.modules.special.formula_convert import FormulaConvertModule
from src.modules.validate.md_cleanup import _clean_paragraph
from src.pipeline.context import PipelineContext
from src.pipeline.tracker import ChangeTracker
from src.shared.engine.ooxml_ops import qn


def test_formula_convert_builds_native_fraction_and_scripts_without_losing_prose():
    doc = Document()
    paragraph = doc.add_paragraph(r"结果为 $\frac{a_1}{b^2}+\alpha$，见正文。")

    tracker = ChangeTracker()
    FormulaConvertModule().apply(doc, ResolvedConfig(), tracker, PipelineContext())

    assert paragraph.text == "结果为 ，见正文。"
    assert paragraph._element.find(f".//{qn('m:oMath')}") is not None
    assert paragraph._element.find(qn("m:oMath")) is not None
    assert paragraph._element.find(f"{qn('w:r')}/{qn('m:oMath')}") is None
    assert paragraph._element.find(f".//{qn('m:f')}") is not None
    assert paragraph._element.find(f".//{qn('m:sSub')}") is not None
    assert paragraph._element.find(f".//{qn('m:sSup')}") is not None
    assert "α" in "".join(paragraph._element.itertext())
    assert tracker.get_by_module("formula_convert")[0].change_type == "convert"


def test_unsupported_formula_is_preserved_and_marked_as_skipped():
    doc = Document()
    source = r"保留 $\unsupported{x}+1$ 原文。"
    paragraph = doc.add_paragraph(source)
    tracker = ChangeTracker()

    FormulaConvertModule().apply(doc, ResolvedConfig(), tracker, PipelineContext())

    assert paragraph.text == source
    assert paragraph._element.find(f".//{qn('m:oMath')}") is None
    records = tracker.get_by_module("formula_convert")
    assert records and records[-1].change_type == "skip"


def test_md_cleanup_never_destroys_unconverted_formula_or_native_omml():
    doc = Document()
    source_paragraph = doc.add_paragraph(r"公式 $\frac{a}{b}$ 与 **说明**")
    assert _clean_paragraph(source_paragraph) is False
    assert source_paragraph.text == r"公式 $\frac{a}{b}$ 与 **说明**"

    native_paragraph = doc.add_paragraph(r"$x^2$ and **note**")
    FormulaConvertModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )
    xml_before = native_paragraph._element.xml
    assert _clean_paragraph(native_paragraph) is False
    assert native_paragraph._element.xml == xml_before


def test_md_cleanup_preserves_layout_spaces_and_url_underscores():
    doc = Document()
    source = (
        "签名：          年    月    日  "
        "http://example.test/moe_833/t20220914_660828.html"
    )
    paragraph = doc.add_paragraph(source)

    assert _clean_paragraph(paragraph) is False
    assert paragraph.text == source


def test_display_formula_uses_inline_valid_omml_and_centers_the_word_paragraph():
    doc = Document()
    paragraph = doc.add_paragraph(r"$$x^2+y^2$$")

    FormulaConvertModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    assert paragraph._element.find(f".//{qn('m:oMath')}") is not None
    assert paragraph._element.find(f".//{qn('m:oMathPara')}") is None
    assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_plain_status_assignments_are_not_converted_as_formulas():
    doc = Document()
    status = doc.add_paragraph("status=ok")
    timeout = doc.add_paragraph("timeout = 30")
    equation = doc.add_paragraph("E=mc^2")

    FormulaConvertModule().apply(
        doc,
        ResolvedConfig(),
        ChangeTracker(),
        PipelineContext(),
    )

    assert status.text == "status=ok"
    assert timeout.text == "timeout = 30"
    assert status._p.find(f".//{qn('m:oMath')}") is None
    assert timeout._p.find(f".//{qn('m:oMath')}") is None
    assert equation._p.find(f".//{qn('m:oMath')}") is not None
