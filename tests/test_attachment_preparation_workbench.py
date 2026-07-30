from pathlib import Path

from docx import Document
from openpyxl import Workbook

from src.config.attachment_materials import AttachmentProcessingMode
from src.config.entity import EntityArchive, EntityProfile
from src.config.material_batch import build_material_batch_items
from src.qt_api import QDialog, QTableWidget, Qt
from src.services.material_attachments import (
    AttachmentRequirementOwner,
    build_attachment_binding,
    build_directory_attachment_binding,
    build_attachment_preparation_report,
)
from src.shared.engine.material_timeline import (
    default_timeline_plan,
    timeline_output_field_keys,
)
from src.ui.bridge import PanelBridge
from src.ui.panels.assets import FIELD_SOURCE_IMPORTED_MAPPING
from src.ui.panels.assets.attachment_preparation_dialog import (
    AttachmentPreparationDialog,
)
from src.ui.panels.assets.batch_import import (
    _load_batch_excel_rows,
    _profiles_from_table_rows,
)
from src.ui.panels.assets_panel import AssetsPanel
from tests.iso_attachment_fixture_factory import build_iso_attachment_fixture


def _template(path: Path) -> Path:
    document = Document()
    document.add_paragraph(
        "{{@text:项目名称}} {{@time:项目开始日期}} "
        "{{@time:项目结束日期}} {{@time:节点_设计策划}}"
    )
    document.save(path)
    return path


def _binding(path: Path):
    return build_attachment_binding(
        role="iso_package",
        source_paths=(path,),
        accepted_types=("docx",),
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
    )


def test_preparation_projection_distinguishes_direct_time_from_timeline_candidate(
    tmp_path: Path,
):
    binding = _binding(_template(tmp_path / "template.docx"))
    profiles = [
        EntityProfile(
            profile_id="one",
            profile_name="项目一",
            fields={
                "项目名称": "管道项目",
                "项目开始日期": "2026-01-01",
                "项目结束日期": "2026-01-11",
            },
            field_sources={
                "项目名称": FIELD_SOURCE_IMPORTED_MAPPING,
                "项目开始日期": FIELD_SOURCE_IMPORTED_MAPPING,
                "项目结束日期": FIELD_SOURCE_IMPORTED_MAPPING,
            },
        )
    ]

    report = build_attachment_preparation_report(binding, profiles)
    by_token = {item.token: item for item in report.requirements}

    assert by_token["{{@time:项目开始日期}}"].owner == "field"
    assert by_token["{{@time:项目开始日期}}"].is_ready
    assert by_token["{{@time:节点_设计策划}}"].owner is AttachmentRequirementOwner.TIMELINE
    assert not by_token["{{@time:节点_设计策划}}"].is_ready


def test_excel_import_preserves_blank_declared_columns(tmp_path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("项目名称", "项目来源"))
    sheet.append(("项目甲", None))
    path = tmp_path / "batch.xlsx"
    workbook.save(path)

    profiles = _profiles_from_table_rows(_load_batch_excel_rows(path))

    assert profiles[0].fields == {"项目名称": "项目甲"}
    assert profiles[0].declared_field_keys == ["项目名称", "项目来源"]
    assert profiles[0].field_sources["项目来源"] == FIELD_SOURCE_IMPORTED_MAPPING


def test_image_binding_is_mapped_before_an_image_value_is_selected(tmp_path: Path):
    document = Document()
    document.add_paragraph("{{@img:LOGO1}}")
    path = tmp_path / "image-template.docx"
    document.save(path)
    binding = _binding(path)
    profile = EntityProfile(profile_id="one", profile_name="项目一")

    unmapped = build_attachment_preparation_report(binding, [profile])
    mapped = build_attachment_preparation_report(
        binding,
        [profile],
        image_token_bindings={"{{@img:LOGO1}}": "logo"},
    )
    logo = tmp_path / "logo.png"
    logo.write_bytes(b"not-rendered-by-this-test")
    ready = build_attachment_preparation_report(
        binding,
        [EntityProfile(profile_id="one", profile_name="项目一", asset_paths={"logo": str(logo)})],
        image_token_bindings={"{{@img:LOGO1}}": "logo"},
    )

    assert unmapped.mapped_requirement_count == 0
    assert mapped.mapped_requirement_count == 1
    assert not mapped.is_ready
    assert ready.mapped_requirement_count == 1
    assert ready.is_ready


def test_workbench_applies_one_timeline_plan_to_every_profile(qapp, tmp_path: Path):
    binding = _binding(_template(tmp_path / "template.docx"))
    profiles = [
        EntityProfile(
            profile_id=f"profile_{index}",
            profile_name=f"项目 {index}",
            fields={
                "项目名称": f"管道项目 {index}",
                "项目开始日期": "2026-01-01",
                "项目结束日期": "2026-01-11",
            },
            field_sources={
                "项目名称": FIELD_SOURCE_IMPORTED_MAPPING,
                "项目开始日期": FIELD_SOURCE_IMPORTED_MAPPING,
                "项目结束日期": FIELD_SOURCE_IMPORTED_MAPPING,
            },
        )
        for index in range(1, 3)
    ]
    dialog = AttachmentPreparationDialog(
        role="iso_package",
        binding=binding,
        profiles=profiles,
    )
    try:
        assert len(dialog._report.timeline_pending_requirements) == 1
        dialog._show_timeline_requirements()
        assert len(dialog._timeline_page.rows) == 1
        dialog._apply_timeline_plan()

        assert dialog._report.is_ready
        prepared = dialog.result_profiles()
        assert all(profile.timeline_plans for profile in prepared)
        assert all(
            "节点_设计策划"
            in next(iter(profile.timeline_plans.values()))["nodes"][0]["outputs"][0]["field"]
            for profile in prepared
        )
    finally:
        dialog._dirty = False
        dialog.close()


def test_workbench_uses_compact_three_phase_shell_with_shared_asset_semantics(
    qapp,
    tmp_path: Path,
):
    binding = _binding(_template(tmp_path / "template.docx"))
    dialog = AttachmentPreparationDialog(
        role="iso_package",
        binding=binding,
        profiles=[EntityProfile(profile_id="one", profile_name="项目一")],
    )
    try:
        assert dialog.windowFlags() & Qt.FramelessWindowHint
        assert dialog.width() <= 1100
        assert dialog._phase_nav.item_count == 3
        assert not dialog.findChildren(QTableWidget)
        assert dialog._source_label.text() == "当前 1 条数据"
        assert dialog._phase_nav.selected_card_id() == "fields"
        assert dialog._phase_cards["fields"]._title.text() == "字段资料"
        assert dialog._phase_cards["fields"].icon_name == "type"
        assert dialog._phase_cards["timeline"]._title.text() == "时间计划"
        assert (
            dialog._phase_cards["timeline"].icon_name
            == "chart-no-axes-gantt"
        )
        assert dialog._phase_cards["source"].icon_name == "file-text"
        assert "check" not in dialog._phase_pages
        assert dialog._complete_button.text() == "请先完成：字段资料"
        assert not dialog._complete_button.isEnabled()
        assert "字段资料还缺 3 项" in dialog._footer_status.text()
        assert "时间计划还缺 1 个节点" in dialog._footer_status.text()
    finally:
        dialog._dirty = False
        dialog.close()


def test_field_phase_edits_canonical_value_inline_without_usage_popover(
    qapp,
    tmp_path: Path,
):
    document = Document()
    document.add_paragraph("{{@text:项目名称}}")
    path = tmp_path / "field-template.docx"
    document.save(path)
    dialog = AttachmentPreparationDialog(
        role="iso_package",
        binding=_binding(path),
        profiles=[EntityProfile(profile_id="one", profile_name="项目一")],
    )
    try:
        assert len(dialog._field_page.rows) == 1
        row = dialog._field_page.rows.widgets()[0]
        assert row.token_label.text() == "{{@text:项目名称}}"
        assert row.usage_label.text() == "1文件·1处"
        assert row.usage_label.toolTip() == ""

        row.value_edit.setText("管道项目")
        row.value_edit.editingFinished.emit()
        qapp.processEvents()

        assert dialog.result_profiles()[0].fields["项目名称"] == "管道项目"
        assert dialog._report.is_ready
        assert dialog._phase_nav.selected_card_id() == "fields"
        assert dialog._complete_button.text() == "完成准备"
        assert dialog._complete_button.isEnabled()
        assert "已准备完成" in dialog._footer_status.text()
    finally:
        dialog._dirty = False
        dialog.close()


def test_workbench_routes_disabled_substitution_to_one_source_problem(
    qapp,
    tmp_path: Path,
):
    path = _template(tmp_path / "passthrough-template.docx")
    binding = build_attachment_binding(
        role="iso_package",
        source_paths=(path,),
        accepted_types=("docx",),
        processing_mode=AttachmentProcessingMode.PASSTHROUGH,
    )
    dialog = AttachmentPreparationDialog(
        role="iso_package",
        binding=binding,
        profiles=[EntityProfile(profile_id="one", profile_name="项目一")],
    )
    try:
        assert dialog._phase_nav.selected_card_id() == "source"
        assert dialog._source_problem_count() == 1
        assert dialog._source_page.meta_label.text() == "发现 1 个需要外部处理的问题。"
        assert dialog._actionable_pending_count(dialog._field_requirements()) == 0
    finally:
        dialog._dirty = False
        dialog.close()


def test_batch_context_keeps_only_package_wide_attachment_binding(
    qapp,
    tmp_path: Path,
):
    binding = _binding(_template(tmp_path / "template.docx"))
    plan = default_timeline_plan()
    plan["input_scope"] = "floating"
    plan["preset"] = {"id": "attachment_workbench_equal", "version": 1}
    plan["nodes"] = [
        {
            "node_id": "node_1",
            "node_no": 1,
            "active": True,
            "label": "节点_设计策划",
            "rule": {"operation": "ratio", "value": "0.5"},
            "outputs": [
                {"field": "节点_设计策划", "format": "yyyy-MM-dd"}
            ],
        }
    ]
    profile = EntityProfile(
        profile_id="one",
        profile_name="项目一",
        fields={
            "项目开始日期": "2026-01-01",
            "项目结束日期": "2026-01-11",
        },
        timeline_plans={"attachment_schedule": plan},
        field_aliases={"项目简称": "项目名称"},
        attachment_bindings={binding.role: binding},
    )
    panel = AssetsPanel(PanelBridge())
    try:
        panel.set_archive(EntityArchive(profiles=[profile]))

        context = panel._shared_batch_context()

        assert context.attachment_bindings == {binding.role: binding}
        assert context.timeline_plans == {}
        assert context.field_aliases == {}

        selection = panel.material_batch_selection()
        item = build_material_batch_items(
            selection.archive,
            profile_ids=selection.profile_ids,
            base_context=selection.base_context,
        )[0]
        assert "节点_设计策划" in timeline_output_field_keys(
            item.context.timeline_plans
        )
        assert item.context.field_aliases == {"项目简称": "项目名称"}
    finally:
        panel.close()


def test_workbench_surfaces_unreadable_source_and_blocks_completion(
    qapp,
    tmp_path: Path,
):
    path = tmp_path / "broken.docx"
    binding = _binding(_template(path))
    path.write_bytes(b"not-a-docx")
    dialog = AttachmentPreparationDialog(
        role="iso_package",
        binding=binding,
        profiles=[EntityProfile(profile_id="one", profile_name="项目一")],
    )
    try:
        assert dialog._report.unreadable_paths == ("broken.docx",)
        assert dialog._phase_nav.item_count == 3
        assert dialog._phase_nav.selected_card_id() == "source"

        dialog._complete_preparation()

        assert dialog.result() != QDialog.Accepted
        assert "1 个源文件无法读取" in dialog._footer_status.text()
    finally:
        dialog._dirty = False
        dialog.close()


def test_iso_fixture_collapses_41_tokens_into_30_columns_and_one_time_group(
    tmp_path: Path,
):
    fixture = build_iso_attachment_fixture(tmp_path)
    binding = build_directory_attachment_binding(
        role="iso_package",
        source_directory=fixture.docx_directory,
        accepted_types=("docx",),
        recursive=True,
        processing_mode=AttachmentProcessingMode.SUBSTITUTE_COPY,
        max_items=None,
    )
    profiles = _profiles_from_table_rows(
        _load_batch_excel_rows(fixture.workbook_path)
    )

    report = build_attachment_preparation_report(binding, profiles)

    assert report.profile_count == 18
    assert report.docx_count == 6
    assert report.token_count == 41
    assert report.occurrence_count == 114
    assert sum(item.owner is AttachmentRequirementOwner.FIELD for item in report.requirements) == 30
    assert sum(
        item.owner is AttachmentRequirementOwner.TIMELINE
        for item in report.requirements
    ) == 11
    assert report.mapped_requirement_count == 41
