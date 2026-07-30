import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from lxml import etree


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.config.style_semantics import apply_style_special_indent
from src.config.template import StyleConfig, TemplateConfig
from src.modules.basic.paragraph_style import _apply_indent, _apply_spacing
from src.qt_api import QApplication
from src.shared.engine.ooxml_ops import qn
from src.shared.ui.font_combo import FontCombo
from src.shared.ui.size_combo import SizeCombo
from src.config.resolver import resolve_template_baseline
from src.ui.panels.template_summary_projection import build_template_detail_summary


def _app():
    return QApplication.instance() or QApplication([])


def test_size_combo_parses_named_size_and_pt_suffix():
    _app()
    combo = SizeCombo()
    try:
        combo.setEditText("小四")
        assert combo.current_pt() == 12.0

        combo.setEditText("12磅")
        assert combo.current_pt() == 12.0
    finally:
        combo.close()


def test_font_combo_preserves_custom_font_name():
    _app()
    combo = FontCombo(lang="cn")
    try:
        combo.set_font_name("思源宋体")
        assert combo.selected_font() == "思源宋体"
    finally:
        combo.close()


def test_paragraph_style_applies_pt_indent_units_as_points():
    document = Document()
    para = document.add_paragraph("正文")
    style = StyleConfig(
        size_pt=12,
        left_indent_chars=24,
        left_indent_unit="pt",
        right_indent_chars=12,
        right_indent_unit="pt",
    )
    apply_style_special_indent(style, "first_line", 18, "pt")

    _apply_indent(para, style, style.size_pt)

    assert round(para.paragraph_format.left_indent.pt, 2) == 24.0
    assert round(para.paragraph_format.right_indent.pt, 2) == 12.0
    assert round(para.paragraph_format.first_line_indent.pt, 2) == 18.0

    ind = para._element.find(qn("w:pPr")).find(qn("w:ind"))
    assert ind.get(qn("w:left")) == "480"
    assert ind.get(qn("w:right")) == "240"
    assert ind.get(qn("w:firstLine")) == "360"
    assert ind.get(qn("w:leftChars")) is None
    assert ind.get(qn("w:rightChars")) is None
    assert ind.get(qn("w:firstLineChars")) is None


def test_paragraph_style_maps_cm_indent_units_to_twips():
    document = Document()
    para = document.add_paragraph("正文")
    style = StyleConfig(
        size_pt=12,
        left_indent_chars=1.0,
        left_indent_unit="cm",
        right_indent_chars=0.5,
        right_indent_unit="cm",
    )

    _apply_indent(para, style, style.size_pt)

    ind = para._element.find(qn("w:pPr")).find(qn("w:ind"))
    assert ind.get(qn("w:left")) == "567"
    assert ind.get(qn("w:right")) == "283"
    assert ind.get(qn("w:leftChars")) is None
    assert ind.get(qn("w:rightChars")) is None


def test_paragraph_style_preserves_char_indent_semantics_in_ooxml():
    document = Document()
    para = document.add_paragraph("正文")
    style = StyleConfig(
        size_pt=12,
        left_indent_chars=2.0,
        left_indent_unit="chars",
        right_indent_chars=1.0,
        right_indent_unit="chars",
    )
    apply_style_special_indent(style, "first_line", 2.0, "chars")

    _apply_indent(para, style, style.size_pt)

    ind = para._element.find(qn("w:pPr")).find(qn("w:ind"))
    assert ind.get(qn("w:leftChars")) == "200"
    assert ind.get(qn("w:rightChars")) == "100"
    assert ind.get(qn("w:firstLineChars")) == "200"
    assert ind.get(qn("w:left")) is None
    assert ind.get(qn("w:right")) is None
    assert ind.get(qn("w:firstLine")) is None


def test_paragraph_style_preserves_hanging_char_indent_semantics_in_ooxml():
    document = Document()
    para = document.add_paragraph("正文")
    style = StyleConfig(size_pt=12)
    apply_style_special_indent(style, "hanging", 2.0, "chars")

    _apply_indent(para, style, style.size_pt)

    ind = para._element.find(qn("w:pPr")).find(qn("w:ind"))
    assert ind.get(qn("w:leftChars")) == "200"
    assert ind.get(qn("w:hangingChars")) == "200"
    assert ind.get(qn("w:firstLine")) is None
    assert ind.get(qn("w:firstLineChars")) is None


def test_paragraph_style_zero_indent_overrides_legacy_char_indent_attrs():
    document = Document()
    para = document.add_paragraph("正文")
    ppr = para._element.get_or_add_pPr()
    ind = etree.SubElement(ppr, qn("w:ind"))
    ind.set(qn("w:leftChars"), "200")
    ind.set(qn("w:firstLineChars"), "200")

    _apply_indent(para, StyleConfig(size_pt=12), 12)

    ind = para._element.find(qn("w:pPr")).find(qn("w:ind"))
    assert ind.get(qn("w:left")) == "0"
    assert ind.get(qn("w:leftChars")) == "0"
    assert ind.get(qn("w:firstLine")) == "0"
    assert ind.get(qn("w:firstLineChars")) == "0"


def test_paragraph_style_syncs_exact_spacing_to_ooxml():
    document = Document()
    para = document.add_paragraph("正文")
    ppr = para._element.get_or_add_pPr()
    spacing = etree.SubElement(ppr, qn("w:spacing"))
    spacing.set(qn("w:beforeAutospacing"), "1")
    spacing.set(qn("w:afterAutospacing"), "1")

    style = StyleConfig(
        line_spacing_type="exact",
        line_spacing_pt=20,
        space_before_pt=6,
        space_after_pt=4,
    )

    _apply_spacing(para, style)

    spacing = para._element.find(qn("w:pPr")).find(qn("w:spacing"))
    assert para.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert spacing.get(qn("w:before")) == "120"
    assert spacing.get(qn("w:after")) == "80"
    assert spacing.get(qn("w:line")) == "400"
    assert spacing.get(qn("w:lineRule")) == "exact"
    assert spacing.get(qn("w:beforeAutospacing")) is None
    assert spacing.get(qn("w:afterAutospacing")) is None


def test_paragraph_style_syncs_multiple_spacing_to_ooxml():
    document = Document()
    para = document.add_paragraph("正文")
    style = StyleConfig(
        line_spacing_type="multiple",
        line_spacing_pt=1.5,
        space_before_pt=0,
        space_after_pt=0,
    )

    _apply_spacing(para, style)

    spacing = para._element.find(qn("w:pPr")).find(qn("w:spacing"))
    assert para.paragraph_format.line_spacing_rule in {
        WD_LINE_SPACING.MULTIPLE,
        WD_LINE_SPACING.ONE_POINT_FIVE,
    }
    assert spacing.get(qn("w:line")) == "360"
    assert spacing.get(qn("w:lineRule")) == "auto"


def test_paragraph_style_syncs_line_and_auto_paragraph_spacing_units_to_ooxml():
    document = Document()
    para = document.add_paragraph("正文")
    style = StyleConfig(
        line_spacing_type="exact",
        line_spacing_pt=20,
        space_before_pt=1.5,
        space_before_unit="lines",
        space_after_pt=0,
        space_after_unit="auto",
    )

    _apply_spacing(para, style)

    spacing = para._element.find(qn("w:pPr")).find(qn("w:spacing"))
    assert spacing.get(qn("w:before")) is None
    assert spacing.get(qn("w:beforeLines")) == "150"
    assert spacing.get(qn("w:after")) is None
    assert spacing.get(qn("w:afterAutospacing")) == "1"


def test_style_preview_summary_reflects_hanging_indent():
    cfg = TemplateConfig()
    body = StyleConfig()
    cfg.styles["body"] = body
    body.font_cn = "黑体"
    body.font_en = "Arial"
    body.size_pt = 12
    body.line_spacing_type = "multiple"
    body.line_spacing_pt = 1.5
    apply_style_special_indent(body, "hanging", 1.5, "chars")

    style_summary = build_template_detail_summary(
        resolve_template_baseline(cfg),
        "tpl_style",
    ).nav_summary

    assert "悬挂" in style_summary
    assert "1.5倍" in style_summary


def test_style_preview_summary_preserves_named_line_spacing_presets():
    cfg = TemplateConfig()
    body = StyleConfig()
    cfg.styles["body"] = body
    body.line_spacing_type = "double"
    body.line_spacing_pt = 2.0

    style_summary = build_template_detail_summary(
        resolve_template_baseline(cfg),
        "tpl_style",
    ).nav_summary

    assert "双倍" in style_summary
