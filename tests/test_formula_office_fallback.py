from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from src.config.resolved import ResolvedConfig
from src.config.document_scope import DocumentScopePolicy
from src.modules.special.formula_convert import FormulaConvertModule
from src.pipeline.runner import Pipeline
from src.pipeline.context import PipelineContext
from src.shared.engine.office_broker_command import (
    MATHTYPE_OFFICE_CHILD_FLAG,
    build_office_broker_child_command,
)
from src.shared.engine.ooxml_ops import qn
from src.shared.io.mathtype_office_fallback import (
    apply_mathml_fallback_payloads_to_docx,
)


def test_pipeline_stage_invokes_opted_in_mathtype_office_fallback(monkeypatch, tmp_path):
    calls: list[tuple[str, int]] = []

    def fake_fallback(path: str, timeout_sec: int):
        calls.append((path, timeout_sec))
        return True, "extracted=1/1, replaced=1", {
            "found": 1,
            "extracted": 1,
            "replaced": 1,
        }

    monkeypatch.setattr(
        "src.shared.io.mathtype_office_fallback.apply_mathtype_office_fallback",
        fake_fallback,
    )
    config = ResolvedConfig()
    config.module_switches["formula_convert"] = True
    config.formula_convert.office_fallback_enabled = True
    config.formula_convert.office_fallback_timeout_sec = 41
    pipeline = Pipeline([FormulaConvertModule()], config)
    stage = tmp_path / "stage.docx"

    pipeline._save_staged_document(Document(), stage)

    assert calls == [(str(stage), 41)]
    assert pipeline._mathtype_fallback_receipts[0]["changed"] is True
    assert pipeline._tracker.get_by_module("formula_convert")[-1].change_type == "convert"


def test_mathml_office_fallback_emits_valid_block_formula_representation(
    monkeypatch,
    tmp_path,
):
    document = Document()
    paragraph = document.add_paragraph()
    object_run = paragraph.add_run()
    object_node = OxmlElement("w:object")
    ole_node = etree.Element("{urn:schemas-microsoft-com:office:office}OLEObject")
    ole_node.set("ProgID", "MathType Equation")
    object_node.append(ole_node)
    object_run._r.append(object_node)
    source = tmp_path / "mathtype-block.docx"
    document.save(source)
    block_flags: list[bool] = []

    def fake_convert(_mathml: str, *, block: bool):
        block_flags.append(block)
        math = etree.Element(qn("m:oMath"))
        math_run = etree.SubElement(math, qn("m:r"))
        etree.SubElement(math_run, qn("m:t")).text = "x"
        return math

    monkeypatch.setattr(
        "src.shared.io.mathtype_office_fallback.convert_mathml_to_omml",
        fake_convert,
    )

    stats = apply_mathml_fallback_payloads_to_docx(
        str(source),
        [{"mathml": "<math><mi>x</mi></math>"}],
    )
    reloaded = Document(source)
    reloaded_paragraph = reloaded.paragraphs[0]

    assert stats["replaced"] == 1
    assert block_flags == [False]
    assert reloaded_paragraph._element.find(f".//{qn('m:oMath')}") is not None
    assert reloaded_paragraph._element.find(f".//{qn('m:oMathPara')}") is None
    assert (
        reloaded_paragraph._element.find(f"{qn('w:pPr')}/{qn('w:jc')}").get(
            qn("w:val")
        )
        == "center"
    )


def test_mathtype_broker_command_supports_source_and_frozen_runtimes(tmp_path):
    document = tmp_path / "formula.docx"
    source_command = build_office_broker_child_command(
        module_name="src.shared.io.mathtype_office_fallback",
        frozen_flag=MATHTYPE_OFFICE_CHILD_FLAG,
        request_path=document,
        executable="python-test",
        frozen=False,
    )
    frozen_command = build_office_broker_child_command(
        module_name="src.shared.io.mathtype_office_fallback",
        frozen_flag=MATHTYPE_OFFICE_CHILD_FLAG,
        request_path=document,
        executable="Alavette-Form_V1.0.exe",
        frozen=True,
    )

    assert source_command == [
        "python-test",
        "-m",
        "src.shared.io.mathtype_office_fallback",
        "--child-request",
        str(document),
    ]
    assert frozen_command == [
        "Alavette-Form_V1.0.exe",
        MATHTYPE_OFFICE_CHILD_FLAG,
        str(document),
    ]


def test_office_fallback_reapplies_equation_layout_after_conversion(monkeypatch, tmp_path):
    postprocessed: list[str] = []
    monkeypatch.setattr(
        "src.shared.io.mathtype_office_fallback.apply_mathtype_office_fallback",
        lambda path, timeout_sec: (
            True,
            "extracted=1/1, replaced=1",
            {"found": 1, "replaced": 1},
        ),
    )

    def fake_postprocess(self, doc, config, tracker, context):
        postprocessed.append("equation_table_format")

    monkeypatch.setattr(
        "src.modules.special.equation_table_format.EquationTableFormatModule.apply",
        fake_postprocess,
    )
    config = ResolvedConfig()
    config.module_switches["formula_convert"] = True
    config.module_switches["equation_table_format"] = True
    config.formula_convert.office_fallback_enabled = True
    stage = tmp_path / "postprocess.docx"

    Pipeline([FormulaConvertModule()], config)._save_staged_document(
        Document(), stage, ctx=PipelineContext()
    )

    assert postprocessed == ["equation_table_format"]


def test_office_fallback_is_blocked_for_partial_document_scope(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "src.shared.io.mathtype_office_fallback.apply_mathtype_office_fallback",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("Office fallback must not mutate a partial-scope document")
        ),
    )
    config = ResolvedConfig()
    config.module_switches["formula_convert"] = True
    config.formula_convert.office_fallback_enabled = True
    pipeline = Pipeline([FormulaConvertModule()], config)
    context = PipelineContext(document_scope=DocumentScopePolicy(mode="body"))

    pipeline._save_staged_document(
        Document(), tmp_path / "partial.docx", ctx=context
    )

    assert context.mathtype_office_fallback[0]["detail"] == (
        "skipped_non_global_document_scope"
    )


def test_main_dispatches_frozen_mathtype_child_before_argument_parser(
    monkeypatch,
    tmp_path,
):
    import main

    document = tmp_path / "formula.docx"
    monkeypatch.setattr(
        "src.shared.io.mathtype_office_fallback.run_mathtype_office_child",
        lambda path: 23 if str(path) == str(document) else 99,
    )
    monkeypatch.setattr(
        main,
        "parse_args",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("argument parser must not run for broker child")
        ),
    )

    assert main.run_app([MATHTYPE_OFFICE_CHILD_FLAG, str(document)]) == 23
