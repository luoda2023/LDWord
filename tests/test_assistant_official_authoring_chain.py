from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from docx import Document

from src.assistant.adapters.content_generation_adapter import (
    AssistantContentGenerationAdapter,
    GeneratedOfficialDraft,
    complete_generated_official_draft_field,
)
from src.assistant.adapters.production_adapter import AssistantProductionAdapter
from src.assistant.adapters.workspace_state_adapter import WorkspaceSnapshot
from src.assistant.application.content_generation_service import (
    AssistantContentGenerationService,
    ContentGenerationRequest,
)
from src.assistant.application.generated_draft_binding import bind_generated_draft
from src.assistant.application.plan_builder import FormDocumentPlanBuilder
from src.assistant.contracts.execution import ExecutionApproval
from src.assistant.contracts.task_plan import (
    ARTIFACT_KIND_OFFICIAL,
    CAPABILITY_EXECUTABLE,
    SOURCE_ROLE_STRUCTURED_SOURCE,
    TASK_OPERATION_AUTHOR,
)
from src.assistant.runtime.provider_contract import (
    PROVIDER_DONE,
    PROVIDER_START,
    PROVIDER_TEXT_DELTA,
    ProviderRequest,
    ProviderStreamEvent,
)
from src.assistant.runtime.providers.openai_compatible import (
    OpenAICompatibleModelGateway,
)
from src.assistant.ui.document_workflow_mixin import (
    AssistantDocumentWorkflowMixin,
)


class _FixedGateway:
    def __init__(self, text: str) -> None:
        self.text = text
        self.request = None

    def stream(self, request):
        self.request = request
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(PROVIDER_TEXT_DELTA, text=self.text)
        yield ProviderStreamEvent(PROVIDER_DONE, text=self.text)

    def cancel(self) -> bool:
        return True


def _workspace(*, input_path: Path | None = None) -> WorkspaceSnapshot:
    return WorkspaceSnapshot(
        mode_id="custom",
        mode_label="通用版",
        scene_id="custom",
        scene_source_type="builtin",
        template_id="default",
        template_source_type="builtin",
        input_path=str(input_path or ""),
        input_name=input_path.name if input_path is not None else "",
        input_exists=bool(input_path is not None and input_path.is_file()),
        material_summary={},
        document_type_id="",
    )


def _approval(preflight) -> ExecutionApproval:
    return ExecutionApproval(
        approval_id="approval-official-author",
        session_id="session-official-author",
        plan_id=preflight.plan_id,
        plan_revision=preflight.plan_revision,
        preflight_hash=preflight.evidence_hash,
        input_hash=preflight.input_hash,
        output_root=preflight.output_root,
        overwrite_policy="deny",
        approved_at="2026-07-31T00:00:00+00:00",
    )


def test_official_check_notice_is_authoring_not_review() -> None:
    plan = FormDocumentPlanBuilder().build(
        query=(
            "请起草一份关于开展2026年第三季度档案检查的通知公文，"
            "并生成规范Word文件"
        ),
        workspace=_workspace(),
        turn_id="turn-official-route",
    )

    assert plan.operation == TASK_OPERATION_AUTHOR
    assert plan.work_mode_id == "official"
    assert plan.capability_ref.status == CAPABILITY_EXECUTABLE
    assert plan.generation_contract.required is True
    assert plan.generation_contract.artifact_kind == ARTIFACT_KIND_OFFICIAL
    assert plan.production_contract.input_role == SOURCE_ROLE_STRUCTURED_SOURCE
    assert plan.production_contract.validator_id == "official_document_draft_v1"
    assert plan.production_contract.document_type_id == "notice"
    assert plan.production_contract.accepted_suffixes == (".json",)
    assert not plan.blocking_issues


def test_official_draft_prefers_materials_and_asks_for_missing_organization(
    tmp_path,
) -> None:
    adapter = AssistantContentGenerationAdapter(tmp_path)
    draft = adapter.compile_official_draft(
        session_id="session-official-fields",
        markdown=(
            '{"fields":{"title":"检查通知","body":"请按要求完成检查。",'
            '"organization":null,"document_no":null,"issue_date":null}}'
        ),
        intent="起草档案检查通知",
        document_type_id="notice",
        authoritative_fields={},
    )

    assert isinstance(draft, GeneratedOfficialDraft)
    assert draft.missing_user_fields == ("organization",)
    assert draft.field_values["document_no"] == "待编"
    assert draft.field_provenance["document_no"] == "auto_placeholder"
    assert draft.field_provenance["issue_date"] == "auto_default"

    completed = complete_generated_official_draft_field(
        draft.source_path,
        field_key="organization",
        value="示例市综合办公室",
    )

    assert completed.missing_user_fields == ()
    assert completed.field_values["organization"] == "示例市综合办公室"
    assert completed.field_provenance["organization"] == "user_confirmed"
    assert Path(completed.preview_path).is_file()

    material_owned = adapter.compile_official_draft(
        session_id="session-official-materials",
        markdown=(
            '{"fields":{"title":"模型标题","body":"模型正文",'
            '"organization":"模型猜测机关"}}'
        ),
        intent="起草通知",
        document_type_id="notice",
        authoritative_fields={
            "title": "材料中的正式标题",
            "organization": "材料中的发文机关",
            "document_no": "示办发〔2026〕8号",
        },
    )

    assert material_owned.field_values["title"] == "材料中的正式标题"
    assert material_owned.field_values["organization"] == "材料中的发文机关"
    assert material_owned.field_values["document_no"] == "示办发〔2026〕8号"
    assert material_owned.field_provenance["organization"] == "user_material"


def test_official_missing_fields_are_presented_in_one_structured_form(
    tmp_path,
) -> None:
    draft = AssistantContentGenerationAdapter(tmp_path).compile_official_draft(
        session_id="session-official-batch-fields",
        markdown='{"fields":{}}',
        intent="起草通知",
        document_type_id="notice",
    )

    assert draft.missing_user_fields == ("title", "body", "organization")
    suggestions = {
        field_key: AssistantDocumentWorkflowMixin._official_field_suggestion(
            field_key,
            draft,
        )
        for field_key in draft.missing_user_fields
    }
    message = AssistantDocumentWorkflowMixin._official_field_question_message(
        draft,
        completion_id="official-batch-1",
        suggestions=suggestions,
    )
    payload = message.blocks[0].data

    assert payload["title"] == "请补充公文必填信息"
    assert payload["field_keys"] == ["title", "body", "organization"]
    assert payload["confirmation_request"]["input_columns"] == 2
    assert payload["confirmation_request"]["compact_heading"] is True
    assert payload["confirmation_request"]["submit_label"] == "保存并继续"
    assert [
        item["id"]
        for item in payload["confirmation_request"]["inputs"]
    ] == ["title", "body", "organization"]
    assert payload["confirmation_request"]["inputs"][1][
        "column_span"
    ] == 2
    assert all(
        item["required"]
        for item in payload["confirmation_request"]["inputs"]
    )


def test_official_authoring_reaches_real_terminal_assembly(tmp_path) -> None:
    plan = FormDocumentPlanBuilder().build(
        query="请根据工作安排起草一份第三季度档案检查通知",
        workspace=_workspace(),
        turn_id="turn-official-e2e",
    )
    plan = replace(
        plan,
        output_policy=replace(
            plan.output_policy,
            output_root=str(tmp_path / "outputs"),
        ),
    )
    gateway = _FixedGateway(
        '{"fields":{'
        '"title":"关于开展第三季度档案检查的通知",'
        '"body":"为规范档案管理，现组织开展检查。请各部门按期完成自查并报送情况。",'
        '"organization":"示例市综合办公室",'
        '"recipient":"各部门"'
        "}}"
    )
    draft = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(tmp_path / "generation")
    ).generate(
        ContentGenerationRequest(
            session_id="session-official-e2e",
            turn_id="turn-official-e2e",
            prompt=plan.intent,
            provider_id="fixed",
            model_id="fixed",
            artifact_kind=plan.generation_contract.artifact_kind,
            prompt_profile_id=plan.generation_contract.prompt_profile_id,
            document_type_id=plan.production_contract.document_type_id,
        ),
        gateway,
    )
    bound = bind_generated_draft(plan, draft)
    adapter = AssistantProductionAdapter()

    preflight = adapter.build_preflight(bound)
    result = adapter.execute_approved_plan(
        bound,
        preflight,
        _approval(preflight),
    )

    assert preflight.ready, preflight.issues
    assert result["status"] == "success", result.get("error_text")
    assert "official_docx" in result["output_paths"]
    assert "internal_review_docx" not in result["output_paths"]
    assert "archive_manifest" not in result["output_paths"]
    assert "archive_manifest_md" not in result["output_paths"]
    official_path = Path(result["output_paths"]["official_docx"])
    assert official_path.is_file()
    assert official_path.name == "关于开展第三季度档案检查的通知.docx"
    text = "\n".join(item.text for item in Document(official_path).paragraphs)
    assert "关于开展第三季度档案检查的通知" in text
    assert "示例市综合办公室" in text
    assert result["official_draft"]["provisional_fields"] == [
        "issue_date",
        "document_no",
    ]
    assert "fields" in gateway.request.system_prompt


def test_existing_official_docx_uses_format_only_preflight_without_materials(
    tmp_path,
) -> None:
    input_path = tmp_path / "notice.docx"
    Document().save(input_path)
    plan = FormDocumentPlanBuilder().build(
        query="将这份通知套用正式公文格式",
        workspace=_workspace(input_path=input_path),
        turn_id="turn-official-transform",
    )
    plan = replace(
        plan,
        output_policy=replace(
            plan.output_policy,
            output_root=str(tmp_path / "outputs"),
        ),
    )

    preflight = AssistantProductionAdapter().build_preflight(plan)

    assert plan.production_contract.terminal_assembler == "generic"
    assert plan.production_contract.validator_id == "docx_preflight"
    assert preflight.ready is True
    assert not any(
        issue.startswith("official_material_") for issue in preflight.issues
    )


def test_provider_retries_clean_eof_before_any_delta() -> None:
    attempts: list[int] = []

    def transport(_request):
        attempts.append(len(attempts) + 1)
        if len(attempts) == 1:
            return []
        return [
            'data: {"choices":[{"delta":{"content":"ok"}}]}\n\n',
            "data: [DONE]\n\n",
        ]

    gateway = OpenAICompatibleModelGateway(
        model="example-model",
        api_key="secret",
        streaming_transport=transport,
        sleep=lambda _seconds: None,
    )
    events = list(
        gateway.stream(
            ProviderRequest(
                request_id="official-clean-eof",
                model="example-model",
                system_prompt="输出测试文本。",
                messages=({"role": "user", "content": "test"},),
            )
        )
    )

    assert attempts == [1, 2]
    assert events[-1].type == PROVIDER_DONE
