import hashlib
import json
from pathlib import Path

from docx import Document
from PIL import Image

from src.config.official_document_profiles import get_official_document_assembly_contract
import src.shared.engine.official_document_sample_verification as official_visual
from src.shared.engine.docx_page_renderer import (
    DocxPageRenderResult,
    png_has_content,
)
from src.shared.engine.official_document_sample_verification import (
    official_sample_entity_data,
    verify_official_document_sample_visual,
)


ROOT = Path(__file__).resolve().parent.parent
OFFICIAL_VISUAL_BASELINE_DIR = (
    ROOT / "config_library" / "masters" / "official" / "builtin" / "visual_baselines"
)


def _all_docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _fake_rendering(page_factory):
    def render(docx_path: Path, output_dir: Path, *, attempt_render=True):
        assert attempt_render is True
        pdf_path = output_dir / f"{docx_path.stem}.pdf"
        pdf_path.write_bytes(b"%PDF-1.4\n% fake\n")
        page_paths, issues = page_factory(pdf_path, output_dir)
        blank_paths = tuple(path for path in page_paths if not png_has_content(path))
        blank_issues = tuple(f"blank_png: {path.name}" for path in blank_paths)
        return DocxPageRenderResult(
            status="png_blank" if blank_paths else "ready",
            docx_path=docx_path,
            pdf_path=pdf_path,
            page_paths=page_paths,
            renderer="fake_renderer",
            issues=(*issues, *blank_issues),
        )

    return render


def test_official_document_sample_verification_covers_common_document_types(tmp_path):
    expected_titles = {
        "notice": "关于开展资料归档检查的通知",
        "letter": "关于商请协助提供归档材料的函",
        "minutes": "专题协调会议纪要",
        "report": "关于年度资料归档工作情况的报告",
        "request": "关于升级电子档案管理系统的请示",
        "approval": "关于电子档案管理系统升级事项的批复",
    }
    expected_masters = {
        "notice": "official_gbt_standard",
        "letter": "official_gbt_letter",
        "minutes": "official_gbt_minutes",
        "report": "official_gbt_upward",
        "request": "official_gbt_upward",
        "approval": "official_gbt_standard",
    }

    for profile_id, title in expected_titles.items():
        result = verify_official_document_sample_visual(
            profile_id,
            tmp_path,
            attempt_render=False,
        )

        assert result.status == "docx_verified"
        assert result.assembly_status == "ok"
        assert result.sample_docx_path is not None
        assert result.sample_docx_path.is_file()
        assert result.manifest_path.is_file()
        assert result.profile_id == profile_id
        assert result.master_id == expected_masters[profile_id]
        assert "official_title" in result.replaced_placeholders
        assert "official_body" in result.replaced_placeholders
        assert result.unresolved_placeholders == ()

        payload = json.loads(result.manifest_path.read_text(encoding="utf-8"))
        assert payload["status"] == "docx_verified"
        assert payload["assembly_status"] == "ok"
        assert payload["sample_docx_path"] == str(result.sample_docx_path)
        assert payload["output_paths"]["official_docx"] == str(result.sample_docx_path)
        assert "internal_review_docx" in payload["output_paths"]
        assert "archive_manifest" in payload["output_paths"]
        sample_text = _all_docx_text(result.sample_docx_path)
        sample_data = official_sample_entity_data(profile_id)
        contract = get_official_document_assembly_contract(profile_id)
        assert title in sample_text
        if "attachment_note" in contract.applicable_material_field_keys:
            assert sample_data["attachment_note"] in sample_text
        if profile_id == "minutes":
            assert f"会议时间：{sample_data['meeting_date']}" in sample_text
            assert f"出席：{sample_data['participants']}" in sample_text

    minutes = official_sample_entity_data("minutes")
    assert minutes["document_type"] == "minutes"
    assert minutes["meeting_date"]
    assert minutes["participants"]


def test_official_document_sample_verification_records_renderer_evidence(
    monkeypatch,
    tmp_path,
):
    def fake_pdf_to_png(pdf_path: Path, output_dir: Path):
        png_path = output_dir / f"{pdf_path.stem}_page-1.png"
        image = Image.new("RGB", (8, 8), "white")
        image.putpixel((3, 3), (0, 0, 0))
        image.save(png_path)
        return (png_path,), ()

    monkeypatch.setattr(
        official_visual,
        "render_docx_pages",
        _fake_rendering(fake_pdf_to_png),
    )

    result = verify_official_document_sample_visual("letter", tmp_path)

    assert result.status == "visual_png_ok"
    assert result.renderer == "fake_renderer"
    assert result.pdf_path is not None
    assert result.pdf_path.is_file()
    assert len(result.png_paths) == 1
    payload = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert payload["status"] == "visual_png_ok"
    assert payload["renderer"] == "fake_renderer"
    assert payload["png_paths"] == [str(result.png_paths[0])]
    assert payload["baseline_status"] == ""


def test_official_document_sample_verification_rejects_trailing_blank_page(
    monkeypatch,
    tmp_path,
):
    def fake_pdf_to_png(pdf_path: Path, output_dir: Path):
        content_path = output_dir / f"{pdf_path.stem}_page-1.png"
        blank_path = output_dir / f"{pdf_path.stem}_page-2.png"
        content = Image.new("RGB", (8, 8), "white")
        content.putpixel((3, 3), (0, 0, 0))
        content.save(content_path)
        Image.new("RGB", (8, 8), "white").save(blank_path)
        return (content_path, blank_path), ()

    monkeypatch.setattr(
        official_visual,
        "render_docx_pages",
        _fake_rendering(fake_pdf_to_png),
    )

    result = verify_official_document_sample_visual("report", tmp_path)

    assert result.status == "png_blank"
    assert result.ok is False
    assert len(result.png_paths) == 2
    assert result.issues == (
        f"blank_png: {result.png_paths[1].name}",
    )


def test_official_document_sample_verification_records_missing_visual_baseline(
    monkeypatch,
    tmp_path,
):
    def fake_pdf_to_png(pdf_path: Path, output_dir: Path):
        png_path = output_dir / f"{pdf_path.stem}_page-1.png"
        image = Image.new("RGB", (8, 8), "white")
        image.putpixel((3, 3), (0, 0, 0))
        image.save(png_path)
        return (png_path,), ()

    monkeypatch.setattr(
        official_visual,
        "render_docx_pages",
        _fake_rendering(fake_pdf_to_png),
    )

    baseline_dir = tmp_path / "baselines"
    result = verify_official_document_sample_visual(
        "letter",
        tmp_path,
        baseline_dir=baseline_dir,
    )

    assert result.status == "visual_baseline_missing"
    assert result.ok is False
    assert result.baseline_status == "baseline_missing"
    assert result.baseline_png_path == baseline_dir / "letter_official_sample_page-1.png"
    assert "visual_baseline_missing" in result.issues[0]
    payload = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert payload["baseline_status"] == "baseline_missing"
    assert payload["baseline_png_path"] == str(result.baseline_png_path)


def test_official_document_sample_verification_matches_visual_baseline(
    monkeypatch,
    tmp_path,
):
    def fake_pdf_to_png(pdf_path: Path, output_dir: Path):
        png_path = output_dir / f"{pdf_path.stem}_page-1.png"
        image = Image.new("RGB", (8, 8), "white")
        image.putpixel((3, 3), (0, 0, 0))
        image.save(png_path)
        return (png_path,), ()

    monkeypatch.setattr(
        official_visual,
        "render_docx_pages",
        _fake_rendering(fake_pdf_to_png),
    )

    baseline_dir = tmp_path / "baselines"
    baseline_dir.mkdir()
    baseline = Image.new("RGB", (8, 8), "white")
    baseline.putpixel((3, 3), (0, 0, 0))
    baseline.save(baseline_dir / "letter_official_sample_page-1.png")

    result = verify_official_document_sample_visual(
        "letter",
        tmp_path,
        baseline_dir=baseline_dir,
    )

    assert result.status == "visual_baseline_ok"
    assert result.ok is True
    assert result.baseline_status == "baseline_ok"
    assert result.baseline_difference_bbox == ()
    assert result.issues == ()
    payload = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert payload["status"] == "visual_baseline_ok"
    assert payload["baseline_status"] == "baseline_ok"


def test_official_document_visual_baseline_manifest_matches_registered_files():
    manifest = json.loads(
        (OFFICIAL_VISUAL_BASELINE_DIR / "visual_baseline_manifest.json").read_text(
            encoding="utf-8"
        )
    )
    decision = json.loads(
        (OFFICIAL_VISUAL_BASELINE_DIR / "official_master_family_decision.json").read_text(
            encoding="utf-8"
        )
    )

    assert manifest["schema_version"] == 2
    assert manifest["renderer"]["docx_to_pdf"] == "word_com"
    assert manifest["rasterizer"]["name"] == "pypdfium2"
    assert set(manifest["profiles"]) == {
        "notice",
        "letter",
        "minutes",
        "report",
        "request",
        "approval",
    }

    master_path = (OFFICIAL_VISUAL_BASELINE_DIR / manifest["master"]["path"]).resolve()
    assert master_path.is_file()
    assert _sha256(master_path) == manifest["master"]["sha256"]
    for master_id, master_evidence in manifest["masters"].items():
        family_path = (
            OFFICIAL_VISUAL_BASELINE_DIR / master_evidence["path"]
        ).resolve()
        assert family_path.is_file(), master_id
        assert _sha256(family_path) == master_evidence["sha256"]

    for profile_id, evidence in manifest["profiles"].items():
        png_path = OFFICIAL_VISUAL_BASELINE_DIR / evidence["png"]
        assert png_path.is_file(), profile_id
        assert _sha256(png_path) == evidence["sha256"]
        with Image.open(png_path) as image:
            assert image.size == (evidence["width"], evidence["height"])
        assert evidence["page_count"] == 1
        assert evidence["verification_status"] == "visual_baseline_ok"
        assert evidence["human_review"] == "passed"

    assert manifest["review"]["master_family_decision"] == (
        "split_master_family_accepted"
    )
    assert decision["status"] == "split_master_family_accepted"
    assert decision["split_candidate_profile_ids"] == [
        "letter",
        "minutes",
        "report",
        "request",
    ]
    assert decision["baseline_statuses"] == {
        "notice": "baseline_ok",
        "letter": "baseline_ok",
        "minutes": "baseline_ok",
        "report": "baseline_ok",
        "request": "baseline_ok",
        "approval": "baseline_ok",
    }


def test_official_document_sample_verification_reports_missing_required_fields(tmp_path):
    result = verify_official_document_sample_visual(
        "notice",
        tmp_path,
        entity_data={"body": ""},
        attempt_render=False,
    )

    assert result.status == "missing_required_fields"
    assert result.sample_docx_path is None
    assert result.missing_required_fields == ("body",)
    payload = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert payload["issues"] == ["assembly_missing_required_fields"]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
