import sys
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.services.document_structure_preview import (
    analyze_document_structure,
    suppression_selectors_before,
)


def test_document_structure_preview_detects_page_start_candidates(tmp_path):
    source = tmp_path / "sample.docx"
    doc = Document()
    for text in ["博士学位论文", "原创性声明", "摘要", "目录"]:
        doc.add_paragraph(text)
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("正文内容")
    doc.save(source)

    items = analyze_document_structure(source)

    assert [(item.section_type, item.label, item.title) for item in items] == [
        ("cover", "封面", "博士学位论文"),
        ("abstract_cn", "中文摘要", "摘要"),
        ("toc", "目录", "目录"),
        ("body", "正文", "第一章 绪论"),
    ]
    assert suppression_selectors_before(items, 3) == [
        "cover",
        "abstract_cn",
        "toc",
    ]
