"""Run representative Assistant-to-Form production chains end to end.

This is an executable regression probe, not a mocked production test.  Fixed
gateways make most provider responses deterministic, while the resulting
DOCX/PDF/package work is performed by the real production runtime.  Pass
``--online-smoke`` to add one request through the currently active provider.
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from dataclasses import replace
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from typing import Callable
from uuid import uuid4

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.application.materials import (  # noqa: E402
    ExecutionMaterialRecord,
    ExecutionMaterialSnapshot,
    ExecutionResource,
)
from src.assistant.adapters.content_generation_adapter import (  # noqa: E402
    AssistantContentGenerationAdapter,
    complete_generated_official_draft_field,
)
from src.assistant.adapters.production_adapter import (  # noqa: E402
    AssistantProductionAdapter,
)
from src.assistant.adapters.workspace_state_adapter import WorkspaceSnapshot  # noqa: E402
from src.assistant.application.content_generation_service import (  # noqa: E402
    AssistantContentGenerationService,
    ContentGenerationRequest,
)
from src.assistant.application.document_job_controller import (  # noqa: E402
    DocumentJobController,
)
from src.assistant.application.generated_draft_binding import (  # noqa: E402
    bind_generated_draft,
)
from src.assistant.contracts.permissions import DisclosureGrant  # noqa: E402
from src.assistant.runtime.provider_contract import (  # noqa: E402
    PROVIDER_DONE,
    PROVIDER_START,
    ProviderStreamEvent,
)
from src.assistant.runtime.providers import ProviderProfileStore, ProviderRouter  # noqa: E402
from src.domain.materials import (  # noqa: E402
    MaterialObjectRef,
    MaterialPackageRef,
    generate_package_id,
    generate_record_id,
)


class FixedGateway:
    def __init__(self, text: str) -> None:
        self.text = text
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(PROVIDER_DONE, text=self.text)

    def cancel(self) -> bool:
        return True


class QuizGateway:
    def __init__(self, *, review_candidate: bool = False) -> None:
        self.review_candidate = review_candidate
        self.requests = []

    def stream(self, request):
        self.requests.append(request)
        if request.metadata["generation_phase"] == "questions":
            text = quiz_questions(review_candidate=self.review_candidate)
        else:
            text = quiz_answers()
        yield ProviderStreamEvent(PROVIDER_START)
        yield ProviderStreamEvent(PROVIDER_DONE, text=text)

    def cancel(self) -> bool:
        return True


def quiz_questions(*, review_candidate: bool = False) -> str:
    lines = [
        "# 六年级语文随堂测验",
        "> 科目：语文　年级：小学六年级　考试时间：25 分钟　满分：30 分",
        "",
    ]
    sections = (
        ("一、基础客观题", range(1, 5), 2),
        ("二、阅读分析题", range(5, 8), 4),
        ("三、综合应用题", range(8, 9), 10),
    )
    for label, numbers, score in sections:
        lines.extend((f"## {label}", ""))
        for number in numbers:
            prompt = f"第 {number} 道完整题目，请结合语文知识作答。（{score} 分）"
            if review_candidate and number == 5:
                prompt += " UnrelatedForeignSentence AnotherForeignPhrase"
            lines.extend(
                (
                    f"{number}. {prompt}",
                    "   difficulty: 中等",
                    f"   knowledge_points: 知识点{(number - 1) % 3 + 1}",
                    "   answer_area_kind: lines",
                    "   answer_lines: 2",
                    "",
                )
            )
    return "\n".join(lines)


def quiz_answers() -> str:
    return "## 答案速查\n" + "\n".join(
        f"{number}. 第 {number} 题参考答案、简要解析与评分标准。"
        for number in range(1, 9)
    )


def workspace(
    input_path: Path | None = None,
    *,
    material_summary: dict[str, object] | None = None,
) -> WorkspaceSnapshot:
    return WorkspaceSnapshot(
        mode_id="custom",
        mode_label="通用",
        scene_id="custom",
        scene_source_type="builtin",
        template_id="default",
        template_source_type="builtin",
        input_path=str(input_path or ""),
        input_name=input_path.name if input_path else "",
        input_exists=bool(input_path and input_path.is_file()),
        material_summary=dict(material_summary or {}),
    )


def with_output(plan, output_root: Path):
    return replace(
        plan,
        output_policy=replace(plan.output_policy, output_root=str(output_root)),
    )


def generation_request(
    plan, *, session_id: str, provider_id: str, model_id: str, **kwargs
):
    return ContentGenerationRequest(
        session_id=session_id,
        turn_id=plan.created_by_turn_id,
        prompt=plan.intent,
        provider_id=provider_id,
        model_id=model_id,
        plan_id=plan.plan_id,
        plan_revision=plan.revision,
        plan_fingerprint=plan.fingerprint,
        capability_id=plan.capability_ref.capability_id,
        artifact_kind=plan.generation_contract.artifact_kind,
        prompt_profile_id=plan.generation_contract.prompt_profile_id,
        scene_id=str(plan.scene_ref.get("id") or ""),
        scale_profile_id=str(plan.scene_ref.get("scale_profile_id") or ""),
        document_type_id=plan.production_contract.document_type_id,
        **kwargs,
    )


def resource(path: Path, media_type: str) -> ExecutionResource:
    payload = path.read_bytes()
    return ExecutionResource(
        MaterialObjectRef(
            object_id="sha256:" + sha256(payload).hexdigest(),
            media_type=media_type,
            original_name=path.name,
            size=len(payload),
        ),
        str(path.resolve()),
    )


def material_snapshot(
    *,
    mode_id: str,
    contract_id: str,
    scene_id: str,
    fields: dict[str, str],
    resources: dict[str, tuple[ExecutionResource, ...]],
    domains: dict[str, str],
) -> ExecutionMaterialSnapshot:
    return ExecutionMaterialSnapshot(
        snapshot_id="",
        run_id="run-" + uuid4().hex,
        package_ref=MaterialPackageRef(
            generate_package_id(),
            "sha256:" + "b" * 64,
        ),
        work_mode_id=mode_id,
        material_contract_id=contract_id,
        recipe_id="document_batch",
        scene_id=scene_id,
        document_type="",
        package_display_name="真实链路测试资料包",
        package_field_values={},
        package_field_owners={},
        groups=(),
        records=(
            ExecutionMaterialRecord(
                record_id=generate_record_id(),
                display_name="测试记录",
                group_id="",
                field_values=fields,
                field_owners={key: "record" for key in fields},
                resources=resources,
                resource_owners={
                    key: tuple("record" for _ in values)
                    for key, values in resources.items()
                },
            ),
        ),
        resource_domains=domains,
    )


def existing_artifacts(result: dict[str, object]) -> list[str]:
    paths: list[str] = []
    for key in (
        "output_paths",
        "compare_paths",
        "material_manifest_paths",
        "material_package_paths",
    ):
        value = result.get(key)
        if isinstance(value, dict):
            paths.extend(str(item) for item in value.values() if str(item))
    value = result.get("report_paths")
    if isinstance(value, (list, tuple)):
        paths.extend(str(item) for item in value if str(item))
    return sorted({path for path in paths if Path(path).exists()})


def execute_flow(
    controller: DocumentJobController,
    plan,
    *,
    session_id: str,
    snapshot: ExecutionMaterialSnapshot | None = None,
    execution_id: str = "",
    cancel_check: Callable[[], bool] | None = None,
) -> tuple[dict[str, object], dict[str, object]]:
    preflight = controller.preflight(plan, material_snapshot=snapshot)
    stages: dict[str, object] = {
        "plan": {
            "mode": plan.work_mode_id,
            "route": plan.scene_ref.get("route_id", ""),
            "operation": plan.operation,
            "capability": plan.capability_ref.status,
        },
        "preflight": {
            "ready": preflight.ready,
            "issues": list(preflight.issues),
            "warnings": list(preflight.warnings),
        },
    }
    if not preflight.ready:
        return {"status": "blocked", "error_text": ";".join(preflight.issues)}, stages
    approval = controller.approve(session_id=session_id, plan=plan, preflight=preflight)
    stages["approval"] = {"approved": True, "approval_id": approval.approval_id}
    result = controller.execute(
        session_id=session_id,
        plan=plan,
        preflight=preflight,
        approval=approval,
        material_snapshot=snapshot,
        execution_id=execution_id,
        cancel_check=cancel_check,
    )
    stages["execution"] = {
        "status": result.get("status"),
        "error_text": result.get("error_text", ""),
        "failed_count": result.get("failed_count", 0),
        "artifact_failure_count": result.get("artifact_failure_count", 0),
    }
    stages["artifact_check"] = {"existing_paths": existing_artifacts(result)}
    return result, stages


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--online-smoke", action="store_true")
    parser.add_argument("--skip-heavy", action="store_true")
    parser.add_argument("--output-root", type=Path)
    args = parser.parse_args()
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    run_root = (
        args.output_root or ROOT / "_codex_work" / f"full_chain_regression_{stamp}"
    ).resolve()
    run_root.mkdir(parents=True, exist_ok=False)
    generation_root = run_root / "generation"
    generation = AssistantContentGenerationService(
        AssistantContentGenerationAdapter(generation_root)
    )
    controller = DocumentJobController(production=AssistantProductionAdapter())
    records: list[dict[str, object]] = []

    def run_case(
        name: str,
        expected: tuple[str, ...],
        action: Callable[[], tuple[dict[str, object], dict[str, object]]],
    ):
        print(f"START {name}", flush=True)
        started = time.monotonic()
        try:
            result, stages = action()
            status = str(result.get("status") or "")
            passed = status in expected
            error = str(result.get("error_text") or "")
        except Exception as exc:  # noqa: BLE001 - probe records the terminal failure
            result = {}
            stages = {}
            status = "exception"
            passed = False
            error = f"{type(exc).__name__}: {exc}"
        record = {
            "scenario": name,
            "passed": passed,
            "expected_statuses": list(expected),
            "status": status,
            "duration_seconds": round(time.monotonic() - started, 3),
            "error": error,
            "stages": stages,
            "artifacts": existing_artifacts(result),
        }
        records.append(record)
        print(
            f"END {name} status={status} passed={passed} duration={record['duration_seconds']}s",
            flush=True,
        )
        return result, stages

    def format_case(index: int):
        case_root = run_root / f"format_{index}"
        case_root.mkdir()
        source = case_root / "source.docx"
        doc = Document()
        doc.add_heading("待排版项目说明", level=1)
        doc.add_paragraph("这是用于真实格式化链路的合成内容。")
        doc.save(source)
        plan = controller.draft_plan(
            query="请把这个 Word 文档统一排版并生成最终文件",
            workspace=workspace(source),
            turn_id=f"turn-format-{index}",
        )
        return execute_flow(
            controller,
            with_output(plan, case_root / "outputs"),
            session_id=f"session-format-{index}",
        )

    for index in (1, 2):
        run_case(
            f"existing_docx_format_{index}",
            ("success",),
            lambda index=index: format_case(index),
        )

    def narrative_case(index: int, gateway=None):
        case_root = run_root / f"narrative_{index}"
        case_root.mkdir()
        plan = controller.draft_plan(
            query="生成一份项目周报，包含进展、风险和下周计划",
            workspace=workspace(),
            turn_id=f"turn-narrative-{index}",
        )
        plan = with_output(plan, case_root / "outputs")
        active_gateway = gateway or FixedGateway(
            "# 项目周报\n\n## 本周进展\n\n已完成需求梳理与联调。\n\n"
            "## 风险\n\n当前存在进度依赖风险。\n\n## 下周计划\n\n完成验收并整理交付材料。"
        )
        draft = generation.generate(
            generation_request(
                plan,
                session_id=f"session-narrative-{index}",
                provider_id="fixed",
                model_id="fixed",
            ),
            active_gateway,
        )
        result, stages = execute_flow(
            controller,
            bind_generated_draft(plan, draft),
            session_id=f"session-narrative-{index}",
        )
        stages["provider"] = {
            "request_count": len(getattr(active_gateway, "requests", ())),
            "kind": "fixed",
        }
        stages["draft"] = {
            "path": draft.document_path,
            "exists": Path(draft.document_path).is_file(),
        }
        return result, stages

    for index in (1, 2):
        run_case(
            f"prompt_narrative_{index}",
            ("success",),
            lambda index=index: narrative_case(index),
        )

    def material_narrative_case():
        case_root = run_root / "material_narrative"
        case_root.mkdir()
        source = case_root / "authorized-brief.docx"
        doc = Document()
        doc.add_heading("授权项目资料", level=1)
        doc.add_paragraph("项目代号为 Aurora，计划在十月完成首轮验收。")
        doc.save(source)
        plan = with_output(
            controller.draft_plan(
                query="根据这份资料生成项目进展报告，不要改动原文件",
                workspace=workspace(source),
                turn_id="turn-material-narrative",
            ),
            case_root / "outputs",
        )
        reference = str(source.resolve())
        grant = DisclosureGrant(
            grant_id="grant-" + uuid4().hex,
            session_id="session-material-narrative",
            provider_id="fixed",
            model_id="fixed",
            allowed_refs=(reference,),
            allowed_fields=("document_text",),
        )
        gateway = FixedGateway(
            "# Aurora 项目进展报告\n\n## 当前进展\n\n项目正在按计划推进。\n\n"
            "## 里程碑\n\n十月完成首轮验收。"
        )
        draft = generation.generate(
            generation_request(
                plan,
                session_id="session-material-narrative",
                provider_id="fixed",
                model_id="fixed",
                context_documents=({"title": source.name, "path": reference},),
                context_refs=(reference,),
                context_fields=("document_text",),
                disclosure_grant=grant,
            ),
            gateway,
        )
        result, stages = execute_flow(
            controller,
            bind_generated_draft(plan, draft),
            session_id="session-material-narrative",
        )
        provider_body = gateway.requests[0].messages[0]["content"]
        stages["provider"] = {
            "request_count": len(gateway.requests),
            "authorized_fact_seen": "Aurora" in provider_body
            and "十月" in provider_body,
            "local_path_redacted": str(case_root) not in provider_body,
        }
        return result, stages

    run_case("authorized_material_narrative", ("success",), material_narrative_case)

    def exam_case(review: bool):
        label = "review" if review else "clean"
        case_root = run_root / f"exam_{label}"
        case_root.mkdir()
        plan = with_output(
            controller.draft_plan(
                query="生成一份小学六年级语文随堂测验，8 道题，包含学生卷和答案卷",
                workspace=workspace(),
                turn_id=f"turn-exam-{label}",
            ),
            case_root / "outputs",
        )
        gateway = QuizGateway(review_candidate=review)
        draft = generation.generate(
            generation_request(
                plan,
                session_id=f"session-exam-{label}",
                provider_id="fixed",
                model_id="fixed",
            ),
            gateway,
        )
        result, stages = execute_flow(
            controller,
            bind_generated_draft(plan, draft),
            session_id=f"session-exam-{label}",
        )
        stages["provider"] = {
            "request_count": len(gateway.requests),
            "phases": [
                request.metadata.get("generation_phase") for request in gateway.requests
            ],
        }
        stages["draft"] = {
            "path": draft.markdown_path,
            "validation": dict(draft.validation_summary),
        }
        return result, stages

    if not args.skip_heavy:
        run_case("exam_clean_full_delivery", ("success",), lambda: exam_case(False))
        run_case(
            "exam_imperfect_review_delivery",
            ("partial_success",),
            lambda: exam_case(True),
        )

    def official_case():
        case_root = run_root / "official_pause_resume"
        case_root.mkdir()
        plan = with_output(
            controller.draft_plan(
                query="起草一份关于开展第三季度档案检查的通知公文并生成 Word",
                workspace=workspace(),
                turn_id="turn-official",
            ),
            case_root / "outputs",
        )
        gateway = FixedGateway(
            '{"fields":{"title":"关于开展第三季度档案检查的通知",'
            '"body":"请各部门按期完成自查并报送情况。","organization":null}}'
        )
        draft = generation.generate(
            generation_request(
                plan,
                session_id="session-official",
                provider_id="fixed",
                model_id="fixed",
            ),
            gateway,
        )
        initial = controller.preflight(bind_generated_draft(plan, draft))
        completed = complete_generated_official_draft_field(
            draft.source_path,
            field_key="organization",
            value="示例市综合办公室",
        )
        result, stages = execute_flow(
            controller,
            bind_generated_draft(plan, completed),
            session_id="session-official",
        )
        stages["pause_resume"] = {
            "initial_ready": initial.ready,
            "initial_issues": list(initial.issues),
            "completed_missing_fields": list(completed.missing_user_fields),
        }
        return result, stages

    run_case("official_missing_field_pause_resume", ("success",), official_case)

    def bidding_case():
        case_root = run_root / "bidding"
        case_root.mkdir()
        logo = case_root / "logo.png"
        seal = case_root / "seal.png"
        Image.new("RGB", (320, 120), "navy").save(logo)
        Image.new("RGBA", (220, 220), (180, 0, 0, 255)).save(seal)
        snapshot = material_snapshot(
            mode_id="bidding",
            contract_id="bid_materials_v1",
            scene_id="bidding",
            fields={
                "company_name": "示例科技有限公司",
                "project_name": "智慧档案平台项目",
                "legal_person": "张示例",
            },
            resources={
                "logo": (resource(logo, "image/png"),),
                "seal": (resource(seal, "image/png"),),
            },
            domains={"logo": "image", "seal": "image"},
        )
        plan = with_output(
            controller.draft_plan(
                query="生成一份投标标书正文，包括项目理解、响应内容、实施方案和承诺事项",
                workspace=workspace(
                    material_summary={
                        "package_id": "bid-main",
                        "material_schema_ids": ["bid_materials_v1"],
                        "field_count": 3,
                        "asset_count": 2,
                    }
                ),
                turn_id="turn-bidding",
            ),
            case_root / "outputs",
        )
        gateway = FixedGateway(
            "# 智慧档案平台项目投标文件\n\n"
            "{{@text:company_name}}\n\n{{@text:project_name}}\n\n{{@text:legal_person}}\n\n"
            "## 项目理解\n\n本项目将建设稳定、可审计的档案平台。\n\n"
            "## 响应内容\n\n我方逐项响应项目需求。\n\n"
            "## 实施方案\n\n项目分为调研、建设、试运行和验收阶段。\n\n"
            "## 承诺事项\n\n我方承诺按期完成交付并提供服务保障。"
        )
        draft = generation.generate(
            generation_request(
                plan,
                session_id="session-bidding",
                provider_id="fixed",
                model_id="fixed",
            ),
            gateway,
        )
        result, stages = execute_flow(
            controller,
            bind_generated_draft(plan, draft),
            session_id="session-bidding",
            snapshot=snapshot,
        )
        return result, stages

    if not args.skip_heavy:
        run_case("bidding_authoring_with_assets", ("success",), bidding_case)

    def qualification_case():
        case_root = run_root / "qualification_archive"
        case_root.mkdir()
        source = case_root / "qualification-index.docx"
        doc = Document()
        doc.add_heading("投标资质材料归档索引", level=1)
        doc.add_paragraph("本文件用于触发附件清单和 ZIP 归档交付。")
        doc.save(source)
        certificate = case_root / "certificate.pdf"
        license_file = case_root / "business_license.pdf"
        certificate.write_bytes(b"%PDF-1.4\n% synthetic qualification certificate\n")
        license_file.write_bytes(b"%PDF-1.4\n% synthetic business license\n")
        snapshot = material_snapshot(
            mode_id="bidding",
            contract_id="qualification_archive_assets_v1",
            scene_id="bidding",
            fields={
                "organization": "示例科技有限公司",
                "package_name": "投标资质归档包",
            },
            resources={
                "certificate": (resource(certificate, "application/pdf"),),
                "business_license": (resource(license_file, "application/pdf"),),
            },
            domains={"certificate": "attachment", "business_license": "attachment"},
        )
        plan = with_output(
            controller.draft_plan(
                query="整理投标资质证书材料并生成归档清单和 ZIP 包",
                workspace=workspace(
                    source,
                    material_summary={
                        "package_id": "qualification-main",
                        "material_schema_ids": ["qualification_archive_assets_v1"],
                        "field_count": 2,
                        "asset_count": 2,
                    },
                ),
                turn_id="turn-qualification",
            ),
            case_root / "outputs",
        )
        return execute_flow(
            controller,
            plan,
            session_id="session-qualification",
            snapshot=snapshot,
        )

    run_case("qualification_manifest_zip_only", ("success",), qualification_case)

    def malformed_provider_case():
        plan = controller.draft_plan(
            query="生成一份项目说明报告",
            workspace=workspace(),
            turn_id="turn-malformed-provider",
        )
        try:
            generation.generate(
                generation_request(
                    plan,
                    session_id="session-malformed",
                    provider_id="fixed",
                    model_id="fixed",
                ),
                FixedGateway("<w:document>forbidden direct OOXML</w:document>"),
            )
        except ValueError as exc:
            return {"status": "rejected", "error_text": str(exc)}, {
                "provider_guard": {"rejected": True}
            }
        return {"status": "unexpected_success", "error_text": ""}, {}

    run_case("malformed_provider_output_guard", ("rejected",), malformed_provider_case)

    def corrupt_input_case():
        case_root = run_root / "corrupt_input"
        case_root.mkdir()
        source = case_root / "corrupt.docx"
        source.write_bytes(b"not-a-docx")
        plan = with_output(
            controller.draft_plan(
                query="把这个 Word 文档统一排版",
                workspace=workspace(source),
                turn_id="turn-corrupt",
            ),
            case_root / "outputs",
        )
        preflight = controller.preflight(plan)
        return (
            {
                "status": "blocked" if not preflight.ready else "unexpected_ready",
                "error_text": ";".join(preflight.issues),
            },
            {"preflight": {"ready": preflight.ready, "issues": list(preflight.issues)}},
        )

    run_case("corrupt_docx_preflight_guard", ("blocked",), corrupt_input_case)

    def mutated_input_case():
        case_root = run_root / "mutated_input"
        case_root.mkdir()
        source = case_root / "source.docx"
        doc = Document()
        doc.add_paragraph("预检时的内容")
        doc.save(source)
        plan = with_output(
            controller.draft_plan(
                query="统一排版这个 Word",
                workspace=workspace(source),
                turn_id="turn-mutated",
            ),
            case_root / "outputs",
        )
        preflight = controller.preflight(plan)
        approval = controller.approve(
            session_id="session-mutated", plan=plan, preflight=preflight
        )
        doc = Document(source)
        doc.add_paragraph("预检后被修改的内容")
        doc.save(source)
        result = controller.execute(
            session_id="session-mutated",
            plan=plan,
            preflight=preflight,
            approval=approval,
        )
        return result, {"input_integrity": {"changed_after_preflight": True}}

    run_case("input_mutation_after_preflight", ("failed",), mutated_input_case)

    def cancellation_case():
        case_root = run_root / "cancelled"
        case_root.mkdir()
        source = case_root / "source.docx"
        Document().save(source)
        plan = with_output(
            controller.draft_plan(
                query="统一排版这个 Word",
                workspace=workspace(source),
                turn_id="turn-cancel",
            ),
            case_root / "outputs",
        )
        return execute_flow(
            controller,
            plan,
            session_id="session-cancel",
            cancel_check=lambda: True,
        )

    run_case("user_cancellation_terminal", ("cancelled",), cancellation_case)

    def idempotent_case():
        case_root = run_root / "idempotent"
        case_root.mkdir()
        source = case_root / "source.docx"
        Document().save(source)
        plan = with_output(
            controller.draft_plan(
                query="统一排版这个 Word",
                workspace=workspace(source),
                turn_id="turn-idempotent",
            ),
            case_root / "outputs",
        )
        execution_id = "regression-" + uuid4().hex
        preflight = controller.preflight(plan)
        approval = controller.approve(
            session_id="session-idempotent", plan=plan, preflight=preflight
        )
        first = controller.execute(
            session_id="session-idempotent",
            plan=plan,
            preflight=preflight,
            approval=approval,
            execution_id=execution_id,
        )
        second = controller.execute(
            session_id="session-idempotent",
            plan=plan,
            preflight=preflight,
            approval=approval,
            execution_id=execution_id,
        )
        second["status"] = (
            "replayed"
            if second.get("assistant_execution_replayed")
            and first.get("status") == "success"
            else "replay_failed"
        )
        return second, {
            "idempotency": {
                "first_status": first.get("status"),
                "replayed": second.get("assistant_execution_replayed", False),
            }
        }

    run_case("idempotent_execution_replay", ("replayed",), idempotent_case)

    if args.online_smoke:

        def online_case():
            store = ProviderProfileStore()
            profile_id = store.active_profile_id()
            profile = store.get(profile_id)
            router = ProviderRouter(profiles=store)
            readiness = router.readiness(profile_id)
            if not readiness.ready or profile.kind == "mock":
                return {
                    "status": "blocked",
                    "error_text": readiness.reason_code or "active_provider_is_mock",
                }, {"provider": {"profile_id": profile_id, "ready": readiness.ready}}
            case_root = run_root / "online_provider"
            case_root.mkdir()
            plan = with_output(
                controller.draft_plan(
                    query="生成一份简短的真实链路测试报告，包含目的、结果和结论",
                    workspace=workspace(),
                    turn_id="turn-online-provider",
                ),
                case_root / "outputs",
            )
            draft = generation.generate(
                generation_request(
                    plan,
                    session_id="session-online-provider",
                    provider_id=profile.profile_id,
                    model_id=profile.model_id,
                ),
                router.resolve(profile.profile_id, timeout_seconds=120),
            )
            result, stages = execute_flow(
                controller,
                bind_generated_draft(plan, draft),
                session_id="session-online-provider",
            )
            stages["provider"] = {
                "profile_id": profile.profile_id,
                "model_id": profile.model_id,
                "ready": True,
                "network_call": True,
            }
            return result, stages

        run_case("active_online_provider_full_chain", ("success",), online_case)

    report = {
        "kind": "assistant_full_chain_regression",
        "schema_version": 1,
        "started_at": stamp,
        "run_root": str(run_root),
        "online_smoke_requested": args.online_smoke,
        "summary": {
            "total": len(records),
            "passed": sum(bool(item["passed"]) for item in records),
            "failed": sum(not bool(item["passed"]) for item in records),
        },
        "scenarios": records,
    }
    report_path = run_root / "full_chain_regression_report.json"
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"REPORT {report_path}", flush=True)
    print(json.dumps(report["summary"], ensure_ascii=False), flush=True)
    return 0 if report["summary"]["failed"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
