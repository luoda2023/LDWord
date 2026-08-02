"""Export deterministic, privacy-safe screenshots for the V1.0 user manual."""

# ruff: noqa: E402

from __future__ import annotations

from pathlib import Path
import shutil
import subprocess
import sys
import time

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.shared.ui.font_engine_policy import (
    configure_application_windows_font_engine,
)

FONT_ENGINE = configure_application_windows_font_engine("freetype").engine

from src.qt_api import QApplication
from src.shared.ui.typography_policy import (
    apply_application_typography,
    register_windows_ui_fonts_for_freetype,
)
from src.ui.main_window import MainWindow


OUTPUT_DIR = ROOT / "docs" / "user" / "screenshots"
SAMPLE_DIR = ROOT / "docs" / "user" / "samples"
SAMPLE_PATH = SAMPLE_DIR / "快速开始示例.docx"


def _settle(app: QApplication, *, rounds: int = 18, delay: float = 0.01) -> None:
    for _index in range(rounds):
        app.processEvents()
        time.sleep(delay)


def _create_sample_document() -> None:
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.core_properties.title = "Alavette Form 快速开始示例"
    document.core_properties.subject = "公开教程演示文件"
    document.core_properties.author = "Alavette Form"
    document.add_heading("项目周报", level=1)
    document.add_paragraph("报告周期：2026 年 7 月 20 日—7 月 26 日")
    document.add_heading("本周完成", level=2)
    document.add_paragraph("完成用户文档结构梳理，并验证发布包构建流程。")
    document.add_paragraph("整理演示资料，全部内容均为虚构示例。")
    document.add_heading("下周计划", level=2)
    document.add_paragraph("完成首发版教程验收，并根据试用反馈修订。")
    document.add_heading("风险与处理", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(
        table.rows[0].cells,
        ("风险", "影响", "处理"),
        strict=True,
    ):
        cell.text = value
    row = table.add_row().cells
    row[0].text = "首次用户不熟悉流程"
    row[1].text = "可能找不到输出文件"
    row[2].text = "提供快速开始和示例文件"
    document.save(SAMPLE_PATH)


def _save_widget(widget, path: Path) -> None:
    pixmap = widget.grab()
    if pixmap.isNull() or not pixmap.save(str(path)):
        raise RuntimeError(f"Unable to save screenshot: {path}")


def _export_workbench(app: QApplication) -> None:
    window = MainWindow(enable_background_services=False)
    window.resize(1224, 824)
    window.show()
    _settle(app)

    workbench = window.panel_stack.widget(0)
    scroll_bar = workbench._detail_scroll.verticalScrollBar()
    scroll_bar.setValue(0)
    _settle(app)
    _save_widget(window, OUTPUT_DIR / "01-workbench-overview.png")

    workbench._on_document_loaded(str(SAMPLE_PATH))
    _settle(app, rounds=35)
    scroll_bar.setValue(0)
    _settle(app)
    _save_widget(window, OUTPUT_DIR / "02-document-selected.png")

    scroll_bar.setValue(scroll_bar.maximum())
    _settle(app)
    _save_widget(window, OUTPUT_DIR / "03-output-and-generate.png")
    window.close()
    _settle(app)


def _export_assistant() -> None:
    source = ROOT / "artifacts" / "assistant_active_conversation_overview_user-docs.png"
    subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts" / "export_assistant_conversation_visual.py"),
            "--suffix",
            "user-docs",
        ],
        cwd=ROOT,
        check=True,
    )
    if not source.is_file():
        raise RuntimeError(f"Assistant screenshot was not produced: {source}")
    shutil.copy2(source, OUTPUT_DIR / "04-ai-assistant-overview.png")


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    _create_sample_document()
    app = QApplication.instance() or QApplication([])
    register_windows_ui_fonts_for_freetype(FONT_ENGINE)
    apply_application_typography(app)
    _export_workbench(app)
    _export_assistant()
    print(OUTPUT_DIR)
    print(SAMPLE_PATH)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
