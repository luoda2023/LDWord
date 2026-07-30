"""Export native visual evidence for the active assistant conversation."""

# ruff: noqa: E402

from __future__ import annotations

import argparse
from pathlib import Path
import re
import sys
from tempfile import TemporaryDirectory
import time

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.assistant.application.session_coordinator import AssistantSessionCoordinator
from src.assistant.adapters.workspace_state_adapter import WorkspaceSnapshot
from src.assistant.application.plan_builder import FormDocumentPlanBuilder
from src.assistant.contracts.messages import AssistantMessage, ROLE_ASSISTANT, ROLE_USER
from src.assistant.runtime.events import EVENT_TEXT_DELTA, AssistantEvent
from src.assistant.runtime.mock_provider import MockModelGateway
from src.assistant.runtime.turn_runner import AssistantTurnRunner
from src.assistant.storage.session_store import AssistantSessionStore
from src.assistant.ui.assistant_panel import AssistantPanel
from src.qt_api import QApplication
from src.shared.ui.typography_policy import (
    apply_application_typography,
    register_windows_ui_fonts_for_freetype,
)
from src.ui.bridge import PanelBridge


_SAFE_SUFFIX = re.compile(r"^[A-Za-z0-9_-]{1,48}$")


def _settle(app: QApplication, rounds: int = 12) -> None:
    for _index in range(rounds):
        app.processEvents()


def _wait_until(app: QApplication, predicate, *, timeout_seconds: float) -> None:
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        app.processEvents()
        if predicate():
            return
        time.sleep(0.01)
    raise RuntimeError("Timed out while exporting Assistant visual evidence")


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export native Assistant conversation screenshots."
    )
    parser.add_argument(
        "--suffix",
        default="native",
        help="Safe filename suffix used to keep DPI/viewport runs isolated.",
    )
    args = parser.parse_args(argv)
    if not _SAFE_SUFFIX.fullmatch(str(args.suffix or "")):
        parser.error("--suffix must contain only letters, numbers, underscores, or hyphens")
    return args


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv)
    app = QApplication.instance() or QApplication([])
    register_windows_ui_fonts_for_freetype("freetype")
    apply_application_typography(app)
    output_root = ROOT / "artifacts"
    output_root.mkdir(parents=True, exist_ok=True)

    with TemporaryDirectory(prefix="alavette-assistant-visual-") as temp_root:
        coordinator = AssistantSessionCoordinator(
            AssistantSessionStore(Path(temp_root) / "assistant")
        )
        panel = AssistantPanel(
            PanelBridge(),
            coordinator=coordinator,
            first_level=True,
        )
        session = coordinator.create_session()
        query = "生成一份七年级数学试卷"
        source_path = Path(temp_root) / "七年级数学资料.docx"
        source_path.write_bytes(b"visual-evidence")
        panel._composer.set_document_path(str(source_path))
        session = coordinator.append_message(
            session,
            AssistantMessage.text(
                role=ROLE_USER,
                text=query,
                source_refs=(
                    {
                        "type": "file",
                        "title": source_path.name,
                        "path": str(source_path),
                    },
                ),
            ),
        )
        session = coordinator.append_message(
            session,
            AssistantMessage.text(
                role=ROLE_ASSISTANT,
                text=(
                    "## 已识别为试卷创作\n\n"
                    "> 当前目标：先生成可校验题稿，再进入本地装配。\n\n"
                    "- AI 先生成结构化 Markdown 题稿\n"
                    "- 本地校验题目、分值和答案结构\n"
                    "- 审批后由试卷装配器生成学生卷与答案卷"
                ),
                source_refs=(
                    {
                        "type": "file",
                        "title": source_path.name,
                        "path": str(source_path),
                    },
                ),
            ),
        )
        plan = FormDocumentPlanBuilder().build(
            query=query,
            workspace=WorkspaceSnapshot(
                mode_id="custom",
                mode_label="通用版",
                scene_id="custom",
                scene_source_type="builtin",
                template_id="default",
                template_source_type="builtin",
                input_path="",
                input_name="",
                input_exists=False,
                material_summary={},
            ),
            turn_id="visual-plan-turn",
        )
        session = coordinator.update_state(
            session,
            active_plan=plan.to_dict(),
            document_job={
                "status": "plan_ready",
                "plan_id": plan.plan_id,
                "plan_revision": plan.revision,
            },
        )
        session = coordinator.append_message(session, panel._plan_message(plan))
        session = coordinator.append_message(
            session,
            AssistantMessage.interaction(
                role=ROLE_ASSISTANT,
                interaction_type="question",
                title="选择试卷交付范围",
                body="选择本次需要生成的交付物。",
                payload={
                    "confirmation_request": {
                        "options": [
                            {"id": "student", "label": "学生卷"},
                            {"id": "answer", "label": "答案卷"},
                        ],
                        "multiple": True,
                        "allow_other": True,
                    }
                },
            ),
        )
        student_path = Path(temp_root) / "七年级数学试卷-学生卷.docx"
        answer_path = Path(temp_root) / "七年级数学试卷-答案卷.docx"
        student_path.write_bytes(b"student")
        answer_path.write_bytes(b"answer")
        session = coordinator.append_message(
            session,
            AssistantMessage.interaction(
                role=ROLE_ASSISTANT,
                interaction_type="artifact",
                title="文档已生成",
                body="全部产物已通过交付完整性检查，可直接打开审阅。",
                payload={
                    "references": [
                        {
                            "type": "file",
                            "artifact_key": "student",
                            "title": student_path.name,
                            "path": str(student_path),
                        },
                        {
                            "type": "file",
                            "artifact_key": "answer_key",
                            "title": answer_path.name,
                            "path": str(answer_path),
                        },
                    ],
                    "actions": [],
                },
            ),
        )
        session = coordinator.update_state(
            session,
            pending_continuation={
                "kind": "question",
                "continuation_id": "visual-question",
            },
            turn_status="waiting_user_question",
        )
        panel._active_session = session
        panel._render_active_session()
        panel.resize(1600, 1180)
        panel.show()
        _settle(app)

        panel._message_scroll.verticalScrollBar().setValue(0)
        _settle(app)
        overview_path = (
            output_root
            / f"assistant_active_conversation_overview_{args.suffix}.png"
        )
        if not panel.grab().save(str(overview_path)):
            raise RuntimeError(f"Unable to save {overview_path}")

        panel._message_scroll.verticalScrollBar().setValue(
            panel._message_scroll.verticalScrollBar().maximum()
        )
        _settle(app)
        conversation_path = (
            output_root / f"assistant_active_conversation_{args.suffix}.png"
        )
        if not panel.grab().save(str(conversation_path)):
            raise RuntimeError(f"Unable to save {conversation_path}")

        turn_id = "visual-live-turn"
        panel._start_turn_preview(session_id=session.session_id, turn_id=turn_id)
        panel._render_active_session()
        panel._on_turn_event(
            AssistantEvent(
                EVENT_TEXT_DELTA,
                turn_id,
                text_delta="正在根据现有材料梳理项目背景、进展与风险…",
            )
        )
        panel._composer.set_busy(True)
        _settle(app)
        streaming_path = (
            output_root
            / f"assistant_active_conversation_streaming_{args.suffix}.png"
        )
        if not panel.grab().save(str(streaming_path)):
            raise RuntimeError(f"Unable to save {streaming_path}")

        panel._clear_turn_preview()
        panel._composer.set_busy(False)
        panel._render_active_session()
        panel.resize(760, 700)
        _settle(app)
        compact_path = (
            output_root
            / f"assistant_active_conversation_compact_{args.suffix}.png"
        )
        if not panel.grab().save(str(compact_path)):
            raise RuntimeError(f"Unable to save {compact_path}")
        panel.close()

        disclosure_source = Path(temp_root) / "标准规划样稿.docx"
        disclosure_document = Document()
        disclosure_document.add_heading("规划标题", level=1)
        disclosure_document.add_paragraph("规划正文")
        disclosure_document.save(disclosure_source)
        disclosure_coordinator = AssistantSessionCoordinator(
            AssistantSessionStore(Path(temp_root) / "assistant-disclosure")
        )
        disclosure_panel = AssistantPanel(
            PanelBridge(),
            coordinator=disclosure_coordinator,
            turn_runner=AssistantTurnRunner(
                MockModelGateway("已完成格式要求分析。")
            ),
            first_level=True,
        )
        disclosure_panel._on_composer_document_selected(
            str(disclosure_source)
        )
        disclosure_panel._send_message(
            "这个是标准的规划文件，帮我看看确定对应的格式要求"
        )
        disclosure_panel.resize(1400, 900)
        disclosure_panel.show()
        _settle(app)
        disclosure_panel._message_scroll.verticalScrollBar().setValue(
            disclosure_panel._message_scroll.verticalScrollBar().maximum()
        )
        _settle(app)
        disclosure_path = (
            output_root / f"assistant_disclosure_{args.suffix}.png"
        )
        if not disclosure_panel.grab().save(str(disclosure_path)):
            raise RuntimeError(f"Unable to save {disclosure_path}")
        disclosure_panel.close()

        approval_source = Path(temp_root) / "待统一格式的项目报告.docx"
        approval_document = Document()
        approval_document.add_heading("项目报告", level=1)
        approval_document.add_paragraph("待统一格式的正文。")
        approval_document.save(approval_source)
        approval_coordinator = AssistantSessionCoordinator(
            AssistantSessionStore(Path(temp_root) / "assistant-approval")
        )
        approval_panel = AssistantPanel(
            PanelBridge(),
            coordinator=approval_coordinator,
            turn_runner=AssistantTurnRunner(
                MockModelGateway("本地路径不会调用模型。")
            ),
            first_level=True,
        )
        approval_panel._on_composer_document_selected(str(approval_source))
        approval_panel._send_message("统一这份 Word 文档格式")
        approval_panel._handle_card_action("preflight", {})
        _wait_until(
            app,
            lambda: approval_panel._preflight_worker is None,
            timeout_seconds=20,
        )
        approval_panel.resize(1400, 900)
        approval_panel.show()
        _settle(app)
        approval_panel._message_scroll.verticalScrollBar().setValue(
            approval_panel._message_scroll.verticalScrollBar().maximum()
        )
        _settle(app)
        approval_path = (
            output_root / f"assistant_execution_approval_{args.suffix}.png"
        )
        if not approval_panel.grab().save(str(approval_path)):
            raise RuntimeError(f"Unable to save {approval_path}")
        approval_panel.close()

    print(overview_path)
    print(conversation_path)
    print(streaming_path)
    print(compact_path)
    print(disclosure_path)
    print(approval_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
