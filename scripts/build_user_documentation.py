"""Build polished DOCX user documentation from the Markdown sources."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import shutil
import sys
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs" / "user"
OUTPUT_DIR = ROOT / "artifacts" / "user-docs" / "v1.0"
SCREENSHOT_DIR = SOURCE_DIR / "screenshots"
SAMPLE_PATH = SOURCE_DIR / "samples" / "快速开始示例.docx"

BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
TEXT = "243247"
MUTED = "637083"
WHITE = "FFFFFF"
WARNING_BG = "FFF4CE"
WARNING_ACCENT = "B7791F"
TIP_BG = "EAF4FF"
TIP_ACCENT = "2E74B5"
IMPORTANT_BG = "FDECEC"
IMPORTANT_ACCENT = "B42318"

INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\))"
)
IMAGE_PATTERN = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<path>[^)]+)\)$")
ORDERED_PATTERN = re.compile(r"^\d+\.\s+(?P<text>.+)$")
BULLET_PATTERN = re.compile(r"^-\s+(?P<text>.+)$")
HEADING_PATTERN = re.compile(r"^(?P<marks>#{1,6})\s+(?P<text>.+)$")


@dataclass(frozen=True)
class BuildSpec:
    source: Path
    output: Path
    short_title: str
    cover_kicker: str
    cover_description: str
    include_toc: bool


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(
    cell,
    *,
    top: int = 80,
    start: int = 120,
    bottom: int = 80,
    end: int = 120,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_width(table, width_dxa: int = 9360) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def _set_font(run, *, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(
        qn("w:eastAsia"),
        "Microsoft YaHei",
    )
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _set_paragraph_panel(
    paragraph,
    *,
    fill: str,
    border_color: str,
    border_side: str = "left",
    border_size: int = 20,
) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    shading = paragraph_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        paragraph_properties.append(shading)
    shading.set(qn("w:fill"), fill)

    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        paragraph_properties.append(borders)
    border = OxmlElement(f"w:{border_side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(border_size))
    border.set(qn("w:space"), "8")
    border.set(qn("w:color"), border_color)
    borders.append(border)


def _set_style_font(style, *, size: float, bold: bool = False, color: str = TEXT) -> None:
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(
        qn("w:eastAsia"),
        "Microsoft YaHei",
    )
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def _configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    _set_style_font(normal, size=11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DEEP_BLUE, 10, 5),
        ("Heading 4", 11, DEEP_BLUE, 8, 4),
    )
    for name, size, color, before, after in heading_tokens:
        style = styles[name]
        _set_style_font(style, size=size, bold=True, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        _set_style_font(style, size=11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.2

    if "Manual Numbered List" not in styles:
        numbered = styles.add_style(
            "Manual Numbered List",
            WD_STYLE_TYPE.PARAGRAPH,
        )
    else:
        numbered = styles["Manual Numbered List"]
    _set_style_font(numbered, size=11)
    numbered.paragraph_format.left_indent = Inches(0.375)
    numbered.paragraph_format.first_line_indent = Inches(-0.188)
    numbered.paragraph_format.space_after = Pt(3)
    numbered.paragraph_format.line_spacing = 1.2

    caption = styles["Caption"]
    _set_style_font(caption, size=9, color=MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    if "Cover Kicker" not in styles:
        kicker = styles.add_style("Cover Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        kicker = styles["Cover Kicker"]
    _set_style_font(kicker, size=11, bold=True, color=BLUE)
    kicker.paragraph_format.space_after = Pt(18)

    if "Cover Title" not in styles:
        cover_title = styles.add_style("Cover Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cover_title = styles["Cover Title"]
    _set_style_font(cover_title, size=26, bold=True, color=TEXT)
    cover_title.paragraph_format.space_after = Pt(14)
    cover_title.paragraph_format.keep_with_next = True

    if "Cover Subtitle" not in styles:
        cover_subtitle = styles.add_style("Cover Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cover_subtitle = styles["Cover Subtitle"]
    _set_style_font(cover_subtitle, size=13, color=MUTED)
    cover_subtitle.paragraph_format.space_after = Pt(24)
    cover_subtitle.paragraph_format.line_spacing = 1.35

    if "Small Metadata" not in styles:
        metadata = styles.add_style("Small Metadata", WD_STYLE_TYPE.PARAGRAPH)
    else:
        metadata = styles["Small Metadata"]
    _set_style_font(metadata, size=9.5, color=MUTED)
    metadata.paragraph_format.space_after = Pt(4)


def _configure_page(document: Document, *, short_title: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    section.different_first_page_header_footer = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(short_title)
    _set_font(run, size=8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Alavette Form V1.0  ·  ")
    _set_font(run, size=8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _add_cover(document: Document, spec: BuildSpec, title: str) -> None:
    paragraph = document.add_paragraph(style="Cover Kicker")
    paragraph.add_run(spec.cover_kicker)
    paragraph = document.add_paragraph(style="Cover Title")
    paragraph.add_run(title)
    paragraph = document.add_paragraph(style="Cover Subtitle")
    paragraph.add_run(spec.cover_description)

    values = (
        ("适用版本", "Alavette Form V1.0"),
        ("适用系统", "Windows 10 / Windows 11"),
        ("文档日期", "2026 年 7 月"),
    )
    for label, value in values:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.08)
        paragraph.paragraph_format.right_indent = Inches(0.08)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_together = True
        _set_paragraph_panel(
            paragraph,
            fill=LIGHT_BLUE,
            border_color=BLUE,
            border_side="bottom",
            border_size=6,
        )
        run = paragraph.add_run(f"{label}　")
        _set_font(run, size=10, bold=True)
        run.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
        run = paragraph.add_run(value)
        _set_font(run, size=10)

    document.add_paragraph("")
    metadata = document.add_paragraph(style="Small Metadata")
    metadata.add_run("这是一份随 V1.0 发布的离线用户文档。")
    metadata = document.add_paragraph(style="Small Metadata")
    metadata.add_run("正式交付前请在 Microsoft Word 中复核生成结果。")
    document.add_page_break()


def _add_toc(document: Document, headings: list[str]) -> None:
    paragraph = document.add_paragraph("内容导航", style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(0)
    intro = document.add_paragraph(
        "本页列出手册的一级章节；PDF 阅读器可通过书签快速跳转。"
    )
    intro.paragraph_format.space_after = Pt(8)
    for heading in headings:
        item = document.add_paragraph()
        item.paragraph_format.left_indent = Inches(0.18)
        item.paragraph_format.space_after = Pt(5)
        run = item.add_run(heading)
        _set_font(run, size=10.5, bold=True)
        run.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
    document.add_page_break()


def _add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            _set_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_font(run, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_font(run)
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
        else:
            label, url = re.match(r"^\[([^\]]+)\]\(([^)]+)\)$", token).groups()
            run = paragraph.add_run(label)
            _set_font(run)
            run.font.color.rgb = RGBColor.from_string(BLUE)
            run.underline = True
            run = paragraph.add_run(f" ({url})")
            _set_font(run, size=9)
            run.font.color.rgb = RGBColor.from_string(MUTED)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_font(run)


def _new_numbering_instance(document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "decimal")
    level.append(number_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.append(level_justification)
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    paragraph_properties.append(tabs)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "540")
    indentation.set(qn("w:hanging"), "270")
    paragraph_properties.append(indentation)
    level.append(paragraph_properties)
    abstract.append(level)
    # Word requires every abstractNum before the first concrete num. Appending
    # an abstractNum after existing num elements makes Word repair the package
    # and can silently turn built-in bullet lists into decimal lists.
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    num.append(abstract_reference)
    numbering.append(num)
    return num_id


def _apply_numbering(paragraph, num_id: int) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    num_properties = paragraph_properties.find(qn("w:numPr"))
    if num_properties is None:
        num_properties = OxmlElement("w:numPr")
        paragraph_properties.append(num_properties)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_properties.extend((level, number))


def _add_callout(document: Document, kind: str, lines: list[str]) -> None:
    labels = {
        "TIP": ("提示", TIP_BG, TIP_ACCENT),
        "WARNING": ("注意", WARNING_BG, WARNING_ACCENT),
        "IMPORTANT": ("重要", IMPORTANT_BG, IMPORTANT_ACCENT),
        "NOTE": ("说明", LIGHT_BLUE, BLUE),
    }
    label, background, accent = labels.get(kind, labels["NOTE"])
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    _set_paragraph_panel(
        paragraph,
        fill=background,
        border_color=accent,
        border_side="left",
        border_size=24,
    )
    run = paragraph.add_run(label)
    _set_font(run, size=10, bold=True)
    run.font.color.rgb = RGBColor.from_string(accent)
    run.add_break()
    _add_inline(paragraph, " ".join(lines))


def _add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    data_rows = rows[2:] if all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]) else rows[1:]
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    _set_table_width(table)
    header = table.rows[0]
    _set_repeat_table_header(header)
    _set_cant_split(header)
    for index, value in enumerate(rows[0]):
        cell = header.cells[index]
        _set_cell_shading(cell, LIGHT_BLUE)
        _set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        _add_inline(paragraph, value)
        for run in paragraph.runs:
            run.bold = True
    for row_values in data_rows:
        row = table.add_row()
        _set_cant_split(row)
        cells = row.cells
        for index, value in enumerate(row_values[: len(cells)]):
            cell = cells[index]
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, value)
    document.add_paragraph("").paragraph_format.space_after = Pt(0)


def _set_image_alt_text(inline_shape, alt_text: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", alt_text)


def _add_image(document: Document, relative_path: str, alt_text: str) -> None:
    image_path = SOURCE_DIR / relative_path
    if not image_path.is_file():
        raise FileNotFoundError(f"Missing manual image: {image_path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(6.35))
    if inline.height > Inches(6.8):
        ratio = Inches(6.8) / inline.height
        inline.height = Inches(6.8)
        inline.width = int(inline.width * ratio)
    _set_image_alt_text(inline, alt_text)
    caption = document.add_paragraph(alt_text, style="Caption")
    caption.paragraph_format.keep_with_next = False


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _render_markdown(document: Document, source: Path, spec: BuildSpec) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    title_match = HEADING_PATTERN.match(lines[0])
    if title_match is None or len(title_match.group("marks")) != 1:
        raise ValueError(f"First line must be one H1 title: {source}")
    title = title_match.group("text")
    _add_cover(document, spec, title)
    if spec.include_toc:
        chapter_headings = [
            match.group("text")
            for line in lines
            if (match := HEADING_PATTERN.match(line))
            and len(match.group("marks")) == 2
        ]
        _add_toc(document, chapter_headings)

    index = 1
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        paragraph = document.add_paragraph()
        _add_inline(paragraph, " ".join(paragraph_lines))
        paragraph_lines.clear()

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue

        image_match = IMAGE_PATTERN.match(stripped)
        if image_match:
            flush_paragraph()
            _add_image(
                document,
                image_match.group("path"),
                image_match.group("alt") or "界面截图",
            )
            index += 1
            continue

        heading_match = HEADING_PATTERN.match(stripped)
        if heading_match:
            flush_paragraph()
            markdown_level = len(heading_match.group("marks"))
            word_level = min(max(markdown_level - 1, 1), 4)
            if (
                spec.include_toc
                and markdown_level == 2
                and heading_match.group("text") == "11. 快速索引"
            ):
                document.add_page_break()
            document.add_paragraph(
                heading_match.group("text"),
                style=f"Heading {word_level}",
            )
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith(">"):
                quote_lines.append(lines[index].lstrip()[1:].strip())
                index += 1
            kind = "NOTE"
            if quote_lines and re.fullmatch(r"\[![A-Z]+\]", quote_lines[0]):
                kind = quote_lines.pop(0)[2:-1]
            _add_callout(document, kind, quote_lines)
            continue

        ordered_match = ORDERED_PATTERN.match(stripped)
        if ordered_match:
            flush_paragraph()
            num_id = _new_numbering_instance(document)
            while index < len(lines):
                item_match = ORDERED_PATTERN.match(lines[index].strip())
                if item_match is None:
                    break
                paragraph = document.add_paragraph(style="Manual Numbered List")
                _apply_numbering(paragraph, num_id)
                _add_inline(paragraph, item_match.group("text"))
                index += 1
            continue

        bullet_match = BULLET_PATTERN.match(stripped)
        if bullet_match:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, bullet_match.group("text"))
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            table_rows: list[list[str]] = []
            while (
                index < len(lines)
                and lines[index].strip().startswith("|")
                and lines[index].strip().endswith("|")
            ):
                table_rows.append(_split_table_row(lines[index]))
                index += 1
            _add_markdown_table(document, table_rows)
            continue

        paragraph_lines.append(stripped)
        index += 1
    flush_paragraph()


def _remove_empty_first_paragraph(document: Document) -> None:
    if not document.paragraphs:
        return
    first = document.paragraphs[0]
    if first.text:
        return
    first._element.getparent().remove(first._element)


def _privacy_audit(path: Path) -> None:
    forbidden = (
        "AppData/Local/Temp",
        "AppData\\Local\\Temp",
        "sk-",
    )
    local_profile_path = re.compile(r"[A-Za-z]:[\\/]Users[\\/][^\\/<]+", re.I)
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith((".xml", ".rels", ".txt")):
                continue
            text = archive.read(name).decode("utf-8", errors="ignore")
            if match := local_profile_path.search(text):
                raise RuntimeError(
                    f"Privacy audit failed for {path.name}: local user path in {name}: "
                    f"{match.group(0)!r}"
                )
            for token in forbidden:
                if token in text:
                    raise RuntimeError(
                        f"Privacy audit failed for {path.name}: {token!r} in {name}"
                    )


def _build(spec: BuildSpec) -> None:
    document = Document()
    _remove_empty_first_paragraph(document)
    _configure_styles(document)
    _configure_page(document, short_title=spec.short_title)
    document.core_properties.title = spec.short_title
    document.core_properties.subject = spec.cover_description
    document.core_properties.author = "Alavette Form"
    document.core_properties.comments = "Generated from docs/user Markdown sources."
    _render_markdown(document, spec.source, spec)
    spec.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(spec.output)
    _privacy_audit(spec.output)
    print(spec.output)


def _build_obsidian_bundle() -> None:
    bundle_name = "Alavette Form V1.0 - Obsidian 用户文档"
    bundle_root = OUTPUT_DIR / bundle_name
    screenshot_target = bundle_root / "screenshots"
    sample_target = bundle_root / "samples"
    screenshot_target.mkdir(parents=True, exist_ok=True)
    sample_target.mkdir(parents=True, exist_ok=True)

    for source_name in (
        "00-开始这里.md",
        "quick-start.md",
        "user-guide.md",
    ):
        shutil.copy2(SOURCE_DIR / source_name, bundle_root / source_name)
    for screenshot in sorted(SCREENSHOT_DIR.glob("*.png")):
        shutil.copy2(screenshot, screenshot_target / screenshot.name)
    shutil.copy2(SAMPLE_PATH, sample_target / SAMPLE_PATH.name)

    archive_path = OUTPUT_DIR / f"{bundle_name}.zip"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(bundle_root.rglob("*")):
            if path.is_file():
                archive.write(path, Path(bundle_name) / path.relative_to(bundle_root))
    print(bundle_root)
    print(archive_path)


def main() -> int:
    required = (
        SOURCE_DIR / "quick-start.md",
        SOURCE_DIR / "user-guide.md",
        SCREENSHOT_DIR / "01-workbench-overview.png",
        SCREENSHOT_DIR / "02-document-selected.png",
        SCREENSHOT_DIR / "03-output-and-generate.png",
        SCREENSHOT_DIR / "04-ai-assistant-overview.png",
        SAMPLE_PATH,
    )
    missing = [path for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError(
            "Missing user-documentation inputs:\n" + "\n".join(map(str, missing))
        )

    specs = (
        BuildSpec(
            source=SOURCE_DIR / "quick-start.md",
            output=OUTPUT_DIR / "Alavette Form V1.0 - 快速开始.docx",
            short_title="Alavette Form V1.0 快速开始",
            cover_kicker="10 分钟完成第一次安全生成",
            cover_description=(
                "从完整解压发布包开始，添加示例 DOCX、完成预检、生成新文档，"
                "并确认原稿没有被覆盖。"
            ),
            include_toc=False,
        ),
        BuildSpec(
            source=SOURCE_DIR / "user-guide.md",
            output=OUTPUT_DIR / "Alavette Form V1.0 - 用户手册.docx",
            short_title="Alavette Form V1.0 用户手册",
            cover_kicker="任务教程 · 操作指南 · 概念参考 · 排障",
            cover_description=(
                "面向文档经办人、模板维护者和支持人员的完整离线手册，"
                "覆盖通用、试卷、论文、公文与可选 AI 流程。"
            ),
            include_toc=True,
        ),
    )
    for spec in specs:
        _build(spec)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SAMPLE_PATH, OUTPUT_DIR / SAMPLE_PATH.name)
    print(OUTPUT_DIR / SAMPLE_PATH.name)
    _build_obsidian_bundle()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
